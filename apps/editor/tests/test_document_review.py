"""Belge incelemesi: biçimden metin çıkarma, sayfalama, belge bağlamı. Dosyalar testin içinde üretilir
(python-docx, pymupdf imajda). Model ve veritabanı yok. Editör imajında çalıştır:

    python /app/tests/test_document_review.py
"""

from __future__ import annotations

import asyncio
import io
import sys
import types
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))


class _Stub(types.ModuleType):
    def __getattr__(self, name):
        if name.startswith("__"):
            raise AttributeError(name)
        sub = _Stub(f"{self.__name__}.{name}")
        sys.modules[sub.__name__] = sub
        return sub


for _mod in ("psycopg", "psycopg.rows", "psycopg.types", "psycopg.types.json", "psycopg_pool"):
    sys.modules.setdefault(_mod, _Stub(_mod))

from editor import document_review as DR  # noqa: E402
from editor.proofing import _doc_context as D  # noqa: E402

TEXT = ["Dolabın gözü açık kaldı.", "Gözüme toz kaçtı, sonra öğretmenin gözüne girdi."]


def _docx() -> bytes:
    from docx import Document
    d = Document()
    for t in TEXT:
        d.add_paragraph(t)
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Tablodaki cümle."
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _odt() -> bytes:
    body = "".join(f"<text:p>{t}</text:p>" for t in TEXT)
    xml = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content '
           'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
           + body + "</office:text></office:body></office:document-content>")
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", xml)
    return b.getvalue()


def _pdf() -> bytes:
    import pymupdf
    doc = pymupdf.open()
    for t in TEXT:
        page = doc.new_page()
        page.insert_text((72, 100), t.replace("ı", "i").replace("ö", "o").replace("ü", "u").replace("ç", "c")
                         .replace("ğ", "g").replace("ş", "s").replace("İ", "I"), fontsize=12)
    return doc.tobytes()


def _all_text(ex: dict) -> str:
    return " ".join(s["text"] for p in ex["pages"] for s in p["spans"])


def test_docx_paragraphs_and_tables():
    ex = DR.extract(_docx(), "Müsvedde.docx")
    assert ex["format"] == "docx" and ex["page_kind"] == "APPROXIMATE"
    assert "Dolabın gözü açık kaldı." in _all_text(ex) and "Tablodaki cümle." in _all_text(ex)
    assert ex["pages"][0]["approximate"] is True and ex["pages"][0]["spans"][0]["source"] == "TEXT_LAYER"


def test_odt_rtf_txt_md():
    assert "Gözüme toz kaçtı" in _all_text(DR.extract(_odt(), "a.odt"))
    rtf = rb"{\rtf1\ansi\ansicpg1254 Dolab\'fdn g\'f6z\'fc a\'e7\'fdk kald\'fd.\par G\'f6z\'fcme toz ka\'e7t\'fd.}"
    assert "Dolabın gözü açık kaldı." in _all_text(DR.extract(rtf, "a.rtf"))
    txt = "\n".join(TEXT[0].split(" ")[:2]) + " " + " ".join(TEXT[0].split(" ")[2:]) + "\n\n" + TEXT[1]
    ex = DR.extract(txt.encode("utf-8"), "a.txt")
    assert [s["text"] for s in ex["pages"][0]["spans"]] == TEXT          # blok içi satır sonu kelime ayrımı
    assert DR.extract(TEXT[0].encode("cp1254"), "b.md")["words"] == 4


def test_pdf_keeps_printed_pages():
    ex = DR.extract(_pdf(), "prova.pdf")
    assert ex["format"] == "pdf" and ex["page_kind"] == "PRINTED" and len(ex["pages"]) == 2
    assert ex["pages"][0]["approximate"] is False and "Dolabin" in ex["pages"][0]["spans"][0]["text"]


def test_rejects_unknown_type_and_empty_text():
    for data, name in ((b"x", "a.exe"), (b"", "a.txt"), (b"   \n\n  ", "a.txt")):
        try:
            DR.extract(data, name)
        except DR.DocumentError:
            continue
        raise AssertionError(name)


def test_paginate_never_splits_a_paragraph():
    paras = ["a " * 100, "b " * 100, "c " * 100, "d " * 400, "e " * 10]
    pages = DR.paginate(paras, 250)
    assert [len(p) for p in pages] == [2, 1, 1, 1]
    assert sum(len(p) for p in pages) == len(paras)


def test_document_context_supplies_pages_profile_and_no_book():
    pages = [{"page_no": 1, "spans": [{"idx": 1, "text": "x", "source": "TEXT_LAYER"}]}]
    tok = D.use({"id": "d1", "pages": pages, "audience": "CHILD", "age_from": 6, "age_to": 8, "title": "T"})
    try:
        assert D.is_document("d1") and not D.is_document("başka")
        assert D.pages("d1") == pages and D.llm_gid("d1") is None and D.book_version("d1") is None
        p = asyncio.run(D.profile("d1"))
        assert p["audience"] == "CHILD" and p["audience_source"] == "EDITOR" and p["age_from"] == 6 and p["pages"] == 1
    finally:
        D.reset(tok)
    assert not D.is_document("d1")


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print("ok", name)
