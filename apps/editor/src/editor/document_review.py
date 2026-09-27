"""Belge incelemesi: editörün yüklediği belgede (doc, docx, pdf, odt, rtf, txt, md) Zeki AI'ın metin denetimleri.

Kitap okumasından (analiz hattı: sayfa görseli, OCR, karakter, olay) ayrıdır; yalnız metin. Akış:
  1. `extract(data, file_name)` — metin çıkarılır, paragraflara ayrılır, sayfalanır (source.read biçimi):
     PDF kendi basılı sayfalarıyla (PRINTED); öbürlerinde sayfa yoktur, paragraf bölünmeden yaklaşık
     `PAGE_WORDS` sözcüklük sayfalar kurulur (APPROXIMATE; ekranda «yaklaşık sayfa»).
  2. `create(...)` — `document_review` satırı, durum QUEUED (kart servisi yükleme ucu).
  3. `consume()` — kuyruğu işler (compose servisi `document-review`): her belge için METİN denetimleri belge
     bağlamında (proofing/_doc_context) koşar, sonuç `document_run` / `document_finding`'e yazılır.
Hiçbir denetim düşmesi belgeyi düşürmez: denetim FAILED yazılır, öbürleri koşar (proofing.run_all ile aynı ilke).
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import re
import time
import traceback
import unicodedata
import zipfile
from pathlib import PurePath

from . import db

# Belgeye uygun denetimler: yalnız metinden çalışanlar. Sıra önemli: word_overuse derlemi okunmuş kitaplardan.
CHECKS = ("word_variety", "sentence_starts", "phrase_repeats", "word_choice", "word_overuse")
FORMATS = ("pdf", "docx", "doc", "odt", "rtf", "txt", "md")
# Basılı bir kitap sayfası ~250 sözcük; sayfası olmayan belgede konum bildirmek için (ayar, kapsam değil).
PAGE_WORDS = 250


class DocumentError(ValueError):
    """Belgenin kendisiyle ilgili (okunamıyor, boş, desteklenmeyen tür): yükleyene gösterilir."""


def _clean(s: str) -> str:
    s = unicodedata.normalize("NFC", s or "").replace("­", "")
    s = "".join(ch for ch in s if ch in "\t\n" or unicodedata.category(ch) not in ("Cc", "Cs", "Co"))
    return " ".join(s.split())


def fmt_of(file_name: str) -> str:
    ext = PurePath(file_name or "").suffix.lower().lstrip(".")
    if ext not in FORMATS:
        raise DocumentError(f"Desteklenmeyen dosya türü: .{ext or '?'} (desteklenen: {', '.join(FORMATS)})")
    return ext


# ------------------------------------------------------------------ biçim → paragraflar
def _pdf_pages(data: bytes) -> list[list[str]]:
    import pymupdf
    from .document import paragraphs_from_layout
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:  # noqa: BLE001
        raise DocumentError("PDF açılamadı") from e
    return [[_clean(p) for p in paragraphs_from_layout(page)] for page in doc]


def _docx_paras(data: bytes) -> list[str]:
    from docx import Document
    try:
        d = Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise DocumentError("Word (.docx) dosyası açılamadı") from e
    out = [p.text for p in d.paragraphs]
    for t in d.tables:                         # tablo hücreleri de metindir
        for row in t.rows:
            for cell in row.cells:
                out += [p.text for p in cell.paragraphs]
    return out


def _doc_paras(data: bytes) -> list[str]:
    import subprocess
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".doc") as f:
        f.write(data)
        f.flush()
        try:
            r = subprocess.run(["antiword", "-w", "0", f.name], capture_output=True, timeout=120)
        except FileNotFoundError as e:
            raise DocumentError("Eski Word (.doc) okuyucusu kurulu değil") from e
    if r.returncode != 0:
        raise DocumentError("Eski Word (.doc) dosyası okunamadı; .docx olarak kaydedip yükleyin")
    return r.stdout.decode("utf-8", "replace").split("\n\n")


def _odt_paras(data: bytes) -> list[str]:
    import xml.etree.ElementTree as ET
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = ET.fromstring(z.read("content.xml"))
    except Exception as e:  # noqa: BLE001
        raise DocumentError("OpenDocument (.odt) dosyası açılamadı") from e
    ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    out = []
    for el in root.iter():
        if el.tag in (ns + "p", ns + "h"):
            parts = []
            for node in el.iter():
                if node.tag == ns + "s":
                    parts.append(" " * int(node.get(ns + "c", "1")))
                elif node.tag in (ns + "tab", ns + "line-break"):
                    parts.append(" ")
                if node.text and node.tag != ns + "note-citation":
                    parts.append(node.text)
                if node is not el and node.tail:
                    parts.append(node.tail)
            out.append("".join(parts))
    return out


def _rtf_paras(data: bytes) -> list[str]:
    from striprtf.striprtf import rtf_to_text
    try:
        text = rtf_to_text(data.decode("latin-1"), encoding="cp1254", errors="replace")
    except Exception as e:  # noqa: BLE001
        raise DocumentError("RTF dosyası okunamadı") from e
    return text.split("\n")


def _text_paras(data: bytes) -> list[str]:
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    # boş satırla ayrılmış bloklar paragraftır; blok içindeki satır sonu kelime ayrımıdır
    return [b.replace("\n", " ") for b in re.split(r"\n\s*\n", text.replace("\r\n", "\n"))]


def paginate(paras: list[str], words_per_page: int = PAGE_WORDS) -> list[list[str]]:
    """Sayfası olmayan belge: paragraflar bölünmeden sırayla sayfalara; bir sayfa ~words_per_page sözcüğe
    ulaşınca yeni sayfa. Tek paragraf sayfadan uzunsa kendi sayfasıdır (bölünmez)."""
    pages, cur, n = [], [], 0
    for p in paras:
        w = len(p.split())
        if cur and n + w > words_per_page:
            pages.append(cur)
            cur, n = [], 0
        cur.append(p)
        n += w
    if cur:
        pages.append(cur)
    return pages


def extract(data: bytes, file_name: str) -> dict:
    """{format, page_kind, pages (source.read biçimi), words}. Boş ya da okunamayan belge DocumentError."""
    if not data:
        raise DocumentError("Dosya boş")
    fmt = fmt_of(file_name)
    if fmt == "pdf":
        raw_pages, kind = _pdf_pages(data), "PRINTED"
    else:
        paras = {"docx": _docx_paras, "doc": _doc_paras, "odt": _odt_paras, "rtf": _rtf_paras,
                 "txt": _text_paras, "md": _text_paras}[fmt](data)
        raw_pages, kind = paginate([_clean(p) for p in paras if _clean(p)]), "APPROXIMATE"
    pages = []
    for i, paras in enumerate(raw_pages, start=1):
        spans = [{"idx": k, "text": t, "source": "TEXT_LAYER", "reading_order": None}
                 for k, t in enumerate((p for p in paras if p), start=1)]
        pages.append({"page_no": i, "approximate": kind == "APPROXIMATE", "spans": spans, "issues": []})
    words = sum(len(s["text"].split()) for p in pages for s in p["spans"])
    if words == 0:
        raise DocumentError("Belgede okunabilir metin yok (taranmış PDF ise önce metin tanıma gerekir)")
    return {"format": fmt, "page_kind": kind, "pages": pages, "words": words}


# ------------------------------------------------------------------ kayıt
def create(data: bytes, file_name: str, title: str | None, uploaded_by: str, audience: str | None = None,
           age_from: int | None = None, age_to: int | None = None) -> dict:
    ex = extract(data, file_name)
    if audience not in (None, "", "CHILD", "YOUNG", "ADULT"):
        raise DocumentError("Okur kitlesi CHILD, YOUNG ya da ADULT olmalı")
    title = (title or "").strip() or PurePath(file_name).stem
    return db.one(
        "INSERT INTO document_review(title, file_name, format, byte_size, sha256, audience, age_from, age_to,"
        " uploaded_by, pages, page_kind, words) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
        " RETURNING id, title, file_name, format, page_kind, words, status, created_at",
        title[:300], PurePath(file_name).name[:300], ex["format"], len(data), hashlib.sha256(data).hexdigest(),
        audience or None, age_from, age_to, uploaded_by[:200], db.J(ex["pages"]), ex["page_kind"], ex["words"])


def _record(doc_id: str, mod, findings: list[dict], stats: dict, started: float) -> dict:
    from .proofing import _clean as clean_finding
    rows = [clean_finding(f) for f in findings]
    with db.tx() as c:
        run = c.execute(
            "INSERT INTO document_run(document_id, check_name, check_version, status, stats, started_at)"
            " VALUES (%s,%s,%s,'SUCCEEDED',%s,to_timestamp(%s)) RETURNING id",
            (doc_id, mod.NAME, str(mod.VERSION), db.J({**stats, "findings": len(rows)}), started)).fetchone()["id"]
        for f in rows:
            c.execute("INSERT INTO document_finding(run_id, document_id, check_name, page_no, severity, quote, message,"
                      " suggestion, details) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                      (run, doc_id, mod.NAME, f["page"], f["severity"], f["quote"], f["message"], f["suggestion"],
                       db.J(f["details"])))
    return {"run_id": str(run), "findings": len(rows)}


async def review(doc_id: str) -> dict:
    """Bir belgenin bütün metin denetimleri, belge bağlamında; denetim başına ayrı kayıt."""
    from .proofing import _doc_context as D
    from .proofing import checks
    doc = db.one("SELECT id, title, pages, audience, age_from, age_to FROM document_review WHERE id=%s", doc_id)
    if doc is None:
        raise KeyError(doc_id)
    mods = checks()
    token = D.use({"id": str(doc["id"]), "pages": doc["pages"], "audience": doc["audience"],
                   "age_from": doc["age_from"], "age_to": doc["age_to"], "title": doc["title"]})
    out = {}
    try:
        for name in CHECKS:
            mod, t0 = mods[name], time.time()
            try:
                res = await mod.run(doc_id)
                findings, stats = res if isinstance(res, tuple) else (res, {})
                out[name] = await asyncio.to_thread(_record, doc_id, mod, findings, stats, t0)
            except Exception as e:  # noqa: BLE001 - kaydedilir, öbür denetimler sürer
                await asyncio.to_thread(
                    db.one, "INSERT INTO document_run(document_id, check_name, check_version, status, error, started_at)"
                    " VALUES (%s,%s,%s,'FAILED',%s,to_timestamp(%s)) RETURNING id",
                    doc_id, name, str(getattr(mod, "VERSION", "?")), (str(e) + "\n" + traceback.format_exc())[-4000:], t0)
                out[name] = {"failed": str(e)[:300]}
    finally:
        D.reset(token)
    return out


def _claim() -> dict | None:
    return db.one("UPDATE document_review SET status='RUNNING', started_at=now(), error=NULL WHERE id = ("
                  " SELECT id FROM document_review WHERE status='QUEUED' ORDER BY created_at"
                  " FOR UPDATE SKIP LOCKED LIMIT 1) RETURNING id, title")


async def consume(poll_sec: float = 10.0) -> None:
    """Kuyruk: QUEUED belgeyi al, incele, DONE/FAILED yaz. Yarıda kalan (RUNNING, süreç öldü) belge başlarken
    yeniden kuyruğa alınır."""
    db.one("UPDATE document_review SET status='QUEUED' WHERE status='RUNNING' RETURNING id")
    while True:
        job = await asyncio.to_thread(_claim)
        if job is None:
            await asyncio.sleep(poll_sec)
            continue
        doc_id = str(job["id"])
        print(f"belge {doc_id} «{job['title']}» başladı", flush=True)
        try:
            out = await review(doc_id)
            failed = [k for k, v in out.items() if "failed" in v]
            await asyncio.to_thread(
                db.one, "UPDATE document_review SET status=%s, finished_at=now(), error=%s WHERE id=%s RETURNING id",
                "DONE", ("koşamayan denetim: " + ", ".join(failed)) if failed else None, doc_id)
            print(f"belge {doc_id} bitti {out}", flush=True)
        except Exception as e:  # noqa: BLE001
            await asyncio.to_thread(
                db.one, "UPDATE document_review SET status='FAILED', finished_at=now(), error=%s WHERE id=%s RETURNING id",
                (str(e) + "\n" + traceback.format_exc())[-4000:], doc_id)
            print(f"belge {doc_id} düştü: {e}", flush=True)


# ------------------------------------------------------------------ okuma (kart servisi)
def report(c, doc_id: str) -> dict | None:
    """Belgenin durumu, denetimleri ve bulguları (son okuma raporuyla aynı biçim)."""
    from .proofing._labels import label_of
    d = c.execute("SELECT id, title, file_name, format, page_kind, words, status, error, audience, age_from, age_to,"
                  " uploaded_by, created_at, started_at, finished_at, jsonb_array_length(pages) AS pages"
                  " FROM ed.document_review WHERE id=%s", (doc_id,)).fetchone()
    if d is None:
        return None
    runs = c.execute("SELECT DISTINCT ON (check_name) id, check_name, check_version, status, error, started_at, finished_at"
                     " FROM ed.document_run WHERE document_id=%s ORDER BY check_name, started_at DESC", (doc_id,)).fetchall()
    rows = c.execute("SELECT id, check_name, page_no, severity, message, quote, suggestion, details FROM ed.document_finding"
                     " WHERE run_id = ANY(%s) ORDER BY page_no NULLS FIRST, severity DESC, created_at",
                     ([r["id"] for r in runs],)).fetchall() if runs else []
    iso = lambda t: t.isoformat() if t else None  # noqa: E731
    count = {}
    for r in rows:
        n = count.setdefault(r["check_name"], [0, 0])
        n[0] += 1
        n[1] += r["severity"] != "INFO"
    return {"document": {**{k: d[k] for k in ("title", "file_name", "format", "page_kind", "words", "status", "error",
                                                "audience", "age_from", "age_to", "uploaded_by", "pages")},
                         "id": str(d["id"]), "created_at": iso(d["created_at"]), "finished_at": iso(d["finished_at"])},
            "checks": [{"name": r["check_name"], "label": label_of(r["check_name"]), "version": r["check_version"],
                        "status": r["status"], "error": r["error"], "started_at": iso(r["started_at"]),
                        "finished_at": iso(r["finished_at"]), "findings": count.get(r["check_name"], [0, 0])[0],
                        "serious": count.get(r["check_name"], [0, 0])[1], "precision": None} for r in runs],
            "findings": [{"id": str(r["id"]), "check": r["check_name"], "label": label_of(r["check_name"]),
                          "page": r["page_no"], "severity": r["severity"], "message": r["message"], "quote": r["quote"],
                          "suggestion": r["suggestion"], "bbox": None, "details": r["details"] or {},
                          "group": (r["details"] or {}).get("group"), "confidence": (r["details"] or {}).get("confidence"),
                          "marks": None, "decision": None} for r in rows]}


if __name__ == "__main__":
    import sys
    if sys.argv[1:2] == ["consume"]:
        asyncio.run(consume())
    elif sys.argv[1:2] == ["review"]:
        print(asyncio.run(review(sys.argv[2])))
    else:
        print("python -m editor.document_review consume | review <belge-id>")
        sys.exit(2)
