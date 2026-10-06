"""M27 Fuar takvimi (Excel): pazarlamanın elle tuttuğu «FUARLAR.xlsx» dosyası portala yüklenir, ekranda okunur.

Dosyanın biçimi sabit sayılmaz: başlık satırı «Fuar Tarihi» sütunundan bulunur; «Fuar Adı», «Gün sayısı», «Fuar Alanı»,
«Düzenleyen», «Katılımcı» başlıkları varsa okunur. Ad sütunu yoksa (Gantt sayfası: ad, başladığı günün sütununa yazılır)
satırın tarih sütunundan önceki ilk metni addır. Yıl, ilk sütundaki 2026/2027 gibi yıl satırlarından; yoksa aylar geriye
gidince bir artar. Birden çok sayfa varsa en çok fuar ve en çok alan okunan sayfa seçilir. «… netleşmedi» başlıklı hücrenin
altındaki adlar tarihi belli olmayan fuarlardır.

Portalın şablonu (`template`): «Fuarlar» sayfasında Fuar Adı · Başlangıç · Bitiş (tarih hücresi) · Fuar Alanı ·
Düzenleyen · Katılımcı (açılır liste), «Tarihi belli değil» sayfasında yalnız ad. Şablon güncel listeyle dolu iner;
okuyucu hem şablonu hem pazarlamanın eski dosyasını okur.

Kiracı başına tek takvim tutulur; yeni yükleme eskisinin yerine geçer. Fuar kartlarına (FAIRS) dokunmaz; CRM'e ve Logo'ya
yazılmaz.
"""

from __future__ import annotations

import io
import json
import re
import threading
import unicodedata
from datetime import date, datetime, timezone
from typing import Any, Optional

import sqlalchemy as sa

_md = sa.MetaData()
_lock = threading.Lock()
_ready: set[int] = set()

PLAN = sa.Table(
    "semantic_events_fair_plan", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("file_name", sa.String(300)),
    sa.Column("sheet", sa.String(120)),
    sa.Column("rows_json", sa.Text, nullable=False),
    sa.Column("pending_json", sa.Text),
    sa.Column("warnings_json", sa.Text),
    sa.Column("uploaded_by", sa.String(120), nullable=False),
    sa.Column("uploaded_at", sa.DateTime(timezone=True), nullable=False),
)

MONTHS = {"ocak": 1, "subat": 2, "mart": 3, "nisan": 4, "mayis": 5, "haziran": 6, "temmuz": 7, "agustos": 8,
          "eylul": 9, "ekim": 10, "kasim": 11, "aralik": 12}
PARTICIPANTS = {"timas": "TİMAŞ", "bayi": "Bayi"}


class PlanError(ValueError):
    pass


def ensure(engine: sa.engine.Engine) -> None:
    with _lock:
        if id(engine) in _ready:
            return
        from semantic_layer.store import schema_stamp
        schema_stamp.create_all(_md, engine)
        _ready.add(id(engine))


def _norm(v: Any) -> str:
    """Büyük/küçük ve Türkçe harf farkını siler: «EYLÜL» → «eylul», «TİMAŞ» → «timas»."""
    s = str(v or "").replace("İ", "i").replace("I", "ı").lower().replace("ı", "i")
    s = unicodedata.normalize("NFKD", s)
    return " ".join("".join(ch for ch in s if not unicodedata.combining(ch)).split())


def _text(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = " ".join(str(v).split()).replace("( ", "(").replace(" )", ")")
    return s or None


def _month(word: str) -> Optional[int]:
    return MONTHS.get(_norm(word))


_RANGE = re.compile(r"^\s*(\d{1,2})\s*([^\d\s-]+)?\s*[-–]\s*(\d{1,2})\s*([^\d\s-]+)\s*$")


def parse_range(v: Any, year: int) -> Optional[tuple[date, date]]:
    """«25 EYLÜL-04 EKİM», «02 - 11 EKİM», «20 KASIM-29 KASIM»; hücre tarih ise tek gün."""
    if isinstance(v, datetime):
        return v.date(), v.date()
    if isinstance(v, date):
        return v, v
    m = _RANGE.match(str(v or ""))
    if not m:
        return None
    d1, w1, d2, w2 = m.groups()
    m2 = _month(w2)
    m1 = _month(w1) if w1 else m2
    if not m1 or not m2:
        return None
    try:
        a = date(year, m1, int(d1))
        b = date(year + (1 if m2 < m1 else 0), m2, int(d2))
    except ValueError:
        return None
    return a, b


def _year_cell(v: Any) -> Optional[int]:
    if isinstance(v, (int, float)) and float(v).is_integer() and 2000 <= int(v) <= 2100:
        return int(v)
    return None


_HEADS = {
    "start": ("baslangic",),
    "end": ("bitis",),
    "date": ("fuar tarihi", "tarih"),
    "name": ("fuar adi", "fuar ismi"),
    "days": ("gun sayisi",),
    "venue": ("fuar alani", "mekan"),
    "organizer": ("duzenleyen",),
    "participant": ("katilimci",),
}


def _header(row: tuple) -> Optional[dict[str, int]]:
    cells = [_norm(c) if isinstance(c, str) else "" for c in row]
    found: dict[str, int] = {}
    for key, words in _HEADS.items():
        for i, c in enumerate(cells):
            if c and any(w in c for w in words):
                found.setdefault(key, i)
                break
    if "start" in found and "end" in found:           # şablon: ayrı başlangıç/bitiş tarih sütunu
        if found.get("date") in (found["start"], found["end"]):
            found.pop("date")
        return found
    found.pop("start", None)
    found.pop("end", None)
    return found if "date" in found else None


_DMY = re.compile(r"^\s*(\d{1,2})[./-](\d{1,2})[./-](\d{4})\s*$")


def _cell_date(v: Any) -> Optional[date]:
    """Tarih hücresi ya da «24.10.2026» / «2026-10-24» metni."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip()
    try:
        if m := _DMY.match(t):
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        return date.fromisoformat(t[:10]) if t else None
    except ValueError:
        return None


def _parse_sheet(rows: list[tuple], today: date) -> dict[str, Any]:
    head_i, cols = next(((i, h) for i, r in enumerate(rows) if (h := _header(r))), (None, None))
    pending = _pending(rows)
    if cols is None:
        return {"items": [], "fields": 0, "pending": pending, "warnings": []}
    split = "start" in cols
    dc = cols["start"] if split else cols["date"]
    year: Optional[int] = None
    last_month: Optional[int] = None
    items: list[dict[str, Any]] = []
    warnings: list[str] = []
    for n, r in enumerate(rows[head_i + 1:], start=head_i + 2):
        y = _year_cell(r[0]) if r else None
        if y:
            year, last_month = y, None
            continue
        raw = r[dc] if dc < len(r) else None
        if raw in (None, ""):
            continue
        if "name" in cols:
            name = _text(r[cols["name"]])
        else:
            name = next((t for c in r[1:dc] if isinstance(c, str) and (t := _text(c))), None)
        if not name:
            warnings.append(f"{n}. satırda tarih var ama fuar adı yok; atlandı.")
            continue
        if split:
            a = _cell_date(raw)
            b = _cell_date(r[cols["end"]] if cols["end"] < len(r) else None) if a else None
            if not a or not b or b < a:
                why = "başlangıç tarihi okunamadı" if not a else "bitiş tarihi okunamadı" if not b else "bitiş başlangıçtan önce"
                warnings.append(f"{n}. satır «{name}»: {why}; atlandı.")
                continue
            rng: Optional[tuple[date, date]] = (a, b)
        else:
            y0 = year or today.year
            rng = parse_range(raw, y0)
        if not split and rng and year is None and last_month and rng[0].month < last_month:
            y0 += 1                                     # yıl satırı yok: ay geriye gitti → ertesi yıl
            year = y0
            rng = parse_range(raw, y0)
        if not rng:
            warnings.append(f"«{name}»: tarih okunamadı ({_text(raw)}).")
            continue
        last_month = rng[0].month
        get = lambda k: _text(r[cols[k]]) if k in cols and cols[k] < len(r) else None  # noqa: E731
        part = get("participant")
        pkey = _norm(part) if part else None
        days = (rng[1] - rng[0]).days + 1
        stated = get("days")
        if stated and (m := re.search(r"\d+", stated)) and int(m.group()) != days:
            warnings.append(f"«{name}»: dosyada {m.group()} gün yazıyor, tarihe göre {days} gün.")
        items.append({
            "name": name, "startsOn": rng[0].isoformat(), "endsOn": rng[1].isoformat(), "days": days,
            "dateText": _text(raw) if not split and not isinstance(raw, (date, datetime)) else None,
            "venue": get("venue"), "organizer": get("organizer"),
            "participant": pkey if pkey in PARTICIPANTS else (pkey or None),
            "participantLabel": PARTICIPANTS.get(pkey or "", part),
        })
    fields = sum(1 for k in ("venue", "organizer", "participant", "days", "start") if k in cols)
    return {"items": items, "fields": fields, "pending": pending, "warnings": warnings}


def _pending(rows: list[tuple]) -> Optional[dict[str, Any]]:
    """«2027 fuarlar tarih netleşmedi» gibi bir hücre: altındaki dolu hücreler (ilk boşluğa kadar)."""
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            if isinstance(c, str) and "netlesmedi" in _norm(c):
                names = []
                for r2 in rows[i + 1:]:
                    t = _text(r2[j]) if j < len(r2) else None
                    if not t:
                        break
                    names.append(t)
                return {"title": _text(c), "items": names}
    return None


def parse_workbook(data: bytes, today: date) -> dict[str, Any]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise PlanError("Dosya Excel (.xlsx) olarak okunamadı.") from None
    best: Optional[tuple[str, dict[str, Any]]] = None
    pending = None
    for name in wb.sheetnames:
        rows = [tuple(r) for r in wb[name].iter_rows(values_only=True)]
        if any(w in _norm(name) for w in ("belli degil", "netles")):   # şablonun «Tarihi belli değil» sayfası
            names = [t for r in rows[1:] if r and (t := _text(r[0]))]
            pending = pending or {"title": name, "items": names}
            continue
        p = _parse_sheet(rows, today)
        pending = pending or p["pending"]
        if best is None or (len(p["items"]), p["fields"]) > (len(best[1]["items"]), best[1]["fields"]):
            best = (name, p)
    if best is None or not best[1]["items"]:
        raise PlanError("Dosyada «Fuar Tarihi» başlıklı bir sütun ve altında okunabilir tarih («25 EYLÜL-04 EKİM») bulunamadı.")
    sheet, p = best
    items = sorted(p["items"], key=lambda x: (x["startsOn"], x["endsOn"], x["name"]))
    return {"sheet": sheet, "items": items, "pending": pending, "warnings": p["warnings"]}


def save(engine: sa.engine.Engine, tenant: str, user: str, file_name: str, parsed: dict[str, Any]) -> None:
    ensure(engine)
    vals = {"file_name": (file_name or "")[:300] or None, "sheet": parsed["sheet"][:120],
            "rows_json": json.dumps(parsed["items"], ensure_ascii=False),
            "pending_json": json.dumps(parsed["pending"], ensure_ascii=False) if parsed["pending"] else None,
            "warnings_json": json.dumps(parsed["warnings"], ensure_ascii=False),
            "uploaded_by": user, "uploaded_at": datetime.now(timezone.utc)}
    with engine.begin() as c:
        c.execute(PLAN.delete().where(PLAN.c.tenant_id == tenant))
        c.execute(PLAN.insert().values(tenant_id=tenant, **vals))


def _phase(it: dict[str, Any], today: date) -> tuple[str, int]:
    a, b = date.fromisoformat(it["startsOn"]), date.fromisoformat(it["endsOn"])
    if b < today:
        return "bitti", (b - today).days
    if a <= today:
        return "suruyor", (b - today).days            # bitişe kalan gün
    return "yaklasan", (a - today).days


def read(engine: sa.engine.Engine, tenant: str, today: date) -> dict[str, Any]:
    ensure(engine)
    with engine.connect() as c:
        r = c.execute(sa.select(PLAN).where(PLAN.c.tenant_id == tenant)).first()
    if r is None:
        return {"items": [], "pending": None, "warnings": [], "fileName": None, "sheet": None, "uploadedBy": None,
                "uploadedAt": None, "today": today.isoformat()}
    items = json.loads(r.rows_json or "[]")
    for i, it in enumerate(items):
        it["id"] = str(i + 1)
        it["phase"], it["daysLeft"] = _phase(it, today)
    at = r.uploaded_at if r.uploaded_at.tzinfo else r.uploaded_at.replace(tzinfo=timezone.utc)
    return {"items": items, "pending": json.loads(r.pending_json) if r.pending_json else None,
            "warnings": json.loads(r.warnings_json or "[]"), "fileName": r.file_name, "sheet": r.sheet,
            "uploadedBy": r.uploaded_by, "uploadedAt": at.isoformat(), "today": today.isoformat()}


TEMPLATE_HEAD = ["Fuar Adı", "Başlangıç", "Bitiş", "Fuar Alanı", "Düzenleyen", "Katılımcı"]
PENDING_SHEET = "Tarihi belli değil"


def template(plan: dict[str, Any]) -> bytes:
    """Doldurulacak şablon: güncel listeyle dolu; tarih sütunları tarih hücresi, katılımcı açılır liste."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Fuarlar"
    ws.append(TEMPLATE_HEAD)
    head_fill = PatternFill("solid", fgColor="5B4BDB")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    for it in plan.get("items") or []:
        ws.append([it["name"], date.fromisoformat(it["startsOn"]), date.fromisoformat(it["endsOn"]), it.get("venue"),
                   it.get("organizer"), PARTICIPANTS.get(it.get("participant") or "", it.get("participantLabel"))])
    last = max(ws.max_row, 2) + 300                    # boş satırlara da biçim ve liste
    for row in range(2, last + 1):
        for col in ("B", "C"):
            ws[f"{col}{row}"].number_format = "DD.MM.YYYY"
    who = DataValidation(type="list", formula1='"TİMAŞ,Bayi"', allow_blank=True, showErrorMessage=True,
                         errorTitle="Katılımcı", error="Listeden seçin: TİMAŞ ya da Bayi.")
    days = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True,
                          showErrorMessage=True, errorTitle="Tarih", error="Tarih girin, ör. 24.10.2026.")
    ws.add_data_validation(who)
    ws.add_data_validation(days)
    who.add(f"F2:F{last}")
    days.add(f"B2:C{last}")
    for col, w in zip("ABCDEF", (36, 13, 13, 34, 24, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    pw = wb.create_sheet(PENDING_SHEET)
    pw.append(["Fuar Adı"])
    pw["A1"].font = Font(bold=True, color="FFFFFF")
    pw["A1"].fill = head_fill
    pw.column_dimensions["A"].width = 40
    for n in ((plan.get("pending") or {}).get("items") or []):
        pw.append([n])

    hw = wb.create_sheet("Nasıl doldurulur")
    hw.column_dimensions["A"].width = 100
    for line in (
        "Fuar takvimi şablonu",
        "",
        "• «Fuarlar» sayfasında her satır bir fuar. Fuar Adı, Başlangıç ve Bitiş zorunlu.",
        "• Başlangıç ve Bitiş tarih olarak girilir: 24.10.2026. Bitiş başlangıçtan önce olamaz.",
        "• Katılımcı listeden seçilir: TİMAŞ (kendi standımız) ya da Bayi.",
        "• Fuar Alanı ve Düzenleyen isteğe bağlı.",
        "• Tarihi henüz belli olmayan fuarları «Tarihi belli değil» sayfasına yalnız adıyla yazın.",
        "• Sütun başlıklarını değiştirmeyin. Satır sırası önemli değil; ekranda tarihe göre dizilir.",
        "• Yüklenen dosya önceki listenin yerine geçer: listede olmayan fuar ekrandan kalkar.",
    ):
        hw.append([line])
    hw["A1"].font = Font(bold=True, size=13)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
