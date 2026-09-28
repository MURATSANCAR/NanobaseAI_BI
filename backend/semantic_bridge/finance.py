"""M45 Finansal raporlama ve analiz: aylık gelir tablosu (hesap eşlemesi muhasebe onaylı, rakamdan Logo fişine
iniş), bütçe–gerçekleşme (M46'dan), kitap/seri/yayınevi/kanal/cari kârlılığı, 13 haftalık nakit, vergi takvimi,
aylık kapanış ve özet.

**Rakamı model üretmez.** Bütün tutarlar Logo okumasının (`finance_sources`) köprü tablolarındaki anlık görüntüsünden
deterministik toplanır. Zeki AI yalnız eşlenmemiş hesaba kapalı kümeden aday satır önerir (`QueuedLlm.choose`);
öneri muhasebe onaylamadan «onaylı» olmaz.

**Gelir tablosu:** tutar = hesabın kâr etkisi (alacak − borç): gelir artı, gider eksi; ara toplamlar alt satırların
toplamıdır. Hesap → satır eşlemesi en özel koddan genele çözülür (`600.01.001` → `600.01` → `600`): önce onaylı/dışlanmış
kayıt, sonra Zeki AI önerisi, sonra Tekdüzen hesap planı kuralı (`DEFAULT_RULES`, 3 haneli grup). Hiçbirine düşmeyen
hesap «Eşlenmemiş hesaplar» satırında toplanır ve net kâra katılır: para sessizce kaybolmaz, uyarı olarak kalır.

**«Yaklaşık» işaretleri:** maliyeti işlenmemiş satış satırları (maliyet Logo'da `OUTCOST = 0`) marja katılmaz; M9 birim
maliyetiyle doldurulan tutar ayrı sütunda «yaklaşık» yazar, birim maliyeti de bilinmeyen satış «maliyet bilinmiyor»
olarak kalır. Telif CRM'deki yürürlükteki sözleşmenin oranından tahakkuk eder (yaklaşık); oranı olmayan kitapta «telif
verisi yok». Vade dağılımı FIFO yaklaşımıdır. Her cevapta veri son günü (`veriSonu`) vardır; Logo kopyası donmuşsa
ekranda o tarih yazar, «bugün» değil.

Kendi tabloları `semantic_finance_*`. Logo'ya, CRM'e yazılmaz. Değişiklikler `semantic_audit`'e (API katmanı).
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import math
import os
import threading
import time
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, Iterable, Optional
from zoneinfo import ZoneInfo

import sqlalchemy as sa

from semantic_bridge import finance_sources as src

log = logging.getLogger("semantic.finance")
TZ = ZoneInfo("Europe/Istanbul")
_md = sa.MetaData()
_MONEY = sa.Float

AY = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
PAGE_SIZE = 100
UNMAPPED = "ESLENMEMIS"
EXCLUDED = "dislandi"

LINES = sa.Table(
    "semantic_finance_lines", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("kod", sa.String(40), primary_key=True),
    sa.Column("ad", sa.String(200), nullable=False),
    sa.Column("sira", sa.Integer, nullable=False),
    sa.Column("ust_kod", sa.String(40)),
    sa.Column("isaret", sa.Integer, nullable=False),            # +1 gelir, −1 gider (gösterim notu)
    sa.Column("tur", sa.String(16), nullable=False),            # gelir | gider | ara_toplam | eslenmemis
    sa.Column("formul_json", sa.Text),                          # ara toplamın topladığı satır kodları
)
ACCOUNT_MAP = sa.Table(
    "semantic_finance_account_map", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("hesap_kodu", sa.String(60), primary_key=True),
    sa.Column("satir_kodu", sa.String(40)),
    sa.Column("durum", sa.String(12), nullable=False),          # oneri | onayli | dislandi
    sa.Column("oneri_satir", sa.String(40)),
    sa.Column("oneri_olasilik", sa.Float),
    sa.Column("oneri_kaynak", sa.String(20)),                   # zeki | hesap-plani
    sa.Column("oneri_json", sa.Text),
    sa.Column("onaylayan", sa.String(120)),
    sa.Column("onay_tarihi", sa.DateTime(timezone=True)),
    sa.Column("not_", sa.String(500)),
    sa.Column("updated_by", sa.String(120)),
    sa.Column("updated_at", sa.DateTime(timezone=True)),
)
ACCOUNTS = sa.Table(
    "semantic_finance_accounts", _md,
    sa.Column("hesap_kodu", sa.String(60), primary_key=True),
    sa.Column("ad", sa.String(300)),
)
ACTUALS = sa.Table(
    "semantic_finance_account_actuals", _md,
    sa.Column("year", sa.Integer, primary_key=True),
    sa.Column("month", sa.Integer, primary_key=True),
    sa.Column("hesap_kodu", sa.String(60), primary_key=True),
    sa.Column("merkez_kodu", sa.String(60), primary_key=True),
    sa.Column("kural", sa.String(12), primary_key=True),        # dahil | yansitma | kapanis
    sa.Column("hesap_adi", sa.String(300)),
    sa.Column("borc", _MONEY, nullable=False),
    sa.Column("alacak", _MONEY, nullable=False),
    sa.Column("satir", sa.Integer, nullable=False),
)
_MEASURES = ("adet", "brut", "net", "maliyet", "maliyetli_net", "maliyetsiz_adet", "maliyetsiz_net")
SALES_MONTH = sa.Table(
    "semantic_finance_sales_month", _md,
    sa.Column("year", sa.Integer, primary_key=True),
    sa.Column("month", sa.Integer, primary_key=True),
    *[sa.Column(m, _MONEY, nullable=False) for m in _MEASURES],
    sa.Column("satis_net", _MONEY, nullable=False),
    sa.Column("iade_net", _MONEY, nullable=False),
    sa.Column("satis_satir", sa.Integer, nullable=False),
    sa.Column("maliyetsiz_satir", sa.Integer, nullable=False),
)
PROFIT_ITEMS = sa.Table(
    "semantic_finance_profit_items", _md,
    sa.Column("year", sa.Integer, primary_key=True),
    sa.Column("month", sa.Integer, primary_key=True),
    sa.Column("stok_kodu", sa.String(60), primary_key=True),
    sa.Column("kanal", sa.String(120), primary_key=True),
    *[sa.Column(m, _MONEY, nullable=False) for m in _MEASURES],
)
PROFIT_CLIENTS = sa.Table(
    "semantic_finance_profit_clients", _md,
    sa.Column("year", sa.Integer, primary_key=True),
    sa.Column("month", sa.Integer, primary_key=True),
    sa.Column("cari_kodu", sa.String(60), primary_key=True),
    sa.Column("kanal", sa.String(120), primary_key=True),
    sa.Column("cari_adi", sa.String(300)),
    *[sa.Column(m, _MONEY, nullable=False) for m in _MEASURES],
    sa.Column("tahmini_maliyet", _MONEY, nullable=False, default=0.0),   # maliyetsiz adet × M9 birim maliyeti
    sa.Column("tahmini_adet", _MONEY, nullable=False, default=0.0),      # birim maliyeti bilinen maliyetsiz adet
)
CLOSES = sa.Table(
    "semantic_finance_closes", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("year", sa.Integer, primary_key=True),
    sa.Column("month", sa.Integer, primary_key=True),
    sa.Column("durum", sa.String(10), nullable=False),          # acik | kapandi
    sa.Column("kapatan", sa.String(120)),
    sa.Column("kapanis_tarihi", sa.DateTime(timezone=True)),
    sa.Column("ozet_hash", sa.String(64)),
    sa.Column("ozet_json", sa.Text),
    sa.Column("not_", sa.Text),
    sa.Column("acan", sa.String(120)),
    sa.Column("acilis_tarihi", sa.DateTime(timezone=True)),
    sa.Column("acilis_gerekce", sa.Text),
)
CASH_RUNS = sa.Table(
    "semantic_finance_cash_runs", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("run_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("veri_son_gunu", sa.String(10)),
    sa.Column("baslangic", sa.String(10), nullable=False),      # ilk haftanın pazartesisi
    sa.Column("params_json", sa.Text, nullable=False),          # açılış bakiyesi, vadesi geçmiş, kaynak durumları
    sa.Column("run_by", sa.String(120)),
)
CASH_LINES = sa.Table(
    "semantic_finance_cash_lines", _md,
    sa.Column("run_id", sa.String(32), primary_key=True),
    sa.Column("hafta", sa.Integer, primary_key=True),           # 1..13
    sa.Column("kalem", sa.String(40), primary_key=True),
    sa.Column("hafta_baslangic", sa.String(10), nullable=False),
    sa.Column("yon", sa.String(8), nullable=False),             # giris | cikis | bilgi
    sa.Column("kaynak", sa.String(200), nullable=False),
    sa.Column("tutar", _MONEY, nullable=False),
    sa.Column("yaklasik", sa.Boolean, nullable=False),
    sa.Column("ayrinti_json", sa.Text),
)
TAX = sa.Table(
    "semantic_finance_tax_calendar", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("beyan", sa.String(200), nullable=False),
    sa.Column("donem", sa.String(40)),
    sa.Column("son_gun", sa.String(10), nullable=False),
    sa.Column("sorumlu", sa.String(200)),
    sa.Column("durum", sa.String(16), nullable=False),          # bekliyor | hazirlaniyor | hazir | verildi
    sa.Column("tutar", _MONEY),                                  # tahmini ödeme (nakit tablosuna çıkış)
    sa.Column("not_", sa.Text),
    sa.Column("hatirlatma_json", sa.Text),                       # gönderilen hatırlatmalar {"7": iso, "2": iso}
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("updated_by", sa.String(120)),
    sa.Column("updated_at", sa.DateTime(timezone=True)),
)
NOTES = sa.Table(
    "semantic_finance_notes", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("year", sa.Integer, nullable=False),
    sa.Column("month", sa.Integer),
    sa.Column("tur", sa.String(10), nullable=False),            # sapma | ozet
    sa.Column("hedef_anahtar", sa.String(200)),
    sa.Column("metin", sa.Text, nullable=False),
    sa.Column("durum", sa.String(10), nullable=False),          # taslak | onayli
    sa.Column("llm_job_id", sa.String(64)),
    sa.Column("hazirlayan", sa.String(120), nullable=False),
    sa.Column("onaylayan", sa.String(120)),
    sa.Column("tarih", sa.DateTime(timezone=True), nullable=False),
)
META = sa.Table(
    "semantic_finance_meta", _md,
    sa.Column("key", sa.String(60), primary_key=True),
    sa.Column("value_json", sa.Text, nullable=False),
    sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
)

#: Gelir tablosu (Tekdüzen hesap planı sırası). (kod, ad, üst, işaret, tür, formül)
DEFAULT_LINES: list[tuple[str, str, Optional[str], int, str, Optional[list[str]]]] = [
    ("BRUT_SATIS", "A. Brüt satışlar", None, 1, "gelir", None),
    ("SATIS_INDIRIM", "B. Satış indirimleri (−)", None, -1, "gider", None),
    ("NET_SATIS", "C. Net satışlar", None, 1, "ara_toplam", ["BRUT_SATIS", "SATIS_INDIRIM"]),
    ("SMM", "D. Satışların maliyeti (−)", None, -1, "gider", None),
    ("BRUT_KAR", "Brüt satış kârı veya zararı", None, 1, "ara_toplam", ["NET_SATIS", "SMM"]),
    ("ARGE", "Araştırma ve geliştirme giderleri (−)", "FAALIYET", -1, "gider", None),
    ("PSD", "Pazarlama, satış ve dağıtım giderleri (−)", "FAALIYET", -1, "gider", None),
    ("GYG", "Genel yönetim giderleri (−)", "FAALIYET", -1, "gider", None),
    ("FAALIYET", "E. Faaliyet giderleri (−)", None, -1, "ara_toplam", ["ARGE", "PSD", "GYG"]),
    ("FAALIYET_KARI", "Faaliyet kârı veya zararı", None, 1, "ara_toplam", ["BRUT_KAR", "FAALIYET"]),
    ("DIGER_GELIR", "F. Diğer faaliyetlerden olağan gelir ve kârlar", None, 1, "gelir", None),
    ("DIGER_GIDER", "G. Diğer faaliyetlerden olağan gider ve zararlar (−)", None, -1, "gider", None),
    ("FINANSMAN", "H. Finansman giderleri (−)", None, -1, "gider", None),
    ("OLAGAN_KAR", "Olağan kâr veya zarar", None, 1, "ara_toplam", ["FAALIYET_KARI", "DIGER_GELIR", "DIGER_GIDER", "FINANSMAN"]),
    ("OLAGANDISI_GELIR", "I. Olağandışı gelir ve kârlar", None, 1, "gelir", None),
    ("OLAGANDISI_GIDER", "J. Olağandışı gider ve zararlar (−)", None, -1, "gider", None),
    (UNMAPPED, "Eşlenmemiş hesaplar (eşleme bekliyor)", None, 1, "eslenmemis", None),
    ("DONEM_KARI", "Dönem kârı veya zararı", None, 1, "ara_toplam", ["OLAGAN_KAR", "OLAGANDISI_GELIR", "OLAGANDISI_GIDER", UNMAPPED]),
    ("VERGI", "K. Dönem kârı vergi ve yasal yükümlülük karşılıkları (−)", None, -1, "gider", None),
    ("NET_KAR", "Dönem net kârı veya zararı", None, 1, "ara_toplam", ["DONEM_KARI", "VERGI"]),
]

#: Tekdüzen hesap planı: 3 haneli grup → satır (ya da dışlanır). Hesap planı kuralı öneridir; muhasebe onaylar.
DEFAULT_RULES: dict[str, str] = {
    **{g: "BRUT_SATIS" for g in ("600", "601", "602")},
    **{g: "SATIS_INDIRIM" for g in ("610", "611", "612")},
    **{g: "SMM" for g in ("620", "621", "622", "623")},
    "630": "ARGE", "631": "PSD", "632": "GYG",
    **{str(g): "DIGER_GELIR" for g in range(640, 650)},
    **{str(g): "DIGER_GIDER" for g in range(653, 660)},
    "660": "FINANSMAN", "661": "FINANSMAN",
    "671": "OLAGANDISI_GELIR", "679": "OLAGANDISI_GELIR",
    "680": "OLAGANDISI_GIDER", "681": "OLAGANDISI_GIDER", "689": "OLAGANDISI_GIDER",
    "691": "VERGI",
    **{g: EXCLUDED for g in ("690", "692", "697", "698")},
    # 7/A maliyet hesapları: gider bir kez 7xx'te sayılır (63x'e yansıtma fişi `yansitma` kuralıyla dışarıda).
    "740": "SMM", "750": "ARGE", "760": "PSD", "770": "GYG", "780": "FINANSMAN",
    # Üretim giderleri stoka gider (151/152); gelir tablosuna satılan malın maliyetiyle girer.
    **{g: EXCLUDED for g in ("710", "720", "730")},
    **{g: EXCLUDED for g in ("711", "721", "731", "741", "751", "761", "771", "781", "791")},
}
RULE_NOTES = {
    **{g: "Dönem sonu kapanış hesabı" for g in ("690", "692", "697", "698")},
    **{g: "Üretim gideri stoka gider; gelir tablosuna satılan malın maliyetiyle girer" for g in ("710", "720", "730")},
    **{g: "Yansıtma hesabı" for g in ("711", "721", "731", "741", "751", "761", "771", "781", "791")},
}

TAX_STATUSES = {"bekliyor": "Bekliyor", "hazirlaniyor": "Hazırlanıyor", "hazir": "Hazır", "verildi": "Verildi"}
GRAINS = {"ay": "Ay", "ceyrek": "Çeyrek", "ytd": "Yıl başından"}
PROFIT_BY = {"kitap": "Kitap", "seri": "Seri / kitaplık", "yayinevi": "Yayınevi", "kanal": "Kanal", "cari": "Cari"}

_ready: set[int] = set()
_lock = threading.Lock()


class FinanceError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def ensure(engine: sa.engine.Engine) -> None:
    key = id(engine)
    with _lock:
        if key in _ready:
            return
        _md.create_all(engine, checkfirst=True)
        _ready.add(key)


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _today() -> date:
    return datetime.now(TZ).date()


def _iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    return (v if v.tzinfo else v.replace(tzinfo=timezone.utc)).isoformat()


def _j(v: Optional[str], default: Any) -> Any:
    try:
        return json.loads(v) if v else default
    except ValueError:
        return default


def _dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str, sort_keys=True)


def _r(v: Optional[float], d: int = 2) -> Optional[float]:
    return None if v is None else round(float(v), d)


def _text(v: Any, limit: int) -> Optional[str]:
    s = " ".join(str(v or "").split())
    return s[:limit] or None


def conf(key: str, default: str) -> str:
    return src._conf(key, default)


# ------------------------------------------------------------------ meta


def meta_get(engine: sa.engine.Engine, key: str) -> dict[str, Any]:
    with engine.connect() as c:
        row = c.execute(sa.select(META).where(META.c.key == key)).first()
    if not row:
        return {}
    return {**_j(row.value_json, {}), "_at": _iso(row.updated_at)}


def meta_set(engine: sa.engine.Engine, key: str, value: dict[str, Any]) -> None:
    now = _now()
    with engine.begin() as c:
        if c.execute(sa.select(META.c.key).where(META.c.key == key)).first():
            c.execute(META.update().where(META.c.key == key).values(value_json=_dump(value), updated_at=now))
        else:
            c.execute(META.insert().values(key=key, value_json=_dump(value), updated_at=now))


def data_end(engine: sa.engine.Engine) -> Optional[date]:
    v = meta_get(engine, "data_end").get("date")
    return date.fromisoformat(v) if v else None


def freshness(engine: sa.engine.Engine) -> dict[str, Any]:
    """Her cevaba eklenen veri tazeliği: veri son günü, maliyeti işlenmiş son gün, son okuma."""
    de = meta_get(engine, "data_end")
    return {"veriSonu": de.get("date"), "maliyetSonu": de.get("costed"), "muhasebeSonu": de.get("ledger"),
            "okundu": de.get("_at")}


# ------------------------------------------------------------------ satırlar ve eşleme


def lines(engine: sa.engine.Engine, tenant: str) -> list[dict[str, Any]]:
    """Gelir tablosu satırları; kiracı için ilk çağrıda Tekdüzen sırasıyla kurulur."""
    with engine.begin() as c:
        rows = c.execute(sa.select(LINES).where(LINES.c.tenant_id == tenant).order_by(LINES.c.sira)).all()
        if not rows:
            c.execute(LINES.insert(), [{"tenant_id": tenant, "kod": k, "ad": ad, "sira": i, "ust_kod": ust, "isaret": s,
                                        "tur": tur, "formul_json": _dump(fm) if fm else None}
                                       for i, (k, ad, ust, s, tur, fm) in enumerate(DEFAULT_LINES)])
            rows = c.execute(sa.select(LINES).where(LINES.c.tenant_id == tenant).order_by(LINES.c.sira)).all()
    return [{"kod": r.kod, "ad": r.ad, "sira": r.sira, "ust": r.ust_kod, "isaret": r.isaret, "tur": r.tur,
             "formul": _j(r.formul_json, None)} for r in rows]


def leaf_codes(ls: list[dict[str, Any]]) -> list[str]:
    return [x["kod"] for x in ls if x["tur"] in ("gelir", "gider")]


def _ancestors(code: str) -> list[str]:
    """`600.01.001` → [`600.01.001`, `600.01`, `600`] (en özelden genele)."""
    out = [code]
    cur = code
    while "." in cur:
        cur = cur.rsplit(".", 1)[0]
        out.append(cur)
    if len(code) > 3 and code[:3] not in out:
        out.append(code[:3])
    return out


def _stored_map(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    with engine.connect() as c:
        return {r.hesap_kodu: r for r in c.execute(sa.select(ACCOUNT_MAP).where(ACCOUNT_MAP.c.tenant_id == tenant)).all()}


def resolve(code: str, stored: dict[str, Any]) -> dict[str, Any]:
    """Hesabın geçerli satırı. `durum`: onayli | dislandi | oneri | yok; `kaynak`: kayıt kodu ya da kural."""
    chain = _ancestors(code)
    for c in chain:
        r = stored.get(c)
        if r is not None and r.durum in ("onayli", EXCLUDED):
            return {"satir": r.satir_kodu if r.durum == "onayli" else EXCLUDED, "durum": r.durum, "kaynak": "onay",
                    "kayit": c, "onaylayan": r.onaylayan, "onayTarihi": _iso(r.onay_tarihi), "not": r.not_}
    for c in chain:
        r = stored.get(c)
        if r is not None and r.oneri_satir:
            return {"satir": r.oneri_satir, "durum": "oneri", "kaynak": r.oneri_kaynak or "zeki", "kayit": c,
                    "olasilik": r.oneri_olasilik}
    rule = DEFAULT_RULES.get(code[:3])
    if rule:
        return {"satir": rule, "durum": "oneri", "kaynak": "hesap-plani", "kayit": code[:3], "not": RULE_NOTES.get(code[:3])}
    return {"satir": None, "durum": "yok", "kaynak": None, "kayit": None}


def _line_of(code: str, stored: dict[str, Any], cache: dict[str, dict[str, Any]]) -> dict[str, Any]:
    if code not in cache:
        cache[code] = resolve(code, stored)
    return cache[code]


# ------------------------------------------------------------------ dönem


def period_months(year: int, month: int, grain: str) -> list[tuple[int, int]]:
    if not 1 <= month <= 12:
        raise FinanceError("Ay 1 ile 12 arasında olmalı.")
    if grain == "ay":
        return [(year, month)]
    if grain == "ceyrek":
        q0 = (month - 1) // 3 * 3 + 1
        return [(year, m) for m in range(q0, q0 + 3)]
    if grain == "ytd":
        return [(year, m) for m in range(1, month + 1)]
    raise FinanceError("Dönem türü ay, çeyrek ya da yıl başından olmalı.")


def previous_period(year: int, month: int, grain: str) -> Optional[list[tuple[int, int]]]:
    if grain == "ay":
        return [(year - 1, 12)] if month == 1 else [(year, month - 1)]
    if grain == "ceyrek":
        q0 = (month - 1) // 3 * 3 + 1
        start = (year * 12 + q0 - 1) - 3
        return [((start + i) // 12, (start + i) % 12 + 1) for i in range(3)]
    return None


def period_label(months: list[tuple[int, int]]) -> str:
    if not months:
        return ""
    (y1, m1), (y2, m2) = months[0], months[-1]
    if (y1, m1) == (y2, m2):
        return f"{AY[m1 - 1]} {y1}"
    if y1 == y2:
        return f"{AY[m1 - 1]}–{AY[m2 - 1]} {y1}"
    return f"{AY[m1 - 1]} {y1} – {AY[m2 - 1]} {y2}"


def _months_cond(table: sa.Table, months: list[tuple[int, int]]):
    by_year: dict[int, list[int]] = {}
    for y, m in months:
        by_year.setdefault(y, []).append(m)
    return sa.or_(*[sa.and_(table.c.year == y, table.c.month.in_(ms)) for y, ms in by_year.items()])


def _loaded_years(engine: sa.engine.Engine) -> set[int]:
    with engine.connect() as c:
        keys = [r.key for r in c.execute(sa.select(META.c.key).where(META.c.key.like("ledger:%"))).all()]
    return {int(k.split(":")[1]) for k in keys if k.split(":")[1].isdigit()}


# ------------------------------------------------------------------ gelir tablosu


def account_totals(engine: sa.engine.Engine, months: list[tuple[int, int]], kural: str = "dahil") -> dict[str, dict[str, Any]]:
    """Hesap → {ad, borc, alacak, etki (alacak − borç), satir}."""
    if not months:
        return {}
    with engine.connect() as c:
        rows = c.execute(sa.select(ACTUALS.c.hesap_kodu, sa.func.max(ACTUALS.c.hesap_adi), sa.func.sum(ACTUALS.c.borc),
                                   sa.func.sum(ACTUALS.c.alacak), sa.func.sum(ACTUALS.c.satir))
                         .where(_months_cond(ACTUALS, months), ACTUALS.c.kural == kural)
                         .group_by(ACTUALS.c.hesap_kodu)).all()
    return {r[0]: {"ad": r[1], "borc": float(r[2] or 0), "alacak": float(r[3] or 0),
                   "etki": float(r[3] or 0) - float(r[2] or 0), "satir": int(r[4] or 0)} for r in rows}


def compute_lines(engine: sa.engine.Engine, tenant: str, months: list[tuple[int, int]],
                  stored: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    """Satır kodu → kâr etkisi; eşleme durumuna göre onaysız tutar; dışlanan ve kural dışı tutarlar."""
    ls = lines(engine, tenant)
    stored = _stored_map(engine, tenant) if stored is None else stored
    cache: dict[str, dict[str, Any]] = {}
    val = {x["kod"]: 0.0 for x in ls}
    unapproved = {x["kod"]: 0.0 for x in ls}
    accounts = {x["kod"]: 0 for x in ls}
    excluded = 0.0
    for code, a in account_totals(engine, months).items():
        res = _line_of(code, stored, cache)
        target = res["satir"]
        if target == EXCLUDED:
            excluded += a["etki"]
            continue
        if not target or target not in val:
            target = UNMAPPED
        val[target] += a["etki"]
        accounts[target] += 1
        if res["durum"] != "onayli":
            unapproved[target] += a["etki"]
    by_code = {x["kod"]: x for x in ls}

    def total(k: str, seen: frozenset = frozenset()) -> float:
        x = by_code.get(k)
        if x is None or k in seen:
            return 0.0
        if x["tur"] == "ara_toplam":
            return sum(total(c, seen | {k}) for c in (x["formul"] or []))
        return val.get(k, 0.0)

    out = {k: total(k) for k in by_code}
    un = {}
    for k, x in by_code.items():
        leaves = _leaves(k, by_code)
        un[k] = sum(unapproved.get(c, 0.0) for c in leaves)
        accounts[k] = sum(accounts.get(c, 0) for c in leaves) if x["tur"] == "ara_toplam" else accounts[k]
    rules = {}
    for kural in ("yansitma", "kapanis"):
        t = account_totals(engine, months, kural)
        rules[kural] = {"hesap": len(t), "borc": _r(sum(v["borc"] for v in t.values())),
                        "alacak": _r(sum(v["alacak"] for v in t.values()))}
    return {"values": out, "unapproved": un, "accounts": accounts, "excluded": excluded, "rules": rules}


def _leaves(k: str, by_code: dict[str, dict[str, Any]], seen: frozenset = frozenset()) -> list[str]:
    x = by_code.get(k)
    if x is None or k in seen:
        return []
    if x["tur"] != "ara_toplam":
        return [k]
    return [c for f in (x["formul"] or []) for c in _leaves(f, by_code, seen | {k})]


def budget_values(engine: sa.engine.Engine, tenant: str, months: list[tuple[int, int]],
                  stored: Optional[dict[str, Any]] = None) -> Optional[dict[str, Any]]:
    """M46 yürürlükteki planından dönem bütçesi: net satış hedefi (planın aylık ciro dağılımı) ve departman gider
    bütçesi (7xx ana hesap → eşlemeyle satır). Onaylı plan yoksa None. Kâr etkisi işaretiyle (gider eksi)."""
    try:
        from semantic_bridge import budget as B
    except Exception:  # noqa: BLE001
        return None
    years = sorted({y for y, _ in months})
    if len(years) != 1:
        return None
    year = years[0]
    ms = [m for _, m in months]
    try:
        B.ensure(engine)
        with engine.connect() as c:
            row = B._approved(c, tenant, year)
            if row is None:
                return None
            tot = B._totals(c, row.id)
            depts = c.execute(sa.select(B.DEPTS).where(B.DEPTS.c.plan_id == row.id)).all()
    except Exception as e:  # noqa: BLE001 — bütçe okunamazsa sütun boş kalır, rapor düşmez
        log.warning("finance: bütçe okunamadı: %s", e)
        return None
    basis = _j(row.basis_json, {})
    wc = B.normalized((basis.get("dagilim") or {}).get("ciro") or [1] * 12)
    stored = _stored_map(engine, tenant) if stored is None else stored
    cache: dict[str, dict[str, Any]] = {}
    val: dict[str, float] = {"NET_SATIS": sum(float(tot.get("ciro") or 0) * wc[m - 1] for m in ms)}
    for d in depts:
        aylar = _j(d.aylar_json, [0.0] * 12)
        amount = sum(float(aylar[m - 1] or 0) for m in ms)
        target = _line_of(d.hesap, stored, cache)["satir"]
        if not target or target == EXCLUDED:
            continue
        val[target] = val.get(target, 0.0) - amount
    ls = {x["kod"]: x for x in lines(engine, tenant)}
    for k in ("FAALIYET",):
        if k in ls:
            val[k] = sum(val.get(c, 0.0) for c in (ls[k]["formul"] or []))
    return {"values": val, "plan": {"id": row.id, "title": row.title, "version": row.version, "scenario": row.scenario}}


def _sales_totals(engine: sa.engine.Engine, months: list[tuple[int, int]]) -> dict[str, float]:
    if not months:
        return {}
    cols = [*_MEASURES, "satis_net", "iade_net", "satis_satir", "maliyetsiz_satir"]
    with engine.connect() as c:
        r = c.execute(sa.select(*[sa.func.sum(SALES_MONTH.c[k]) for k in cols], sa.func.count())
                      .where(_months_cond(SALES_MONTH, months))).first()
    if not r or not r[-1]:
        return {}
    return {k: float(v or 0) for k, v in zip(cols, r[:-1])}


def cost_coverage(s: dict[str, float]) -> dict[str, Any]:
    """Maliyeti işlenmiş satış payı (net tutara göre) ve «yaklaşık» notu."""
    net = s.get("net") or 0.0
    un = s.get("maliyetsiz_net") or 0.0
    share = (net - un) / net if net else None
    return {"maliyetliPay": _r(share, 4), "maliyetsizNet": _r(un), "maliyetsizSatir": int(s.get("maliyetsiz_satir") or 0),
            "satisSatir": int(s.get("satis_satir") or 0), "yaklasik": bool(un) and abs(un) > 0.005}


def _complete(months: list[tuple[int, int]], end: Optional[date]) -> bool:
    if not end:
        return False
    y, m = months[-1]
    last = (date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1))
    return end >= last


def pnl(engine: sa.engine.Engine, tenant: str, year: int, month: int, grain: str = "ay",
        compare: Iterable[str] = ("onceki", "gecen-yil", "butce")) -> dict[str, Any]:
    months = period_months(int(year), int(month), grain)
    ls = lines(engine, tenant)
    stored = _stored_map(engine, tenant)
    loaded = _loaded_years(engine)
    end = data_end(engine)
    cur = compute_lines(engine, tenant, months, stored)
    cols: dict[str, dict[str, Any]] = {"donem": {"label": period_label(months), "values": cur["values"],
                                                 "complete": _complete(months, end), "loaded": int(year) in loaded}}
    cmp = set(compare)
    if "onceki" in cmp:
        pm = previous_period(int(year), int(month), grain)
        if pm:
            cols["onceki"] = {"label": period_label(pm), "loaded": all(y in loaded for y, _ in pm),
                              "values": compute_lines(engine, tenant, pm, stored)["values"]}
    if "gecen-yil" in cmp:
        ly = [(y - 1, m) for y, m in months]
        cols["gecenYil"] = {"label": period_label(ly), "loaded": int(year) - 1 in loaded,
                            "values": compute_lines(engine, tenant, ly, stored)["values"]}
    if "butce" in cmp:
        b = budget_values(engine, tenant, months, stored)
        cols["butce"] = ({"label": "Bütçe", "values": b["values"], "plan": b["plan"]} if b
                         else {"label": "Bütçe", "values": {}, "plan": None, "note": "Bu yıl için yürürlükte bütçe planı yok."})
    sales = _sales_totals(engine, months)
    cov = cost_coverage(sales)
    net_sales_ledger = cur["values"].get("NET_SATIS", 0.0)
    rows = []
    for x in ls:
        k = x["kod"]
        v = {c: (_r(col["values"].get(k)) if k in col["values"] else None) for c, col in cols.items()}
        notes = []
        if k in ("SMM", "BRUT_KAR") and cov["yaklasik"]:
            notes.append(f"Satışların %{_pct(1 - (cov['maliyetliPay'] or 0))}'inde maliyet işlenmemiş; satılan malın "
                         f"maliyeti eksik okunur (yaklaşık).")
        if cur["unapproved"].get(k):
            notes.append("Eşlemesi onaylanmamış hesap içerir.")
        rows.append({**x, "values": v, "onaysiz": _r(cur["unapproved"].get(k)), "hesap": cur["accounts"].get(k, 0),
                     "notlar": notes, "yaklasik": bool(notes and k in ("SMM", "BRUT_KAR"))})
    close = close_info(engine, tenant, int(year), int(month)) if grain == "ay" else None
    if close and close.get("durum") == "kapandi":
        close["degisti"] = close.get("ozetHash") != _hash(cur["values"])
    un_total = sum(abs(cur["unapproved"].get(c, 0.0)) for c in leaf_codes(ls) + [UNMAPPED])
    return {
        "year": int(year), "month": int(month), "grain": grain, "months": [list(m) for m in months],
        "columns": {k: {kk: vv for kk, vv in col.items() if kk != "values"} for k, col in cols.items()},
        "rows": rows, "dislanan": _r(cur["excluded"]), "kurallar": cur["rules"],
        "esleme": {"onaysizTutar": _r(un_total), "eslenmemis": cur["accounts"].get(UNMAPPED, 0),
                   "eslenmemisTutar": _r(cur["values"].get(UNMAPPED))},
        "maliyet": cov, "faturaNetSatis": _r(sales.get("net")) if sales else None,
        "mutabakatFarki": _r(net_sales_ledger - sales["net"]) if sales else None,
        "mizan": trial_status(engine, int(year), [m for _, m in months]),
        "kapanis": close, **freshness(engine),
    }


def _pct(x: float) -> str:
    return f"{x * 100:.1f}".replace(".", ",")


def _hash(values: dict[str, float]) -> str:
    return hashlib.sha256(_dump({k: round(v, 2) for k, v in sorted(values.items())}).encode()).hexdigest()


def trial_status(engine: sa.engine.Engine, year: int, months: list[int]) -> dict[str, Any]:
    """Mizan denkliği (bütün hesaplar): |Σ borç − Σ alacak| < 0,01 ₺ ise yeşil."""
    t = meta_get(engine, f"trial:{year}").get("months") or {}
    if not t:
        return {"durum": "okunmadi"}
    b = sum(float((t.get(str(m)) or {}).get("borc") or 0) for m in months)
    a = sum(float((t.get(str(m)) or {}).get("alacak") or 0) for m in months)
    diff = b - a
    return {"durum": "denk" if abs(diff) < 0.01 else "fark", "borc": _r(b), "alacak": _r(a), "fark": _r(diff),
            "yil": {"borc": _r(sum(float(v.get("borc") or 0) for v in t.values())),
                    "alacak": _r(sum(float(v.get("alacak") or 0) for v in t.values()))}}


def line_accounts(engine: sa.engine.Engine, tenant: str, kod: str, year: int, month: int, grain: str) -> dict[str, Any]:
    """Bir satırın hesapları (alt satırlarıyla), dönem tutarı ve eşleme durumu."""
    ls = {x["kod"]: x for x in lines(engine, tenant)}
    if kod not in ls:
        raise FinanceError("Satır bulunamadı.", 404)
    leaves = set(_leaves(kod, ls))
    months = period_months(int(year), int(month), grain)
    stored = _stored_map(engine, tenant)
    cache: dict[str, dict[str, Any]] = {}
    items = []
    for code, a in sorted(account_totals(engine, months).items()):
        res = _line_of(code, stored, cache)
        target = res["satir"] if res["satir"] in ls else (None if res["satir"] == EXCLUDED else UNMAPPED)
        if res["satir"] == EXCLUDED or target not in leaves:
            continue
        items.append({"hesap": code, "ad": a["ad"], "borc": _r(a["borc"]), "alacak": _r(a["alacak"]), "etki": _r(a["etki"]),
                      "satirSayisi": a["satir"], "satir": target, "esleme": res})
    items.sort(key=lambda i: -abs(i["etki"] or 0))
    return {"kod": kod, "ad": ls[kod]["ad"], "donem": period_label(months), "items": items,
            "toplam": _r(sum(i["etki"] or 0 for i in items)), **freshness(engine)}


def reconciliation(engine: sa.engine.Engine, tenant: str, year: int, month: int, grain: str = "ay") -> dict[str, Any]:
    """Muhasebedeki net satış (600–612, eşlemeyle) ↔ fatura satırlarından net satış (LINENET)."""
    months = period_months(int(year), int(month), grain)
    cur = compute_lines(engine, tenant, months)
    s = _sales_totals(engine, months)
    ledger = cur["values"].get("NET_SATIS", 0.0)
    brut = cur["values"].get("BRUT_SATIS", 0.0)
    ind = cur["values"].get("SATIS_INDIRIM", 0.0)
    inv = s.get("net") if s else None
    return {
        "donem": period_label(months),
        "muhasebe": {"brutSatis": _r(brut), "satisIndirimleri": _r(ind), "netSatis": _r(ledger)},
        "fatura": ({"satis": _r(s.get("satis_net")), "iade": _r(-(s.get("iade_net") or 0)), "netSatis": _r(inv),
                    "iskontoOncesi": _r(s.get("brut")), "iskonto": _r((s.get("brut") or 0) - (s.get("net") or 0))} if s else None),
        "fark": _r(ledger - inv) if inv is not None else None,
        "olasiNedenler": [
            "Fatura satırı tanımı yalnız malzeme satırlarını okur; hizmet ve masraf satırlarının satış hesabı "
            "muhasebede vardır.",
            "Faturası kesilip muhasebeleşmemiş ya da ayrı tarihte muhasebeleşen belge.",
            "Satış hesabına elle kaydedilen düzeltme ve mahsup fişleri.",
            "Satış indirimleri (611) fatura satırında net tutardan düşülmüş gelir; muhasebede ayrı hesaptadır.",
        ],
        **freshness(engine),
    }


# ------------------------------------------------------------------ hesap eşlemesi


def account_map(engine: sa.engine.Engine, tenant: str, year: Optional[int] = None) -> dict[str, Any]:
    """Bütün 6/7 hesapları (Logo hesap planı + hareketi olanlar), geçerli eşlemeleri ve bu yılın tutarı."""
    stored = _stored_map(engine, tenant)
    end = data_end(engine)
    year = int(year or (end.year if end else _today().year))
    totals = account_totals(engine, [(year, m) for m in range(1, 13)])
    with engine.connect() as c:
        names = {r.hesap_kodu: r.ad for r in c.execute(sa.select(ACCOUNTS)).all()}
    codes = sorted(set(names) | set(totals))
    ls = lines(engine, tenant)
    items = []
    counts = {"onayli": 0, "dislandi": 0, "oneri": 0, "yok": 0}
    amounts = {"onayli": 0.0, "dislandi": 0.0, "oneri": 0.0, "yok": 0.0}
    for code in codes:
        res = resolve(code, stored)
        a = totals.get(code)
        row = stored.get(code)
        items.append({"hesap": code, "ad": names.get(code) or (a or {}).get("ad"), "grup": code[:3],
                      "etki": _r((a or {}).get("etki")), "hareketli": a is not None, "esleme": res,
                      "kayit": ({"durum": row.durum, "satir": row.satir_kodu, "oneriSatir": row.oneri_satir,
                                 "olasilik": row.oneri_olasilik, "kaynak": row.oneri_kaynak,
                                 "oneri": _j(row.oneri_json, None), "not": row.not_, "updatedBy": row.updated_by,
                                 "updatedAt": _iso(row.updated_at)} if row else None)})
        if a is not None:
            counts[res["durum"]] += 1
            amounts[res["durum"]] += abs(a["etki"])
    return {"year": year, "lines": ls, "items": items, "counts": counts, "amounts": {k: _r(v) for k, v in amounts.items()},
            "suggest": meta_get(engine, f"suggest:{tenant}"), **freshness(engine)}


def _valid_target(engine: sa.engine.Engine, tenant: str, satir: Any) -> str:
    satir = str(satir or "").strip()
    if satir == EXCLUDED:
        return satir
    if satir not in leaf_codes(lines(engine, tenant)):
        raise FinanceError("Satır gelir tablosunun bir alt satırı ya da «dışla» olmalı.")
    return satir


def set_mapping(engine: sa.engine.Engine, tenant: str, user: str, code: str, body: dict[str, Any]) -> dict[str, Any]:
    """Muhasebe kararı: hesabı (ya da grubu) bir satıra onaylar ya da dışlar."""
    code = src.check_code(code)
    target = _valid_target(engine, tenant, body.get("satir"))
    note = _text(body.get("not"), 500)
    now = _now()
    vals = {"satir_kodu": None if target == EXCLUDED else target, "durum": "dislandi" if target == EXCLUDED else "onayli",
            "onaylayan": user, "onay_tarihi": now, "not_": note, "updated_by": user, "updated_at": now}
    with engine.begin() as c:
        cur = c.execute(sa.select(ACCOUNT_MAP).where(ACCOUNT_MAP.c.tenant_id == tenant, ACCOUNT_MAP.c.hesap_kodu == code)).first()
        if cur:
            c.execute(ACCOUNT_MAP.update().where(ACCOUNT_MAP.c.tenant_id == tenant, ACCOUNT_MAP.c.hesap_kodu == code).values(**vals))
        else:
            c.execute(ACCOUNT_MAP.insert().values(tenant_id=tenant, hesap_kodu=code, **vals))
    return {"hesap": code, "satir": target, "durum": vals["durum"], "onceki": ({"durum": cur.durum, "satir": cur.satir_kodu} if cur else None)}


def reset_mapping(engine: sa.engine.Engine, tenant: str, code: str) -> dict[str, Any]:
    """Onayı geri alır; öneri (Zeki AI) varsa kayıt öneri olarak kalır."""
    code = src.check_code(code)
    with engine.begin() as c:
        cur = c.execute(sa.select(ACCOUNT_MAP).where(ACCOUNT_MAP.c.tenant_id == tenant, ACCOUNT_MAP.c.hesap_kodu == code)).first()
        if not cur:
            raise FinanceError("Bu hesap için kayıtlı eşleme yok.", 404)
        if cur.oneri_satir:
            c.execute(ACCOUNT_MAP.update().where(ACCOUNT_MAP.c.tenant_id == tenant, ACCOUNT_MAP.c.hesap_kodu == code)
                      .values(durum="oneri", satir_kodu=None, onaylayan=None, onay_tarihi=None, updated_at=_now()))
        else:
            c.execute(ACCOUNT_MAP.delete().where(ACCOUNT_MAP.c.tenant_id == tenant, ACCOUNT_MAP.c.hesap_kodu == code))
    return {"hesap": code, "onceki": {"durum": cur.durum, "satir": cur.satir_kodu}}


def approve_suggestions(engine: sa.engine.Engine, tenant: str, user: str, codes: Iterable[str]) -> dict[str, Any]:
    """Seçilen hesapların (ya da grupların) geçerli önerisini onaylar. Önerisi olmayan kod atlanır ve söylenir."""
    stored = _stored_map(engine, tenant)
    done, skipped = [], []
    for raw in codes:
        try:
            code = src.check_code(raw)
        except src.SourceError:
            skipped.append({"hesap": str(raw)[:60], "neden": "Hesap kodu geçersiz."})
            continue
        res = resolve(code, stored)
        if res["durum"] in ("onayli", "dislandi"):
            skipped.append({"hesap": code, "neden": "Zaten karar verilmiş."})
            continue
        if not res["satir"]:
            skipped.append({"hesap": code, "neden": "Önerisi yok; satırı elle seçin."})
            continue
        set_mapping(engine, tenant, user, code, {"satir": res["satir"], "not": f"Öneri onaylandı ({res['kaynak']})"})
        done.append({"hesap": code, "satir": res["satir"]})
    return {"approved": done, "skipped": skipped}


def suggestion_targets(engine: sa.engine.Engine, tenant: str, codes: Optional[Iterable[str]] = None) -> list[dict[str, Any]]:
    """Zeki AI'a sorulacak hesaplar: istenenler ya da kuralı ve kararı olmayan hareketli hesaplar."""
    data = account_map(engine, tenant)
    want = {str(c) for c in (codes or [])}
    out = []
    for i in data["items"]:
        if want:
            if i["hesap"] in want and i["esleme"]["durum"] not in ("onayli", "dislandi"):
                out.append(i)
        elif i["esleme"]["durum"] == "yok" and i["hareketli"]:
            out.append(i)
    return out


def suggest_with_model(engine: sa.engine.Engine, tenant: str, llm: Any, items: list[dict[str, Any]],
                       min_prob: float, min_margin: float) -> dict[str, Any]:
    """Kapalı küme seçim: her hesap için gelir tablosunun alt satırlarından biri ya da «gelir tablosu dışı».
    Eşik altı öneri «belirsiz» kaydedilir (satırı boş, olasılıklar kanıt olarak). Rakam üretilmez."""
    ls = lines(engine, tenant)
    leaves = [x for x in ls if x["tur"] in ("gelir", "gider")]
    labels = [x["ad"] for x in leaves] + ["Gelir tablosu dışı (dışla)"]
    codes = [x["kod"] for x in leaves] + [EXCLUDED]
    done = uncertain = failed = 0
    for i in items:
        prompt = ("Tekdüzen Hesap Planı'na göre çalışan bir yayınevinin muhasebe hesabı aşağıda. Bu hesabın bakiyesi "
                  "aylık gelir tablosunun hangi satırına gider?\n\n"
                  f"Hesap kodu: {i['hesap']}\nHesap adı: {i.get('ad') or '(adsız)'}\n"
                  f"Bu yılki kâr etkisi (alacak − borç) işareti: {'artı' if (i.get('etki') or 0) >= 0 else 'eksi'}")
        try:
            r = llm.choose(prompt, labels)
        except Exception as e:  # noqa: BLE001 — model yoksa kalan hesaplar sonraki denemeye kalır
            log.warning("finance: eşleme önerisi alınamadı: %s", e)
            failed += 1
            continue
        ok = bool(r.choice) and r.confident(min_prob, min_margin=min_margin)
        choice = codes[r.index] if ok and r.index is not None else None
        ev = r.as_dict() if hasattr(r, "as_dict") else {"choice": r.choice, "probability": r.probability}
        now = _now()
        vals = {"oneri_satir": choice, "oneri_olasilik": r.probability, "oneri_kaynak": "zeki",
                "oneri_json": _dump({"secenekler": dict(zip(codes, labels)), "sonuc": ev, "esik": [min_prob, min_margin]}),
                "updated_by": "Zeki AI", "updated_at": now}
        with engine.begin() as c:
            cur = c.execute(sa.select(ACCOUNT_MAP.c.durum).where(ACCOUNT_MAP.c.tenant_id == tenant,
                                                                  ACCOUNT_MAP.c.hesap_kodu == i["hesap"])).first()
            if cur and cur.durum in ("onayli", EXCLUDED):
                continue
            if cur:
                c.execute(ACCOUNT_MAP.update().where(ACCOUNT_MAP.c.tenant_id == tenant, ACCOUNT_MAP.c.hesap_kodu == i["hesap"]).values(**vals))
            else:
                c.execute(ACCOUNT_MAP.insert().values(tenant_id=tenant, hesap_kodu=i["hesap"], durum="oneri", **vals))
        if choice:
            done += 1
        else:
            uncertain += 1
    return {"istenen": len(items), "oneri": done, "belirsiz": uncertain, "hata": failed}


# ------------------------------------------------------------------ kapanış


def close_info(engine: sa.engine.Engine, tenant: str, year: int, month: int) -> Optional[dict[str, Any]]:
    with engine.connect() as c:
        r = c.execute(sa.select(CLOSES).where(CLOSES.c.tenant_id == tenant, CLOSES.c.year == year, CLOSES.c.month == month)).first()
    if not r:
        return None
    return {"durum": r.durum, "kapatan": r.kapatan, "kapanisTarihi": _iso(r.kapanis_tarihi), "ozetHash": r.ozet_hash,
            "ozet": _j(r.ozet_json, None), "not": r.not_, "acan": r.acan, "acilisTarihi": _iso(r.acilis_tarihi),
            "acilisGerekce": r.acilis_gerekce}


def close_month(engine: sa.engine.Engine, tenant: str, user: str, year: int, month: int, note: Any = None) -> dict[str, Any]:
    """Ayı «kapandı» işaretler; o anki gelir tablosu değerleri saklanır, sonraki değişiklik fark olarak görünür."""
    end = data_end(engine)
    months = period_months(int(year), int(month), "ay")
    if not _complete(months, end):
        raise FinanceError("Ayın verisi tamamlanmadan (veri son günü ay sonundan önce) kapanış yapılamaz.")
    if int(year) not in _loaded_years(engine):
        raise FinanceError("Bu yılın muhasebe verisi henüz okunmadı.")
    cur = close_info(engine, tenant, int(year), int(month))
    if cur and cur["durum"] == "kapandi":
        raise FinanceError("Bu ay zaten kapandı.", 409)
    vals = compute_lines(engine, tenant, months)["values"]
    row = {"durum": "kapandi", "kapatan": user, "kapanis_tarihi": _now(), "ozet_hash": _hash(vals),
           "ozet_json": _dump({k: round(v, 2) for k, v in vals.items()}), "not_": _text(note, 2000)}
    with engine.begin() as c:
        if cur:
            c.execute(CLOSES.update().where(CLOSES.c.tenant_id == tenant, CLOSES.c.year == int(year),
                                            CLOSES.c.month == int(month)).values(**row))
        else:
            c.execute(CLOSES.insert().values(tenant_id=tenant, year=int(year), month=int(month), **row))
    return close_info(engine, tenant, int(year), int(month)) or {}


def reopen_month(engine: sa.engine.Engine, tenant: str, user: str, year: int, month: int, reason: Any) -> dict[str, Any]:
    reason = _text(reason, 2000)
    if not reason:
        raise FinanceError("Yeniden açmanın gerekçesi zorunlu.")
    cur = close_info(engine, tenant, int(year), int(month))
    if not cur or cur["durum"] != "kapandi":
        raise FinanceError("Bu ay kapalı değil.", 409)
    with engine.begin() as c:
        c.execute(CLOSES.update().where(CLOSES.c.tenant_id == tenant, CLOSES.c.year == int(year), CLOSES.c.month == int(month))
                  .values(durum="acik", acan=user, acilis_tarihi=_now(), acilis_gerekce=reason))
    return close_info(engine, tenant, int(year), int(month)) or {}


# ------------------------------------------------------------------ kârlılık


def snapshot_unit_costs(snap: Optional[dict], codes: Iterable[str]) -> dict[str, dict[str, Any]]:
    """M9 sağlayıcısı yokken aynı tanım: M9 Logo görüntüsünde kitabın maliyeti girilmiş satışlarının en son yılı,
    maliyet ÷ maliyetli adet (`pricing/cost_provider.actual_costs` ile birebir)."""
    if not snap:
        return {}
    sales = snap.get("sales") or {}
    out = {}
    for code in codes:
        for y in sorted(sales.get(code) or {}, reverse=True):
            v = (sales.get(code) or {}).get(y) or {}
            qty = float(v.get("costedQty") or 0)
            cost = float(v.get("cogs") or 0) / qty if qty > 0 else 0.0
            if cost > 0:
                out[code] = {"maliyet": round(cost, 4), "kaynak": "gerceklesen", "yil": int(y)}
                break
    return out


def royalty_terms(snap: Optional[dict], code: str) -> Optional[dict[str, Any]]:
    """CRM'deki yürürlükteki sözleşmenin telif oranı (M9 görüntüsünden). Yoksa None."""
    b = ((snap or {}).get("books") or {}).get(code) or {}
    r = b.get("royalty")
    if not r or not r.get("rate"):
        return None
    vat = float(b.get("vat") or 0)
    vat = vat / 100.0 if vat > 1 else vat
    price = float(b.get("price") or 0)
    return {"rate": float(r["rate"]), "basis": r.get("basis") or "kapak", "on": r.get("on") or "satis",
            "netPrice": price / (1 + vat) if price else None}


def _prints_in(snap: Optional[dict], code: str, months: list[tuple[int, int]]) -> float:
    want = {f"{y:04d}-{m:02d}" for y, m in months}
    return sum(float(p.get("qty") or 0) for p in ((snap or {}).get("prints") or {}).get(code) or []
               if (p.get("date") or "")[:7] in want)


def _book_info(engine: sa.engine.Engine) -> dict[str, dict[str, Any]]:
    """M46'nın CRM kitap kartı önbelleği (ad, yayınevi, kitaplık). Tablo yoksa boş."""
    try:
        from semantic_bridge import budget as B

        B.ensure(engine)
        with engine.connect() as c:
            return {r.stok_kodu: {"ad": r.ad, "yayinevi": r.yayinevi, "kitaplik": r.kitaplik}
                    for r in c.execute(sa.select(B.BOOKINFO)).all()}
    except Exception as e:  # noqa: BLE001
        log.warning("finance: kitap kartları okunamadı: %s", e)
        return {}


def _acc() -> dict[str, float]:
    return {k: 0.0 for k in (*_MEASURES, "smm_tahmini", "tahmini_adet", "bilinmeyen_net", "telif", "telif_net",
                             "telifsiz_net")}


def profitability(engine: sa.engine.Engine, *, by: str, year: int, frm: int = 1, to: int = 12, q: str = "",
                  sort: str = "net", page: int = 0, unit_costs: Callable[[list[str]], dict[str, dict[str, Any]]] = lambda c: {},
                  snapshot: Optional[dict] = None, with_royalty: bool = True, all_rows: bool = False) -> dict[str, Any]:
    """Kitap · seri · yayınevi · kanal · cari kârlılığı. Kesin katkı = maliyetli satırların net satışı − Logo maliyeti.
    Yaklaşık katkı = (net − maliyeti bilinmeyen) − Logo maliyeti − M9 birim maliyetiyle doldurulan − telif (yaklaşık)."""
    if by not in PROFIT_BY:
        raise FinanceError("Kırılım kitap, seri, yayınevi, kanal ya da cari olmalı.")
    if not (1 <= int(frm) <= int(to) <= 12):
        raise FinanceError("Ay aralığı 1–12 olmalı.")
    months = [(int(year), m) for m in range(int(frm), int(to) + 1)]
    groups: dict[str, dict[str, Any]] = {}
    names: dict[str, str] = {}
    telif_status = {"var": 0, "yok": 0}
    if by == "cari":
        with engine.connect() as c:
            rows = c.execute(sa.select(PROFIT_CLIENTS).where(_months_cond(PROFIT_CLIENTS, months))).all()
        for r in rows:
            g = groups.setdefault(r.cari_kodu, _acc())
            names[r.cari_kodu] = r.cari_adi or r.cari_kodu
            g.setdefault("_kanal", set()).add(r.kanal)
            for k in _MEASURES:
                g[k] += float(getattr(r, k) or 0)
            g["smm_tahmini"] += float(r.tahmini_maliyet or 0)
            g["tahmini_adet"] += float(r.tahmini_adet or 0)
            # maliyetsiz net'in birim maliyeti bilinen payı adet oranıyla (yaklaşık); gerisi bilinmiyor
            un_qty = float(r.maliyetsiz_adet or 0)
            share = (float(r.tahmini_adet or 0) / un_qty) if un_qty else 0.0
            g["bilinmeyen_net"] += float(r.maliyetsiz_net or 0) * (1 - max(0.0, min(1.0, share)))
            g["telifsiz_net"] += float(r.net or 0)
        with_royalty = False
    else:
        with engine.connect() as c:
            rows = c.execute(sa.select(PROFIT_ITEMS).where(_months_cond(PROFIT_ITEMS, months))).all()
        info = _book_info(engine) if by in ("kitap", "seri", "yayinevi") else {}
        codes = sorted({r.stok_kodu for r in rows if (r.maliyetsiz_adet or 0) != 0})
        costs = unit_costs(codes) if codes else {}
        book_net: dict[str, float] = {}
        for r in rows:
            book_net[r.stok_kodu] = book_net.get(r.stok_kodu, 0.0) + float(r.net or 0)
        terms = {code: royalty_terms(snapshot, code) for code in book_net} if with_royalty else {}
        printed = {code: _prints_in(snapshot, code, months) for code, t in terms.items() if t and t["on"] == "baski"}
        for code, t in terms.items():
            telif_status["var" if t else "yok"] += 1
        for r in rows:
            b = info.get(r.stok_kodu) or {}
            key = {"kitap": r.stok_kodu, "seri": b.get("kitaplik") or "Seri/kitaplık belirsiz",
                   "yayinevi": b.get("yayinevi") or "Yayınevi belirsiz", "kanal": r.kanal}[by]
            if by == "kitap":
                names[key] = b.get("ad") or r.stok_kodu
            g = groups.setdefault(key, _acc())
            for k in _MEASURES:
                g[k] += float(getattr(r, k) or 0)
            un_qty = float(r.maliyetsiz_adet or 0)
            uc = costs.get(r.stok_kodu)
            if un_qty and uc and uc.get("maliyet") is not None:
                g["smm_tahmini"] += un_qty * float(uc["maliyet"])
                g["tahmini_adet"] += un_qty
            elif un_qty or (r.maliyetsiz_net or 0):
                g["bilinmeyen_net"] += float(r.maliyetsiz_net or 0)
            t = terms.get(r.stok_kodu)
            if t:
                if t["on"] == "baski":
                    bn = book_net.get(r.stok_kodu) or 0.0
                    share = (float(r.net or 0) / bn) if bn else 0.0
                    amount = t["rate"] * printed.get(r.stok_kodu, 0.0) * (t["netPrice"] or 0) * share
                elif t["basis"] == "net":
                    amount = t["rate"] * float(r.net or 0)
                else:
                    amount = t["rate"] * float(r.adet or 0) * (t["netPrice"] or 0)
                g["telif"] += amount
                g["telif_net"] += float(r.net or 0)
            else:
                g["telifsiz_net"] += float(r.net or 0)
    needle = (q or "").strip().casefold()
    items = []
    for key, g in groups.items():
        name = names.get(key) or key
        if needle and needle not in key.casefold() and needle not in str(name).casefold():
            continue
        items.append(_profit_row(key, name, g, with_royalty))
    order = {"net": lambda i: -(i["net"] or 0), "katki": lambda i: -(i["katkiYaklasik"] if i["katkiYaklasik"] is not None else -1e18),
             "marj": lambda i: (i["marjYaklasik"] if i["marjYaklasik"] is not None else 9e9),
             "iskonto": lambda i: -(i["iskontoOrani"] or 0), "ad": lambda i: str(i["ad"]).casefold()}.get(sort)
    items.sort(key=order or (lambda i: -(i["net"] or 0)))
    tot = _acc()
    for g in groups.values():
        for k in tot:
            tot[k] += g.get(k, 0.0)
    total = len(items)
    page = max(0, int(page))
    return {
        "by": by, "byLabel": PROFIT_BY[by], "year": int(year), "from": int(frm), "to": int(to),
        "donem": period_label(months), "items": items if all_rows else items[page * PAGE_SIZE:(page + 1) * PAGE_SIZE],
        "total": total, "page": page, "pageSize": PAGE_SIZE, "toplam": _profit_row("toplam", "Toplam", tot, with_royalty),
        "telif": {"hesaplandi": with_royalty, "sozlesmeli": telif_status["var"], "sozlesmesiz": telif_status["yok"],
                  "kaynak": "CRM'deki yürürlükteki sözleşmenin oranından tahakkuk (yaklaşık)" if with_royalty else
                  ("Cari kırılımında telif hesaplanmaz (kitap bazında bilinir)." if by == "cari" else "Telif düşülmedi.")},
        **freshness(engine),
    }


def _profit_row(key: str, name: str, g: dict[str, Any], with_royalty: bool) -> dict[str, Any]:
    net = g["net"]
    kesin = g["maliyetli_net"] - g["maliyet"]
    known_net = net - g["bilinmeyen_net"]
    telif = g["telif"] if with_royalty else 0.0
    approx = known_net - g["maliyet"] - g["smm_tahmini"] - telif
    brut = g["brut"]
    return {
        "key": key, "ad": name, "kanallar": sorted(g.get("_kanal") or []) or None,
        "adet": _r(g["adet"], 1), "iskontoOncesi": _r(brut), "iskonto": _r(brut - net),
        "iskontoOrani": _r((brut - net) / brut, 4) if brut > 0 else None, "net": _r(net),
        "maliyetLogo": _r(g["maliyet"]), "maliyetliNet": _r(g["maliyetli_net"]),
        "katkiKesin": _r(kesin), "marjKesin": _r(kesin / g["maliyetli_net"], 4) if g["maliyetli_net"] > 0 else None,
        "maliyetsizNet": _r(g["maliyetsiz_net"]), "maliyetTahmini": _r(g["smm_tahmini"]),
        "maliyetBilinmeyenNet": _r(g["bilinmeyen_net"]),
        "telif": _r(telif) if with_royalty else None, "telifsizNet": _r(g["telifsiz_net"]) if with_royalty else None,
        "katkiYaklasik": _r(approx) if known_net else None,
        "marjYaklasik": _r(approx / known_net, 4) if known_net > 0 else None,
        "kapsam": _r(known_net / net, 4) if net else None,
        "yaklasik": bool(abs(g["maliyetsiz_net"]) > 0.005 or (with_royalty and telif)),
    }


def profitability_csv(data: dict[str, Any]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    royalty = data["telif"]["hesaplandi"]
    head = [data["byLabel"], "Ad", "Net adet", "İskonto öncesi", "İskonto", "Net satış", "Logo maliyeti (maliyetli satırlar)",
            "Maliyetli satırların net satışı", "Kesin katkı", "Kesin marj", "Maliyetsiz net satış",
            "Yaklaşık maliyet (fiyatlama birim maliyeti)", "Maliyeti bilinmeyen net satış"]
    if royalty:
        head += ["Telif (yaklaşık)", "Telif verisi olmayan net satış"]
    head += ["Yaklaşık katkı", "Yaklaşık marj", "Kapsam"]
    w.writerow(head)

    def n(v: Any, d: int = 2) -> str:
        return "" if v is None else f"{float(v):.{d}f}".replace(".", ",")

    for i in data["items"] + [data["toplam"]]:
        row = [i["key"], i["ad"], n(i["adet"], 1), n(i["iskontoOncesi"]), n(i["iskonto"]), n(i["net"]), n(i["maliyetLogo"]),
               n(i["maliyetliNet"]), n(i["katkiKesin"]), n(i["marjKesin"], 4), n(i["maliyetsizNet"]), n(i["maliyetTahmini"]),
               n(i["maliyetBilinmeyenNet"])]
        if royalty:
            row += [n(i["telif"]), n(i["telifsizNet"])]
        row += [n(i["katkiYaklasik"]), n(i["marjYaklasik"], 4), n(i["kapsam"], 4)]
        w.writerow(row)
    w.writerow([])
    w.writerow([f"Veri son günü: {data.get('veriSonu') or '—'}; maliyeti işlenmiş son satış: {data.get('maliyetSonu') or '—'}. "
                "«Yaklaşık» sütunlar tahmindir."])
    return "﻿" + buf.getvalue()


# ------------------------------------------------------------------ nakit


def monday(d: date) -> date:
    return d - timedelta(days=d.weekday())


def weeks_from(start: date, n: int) -> list[date]:
    return [start + timedelta(days=7 * i) for i in range(n)]


def _week_index(d: date, start: date, n: int) -> Optional[int]:
    i = (d - start).days // 7
    return i if 0 <= i < n else None


def build_cash(inputs: dict[str, Any], weeks: int = 13) -> dict[str, Any]:
    """Saf hesap: kaynak okumalarından 13 haftalık tablo. `inputs`:
    `asof` (veri son günü), `position` {grup: bakiye}, `receivables`/`payables` [(vade, tutar, cari)],
    `cheques` [(tur, durum, vade, tutar)], `crm` [(vade, tutar)], `royalty` [(vade, tutar, para)],
    `tax` [(son_gun, tutar, beyan)], `budget` [(ay_başı, aylık tutar)] (bilgi), `errors` {kaynak: metin}."""
    asof: date = inputs["asof"]
    start = monday(asof)
    ws = weeks_from(start, weeks)
    lines_: dict[str, dict[str, Any]] = {}
    overdue: dict[str, float] = {}

    def add(kalem: str, yon: str, kaynak: str, approx: bool, d: Optional[date], amount: float, detail: Any = None) -> None:
        if not amount or d is None:
            return
        spec = lines_.setdefault(kalem, {"kalem": kalem, "yon": yon, "kaynak": kaynak, "yaklasik": approx,
                                         "haftalar": [0.0] * weeks, "ayrinti": [[] for _ in range(weeks)]})
        if d < asof:
            overdue[kalem] = overdue.get(kalem, 0.0) + amount
            return
        i = _week_index(d, start, weeks)
        if i is None:
            return
        spec["haftalar"][i] += amount
        if detail is not None and len(spec["ayrinti"][i]) < 50:
            spec["ayrinti"][i].append(detail)

    for d, amt, cari in inputs.get("receivables") or []:
        add("alacak", "giris", "Müşteri alacakları — ödeme planı vadeleri, FIFO yaklaşımı (Logo)", True, d, amt, {"vade": d.isoformat(), "cari": cari})
    for tur, durum, d, amt in inputs.get("cheques") or []:
        yon = src.cheque_direction(tur, durum)
        if yon == "giris":
            add("cek-giris", "giris", "Alınan çek ve senetler — vade (Logo)", False, d, amt, {"vade": d.isoformat() if d else None, "tur": tur, "durum": durum})
        elif yon == "cikis":
            add("cek-cikis", "cikis", "Verilen çek ve senetler — vade (Logo)", False, d, amt, {"vade": d.isoformat() if d else None, "tur": tur, "durum": durum})
    for d, amt in inputs.get("crm") or []:
        add("crm-tahsilat", "giris", "CRM'de onay bekleyen tahsilat (Logo'ya düşmemiş)", True, d, amt)
    for d, amt, cari in inputs.get("payables") or []:
        add("satici", "cikis", "Satıcı borçları — ödeme planı vadeleri, FIFO yaklaşımı (Logo)", True, d, amt, {"vade": d.isoformat(), "cari": cari})
    other_ccy: dict[str, float] = {}
    for d, amt, ccy in inputs.get("royalty") or []:
        if (ccy or "TRY").upper() not in ("TRY", "TL"):
            other_ccy[ccy] = other_ccy.get(ccy, 0.0) + amt
            continue
        add("telif", "cikis", "Sözleşme ödeme takvimi (telif, avans, hakediş)", False, d, amt)
    for d, amt, beyan in inputs.get("tax") or []:
        add("vergi", "cikis", "Vergi takvimi (girilen tahmini ödeme)", True, d, amt, {"beyan": beyan})
    budget_rows = []
    for m_start, monthly in inputs.get("budget") or []:
        days = ((date(m_start.year + (m_start.month == 12), m_start.month % 12 + 1, 1)) - m_start).days
        for k in range(days):
            d = m_start + timedelta(days=k)
            if d < asof:
                continue
            i = _week_index(d, start, weeks)
            if i is not None:
                budget_rows.append((i, monthly / days))
    if budget_rows:
        spec = lines_.setdefault("butce", {"kalem": "butce", "yon": "bilgi", "kaynak": "Bütçe gider temposu (departman bütçesi; satıcı borcuyla örtüşür, toplama katılmaz)",
                                           "yaklasik": True, "haftalar": [0.0] * weeks, "ayrinti": [[] for _ in range(weeks)]})
        for i, v in budget_rows:
            spec["haftalar"][i] += v
    opening = sum(float((inputs.get("position") or {}).get(g) or 0) for g in ("100", "102"))
    week_rows = []
    bal = opening
    for i, w in enumerate(ws):
        gin = sum(v["haftalar"][i] for v in lines_.values() if v["yon"] == "giris")
        out = sum(v["haftalar"][i] for v in lines_.values() if v["yon"] == "cikis")
        close = bal + gin - out
        week_rows.append({"hafta": i + 1, "baslangic": w.isoformat(), "bitis": (w + timedelta(days=6)).isoformat(),
                          "acilis": round(bal, 2), "giris": round(gin, 2), "cikis": round(out, 2), "net": round(gin - out, 2),
                          "kapanis": round(close, 2), "acik": close < 0,
                          "enBuyukCikis": max(((v["kaynak"], v["haftalar"][i]) for v in lines_.values() if v["yon"] == "cikis"),
                                              key=lambda x: x[1], default=(None, 0.0))[0] if out else None,
                          "kismi": i == 0 and asof > w})
        bal = close
    return {"baslangic": start.isoformat(), "asof": asof.isoformat(), "acilisBakiye": round(opening, 2),
            "pozisyon": {k: round(float(v), 2) for k, v in (inputs.get("position") or {}).items()},
            "vadesiGecmis": {k: round(v, 2) for k, v in overdue.items()}, "dovizTelif": {k: round(v, 2) for k, v in other_ccy.items()},
            "kalemler": [{**{k: v for k, v in s.items() if k != "ayrinti"}, "haftalar": [round(x, 2) for x in s["haftalar"]],
                          "toplam": round(sum(s["haftalar"]), 2), "ayrinti": s["ayrinti"]} for s in lines_.values()],
            "haftalar": week_rows, "hatalar": inputs.get("errors") or {}}


#: Olasılıklı nakit bandı (öneri 7). Vadesi belli kalemler tablodaki kuraldan gelir (işaretiyle); vadesi belirsiz
#: tahsilat/ödeme (FIFO yaklaşımlı alacak, CRM onay bekleyen, satıcı borcu) yerine geçmiş haftalık gerçekleşen
#: tahsilat/ödeme serisinin tahmini (p10/p50/p90) konur. «En kötü %10» kasa çizgisi = kesin kalemler + tahsilat p10 −
#: ödeme p90; haftalık aralıklar toplanır (temkinli, geniş). Kantil yoksa bant yoktur — uydurulmaz.
BAND_CERTAIN = {"cek-giris": 1.0, "cek-cikis": -1.0, "telif": -1.0, "vergi": -1.0}
BAND_REPLACED = ("alacak", "crm-tahsilat", "satici")
BAND_HISTORY_WEEKS = 104
BAND_MIN_WEEKS = 26


def weekly_flows(daily: dict[date, tuple[float, float]], start: date, weeks: int) -> tuple[list[float], list[float], date]:
    """`start` haftasından önceki `weeks` tam haftanın tahsilat ve ödeme toplamı (pazartesi başlangıçlı). Verinin
    başladığı haftadan öncesi atılır. Dönen: (tahsilat, ödeme, ilk hafta)."""
    first_day = min(daily) if daily else start
    first = max(start - timedelta(days=7 * weeks), monday(first_day))
    n = max(0, (start - first).days // 7)
    tah, od = [0.0] * n, [0.0] * n
    for d, (t, o) in daily.items():
        i = (d - first).days // 7
        if 0 <= i < n:
            tah[i] += t
            od[i] += o
    return tah, od, first


def cash_band(result: dict[str, Any], forecast: dict[str, dict[str, list[float]]], history_weeks: int) -> dict[str, Any]:
    """Saf hesap: 13 haftalık tablo + tahsilat/ödeme kantilleri → hafta hafta kapanış bandı."""
    weeks = result.get("haftalar") or []
    lines = {s["kalem"]: s for s in result.get("kalemler") or []}
    opening = float(result.get("acilisBakiye") or 0)
    tah, od = forecast["tahsilat"], forecast["odeme"]
    lo = mid = hi = opening
    rows = []
    for i, w in enumerate(weeks):
        certain = sum(sign * float((lines.get(k) or {}).get("haftalar", [0.0] * len(weeks))[i] or 0) for k, sign in BAND_CERTAIN.items())
        t = {q: round(tah[q][i], 2) for q in ("p10", "p50", "p90")}
        o = {q: round(od[q][i], 2) for q in ("p10", "p50", "p90")}
        lo += certain + t["p10"] - o["p90"]
        mid += certain + t["p50"] - o["p50"]
        hi += certain + t["p90"] - o["p10"]
        rows.append({"hafta": w["hafta"], "baslangic": w["baslangic"], "kesin": round(certain, 2), "tahsilat": t, "odeme": o,
                     "kapanis": {"kotu": round(lo, 2), "orta": round(mid, 2), "iyi": round(hi, 2)}, "kuralKapanis": w["kapanis"]})
    worst = next((r for r in rows if r["kapanis"]["kotu"] < 0), None)
    return {"var": True, "etiket": "tahmin", "gecmisHafta": history_weeks, "haftalar": rows,
            "enKotuAcik": {"hafta": worst["hafta"], "baslangic": worst["baslangic"], "kapanis": worst["kapanis"]["kotu"]} if worst else None,
            "kesinKalemler": [k for k in BAND_CERTAIN if k in lines], "yerineGecen": [k for k in BAND_REPLACED if k in lines],
            "not": ("Vadesi belli kalemler (çek/senet, sözleşme ödemesi, vergi) kuraldan; müşteri tahsilatı ve satıcı ödemesi "
                    "geçmiş haftalık gerçekleşenin tahmininden. Haftalık aralıklar toplandı: bant temkinli (geniş) okunur.")}


def build_band(result: dict[str, Any], daily: dict[date, tuple[float, float]],
               forecaster: Callable[[dict[str, list[float]], int], dict[str, dict[str, list[float]]]]) -> dict[str, Any]:
    """Geçmiş + tahmin servisi → bant; yetersiz geçmişte ya da servis yoksa {"var": False, "neden"}."""
    start = date.fromisoformat(result["baslangic"])
    tah, od, _first = weekly_flows(daily, start, BAND_HISTORY_WEEKS)
    if len(tah) < BAND_MIN_WEEKS or not any(tah) or not any(od):
        return {"var": False, "neden": f"Tahsilat/ödeme geçmişi {len(tah)} hafta; bant için en az {BAND_MIN_WEEKS} hafta gerekir."}
    horizon = len(result.get("haftalar") or [])
    try:
        fc = forecaster({"tahsilat": tah, "odeme": od}, horizon)
    except Exception as e:  # noqa: BLE001 — servis yoksa bant yok, tablo kuraldan kalır
        log.info("finance: nakit bandı tahmini alınamadı: %s", e)
        return {"var": False, "neden": "Tahmin servisine ulaşılamadı; olasılıklı bant bu kurulumda gösterilmiyor."}
    if not all(len(fc.get(k, {}).get(q) or []) >= horizon for k in ("tahsilat", "odeme") for q in ("p10", "p50", "p90")):
        return {"var": False, "neden": "Tahmin aralığı (p10–p90) dönmedi; bant gösterilmiyor."}
    return cash_band(result, fc, len(tah))


def save_cash_run(engine: sa.engine.Engine, tenant: str, user: Optional[str], result: dict[str, Any]) -> str:
    rid = uuid.uuid4().hex
    params = {k: result[k] for k in ("acilisBakiye", "pozisyon", "vadesiGecmis", "dovizTelif", "hatalar", "asof")}
    params["bant"] = result.get("bant")
    params["haftalar"] = result["haftalar"]
    with engine.begin() as c:
        c.execute(CASH_RUNS.insert().values(id=rid, tenant_id=tenant, run_at=_now(), veri_son_gunu=result["asof"],
                                            baslangic=result["baslangic"], params_json=_dump(params), run_by=user))
        rows = []
        start = date.fromisoformat(result["baslangic"])
        for s in result["kalemler"]:
            for i, v in enumerate(s["haftalar"]):
                rows.append({"run_id": rid, "hafta": i + 1, "kalem": s["kalem"][:40],
                             "hafta_baslangic": (start + timedelta(days=7 * i)).isoformat(), "yon": s["yon"],
                             "kaynak": s["kaynak"][:200], "tutar": v, "yaklasik": bool(s["yaklasik"]),
                             "ayrinti_json": _dump(s["ayrinti"][i]) if s["ayrinti"][i] else None})
        if rows:
            c.execute(CASH_LINES.insert(), rows)
    return rid


def cash(engine: sa.engine.Engine, tenant: str, include_budget: bool = False) -> dict[str, Any]:
    """En son kurulan 13 haftalık tablo. `include_budget`: bütçe temposu toplama katılır (satıcı borcuyla örtüşür)."""
    with engine.connect() as c:
        run = c.execute(sa.select(CASH_RUNS).where(CASH_RUNS.c.tenant_id == tenant).order_by(CASH_RUNS.c.run_at.desc())).first()
        if not run:
            return {"run": None, **freshness(engine)}
        rows = c.execute(sa.select(CASH_LINES).where(CASH_LINES.c.run_id == run.id).order_by(CASH_LINES.c.kalem, CASH_LINES.c.hafta)).all()
    params = _j(run.params_json, {})
    items: dict[str, dict[str, Any]] = {}
    n = max([r.hafta for r in rows] or [13])
    for r in rows:
        it = items.setdefault(r.kalem, {"kalem": r.kalem, "yon": r.yon, "kaynak": r.kaynak, "yaklasik": r.yaklasik,
                                        "haftalar": [0.0] * n, "ayrinti": [[] for _ in range(n)]})
        it["haftalar"][r.hafta - 1] = round(float(r.tutar or 0), 2)
        it["ayrinti"][r.hafta - 1] = _j(r.ayrinti_json, [])
    weeks = params.get("haftalar") or []
    if include_budget and "butce" in items:
        bal = float(params.get("acilisBakiye") or 0)
        for i, w in enumerate(weeks):
            extra = items["butce"]["haftalar"][i]
            w = dict(w, cikis=round(w["cikis"] + extra, 2))
            w["net"] = round(w["giris"] - w["cikis"], 2)
            w["acilis"] = round(bal, 2)
            bal = bal + w["net"]
            w["kapanis"] = round(bal, 2)
            w["acik"] = bal < 0
            weeks[i] = w
    for it in items.values():
        it["toplam"] = round(sum(it["haftalar"]), 2)
    order = {"giris": 0, "cikis": 1, "bilgi": 2}
    return {"run": {"id": run.id, "at": _iso(run.run_at), "by": run.run_by, "veriSonu": run.veri_son_gunu,
                    "baslangic": run.baslangic},
            "acilisBakiye": params.get("acilisBakiye"), "pozisyon": params.get("pozisyon") or {},
            "vadesiGecmis": params.get("vadesiGecmis") or {}, "dovizTelif": params.get("dovizTelif") or {},
            "hatalar": params.get("hatalar") or {}, "butceDahil": bool(include_budget),
            "kalemler": sorted(items.values(), key=lambda x: (order.get(x["yon"], 3), -abs(x["toplam"]))),
            "haftalar": weeks, "acikHafta": next((w for w in weeks if w.get("acik")), None),
            "bant": params.get("bant") or {"var": False, "neden": "Bu tablo olasılıklı bant eklenmeden önce kuruldu; yeniden kurun."},
            **freshness(engine)}


def cash_history(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    """Geçmiş tahminlerin haftalık neti ↔ gerçekleşen kasa+banka hareketi (100, 102). Veri son gününden sonraki
    haftalar gerçekleşmemiş sayılır."""
    end = data_end(engine)
    daily: dict[str, dict[str, float]] = {}
    for y in sorted(_loaded_years(engine)):
        daily.update(meta_get(engine, f"bank:{y}").get("days") or {})
    with engine.connect() as c:
        runs = c.execute(sa.select(CASH_RUNS).where(CASH_RUNS.c.tenant_id == tenant).order_by(CASH_RUNS.c.run_at.desc())).all()
    out = []
    errors = []
    for run in runs:
        weeks = _j(run.params_json, {}).get("haftalar") or []
        for w in weeks:
            s, e = date.fromisoformat(w["baslangic"]), date.fromisoformat(w["bitis"])
            if not end or e > end:
                continue
            gin = sum(float((daily.get((s + timedelta(days=k)).isoformat()) or {}).get("giris") or 0) for k in range(7))
            cik = sum(float((daily.get((s + timedelta(days=k)).isoformat()) or {}).get("cikis") or 0) for k in range(7))
            real = gin - cik
            out.append({"run": run.id, "runAt": _iso(run.run_at), "hafta": w["hafta"], "baslangic": w["baslangic"],
                        "tahminNet": w["net"], "gercekNet": round(real, 2), "gercekGiris": round(gin, 2), "gercekCikis": round(cik, 2),
                        "sapma": round(real - w["net"], 2)})
            errors.append(abs(real - w["net"]))
    return {"items": out, "ortalamaMutlakSapma": round(sum(errors) / len(errors), 2) if errors else None,
            "not": None if out else "Karşılaştırılacak geçmiş hafta yok: tahminin haftaları veri son gününden sonra.",
            **freshness(engine)}


# ------------------------------------------------------------------ vergi takvimi


def _check_day(v: Any, label: str) -> str:
    try:
        return date.fromisoformat(str(v)[:10]).isoformat()
    except (TypeError, ValueError):
        raise FinanceError(f"{label} YYYY-AA-GG biçiminde bir tarih olmalı.") from None


def _tax_dict(r: Any, today: date) -> dict[str, Any]:
    left = (date.fromisoformat(r.son_gun) - today).days
    return {"id": r.id, "beyan": r.beyan, "donem": r.donem, "sonGun": r.son_gun, "sorumlu": r.sorumlu, "durum": r.durum,
            "durumLabel": TAX_STATUSES.get(r.durum, r.durum), "tutar": r.tutar, "not": r.not_, "kalanGun": left,
            "gecikti": left < 0 and r.durum != "verildi", "hatirlatma": _j(r.hatirlatma_json, {}),
            "createdBy": r.created_by, "updatedBy": r.updated_by, "updatedAt": _iso(r.updated_at)}


def tax_list(engine: sa.engine.Engine, tenant: str, year: Optional[int] = None) -> dict[str, Any]:
    with engine.connect() as c:
        stmt = sa.select(TAX).where(TAX.c.tenant_id == tenant)
        if year:
            stmt = stmt.where(TAX.c.son_gun >= f"{int(year)}-01-01", TAX.c.son_gun <= f"{int(year)}-12-31")
        rows = c.execute(stmt.order_by(TAX.c.son_gun, TAX.c.beyan)).all()
    today = _today()
    items = [_tax_dict(r, today) for r in rows]
    return {"items": items, "statuses": TAX_STATUSES, "today": today.isoformat(),
            "yaklasan": [i for i in items if 0 <= i["kalanGun"] <= 14 and i["durum"] != "verildi"],
            "geciken": [i for i in items if i["gecikti"]]}


def _tax_values(body: dict[str, Any], partial: bool) -> dict[str, Any]:
    vals: dict[str, Any] = {}
    if not partial or "beyan" in body:
        b = _text(body.get("beyan"), 200)
        if not b:
            raise FinanceError("Beyan adı zorunlu.")
        vals["beyan"] = b
    if not partial or "sonGun" in body:
        vals["son_gun"] = _check_day(body.get("sonGun"), "Son gün")
    for k, col, lim in (("donem", "donem", 40), ("sorumlu", "sorumlu", 200), ("not", "not_", 2000)):
        if k in body:
            vals[col] = _text(body.get(k), lim)
    if "durum" in body or not partial:
        d = str(body.get("durum") or "bekliyor")
        if d not in TAX_STATUSES:
            raise FinanceError("Durum bekliyor, hazırlanıyor, hazır ya da verildi olmalı.")
        vals["durum"] = d
    if "tutar" in body:
        t = body.get("tutar")
        if t in (None, ""):
            vals["tutar"] = None
        else:
            try:
                x = float(str(t).replace(".", "").replace(",", ".")) if isinstance(t, str) else float(t)
            except ValueError:
                raise FinanceError("Tutar sayı olmalı.") from None
            if x < 0 or math.isnan(x) or math.isinf(x):
                raise FinanceError("Tutar sıfır ya da artı olmalı.")
            vals["tutar"] = x
    return vals


def tax_create(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    vals = _tax_values(body, partial=False)
    rid = uuid.uuid4().hex
    with engine.begin() as c:
        c.execute(TAX.insert().values(id=rid, tenant_id=tenant, created_by=user, created_at=_now(), **vals))
        r = c.execute(sa.select(TAX).where(TAX.c.id == rid)).first()
    return _tax_dict(r, _today())


def tax_update(engine: sa.engine.Engine, tenant: str, user: str, tid: str, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    vals = _tax_values(body, partial=True)
    with engine.begin() as c:
        r = c.execute(sa.select(TAX).where(TAX.c.id == tid, TAX.c.tenant_id == tenant)).first()
        if not r:
            raise FinanceError("Kayıt bulunamadı.", 404)
        diff = {k: {"once": getattr(r, k), "sonra": v} for k, v in vals.items() if getattr(r, k) != v}
        if "son_gun" in vals and vals["son_gun"] != r.son_gun:
            vals["hatirlatma_json"] = None
        c.execute(TAX.update().where(TAX.c.id == tid).values(**vals, updated_by=user, updated_at=_now()))
        r = c.execute(sa.select(TAX).where(TAX.c.id == tid)).first()
    return _tax_dict(r, _today()), diff


def tax_delete(engine: sa.engine.Engine, tenant: str, tid: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = c.execute(sa.select(TAX).where(TAX.c.id == tid, TAX.c.tenant_id == tenant)).first()
        if not r:
            raise FinanceError("Kayıt bulunamadı.", 404)
        c.execute(TAX.delete().where(TAX.c.id == tid))
    return {"id": tid, "beyan": r.beyan, "sonGun": r.son_gun}


def tax_copy_year(engine: sa.engine.Engine, tenant: str, user: str, src_year: int, dst_year: int) -> dict[str, Any]:
    """Bir yılın takvimini bir sonraki yıla kopyalar (aynı gün/ay; ayın son günü korunur). Durum «bekliyor», tutar
    boş. Mevzuat değişikliği kopyalanmaz: yeni yılın günlerini muhasebe denetler."""
    if dst_year == src_year:
        raise FinanceError("Kaynak ve hedef yıl farklı olmalı.")
    items = tax_list(engine, tenant, src_year)["items"]
    existing = {(i["beyan"], i["sonGun"]) for i in tax_list(engine, tenant, dst_year)["items"]}
    made = 0
    for i in items:
        d = date.fromisoformat(i["sonGun"])
        y = d.year + (dst_year - src_year)
        last = (date(y + (d.month == 12), d.month % 12 + 1, 1) - timedelta(days=1)).day
        nd = date(y, d.month, min(d.day, last))
        if (i["beyan"], nd.isoformat()) in existing:
            continue
        tax_create(engine, tenant, user, {"beyan": i["beyan"], "donem": None, "sonGun": nd.isoformat(), "sorumlu": i["sorumlu"],
                                          "durum": "bekliyor", "not": f"{src_year} takviminden kopyalandı; günü denetleyin."})
        made += 1
    return {"kopyalanan": made, "atlanan": len(items) - made}


def tax_due_reminders(engine: sa.engine.Engine, tenant: str, today: Optional[date] = None) -> list[dict[str, Any]]:
    """Son güne 7 ve 2 gün kala gönderilmemiş hatırlatmalar (verilen beyan hariç)."""
    today = today or _today()
    out = []
    with engine.connect() as c:
        rows = c.execute(sa.select(TAX).where(TAX.c.tenant_id == tenant, TAX.c.durum != "verildi",
                                              TAX.c.son_gun >= today.isoformat())).all()
    for r in rows:
        left = (date.fromisoformat(r.son_gun) - today).days
        sent = _j(r.hatirlatma_json, {})
        for k in (7, 2):
            if left <= k and str(k) not in sent:
                out.append({"id": r.id, "esik": k, "beyan": r.beyan, "donem": r.donem, "sonGun": r.son_gun, "kalanGun": left,
                            "sorumlu": r.sorumlu, "durum": TAX_STATUSES.get(r.durum, r.durum)})
                break
    return out


def mark_reminded(engine: sa.engine.Engine, items: list[dict[str, Any]]) -> None:
    now = _now().isoformat()
    with engine.begin() as c:
        for i in items:
            r = c.execute(sa.select(TAX.c.hatirlatma_json).where(TAX.c.id == i["id"])).first()
            sent = _j(r.hatirlatma_json if r else None, {})
            for k in (7, 2):
                if k >= i["esik"]:
                    sent.setdefault(str(k), now)
            c.execute(TAX.update().where(TAX.c.id == i["id"]).values(hatirlatma_json=_dump(sent)))


# ------------------------------------------------------------------ sapma notları


def notes(engine: sa.engine.Engine, tenant: str, year: int, tur: str = "sapma") -> list[dict[str, Any]]:
    with engine.connect() as c:
        rows = c.execute(sa.select(NOTES).where(NOTES.c.tenant_id == tenant, NOTES.c.year == int(year), NOTES.c.tur == tur)
                         .order_by(NOTES.c.tarih.desc())).all()
    return [{"id": r.id, "year": r.year, "month": r.month, "tur": r.tur, "anahtar": r.hedef_anahtar, "metin": r.metin,
             "durum": r.durum, "hazirlayan": r.hazirlayan, "onaylayan": r.onaylayan, "tarih": _iso(r.tarih)} for r in rows]


def save_note(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    """Sapma açıklaması (CFO'nun elle yazdığı neden cümlesi). Aynı anahtarın notu güncellenir."""
    key = _text(body.get("anahtar"), 200)
    text = _text(body.get("metin"), 4000)
    try:
        year = int(body.get("year"))
    except (TypeError, ValueError):
        raise FinanceError("Yıl zorunlu.") from None
    if not key or not text:
        raise FinanceError("Sapma anahtarı ve açıklama zorunlu.")
    month = body.get("month")
    with engine.begin() as c:
        cur = c.execute(sa.select(NOTES).where(NOTES.c.tenant_id == tenant, NOTES.c.year == year, NOTES.c.tur == "sapma",
                                               NOTES.c.hedef_anahtar == key)).first()
        if cur:
            c.execute(NOTES.update().where(NOTES.c.id == cur.id).values(metin=text, hazirlayan=user, tarih=_now()))
            nid = cur.id
        else:
            nid = uuid.uuid4().hex
            c.execute(NOTES.insert().values(id=nid, tenant_id=tenant, year=year, month=int(month) if month else None, tur="sapma",
                                            hedef_anahtar=key, metin=text, durum="taslak", hazirlayan=user, tarih=_now()))
    return {"id": nid, "anahtar": key, "metin": text}


# ------------------------------------------------------------------ bütçe–gerçekleşme (M46 okur)


def budget_view(engine: sa.engine.Engine, tenant: str, year: int) -> dict[str, Any]:
    """M46'nın yürürlükteki planı: şirket satış oranı, departman × hesap kullanımı, açık sapmalar, sapma notları.
    Sayılar M46'dan (aynı tanım); bu modül yeniden hesaplamaz."""
    from semantic_bridge import budget as B

    B.ensure(engine)
    try:
        tr = B.tracking(engine, tenant, int(year))
    except B.BudgetError as e:
        raise FinanceError(str(e), e.status) from e
    plan = tr.get("plan")
    if not plan:
        return {"year": int(year), "plan": None, "mesaj": "Bütçe onaylanmadı: bu yıl için yürürlükte plan yok.", **freshness(engine)}
    dep = B.departments(engine, tenant, plan["id"])
    dev = B.deviations(engine, tenant, int(year), status="acik")
    ns = {n["anahtar"]: n for n in notes(engine, tenant, int(year))}
    items = []
    for d in dep["items"]:
        iz = d.get("izleme") or {}
        key = f"gider|{d['merkezKodu']}|{d['hesap']}"
        items.append({"anahtar": key, "merkezKodu": d["merkezKodu"], "merkezAdi": d["merkezAdi"], "hesap": d["hesap"],
                      "hesapAdi": d["hesapAdi"], "yillik": d["yillik"], "butceDonem": iz.get("butceDonem"),
                      "gercek": iz.get("gercek"), "kullanim": iz.get("kullanim"), "durum": iz.get("durum"),
                      "sapma": _r((iz.get("gercek") or 0) - (iz.get("butceDonem") or 0)) if iz else None,
                      "not": ns.get(key)})
    items.sort(key=lambda i: -((i["sapma"] or 0)))
    devs = []
    for a in dev["items"]:
        key = f"{a['kind']}|{a['scope']}|{a['key']}"
        devs.append({**a, "anahtar": key, "not": ns.get(key)})
    return {"year": int(year), "plan": {k: plan[k] for k in ("id", "title", "version", "scenarioLabel", "statusLabel", "decidedAt", "decidedBy") if k in plan},
            "asof": tr.get("asof"), "esik": tr.get("esik"), "sirket": tr.get("sirket"), "gider": tr.get("gider"),
            "aylar": tr.get("aylar"), "departmanlar": items, "sapmalar": devs, "sapmaToplam": dev["total"],
            **freshness(engine)}


# ------------------------------------------------------------------ özet


def summary(engine: sa.engine.Engine, tenant: str, *, with_cash: bool) -> dict[str, Any]:
    """Özet kartlar: net satış (yıl başından, geçen yılın aynı dönemiyle), brüt kâr (maliyetli satırlar), faaliyet
    gideri, nakit pozisyonu, vadesi geçmiş alacak, bütçe durumu; «bu ay dikkat» listesi. Her kartta veri son günü."""
    end = data_end(engine)
    fr = freshness(engine)
    if not end:
        return {"cards": [], "dikkat": [], "hazir": False, **fr}
    y, m = end.year, end.month
    ytd = [(y, k) for k in range(1, m + 1)]
    s = _sales_totals(engine, ytd)
    prev = meta_get(engine, f"sales_prev:{y}")
    cov = cost_coverage(s)
    cards = []
    net = s.get("net")
    cards.append({"id": "net-satis", "label": "Net satış (yıl başından)", "value": _r(net), "unit": "₺",
                  "compare": _r(prev.get("net")), "compareLabel": f"Geçen yıl aynı dönem (1 Ocak – {end.day} {AY[end.month - 1]} {y - 1})",
                  "change": _r((net - prev["net"]) / prev["net"], 4) if net is not None and prev.get("net") else None,
                  "note": "Faturalı satış satırları; satır iskontosu ve iade düşülmüş.", "yaklasik": False, "sekme": "gelir"})
    mnet, mcost = s.get("maliyetli_net") or 0.0, s.get("maliyet") or 0.0
    cards.append({"id": "brut-kar", "label": "Brüt kâr (maliyeti işlenmiş satışlar)", "value": _r(mnet - mcost), "unit": "₺",
                  "ratio": _r((mnet - mcost) / mnet, 4) if mnet else None,
                  "note": f"Satışların %{_pct(cov['maliyetliPay'] or 0)}'inde maliyet işlenmiş; maliyeti işlenmemiş satış marja katılmadı."
                  + (f" Maliyeti işlenmiş son satış: {fr['maliyetSonu']}." if fr.get("maliyetSonu") else ""),
                  "yaklasik": cov["yaklasik"], "sekme": "karlilik"})
    loaded = y in _loaded_years(engine)
    if loaded:
        vals = compute_lines(engine, tenant, ytd)["values"]
        cards.append({"id": "faaliyet", "label": "Faaliyet giderleri (yıl başından)", "value": _r(-vals.get("FAALIYET", 0.0)),
                      "unit": "₺", "note": "Muhasebe (Ar-Ge, pazarlama-satış-dağıtım, genel yönetim), eşlemeyle.",
                      "yaklasik": False, "sekme": "gelir"})
    if with_cash:
        pos = meta_get(engine, f"position:{y}")
        if pos.get("groups"):
            g = pos["groups"]
            cards.append({"id": "nakit", "label": "Kasa ve banka", "value": _r(float(g.get("100") or 0) + float(g.get("102") or 0)),
                          "unit": "₺", "note": f"{end.isoformat()} itibarıyla muhasebe bakiyesi (100 + 102).", "yaklasik": False,
                          "sekme": "nakit"})
        c = cash(engine, tenant)
        if c.get("run"):
            od = float((c.get("vadesiGecmis") or {}).get("alacak") or 0)
            cards.append({"id": "vadesi-gecmis", "label": "Vadesi geçmiş alacak", "value": _r(od), "unit": "₺",
                          "note": "Ödeme kapama kullanılmadığı için FIFO yaklaşımı.", "yaklasik": True, "sekme": "nakit"})
    bud = None
    try:
        from semantic_bridge import budget as B

        B.ensure(engine)
        tr = B.tracking(engine, tenant, y)
        if tr.get("plan") and tr.get("sirket"):
            bud = tr
            cards.append({"id": "butce", "label": "Satış hedefine göre (bugüne beklenen)", "value": tr["sirket"].get("oran"),
                          "unit": "oran", "state": tr["sirket"].get("durum"),
                          "note": f"Gider bütçesi kullanımı %{_pct(tr['gider'].get('kullanim') or 0)}" if (tr.get("gider") or {}).get("kullanim") is not None else None,
                          "yaklasik": False, "sekme": "butce"})
        else:
            cards.append({"id": "butce", "label": "Bütçe", "value": None, "unit": "metin", "note": "Bütçe onaylanmadı.",
                          "yaklasik": False, "sekme": "butce"})
    except Exception as e:  # noqa: BLE001 — bütçe okunamazsa kart düşer, özet düşmez
        log.warning("finance: bütçe özeti okunamadı: %s", e)
    dikkat = []
    if bud:
        try:
            from semantic_bridge import budget as B

            for a in B.deviations(engine, tenant, y, status="acik")["items"][:3]:
                dikkat.append({"tur": "sapma", "metin": f"{a['label']}: beklenenin %{_pct(a['ratio'] or 0)}'i" if a["kind"] == "satis"
                               else f"{a['label']}: bütçenin %{_pct(a['ratio'] or 0)}'i kullanıldı", "tutar": a.get("gap"), "sekme": "butce"})
        except Exception:  # noqa: BLE001
            pass
    if with_cash:
        c = cash(engine, tenant)
        if c.get("acikHafta"):
            w = c["acikHafta"]
            dikkat.append({"tur": "nakit", "metin": f"{w['hafta']}. hafta ({w['baslangic']}) nakit açığı; en büyük çıkış: {w.get('enBuyukCikis') or '—'}",
                           "tutar": w["kapanis"], "sekme": "nakit"})
    tl = tax_list(engine, tenant)
    for t in tl["geciken"][:3] + tl["yaklasan"][:3]:
        dikkat.append({"tur": "vergi", "metin": f"{t['beyan']}{' (' + t['donem'] + ')' if t['donem'] else ''}: son gün {t['sonGun']}"
                       + (" — gecikti" if t["gecikti"] else f", {t['kalanGun']} gün kaldı"), "tutar": t["tutar"], "sekme": "vergi"})
    return {"cards": cards, "dikkat": dikkat, "hazir": True, "donem": period_label(ytd), **fr}


def summary_text(data: dict[str, Any], link: str = "") -> str:
    """Sabah özeti e-postası (düz metin). Rakamlar özet kartlarından; model yok."""
    lines_ = [f"Finansal özet — veri {data.get('veriSonu') or '—'} tarihine kadar ({data.get('donem') or ''})", ""]
    for c in data.get("cards") or []:
        v = c.get("value")
        if c.get("unit") == "₺" and v is not None:
            txt = f"{_tr_money(v)} ₺"
        elif c.get("unit") == "oran" and v is not None:
            txt = f"%{_pct(v)}"
        else:
            txt = c.get("note") or "—"
        extra = f" (yaklaşık)" if c.get("yaklasik") else ""
        lines_.append(f"- {c['label']}: {txt}{extra}")
    if data.get("dikkat"):
        lines_ += ["", "Bu ay dikkat:"] + [f"- {d['metin']}" for d in data["dikkat"]]
    lines_ += ["", f"Ayrıntı: {link}" if link else "Ayrıntı Finansal raporlar ekranında."]
    return "\n".join(lines_)


def _tr_money(v: float) -> str:
    return f"{v:,.0f}".replace(",", ".")


# ------------------------------------------------------------------ aylık finansal yorum taslağı (Zeki AI, CFO onayı)
#
# Olgular kodla hesaplanır (gelir tablosu, yıl başından net satış, maliyet kapsamı, kanal kârlılığı, «bu ay dikkat»);
# değişim oranları da burada hesaplanır, model hesap yapmaz. Zeki AI yalnız bu olguları 5–8 cümleyle anlatır
# (`zeki_text.interpret`: her sayı olgularda olmalı, yoksa kural metni). Metin `semantic_finance_notes`'ta
# `tur='ozet'`, `hedef_anahtar='ozet|YYYY-AA'` satırında «taslak» durur; CFO düzeltir ve açıkça verilen
# `ozellik:finans.yorum-onay` ile onaylar. Onaylı metin DYK kurul panelinde M45 göstergesinin ayrıntısına gider.

COMMENT_LINES = (("NET_SATIS", "Net satışlar", False), ("BRUT_KAR", "Brüt satış kârı", True),
                 ("FAALIYET", "Faaliyet giderleri", False), ("FAALIYET_KARI", "Faaliyet kârı", True),
                 ("NET_KAR", "Dönem net kârı", True))


def _money_tr(v: float) -> str:
    return f"{_tr_money(abs(v))} ₺"


def _chg(cur: Optional[float], ref: Optional[float], magnitude: bool = False) -> Optional[str]:
    """Değişim oranı yazısı («+%12,3» / «−%4,0»). `magnitude`: gider satırında tutarın büyüklüğü karşılaştırılır
    (gider −120 ← −100 «+%20» artış demektir)."""
    if cur is None or ref is None or abs(ref) < 0.005:
        return None
    ch = (abs(cur) - abs(ref)) / abs(ref) if magnitude else (cur - ref) / abs(ref)
    return ("+" if ch >= 0 else "−") + "%" + _pct(abs(ch))


def comment_key(year: int, month: int) -> str:
    return f"ozet|{int(year):04d}-{int(month):02d}"


def comment_facts(engine: sa.engine.Engine, tenant: str, year: int, month: int, *, with_cash: bool,
                  unit_costs: Callable[[list[str]], dict[str, dict[str, Any]]] = lambda c: {},
                  snapshot: Optional[dict] = None) -> list[str]:
    """Ayın yorum olguları (Türkçe cümleler, rakamlar yazılı biçimiyle). Model yalnız bunları görür."""
    if not 1 <= int(month) <= 12:
        raise FinanceError("Ay 1 ile 12 arasında olmalı.")
    y, m = int(year), int(month)
    end = data_end(engine)
    if end is None:
        raise FinanceError("Logo verisi henüz okunmadı; yorum taslağı hazırlanamaz.", 409)
    facts: list[str] = []
    label = f"{AY[m - 1]} {y}"
    if (y, m) > (end.year, end.month):
        raise FinanceError(f"{label} için veri yok (veri {end.isoformat()} tarihinde bitiyor).", 409)
    if (y, m) == (end.year, end.month):
        facts.append(f"{label} tamamlanmadı: veri {end.isoformat()} tarihine kadar.")
    p = pnl(engine, tenant, y, m, "ay", ("onceki", "gecen-yil", "butce"))
    rows = {r["kod"]: r for r in p["rows"]}
    cols = p["columns"]
    if cols.get("donem", {}).get("loaded"):
        for kod, name, signed in COMMENT_LINES:
            r = rows.get(kod)
            if not r or r["values"].get("donem") is None:
                continue
            cur = r["values"]["donem"]
            word = name if (not signed or cur >= 0) else name.replace("kârı", "zararı")
            parts = [f"{label} {word.lower()}: {_money_tr(cur)}"]
            for col, lab in (("onceki", cols.get("onceki", {}).get("label")), ("gecenYil", cols.get("gecenYil", {}).get("label"))):
                ref = r["values"].get(col)
                if ref is None or not cols.get(col, {}).get("loaded", True):
                    continue
                ch = _chg(cur, ref, magnitude=not signed)
                parts.append(f"{lab}: {_money_tr(ref)}" + (f" ({ch})" if ch else ""))
            b = r["values"].get("butce")
            if b is not None and (cols.get("butce") or {}).get("plan"):
                parts.append(f"bütçe: {_money_tr(b)}")
            facts.append("; ".join(parts) + ".")
            if r.get("yaklasik"):
                facts.append(f"{name} yaklaşıktır: " + " ".join(r.get("notlar") or []))
    else:
        facts.append(f"{label} muhasebe fişleri henüz okunmadı; gelir tablosu satırları yok.")
    cov = p.get("maliyet") or {}
    if cov.get("maliyetliPay") is not None:
        facts.append(f"{label} satışlarının %{_pct(cov['maliyetliPay'])}'inde maliyet işlenmiş.")
    if p.get("mutabakatFarki") is not None and abs(p["mutabakatFarki"]) >= 1:
        facts.append(f"Muhasebedeki net satış ile faturalı satış arasında {_money_tr(p['mutabakatFarki'])} fark var.")
    ytd = _sales_totals(engine, [(y, k) for k in range(1, m + 1)])
    if ytd.get("net") is not None:
        prev = meta_get(engine, f"sales_prev:{y}")
        ch = _chg(ytd["net"], prev.get("net")) if prev.get("net") and (y, m) == (end.year, end.month) else None
        facts.append(f"Yıl başından net satış (faturalı satır): {_money_tr(ytd['net'])}"
                     + (f"; geçen yılın aynı dönemine göre {ch}" if ch else "") + ".")
    try:
        k = profitability(engine, by="kanal", year=y, frm=m, to=m, unit_costs=unit_costs, snapshot=snapshot,
                          with_royalty=False, all_rows=True)
        top = [r for r in (k.get("items") or []) if (r.get("net") or 0) > 0][:3]
        if top:
            facts.append(f"{label} en yüksek net satışlı kanallar: " + "; ".join(
                f"{r['ad']} {_money_tr(r['net'])}" + (f" (kesin marj %{_pct(r['marjKesin'])})" if r.get("marjKesin") is not None else "")
                for r in top) + ".")
    except FinanceError as e:
        log.info("finance: yorum için kanal kârlılığı okunamadı: %s", e)
    s = summary(engine, tenant, with_cash=with_cash)
    for d in (s.get("dikkat") or [])[:4]:
        facts.append(f"Dikkat: {d['metin']}" + (f" ({_money_tr(d['tutar'])})" if d.get("tutar") is not None else "") + ".")
    return facts


def comment_rule_text(facts: list[str]) -> str:
    """Modelsiz yorum: olguların ilk sekizi, olduğu gibi (kural metni)."""
    return " ".join(facts[:8])


COMMENT_TASK = ("Aşağıdaki olgulardan bir yayınevinin mali işler direktörü için aylık finansal yorum taslağı yaz. Sıra: "
                "satışlar, kârlılık, giderler, dikkat edilecekler. Olgularda yazan karşılaştırmaları kullan; neden "
                "uydurma, tahmin yapma. «Yaklaşık» denen rakamı yaklaşık diye an.")


def comment_draft(engine: sa.engine.Engine, tenant: str, user: str, year: int, month: int, llm: Any, *, with_cash: bool,
                  unit_costs: Callable[[list[str]], dict[str, dict[str, Any]]] = lambda c: {},
                  snapshot: Optional[dict] = None) -> dict[str, Any]:
    """Taslak üretir ve «taslak» olarak yazar (onaylı metin varsa yeni taslak onu geri taslağa çevirir). Dönen:
    `comment_get` biçimi + `kaynak`/`neden`."""
    from semantic_bridge import zeki_text as Z

    facts = comment_facts(engine, tenant, year, month, with_cash=with_cash, unit_costs=unit_costs, snapshot=snapshot)
    res = Z.interpret(facts, comment_rule_text(facts), llm=llm, task=COMMENT_TASK, min_sentences=5, max_sentences=8,
                      max_chars=2400, max_tokens=900)
    _comment_write(engine, tenant, user, year, month, res.metin)
    meta_set(engine, _comment_meta_key(tenant, year, month),
             {"kaynak": res.kaynak, "neden": res.neden, "olgular": facts, "at": _now().isoformat(), "by": user})
    return comment_get(engine, tenant, year, month)


def _comment_meta_key(tenant: str, year: int, month: int) -> str:
    return f"yorum:{tenant[:30]}:{int(year):04d}-{int(month):02d}"


def _comment_row(c: Any, tenant: str, year: int, month: int) -> Any:
    return c.execute(sa.select(NOTES).where(NOTES.c.tenant_id == tenant, NOTES.c.tur == "ozet",
                                            NOTES.c.hedef_anahtar == comment_key(year, month))).first()


def _comment_write(engine: sa.engine.Engine, tenant: str, user: str, year: int, month: int, text: str) -> str:
    with engine.begin() as c:
        cur = _comment_row(c, tenant, year, month)
        if cur:
            c.execute(NOTES.update().where(NOTES.c.id == cur.id).values(metin=text, durum="taslak", hazirlayan=user,
                                                                         onaylayan=None, tarih=_now()))
            return cur.id
        nid = uuid.uuid4().hex
        c.execute(NOTES.insert().values(id=nid, tenant_id=tenant, year=int(year), month=int(month), tur="ozet",
                                        hedef_anahtar=comment_key(year, month), metin=text, durum="taslak",
                                        hazirlayan=user, tarih=_now()))
        return nid


def comment_get(engine: sa.engine.Engine, tenant: str, year: int, month: int) -> dict[str, Any]:
    """Ayın yorumu: metin, durum (taslak | onayli), kaynak (zeki | kural | insan), hazırlayan/onaylayan, olgularda
    olmayan sayılar (insan düzeltmesinde gözden geçirme için; engel değil)."""
    from semantic_bridge import zeki_text as Z

    with engine.connect() as c:
        r = _comment_row(c, tenant, year, month)
    meta = meta_get(engine, _comment_meta_key(tenant, year, month))
    base = {"year": int(year), "month": int(month), "donem": f"{AY[int(month) - 1]} {int(year)}"}
    if not r:
        return {**base, "metin": None, "durum": None}
    facts = meta.get("olgular") or []
    return {**base, "id": r.id, "metin": r.metin, "durum": r.durum, "hazirlayan": r.hazirlayan, "onaylayan": r.onaylayan,
            "tarih": _iso(r.tarih), "kaynak": meta.get("kaynak"), "neden": meta.get("neden"), "olgular": facts,
            "olguDisiSayilar": Z.unsupported(r.metin, facts) if facts else []}


def comment_save(engine: sa.engine.Engine, tenant: str, user: str, year: int, month: int, text: Any) -> tuple[dict[str, Any], dict[str, Any]]:
    """CFO düzeltmesi: metin taslak olarak kalır (onaylıysa onayı düşer). Dönen: (yorum, fark)."""
    t = _text(text, 6000)
    if not t:
        raise FinanceError("Yorum metni boş olamaz.")
    before = comment_get(engine, tenant, year, month)
    if before.get("metin") is None:
        raise FinanceError("Bu ay için yorum taslağı yok; önce taslak hazırlayın.", 404)
    _comment_write(engine, tenant, user, year, month, t)
    meta = meta_get(engine, _comment_meta_key(tenant, year, month))
    if meta:
        meta_set(engine, _comment_meta_key(tenant, year, month), {**meta, "kaynak": "insan", "neden": None})
    return comment_get(engine, tenant, year, month), {"metin": {"eski": (before.get("metin") or "")[:300], "yeni": t[:300]},
                                                      "durum": {"eski": before.get("durum"), "yeni": "taslak"}}


def comment_approve(engine: sa.engine.Engine, tenant: str, user: str, year: int, month: int) -> dict[str, Any]:
    with engine.begin() as c:
        r = _comment_row(c, tenant, year, month)
        if not r:
            raise FinanceError("Onaylanacak yorum taslağı yok.", 404)
        if r.durum == "onayli":
            raise FinanceError("Yorum zaten onaylı.", 409)
        c.execute(NOTES.update().where(NOTES.c.id == r.id).values(durum="onayli", onaylayan=user))
    return comment_get(engine, tenant, year, month)


def approved_comment(engine: sa.engine.Engine, tenant: str) -> Optional[dict[str, Any]]:
    """En son onaylı aylık yorum (DYK kurul paneli okur)."""
    with engine.connect() as c:
        r = c.execute(sa.select(NOTES).where(NOTES.c.tenant_id == tenant, NOTES.c.tur == "ozet", NOTES.c.durum == "onayli")
                      .order_by(NOTES.c.year.desc(), NOTES.c.month.desc())).first()
    if not r:
        return None
    return {"metin": r.metin, "donem": f"{AY[int(r.month or 1) - 1]} {r.year}", "onaylayan": r.onaylayan, "tarih": _iso(r.tarih)}


# ------------------------------------------------------------------ Excel


def pnl_xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Gelir tablosu"
    ws.append([f"Gelir tablosu — {data['columns']['donem']['label']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Veri son günü: {data.get('veriSonu') or '—'}. Tutar = kâr etkisi (alacak − borç): gelir artı, gider eksi."])
    ws.append([])
    keys = list(data["columns"].keys())
    ws.append(["Satır"] + [data["columns"][k]["label"] for k in keys] + ["Not"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in data["rows"]:
        ws.append([r["ad"]] + [r["values"].get(k) for k in keys] + ["; ".join(r["notlar"])])
        if r["tur"] == "ara_toplam":
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, min_col=2, max_col=1 + len(keys)):
        for c in row:
            c.number_format = "#,##0.00;[Red]-#,##0.00"
    ws.column_dimensions["A"].width = 58
    for i in range(len(keys)):
        ws.column_dimensions[chr(66 + i)].width = 20
    ws.append([])
    m = data.get("mizan") or {}
    ws.append([f"Mizan denkliği: {m.get('durum')}; fark {m.get('fark')}"])
    ws.append([f"Rapor dışı (dışlanan hesaplar): {data.get('dislanan')}; eşleme onayı bekleyen tutar: {data['esleme']['onaysizTutar']}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ Logo okuması


class Refresher:
    """Logo/CRM okumasını arka planda yapar; aynı anda tek okuma. Nakit tablosu da buradan kurulur."""

    def __init__(self, engine_fn: Callable[[], sa.engine.Engine], tenant_fn: Callable[[], str], logo_file: Callable[[], str],
                 crm_file: Callable[[], str], unit_costs: Callable[[list[str]], dict[str, dict[str, Any]]]):
        self._engine = engine_fn
        self._tenant = tenant_fn
        self._logo = logo_file
        self._crm = crm_file
        self._unit_costs = unit_costs
        self._thread: Optional[threading.Thread] = None
        self._guard = threading.Lock()
        self.state: dict[str, Any] = {"running": False, "step": None, "startedAt": None, "error": None}

    def status(self) -> dict[str, Any]:
        engine = self._engine()
        ensure(engine)
        years: dict[str, Any] = {}
        with engine.connect() as c:
            for r in c.execute(sa.select(META).where(META.c.key.like("ledger:%"))).all():
                years[r.key.split(":")[1]] = {**_j(r.value_json, {}), "at": _iso(r.updated_at)}
        return {**self.state, **freshness(engine), "years": dict(sorted(years.items())),
                "cash": meta_get(engine, "cash_run")}

    def running(self) -> bool:
        return bool(self._thread and self._thread.is_alive())

    def years(self, engine: sa.engine.Engine) -> list[int]:
        end = data_end(engine)
        top = end.year if end else _today().year
        return [top - 1, top]

    def due(self, engine: sa.engine.Engine, now: Optional[float] = None) -> list[int]:
        """Tazelenmesi gereken yıllar: son yıl `FINANCE_ACTUALS_TTL` (1 sa), önceki yıl 7 gün."""
        now = now or time.time()
        ys = self.years(engine)
        top = max(ys)
        ttl_now = int(conf("FINANCE_ACTUALS_TTL", "3600"))
        ttl_past = int(conf("FINANCE_ACTUALS_PAST_TTL", str(7 * 86400)))
        out = []
        for y in ys:
            at = meta_get(engine, f"ledger:{y}").get("_at")
            age = now - datetime.fromisoformat(at).timestamp() if at else None
            if age is None or age > (ttl_now if y >= top else ttl_past):
                out.append(y)
        return out

    def start(self, years: Optional[list[int]] = None, cash: bool = True, user: Optional[str] = None) -> bool:
        with self._guard:
            if self.running():
                return False
            self.state.update(running=True, step="Başlıyor", startedAt=time.time(), error=None)
            self._thread = threading.Thread(target=self.run, args=(years, cash, user), daemon=True, name="finance-refresh")
            self._thread.start()
            return True

    def run(self, years: Optional[list[int]] = None, cash_too: bool = False, user: Optional[str] = None) -> dict[str, Any]:
        engine = self._engine()
        ensure(engine)
        done: dict[str, Any] = {}
        try:
            self.state.update(running=True, step="Logo dönemleri")
            logo = src.runner(self._logo())
            firms = src.firms_by_year(logo)
            end, costed = src.read_data_end(logo, firms)
            if end:
                meta_set(engine, "data_end", {**meta_get(engine, "data_end"), "date": end.isoformat(),
                                              "costed": costed.isoformat() if costed else None})
            want = self.years(engine) if years is None else years
            for y in want:
                if y not in firms:
                    continue
                done[str(y)] = self._read_year(engine, logo, firms, y, end)
            self.state["step"] = "Hesap planı"
            last = firms[max(firms)]
            accts = [{"hesap_kodu": str(r["hesap"]).strip()[:60], "ad": (src.clean(r.get("ad")) or "")[:300] or None}
                     for r in logo(src.accounts_sql(last)) if r.get("hesap")]
            with engine.begin() as c:
                c.execute(ACCOUNTS.delete())
                seen = set()
                rows = [a for a in accts if not (a["hesap_kodu"] in seen or seen.add(a["hesap_kodu"]))]
                for i in range(0, len(rows), 5000):
                    c.execute(ACCOUNTS.insert(), rows[i:i + 5000])
            if cash_too and end:
                done["nakit"] = self._build_cash(engine, logo, firms, end, user)
            self.state.update(running=False, step=None, error=None, finishedAt=time.time())
            return {"ok": True, "done": done, "dataEnd": end.isoformat() if end else None}
        except Exception as e:  # noqa: BLE001 — eski anlık görüntü kalır, hata ekranda
            log.warning("finance refresh failed: %s", e)
            msg = str(e) if isinstance(e, (src.SourceError, FinanceError)) else f"Okuma hata verdi: {str(e)[:200]}"
            self.state.update(running=False, step=None, error=msg, finishedAt=time.time())
            return {"ok": False, "error": msg, "done": done}

    def _replace(self, engine: sa.engine.Engine, table: sa.Table, year: int, rows: list[dict[str, Any]]) -> None:
        with engine.begin() as c:
            c.execute(table.delete().where(table.c.year == year))
            for i in range(0, len(rows), 5000):
                c.execute(table.insert(), rows[i:i + 5000])

    def _read_year(self, engine: sa.engine.Engine, logo: Any, firms: dict[int, str], y: int, end: Optional[date]) -> dict[str, Any]:
        firm = firms[y]
        f = src.f
        self.state["step"] = f"{y} muhasebe"
        t0 = time.monotonic()
        merged: dict[tuple, dict[str, Any]] = {}
        for r in logo(src.account_actuals_sql(firm, y)):
            code = str(r.get("hesap") or "").strip()[:60]
            if not code:
                continue
            key = (int(r["ay"]), code, (str(r.get("merkez_kodu") or src.NO_CENTER).strip() or src.NO_CENTER)[:60], str(r.get("kural") or "dahil"))
            cur = merged.setdefault(key, {"year": y, "month": key[0], "hesap_kodu": key[1], "merkez_kodu": key[2], "kural": key[3],
                                          "hesap_adi": (src.clean(r.get("hesap_adi")) or "")[:300] or None, "borc": 0.0, "alacak": 0.0, "satir": 0})
            cur["borc"] += f(r.get("borc"))
            cur["alacak"] += f(r.get("alacak"))
            cur["satir"] += int(f(r.get("satir")))
        self._replace(engine, ACTUALS, y, list(merged.values()))
        trial = {str(int(r["ay"])): {"borc": f(r.get("borc")), "alacak": f(r.get("alacak")), "satir": int(f(r.get("satir"))),
                                     "son": src.to_day(r.get("son"))} for r in logo(src.trial_sql(firm, y))}
        meta_set(engine, f"trial:{y}", {"months": trial, "firm": firm})
        ledger_last = max((v["son"] for v in trial.values() if v.get("son")), default=None)
        if end and y == end.year and ledger_last:
            meta_set(engine, "data_end", {**meta_get(engine, "data_end"), "ledger": ledger_last})
        ms1 = int((time.monotonic() - t0) * 1000)
        self.state["step"] = f"{y} satış"
        rows = []
        for r in logo(src.sales_month_sql(firm, y)):
            rows.append({"year": y, "month": int(r["ay"]), **{k: f(r.get(k)) for k in _MEASURES},
                         "satis_net": f(r.get("satis_net")), "iade_net": f(r.get("iade_net")),
                         "satis_satir": int(f(r.get("satis_satir"))), "maliyetsiz_satir": int(f(r.get("maliyetsiz_satir")))})
        self._replace(engine, SALES_MONTH, y, rows)
        if end and y == end.year:
            ps = logo(src.sales_period_sql(firms.get(y - 1, firm) if (y - 1) in firms else firm, date(y - 1, 1, 1),
                                           _same_day_last_year(end) + timedelta(days=1))) if (y - 1) in firms else []
            meta_set(engine, f"sales_prev:{y}", {k: f((ps[0] if ps else {}).get(k)) for k in _MEASURES} if ps else {})
            pos = {str(r["grup"]): f(r.get("borc")) - f(r.get("alacak")) for r in logo(src.position_sql(firm, y, end))}
            meta_set(engine, f"position:{y}", {"groups": pos, "asof": end.isoformat()})
        self.state["step"] = f"{y} kârlılık"
        items = []
        for r in logo(src.profit_items_sql(firm, y)):
            code = str(r.get("stok_kodu") or "").strip()
            if not code:
                continue
            items.append({"year": y, "month": int(r["ay"]), "stok_kodu": code[:60], "kanal": (src.clean(r.get("kanal")) or "Grup kodu boş")[:120],
                          **{k: f(r.get(k)) for k in _MEASURES}})
        items = _merge(items, ("year", "month", "stok_kodu", "kanal"), _MEASURES)
        self._replace(engine, PROFIT_ITEMS, y, items)
        clients = {}
        for r in logo(src.profit_clients_sql(firm, y)):
            key = (int(r["ay"]), str(r.get("cari_kodu") or "#YOK").strip()[:60], (src.clean(r.get("kanal")) or "Grup kodu boş")[:120])
            cur = clients.setdefault(key, {"year": y, "month": key[0], "cari_kodu": key[1], "kanal": key[2],
                                           "cari_adi": (src.clean(r.get("cari_adi")) or "")[:300] or None,
                                           **{k: 0.0 for k in _MEASURES}, "tahmini_maliyet": 0.0, "tahmini_adet": 0.0})
            for k in _MEASURES:
                cur[k] += f(r.get(k))
        unc = logo(src.client_uncosted_sql(firm, y))
        codes = sorted({str(r.get("stok_kodu") or "").strip() for r in unc if r.get("stok_kodu")})
        costs = self._unit_costs(codes) if codes else {}
        for r in unc:
            code = str(r.get("stok_kodu") or "").strip()
            uc = costs.get(code)
            if not uc or uc.get("maliyet") is None:
                continue
            key = (int(r["ay"]), str(r.get("cari_kodu") or "#YOK").strip()[:60], (src.clean(r.get("kanal")) or "Grup kodu boş")[:120])
            if key in clients:
                clients[key]["tahmini_maliyet"] += f(r.get("adet")) * float(uc["maliyet"])
                clients[key]["tahmini_adet"] += f(r.get("adet"))
        self._replace(engine, PROFIT_CLIENTS, y, list(clients.values()))
        if end and y == end.year:
            self.state["step"] = f"{y} kasa ve banka hareketleri"
            days = {src.to_day(r.get("gun")): {"giris": f(r.get("giris")), "cikis": f(r.get("cikis"))}
                    for r in logo(src.bank_daily_sql(firm, y)) if r.get("gun")}
            meta_set(engine, f"bank:{y}", {"days": days})
        info = {"firm": firm, "hesapSatir": len(merged), "satisAy": len(rows), "kitapKanal": len(items), "cariKanal": len(clients),
                "muhasebeMs": ms1, "net": round(sum(r["net"] for r in rows), 2)}
        meta_set(engine, f"ledger:{y}", info)
        return info

    def _build_cash(self, engine: sa.engine.Engine, logo: Any, firms: dict[int, str], end: date, user: Optional[str]) -> dict[str, Any]:
        self.state["step"] = "Nakit tablosu"
        firm = firms[end.year]
        tenant = self._tenant()
        errors: dict[str, str] = {}
        f = src.f

        def safe(name: str, fn: Callable[[], Any], default: Any) -> Any:
            try:
                return fn()
            except Exception as e:  # noqa: BLE001 — bir kaynak okunamazsa tablo o satırsız kurulur, neden yazılır
                log.warning("finance cash source %s failed: %s", name, e)
                errors[name] = str(e)[:300] if isinstance(e, src.SourceError) else f"Okunamadı: {str(e)[:200]}"
                return default

        position = safe("pozisyon", lambda: {str(r["grup"]): f(r.get("borc")) - f(r.get("alacak"))
                                             for r in logo(src.position_sql(firm, end.year, end))}, {})
        rec = safe("alacak", lambda: [(src.day(r["vade"]), f(r.get("tutar")), int(f(r.get("cari"))))
                                      for r in logo(src.fifo_due_sql(firm, end.year, end, payable=False)) if src.day(r.get("vade"))], [])
        pay = safe("satici", lambda: [(src.day(r["vade"]), f(r.get("tutar")), int(f(r.get("cari"))))
                                      for r in logo(src.fifo_due_sql(firm, end.year, end, payable=True)) if src.day(r.get("vade"))], [])
        chq = safe("cek", lambda: [(r.get("tur"), r.get("durum"), src.day(r.get("vade")), f(r.get("tutar")))
                                   for r in logo(src.cheques_sql(firm))], [])

        def crm_rows():
            run = src.runner(self._crm())
            return [(src.day(r["vade"]), f(r.get("tutar"))) for r in run(src.crm_pending_collections_sql()) if src.day(r.get("vade"))]

        crm = safe("crm", crm_rows, [])

        def royalty_rows():
            from semantic_bridge import contracts as C

            out = []
            for p in C.due_list(engine, tenant, status="planlandi")["items"]:
                d = src.day(p.get("dueOn"))
                if d:
                    out.append((d, float(p.get("amount") or 0), p.get("currency") or "TRY"))
            return out

        royalty = safe("telif", royalty_rows, [])
        tax = [(date.fromisoformat(t["sonGun"]), float(t["tutar"]), t["beyan"]) for t in tax_list(engine, tenant)["items"]
               if t.get("tutar") and t["durum"] != "verildi"]

        def budget_rows():
            from semantic_bridge import budget as B

            out = []
            for y in (end.year, end.year + 1):
                with engine.connect() as c:
                    row = B._approved(c, tenant, y)
                    if row is None:
                        continue
                    depts = c.execute(sa.select(B.DEPTS).where(B.DEPTS.c.plan_id == row.id)).all()
                for m in range(1, 13):
                    # Nakitte üretim giderleri (baskı, telif merkezleri) de çıkıştır: bütün departman satırları.
                    total = sum(float(_j(d.aylar_json, [0.0] * 12)[m - 1] or 0) for d in depts)
                    if total:
                        out.append((date(y, m, 1), total))
            return out

        budget = safe("butce", budget_rows, [])
        result = build_cash({"asof": end, "position": position, "receivables": rec, "payables": pay, "cheques": chq, "crm": crm,
                             "royalty": royalty, "tax": tax, "budget": budget, "errors": errors})
        self.state["step"] = "Nakit bandı (tahmin)"
        try:
            result["bant"] = self._cash_band(logo, firms, end, result)
        except Exception as e:  # noqa: BLE001 — bant okunamazsa tablo kuraldan kalır
            log.warning("finance cash band failed: %s", e)
            result["bant"] = {"var": False, "neden": f"Tahsilat/ödeme geçmişi okunamadı: {str(e)[:200]}"}
        rid = save_cash_run(engine, tenant, user, result)
        meta_set(engine, "cash_run", {"id": rid, "asof": end.isoformat(), "errors": errors})
        return {"run": rid, "hatalar": errors}

    def _cash_band(self, logo: Any, firms: dict[int, str], end: date, result: dict[str, Any]) -> dict[str, Any]:
        """Geçmiş 104 haftanın günlük tahsilat/ödemesi (her yıl kendi kopyasından, kendi tarihleriyle) → bant."""
        from semantic_bridge import forecast_client as fc

        start = date.fromisoformat(result["baslangic"])
        h0 = start - timedelta(days=7 * BAND_HISTORY_WEEKS)
        daily: dict[date, tuple[float, float]] = {}
        for y in range(h0.year, start.year + 1):
            if y not in firms:
                continue
            a, b = max(h0, date(y, 1, 1)), min(start, date(y + 1, 1, 1))
            if a >= b:
                continue
            for r in logo(src.cash_flows_daily_sql(firms[y], a, b)):
                d = src.day(r.get("gun"))
                if d:
                    t, o = daily.get(d, (0.0, 0.0))
                    daily[d] = (t + src.f(r.get("tahsilat")), o + src.f(r.get("odeme")))
        return build_band(result, daily, lambda series, h: fc.forecast_series(series, h, start=h0.strftime("%Y-%m")))

    def rebuild_cash(self, user: Optional[str] = None) -> bool:
        """Yalnız nakit tablosu (muhasebe ve satış okuması olmadan)."""
        with self._guard:
            if self.running():
                return False
            self.state.update(running=True, step="Nakit tablosu", startedAt=time.time(), error=None)

            def go():
                engine = self._engine()
                try:
                    logo = src.runner(self._logo())
                    firms = src.firms_by_year(logo)
                    end, costed = src.read_data_end(logo, firms)
                    if not end:
                        raise FinanceError("Logo'da satış verisi yok.")
                    self._build_cash(engine, logo, firms, end, user)
                    self.state.update(running=False, step=None, error=None, finishedAt=time.time())
                except Exception as e:  # noqa: BLE001
                    log.warning("finance cash failed: %s", e)
                    msg = str(e) if isinstance(e, (src.SourceError, FinanceError)) else f"Nakit tablosu kurulamadı: {str(e)[:200]}"
                    self.state.update(running=False, step=None, error=msg, finishedAt=time.time())

            self._thread = threading.Thread(target=go, daemon=True, name="finance-cash")
            self._thread.start()
            return True


def cash_due(engine: sa.engine.Engine, tenant: str, now: Optional[datetime] = None) -> bool:
    """Pazartesi 07:00'den sonra bu hafta nakit tablosu kurulmadıysa (ya da hiç yoksa) kurulmalı."""
    now = (now or datetime.now(TZ)).astimezone(TZ)
    with engine.connect() as c:
        last = c.execute(sa.select(sa.func.max(CASH_RUNS.c.run_at)).where(CASH_RUNS.c.tenant_id == tenant)).scalar()
    if last is None:
        return True
    last = (last if last.tzinfo else last.replace(tzinfo=timezone.utc)).astimezone(TZ)
    week_start = datetime.combine(monday(now.date()), datetime.min.time(), tzinfo=TZ).replace(hour=7)
    return now >= week_start and last < week_start


def summary_due(engine: sa.engine.Engine, now: Optional[datetime] = None) -> bool:
    """İş günü 08:30'dan sonra bugün özet e-postası gitmediyse."""
    now = (now or datetime.now(TZ)).astimezone(TZ)
    if now.weekday() >= 5 or (now.hour, now.minute) < (8, 30):
        return False
    return meta_get(engine, "summary_mail").get("day") != now.date().isoformat()


def _same_day_last_year(d: date) -> date:
    try:
        return d.replace(year=d.year - 1)
    except ValueError:  # 29 Şubat
        return d.replace(year=d.year - 1, day=28)


def _merge(rows: list[dict[str, Any]], keys: tuple[str, ...], measures: Iterable[str]) -> list[dict[str, Any]]:
    """Aynı anahtara düşen satırları toplar (kırpılan kanal adı iki satırı birleştirebilir)."""
    out: dict[tuple, dict[str, Any]] = {}
    for r in rows:
        k = tuple(r[x] for x in keys)
        if k in out:
            for m in measures:
                out[k][m] += r[m]
        else:
            out[k] = dict(r)
    return list(out.values())
