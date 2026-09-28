"""M52 Tedarik ve baskı uçları: /api/v1/supply/*.

Sayfa kapısı `access.RULES` (`sayfa:tedarik*`); işlem yetkileri `FEATURE_RULES` (kapasite, öneri kararı, taslak,
dışa aktarma). Açıkça verilen yetkiler burada denetlenir: tedarikçi borç/ödeme/fatura tutarı `ozellik:tedarik.borc`,
birim maliyet ve fiyat `ozellik:tedarik.maliyet`, fatura–kart eşleşmesi ve matbaa–cari eşlemesi `ozellik:tedarik.eslesme`.
Yetkisi olmayana alan boş gider (ekran «yetkiniz yok» yazar), tutar sızmaz.

Zamanlayıcı (`timas-supply.timer`, her gece 03:30; `timas-supply-weekly.timer` pazartesi 08:00 ödeme listesi) yalnız
`POST /api/v1/supply/run-due`'yu çağırır. CRM'e, Logo'ya ve matbaaya yazma/gönderim yoktur.
"""
from __future__ import annotations

import io
import logging
from datetime import date, datetime
from typing import Any, Callable, Optional

from fastapi import HTTPException, Request
from fastapi.responses import Response

from semantic_bridge import supply as S
from semantic_bridge import supply_sources as src
from semantic_bridge import supply_store as store
from semantic_bridge import supply_suggest as G

log = logging.getLogger("semantic.supply.api")

P = "/api/v1/supply"
DEBT_KEYS = ("bakiye", "vadesiGecmis", "plansiz", "gelmemis", "kovalar", "gelecek", "katalogVadesiGecmis", "alisBuYil",
             "alis12", "fatura12")
COST_KEYS = ("unitPrice", "unitRecent", "unitPrevious", "unitTrend", "unitSamples")
EXPORTS = {"yuk": "Baskı yükü", "kagit": "Kağıt ihtiyacı", "borc": "Tedarikçi borcu", "odeme": "Ödeme planı",
           "faturasiz": "Faturası görünmeyen baskılar", "maliyet": "Birim baskı maliyeti"}


def _recipients(raw: str) -> list[str]:
    return [x.strip() for x in (raw or "").replace(";", ",").split(",") if "@" in x]


def xlsx(title: str, notes: list[str], columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> bytes:
    """Tek sayfalık Excel: başlık, notlar (FIFO notu gibi), kolonlar; sayı hücreleri sayı olarak."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    r = 2
    for n in notes:
        ws.cell(row=r, column=1, value=n)
        r += 1
    r += 1
    fill = PatternFill("solid", fgColor="EDE9FE")
    for i, (_, label) in enumerate(columns, start=1):
        c = ws.cell(row=r, column=i, value=label)
        c.font, c.fill = Font(bold=True), fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        r += 1
        for i, (key, _) in enumerate(columns, start=1):
            v = row.get(key)
            if isinstance(v, (list, dict)):
                v = ", ".join(f"{k}: {x}" for k, x in v.items()) if isinstance(v, dict) else ", ".join(map(str, v))
            ws.cell(row=r, column=i, value=v)
    for i in range(1, len(columns) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def register(app: Any, deps: dict[str, Any]) -> S.Service:
    """app.py'de bağlanır. `deps`:
    auth(request) → (engine, tenant, user, display) · can(user, key) · is_admin(user) · audit(engine, user, action, kind,
    id, title, detail) · conf(key, default) · fresh() · require_caller(request) · runtime() → (engine, tenant) ·
    production → M12 `production.Service` · logo_file() / crm_file() → bağlantı dosyası · llm(priority) → LLM kapısı ·
    send_mail(subject, text, to) · report_data() → Baskı Öneri önbelleği (isteğe bağlı) · decisions(engine, tenant) → M10
    onaylı ilk baskı kararları (isteğe bağlı) · unit_costs(codes) → M9 birim maliyet (isteğe bağlı)."""
    from semantic_bridge import production as production_mod
    from semantic_bridge import production_plan as plan_mod

    auth, can, is_admin, audit, conf, fresh = (deps[k] for k in ("auth", "can", "is_admin", "audit", "conf", "fresh"))
    production = deps["production"]
    settings = lambda: S.settings_from(conf)  # noqa: E731
    source = S.Source(production, deps["logo_file"], deps["crm_file"], lambda: conf("CRM_SCHEMA"), settings)
    svc = S.Service(source, settings, match_costs=production_mod.match_costs, printer_stats=plan_mod.printer_stats,
                    report_data=deps.get("report_data"), decisions=deps.get("decisions"), unit_costs=deps.get("unit_costs"))

    def ctx(request: Request) -> tuple[Any, str, str, str]:
        engine, tenant, user, display = auth(request)
        store.ensure(engine)
        return engine, tenant, user, display

    def allowed(user: str, key: str) -> bool:
        return bool(is_admin(user) or can(user, key))

    def need(user: str, key: str, what: str) -> None:
        if not allowed(user, key):
            raise HTTPException(status_code=403, detail={"code": "FORBIDDEN", "message": f"{what} rolünüzde yok."})

    def call(fn, *a, **kw):
        try:
            return fn(*a, **kw)
        except store.SupplyError as e:
            raise HTTPException(status_code=e.status, detail={"code": "SUPPLY", "message": str(e)}) from e
        except production_mod.ProductionError as e:
            raise HTTPException(status_code=e.status, detail={"code": "SUPPLY", "message": str(e)}) from e
        except HTTPException:
            raise
        except src.SourceError as e:
            raise HTTPException(status_code=503, detail={"code": "DATA_SOURCE_UNAVAILABLE", "retryable": True,
                                                         "message": str(e)}) from e
        except Exception as e:  # noqa: BLE001 — kaynak düştüyse kişiye düz cümle, ayrıntı günlükte
            log.exception("supply: okuma başarısız")
            state = str(getattr(e, "args", [""])[0])
            if state in ("08S01", "08001", "HYT00", "HYT01") or "timeout" in str(e).lower():
                raise HTTPException(status_code=503, detail={"code": "DATA_SOURCE_UNAVAILABLE", "retryable": True,
                                    "message": "CRM ya da Logo şu anda yanıt vermiyor; birazdan tekrar deneyin."}) from e
            raise HTTPException(status_code=502, detail={"code": "SUPPLY", "message": "Tedarik verisi okunamadı."}) from e

    def printers() -> list[str]:
        try:
            snap = production.source.snapshot(False)
            return sorted(set(snap["options"].get("new_matbaa", {}).values()), key=production_mod._fold)
        except Exception as e:  # noqa: BLE001
            log.info("supply: matbaa listesi okunamadı: %s", e)
            return []

    def strip_cost(card: dict[str, Any]) -> dict[str, Any]:
        return {k: v for k, v in card.items() if k not in ("birimFiyat", "faturaTutari", "m9Maliyet")}

    def strip_karne(k: dict[str, Any]) -> dict[str, Any]:
        return {x: v for x, v in k.items() if x not in COST_KEYS}

    # ------------------------------------------------------------------ genel

    @app.get(f"{P}/meta")
    def supply_meta(request: Request) -> dict[str, Any]:
        engine, tenant, user, display = ctx(request)
        s = settings()
        return {"settings": {k: s[k] for k in ("loadMonths", "paperMonths", "paperMeasure", "paperLeadDays", "paperBuyer",
                                               "overloadRatio", "unbilledGraceDays", "supplierPrefix", "printerSpecodes",
                                               "paperSpecodes", "trendMonths", "pageBands", "agingAsOf")},
                "printers": printers(), "kinds": store.KINDS, "states": store.STATES, "linkMethods": store.LINK_METHODS,
                "exports": EXPORTS, "fifoNote": S.FIFO_NOTE, "referenceNote": S.REFERENCE_NOTE,
                "me": {"username": user, "display": display, "admin": is_admin(user),
                       "canDebt": allowed(user, "ozellik:tedarik.borc"), "canCost": allowed(user, "ozellik:tedarik.maliyet"),
                       "canCapacity": allowed(user, "ozellik:tedarik.kapasite"),
                       "canDecide": allowed(user, "ozellik:tedarik.oneri-karar"),
                       "canMatch": allowed(user, "ozellik:tedarik.eslesme"),
                       "canExport": allowed(user, "ozellik:veri.disa-aktar")}}

    @app.get(f"{P}/overview")
    def supply_overview(request: Request) -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        return call(svc.overview, engine, tenant, debt=allowed(user, "ozellik:tedarik.borc"),
                    cost=allowed(user, "ozellik:tedarik.maliyet"), fresh=fresh())

    @app.get(f"{P}/sources")
    def supply_sources(request: Request) -> dict[str, Any]:
        ctx(request)
        return {"sources": [{"id": sid, "connection": conn, "title": t, "description": d, "sql": src.sql_text(sid)}
                            for sid, conn, t, d in src.SOURCES],
                "notes": [S.FIFO_NOTE, S.REFERENCE_NOTE]}

    # ------------------------------------------------------------------ yük ve çakışma

    @app.get(f"{P}/load")
    def supply_load(request: Request, aylar: int = 0) -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        return call(svc.load, engine, tenant, aylar or None, fresh(), allowed(user, "ozellik:tedarik.maliyet"))

    @app.get(f"{P}/conflicts")
    def supply_conflicts(request: Request, aylar: int = 0) -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        t = call(svc.load, engine, tenant, aylar or None, fresh(), allowed(user, "ozellik:tedarik.maliyet"))
        return {"items": t["cakismalar"], "referansNotu": t["referansNotu"], "asOf": t["asOf"]}

    @app.get(f"{P}/incoming")
    def supply_incoming(request: Request, aylar: int = 6) -> dict[str, Any]:
        """Gelecek depo girişleri (M43 depo ve stok da okur; sayfa anahtarını RULES satırına ekler)."""
        engine, tenant, _, _ = ctx(request)
        return call(svc.incoming, engine, tenant, max(1, min(24, int(aylar or 6))), fresh())

    # ------------------------------------------------------------------ kağıt

    @app.get(f"{P}/paper")
    def supply_paper(request: Request, aylar: int = 0) -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        out = call(svc.paper, engine, tenant, aylar or None, fresh())
        if not allowed(user, "ozellik:tedarik.maliyet"):
            out["fiyat"] = None
        return out

    # ------------------------------------------------------------------ tedarikçiler, ödeme, fatura

    @app.get(f"{P}/suppliers")
    def supply_suppliers(request: Request) -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        out = call(svc.suppliers, engine, tenant, fresh())
        debt, cost = allowed(user, "ozellik:tedarik.borc"), allowed(user, "ozellik:tedarik.maliyet")
        for it in out["items"]:
            if not debt:
                for k in DEBT_KEYS:
                    it.pop(k, None)
            if not cost:
                it["karne"] = [strip_karne(k) for k in it["karne"]]
        out["borcGorunur"], out["maliyetGorunur"] = debt, cost
        return out

    @app.get(f"{P}/suppliers/{{cari}}")
    def supply_supplier(cari: str, request: Request) -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        # Kod Logo'daki tedarikçi listesinde birebir aranır; listede yoksa 404 (SQL'e yalnız bulunan kod girer).
        out = call(svc.supplier, engine, tenant, str(cari or "").strip(), fresh())
        debt, cost = allowed(user, "ozellik:tedarik.borc"), allowed(user, "ozellik:tedarik.maliyet")
        if not debt:
            out["yaslandirma"], out["alis"], out["faturalar"] = None, None, None
        if not cost:
            out["acikIsler"] = [strip_cost(c) for c in out["acikIsler"]]
            out["bitenIsler"] = [strip_cost(c) for c in out["bitenIsler"]]
            out["karne"] = [strip_karne(k) for k in out["karne"]]
        out["borcGorunur"], out["maliyetGorunur"] = debt, cost
        return out

    @app.get(f"{P}/payments")
    def supply_payments(request: Request, gun: int = 30, tur: str = "") -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        need(user, "ozellik:tedarik.borc", "Tedarikçi borç ve ödeme bilgisi")
        if tur not in ("", "matbaa", "kagit", "diger"):
            raise HTTPException(status_code=400, detail={"code": "SUPPLY", "message": "Tür matbaa, kagit ya da diger olmalı."})
        return call(svc.payments, engine, tenant, max(1, min(366, int(gun or 30))), tur, fresh())

    @app.get(f"{P}/unbilled")
    def supply_unbilled(request: Request) -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        out = call(svc.unbilled, engine, tenant, fresh())
        if not allowed(user, "ozellik:tedarik.borc"):
            for f in out["faturalar"]:
                f.pop("tutar", None)
            out["tutarGorunur"] = False
        else:
            out["tutarGorunur"] = True
        return out

    @app.get(f"{P}/cost-trend")
    def supply_cost_trend(request: Request, kirilim: str = "") -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        need(user, "ozellik:tedarik.maliyet", "Birim baskı maliyeti")
        return call(svc.cost, engine, tenant, kirilim, fresh())

    # ------------------------------------------------------------------ öneriler ve taslaklar

    @app.get(f"{P}/suggestions")
    def supply_suggestions(request: Request, tur: str = "", durum: str = "") -> dict[str, Any]:
        engine, tenant, _, _ = ctx(request)
        if tur and tur not in store.KINDS:
            raise HTTPException(status_code=400, detail={"code": "SUPPLY", "message": "Öneri türü geçersiz."})
        if durum and durum not in store.STATES:
            raise HTTPException(status_code=400, detail={"code": "SUPPLY", "message": "Durum geçersiz."})
        return {"items": store.list_suggestions(engine, tenant, tur, durum)}

    @app.post(f"{P}/suggestions/{{sid}}/decision")
    def supply_suggestion_decision(sid: str, body: dict[str, Any], request: Request) -> dict[str, Any]:
        """Karar portal kaydıdır: CRM kartı ve M12 kaydı değişmez; kabul edilen öneriyi insan uygular."""
        engine, tenant, user, display = ctx(request)
        out = call(store.decide_suggestion, engine, tenant, user, display, sid, body)
        audit(engine, user, "decide", "supply_suggestion", out["id"], (out["baslik"] or "")[:200],
              {"tur": out["tur"], "karar": out["durum"], "not": out["kararNotu"]})
        return out

    @app.post(f"{P}/drafts", status_code=201)
    def supply_draft(body: dict[str, Any], request: Request) -> dict[str, Any]:
        """Sipariş formu / teknik şartname (kalıp) ya da matbaaya gecikme yazısı (Zeki AI, sayı denetimli) taslağı.
        Gönderim yok: kişi kopyalayıp kendisi gönderir."""
        engine, tenant, user, _ = ctx(request)
        tur = str(body.get("tur") or "")
        if tur not in ("sartname", "eskalasyon"):
            raise HTTPException(status_code=400, detail={"code": "SUPPLY", "message": "Taslak türü «sartname» ya da «eskalasyon» olmalı."})
        cid = call(store.card_id, body.get("kartId"))
        snap = call(svc.snap, engine, tenant, False)
        card = next((c for c in snap["cards"] if c.get("id") == cid), None)
        if card is None:
            raise HTTPException(status_code=404, detail={"code": "SUPPLY", "message": "Üretim kartı bulunamadı."})
        if tur == "sartname":
            text = G.spec_text(card, snap["tech"].get(cid) or {}, snap["paperNames"], snap["options"])
            how = "kalip"
        else:
            facts = G.escalation_facts(card)
            llm = None
            try:
                llm = deps["llm"](None)
            except Exception as e:  # noqa: BLE001
                log.info("supply: model yok: %s", e)
            model = G.model_text(llm, G.ESCALATION_SYSTEM, facts)
            text, how = (model, "zeki") if model else (G.escalation_template(card), "kalip")
        title = f"{store.KINDS[tur]} — {card.get('bookTitle') or card.get('name') or cid}"
        out = call(store.add_draft, engine, tenant, user, tur=tur, kart=cid, baslik=title,
                   payload={"yontem": how, "matbaa": card.get("printer"), "kitap": card.get("bookTitle")}, metin=text)
        audit(engine, user, "create", "supply_draft", out["id"], title[:200], {"tur": tur, "kart": cid, "yontem": how})
        return out

    # ------------------------------------------------------------------ kapasite

    @app.get(f"{P}/capacity")
    def supply_capacity(request: Request) -> dict[str, Any]:
        engine, tenant, _, _ = ctx(request)
        return {"items": store.list_capacity(engine, tenant), "printers": printers(), "referansNotu": S.REFERENCE_NOTE}

    @app.put(f"{P}/capacity")
    def supply_capacity_put(body: dict[str, Any], request: Request) -> dict[str, Any]:
        engine, tenant, user, display = ctx(request)
        plist = printers() or None
        out = call(store.add_capacity, engine, tenant, user, display, body, plist)
        audit(engine, user, "create", "supply_capacity", out["id"], out["matbaa"],
              {"ay": out["ay"], "adet": out["kapasiteAdet"], "forma": out["kapasiteForma"]})
        return out

    @app.delete(f"{P}/capacity/{{rid}}")
    def supply_capacity_delete(rid: str, request: Request) -> dict[str, Any]:
        engine, tenant, user, _ = ctx(request)
        out = call(store.delete_capacity, engine, tenant, user, rid)
        audit(engine, user, "delete", "supply_capacity", out["id"], out["matbaa"], {"ay": out["ay"]})
        return {"ok": True}

    # ------------------------------------------------------------------ matbaa ↔ cari, fatura ↔ kart

    @app.get(f"{P}/supplier-map")
    def supply_supplier_map(request: Request) -> dict[str, Any]:
        engine, tenant, _, _ = ctx(request)
        snap = call(svc.snap, engine, tenant, fresh())
        return {"items": call(svc.mapping, engine, tenant, snap), "printers": printers()}

    @app.put(f"{P}/supplier-map")
    def supply_supplier_map_put(body: dict[str, Any], request: Request) -> dict[str, Any]:
        engine, tenant, user, display = ctx(request)
        need(user, "ozellik:tedarik.eslesme", "Matbaa–cari eşlemesi")
        out = call(store.set_supplier_map, engine, tenant, user, display, body, printers() or None)
        audit(engine, user, "update", "supply_map", None, out["matbaa"], out)
        return out

    @app.get(f"{P}/invoice-links")
    def supply_invoice_links(request: Request) -> dict[str, Any]:
        engine, tenant, _, _ = ctx(request)
        return {"items": store.list_links(engine, tenant)}

    @app.post(f"{P}/invoice-links")
    def supply_invoice_link(body: dict[str, Any], request: Request) -> dict[str, Any]:
        """Fatura satırı ↔ kart bağını onaylama / reddetme ya da elle bağlama («hiçbiri» = kartId boş)."""
        engine, tenant, user, _ = ctx(request)
        need(user, "ozellik:tedarik.eslesme", "Fatura–kart eşleşmesini onaylama")
        snap = call(svc.snap, engine, tenant, False)
        out = call(store.decide_link, engine, tenant, user, body, {c["id"] for c in snap["cards"]})
        audit(engine, user, "decide", "supply_link", out["id"], f"{out['firma']}:{out['satirRef']}",
              {"kart": out["kartId"], "durum": out["durum"], "yontem": out["yontem"]})
        return out

    # ------------------------------------------------------------------ dışa aktarma

    @app.get(f"{P}/export/{{liste}}.xlsx")
    def supply_export(liste: str, request: Request) -> Response:
        engine, tenant, user, _ = ctx(request)
        if liste not in EXPORTS:
            raise HTTPException(status_code=404, detail={"code": "SUPPLY", "message": "Böyle bir liste yok."})
        if liste in ("borc", "odeme", "faturasiz"):
            need(user, "ozellik:tedarik.borc", "Tedarikçi borç ve ödeme bilgisi")
        if liste == "maliyet":
            need(user, "ozellik:tedarik.maliyet", "Birim baskı maliyeti")
        stamp = f"Hazırlanma: {datetime.now(S.TZ):%d.%m.%Y %H:%M} · kaynak Logo + CRM (yalnız okuma)"
        if liste == "yuk":
            t = call(svc.load, engine, tenant, None, fresh(), False)
            rows = [dict(c, ayAdi=S.month_label(k), hucreMatbaa=r["matbaa"]) for r in t["satirlar"]
                    for k, cell in r["hucreler"].items() for c in cell["cards"]]
            data = xlsx(EXPORTS[liste], [stamp, S.REFERENCE_NOTE],
                        [("ayAdi", "Ay"), ("hucreMatbaa", "Matbaa"), ("kitap", "Kitap"), ("stokKodu", "Stok kodu"),
                         ("baskiNo", "Baskı no"), ("adet", "Adet"), ("forma", "Forma"), ("formaYuku", "Forma × adet"),
                         ("asamaAdi", "Aşama"), ("oncelik", "Öncelik"), ("planBaski", "Baskı (plan)"),
                         ("planDosya", "Dosya teslimi (plan)"), ("yayin", "Hedef yayın")], rows)
        elif liste == "kagit":
            p = call(svc.paper, engine, tenant, None, fresh())
            data = xlsx(EXPORTS[liste], [stamp, f"Ölçü: {'brüt kg (fire dahil)' if p['olcu'] == 'brut' else 'net kg'}; "
                                                f"CRM «toplam kağıt ihtiyacı» kolonunun birimi ölçülecek."],
                        [("ayAdi", "Baskı ayı"), ("cinsAdi", "Kağıt cinsi"), ("gramaj", "Gramaj"), ("kg", "Kg"),
                         ("digerKg", "Diğer ölçü (kg)"), ("toplam", "CRM toplam kolonu"), ("kart", "Kart"),
                         ("parcalar", "Parçalar (kg)"), ("ebatlar", "Ebatlar (kg)")], p["satirlar"])
        elif liste == "borc":
            out = call(svc.suppliers, engine, tenant, fresh())
            rows = [dict(it, **{f"k{i}": it["kovalar"][k] for i, (k, _, _) in enumerate(S.PAST)},
                         **{f"g{i}": it["gelecek"][k] for i, (k, _, _) in enumerate(S.AHEAD)}) for it in out["items"]]
            data = xlsx(EXPORTS[liste], [stamp, S.FIFO_NOTE],
                        [("kod", "Cari kodu"), ("unvan", "Ünvan"), ("tur", "Tür"), ("bakiye", "Bakiye"),
                         ("vadesiGecmis", "Vadesi geçmiş (FIFO)"), ("k0", "1–30 gün"), ("k1", "31–60 gün"), ("k2", "61–90 gün"),
                         ("k3", "90+ gün"), ("plansiz", "Vade planı olmayan"), ("gelmemis", "Vadesi gelmemiş"),
                         ("g0", "0–30 gün içinde"), ("g1", "31–60 gün içinde"), ("g2", "61–90 gün içinde"), ("g3", "90+ gün sonra"),
                         ("alis12", "Son 12 ay alış (KDV dahil)"), ("acikIs", "Açık iş")], rows)
        elif liste == "odeme":
            p = call(svc.payments, engine, tenant, 90, "", fresh())
            data = xlsx(EXPORTS[liste], [stamp, S.FIFO_NOTE],
                        [("vade", "Vade"), ("kod", "Cari kodu"), ("unvan", "Ünvan"), ("tur", "Tür"), ("acik", "Açık tutar"),
                         ("tutar", "Plan tutarı"), ("faturaNo", "Fatura no"), ("faturaTarihi", "Fatura tarihi")], p["satirlar"])
        elif liste == "faturasiz":
            u = call(svc.unbilled, engine, tenant, fresh())
            rows = [dict(c, satir="Kart: fatura görünmüyor") for c in u["kartlar"]] + \
                   [dict(f, satir="Fatura: kart görünmüyor", kitap=f.get("stok")) for f in u["faturalar"]]
            data = xlsx(EXPORTS[liste], [stamp, f"Bekleme süresi {u['bekleme']['gun']} gün ({u['bekleme']['kaynak']})."],
                        [("satir", "Satır"), ("kitap", "Kitap / stok kodu"), ("matbaa", "Matbaa"), ("cari", "Faturayı kesen"),
                         ("adet", "Adet"), ("depo", "Depo girişi"), ("bekleyenGun", "Bekleyen gün"), ("faturaNo", "Fatura no"),
                         ("tarih", "Fatura tarihi"), ("tutar", "Tutar (KDV hariç)")], rows)
        else:
            k = request.query_params.get("kirilim", "")
            c = call(svc.cost, engine, tenant, k, fresh())
            rows = [{"grup": g["grup"], "ay": S.month_label(m), **v} for g in c["gruplar"] for m, v in sorted(g["aylar"].items())]
            data = xlsx(EXPORTS[liste], [stamp, "Birim = Logo matbaa baskı faturasında adet başı bedel (KDV hariç)."],
                        [("grup", "Grup"), ("ay", "Ay"), ("agirlikliBirim", "Ağırlıklı birim"), ("ortancaBirim", "Ortanca birim"),
                         ("sayfa100", "100 sayfa başı (ortanca)"), ("is", "İş"), ("adet", "Adet"), ("tutar", "Tutar")], rows)
        audit(engine, user, "export", "supply_export", liste, EXPORTS[liste], None)
        name = f"tedarik-{liste}-{date.today().isoformat()}.xlsx"
        return Response(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{name}"'})

    # ------------------------------------------------------------------ zamanlayıcı

    @app.post(f"{P}/run-due")
    def supply_run_due(request: Request, haftalik: int = 0) -> dict[str, Any]:
        """Gece: öneriler + fatura adayları + eşik aşımı e-postası; `haftalik=1` (pazartesi 08:00) 30 günlük ödeme listesi."""
        deps["require_caller"](request)
        engine, tenant = deps["runtime"]()
        store.ensure(engine)
        from semantic_layer.runtime.llm_queue import BATCH

        llm = None
        try:
            llm = deps["llm"](BATCH)
        except Exception as e:  # noqa: BLE001 — model yoksa kural metni ve kural adayı
            log.info("supply: model yok: %s", e)
        link = (conf("ALERT_LINK") or "").split("/uyarilar")[0]
        # Değişiklik kaydına yazılmaz: kayıttaki ad Yönetim → Kişiler'de kişi gibi görünür (test/sistem adı bırakılmaz).
        return call(G.run_due, svc, engine, tenant, llm=llm, send_mail=deps.get("send_mail"),
                    recipients=_recipients(conf("SUPPLY_ALERT_RECIPIENTS")), link=link, weekly=bool(haftalik),
                    payment_recipients=_recipients(conf("SUPPLY_PAYMENT_RECIPIENTS")))

    return svc
