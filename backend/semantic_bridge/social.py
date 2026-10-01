"""M22 Sosyal medya yönetimi: hesaplar, tek içerik takvimi, gönderi durum makinesi, fırsat kuralları, performans
içe aktarma ve aylık rapor.

**Otomatik yayın yok** (kullanıcı kararı 2026-09-28): portal yalnız yayına hazır paket üretir (görsel + metin +
etiket); paylaşımı uzman kendi hesabından yapar ve bağlantıyı yapıştırıp «yayınlandı» işaretler. Sosyal medya
platformlarına hiçbir istek gitmez; içgörü verisi platformun dışa aktarım dosyasından ya da elle girilir.

**Gönderi durumu:** `fikir` → `taslak` → (onaya gönder) `onayda` → (onay) `onayli` → (bağlantıyla) `yayinlandi`.
`onayda` iken geri gönderme (gerekçe zorunlu) ve geri çekme gönderiyi `taslak`a döndürür. Onaylı gönderinin metni,
etiketi, hesabı ya da görseli değişirse onay düşer (`taslak`); yalnız saati değişirse onay kalır ve geçmişe yazılır.
Gönderen onaylayamaz (iki göz). `iptal` yayınlanmamış her gönderiye gerekçeyle verilir.

Tablolar `semantic_social_*` (ilk kullanımda kurulur, `budget.py` gibi SQLAlchemy Core). Saat İstanbul yerel saati
olarak `YYYY-AA-GG SS:DD` metni saklanır (takvim gün sınırı İstanbul günü; sunucu saat dilimi fark etmez).
CRM'e, Logo'ya, T-soft'a ve sosyal medya platformlarına yazılmaz.
"""
from __future__ import annotations

import json
import math
import re
import statistics
import threading
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, Iterable, Optional
from zoneinfo import ZoneInfo

import sqlalchemy as sa

TZ = ZoneInfo("Europe/Istanbul")
_md = sa.MetaData()

ACCOUNTS = sa.Table(
    "semantic_social_accounts", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("platform", sa.String(16), nullable=False),
    sa.Column("handle", sa.String(200), nullable=False),              # @kullanici ya da kanal adı
    sa.Column("ad", sa.String(200)),                                   # ekranda görünen ad (ör. «Timaş Çocuk · Instagram»)
    sa.Column("imprint_crm_id", sa.String(40)),                        # CRM new_markaBase.new_markaId
    sa.Column("imprint_ad", sa.String(200)),
    sa.Column("owner_user", sa.String(120)),                           # hesabı yöneten AD hesabı
    sa.Column("ton", sa.Text),                                         # hesabın dili (Zeki AI taslağına girer)
    sa.Column("renk", sa.String(9)),
    sa.Column("active", sa.Boolean, nullable=False, default=True),
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("updated_by", sa.String(120)),
    sa.Column("updated_at", sa.DateTime(timezone=True)),
)
POSTS = sa.Table(
    "semantic_social_posts", _md,
    sa.Column("id", sa.String(24), primary_key=True),                  # SM-<yıl>-<sıra>
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("account_id", sa.String(32), index=True),
    sa.Column("crm_book_id", sa.String(40), index=True),
    sa.Column("stok_kodu", sa.String(60), index=True),
    sa.Column("kitap_ad", sa.String(400)),
    sa.Column("occasion_key", sa.String(120)),                         # özel gün anahtarı (seasons.slug)
    sa.Column("occasion_ad", sa.String(300)),
    sa.Column("plan_ref", sa.String(40)),                              # M18 aylık plan satırı (plan → gönderi aktarımı)
    sa.Column("kind", sa.String(16)),                                  # içerik türü (KINDS)
    sa.Column("kind_source", sa.String(10)),                           # kullanici | zeki
    sa.Column("kind_prob", sa.Float),
    sa.Column("planned_at", sa.String(16), index=True),                # İstanbul yerel: YYYY-AA-GG SS:DD
    sa.Column("status", sa.String(12), nullable=False),
    sa.Column("text", sa.Text),
    sa.Column("hashtags", sa.Text),
    sa.Column("asset_refs_json", sa.Text),
    sa.Column("draft_json", sa.Text),                                  # Zeki AI seçenekleri ve düşen cümleler
    sa.Column("published_url", sa.String(1000)),
    sa.Column("published_at", sa.DateTime(timezone=True)),
    sa.Column("published_by", sa.String(120)),
    sa.Column("submitted_by", sa.String(120)),
    sa.Column("submitted_at", sa.DateTime(timezone=True)),
    sa.Column("approved_by", sa.String(120)),
    sa.Column("approved_at", sa.DateTime(timezone=True)),
    sa.Column("note", sa.Text),                                        # geri gönderme / iptal gerekçesi
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("updated_by", sa.String(120)),
    sa.Column("updated_at", sa.DateTime(timezone=True)),
    sa.Index("ix_semantic_social_posts_cal", "tenant_id", "planned_at", "status"),
)
EVENTS = sa.Table(
    "semantic_social_post_events", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("post_id", sa.String(24), nullable=False, index=True),
    sa.Column("at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("user", sa.String(120), nullable=False),
    sa.Column("action", sa.String(40), nullable=False),
    sa.Column("note", sa.Text),
    sa.Column("detail_json", sa.Text),
)
METRICS = sa.Table(
    "semantic_social_metrics", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("import_id", sa.String(32), index=True),                 # boşsa elle girildi
    sa.Column("account_id", sa.String(32), nullable=False, index=True),
    sa.Column("post_id", sa.String(24), index=True),
    sa.Column("day", sa.String(10), nullable=False, index=True),
    sa.Column("post_url", sa.String(1000)),
    sa.Column("impressions", sa.Float),
    sa.Column("reach", sa.Float),
    sa.Column("likes", sa.Float),
    sa.Column("comments", sa.Float),
    sa.Column("shares", sa.Float),
    sa.Column("saves", sa.Float),
    sa.Column("followers", sa.Float),
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
)
IMPORTS = sa.Table(
    "semantic_social_imports", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("account_id", sa.String(32), nullable=False),
    sa.Column("file_name", sa.String(300), nullable=False),
    sa.Column("rows", sa.Integer, nullable=False),
    sa.Column("matched", sa.Integer, nullable=False, default=0),
    sa.Column("summary_json", sa.Text),                                # tanınan kolonlar, atlanan satırlar
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
)
JOBS = sa.Table(
    "semantic_social_jobs", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("post_id", sa.String(24), index=True),                   # boşsa rapor yorumu
    sa.Column("tur", sa.String(16), nullable=False),                   # draft | report
    sa.Column("hedef", sa.String(40)),                                 # rapor ayı
    sa.Column("status", sa.String(12), nullable=False),                # bekliyor | calisiyor | bitti | hata
    sa.Column("step", sa.String(200)),
    sa.Column("error", sa.Text),
    sa.Column("result_json", sa.Text),
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("finished_at", sa.DateTime(timezone=True)),
)
META = sa.Table(
    "semantic_social_meta", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("key", sa.String(80), primary_key=True),
    sa.Column("value_json", sa.Text, nullable=False),
    sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
)

PLATFORMS = {"instagram": "Instagram", "facebook": "Facebook", "x": "X", "linkedin": "LinkedIn", "tiktok": "TikTok",
             "youtube": "YouTube"}
#: Gönderi metni karakter sınırı (platformların yayımladığı sınırlar; Yönetim → Sosyal medya'dan değişir).
DEFAULT_LIMITS = {"instagram": 2200, "facebook": 63206, "x": 280, "linkedin": 3000, "tiktok": 2200, "youtube": 5000}
#: Instagram tek gönderide en çok 30 etiket kabul eder; diğerlerinde sınır yok sayılır (ayardan değişir).
DEFAULT_TAG_LIMITS = {"instagram": 30}
STATUSES = {"fikir": "Fikir", "taslak": "Taslak", "onayda": "Onay bekliyor", "onayli": "Onaylı",
            "yayinlandi": "Yayınlandı", "iptal": "İptal"}
OPEN = ("fikir", "taslak")
#: İçerik türü (rapor gruplaması). Zeki AI kapalı seçimle etiketler; kullanıcı her zaman değiştirebilir.
KINDS = {"kapak": "Kapak", "alinti": "Alıntı", "yazar": "Yazar", "etkinlik": "Etkinlik", "ozel-gun": "Özel gün",
         "kampanya": "Kampanya", "diger": "Diğer"}
ASSET_KINDS = ("studio", "creative")
METRIC_COLS = ("impressions", "reach", "likes", "comments", "shares", "saves", "followers")
METRIC_LABELS = {"impressions": "Gösterim", "reach": "Erişim", "likes": "Beğeni", "comments": "Yorum",
                 "shares": "Paylaşım", "saves": "Kaydetme", "followers": "Takipçi"}

_ready: set[int] = set()
_lock = threading.Lock()
_URL = re.compile(r"^https?://[^\s]{3,}$", re.I)
_HANDLE = re.compile(r"^[@A-Za-z0-9._\-ÇĞİÖŞÜçğıöşü ]{1,200}$")
_COLOR = re.compile(r"^#[0-9a-fA-F]{6}$")
_SID = re.compile(r"^s_[0-9a-f]{8}$")
_JOB = re.compile(r"^[A-Za-z0-9_\-]{1,80}$")


class SocialError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def ensure(engine: sa.engine.Engine) -> None:
    with _lock:
        if id(engine) in _ready:
            return
        from semantic_layer.store import schema_stamp
        schema_stamp.create_all(_md, engine)
        _ready.add(id(engine))


# ------------------------------------------------------------------ yardımcılar


def now() -> datetime:
    return datetime.now(timezone.utc)


def today() -> date:
    return datetime.now(TZ).date()


def iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    return (v if v.tzinfo else v.replace(tzinfo=timezone.utc)).isoformat()


def loads(v: Optional[str], default: Any) -> Any:
    try:
        return json.loads(v) if v else default
    except ValueError:
        return default


def dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def text(v: Any, limit: int) -> Optional[str]:
    s = str(v or "").strip()
    return s[:limit] or None


def one_line(v: Any, limit: int) -> Optional[str]:
    s = " ".join(str(v or "").split())
    return s[:limit] or None


def _uid() -> str:
    return uuid.uuid4().hex


def planned(v: Any, *, allow_none: bool = True) -> Optional[str]:
    """«2026-10-05», «2026-10-05T14:30», «2026-10-05 14:30(:00)» → «2026-10-05 14:30». Saat yoksa 10:00 (ayar)."""
    if v is None or str(v).strip() == "":
        if allow_none:
            return None
        raise SocialError("Paylaşım tarihi gerekli.")
    s = str(v).strip().replace("T", " ")
    try:
        if len(s) <= 10:
            d = date.fromisoformat(s[:10])
            return f"{d.isoformat()} 10:00"
        dt = datetime.fromisoformat(s[:16])
    except ValueError:
        raise SocialError("Paylaşım zamanı YYYY-AA-GG SS:DD biçiminde olmalı.") from None
    return dt.strftime("%Y-%m-%d %H:%M")


def day(v: Any, label: str) -> date:
    try:
        return date.fromisoformat(str(v)[:10])
    except (TypeError, ValueError):
        raise SocialError(f"{label} YYYY-AA-GG biçiminde olmalı.") from None


def hashtags(v: Any) -> Optional[str]:
    """Etiketler: boşluk/virgülle ayrılmış; her biri tek # ile, yinelenmeden, sırası korunur."""
    if v is None:
        return None
    raw = v if isinstance(v, list) else re.split(r"[\s,;]+", str(v))
    out: list[str] = []
    seen = set()
    for t in raw:
        t = str(t or "").strip().lstrip("#").strip()
        t = re.sub(r"[^\w]", "", t)
        if not t:
            continue
        k = t.casefold()
        if k in seen:
            continue
        seen.add(k)
        out.append("#" + t[:100])
    return " ".join(out) or None


def tag_count(v: Optional[str]) -> int:
    return len([t for t in (v or "").split() if t.startswith("#")])


def norm_url(u: Any) -> str:
    """Bağlantı eşleştirmesi: şema, www, sondaki / ve sorgu atılır; harf küçük."""
    s = str(u or "").strip().lower()
    s = re.sub(r"^https?://", "", s)
    s = re.sub(r"^www\.|^m\.", "", s)
    s = s.split("?")[0].split("#")[0].rstrip("/")
    return s


# ------------------------------------------------------------------ ayarlar


def settings(conf: Callable[[str], str]) -> dict[str, Any]:
    def num(key: str, default: int) -> int:
        raw = (conf(key) or "").strip()
        try:
            v = int(float(raw)) if raw else default
        except ValueError:
            v = default
        return v if v >= 0 else default

    def lst(key: str) -> list[str]:
        return [x.strip() for x in (conf(key) or "").replace(";", ",").split(",") if x.strip()]

    limits = dict(DEFAULT_LIMITS)
    tag_limits = dict(DEFAULT_TAG_LIMITS)
    raw = loads((conf("SOCIAL_PLATFORM_LIMITS") or "").strip(), {})
    if isinstance(raw, dict):
        for k, v in raw.items():
            key = str(k)
            if key.endswith(":etiket") and key[:-7] in PLATFORMS and isinstance(v, (int, float)) and v > 0:
                tag_limits[key[:-7]] = int(v)
            elif key in PLATFORMS and isinstance(v, (int, float)) and v > 0:
                limits[key] = int(v)
    return {
        "recipients": [x for x in lst("SOCIAL_ALERT_RECIPIENTS") if "@" in x],
        "opportunityDays": num("SOCIAL_OPPORTUNITY_DAYS", 30) or 30,
        "leadDays": num("SOCIAL_OCCASION_LEAD_DAYS", 14) or 14,
        "backlistMinAgeDays": num("SOCIAL_BACKLIST_MIN_AGE_DAYS", 365),
        "backlistQuietDays": num("SOCIAL_BACKLIST_QUIET_DAYS", 90) or 90,
        "limits": limits,
        "tagLimits": tag_limits,
        "claims": lst("SOCIAL_BANNED_CLAIMS"),
        "defaultHour": (conf("SOCIAL_DEFAULT_HOUR") or "10:00").strip()[:5] or "10:00",
        "minProb": 0.70,
        "minMargin": 0.30,
    }


# ------------------------------------------------------------------ meta


def meta_stmt(tenant: str, key: str):
    return sa.select(META).where(META.c.tenant_id == tenant, META.c.key == key)


def meta_get(engine: sa.engine.Engine, tenant: str, key: str) -> dict[str, Any]:
    with engine.connect() as c:
        row = c.execute(meta_stmt(tenant, key)).first()
    return {**loads(row.value_json, {}), "_at": iso(row.updated_at)} if row else {}


def meta_prune(engine: sa.engine.Engine, tenant: str, prefix: str, keep: str, older_than: datetime) -> int:
    """`prefix` ile başlayan, `keep` dışındaki ve `older_than`dan önce yazılmış kayıtlar silinir (tarih aralığına bağlı
    saklanan CRM okumaları her gün yeni anahtar açar; eskileri birikmez)."""
    with engine.begin() as c:
        res = c.execute(META.delete().where(META.c.tenant_id == tenant, META.c.key.like(prefix + "%"), META.c.key != keep,
                                            META.c.updated_at < older_than))
    return int(res.rowcount or 0)


def meta_set(engine: sa.engine.Engine, tenant: str, key: str, value: dict[str, Any]) -> None:
    with engine.begin() as c:
        cond = (META.c.tenant_id == tenant, META.c.key == key)
        if c.execute(sa.select(META.c.key).where(*cond)).first():
            c.execute(META.update().where(*cond).values(value_json=dump(value), updated_at=now()))
        else:
            c.execute(META.insert().values(tenant_id=tenant, key=key, value_json=dump(value), updated_at=now()))


# ------------------------------------------------------------------ hesaplar


def account_dict(r: Any) -> dict[str, Any]:
    return {"id": r.id, "platform": r.platform, "platformAdi": PLATFORMS.get(r.platform, r.platform), "handle": r.handle,
            "ad": r.ad or f"{r.imprint_ad or r.handle} · {PLATFORMS.get(r.platform, r.platform)}",
            "imprintCrmId": r.imprint_crm_id, "imprintAd": r.imprint_ad, "sahip": r.owner_user, "ton": r.ton,
            "renk": r.renk, "aktif": bool(r.active), "olusturan": r.created_by, "olusturma": iso(r.created_at),
            "guncelleyen": r.updated_by, "guncelleme": iso(r.updated_at)}


def accounts_stmt(tenant: str, include_inactive: bool = True):
    cond = [ACCOUNTS.c.tenant_id == tenant]
    if not include_inactive:
        cond.append(ACCOUNTS.c.active.is_(True))
    return sa.select(ACCOUNTS).where(*cond).order_by(ACCOUNTS.c.imprint_ad, ACCOUNTS.c.platform, ACCOUNTS.c.handle)


def list_accounts(engine: sa.engine.Engine, tenant: str, *, include_inactive: bool = True) -> list[dict[str, Any]]:
    with engine.connect() as c:
        rows = c.execute(accounts_stmt(tenant, include_inactive)).all()
    return [account_dict(r) for r in rows]


def _account_row(c: Any, tenant: str, aid: Any) -> Any:
    r = c.execute(sa.select(ACCOUNTS).where(ACCOUNTS.c.tenant_id == tenant, ACCOUNTS.c.id == str(aid or "")[:32])).first()
    if not r:
        raise SocialError("Hesap bulunamadı.", 404)
    return r


def _clean_account(body: dict[str, Any], partial: bool) -> dict[str, Any]:
    vals: dict[str, Any] = {}
    if not partial or "platform" in body:
        p = str(body.get("platform") or "").strip().lower()
        if p not in PLATFORMS:
            raise SocialError("Platform Instagram, Facebook, X, LinkedIn, TikTok ya da YouTube olmalı.")
        vals["platform"] = p
    if not partial or "handle" in body:
        h = one_line(body.get("handle"), 200)
        if not h or not _HANDLE.match(h):
            raise SocialError("Hesap adı gerekli (ör. @timasyayinlari).")
        vals["handle"] = h
    for key, col, lim in (("ad", "ad", 200), ("imprintAd", "imprint_ad", 200), ("sahip", "owner_user", 120)):
        if key in body:
            vals[col] = one_line(body.get(key), lim)
    if "imprintCrmId" in body:
        v = str(body.get("imprintCrmId") or "").strip().strip("{}").lower()
        if v and not re.fullmatch(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", v):
            raise SocialError("Marka kimliği geçersiz.")
        vals["imprint_crm_id"] = v or None
    if "ton" in body:
        vals["ton"] = text(body.get("ton"), 2000)
    if "renk" in body:
        v = str(body.get("renk") or "").strip()
        if v and not _COLOR.match(v):
            raise SocialError("Renk #RRGGBB biçiminde olmalı.")
        vals["renk"] = v or None
    if "aktif" in body:
        vals["active"] = bool(body.get("aktif"))
    return vals


def create_account(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    vals = _clean_account(body, partial=False)
    with engine.begin() as c:
        dup = c.execute(sa.select(ACCOUNTS.c.id).where(
            ACCOUNTS.c.tenant_id == tenant, ACCOUNTS.c.platform == vals["platform"],
            sa.func.lower(ACCOUNTS.c.handle) == vals["handle"].lower())).first()
        if dup:
            raise SocialError("Bu platformda aynı hesap zaten tanımlı.", 409)
        aid = _uid()
        c.execute(ACCOUNTS.insert().values(id=aid, tenant_id=tenant, active=vals.pop("active", True), created_by=user,
                                           created_at=now(), **vals))
        return account_dict(_account_row(c, tenant, aid))


def update_account(engine: sa.engine.Engine, tenant: str, user: str, aid: str, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    vals = _clean_account(body, partial=True)
    if not vals:
        raise SocialError("Değiştirilecek alan yok.")
    with engine.begin() as c:
        r = _account_row(c, tenant, aid)
        old = account_dict(r)
        c.execute(ACCOUNTS.update().where(ACCOUNTS.c.id == r.id).values(updated_by=user, updated_at=now(), **vals))
        return old, account_dict(_account_row(c, tenant, r.id))


# ------------------------------------------------------------------ gönderi okuma


def event(c: Any, post_id: str, user: str, action: str, note: Any = None, detail: Any = None) -> None:
    c.execute(EVENTS.insert().values(id=_uid(), post_id=post_id, at=now(), user=(user or "sistem")[:120], action=action[:40],
                                     note=text(note, 4000), detail_json=None if detail is None else dump(detail)[:20000]))


def events_stmt(post_id: str):
    return sa.select(EVENTS).where(EVENTS.c.post_id == post_id).order_by(EVENTS.c.at.desc())


def events(engine: sa.engine.Engine, post_id: str) -> list[dict[str, Any]]:
    with engine.connect() as c:
        rows = c.execute(events_stmt(post_id)).all()
    return [{"id": r.id, "zaman": iso(r.at), "kim": r.user, "ne": r.action, "not": r.note, "ayrinti": loads(r.detail_json, None)}
            for r in rows]


def post_stmt(tenant: str, pid: Any):
    return sa.select(POSTS).where(POSTS.c.tenant_id == tenant, POSTS.c.id == str(pid or "")[:24])


def account_stmt(aid: Any):
    return sa.select(ACCOUNTS).where(ACCOUNTS.c.id == aid)


def _post_row(c: Any, tenant: str, pid: Any, *, lock: bool = False) -> Any:
    q = post_stmt(tenant, pid)
    if lock and c.engine.dialect.name == "postgresql":
        q = q.with_for_update()
    r = c.execute(q).first()
    if not r:
        raise SocialError("Gönderi bulunamadı.", 404)
    return r


def warnings(p: dict[str, Any], acc: Optional[dict[str, Any]], st: dict[str, Any]) -> list[dict[str, str]]:
    """Gönderi kartındaki uyarılar (engel değil; onaya gönderirken sınır aşımı ayrıca durdurur)."""
    out: list[dict[str, str]] = []
    plat = (acc or {}).get("platform")
    if not p.get("accountId"):
        out.append({"kod": "hesap-yok", "metin": "Hesap seçilmedi."})
    if p["status"] not in ("yayinlandi", "iptal"):
        if not p.get("assets"):
            out.append({"kod": "gorsel-yok", "metin": "Görsel eklenmedi."})
        if not p.get("text"):
            out.append({"kod": "metin-yok", "metin": "Metin yok."})
    lim = st["limits"].get(plat or "")
    n = len(p.get("text") or "") + (len(p.get("hashtags") or "") + 1 if p.get("hashtags") else 0)
    if lim and n > lim:
        out.append({"kod": "uzun", "metin": f"Metin ve etiketler {n} karakter; {PLATFORMS.get(plat, plat)} sınırı {lim}."})
    tl = st["tagLimits"].get(plat or "")
    if tl and tag_count(p.get("hashtags")) > tl:
        out.append({"kod": "etiket", "metin": f"{tag_count(p.get('hashtags'))} etiket; {PLATFORMS.get(plat, plat)} en çok {tl} kabul eder."})
    if p["status"] not in ("yayinlandi", "iptal") and p.get("plannedAt"):
        try:
            when = datetime.strptime(p["plannedAt"], "%Y-%m-%d %H:%M").replace(tzinfo=TZ)
        except ValueError:
            when = None
        if when and when < datetime.now(TZ):
            out.append({"kod": "gecmis", "metin": "Paylaşım zamanı geçti; yayınlandıysa bağlantıyı girin ya da yeni zaman verin."})
    return out


def post_dict(r: Any, acc: Optional[dict[str, Any]] = None, st: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    out = {
        "id": r.id, "accountId": r.account_id, "crmBookId": r.crm_book_id, "stokKodu": r.stok_kodu, "kitapAd": r.kitap_ad,
        "occasionKey": r.occasion_key, "occasionAd": r.occasion_ad, "planRef": r.plan_ref, "kind": r.kind,
        "kindAdi": KINDS.get(r.kind or "", None), "kindSource": r.kind_source, "kindProb": r.kind_prob,
        "plannedAt": r.planned_at, "status": r.status, "statusAdi": STATUSES.get(r.status, r.status), "text": r.text,
        "hashtags": r.hashtags, "assets": loads(r.asset_refs_json, []), "draft": loads(r.draft_json, None),
        "publishedUrl": r.published_url, "publishedAt": iso(r.published_at), "publishedBy": r.published_by,
        "submittedBy": r.submitted_by, "submittedAt": iso(r.submitted_at), "approvedBy": r.approved_by,
        "approvedAt": iso(r.approved_at), "note": r.note, "createdBy": r.created_by, "createdAt": iso(r.created_at),
        "updatedBy": r.updated_by, "updatedAt": iso(r.updated_at),
    }
    if acc is not None:
        out["account"] = acc
    if st is not None:
        out["uyarilar"] = warnings(out, acc, st)
    return out


def get_post(engine: sa.engine.Engine, tenant: str, pid: str, st: dict[str, Any]) -> dict[str, Any]:
    with engine.connect() as c:
        r = _post_row(c, tenant, pid)
        acc = None
        if r.account_id:
            a = c.execute(account_stmt(r.account_id)).first()
            acc = account_dict(a) if a else None
    return post_dict(r, acc, st)


def posts_stmt(tenant: str, *, frm: Optional[date] = None, to: Optional[date] = None, account: str = "", status: str = "",
               stok: str = "", unscheduled: bool = False, occasion: str = ""):
    """Gönderi okuması (takvim, liste, fırsat, kitap içeriği)."""
    cond = [POSTS.c.tenant_id == tenant]
    if frm is not None:
        cond.append(POSTS.c.planned_at >= f"{frm.isoformat()} 00:00")
    if to is not None:
        cond.append(POSTS.c.planned_at <= f"{to.isoformat()} 23:59")
    if unscheduled:
        cond.append(POSTS.c.planned_at.is_(None))
    if account:
        cond.append(POSTS.c.account_id.in_([a for a in account.split(",") if a]))
    if status:
        cond.append(POSTS.c.status.in_([s for s in status.split(",") if s]))
    if stok:
        cond.append(POSTS.c.stok_kodu.in_([s for s in stok.split(",") if s]))
    if occasion:
        cond.append(POSTS.c.occasion_key == occasion)
    return sa.select(POSTS).where(*cond).order_by(POSTS.c.planned_at, POSTS.c.id)


def list_posts(engine: sa.engine.Engine, tenant: str, st: dict[str, Any], *, frm: Optional[date] = None,
               to: Optional[date] = None, account: str = "", status: str = "", stok: str = "", unscheduled: bool = False,
               occasion: str = "") -> list[dict[str, Any]]:
    q = posts_stmt(tenant, frm=frm, to=to, account=account, status=status, stok=stok, unscheduled=unscheduled,
                   occasion=occasion)
    with engine.connect() as c:
        rows = c.execute(q).all()
        accs = {a.id: account_dict(a) for a in c.execute(accounts_stmt(tenant)).all()}
    return [post_dict(r, accs.get(r.account_id or ""), st) for r in rows]


def calendar(engine: sa.engine.Engine, tenant: str, st: dict[str, Any], frm: date, to: date, account: str = "") -> dict[str, Any]:
    """Aralıktaki gönderiler, durum sayıları (`SELECT status, COUNT(*) … GROUP BY status` ile aynı) ve tarihsiz fikirler."""
    if to < frm:
        raise SocialError("Bitiş başlangıçtan önce olamaz.")
    items = list_posts(engine, tenant, st, frm=frm, to=to, account=account)
    counts: dict[str, int] = {}
    for p in items:
        counts[p["status"]] = counts.get(p["status"], 0) + 1
    ideas = list_posts(engine, tenant, st, unscheduled=True, account=account, status="fikir,taslak,onayda,onayli")
    return {"from": frm.isoformat(), "to": to.isoformat(), "items": items, "counts": counts, "unscheduled": ideas,
            "total": len(items)}


def pending(engine: sa.engine.Engine, tenant: str, st: dict[str, Any]) -> list[dict[str, Any]]:
    return list_posts(engine, tenant, st, status="onayda")


# ------------------------------------------------------------------ gönderi yazma


def next_id(c: Any, tenant: str, year: int) -> str:
    prefix = f"SM-{year}-"
    ids = c.execute(sa.select(POSTS.c.id).where(POSTS.c.tenant_id == tenant, POSTS.c.id.like(prefix + "%"))).scalars().all()
    n = max((int(i[len(prefix):]) for i in ids if i[len(prefix):].isdigit()), default=0) + 1
    return f"{prefix}{n:04d}"


def clean_assets(v: Any) -> list[dict[str, Any]]:
    """Görsel bağları: stüdyo pazarlama kiti görseli (`{"tip":"studio","job","sid"}`) ya da M19 içerik arşivi varlığı
    (`{"tip":"creative","id"}`). Dosyanın kendisi kaynağında durur; pakette oradan okunur."""
    if v is None:
        return []
    if not isinstance(v, list):
        raise SocialError("Görsel listesi gerekli.")
    out, seen = [], set()
    for a in v:
        if not isinstance(a, dict):
            raise SocialError("Görsel bağı geçersiz.")
        tip = str(a.get("tip") or "")
        if tip == "studio":
            job, sid = str(a.get("job") or ""), str(a.get("sid") or "")
            if not _JOB.match(job) or not _SID.match(sid):
                raise SocialError("Stüdyo görseli bağı geçersiz.")
            ref = {"tip": "studio", "job": job, "sid": sid}
        elif tip == "creative":
            aid = str(a.get("id") or "")
            if not re.fullmatch(r"[A-Za-z0-9_\-]{1,64}", aid):
                raise SocialError("İçerik arşivi bağı geçersiz.")
            ref = {"tip": "creative", "id": aid}
        else:
            raise SocialError("Görsel kaynağı stüdyo ya da içerik arşivi olmalı.")
        key = dump(ref)
        if key in seen:
            continue
        seen.add(key)
        ref.update({k: one_line(a.get(k), 300) for k in ("ad", "boyut", "onizleme") if a.get(k)})
        ref["onayli"] = bool(a.get("onayli"))
        out.append(ref)
    return out


def _clean_post(body: dict[str, Any], st: dict[str, Any]) -> dict[str, Any]:
    vals: dict[str, Any] = {}
    if "accountId" in body:
        vals["account_id"] = str(body.get("accountId") or "")[:32] or None
    if "plannedAt" in body:
        v = body.get("plannedAt")
        if v and len(str(v).strip()) <= 10:
            v = f"{str(v).strip()[:10]} {st.get('defaultHour') or '10:00'}"
        vals["planned_at"] = planned(v)
    if "kind" in body:
        k = str(body.get("kind") or "").strip()
        if k and k not in KINDS:
            raise SocialError("İçerik türü tanınmıyor.")
        vals["kind"] = k or None
        vals["kind_source"] = "kullanici" if k else None
        vals["kind_prob"] = None
    if "text" in body:
        vals["text"] = text(body.get("text"), 70000)
    if "hashtags" in body:
        vals["hashtags"] = hashtags(body.get("hashtags"))
    if "assets" in body:
        vals["asset_refs_json"] = dump(clean_assets(body.get("assets")))
    for key, col, lim in (("stokKodu", "stok_kodu", 60), ("kitapAd", "kitap_ad", 400), ("occasionKey", "occasion_key", 120),
                          ("occasionAd", "occasion_ad", 300), ("planRef", "plan_ref", 40)):
        if key in body:
            vals[col] = one_line(body.get(key), lim)
    if "crmBookId" in body:
        v = str(body.get("crmBookId") or "").strip().strip("{}").lower()
        if v and not re.fullmatch(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", v):
            raise SocialError("Kitap kimliği geçersiz.")
        vals["crm_book_id"] = v or None
    return vals


def _check_account(c: Any, tenant: str, aid: Optional[str]) -> None:
    if aid:
        r = _account_row(c, tenant, aid)
        if not r.active:
            raise SocialError("Hesap pasif; önce hesabı etkinleştirin.", 409)


def create_post(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any], st: dict[str, Any]) -> dict[str, Any]:
    vals = _clean_post(body, st)
    with engine.begin() as c:
        _check_account(c, tenant, vals.get("account_id"))
        pid = next_id(c, tenant, today().year)
        t = now()
        status = "taslak" if vals.get("text") else "fikir"
        c.execute(POSTS.insert().values(id=pid, tenant_id=tenant, status=status, created_by=user, created_at=t,
                                        updated_by=user, updated_at=t, **vals))
        event(c, pid, user, "olusturuldu", None, {"durum": status, "kitap": vals.get("stok_kodu"), "ozelGun": vals.get("occasion_key"),
                                                  "zaman": vals.get("planned_at")})
    return get_post(engine, tenant, pid, st)


#: Değişince onayı düşüren alanlar (içerik); yalnız saat değişirse onay kalır.
_CONTENT = ("account_id", "text", "hashtags", "asset_refs_json", "crm_book_id", "stok_kodu")


def update_post(engine: sa.engine.Engine, tenant: str, user: str, pid: str, body: dict[str, Any], st: dict[str, Any]) -> dict[str, Any]:
    vals = _clean_post(body, st)
    if "publishedUrl" in body:
        u = one_line(body.get("publishedUrl"), 1000)
        if u and not _URL.match(u):
            raise SocialError("Yayın bağlantısı http(s):// ile başlamalı.")
        vals["published_url"] = u
    if not vals:
        raise SocialError("Değiştirilecek alan yok.")
    with engine.begin() as c:
        r = _post_row(c, tenant, pid, lock=True)
        if r.status == "iptal":
            raise SocialError("İptal edilmiş gönderi değiştirilemez.", 409)
        if r.status == "yayinlandi" and set(vals) - {"published_url", "kind", "kind_source", "kind_prob"}:
            raise SocialError("Yayınlanmış gönderide yalnız bağlantı ve içerik türü değişir.", 409)
        if r.status != "yayinlandi" and "published_url" in vals:
            raise SocialError("Bağlantı «Yayınlandı» işaretlenirken girilir.", 409)
        if r.status == "onayda" and set(vals) - {"kind", "kind_source", "kind_prob"}:
            raise SocialError("Onaydaki gönderi değiştirilemez; önce onaydan geri çekin.", 409)
        if "account_id" in vals:
            _check_account(c, tenant, vals["account_id"])
        changed = {k: v for k, v in vals.items() if getattr(r, k) != v}
        if not changed:
            return get_post(engine, tenant, pid, st)
        extra: dict[str, Any] = {}
        if r.status == "onayli" and set(changed) & set(_CONTENT):
            extra = {"status": "taslak", "approved_by": None, "approved_at": None}
        elif r.status == "fikir" and changed.get("text"):
            extra = {"status": "taslak"}
        c.execute(POSTS.update().where(POSTS.c.id == r.id).values(updated_by=user, updated_at=now(), **changed, **extra))
        old = {k: getattr(r, k) for k in changed}
        event(c, r.id, user, "onay-dustu" if extra.get("status") == "taslak" and r.status == "onayli" else "duzenlendi", None,
              {"alanlar": sorted(changed), "eski": {k: v for k, v in old.items() if k not in ("text", "draft_json")},
               "durum": extra.get("status") or r.status})
    return get_post(engine, tenant, pid, st)


def _limit_errors(r: Any, platform: Optional[str], st: dict[str, Any]) -> Optional[str]:
    lim = st["limits"].get(platform or "")
    n = len(r.text or "") + (len(r.hashtags or "") + 1 if r.hashtags else 0)
    if lim and n > lim:
        return f"Metin ve etiketler {n} karakter; {PLATFORMS.get(platform, platform)} sınırı {lim}. Kısaltın."
    tl = st["tagLimits"].get(platform or "")
    if tl and tag_count(r.hashtags) > tl:
        return f"{PLATFORMS.get(platform, platform)} en çok {tl} etiket kabul eder."
    return None


def transition(engine: sa.engine.Engine, tenant: str, user: str, pid: str, action: str, st: dict[str, Any],
               note: Any = None, url: Any = None) -> dict[str, Any]:
    """submit | withdraw | approve | reject | published | cancel | reopen. Yetki uç katmanında; iki göz burada."""
    note_t = text(note, 4000)
    with engine.begin() as c:
        r = _post_row(c, tenant, pid, lock=True)
        t = now()
        vals: dict[str, Any]
        if action == "submit":
            if r.status not in OPEN:
                raise SocialError("Yalnız fikir ya da taslak gönderi onaya gönderilir.", 409)
            if not r.account_id:
                raise SocialError("Önce hesap seçin.")
            if not r.text:
                raise SocialError("Önce metni yazın.")
            if not r.planned_at:
                raise SocialError("Önce paylaşım zamanını girin.")
            acc = _account_row(c, tenant, r.account_id)
            err = _limit_errors(r, acc.platform, st)
            if err:
                raise SocialError(err)
            vals = {"status": "onayda", "submitted_by": user, "submitted_at": t, "approved_by": None, "approved_at": None,
                    "note": None}
        elif action == "withdraw":
            if r.status != "onayda":
                raise SocialError("Gönderi onayda değil.", 409)
            vals = {"status": "taslak"}
        elif action in ("approve", "reject"):
            if r.status != "onayda":
                raise SocialError("Gönderi onay beklemiyor.", 409)
            if (r.submitted_by or "").lower() == user.lower():
                raise SocialError("Onaya gönderen kişi aynı gönderiyi onaylayamaz; başka bir yetkili onaylamalı.", 409)
            if action == "reject":
                if not note_t:
                    raise SocialError("Geri gönderme gerekçesi yazın.")
                vals = {"status": "taslak", "note": note_t}
            else:
                vals = {"status": "onayli", "approved_by": user, "approved_at": t, "note": note_t}
        elif action == "published":
            if r.status != "onayli":
                raise SocialError("Yalnız onaylı gönderi yayınlandı işaretlenir.", 409)
            u = one_line(url, 1000)
            if not u or not _URL.match(u):
                raise SocialError("Paylaşımın bağlantısını girin (http:// ya da https://).")
            vals = {"status": "yayinlandi", "published_url": u, "published_at": t, "published_by": user}
        elif action == "cancel":
            if r.status in ("yayinlandi", "iptal"):
                raise SocialError("Yayınlanmış ya da iptal edilmiş gönderi iptal edilemez.", 409)
            if not note_t:
                raise SocialError("İptal gerekçesi yazın.")
            vals = {"status": "iptal", "note": note_t}
        elif action == "reopen":
            if r.status != "iptal":
                raise SocialError("Yalnız iptal edilmiş gönderi yeniden açılır.", 409)
            vals = {"status": "taslak" if r.text else "fikir", "approved_by": None, "approved_at": None}
        else:
            raise SocialError("İşlem tanınmıyor.", 404)
        c.execute(POSTS.update().where(POSTS.c.id == r.id).values(updated_by=user, updated_at=t, **vals))
        event(c, r.id, user, action, note_t or (one_line(url, 1000) if action == "published" else None),
              {"eski": r.status, "yeni": vals["status"]})
    return get_post(engine, tenant, pid, st)


def delete_post(engine: sa.engine.Engine, tenant: str, pid: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = _post_row(c, tenant, pid, lock=True)
        if r.status not in OPEN:
            raise SocialError("Yalnız fikir ya da taslak gönderi silinir; diğerleri iptal edilir.", 409)
        for tbl in (EVENTS, JOBS):
            c.execute(tbl.delete().where(tbl.c.post_id == r.id))
        c.execute(METRICS.update().where(METRICS.c.post_id == r.id).values(post_id=None))
        c.execute(POSTS.delete().where(POSTS.c.id == r.id))
    return {"id": r.id, "kitapAd": r.kitap_ad}


def set_draft(engine: sa.engine.Engine, pid: str, draft: dict[str, Any], user: str) -> None:
    with engine.begin() as c:
        c.execute(POSTS.update().where(POSTS.c.id == pid).values(draft_json=dump(draft)))
        event(c, pid, user, "zeki-taslak", None, {"secenek": len(draft.get("secenekler") or []), "dusen": draft.get("dusen")})


def set_kind(engine: sa.engine.Engine, pid: str, kind: str, prob: Optional[float]) -> None:
    with engine.begin() as c:
        c.execute(POSTS.update().where(POSTS.c.id == pid, POSTS.c.kind.is_(None)).values(kind=kind, kind_source="zeki", kind_prob=prob))
        event(c, pid, "sistem", "zeki-tur", None, {"tur": kind, "olasilik": prob})


# ------------------------------------------------------------------ iş kuyruğu (Zeki AI)


def job_create(engine: sa.engine.Engine, tenant: str, user: str, tur: str, post_id: Optional[str] = None,
               hedef: Optional[str] = None) -> dict[str, Any]:
    with engine.begin() as c:
        cond = [JOBS.c.tenant_id == tenant, JOBS.c.tur == tur, JOBS.c.status.in_(("bekliyor", "calisiyor"))]
        cond.append(JOBS.c.post_id == post_id if post_id else JOBS.c.hedef == hedef)
        if c.execute(sa.select(JOBS.c.id).where(*cond)).first():
            raise SocialError("Aynı iş zaten sürüyor; bitince yeniden deneyin.", 409)
        jid = _uid()
        c.execute(JOBS.insert().values(id=jid, tenant_id=tenant, post_id=post_id, tur=tur, hedef=hedef, status="bekliyor",
                                       created_by=user, created_at=now()))
    return job_get(engine, jid)


def job_update(engine: sa.engine.Engine, jid: str, **vals: Any) -> None:
    if "result" in vals:
        vals["result_json"] = dump(vals.pop("result"))
    if vals.get("status") in ("bitti", "hata"):
        vals["finished_at"] = now()
    with engine.begin() as c:
        c.execute(JOBS.update().where(JOBS.c.id == jid).values(**vals))


def _job(j: Any) -> dict[str, Any]:
    return {"id": j.id, "postId": j.post_id, "tur": j.tur, "hedef": j.hedef, "durum": j.status, "adim": j.step, "hata": j.error,
            "sonuc": loads(j.result_json, None), "olusturan": j.created_by, "olusturma": iso(j.created_at), "bitis": iso(j.finished_at)}


def job_get(engine: sa.engine.Engine, jid: str) -> dict[str, Any]:
    with engine.connect() as c:
        j = c.execute(sa.select(JOBS).where(JOBS.c.id == jid)).first()
    if not j:
        raise SocialError("İş bulunamadı.", 404)
    return _job(j)


def jobs_stmt(tenant: str, *, post_id: Optional[str] = None, hedef: Optional[str] = None):
    cond = [JOBS.c.tenant_id == tenant]
    cond.append(JOBS.c.post_id == post_id if post_id else sa.and_(JOBS.c.post_id.is_(None), JOBS.c.hedef == hedef))
    return sa.select(JOBS).where(*cond).order_by(JOBS.c.created_at.desc())


def jobs_of(engine: sa.engine.Engine, tenant: str, *, post_id: Optional[str] = None, hedef: Optional[str] = None) -> list[dict[str, Any]]:
    with engine.connect() as c:
        rows = c.execute(jobs_stmt(tenant, post_id=post_id, hedef=hedef)).all()
    return [_job(j) for j in rows]


def fail_stale_jobs(engine: sa.engine.Engine) -> int:
    with engine.begin() as c:
        res = c.execute(JOBS.update().where(JOBS.c.status.in_(("bekliyor", "calisiyor"))).values(
            status="hata", error="Servis yeniden başladı; iş yarıda kaldı. Yeniden başlatın.", finished_at=now()))
    return int(res.rowcount or 0)


# ------------------------------------------------------------------ fırsat kuralları (saf işlevler)


def occasion_items(days: list[dict[str, Any]], books: dict[str, list[dict[str, Any]]], ref: date, horizon: int,
                   lead: int, planned_books: dict[str, set[str]], resolve: Callable[[dict[str, Any]], dict[str, Any]],
                   next_occurrence: Callable[[dict[str, Any], date], Any]) -> list[dict[str, Any]]:
    """Önümüzdeki `horizon` gün içinde başlayan (ya da süren) özel günler, bağlı kitaplar ve takvimdeki kitap sayısı.

    `planned_books[gün]`: o güne bağlı ya da gün penceresinde (hazırlık başı → gün sonu) takvimde gönderisi olan kitap
    kimlikleri. «Uyarı»: güne `lead` gün ya da daha az kaldı ve bağlı kitaplardan hiçbiri takvimde değil."""
    out = []
    for d in days:
        how = resolve(d)
        occ = next_occurrence(how, ref) if how.get("method") != "unknown" else None
        if not occ:
            continue
        start, end = occ
        left = (start - ref).days
        if left > horizon:
            continue
        linked = books.get(d["key"]) or []
        ids = {b["bookId"] for b in linked}
        on_cal = planned_books.get(d["key"], set()) & ids if ids else planned_books.get(d["key"], set())
        out.append({
            "key": d["key"], "ad": d["name"], "baslangic": start.isoformat(), "bitis": end.isoformat(), "kalanGun": max(0, left),
            "suruyor": start <= ref <= end, "yontem": how.get("why"), "kesinlik": how.get("precision"),
            "kaynak": d.get("source"), "kitapSayisi": len(ids), "takvimde": len(on_cal),
            "kitaplar": sorted(({**b, "takvimde": b["bookId"] in on_cal} for b in linked), key=lambda b: (b["takvimde"], b.get("ad") or "")),
            "uyari": left <= lead and bool(ids) and not on_cal,
        })
    out.sort(key=lambda x: (x["baslangic"], x["ad"]))
    return out


def backlist_rank(sales: dict[str, float], info: dict[str, dict[str, Any]], last_post: dict[str, str], ref: date,
                  min_age_days: int, quiet_days: int, exclude_status: Iterable[str] = ()) -> list[dict[str, Any]]:
    """Son 12 ayın net adedine göre çok satan backlist: ilk yayını `min_age_days`tan eski, son `quiet_days` gündür
    takvimde (yayınlanmış ya da planlı) gönderisi olmayan kitaplar. Net adedi sıfır ya da eksi olan girmez."""
    skip = {s.strip().casefold() for s in exclude_status if s.strip()}
    out = []
    for code, qty in sales.items():
        if qty is None or qty <= 0:
            continue
        b = info.get(code) or {}
        ilk = b.get("ilk_yayin")
        if ilk:
            try:
                if (ref - date.fromisoformat(str(ilk)[:10])).days < min_age_days:
                    continue
            except ValueError:
                pass
        elif min_age_days:
            continue                     # ilk yayını bilinmeyen kitap backlist sayılmaz (yeni olabilir)
        if skip and str(b.get("statu") or "").strip().casefold() in skip:
            continue
        last = last_post.get(code)
        if last and (ref - date.fromisoformat(last[:10])).days < quiet_days:
            continue
        out.append({"stokKodu": code, "ad": b.get("ad"), "yazar": b.get("yazar"), "yayinevi": b.get("yayinevi"),
                    "ilkYayin": str(ilk)[:10] if ilk else None, "adet12": round(qty, 2), "sonGonderi": last[:10] if last else None})
    out.sort(key=lambda x: (-x["adet12"], x["stokKodu"]))
    for i, x in enumerate(out, 1):
        x["sira"] = i
    return out


def month_window(end: date) -> tuple[int, int]:
    """`end` ayı dahil geriye 12 ay: (başlangıç, bitiş) ay indeksleri (yıl*12 + ay-1)."""
    e = end.year * 12 + end.month - 1
    return e - 11, e


# ------------------------------------------------------------------ içe aktarma


#: Platform dışa aktarım başlıkları (katlanmış, noktalama atılmış) → ölçü. Gerçek dosyalarla ölçülecek; tanınmayan
#: kolon atlanır ve içe aktarma özetinde listelenir.
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "day": ("date", "day", "tarih", "gun", "publish time", "published", "yayinlanma zamani", "yayin tarihi", "time",
            "created", "olusturma zamani", "paylasim tarihi", "video publish time"),
    "post_url": ("permalink", "post link", "link", "url", "baglanti", "kalici baglanti", "tweet permalink", "gonderi baglantisi",
                 "content link"),
    "impressions": ("impressions", "gosterim", "gosterimler", "views", "goruntulenme", "goruntulemeler", "video views", "izlenme"),
    "reach": ("reach", "erisim", "accounts reached", "erisilen hesaplar", "unique viewers", "tekil izleyici"),
    "likes": ("likes", "begeni", "begeniler", "reactions", "tepkiler", "tepki"),
    "comments": ("comments", "yorum", "yorumlar", "replies", "yanitlar"),
    "shares": ("shares", "paylasim", "paylasimlar", "retweets", "reposts", "yeniden paylasim"),
    "saves": ("saves", "kaydetme", "kaydetmeler", "kaydedilenler", "bookmarks", "yer isaretleri"),
    "followers": ("followers", "takipci", "takipciler", "follows", "takip", "new followers", "subscribers gained", "abone"),
}
_FOLD = str.maketrans("İIıŞşĞğÜüÖöÇçÂâÎîÛû", "iiissgguuooccaaiiuu")


def fold(s: Any) -> str:
    return " ".join(re.sub(r"[^\w]+", " ", str(s or "").translate(_FOLD).lower()).split())


def map_headers(headers: list[Any]) -> dict[str, int]:
    """Kolon adı → ölçü. Tam eşleşme önce; yoksa başlığın başında geçen takma ad. Her ölçüye ilk uyan kolon."""
    out: dict[str, int] = {}
    folded = [fold(h) for h in headers]
    for key, names in HEADER_ALIASES.items():
        for i, h in enumerate(folded):
            if h in names:
                out.setdefault(key, i)
                break
        if key in out:
            continue
        for i, h in enumerate(folded):
            if any(h.startswith(n + " ") for n in names) and i not in out.values():
                out[key] = i
                break
    return out


def _num(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if isinstance(v, float) and math.isnan(v) else float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if not s or s in ("-", "—"):
        return None
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):          # 1.234,5
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", s):        # 1,234.5
        s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _day_of(v: Any) -> Optional[str]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).date().isoformat()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s[:19].replace("Z", "")).date().isoformat()
    except ValueError:
        pass
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(1)), int(m.group(2))).isoformat()
        except ValueError:
            return None
    return None


def parse_rows(rows: list[list[Any]], default_day: Optional[str] = None) -> dict[str, Any]:
    """İlk dolu satır başlıktır (platform dosyalarında başlıktan önce açıklama satırı olabilir: ilk ölçü tanınan satır
    başlık sayılır). Dönen: satırlar, tanınan/atlanan kolonlar, atlanan satır sayısı ve nedeni."""
    head_i, cols = None, {}
    for i, r in enumerate(rows[:20]):
        m = map_headers(list(r))
        if any(k in m for k in METRIC_COLS):
            head_i, cols = i, m
            break
    if head_i is None:
        seen = [str(x) for x in (rows[0] if rows else []) if x not in (None, "")]
        raise SocialError("Dosyada ölçü kolonu tanınmadı (gösterim, erişim, beğeni, yorum, paylaşım, kaydetme, takipçi). "
                          f"Başlıklar: {', '.join(seen[:20]) or 'yok'}.")
    headers = [str(x or "") for x in rows[head_i]]
    out, skipped = [], {"tarihsiz": 0, "bos": 0}
    for r in rows[head_i + 1:]:
        cells = list(r) + [None] * (len(headers) - len(r))
        vals = {k: _num(cells[i]) for k, i in cols.items() if k in METRIC_COLS}
        if all(v is None for v in vals.values()):
            skipped["bos"] += 1
            continue
        d = _day_of(cells[cols["day"]]) if "day" in cols else None
        d = d or default_day
        if not d:
            skipped["tarihsiz"] += 1
            continue
        url = one_line(cells[cols["post_url"]], 1000) if "post_url" in cols else None
        out.append({"day": d, "post_url": url if url and _URL.match(url) else None, **vals})
    return {"rows": out, "cols": {k: headers[i] for k, i in cols.items()},
            "ignored": [h for i, h in enumerate(headers) if h and i not in cols.values()], "skipped": skipped}


def extract_rows(filename: str, data: bytes) -> list[list[Any]]:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "xlsx":
        import io

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    if ext in ("csv", "txt"):
        import csv
        import io

        body = None
        for enc in ("utf-8-sig", "utf-16", "cp1254", "latin-1"):
            try:
                body = data.decode(enc)
                if enc == "utf-16" and "\x00" in body:
                    continue
                break
            except UnicodeDecodeError:
                continue
        body = body if body is not None else data.decode("utf-8", "replace")
        try:
            dialect = csv.Sniffer().sniff(body[:4096], delimiters=";,\t|")
            return [row for row in csv.reader(io.StringIO(body), dialect)]
        except csv.Error:
            return [row for row in csv.reader(io.StringIO(body))]
    raise SocialError("Yalnız .csv ya da .xlsx dosyası içe aktarılır.")


def import_metrics(engine: sa.engine.Engine, tenant: str, user: str, account_id: str, file_name: str, data: bytes,
                   default_day: Optional[str] = None) -> dict[str, Any]:
    if not data:
        raise SocialError("Dosya boş.")
    dd = day(default_day, "Varsayılan gün").isoformat() if default_day else None
    parsed = parse_rows(extract_rows(file_name or "", data), dd)
    if not parsed["rows"]:
        raise SocialError("Dosyada içe aktarılacak satır bulunamadı" + (" (tarih kolonu yok; «gün» girin)." if parsed["skipped"]["tarihsiz"] else "."))
    with engine.begin() as c:
        _account_row(c, tenant, account_id)
        posts = c.execute(sa.select(POSTS.c.id, POSTS.c.published_url).where(
            POSTS.c.tenant_id == tenant, POSTS.c.account_id == account_id, POSTS.c.published_url.is_not(None))).all()
        by_url = {norm_url(p.published_url): p.id for p in posts if p.published_url}
        iid = _uid()
        matched = 0
        t = now()
        for r in parsed["rows"]:
            pid = by_url.get(norm_url(r["post_url"])) if r.get("post_url") else None
            matched += 1 if pid else 0
            c.execute(METRICS.insert().values(id=_uid(), tenant_id=tenant, import_id=iid, account_id=account_id, post_id=pid,
                                              day=r["day"], post_url=r.get("post_url"), created_by=user, created_at=t,
                                              **{k: r.get(k) for k in METRIC_COLS}))
        summary = {"kolonlar": parsed["cols"], "atlananKolonlar": parsed["ignored"], "atlananSatir": parsed["skipped"],
                   "toplam": {k: round(sum(r.get(k) or 0 for r in parsed["rows"]), 2) for k in METRIC_COLS}}
        c.execute(IMPORTS.insert().values(id=iid, tenant_id=tenant, account_id=account_id, file_name=(file_name or "dosya")[:300],
                                          rows=len(parsed["rows"]), matched=matched, summary_json=dump(summary),
                                          created_by=user, created_at=t))
    return {"id": iid, "satir": len(parsed["rows"]), "eslesen": matched, **summary}


def imports_stmt(tenant: str):
    return sa.select(IMPORTS).where(IMPORTS.c.tenant_id == tenant).order_by(IMPORTS.c.created_at.desc())


def list_imports(engine: sa.engine.Engine, tenant: str) -> list[dict[str, Any]]:
    with engine.connect() as c:
        rows = c.execute(imports_stmt(tenant)).all()
    return [{"id": r.id, "accountId": r.account_id, "dosya": r.file_name, "satir": r.rows, "eslesen": r.matched,
             "ozet": loads(r.summary_json, {}), "kim": r.created_by, "zaman": iso(r.created_at)} for r in rows]


def delete_import(engine: sa.engine.Engine, tenant: str, iid: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = c.execute(sa.select(IMPORTS).where(IMPORTS.c.tenant_id == tenant, IMPORTS.c.id == str(iid)[:32])).first()
        if not r:
            raise SocialError("İçe aktarma bulunamadı.", 404)
        c.execute(METRICS.delete().where(METRICS.c.import_id == r.id))
        c.execute(IMPORTS.delete().where(IMPORTS.c.id == r.id))
    return {"id": r.id, "dosya": r.file_name, "satir": r.rows}


def add_manual_metric(engine: sa.engine.Engine, tenant: str, user: str, pid: str, body: dict[str, Any]) -> dict[str, Any]:
    """Tek gönderinin içgörüsü elle (platform ekranından okunarak). Aynı gün ikinci kez girilirse eskisinin yerine geçer."""
    vals = {k: _num(body.get(k)) for k in METRIC_COLS}
    if all(v is None for v in vals.values()):
        raise SocialError("En az bir ölçü girin.")
    if any(v is not None and v < 0 for v in vals.values()):
        raise SocialError("Ölçüler eksi olamaz.")
    d = day(body.get("day") or today().isoformat(), "Gün").isoformat()
    with engine.begin() as c:
        r = _post_row(c, tenant, pid)
        if r.status != "yayinlandi" or not r.account_id:
            raise SocialError("İçgörü yalnız yayınlanmış gönderiye girilir.", 409)
        c.execute(METRICS.delete().where(METRICS.c.post_id == r.id, METRICS.c.day == d, METRICS.c.import_id.is_(None)))
        c.execute(METRICS.insert().values(id=_uid(), tenant_id=tenant, import_id=None, account_id=r.account_id, post_id=r.id,
                                          day=d, post_url=r.published_url, created_by=user, created_at=now(), **vals))
        event(c, r.id, user, "icgoru", None, {"gun": d, **{k: v for k, v in vals.items() if v is not None}})
    return {"postId": pid, "day": d, **vals}


def metrics_stmt(tenant: str, pid: str):
    return sa.select(METRICS).where(METRICS.c.tenant_id == tenant, METRICS.c.post_id == pid).order_by(METRICS.c.day.desc())


def post_metrics(engine: sa.engine.Engine, tenant: str, pid: str) -> list[dict[str, Any]]:
    with engine.connect() as c:
        rows = c.execute(metrics_stmt(tenant, pid)).all()
    return [{"id": r.id, "day": r.day, "kaynak": "dosya" if r.import_id else "elle", **{k: getattr(r, k) for k in METRIC_COLS}}
            for r in rows]


# ------------------------------------------------------------------ rapor


def _month(v: str) -> tuple[date, date]:
    try:
        y, m = (int(x) for x in str(v).split("-")[:2])
        start = date(y, m, 1)
    except (TypeError, ValueError):
        raise SocialError("Ay YYYY-AA biçiminde olmalı.") from None
    end = date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)
    return start, end


def _eng(r: dict[str, Any]) -> float:
    return sum((r.get(k) or 0) for k in ("likes", "comments", "shares", "saves"))


def report_stmts(tenant: str, month: str) -> tuple[Any, Any, Any, Any]:
    """Ay raporu okumaları: hesaplar, aydaki ölçü satırları, aya planlanan gönderiler, ölçüsü olan gönderilerin türü."""
    start, end = _month(month)
    mcond = [METRICS.c.tenant_id == tenant, METRICS.c.day >= start.isoformat(), METRICS.c.day <= end.isoformat()]
    linked = sa.select(METRICS.c.post_id).where(*mcond, METRICS.c.post_id.is_not(None))
    return (accounts_stmt(tenant),
            sa.select(METRICS).where(*mcond),
            sa.select(POSTS).where(POSTS.c.tenant_id == tenant, POSTS.c.planned_at >= f"{start.isoformat()} 00:00",
                                   POSTS.c.planned_at <= f"{end.isoformat()} 23:59"),
            sa.select(POSTS.c.id, POSTS.c.kind).where(POSTS.c.id.in_(linked)))


def report(engine: sa.engine.Engine, tenant: str, month: str) -> dict[str, Any]:
    """Ay raporu: ölçüler gün sütununa göre ayda (`day BETWEEN ay başı AND ay sonu`), gönderi sayıları paylaşım
    zamanına göre. Etkileşim = beğeni + yorum + paylaşım + kaydetme; oran = etkileşim ÷ erişim (erişim yoksa boş).
    Kişi adı yok (KVKK): rapor hesap, içerik türü ve gönderi düzeyinde."""
    start, end = _month(month)
    aq, mq, pq, kq = report_stmts(tenant, month)
    with engine.connect() as c:
        accs = {a.id: account_dict(a) for a in c.execute(aq).all()}
        mrows = c.execute(mq).all()
        prows = c.execute(pq).all()
        kinds = {p.id: (p.kind or "") for p in c.execute(kq).all()}

    def blank() -> dict[str, float]:
        return {k: 0.0 for k in METRIC_COLS if k != "followers"}

    total, by_acc, by_kind = blank(), {}, {}
    for m in mrows:
        vals = {k: getattr(m, k) or 0.0 for k in METRIC_COLS if k != "followers"}
        for k, v in vals.items():
            total[k] += v
        a = by_acc.setdefault(m.account_id, {**blank(), "gonderi": set(), "takipci": None, "takipciGun": None})
        for k, v in vals.items():
            a[k] += v
        if m.post_id:
            a["gonderi"].add(m.post_id)
        if m.followers is not None and (a["takipciGun"] is None or m.day >= a["takipciGun"]):
            a["takipci"], a["takipciGun"] = m.followers, m.day
        if m.post_id:
            kk = kinds.get(m.post_id) or "etiketsiz"
            b = by_kind.setdefault(kk, {**blank(), "gonderi": set()})
            for k, v in vals.items():
                b[k] += v
            b["gonderi"].add(m.post_id)

    def fin(x: dict[str, Any]) -> dict[str, Any]:
        e = _eng(x)
        out = {k: round(v, 2) for k, v in x.items() if k in METRIC_COLS}
        out.update(etkilesim=round(e, 2), oran=(round(e / x["reach"], 4) if x.get("reach") else None))
        if "gonderi" in x:
            out["gonderi"] = len(x["gonderi"])
            out["gonderiBasina"] = round(e / len(x["gonderi"]), 2) if x["gonderi"] else None
        return out

    status_counts: dict[str, int] = {}
    waits = []
    for p in prows:
        status_counts[p.status] = status_counts.get(p.status, 0) + 1
        if p.submitted_at and p.approved_at:
            a, b = p.submitted_at, p.approved_at
            a = a if a.tzinfo else a.replace(tzinfo=timezone.utc)
            b = b if b.tzinfo else b.replace(tzinfo=timezone.utc)
            waits.append((b - a).total_seconds() / 3600.0)
    return {
        "ay": f"{start.year}-{start.month:02d}", "baslangic": start.isoformat(), "bitis": end.isoformat(),
        "toplam": fin(total), "olcuSatiri": len(mrows), "gonderiliOlcu": sum(1 for m in mrows if m.post_id),
        "hesaplar": sorted(({"accountId": k, "hesap": (accs.get(k) or {}).get("ad") or k, "platform": (accs.get(k) or {}).get("platform"),
                             **fin(v), "takipci": v["takipci"], "takipciGun": v["takipciGun"]} for k, v in by_acc.items()),
                           key=lambda x: -x["etkilesim"]),
        "turler": sorted(({"tur": k, "turAdi": KINDS.get(k, "Türü girilmemiş"), **fin(v)} for k, v in by_kind.items()),
                         key=lambda x: -(x.get("gonderiBasina") or 0)),
        "gonderiDurum": status_counts, "gonderiSayisi": len(prows),
        "onaySuresiSaat": round(statistics.median(waits), 1) if waits else None, "onaySayisi": len(waits),
    }


def report_facts(rep: dict[str, Any]) -> list[str]:
    """Rapor yorumunda modelin kullanabileceği sayılar (başka sayı yazılırsa cümle düşer)."""
    t = rep["toplam"]
    f = [f"Ay: {rep['ay']}", f"Gönderi sayısı: {rep['gonderiSayisi']}", f"Toplam erişim: {int(t.get('reach') or 0)}",
         f"Toplam gösterim: {int(t.get('impressions') or 0)}", f"Toplam etkileşim: {int(t.get('etkilesim') or 0)}"]
    if t.get("oran") is not None:
        f.append(f"Etkileşim oranı: %{round(t['oran'] * 100, 1)}")
    for x in rep["turler"]:
        f.append(f"{x['turAdi']}: {x['gonderi']} gönderi, gönderi başına {int(x.get('gonderiBasina') or 0)} etkileşim")
    for x in rep["hesaplar"]:
        f.append(f"{x['hesap']}: {int(x['etkilesim'])} etkileşim")
    if rep.get("onaySuresiSaat") is not None:
        f.append(f"Onay süresi medyanı: {rep['onaySuresiSaat']} saat")
    return f


# ------------------------------------------------------------------ zamanlayıcı özeti


def due_tomorrow(engine: sa.engine.Engine, tenant: str, st: dict[str, Any], ref: date) -> list[dict[str, Any]]:
    """Yarın paylaşılacak ama onaylı değil ya da görselsiz gönderiler."""
    t = ref + timedelta(days=1)
    rows = list_posts(engine, tenant, st, frm=t, to=t, status="fikir,taslak,onayda,onayli")
    out = []
    for p in rows:
        why = []
        if p["status"] != "onayli":
            why.append(f"durum {p['statusAdi'].lower()}")
        if not p["assets"]:
            why.append("görsel yok")
        if why:
            out.append({**p, "neden": why})
    return out


def digest_text(dg: dict[str, Any], link: str) -> Optional[str]:
    parts = []
    if dg["yarin"]:
        parts.append("Yarın paylaşılacak ama hazır olmayan gönderiler:")
        for p in dg["yarin"]:
            parts.append(f"  - {p['plannedAt']} · {(p.get('account') or {}).get('ad') or 'hesap yok'} · "
                         f"{p.get('kitapAd') or p.get('occasionAd') or p['id']} ({', '.join(p['neden'])})")
    if dg["ozelGun"]:
        parts.append("Yaklaşan özel günler (bağlı kitaplardan hiçbiri takvimde değil):")
        for o in dg["ozelGun"]:
            parts.append(f"  - {o['ad']}: {o['kalanGun']} gün kaldı, {o['kitapSayisi']} bağlı kitap")
    if dg["onayda"]:
        parts.append(f"Onay bekleyen gönderi: {dg['onayda']}")
    if not parts:
        return None
    return "\n".join(parts) + (f"\n\nTakvim: {link}" if link else "") + \
        "\n\nPortal hiçbir sosyal medya hesabına kendiliğinden paylaşım yapmaz; paylaşımı ekip yapar."


# ------------------------------------------------------------------ Zeki AI çıktısının ayrıştırılması


def parse_options(raw: str) -> list[dict[str, str]]:
    """«SEÇENEK n» başlıklarıyla ayrılmış seçenekler; her birinde «ETİKETLER:» satırı (yoksa # ile başlayan son satır)."""
    parts = re.split(r"(?im)^[ \t*#]*se[çc]enek[ \t]*\d+[ \t]*\**[ \t]*[:.)\-–]?[ \t]*\**[ \t]*", raw or "")
    parts = [p.strip() for p in parts if p.strip()]
    if len(parts) <= 1:
        parts = [p.strip() for p in re.split(r"\n\s*-{3,}\s*\n", raw or "") if p.strip()]
    out = []
    for p in parts:
        lines = p.splitlines()
        m = re.search(r"(?im)^[ \t]*\**[ \t]*(?:et[iİı]ket(?:ler)?|hashtag(?:ler)?)[ \t]*\**[ \t]*:[ \t]*(.*)$", p)
        if m:
            tags, body = m.group(1), p[:m.start()].strip()
        elif len(lines) > 1 and lines[-1].strip().startswith("#"):
            tags, body = lines[-1], "\n".join(lines[:-1]).strip()
        else:
            tags, body = "", p
        if body:
            out.append({"metin": body, "etiketler": tags})
    return out


# ------------------------------------------------------------------ yayına hazır paket ve rapor PDF'i


def _safe_name(s: Any) -> str:
    t = str(s or "").translate(_FOLD).lower()
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")[:60] or "gorsel"


def package_zip(post: dict[str, Any], fetch: Callable[[dict[str, Any]], Optional[tuple[bytes, str]]], st: dict[str, Any]) -> bytes:
    """Onaylı gönderinin paketi: metin (etiketlerle), bilgi dosyası ve görseller. Görsel kaynağından okunamazsa
    OKUBENI'de yazar; paket yine iner. Portal hiçbir platforma paylaşım yapmaz."""
    import io
    import zipfile

    acc = post.get("account") or {}
    lim = st["limits"].get(acc.get("platform") or "")
    buf = io.BytesIO()
    notes: list[str] = []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        body = (post.get("text") or "").rstrip()
        if post.get("hashtags"):
            body += "\n\n" + post["hashtags"]
        z.writestr("metin.txt", body + "\n")
        info = [f"Gönderi: {post['id']} ({post['statusAdi']})",
                f"Hesap: {acc.get('ad') or '—'} ({acc.get('platformAdi') or '—'})",
                f"Paylaşım zamanı: {post.get('plannedAt') or '—'}",
                f"Kitap: {post.get('kitapAd') or '—'} {post.get('stokKodu') or ''}".rstrip(),
                f"Özel gün: {post.get('occasionAd') or '—'}", f"İçerik türü: {post.get('kindAdi') or '—'}",
                f"Onaya gönderen: {post.get('submittedBy') or '—'}",
                f"Onaylayan: {post.get('approvedBy') or '—'} {post.get('approvedAt') or ''}".rstrip(),
                f"Karakter: {len(body)}" + (f" / sınır {lim}" if lim else "")]
        if post.get("publishedUrl"):
            info.append(f"Yayın bağlantısı: {post['publishedUrl']}")
        z.writestr("bilgi.txt", "\n".join(info) + "\n")
        for i, a in enumerate(post.get("assets") or [], 1):
            label = a.get("ad") or a.get("sid") or a.get("id")
            try:
                got = fetch(a)
            except Exception as e:  # noqa: BLE001 — kaynak kapalıysa paket yine iner
                notes.append(f"{i}. görsel okunamadı ({label}): {str(e)[:120]}")
                continue
            if got is None:
                notes.append(f"{i}. görsel kaynağında bulunamadı ({label}).")
                continue
            data, mime = got
            ext = {"image/png": "png", "image/jpeg": "jpg", "image/webp": "webp", "application/pdf": "pdf"}.get(mime, "bin")
            z.writestr(f"{i:02d}-{_safe_name(a.get('ad'))}.{ext}", data)
        readme = ["Yayına hazır paket: onaylı metin, etiketler ve görseller.",
                  "Paylaşımı ekip kendi hesabından yapar; portal hiçbir sosyal medya hesabına kendiliğinden gönderim yapmaz.",
                  "Paylaştıktan sonra portalda gönderiyi «Yayınlandı» işaretleyip bağlantısını girin."]
        z.writestr("OKUBENI.txt", "\n".join(readme + ([""] + notes if notes else [])) + "\n")
    return buf.getvalue()


def report_pdf(rep: dict[str, Any], yorum: Optional[str], user: str = "") -> bytes:
    try:
        from fpdf import FPDF
    except ImportError as e:  # pragma: no cover
        raise SocialError("PDF üretici sunucuda kurulu değil.", 503) from e
    from semantic_bridge.editorial_export import _fold, _resolve_fonts

    regular, bold = _resolve_fonts()
    fam = "Body" if regular else "Helvetica"
    T = (lambda s: str(s if s is not None else "")) if regular else (lambda s: _fold(str(s if s is not None else "")))
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_title(T(f"Sosyal medya raporu {rep['ay']}"))
    pdf.set_author("Zeki AI")
    pdf.set_creator("Zeki AI")
    pdf.set_margins(14, 14, 14)
    pdf.set_auto_page_break(auto=True, margin=14)
    if regular:
        pdf.add_font(fam, "", str(regular))
        pdf.add_font(fam, "B", str(bold or regular))
    pdf.add_page()
    W = pdf.w - pdf.l_margin - pdf.r_margin

    def para(s: str, size: float = 9.5, style: str = "", gap: float = 1.5) -> None:
        pdf.set_font(fam, style, size)
        pdf.multi_cell(W, size * 0.5, T(s), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    def row(cells: list[tuple[str, float]], style: str = "") -> None:
        pdf.set_font(fam, style, 8.5)
        for txt, w in cells:
            t = T(txt)
            while pdf.get_string_width(t) > w - 1.5 and len(t) > 1:
                t = t[:-2] + "…" if len(t) > 2 else t[:-1]
            pdf.cell(w, 5.5, t, border="B")
        pdf.ln(5.5)

    def n(v: Any) -> str:
        return "—" if v is None else f"{v:,.0f}".replace(",", ".")

    def pct(v: Any) -> str:
        return "—" if v is None else f"%{v * 100:.1f}".replace(".", ",")

    t = rep["toplam"]
    para("Timaş Yayınları · sosyal medya aylık raporu", 8.5)
    para(f"{rep['ay']} ({rep['baslangic']} – {rep['bitis']})", 15, "B", 1)
    para(f"Gönderi: {rep['gonderiSayisi']} · erişim {n(t.get('reach'))} · gösterim {n(t.get('impressions'))} · etkileşim "
         f"{n(t.get('etkilesim'))} · oran {pct(t.get('oran'))}"
         + (f" · onay süresi medyanı {str(rep['onaySuresiSaat']).replace('.', ',')} saat" if rep.get("onaySuresiSaat") is not None else ""))
    if yorum:
        para("Zeki AI yorumu", 11, "B", 0.5)
        para(yorum, 9)
    para("İçerik türüne göre", 11, "B", 0.5)
    row([("Tür", 44), ("Gönderi", 20), ("Erişim", 28), ("Etkileşim", 28), ("Gönderi başına", 30), ("Oran", 22)], "B")
    for x in rep["turler"]:
        row([(x["turAdi"], 44), (n(x.get("gonderi")), 20), (n(x.get("reach")), 28), (n(x.get("etkilesim")), 28),
             (n(x.get("gonderiBasina")), 30), (pct(x.get("oran")), 22)])
    pdf.ln(2)
    para("Hesaba göre", 11, "B", 0.5)
    row([("Hesap", 62), ("Erişim", 28), ("Gösterim", 28), ("Etkileşim", 28), ("Takipçi", 26)], "B")
    for x in rep["hesaplar"]:
        row([(x["hesap"], 62), (n(x.get("reach")), 28), (n(x.get("impressions")), 28), (n(x.get("etkilesim")), 28),
             (n(x.get("takipci")), 26)])
    pdf.ln(3)
    para("Ölçüler platformların dışa aktarım dosyasından ya da elle girildi (portal platformlara bağlanmaz). Etkileşim = "
         "beğeni + yorum + paylaşım + kaydetme; oran = etkileşim ÷ erişim.", 8)
    para(f"Hazırlayan: {user or '—'} · Zeki AI", 8)
    return bytes(pdf.output())
