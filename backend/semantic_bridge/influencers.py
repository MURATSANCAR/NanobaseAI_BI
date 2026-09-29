"""M23 İnfluencer ve işbirliği yönetimi: içerik üreticisi kayıt defteri, işbirliği panosu, gerekçeli aday sırası,
brief/iletişim/sadakat taslağı, sonuç ve CPE raporu, ödeme listesi. Hepsi kendi tablolarımızda (`semantic_infl_*`);
CRM'de influencer varlığı yok, CRM ve Logo'ya yazılmaz, T-soft'a dokunulmaz.

- **Kayıt defteri** (`semantic_infl_people` + `semantic_infl_accounts`): kişi, platform hesapları (platform + kullanıcı
  adı kiracıda tekil), konu etiketleri, kitlenin yaş grupları, ücret aralığı, iletişim tercihi, «iletişim kurulmasın»,
  reşit olmayan içerik üreticisi işareti (veli onayı hukuka sorulacak). İsteğe bağlı CRM kişisi bağı.
- **Hesap sayıları** (`semantic_infl_snapshots`): takipçi, gönderi, ortalama beğeni/yorum — ilk sürümde **elle ya da
  dosyayla**; resmî API ile periyodik alma ikinci sürümdedir (`INFLUENCER_API_ENABLED`). Kazıma yok, sahte takipçi
  skoru yok. «Ani sıçrama» yalnız kendi anlık görüntülerimizden hesaplanır.
- **İşbirliği** (`semantic_infl_collabs`): teklif → kitap gönderildi → içerik bekleniyor → içerik onayda → yayında →
  rapor → ödeme → kapandı (ya da vazgeçildi). Kapılar: teklif seçim onayı olmadan ilerlemez (öneren onaylayamaz);
  yayında bağlantı ister; rapor yasal etiket işaretini (Reklam Kurulu sosyal medya etkileyicileri kılavuzu) ve sonuç
  sayısını ister; ödeme aşaması ücretli işte ödeme satırını (`hazir`) açar; ödeme yapılınca iş kapanır.
- **Ödeme** (`semantic_infl_payouts`): M8 hakediş deseni (iki göz, Logo belge numarasıyla ödendi); tutar aşamaya
  girildiği anki ücrettir (dondurulur). Ödeme Logo'da yapılır; burada yalnız belge numarası ve tarih tutulur.
- **Aday sırası**: puan kuralla (konu uyumu, yaş grubu, geçmiş etkileşim, ilişki puanı, son işbirliği tazeliği,
  bütçe); Zeki AI yalnız gerekçe cümlesini yazar ve cümle `marketing.guard` denetiminden geçer. Tavan yok: listenin
  tamamı döner; iletişim kurulmayacak ve bu kitabı daha önce almış kişiler ayrı listede nedeniyle görünür.
- **Taslak** (`semantic_infl_drafts`): brief, iletişim e-postası, sadakat teklifi. Yasal etiket maddesi şablonda sabittir,
  modele bırakılmaz. Otomatik dış gönderim yok: onaylanan taslağı insan kendi e-postasıyla gönderir, portal yalnız
  «gönderildi» kaydını tutar.
- **Ücret görünürlüğü** (KVKK): kişi ücret aralığı, işbirliği ücreti, harcama ve CPE yalnız `ozellik:isbirligi.onay`
  ya da `ozellik:isbirligi.odeme` sahibine döner; diğerlerinde alan boştur (`None`).
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import threading
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal
from statistics import median
from typing import Any, Callable, Iterable, Optional

import sqlalchemy as sa

from semantic_bridge import freelance as fl

log = logging.getLogger("semantic.influencers")
_md = sa.MetaData()

PLATFORMS = {"instagram": "Instagram", "youtube": "YouTube", "tiktok": "TikTok", "x": "X", "blog": "Blog",
             "podcast": "Podcast", "diger": "Diğer"}
TOPICS = {"cocuk": "Çocuk", "genc": "Genç", "edebiyat": "Edebiyat", "tarih": "Tarih ve biyografi",
          "kisisel-gelisim": "Kişisel gelişim", "din-tasavvuf": "Din ve tasavvuf", "bilim": "Bilim", "diger": "Diğer"}
#: Kitlenin yaş grupları: anahtar → (ad, alt yaş, üst yaş).
AGE_GROUPS: dict[str, tuple[str, int, int]] = {
    "0-6": ("Okul öncesi (0–6)", 0, 6), "7-10": ("Çocuk (7–10)", 7, 10), "11-14": ("İlk gençlik (11–14)", 11, 14),
    "15-18": ("Genç (15–18)", 15, 18), "18+": ("Yetişkin", 19, 120),
}
KINDS = {"hediye": "Hediye kitap", "ucretli": "Ücretli", "karsilikli": "Karşılıklı"}
STAGES = {"teklif": "Teklif", "gonderildi": "Kitap gönderildi", "icerik-bekleniyor": "İçerik bekleniyor",
          "onayda": "İçerik onayda", "yayinda": "Yayında", "rapor": "Rapor", "odeme": "Ödeme", "kapali": "Kapandı",
          "vazgecildi": "Vazgeçildi"}
#: Aşama sırası (vazgeçildi dışarıda).
FLOW = ("teklif", "gonderildi", "icerik-bekleniyor", "onayda", "yayinda", "rapor", "odeme", "kapali")
#: Panoda sütun olan açık aşamalar.
BOARD = FLOW[:-1]
CLOSED = ("kapali", "vazgecildi")
PAYOUT_STATES = {"hazir": "Hazır", "onayli": "Onaylı", "odendi": "Ödendi", "iptal": "İptal"}
DRAFT_KINDS = {"brief": "Brief", "iletisim": "İletişim e-postası", "sadakat": "Sadakat teklifi"}
CONTACT_PREFS = {"eposta": "E-posta", "dm": "Doğrudan mesaj", "telefon": "Telefon", "ajans": "Ajans / menajer"}

#: Konu etiketinin kitaptaki izleri (CRM tür metni, raf türü, hedef kitle). Katlanmış küçük harf; ayar ile genişler.
TOPIC_WORDS: dict[str, tuple[str, ...]] = {
    "cocuk": ("cocuk", "masal", "okul oncesi", "boyama", "resimli", "etkinlik kitabi"),
    "genc": ("genc", "genclik", "ergen", "young adult"),
    "edebiyat": ("roman", "oyku", "hikaye", "siir", "edebiyat", "deneme", "novella"),
    "tarih": ("tarih", "biyografi", "ani", "anilar", "hatirat", "osmanli", "otobiyografi"),
    "kisisel-gelisim": ("kisisel gelisim", "psikoloji", "motivasyon", "is dunyasi", "yonetim", "ebeveyn", "aile", "iliski"),
    "din-tasavvuf": ("din", "dini", "tasavvuf", "islam", "ilahiyat", "kuran", "hadis", "siyer", "peygamber"),
    "bilim": ("bilim", "populer bilim", "teknoloji", "doga", "uzay"),
}
DEFAULT_WEIGHTS = {"konu": 35, "yas": 15, "performans": 20, "iliski": 10, "tazelik": 10, "butce": 10}

#: Yasal etiket maddesi: şablonda sabit, modele yazdırılmaz.
DISCLOSURE_CLAUSE = ("Yasal etiket: Bu bir işbirliğidir. Paylaşımın başında ve açıkça görülecek biçimde "
                     "«#işbirliği» ya da «#reklam» etiketi ve Timaş Yayınları ile işbirliği yapıldığı belirtilmelidir "
                     "(Reklam Kurulu, sosyal medya etkileyicileri kılavuzu).")
REVIEW_CLAUSE = ("İçerik yayından önce onaya gönderilir; kitabın sonunu ya da sürprizlerini açık eden anlatım "
                 "(spoiler) ve kitapta olmayan bilgi kullanılmaz.")


def _id() -> sa.Column:
    return sa.Column("id", sa.String(32), primary_key=True)


def _ts(name: str, nullable: bool = True) -> sa.Column:
    return sa.Column(name, sa.DateTime(timezone=True), nullable=nullable)


PEOPLE = sa.Table(
    "semantic_infl_people", _md, _id(),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("crm_contact_id", sa.String(40)),
    sa.Column("name", sa.String(200), nullable=False),
    sa.Column("email", sa.String(200)),
    sa.Column("phone", sa.String(40)),
    sa.Column("city", sa.String(80)),
    sa.Column("topics_json", sa.Text, nullable=False, default="[]"),
    sa.Column("age_groups_json", sa.Text, nullable=False, default="[]"),
    sa.Column("fee_min", sa.Numeric(14, 2)),
    sa.Column("fee_max", sa.Numeric(14, 2)),
    sa.Column("currency", sa.String(3), nullable=False, default="TRY"),
    sa.Column("contact_pref", sa.String(12)),
    sa.Column("notes", sa.Text),
    sa.Column("do_not_contact", sa.Boolean, nullable=False, default=False),
    sa.Column("minor", sa.Boolean, nullable=False, default=False),
    sa.Column("created_by", sa.String(120), nullable=False),
    _ts("created_at", False),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at"),
)
ACCOUNTS = sa.Table(
    "semantic_infl_accounts", _md, _id(),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("person_id", sa.String(32), nullable=False, index=True),
    sa.Column("platform", sa.String(12), nullable=False),
    sa.Column("handle", sa.String(120), nullable=False),
    sa.Column("url", sa.String(500)),
    sa.Column("verified_by_api", sa.Boolean, nullable=False, default=False),
    _ts("created_at", False),
)
SNAPSHOTS = sa.Table(
    "semantic_infl_snapshots", _md,
    sa.Column("account_id", sa.String(32), primary_key=True),
    sa.Column("day", sa.Date, primary_key=True),
    sa.Column("followers", sa.BigInteger),
    sa.Column("posts", sa.Integer),
    sa.Column("avg_likes", sa.Float),
    sa.Column("avg_comments", sa.Float),
    sa.Column("source", sa.String(8), nullable=False, default="elle"),     # api | elle
    sa.Column("by", sa.String(120)),
    _ts("at", False),
)
COLLABS = sa.Table(
    "semantic_infl_collabs", _md, _id(),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("no", sa.Integer, nullable=False),
    sa.Column("person_id", sa.String(32), nullable=False, index=True),
    sa.Column("crm_book_id", sa.String(40), index=True),
    sa.Column("book_code", sa.String(60)),
    sa.Column("book_title", sa.String(300)),
    sa.Column("kind", sa.String(12), nullable=False),
    sa.Column("stage", sa.String(20), nullable=False, default="teklif"),
    sa.Column("fee", sa.Numeric(14, 2), nullable=False, default=0),
    sa.Column("crm_order_no", sa.String(200)),
    sa.Column("due_publish", sa.Date),
    sa.Column("published_url", sa.String(1000)),
    sa.Column("published_at", sa.Date),
    sa.Column("disclosure_ok", sa.Boolean),
    sa.Column("reach", sa.BigInteger),
    sa.Column("engagement", sa.BigInteger),
    sa.Column("result_note", sa.Text),
    sa.Column("note", sa.Text),
    sa.Column("approved_by", sa.String(120)),
    _ts("approved_at"),
    _ts("stage_at"),
    sa.Column("created_day", sa.Date, nullable=False),
    sa.Column("created_by", sa.String(120), nullable=False),
    _ts("created_at", False),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at"),
)
EVENTS = sa.Table(
    "semantic_infl_events", _md,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("collab_id", sa.String(32), nullable=False, index=True),
    _ts("at", False),
    sa.Column("user", sa.String(120), nullable=False),
    sa.Column("action", sa.String(24), nullable=False),
    sa.Column("note", sa.Text),
)
PAYOUTS = sa.Table(
    "semantic_infl_payouts", _md, _id(),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("no", sa.Integer, nullable=False),
    sa.Column("collab_id", sa.String(32), nullable=False, index=True),
    sa.Column("amount", sa.Numeric(14, 2), nullable=False),
    sa.Column("status", sa.String(10), nullable=False, default="hazir"),
    sa.Column("logo_doc_no", sa.String(200)),
    sa.Column("note", sa.Text),
    sa.Column("created_by", sa.String(120), nullable=False),
    _ts("created_at", False),
    sa.Column("approved_by", sa.String(120)),
    _ts("approved_at"),
    sa.Column("paid_at", sa.Date),
    sa.Column("paid_by", sa.String(120)),
)
DRAFTS = sa.Table(
    "semantic_infl_drafts", _md, _id(),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("collab_id", sa.String(32), nullable=False, index=True),
    sa.Column("kind", sa.String(12), nullable=False),
    sa.Column("subject", sa.String(300)),
    sa.Column("body", sa.Text, nullable=False),
    sa.Column("source", sa.String(8), nullable=False),                     # zeki | kural | elle
    sa.Column("dropped_json", sa.Text, nullable=False, default="[]"),
    sa.Column("status", sa.String(10), nullable=False, default="taslak"),  # taslak | onayli
    sa.Column("created_by", sa.String(120), nullable=False),
    _ts("created_at", False),
    sa.Column("approved_by", sa.String(120)),
    _ts("approved_at"),
    sa.Column("sent_by", sa.String(120)),
    _ts("sent_at"),
)
REMINDERS = sa.Table(
    "semantic_infl_reminders", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("key", sa.String(160), primary_key=True),
    _ts("sent_at", False),
)

_ready: set[int] = set()
_lock = threading.Lock()


class InfluencerError(fl.FreelanceError):
    """Kişiye gösterilecek düz Türkçe hata (M8 yardımcılarının hatasıyla aynı aile)."""


def ensure(engine: sa.engine.Engine) -> None:
    with _lock:
        if id(engine) in _ready:
            return
        from semantic_layer.store import schema_stamp
        schema_stamp.create_all(_md, engine)
        _ready.add(id(engine))


# ------------------------------------------------------------------ ortak yardımcılar (M8'den)

_now, _today, _iso, _loads, _text, _tl = fl._now, fl._today, fl._iso, fl._loads, fl._text, fl._tl


def _new() -> str:
    return uuid.uuid4().hex


def _money(v: Any, what: str = "Tutar") -> Decimal:
    return fl._money(v, what)


def _day(v: Any, what: str) -> Optional[date]:
    return fl._day(v, what)


def _dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False)


def _count(v: Any, what: str) -> Optional[int]:
    if v in (None, ""):
        return None
    try:
        n = int(float(str(v).replace(".", "").replace(",", ".")) if isinstance(v, str) else float(v))
    except (TypeError, ValueError) as e:
        raise InfluencerError(f"{what} tam sayı olmalı.") from e
    if n < 0 or n > 10**12:
        raise InfluencerError(f"{what} 0 ile 1 trilyon arasında olmalı.")
    return n


_FOLD = str.maketrans("İIıŞşĞğÜüÖöÇçÂâÎîÛû", "iiissgguuooccaaiiuu")


def fold(s: Any) -> str:
    return " ".join(str(s or "").translate(_FOLD).lower().split())


def _keys(v: Any, allowed: dict[str, Any], what: str) -> list[str]:
    items = v if isinstance(v, list) else [x for x in str(v or "").replace(";", ",").split(",")]
    out: list[str] = []
    for x in items:
        k = str(x).strip()
        if not k:
            continue
        if k not in allowed:
            # Etiket adıyla da gelebilir (CSV): «Tarih ve biyografi» → tarih.
            by_label = {fold(lbl if isinstance(lbl, str) else lbl[0]): key for key, lbl in allowed.items()}
            k = by_label.get(fold(k), "")
            if not k:
                raise InfluencerError(f"{what} «{x}» tanımlı değil.")
        if k not in out:
            out.append(k)
    return out


_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_URL = re.compile(r"^https?://\S{3,}$", re.I)
_HANDLE = re.compile(r"^[0-9a-z._\-]{1,80}$")


def norm_handle(v: Any) -> Optional[str]:
    s = str(v or "").strip()
    s = re.sub(r"^https?://(www\.|m\.)?[^/]+/", "", s, flags=re.I).split("?")[0].strip("/")
    parts = [p for p in s.split("/") if p]
    if parts and parts[0] in ("c", "channel", "user") and len(parts) > 1:
        parts = parts[1:]
    s = (parts[0] if parts else "").lstrip("@").lower()
    return s if _HANDLE.match(s) else None


# ------------------------------------------------------------------ ayarlar


def settings() -> dict[str, Any]:
    from semantic_bridge.admin import conf

    def num(key: str, default: float) -> float:
        try:
            return float(str(conf(key) or default).replace(",", "."))
        except ValueError:
            return default

    weights = dict(DEFAULT_WEIGHTS)
    try:
        raw = json.loads(conf("INFLUENCER_RANK_WEIGHTS") or "{}")
        for k, v in (raw.items() if isinstance(raw, dict) else []):
            if k in weights and float(v) >= 0:
                weights[k] = float(v)
    except (ValueError, TypeError):
        log.warning("INFLUENCER_RANK_WEIGHTS okunamadı; varsayılan ağırlıklar")
    words = {k: tuple(v) for k, v in TOPIC_WORDS.items()}
    try:
        extra = json.loads(conf("INFLUENCER_TOPIC_WORDS") or "{}")
        for k, v in (extra.items() if isinstance(extra, dict) else []):
            if k in TOPICS and isinstance(v, list):
                words[k] = tuple(dict.fromkeys(words.get(k, ()) + tuple(fold(x) for x in v if str(x).strip())))
    except (ValueError, TypeError):
        log.warning("INFLUENCER_TOPIC_WORDS okunamadı")
    budget = num("INFLUENCER_MONTHLY_BUDGET", 0.0)
    kinds = [k.strip() for k in (conf("INFLUENCER_DISCLOSURE_KINDS") or "hediye,ucretli,karsilikli").split(",") if k.strip() in KINDS]
    return {
        "weights": weights, "topicWords": words,
        "cooldownDays": int(num("INFLUENCER_COOLDOWN_DAYS", 60)),
        "jumpPct": num("INFLUENCER_JUMP_PCT", 30.0),
        "contentWaitDays": int(num("INFLUENCER_CONTENT_WAIT_DAYS", 3)),
        "linkGraceDays": int(num("INFLUENCER_LINK_GRACE_DAYS", 2)),
        "payoutDay": max(1, min(28, int(num("INFLUENCER_PAYOUT_DAY", 25)))),
        "monthlyBudget": budget if budget > 0 else None,
        "budgetWarnPct": num("INFLUENCER_BUDGET_WARN_PCT", 90.0),
        "disclosureKinds": kinds,
        "giftNeedsApproval": (conf("INFLUENCER_GIFT_NEEDS_APPROVAL") or "1").strip() not in ("0", "false", "hayir", "hayır"),
        "apiEnabled": (conf("INFLUENCER_API_ENABLED") or "0").strip() in ("1", "true", "evet"),
        "explainProb": num("INFLUENCER_TOPIC_MIN_PROB", 0.70),
    }


def meta(cfg: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    cfg = cfg or settings()
    return {"platformlar": PLATFORMS, "konular": TOPICS, "yasGruplari": {k: v[0] for k, v in AGE_GROUPS.items()},
            "turler": KINDS, "asamalar": STAGES, "pano": list(BOARD), "odemeDurumlari": PAYOUT_STATES,
            "taslakTurleri": DRAFT_KINDS, "iletisimTercihleri": CONTACT_PREFS,
            "ayarlar": {k: cfg[k] for k in ("weights", "cooldownDays", "jumpPct", "contentWaitDays", "linkGraceDays",
                                              "payoutDay", "disclosureKinds", "giftNeedsApproval", "apiEnabled")},
            "yasalEtiket": DISCLOSURE_CLAUSE}


# ------------------------------------------------------------------ kişi


def _accounts_in(v: Any) -> list[dict[str, Any]]:
    out, seen = [], set()
    for a in v if isinstance(v, list) else []:
        platform = str((a or {}).get("platform") or "").strip()
        if platform not in PLATFORMS:
            raise InfluencerError(f"Platform «{platform}» tanımlı değil.")
        h = norm_handle(a.get("handle") or a.get("url"))
        if not h:
            raise InfluencerError(f"{PLATFORMS[platform]} kullanıcı adı geçerli değil.")
        url = _text(a.get("url"), 500)
        if url and not _URL.match(url):
            raise InfluencerError("Hesap bağlantısı http(s):// ile başlamalı.")
        if (platform, h) in seen:
            continue
        seen.add((platform, h))
        out.append({"platform": platform, "handle": h, "url": url})
    return out


def _person_values(body: dict[str, Any], partial: bool) -> dict[str, Any]:
    v: dict[str, Any] = {}
    if not partial or "name" in body:
        name = _text(body.get("name"), 200)
        if not name:
            raise InfluencerError("Ad yazılmalı.")
        v["name"] = name
    if "email" in body:
        email = _text(body.get("email"), 200)
        if email and not _EMAIL.match(email):
            raise InfluencerError("E-posta adresi geçerli değil.")
        v["email"] = email
    for k, lim in (("phone", 40), ("city", 80)):
        if k in body:
            v[k] = _text(body.get(k), lim)
    if "notes" in body:
        v["notes"] = _text(body.get("notes"), 4000)
    if "topics" in body:
        v["topics_json"] = _dump(_keys(body.get("topics"), TOPICS, "Konu"))
    if "ageGroups" in body:
        v["age_groups_json"] = _dump(_keys(body.get("ageGroups"), {k: x[0] for k, x in AGE_GROUPS.items()}, "Yaş grubu"))
    if "contactPref" in body:
        cp = str(body.get("contactPref") or "").strip() or None
        if cp and cp not in CONTACT_PREFS:
            raise InfluencerError("İletişim tercihi tanımlı değil.")
        v["contact_pref"] = cp
    for k, col in (("doNotContact", "do_not_contact"), ("minor", "minor")):
        if k in body:
            v[col] = bool(body.get(k))
    if "crmContactId" in body:
        cid = str(body.get("crmContactId") or "").strip().strip("{}").lower() or None
        if cid and not fl._CRM_ID.match(cid):
            raise InfluencerError("CRM kişi kimliği geçerli değil.")
        v["crm_contact_id"] = cid
    for k, col in (("feeMin", "fee_min"), ("feeMax", "fee_max")):
        if k in body:
            v[col] = None if body.get(k) in (None, "") else _money(body.get(k), "Ücret")
    return v


def _check_fee_range(lo: Any, hi: Any) -> None:
    if lo is not None and hi is not None and Decimal(lo) > Decimal(hi):
        raise InfluencerError("Ücret aralığında alt sınır üst sınırdan büyük olamaz.")


def _person_row(conn: sa.Connection, tenant: str, person_id: str) -> Any:
    row = conn.execute(sa.select(PEOPLE).where(PEOPLE.c.id == person_id, PEOPLE.c.tenant_id == tenant)).first()
    if row is None:
        raise InfluencerError("İçerik üreticisi bulunamadı.", 404)
    return row


def _set_accounts(conn: sa.Connection, tenant: str, person_id: str, accounts: list[dict[str, Any]]) -> None:
    """Hesapları istenen listeye getirir; aynı platform + kullanıcı adı kaydı (ve anlık görüntüleri) korunur."""
    for a in accounts:
        other = conn.execute(sa.select(ACCOUNTS.c.person_id, PEOPLE.c.name).join(PEOPLE, PEOPLE.c.id == ACCOUNTS.c.person_id)
                             .where(ACCOUNTS.c.tenant_id == tenant, ACCOUNTS.c.platform == a["platform"],
                                    ACCOUNTS.c.handle == a["handle"], ACCOUNTS.c.person_id != person_id)).first()
        if other is not None:
            raise InfluencerError(f"{PLATFORMS[a['platform']]} @{a['handle']} zaten «{other.name}» kaydında.", 409)
    have = {(r.platform, r.handle): r for r in conn.execute(sa.select(ACCOUNTS).where(ACCOUNTS.c.person_id == person_id))}
    want = {(a["platform"], a["handle"]): a for a in accounts}
    gone = [r.id for k, r in have.items() if k not in want]
    if gone:
        conn.execute(sa.delete(SNAPSHOTS).where(SNAPSHOTS.c.account_id.in_(gone)))
        conn.execute(sa.delete(ACCOUNTS).where(ACCOUNTS.c.id.in_(gone)))
    now = _now()
    for k, a in want.items():
        if k in have:
            if a["url"] != have[k].url:
                conn.execute(sa.update(ACCOUNTS).where(ACCOUNTS.c.id == have[k].id).values(url=a["url"]))
        else:
            conn.execute(sa.insert(ACCOUNTS).values(id=_new(), tenant_id=tenant, person_id=person_id, platform=a["platform"],
                                                    handle=a["handle"], url=a["url"], verified_by_api=False, created_at=now))


def create_person(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    v = _person_values(body, partial=False)
    _check_fee_range(v.get("fee_min"), v.get("fee_max"))
    accounts = _accounts_in(body.get("accounts"))
    pid = _new()
    with engine.begin() as conn:
        conn.execute(sa.insert(PEOPLE).values(id=pid, tenant_id=tenant, created_by=user, created_at=_now(),
                                              topics_json=v.pop("topics_json", "[]"),
                                              age_groups_json=v.pop("age_groups_json", "[]"), currency="TRY",
                                              do_not_contact=v.pop("do_not_contact", False), minor=v.pop("minor", False), **v))
        _set_accounts(conn, tenant, pid, accounts)
    return {"id": pid, "name": v["name"]}


def update_person(engine: sa.engine.Engine, tenant: str, user: str, person_id: str, body: dict[str, Any],
                  can_fee: bool) -> tuple[dict[str, Any], list[str]]:
    v = _person_values(body, partial=True)
    if ("fee_min" in v or "fee_max" in v) and not can_fee:
        raise InfluencerError("Ücret aralığını yalnız işbirliği onay ya da ödeme yetkisi olan değiştirir.", 403)
    with engine.begin() as conn:
        row = _person_row(conn, tenant, person_id)
        _check_fee_range(v.get("fee_min", row.fee_min), v.get("fee_max", row.fee_max))
        changed = sorted(k for k in v if getattr(row, k) != v[k])
        if v:
            conn.execute(sa.update(PEOPLE).where(PEOPLE.c.id == person_id).values(updated_by=user, updated_at=_now(), **v))
        if "accounts" in body:
            _set_accounts(conn, tenant, person_id, _accounts_in(body.get("accounts")))
            changed.append("accounts")
    return {"id": person_id, "name": v.get("name", row.name)}, changed


def _fee(v: Any, can_fee: bool) -> Optional[float]:
    return (float(v) if v is not None else None) if can_fee else None


def _person_out(p: Any, can_fee: bool) -> dict[str, Any]:
    return {"id": p.id, "name": p.name, "email": p.email, "phone": p.phone, "city": p.city,
            "crmContactId": p.crm_contact_id, "topics": _loads(p.topics_json, []), "ageGroups": _loads(p.age_groups_json, []),
            "feeMin": _fee(p.fee_min, can_fee), "feeMax": _fee(p.fee_max, can_fee),
            "feeSet": p.fee_min is not None or p.fee_max is not None, "currency": p.currency, "contactPref": p.contact_pref,
            "notes": p.notes, "doNotContact": bool(p.do_not_contact), "minor": bool(p.minor), "createdBy": p.created_by,
            "createdAt": _iso(p.created_at), "updatedBy": p.updated_by, "updatedAt": _iso(p.updated_at)}


def _latest(snaps: list[Any]) -> Optional[dict[str, Any]]:
    if not snaps:
        return None
    s = max(snaps, key=lambda x: x.day)
    return {"day": s.day.isoformat(), "followers": s.followers, "posts": s.posts, "avgLikes": s.avg_likes,
            "avgComments": s.avg_comments, "source": s.source}


def jumps(snaps: Iterable[Any], pct: float, window_days: int = 45) -> list[dict[str, Any]]:
    """Ardışık iki anlık görüntü arasında (en çok `window_days` gün) takipçi `pct` yüzdeden fazla arttıysa."""
    by_acc: dict[str, list[Any]] = {}
    for s in snaps:
        if s.followers is not None:
            by_acc.setdefault(s.account_id, []).append(s)
    out = []
    for acc, rows in by_acc.items():
        rows.sort(key=lambda x: x.day)
        for a, b in zip(rows, rows[1:]):
            if a.followers and (b.day - a.day).days <= window_days:
                rise = (b.followers - a.followers) / a.followers * 100
                if rise > pct:
                    out.append({"accountId": acc, "from": a.day.isoformat(), "to": b.day.isoformat(),
                                "before": a.followers, "after": b.followers, "pct": round(rise, 1)})
    return out


def relation(collabs: Iterable[Any], today: date) -> dict[str, Any]:
    """İlişki puanı (0–100, M7 ısı deseni): yayına ulaşmış son işbirliğinin yakınlığı en çok 40 (365 günde sıfırlanır),
    son 12 aydaki yayına ulaşmış işbirliği sayısı en çok 30 (işbirliği başı 10), sonuçlanma oranı en çok 30
    (yayına ulaşan ÷ yayına ulaşan + vazgeçilen). Hiç yayın yoksa 0."""
    done, dropped = [], 0
    for c in collabs:
        if c.stage == "vazgecildi":
            dropped += 1
        elif FLOW.index(c.stage) >= FLOW.index("yayinda"):
            done.append(c.published_at or c.due_publish or c.created_day)
    if not done:
        return {"score": 0, "band": "yok", "parts": {"yakinlik": 0, "siklik": 0, "sonuc": 0}, "last": None,
                "published": 0, "dropped": dropped}
    last = max(done)
    days = max(0, (today - last).days)
    recency = round(40 * max(0.0, 1 - days / 365))
    freq = min(30, 10 * sum(1 for d in done if (today - d).days <= 365))
    outcome = round(30 * len(done) / (len(done) + dropped))
    score = recency + freq + outcome
    return {"score": score, "band": "soguk" if score <= 33 else "ilik" if score <= 66 else "sicak",
            "parts": {"yakinlik": recency, "siklik": freq, "sonuc": outcome}, "last": last.isoformat(),
            "daysSince": days, "published": len(done), "dropped": dropped}


def people_stmts(tenant: str, ids: Optional[list[str]] = None) -> tuple[Any, Any, Any, Any]:
    """Kayıt defteri okumaları: (kişiler, hesapları, hesap ölçümleri, işbirlikleri); kişi süzgeci alt sorguyla."""
    who = sa.select(PEOPLE.c.id).where(PEOPLE.c.tenant_id == tenant)
    people = sa.select(PEOPLE).where(PEOPLE.c.tenant_id == tenant)
    if ids is not None:
        who = who.where(PEOPLE.c.id.in_(ids))
        people = people.where(PEOPLE.c.id.in_(ids))
    accs = sa.select(ACCOUNTS).where(ACCOUNTS.c.person_id.in_(who)).order_by(ACCOUNTS.c.platform)
    snaps = sa.select(SNAPSHOTS).where(SNAPSHOTS.c.account_id.in_(sa.select(ACCOUNTS.c.id).where(ACCOUNTS.c.person_id.in_(who))))
    cols = sa.select(COLLABS).where(COLLABS.c.tenant_id == tenant, COLLABS.c.person_id.in_(who))
    return people.order_by(PEOPLE.c.name), accs, snaps, cols


def _load_people(conn: sa.Connection, tenant: str, ids: Optional[list[str]] = None) -> tuple[list[Any], dict[str, list[Any]],
                                                                                              dict[str, list[Any]], dict[str, list[Any]]]:
    pq, aq, sq, cq = people_stmts(tenant, ids)
    people = conn.execute(pq).all()
    accs: dict[str, list[Any]] = {}
    snaps: dict[str, list[Any]] = {}
    cols: dict[str, list[Any]] = {}
    if people:
        owner = {}
        for a in conn.execute(aq).all():
            accs.setdefault(a.person_id, []).append(a)
            owner[a.id] = a.person_id
        if owner:
            for s in conn.execute(sq):
                if s.account_id in owner:
                    snaps.setdefault(owner[s.account_id], []).append(s)
        for c in conn.execute(cq):
            cols.setdefault(c.person_id, []).append(c)
    return people, accs, snaps, cols


def _account_out(a: Any, snaps: list[Any]) -> dict[str, Any]:
    mine = [s for s in snaps if s.account_id == a.id]
    return {"id": a.id, "platform": a.platform, "platformAdi": PLATFORMS.get(a.platform, a.platform), "handle": a.handle,
            "url": a.url, "verifiedByApi": bool(a.verified_by_api), "latest": _latest(mine)}


def list_people(engine: sa.engine.Engine, tenant: str, *, can_fee: bool, q: str = "", platform: str = "", topic: str = "",
                age: str = "", idle_days: Optional[int] = None, include_dnc: bool = True) -> dict[str, Any]:
    today = _today()
    jump_pct = settings()["jumpPct"]
    with engine.connect() as conn:
        people, accs, snaps, cols = _load_people(conn, tenant)
    fq = fold(q)
    items = []
    for p in people:
        a = accs.get(p.id, [])
        if fq and fq not in fold(p.name) and not any(fq in a_.handle for a_ in a) and fq not in fold(p.city):
            continue
        if platform and not any(x.platform == platform for x in a):
            continue
        topics = _loads(p.topics_json, [])
        ages = _loads(p.age_groups_json, [])
        if topic and topic not in topics:
            continue
        if age and age not in ages:
            continue
        if not include_dnc and p.do_not_contact:
            continue
        mine = cols.get(p.id, [])
        rel = relation(mine, today)
        last = max((c.created_day for c in mine), default=None)
        if idle_days is not None and last is not None and (today - last).days < idle_days:
            continue
        followers = sum((_latest([s for s in snaps.get(p.id, []) if s.account_id == x.id]) or {}).get("followers") or 0 for x in a)
        items.append({**_person_out(p, can_fee), "accounts": [_account_out(x, snaps.get(p.id, [])) for x in a],
                      "relation": rel, "lastCollab": last.isoformat() if last else None,
                      "openCollabs": sum(1 for c in mine if c.stage not in CLOSED), "collabs": len(mine),
                      "followers": followers or None,
                      "jumps": jumps(snaps.get(p.id, []), jump_pct)})
    return {"items": items, "total": len(items)}


def person_payouts_stmt(tenant: str, person_id: str):
    return sa.select(PAYOUTS).where(PAYOUTS.c.collab_id.in_(
        sa.select(COLLABS.c.id).where(COLLABS.c.tenant_id == tenant, COLLABS.c.person_id == person_id)))


def get_person(engine: sa.engine.Engine, tenant: str, person_id: str, can_fee: bool) -> dict[str, Any]:
    today = _today()
    cfg = settings()
    with engine.connect() as conn:
        _person_row(conn, tenant, person_id)
        people, accs, snaps, cols = _load_people(conn, tenant, [person_id])
        pays = conn.execute(person_payouts_stmt(tenant, person_id)).all()
    p = people[0]
    mine = sorted(cols.get(person_id, []), key=lambda c: c.created_at, reverse=True)
    paid = {x.collab_id: x for x in pays if x.status != "iptal"}
    history = [dict(_collab_out(c, p.name, can_fee, cfg), payout=_payout_brief(paid.get(c.id), can_fee)) for c in mine]
    s_all = snaps.get(person_id, [])
    snaps_out = [{"accountId": s.account_id, "day": s.day.isoformat(), "followers": s.followers, "posts": s.posts,
                  "avgLikes": s.avg_likes, "avgComments": s.avg_comments, "source": s.source, "by": s.by}
                 for s in sorted(s_all, key=lambda x: (x.account_id, x.day))]
    fees = [float(c.fee) for c in mine if c.stage != "vazgecildi"]
    eng = sum(c.engagement or 0 for c in mine if c.stage != "vazgecildi")
    orders = sorted({n.strip() for c in mine for n in (c.crm_order_no or "").replace(";", ",").split(",") if n.strip()})
    return {**_person_out(p, can_fee), "accounts": [_account_out(a, s_all) for a in accs.get(person_id, [])],
            "snapshots": snaps_out, "jumps": jumps(s_all, cfg["jumpPct"]), "relation": relation(mine, today),
            "collabs": history, "books": sorted({c.book_title for c in mine if c.book_title and c.stage != "vazgecildi"}),
            "orderNos": orders,
            "totals": {"collabs": len(mine), "published": sum(1 for c in mine if c.stage in FLOW and FLOW.index(c.stage) >= 4),
                       "engagement": eng, "spend": round(sum(fees), 2) if can_fee else None,
                       "cpe": (round(sum(fees) / eng, 4) if eng and can_fee else None)}}


def add_snapshot(engine: sa.engine.Engine, tenant: str, user: str, account_id: str, body: dict[str, Any]) -> dict[str, Any]:
    day = _day(body.get("day"), "Ölçüm") or _today()
    if day > _today():
        raise InfluencerError("Ölçüm tarihi ileri bir gün olamaz.")
    vals = {"followers": _count(body.get("followers"), "Takipçi"), "posts": _count(body.get("posts"), "Gönderi"),
            "avg_likes": _count(body.get("avgLikes"), "Ortalama beğeni"),
            "avg_comments": _count(body.get("avgComments"), "Ortalama yorum")}
    if all(v is None for v in vals.values()):
        raise InfluencerError("En az bir sayı girilmeli.")
    with engine.begin() as conn:
        acc = conn.execute(sa.select(ACCOUNTS).where(ACCOUNTS.c.id == account_id, ACCOUNTS.c.tenant_id == tenant)).first()
        if acc is None:
            raise InfluencerError("Hesap bulunamadı.", 404)
        conn.execute(sa.delete(SNAPSHOTS).where(SNAPSHOTS.c.account_id == account_id, SNAPSHOTS.c.day == day))
        conn.execute(sa.insert(SNAPSHOTS).values(account_id=account_id, day=day, source="elle", by=user, at=_now(), **vals))
    return {"accountId": account_id, "personId": acc.person_id, "day": day.isoformat(), **vals}


# ------------------------------------------------------------------ CSV içe aktarma

_CSV_COLS = {
    "name": ("ad", "ad soyad", "isim", "name", "icerik uretici", "kisi"),
    "email": ("e-posta", "eposta", "email", "mail"),
    "phone": ("telefon", "tel", "phone"),
    "city": ("sehir", "il", "city"),
    "platform": ("platform", "mecra"),
    "handle": ("kullanici adi", "hesap", "handle", "kullanici"),
    "url": ("baglanti", "url", "link", "adres"),
    "topics": ("konu", "konular", "topic"),
    "ageGroups": ("yas grubu", "yas gruplari", "yas"),
    "followers": ("takipci", "followers", "abone"),
    "notes": ("not", "notlar", "aciklama"),
}


def import_csv(engine: sa.engine.Engine, tenant: str, user: str, text: str) -> dict[str, Any]:
    """Satır başına bir hesap. Aynı platform + kullanıcı adı varsa kişi yeniden açılmaz; takipçi verildiyse bugünün
    anlık görüntüsü yazılır. Okunamayan satır nedeniyle listelenir (tavan yok)."""
    text = (text or "").lstrip("﻿")
    if not text.strip():
        raise InfluencerError("Dosya boş.")
    try:
        delimiter = csv.Sniffer().sniff(text[:4000], delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";"
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if len(rows) < 2:
        raise InfluencerError("Başlık satırı ve en az bir kayıt olmalı.")
    header = [fold(h) for h in rows[0]]
    index: dict[str, int] = {}
    for key, names in _CSV_COLS.items():
        for i, h in enumerate(header):
            if h in names and key not in index:
                index[key] = i
    if "name" not in index and "handle" not in index:
        raise InfluencerError("Başlıkta «Ad» ya da «Kullanıcı adı» kolonu bulunmalı.")
    added = updated = snaps = 0
    errors: list[dict[str, Any]] = []
    for n, r in enumerate(rows[1:], start=2):
        if not any(c.strip() for c in r):
            continue
        get = lambda k: (r[index[k]].strip() if k in index and index[k] < len(r) else "")  # noqa: E731
        try:
            platform_raw = get("platform")
            platform = (_keys([platform_raw], PLATFORMS, "Platform") or [""])[0] if platform_raw else ""
            h = norm_handle(get("handle") or get("url")) if (get("handle") or get("url")) else None
            if (get("handle") or get("url")) and not h:
                raise InfluencerError("kullanıcı adı okunamadı")
            if h and not platform:
                raise InfluencerError("platform yazılmamış")
            with engine.begin() as conn:
                acc = None
                if h:
                    acc = conn.execute(sa.select(ACCOUNTS).where(ACCOUNTS.c.tenant_id == tenant, ACCOUNTS.c.platform == platform,
                                                                 ACCOUNTS.c.handle == h)).first()
                if acc is None:
                    body = {"name": get("name") or (f"@{h}" if h else ""), "email": get("email"), "phone": get("phone"),
                            "city": get("city"), "notes": get("notes"), "topics": get("topics"), "ageGroups": get("ageGroups")}
                    v = _person_values(body, partial=False)
                    pid = _new()
                    conn.execute(sa.insert(PEOPLE).values(id=pid, tenant_id=tenant, created_by=user, created_at=_now(),
                                                          currency="TRY", do_not_contact=False, minor=False,
                                                          topics_json=v.pop("topics_json", "[]"),
                                                          age_groups_json=v.pop("age_groups_json", "[]"), **v))
                    if h:
                        aid = _new()
                        conn.execute(sa.insert(ACCOUNTS).values(id=aid, tenant_id=tenant, person_id=pid, platform=platform,
                                                                handle=h, url=_text(get("url"), 500) if _URL.match(get("url")) else None,
                                                                verified_by_api=False, created_at=_now()))
                        acc = conn.execute(sa.select(ACCOUNTS).where(ACCOUNTS.c.id == aid)).first()
                    added += 1
                else:
                    updated += 1
                followers = _count(get("followers"), "Takipçi") if get("followers") else None
                if followers is not None and acc is not None:
                    conn.execute(sa.delete(SNAPSHOTS).where(SNAPSHOTS.c.account_id == acc.id, SNAPSHOTS.c.day == _today()))
                    conn.execute(sa.insert(SNAPSHOTS).values(account_id=acc.id, day=_today(), followers=followers,
                                                             source="elle", by=user, at=_now()))
                    snaps += 1
        except fl.FreelanceError as e:
            errors.append({"satir": n, "neden": str(e)})
    return {"eklenen": added, "mevcut": updated, "olcum": snaps, "okunamayan": errors, "satir": len(rows) - 1}


# ------------------------------------------------------------------ işbirliği


def collab_stmts(tenant: str, collab_id: str) -> tuple[Any, Any, Any, Any]:
    """İşbirliği kartı: (işbirliği, olaylar, taslaklar, ödeme satırı)."""
    return (sa.select(COLLABS).where(COLLABS.c.id == collab_id, COLLABS.c.tenant_id == tenant),
            sa.select(EVENTS).where(EVENTS.c.collab_id == collab_id).order_by(EVENTS.c.at, EVENTS.c.id),
            sa.select(DRAFTS).where(DRAFTS.c.collab_id == collab_id).order_by(DRAFTS.c.created_at.desc()),
            sa.select(PAYOUTS).where(PAYOUTS.c.collab_id == collab_id, PAYOUTS.c.status != "iptal"))


def _collab_row(conn: sa.Connection, tenant: str, collab_id: str) -> Any:
    row = conn.execute(collab_stmts(tenant, collab_id)[0]).first()
    if row is None:
        raise InfluencerError("İşbirliği bulunamadı.", 404)
    return row


def needs_approval(kind: str, cfg: dict[str, Any]) -> bool:
    return kind != "hediye" or cfg["giftNeedsApproval"]


def period_day(c: Any) -> date:
    """Raporun dönem günü: yayın günü, yoksa planlanan yayın günü, yoksa kayıt günü."""
    return c.published_at or c.due_publish or c.created_day


def _collab_out(c: Any, person_name: Optional[str], can_fee: bool, cfg: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    cfg = cfg or settings()
    today = _today()
    late = bool(c.due_publish and not c.published_url and c.stage not in CLOSED
                and (today - c.due_publish).days >= cfg["linkGraceDays"])
    return {"id": c.id, "no": c.no, "code": f"İŞB-{c.no}", "personId": c.person_id, "personName": person_name,
            "crmBookId": c.crm_book_id, "bookCode": c.book_code, "bookTitle": c.book_title, "kind": c.kind,
            "kindLabel": KINDS.get(c.kind, c.kind), "stage": c.stage, "stageLabel": STAGES.get(c.stage, c.stage),
            "fee": _fee(c.fee, can_fee), "feeSet": bool(c.fee and Decimal(c.fee) > 0), "crmOrderNo": c.crm_order_no,
            "duePublish": _iso(c.due_publish), "publishedUrl": c.published_url, "publishedAt": _iso(c.published_at),
            "disclosureOk": c.disclosure_ok, "reach": c.reach, "engagement": c.engagement, "resultNote": c.result_note,
            "note": c.note, "needsApproval": needs_approval(c.kind, cfg), "approvedBy": c.approved_by,
            "approvedAt": _iso(c.approved_at), "waitingApproval": c.stage == "teklif" and needs_approval(c.kind, cfg) and not c.approved_by,
            "stageAt": _iso(c.stage_at), "createdBy": c.created_by, "createdAt": _iso(c.created_at),
            "daysToPublish": (c.due_publish - today).days if c.due_publish else None, "linkLate": late,
            "cpe": (round(float(c.fee) / c.engagement, 4) if can_fee and c.engagement and c.fee else None)}


def _event(conn: sa.Connection, collab_id: str, user: str, action: str, note: Optional[str] = None) -> None:
    conn.execute(sa.insert(EVENTS).values(collab_id=collab_id, at=_now(), user=user, action=action, note=note))


def create_collab(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    cfg = settings()
    person_id = str(body.get("personId") or "")
    kind = str(body.get("kind") or "hediye")
    if kind not in KINDS:
        raise InfluencerError("İşbirliği türü tanımlı değil.")
    fee = _money(body.get("fee"), "Ücret")
    if kind == "ucretli" and fee <= 0:
        raise InfluencerError("Ücretli işbirliğinde ücret girilmeli.")
    book_id = str(body.get("crmBookId") or "").strip().strip("{}").lower() or None
    if book_id and not fl._CRM_ID.match(book_id):
        raise InfluencerError("CRM kitap kimliği geçerli değil.")
    title = _text(body.get("bookTitle"), 300)
    if not book_id and not title:
        raise InfluencerError("Kitap seçilmeli ya da adı yazılmalı.")
    due = _day(body.get("duePublish"), "Yayın")
    now, today = _now(), _today()
    with engine.begin() as conn:
        person = _person_row(conn, tenant, person_id)
        if person.do_not_contact:
            raise InfluencerError(f"«{person.name}» için «iletişim kurulmasın» işaretli.", 409)
        if book_id and not body.get("repeat"):
            dup = conn.execute(sa.select(COLLABS.c.no, COLLABS.c.stage).where(
                COLLABS.c.tenant_id == tenant, COLLABS.c.person_id == person_id, COLLABS.c.crm_book_id == book_id,
                COLLABS.c.stage != "vazgecildi")).first()
            if dup is not None:
                raise InfluencerError(f"Bu kitap bu kişiyle İŞB-{dup.no} işbirliğinde zaten var ({STAGES[dup.stage]}). "
                                      "Yine de açmak için «tekrar» işaretleyin.", 409)
        no = (conn.execute(sa.select(sa.func.max(COLLABS.c.no)).where(COLLABS.c.tenant_id == tenant)).scalar() or 0) + 1
        cid = _new()
        conn.execute(sa.insert(COLLABS).values(
            id=cid, tenant_id=tenant, no=no, person_id=person_id, crm_book_id=book_id, book_code=_text(body.get("bookCode"), 60),
            book_title=title, kind=kind, stage="teklif", fee=fee, crm_order_no=_text(body.get("crmOrderNo"), 200),
            due_publish=due, note=_text(body.get("note"), 4000), stage_at=now, created_day=today, created_by=user,
            created_at=now))
        _event(conn, cid, user, "olusturuldu", f"{KINDS[kind]} · {title or book_id}")
    return {"id": cid, "no": no, "personName": person.name, "bookTitle": title,
            "waitingApproval": needs_approval(kind, cfg)}


def _gate(c: dict[str, Any], target: str, cfg: dict[str, Any]) -> None:
    """Aşamaya girişin koşulları (birleşmiş değerlerle)."""
    ti = FLOW.index(target)
    if ti > 0 and needs_approval(c["kind"], cfg) and not c["approved_by"]:
        raise InfluencerError("Teklif onaylanmadan işbirliği ilerlemez.", 409)
    if ti >= FLOW.index("yayinda") and not c["published_url"]:
        raise InfluencerError("«Yayında» için paylaşım bağlantısı girilmeli.")
    if ti >= FLOW.index("rapor"):
        if c["kind"] in cfg["disclosureKinds"] and c["disclosure_ok"] is not True:
            raise InfluencerError("Paylaşımda işbirliği/reklam etiketi yoksa rapora geçilmez; içerik üreticisinden "
                                  "düzeltme isteyin ve «yasal etiket var» işaretleyin.", 409)
        if c["reach"] is None and c["engagement"] is None:
            raise InfluencerError("Rapor için erişim ya da etkileşim sayısı girilmeli.")
    if target == "odeme" and not (c["fee"] and Decimal(c["fee"]) > 0):
        raise InfluencerError("Ücreti olmayan işbirliği ödeme aşamasına girmez; rapordan kapatılır.", 409)


def update_collab(engine: sa.engine.Engine, tenant: str, user: str, collab_id: str, body: dict[str, Any],
                  can_fee: bool) -> tuple[dict[str, Any], dict[str, Any]]:
    cfg = settings()
    now = _now()
    with engine.begin() as conn:
        row = _collab_row(conn, tenant, collab_id)
        cur = dict(row._mapping)
        v: dict[str, Any] = {}
        if "fee" in body:
            fee = _money(body.get("fee"), "Ücret")
            if Decimal(fee) != Decimal(row.fee or 0):
                if row.approved_by and not can_fee:
                    raise InfluencerError("Onaylanmış işbirliğinin ücretini yalnız onay yetkisi olan değiştirir.", 403)
                if row.stage in ("odeme", "kapali"):
                    raise InfluencerError("Ödeme aşamasındaki işin ücreti değişmez; ödeme satırını iptal edin.", 409)
                v["fee"] = fee
        if "kind" in body:
            k = str(body.get("kind") or "")
            if k not in KINDS:
                raise InfluencerError("İşbirliği türü tanımlı değil.")
            if k != row.kind and row.stage != "teklif":
                raise InfluencerError("Tür yalnız teklif aşamasında değişir.", 409)
            v["kind"] = k
        for key, col, lim in (("crmOrderNo", "crm_order_no", 200), ("resultNote", "result_note", 4000), ("note", "note", 4000),
                              ("bookTitle", "book_title", 300)):
            if key in body:
                v[col] = _text(body.get(key), lim)
        if "publishedUrl" in body:
            url = _text(body.get("publishedUrl"), 1000)
            if url and not _URL.match(url):
                raise InfluencerError("Paylaşım bağlantısı http(s):// ile başlamalı.")
            v["published_url"] = url
            if url and not row.published_at and "publishedAt" not in body:
                v["published_at"] = _today()
        if "publishedAt" in body:
            d = _day(body.get("publishedAt"), "Yayın")
            if d and d > _today():
                raise InfluencerError("Yayın tarihi ileri bir gün olamaz.")
            v["published_at"] = d
        if "duePublish" in body:
            v["due_publish"] = _day(body.get("duePublish"), "Planlanan yayın")
        if "disclosureOk" in body:
            v["disclosure_ok"] = None if body.get("disclosureOk") is None else bool(body.get("disclosureOk"))
        for key, col, what in (("reach", "reach", "Erişim"), ("engagement", "engagement", "Etkileşim")):
            if key in body:
                v[col] = _count(body.get(key), what)
        merged = {**cur, **v}
        if merged["kind"] == "ucretli" and not (merged["fee"] and Decimal(merged["fee"]) > 0):
            raise InfluencerError("Ücretli işbirliğinde ücret girilmeli.")
        target = body.get("stage")
        payout_note = None
        if target and target != row.stage:
            if target not in STAGES:
                raise InfluencerError("Aşama tanımlı değil.")
            if row.stage in CLOSED:
                raise InfluencerError("Kapanmış işbirliği yeniden açılmaz; yeni işbirliği açın.", 409)
            open_pay = conn.execute(sa.select(PAYOUTS).where(PAYOUTS.c.collab_id == collab_id, PAYOUTS.c.status != "iptal")).first()
            if target == "vazgecildi":
                if open_pay is not None:
                    raise InfluencerError("Ödeme satırı olan iş vazgeçilemez; önce ödeme satırını iptal edin.", 409)
                if not _text(body.get("reason") or body.get("note"), 2000):
                    raise InfluencerError("Vazgeçme nedeni yazılmalı.")
            elif target == "kapali":
                if merged["fee"] and Decimal(merged["fee"]) > 0:
                    raise InfluencerError("Ücretli iş, ödemesi yapılınca kendiliğinden kapanır.", 409)
                if row.stage != "rapor":
                    raise InfluencerError("Ücretsiz iş rapordan kapatılır.", 409)
            else:
                if open_pay is not None and FLOW.index(target) < FLOW.index("odeme"):
                    raise InfluencerError("Ödeme satırı olan iş geri alınmaz; önce ödeme satırını iptal edin.", 409)
                _gate(merged, target, cfg)
            v.update(stage=target, stage_at=now)
            if target == "odeme":
                no = (conn.execute(sa.select(sa.func.max(PAYOUTS.c.no)).where(PAYOUTS.c.tenant_id == tenant)).scalar() or 0) + 1
                conn.execute(sa.insert(PAYOUTS).values(id=_new(), tenant_id=tenant, no=no, collab_id=collab_id,
                                                       amount=Decimal(merged["fee"]), status="hazir", created_by=user,
                                                       created_at=now))
                payout_note = f"Ödeme satırı #{no} hazır"
        if not v:
            return _collab_out(row, None, can_fee, cfg), {}
        conn.execute(sa.update(COLLABS).where(COLLABS.c.id == collab_id).values(updated_by=user, updated_at=now, **v))
        diff = {k: (str(cur.get(k)) if cur.get(k) is not None else None, str(val) if val is not None else None)
                for k, val in v.items() if k not in ("stage_at",) and cur.get(k) != val}
        if "fee" in diff and not can_fee:
            diff["fee"] = ("gizli", "gizli")
        if "stage" in v:
            _event(conn, collab_id, user, "asama", f"{STAGES[row.stage]} → {STAGES[v['stage']]}"
                   + (f": {_text(body.get('reason') or body.get('note'), 2000)}" if v["stage"] == "vazgecildi" else ""))
        rest = sorted(k for k in diff if k not in ("stage", "approved_by", "approved_at"))
        if rest:
            _event(conn, collab_id, user, "guncellendi", ", ".join(rest))
        if payout_note:
            _event(conn, collab_id, user, "odeme", payout_note)
        new = _collab_row(conn, tenant, collab_id)
        person = _person_row(conn, tenant, new.person_id)
    return _collab_out(new, person.name, can_fee, cfg), diff


def approve_collab(engine: sa.engine.Engine, tenant: str, user: str, collab_id: str, body: dict[str, Any]) -> dict[str, Any]:
    """Seçim/teklif onayı (yetki ucun içinde denetlenir). Teklifi açan onaylayamaz. Ret işi vazgeçer, neden şart."""
    decision = str(body.get("decision") or "onay")
    note = _text(body.get("note"), 2000)
    now = _now()
    with engine.begin() as conn:
        row = _collab_row(conn, tenant, collab_id)
        if row.stage != "teklif":
            raise InfluencerError("Yalnız teklif aşamasındaki işbirliği onaylanır.", 409)
        if (row.created_by or "").lower() == user.lower():
            raise InfluencerError("Açtığınız teklifi başka biri onaylamalı.", 409)
        if decision == "onay":
            if row.approved_by:
                raise InfluencerError("Teklif zaten onaylı.", 409)
            conn.execute(sa.update(COLLABS).where(COLLABS.c.id == collab_id).values(approved_by=user, approved_at=now))
            _event(conn, collab_id, user, "onay", note)
        elif decision == "ret":
            if not note:
                raise InfluencerError("Ret nedeni yazılmalı.")
            conn.execute(sa.update(COLLABS).where(COLLABS.c.id == collab_id).values(
                stage="vazgecildi", stage_at=now, approved_by=None, approved_at=None, updated_by=user, updated_at=now))
            _event(conn, collab_id, user, "ret", note)
        else:
            raise InfluencerError("Karar «onay» ya da «ret» olmalı.")
    return {"id": collab_id, "no": row.no, "decision": decision}


def get_collab(engine: sa.engine.Engine, tenant: str, collab_id: str, can_fee: bool) -> dict[str, Any]:
    with engine.connect() as conn:
        c = _collab_row(conn, tenant, collab_id)
        person = _person_row(conn, tenant, c.person_id)
        _cq, eq, dq, pq = collab_stmts(tenant, collab_id)
        events = conn.execute(eq).all()
        drafts = conn.execute(dq).all()
        pay = conn.execute(pq).first()
    return {**_collab_out(c, person.name, can_fee), "person": {"id": person.id, "name": person.name, "email": person.email,
                                                                "doNotContact": bool(person.do_not_contact), "minor": bool(person.minor)},
            "events": [{"at": _iso(e.at), "user": e.user, "action": e.action, "note": e.note} for e in events],
            "drafts": [_draft_out(d) for d in drafts], "payout": _payout_brief(pay, can_fee)}


def book_collabs_stmt(tenant: str, bid: str):
    return sa.select(COLLABS, PEOPLE.c.name.label("pname")).join(PEOPLE, PEOPLE.c.id == COLLABS.c.person_id) \
        .where(COLLABS.c.tenant_id == tenant, sa.func.lower(COLLABS.c.crm_book_id) == bid).order_by(COLLABS.c.no.desc())


def book_collabs(engine: sa.engine.Engine, tenant: str, crm_book_id: str, can_fee: bool) -> dict[str, Any]:
    """Kitap sayfasındaki «İşbirlikleri» bölümü: bu kitabın bütün işbirlikleri (vazgeçilenler dahil)."""
    bid = str(crm_book_id or "").strip().strip("{}").lower()
    if not fl._CRM_ID.match(bid):
        raise InfluencerError("CRM kitap kimliği geçerli değil.")
    cfg = settings()
    with engine.connect() as conn:
        rows = conn.execute(book_collabs_stmt(tenant, bid)).all()
    items = [_collab_out(c, c.pname, can_fee, cfg) for c in rows]
    live = [c for c in rows if c.stage != "vazgecildi"]
    return {"items": items, "published": sum(1 for c in live if FLOW.index(c.stage) >= FLOW.index("yayinda")),
            "engagement": sum(c.engagement or 0 for c in live), "reach": sum(c.reach or 0 for c in live)}


# ------------------------------------------------------------------ pano ve hatırlatmalar


def _month_bounds(d: date) -> tuple[date, date]:
    start = d.replace(day=1)
    nxt = (start + timedelta(days=32)).replace(day=1)
    return start, nxt - timedelta(days=1)


def reminders_stmts(tenant: str, today: date) -> tuple[Any, Any, Any, Any]:
    """Hatırlatma okumaları: açık işbirlikleri, açık ödeme sayısı, onaylı harcama satırları, son 60 gün ölçümleri."""
    return (sa.select(COLLABS, PEOPLE.c.name.label("pname")).join(PEOPLE, PEOPLE.c.id == COLLABS.c.person_id)
            .where(COLLABS.c.tenant_id == tenant, COLLABS.c.stage.notin_(CLOSED)),
            sa.select(sa.func.count()).select_from(PAYOUTS).where(PAYOUTS.c.tenant_id == tenant,
                                                                  PAYOUTS.c.status.in_(("hazir", "onayli"))),
            sa.select(COLLABS).where(COLLABS.c.tenant_id == tenant, COLLABS.c.stage != "vazgecildi", COLLABS.c.approved_by.isnot(None)),
            sa.select(SNAPSHOTS, ACCOUNTS.c.person_id, ACCOUNTS.c.handle, ACCOUNTS.c.platform)
            .join(ACCOUNTS, ACCOUNTS.c.id == SNAPSHOTS.c.account_id)
            .where(ACCOUNTS.c.tenant_id == tenant, SNAPSHOTS.c.day >= today - timedelta(days=60)))


def reminders(engine: sa.engine.Engine, tenant: str, cfg: dict[str, Any], today: Optional[date] = None) -> list[dict[str, Any]]:
    """Bugünün hatırlatmaları (anahtar bir kez gönderilir). Sorumlu = işbirliğini açan kişi."""
    today = today or _today()
    out: list[dict[str, Any]] = []
    rq, oq, sq, nq = reminders_stmts(tenant, today)
    with engine.connect() as conn:
        rows = conn.execute(rq).all()
        open_pays = conn.execute(oq).scalar() or 0
        m0, m1 = _month_bounds(today)
        spend_rows = conn.execute(sq).all()
        snaps = conn.execute(nq).all()
    for c in rows:
        who, label = c.created_by, f"İŞB-{c.no} {c.pname} · {c.book_title or 'kitap'}"
        idx = FLOW.index(c.stage)
        if c.due_publish and idx < FLOW.index("yayinda") and 0 <= (c.due_publish - today).days <= 1:
            when = "bugün" if c.due_publish == today else "yarın"
            out.append({"key": f"yayin:{c.id}:{c.due_publish}", "kind": "yayin-yaklasti", "collabId": c.id, "to": who,
                        "text": f"{label}: yayın tarihi {when}."})
        if c.due_publish and not c.published_url and (today - c.due_publish).days >= cfg["linkGraceDays"]:
            out.append({"key": f"baglanti:{c.id}:{c.due_publish}", "kind": "baglanti-yok", "collabId": c.id, "to": who,
                        "text": f"{label}: yayın tarihi {(today - c.due_publish).days} gün önce geçti, paylaşım bağlantısı girilmedi."})
        if c.stage == "onayda" and c.stage_at is not None:
            since = (today - (c.stage_at + timedelta(hours=3)).date()).days
            if since >= cfg["contentWaitDays"]:
                out.append({"key": f"onayda:{c.id}:{c.stage_at.date()}", "kind": "icerik-onay-bekliyor", "collabId": c.id,
                            "to": who, "text": f"{label}: içerik taslağı {since} gündür onay bekliyor."})
        if c.stage == "teklif" and needs_approval(c.kind, cfg) and not c.approved_by:
            out.append({"key": f"teklif:{c.id}", "kind": "teklif-onay-bekliyor", "collabId": c.id, "to": "mudur",
                        "text": f"{label}: teklif onayınızı bekliyor ({KINDS[c.kind]})."})
    if cfg["monthlyBudget"]:
        spent = sum(float(c.fee) for c in spend_rows if m0 <= period_day(c) <= m1)
        if spent >= cfg["monthlyBudget"] * cfg["budgetWarnPct"] / 100:
            out.append({"key": f"butce:{m0:%Y-%m}", "kind": "butce", "collabId": None, "to": "mudur",
                        "text": f"Bu ay işbirliği harcaması aylık bütçenin %{round(spent / cfg['monthlyBudget'] * 100)}'ine ulaştı."})
    if open_pays and today.day >= cfg["payoutDay"]:
        out.append({"key": f"odeme:{today:%Y-%m}", "kind": "odeme-listesi", "collabId": None, "to": "muhasebe",
                    "text": f"İşbirliği ödeme listesi hazır: {open_pays} satır onay ya da ödeme bekliyor."})
    by_acc: dict[str, Any] = {}
    for s in snaps:
        by_acc.setdefault(s.account_id, s)
    for j in jumps(snaps, cfg["jumpPct"]):
        s = by_acc[j["accountId"]]
        out.append({"key": f"sicrama:{j['accountId']}:{j['to']}", "kind": "takipci-sicramasi", "collabId": None,
                    "personId": s.person_id, "to": "sorumlu",
                    "text": f"{PLATFORMS.get(s.platform, s.platform)} @{s.handle}: takipçi {j['from']} → {j['to']} arasında %{j['pct']} arttı (kendi ölçümlerimiz)."})
    return out


def unsent(engine: sa.engine.Engine, tenant: str, items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    keys = [i["key"] for i in items]
    if not keys:
        return []
    with engine.connect() as conn:
        sent = {r.key for r in conn.execute(sa.select(REMINDERS.c.key).where(REMINDERS.c.tenant_id == tenant, REMINDERS.c.key.in_(keys)))}
    return [i for i in items if i["key"] not in sent]


def mark_sent(engine: sa.engine.Engine, tenant: str, keys: list[str]) -> None:
    if not keys:
        return
    with engine.begin() as conn:
        conn.execute(sa.insert(REMINDERS), [{"tenant_id": tenant, "key": k, "sent_at": _now()} for k in keys])


def reminder_text(items: list[dict[str, Any]], link: str) -> str:
    lines = ["İşbirlikleri — bugünün hatırlatmaları", ""] + [f"- {i['text']}" for i in items]
    if link:
        lines += ["", f"Pano: {link}"]
    return "\n".join(lines) + "\n"


def board_stmt(tenant: str):
    return sa.select(COLLABS, PEOPLE.c.name.label("pname")).join(PEOPLE, PEOPLE.c.id == COLLABS.c.person_id) \
        .where(COLLABS.c.tenant_id == tenant).order_by(COLLABS.c.due_publish, COLLABS.c.no)


def board(engine: sa.engine.Engine, tenant: str, user: str, can_fee: bool, *, mine: bool = False) -> dict[str, Any]:
    cfg = settings()
    today = _today()
    m0, m1 = _month_bounds(today)
    with engine.connect() as conn:
        rows = conn.execute(board_stmt(tenant)).all()
    cards = [_collab_out(c, c.pname, can_fee, cfg) for c in rows
             if c.stage not in CLOSED and (not mine or (c.created_by or "").lower() == user.lower())]
    columns = [{"stage": s, "label": STAGES[s], "items": [x for x in cards if x["stage"] == s]} for s in BOARD]
    month = [c for c in rows if c.stage != "vazgecildi" and m0 <= period_day(c) <= m1]
    spent = sum(float(c.fee) for c in month if c.approved_by or not needs_approval(c.kind, cfg))
    rem = reminders(engine, tenant, cfg, today)
    closed_recent = [_collab_out(c, c.pname, can_fee, cfg) for c in rows
                     if c.stage in CLOSED and c.stage_at and (today - (c.stage_at + timedelta(hours=3)).date()).days <= 30]
    return {"columns": columns, "open": len(cards), "closedRecent": closed_recent,
            "waitingApproval": sum(1 for x in cards if x["waitingApproval"]),
            "linkLate": sum(1 for x in cards if x["linkLate"]),
            "publishSoon": sum(1 for x in cards if x["daysToPublish"] is not None and 0 <= x["daysToPublish"] <= 7
                               and FLOW.index(x["stage"]) < FLOW.index("yayinda")),
            "month": {"from": m0.isoformat(), "to": m1.isoformat(), "collabs": len(month),
                      "spend": round(spent, 2) if can_fee else None,
                      "budget": cfg["monthlyBudget"] if can_fee else None},
            "reminders": rem}


# ------------------------------------------------------------------ aday sırası


def _age_groups_of(lo: Optional[int], hi: Optional[int]) -> list[str]:
    if lo is None and hi is None:
        return []
    lo = 0 if lo is None else lo
    hi = 120 if hi is None or hi <= 0 else hi
    return [k for k, (_, a, b) in AGE_GROUPS.items() if a <= hi and b >= lo]


def book_profile(book: dict[str, Any], cfg: dict[str, Any]) -> dict[str, Any]:
    """Kitabın konu etiketleri (CRM tür metni + raf türü + hedef kitle, kelime kuralıyla) ve yaş grupları."""
    hay = " | ".join(fold(book.get(k)) for k in ("turler", "raf", "hedefKitle") if book.get(k))
    topics = []
    for t, words in cfg["topicWords"].items():
        # Kısa kelime (≤3 harf) tam kelime olarak aranır («din» «dinozor»da eşleşmesin); uzunlar ek alabilir.
        if any(re.search(r"(?<![a-z0-9])" + re.escape(w) + (r"(?![a-z0-9])" if len(w) <= 3 else ""), hay) for w in words):
            topics.append(t)
    yas = book.get("yas") or [None, None]
    return {"topics": topics, "topicSource": "kural" if topics else None, "ageGroups": _age_groups_of(yas[0], yas[1]),
            "evidence": hay or None}


def rank(book: dict[str, Any], profile: dict[str, Any], people: list[Any], accs: dict[str, list[Any]],
         snaps: dict[str, list[Any]], cols: dict[str, list[Any]], cfg: dict[str, Any], *, budget: Optional[float],
         can_fee: bool, today: Optional[date] = None) -> dict[str, Any]:
    """Kural puanı (0–100): ağırlıklar ayardan. Bileşenler 0–1; ağırlıkla çarpılır. Rakamı model üretmez."""
    today = today or _today()
    w = cfg["weights"]
    total_w = sum(w.values()) or 1.0
    book_id = (book.get("kitapId") or "").lower()
    # Geçmiş performans tabanı: sonucu girilmiş işbirliklerinin kişi başı ortalama etkileşimi.
    per_person: dict[str, float] = {}
    for pid, cs in cols.items():
        vals = [c.engagement for c in cs if c.engagement is not None and c.stage != "vazgecildi"]
        if vals:
            per_person[pid] = sum(vals) / len(vals)
    base = median(per_person.values()) if per_person else None
    items, excluded = [], []
    for p in people:
        mine = cols.get(p.id, [])
        if p.do_not_contact:
            excluded.append({"personId": p.id, "name": p.name, "reason": "İletişim kurulmasın işaretli."})
            continue
        same = [c for c in mine if book_id and (c.crm_book_id or "").lower() == book_id and c.stage != "vazgecildi"]
        if same:
            excluded.append({"personId": p.id, "name": p.name,
                             "reason": f"Bu kitap İŞB-{same[0].no} ile zaten işbirliğinde ({STAGES[same[0].stage]})."})
            continue
        topics = _loads(p.topics_json, [])
        ages = _loads(p.age_groups_json, [])
        parts: dict[str, float] = {}
        why: list[str] = []
        hit = [t for t in topics if t in profile["topics"]]
        if not profile["topics"]:
            parts["konu"] = 0.5
            why.append("kitabın konu etiketi CRM'den çıkmadı")
        elif hit:
            parts["konu"] = 1.0
            why.append("konu uyumu: " + ", ".join(TOPICS[t] for t in hit))
        else:
            parts["konu"] = 0.0
            why.append("konusu kitapla örtüşmüyor" if topics else "konu etiketi girilmemiş")
        age_hit = [a for a in ages if a in profile["ageGroups"]]
        if not profile["ageGroups"] or not ages:
            parts["yas"] = 0.5
        else:
            parts["yas"] = 1.0 if age_hit else 0.0
            why.append("yaş grubu uyuyor" if age_hit else "kitlesinin yaşı kitaba uymuyor")
        mine_eng = per_person.get(p.id)
        if mine_eng is None or not base:
            parts["performans"] = 0.5
            why.append("geçmiş sonuç kaydı yok")
        else:
            ratio = mine_eng / base
            parts["performans"] = min(1.0, ratio / 2)
            n = sum(1 for c in mine if c.engagement is not None and c.stage != "vazgecildi")
            why.append(f"{n} işbirliğinde ortalama etkileşim, kayıtlı içerik üreticilerinin ortancasının {ratio:.1f} katı")
        rel = relation(mine, today)
        parts["iliski"] = rel["score"] / 100
        if rel["published"]:
            why.append(f"ilişki puanı {rel['score']}")
        last = max((c.created_day for c in mine if c.stage != "vazgecildi"), default=None)
        if last is not None and (today - last).days < cfg["cooldownDays"]:
            parts["tazelik"] = 0.0
            why.append(f"son işbirliği {(today - last).days} gün önce")
        else:
            parts["tazelik"] = 1.0
        if budget is None or (p.fee_min is None and p.fee_max is None):
            parts["butce"] = 0.5
        else:
            lo = float(p.fee_min if p.fee_min is not None else p.fee_max)
            parts["butce"] = 1.0 if lo <= budget else 0.0
            why.append("ücret aralığı bütçeye uyuyor" if lo <= budget else "ücret aralığı bütçeyi aşıyor")
        score = round(sum(parts[k] * w.get(k, 0) for k in parts) / total_w * 100, 1)
        a = accs.get(p.id, [])
        followers = 0
        for acc in a:
            lt = _latest([s for s in snaps.get(p.id, []) if s.account_id == acc.id])
            followers += (lt or {}).get("followers") or 0
        items.append({"personId": p.id, "name": p.name, "score": score, "parts": {k: round(v, 3) for k, v in parts.items()},
                      "reason": "; ".join(why) + ".", "topics": topics, "ageGroups": ages,
                      "accounts": [{"platform": x.platform, "handle": x.handle} for x in a], "followers": followers or None,
                      "feeMin": _fee(p.fee_min, can_fee), "feeMax": _fee(p.fee_max, can_fee), "relation": rel,
                      "lastCollab": last.isoformat() if last else None, "minor": bool(p.minor)})
    items.sort(key=lambda x: (-x["score"], fold(x["name"])))
    return {"book": {k: book.get(k) for k in ("kitapId", "stokKodu", "ad", "yazar", "turler", "raf", "hedefKitle", "yas")},
            "profile": profile, "weights": w, "items": items, "excluded": excluded, "total": len(items)}


def candidates(engine: sa.engine.Engine, tenant: str, book: dict[str, Any], profile: dict[str, Any], *,
               budget: Optional[float], can_fee: bool) -> dict[str, Any]:
    cfg = settings()
    with engine.connect() as conn:
        people, accs, snaps, cols = _load_people(conn, tenant)
        # Performans tabanı bütün kişilerin geçmişinden: _load_people zaten hepsini getirir.
    return rank(book, profile, people, accs, snaps, cols, cfg, budget=budget, can_fee=can_fee)


# ------------------------------------------------------------------ taslak (brief / iletişim / sadakat)


def _draft_out(d: Any) -> dict[str, Any]:
    return {"id": d.id, "kind": d.kind, "kindLabel": DRAFT_KINDS.get(d.kind, d.kind), "subject": d.subject, "body": d.body,
            "source": d.source, "dropped": _loads(d.dropped_json, []), "status": d.status, "createdBy": d.created_by,
            "createdAt": _iso(d.created_at), "approvedBy": d.approved_by, "approvedAt": _iso(d.approved_at),
            "sentBy": d.sent_by, "sentAt": _iso(d.sent_at)}


def draft_facts(engine: sa.engine.Engine, tenant: str, collab_id: str, book: Optional[dict[str, Any]]) -> dict[str, Any]:
    with engine.connect() as conn:
        c = _collab_row(conn, tenant, collab_id)
        p = _person_row(conn, tenant, c.person_id)
        accs = conn.execute(sa.select(ACCOUNTS).where(ACCOUNTS.c.person_id == p.id)).all()
        past = conn.execute(sa.select(COLLABS).where(COLLABS.c.tenant_id == tenant, COLLABS.c.person_id == p.id,
                                                      COLLABS.c.id != collab_id, COLLABS.c.stage.notin_(("vazgecildi", "teklif")))
                            .order_by(COLLABS.c.created_at)).all()
    b = book or {}
    return {"collab": c, "person": p, "kisi": p.name, "kitap": b.get("ad") or c.book_title, "yazar": b.get("yazar"),
            "tur": KINDS[c.kind], "kind": c.kind, "yayin": c.due_publish.isoformat() if c.due_publish else None,
            "hesaplar": [f"{PLATFORMS.get(a.platform, a.platform)} @{a.handle}" for a in accs],
            "gecmis": [x.book_title for x in past if x.book_title], "ozet": b.get("ozet"), "oneCikan": b.get("oneCikan"),
            "turler": b.get("turler"), "hedefKitle": b.get("hedefKitle")}


def _rule_body(kind: str, f: dict[str, Any]) -> str:
    kitap = f["kitap"] or "kitabımız"
    if kind == "brief":
        return (f"{kitap}{' (' + f['yazar'] + ')' if f['yazar'] else ''} için kendi üslubunuzla, okurlarınıza kitabı "
                "neden önerdiğinizi anlatan bir paylaşım bekliyoruz.")
    if kind == "iletisim":
        return (f"Merhaba {f['kisi']},\n\nTimaş Yayınları olarak {kitap} kitabımızı sizinle paylaşmak isteriz. "
                "Uygun görürseniz kitabı adresinize gönderelim.")
    gecmis = ", ".join(f["gecmis"]) if f["gecmis"] else ""
    return (f"Merhaba {f['kisi']},\n\n" + (f"Daha önce birlikte çalıştığımız {gecmis} için teşekkür ederiz. " if gecmis else "")
            + f"Yeni kitabımız {kitap} için yeniden birlikte çalışmayı isteriz.")


def render_draft(kind: str, f: dict[str, Any], chat: Optional[Callable[[list[dict[str, str]]], str]],
                 extra_claims: Iterable[str] = ()) -> dict[str, Any]:
    """Model yalnız serbest paragrafı yazar; alıntı/rakam/iddia/teknoloji adı denetiminden (`marketing.guard`) geçer.
    Kitap künyesi, tarih ve yasal etiket maddesi şablondadır."""
    from semantic_bridge.marketing import guard

    if kind not in DRAFT_KINDS:
        raise InfluencerError("Taslak türü tanımlı değil.")
    facts = [x for x in (f["kisi"], f["kitap"], f["yazar"], f["tur"], f["yayin"], *f["hesaplar"], *f["gecmis"]) if x]
    sources = [x for x in (f["ozet"], f["oneCikan"], f["turler"], f["hedefKitle"]) if x]
    source, dropped, body = "kural", [], _rule_body(kind, f)
    if chat is not None:
        task = {"brief": "İçerik üreticisine gidecek brief'in 'Kitap hakkında ve beklenen içerik' paragrafını yaz (4-6 cümle).",
                "iletisim": "İçerik üreticisine ilk iletişim e-postasının gövdesini yaz (selamlama dahil, 4-6 cümle, samimi ve kısa).",
                "sadakat": "Daha önce çalıştığımız içerik üreticisine yeniden işbirliği teklif eden e-posta gövdesini yaz "
                           "(selamlama dahil, geçmiş kitaplara teşekkür, 4-6 cümle)."}[kind]
        lines = [f"Kişi: {f['kisi']}", f"Kitap: {f['kitap']}"] + ([f"Yazar: {f['yazar']}"] if f["yazar"] else []) + [
            f"İşbirliği türü: {f['tur']}"] + ([f"Planlanan yayın: {f['yayin']}"] if f["yayin"] else []) + (
            [f"Hesaplar: {', '.join(f['hesaplar'])}"] if f["hesaplar"] else []) + (
            [f"Birlikte çalışılan önceki kitaplar: {', '.join(f['gecmis'])}"] if f["gecmis"] else []) + (
            [f"Kitabın tanıtım metni: {f['ozet'][:1500]}"] if f["ozet"] else []) + (
            [f"Kitap neden önemli: {f['oneCikan'][:800]}"] if f["oneCikan"] else [])
        messages = [
            {"role": "system", "content": "Sen Timaş Yayınları'nın işbirliği ekibine yardım eden Zeki AI'sın. Türkçe yaz. "
                                          "Yalnız verilen bilgileri kullan; rakam, tarih, ödül, satış bilgisi ya da alıntı uydurma. "
                                          "Reklam/işbirliği etiketi maddesini yazma, şablonda var. Ücretten söz etme."},
            {"role": "user", "content": task + "\n\n" + "\n".join(lines)},
        ]
        try:
            raw = (chat(messages) or "").strip()
        except Exception as e:  # noqa: BLE001 — model yoksa kural metni kalır
            log.warning("işbirliği taslağı modelden alınamadı: %s", e)
            raw = ""
        if raw:
            checked = guard.check(raw, sources, facts, extra_claims)
            dropped = checked["dusen"]
            if checked["metin"].strip():
                body, source = checked["metin"].strip(), "zeki"
    kitap = f["kitap"] or "Kitap"
    if kind == "brief":
        head = [f"Brief — {kitap}", "", f"İçerik üreticisi: {f['kisi']}"] + (
            [f"Hesaplar: {', '.join(f['hesaplar'])}"] if f["hesaplar"] else []) + (
            [f"Yazar: {f['yazar']}"] if f["yazar"] else []) + [f"İşbirliği türü: {f['tur']}"] + (
            [f"Planlanan yayın tarihi: {f['yayin']}"] if f["yayin"] else [])
        text = "\n".join(head) + "\n\nKitap hakkında ve beklenen içerik\n" + body + "\n\nKurallar\n- " + DISCLOSURE_CLAUSE + "\n- " + REVIEW_CLAUSE + "\n"
        subject = f"Brief: {kitap}"
    else:
        tail = "\n\n" + DISCLOSURE_CLAUSE if f["kind"] in ("ucretli", "karsilikli") else ""
        text = body + tail + "\n\nSevgiler,\nTimaş Yayınları\n"
        subject = f"Timaş Yayınları · {kitap}" if kind == "iletisim" else f"Yeniden birlikte: {kitap}"
    return {"subject": subject, "body": text, "source": source, "dropped": dropped}


def save_draft(engine: sa.engine.Engine, tenant: str, user: str, collab_id: str, kind: str, d: dict[str, Any]) -> dict[str, Any]:
    did = _new()
    with engine.begin() as conn:
        _collab_row(conn, tenant, collab_id)
        conn.execute(sa.insert(DRAFTS).values(id=did, tenant_id=tenant, collab_id=collab_id, kind=kind, subject=d["subject"][:300],
                                              body=d["body"], source=d["source"], dropped_json=_dump(d.get("dropped") or []),
                                              status="taslak", created_by=user, created_at=_now()))
        _event(conn, collab_id, user, "taslak", f"{DRAFT_KINDS[kind]} ({'Zeki AI' if d['source'] == 'zeki' else 'şablon'})")
        row = conn.execute(sa.select(DRAFTS).where(DRAFTS.c.id == did)).first()
    return _draft_out(row)


def _draft_row(conn: sa.Connection, tenant: str, collab_id: str, draft_id: str) -> Any:
    row = conn.execute(sa.select(DRAFTS).where(DRAFTS.c.id == draft_id, DRAFTS.c.collab_id == collab_id,
                                               DRAFTS.c.tenant_id == tenant)).first()
    if row is None:
        raise InfluencerError("Taslak bulunamadı.", 404)
    return row


def edit_draft(engine: sa.engine.Engine, tenant: str, user: str, collab_id: str, draft_id: str, body: dict[str, Any]) -> dict[str, Any]:
    text = str(body.get("body") or "").strip()[:20000]
    if not text:
        raise InfluencerError("Taslak boş olamaz.")
    with engine.begin() as conn:
        d = _draft_row(conn, tenant, collab_id, draft_id)
        if d.status != "taslak":
            raise InfluencerError("Onaylanmış taslak değişmez; yeni taslak açın.", 409)
        conn.execute(sa.update(DRAFTS).where(DRAFTS.c.id == draft_id).values(
            body=text, subject=_text(body.get("subject"), 300) or d.subject, source="elle" if d.source != "elle" else d.source))
        row = conn.execute(sa.select(DRAFTS).where(DRAFTS.c.id == draft_id)).first()
    return _draft_out(row)


def approve_draft(engine: sa.engine.Engine, tenant: str, user: str, collab_id: str, draft_id: str) -> dict[str, Any]:
    with engine.begin() as conn:
        d = _draft_row(conn, tenant, collab_id, draft_id)
        if d.status != "taslak":
            raise InfluencerError("Taslak zaten onaylı.", 409)
        if d.kind == "brief" and DISCLOSURE_CLAUSE not in d.body:
            raise InfluencerError("Brief'ten yasal etiket maddesi silinmiş; madde olmadan onaylanmaz.", 409)
        conn.execute(sa.update(DRAFTS).where(DRAFTS.c.id == draft_id).values(status="onayli", approved_by=user, approved_at=_now()))
        _event(conn, collab_id, user, "taslak-onay", DRAFT_KINDS[d.kind])
        row = conn.execute(sa.select(DRAFTS).where(DRAFTS.c.id == draft_id)).first()
    return _draft_out(row)


def mark_mailed(engine: sa.engine.Engine, tenant: str, user: str, collab_id: str, draft_id: str) -> dict[str, Any]:
    """Onaylı taslağı insan kendi e-postasıyla gönderdi: kayıt. Portal dışarıya e-posta göndermez (ilk sürüm kararı)."""
    with engine.begin() as conn:
        d = _draft_row(conn, tenant, collab_id, draft_id)
        c = _collab_row(conn, tenant, collab_id)
        p = _person_row(conn, tenant, c.person_id)
        if d.status != "onayli":
            raise InfluencerError("Yalnız onaylı taslak gönderildi işaretlenir.", 409)
        if p.do_not_contact:
            raise InfluencerError(f"«{p.name}» için «iletişim kurulmasın» işaretli.", 409)
        conn.execute(sa.update(DRAFTS).where(DRAFTS.c.id == draft_id).values(sent_by=user, sent_at=_now()))
        _event(conn, collab_id, user, "eposta", f"{DRAFT_KINDS[d.kind]} elle gönderildi" + (f" ({p.email})" if p.email else ""))
        row = conn.execute(sa.select(DRAFTS).where(DRAFTS.c.id == draft_id)).first()
    return {**_draft_out(row), "email": p.email}


# ------------------------------------------------------------------ ödeme


def _payout_brief(p: Any, can_fee: bool) -> Optional[dict[str, Any]]:
    if p is None:
        return None
    return {"id": p.id, "no": p.no, "status": p.status, "statusLabel": PAYOUT_STATES.get(p.status, p.status),
            "amount": _fee(p.amount, can_fee), "logoDocNo": p.logo_doc_no, "paidAt": _iso(p.paid_at)}


def payout_month(month: str = "") -> tuple[date, date]:
    try:
        m = datetime.strptime(month, "%Y-%m").date() if month else _today().replace(day=1)
    except ValueError as e:
        raise InfluencerError("Ay YYYY-AA biçiminde olmalı.") from e
    return _month_bounds(m)


def payouts_stmt(tenant: str, m0: date, m1: date):
    """Açık ödeme satırları (hazır, onaylı) + ödeme günü ayda olan ödenenler."""
    return sa.select(PAYOUTS, COLLABS.c.no.label("cno"), COLLABS.c.book_title, COLLABS.c.kind, COLLABS.c.stage,
                     COLLABS.c.published_url, COLLABS.c.created_by.label("owner"), PEOPLE.c.name.label("pname"),
                     PEOPLE.c.id.label("pid"), PEOPLE.c.email) \
        .join(COLLABS, COLLABS.c.id == PAYOUTS.c.collab_id).join(PEOPLE, PEOPLE.c.id == COLLABS.c.person_id) \
        .where(PAYOUTS.c.tenant_id == tenant, sa.or_(PAYOUTS.c.status.in_(("hazir", "onayli")),
                                                     sa.and_(PAYOUTS.c.status == "odendi", PAYOUTS.c.paid_at >= m0,
                                                             PAYOUTS.c.paid_at <= m1))) \
        .order_by(PAYOUTS.c.status, PAYOUTS.c.no)


def list_payouts(engine: sa.engine.Engine, tenant: str, month: str = "") -> dict[str, Any]:
    """Açık satırlar (hazır, onaylı) her zaman; ödenenler ödeme günü `month` (YYYY-AA, boşsa bu ay) içinde olanlar."""
    try:
        m = datetime.strptime(month, "%Y-%m").date() if month else _today().replace(day=1)
    except ValueError as e:
        raise InfluencerError("Ay YYYY-AA biçiminde olmalı.") from e
    m0, m1 = _month_bounds(m)
    with engine.connect() as conn:
        rows = conn.execute(payouts_stmt(tenant, m0, m1)).all()
    items = [{"id": r.id, "no": r.no, "collabId": r.collab_id, "collabCode": f"İŞB-{r.cno}", "personId": r.pid,
              "personName": r.pname, "email": r.email, "bookTitle": r.book_title, "kindLabel": KINDS.get(r.kind, r.kind),
              "stage": r.stage, "publishedUrl": r.published_url, "amount": float(r.amount), "status": r.status,
              "statusLabel": PAYOUT_STATES[r.status], "logoDocNo": r.logo_doc_no, "note": r.note, "createdBy": r.created_by,
              "createdAt": _iso(r.created_at), "approvedBy": r.approved_by, "approvedAt": _iso(r.approved_at),
              "paidAt": _iso(r.paid_at), "paidBy": r.paid_by, "owner": r.owner} for r in rows]
    totals = {s: round(sum(i["amount"] for i in items if i["status"] == s), 2) for s in ("hazir", "onayli", "odendi")}
    return {"month": m0.strftime("%Y-%m"), "items": items, "totals": totals}


def decide_payout(engine: sa.engine.Engine, tenant: str, user: str, payout_id: str, body: dict[str, Any], *,
                  can_approve: bool, can_pay: bool) -> dict[str, Any]:
    """onayla (hazır→onaylı; işi açan ya da satırı hazırlayan onaylayamaz), ode (onaylı→ödendi; Logo belge no şart,
    iş kapanır), geri (onaylı→hazır; not şart), iptal (hazır→iptal; iş rapora döner)."""
    action = str(body.get("action") or "")
    now = _now()
    with engine.begin() as conn:
        p = conn.execute(sa.select(PAYOUTS).where(PAYOUTS.c.id == payout_id, PAYOUTS.c.tenant_id == tenant)).first()
        if p is None:
            raise InfluencerError("Ödeme satırı bulunamadı.", 404)
        c = _collab_row(conn, tenant, p.collab_id)
        if action == "onayla":
            if not (can_approve or can_pay):
                raise InfluencerError("Ödeme onayı yetkiniz yok.", 403)
            if p.status != "hazir":
                raise InfluencerError("Yalnız hazır satır onaylanır.", 409)
            if user.lower() in {(p.created_by or "").lower(), (c.created_by or "").lower()}:
                raise InfluencerError("İşi açan ya da ödeme satırını hazırlayan onaylayamaz; başka biri onaylamalı.", 409)
            conn.execute(sa.update(PAYOUTS).where(PAYOUTS.c.id == payout_id).values(status="onayli", approved_by=user, approved_at=now))
            _event(conn, c.id, user, "odeme", f"Ödeme #{p.no} onaylandı ({_tl(p.amount)})")
        elif action == "ode":
            if not can_pay:
                raise InfluencerError("Ödeme kaydı yetkiniz yok.", 403)
            if p.status != "onayli":
                raise InfluencerError("Yalnız onaylı satır ödendi işaretlenir.", 409)
            ref = _text(body.get("logoDocNo"), 200)
            if not ref:
                raise InfluencerError("Ödemenin Logo belge numarası yazılmalı.")
            paid = _day(body.get("paidAt"), "Ödeme") or _today()
            if paid > _today():
                raise InfluencerError("Ödeme tarihi ileri bir gün olamaz.")
            conn.execute(sa.update(PAYOUTS).where(PAYOUTS.c.id == payout_id).values(status="odendi", logo_doc_no=ref,
                                                                                   paid_at=paid, paid_by=user))
            conn.execute(sa.update(COLLABS).where(COLLABS.c.id == c.id).values(stage="kapali", stage_at=now, updated_by=user,
                                                                               updated_at=now))
            _event(conn, c.id, user, "odeme", f"Ödeme #{p.no} ödendi: {paid.isoformat()} · {ref}; iş kapandı")
        elif action == "geri":
            if not (can_approve or can_pay):
                raise InfluencerError("Ödeme onayı yetkiniz yok.", 403)
            if p.status != "onayli":
                raise InfluencerError("Yalnız onaylı (ödenmemiş) satır geri alınır.", 409)
            note = _text(body.get("note"), 2000)
            if not note:
                raise InfluencerError("Geri alma nedeni yazılmalı.")
            conn.execute(sa.update(PAYOUTS).where(PAYOUTS.c.id == payout_id).values(status="hazir", approved_by=None,
                                                                                   approved_at=None, note=note))
            _event(conn, c.id, user, "odeme", f"Ödeme #{p.no} onayı geri alındı: {note}")
        elif action == "iptal":
            if p.status != "hazir":
                raise InfluencerError("Yalnız hazır satır iptal edilir.", 409)
            note = _text(body.get("note"), 2000)
            if not note:
                raise InfluencerError("İptal nedeni yazılmalı.")
            conn.execute(sa.update(PAYOUTS).where(PAYOUTS.c.id == payout_id).values(status="iptal", note=note))
            conn.execute(sa.update(COLLABS).where(COLLABS.c.id == c.id).values(stage="rapor", stage_at=now, updated_by=user,
                                                                               updated_at=now))
            _event(conn, c.id, user, "odeme", f"Ödeme #{p.no} iptal: {note}; iş rapora döndü")
        else:
            raise InfluencerError("İşlem «onayla», «ode», «geri» ya da «iptal» olmalı.")
    return {"id": payout_id, "no": p.no, "action": action, "collabId": c.id, "amount": float(p.amount)}


# ------------------------------------------------------------------ rapor


def report_stmt(tenant: str):
    """Vazgeçilmemiş işbirlikleri (dönem süzgeci yayın / plan / kayıt gününe göre hesapta)."""
    return sa.select(COLLABS, PEOPLE.c.name.label("pname")).join(PEOPLE, PEOPLE.c.id == COLLABS.c.person_id) \
        .where(COLLABS.c.tenant_id == tenant, COLLABS.c.stage != "vazgecildi")


def report(engine: sa.engine.Engine, tenant: str, frm: date, to: date, can_fee: bool) -> dict[str, Any]:
    """Dönem (yayın günü, yoksa planlanan yayın, yoksa kayıt günü) içindeki vazgeçilmemiş işbirlikleri. Harcama =
    işbirliği ücreti toplamı; CPE = harcama ÷ etkileşim toplamı (etkileşim 0 ise yok). Ücret görünmeyen kullanıcıda
    harcama ve CPE boştur."""
    if to < frm:
        raise InfluencerError("Dönem sonu başlangıçtan önce olamaz.")
    with engine.connect() as conn:
        rows = conn.execute(report_stmt(tenant)).all()
    rows = [r for r in rows if frm <= period_day(r) <= to]

    def agg(group: list[Any]) -> dict[str, Any]:
        spend = sum(Decimal(r.fee or 0) for r in group)
        eng = sum(r.engagement or 0 for r in group)
        return {"collabs": len(group), "published": sum(1 for r in group if FLOW.index(r.stage) >= FLOW.index("yayinda")),
                "reach": sum(r.reach or 0 for r in group), "engagement": eng,
                "spend": float(spend) if can_fee else None,
                "cpe": (float((spend / eng).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)) if eng and can_fee else None),
                "disclosureMissing": sum(1 for r in group if FLOW.index(r.stage) >= FLOW.index("yayinda") and r.disclosure_ok is not True)}

    by_person: dict[str, list[Any]] = {}
    by_book: dict[str, list[Any]] = {}
    for r in rows:
        by_person.setdefault(r.person_id, []).append(r)
        by_book.setdefault(r.crm_book_id or f"ad:{r.book_title}", []).append(r)
    persons = sorted(({"personId": k, "name": v[0].pname, **agg(v)} for k, v in by_person.items()),
                     key=lambda x: (-(x["engagement"] or 0), fold(x["name"])))
    books = sorted(({"crmBookId": v[0].crm_book_id, "bookTitle": v[0].book_title, **agg(v)} for k, v in by_book.items()),
                   key=lambda x: (-(x["engagement"] or 0), fold(x["bookTitle"])))
    detail = [dict(_collab_out(r, r.pname, can_fee), periodDay=period_day(r).isoformat())
              for r in sorted(rows, key=lambda x: (period_day(x), x.no))]
    return {"from": frm.isoformat(), "to": to.isoformat(), "total": agg(rows), "people": persons, "books": books,
            "items": detail}


def report_xlsx(rep: dict[str, Any], crm: Optional[dict[str, Any]], can_fee: bool) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    t = rep["total"]
    ws.append(["İşbirlikleri raporu", f"{rep['from']} – {rep['to']}"])
    ws.append([])
    ws.append(["İşbirliği", t["collabs"]])
    ws.append(["Yayına ulaşan", t["published"]])
    ws.append(["Erişim", t["reach"]])
    ws.append(["Etkileşim", t["engagement"]])
    if can_fee:
        ws.append(["Harcama (₺, KDV hariç)", t["spend"]])
        ws.append(["Etkileşim başı maliyet (₺)", t["cpe"]])
    ws.append(["Yasal etiketi işaretlenmemiş yayın", t["disclosureMissing"]])
    if crm is not None and can_fee:
        ws.append([])
        ws.append(["CRM pazarlama bütçe modülü «Influencer» kayıtları (ayrı kaynak, toplanmaz)", crm.get("toplam"), f"{crm.get('kayit')} kayıt"])
    ws["A1"].font = Font(bold=True)
    for title, key, first in (("Kişiler", "people", "name"), ("Kitaplar", "books", "bookTitle")):
        s = wb.create_sheet(title)
        head = ["Ad" if key == "people" else "Kitap", "İşbirliği", "Yayında", "Erişim", "Etkileşim"] + (["Harcama (₺)", "CPE (₺)"] if can_fee else [])
        s.append(head)
        for r in rep[key]:
            s.append([r[first], r["collabs"], r["published"], r["reach"], r["engagement"]] + ([r["spend"], r["cpe"]] if can_fee else []))
    s = wb.create_sheet("İşbirlikleri")
    head = ["No", "Dönem günü", "Kişi", "Kitap", "Tür", "Aşama", "Bağlantı", "Yasal etiket", "Erişim", "Etkileşim"] + (["Ücret (₺)", "CPE (₺)"] if can_fee else [])
    s.append(head)
    for r in rep["items"]:
        s.append([r["code"], r["periodDay"], r["personName"], r["bookTitle"], r["kindLabel"], r["stageLabel"], r["publishedUrl"],
                  {True: "var", False: "yok", None: "işaretlenmedi"}[r["disclosureOk"]], r["reach"], r["engagement"]]
                 + ([r["fee"], r["cpe"]] if can_fee else []))
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def payouts_xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Ödeme listesi"
    ws.append(["No", "İşbirliği", "Kişi", "E-posta", "Kitap", "Tür", "Tutar (₺, KDV hariç)", "Durum", "Onaylayan",
               "Ödeme günü", "Logo belge no", "Paylaşım"])
    for i in data["items"]:
        ws.append([i["no"], i["collabCode"], i["personName"], i["email"], i["bookTitle"], i["kindLabel"], i["amount"],
                   i["statusLabel"], i["approvedBy"], i["paidAt"], i["logoDocNo"], i["publishedUrl"]])
    ws.append([])
    for s in ("hazir", "onayli", "odendi"):
        ws.append(["", "", "", "", "", f"Toplam {PAYOUT_STATES[s].lower()}", data["totals"][s]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
