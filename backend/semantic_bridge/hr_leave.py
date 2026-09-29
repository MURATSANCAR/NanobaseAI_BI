"""M60 İzin yönetimi (İnsan Kaynakları modülünün içinde): izin türleri, resmî tatil ve çalışma takvimi, gün hesabı,
talep → yönetici onayı → (ayara göre) İK onayı akışı, bakiye defteri, yıllık izin hakedişi, ekip takvimi, bordro listesi.

Plan: docs/analiz/kullanici-ihtiyaclari/M60-izin-yonetimi-ve-ik-bildirimleri.md.

Kurallar (bağlayıcı):
- **Kişi = özlük kaydı** (`semantic_hr_people`, `hr_portal.person_of`). Yönetici zinciri `yonetici_id_no`; kendisi ya da boşsa
  talep doğrudan İK'ya düşer. Onaylayan = talep anındaki yöneticinin özlük kaydına bağlı portal hesabı ya da İK.
- **Bakiye bir defterdir** (`semantic_hr_leave_ledger`): açılış, hakediş, kullanım, iade, düzeltme hareketlerinin toplamı.
  Ayrı «bakiye» kolonu yok. Yıllık izin bakiyesi yalnız `from_balance` türlerinden (başlangıçta «yillik»).
- **Gün hesabı tek yer** (`count_days`): iş günü türünde kişinin çalışma takviminde çalışılan ve resmî tatil olmayan günler
  (arife yarım), takvim günü türünde (doğum izni, rapor) aralıktaki bütün günler. Ekran yalnız gösterir.
- **Yasal değerler ayardır**, kodda sabit değildir: türlerin gün sınırları ve hakediş basamakları İK/hukuk onayıyla
  (`approved`) açılır; onaysız türle talep açılmaz. Başlangıç değerleri 4857 sayılı Kanun md. 53–56 ve ek md. 2'ye göre.
- **Hassas tür** (rapor, doğum): yönetici ve rehber yalnız «izinli» görür; tür adı ve belge yalnız kişinin kendisine ve
  `ozellik:ik.ozluk-hassas` yetkisine. E-postada tür yazılmaz.
- Hakediş gece işi yalnız **başlangıç tarihinden sonraki** yıldönümlerini yazar (`leave.accrualFrom`, ilk koşuda o gün):
  geçmiş yılların bakiyesi açılış Excel'iyle girilir; aynı yıldönümü iki kez yazılmaz.
"""
from __future__ import annotations

import io
import threading
import weakref
from datetime import date, datetime, timedelta
from typing import Any, Callable, Iterable, Optional

import sqlalchemy as sa

from semantic_bridge import hr_core as H
from semantic_bridge import hr_portal as PT
from semantic_bridge.hr_core import HrError, clean, dump, iso, load, new_id, now

_md = sa.MetaData()
_ready: "weakref.WeakSet[sa.engine.Engine]" = weakref.WeakSet()
_lock = threading.Lock()

F_LEAVE = "ozellik:ik.izin-yonet"
F_LEAVE_SET = "ozellik:ik.izin-ayar"

STATUS = {"yonetici": "Yönetici onayı bekliyor", "ik": "İK onayı bekliyor", "onaylandi": "Onaylandı", "reddedildi": "Reddedildi",
          "geri_alindi": "Geri alındı", "iptal": "İptal edildi (bakiye iade)"}
OPEN = ("yonetici", "ik")
REASONS = {"acilis": "Açılış bakiyesi", "hakedis": "Yıllık hakediş", "kullanim": "Kullanım", "iade": "İptal iadesi",
           "duzeltme": "İK düzeltmesi"}
HALF = {"": "Tam gün", "sabah": "Yarım gün (sabah)", "ogleden_sonra": "Yarım gün (öğleden sonra)"}
COUNT_MODES = {"is_gunu": "İş günü", "takvim_gunu": "Takvim günü"}
WEEKDAYS = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]

LEAVE_DEFAULTS: dict[str, Any] = {
    "leave.flow": "yonetici",                 # yonetici | yonetici_ik
    "leave.advance": False,                   # yıllık izin bakiyesiz (eksiye) alınabilir mi
    "leave.tiers": [{"minYears": 1, "days": 14}, {"minYears": 6, "days": 20}, {"minYears": 15, "days": 26}],
    "leave.ageRule": {"maxYoungAge": 18, "minOldAge": 50, "days": 20},
    "leave.startField": "f_ise_giris_tarihi",
    "leave.conflictPct": 50,
    "leave.reminderWorkdays": 2,
    "leave.accrualFrom": None,
}

#: Başlangıç izin türleri — hepsi «onaylanmadı» gelir; İK/hukuk gün sayısını doğrulayıp açar.
SEED_TYPES: list[dict[str, Any]] = [
    {"key": "yillik", "label": "Yıllık ücretli izin", "paid": True, "from_balance": True, "half_day": True, "count_mode": "is_gunu",
     "legal": "4857 md. 53–56: hizmet 1–5 yıl 14, 5–15 yıl 20, 15+ yıl 26 gün; 18 yaş altı ve 50 yaş üstü en az 20 gün; bölünürse bir parça en az 10 gün."},
    {"key": "evlilik", "label": "Evlilik izni", "paid": True, "max_request": 3, "count_mode": "is_gunu", "legal": "4857 ek md. 2: 3 gün ücretli."},
    {"key": "es_dogum", "label": "Eşinin doğumu (babalık) izni", "paid": True, "max_request": 5, "count_mode": "is_gunu", "legal": "4857 ek md. 2: 5 gün ücretli."},
    {"key": "olum", "label": "Ölüm izni", "paid": True, "max_request": 3, "count_mode": "is_gunu",
     "legal": "4857 ek md. 2: anne, baba, eş, kardeş ya da çocuğun ölümünde 3 gün ücretli."},
    {"key": "evlat_edinme", "label": "Evlat edinme izni", "paid": True, "max_request": 3, "count_mode": "is_gunu", "legal": "4857 ek md. 2: 3 gün ücretli."},
    {"key": "engelli_cocuk", "label": "Engelli / kronik hasta çocuk tedavi izni", "paid": True, "max_year": 10, "count_mode": "is_gunu", "needs_doc": True,
     "legal": "4857 ek md. 2: raporla, yılda toplu ya da bölümler hâlinde 10 güne kadar ücretli."},
    {"key": "rapor", "label": "Raporlu (sağlık) izni", "paid": False, "payroll": True, "sensitive": True, "needs_doc": True, "count_mode": "takvim_gunu",
     "legal": "Hekim raporu; ödeme SGK geçici iş göremezlik kuralına göre (bordro)."},
    {"key": "dogum", "label": "Doğum (analık) izni", "paid": False, "payroll": True, "sensitive": True, "needs_doc": True, "count_mode": "takvim_gunu",
     "max_request": 112, "legal": "4857 md. 74: doğumdan önce 8, sonra 8 hafta (çoğul gebelikte +2 hafta)."},
    {"key": "yol", "label": "Yol izni (ücretsiz)", "paid": False, "payroll": True, "max_request": 4, "count_mode": "is_gunu",
     "legal": "4857 md. 57: yıllık iznini başka yerde geçirecek işçiye istemi hâlinde toplam 4 güne kadar ücretsiz yol izni."},
    {"key": "ucretsiz", "label": "Ücretsiz izin", "paid": False, "payroll": True, "count_mode": "is_gunu", "legal": "Tarafların anlaşmasıyla; bordroda gün düşer."},
    {"key": "idari", "label": "İdari izin (ücretli)", "paid": True, "count_mode": "is_gunu", "legal": "İşverenin verdiği ücretli izin; bakiyeden düşmez."},
]

#: «Sabit resmî tatilleri ekle» düğmesinin eklediği günler (2429 sayılı Kanun); dini bayramlar her yıl İK'ca girilir.
FIXED_HOLIDAYS = [((1, 1), "Yılbaşı", False), ((4, 23), "Ulusal Egemenlik ve Çocuk Bayramı", False), ((5, 1), "Emek ve Dayanışma Günü", False),
                  ((5, 19), "Atatürk'ü Anma, Gençlik ve Spor Bayramı", False), ((7, 15), "Demokrasi ve Millî Birlik Günü", False),
                  ((8, 30), "Zafer Bayramı", False), ((10, 28), "Cumhuriyet Bayramı arifesi", True), ((10, 29), "Cumhuriyet Bayramı", False)]


def _ts(name: str, **kw: Any) -> sa.Column:
    return sa.Column(name, sa.DateTime(timezone=True), **kw)


TYPES = sa.Table(
    "semantic_hr_leave_types", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("key", sa.String(40), primary_key=True),
    sa.Column("label", sa.String(120), nullable=False),
    sa.Column("paid", sa.Boolean, nullable=False, default=True),
    sa.Column("from_balance", sa.Boolean, nullable=False, default=False),
    sa.Column("payroll", sa.Boolean, nullable=False, default=False),        # aylık bordro listesine girer
    sa.Column("sensitive", sa.Boolean, nullable=False, default=False),
    sa.Column("needs_doc", sa.Boolean, nullable=False, default=False),
    sa.Column("half_day", sa.Boolean, nullable=False, default=False),
    sa.Column("count_mode", sa.String(16), nullable=False, default="is_gunu"),
    sa.Column("max_request", sa.Float),                                     # bir talepte en çok gün
    sa.Column("max_year", sa.Float),                                        # takvim yılında en çok gün
    sa.Column("legal", sa.Text),
    sa.Column("approved", sa.Boolean, nullable=False, default=False),       # İK/hukuk gün sayısını doğruladı
    sa.Column("approved_by", sa.String(120)),
    _ts("approved_at"),
    sa.Column("active", sa.Boolean, nullable=False, default=True),
    sa.Column("sort", sa.Integer, nullable=False, default=0),
    sa.Column("builtin", sa.Boolean, nullable=False, default=False),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at"),
)

HOLIDAYS = sa.Table(
    "semantic_hr_holidays", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("day", sa.Date, primary_key=True),
    sa.Column("name", sa.String(120), nullable=False),
    sa.Column("half", sa.Boolean, nullable=False, default=False),
    sa.Column("updated_by", sa.String(120)),
)

CALENDARS = sa.Table(
    "semantic_hr_work_calendars", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("name", sa.String(120), nullable=False),
    sa.Column("days", sa.String(7), nullable=False),                        # «12345»: Pzt=1 … Paz=7
    sa.Column("match_field", sa.String(40)),                                # sube | ofis_lokasyon | departman
    sa.Column("match_values", sa.Text),                                     # JSON liste
    sa.Column("is_default", sa.Boolean, nullable=False, default=False),
    sa.Column("updated_by", sa.String(120)),
)

REQUESTS = sa.Table(
    "semantic_hr_leave_requests", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("person_id", sa.String(40), nullable=False, index=True),
    sa.Column("username", sa.String(120)),
    sa.Column("type_key", sa.String(40), nullable=False),
    sa.Column("start", sa.Date, nullable=False),
    sa.Column("end", sa.Date, nullable=False),
    sa.Column("half", sa.String(16)),
    sa.Column("days", sa.Float, nullable=False),
    sa.Column("deputy_person_id", sa.String(40)),
    sa.Column("contact", sa.String(120)),
    sa.Column("note", sa.Text),
    sa.Column("status", sa.String(16), nullable=False),
    sa.Column("manager_id_no", sa.String(60)),
    sa.Column("needs_hr", sa.Boolean, nullable=False, default=False),
    sa.Column("manager_decided_by", sa.String(120)),
    _ts("manager_decided_at"),
    sa.Column("decided_by", sa.String(120)),
    _ts("decided_at"),
    sa.Column("reject_reason", sa.Text),
    sa.Column("last_reminded", sa.Date),
    sa.Column("created_by", sa.String(120)),
    _ts("created_at", nullable=False),
    _ts("updated_at", nullable=False),
)

EVENTS = sa.Table(
    "semantic_hr_leave_events", _md,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("request_id", sa.String(40), nullable=False, index=True),
    _ts("at", nullable=False),
    sa.Column("actor", sa.String(120), nullable=False),
    sa.Column("action", sa.String(24), nullable=False),
    sa.Column("note", sa.Text),
)

LEDGER = sa.Table(
    "semantic_hr_leave_ledger", _md,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("person_id", sa.String(40), nullable=False, index=True),
    sa.Column("type_key", sa.String(40), nullable=False),
    sa.Column("days", sa.Float, nullable=False),
    sa.Column("reason", sa.String(16), nullable=False),
    sa.Column("request_id", sa.String(40)),
    sa.Column("period", sa.String(16)),                                     # hakediş: yıldönümü tarihi
    sa.Column("on_date", sa.Date, nullable=False),
    sa.Column("note", sa.Text),
    sa.Column("actor", sa.String(120)),
    _ts("at", nullable=False),
)

FILES = sa.Table(
    "semantic_hr_leave_files", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("request_id", sa.String(40), nullable=False, index=True),
    sa.Column("filename", sa.String(300), nullable=False),
    sa.Column("mime", sa.String(120)),
    sa.Column("size", sa.Integer, nullable=False, default=0),
    sa.Column("blob", sa.LargeBinary, nullable=False),
    sa.Column("uploaded_by", sa.String(120)),
    _ts("uploaded_at", nullable=False),
)

TABLE_LABELS = {"semantic_hr_leave_types": "izin türleri", "semantic_hr_holidays": "resmî tatiller",
                "semantic_hr_work_calendars": "çalışma takvimleri", "semantic_hr_leave_requests": "izin talepleri",
                "semantic_hr_leave_events": "izin talebi geçmişi", "semantic_hr_leave_ledger": "izin bakiye defteri",
                "semantic_hr_leave_files": "izin belgeleri"}


def ensure(engine: sa.engine.Engine) -> None:
    with _lock:
        if engine in _ready:
            return
        PT.ensure(engine)
        _md.create_all(engine, checkfirst=True)
        _ready.add(engine)


def delete_person_rows(c: sa.engine.Connection, tenant: str, pid: str, arc: H.Archiver) -> dict[str, int]:
    """Kişi silinirken (`hr_portal.delete_person`, aynı işlem) kişinin izin talepleri, onların geçmişi ve belgeleri (sağlık
    raporu olabilir) ve bakiye defteri pasife alınır (arşive taşınır); başkasının talebinde vekil olarak anılışı boşaltılır,
    eski değer arşive yazılır (talep kalır)."""
    mine = sa.select(REQUESTS.c.id).where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.person_id == pid)
    nf = arc.move(FILES, FILES.c.tenant_id == tenant, FILES.c.request_id.in_(mine))
    ne = arc.move(EVENTS, EVENTS.c.tenant_id == tenant, EVENTS.c.request_id.in_(mine))
    nl = arc.move(LEDGER, LEDGER.c.tenant_id == tenant, LEDGER.c.person_id == pid)
    nr = arc.move(REQUESTS, REQUESTS.c.tenant_id == tenant, REQUESTS.c.person_id == pid)
    dep = [r.id for r in c.execute(sa.select(REQUESTS.c.id).where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.deputy_person_id == pid))]
    for rid in dep:
        arc.note(REQUESTS.name, rid, {"id": rid, "deputy_person_id": pid}, reason="vekil_bosaltildi")
    if dep:
        c.execute(REQUESTS.update().where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.id.in_(dep))
                  .values(deputy_person_id=None, updated_at=now()))
    return {"requests": nr, "events": ne, "files": nf, "ledger": nl, "deputyCleared": len(dep)}


def register_hooks() -> None:
    PT.register_person_cleanup(PT.PersonCleanup("izin", ensure, delete_person_rows))


_seeded: set[tuple[int, str]] = set()


def _seed(engine: sa.engine.Engine, tenant: str) -> None:
    ensure(engine)
    if (id(engine), tenant) in _seeded:
        return
    with _lock, engine.begin() as c:
        have = {r.key for r in c.execute(sa.select(TYPES.c.key).where(TYPES.c.tenant_id == tenant)).all()}
        rows = []
        for i, t in enumerate(SEED_TYPES):
            if t["key"] in have:
                continue
            rows.append(dict(tenant_id=tenant, key=t["key"], label=t["label"], paid=t.get("paid", True), from_balance=t.get("from_balance", False),
                             payroll=t.get("payroll", False), sensitive=t.get("sensitive", False), needs_doc=t.get("needs_doc", False),
                             half_day=t.get("half_day", False), count_mode=t.get("count_mode", "is_gunu"), max_request=t.get("max_request"),
                             max_year=t.get("max_year"), legal=t.get("legal"), approved=False, active=True, sort=(i + 1) * 10, builtin=True,
                             updated_at=now()))
        if rows:
            c.execute(TYPES.insert(), rows)
        if not c.execute(sa.select(CALENDARS.c.id).where(CALENDARS.c.tenant_id == tenant)).first():
            c.execute(CALENDARS.insert().values(id=new_id("tkv"), tenant_id=tenant, name="Hafta içi (Pazartesi–Cuma)", days="12345",
                                                match_field=None, match_values=dump([]), is_default=True))
    _seeded.add((id(engine), tenant))


# ------------------------------------------------------------------ ayarlar


def settings(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    ensure(engine)
    with engine.connect() as c:
        rows = c.execute(sa.select(PT.PSETTINGS).where(PT.PSETTINGS.c.tenant_id == tenant, PT.PSETTINGS.c.key.like("leave.%"))).all()
    out = {k: (v.copy() if isinstance(v, (list, dict)) else v) for k, v in LEAVE_DEFAULTS.items()}
    for r in rows:
        if r.key in out:
            out[r.key] = load(r.value_json, out[r.key])
    return out


def _put_setting(c: Any, tenant: str, key: str, value: Any, actor: str) -> None:
    c.execute(PT.PSETTINGS.delete().where(PT.PSETTINGS.c.tenant_id == tenant, PT.PSETTINGS.c.key == key))
    c.execute(PT.PSETTINGS.insert().values(tenant_id=tenant, key=key, value_json=dump(value), updated_by=actor, updated_at=now()))


def save_settings(engine: sa.engine.Engine, tenant: str, actor: str, body: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    cur = settings(engine, tenant)
    changed: list[str] = []
    with engine.begin() as c:
        for k in LEAVE_DEFAULTS:
            short = k.split(".", 1)[1]
            if short not in body or k == "leave.accrualFrom":
                continue
            v = body[short]
            if k == "leave.flow" and v not in ("yonetici", "yonetici_ik"):
                raise HrError("Onay akışı geçersiz.")
            if k == "leave.advance":
                v = bool(v)
            if k == "leave.tiers":
                try:
                    v = sorted(({"minYears": int(x["minYears"]), "days": float(x["days"])} for x in v), key=lambda x: x["minYears"])
                except (TypeError, ValueError, KeyError):
                    raise HrError("Hakediş basamakları: her satırda «en az yıl» ve «gün» sayı olmalı.") from None
                if not v or v[0]["minYears"] < 1 or any(x["days"] <= 0 for x in v):
                    raise HrError("Hakediş basamakları en az 1. yıldan başlamalı, gün sayısı sıfırdan büyük olmalı.")
            if k == "leave.ageRule":
                try:
                    v = {"maxYoungAge": int(v["maxYoungAge"]), "minOldAge": int(v["minOldAge"]), "days": float(v["days"])}
                except (TypeError, ValueError, KeyError):
                    raise HrError("Yaş kuralı sayı olmalı.") from None
            if k == "leave.startField" and v not in ("f_ise_giris_tarihi", "s_ise_giris_tarihi"):
                raise HrError("Kıdem başlangıcı ilk ya da son işe giriş tarihi olmalı.")
            if k in ("leave.conflictPct", "leave.reminderWorkdays"):
                try:
                    v = int(v)
                except (TypeError, ValueError):
                    raise HrError("Sayı girin.") from None
                if not (1 <= v <= 100):
                    raise HrError("Değer 1 ile 100 arasında olmalı.")
            if v != cur[k]:
                _put_setting(c, tenant, k, v, actor)
                changed.append(short)
    return settings(engine, tenant), changed


def mark_accrual_start(engine: sa.engine.Engine, tenant: str, on: date) -> date:
    st = settings(engine, tenant)
    if st["leave.accrualFrom"]:
        return date.fromisoformat(st["leave.accrualFrom"])
    with engine.begin() as c:
        _put_setting(c, tenant, "leave.accrualFrom", on.isoformat(), "ZEKİ AI")
    return on


# ------------------------------------------------------------------ türler, tatiller, takvimler


def _type_out(r: Any) -> dict[str, Any]:
    return {"key": r.key, "label": r.label, "paid": bool(r.paid), "fromBalance": bool(r.from_balance), "payroll": bool(r.payroll),
            "sensitive": bool(r.sensitive), "needsDoc": bool(r.needs_doc), "halfDay": bool(r.half_day), "countMode": r.count_mode,
            "maxRequest": r.max_request, "maxYear": r.max_year, "legal": r.legal, "approved": bool(r.approved),
            "approvedBy": r.approved_by, "approvedAt": iso(r.approved_at), "active": bool(r.active), "sort": int(r.sort or 0),
            "builtin": bool(r.builtin)}


def types(engine: sa.engine.Engine, tenant: str) -> list[dict[str, Any]]:
    _seed(engine, tenant)
    with engine.connect() as c:
        return [_type_out(r) for r in c.execute(sa.select(TYPES).where(TYPES.c.tenant_id == tenant).order_by(TYPES.c.sort, TYPES.c.key)).all()]


def save_type(engine: sa.engine.Engine, tenant: str, actor: str, body: dict[str, Any], key: Optional[str] = None) -> dict[str, Any]:
    _seed(engine, tenant)
    with engine.begin() as c:
        before = None
        if key:
            before = c.execute(sa.select(TYPES).where(TYPES.c.tenant_id == tenant, TYPES.c.key == key)).first()
            if before is None:
                raise HrError("İzin türü bulunamadı.", 404)
        vals: dict[str, Any] = {}
        if before is None:
            k = clean(body.get("key"), 40).lower()
            if not PT._KEY_RE.match(k) or c.execute(sa.select(TYPES.c.key).where(TYPES.c.tenant_id == tenant, TYPES.c.key == k)).first():
                raise HrError("Anahtar geçersiz ya da kullanılıyor (küçük harf, rakam, alt çizgi).")
            vals.update(key=k, builtin=False, sort=10_000)
        if "label" in body or before is None:
            vals["label"] = clean(body.get("label"), 120)
            if not vals["label"]:
                raise HrError("İzin türünün adı boş olamaz.")
        for src, col in (("paid", "paid"), ("fromBalance", "from_balance"), ("payroll", "payroll"), ("sensitive", "sensitive"),
                         ("needsDoc", "needs_doc"), ("halfDay", "half_day"), ("active", "active")):
            if src in body:
                vals[col] = bool(body[src])
        if "countMode" in body:
            if body["countMode"] not in COUNT_MODES:
                raise HrError("Sayım biçimi geçersiz.")
            vals["count_mode"] = body["countMode"]
        for src, col in (("maxRequest", "max_request"), ("maxYear", "max_year")):
            if src in body:
                raw = body[src]
                if raw in (None, ""):
                    vals[col] = None
                else:
                    try:
                        vals[col] = float(raw)
                    except (TypeError, ValueError):
                        raise HrError("Gün sınırı sayı olmalı.") from None
                    if vals[col] <= 0:
                        raise HrError("Gün sınırı sıfırdan büyük olmalı.")
        if "legal" in body:
            vals["legal"] = str(body.get("legal") or "").strip()[:2000] or None
        if "sort" in body:
            try:
                vals["sort"] = int(body["sort"])
            except (TypeError, ValueError):
                raise HrError("Sıra sayı olmalı.") from None
        # Gün sayısı ya da kural değişirse onay düşer: kural değişikliği yeniden doğrulanmalı.
        rule_cols = {"paid", "from_balance", "count_mode", "max_request", "max_year"}
        if before is not None and any(col in vals and getattr(before, col) != vals[col] for col in rule_cols):
            vals.update(approved=False, approved_by=None, approved_at=None)
        if "approved" in body:
            ok = bool(body["approved"])
            vals.update(approved=ok, approved_by=actor if ok else None, approved_at=now() if ok else None)
        vals.update(updated_by=actor, updated_at=now())
        if before is None:
            c.execute(TYPES.insert().values(tenant_id=tenant, approved=vals.pop("approved", False), active=vals.pop("active", True), **vals))
            key = vals["key"]
        else:
            c.execute(TYPES.update().where(TYPES.c.tenant_id == tenant, TYPES.c.key == key).values(**vals))
        r = c.execute(sa.select(TYPES).where(TYPES.c.tenant_id == tenant, TYPES.c.key == key)).first()
    return _type_out(r)


def holidays(engine: sa.engine.Engine, tenant: str, year: Optional[int] = None) -> list[dict[str, Any]]:
    ensure(engine)
    stmt = sa.select(HOLIDAYS).where(HOLIDAYS.c.tenant_id == tenant)
    if year:
        stmt = stmt.where(HOLIDAYS.c.day >= date(year, 1, 1), HOLIDAYS.c.day <= date(year, 12, 31))
    with engine.connect() as c:
        return [{"day": r.day.isoformat(), "name": r.name, "half": bool(r.half)} for r in c.execute(stmt.order_by(HOLIDAYS.c.day)).all()]


def save_holidays(engine: sa.engine.Engine, tenant: str, actor: str, days: list[dict[str, Any]]) -> int:
    ensure(engine)
    n = 0
    with engine.begin() as c:
        for d in days:
            day = H.parse_date((d or {}).get("day"), "Tatil")
            if day is None:
                raise HrError("Tatil günü boş.")
            c.execute(HOLIDAYS.delete().where(HOLIDAYS.c.tenant_id == tenant, HOLIDAYS.c.day == day))
            if (d or {}).get("delete"):
                n += 1
                continue
            name = clean(d.get("name"), 120)
            if not name:
                raise HrError("Tatilin adı boş olamaz.")
            c.execute(HOLIDAYS.insert().values(tenant_id=tenant, day=day, name=name, half=bool(d.get("half")), updated_by=actor))
            n += 1
    return n


def add_fixed_holidays(engine: sa.engine.Engine, tenant: str, actor: str, year: int) -> int:
    if not 2000 <= year <= 2100:
        raise HrError("Yıl geçersiz.")
    existing = {h["day"] for h in holidays(engine, tenant, year)}
    new = [{"day": date(year, m, d).isoformat(), "name": name, "half": half} for (m, d), name, half in FIXED_HOLIDAYS
           if date(year, m, d).isoformat() not in existing]
    return save_holidays(engine, tenant, actor, new) if new else 0


def calendars(engine: sa.engine.Engine, tenant: str) -> list[dict[str, Any]]:
    _seed(engine, tenant)
    with engine.connect() as c:
        rows = c.execute(sa.select(CALENDARS).where(CALENDARS.c.tenant_id == tenant).order_by(CALENDARS.c.is_default.desc(), CALENDARS.c.name)).all()
    return [{"id": r.id, "name": r.name, "days": r.days, "matchField": r.match_field, "matchValues": load(r.match_values, []),
             "isDefault": bool(r.is_default)} for r in rows]


def save_calendars(engine: sa.engine.Engine, tenant: str, actor: str, items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Takvim listesinin tamamı yazılır; biri varsayılandır (eşleşmeyen herkes)."""
    _seed(engine, tenant)
    if not items:
        raise HrError("En az bir çalışma takvimi olmalı.")
    if sum(1 for x in items if x.get("isDefault")) != 1:
        raise HrError("Tam olarak bir takvim varsayılan olmalı.")
    rows = []
    for x in items:
        days = "".join(sorted({ch for ch in str(x.get("days") or "") if ch in "1234567"}))
        if not days:
            raise HrError("Takvimde en az bir çalışma günü olmalı.")
        mf = x.get("matchField") or None
        if mf not in (None, "sube", "ofis_lokasyon", "departman", "firma"):
            raise HrError("Takvim eşleşme alanı geçersiz.")
        vals = [clean(v, 80) for v in (x.get("matchValues") or []) if clean(v, 80)]
        if not x.get("isDefault") and (not mf or not vals):
            raise HrError(f"«{clean(x.get('name'), 60)}» takvimi için eşleşen alan ve değer girin (ör. Şube = Depo).")
        rows.append(dict(id=x.get("id") or new_id("tkv"), tenant_id=tenant, name=clean(x.get("name"), 120) or "Takvim", days=days,
                         match_field=None if x.get("isDefault") else mf, match_values=dump([] if x.get("isDefault") else vals),
                         is_default=bool(x.get("isDefault")), updated_by=actor))
    with engine.begin() as c:
        c.execute(CALENDARS.delete().where(CALENDARS.c.tenant_id == tenant))
        c.execute(CALENDARS.insert(), rows)
    return calendars(engine, tenant)


def _calendar_for(cals: list[dict[str, Any]], data: dict[str, Any]) -> dict[str, Any]:
    for cal in cals:
        if not cal["isDefault"] and cal["matchField"] and str(data.get(cal["matchField"]) or "") in cal["matchValues"]:
            return cal
    return next((x for x in cals if x["isDefault"]), {"days": "12345", "name": "Hafta içi"})


# ------------------------------------------------------------------ gün hesabı


def count_days(start: date, end: date, *, mode: str, workdays: str, hol: dict[date, bool], half: str = "") -> float:
    """İş günü: takvimde çalışılan ve tatil olmayan günler (yarım tatil 0,5); takvim günü: aradaki bütün günler.
    Yarım gün izin yalnız tek günlük talepte, o gün iş günüyse 0,5."""
    if end < start:
        return 0.0
    if mode == "takvim_gunu":
        return float((end - start).days + 1)
    total = 0.0
    d = start
    while d <= end:
        if str(d.isoweekday()) in workdays:
            if d in hol:
                total += 0.5 if hol[d] else 0.0
            else:
                total += 1.0
        d += timedelta(days=1)
    if half and start == end:
        total = min(total, 0.5)
    return total


def _holiday_map(engine: sa.engine.Engine, tenant: str, start: date, end: date) -> dict[date, bool]:
    with engine.connect() as c:
        rows = c.execute(sa.select(HOLIDAYS.c.day, HOLIDAYS.c.half).where(
            HOLIDAYS.c.tenant_id == tenant, HOLIDAYS.c.day >= start, HOLIDAYS.c.day <= end)).all()
    return {r.day: bool(r.half) for r in rows}


def _next_workday(after: date, workdays: str, hol: dict[date, bool]) -> date:
    d = after + timedelta(days=1)
    for _ in range(60):
        if str(d.isoweekday()) in workdays and not (d in hol and not hol[d]):
            return d
        d += timedelta(days=1)
    return d


# ------------------------------------------------------------------ kişi, yönetici, bakiye


def _person(engine: sa.engine.Engine, tenant: str, pid: str) -> Any:
    with engine.connect() as c:
        r = c.execute(sa.select(PT.PEOPLE).where(PT.PEOPLE.c.id == pid, PT.PEOPLE.c.tenant_id == tenant)).first()
    if r is None:
        raise HrError("Personel kaydı bulunamadı.", 404)
    return r


def _people(engine: sa.engine.Engine, tenant: str) -> list[Any]:
    with engine.connect() as c:
        return c.execute(sa.select(PT.PEOPLE).where(PT.PEOPLE.c.tenant_id == tenant)).all()


def _me(engine: sa.engine.Engine, tenant: str, user: str) -> Any:
    r = PT.person_of(engine, tenant, user)
    if r is None:
        raise HrError("Portal hesabınız bir özlük kaydına bağlı değil; İK kaydınızı bağlayınca izin isteyebilirsiniz.", 409)
    return r


def _manager_of(people: list[Any], person: Any) -> Optional[Any]:
    d = load(person.data_json, {})
    mgr = str(d.get("yonetici_id_no") or "")
    if not mgr or mgr == person.id_no:
        return None
    return next((p for p in people if p.id_no == mgr and p.durum == "Aktif"), None)


def balance(engine: sa.engine.Engine, tenant: str, person_id: str, type_key: str = "yillik") -> float:
    with engine.connect() as c:
        v = c.execute(sa.select(sa.func.coalesce(sa.func.sum(LEDGER.c.days), 0.0)).where(
            LEDGER.c.tenant_id == tenant, LEDGER.c.person_id == person_id, LEDGER.c.type_key == type_key)).scalar()
    return round(float(v or 0), 2)


def _years_completed(start: date, on: date) -> int:
    y = on.year - start.year - ((on.month, on.day) < (start.month, start.day))
    return max(0, y)


def _anniv(start: date, year: int) -> date:
    try:
        return start.replace(year=year)
    except ValueError:                         # 29 Şubat
        return date(year, 3, 1)


def entitlement(st: dict[str, Any], years: int, age: Optional[int]) -> float:
    """Tamamlanan hizmet yılı ve yaşa göre yıllık izin günü (ayardaki basamaklar)."""
    days = 0.0
    for t in st["leave.tiers"]:
        if years >= t["minYears"]:
            days = float(t["days"])
    ar = st["leave.ageRule"]
    if days and age is not None and (age <= ar["maxYoungAge"] or age >= ar["minOldAge"]):
        days = max(days, float(ar["days"]))
    return days


def _start_date(st: dict[str, Any], data: dict[str, Any]) -> Optional[date]:
    return PT.to_date(data.get(st["leave.startField"])) or PT.to_date(data.get("f_ise_giris_tarihi")) or PT.to_date(data.get("s_ise_giris_tarihi"))


def summary_for(engine: sa.engine.Engine, tenant: str, person: Any, today: Optional[date] = None) -> dict[str, Any]:
    today = today or date.today()
    st = settings(engine, tenant)
    data = load(person.data_json, {})
    y0 = date(today.year, 1, 1)
    with engine.connect() as c:
        led = c.execute(sa.select(LEDGER).where(LEDGER.c.tenant_id == tenant, LEDGER.c.person_id == person.id, LEDGER.c.type_key == "yillik")).all()
        pend = c.execute(sa.select(sa.func.coalesce(sa.func.sum(REQUESTS.c.days), 0.0)).where(
            REQUESTS.c.tenant_id == tenant, REQUESTS.c.person_id == person.id, REQUESTS.c.type_key == "yillik",
            REQUESTS.c.status.in_(OPEN))).scalar()
    bal = round(sum(float(x.days) for x in led), 2)
    used_year = round(-sum(float(x.days) for x in led if x.reason == "kullanim" and x.on_date >= y0)
                      - sum(float(x.days) for x in led if x.reason == "iade" and x.on_date >= y0), 2)
    accrued_year = round(sum(float(x.days) for x in led if x.reason == "hakedis" and x.on_date >= y0), 2)
    start = _start_date(st, data)
    nxt = None
    if start:
        n = _years_completed(start, today)
        ad = _anniv(start, start.year + n + 1)
        birth = PT.to_date(data.get("dogum_tarihi"))
        age = _years_completed(birth, ad) if birth else None
        nxt = {"date": ad.isoformat(), "days": entitlement(st, n + 1, age), "years": n + 1}
    return {"balance": bal, "pending": round(float(pend or 0), 2), "available": round(bal - float(pend or 0), 2),
            "usedYear": used_year, "accruedYear": accrued_year, "nextAccrual": nxt, "startDate": start.isoformat() if start else None,
            "hasOpening": any(x.reason == "acilis" for x in led)}


# ------------------------------------------------------------------ talepler


def _req_out(r: Any, *, show_type: bool, type_label: dict[str, str], people_by_id: dict[str, Any]) -> dict[str, Any]:
    p = people_by_id.get(r.person_id)
    dep = people_by_id.get(r.deputy_person_id) if r.deputy_person_id else None
    return {"id": r.id, "personId": r.person_id, "adSoyad": p.ad_soyad if p else "—",
            "departman": load(p.data_json, {}).get("departman") if p else None,
            "typeKey": r.type_key if show_type else None, "typeLabel": type_label.get(r.type_key, r.type_key) if show_type else "İzin",
            "start": r.start.isoformat(), "end": r.end.isoformat(), "half": r.half or "", "days": float(r.days),
            "deputy": dep.ad_soyad if dep else None, "deputyId": r.deputy_person_id, "contact": r.contact if show_type else None,
            "note": r.note if show_type else None, "status": r.status, "statusLabel": STATUS.get(r.status, r.status),
            "needsHr": bool(r.needs_hr), "managerDecidedBy": r.manager_decided_by, "decidedBy": r.decided_by,
            "decidedAt": iso(r.decided_at), "rejectReason": r.reject_reason, "createdAt": iso(r.created_at)}


def _ctx(engine: sa.engine.Engine, tenant: str) -> tuple[dict[str, dict[str, Any]], dict[str, Any], list[Any]]:
    ts = {t["key"]: t for t in types(engine, tenant)}
    people = _people(engine, tenant)
    return ts, {p.id: p for p in people}, people


def calc(engine: sa.engine.Engine, tenant: str, person: Any, body: dict[str, Any], *, exclude: str = "") -> dict[str, Any]:
    """Talebi göndermeden önce gün hesabı ve uyarılar. Hata değil uyarı: kişi görür, karar İK/yöneticinin."""
    ts = {t["key"]: t for t in types(engine, tenant)}
    t = ts.get(str(body.get("type") or ""))
    if t is None or not t["active"]:
        raise HrError("İzin türünü seçin.")
    start = H.parse_date(body.get("start"), "Başlangıç")
    end = H.parse_date(body.get("end"), "Bitiş") or start
    if start is None:
        raise HrError("Başlangıç tarihini seçin.")
    if end < start:
        raise HrError("Bitiş, başlangıçtan önce olamaz.")
    if (end - start).days > 400:
        raise HrError("Bir talep en çok bir yılı kapsar.")
    half = str(body.get("half") or "")
    if half and (half not in HALF or not t["halfDay"] or start != end):
        raise HrError("Yarım gün yalnız tek günlük talepte ve yarım güne izin veren türde seçilir.")
    data = load(person.data_json, {})
    cal = _calendar_for(calendars(engine, tenant), data)
    hol = _holiday_map(engine, tenant, start, end)
    days = count_days(start, end, mode=t["countMode"], workdays=cal["days"], hol=hol, half=half)
    warnings: list[str] = []
    errors: list[str] = []
    if days <= 0:
        errors.append("Seçilen aralıkta iş günü yok (hafta sonu ya da resmî tatil).")
    if t["maxRequest"] and days > t["maxRequest"]:
        errors.append(f"Bu türde bir talep en çok {t['maxRequest']:g} gün olabilir.")
    with engine.connect() as c:
        overlap = c.execute(sa.select(REQUESTS.c.id).where(
            REQUESTS.c.tenant_id == tenant, REQUESTS.c.person_id == person.id, REQUESTS.c.id != exclude,
            REQUESTS.c.status.in_(OPEN + ("onaylandi",)), REQUESTS.c.start <= end, REQUESTS.c.end >= start)).first()
        year_used = c.execute(sa.select(sa.func.coalesce(sa.func.sum(REQUESTS.c.days), 0.0)).where(
            REQUESTS.c.tenant_id == tenant, REQUESTS.c.person_id == person.id, REQUESTS.c.type_key == t["key"],
            REQUESTS.c.status.in_(OPEN + ("onaylandi",)), REQUESTS.c.start >= date(start.year, 1, 1),
            REQUESTS.c.start <= date(start.year, 12, 31))).scalar()
    if overlap:
        errors.append("Bu tarihlerle çakışan bekleyen ya da onaylı bir izniniz var.")
    if t["maxYear"] and float(year_used or 0) + days > t["maxYear"]:
        errors.append(f"Bu türde yılda en çok {t['maxYear']:g} gün; {start.year} için kalan {max(0.0, t['maxYear'] - float(year_used or 0)):g} gün.")
    avail = None
    if t["fromBalance"]:
        s = summary_for(engine, tenant, person)
        avail = s["available"]
        if avail - days < 0:
            if settings(engine, tenant)["leave.advance"]:
                warnings.append(f"Kullanılabilir bakiyeniz {avail:g} gün; talep onaylanırsa bakiye eksiye düşer (avans izin).")
            else:
                errors.append(f"Kullanılabilir yıllık izin bakiyeniz {avail:g} gün; talep {days:g} gün.")
        if t["key"] == "yillik" and 0 < days < 10:
            warnings.append("Yıllık izin bölünerek kullanılırsa parçalardan biri en az 10 gün olmalı (4857 md. 56); İK bu yılın parçalarına bakar.")
    if t["needsDoc"]:
        warnings.append("Bu türde belge gerekir (ör. rapor); talebi açtıktan sonra belgeyi yükleyin.")
    if not t["approved"]:
        errors.append("Bu izin türünün gün sayısı İK tarafından henüz onaylanmadı; şimdilik talep açılamaz.")
    return {"days": days, "calendar": cal["name"], "holidays": [{"day": d.isoformat(), "half": h} for d, h in sorted(hol.items())],
            "available": avail, "warnings": warnings, "errors": errors, "type": t}


def _event(c: Any, tenant: str, rid: str, actor: str, action: str, note: Optional[str] = None) -> None:
    c.execute(EVENTS.insert().values(tenant_id=tenant, request_id=rid, at=now(), actor=actor[:120], action=action[:24], note=note))


def create(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    me = _me(engine, tenant, user)
    if me.durum != "Aktif":
        raise HrError("Pasif kayıtla izin talebi açılamaz.")
    r = calc(engine, tenant, me, body)
    if r["errors"]:
        raise HrError(" ".join(r["errors"]))
    st = settings(engine, tenant)
    ts, by_id, people = _ctx(engine, tenant)
    mgr = _manager_of(people, me)
    deputy = str(body.get("deputyId") or "") or None
    if deputy and (deputy == me.id or deputy not in by_id or by_id[deputy].durum != "Aktif"):
        raise HrError("Vekil geçersiz.")
    rid = new_id("izn")
    status = "yonetici" if mgr else "ik"
    with engine.begin() as c:
        c.execute(REQUESTS.insert().values(
            id=rid, tenant_id=tenant, person_id=me.id, username=user, type_key=r["type"]["key"], start=H.parse_date(body.get("start"), "Başlangıç"),
            end=H.parse_date(body.get("end"), "Bitiş") or H.parse_date(body.get("start"), "Başlangıç"), half=str(body.get("half") or "") or None,
            days=r["days"], deputy_person_id=deputy, contact=clean(body.get("contact"), 120) or None,
            note=str(body.get("note") or "").strip()[:2000] or None, status=status, manager_id_no=mgr.id_no if mgr else None,
            needs_hr=(st["leave.flow"] == "yonetici_ik") or mgr is None, created_by=user, created_at=now(), updated_at=now()))
        _event(c, tenant, rid, user, "acildi")
    return {"id": rid, "status": status, "manager": mgr, "person": me, "type": r["type"], "days": r["days"]}


def _load_req(engine: sa.engine.Engine, tenant: str, rid: str) -> Any:
    with engine.connect() as c:
        r = c.execute(sa.select(REQUESTS).where(REQUESTS.c.id == rid, REQUESTS.c.tenant_id == tenant)).first()
    if r is None:
        raise HrError("İzin talebi bulunamadı.", 404)
    return r


def is_manager_of(engine: sa.engine.Engine, tenant: str, user: str, req: Any) -> bool:
    me = PT.person_of(engine, tenant, user)
    return bool(me is not None and req.manager_id_no and me.id_no == req.manager_id_no and me.id != req.person_id)


def decide(engine: sa.engine.Engine, tenant: str, user: str, can_hr: bool, rid: str, body: dict[str, Any]) -> dict[str, Any]:
    """Yönetici ya da İK kararı. Dönen: güncel satır + olay (bildirim için)."""
    req = _load_req(engine, tenant, rid)
    action = str(body.get("action") or "")
    if action not in ("onayla", "reddet"):
        raise HrError("Karar geçersiz.")
    reason = str(body.get("reason") or "").strip()[:1000] or None
    if action == "reddet" and not reason:
        raise HrError("Ret gerekçesi yazın; çalışan bunu görür.")
    me = PT.person_of(engine, tenant, user)
    if me is not None and me.id == req.person_id:
        raise HrError("Kendi izin talebinizi onaylayamazsınız.", 403)
    mgr = is_manager_of(engine, tenant, user, req)
    if req.status == "yonetici":
        if not (mgr or can_hr):
            raise HrError("Bu talebi yalnız çalışanın yöneticisi ya da İK karara bağlar.", 403)
    elif req.status == "ik":
        if not can_hr:
            raise HrError("Bu talep İK onayı bekliyor.", 403)
    else:
        raise HrError("Talep karara bağlanmış.")
    vals: dict[str, Any] = {"updated_at": now()}
    ev = "reddedildi"
    if action == "reddet":
        vals.update(status="reddedildi", decided_by=user, decided_at=now(), reject_reason=reason)
    elif req.status == "yonetici" and req.needs_hr and not can_hr:
        vals.update(status="ik", manager_decided_by=user, manager_decided_at=now())
        ev = "yonetici_onayladi"
    else:
        if req.status == "yonetici":
            vals.update(manager_decided_by=user, manager_decided_at=now())
        vals.update(status="onaylandi", decided_by=user, decided_at=now())
        ev = "onaylandi"
    ts = {t["key"]: t for t in types(engine, tenant)}
    with engine.begin() as c:
        # Yarış: iki onaylayan aynı anda basarsa yalnız ilki geçer.
        n = c.execute(REQUESTS.update().where(REQUESTS.c.id == rid, REQUESTS.c.status == req.status).values(**vals)).rowcount
        if not n:
            raise HrError("Talep bu arada değişti; sayfayı yenileyin.", 409)
        _event(c, tenant, rid, user, ev, reason)
        if ev == "onaylandi" and ts.get(req.type_key, {}).get("fromBalance"):
            c.execute(LEDGER.insert().values(tenant_id=tenant, person_id=req.person_id, type_key=req.type_key, days=-float(req.days),
                                             reason="kullanim", request_id=rid, on_date=req.start, actor=user, at=now()))
    return {"event": ev, "request": _load_req(engine, tenant, rid)}


def cancel(engine: sa.engine.Engine, tenant: str, user: str, can_hr: bool, rid: str, today: Optional[date] = None) -> dict[str, Any]:
    """Bekleyen talebi kişi geri alır; onaylı izni kişi başlamadan, İK her zaman iptal eder (bakiye iade edilir)."""
    today = today or date.today()
    req = _load_req(engine, tenant, rid)
    me = PT.person_of(engine, tenant, user)
    own = req.username == user or (me is not None and me.id == req.person_id)
    if not (own or can_hr):
        raise HrError("Bu talebi yalnız sahibi ya da İK geri alır.", 403)
    ts = {t["key"]: t for t in types(engine, tenant)}
    if req.status in OPEN:
        new, ev = "geri_alindi", "geri_alindi"
    elif req.status == "onaylandi":
        if not can_hr and req.start <= today:
            raise HrError("Başlamış izni İK iptal eder.")
        new, ev = "iptal", "iptal"
    else:
        raise HrError("Bu talep geri alınamaz.")
    with engine.begin() as c:
        n = c.execute(REQUESTS.update().where(REQUESTS.c.id == rid, REQUESTS.c.status == req.status).values(
            status=new, updated_at=now())).rowcount
        if not n:
            raise HrError("Talep bu arada değişti; sayfayı yenileyin.", 409)
        _event(c, tenant, rid, user, ev)
        if new == "iptal" and ts.get(req.type_key, {}).get("fromBalance"):
            c.execute(LEDGER.insert().values(tenant_id=tenant, person_id=req.person_id, type_key=req.type_key, days=float(req.days),
                                             reason="iade", request_id=rid, on_date=today, actor=user, at=now()))
    return {"event": ev, "request": _load_req(engine, tenant, rid)}


def my_view(engine: sa.engine.Engine, tenant: str, user: str) -> dict[str, Any]:
    me = PT.person_of(engine, tenant, user)
    ts = types(engine, tenant)
    usable = [t for t in ts if t["active"] and t["approved"]]
    today = date.today()
    out: dict[str, Any] = {"linked": me is not None, "types": usable, "pendingTypes": [t["label"] for t in ts if t["active"] and not t["approved"]],
                           "halfOptions": HALF, "status": STATUS,
                           "holidays": [h for h in holidays(engine, tenant) if h["day"] >= today.isoformat()][:12]}
    if me is None:
        return out
    tsd, by_id, people = _ctx(engine, tenant)
    with engine.connect() as c:
        rows = c.execute(sa.select(REQUESTS).where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.person_id == me.id)
                         .order_by(REQUESTS.c.start.desc())).all()
        files = c.execute(sa.select(FILES.c.id, FILES.c.request_id, FILES.c.filename, FILES.c.size).where(
            FILES.c.tenant_id == tenant, FILES.c.request_id.in_([r.id for r in rows] or [""]))).all()
    labels = {k: t["label"] for k, t in tsd.items()}
    reqs = []
    for r in rows:
        o = _req_out(r, show_type=True, type_label=labels, people_by_id=by_id)
        o["files"] = [{"id": f.id, "filename": f.filename, "size": f.size} for f in files if f.request_id == r.id]
        o["needsDoc"] = bool(tsd.get(r.type_key, {}).get("needsDoc"))
        reqs.append(o)
    mgr = _manager_of(people, me)
    out.update(summary=summary_for(engine, tenant, me), requests=reqs, manager=mgr.ad_soyad if mgr else None,
               deputies=[{"id": p.id, "adSoyad": p.ad_soyad, "departman": load(p.data_json, {}).get("departman")}
                         for p in sorted(people, key=lambda x: x.ad_soyad.casefold()) if p.durum == "Aktif" and p.id != me.id],
               calendar=_calendar_for(calendars(engine, tenant), load(me.data_json, {}))["name"])
    return out


# ------------------------------------------------------------------ ekip


def team_view(engine: sa.engine.Engine, tenant: str, user: str, month: str, can_sensitive: bool) -> dict[str, Any]:
    me = PT.person_of(engine, tenant, user)
    try:
        m0 = date.fromisoformat(f"{month}-01") if month else date.today().replace(day=1)
    except ValueError:
        raise HrError("Ay YYYY-AA biçiminde olmalı.") from None
    m1 = (m0.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    tsd, by_id, people = _ctx(engine, tenant)
    if me is None:
        return {"linked": False, "members": [], "queue": [], "requests": [], "month": m0.isoformat()[:7], "days": [], "conflicts": []}
    members = [p for p in people if p.durum == "Aktif" and load(p.data_json, {}).get("yonetici_id_no") == me.id_no and p.id != me.id]
    ids = [p.id for p in members]
    with engine.connect() as c:
        queue = c.execute(sa.select(REQUESTS).where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.manager_id_no == me.id_no,
                                                    REQUESTS.c.status == "yonetici").order_by(REQUESTS.c.created_at)).all()
        rows = c.execute(sa.select(REQUESTS).where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.person_id.in_(ids or [""]),
                                                   REQUESTS.c.status.in_(OPEN + ("onaylandi",)), REQUESTS.c.start <= m1,
                                                   REQUESTS.c.end >= m0)).all()
    labels = {k: t["label"] for k, t in tsd.items()}

    def show(r: Any) -> bool:
        return can_sensitive or not tsd.get(r.type_key, {}).get("sensitive")

    reqs = [_req_out(r, show_type=show(r), type_label=labels, people_by_id=by_id) for r in rows]
    st = settings(engine, tenant)
    hol = _holiday_map(engine, tenant, m0, m1)
    days, conflicts = [], []
    d = m0
    while d <= m1:
        off = [r for r in rows if r.start <= d <= r.end]
        days.append({"day": d.isoformat(), "weekday": d.isoweekday(), "holiday": d in hol, "off": [r.person_id for r in off]})
        if members and str(d.isoweekday()) in "12345" and d not in hol and len({r.person_id for r in off}) * 100 >= st["leave.conflictPct"] * len(members) and len(off) > 1:
            conflicts.append({"day": d.isoformat(), "off": len({r.person_id for r in off}), "team": len(members)})
        d += timedelta(days=1)
    return {"linked": True, "month": m0.isoformat()[:7], "conflictPct": st["leave.conflictPct"],
            "members": [{"id": p.id, "adSoyad": p.ad_soyad, "unvan": load(p.data_json, {}).get("unvan"),
                         "balance": summary_for(engine, tenant, p)["available"]} for p in sorted(members, key=lambda x: x.ad_soyad.casefold())],
            "queue": [_req_out(r, show_type=show(r), type_label=labels, people_by_id=by_id) for r in queue],
            "requests": reqs, "days": days, "conflicts": conflicts}


def team_waiting(engine: sa.engine.Engine, tenant: str, manager_id_no: str) -> int:
    ensure(engine)
    with engine.connect() as c:
        return int(c.execute(sa.select(sa.func.count()).select_from(REQUESTS).where(
            REQUESTS.c.tenant_id == tenant, REQUESTS.c.manager_id_no == manager_id_no, REQUESTS.c.status == "yonetici")).scalar() or 0)


def on_leave_today(engine: sa.engine.Engine, tenant: str, today: Optional[date] = None) -> list[dict[str, Any]]:
    """Bugün onaylı izinde olan aktif kişiler: ad, departman, dönüş günü (tür yok)."""
    today = today or date.today()
    ensure(engine)
    with engine.connect() as c:
        rows = c.execute(sa.select(REQUESTS.c.person_id, REQUESTS.c.end).where(
            REQUESTS.c.tenant_id == tenant, REQUESTS.c.status == "onaylandi", REQUESTS.c.start <= today, REQUESTS.c.end >= today)).all()
    if not rows:
        return []
    by_id = {p.id: p for p in _people(engine, tenant)}
    cals = calendars(engine, tenant)
    hol = _holiday_map(engine, tenant, today, today + timedelta(days=90))
    out = {}
    for r in rows:
        p = by_id.get(r.person_id)
        if p is None or p.durum != "Aktif":
            continue
        d = load(p.data_json, {})
        back = _next_workday(r.end, _calendar_for(cals, d)["days"], hol)
        if r.person_id not in out or out[r.person_id]["back"] < back.isoformat():
            out[r.person_id] = {"id": p.id, "adSoyad": p.ad_soyad, "departman": d.get("departman"), "back": back.isoformat()}
    return sorted(out.values(), key=lambda x: x["adSoyad"].casefold())


# ------------------------------------------------------------------ İK yönetimi


def admin_requests(engine: sa.engine.Engine, tenant: str, can_sensitive: bool, *, status: str = "acik", type_key: str = "",
                   month: str = "", q: str = "") -> list[dict[str, Any]]:
    tsd, by_id, _ = _ctx(engine, tenant)
    stmt = sa.select(REQUESTS).where(REQUESTS.c.tenant_id == tenant)
    if status == "acik":
        stmt = stmt.where(REQUESTS.c.status.in_(OPEN))
    elif status and status != "hepsi":
        stmt = stmt.where(REQUESTS.c.status == status)
    if type_key:
        stmt = stmt.where(REQUESTS.c.type_key == type_key)
    if month:
        try:
            m0 = date.fromisoformat(f"{month}-01")
        except ValueError:
            raise HrError("Ay YYYY-AA biçiminde olmalı.") from None
        m1 = (m0.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
        stmt = stmt.where(REQUESTS.c.start <= m1, REQUESTS.c.end >= m0)
    with engine.connect() as c:
        rows = c.execute(stmt.order_by(REQUESTS.c.start.desc())).all()
    labels = {k: t["label"] for k, t in tsd.items()}
    qq = (q or "").strip().casefold()
    out = []
    for r in rows:
        o = _req_out(r, show_type=can_sensitive or not tsd.get(r.type_key, {}).get("sensitive"), type_label=labels, people_by_id=by_id)
        if qq and qq not in f"{o['adSoyad']} {o['departman'] or ''}".casefold():
            continue
        out.append(o)
    return out


def balances(engine: sa.engine.Engine, tenant: str) -> list[dict[str, Any]]:
    out = []
    for p in _people(engine, tenant):
        if p.durum != "Aktif":
            continue
        s = summary_for(engine, tenant, p)
        d = load(p.data_json, {})
        out.append({"id": p.id, "idNo": p.id_no, "adSoyad": p.ad_soyad, "departman": d.get("departman"), **s})
    return sorted(out, key=lambda x: x["adSoyad"].casefold())


def ledger(engine: sa.engine.Engine, tenant: str, person_id: str) -> list[dict[str, Any]]:
    ensure(engine)
    with engine.connect() as c:
        rows = c.execute(sa.select(LEDGER).where(LEDGER.c.tenant_id == tenant, LEDGER.c.person_id == person_id)
                         .order_by(LEDGER.c.on_date.desc(), LEDGER.c.id.desc())).all()
    return [{"id": r.id, "typeKey": r.type_key, "days": float(r.days), "reason": r.reason, "reasonLabel": REASONS.get(r.reason, r.reason),
             "requestId": r.request_id, "onDate": r.on_date.isoformat(), "note": r.note, "actor": r.actor, "at": iso(r.at)} for r in rows]


def adjust(engine: sa.engine.Engine, tenant: str, actor: str, body: dict[str, Any]) -> dict[str, Any]:
    pid = str(body.get("personId") or "")
    _person(engine, tenant, pid)
    try:
        days = float(body.get("days"))
    except (TypeError, ValueError):
        raise HrError("Gün sayı olmalı (eksi değer bakiyeden düşer).") from None
    if days == 0 or abs(days) > 400 or (days * 2) != int(days * 2):
        raise HrError("Gün sıfır olamaz; yarım gün katları girin (ör. 1,5).")
    note = str(body.get("note") or "").strip()[:500]
    if not note:
        raise HrError("Düzeltmenin nedenini yazın (defterde görünür).")
    on = H.parse_date(body.get("onDate"), "Tarih") or date.today()
    with engine.begin() as c:
        c.execute(LEDGER.insert().values(tenant_id=tenant, person_id=pid, type_key="yillik", days=days, reason="duzeltme", on_date=on,
                                         note=note, actor=actor, at=now()))
    return {"balance": balance(engine, tenant, pid)}


def import_opening(engine: sa.engine.Engine, tenant: str, actor: str, data: bytes, *, apply: bool) -> dict[str, Any]:
    """Açılış bakiyesi Excel'i: başlık `id_no`, `gun` (zorunlu), `tarih`, `aciklama`. Kişi başına tek açılış satırı:
    varsa yenisi yazılmaz (hata). `apply=False` hiçbir şey yazmaz."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise HrError("Dosya Excel (.xlsx) olarak okunamadı.") from None
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    head = [str(h or "").strip().casefold() for h in (next(it, None) or [])]
    idx = {k: head.index(k) for k in ("id_no", "gun", "tarih", "aciklama") if k in head}
    if "id_no" not in idx or "gun" not in idx:
        raise HrError("İlk satırda «id_no» ve «gun» başlıkları olmalı (isteğe bağlı: «tarih», «aciklama»).")
    people = {p.id_no: p for p in _people(engine, tenant)}
    with engine.connect() as c:
        has_open = {r.person_id for r in c.execute(sa.select(LEDGER.c.person_id).where(
            LEDGER.c.tenant_id == tenant, LEDGER.c.reason == "acilis")).all()}
    rows, seen = [], set()
    for n, raw in enumerate(it, start=2):
        if not any(v not in (None, "") for v in raw):
            continue
        get = lambda k: raw[idx[k]] if k in idx and idx[k] < len(raw) else None  # noqa: E731
        idn = clean(get("id_no"), 60)
        if isinstance(get("id_no"), float) and get("id_no").is_integer():
            idn = str(int(get("id_no")))
        errs = []
        p = people.get(idn)
        if p is None:
            errs.append(f"«{idn}» personel numarası özlük kaydında yok.")
        elif p.id in has_open:
            errs.append("Bu kişinin açılış bakiyesi zaten var; farkı İK düzeltmesiyle girin.")
        if idn in seen:
            errs.append("Dosyada iki kez var.")
        seen.add(idn)
        try:
            days = float(str(get("gun")).replace(",", "."))
            if (days * 2) != int(days * 2):
                raise ValueError
        except (TypeError, ValueError):
            days = 0.0
            errs.append("Gün sayı olmalı (yarım gün katları).")
        on = PT.to_date(get("tarih")) or date.today()
        rows.append({"row": n, "idNo": idn, "adSoyad": p.ad_soyad if p else None, "days": days, "onDate": on.isoformat(),
                     "note": clean(get("aciklama"), 300) or None, "errors": errs, "_pid": p.id if p else None})
    bad = sum(1 for r in rows if r["errors"])
    applied = 0
    if apply:
        if bad:
            raise HrError(f"{bad} satırda hata var; hiçbir satır yazılmadı.")
        with engine.begin() as c:
            for r in rows:
                c.execute(LEDGER.insert().values(tenant_id=tenant, person_id=r["_pid"], type_key="yillik", days=r["days"], reason="acilis",
                                                 on_date=date.fromisoformat(r["onDate"]), note=r["note"] or "Açılış bakiyesi (Excel)",
                                                 actor=actor, at=now()))
                applied += 1
    for r in rows:
        r.pop("_pid", None)
    return {"rows": rows, "counts": {"ok": len(rows) - bad, "error": bad}, "applied": applied}


def payroll_xlsx(engine: sa.engine.Engine, tenant: str, month: str, can_sensitive: bool) -> bytes:
    """Ayın bordroyu etkileyen onaylı izinleri (bordro işaretli türler): kişi, tür, ay içindeki gün."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    try:
        m0 = date.fromisoformat(f"{month}-01")
    except ValueError:
        raise HrError("Ay YYYY-AA biçiminde olmalı.") from None
    m1 = (m0.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    tsd, by_id, _ = _ctx(engine, tenant)
    keys = [k for k, t in tsd.items() if t["payroll"]]
    with engine.connect() as c:
        rows = c.execute(sa.select(REQUESTS).where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.status == "onaylandi",
                                                   REQUESTS.c.type_key.in_(keys or [""]), REQUESTS.c.start <= m1,
                                                   REQUESTS.c.end >= m0).order_by(REQUESTS.c.start)).all()
    cals = calendars(engine, tenant)
    hol = _holiday_map(engine, tenant, m0, m1)
    wb = Workbook()
    ws = wb.active
    ws.title = f"izin-{month}"
    ws.append(["id_no", "ad_soyad", "departman", "izin_turu", "ucretli", "baslangic", "bitis", "ay_icindeki_gun", "talep_toplam_gun"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        p = by_id.get(r.person_id)
        d = load(p.data_json, {}) if p else {}
        t = tsd.get(r.type_key, {})
        a, b = max(r.start, m0), min(r.end, m1)
        inside = count_days(a, b, mode=t.get("countMode", "is_gunu"), workdays=_calendar_for(cals, d)["days"], hol=hol,
                            half=r.half or "" if r.start == r.end else "")
        label = t.get("label", r.type_key) if (can_sensitive or not t.get("sensitive")) else "Hassas izin (ayrıntı İK'da)"
        ws.append([p.id_no if p else "", p.ad_soyad if p else "", d.get("departman"), label, "evet" if t.get("paid") else "hayır",
                   a, b, inside, float(r.days)])
    for col in ("F", "G"):
        for cell in ws[col][1:]:
            cell.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stats(engine: sa.engine.Engine, tenant: str, today: Optional[date] = None) -> dict[str, Any]:
    """İK ana sayfası için: bugün izinde, onay bekleyen, birikmiş (30+ gün) izin, 30 gün içinde hakediş."""
    today = today or date.today()
    ensure(engine)
    with engine.connect() as c:
        waiting = c.execute(sa.select(sa.func.count()).select_from(REQUESTS).where(
            REQUESTS.c.tenant_id == tenant, REQUESTS.c.status.in_(OPEN))).scalar() or 0
    bal = balances(engine, tenant)
    return {"onLeaveToday": len(on_leave_today(engine, tenant, today)), "waiting": int(waiting),
            "highBalance": sorted(({"id": b["id"], "adSoyad": b["adSoyad"], "balance": b["balance"]} for b in bal if b["balance"] >= 30),
                                  key=lambda x: -x["balance"])[:10],
            "accrualSoon": sorted(({"id": b["id"], "adSoyad": b["adSoyad"], "date": b["nextAccrual"]["date"], "days": b["nextAccrual"]["days"]}
                                   for b in bal if b["nextAccrual"] and b["nextAccrual"]["date"] <= (today + timedelta(days=30)).isoformat()),
                                  key=lambda x: x["date"])[:10]}


# ------------------------------------------------------------------ belgeler


def add_file(engine: sa.engine.Engine, tenant: str, user: str, can_hr_sensitive: bool, rid: str, filename: str, data: bytes,
             max_mb: int) -> dict[str, Any]:
    req = _load_req(engine, tenant, rid)
    me = PT.person_of(engine, tenant, user)
    if not ((me is not None and me.id == req.person_id) or can_hr_sensitive):
        raise HrError("Belgeyi talebin sahibi ya da hassas yetkili İK yükler.", 403)
    if not data:
        raise HrError("Dosya boş.")
    if len(data) > max_mb * 1024 * 1024:
        raise HrError(f"Dosya {max_mb} MB sınırını aşıyor.", 413)
    name = clean(filename, 300) or "belge"
    mime = PT._mime(name)
    fid = new_id("izb")
    with engine.begin() as c:
        c.execute(FILES.insert().values(id=fid, tenant_id=tenant, request_id=rid, filename=name, mime=mime, size=len(data), blob=data,
                                        uploaded_by=user, uploaded_at=now()))
        _event(c, tenant, rid, user, "belge")
    return {"id": fid, "filename": name, "size": len(data)}


def file_blob(engine: sa.engine.Engine, tenant: str, user: str, can_hr_sensitive: bool, can_hr: bool, rid: str,
              fid: str) -> tuple[bytes, str, str, str]:
    req = _load_req(engine, tenant, rid)
    me = PT.person_of(engine, tenant, user)
    tsd = {t["key"]: t for t in types(engine, tenant)}
    sensitive = bool(tsd.get(req.type_key, {}).get("sensitive"))
    own = me is not None and me.id == req.person_id
    if not (own or can_hr_sensitive or (can_hr and not sensitive)):
        raise HrError("Belgeyi görme yetkiniz yok.", 403)
    with engine.connect() as c:
        r = c.execute(sa.select(FILES).where(FILES.c.id == fid, FILES.c.request_id == rid, FILES.c.tenant_id == tenant)).first()
    if r is None:
        raise HrError("Belge bulunamadı.", 404)
    return bytes(r.blob), r.filename, r.mime or "application/octet-stream", req.person_id


# ------------------------------------------------------------------ gece işi


def _workdays_between(a: date, b: date) -> int:
    n, d = 0, a
    while d < b:
        d += timedelta(days=1)
        if d.isoweekday() <= 5:
            n += 1
    return n


def run_due(engine: sa.engine.Engine, tenant: str, today: Optional[date] = None) -> dict[str, Any]:
    """Hakediş (başlangıç tarihinden sonraki yıldönümleri, tekrar yazmaz) + onay hatırlatması adayları."""
    today = today or date.today()
    _seed(engine, tenant)
    st = settings(engine, tenant)
    since = mark_accrual_start(engine, tenant, today)
    written = []
    with engine.connect() as c:
        have = {(r.person_id, r.period) for r in c.execute(sa.select(LEDGER.c.person_id, LEDGER.c.period).where(
            LEDGER.c.tenant_id == tenant, LEDGER.c.reason == "hakedis")).all()}
    rows = []
    for p in _people(engine, tenant):
        if p.durum != "Aktif":
            continue
        d = load(p.data_json, {})
        start = _start_date(st, d)
        if not start:
            continue
        birth = PT.to_date(d.get("dogum_tarihi"))
        for y in range(max(start.year + 1, since.year), today.year + 1):
            ad = _anniv(start, y)
            if not (since <= ad <= today) or (p.id, ad.isoformat()) in have:
                continue
            years = _years_completed(start, ad)
            days = entitlement(st, years, _years_completed(birth, ad) if birth else None)
            if days <= 0:
                continue
            rows.append(dict(tenant_id=tenant, person_id=p.id, type_key="yillik", days=days, reason="hakedis", period=ad.isoformat(),
                             on_date=ad, note=f"{years}. hizmet yılı", actor="ZEKİ AI", at=now()))
            written.append({"personId": p.id, "date": ad.isoformat(), "days": days, "years": years})
    if rows:
        with engine.begin() as c:
            c.execute(LEDGER.insert(), rows)
    with engine.connect() as c:
        waiting = c.execute(sa.select(REQUESTS).where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.status.in_(OPEN),
                                                      sa.or_(REQUESTS.c.last_reminded.is_(None), REQUESTS.c.last_reminded < today))).all()
    due = [r for r in waiting if _workdays_between(r.created_at.date() if isinstance(r.created_at, datetime) else r.created_at, today)
           >= int(st["leave.reminderWorkdays"])]
    return {"accrued": written, "remind": due, "accrualFrom": since.isoformat()}


def mark_reminded(engine: sa.engine.Engine, tenant: str, ids: Iterable[str], today: date) -> None:
    ids = list(ids)
    if not ids:
        return
    with engine.begin() as c:
        c.execute(REQUESTS.update().where(REQUESTS.c.tenant_id == tenant, REQUESTS.c.id.in_(ids)).values(last_reminded=today))


# ------------------------------------------------------------------ e-posta metinleri (hassas veri yok)


def _span(start: date, end: date, days: float) -> str:
    fmt = lambda d: f"{d.day:02d}.{d.month:02d}.{d.year}"  # noqa: E731
    rng = fmt(start) if start == end else f"{fmt(start)} – {fmt(end)}"
    return f"{rng} ({days:g} gün)"


def mail_text(event: str, req: Any, person_name: str, type_label: str, sensitive: bool, reason: Optional[str] = None) -> tuple[str, str]:
    """(konu, gövde). Hassas türde tür adı yazılmaz; `{link}` kuyrukta portal adresiyle değişir."""
    what = "izin" if sensitive else type_label.lower()
    span = _span(req.start, req.end, float(req.days))
    if event == "yeni":
        return (f"İzin talebi onayınızı bekliyor: {person_name}",
                f"{person_name} {span} için {what} talebinde bulundu.\n\nKarar vermek için: {{link}}/ik/izin/ekip")
    if event == "ik":
        return (f"İzin talebi İK onayı bekliyor: {person_name}",
                f"{person_name} için {span} {what} talebi İK onayı bekliyor.\n\nİK yönetimi › İzinler: {{link}}/ik/yonetim?sekme=izinler")
    if event == "onaylandi":
        return ("İzin talebiniz onaylandı", f"{span} {what} talebiniz onaylandı.\n\nİzinlerim: {{link}}/ik/izin")
    if event == "reddedildi":
        return ("İzin talebiniz reddedildi", f"{span} {what} talebiniz reddedildi.\nGerekçe: {reason or '—'}\n\nİzinlerim: {{link}}/ik/izin")
    if event == "vekil":
        return (f"Vekâlet: {person_name} izinde", f"{person_name} {span} izinde; vekil olarak siz belirlendiniz.\n\n{{link}}/ik")
    if event == "iptal":
        return (f"Onaylı izin iptal edildi: {person_name}", f"{person_name} için {span} onaylı izin iptal edildi.\n\n{{link}}/ik/izin/ekip")
    if event == "hatirlatma":
        return ("Onay bekleyen izin talepleri", f"Onayınızı bekleyen izin talepleri var.\n\n{{link}}/ik/izin/ekip")
    return ("İzin talebi", f"{span}\n\n{{link}}/ik/izin")
