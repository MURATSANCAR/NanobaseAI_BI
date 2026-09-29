"""İK personel portalı: özlük kaydı (personel kartı + belgeler), İK ana sayfası ve çalışanın kendi ekranları.

Eski «Timaş Personel Portal» (AppSheet) uygulamasının karşılığı; alan listesi İK'nın `personel_data.xlsx` dosyasından
(`hr_portal_fields.json`, 69 alan). İK-0 çalışan kaydından (`semantic_hr_employees`, CRM ∩ AD) ayrıdır: o kayıt
işe alım/performans/eğitim modüllerinin kimlik listesidir; bu kayıt İK'nın elle tuttuğu özlük dosyasıdır.

- **Alanlar** (`semantic_hr_person_fields`): ad, tür (metin, uzun metin, seçim, tarih, e-posta, telefon, sayı,
  dosya), seçenekler, zorunlu, üç sayfada görünürlük (yönetici paneli, profilim, personel rehberi), grup, sıra,
  hassas. İlk açılışta Excel'deki listeyle dolar; İK ekrandan değiştirir, yeni alan ekler. Excel'deki alan silinmez,
  kapatılır.
- **Personel** (`semantic_hr_people`): personel no (`id_no`) tenant içinde tektir; bütün alan değerleri `data_json`'da,
  süzgeç için birkaç alan kolonda. Belgeler `semantic_hr_people_files` (bayt veritabanında; aday dosyası gibi).
- **Hassas alan** (T.C. kimlik, IBAN, adres, sağlık, yakın bilgisi, belgeler) yalnız `ozellik:ik.ozluk-hassas`
  yetkisiyle görünür ve yazılır; yetkisi olmayana alan hiç gönderilmez (maskeli değer de değil). Kişi kendi kaydının
  «profilim» alanlarını görür. Her kart açılışı, belge indirme ve dışa aktarma erişim kaydına düşer.
- **Portal içeriği:** duyurular, evrak deposu (form/rehber dosyaları), evrak talebi, yemek listesi, sık sorulan
  sorular. Hiçbiri dışarı gönderim yapmaz; evrak talebi İK'nın kuyruğuna düşer, İK durumunu işaretler.
"""
from __future__ import annotations

import io
import json
import re
import threading
import weakref
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Optional

import sqlalchemy as sa

from semantic_bridge import hr_core as H
from semantic_bridge.hr_core import HrError, clean, dump, iso, load, new_id, now

_md = sa.MetaData()
_ready: "weakref.WeakSet[sa.engine.Engine]" = weakref.WeakSet()
_lock = threading.Lock()

FIELD_TYPES = {"text": "Metin", "long_text": "Uzun metin", "enum": "Seçim listesi", "date": "Tarih", "email": "E-posta",
               "phone": "Telefon", "number": "Sayı", "file": "Belge / dosya"}
GROUPS = {"kimlik": "Kimlik", "is": "İş bilgileri", "sgk": "SGK ve çalışma izni", "iletisim": "İletişim",
          "egitim": "Öğrenim, askerlik, ehliyet", "aile": "Aile", "odeme": "Banka ve yan haklar", "belgeler": "Belgeler",
          "diger": "Diğer"}
PEOPLE_STATUS = ("Aktif", "Pasif")
REQUEST_STATUS = {"bekliyor": "Bekliyor", "hazirlaniyor": "Hazırlanıyor", "hazir": "Hazır / teslim edildi", "red": "Karşılanamadı"}
DELIVERY = {"eposta": "E-posta (dijital)", "elden": "Islak imzalı (elden)"}
DEFAULT_SETTINGS: dict[str, Any] = {
    "docTypes": ["Çalışma belgesi", "SGK hizmet dökümü", "Bordro", "Banka için maaş yazısı", "Vize için yazı", "Diğer"],
    "postCategories": ["Genel", "Etkinlik", "İşe Giriş", "Kutlama"],
    "docCategories": ["Formlar", "Rehberler"],
    "faqCategories": ["İzin Süreçleri", "Ücretler", "Çalışma Düzeni", "İnsan Kaynakları ve Eğitim"],
    "expiryDays": 60,
}
#: Kolona da yazılan alanlar (süzgeç ve eşleşme); değerleri yine data_json'dadır.
_COLUMNS = ("id_no", "ad_soyad", "durum", "mail_adresi")
_SEED = Path(__file__).with_name("hr_portal_fields.json")
_KEY_RE = re.compile(r"^[a-z][a-z0-9_]{1,59}$")


def _ts(name: str, **kw: Any) -> sa.Column:
    return sa.Column(name, sa.DateTime(timezone=True), **kw)


FIELDS = sa.Table(
    "semantic_hr_person_fields", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("key", sa.String(60), primary_key=True),
    sa.Column("label", sa.String(120), nullable=False),
    sa.Column("type", sa.String(16), nullable=False),
    sa.Column("options_json", sa.Text),
    sa.Column("required", sa.Boolean, nullable=False, default=False),
    sa.Column("show_admin", sa.Boolean, nullable=False, default=True),
    sa.Column("show_profile", sa.Boolean, nullable=False, default=False),
    sa.Column("show_directory", sa.Boolean, nullable=False, default=False),
    sa.Column("group_key", sa.String(20), nullable=False, default="diger"),
    sa.Column("sensitive", sa.Boolean, nullable=False, default=False),
    sa.Column("sort", sa.Integer, nullable=False, default=0),
    sa.Column("builtin", sa.Boolean, nullable=False, default=False),   # Excel'den gelen: silinmez, kapatılır
    sa.Column("active", sa.Boolean, nullable=False, default=True),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at"),
)

PEOPLE = sa.Table(
    "semantic_hr_people", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("id_no", sa.String(60), nullable=False),
    sa.Column("username", sa.String(120), index=True),                  # portal hesabı (AD); Profilim eşleşmesi
    sa.Column("ad_soyad", sa.String(200), nullable=False),
    sa.Column("durum", sa.String(16), nullable=False, default="Aktif"),
    sa.Column("mail_adresi", sa.String(200)),
    sa.Column("data_json", sa.Text, nullable=False),
    sa.Column("created_by", sa.String(120)),
    _ts("created_at", nullable=False),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at", nullable=False),
    sa.UniqueConstraint("tenant_id", "id_no", name="uq_hr_people_idno"),
)

PFILES = sa.Table(
    "semantic_hr_people_files", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("person_id", sa.String(40), nullable=False, index=True),
    sa.Column("field_key", sa.String(60), nullable=False),
    sa.Column("filename", sa.String(300), nullable=False),
    sa.Column("mime", sa.String(120)),
    sa.Column("size", sa.Integer, nullable=False, default=0),
    sa.Column("blob", sa.LargeBinary, nullable=False),
    sa.Column("uploaded_by", sa.String(120)),
    _ts("uploaded_at", nullable=False),
)

POSTS = sa.Table(
    "semantic_hr_posts", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("category", sa.String(60), nullable=False),
    sa.Column("title", sa.String(200), nullable=False),
    sa.Column("body", sa.Text, nullable=False),
    sa.Column("publish_date", sa.Date, nullable=False),
    sa.Column("image", sa.LargeBinary),
    sa.Column("image_mime", sa.String(60)),
    sa.Column("active", sa.Boolean, nullable=False, default=True),
    sa.Column("created_by", sa.String(120)),
    _ts("created_at", nullable=False),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at", nullable=False),
)

DOCS = sa.Table(
    "semantic_hr_docs", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("category", sa.String(60), nullable=False),
    sa.Column("code", sa.String(40)),
    sa.Column("title", sa.String(200), nullable=False),
    sa.Column("filename", sa.String(300), nullable=False),
    sa.Column("mime", sa.String(120)),
    sa.Column("size", sa.Integer, nullable=False, default=0),
    sa.Column("blob", sa.LargeBinary, nullable=False),
    sa.Column("uploaded_by", sa.String(120)),
    _ts("uploaded_at", nullable=False),
)

REQUESTS = sa.Table(
    "semantic_hr_doc_requests", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("username", sa.String(120), nullable=False, index=True),
    sa.Column("display", sa.String(200)),
    sa.Column("mail", sa.String(200)),
    sa.Column("doc_type", sa.String(120), nullable=False),
    sa.Column("delivery", sa.String(16), nullable=False),
    sa.Column("note", sa.Text),
    sa.Column("status", sa.String(16), nullable=False, default="bekliyor"),
    sa.Column("answer", sa.Text),
    _ts("created_at", nullable=False),
    sa.Column("handled_by", sa.String(120)),
    _ts("handled_at"),
)

REQUEST_FILES = sa.Table(
    "semantic_hr_doc_request_files", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("request_id", sa.String(40), nullable=False, index=True),
    sa.Column("filename", sa.String(300), nullable=False),
    sa.Column("mime", sa.String(120)),
    sa.Column("size", sa.Integer, nullable=False, default=0),
    sa.Column("blob", sa.LargeBinary, nullable=False),
    sa.Column("uploaded_by", sa.String(120)),
    _ts("uploaded_at", nullable=False),
)

MENU = sa.Table(
    "semantic_hr_menu", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("day", sa.Date, primary_key=True),
    sa.Column("items", sa.Text, nullable=False),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at", nullable=False),
)

FAQ = sa.Table(
    "semantic_hr_faq", _md,
    sa.Column("id", sa.String(40), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("category", sa.String(80), nullable=False),
    sa.Column("question", sa.String(300), nullable=False),
    sa.Column("answer", sa.Text, nullable=False),
    sa.Column("sort", sa.Integer, nullable=False, default=0),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at", nullable=False),
)

PSETTINGS = sa.Table(
    "semantic_hr_portal_settings", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("key", sa.String(40), primary_key=True),
    sa.Column("value_json", sa.Text, nullable=False),
    sa.Column("updated_by", sa.String(120)),
    _ts("updated_at", nullable=False),
)

TABLE_LABELS = {
    "semantic_hr_person_fields": "personel alan tanımları", "semantic_hr_people": "personel özlük kaydı",
    "semantic_hr_people_files": "personel belgeleri", "semantic_hr_posts": "şirket içi duyurular",
    "semantic_hr_docs": "evrak deposu", "semantic_hr_doc_requests": "evrak talepleri", "semantic_hr_doc_request_files": "hazır evraklar", "semantic_hr_menu": "yemek listesi",
    "semantic_hr_faq": "sık sorulan sorular", "semantic_hr_portal_settings": "İK portal ayarları",
}


def ensure(engine: sa.engine.Engine) -> None:
    with _lock:
        if engine in _ready:
            return
        H.ensure(engine)
        from semantic_layer.store import schema_stamp
        schema_stamp.create_all(_md, engine)   # sürüm damgası: tanım değişmediyse açılışta veritabanına sorulmaz
        _ready.add(engine)


def seed_fields() -> list[dict[str, Any]]:
    return json.loads(_SEED.read_text(encoding="utf-8"))


_seeded: set[tuple[int, str]] = set()


def _ensure_fields(engine: sa.engine.Engine, tenant: str) -> None:
    """Excel listesindeki alan tenant'ta yoksa ekler (İK'nın değiştirdiği alana dokunmaz). Kilitli: ilk açılışta
    paralel istekler aynı satırı iki kez eklemesin."""
    ensure(engine)
    if (id(engine), tenant) in _seeded:
        return
    with _lock, engine.begin() as c:
        have = {r.key for r in c.execute(sa.select(FIELDS.c.key).where(FIELDS.c.tenant_id == tenant)).all()}
        rows = [dict(tenant_id=tenant, key=f["key"], label=f["label"], type=f["type"], options_json=dump(f["options"]),
                     required=f["required"], show_admin=f["showAdmin"], show_profile=f["showProfile"],
                     show_directory=f["showDirectory"], group_key=f["group"], sensitive=f["sensitive"], sort=f["sort"],
                     builtin=True, active=True, updated_by=None, updated_at=now())
                for f in seed_fields() if f["key"] not in have]
        if rows:
            c.execute(FIELDS.insert(), rows)
    _seeded.add((id(engine), tenant))


# ------------------------------------------------------------------ yetki

F_VIEW = "ozellik:ik.ozluk-gor"
F_EDIT = "ozellik:ik.ozluk-duzenle"
F_SENS = "ozellik:ik.ozluk-hassas"
F_PORTAL = "ozellik:ik.portal-yonet"
F_FIELDS = "ozellik:ik.alan-ayar"


def rights(who: H.Who) -> dict[str, bool]:
    return {"view": who.can(F_VIEW, F_EDIT), "edit": who.can(F_EDIT), "sensitive": who.can(F_SENS),
            "portal": who.can(F_PORTAL), "fields": who.can(F_FIELDS),
            "leave": who.can("ozellik:ik.izin-yonet"), "leaveSettings": who.can("ozellik:ik.izin-ayar")}


# ------------------------------------------------------------------ alanlar


def _field_out(r: Any) -> dict[str, Any]:
    return {"key": r.key, "label": r.label, "type": r.type, "options": load(r.options_json, []), "required": bool(r.required),
            "showAdmin": bool(r.show_admin), "showProfile": bool(r.show_profile), "showDirectory": bool(r.show_directory),
            "group": r.group_key, "sensitive": bool(r.sensitive), "sort": int(r.sort or 0), "builtin": bool(r.builtin),
            "active": bool(r.active), "updatedBy": r.updated_by, "updatedAt": iso(r.updated_at)}


def fields(engine: sa.engine.Engine, tenant: str, *, include_inactive: bool = False) -> list[dict[str, Any]]:
    _ensure_fields(engine, tenant)
    with engine.connect() as c:
        rows = c.execute(sa.select(FIELDS).where(FIELDS.c.tenant_id == tenant).order_by(FIELDS.c.sort, FIELDS.c.key)).all()
    out = [_field_out(r) for r in rows]
    return out if include_inactive else [f for f in out if f["active"]]


def save_field(engine: sa.engine.Engine, tenant: str, actor: str, who_rights: dict[str, bool], body: dict[str, Any],
               key: Optional[str] = None) -> tuple[dict[str, Any], dict[str, Any]]:
    """Alan düzeltir (key) ya da yeni alan ekler. Excel'den gelen alanın türü ve anahtarı değişmez. Hassas alanı ve
    hassas işaretini yalnız «Hassas özlük verisi» yetkisi olan değiştirir (işareti kaldırıp alanı açmak olmasın)."""
    _ensure_fields(engine, tenant)
    with engine.begin() as c:
        before = None
        if key:
            before = c.execute(sa.select(FIELDS).where(FIELDS.c.tenant_id == tenant, FIELDS.c.key == key)).first()
            if before is None:
                raise HrError("Alan bulunamadı.", 404)
        if not who_rights["sensitive"] and ((before is not None and before.sensitive) or "sensitive" in body and
                                            bool(body.get("sensitive")) != bool(before.sensitive if before else False)):
            raise HrError("Hassas alanı ya da hassas işaretini değiştirmek için «Hassas özlük verisi» yetkisi gerekir.", 403)
        vals: dict[str, Any] = {}
        if before is None:
            k = clean(body.get("key"), 60).lower()
            if not _KEY_RE.match(k):
                raise HrError("Alan anahtarı küçük harf, rakam ve alt çizgiden oluşmalı (ör. yan_hak_notu).")
            if c.execute(sa.select(FIELDS.c.key).where(FIELDS.c.tenant_id == tenant, FIELDS.c.key == k)).first():
                raise HrError("Bu anahtarla bir alan zaten var.")
            if k in ("id", "username"):
                raise HrError("Bu anahtar ayrılmış; başka bir ad seçin.")
            t = str(body.get("type") or "")
            if t not in FIELD_TYPES:
                raise HrError("Alan türü geçersiz.")
            vals.update(key=k, type=t, builtin=False)
        if "label" in body or before is None:
            lab = clean(body.get("label"), 120)
            if not lab:
                raise HrError("Alanın görünen adı boş olamaz.")
            vals["label"] = lab
        ftype = vals.get("type") or before.type
        if "options" in body:
            opts = [clean(x, 80) for x in (body.get("options") or []) if clean(x, 80)]
            if len({o.casefold() for o in opts}) != len(opts):
                raise HrError("Seçenek listesinde aynı değer iki kez var.")
            vals["options_json"] = dump(opts)
        if ftype == "enum" and not load(vals.get("options_json") or (before.options_json if before else None), []):
            raise HrError("Seçim listesi türündeki alanın en az bir seçeneği olmalı.")
        for src, col in (("required", "required"), ("showAdmin", "show_admin"), ("showProfile", "show_profile"),
                         ("showDirectory", "show_directory"), ("sensitive", "sensitive"), ("active", "active")):
            if src in body:
                vals[col] = bool(body.get(src))
        if key in ("id_no", "ad_soyad", "durum") and (vals.get("active") is False or vals.get("required") is False):
            raise HrError("Personel no, ad soyad ve durum her kayıtta zorunludur; kapatılamaz.")
        if "group" in body:
            g = str(body.get("group") or "")
            if g not in GROUPS:
                raise HrError("Grup geçersiz.")
            vals["group_key"] = g
        if "sort" in body:
            try:
                vals["sort"] = int(body.get("sort"))
            except (TypeError, ValueError):
                raise HrError("Sıra bir sayı olmalı.") from None
        vals.update(updated_by=actor, updated_at=now())
        if before is None:
            vals.setdefault("sort", 10_000)
            vals.setdefault("group_key", "diger")
            vals.setdefault("options_json", dump([]))
            c.execute(FIELDS.insert().values(tenant_id=tenant, **vals))
            diff = {"yeni": vals["key"]}
            k = vals["key"]
        else:
            k = key
            diff = {col: {"once": getattr(before, col), "sonra": v} for col, v in vals.items()
                    if col not in ("updated_by", "updated_at") and getattr(before, col) != v}
            c.execute(FIELDS.update().where(FIELDS.c.tenant_id == tenant, FIELDS.c.key == k).values(**vals))
        row = c.execute(sa.select(FIELDS).where(FIELDS.c.tenant_id == tenant, FIELDS.c.key == k)).first()
    return _field_out(row), diff


def delete_field(engine: sa.engine.Engine, tenant: str, who_rights: dict[str, bool], key: str) -> int:
    """Sonradan eklenen alanı siler; kayıtlardaki değerleri de temizler. Dönen: değeri silinen kayıt sayısı."""
    ensure(engine)
    with engine.begin() as c:
        r = c.execute(sa.select(FIELDS).where(FIELDS.c.tenant_id == tenant, FIELDS.c.key == key)).first()
        if r is None:
            raise HrError("Alan bulunamadı.", 404)
        if r.builtin:
            raise HrError("Excel listesinden gelen alan silinmez; gerekmiyorsa kapatın.")
        if r.sensitive and not who_rights["sensitive"]:
            raise HrError("Hassas alanı silmek için «Hassas özlük verisi» yetkisi gerekir.", 403)
        n = 0
        for p in c.execute(sa.select(PEOPLE.c.id, PEOPLE.c.data_json).where(PEOPLE.c.tenant_id == tenant)).all():
            d = load(p.data_json, {})
            if key in d:
                d.pop(key)
                c.execute(PEOPLE.update().where(PEOPLE.c.id == p.id).values(data_json=dump(d)))
                n += 1
        c.execute(PFILES.delete().where(PFILES.c.tenant_id == tenant, PFILES.c.field_key == key))
        c.execute(FIELDS.delete().where(FIELDS.c.tenant_id == tenant, FIELDS.c.key == key))
    return n


# ------------------------------------------------------------------ ayarlar (listeler)


def settings(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    ensure(engine)
    with engine.connect() as c:
        rows = c.execute(sa.select(PSETTINGS).where(PSETTINGS.c.tenant_id == tenant)).all()
    out = dict(DEFAULT_SETTINGS)
    for r in rows:
        if r.key in out:
            out[r.key] = load(r.value_json, out[r.key])
    return out


def save_settings(engine: sa.engine.Engine, tenant: str, actor: str, body: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    ensure(engine)
    changed: list[str] = []
    cur = settings(engine, tenant)
    with engine.begin() as c:
        for k, default in DEFAULT_SETTINGS.items():
            if k not in body:
                continue
            if isinstance(default, list):
                v = [clean(x, 120) for x in (body.get(k) or []) if clean(x, 120)]
                v = list(dict.fromkeys(v))
                if not v:
                    raise HrError("Liste boş bırakılamaz; en az bir değer girin.")
            else:
                try:
                    v = int(body.get(k))
                except (TypeError, ValueError):
                    raise HrError("Gün sayısı bir sayı olmalı.") from None
                if not 1 <= v <= 365:
                    raise HrError("Gün sayısı 1 ile 365 arasında olmalı.")
            if v == cur[k]:
                continue
            c.execute(PSETTINGS.delete().where(PSETTINGS.c.tenant_id == tenant, PSETTINGS.c.key == k))
            c.execute(PSETTINGS.insert().values(tenant_id=tenant, key=k, value_json=dump(v), updated_by=actor, updated_at=now()))
            changed.append(k)
    return settings(engine, tenant), changed


# ------------------------------------------------------------------ değer doğrulama

_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_PHONE = re.compile(r"^[+0-9 ()\-/]{3,30}$")


def tc_ok(v: str) -> bool:
    """T.C. kimlik no: 11 hane, ilk hane 0 değil, 10. ve 11. hane algoritması."""
    if not re.fullmatch(r"[1-9][0-9]{10}", v):
        return False
    d = [int(x) for x in v]
    if ((sum(d[0:9:2]) * 7 - sum(d[1:8:2])) % 10) != d[9]:
        return False
    return sum(d[:10]) % 10 == d[10]


def iban_ok(v: str) -> bool:
    s = v.replace(" ", "").upper()
    if not re.fullmatch(r"[A-Z]{2}[0-9]{2}[A-Z0-9]{10,30}", s):
        return False
    if s.startswith("TR") and len(s) != 26:
        return False
    num = "".join(str(int(ch, 36)) for ch in s[4:] + s[:4])
    return int(num) % 97 == 1


def to_date(v: Any) -> Optional[date]:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", s)
    try:
        if m:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        return date.fromisoformat(s[:10])
    except ValueError:
        return None


def normalize(f: dict[str, Any], raw: Any) -> tuple[Any, Optional[str]]:
    """Bir alan değerini türüne göre düzenler. Dönen: (değer ya da None, hata metni)."""
    lab = f["label"]
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None, None
    t = f["type"]
    if t == "date":
        d = to_date(raw)
        return (d.isoformat(), None) if d else (None, f"{lab}: tarih okunamadı (GG.AA.YYYY).")
    if t == "number":
        s = str(raw).strip().replace("%", "").replace(",", ".")
        try:
            n = float(s)
        except ValueError:
            return None, f"{lab}: sayı olmalı."
        return (int(n) if n.is_integer() else n), None
    s = str(raw).strip() if t == "long_text" else clean(raw, 500)
    if isinstance(raw, float) and raw.is_integer():
        s = str(int(raw))                       # Excel sayıya çevirdiyse (dahili no, sicil no) «.0» düşer
    if t == "long_text":
        s = s[:4000]
    if t == "enum":
        opts = f["options"]
        hit = next((o for o in opts if o.casefold() == s.casefold()), None)
        return (hit, None) if hit else (None, f"{lab}: «{s}» listede yok ({', '.join(opts)}).")
    if t == "email":
        s = s.lower()
        return (s, None) if _EMAIL.match(s) else (None, f"{lab}: e-posta adresi geçersiz.")
    if t == "phone":
        return (s, None) if _PHONE.match(s) else (None, f"{lab}: telefon numarası geçersiz.")
    if f["key"] == "tc_kimlik_no":
        s = s.replace(" ", "")
        return (s, None) if tc_ok(s) else (None, "T.C. kimlik no geçersiz (11 hane ve doğrulama hanesi).")
    if f["key"] == "banka_iban_no":
        s = s.replace(" ", "").upper()
        return (s, None) if iban_ok(s) else (None, "IBAN geçersiz (TR ile başlayan 26 karakter, doğrulama hanesi).")
    return s, None


# ------------------------------------------------------------------ personel


def _visible(fs: list[dict[str, Any]], who_rights: dict[str, bool]) -> list[dict[str, Any]]:
    """İK'nın bu kişiye gösterebileceği yönetici paneli alanları."""
    return [f for f in fs if f["showAdmin"] and (who_rights["sensitive"] or not f["sensitive"])]


def _person_out(r: Any, fs: list[dict[str, Any]], files: list[dict[str, Any]], *, full: bool) -> dict[str, Any]:
    data = load(r.data_json, {})
    keys = {f["key"] for f in fs}
    out = {"id": r.id, "idNo": r.id_no, "adSoyad": r.ad_soyad, "durum": r.durum, "username": r.username,
           "updatedBy": r.updated_by, "updatedAt": iso(r.updated_at), "createdAt": iso(r.created_at)}
    if full:
        out["data"] = {k: v for k, v in data.items() if k in keys}
        out["files"] = [x for x in files if x["field"] in keys]
    else:
        out["data"] = {k: data.get(k) for k in ("firma", "sube", "ofis_lokasyon", "departman", "ekip", "unvan", "calisma_sekli",
                                                 "s_ise_giris_tarihi", "mail_adresi", "dahili_no") if k in keys}
    return out


def _file_out(r: Any) -> dict[str, Any]:
    return {"id": r.id, "field": r.field_key, "filename": r.filename, "mime": r.mime, "size": int(r.size or 0),
            "uploadedBy": r.uploaded_by, "uploadedAt": iso(r.uploaded_at)}


def _files_of(c: Any, tenant: str, pids: Iterable[str]) -> dict[str, list[dict[str, Any]]]:
    ids = list(pids)
    out: dict[str, list[dict[str, Any]]] = {}
    if not ids:
        return out
    cols = [PFILES.c.id, PFILES.c.person_id, PFILES.c.field_key, PFILES.c.filename, PFILES.c.mime, PFILES.c.size,
            PFILES.c.uploaded_by, PFILES.c.uploaded_at]
    for r in c.execute(sa.select(*cols).where(PFILES.c.tenant_id == tenant, PFILES.c.person_id.in_(ids))
                       .order_by(PFILES.c.uploaded_at.desc())).all():
        out.setdefault(r.person_id, []).append(_file_out(r))
    return out


def _all_people(c: Any, tenant: str) -> list[Any]:
    return c.execute(sa.select(PEOPLE).where(PEOPLE.c.tenant_id == tenant)).all()


def list_people(engine: sa.engine.Engine, tenant: str, who_rights: dict[str, bool], *, q: str = "", durum: str = "Aktif",
                filters: Optional[dict[str, str]] = None) -> dict[str, Any]:
    fs = _visible(fields(engine, tenant), who_rights)
    keys = {f["key"] for f in fs}
    qq = (q or "").strip().casefold()
    with engine.connect() as c:
        rows = _all_people(c, tenant)
        files = _files_of(c, tenant, [r.id for r in rows])
    need_docs = [f["key"] for f in fs if f["type"] == "file" and f["required"]]
    req = [f for f in fs if f["required"] and f["type"] != "file"]
    items = []
    for r in rows:
        if durum and durum != "hepsi" and r.durum != durum:
            continue
        data = load(r.data_json, {})
        if filters and any(v and str(data.get(k) or "") != v for k, v in filters.items() if k in keys):
            continue
        if qq and qq not in " ".join(str(x) for x in (r.ad_soyad, r.id_no, r.username or "", data.get("unvan") or "",
                                                         data.get("departman") or "", data.get("ekip") or "",
                                                         data.get("mail_adresi") or "")).casefold():
            continue
        o = _person_out(r, fs, [], full=False)
        have = {x["field"] for x in files.get(r.id, [])}
        o["missing"] = [f["key"] for f in req if data.get(f["key"]) in (None, "")] + [k for k in need_docs if k not in have]
        o["hasPhoto"] = "fotograf" in have and "fotograf" in keys
        items.append(o)
    items.sort(key=lambda x: x["adSoyad"].casefold())
    return {"items": items, "total": len(items)}


def get_person(engine: sa.engine.Engine, tenant: str, who_rights: dict[str, bool], pid: str) -> dict[str, Any]:
    fs = _visible(fields(engine, tenant), who_rights)
    with engine.connect() as c:
        r = c.execute(sa.select(PEOPLE).where(PEOPLE.c.id == pid, PEOPLE.c.tenant_id == tenant)).first()
        if r is None:
            raise HrError("Personel kaydı bulunamadı.", 404)
        files = _files_of(c, tenant, [pid]).get(pid, [])
    return _person_out(r, fs, files, full=True)


def _validate(fs: list[dict[str, Any]], data: dict[str, Any], known_ids: set[str], own_id: str) -> list[str]:
    errs: list[str] = []
    for f in fs:
        if f["type"] == "file":
            continue
        if f["required"] and data.get(f["key"]) in (None, ""):
            errs.append(f"{f['label']} zorunlu.")
    mgr = data.get("yonetici_id_no")
    if mgr and mgr != own_id and mgr not in known_ids:
        errs.append(f"Yöneticisi: «{mgr}» personel numarasıyla kayıt yok.")
    return errs


def save_person(engine: sa.engine.Engine, tenant: str, actor: str, who_rights: dict[str, bool], body: dict[str, Any],
                pid: Optional[str] = None) -> tuple[dict[str, Any], list[str]]:
    """Personel kartını yazar. Yalnız kişinin görebildiği alanlar değişir; hassas alan yetkisizse dokunulmaz.
    Dönen: (kart, değişen alan anahtarları — değer değil)."""
    all_fs = fields(engine, tenant)
    fs = _visible(all_fs, who_rights)
    by_key = {f["key"]: f for f in fs}
    incoming = body.get("data") or {}
    if not isinstance(incoming, dict):
        raise HrError("Gövde geçersiz.")
    with engine.begin() as c:
        before = None
        if pid:
            before = c.execute(sa.select(PEOPLE).where(PEOPLE.c.id == pid, PEOPLE.c.tenant_id == tenant)).first()
            if before is None:
                raise HrError("Personel kaydı bulunamadı.", 404)
        data = load(before.data_json, {}) if before else {}
        errs: list[str] = []
        changed: list[str] = []
        for k, raw in incoming.items():
            f = by_key.get(k)
            if f is None or f["type"] == "file":
                continue                              # görünmeyen/hassas alan yetkisizse sessizce yazılmaz
            v, e = normalize(f, raw)
            if e:
                errs.append(e)
                continue
            if data.get(k) != v:
                changed.append(k)
                if v is None:
                    data.pop(k, None)
                else:
                    data[k] = v
        if before is not None and "mail_adresi" in changed and not who_rights["sensitive"] and not before.username:
            # Portal hesabı bağlı olmayan kayıtta e-postanın @ öncesi Profilim eşleşmesidir; değiştirmek hassas yetki ister.
            raise HrError("Portal hesabı bağlı olmayan kaydın e-postasını değiştirmek «Hassas özlük verisi» yetkisi ister "
                          "(e-posta Profilim eşleşmesini belirler).", 403)
        id_no = str(data.get("id_no") or "")
        others = {r.id_no for r in c.execute(sa.select(PEOPLE.c.id_no).where(
            PEOPLE.c.tenant_id == tenant, PEOPLE.c.id != (pid or ""))).all()}
        if id_no and id_no in others:
            errs.append(f"«{id_no}» personel numarası başka bir kayıtta var.")
        # Zorunlu alan denetimi yalnız kişinin görebildiği alanlarda (hassas zorunlu alanı göremeyen onu dolduramaz).
        errs += _validate(fs, data, others, id_no)
        if errs:
            raise HrError(" ".join(dict.fromkeys(errs)))
        vals: dict[str, Any] = dict(id_no=id_no, ad_soyad=str(data.get("ad_soyad") or "")[:200],
                                    durum=str(data.get("durum") or "Aktif"), mail_adresi=(data.get("mail_adresi") or None),
                                    data_json=dump(data), updated_by=actor, updated_at=now())
        if "username" in body and who_rights["sensitive"]:
            # Portal hesabı Profilim'i (hassas alanlar ve belgeler dahil) açar: bağlamak hassas yetki ister.
            u = clean(body.get("username"), 120).lower() or None
            if u and c.execute(sa.select(PEOPLE.c.id).where(PEOPLE.c.tenant_id == tenant, PEOPLE.c.username == u,
                                                            PEOPLE.c.id != (pid or ""))).first():
                raise HrError(f"«{u}» portal hesabı başka bir personel kaydına bağlı.")
            if (before.username if before else None) != u:
                changed.append("username")
            vals["username"] = u
        if before is None:
            pid = new_id("prs")
            c.execute(PEOPLE.insert().values(id=pid, tenant_id=tenant, created_by=actor, created_at=now(), **vals))
        elif changed:
            c.execute(PEOPLE.update().where(PEOPLE.c.id == pid).values(**vals))
    return get_person(engine, tenant, who_rights, pid), changed


def delete_person(engine: sa.engine.Engine, tenant: str, pid: str) -> dict[str, Any]:
    ensure(engine)
    with engine.begin() as c:
        r = c.execute(sa.select(PEOPLE.c.id, PEOPLE.c.id_no).where(PEOPLE.c.id == pid, PEOPLE.c.tenant_id == tenant)).first()
        if r is None:
            raise HrError("Personel kaydı bulunamadı.", 404)
        nf = c.execute(PFILES.delete().where(PFILES.c.tenant_id == tenant, PFILES.c.person_id == pid)).rowcount
        c.execute(PEOPLE.delete().where(PEOPLE.c.id == pid))
    return {"idNo": r.id_no, "files": int(nf or 0)}


# ------------------------------------------------------------------ belgeler

_MIME = {".pdf": "application/pdf", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp",
         ".heic": "image/heic", ".doc": "application/msword",
         ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         ".xls": "application/vnd.ms-excel", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         ".odt": "application/vnd.oasis.opendocument.text", ".txt": "text/plain"}
FILE_ACCEPT = ",".join(_MIME)
IMAGE_EXT = (".jpg", ".jpeg", ".png", ".webp")


def _mime(name: str, only: Optional[tuple[str, ...]] = None) -> str:
    ext = Path(name.lower()).suffix
    if ext not in _MIME or (only and ext not in only):
        allowed = ", ".join(only or tuple(_MIME))
        raise HrError(f"Bu dosya türü alınmıyor; kabul edilen: {allowed}.")
    return _MIME[ext]


def add_file(engine: sa.engine.Engine, tenant: str, actor: str, who_rights: dict[str, bool], pid: str, field_key: str,
             filename: str, data: bytes, max_mb: int) -> dict[str, Any]:
    fs = {f["key"]: f for f in _visible(fields(engine, tenant), who_rights)}
    f = fs.get(field_key)
    if f is None or f["type"] != "file":
        raise HrError("Bu belge alanı yok ya da yetkiniz dışında.", 403)
    name = clean(filename, 300) or "belge"
    if not data:
        raise HrError("Dosya boş.")
    if len(data) > max_mb * 1024 * 1024:
        raise HrError(f"Dosya {max_mb} MB sınırını aşıyor.", 413)
    mime = _mime(name, IMAGE_EXT if field_key == "fotograf" else None)
    with engine.begin() as c:
        if c.execute(sa.select(PEOPLE.c.id).where(PEOPLE.c.id == pid, PEOPLE.c.tenant_id == tenant)).first() is None:
            raise HrError("Personel kaydı bulunamadı.", 404)
        if field_key == "fotograf":                  # tek fotoğraf: yenisi eskisinin yerine
            c.execute(PFILES.delete().where(PFILES.c.tenant_id == tenant, PFILES.c.person_id == pid, PFILES.c.field_key == "fotograf"))
        fid = new_id("prd")
        c.execute(PFILES.insert().values(id=fid, tenant_id=tenant, person_id=pid, field_key=field_key, filename=name, mime=mime,
                                         size=len(data), blob=data, uploaded_by=actor, uploaded_at=now()))
        c.execute(PEOPLE.update().where(PEOPLE.c.id == pid).values(updated_by=actor, updated_at=now()))
        r = c.execute(sa.select(PFILES).where(PFILES.c.id == fid)).first()
    return _file_out(r)


def file_blob(engine: sa.engine.Engine, tenant: str, pid: str, fid: str, allowed_fields: set[str]) -> tuple[bytes, str, str, str]:
    ensure(engine)
    with engine.connect() as c:
        r = c.execute(sa.select(PFILES).where(PFILES.c.id == fid, PFILES.c.person_id == pid, PFILES.c.tenant_id == tenant)).first()
    if r is None or r.field_key not in allowed_fields:
        raise HrError("Belge bulunamadı.", 404)
    return bytes(r.blob), r.filename, r.mime or "application/octet-stream", r.field_key


def delete_file(engine: sa.engine.Engine, tenant: str, who_rights: dict[str, bool], pid: str, fid: str) -> str:
    keys = {f["key"] for f in _visible(fields(engine, tenant), who_rights) if f["type"] == "file"}
    with engine.begin() as c:
        r = c.execute(sa.select(PFILES.c.field_key, PFILES.c.filename).where(
            PFILES.c.id == fid, PFILES.c.person_id == pid, PFILES.c.tenant_id == tenant)).first()
        if r is None or r.field_key not in keys:
            raise HrError("Belge bulunamadı.", 404)
        c.execute(PFILES.delete().where(PFILES.c.id == fid))
    return r.field_key


def photo(engine: sa.engine.Engine, tenant: str, pid: str) -> Optional[tuple[bytes, str]]:
    ensure(engine)
    with engine.connect() as c:
        r = c.execute(sa.select(PFILES.c.blob, PFILES.c.mime).where(
            PFILES.c.tenant_id == tenant, PFILES.c.person_id == pid, PFILES.c.field_key == "fotograf")).first()
    return (bytes(r.blob), r.mime or "image/jpeg") if r else None


# ------------------------------------------------------------------ Profilim, rehber, doğum günleri


def person_of(engine: sa.engine.Engine, tenant: str, user: str) -> Optional[Any]:
    """Oturumdaki kişinin özlük kaydı: portal hesabı bağlıysa o, değilse e-postasının @ öncesi hesap adıyla."""
    ensure(engine)
    u = (user or "").strip().lower()
    if not u:
        return None
    with engine.connect() as c:
        r = c.execute(sa.select(PEOPLE).where(PEOPLE.c.tenant_id == tenant, PEOPLE.c.username == u)).first()
        if r is None:
            rows = c.execute(sa.select(PEOPLE).where(PEOPLE.c.tenant_id == tenant, PEOPLE.c.username.is_(None),
                                                     sa.func.lower(PEOPLE.c.mail_adresi).startswith(u + "@", autoescape=True))).all()
            r = rows[0] if len(rows) == 1 else None
    return r


def my_profile(engine: sa.engine.Engine, tenant: str, user: str) -> dict[str, Any]:
    fs = [f for f in fields(engine, tenant) if f["showProfile"]]
    r = person_of(engine, tenant, user)
    if r is None:
        return {"person": None, "fields": fs, "groups": GROUPS}
    with engine.connect() as c:
        files = _files_of(c, tenant, [r.id]).get(r.id, [])
    out = _person_out(r, fs, files, full=True)
    by_id = {}
    mgr = out["data"].get("yonetici_id_no") or load(r.data_json, {}).get("yonetici_id_no")
    if mgr:
        with engine.connect() as c:
            m = c.execute(sa.select(PEOPLE.c.ad_soyad).where(PEOPLE.c.tenant_id == tenant, PEOPLE.c.id_no == mgr)).first()
        by_id["manager"] = m.ad_soyad if m else None
    return {"person": out, "fields": fs, "groups": GROUPS, **by_id}


def directory(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    fs = [f for f in fields(engine, tenant) if f["showDirectory"] and f["type"] != "file" and not f["sensitive"]]
    photo_ok = any(f["key"] == "fotograf" and f["showDirectory"] for f in fields(engine, tenant))
    keys = [f["key"] for f in fs]
    with engine.connect() as c:
        rows = c.execute(sa.select(PEOPLE).where(PEOPLE.c.tenant_id == tenant, PEOPLE.c.durum == "Aktif")).all()
        with_photo = {r.person_id for r in c.execute(sa.select(PFILES.c.person_id).where(
            PFILES.c.tenant_id == tenant, PFILES.c.field_key == "fotograf")).all()} if photo_ok else set()
    items = []
    for r in rows:
        d = load(r.data_json, {})
        items.append({"id": r.id, "adSoyad": r.ad_soyad, "data": {k: d.get(k) for k in keys if d.get(k) not in (None, "")},
                      "hasPhoto": r.id in with_photo})
    items.sort(key=lambda x: (str(x["data"].get("departman") or "~").casefold(), x["adSoyad"].casefold()))
    return {"items": items, "fields": fs, "total": len(items)}


def directory_photo_allowed(engine: sa.engine.Engine, tenant: str) -> bool:
    return any(f["key"] == "fotograf" and f["showDirectory"] for f in fields(engine, tenant))


def birthdays(engine: sa.engine.Engine, tenant: str, today: Optional[date] = None) -> dict[str, Any]:
    """Aktif personelin doğum günü (gün ve ay; yıl ve yaş gönderilmez). Bu ay ve önümüzdeki 7 gün."""
    today = today or date.today()
    f = next((x for x in fields(engine, tenant) if x["key"] == "dogum_tarihi"), None)
    if f is None or f["sensitive"]:
        # Alan kapalı ya da hassas işaretliyse gün/ay da kimseye gitmez.
        return {"today": today.isoformat(), "items": [], "thisMonth": [], "upcoming": [], "off": True}
    with engine.connect() as c:
        rows = c.execute(sa.select(PEOPLE.c.id, PEOPLE.c.ad_soyad, PEOPLE.c.data_json).where(
            PEOPLE.c.tenant_id == tenant, PEOPLE.c.durum == "Aktif")).all()
    out = []
    for r in rows:
        d = load(r.data_json, {})
        b = to_date(d.get("dogum_tarihi"))
        if not b:
            continue
        try:
            nxt = b.replace(year=today.year)
        except ValueError:                              # 29 Şubat
            nxt = date(today.year, 3, 1)
        if nxt < today:
            try:
                nxt = b.replace(year=today.year + 1)
            except ValueError:
                nxt = date(today.year + 1, 3, 1)
        out.append({"id": r.id, "adSoyad": r.ad_soyad, "departman": d.get("departman"), "unvan": d.get("unvan"),
                    "day": b.day, "month": b.month, "inDays": (nxt - today).days})
    out.sort(key=lambda x: (x["month"], x["day"], x["adSoyad"].casefold()))
    return {"today": today.isoformat(), "items": out,
            "thisMonth": [x for x in out if x["month"] == today.month],
            "upcoming": sorted([x for x in out if x["inDays"] <= 7], key=lambda x: x["inDays"])}


# ------------------------------------------------------------------ İK ana sayfası (sayılar)


def _count(rows: Iterable[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    agg: dict[str, int] = {}
    for d in rows:
        v = str(d.get(key) or "").strip() or "Girilmemiş"
        agg[v] = agg.get(v, 0) + 1
    return sorted(({"label": k, "value": v} for k, v in agg.items()), key=lambda x: (-x["value"], x["label"]))


def _years(a: date, b: date) -> float:
    return (b - a).days / 365.25


def stats(engine: sa.engine.Engine, tenant: str, who_rights: dict[str, bool], today: Optional[date] = None) -> dict[str, Any]:
    """İK'nın ana sayfa sayıları. Hepsi portal kaydından Python'da; kişi adı yalnız yaklaşan tarih listelerinde."""
    today = today or date.today()
    st = settings(engine, tenant)
    horizon = today + timedelta(days=int(st["expiryDays"]))
    all_fs = fields(engine, tenant)
    fs = _visible(all_fs, who_rights)
    keys = {f["key"] for f in fs}
    with engine.connect() as c:
        rows = _all_people(c, tenant)
        file_rows = c.execute(sa.select(PFILES.c.person_id, PFILES.c.field_key).where(PFILES.c.tenant_id == tenant)).all()
        pending = c.execute(sa.select(sa.func.count()).select_from(REQUESTS).where(
            REQUESTS.c.tenant_id == tenant, REQUESTS.c.status.in_(("bekliyor", "hazirlaniyor")))).scalar() or 0
    have: dict[str, set[str]] = {}
    for fr in file_rows:
        have.setdefault(fr.person_id, set()).add(fr.field_key)
    people = [(r, load(r.data_json, {})) for r in rows]
    active = [(r, d) for r, d in people if r.durum == "Aktif"]
    ad = [d for _, d in active]

    def within(d: dict[str, Any], k: str, a: date, b: date) -> bool:
        x = to_date(d.get(k))
        return bool(x and a <= x <= b)

    month0 = today.replace(day=1)
    year0 = today.replace(month=1, day=1)
    start_key = "s_ise_giris_tarihi"
    joined_m = [d for _, d in people if within(d, start_key, month0, today)]
    joined_y = [d for _, d in people if within(d, start_key, year0, today)]
    left_m = [d for _, d in people if within(d, "isten_cikis_tarihi", month0, today)]
    left_y = [d for _, d in people if within(d, "isten_cikis_tarihi", year0, today)]
    tenure = [_years(x, today) for x in (to_date(d.get(start_key)) or to_date(d.get("f_ise_giris_tarihi")) for d in ad) if x and x <= today]
    ages = [_years(x, today) for x in (to_date(d.get("dogum_tarihi")) for d in ad) if x and x <= today] if "dogum_tarihi" in keys else []

    def expiring(k: str, label: str) -> list[dict[str, Any]]:
        if k not in keys:
            return []
        out = []
        for r, d in active:
            x = to_date(d.get(k))
            if x and x <= horizon:
                out.append({"id": r.id, "adSoyad": r.ad_soyad, "what": label, "date": x.isoformat(), "inDays": (x - today).days})
        return out

    upcoming = sorted(expiring("calisma_izni_bitis_tarihi", "Çalışma izni bitiyor")
                      + expiring("askerlik_tecil_suresi", "Askerlik tecili bitiyor"), key=lambda x: x["date"])
    anniv = []
    for r, d in active:
        s = to_date(d.get(start_key)) or to_date(d.get("f_ise_giris_tarihi"))
        if s and s.month == today.month and s.year < today.year:
            anniv.append({"id": r.id, "adSoyad": r.ad_soyad, "years": today.year - s.year, "day": s.day})
    anniv.sort(key=lambda x: x["day"])
    req = [f for f in fs if f["required"] and f["type"] != "file"]
    docs = [f for f in fs if f["type"] == "file" and f["key"] != "fotograf"]
    missing_req = sum(1 for r, d in active if any(d.get(f["key"]) in (None, "") for f in req))
    doc_gaps = sorted(({"key": f["key"], "label": f["label"], "missing": sum(1 for r, _ in active if f["key"] not in have.get(r.id, set()))}
                       for f in docs), key=lambda x: (-x["missing"], x["label"]))
    avg_head = (len(active) + len(left_y)) or 1
    return {
        "today": today.isoformat(),
        "headcount": {"active": len(active), "passive": len(people) - len(active), "total": len(people)},
        "joined": {"month": len(joined_m), "year": len(joined_y)},
        "left": {"month": len(left_m), "year": len(left_y), "turnoverYearPct": round(100 * len(left_y) / avg_head, 1)},
        "avgTenureYears": round(sum(tenure) / len(tenure), 1) if tenure else None,
        "avgAgeYears": round(sum(ages) / len(ages), 1) if ages else None,
        "by": {k: _count(ad, k) for k in ("firma", "sube", "ofis_lokasyon", "departman", "calisma_sekli", "cinsiyet") if k in keys},
        "upcoming": upcoming, "anniversaries": anniv, "expiryDays": int(st["expiryDays"]),
        "completeness": {"activeMissingRequired": missing_req, "docs": doc_gaps},
        "pendingRequests": int(pending),
    }


# ------------------------------------------------------------------ Excel içe / dışa aktarma


def _header_map(header: list[Any], fs: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    by = {}
    for f in fs:
        by[f["key"].casefold()] = f
        by[f["label"].casefold()] = f
    out = {}
    for i, h in enumerate(header):
        f = by.get(str(h or "").strip().casefold())
        if f and f["type"] != "file":
            out[i] = f
    return out


def import_xlsx(engine: sa.engine.Engine, tenant: str, actor: str, who_rights: dict[str, bool], data: bytes, *,
                apply: bool) -> dict[str, Any]:
    """İK'nın Excel'i (başlık satırı alan anahtarı ya da görünen adı; «crm_data» sayfası ya da ilk uyan sayfa).
    Personel no ile eşleşir: yoksa yeni, varsa yalnız dolu hücreler güncellenir (boş hücre var olanı silmez).
    `apply=False` hiçbir şey yazmaz; satır satır ne olacağını döner."""
    from openpyxl import load_workbook

    fs = [f for f in _visible(fields(engine, tenant), who_rights)]
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — bozuk ya da Excel değil
        raise HrError("Dosya Excel (.xlsx) olarak okunamadı.") from None
    sheet, hmap, rows, header = None, {}, [], []
    names = sorted(wb.sheetnames, key=lambda n: n != "crm_data")
    for n in names:
        ws = wb[n]
        it = ws.iter_rows(values_only=True)
        head = next(it, None)
        if not head:
            continue
        m = _header_map(list(head), fs)
        if any(f["key"] == "id_no" for f in m.values()):
            sheet, hmap, rows, header = n, m, [list(r) for r in it], list(head)
            break
    if sheet is None:
        raise HrError("Excel'de «id_no» (personel no) başlıklı bir sayfa bulunamadı. Boş şablonu indirip onu doldurun.")
    with engine.connect() as c:
        existing = {r.id_no: r for r in _all_people(c, tenant)}
    file_ids = set()
    parsed = []
    for n, raw in enumerate(rows, start=2):
        if not any(v not in (None, "") for v in raw):
            continue
        vals: dict[str, Any] = {}
        errs: list[str] = []
        for i, f in hmap.items():
            v, e = normalize(f, raw[i] if i < len(raw) else None)
            if e:
                errs.append(e)
            elif v is not None:
                vals[f["key"]] = v
        idn = str(vals.get("id_no") or "")
        if not idn:
            errs.append("Personel no boş.")
        elif idn in file_ids:
            errs.append(f"«{idn}» dosyada iki kez var.")
        file_ids.add(idn)
        parsed.append((n, idn, vals, errs))
    known = set(existing) | file_ids
    out_rows = []
    counts = {"new": 0, "update": 0, "same": 0, "error": 0}
    for n, idn, vals, errs in parsed:
        before = existing.get(idn)
        merged = {**(load(before.data_json, {}) if before else {}), **vals}
        if not errs:
            errs = _validate(fs, merged, known, idn)
        changed = [k for k, v in vals.items() if not before or load(before.data_json, {}).get(k) != v]
        if (before is not None and "mail_adresi" in changed and not who_rights["sensitive"] and not before.username
                and not errs):
            errs = ["E-posta değişikliği (portal hesabı bağlı olmayan kayıt) «Hassas özlük verisi» yetkisi ister."]
        kind = "error" if errs else ("new" if before is None else ("update" if changed else "same"))
        counts[kind] += 1
        out_rows.append({"row": n, "idNo": idn, "adSoyad": merged.get("ad_soyad"), "kind": kind, "errors": errs,
                         "changed": changed if before else []})
    applied = 0
    if apply:
        if counts["error"]:
            raise HrError(f"{counts['error']} satırda hata var; önce düzeltin, sonra yeniden yükleyin. Hiçbir satır yazılmadı.")
        with engine.begin() as c:
            for (n, idn, vals, _), o in zip(parsed, out_rows):
                if o["kind"] == "same":
                    continue
                before = existing.get(idn)
                merged = {**(load(before.data_json, {}) if before else {}), **vals}
                cols = dict(id_no=idn, ad_soyad=str(merged.get("ad_soyad") or "")[:200], durum=str(merged.get("durum") or "Aktif"),
                            mail_adresi=merged.get("mail_adresi") or None, data_json=dump(merged), updated_by=actor, updated_at=now())
                if before is None:
                    c.execute(PEOPLE.insert().values(id=new_id("prs"), tenant_id=tenant, created_by=actor, created_at=now(), **cols))
                else:
                    c.execute(PEOPLE.update().where(PEOPLE.c.id == before.id).values(**cols))
                applied += 1
    return {"sheet": sheet, "columns": [f["key"] for f in hmap.values()],
            "ignoredColumns": [str(h) for i, h in enumerate(header) if h not in (None, "") and i not in hmap],
            "counts": counts, "rows": out_rows, "applied": applied}


def export_xlsx(engine: sa.engine.Engine, tenant: str, who_rights: dict[str, bool], *, template: bool = False,
                durum: str = "hepsi") -> bytes:
    """Başlık = alan anahtarı (içe aktarmayla aynı biçim), ikinci satır yok. Yetkiyle görünmeyen alan kolon olarak yok."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    fs = [f for f in _visible(fields(engine, tenant), who_rights) if f["type"] != "file"]
    wb = Workbook()
    ws = wb.active
    ws.title = "crm_data"
    ws.append([f["key"] for f in fs])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    if not template:
        with engine.connect() as c:
            rows = _all_people(c, tenant)
        for r in sorted(rows, key=lambda x: x.ad_soyad.casefold()):
            if durum != "hepsi" and r.durum != durum:
                continue
            d = load(r.data_json, {})
            line = []
            for f in fs:
                v = d.get(f["key"])
                if f["type"] == "date" and v:
                    v = to_date(v)
                line.append(v)
            ws.append(line)
        for col in ws.iter_cols(min_row=2):
            for cell in col:
                if isinstance(cell.value, str):
                    cell.number_format = "@"
                elif isinstance(cell.value, date):
                    cell.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, f in enumerate(fs, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(40, len(f["key"]) + 4))
    labels = wb.create_sheet("alanlar")
    labels.append(["alan", "görünen ad", "tür", "seçenekler", "zorunlu"])
    for cell in labels[1]:
        cell.font = Font(bold=True)
    for f in fs:
        labels.append([f["key"], f["label"], FIELD_TYPES[f["type"]], ", ".join(f["options"]), "evet" if f["required"] else ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ duyurular


def _post_out(r: Any, *, full: bool = True) -> dict[str, Any]:
    body = r.body or ""
    return {"id": r.id, "category": r.category, "title": r.title, "body": body if full else body[:280],
            "publishDate": iso(r.publish_date), "hasImage": r.image_mime is not None, "active": bool(r.active),
            "createdBy": r.created_by, "updatedAt": iso(r.updated_at)}


_POST_COLS = [POSTS.c.id, POSTS.c.category, POSTS.c.title, POSTS.c.body, POSTS.c.publish_date, POSTS.c.image_mime,
              POSTS.c.active, POSTS.c.created_by, POSTS.c.updated_at]


def list_posts(engine: sa.engine.Engine, tenant: str, *, all_: bool = False, limit: int = 0,
               today: Optional[date] = None) -> list[dict[str, Any]]:
    """Yayındaki duyurular (en yeni önce). İK tümünü (ileri tarihli ve kapatılmış dahil) görür."""
    ensure(engine)
    stmt = sa.select(*_POST_COLS).where(POSTS.c.tenant_id == tenant)
    if not all_:
        stmt = stmt.where(POSTS.c.active.is_(True), POSTS.c.publish_date <= (today or date.today()))
    stmt = stmt.order_by(POSTS.c.publish_date.desc(), POSTS.c.created_at.desc())
    if limit:
        stmt = stmt.limit(limit)
    with engine.connect() as c:
        return [_post_out(r) for r in c.execute(stmt).all()]


def save_post(engine: sa.engine.Engine, tenant: str, actor: str, body: dict[str, Any], post_id: Optional[str] = None) -> dict[str, Any]:
    ensure(engine)
    st = settings(engine, tenant)
    with engine.begin() as c:
        before = None
        if post_id:
            before = c.execute(sa.select(*_POST_COLS).where(POSTS.c.id == post_id, POSTS.c.tenant_id == tenant)).first()
            if before is None:
                raise HrError("Duyuru bulunamadı.", 404)
        vals: dict[str, Any] = {}
        if "title" in body or before is None:
            vals["title"] = clean(body.get("title"), 200)
            if not vals["title"]:
                raise HrError("Başlık boş olamaz.")
        if "body" in body or before is None:
            vals["body"] = str(body.get("body") or "").strip()[:20_000]
            if not vals["body"]:
                raise HrError("Duyuru metni boş olamaz.")
        if "category" in body or before is None:
            cat = clean(body.get("category"), 60)
            if cat not in st["postCategories"]:
                raise HrError("Kategori listede yok; listeyi Ayarlar sekmesinden düzenleyin.")
            vals["category"] = cat
        if "publishDate" in body or before is None:
            vals["publish_date"] = H.parse_date(body.get("publishDate"), "Yayın") or date.today()
        if "active" in body:
            vals["active"] = bool(body.get("active"))
        vals.update(updated_by=actor, updated_at=now())
        if before is None:
            post_id = new_id("dyr")
            c.execute(POSTS.insert().values(id=post_id, tenant_id=tenant, created_by=actor, created_at=now(),
                                            active=vals.pop("active", True), **vals))
        else:
            c.execute(POSTS.update().where(POSTS.c.id == post_id).values(**vals))
        r = c.execute(sa.select(*_POST_COLS).where(POSTS.c.id == post_id)).first()
    return _post_out(r)


def set_post_image(engine: sa.engine.Engine, tenant: str, actor: str, post_id: str, filename: str, data: Optional[bytes],
                   max_mb: int) -> dict[str, Any]:
    ensure(engine)
    mime = None
    if data:
        if len(data) > max_mb * 1024 * 1024:
            raise HrError(f"Görsel {max_mb} MB sınırını aşıyor.", 413)
        mime = _mime(filename or "", IMAGE_EXT)
    with engine.begin() as c:
        n = c.execute(POSTS.update().where(POSTS.c.id == post_id, POSTS.c.tenant_id == tenant).values(
            image=data or None, image_mime=mime, updated_by=actor, updated_at=now())).rowcount
        if not n:
            raise HrError("Duyuru bulunamadı.", 404)
        r = c.execute(sa.select(*_POST_COLS).where(POSTS.c.id == post_id)).first()
    return _post_out(r)


def post_image(engine: sa.engine.Engine, tenant: str, post_id: str, *, all_: bool) -> tuple[bytes, str]:
    ensure(engine)
    with engine.connect() as c:
        r = c.execute(sa.select(POSTS.c.image, POSTS.c.image_mime, POSTS.c.active, POSTS.c.publish_date).where(
            POSTS.c.id == post_id, POSTS.c.tenant_id == tenant)).first()
    if r is None or r.image is None or (not all_ and (not r.active or r.publish_date > date.today())):
        raise HrError("Görsel bulunamadı.", 404)
    return bytes(r.image), r.image_mime or "image/jpeg"


def delete_post(engine: sa.engine.Engine, tenant: str, post_id: str) -> str:
    ensure(engine)
    with engine.begin() as c:
        r = c.execute(sa.select(POSTS.c.title).where(POSTS.c.id == post_id, POSTS.c.tenant_id == tenant)).first()
        if r is None:
            raise HrError("Duyuru bulunamadı.", 404)
        c.execute(POSTS.delete().where(POSTS.c.id == post_id))
    return r.title


# ------------------------------------------------------------------ evrak deposu


def _doc_out(r: Any) -> dict[str, Any]:
    return {"id": r.id, "category": r.category, "code": r.code, "title": r.title, "filename": r.filename, "mime": r.mime,
            "size": int(r.size or 0), "uploadedBy": r.uploaded_by, "uploadedAt": iso(r.uploaded_at)}


_DOC_COLS = [DOCS.c.id, DOCS.c.category, DOCS.c.code, DOCS.c.title, DOCS.c.filename, DOCS.c.mime, DOCS.c.size,
             DOCS.c.uploaded_by, DOCS.c.uploaded_at]


def list_docs(engine: sa.engine.Engine, tenant: str) -> list[dict[str, Any]]:
    ensure(engine)
    with engine.connect() as c:
        rows = c.execute(sa.select(*_DOC_COLS).where(DOCS.c.tenant_id == tenant).order_by(DOCS.c.category, DOCS.c.code, DOCS.c.title)).all()
    return [_doc_out(r) for r in rows]


def add_doc(engine: sa.engine.Engine, tenant: str, actor: str, meta: dict[str, str], filename: str, data: bytes,
            max_mb: int) -> dict[str, Any]:
    ensure(engine)
    st = settings(engine, tenant)
    title = clean(meta.get("title"), 200)
    cat = clean(meta.get("category"), 60)
    if not title:
        raise HrError("Evrakın adı boş olamaz.")
    if cat not in st["docCategories"]:
        raise HrError("Kategori listede yok; listeyi Ayarlar sekmesinden düzenleyin.")
    if not data:
        raise HrError("Dosya boş.")
    if len(data) > max_mb * 1024 * 1024:
        raise HrError(f"Dosya {max_mb} MB sınırını aşıyor.", 413)
    name = clean(filename, 300) or "evrak"
    mime = _mime(name)
    did = new_id("evr")
    with engine.begin() as c:
        c.execute(DOCS.insert().values(id=did, tenant_id=tenant, category=cat, code=clean(meta.get("code"), 40) or None, title=title,
                                       filename=name, mime=mime, size=len(data), blob=data, uploaded_by=actor, uploaded_at=now()))
        r = c.execute(sa.select(*_DOC_COLS).where(DOCS.c.id == did)).first()
    return _doc_out(r)


def doc_blob(engine: sa.engine.Engine, tenant: str, did: str) -> tuple[bytes, str, str]:
    ensure(engine)
    with engine.connect() as c:
        r = c.execute(sa.select(DOCS.c.blob, DOCS.c.filename, DOCS.c.mime).where(DOCS.c.id == did, DOCS.c.tenant_id == tenant)).first()
    if r is None:
        raise HrError("Evrak bulunamadı.", 404)
    return bytes(r.blob), r.filename, r.mime or "application/octet-stream"


def delete_doc(engine: sa.engine.Engine, tenant: str, did: str) -> str:
    ensure(engine)
    with engine.begin() as c:
        r = c.execute(sa.select(DOCS.c.title).where(DOCS.c.id == did, DOCS.c.tenant_id == tenant)).first()
        if r is None:
            raise HrError("Evrak bulunamadı.", 404)
        c.execute(DOCS.delete().where(DOCS.c.id == did))
    return r.title


# ------------------------------------------------------------------ evrak talebi


def _req_out(r: Any) -> dict[str, Any]:
    return {"id": r.id, "username": r.username, "display": r.display, "mail": r.mail, "docType": r.doc_type,
            "delivery": r.delivery, "deliveryLabel": DELIVERY.get(r.delivery, r.delivery), "note": r.note,
            "status": r.status, "statusLabel": REQUEST_STATUS.get(r.status, r.status), "answer": r.answer,
            "createdAt": iso(r.created_at), "handledBy": r.handled_by, "handledAt": iso(r.handled_at)}


def create_request(engine: sa.engine.Engine, tenant: str, user: str, display: str, body: dict[str, Any]) -> dict[str, Any]:
    ensure(engine)
    st = settings(engine, tenant)
    typ = clean(body.get("docType"), 120)
    if typ not in st["docTypes"]:
        raise HrError("Evrak tipini listeden seçin.")
    dl = str(body.get("delivery") or "")
    if dl not in DELIVERY:
        raise HrError("Teslim şeklini seçin.")
    mail = clean(body.get("mail"), 200).lower()
    if dl == "eposta" and not _EMAIL.match(mail):
        raise HrError("E-posta ile teslim için geçerli bir e-posta adresi girin.")
    rid = new_id("evt")
    with engine.begin() as c:
        c.execute(REQUESTS.insert().values(id=rid, tenant_id=tenant, username=user, display=clean(display, 200) or user,
                                           mail=mail or None, doc_type=typ, delivery=dl,
                                           note=str(body.get("note") or "").strip()[:2000] or None, status="bekliyor",
                                           created_at=now()))
        r = c.execute(sa.select(REQUESTS).where(REQUESTS.c.id == rid)).first()
    return _req_out(r)


def list_requests(engine: sa.engine.Engine, tenant: str, *, user: str = "", status: str = "") -> list[dict[str, Any]]:
    """Talepler (en yeni önce); İK'nın eklediği hazır evrakların adı ve boyutu dahil (bayt değil)."""
    ensure(engine)
    stmt = sa.select(REQUESTS).where(REQUESTS.c.tenant_id == tenant)
    if user:
        stmt = stmt.where(REQUESTS.c.username == user)
    if status == "acik":
        stmt = stmt.where(REQUESTS.c.status.in_(("bekliyor", "hazirlaniyor")))
    elif status:
        stmt = stmt.where(REQUESTS.c.status == status)
    with engine.connect() as c:
        rows = c.execute(stmt.order_by(REQUESTS.c.created_at.desc())).all()
        files = c.execute(sa.select(REQUEST_FILES.c.id, REQUEST_FILES.c.request_id, REQUEST_FILES.c.filename, REQUEST_FILES.c.size)
                          .where(REQUEST_FILES.c.tenant_id == tenant, REQUEST_FILES.c.request_id.in_([r.id for r in rows] or [""]))).all()
    out = []
    for r in rows:
        o = _req_out(r)
        o["files"] = [{"id": f.id, "filename": f.filename, "size": int(f.size or 0)} for f in files if f.request_id == r.id]
        out.append(o)
    return out


def get_request(engine: sa.engine.Engine, tenant: str, rid: str) -> dict[str, Any]:
    ensure(engine)
    with engine.connect() as c:
        r = c.execute(sa.select(REQUESTS).where(REQUESTS.c.id == rid, REQUESTS.c.tenant_id == tenant)).first()
    if r is None:
        raise HrError("Talep bulunamadı.", 404)
    return _req_out(r)


def add_request_file(engine: sa.engine.Engine, tenant: str, actor: str, rid: str, filename: str, data: bytes, max_mb: int) -> dict[str, Any]:
    """İK'nın hazırladığı belge talebe eklenir; çalışan portaldan indirir (e-postaya ek olarak gitmez)."""
    get_request(engine, tenant, rid)
    if not data:
        raise HrError("Dosya boş.")
    if len(data) > max_mb * 1024 * 1024:
        raise HrError(f"Dosya {max_mb} MB sınırını aşıyor.", 413)
    name = clean(filename, 300) or "evrak"
    mime = _mime(name)
    fid = new_id("evd")
    with engine.begin() as c:
        c.execute(REQUEST_FILES.insert().values(id=fid, tenant_id=tenant, request_id=rid, filename=name, mime=mime, size=len(data),
                                                blob=data, uploaded_by=actor, uploaded_at=now()))
    return {"id": fid, "filename": name, "size": len(data)}


def request_file(engine: sa.engine.Engine, tenant: str, rid: str, fid: str, *, user: str = "") -> tuple[bytes, str, str]:
    """`user` verilirse yalnız o kişinin talebinin belgesi döner."""
    ensure(engine)
    with engine.connect() as c:
        req = c.execute(sa.select(REQUESTS.c.username).where(REQUESTS.c.id == rid, REQUESTS.c.tenant_id == tenant)).first()
        r = c.execute(sa.select(REQUEST_FILES).where(REQUEST_FILES.c.id == fid, REQUEST_FILES.c.request_id == rid,
                                                     REQUEST_FILES.c.tenant_id == tenant)).first()
    if req is None or r is None or (user and req.username != user):
        raise HrError("Belge bulunamadı.", 404)
    return bytes(r.blob), r.filename, r.mime or "application/octet-stream"


def delete_request_file(engine: sa.engine.Engine, tenant: str, rid: str, fid: str) -> str:
    ensure(engine)
    with engine.begin() as c:
        r = c.execute(sa.select(REQUEST_FILES.c.filename).where(REQUEST_FILES.c.id == fid, REQUEST_FILES.c.request_id == rid,
                                                                REQUEST_FILES.c.tenant_id == tenant)).first()
        if r is None:
            raise HrError("Belge bulunamadı.", 404)
        c.execute(REQUEST_FILES.delete().where(REQUEST_FILES.c.id == fid))
    return r.filename


def handle_request(engine: sa.engine.Engine, tenant: str, actor: str, rid: str, body: dict[str, Any]) -> dict[str, Any]:
    ensure(engine)
    stt = str(body.get("status") or "")
    if stt not in REQUEST_STATUS:
        raise HrError("Durum geçersiz.")
    answer = str(body.get("answer") or "").strip()[:2000] or None
    if stt == "red" and not answer:
        raise HrError("Karşılanamayan talebe kısa bir gerekçe yazın; çalışan bunu görür.")
    with engine.begin() as c:
        n = c.execute(REQUESTS.update().where(REQUESTS.c.id == rid, REQUESTS.c.tenant_id == tenant).values(
            status=stt, answer=answer, handled_by=actor, handled_at=now())).rowcount
        if not n:
            raise HrError("Talep bulunamadı.", 404)
        r = c.execute(sa.select(REQUESTS).where(REQUESTS.c.id == rid)).first()
    return _req_out(r)


def cancel_request(engine: sa.engine.Engine, tenant: str, user: str, rid: str) -> None:
    """Çalışan, İK henüz ele almadıysa kendi talebini geri alır."""
    ensure(engine)
    with engine.begin() as c:
        r = c.execute(sa.select(REQUESTS.c.status).where(REQUESTS.c.id == rid, REQUESTS.c.tenant_id == tenant,
                                                         REQUESTS.c.username == user)).first()
        if r is None:
            raise HrError("Talep bulunamadı.", 404)
        if r.status != "bekliyor":
            raise HrError("İK talebi ele aldı; geri almak için İK ile görüşün.")
        c.execute(REQUESTS.delete().where(REQUESTS.c.id == rid))


# ------------------------------------------------------------------ yemek listesi


def menu(engine: sa.engine.Engine, tenant: str, start: date, end: date) -> list[dict[str, Any]]:
    ensure(engine)
    with engine.connect() as c:
        rows = c.execute(sa.select(MENU).where(MENU.c.tenant_id == tenant, MENU.c.day >= start, MENU.c.day <= end)
                         .order_by(MENU.c.day)).all()
    return [{"day": r.day.isoformat(), "items": [x for x in r.items.split("\n") if x.strip()], "updatedBy": r.updated_by}
            for r in rows]


def save_menu(engine: sa.engine.Engine, tenant: str, actor: str, body: dict[str, Any]) -> int:
    """Gün gün menü: `days: [{day, items: [..]}]`; boş liste o günün menüsünü siler."""
    ensure(engine)
    days = body.get("days")
    if not isinstance(days, list) or not days:
        raise HrError("En az bir gün gönderin.")
    n = 0
    with engine.begin() as c:
        for d in days:
            day = H.parse_date((d or {}).get("day"), "Menü")
            if day is None:
                raise HrError("Menü günü boş.")
            items = [clean(x, 120) for x in ((d or {}).get("items") or []) if clean(x, 120)]
            c.execute(MENU.delete().where(MENU.c.tenant_id == tenant, MENU.c.day == day))
            if items:
                c.execute(MENU.insert().values(tenant_id=tenant, day=day, items="\n".join(items), updated_by=actor, updated_at=now()))
            n += 1
    return n


# ------------------------------------------------------------------ sık sorulan sorular


def _faq_out(r: Any) -> dict[str, Any]:
    return {"id": r.id, "category": r.category, "question": r.question, "answer": r.answer, "sort": int(r.sort or 0),
            "updatedBy": r.updated_by, "updatedAt": iso(r.updated_at)}


def list_faq(engine: sa.engine.Engine, tenant: str) -> list[dict[str, Any]]:
    ensure(engine)
    order = {c: i for i, c in enumerate(settings(engine, tenant)["faqCategories"])}
    with engine.connect() as c:
        rows = [_faq_out(r) for r in c.execute(sa.select(FAQ).where(FAQ.c.tenant_id == tenant)).all()]
    rows.sort(key=lambda x: (order.get(x["category"], 999), x["category"], x["sort"], x["question"].casefold()))
    return rows


def save_faq(engine: sa.engine.Engine, tenant: str, actor: str, body: dict[str, Any], fid: Optional[str] = None) -> dict[str, Any]:
    ensure(engine)
    st = settings(engine, tenant)
    cat = clean(body.get("category"), 80)
    q = clean(body.get("question"), 300)
    a = str(body.get("answer") or "").strip()[:8000]
    if cat not in st["faqCategories"]:
        raise HrError("Kategori listede yok; listeyi Ayarlar sekmesinden düzenleyin.")
    if not q or not a:
        raise HrError("Soru ve cevap boş olamaz.")
    try:
        sort = int(body.get("sort") or 0)
    except (TypeError, ValueError):
        sort = 0
    vals = dict(category=cat, question=q, answer=a, sort=sort, updated_by=actor, updated_at=now())
    with engine.begin() as c:
        if fid:
            if not c.execute(FAQ.update().where(FAQ.c.id == fid, FAQ.c.tenant_id == tenant).values(**vals)).rowcount:
                raise HrError("Soru bulunamadı.", 404)
        else:
            fid = new_id("sss")
            c.execute(FAQ.insert().values(id=fid, tenant_id=tenant, **vals))
        r = c.execute(sa.select(FAQ).where(FAQ.c.id == fid)).first()
    return _faq_out(r)


def delete_faq(engine: sa.engine.Engine, tenant: str, fid: str) -> str:
    ensure(engine)
    with engine.begin() as c:
        r = c.execute(sa.select(FAQ.c.question).where(FAQ.c.id == fid, FAQ.c.tenant_id == tenant)).first()
        if r is None:
            raise HrError("Soru bulunamadı.", 404)
        c.execute(FAQ.delete().where(FAQ.c.id == fid))
    return r.question
