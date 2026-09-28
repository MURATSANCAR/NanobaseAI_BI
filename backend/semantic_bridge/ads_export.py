"""M21 Reklam raporu: PDF (tek sayfa hedefli aylık rapor) ve Excel (özet, kanallar, kampanyalar, kitaplar, günlük).

PDF yazı tipi çözümü projedeki tek sunucu tarafı PDF yoluyla aynıdır (`editorial_export._resolve_fonts`, fpdf2 +
DejaVu). Metinde teknoloji adı geçmez; ürün adı «Zeki AI». Satış verisinin hangi güne kadar geldiği her çıktıda yazılır.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any, Optional

from semantic_bridge import ads as A

PRODUCT = "Zeki AI"


def _money(v: Any) -> str:
    if v is None:
        return "—"
    return f"{float(v):,.0f}".replace(",", ".") + " TL"


def _num(v: Any, digits: int = 0) -> str:
    if v is None:
        return "—"
    s = f"{float(v):,.{digits}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _ratio(v: Any) -> str:
    return "—" if v is None else _num(v, 2)


def _day(v: Optional[str]) -> str:
    if not v:
        return "—"
    d = date.fromisoformat(v[:10])
    return f"{d.day:02d}.{d.month:02d}.{d.year}"


def data_note(ov: dict[str, Any]) -> str:
    end = ov.get("veriSonu")
    vd = ov.get("verimDonemi")
    if not end:
        return "Logo satış verisi okunamadı: e-ticaret cirosu ve pazarlama verimi hesaplanmadı."
    s = f"Satış verisi {_day(end)} tarihine kadar."
    if vd:
        s += f" Pazarlama verimi {_day(vd['bas'])}–{_day(vd['bit'])} günlerinin harcama ve cirosuyla hesaplandı."
    else:
        s += " Seçilen dönemde satış verisi yok; pazarlama verimi hesaplanmadı."
    return s


def report_pdf(ov: dict[str, Any], suggestions: list[dict[str, Any]], comment: Optional[str], user: str = "") -> bytes:
    try:
        from fpdf import FPDF
    except ImportError as e:  # pragma: no cover
        raise A.AdsError("PDF üretici sunucuda kurulu değil.", 503) from e
    from semantic_bridge.editorial_export import _fold, _resolve_fonts

    regular, bold = _resolve_fonts()
    fam = "Body" if regular else "Helvetica"
    T = (lambda s: str(s or "")) if regular else (lambda s: _fold(str(s or "")))
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_title(T("Dijital pazarlama ve reklam raporu"))
    pdf.set_author(PRODUCT)
    pdf.set_creator(PRODUCT)
    pdf.set_margins(14, 14, 14)
    pdf.set_auto_page_break(auto=True, margin=14)
    if regular:
        pdf.add_font(fam, "", str(regular))
        pdf.add_font(fam, "B", str(bold or regular))
    pdf.add_page()
    W = pdf.w - pdf.l_margin - pdf.r_margin

    def para(s: str, size: float = 9.5, style: str = "", gap: float = 1.5) -> None:
        pdf.set_font(fam, style, size)
        pdf.multi_cell(W, size * 0.5, T(s), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    def head(s: str) -> None:
        pdf.ln(2)
        para(s, 12, "B", 1)

    def row(cells: list[tuple[str, float]], style: str = "") -> None:
        pdf.set_font(fam, style, 8.5)
        for txt, w in cells:
            t = T(txt)
            while pdf.get_string_width(t) > w - 1.5 and len(t) > 1:
                t = t[:-2] + "…" if len(t) > 2 else t[:-1]
            pdf.cell(w, 5.5, t, border="B")
        pdf.ln(5.5)

    d = ov["donem"]
    g = ov["gosterge"]
    para("Timaş Yayınları · Dijital pazarlama ve reklam", 8.5)
    para(f"Reklam raporu: {_day(d['bas'])} – {_day(d['bit'])}" + (f" · {A.PLATFORMS.get(d['kanal'], d['kanal'])}" if d.get("kanal") else ""), 15, "B", 1)
    para(data_note(ov), 8.5)

    head("Göstergeler")
    row([("Harcama", 36), ("Tıklama", 26), ("TBM", 24), ("Platform ROAS", 30), ("E-ticaret net ciro", 36), ("Pazarlama verimi", 30)], "B")
    row([(_money(g["harcama"]), 36), (_num(g.get("tiklama")), 26), (_num(g.get("tbm"), 2) + " TL" if g.get("tbm") is not None else "—", 24),
         (_ratio(g.get("platformRoas")), 30), (_money(g.get("eticaretCiro")), 36), (_ratio(g.get("verim")), 30)])
    para("Platform ROAS platformun kendi bildirdiği dönüşüm değeridir; platformlar aynı satışı birden çok kez sahiplenebilir. "
         "Pazarlama verimi = Logo e-ticaret kanalı net cirosu ÷ reklam harcaması.", 8)
    if ov.get("digerParaBirimi"):
        para("TL dışı harcama (toplama katılmadı): " + ", ".join(f"{k} {_num(v, 2)}" for k, v in ov["digerParaBirimi"].items()), 8)
    b = ov.get("bagsiz") or {}
    para(f"Kitaba bağlı olmayan harcama: {_money(b.get('harcama'))}" + (f" (%{_num((b.get('pay') or 0) * 100, 1)})" if b.get("pay") is not None else ""), 9)

    if comment:
        head("Zeki AI yorumu")
        para(comment, 9)

    head("Kanallar")
    row([("Kanal", 50), ("Harcama", 30), ("Pay", 18), ("Tıklama", 24), ("TBM", 24), ("Platform ROAS", 36)], "B")
    for ch in ov["kanallar"]:
        row([(ch["kanalAdi"], 50), (_money(ch["harcama"]), 30), (f"%{_num((ch.get('pay') or 0) * 100, 1)}", 18), (_num(ch.get("tiklama")), 24),
             (_num(ch.get("tbm"), 2), 24), (_ratio(ch.get("platformRoas")), 36)])

    head("Kitaplar (bağlı kampanyalar)")
    if not ov["kitaplar"]:
        para("Dönemde kitaba bağlı kampanya harcaması yok.")
    else:
        row([("Kitap", 70), ("Harcama", 28), ("E-ticaret ciro", 32), ("Verim", 18), ("Stok (gün)", 34)], "B")
        for k in ov["kitaplar"]:
            s = k.get("stok") or {}
            stock = "—" if not s else (f"{_num(s.get('bakiye'))} adet" + (f" · {_num(s.get('gun'))} gün" if s.get("gun") is not None else ""))
            row([(k.get("ad") or k["stokKodu"], 70), (_money(k["harcama"]), 28), (_money((k.get("satis") or {}).get("eticaretCiro")), 32),
                 (_ratio(k.get("verim")), 18), (stock, 34)])

    head("Kampanyalar (harcamaya göre)")
    row([("Kampanya", 72), ("Kanal", 26), ("Harcama", 28), ("Tıklama", 22), ("Platform ROAS", 34)], "B")
    for c in ov["kampanyalar"]:
        row([(c["ad"], 72), (c.get("platformAdi") or "—", 26), (_money(c["harcama"]), 28), (_num(c.get("tiklama")), 22),
             (_ratio(c.get("platformRoas")), 34)])

    open_ = [s for s in suggestions if s["durum"] in A.OPEN]
    head("Açık öneri ve uyarılar")
    if not open_:
        para("Açık öneri yok.")
    for s in open_:
        para(f"{s['turAdi']} · {s['durumAdi']}: {s['gerekce']}", 8.5, "", 1)

    pdf.ln(2)
    para("Reklam platformlarında hiçbir değişiklik portaldan yapılmaz; öneriler onaylandıktan sonra ekip tarafından uygulanır.", 8)
    para(f"Hazırlayan: {user or '—'} · {PRODUCT}", 8)
    return bytes(pdf.output())


def report_xlsx(ov: dict[str, Any], suggestions: list[dict[str, Any]], comment: Optional[str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    d, g = ov["donem"], ov["gosterge"]
    rows = [
        ["Dönem", d["bas"], d["bit"]],
        ["Kanal", A.PLATFORMS.get(d.get("kanal") or "", "Hepsi")],
        ["Satış verisi sonu", ov.get("veriSonu") or "—"],
        ["Verim dönemi", (ov.get("verimDonemi") or {}).get("bas") or "—", (ov.get("verimDonemi") or {}).get("bit") or "—"],
        [],
        ["Harcama (TL)", g["harcama"]], ["Gösterim", g.get("gosterim")], ["Tıklama", g.get("tiklama")], ["TBM (TL)", g.get("tbm")],
        ["Dönüşüm", g.get("donusum")], ["Dönüşüm değeri (platform, TL)", g.get("donusumDegeri")], ["Platform ROAS", g.get("platformRoas")],
        ["E-ticaret net ciro (Logo, TL)", g.get("eticaretCiro")], ["Verim dönemindeki harcama (TL)", g.get("harcamaVeriIcinde")],
        ["Pazarlama verimi", g.get("verim")], ["Kitaba bağsız harcama (TL)", (ov.get("bagsiz") or {}).get("harcama")],
        [], ["Not", data_note(ov)],
    ]
    if comment:
        rows.append(["Zeki AI yorumu", comment])
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18

    def sheet(title: str, head: list[str], body: list[list[Any]]) -> None:
        s = wb.create_sheet(title)
        s.append(head)
        for c in s[1]:
            c.font = Font(bold=True)
        for r in body:
            s.append(r)
        s.freeze_panes = "A2"

    sheet("Kanallar", ["Kanal", "Harcama (TL)", "Pay", "Gösterim", "Tıklama", "TBM (TL)", "Dönüşüm", "Dönüşüm değeri", "Platform ROAS"],
          [[c["kanalAdi"], c["harcama"], c.get("pay"), c.get("gosterim"), c.get("tiklama"), c.get("tbm"), c.get("donusum"),
            c.get("donusumDegeri"), c.get("platformRoas")] for c in ov["kanallar"]])
    sheet("Kampanyalar", ["Kampanya", "Kanal", "Hesap", "Kitap", "Stok kodu", "Bağ", "Harcama (TL)", "Gösterim", "Tıklama", "TBM (TL)",
                          "Dönüşüm", "Dönüşüm değeri", "Platform ROAS", "Son gün"],
          [[c["ad"], c.get("platformAdi"), c.get("hesap"), c.get("kitapAdi"), c.get("stokKodu"), c.get("bagAdi"), c["harcama"], c.get("gosterim"),
            c.get("tiklama"), c.get("tbm"), c.get("donusum"), c.get("donusumDegeri"), c.get("platformRoas"), c.get("sonGun")] for c in ov["kampanyalar"]])
    sheet("Kitaplar", ["Stok kodu", "Kitap", "Kampanya", "Harcama (TL)", "Verim dönemindeki harcama", "E-ticaret net ciro", "Toplam net ciro",
                       "Pazarlama verimi", "Stok bakiyesi", "Günlük satış", "Stok (gün)", "M15 plan reklam satırı (TL)", "Yayın durumu"],
          [[k["stokKodu"], k.get("ad"), k.get("kampanya"), k["harcama"], k.get("harcamaVeriIcinde"), (k.get("satis") or {}).get("eticaretCiro"),
            (k.get("satis") or {}).get("toplamCiro"), k.get("verim"), (k.get("stok") or {}).get("bakiye"), (k.get("stok") or {}).get("gunlukSatis"),
            (k.get("stok") or {}).get("gun"), (k.get("m15") or {}).get("tutar"), k.get("yayinDurumu")] for k in ov["kitaplar"]])
    sheet("Günlük", ["Gün", "Harcama (TL)", "E-ticaret net ciro (TL)"], [[x["gun"], x["harcama"], x["eticaretCiro"]] for x in ov["gunluk"]])
    sheet("Öneriler", ["Tür", "Durum", "Kampanya", "Kitap", "Gerekçe", "Karar veren", "Karar notu", "Uygulayan", "Zaman"],
          [[s["turAdi"], s["durumAdi"], s.get("kampanya"), s.get("kitapAdi"), s["gerekce"], s.get("karar"), s.get("kararNotu"), s.get("uygulayan"),
            s.get("zaman")] for s in suggestions])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
