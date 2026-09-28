"""M42 bildirimleri ve dışa aktarım.

- **Haftalık uyarı** (pazartesi): platformun dönem iade oranı eşiği (`CHANNEL_RETURN_THRESHOLD`, varsayılan %10 — ölçülecek)
  aşarsa ya da iskonto oranı geçen yılın aynı dönemine göre `CHANNEL_DISCOUNT_RISE_PTS` (2 puan) üstü artarsa. Portal
  kaydı `meta:alerts`; alıcı (`CHANNEL_REPORT_RECIPIENTS`) varsa tek özet e-postası.
- **Aylık karne** (ayın 2'si): önceki ay sonuna kadar yıl içi karne tablosu; Zeki AI rakamsız 5–8 cümlelik yorum ekler
  (rakam içeren satır atılır; rakamlar tablodan). Aynı ay ikinci kez gönderilmez (`meta:report:<ay>`).
- **Excel**: karne, kitaplar, iadeler, matris, eşleme.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import date, timedelta
from typing import Any, Callable, Optional

from semantic_bridge.channels import mapping as M
from semantic_bridge.channels import scorecard as SC
from semantic_bridge.channels import store as S

log = logging.getLogger("semantic.channels.report")

KARNE_PROMPT = (
    "Aşağıda bir yayınevinin e-ticaret kanallarının (pazar yerleri ve kendi sitesi) yıl içi karnesi var. Yönetim için "
    "5–8 cümlelik bir «kanallarda ne oldu» özeti yaz: hangi kanal büyüdü ya da küçüldü, iade ve iskonto nerede dikkat "
    "istiyor. Rakam, yüzde, tutar yazma (rakamlar tabloda); kanıtsız yorum ekleme; kanala yapılan satışı «kanalın "
    "sattığı» diye anlatma.\n\n{facts}"
)


def _pct(x: Optional[float]) -> str:
    return "—" if x is None else f"%{x * 100:.1f}".replace(".", ",")


def _money(x: Optional[float]) -> str:
    return "—" if x is None else f"{x:,.0f} TL".replace(",", ".")


def alerts(engine: Any, tenant: str, st: dict[str, Any]) -> list[dict[str, Any]]:
    card = SC.scorecard(engine, tenant)
    out = []
    for x in card["platforms"]:
        if x["platform"] == M.UNMAPPED:
            continue
        d, prev = x["donem"], x["gecenYil"]
        if d["iadeOrani"] is not None and d["iadeOrani"] > st["returnThreshold"] and d["satisCiro"] > 0:
            out.append({"platform": x["platform"], "label": x["label"], "tur": "iade", "deger": d["iadeOrani"], "esik": st["returnThreshold"],
                        "metin": f"{x['label']}: iade oranı {_pct(d['iadeOrani'])} (eşik {_pct(st['returnThreshold'])})"})
        if prev and d["iskontoOrani"] is not None and prev["iskontoOrani"] is not None:
            rise = (d["iskontoOrani"] - prev["iskontoOrani"]) * 100
            if rise > st["discountRise"]:
                out.append({"platform": x["platform"], "label": x["label"], "tur": "iskonto", "deger": d["iskontoOrani"],
                            "onceki": prev["iskontoOrani"],
                            "metin": f"{x['label']}: iskonto oranı {_pct(prev['iskontoOrani'])} → {_pct(d['iskontoOrani'])} "
                                     f"(+{rise:.1f} puan)".replace(".", ",")})
    return out


def karne_table(card: dict[str, Any], with_margin: bool) -> str:
    p = card["period"]
    lines = [f"Dönem: {p['yil']} Ocak–{p['ayAdi']} (Logo verisi {p['veriSonu'] or '—'} tarihine kadar; kanala satış, sell-in)", ""]
    for x in card["platforms"]:
        d = x["donem"]
        bits = [f"net ciro {_money(d['netCiro'])}", f"geçen yıla göre {_pct(x['degisim'])}", f"iade {_pct(d['iadeOrani'])}",
                f"iskonto {_pct(d['iskontoOrani'])}"]
        if with_margin:
            bits.append(f"brüt marj {_pct(d['marj'])} (maliyetsiz ciro {_money(d['maliyetsizCiro'])})")
        if x.get("hedef"):
            bits.append(f"CRM hedef gerçekleşme {_pct(x['hedef']['oran'])}")
        lines.append(f"- {x['label']}: " + ", ".join(bits))
    t = card["toplam"]
    lines += ["", f"E-ticaret toplamı {_money(t['eticaret']['netCiro'])}; şirket içindeki payı {_pct(t['eticaretPay'])}; "
                  f"D2C'nin e-ticaret içindeki payı {_pct(t['d2cPay'])}."]
    return "\n".join(lines)


def monthly(engine: Any, tenant: str, st: dict[str, Any], llm: Any, send: Callable[[str, str, list[str]], str],
            today: Optional[date] = None, force: bool = False) -> dict[str, Any]:
    today = today or date.today()
    last = today.replace(day=1) - timedelta(days=1)
    key = f"report:{last.year}-{last.month:02d}"
    if not force and S.meta_get(engine, tenant, key):
        return {"skipped": "bu ayın karnesi gönderildi"}
    if not st["recipients"]:
        return {"skipped": "alıcı yok (CHANNEL_REPORT_RECIPIENTS)"}
    card = SC.scorecard(engine, tenant, last.year, last.month)
    table = karne_table(card, with_margin=True)
    note = None
    if llm is not None:
        try:
            text = (llm.chat([{"role": "user", "content": KARNE_PROMPT.format(facts=table)}], max_tokens=500, temperature=0.2) or "").strip()
            note = re.sub(r"[^\n]*\d[^\n]*\n?", "", text).strip() or None
        except Exception as e:  # noqa: BLE001 — model yoksa karne yorumsuz gider
            log.warning("channels: karne yorumu yazılamadı: %s", e)
    body = table + ("\n\nZeki AI yorumu:\n" + note if note else "")
    status = send(f"Kanal karnesi — {SC.AY[last.month - 1]} {last.year}", body, st["recipients"])
    S.meta_set(engine, tenant, key, {"status": status, "alici": len(st["recipients"]), "yorum": bool(note)})
    return {"status": status, "yorum": bool(note)}


def weekly(engine: Any, tenant: str, st: dict[str, Any], send: Callable[[str, str, list[str]], str],
           today: Optional[date] = None, force: bool = False) -> dict[str, Any]:
    today = today or date.today()
    iso = today.isocalendar()
    key = f"weekly:{iso[0]}-{iso[1]:02d}"
    if not force and S.meta_get(engine, tenant, key):
        return {"skipped": "bu hafta gönderildi"}
    items = alerts(engine, tenant, st)
    S.meta_set(engine, tenant, "alerts", {"items": items, "tarih": today.isoformat()})
    status = "no_alert"
    if items and st["recipients"]:
        status = send("Kanal uyarıları", "\n".join(f"- {x['metin']}" for x in items), st["recipients"])
    elif items:
        status = "no_recipient"
    S.meta_set(engine, tenant, key, {"status": status, "uyari": len(items)})
    return {"status": status, "uyari": len(items)}


# ------------------------------------------------------------------ Excel


def xlsx(title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]], note: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="555555")
        start = 3
    for j, (_, label) in enumerate(columns, 1):
        c = ws.cell(row=start, column=j, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="6D28D9")
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(48, len(label) + 4))
    for i, r in enumerate(rows, start + 1):
        for j, (key, _) in enumerate(columns, 1):
            v = r
            for part in key.split("."):
                v = v.get(part) if isinstance(v, dict) else None
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
