"""Platform panelinden indirilen satış raporunun (Excel/CSV) yüklenmesi — kanalın son tüketiciye sattığı adet
(sell-through) ve kanal stoğu. Kullanıcı kararı (2026-09-28): Trendyol/Amazon satış modeli sonraya; M40–M42 yalnız okuma +
Excel yükleme. M40/M41 kendi rapor türleri için bu ayrıştırıcıyı kullanır.

Şema: başlık satırı ilk 20 satırda aranır; tanınan kolonlar yalnız ürün kimliği (barkod / stok kodu / ürün adı), adet,
tutar ve kanal stoğudur. **Başka hiçbir kolon içeri alınmaz** (müşteri adı, adres, telefon gibi kişisel alanlar dahil);
atlanan kolon adları yükleme kaydına yazılır, değerleri yazılmaz.
"""
from __future__ import annotations

import csv
import io
import json
import re
from typing import Any, Optional

import sqlalchemy as sa

from semantic_bridge.channels import mapping as M
from semantic_bridge.channels import store as S

MAX_MB = 25
#: Kolon adı eşleşmeleri (Türkçe harfler sadeleştirilmiş, küçük harf). Önce tam eşleşme, sonra «içerir».
COLUMNS: dict[str, tuple[str, ...]] = {
    "barkod": ("barkod", "barcode", "ean", "isbn", "urun barkodu", "barkod no"),
    "stok_kodu": ("stok kodu", "satici stok kodu", "urun kodu", "sku", "model kodu", "satici urun kodu", "stock code"),
    "ad": ("urun adi", "urun", "baslik", "kitap adi", "product name", "title", "urun ismi"),
    "adet": ("satis adedi", "satilan adet", "net satis adedi", "adet", "miktar", "quantity", "satis miktari", "units sold",
             "siparis adedi"),
    "tutar": ("satis tutari", "net satis tutari", "tutar", "ciro", "toplam tutar", "revenue", "sales"),
    "kanal_stok": ("stok", "stok adedi", "mevcut stok", "kalan stok", "stock", "depo stogu", "satilabilir stok"),
}
#: Bu sözcükleri içeren kolon kişisel veri olabilir: adı kayda «kişisel olabilir» diye düşer.
PERSONAL = ("musteri", "alici", "adres", "telefon", "e-posta", "eposta", "email", "tc", "kimlik", "ad soyad", "isim")


class ImportError_(ValueError):
    pass


def _cell(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def _num(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "").replace("₺", "").replace("TL", "")
    if not t:
        return None
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def read_rows(filename: str, data: bytes) -> list[list[Any]]:
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ImportError_(f"Excel dosyası açılamadı: {str(e)[:120]}") from None
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if ext in ("csv", "txt"):
        for enc in ("utf-8-sig", "cp1254", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        # Ayırıcı: ilk 4 KB'ta en sık geçen aday (Türkçe dışa aktarımda «;» ve ondalık «,» birlikte gelir; Sniffer yanılır).
        head = text[:4096]
        delim = max((";", "\t", "|", ","), key=lambda d: head.count(d))
        return [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    raise ImportError_("Yalnız .xlsx ya da .csv dosyası yüklenir.")


def _match(header: str) -> Optional[str]:
    h = M.fold(header)
    if not h:
        return None
    for key, names in COLUMNS.items():
        if h in names:
            return key
    for key, names in COLUMNS.items():
        if key == "kanal_stok":
            continue  # «stok» kısa ve «stok kodu»yla karışır: yalnız tam eşleşme
        if any(n in h for n in names if len(n) >= 4):
            return key
    return None


def parse(filename: str, data: bytes) -> dict[str, Any]:
    rows = read_rows(filename, data)
    head_i, cols = None, {}
    for i, r in enumerate(rows[:20]):
        found: dict[str, int] = {}
        for j, v in enumerate(r):
            k = _match(_cell(v))
            if k and k not in found:
                found[k] = j
        if "adet" in found and ({"barkod", "stok_kodu"} & set(found)):
            head_i, cols = i, found
            break
    if head_i is None:
        raise ImportError_("Başlık satırı bulunamadı: dosyada adet kolonu ve barkod ya da stok kodu kolonu olmalı.")
    header = [_cell(v) for v in rows[head_i]]
    used = set(cols.values())
    skipped = [h for j, h in enumerate(header) if h and j not in used]
    personal = [h for h in skipped if any(w in M.fold(h) for w in PERSONAL)]
    out = []
    for r in rows[head_i + 1:]:
        def get(k: str) -> Any:
            j = cols.get(k)
            return r[j] if j is not None and j < len(r) else None
        adet = _num(get("adet"))
        bk = _cell(get("barkod")).replace(" ", "")
        sk = _cell(get("stok_kodu"))
        if adet is None or not (bk or sk):
            continue
        out.append({"barkod": bk[:40] or None, "stok_kodu": sk[:60] or None, "ad": _cell(get("ad"))[:400] or None,
                    "adet": adet, "tutar": _num(get("tutar")), "kanal_stok": _num(get("kanal_stok"))})
    return {"rows": out, "columns": {k: header[j] for k, j in cols.items()}, "skipped": skipped, "personal": personal}


def store_import(engine: sa.engine.Engine, tenant: str, user: str, platform: str, filename: str, data: bytes,
                 donem_bas: str = "", donem_bit: str = "") -> dict[str, Any]:
    if platform not in M.PLATFORMS or platform == "degil":
        raise ImportError_("Platform seçilmeli.")
    for d in (donem_bas, donem_bit):
        if d and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", d):
            raise ImportError_("Dönem tarihi YYYY-AA-GG olmalı.")
    if donem_bas and donem_bit and donem_bas > donem_bit:
        raise ImportError_("Dönem başlangıcı bitişten sonra olamaz.")
    if not data:
        raise ImportError_("Dosya boş.")
    parsed = parse(filename or "", data)
    with engine.connect() as c:
        bmap = {r.barkod: r.stok_kodu for r in c.execute(sa.select(S.BARCODES).where(S.BARCODES.c.tenant_id == tenant)).all()}
        known = {r[0] for r in c.execute(sa.select(S.BOOKS.c.stok_kodu).where(S.BOOKS.c.tenant_id == tenant)).all()}
    iid = S.new_id()
    items = []
    matched = 0
    for i, r in enumerate(parsed["rows"]):
        code = r["stok_kodu"] if r["stok_kodu"] in known else bmap.get(r["barkod"] or "") or (r["barkod"] if r["barkod"] in known else None)
        if code:
            matched += 1
        items.append({"import_id": iid, "sira": i + 1, "stok_kodu": code, "barkod": r["barkod"], "ad": r["ad"],
                      "adet": r["adet"], "tutar": r["tutar"], "kanal_stok": r["kanal_stok"]})
    with engine.begin() as c:
        c.execute(S.IMPORTS.insert().values(
            id=iid, tenant_id=tenant, platform=platform, tur="sell-through", dosya_adi=(filename or "")[:300],
            donem_bas=donem_bas or None, donem_bit=donem_bit or None, satir=len(items), eslesen=matched,
            atlanan_kolonlar=json.dumps({"atlanan": parsed["skipped"], "kisiselOlabilir": parsed["personal"],
                                         "taninan": parsed["columns"]}, ensure_ascii=False),
            yukleyen=user, tarih=S.now()))
        for i in range(0, len(items), 5000):
            c.execute(S.IMPORT_ROWS.insert(), items[i:i + 5000])
    return get(engine, tenant, iid)


def _view(r: Any) -> dict[str, Any]:
    return {"id": r.id, "platform": r.platform, "tur": r.tur, "dosya": r.dosya_adi, "donemBas": r.donem_bas, "donemBit": r.donem_bit,
            "satir": r.satir, "eslesen": r.eslesen, "kolonlar": S.jload(r.atlanan_kolonlar, {}), "yukleyen": r.yukleyen,
            "tarih": S.iso(r.tarih)}


def list_imports(engine: sa.engine.Engine, tenant: str, platform: str = "") -> list[dict[str, Any]]:
    stmt = sa.select(S.IMPORTS).where(S.IMPORTS.c.tenant_id == tenant)
    if platform:
        stmt = stmt.where(S.IMPORTS.c.platform == platform)
    with engine.connect() as c:
        return [_view(r) for r in c.execute(stmt.order_by(S.IMPORTS.c.tarih.desc())).all()]


def get(engine: sa.engine.Engine, tenant: str, iid: str, with_rows: bool = False) -> dict[str, Any]:
    with engine.connect() as c:
        r = c.execute(sa.select(S.IMPORTS).where(S.IMPORTS.c.tenant_id == tenant, S.IMPORTS.c.id == iid)).first()
        if not r:
            raise ImportError_("Yükleme bulunamadı.")
        out = _view(r)
        if with_rows:
            rows = c.execute(sa.select(S.IMPORT_ROWS).where(S.IMPORT_ROWS.c.import_id == iid).order_by(S.IMPORT_ROWS.c.sira)).all()
            out["rows"] = [{"sira": x.sira, "stokKodu": x.stok_kodu, "barkod": x.barkod, "ad": x.ad, "adet": x.adet,
                            "tutar": x.tutar, "kanalStok": x.kanal_stok} for x in rows]
    return out


def delete(engine: sa.engine.Engine, tenant: str, iid: str) -> dict[str, Any]:
    out = get(engine, tenant, iid)
    with engine.begin() as c:
        c.execute(S.IMPORT_ROWS.delete().where(S.IMPORT_ROWS.c.import_id == iid))
        c.execute(S.IMPORTS.delete().where(S.IMPORTS.c.tenant_id == tenant, S.IMPORTS.c.id == iid))
    return out


def sell_through(engine: sa.engine.Engine, tenant: str, iid: str, sell_in: dict[str, float]) -> dict[str, Any]:
    """Yüklenen kanal satışı (sell-through) ile TİMAŞ'ın aynı döneme ait kanala satışı (sell-in, aylık önbellekten)."""
    out = get(engine, tenant, iid, with_rows=True)
    agg: dict[str, dict[str, Any]] = {}
    unmatched = 0
    for r in out.pop("rows"):
        if not r["stokKodu"]:
            unmatched += 1
            continue
        a = agg.setdefault(r["stokKodu"], {"stokKodu": r["stokKodu"], "ad": r["ad"], "kanalSatis": 0.0, "kanalStok": None})
        a["kanalSatis"] += float(r["adet"] or 0)
        if r["kanalStok"] is not None:
            a["kanalStok"] = (a["kanalStok"] or 0.0) + float(r["kanalStok"])
    names = S.book_names(engine, tenant, list(agg))
    items = []
    for code, a in agg.items():
        si = sell_in.get(code, 0.0)
        items.append({**a, "ad": names.get(code) or a["ad"] or "", "kanalaSatis": round(si, 2),
                      "oran": (a["kanalSatis"] / si) if si > 0 else None})
    items.sort(key=lambda x: -x["kanalSatis"])
    return {**out, "items": items, "eslesmeyen": unmatched}
