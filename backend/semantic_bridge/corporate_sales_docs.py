"""M32 teklif belgesi: PDF (kuruma gidecek biçim) ve Excel (teklif + sipariş satırları).

Belgede yalnız kurumun görmesi gerekenler vardır: kitap, adet, liste fiyatı, indirim, net tutar. Maliyet ve marj iç bilgidir,
belgeye yazılmaz (ekranda kalır). Onaylanmamış teklifin (taslak, onay bekliyor) her sayfasında «TASLAK» yazar. Belge hiçbir yere
gönderilmez; indirilir, gönderimi temsilci yapar. PDF yolu `editorial_export` ile aynıdır (fpdf2 + Unicode TTF).
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from typing import Any, Optional

from semantic_bridge.editorial_export import _fold, _resolve_fonts

INK = (24, 22, 46)
MUTED = (110, 110, 135)
ACCENT = (124, 92, 255)
LINE = (226, 226, 240)
SOFT = (246, 244, 255)
DRAFT = (185, 28, 28)
PAGE_W = 210.0
MARGIN = 16.0
CONTENT_W = PAGE_W - 2 * MARGIN
FINAL = ("hazir", "gonderildi", "kabul", "ret")


def money(v: Optional[float]) -> str:
    if v is None:
        return "—"
    s = f"{float(v):,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".") + " TL"


def pct(v: Optional[float]) -> str:
    if v is None:
        return "—"
    s = f"{float(v) * 100:.1f}".rstrip("0").rstrip(".")
    return "%" + s.replace(".", ",")


def quote_no(q: dict[str, Any]) -> str:
    return f"KT-{q['firsatId'][:6].upper()}-{q['surum']}"


def _issue_day(q: dict[str, Any]) -> date:
    for k in ("gonderildiAt", "onayAt", "gonderimAt", "createdAt"):
        v = q.get(k)
        if v:
            try:
                return datetime.fromisoformat(v).date()
            except ValueError:
                continue
    return date.today()


def _vat_note(lines: list[dict[str, Any]]) -> str:
    flags = {l.get("kdvDahil") for l in lines}
    if flags == {True}:
        return "Fiyatlar KDV dahildir."
    if flags == {False}:
        return "Fiyatlara KDV dahil değildir."
    return "KDV durumu kalem bazında fiyat listesine göredir."


def pdf(q: dict[str, Any], company: str) -> bytes:
    try:
        from fpdf import FPDF
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("PDF üretici sunucuda kurulu değil.") from e
    regular, bold = _resolve_fonts()
    uni = regular is not None
    fam = "TeklifSans" if uni else "Helvetica"
    T = (lambda s: s) if uni else _fold
    opp = q.get("firsat") or {}
    lines = q.get("kalemler") or []
    draft = q.get("durum") not in FINAL
    issued = _issue_day(q)
    valid = issued + timedelta(days=int(q.get("gecerlilikGun") or 30))

    class QuotePdf(FPDF):
        def header(self) -> None:
            self.set_font(fam, "B", 9)
            self.set_text_color(*ACCENT)
            self.set_xy(MARGIN, 9)
            self.cell(CONTENT_W / 2, 5, T(company), align="L")
            self.set_font(fam, "", 8.5)
            self.set_text_color(*MUTED)
            self.cell(CONTENT_W / 2, 5, T(f"Teklif {quote_no(q)}"), align="R")
            self.set_draw_color(*ACCENT)
            self.set_line_width(0.5)
            self.line(MARGIN, 15.5, PAGE_W - MARGIN, 15.5)
            if draft:
                self.set_font(fam, "B", 9)
                self.set_text_color(*DRAFT)
                self.set_xy(MARGIN, 17)
                self.cell(CONTENT_W, 5, T("TASLAK — onaylanmamış teklif, kuruma gönderilmez"), align="C")
            self.set_y(24)
            self.set_text_color(*INK)

        def footer(self) -> None:
            self.set_y(-13)
            self.set_font(fam, "", 7.5)
            self.set_text_color(*MUTED)
            self.cell(CONTENT_W * 0.75, 5, T("Teklif; muhasebe ve hukuk onayıyla kesinleşir. Stok durumu teklif tarihindeki durumdur."), align="L")
            self.cell(CONTENT_W * 0.25, 5, T(f"Sayfa {self.page_no()} / {{nb}}"), align="R")

    p = QuotePdf(orientation="P", unit="mm", format="A4")
    p.alias_nb_pages()
    p.set_title(f"Teklif {quote_no(q)}")
    p.set_author(company)
    p.set_creator(company)
    p.set_margins(MARGIN, 24, MARGIN)
    p.set_auto_page_break(auto=True, margin=18)
    if uni:
        p.add_font(fam, "", str(regular))
        p.add_font(fam, "B", str(bold or regular))
    p.add_page()

    p.set_font(fam, "B", 16)
    p.multi_cell(CONTENT_W, 8, T(opp.get("kurum") or "Kurum"), align="L")
    p.set_font(fam, "", 10)
    p.set_text_color(*MUTED)
    p.multi_cell(CONTENT_W, 5.5, T(f"{opp.get('ad') or ''}" + (f" · {opp['tema']}" if opp.get("tema") else "")), align="L")
    p.multi_cell(CONTENT_W, 5.5, T(f"Teklif tarihi {issued.strftime('%d.%m.%Y')} · Geçerlilik {valid.strftime('%d.%m.%Y')}"
                                   + (f" · {q['paketAdet']} paket" if q.get("paketAdet") else "")), align="L")
    p.set_text_color(*INK)
    p.ln(3)

    if q.get("mektup"):
        p.set_font(fam, "", 10)
        for para in [x.strip() for x in str(q["mektup"]).split("\n\n") if x.strip()]:
            p.multi_cell(CONTENT_W, 5.4, T(para), align="L")
            p.ln(2)
        p.ln(1)

    cols = [("#", 8, "C"), ("Kitap", 62, "L"), ("Yazar", 30, "L"), ("Adet", 14, "R"), ("Liste fiyatı", 20, "R"),
            ("İndirim", 14, "R"), ("Tutar", 30, "R")]

    def head() -> None:
        p.set_font(fam, "B", 8.5)
        p.set_fill_color(*SOFT)
        p.set_text_color(*MUTED)
        for name, w, al in cols:
            p.cell(w, 7, T(name), border=0, align=al, fill=True)
        p.ln(7)
        p.set_text_color(*INK)

    head()
    p.set_font(fam, "", 8.5)
    for i, l in enumerate(lines, 1):
        title = (l.get("ad") or l["stok"])
        if p.get_y() > 265:
            p.add_page()
            head()
            p.set_font(fam, "", 8.5)
        vals = [str(i), title[:60], (l.get("yazar") or "")[:28], f"{int(l['adet']):,}".replace(",", "."),
                money(l.get("listeFiyati")), pct(l.get("indirim")), money(l.get("netTutar"))]
        for (name, w, al), v in zip(cols, vals):
            p.cell(w, 6.2, T(v), border="B", align=al)
        p.ln(6.2)
    p.ln(3)
    p.set_font(fam, "", 9.5)
    total_w = CONTENT_W - 60
    for label, value, strong in (("Liste fiyatıyla toplam", money(q.get("toplamListe")), False),
                                 ("İndirim", pct(q.get("indirimOrani")), False),
                                 ("Teklif tutarı", money(q.get("toplamNet")), True)):
        p.set_font(fam, "B" if strong else "", 11 if strong else 9.5)
        p.cell(total_w, 6.5, "", align="R")
        p.cell(30, 6.5, T(label), align="L")
        p.cell(30, 6.5, T(value), align="R")
        p.ln(6.5)
    p.ln(3)
    p.set_font(fam, "", 8.5)
    p.set_text_color(*MUTED)
    p.multi_cell(CONTENT_W, 4.8, T(_vat_note(lines) + " Liste fiyatı yayınevinin geçerli satış fiyat listesidir."), align="L")
    if q.get("notlar"):
        p.ln(1)
        p.multi_cell(CONTENT_W, 4.8, T(str(q["notlar"])), align="L")
    return bytes(p.output())


def xlsx(q: dict[str, Any], company: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    opp = q.get("firsat") or {}
    lines = q.get("kalemler") or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Teklif"
    issued = _issue_day(q)
    ws["A1"] = company
    ws["A1"].font = Font(bold=True, size=12, color="7C5CFF")
    ws["A2"] = f"Teklif {quote_no(q)} · {opp.get('kurum') or ''}"
    ws["A2"].font = Font(bold=True, size=14)
    ws["A3"] = (f"{opp.get('ad') or ''} · Teklif tarihi {issued.strftime('%d.%m.%Y')} · Geçerlilik {int(q.get('gecerlilikGun') or 30)} gün"
                + (" · TASLAK (onaylanmadı)" if q.get("durum") not in FINAL else ""))
    head = ["Sıra", "Stok kodu", "Kitap", "Yazar", "Adet", "Liste fiyatı", "İndirim", "Net birim", "Tutar"]
    r0 = 5
    for j, h in enumerate(head, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = Font(bold=True, color="605E5C")
        c.fill = PatternFill("solid", fgColor="F6F4FF")
    for i, l in enumerate(lines, 1):
        row = [i, l["stok"], l.get("ad"), l.get("yazar"), int(l["adet"]), l.get("listeFiyati"), l.get("indirim"), l.get("netBirim"),
               l.get("netTutar")]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r0 + i, column=j, value=v)
            if j in (6, 8, 9):
                c.number_format = '#,##0.00 "₺"'
            if j == 7:
                c.number_format = "0.0%"
    end = r0 + len(lines)
    for k, (label, val, fmt) in enumerate((("Liste fiyatıyla toplam", q.get("toplamListe"), '#,##0.00 "₺"'),
                                           ("İndirim", q.get("indirimOrani"), "0.0%"),
                                           ("Teklif tutarı", q.get("toplamNet"), '#,##0.00 "₺"')), 2):
        ws.cell(row=end + k, column=8, value=label).font = Font(bold=True)
        c = ws.cell(row=end + k, column=9, value=val)
        c.number_format = fmt
        c.font = Font(bold=True)
    ws.cell(row=end + 6, column=1, value=_vat_note(lines)).alignment = Alignment(wrap_text=False)
    for j, w in enumerate((6, 16, 46, 26, 9, 14, 10, 14, 16), 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    ws2 = wb.create_sheet("Sipariş satırları")
    ws2["A1"] = "Kabul edilen teklifin CRM'de sipariş olarak açılması için satırlar (portal CRM'e yazmaz)."
    for j, h in enumerate(["Stok kodu", "Kitap", "Adet", "Net birim fiyat", "İndirim", "Fiyat listesi"], 1):
        ws2.cell(row=3, column=j, value=h).font = Font(bold=True)
    for i, l in enumerate(lines, 4):
        for j, v in enumerate([l["stok"], l.get("ad"), int(l["adet"]), l.get("netBirim"), l.get("indirim"), l.get("fiyatListe")], 1):
            c = ws2.cell(row=i, column=j, value=v)
            if j == 4:
                c.number_format = "#,##0.00"
            if j == 5:
                c.number_format = "0.0%"
    for j, w in enumerate((16, 46, 9, 16, 10, 16), 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
