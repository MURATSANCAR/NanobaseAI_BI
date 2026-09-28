"""M21 Reklam: dışa aktarım dosyasının okunması ve Logo/CRM sorguları (yalnız okuma).

**Dosya.** Platformların dışa aktarım biçimleri farklıdır (Google Ads CSV'si UTF-16 ve sekmeli gelebilir, başında rapor
adı/tarih aralığı satırları ve sonunda «Toplam» satırları olur; Meta CSV'si virgüllüdür, «Raporlama başlangıcı» kolonuyla
günlük kırılım verir; Excel de yüklenebilir). Başlık satırı, bilinen kolon adlarından en az ikisinin geçtiği ilk satırdır.
Kolon eşlemesi önce hesabın kayıtlı eşlemesinden, sonra kolon adı sözlüğünden önerilir; zorunlu alan (gün, kampanya,
harcama) bulunamazsa Zeki AI kolonlar arasından kapalı seçim yapar. Eşleme insan onayıyla kaydedilir.

Bir veri satırı okunamazsa (tarih ya da sayı çözülemedi) yükleme yapılmaz; hatalı satır sayısı ve örnekleri gösterilir.
Böylece dosyanın harcama toplamı yazılan toplamla her zaman aynıdır (kabul 3).

**Logo** (`SEMANTIC_CONNECTION_FILE`): e-ticaret kanalı günlük net cirosu, bağlı kitapların günlük cirosu (e-ticaret ve
bütün kanallar), güncel kopyada stok bakiyesi. Satış satırı tanımı M46 ile aynı (`budget_sources` başlığı); kanal cari
kartın `SPECODE2`'si. Yıllar ayrı firma numarasıdır (`budget_sources.firms_by_year`).

**CRM** (`SEMANTIC_CRM_CONNECTION_FILE`): kitap listesi (eşleştirme; yayıncılık durumu etiketiyle), reklam planları
(`new_reklamplaniBase`, mecra/tip adları, bağlı kitaplar), pazarlama bütçe modülü kayıtları. Yazma yok.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import threading
import time
from datetime import date, datetime, timedelta
from typing import Any, Callable, Optional

from semantic_bridge import ads as A
from semantic_bridge import budget_sources as bsrc

Runner = Callable[[str], list[dict[str, Any]]]
_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_CHANNEL = re.compile(r"^[A-Z0-9 _\-.İŞĞÜÖÇ]{1,40}$")
_CODE = re.compile(r"^[0-9A-Za-z._/\- ]{1,60}$")
CRM_TTL = 1800


class SourceError(RuntimeError):
    pass


# ================================================================== dosya


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    if data[:3] == b"\xef\xbb\xbf":
        return data[3:].decode("utf-8")
    # BOM'suz UTF-16: her ikinci bayt sıfır
    if len(data) > 4 and data[1:2] == b"\x00" and data[3:4] == b"\x00":
        return data.decode("utf-16-le")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1254")


def read_table(name: str, data: bytes) -> list[list[Any]]:
    """Dosyanın bütün satırları (hücre listesi). Excel ise ilk dolu sayfa."""
    if not data:
        raise A.AdsError("Dosya boş.")
    low = (name or "").lower()
    if data[:2] == b"PK" or low.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as e:  # pragma: no cover
            raise A.AdsError("Excel okuyucu sunucuda kurulu değil; dosyayı CSV olarak yükleyin.", 503) from e
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 — bozuk dosya
            raise A.AdsError(f"Excel dosyası okunamadı: {str(e)[:120]}") from None
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(v not in (None, "") for v in r)]
            if rows:
                return rows
        raise A.AdsError("Excel dosyasında veri yok.")
    if low.endswith(".xls"):
        raise A.AdsError("Eski Excel (.xls) biçimi okunmuyor; .xlsx ya da CSV olarak kaydedip yükleyin.")
    text = _decode(data)
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        raise A.AdsError("Dosyada veri yok.")
    delim = delimiter(lines[:40])
    return [row for row in csv.reader(io.StringIO("\n".join(lines)), delimiter=delim) if any(c.strip() for c in row)]


def delimiter(lines: list[str]) -> str:
    """Ayırıcı: aynı (sıfırdan büyük) sayıda geçtiği satırı en çok olan aday; eşitlikte satır başına sayısı çok olan.
    Rapor başlığı satırları (ayırıcısız) ve kampanya adındaki «|» gibi tek tük geçişler kararı bozmaz."""
    best, score = ",", (-1, -1)
    for d in ("\t", ";", ",", "|"):
        counts: dict[int, int] = {}
        for ln in lines:
            n = ln.count(d)
            if n:
                counts[n] = counts.get(n, 0) + 1
        if not counts:
            continue
        n, lines_with = max(counts.items(), key=lambda kv: (kv[1], kv[0]))
        if (lines_with, n) > score:
            best, score = d, (lines_with, n)
    return best


def norm_header(h: Any) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", A.fold(h)).split())


#: Kolon adı sözlüğü (katlanmış, noktalama sökülmüş). Tam eşleşme önce; sonra önek (birim başı ölçüler hariç).
SYNONYMS: dict[str, tuple[str, ...]] = {
    "day": ("gun", "day", "date", "tarih", "reporting starts", "raporlama baslangici", "rapor baslangici", "stat time day",
            "by day", "donem baslangici", "gunluk"),
    "day_end": ("reporting ends", "raporlama bitisi", "rapor bitisi", "donem bitisi"),
    "campaign": ("campaign", "kampanya", "campaign name", "kampanya adi", "kampanya ismi"),
    "campaign_id": ("campaign id", "kampanya kimligi", "kampanya id", "campaign no", "kampanya no"),
    "spend": ("cost", "maliyet", "spend", "harcama", "amount spent", "harcanan tutar", "harcanan", "toplam maliyet", "total cost",
              "toplam harcama"),
    "impressions": ("impr", "impressions", "impression", "gosterim", "gosterimler", "gosterim sayisi"),
    "clicks": ("clicks", "click", "tiklama", "tiklamalar", "link clicks", "baglanti tiklamalari", "tiklama sayisi"),
    "conversions": ("conversions", "conversion", "donusum", "donusumler", "results", "sonuclar", "purchases", "satin almalar",
                    "satin alma"),
    "conv_value": ("conv value", "conversion value", "donusum degeri", "purchase conversion value", "purchases conversion value",
                   "satin alma donusum degeri", "all conv value", "total purchase value", "website purchases conversion value",
                   "conversions value", "revenue", "gelir"),
    "currency": ("currency", "currency code", "para birimi", "para birimi kodu"),
    "status": ("campaign state", "campaign status", "kampanya durumu", "status", "durum", "delivery", "yayin durumu", "yayin"),
}
#: Birim başı / oran ölçüleri: «Ort. TBM», «Tıklama oranı», «Dönüşüm başına maliyet» harcama ya da tıklama değildir.
_PER_UNIT_WORDS = frozenset(("per", "avg", "ort", "ctr", "cpc", "cpm", "cpa", "roas", "tbm", "bgbm", "ebm"))
_PER_UNIT_STEMS = ("basina", "ortalama", "oran", "rate", "yuzde", "percent", "frekans", "frequency")


def match(header: Any) -> tuple[Optional[str], bool]:
    """(alan, tam eşleşme mi). Önek eşleşmesi birim başı ölçülerde yapılmaz."""
    raw = str(header or "")
    h = norm_header(raw)
    if not h:
        return None, False
    for f, words in SYNONYMS.items():
        if h in words:
            return f, True
    if "/" in raw or "%" in raw or any(t in _PER_UNIT_WORDS or t.startswith(_PER_UNIT_STEMS) for t in h.split()):
        return None, False
    for f, words in SYNONYMS.items():
        for w in words:
            if len(w) >= 4 and (h.startswith(w + " ") or h.endswith(" " + w)):
                return f, False
    return None, False


def field_of(header: Any) -> Optional[str]:
    return match(header)[0]


def find_header(table: list[list[Any]]) -> int:
    for i, row in enumerate(table[:25]):
        hits = {field_of(c) for c in row} - {None}
        if len(hits) >= 2 and ("campaign" in hits or "spend" in hits):
            return i
    return 0


def suggest_mapping(headers: list[str], saved: Optional[dict[str, str]]) -> tuple[dict[str, str], str]:
    """Önce hesabın kayıtlı eşlemesi (bütün kolonları dosyada varsa), sonra sözlük: tam eşleşen kolon önce, önekle
    eşleşen yalnız alan boş kaldıysa («Kampanya türü» «Kampanya»nın yerini almaz)."""
    hs = [str(h) for h in headers]
    if saved and all(v in hs for v in saved.values()):
        return dict(saved), "kayitli"
    out: dict[str, str] = {}
    for exact in (True, False):
        for h in hs:
            f, ex = match(h)
            if f and ex == exact and f not in out and h not in out.values():
                out[f] = h
    return out, "sozluk"


MODEL_NONE = "Bu dosyada yok"


def model_mapping(llm: Any, headers: list[str], sample: list[list[Any]], fields: list[str], st: dict[str, Any]) -> dict[str, Any]:
    """Sözlüğün bulamadığı zorunlu alanlar için Zeki AI: kolon adları + ilk satırlardan örnek değerlerle kapalı seçim.
    Olasılık eşiği bağ eşiğiyle aynı (`ADS_LINK_MIN_PROB/MARGIN`); altıysa öneri yazılmaz."""
    out: dict[str, Any] = {}
    if llm is None or not hasattr(llm, "choose") or not headers:
        return out
    cols = []
    for i, h in enumerate(headers):
        vals = [str(r[i]) for r in sample[:3] if i < len(r) and r[i] not in (None, "")]
        cols.append(f"- {h}: {', '.join(v[:40] for v in vals) or 'boş'}")
    for f in fields:
        prompt = ("Bir reklam platformunun dışa aktardığı rapor dosyasının kolonları ve ilk satırlardaki örnek değerleri:\n"
                  + "\n".join(cols) + f"\n\nHangi kolon «{A.FIELDS[f]}» bilgisini taşıyor?"
                  + (" (Harcama: reklama ödenen toplam tutar; birim başı maliyet değil.)" if f == "spend" else "")
                  + (" (Gün: satırın ait olduğu tarih.)" if f == "day" else ""))
        ch = llm.choose(prompt, [*headers, MODEL_NONE])
        if ch.choice and ch.choice != MODEL_NONE and ch.confident(st["linkMinProb"], st["linkMinMargin"]):
            out[f] = {"kolon": ch.choice, "olasilik": ch.probability}
    return out


_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
    "oca": 1, "sub": 2, "nis": 4, "haz": 6, "tem": 7, "agu": 8, "eyl": 9, "eki": 10, "kas": 11, "ara": 12,
}


def parse_day(v: Any, day_first: bool = True) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < float(v) < 80000:        # Excel gün sayısı
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s) or re.match(r"^(\d{4})/(\d{1,2})/(\d{1,2})", s)
    if m:
        return _mk(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{4})(\d{2})(\d{2})$", s)
    if m:
        return _mk(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})", s)
    if m:
        return _mk(int(m[3]), int(m[2]), int(m[1]))
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if m:
        a, b = int(m[1]), int(m[2])
        d_, mo = (a, b) if day_first else (b, a)
        if mo > 12 and d_ <= 12:
            d_, mo = mo, d_
        return _mk(int(m[3]), mo, d_)
    f = A.fold(s).replace(",", " ")
    m = re.match(r"^([a-z]+)\.?\s+(\d{1,2})\s+(\d{4})", f)                 # Sep 1 2026
    if m and m[1][:3] in _MONTHS:
        return _mk(int(m[3]), _MONTHS[m[1][:3]], int(m[2]))
    m = re.match(r"^(\d{1,2})\s+([a-z]+)\.?\s+(\d{4})", f)                 # 1 Eyl 2026 / 1 Eylül 2026
    if m and m[2][:3] in _MONTHS:
        return _mk(int(m[3]), _MONTHS[m[2][:3]], int(m[1]))
    return None


def _mk(y: int, m: int, d: int) -> Optional[date]:
    try:
        return date(y, m, d)
    except ValueError:
        return None


_EMPTY = {"", "-", "--", "—", "–", "n/a", "na", "yok", "none", "null"}


def decimal_sep(table: list[list[Any]], cols: list[int]) -> str:
    """Dosyanın ondalık ayırıcısı: «12,5» / «12.50» biçimli hücrelerin çoğunluğu. Karar veremezse virgül (Türkçe)."""
    comma = dot = 0
    for row in table:
        for i in cols:
            if i >= len(row) or not isinstance(row[i], str):
                continue
            s = row[i].strip()
            if re.search(r"\d,\d{1,2}$", s) or re.search(r"\d\.\d{3},\d", s):
                comma += 1
            elif re.search(r"\d\.\d{1,2}$", s) or re.search(r"\d,\d{3}\.\d", s):
                dot += 1
    return "." if dot > comma else ","


def parse_num(v: Any, dec: str) -> tuple[Optional[float], bool]:
    """(sayı, geçerli mi). Boş/tire → (None, True)."""
    if v is None:
        return None, True
    if isinstance(v, bool):
        return None, False
    if isinstance(v, (int, float)):
        return float(v), True
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if s.lower() in _EMPTY:
        return None, True
    s = re.sub(r"(?i)(try|tl|usd|eur|₺|\$|€|%)", "", s)
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if dec == ",":
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        n = float(s)
    except ValueError:
        return None, False
    return (-n if neg else n), True


def apply_mapping(table: list[list[Any]], header_row: int, mapping: dict[str, str], day_first: bool) -> dict[str, Any]:
    """Eşlemeyle satırları çözer. Dönen: rows (normal biçim), hatalar (sayı + örnekler), atlanan özet satırı sayısı."""
    headers = [str(h if h is not None else "") for h in table[header_row]]
    idx: dict[str, int] = {}
    for f, col in mapping.items():
        if col not in headers:
            raise A.AdsError(f"Eşlemedeki «{col}» kolonu dosyada yok.")
        idx[f] = headers.index(col)
    data = table[header_row + 1:]
    num_cols = [idx[f] for f in ("spend", "impressions", "clicks", "conversions", "conv_value") if f in idx]
    dec = decimal_sep(data, num_cols)
    rows: list[dict[str, Any]] = []
    errors: list[str] = []
    skipped_totals = 0
    not_daily = 0
    read = 0

    def cell(r: list[Any], f: str) -> Any:
        i = idx.get(f)
        return r[i] if i is not None and i < len(r) else None

    for n, r in enumerate(data, start=header_row + 2):
        camp = str(cell(r, "campaign") or "").strip()
        dv = cell(r, "day")
        first = A.fold(next((c for c in r if c not in (None, "")), ""))
        # Özet satırı: kampanya adı boş (Meta'nın üstteki «bütün sonuçlar» satırı) ya da «Toplam/Total» ile başlıyor
        # (Google Ads'in alttaki toplamları). Sayısı ekranda «özet satırı» diye görünür; kampanyaya yazılamaz.
        if not camp or A.fold(camp).startswith(("total", "toplam", "genel toplam")) or first.startswith(("total", "toplam")):
            skipped_totals += 1
            continue
        read += 1
        d = parse_day(dv, day_first)
        if d is None:
            errors.append(f"{n}. satır: gün okunamadı ({str(dv)[:30]!r})")
            continue
        if "day_end" in idx:
            de = parse_day(cell(r, "day_end"), day_first)
            if de is not None and de != d:
                not_daily += 1
                errors.append(f"{n}. satır: günlük değil ({d.isoformat()} – {de.isoformat()})")
                continue
        out: dict[str, Any] = {"day": d.isoformat(), "campaign": camp[:400],
                               "campaign_id": str(cell(r, "campaign_id") or "").strip()[:120] or None}
        bad = False
        for f in ("spend", "impressions", "clicks", "conversions", "conv_value"):
            if f not in idx:
                out[f] = None
                continue
            val, ok = parse_num(cell(r, f), dec)
            if not ok:
                errors.append(f"{n}. satır: «{A.FIELDS[f]}» sayı değil ({str(cell(r, f))[:30]!r})")
                bad = True
                break
            out[f] = val
        if bad:
            continue
        out["spend"] = out["spend"] or 0.0
        if out["spend"] < 0:
            errors.append(f"{n}. satır: harcama eksi")
            continue
        cur = str(cell(r, "currency") or "").strip().upper()
        out["currency"] = {"TL": "TRY", "₺": "TRY"}.get(cur, cur) or None
        out["status"] = str(cell(r, "status") or "").strip()[:80] or None
        rows.append(out)
    return {"rows": rows, "hataSayisi": len(errors), "hatalar": errors[:20], "ozetSatiri": skipped_totals,
            "gunlukDegil": not_daily, "okunan": read, "ondalik": dec}


def preview(name: str, data: bytes, saved: Optional[dict[str, str]], day_first: bool,
            override: Optional[dict[str, str]] = None, header_row: Optional[int] = None) -> dict[str, Any]:
    """Başlık satırı, kolonlar, ilk 10 satır, önerilen eşleme ve (zorunlu alanlar tamsa) deneme okuması. `override`:
    ekranda elle düzeltilen eşleme; `header_row` (1'den): elle seçilen başlık satırı."""
    table = read_table(name, data)
    h = header_row - 1 if header_row and 0 < header_row <= len(table) else find_header(table)
    headers = [str(x if x is not None else "") for x in table[h]]
    if override:
        mapping = {k: str(v) for k, v in override.items() if k in A.FIELDS and v and str(v) in headers}
        source = "elle"
    else:
        mapping, source = suggest_mapping(headers, saved)
    sample = [[("" if v is None else (v.isoformat() if isinstance(v, (date, datetime)) else v)) for v in r]
              for r in table[h + 1:h + 11]]
    out: dict[str, Any] = {"baslikSatiri": h + 1, "kolonlar": headers, "ornek": sample, "eslem": mapping, "eslemKaynagi": source,
                           "eksik": [f for f in A.REQUIRED if f not in mapping], "satirSayisi": max(0, len(table) - h - 1)}
    if not out["eksik"]:
        out["deneme"] = trial(table, h, mapping, day_first)
    return out


def trial(table: list[list[Any]], header_row: int, mapping: dict[str, str], day_first: bool) -> dict[str, Any]:
    """Eşlemeyle deneme okuması: yazılacak satır, harcama toplamı (para birimi başına), dönem, hatalar."""
    res = apply_mapping(table, header_row, mapping, day_first)
    totals: dict[str, float] = {}
    for r in res["rows"]:
        totals[r["currency"] or "?"] = totals.get(r["currency"] or "?", 0.0) + r["spend"]
    days = sorted({r["day"] for r in res["rows"]})
    return {"satir": len(res["rows"]), "hataSayisi": res["hataSayisi"], "hatalar": res["hatalar"],
            "ozetSatiri": res["ozetSatiri"], "gunlukDegil": res["gunlukDegil"], "ondalik": res["ondalik"],
            "kampanya": len({r["campaign_id"] or A.fold(r["campaign"]) for r in res["rows"]}),
            "toplam": {k: round(v, 2) for k, v in totals.items()}, "bas": days[0] if days else None,
            "bit": days[-1] if days else None}


# ================================================================== Logo


def _firm(v: str) -> str:
    if not re.fullmatch(r"\d{3}", v or ""):
        raise SourceError("Geçersiz Logo firma numarası.")
    return v


def _in(values: list[str], rx: re.Pattern[str], what: str) -> str:
    out = []
    for v in values:
        if not rx.match(v):
            raise SourceError(f"Geçersiz {what}: {v!r}")
        out.append("N'" + v.replace("'", "''") + "'")
    if not out:
        raise SourceError(f"{what} listesi boş.")
    return ", ".join(out)


def ecom_daily_sql(firm: str, frm: date, to: date, channels: list[str]) -> str:
    """E-ticaret kanalı günlük net ciro ve adet (faturalı satır, iade eksi; M46 satış satırı tanımı)."""
    f = _firm(firm)
    return f"""
-- E-ticaret kanalı (cari özel kod 2) günlük net ciro: faturalı satış satırı, iade eksi, net ciro = LINENET.
SELECT CAST(S.DATE_ AS date) AS gun,
  SUM(CASE WHEN S.TRCODE IN (7,8,9) THEN S.LINENET ELSE -S.LINENET END) AS ciro,
  SUM(CASE WHEN S.TRCODE IN (7,8,9) THEN S.AMOUNT ELSE -S.AMOUNT END) AS adet
FROM dbo.LG_{f}_01_STLINE AS S
JOIN dbo.LG_{f}_CLCARD AS C ON C.LOGICALREF = S.CLIENTREF
WHERE S.CANCELLED = 0 AND S.LINETYPE = 0 AND S.INVOICEREF <> 0 AND S.TRCODE IN (2,3,7,8,9)
  AND C.SPECODE2 IN ({_in(channels, _CHANNEL, 'kanal kodu')})
  AND S.DATE_ >= '{frm.isoformat()}' AND S.DATE_ < '{(to + timedelta(days=1)).isoformat()}'
GROUP BY CAST(S.DATE_ AS date)""".strip()


def book_daily_sql(firm: str, frm: date, to: date, codes: list[str], channels: list[str]) -> str:
    """Bağlı kitapların günlük cirosu: e-ticaret kanalı ve bütün kanallar ayrı kolonda (tek okuma)."""
    f = _firm(firm)
    ch = _in(channels, _CHANNEL, "kanal kodu")
    sign = "(CASE WHEN S.TRCODE IN (7,8,9) THEN 1 ELSE -1 END)"
    return f"""
-- Kitap × gün: e-ticaret kanalı ve toplam net ciro/adet (faturalı satır, iade eksi).
SELECT I.CODE AS stok_kodu, CAST(S.DATE_ AS date) AS gun,
  SUM(CASE WHEN C.SPECODE2 IN ({ch}) THEN {sign} * S.LINENET ELSE 0 END) AS eticaret_ciro,
  SUM(CASE WHEN C.SPECODE2 IN ({ch}) THEN {sign} * S.AMOUNT ELSE 0 END) AS eticaret_adet,
  SUM({sign} * S.LINENET) AS toplam_ciro,
  SUM({sign} * S.AMOUNT) AS toplam_adet
FROM dbo.LG_{f}_01_STLINE AS S
JOIN dbo.LG_{f}_ITEMS AS I ON I.LOGICALREF = S.STOCKREF
LEFT JOIN dbo.LG_{f}_CLCARD AS C ON C.LOGICALREF = S.CLIENTREF
WHERE S.CANCELLED = 0 AND S.LINETYPE = 0 AND S.INVOICEREF <> 0 AND S.TRCODE IN (2,3,7,8,9)
  AND I.CODE IN ({_in(codes, _CODE, 'stok kodu')})
  AND S.DATE_ >= '{frm.isoformat()}' AND S.DATE_ < '{(to + timedelta(days=1)).isoformat()}'
GROUP BY I.CODE, CAST(S.DATE_ AS date)""".strip()


def stock_sql(firm: str, codes: list[str]) -> str:
    """Güncel kopyada stok bakiyesi (katalog «stok bakiyesi»: IOCODE 1/2 giriş − 3/4 çıkış, tarih süzgeci yok).
    Hareketi olmayan malzeme 0 (LEFT JOIN); Logo'da kartı olmayan kod dönmez (stok bilinmiyor sayılır)."""
    f = _firm(firm)
    return f"""
-- Stok bakiyesi (durum ölçüsü, güncel kopya, tarihsiz).
SELECT I.CODE AS stok_kodu,
  ISNULL(SUM(CASE WHEN S.IOCODE IN (1,2) THEN S.AMOUNT WHEN S.IOCODE IN (3,4) THEN -S.AMOUNT ELSE 0 END), 0) AS bakiye
FROM dbo.LG_{f}_ITEMS AS I
LEFT JOIN dbo.LG_{f}_01_STLINE AS S ON S.STOCKREF = I.LOGICALREF AND S.LINETYPE = 0 AND S.CANCELLED = 0 AND S.IOCODE IN (1,2,3,4)
WHERE I.CODE IN ({_in(codes, _CODE, 'stok kodu')})
GROUP BY I.CODE""".strip()


def _years(frm: date, to: date) -> list[tuple[int, date, date]]:
    return [(y, max(frm, date(y, 1, 1)), min(to, date(y, 12, 31))) for y in range(frm.year, to.year + 1)]


class Logo:
    """Logo okumaları; yıl → firma eşlemesi her yenilemede `L_CAPIPERIOD`'dan."""

    def __init__(self, path: Callable[[], str], runner: Callable[[str], Runner] = bsrc.runner):
        self.path = path
        self._runner = runner

    def run(self) -> Runner:
        try:
            return self._runner(self.path())
        except bsrc.SourceError as e:
            raise SourceError(f"Logo okunamıyor: {e}") from None

    @staticmethod
    def _call(run: Runner, sql: str) -> list[dict[str, Any]]:
        try:
            return run(sql)
        except bsrc.SourceError as e:
            raise SourceError(str(e)) from None

    def context(self) -> tuple[Runner, dict[int, str], Optional[date]]:
        run = self.run()
        try:
            firms = bsrc.firms_by_year(run)
            end = bsrc.read_data_end(run, firms) if firms else None
        except bsrc.SourceError as e:
            raise SourceError(str(e)) from None
        return run, firms, end

    def ecom_daily(self, run: Runner, firms: dict[int, str], frm: date, to: date, channels: list[str]) -> list[dict[str, Any]]:
        out = []
        for y, a, b in _years(frm, to):
            if y not in firms:
                continue
            for r in self._call(run, ecom_daily_sql(firms[y], a, b, channels)):
                d = bsrc._day(r.get("gun"))
                if d:
                    out.append({"day": d.isoformat(), "ciro": round(float(r.get("ciro") or 0), 4), "adet": float(r.get("adet") or 0)})
        return out

    def book_daily(self, run: Runner, firms: dict[int, str], frm: date, to: date, codes: list[str],
                   channels: list[str]) -> list[dict[str, Any]]:
        out = []
        for y, a, b in _years(frm, to):
            if y not in firms:
                continue
            for i in range(0, len(codes), 400):
                for r in self._call(run, book_daily_sql(firms[y], a, b, codes[i:i + 400], channels)):
                    d = bsrc._day(r.get("gun"))
                    code = str(r.get("stok_kodu") or "").strip()
                    if d and code:
                        out.append({"stok_kodu": code[:60], "day": d.isoformat(),
                                    "eticaret_ciro": round(float(r.get("eticaret_ciro") or 0), 4),
                                    "eticaret_adet": float(r.get("eticaret_adet") or 0),
                                    "toplam_ciro": round(float(r.get("toplam_ciro") or 0), 4), "toplam_adet": float(r.get("toplam_adet") or 0)})
        return out

    def stock(self, run: Runner, firms: dict[int, str], codes: list[str]) -> list[dict[str, Any]]:
        if not firms or not codes:
            return []
        firm = firms[max(firms)]
        out = []
        for i in range(0, len(codes), 400):
            for r in self._call(run, stock_sql(firm, codes[i:i + 400])):
                code = str(r.get("stok_kodu") or "").strip()
                if code:
                    out.append({"stok_kodu": code[:60], "bakiye": float(r.get("bakiye") or 0)})
        return out


# ================================================================== CRM


def prefix(schema: str) -> str:
    db, _, sch = (schema or "").strip().rpartition(".")
    for part in (db, sch):
        if part and not _NAME.match(part):
            raise SourceError(f"CRM şeması «{schema}» geçerli bir ad değil.")
    if not sch:
        raise SourceError("CRM şeması girilmemiş; CRM okunamıyor.")
    return (f"{db}." if db else "") + f"{sch}."


def _utc(d: date) -> str:
    """İstanbul günü başlangıcının UTC karşılığı (CRM tarihi UTC saklar)."""
    return (datetime(d.year, d.month, d.day) - timedelta(hours=3)).strftime("%Y-%m-%d %H:%M:%S")


def books_sql(schema: str) -> str:
    p = prefix(schema)
    return f"""
-- Eşleştirme için kitap kartları (etkin, stok kodlu) ve yayıncılık durumu etiketi (yalnız okuma).
SELECT k.new_kitapId AS kitap_id, k.new_StokKodu AS stok_kodu, k.new_name AS ad, k.new_ean13 AS ean,
       k.new_yazartext AS yazar, sm.Value AS durum
FROM {p}new_kitapBase AS k
LEFT JOIN {p}StringMapBase AS sm ON sm.AttributeName = 'new_kitap_yayincilikstatusu'
      AND sm.AttributeValue = k.new_kitap_yayincilikstatusu AND sm.LangId = 1055
      AND sm.ObjectTypeCode = (SELECT ObjectTypeCode FROM {p}EntityView WHERE Name = 'new_kitap')
WHERE k.statecode = 0 AND k.new_StokKodu IS NOT NULL AND k.new_StokKodu <> ''""".strip()


def ad_plans_sql(schema: str) -> str:
    p = prefix(schema)
    return f"""
-- CRM reklam planları (new_reklamplani), etkin kayıtlar; mecra ve tip adları kendi tablolarından (yalnız okuma).
SELECT r.new_reklamplaniId AS id, r.new_name AS ad, COALESCE(r.new_tutar_Base, r.new_Tutar) AS tutar,
       CAST(DATEADD(HOUR, 3, r.new_ReklamTarihi) AS date) AS bas, CAST(DATEADD(HOUR, 3, r.new_reklambitistarihi) AS date) AS bit,
       CAST(r.statuscode AS int) AS durum_kodu, m.new_name AS mecra, t.new_name AS tip,
       CAST(ISNULL(r.new_iptal, 0) AS int) AS iptal, CAST(ISNULL(r.new_faturasigirildi, 0) AS int) AS fatura,
       r.new_faturanumarasi AS fatura_no, CAST(ISNULL(r.new_reklamteslimedildimi, 0) AS int) AS teslim,
       CAST(DATEADD(HOUR, 3, r.new_OnayTarihi) AS date) AS onay
FROM {p}new_reklamplaniBase AS r
LEFT JOIN {p}new_reklammecrasiBase AS m ON m.new_reklammecrasiId = r.new_ReklamMecrasi
LEFT JOIN {p}new_reklamtipiBase AS t ON t.new_reklamtipiId = r.new_ReklamTipi
WHERE r.statecode = 0""".strip()


def ad_plan_books_sql(schema: str) -> str:
    p = prefix(schema)
    return f"""
-- Reklam planı ↔ kitap bağı (N:N).
SELECT l.new_reklamplaniid AS plan_id, k.new_StokKodu AS stok_kodu, k.new_name AS ad
FROM {p}new_new_reklamplani_new_kitapBase AS l
JOIN {p}new_kitapBase AS k ON k.new_kitapId = l.new_kitapid""".strip()


def budget_records_sql(schema: str, frm: date, to: date) -> str:
    p = prefix(schema)
    return f"""
-- CRM «Pazarlama Bütçe Modülü» kayıtları: dönemle kesişen etkin kayıtlar (başlangıç < dönem sonu, bitiş ≥ dönem başı).
SELECT m.new_pazarlamamoduluId AS id, m.new_name AS ad, CAST(m.new_pazarlamatipi AS int) AS tip,
       COALESCE(m.new_tutar_Base, m.new_tutar) AS tutar, m.new_tutar AS tutar_kayit,
       CAST(DATEADD(HOUR, 3, m.new_baslangictarihi) AS date) AS bas, CAST(DATEADD(HOUR, 3, m.new_bitistarihi) AS date) AS bit,
       CAST(m.new_mecratipi4 AS int) AS mecra4, CAST(m.new_mecratipi5 AS int) AS mecra5
FROM {p}new_pazarlamamoduluBase AS m
WHERE m.statecode = 0 AND m.new_baslangictarihi < '{_utc(to + timedelta(days=1))}' AND m.new_bitistarihi >= '{_utc(frm)}'""".strip()


#: CRM seçenek adları (`new_pazarlamatipi`, `new_mecratipi4/5`; CRM'in Türkçe etiket tablosundan, 2026-09-09).
CRM_TYPES = {1: "Basın", 2: "Medya", 3: "Promosyon", 4: "Sosyal Medya", 5: "Dijital Pazarlama", 6: "Satış Kampanyası"}
CRM_MEDIA4 = {1: "Facebook", 2: "Instagram", 3: "Twitter", 4: "Pinterest", 5: "Linkedin", 6: "Influencer", 7: "Video"}
CRM_MEDIA5 = {1: "Adwords", 2: "Seo", 3: "TSınav", 4: "Timaş Okul", 5: "Sizbiz TV", 6: "TKitab", 7: "TLand", 8: "Website"}
AD_PLAN_STATUS = {100000004: "Taslak", 100000000: "Planlandı", 100000002: "Tamamlandı", 100000003: "İptal Edildi", 1: "Etkin", 2: "Etkin değil"}
#: Reklam sayılan CRM pazarlama tipleri (Sosyal Medya, Dijital Pazarlama).
AD_TYPES = (4, 5)


def _s(v: Any) -> Optional[str]:
    return bsrc._clean(v)


def _d(v: Any) -> Optional[str]:
    d = bsrc._day(v)
    if d is None or d.year < 1950:
        return None
    return d.isoformat()


class Crm:
    """CRM okumaları. Kitap listesi 30 dk bellekte (eşleştirme her kampanyada aynı listeyi kullanır)."""

    def __init__(self, schema: Callable[[], str], runner: Callable[[], Runner]):
        self.schema = schema
        self.runner = runner
        self._cache: dict[Any, tuple[float, Any]] = {}
        self._lock = threading.Lock()

    def _run(self, sql: str) -> list[dict[str, Any]]:
        try:
            return self.runner()(sql)
        except bsrc.SourceError as e:
            raise SourceError(f"CRM okunamadı: {e}") from None

    def _cached(self, key: Any, fresh: bool, fn: Callable[[], Any], ttl: int = CRM_TTL) -> Any:
        with self._lock:
            hit = self._cache.get(key)
        if hit and not fresh and time.monotonic() - hit[0] < ttl:
            return hit[1]
        val = fn()
        with self._lock:
            self._cache[key] = (time.monotonic(), val)
        return val

    def books(self, off_sale: list[str], fresh: bool = False) -> list[dict[str, Any]]:
        def load() -> list[dict[str, Any]]:
            seen: dict[str, dict[str, Any]] = {}
            for r in self._run(books_sql(self.schema())):
                code = _s(r.get("stok_kodu"))
                if not code or code in seen:
                    continue
                durum = _s(r.get("durum"))
                seen[code] = {"kitapId": (_s(r.get("kitap_id")) or "").lower() or None, "stokKodu": code, "ad": _s(r.get("ad")),
                              "ean": _s(r.get("ean")), "yazar": _s(r.get("yazar")), "durum": durum}
            return list(seen.values())
        rows = self._cached("books", fresh, load)
        codes = {c.upper() for c in off_sale}
        return [{**b, "satisDisi": bool(b["durum"]) and (b["durum"] or "").split(" ", 1)[0].upper() in codes} for b in rows]

    def ad_plans(self, fresh: bool = False) -> list[dict[str, Any]]:
        def load() -> list[dict[str, Any]]:
            books: dict[str, list[dict[str, Any]]] = {}
            for r in self._run(ad_plan_books_sql(self.schema())):
                pid = (_s(r.get("plan_id")) or "").lower()
                if pid:
                    books.setdefault(pid, []).append({"stokKodu": _s(r.get("stok_kodu")), "ad": _s(r.get("ad"))})
            out = []
            for r in self._run(ad_plans_sql(self.schema())):
                pid = (_s(r.get("id")) or "").lower()
                code = r.get("durum_kodu")
                out.append({"id": pid, "ad": _s(r.get("ad")), "tutar": bsrc._num(r.get("tutar")), "bas": _d(r.get("bas")),
                            "bit": _d(r.get("bit")), "durum": AD_PLAN_STATUS.get(int(code), str(code)) if code is not None else None,
                            "mecra": _s(r.get("mecra")), "tip": _s(r.get("tip")), "iptal": bool(r.get("iptal")),
                            "faturaGirildi": bool(r.get("fatura")), "faturaNo": _s(r.get("fatura_no")),
                            "teslim": bool(r.get("teslim")), "onay": _d(r.get("onay")), "kitaplar": books.get(pid, [])})
            return out
        return self._cached("plans", fresh, load, 300)

    def budget_records(self, frm: date, to: date, fresh: bool = False) -> list[dict[str, Any]]:
        def load() -> list[dict[str, Any]]:
            out = []
            for r in self._run(budget_records_sql(self.schema(), frm, to)):
                tip = r.get("tip")
                m4, m5 = r.get("mecra4"), r.get("mecra5")
                out.append({"id": (_s(r.get("id")) or "").lower(), "ad": _s(r.get("ad")), "tip": tip,
                            "tipAdi": CRM_TYPES.get(int(tip)) if tip is not None else None, "tutar": bsrc._num(r.get("tutar")) or 0.0,
                            "tutarKayit": bsrc._num(r.get("tutar_kayit")), "baslangic": _d(r.get("bas")), "bitis": _d(r.get("bit")),
                            "mecra": ", ".join(x for x in (CRM_MEDIA4.get(int(m4)) if m4 is not None else None,
                                                           CRM_MEDIA5.get(int(m5)) if m5 is not None else None) if x) or None,
                            "reklam": tip is not None and int(tip) in AD_TYPES})
            return out
        return self._cached(("budget", frm, to), fresh, load, 300)


def crm_runner(path: Callable[[], str]) -> Callable[[], Runner]:
    def make() -> Runner:
        try:
            return bsrc.runner(path())
        except bsrc.SourceError as e:
            raise SourceError(f"CRM okunamıyor: {e}") from None
    return make
