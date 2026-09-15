"""Panodan Excel: her kart bir sayfa, başta bir özet sayfası.

Veri kartın son ekrana gelen sonucu değildir; SQL indirme anında tam koşulur, satır tavanı
yoktur. Her sayfada biçimli bir Excel tablosu (filtre, bantlı satır, toplam), sayı biçimi ve
kartın grafiği Excel'in kendi grafiği olarak durur; kişi Excel'de değiştirebilir.
"""

from __future__ import annotations

import io
import json
import re
from datetime import datetime
from typing import Any, Callable, Iterable

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

#: (kolonlar, satırlar) döndürür; tam veri.
Fetch = Callable[[str], tuple[list[dict[str, Any]], list[dict[str, Any]]]]

INK = "252423"
MUTED = "605E5C"
ACCENT = "118DFF"
PALETTE = ["118DFF", "12239E", "E66C37", "6B007B", "E044A7", "744EC2", "D9B300", "D64550"]
#: Bundan çok kategoride grafik okunmaz; sayfaya bunun yerine açık bir not yazılır.
CHART_MAX_ROWS = 60

_HEAD_ROW = 5  # başlık, soru, zaman, boşluk sonrası tablo başlığı
_INVALID_SHEET = re.compile(r"[\[\]:*?/\\]")


def humanize(name: str) -> str:
    s = re.sub(r"^d(?=\d{4})", "", str(name)).replace("_", " ").strip()
    return s[:1].upper() + s[1:] if s else s


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and v == v


def _sheet_name(title: str, used: set[str]) -> str:
    base = _INVALID_SHEET.sub(" ", title).strip()[:28] or "Kart"
    name, i = base, 2
    while name.lower() in used:
        name = f"{base[:25]} {i}"
        i += 1
    used.add(name.lower())
    return name


def _cell_value(v: Any) -> Any:
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, str):
        m = re.fullmatch(r"(\d{4}-\d{2}-\d{2})(?:[T ](\d{2}:\d{2}(?::\d{2})?)(?:\.\d+)?Z?)?", v)
        if m:
            try:
                return datetime.fromisoformat(v.replace("Z", "").split(".")[0])
            except ValueError:
                return v
    return v


def _chart(kind: str, ws, first_row: int, last_row: int, label_idx: int, num_idx: list[int], title: str):
    if kind in ("kpi", "table", "treemap") or not num_idx or last_row < first_row:
        return None
    cats = Reference(ws, min_col=label_idx, min_row=first_row, max_row=last_row)
    if kind in ("pie", "donut"):
        ch = DoughnutChart(holeSize=55) if kind == "donut" else PieChart()
        ch.add_data(Reference(ws, min_col=num_idx[0], min_row=first_row - 1, max_row=last_row), titles_from_data=True)
        ch.set_categories(cats)
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
        ch.dataLabels.showVal = False
        ch.dataLabels.showCatName = False
        ch.legend.position = "r"
    elif kind == "scatter" and len(num_idx) >= 2:
        ch = ScatterChart()
        ch.style = 13
        xs = Reference(ws, min_col=num_idx[0], min_row=first_row, max_row=last_row)
        ys = Reference(ws, min_col=num_idx[1], min_row=first_row - 1, max_row=last_row)
        s = Series(ys, xs, title_from_data=True)
        s.marker.symbol = "circle"
        s.graphicalProperties.line.noFill = True
        ch.series.append(s)
        ch.legend = None
    else:
        if kind == "line":
            ch = LineChart()
        elif kind == "area":
            ch = AreaChart()
        else:
            ch = BarChart()
            ch.type = "bar" if kind == "bar" else "col"
            ch.gapWidth = 60
            ch.overlap = 0 if len(num_idx) > 1 else None
        for i in num_idx:
            ch.add_data(Reference(ws, min_col=i, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
        ch.set_categories(cats)
        ch.y_axis.numFmt = "#,##0"
        ch.y_axis.majorGridlines.spPr = None
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        if kind == "bar":
            ch.x_axis.scaling.orientation = "maxMin"  # ilk satır üstte, ekrandaki sırayla
        if len(num_idx) == 1:
            ch.legend = None
            if last_row - first_row < 24:
                ch.dataLabels = DataLabelList()
                ch.dataLabels.showVal = True
                ch.dataLabels.numFmt = '#,##0,,"Mn";-#,##0,,"Mn";0'
        else:
            ch.legend.position = "t"
    for i, s in enumerate(ch.series):
        color = PALETTE[i % len(PALETTE)]
        if kind in ("pie", "donut"):
            from openpyxl.chart.series import DataPoint

            for p in range(last_row - first_row + 1):
                pt = DataPoint(idx=p)
                pt.graphicalProperties.solidFill = PALETTE[p % len(PALETTE)]
                pt.graphicalProperties.line.solidFill = "FFFFFF"
                s.dPt.append(pt)
        elif kind == "line":
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 28000
            s.smooth = False
        elif kind != "scatter":
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.solidFill = color
    ch.title = title[:80]
    ch.height = 9.5
    ch.width = 18
    return ch


def _card_sheet(ws, card: dict[str, Any], cols: list[dict[str, Any]], rows: list[dict[str, Any]], stamp: str) -> int:
    names = [str(c.get("name")) for c in cols] or (list(rows[0].keys()) if rows else [])
    nums = [n for n in names if any(_is_num(r.get(n)) for r in rows) and all(r.get(n) is None or _is_num(r.get(n)) for r in rows)]
    label = next((n for n in names if n not in nums), names[0] if names else "")

    ws.sheet_view.showGridLines = False
    ws["A1"] = card.get("title") or "Kart"
    ws["A1"].font = Font(size=16, bold=True, color=INK)
    question = card.get("question") or ""
    if question and question != card.get("title"):
        ws["A2"] = f"Soru: {question}"
    if card.get("note"):
        ws["A2"] = f"{ws['A2'].value + '  ·  ' if ws['A2'].value else ''}{card['note']}"
    ws["A2"].font = Font(size=10, italic=True, color=MUTED)
    ws["A3"] = f"{len(rows):,} satır · sorgu {stamp}".replace(",", ".")
    ws["A3"].font = Font(size=10, color=MUTED)

    if not names:
        ws["A5"] = "Sonuç boş"
        return 0

    thin = Side(style="thin", color="E1DFDD")
    for j, n in enumerate(names, start=1):
        c = ws.cell(row=_HEAD_ROW, column=j, value=humanize(n))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
        c.alignment = Alignment(horizontal="right" if n in nums else "left", vertical="center")
    ws.row_dimensions[_HEAD_ROW].height = 22

    decimals = {n: any(_is_num(r.get(n)) and float(r.get(n)) != int(r.get(n)) for r in rows) for n in nums}
    widths = [max(10, len(humanize(n)) + 4) for n in names]
    for i, r in enumerate(rows, start=_HEAD_ROW + 1):
        for j, n in enumerate(names, start=1):
            v = r.get(n)
            if n == label and (v is None or v == ""):
                v = "(Boş)"
            v = _cell_value(v)
            c = ws.cell(row=i, column=j, value=v)
            c.border = Border(bottom=thin)
            if n in nums:
                c.number_format = "#,##0.00" if decimals[n] else "#,##0"
                c.alignment = Alignment(horizontal="right")
            elif isinstance(v, datetime):
                c.number_format = "dd.mm.yyyy" if (v.hour, v.minute, v.second) == (0, 0, 0) else "dd.mm.yyyy hh:mm"
            text = f"{v:,.2f}" if _is_num(v) else str(v if v is not None else "")
            widths[j - 1] = min(60, max(widths[j - 1], len(text) + 3))
    last = _HEAD_ROW + len(rows)

    # Excel tablosu: filtre ve bantlı satır Excel'in kendi özelliği olarak gelir.
    if rows:
        ref = f"A{_HEAD_ROW}:{get_column_letter(len(names))}{last}"
        safe = re.sub(r"[^A-Za-z0-9_]", "_", ws.title)
        t = Table(displayName=f"T_{safe}_{ws.parent.index(ws)}"[:250], ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
        if nums and len(rows) > 1:
            total = last + 1
            ws.cell(row=total, column=1, value="Toplam").font = Font(bold=True, color=INK)
            for j, n in enumerate(names, start=1):
                if n in nums:
                    col = get_column_letter(j)
                    c = ws.cell(row=total, column=j, value=f"=SUBTOTAL(109,{col}{_HEAD_ROW + 1}:{col}{last})")
                    c.font = Font(bold=True, color=INK)
                    c.number_format = "#,##0.00" if decimals[n] else "#,##0"
                    c.alignment = Alignment(horizontal="right")
                    c.border = Border(top=Side(style="medium", color=INK))
            ws.cell(row=total, column=1).border = Border(top=Side(style="medium", color=INK))

    for j, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    ws.freeze_panes = ws.cell(row=_HEAD_ROW + 1, column=1)

    kind = card.get("chart") or "table"
    anchor = f"{get_column_letter(len(names) + 2)}{_HEAD_ROW}"
    if kind == "kpi" and rows and nums:
        box = ws[anchor]
        box.value = rows[0].get(nums[0])
        box.number_format = "#,##0"
        box.font = Font(size=28, bold=True, color=INK)
        ws.cell(row=box.row + 1, column=box.column, value=humanize(nums[0])).font = Font(size=10, color=MUTED)
        ws.column_dimensions[box.column_letter].width = 26
    elif len(rows) > CHART_MAX_ROWS and kind not in ("kpi", "table"):
        ws[anchor] = f"Grafik eklenmedi: {len(rows)} satır tek grafikte okunmuyor (en çok {CHART_MAX_ROWS})."
        ws[anchor].font = Font(size=10, italic=True, color=MUTED)
    else:
        idx = {n: j for j, n in enumerate(names, start=1)}
        ch = _chart(kind, ws, _HEAD_ROW + 1, last, idx[label], [idx[n] for n in nums if n != label], card.get("title") or "")
        if ch is not None:
            ws.add_chart(ch, anchor)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return len(rows)


def build(cards: Iterable[dict[str, Any]], fetch: Fetch, *, user: str, now: datetime | None = None) -> bytes:
    """Kartlardan çalışma kitabı. Bir kartın SQL'i koşmazsa o sayfa hatayı yazar, diğerleri etkilenmez."""
    now = now or datetime.now()
    stamp = now.strftime("%d.%m.%Y %H:%M")
    wb = Workbook()
    summary = wb.active
    summary.title = "Özet"
    summary.sheet_view.showGridLines = False
    summary["A1"] = "Panom"
    summary["A1"].font = Font(size=18, bold=True, color=INK)
    summary["A2"] = f"{user} · {stamp}"
    summary["A2"].font = Font(size=10, color=MUTED)
    for j, h in enumerate(["Kart", "Grafik", "Satır", "Durum"], start=1):
        c = summary.cell(row=4, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
    used = {"özet"}
    row = 5
    for card in cards:
        ws = wb.create_sheet(_sheet_name(card.get("title") or "Kart", used))
        try:
            cols, rows = fetch(card["sql"])
            n = _card_sheet(ws, card, cols, rows, stamp)
            status = "Tamam"
        except Exception as e:  # noqa: BLE001 — bir kartın hatası kitabı düşürmez
            ws["A1"] = card.get("title") or "Kart"
            ws["A1"].font = Font(size=16, bold=True, color=INK)
            ws["A3"] = f"Veri alınamadı: {e}"
            ws["A3"].font = Font(color="D64550")
            n, status = 0, "Veri alınamadı"
        link = summary.cell(row=row, column=1, value=card.get("title") or "Kart")
        link.hyperlink = f"#'{ws.title}'!A1"
        link.font = Font(color=ACCENT, underline="single")
        summary.cell(row=row, column=2, value=card.get("chart") or "table")
        summary.cell(row=row, column=3, value=n).number_format = "#,##0"
        summary.cell(row=row, column=4, value=status)
        row += 1
    summary.column_dimensions["A"].width = 60
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["C"].width = 12
    summary.column_dimensions["D"].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
