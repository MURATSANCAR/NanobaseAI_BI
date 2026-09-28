"""M29 İlk dağılım yönetimi: depoya giren kitabın bölge × kanal × müşteri dağılım önerisi, düzeltme, iki göz onayı,
sevk listesi (Excel), ilk 8 hafta takibi ve uyarılar.

**Akış (K2 — ZEKİ AI önerir, satış müdürü düzeltir, lojistik onaylar):**
1. *Dağılım bekleyen kitaplar* (K1): Logo'da gerçek üretimden giriş (`distribution_sources`) ve M12'nin üretim kartı
   (CRM depo girişi — Logo kopyası donmuşken son girişler buradan). Günde iki kez `run-due` tazeler.
2. *Öneri*: benzer kitaplar seçilir — yeni kitapta M10'un emsal puanı (CRM emsali, yazar, dizi, kitaplık, yayınevi, fiyat,
   sayfa; `management/ilk_baski_model.similarity`), ZEKİ AI aday listesini «benzer / az / değil» diye ayıklar (rakam
   üretmez); baskı tekrarında kitabın kendi son 8 haftası. Her benzer kitabın ilk faturalı satış gününden 8 hafta
   (`DIST_WINDOW_DAYS`) içindeki cari bazında net adedi (satış − iade) Logo'dan okunur; cari payı = benzer kitapların
   kendi toplamlarına göre paylarının puan ağırlıklı ortalaması. Toplam adet = M46 onaylı hedefinin depo girişinden
   itibaren ilk iki ayı; hedef yoksa benzer kitapların 8 haftalık net adedinin ortancası. Üst sınır = stok bakiyesi −
   rezerv (`DIST_RESERVE_SHARE`, varsayılan %20). Adetler en büyük kalan yöntemiyle yuvarlanır (toplam tam tutar).
   CRM'de «Dağılım Durumu Göster» işaretli cariler payı sıfır da olsa listede kalır (liste kesilmez).
3. *Düzeltme*: müşteri satırı ya da bölge × kanal hücresi (hücre içindeki carilere oranla); elle düzeltilen satır işaretli,
   gerekçe yazılır. Rezerv plan başlığında.
4. *Onay (iki göz)*: taslak → onayda → onaylı; gönderen onaylayamaz, geri gönderme gerekçeli taslağa döner. Değişmez:
   Σ adet + rezerv ≤ stok bakiyesi (stok bilinmiyorsa baskı adedi); aşan plan onaya gidemez (409). Onaylanan plan kitabın
   yürürlükteki planıdır, önceki onaylı plan arşive geçer; «Revize et» gerekçeli yeni sürüm açar.
5. *Sevk listesi*: yalnız onaylı plan, Excel; CRM «Dağılım» siparişine elle yazmaya uygun kolonlar. CRM'e ve Logo'ya yazılmaz.
6. *Takip* (K1): onay gününden itibaren 8 hafta, cari × hafta sevk / faturalanan / iade (Logo). Logo verisinin bittiği gün
   her ekranda yazılır; gecikme ve «hiç satmadı» uyarısı bugüne değil verinin bittiği güne göre değerlendirilir (donmuş
   kopyada sahte uyarı açılmasın).

Uyarılar: `plan_yok` (depoya girdi, plan yok), `sevk_gecikti` (onaylı satır `DIST_SHIP_DAYS` iş gününde hiç sevk
edilmedi), `hic_satmadi` (onaydan `DIST_NO_SALE_DAYS` gün sonra bölgede sevk var, faturalanan net ≤ 0), `tukendi` (bilgi:
carinin faturalanan adedi sevkinin %90'ına ulaştı — yeniden sipariş önerisinin girdisi, sonraki sürüm).

Her yazma `semantic_audit`'e (`dist_plan`, `dist_line`, `dist_export`). Ayrıntı günlükte (2026-09-28, M29).
"""
from __future__ import annotations

import io
import json
import logging
import math
import re
import statistics
import threading
import time
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, Iterable, Optional
from zoneinfo import ZoneInfo

import sqlalchemy as sa

from semantic_bridge import distribution_sources as src

log = logging.getLogger("semantic.distribution")
TZ = ZoneInfo("Europe/Istanbul")
_md = sa.MetaData()

BOOKS = sa.Table(
    "semantic_dist_books", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("stok_kodu", sa.String(60), primary_key=True),
    sa.Column("ad", sa.String(400)),
    sa.Column("yayinevi", sa.String(200)),
    sa.Column("depo_giris_tarihi", sa.String(10)),
    sa.Column("baski_adedi", sa.Float),
    sa.Column("baski_no", sa.Integer),
    sa.Column("ilk_baski", sa.Boolean),
    sa.Column("kaynak", sa.String(12), nullable=False),        # logo | m12 | logo+m12
    sa.Column("stok_bakiye", sa.Float),
    sa.Column("asof", sa.DateTime(timezone=True), nullable=False),
)
PLANS = sa.Table(
    "semantic_dist_plans", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("stok_kodu", sa.String(60), nullable=False),
    sa.Column("ad", sa.String(400)),
    sa.Column("surum", sa.Integer, nullable=False),
    sa.Column("durum", sa.String(12), nullable=False),         # taslak | onayda | onayli | arsiv
    sa.Column("depo_giris_tarihi", sa.String(10)),
    sa.Column("baski_adedi", sa.Float),
    sa.Column("stok_bakiye", sa.Float),
    sa.Column("toplam_adet", sa.Float, nullable=False, default=0),
    sa.Column("onerilen_toplam", sa.Float, nullable=False, default=0),
    sa.Column("rezerv_adet", sa.Float, nullable=False, default=0),
    sa.Column("hedef_plan_id", sa.String(32)),                 # M46 yürürlükteki planı
    sa.Column("hedef_json", sa.Text),
    sa.Column("basis_json", sa.Text, nullable=False),
    sa.Column("gerekce", sa.Text),                             # öneri gerekçesi (ZEKİ AI ya da kural metni)
    sa.Column("gerekce_kaynak", sa.String(8)),                 # model | kural
    sa.Column("note", sa.Text),                                # hazırlayanın notu / düzeltme gerekçesi
    sa.Column("revision_of", sa.String(32)),
    sa.Column("revision_reason", sa.Text),
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("updated_by", sa.String(120)),
    sa.Column("updated_at", sa.DateTime(timezone=True)),
    sa.Column("submitted_by", sa.String(120)),
    sa.Column("submitted_at", sa.DateTime(timezone=True)),
    sa.Column("decided_by", sa.String(120)),
    sa.Column("decided_at", sa.DateTime(timezone=True)),
    sa.Column("decision_note", sa.Text),
    sa.Index("ix_semantic_dist_plans_book", "tenant_id", "stok_kodu"),
)
LINES = sa.Table(
    "semantic_dist_plan_lines", _md,
    sa.Column("plan_id", sa.String(32), primary_key=True),
    sa.Column("line_no", sa.Integer, primary_key=True),
    sa.Column("logo_cari_kodu", sa.String(60)),
    sa.Column("crm_account_id", sa.String(40)),
    sa.Column("cari_unvan", sa.String(300)),
    sa.Column("il", sa.String(80)),
    sa.Column("bolge", sa.String(40), nullable=False),
    sa.Column("kanal", sa.String(60), nullable=False),
    sa.Column("kanal_kodu", sa.String(60)),
    sa.Column("bmt_ad", sa.String(200)),
    sa.Column("bmt_hesap", sa.String(120), index=True),
    sa.Column("dagilim_carisi", sa.Boolean, nullable=False, default=False),
    sa.Column("pay", sa.Float, nullable=False, default=0),
    sa.Column("gecmis_net", sa.Float),                         # benzer kitaplarda 8 haftalık ağırlıklı net adet
    sa.Column("gecmis_iade_orani", sa.Float),
    sa.Column("onerilen_adet", sa.Float, nullable=False, default=0),
    sa.Column("adet", sa.Float, nullable=False, default=0),
    sa.Column("elle_duzeltildi", sa.Boolean, nullable=False, default=False),
    sa.Column("gerekce", sa.String(500)),
    sa.Column("edited_by", sa.String(120)),
    sa.Column("edited_at", sa.DateTime(timezone=True)),
)
COMPS = sa.Table(
    "semantic_dist_comps", _md,
    sa.Column("plan_id", sa.String(32), primary_key=True),
    sa.Column("comp_stok_kodu", sa.String(60), primary_key=True),
    sa.Column("ad", sa.String(400)),
    sa.Column("yazar", sa.String(300)),
    sa.Column("yayinevi", sa.String(200)),
    sa.Column("pencere_bas", sa.String(10)),
    sa.Column("pencere_bit", sa.String(10)),
    sa.Column("puan", sa.Float),
    sa.Column("benzerlik_olasiligi", sa.Float),                # ZEKİ AI P(benzer)+P(az benzer); yoksa puanın 0–1 ölçeği
    sa.Column("model_karar", sa.String(12)),                   # benzer | az | degil | None (model yok)
    sa.Column("gerekce", sa.String(500)),
    sa.Column("secildi", sa.Boolean, nullable=False, default=True),
    sa.Column("agirlik", sa.Float),
    sa.Column("satis_adet", sa.Float),
    sa.Column("iade_adet", sa.Float),
    sa.Column("net_ciro", sa.Float),
    sa.Column("musteri", sa.Integer),
    sa.Column("kanallar_json", sa.Text),
    sa.Column("crm_dagilim_adet", sa.Float),
    sa.Column("crm_dagilim_siparis", sa.Integer),
)
TRACKING = sa.Table(
    "semantic_dist_tracking", _md,
    sa.Column("plan_id", sa.String(32), primary_key=True),
    sa.Column("logo_cari_kodu", sa.String(60), primary_key=True),
    sa.Column("hafta", sa.Integer, primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("sevk_adet", sa.Float, nullable=False, default=0),
    sa.Column("satis_adet", sa.Float, nullable=False, default=0),   # faturalanan
    sa.Column("iade_adet", sa.Float, nullable=False, default=0),
    sa.Column("asof", sa.DateTime(timezone=True), nullable=False),
)
ALERTS = sa.Table(
    "semantic_dist_alerts", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("plan_id", sa.String(32)),
    sa.Column("stok_kodu", sa.String(60)),
    sa.Column("tur", sa.String(16), nullable=False),          # plan_yok | sevk_gecikti | hic_satmadi | tukendi
    sa.Column("anahtar", sa.String(200), nullable=False),
    sa.Column("etiket", sa.String(400)),
    sa.Column("detay_json", sa.Text),
    sa.Column("bmt_hesap", sa.String(120)),
    sa.Column("durum", sa.String(10), nullable=False),         # acik | kapandi | bilgi
    sa.Column("ilk_zaman", sa.DateTime(timezone=True), nullable=False),
    sa.Column("son_zaman", sa.DateTime(timezone=True), nullable=False),
    sa.Column("kapanis", sa.DateTime(timezone=True)),
    sa.Column("bildirim", sa.DateTime(timezone=True)),
    sa.Index("ix_semantic_dist_alerts_open", "tenant_id", "durum"),
)
META = sa.Table(
    "semantic_dist_meta", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("key", sa.String(60), primary_key=True),
    sa.Column("value_json", sa.Text, nullable=False),
    sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
)

STATUSES = {"taslak": "Taslak", "onayda": "Onay bekliyor", "onayli": "Onaylı", "arsiv": "Arşiv"}
#: Kitap listesindeki durum (plan durumu + sevk).
BOOK_STATES = {"yok": "Plan yok", "taslak": "Taslak", "onayda": "Onayda", "onayli": "Onaylı", "sevkte": "Sevkte"}
ALERT_KINDS = {"plan_yok": "Depoya girdi, plan yok", "sevk_gecikti": "Sevk edilmedi", "hic_satmadi": "Hiç satmadı",
               "tukendi": "Tükeniyor"}
PAGE_SIZE = 100

# ------------------------------------------------------------------ ayarlar (Yönetim ekranı > ortam > varsayılan)

DEFAULTS: dict[str, str] = {
    "DIST_WINDOW_DAYS": "56",        # benzer kitabın ilk satışından okunan pencere (8 hafta)
    "DIST_TRACK_WEEKS": "8",         # onaydan sonra izlenen hafta
    "DIST_BOOKS_DAYS": "90",         # «dağılım bekleyen» listesine giren depo girişi penceresi
    "DIST_RESERVE_SHARE": "0.20",    # depoda tutulan pay (iş kuralı sorulacak; ölçüm kabul betiğinde)
    "DIST_COMPS": "8",               # öneriye giren en çok benzer kitap
    "DIST_MIN_QTY": "0",             # bu adedin altındaki satır sıfırlanır, adet diğerlerine dağılır
    "DIST_SHIP_DAYS": "5",           # onaydan sonra sevk için iş günü
    "DIST_NO_SALE_DAYS": "28",       # «hiç satmadı» değerlendirmesi için gün
    "DIST_SOLD_OUT_SHARE": "0.90",   # faturalanan / sevk ≥ bu pay → tükeniyor
    "DIST_MODEL": "1",               # ZEKİ AI ayıklaması ve gerekçe metni
}


def setting(key: str) -> str:
    try:
        from semantic_bridge.admin import conf

        v = conf(key, DEFAULTS.get(key, ""))
    except Exception:  # noqa: BLE001 — yönetim modülü yoksa ortam
        import os

        v = os.environ.get(key, DEFAULTS.get(key, ""))
    return str(v if v not in (None, "") else DEFAULTS.get(key, ""))


def num_setting(key: str) -> float:
    try:
        return float(setting(key).replace(",", "."))
    except ValueError:
        return float(DEFAULTS[key])


def params() -> dict[str, Any]:
    return {"pencere": int(num_setting("DIST_WINDOW_DAYS")), "takipHafta": int(num_setting("DIST_TRACK_WEEKS")),
            "listeGun": int(num_setting("DIST_BOOKS_DAYS")), "rezervPay": num_setting("DIST_RESERVE_SHARE"),
            "benzer": int(num_setting("DIST_COMPS")), "enAz": num_setting("DIST_MIN_QTY"),
            "sevkGun": int(num_setting("DIST_SHIP_DAYS")), "satisGun": int(num_setting("DIST_NO_SALE_DAYS")),
            "tukenmePay": num_setting("DIST_SOLD_OUT_SHARE"), "model": setting("DIST_MODEL").strip() not in ("0", "false")}


# ------------------------------------------------------------------ altyapı

_ready: set[int] = set()
_lock = threading.Lock()


class DistError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def ensure(engine: sa.engine.Engine) -> None:
    with _lock:
        if id(engine) in _ready:
            return
        _md.create_all(engine, checkfirst=True)
        _ready.add(id(engine))


def _now() -> datetime:
    return datetime.now(timezone.utc)


def today() -> date:
    return datetime.now(TZ).date()


def _iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    return (v if v.tzinfo else v.replace(tzinfo=timezone.utc)).isoformat()


def _j(v: Optional[str], default: Any) -> Any:
    try:
        return json.loads(v) if v else default
    except ValueError:
        return default


def _dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _text(v: Any, limit: int) -> Optional[str]:
    s = " ".join(str(v or "").split())
    return s[:limit] or None


def _qty(v: Any, label: str = "Adet") -> float:
    if v is None or v == "":
        raise DistError(f"{label} boş olamaz.")
    try:
        if isinstance(v, str):
            t = v.strip().replace(" ", "")
            if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", t):   # Türkçe yazım: 1.250 ya da 1.250,0
                t = t.replace(".", "")
            n = float(t.replace(",", "."))
        else:
            n = float(v)
    except (TypeError, ValueError):
        raise DistError(f"{label} sayı olmalı.") from None
    if math.isnan(n) or math.isinf(n):
        raise DistError(f"{label} sayı olmalı.")
    if n < 0:
        raise DistError(f"{label} eksi olamaz.")
    if n != int(n):
        raise DistError(f"{label} tam sayı olmalı.")
    return float(int(n))


def _day(v: Any) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10]) if v else None
    except ValueError:
        return None


def meta_get(engine: sa.engine.Engine, tenant: str, key: str) -> dict[str, Any]:
    with engine.connect() as c:
        row = c.execute(sa.select(META).where(META.c.tenant_id == tenant, META.c.key == key)).first()
    return {**_j(row.value_json, {}), "_at": _iso(row.updated_at)} if row else {}


def meta_set(engine: sa.engine.Engine, tenant: str, key: str, value: dict[str, Any]) -> None:
    now = _now()
    with engine.begin() as c:
        cond = (META.c.tenant_id == tenant, META.c.key == key)
        if c.execute(sa.select(META.c.key).where(*cond)).first():
            c.execute(META.update().where(*cond).values(value_json=_dump(value), updated_at=now))
        else:
            c.execute(META.insert().values(tenant_id=tenant, key=key, value_json=_dump(value), updated_at=now))


def data_end(engine: sa.engine.Engine, tenant: str) -> Optional[date]:
    return _day(meta_get(engine, tenant, "logo").get("veriSonu"))


# ------------------------------------------------------------------ bölge ve kanal

_TR = str.maketrans("İIıŞşĞğÜüÖöÇçÂâÎîÛû", "iiissgguuooccaaiiuu")


def fold(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s or "").translate(_TR).lower()).strip()


#: Türkiye'nin 7 coğrafi bölgesi, İstanbul ayrı (iş gününde «İstanbul %38, Ege %14» diye konuşuluyor). 81 il.
REGIONS: dict[str, list[str]] = {
    "İstanbul": ["istanbul"],
    "Marmara": ["edirne", "kirklareli", "tekirdag", "kocaeli", "sakarya", "yalova", "bursa", "bilecik", "balikesir",
                "canakkale"],
    "Ege": ["izmir", "manisa", "aydin", "denizli", "mugla", "usak", "kutahya", "afyonkarahisar"],
    "Akdeniz": ["antalya", "isparta", "burdur", "mersin", "adana", "hatay", "osmaniye", "kahramanmaras"],
    "İç Anadolu": ["ankara", "konya", "karaman", "aksaray", "nigde", "nevsehir", "kirsehir", "kirikkale", "kayseri", "sivas",
                   "yozgat", "cankiri", "eskisehir"],
    "Karadeniz": ["zonguldak", "karabuk", "bartin", "kastamonu", "sinop", "samsun", "amasya", "tokat", "corum", "ordu",
                  "giresun", "trabzon", "rize", "artvin", "gumushane", "bayburt", "bolu", "duzce"],
    "Doğu Anadolu": ["erzurum", "erzincan", "kars", "ardahan", "igdir", "agri", "van", "mus", "bitlis", "bingol", "tunceli",
                     "elazig", "malatya", "hakkari"],
    "Güneydoğu Anadolu": ["gaziantep", "kilis", "sanliurfa", "diyarbakir", "mardin", "batman", "siirt", "sirnak", "adiyaman"],
}
REGION_ORDER = list(REGIONS) + ["Yurt dışı", "İli belirsiz"]
_CITY: dict[str, str] = {c: r for r, cs in REGIONS.items() for c in cs}
_CITY_ALIAS = {"icel": "mersin", "afyon": "afyonkarahisar", "maras": "kahramanmaras", "urfa": "sanliurfa", "antep": "gaziantep",
               "izmit": "kocaeli", "adapazari": "sakarya", "dersim": "tunceli"}
FOREIGN_CHANNELS = {"yurtdisi", "yurt disi", "ihracat"}


def region_of(city: Any, channel_code: Any = None) -> str:
    if fold(channel_code) in FOREIGN_CHANNELS:
        return "Yurt dışı"
    s = fold(city)
    if not s:
        return "İli belirsiz"
    head = re.split(r"[/,\-()]", s)[0].strip()
    for token in (head, head.split(" ")[0]):
        token = _CITY_ALIAS.get(token, token)
        if token in _CITY:
            return _CITY[token]
    for name, reg in _CITY.items():
        if head.startswith(name):
            return reg
    return "İli belirsiz"


def channel_label(code: Any) -> str:
    from semantic_bridge.management.ilk_baski_model import channel_label as cl

    return cl(code) if str(code or "").strip() else "Kanal kodu boş"


# ------------------------------------------------------------------ saf hesaplar


def allocate(total: float, shares: dict[str, float], min_qty: float = 0) -> dict[str, int]:
    """Toplamı paylara en büyük kalan yöntemiyle tam sayı olarak dağıtır; Σ sonuç = toplam (pay varsa). `min_qty`
    altına düşen anahtar sıfırlanır, adedi kalanlara yeniden dağıtılır."""
    total = int(max(0, round(total)))
    live = {k: v for k, v in shares.items() if v and v > 0}
    out = {k: 0 for k in shares}
    while live and total > 0:
        s = sum(live.values())
        raw = {k: total * v / s for k, v in live.items()}
        base = {k: int(math.floor(x)) for k, x in raw.items()}
        left = total - sum(base.values())
        for k, _ in sorted(raw.items(), key=lambda kv: (-(kv[1] - base[kv[0]]), -live[kv[0]], kv[0]))[:left]:
            base[k] += 1
        small = [k for k, v in base.items() if min_qty and 0 < v < min_qty]
        if small and len(small) < len(live):
            smallest = min(small, key=lambda k: (live[k], k))
            live.pop(smallest)
            continue
        out.update(base)
        break
    return out


def comp_shares(rows: list[dict[str, Any]], weights: dict[str, float]) -> tuple[dict[str, float], dict[str, dict[str, float]], dict[str, float]]:
    """Benzer kitap × cari satırlarından cari payı. Her kitabın dağılımı kendi net toplamına göre paya çevrilir, paylar
    kitap ağırlığıyla ortalanır (çok satan tek kitap öneriyi ele geçirmesin). Döner: (pay, cari → {net, satis, iade},
    kitap → net toplam)."""
    by: dict[str, dict[str, dict[str, float]]] = {}
    for r in rows:
        k, c = r["comp"], r["cari_kodu"]
        if not c:
            continue
        cur = by.setdefault(k, {}).setdefault(c, {"satis": 0.0, "iade": 0.0})
        cur["satis"] += r["satis"]
        cur["iade"] += r["iade"]
    totals: dict[str, float] = {}
    for k, cs in by.items():
        totals[k] = sum(max(0.0, v["satis"] - v["iade"]) for v in cs.values())
    share: dict[str, float] = {}
    stats: dict[str, dict[str, float]] = {}
    wsum = sum(weights.get(k, 0) for k in by if totals.get(k, 0) > 0)
    for k, cs in by.items():
        n, w = totals.get(k, 0), weights.get(k, 0)
        for c, v in cs.items():
            st = stats.setdefault(c, {"net": 0.0, "satis": 0.0, "iade": 0.0})
            st["satis"] += v["satis"]
            st["iade"] += v["iade"]
            if n > 0 and wsum > 0:
                st["net"] += w * max(0.0, v["satis"] - v["iade"]) / wsum
                share[c] = share.get(c, 0.0) + (w / wsum) * max(0.0, v["satis"] - v["iade"]) / n
    return share, stats, totals


def weighted_median(values: list[float]) -> Optional[float]:
    vals = [v for v in values if v and v > 0]
    return float(statistics.median(vals)) if vals else None


def matrix(lines: Iterable[Any]) -> dict[str, Any]:
    """Bölge × kanal: adet, önerilen, müşteri (adedi olan) sayısı."""
    cells: dict[tuple[str, str], dict[str, float]] = {}
    regions: dict[str, float] = {}
    channels: dict[str, float] = {}
    for ln in lines:
        b, k = ln["bolge"], ln["kanal"]
        cell = cells.setdefault((b, k), {"adet": 0.0, "onerilen": 0.0, "musteri": 0, "satir": 0})
        cell["adet"] += ln["adet"]
        cell["onerilen"] += ln["onerilen_adet"]
        cell["satir"] += 1
        cell["musteri"] += 1 if ln["adet"] > 0 else 0
        regions[b] = regions.get(b, 0.0) + ln["adet"]
        channels[k] = channels.get(k, 0.0) + ln["adet"]
    reg_order = sorted(regions, key=lambda r: (REGION_ORDER.index(r) if r in REGION_ORDER else 99, r))
    ch_order = sorted(channels, key=lambda k: (-channels[k], k))
    total = sum(regions.values())
    return {
        "bolgeler": [{"bolge": r, "adet": regions[r], "pay": (regions[r] / total if total else 0)} for r in reg_order],
        "kanallar": [{"kanal": k, "adet": channels[k], "pay": (channels[k] / total if total else 0)} for k in ch_order],
        "hucreler": [{"bolge": b, "kanal": k, **v} for (b, k), v in sorted(cells.items(), key=lambda x: (
            reg_order.index(x[0][0]), ch_order.index(x[0][1])))],
        "toplam": total,
    }


def business_days_between(a: date, b: date) -> int:
    """a'dan (hariç) b'ye (dahil) hafta içi gün sayısı."""
    if b <= a:
        return 0
    n, d = 0, a
    while d < b:
        d += timedelta(days=1)
        if d.weekday() < 5:
            n += 1
    return n


_NUM = re.compile(r"\d[\d.,]*")


def numbers_in(text: str) -> set[str]:
    out = set()
    for m in _NUM.findall(text or ""):
        s = m.rstrip(".,")
        norm = s.replace(".", "").replace(",", ".") if "," in s or re.fullmatch(r"\d{1,3}(\.\d{3})+", s) else s
        try:
            out.add(f"{float(norm):g}")
        except ValueError:
            pass
    return out


def safe_model_text(text: str, facts_text: str) -> bool:
    """Model metnindeki her sayı verilen olgularda geçmeli (rakamı model üretmez). Denetim: `zeki_text`."""
    from semantic_bridge import zeki_text as Z

    return bool((text or "").strip()) and Z.numbers_ok(text, facts_text)


# ------------------------------------------------------------------ kaynaklar


class Sources:
    """Köprünün verdiği bağlantılar ve kardeş modüller. Hepsi tembel; yoksa ilgili adım atlanır, yanıt söyler."""

    def __init__(self, logo_run: Callable[[], src.Runner], crm_run: Callable[[], src.Runner], schema: Callable[[], str],
                 m12: Callable[[], Any] = lambda: None, m10: Callable[[], Any] = lambda: None,
                 llm: Callable[[], Any] = lambda: None):
        self.logo_run, self.crm_run, self.schema = logo_run, crm_run, schema
        self.m12, self.m10, self.llm = m12, m10, llm
        self._logo: Optional[src.Logo] = None
        self._accounts: Optional[tuple[float, list[dict[str, Any]]]] = None

    def logo(self) -> src.Logo:
        if self._logo is None:
            self._logo = src.Logo(self.logo_run())
        return self._logo

    def accounts(self) -> list[dict[str, Any]]:
        if self._accounts and time.monotonic() - self._accounts[0] < 600:
            return self._accounts[1]
        rows = src.read_accounts(self.crm_run(), self.schema())
        self._accounts = (time.monotonic(), rows)
        return rows

    def dist_orders(self, codes: list[str]) -> dict[str, dict[str, Any]]:
        return src.read_dist_orders(self.crm_run(), self.schema(), codes) if codes else {}


# ------------------------------------------------------------------ dağılım bekleyen kitaplar


def refresh_books(engine: sa.engine.Engine, tenant: str, S: Sources, engine_tenant_cards: Optional[Callable[[], list]] = None) -> dict[str, Any]:
    """Logo gerçek girişleri + M12 kartları → `semantic_dist_books` (kiracının anlık görüntüsü)."""
    p = params()
    since = today() - timedelta(days=p["listeGun"])
    logo = S.logo()
    warnings: list[str] = []
    entries = {e["stok_kodu"]: e for e in logo.depot_entries(since)}
    cards: list[dict[str, Any]] = []
    try:
        cards = engine_tenant_cards() if engine_tenant_cards else []
    except Exception as e:  # noqa: BLE001
        warnings.append(f"Üretim kartları okunamadı: {str(e)[:160]}")
    prefixes = src.book_prefixes()
    m12: dict[str, dict[str, Any]] = {}
    for c in cards:
        code = str(c.get("stockCode") or "").strip()
        d = _day((c.get("actual") or {}).get("depo", {}).get("day") if isinstance((c.get("actual") or {}).get("depo"), dict) else None)
        if not code or not d or d < since or not code.startswith(prefixes) or c.get("stage") == "iptal":
            continue
        cur = m12.get(code)
        if cur is None or d > cur["depo"]:
            m12[code] = {"depo": d, "qty": c.get("qty"), "printNo": c.get("printNo"), "firstPrint": c.get("firstPrint"),
                         "title": c.get("bookTitle")}
    codes = sorted(set(entries) | set(m12))
    first = logo.first_entries(codes) if codes else {}
    stock = logo.stock(codes) if codes else {}
    pub: dict[str, str] = {}
    try:
        from semantic_bridge import budget as B

        B.ensure(engine)
        with engine.connect() as c:
            for r in c.execute(sa.select(B.BOOKINFO.c.stok_kodu, B.BOOKINFO.c.yayinevi).where(B.BOOKINFO.c.stok_kodu.in_(codes))):
                if r.yayinevi:
                    pub[r.stok_kodu] = r.yayinevi
    except Exception as e:  # noqa: BLE001
        warnings.append(f"Kitap kartı (yayınevi) okunamadı: {str(e)[:120]}")
    now = _now()
    rows = []
    for code in codes:
        e, m = entries.get(code), m12.get(code)
        giris = e["ilk"] if e and e.get("ilk") else (m["depo"] if m else None)
        if m and e and e.get("ilk"):
            first_print = bool(m.get("firstPrint"))
        elif m:
            first_print = bool(m.get("firstPrint"))
        else:
            f0 = first.get(code)
            first_print = not (f0 and giris and f0 < giris - timedelta(days=7))
        rows.append(dict(
            tenant_id=tenant, stok_kodu=code[:60], ad=((e or {}).get("ad") or (m or {}).get("title") or "")[:400] or None,
            yayinevi=(pub.get(code) or "")[:200] or None, depo_giris_tarihi=giris.isoformat() if giris else None,
            baski_adedi=(e or {}).get("adet") or (m or {}).get("qty"), baski_no=(m or {}).get("printNo"),
            ilk_baski=first_print, kaynak="logo+m12" if (e and m) else ("logo" if e else "m12"),
            stok_bakiye=stock.get(code), asof=now))
    with engine.begin() as c:
        c.execute(BOOKS.delete().where(BOOKS.c.tenant_id == tenant))
        if rows:
            c.execute(BOOKS.insert(), rows)
    end = logo.data_end()
    depot_end = logo.depot_end()
    meta_set(engine, tenant, "logo", {"veriSonu": end.isoformat() if end else None,
                                      "depoSonu": depot_end.isoformat() if depot_end else None,
                                      "kitap": len(rows), "logoGiris": len(entries), "uretimKarti": len(m12),
                                      "uyarilar": warnings, "pencere": since.isoformat()})
    return {"kitap": len(rows), "logo": len(entries), "m12": len(m12), "uyarilar": warnings}


def _latest_plans(c: Any, tenant: str, codes: Optional[list[str]] = None) -> dict[str, Any]:
    """Kitap başına arşivde olmayan en son plan (onaylı varsa o; açık taslak/onayda varsa o öne)."""
    q = sa.select(PLANS).where(PLANS.c.tenant_id == tenant, PLANS.c.durum != "arsiv")
    if codes is not None:
        q = q.where(PLANS.c.stok_kodu.in_(codes))
    out: dict[str, Any] = {}
    rank = {"taslak": 3, "onayda": 2, "onayli": 1}
    for r in c.execute(q):
        cur = out.get(r.stok_kodu)
        if cur is None or (rank[r.durum], r.surum) > (rank[cur.durum], cur.surum):
            out[r.stok_kodu] = r
    return out


def _shipped(c: Any, plan_ids: list[str]) -> dict[str, float]:
    if not plan_ids:
        return {}
    q = (sa.select(TRACKING.c.plan_id, sa.func.sum(TRACKING.c.sevk_adet)).where(TRACKING.c.plan_id.in_(plan_ids))
         .group_by(TRACKING.c.plan_id))
    return {r[0]: float(r[1] or 0) for r in c.execute(q)}


def list_books(engine: sa.engine.Engine, tenant: str, durum: str = "", q: str = "") -> dict[str, Any]:
    """«Dağılım bekleyen kitaplar» (depoya giriş penceresi) ve «İzlenen kitaplar» (onaydan sonraki takip penceresi)."""
    p = params()
    with engine.connect() as c:
        books = c.execute(sa.select(BOOKS).where(BOOKS.c.tenant_id == tenant)).all()
        plans = _latest_plans(c, tenant)
        approved = [r for r in plans.values() if r.durum == "onayli"]
        shipped = _shipped(c, [r.id for r in approved])
    needle = fold(q)
    items = []
    for b in books:
        pl = plans.get(b.stok_kodu)
        state = "yok" if pl is None else ("sevkte" if pl.durum == "onayli" and shipped.get(pl.id, 0) > 0 else pl.durum)
        if durum and durum != state and not (durum == "bekleyen" and state in ("yok", "taslak", "onayda")):
            continue
        if needle and needle not in fold(f"{b.stok_kodu} {b.ad or ''} {b.yayinevi or ''}"):
            continue
        items.append({"stokKodu": b.stok_kodu, "ad": b.ad, "yayinevi": b.yayinevi, "depoGiris": b.depo_giris_tarihi,
                      "baskiAdedi": b.baski_adedi, "baskiNo": b.baski_no, "ilkBaski": b.ilk_baski, "kaynak": b.kaynak,
                      "stok": b.stok_bakiye, "durum": state, "durumEtiket": BOOK_STATES[state],
                      "plan": None if pl is None else {"id": pl.id, "surum": pl.surum, "durum": pl.durum,
                                                       "toplam": pl.toplam_adet, "rezerv": pl.rezerv_adet}})
    items.sort(key=lambda x: (x["depoGiris"] or "", x["stokKodu"]), reverse=True)
    limit_day = today() - timedelta(days=7 * p["takipHafta"] + 7)
    tracked = []
    for r in sorted(approved, key=lambda r: r.decided_at or _now(), reverse=True):
        if r.decided_at and (r.decided_at.date() if r.decided_at.tzinfo is None else r.decided_at.astimezone(TZ).date()) < limit_day:
            continue
        tracked.append(tracking_summary(engine, tenant, r))
    lm = meta_get(engine, tenant, "logo")
    return {"asof": lm.get("_at"), "veriSonu": lm.get("veriSonu"), "depoSonu": lm.get("depoSonu"),
            "uyarilar": lm.get("uyarilar") or [], "pencereGun": p["listeGun"], "items": items, "izlenen": tracked,
            "durumlar": BOOK_STATES}


# ------------------------------------------------------------------ öneri


def _target(engine: sa.engine.Engine, tenant: str, code: str, entry: date) -> dict[str, Any]:
    """M46 yürürlükteki planında kitabın hedefi ve depo girişinden itibaren ilk iki ayın adedi."""
    try:
        from semantic_bridge import budget as B

        B.ensure(engine)
        t = B.approved_targets(engine, tenant, entry.year, codes=[code], with_actuals=False)
    except Exception as e:  # noqa: BLE001
        return {"var": False, "hata": str(e)[:160]}
    if not t.get("plan") or not t.get("items"):
        return {"var": False, "yil": entry.year, "planVar": bool(t.get("plan"))}
    it = t["items"][0]
    months = [m for m in it["aylik"] if entry.month <= m["ay"] <= entry.month + 1]
    return {"var": True, "yil": entry.year, "planId": t["plan"]["id"], "planBaslik": t["plan"]["title"],
            "yillikAdet": it["hedef"]["adet"], "ikiAyAdet": round(sum(m["adet"] for m in months)),
            "aylar": [m["ay"] for m in months]}


def _candidates(engine: sa.engine.Engine, S: Sources, code: str, name: Optional[str], entry: date,
                first_year: int, k: int) -> tuple[list[dict[str, Any]], str]:
    """Yeni kitap için aday benzer kitaplar (puanlı). Önce M10'un emsal veri kümesi, yoksa bütçe kitap kartı."""
    from semantic_bridge.management import ilk_baski_model as M

    store = S.m10()
    eng = None
    if store is not None:
        try:
            eng = store.load()[1]
        except Exception as e:  # noqa: BLE001
            log.warning("dist: emsal veri kümesi okunamadı: %s", e)
    min_launch = M.mi(first_year, 2)
    if eng is not None:
        ds = eng.ds
        t = ds.books.get(code) or M.Book(code=code, name=name or code).finish()
        t_launch = M.mi(entry.year, entry.month)
        p = M.params_for(6)
        em = set(ds.emsal.get(code, []))
        scored = []
        for c in M.pool_for(ds, ds.end, 2):
            a = ds.books[c]
            if c == code or a.launch is None or a.launch < min_launch or a.launch >= t_launch:
                continue
            s, why = M.similarity(ds, t, t_launch, a, p, em)
            scored.append({"code": c, "ad": a.name, "yazar": a.authors_text or None, "yayinevi": a.publisher or None,
                           "kitaplik": a.library or None, "puan": s, "neden": why})
        scored.sort(key=lambda x: (-x["puan"], x["code"]))
        return scored[: 2 * k], "emsal"
    # Yedek: bütçenin CRM kitap kartı önbelleği (yazar, kitaplık, yayınevi, ilk yayın).
    from semantic_bridge import budget as B

    B.ensure(engine)
    with engine.connect() as c:
        rows = c.execute(sa.select(B.BOOKINFO)).all()
    me = next((r for r in rows if r.stok_kodu == code), None)
    if me is None:
        return [], "yok"
    authors = {fold(x) for x in re.split(r"[,;/]", me.yazar or "") if x.strip()}
    out = []
    lo = date(first_year, 2, 1)
    for r in rows:
        d = _day(r.ilk_yayin)
        if r.stok_kodu == code or not d or d < lo or d >= entry - timedelta(days=60) or not r.in_logo:
            continue
        s, why = 0.05, []
        if authors and authors & {fold(x) for x in re.split(r"[,;/]", r.yazar or "") if x.strip()}:
            s += 2; why.append("aynı yazar")
        if me.kitaplik and fold(me.kitaplik) == fold(r.kitaplik):
            s += 1; why.append("aynı kitaplık")
        if me.yayinevi and fold(me.yayinevi) == fold(r.yayinevi):
            s += 0.5; why.append("aynı yayınevi")
        if not why:
            continue
        s *= math.exp(-max(0, (entry - d).days) / 365 / 3)
        out.append({"code": r.stok_kodu, "ad": r.ad, "yazar": r.yazar, "yayinevi": r.yayinevi, "kitaplik": r.kitaplik,
                    "puan": s, "neden": why})
    out.sort(key=lambda x: (-x["puan"], x["code"]))
    return out[: 2 * k], "kitap-karti"


#: Kapalı küme seçimin seçenekleri → karar; otomatik kabul eşiği (LLM kapısı önerisi, docs/analiz/llm-choose.md).
SCREEN_CHOICES = {"benzer": "benzer", "az benzer": "az", "benzemez": "degil"}
SCREEN_MIN_PROB, SCREEN_MIN_MARGIN = 0.90, 0.50


def _describe(c: dict[str, Any]) -> str:
    return (f"{c.get('ad') or c.get('code')} | yazar: {c.get('yazar') or '-'} | kitaplık: {c.get('kitaplik') or '-'} | "
            f"yayınevi: {c.get('yayinevi') or '-'}")


def _model_screen(llm: Any, target: dict[str, Any], cands: list[dict[str, Any]]) -> tuple[dict[str, str], dict[str, float]]:
    """ZEKİ AI: her aday için «benzer / az / değil» ve «en az az benzer» olasılığı. Kapalı küme; rakam istenmez.

    Kapı `choose` verirse aday başına tek token + olasılık (`QueuedLlm.choose`): karar yalnız p ≥ 0,90 ve marj ≥ 0,50
    ise kabul edilir, altı «emin değil» sayılır ve aday puanıyla kalır. `choose` yoksa tek çağrıda metin listesi.
    Model cevap veremezse ayıklama yapılmaz (boş)."""
    if llm is None or not cands:
        return {}, {}
    if hasattr(llm, "choose"):
        verdict: dict[str, str] = {}
        probs: dict[str, float] = {}
        system = ("Bir yayınevinin satış planlama yardımcısısın. Yeni kitabın ilk dağılımı için eski bir kitabın okur kitlesi "
                  "ve satış kanalı bakımından karşılaştırılabilir olup olmadığına karar verirsin.")
        for c in cands:
            prompt = (f"Yeni kitap: {_describe({**target, 'code': target['code']})}\n"
                      f"Eski kitap: {_describe(c)} | ortak özellik: {', '.join(c.get('neden') or []) or '-'}\n"
                      "Eski kitabın ilk haftalardaki müşteri ve bölge dağılımı yeni kitaba örnek alınabilir mi?")
            try:
                res = llm.choose(prompt, list(SCREEN_CHOICES), system=system)
            except Exception as e:  # noqa: BLE001 — model yok: ayıklama yapılmaz
                log.warning("dist: aday ayıklaması yapılamadı: %s", e)
                return {}, {}
            if res.probs:
                probs[c["code"]] = round(float(res.probs.get("benzer", 0.0)) + float(res.probs.get("az benzer", 0.0)), 4)
            if res.choice in SCREEN_CHOICES and res.confident(SCREEN_MIN_PROB, SCREEN_MIN_MARGIN):
                verdict[c["code"]] = SCREEN_CHOICES[res.choice]
        return verdict, probs
    lines = [f"{c['code']} | {c.get('ad') or ''} | yazar: {c.get('yazar') or '-'} | kitaplık: {c.get('kitaplik') or '-'} | "
             f"yayınevi: {c.get('yayinevi') or '-'} | ortak: {', '.join(c.get('neden') or []) or '-'}" for c in cands]
    msg = [
        {"role": "system", "content": "Bir yayınevinin satış planlama yardımcısısın. Yeni kitabın ilk dağılımı için hangi "
                                      "eski kitapların okur kitlesi ve satış kanalı bakımından karşılaştırılabilir olduğuna karar verirsin. "
                                      "Yalnız verilen listeden, verilen biçimde cevap verirsin."},
        {"role": "user", "content": (
            f"Yeni kitap: {target.get('ad') or target['code']} | yazar: {target.get('yazar') or '-'} | "
            f"kitaplık: {target.get('kitaplik') or '-'} | yayınevi: {target.get('yayinevi') or '-'}\n\nAdaylar:\n"
            + "\n".join(lines)
            + "\n\nHer aday için tek satır yaz: `KOD: benzer` ya da `KOD: az` ya da `KOD: değil`. Başka bir şey yazma.")},
    ]
    try:
        raw = llm.chat(msg, max_tokens=40 * len(cands) + 64, temperature=0.0)
    except Exception as e:  # noqa: BLE001
        log.warning("dist: aday ayıklaması yapılamadı: %s", e)
        return {}, {}
    # Satırda geçen aday kodu (uzun kod önce; kodun parçası olan kısa kod eşleşmesin) ve karar sözcüğü.
    pats = [(c["code"], re.compile(r"(?<![0-9A-Za-z.])" + re.escape(c["code"]) + r"(?![0-9A-Za-z])"))
            for c in sorted(cands, key=lambda x: -len(x["code"]))]
    out: dict[str, str] = {}
    for line in str(raw or "").splitlines():
        code = next((k for k, rx in pats if rx.search(line)), None)
        if not code or code in out:
            continue
        tail = fold(line.split(code, 1)[1])
        if re.search(r"\bdegil\b", tail):
            out[code] = "degil"
        elif re.search(r"\baz\b", tail):
            out[code] = "az"
        elif "benzer" in tail:
            out[code] = "benzer"
    return out, {}


def _rationale(llm: Any, facts: str, fallback: str) -> tuple[str, str]:
    """Gerekçe metni. ZEKİ AI 3–5 cümle yazar; metinde olgularda olmayan bir sayı varsa kural metni kullanılır."""
    if llm is None:
        return fallback, "kural"
    msg = [
        {"role": "system", "content": "Bir yayınevinin satış müdürüne ilk dağılım önerisinin gerekçesini Türkçe, 3–5 cümleyle "
                                      "yazarsın. Yalnız verilen olgulardaki sayıları kullanırsın; yeni sayı, yüzde ya da tahmin yazmazsın. "
                                      "Teknoloji ya da model adı yazmazsın."},
        {"role": "user", "content": "Olgular:\n" + facts + "\n\nGerekçeyi yaz."},
    ]
    try:
        text = str(llm.chat(msg, max_tokens=500, temperature=0.2) or "").strip()
    except Exception as e:  # noqa: BLE001
        log.warning("dist: gerekçe yazılamadı: %s", e)
        return fallback, "kural"
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S).strip()
    if not safe_model_text(text, facts):
        log.info("dist: model gerekçesi olgu dışı sayı içeriyor, kural metni kullanıldı")
        return fallback, "kural"
    return text[:2000], "model"


def _tr(v: Optional[float]) -> str:
    if v is None:
        return "—"
    return f"{int(round(v)):,}".replace(",", ".")


def _pct(v: float) -> str:
    return f"%{round(v * 100):d}"


def build_proposal(engine: sa.engine.Engine, tenant: str, S: Sources, code: str) -> dict[str, Any]:
    """Öneriyi kurar (yazmaz): satırlar, benzer kitaplar, toplam, gerekçe olguları."""
    p = params()
    with engine.connect() as c:
        book = c.execute(sa.select(BOOKS).where(BOOKS.c.tenant_id == tenant, BOOKS.c.stok_kodu == code)).first()
    if book is None:
        raise DistError("Kitap dağılım bekleyen listesinde değil; önce listeyi yenileyin.", 404)
    entry = _day(book.depo_giris_tarihi) or today()
    logo = S.logo()
    first_year = logo.first_year()
    logo_end = logo.data_end() or today()
    comps_meta: dict[str, dict[str, Any]] = {}
    source = "kendi"
    model_used = False
    warnings: list[str] = []
    windows: list[tuple[str, date, date]] = []
    if book.ilk_baski is False:
        # Baskı tekrarı: kitabın depo girişinden önceki kendi son penceresi.
        windows = [(code, entry - timedelta(days=p["pencere"]), entry)]
        comps_meta[code] = {"code": code, "ad": book.ad, "yazar": None, "yayinevi": book.yayinevi, "puan": 1.0,
                            "neden": ["kitabın kendi son satışı"], "karar": None, "secildi": True}
    else:
        cands, source = _candidates(engine, S, code, book.ad, entry, first_year, p["benzer"])
        if not cands:
            warnings.append("Benzer kitap bulunamadı; öneri boş, adetler elle girilir.")
        verdict: dict[str, str] = {}
        model_probs: dict[str, float] = {}
        if cands and p["model"]:
            target = {"code": code, "ad": book.ad, "yayinevi": book.yayinevi}
            store = S.m10()
            try:
                eng = store.load()[1] if store is not None else None
                tb = eng.ds.books.get(code) if eng is not None else None
                if tb is not None:
                    target.update(yazar=tb.authors_text or None, kitaplik=tb.library or None)
            except Exception:  # noqa: BLE001 — hedef kitabın kartı yoksa adıyla sorulur
                pass
            verdict, model_probs = _model_screen(S.llm(), target, cands)
            model_used = bool(verdict or model_probs)
        kept = [c for c in cands if verdict.get(c["code"]) != "degil"]
        if verdict and len(kept) < 2:
            warnings.append("ZEKİ AI ayıklaması yeterli benzer kitap bırakmadı; puan sırası kullanıldı.")
            verdict, kept, model_used = {}, cands, False
        for c in cands:
            comps_meta[c["code"]] = {**c, "karar": verdict.get(c["code"]), "olasilik": model_probs.get(c["code"]), "secildi": False}
        firsts = logo.first_sales([c["code"] for c in kept], first_year) if kept else {}
        lo = date(first_year, 1, 1)
        for c in kept:
            d0 = firsts.get(c["code"])
            if not d0 or d0 < lo + timedelta(days=15) or d0 + timedelta(days=p["pencere"]) > logo_end:
                comps_meta[c["code"]]["neden"] = list(comps_meta[c["code"]]["neden"]) + ["penceresi okunamadı"]
                continue
            windows.append((c["code"], d0, d0 + timedelta(days=p["pencere"])))
            if len(windows) >= p["benzer"]:
                break
    rows = logo.comp_windows(windows) if windows else []
    weights: dict[str, float] = {}
    for k, _, _ in windows:
        m = comps_meta[k]
        weights[k] = max(0.0, float(m["puan"] or 0)) * (0.5 if m.get("karar") == "az" else 1.0)
    share, stats, totals = comp_shares(rows, weights)
    for k, a, b in windows:
        m = comps_meta[k]
        m["pencere"] = (a.isoformat(), b.isoformat())
        m["secildi"] = totals.get(k, 0) > 0
        if not m["secildi"]:
            m["neden"] = list(m["neden"]) + ["8 haftada net satış yok"]
    per_comp: dict[str, dict[str, Any]] = {}
    for r in rows:
        pc = per_comp.setdefault(r["comp"], {"satis": 0.0, "iade": 0.0, "ciro": 0.0, "cariler": set(), "kanal": {}})
        pc["satis"] += r["satis"]
        pc["iade"] += r["iade"]
        pc["ciro"] += r["ciro"]
        pc["cariler"].add(r["cari_kodu"])
    # Cari öznitelikleri (güncel kopya) ve CRM.
    accounts: list[dict[str, Any]] = []
    try:
        accounts = S.accounts()
    except Exception as e:  # noqa: BLE001
        warnings.append(f"CRM carileri okunamadı; BMT ve dağılım listesi eksik: {str(e)[:120]}")
    by_code = {a["cari_kodu"]: a for a in accounts if a.get("cari_kodu")}
    dist_accounts = [a for a in accounts if a.get("dagilim")]
    logo_codes = sorted(set(share) | set(stats) | {a["cari_kodu"] for a in dist_accounts if a.get("cari_kodu")})
    clients = logo.clients(logo_codes) if logo_codes else {}
    for r in rows:
        ch = channel_label((clients.get(r["cari_kodu"]) or {}).get("kanal"))
        pc = per_comp[r["comp"]]
        pc["kanal"][ch] = pc["kanal"].get(ch, 0.0) + max(0.0, r["satis"] - r["iade"])
    try:
        dist_hist = S.dist_orders([k for k, _, _ in windows])
    except Exception as e:  # noqa: BLE001
        dist_hist = {}
        warnings.append(f"CRM geçmiş dağılım siparişleri okunamadı: {str(e)[:120]}")
    # Toplam adet.
    target = _target(engine, tenant, code, entry)
    comps_median = weighted_median([totals.get(k, 0) for k, _, _ in windows])
    stock = book.stok_bakiye
    stock_ref = stock if stock is not None else book.baski_adedi
    reserve = float(round((stock_ref or 0) * p["rezervPay"]))
    available = max(0.0, (stock_ref or 0) - reserve) if stock_ref is not None else None
    if target.get("var") and target.get("ikiAyAdet"):
        wanted, basis = float(target["ikiAyAdet"]), "hedef"
    elif comps_median:
        wanted, basis = float(round(comps_median)), "benzer"
    else:
        wanted, basis = 0.0, "yok"
    total = min(wanted, available) if available is not None else wanted
    alloc = allocate(total, share, p["enAz"])
    # Satırlar: öneri payı olan cariler + CRM dağılım carileri.
    line_rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for cari in sorted(set(share) | {a["cari_kodu"] for a in dist_accounts if a.get("cari_kodu")},
                       key=lambda k: (-share.get(k, 0.0), k)):
        if cari in seen:
            continue
        seen.add(cari)
        cl, acc, st = clients.get(cari) or {}, by_code.get(cari) or {}, stats.get(cari) or {}
        il = cl.get("il") or acc.get("il")
        sat = st.get("satis", 0.0)
        line_rows.append({
            "logo_cari_kodu": cari[:60], "crm_account_id": acc.get("id"),
            "cari_unvan": (cl.get("unvan") or acc.get("ad") or cari)[:300], "il": (il or "")[:80] or None,
            "bolge": region_of(il, cl.get("kanal")), "kanal": channel_label(cl.get("kanal"))[:60],
            "kanal_kodu": (cl.get("kanal") or "")[:60] or None, "bmt_ad": (acc.get("bmt_ad") or "")[:200] or None,
            "bmt_hesap": (acc.get("bmt_hesap") or "")[:120] or None, "dagilim_carisi": bool(acc.get("dagilim")),
            "pay": share.get(cari, 0.0), "gecmis_net": round(st.get("net", 0.0), 2) if st else None,
            "gecmis_iade_orani": (st.get("iade", 0.0) / sat) if sat > 0 else None,
            "onerilen_adet": float(alloc.get(cari, 0)), "adet": float(alloc.get(cari, 0)),
        })
    for a in dist_accounts:
        if a.get("cari_kodu"):
            continue
        line_rows.append({
            "logo_cari_kodu": None, "crm_account_id": a.get("id"), "cari_unvan": (a.get("ad") or "")[:300],
            "il": (a.get("il") or "")[:80] or None, "bolge": region_of(a.get("il")), "kanal": "Logo kodu yok",
            "kanal_kodu": None, "bmt_ad": a.get("bmt_ad"), "bmt_hesap": a.get("bmt_hesap"), "dagilim_carisi": True,
            "pay": 0.0, "gecmis_net": None, "gecmis_iade_orani": None, "onerilen_adet": 0.0, "adet": 0.0})
    comp_rows = []
    for k, m in comps_meta.items():
        pc = per_comp.get(k) or {}
        h = dist_hist.get(k) or {}
        comp_rows.append({
            "comp_stok_kodu": k[:60], "ad": (m.get("ad") or "")[:400] or None, "yazar": (m.get("yazar") or "")[:300] or None,
            "yayinevi": (m.get("yayinevi") or "")[:200] or None,
            "pencere_bas": (m.get("pencere") or (None, None))[0], "pencere_bit": (m.get("pencere") or (None, None))[1],
            "puan": round(float(m.get("puan") or 0), 4), "model_karar": m.get("karar"),
            "gerekce": ", ".join(m.get("neden") or [])[:500] or None, "secildi": bool(m.get("secildi")),
            "agirlik": weights.get(k), "satis_adet": pc.get("satis"), "iade_adet": pc.get("iade"),
            "net_ciro": round(pc["ciro"], 2) if pc else None, "musteri": len(pc.get("cariler") or []) if pc else None,
            "kanallar_json": _dump(pc.get("kanal") or {}), "crm_dagilim_adet": h.get("adet"),
            "crm_dagilim_siparis": h.get("siparis")})
    # Benzerlik olasılığı: ZEKİ AI'ın «benzer» + «az benzer» olasılığı (kapı verdiyse); yoksa puanın 0–1 ölçeği.
    top = max((c["puan"] for c in comp_rows), default=0) or 1
    for c in comp_rows:
        pm = (comps_meta.get(c["comp_stok_kodu"]) or {}).get("olasilik")
        c["benzerlik_olasiligi"] = pm if pm is not None else (round(c["puan"] / top, 4) if c["puan"] else 0.0)
    used = [c for c in comp_rows if c["secildi"]]
    mx = matrix([{**ln} for ln in line_rows])
    facts_lines = [
        f"Kitap: {book.ad or code} ({'ilk baskı' if book.ilk_baski is not False else 'baskı tekrarı'}), depoya giriş {entry.strftime('%d.%m.%Y')}.",
        f"Stok bakiyesi: {_tr(stock)} adet; rezerv: {_tr(reserve)} adet; dağıtılabilir: {_tr(available)} adet.",
    ]
    if target.get("var"):
        facts_lines.append(f"Onaylı yıllık satış hedefi {_tr(target['yillikAdet'])} adet; depo girişinden itibaren ilk iki ayın hedefi {_tr(target['ikiAyAdet'])} adet.")
    if comps_median:
        facts_lines.append(f"Benzer kitapların ilk {p['pencere']} gündeki net satış ortancası {_tr(comps_median)} adet.")
    for cr in sorted(used, key=lambda x: -(x["agirlik"] or 0))[:5]:
        n = (cr["satis_adet"] or 0) - (cr["iade_adet"] or 0)
        facts_lines.append(f"Benzer kitap: {cr['ad'] or cr['comp_stok_kodu']} — ilk {p['pencere']} günde {_tr(n)} adet net, "
                           f"{_tr(cr['musteri'])} müşteri; ortak özellik: {cr['gerekce'] or '-'}.")
    facts_lines.append(f"Önerilen toplam: {_tr(total)} adet, {_tr(sum(1 for x in line_rows if x['adet'] > 0))} müşteriye.")
    for rg in mx["bolgeler"][:4]:
        if rg["adet"] > 0:
            facts_lines.append(f"Bölge {rg['bolge']}: {_tr(rg['adet'])} adet ({_pct(rg['pay'])}).")
    for ch in mx["kanallar"][:4]:
        if ch["adet"] > 0:
            facts_lines.append(f"Kanal {ch['kanal']}: {_tr(ch['adet'])} adet ({_pct(ch['pay'])}).")
    facts = "\n".join(facts_lines)
    fallback = _fallback_text(book, entry, used, total, basis, target, comps_median, reserve, mx, p)
    text, text_src = _rationale(S.llm() if p["model"] else None, facts, fallback)
    return {
        "book": book, "lines": line_rows, "comps": comp_rows, "total": total, "wanted": wanted, "reserve": reserve,
        "stock": stock, "target": target, "text": text, "textSource": text_src,
        "basis": {"yontem": "kendi" if book.ilk_baski is False else source, "toplamKaynak": basis,
                  "benzerOrtanca": comps_median, "istenen": wanted, "dagitilabilir": available,
                  "stokKaynak": "logo" if stock is not None else ("baski" if book.baski_adedi is not None else None),
                  "pencereGun": p["pencere"], "rezervPay": p["rezervPay"], "enAz": p["enAz"], "modelAyiklama": model_used,
                  "benzerSayisi": len(used), "veriSonu": logo_end.isoformat(), "uyarilar": warnings,
                  "olgular": facts},
    }


def _fallback_text(book: Any, entry: date, used: list[dict[str, Any]], total: float, basis: str, target: dict[str, Any],
                   comps_median: Optional[float], reserve: float, mx: dict[str, Any], p: dict[str, Any]) -> str:
    if book.ilk_baski is False:
        s1 = f"Baskı tekrarı: dağılım kitabın depo girişinden önceki {p['pencere']} gündeki kendi müşteri dağılımından kuruldu."
    elif used:
        names = ", ".join((c["ad"] or c["comp_stok_kodu"]) for c in sorted(used, key=lambda x: -(x["agirlik"] or 0))[:3])
        s1 = f"Öneri {len(used)} benzer kitabın ilk {p['pencere']} gündeki müşteri dağılımından kuruldu ({names})."
    else:
        s1 = "Benzer kitap bulunamadı; adetler elle girilmeli."
    if basis == "hedef":
        s2 = f"Toplam {_tr(total)} adet: onaylı satış hedefinin depo girişinden itibaren ilk iki ayı ({_tr(target.get('ikiAyAdet'))} adet)"
    elif basis == "benzer":
        s2 = f"Toplam {_tr(total)} adet: benzer kitapların ilk {p['pencere']} gündeki net satış ortancası ({_tr(comps_median)} adet)"
    else:
        s2 = "Toplam adet için hedef ya da benzer kitap satışı yok"
    s2 += f"; depoda {_tr(reserve)} adet rezerv kalır."
    regs = [f"{r['bolge']} {_pct(r['pay'])}" for r in mx["bolgeler"] if r["adet"] > 0][:3]
    s3 = f"En büyük paylar: {', '.join(regs)}." if regs else ""
    return " ".join(x for x in (s1, s2, s3) if x)


# ------------------------------------------------------------------ plan yazma


def _plan_row(c: Any, tenant: str, plan_id: str, *, lock: bool = False) -> Any:
    q = sa.select(PLANS).where(PLANS.c.tenant_id == tenant, PLANS.c.id == str(plan_id)[:32])
    if lock and c.engine.dialect.name == "postgresql":
        q = q.with_for_update()
    row = c.execute(q).first()
    if not row:
        raise DistError("Plan bulunamadı.", 404)
    return row


def _next_version(c: Any, tenant: str, code: str) -> int:
    v = c.execute(sa.select(sa.func.max(PLANS.c.surum)).where(PLANS.c.tenant_id == tenant, PLANS.c.stok_kodu == code)).scalar()
    return int(v or 0) + 1


def _write_lines(c: Any, plan_id: str, lines: list[dict[str, Any]]) -> None:
    if lines:
        c.execute(LINES.insert(), [{"plan_id": plan_id, "line_no": i + 1, "elle_duzeltildi": False, **ln}
                                   for i, ln in enumerate(lines)])


def _sync_total(c: Any, plan_id: str) -> float:
    t = c.execute(sa.select(sa.func.coalesce(sa.func.sum(LINES.c.adet), 0)).where(LINES.c.plan_id == plan_id)).scalar()
    c.execute(PLANS.update().where(PLANS.c.id == plan_id).values(toplam_adet=float(t or 0)))
    return float(t or 0)


def generate(engine: sa.engine.Engine, tenant: str, user: str, code: str, S: Sources) -> dict[str, Any]:
    """ZEKİ AI önerisi: kitap için taslak plan. Açık (taslak/onayda) plan varsa 409; onaylı varsa revize edilir."""
    code = str(code or "").strip()[:60]
    if not code:
        raise DistError("Stok kodu gerekli.")
    with engine.connect() as c:
        open_ = c.execute(sa.select(PLANS.c.id, PLANS.c.durum).where(
            PLANS.c.tenant_id == tenant, PLANS.c.stok_kodu == code, PLANS.c.durum.in_(("taslak", "onayda", "onayli")))).all()
    for r in open_:
        if r.durum in ("taslak", "onayda"):
            raise DistError(f"Bu kitabın açık bir planı var ({STATUSES[r.durum]}); onu açın.", 409)
        raise DistError("Bu kitabın onaylı planı var; değişiklik için «Revize et».", 409)
    prop = build_proposal(engine, tenant, S, code)
    book = prop["book"]
    now = _now()
    pid = uuid.uuid4().hex
    with engine.begin() as c:
        v = _next_version(c, tenant, code)
        c.execute(PLANS.insert().values(
            id=pid, tenant_id=tenant, stok_kodu=code, ad=book.ad, surum=v, durum="taslak",
            depo_giris_tarihi=book.depo_giris_tarihi, baski_adedi=book.baski_adedi, stok_bakiye=book.stok_bakiye,
            toplam_adet=0, onerilen_toplam=prop["total"], rezerv_adet=prop["reserve"],
            hedef_plan_id=prop["target"].get("planId"), hedef_json=_dump(prop["target"]), basis_json=_dump(prop["basis"]),
            gerekce=prop["text"], gerekce_kaynak=prop["textSource"], created_by=user, created_at=now, updated_by=user,
            updated_at=now))
        _write_lines(c, pid, prop["lines"])
        if prop["comps"]:
            c.execute(COMPS.insert(), [{"plan_id": pid, **cr} for cr in prop["comps"]])
        _sync_total(c, pid)
    return plan_detail(engine, tenant, pid)


def _editable(row: Any) -> None:
    if row.durum != "taslak":
        raise DistError("Yalnız taslak plan değiştirilebilir; onaylı planı «Revize et» ile yeni sürüme alın.", 409)


def update_plan(engine: sa.engine.Engine, tenant: str, user: str, plan_id: str, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    vals: dict[str, Any] = {}
    if "rezerv" in body:
        vals["rezerv_adet"] = _qty(body.get("rezerv"), "Rezerv")
    if "note" in body:
        vals["note"] = _text(body.get("note"), 4000)
    if not vals:
        raise DistError("Değiştirilecek alan yok.")
    with engine.begin() as c:
        row = _plan_row(c, tenant, plan_id, lock=True)
        _editable(row)
        diff = {k: {"eski": getattr(row, k), "yeni": v} for k, v in vals.items() if getattr(row, k) != v}
        c.execute(PLANS.update().where(PLANS.c.id == row.id).values(**vals, updated_by=user, updated_at=_now()))
    return plan_detail(engine, tenant, plan_id), diff


def update_line(engine: sa.engine.Engine, tenant: str, user: str, plan_id: str, line_no: int, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    vals: dict[str, Any] = {}
    if "adet" in body:
        vals["adet"] = _qty(body.get("adet"))
    if "gerekce" in body:
        vals["gerekce"] = _text(body.get("gerekce"), 500)
    if not vals:
        raise DistError("Değiştirilecek alan yok.")
    with engine.begin() as c:
        row = _plan_row(c, tenant, plan_id, lock=True)
        _editable(row)
        ln = c.execute(sa.select(LINES).where(LINES.c.plan_id == row.id, LINES.c.line_no == int(line_no))).first()
        if ln is None:
            raise DistError("Satır bulunamadı.", 404)
        diff = {k: {"eski": getattr(ln, k), "yeni": v} for k, v in vals.items() if getattr(ln, k) != v}
        if "adet" in vals:
            vals["elle_duzeltildi"] = vals["adet"] != ln.onerilen_adet
        c.execute(LINES.update().where(LINES.c.plan_id == row.id, LINES.c.line_no == ln.line_no).values(
            **vals, edited_by=user, edited_at=_now()))
        _sync_total(c, row.id)
        c.execute(PLANS.update().where(PLANS.c.id == row.id).values(updated_by=user, updated_at=_now()))
        out = c.execute(sa.select(LINES).where(LINES.c.plan_id == row.id, LINES.c.line_no == ln.line_no)).first()
    return _line_dict(out), {"satir": ln.cari_unvan, **diff}


def update_cell(engine: sa.engine.Engine, tenant: str, user: str, plan_id: str, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """Bölge × kanal hücresinin toplamını değiştirir; hücredeki carilere mevcut adetleri (hepsi sıfırsa önerilen payları)
    oranında dağıtılır."""
    bolge, kanal = str(body.get("bolge") or ""), str(body.get("kanal") or "")
    target = _qty(body.get("adet"))
    why = _text(body.get("gerekce"), 500)
    with engine.begin() as c:
        row = _plan_row(c, tenant, plan_id, lock=True)
        _editable(row)
        lines = c.execute(sa.select(LINES).where(LINES.c.plan_id == row.id, LINES.c.bolge == bolge, LINES.c.kanal == kanal)).all()
        if not lines:
            raise DistError("Bu bölge ve kanalda müşteri yok.", 404)
        weights = {str(ln.line_no): ln.adet for ln in lines}
        if sum(weights.values()) <= 0:
            weights = {str(ln.line_no): ln.pay for ln in lines}
        if sum(weights.values()) <= 0:
            weights = {str(ln.line_no): 1.0 for ln in lines if ln.logo_cari_kodu}
        if not weights and target > 0:
            raise DistError("Bu hücrede Logo kodu olan müşteri yok; adet satırdan girilmeli.", 409)
        alloc = allocate(target, weights)
        before = sum(ln.adet for ln in lines)
        now = _now()
        for ln in lines:
            new = float(alloc.get(str(ln.line_no), 0))
            if new != ln.adet:
                c.execute(LINES.update().where(LINES.c.plan_id == row.id, LINES.c.line_no == ln.line_no).values(
                    adet=new, elle_duzeltildi=new != ln.onerilen_adet, gerekce=why or ln.gerekce, edited_by=user, edited_at=now))
        _sync_total(c, row.id)
        c.execute(PLANS.update().where(PLANS.c.id == row.id).values(updated_by=user, updated_at=now))
    return plan_detail(engine, tenant, plan_id), {"bolge": bolge, "kanal": kanal, "eski": before, "yeni": target, "gerekce": why}


def delete_plan(engine: sa.engine.Engine, tenant: str, plan_id: str) -> dict[str, Any]:
    with engine.begin() as c:
        row = _plan_row(c, tenant, plan_id, lock=True)
        if row.durum != "taslak":
            raise DistError("Yalnız taslak plan silinebilir.", 409)
        for t in (LINES, COMPS, TRACKING):
            c.execute(t.delete().where(t.c.plan_id == row.id))
        c.execute(PLANS.delete().where(PLANS.c.id == row.id))
    return {"id": row.id, "title": f"{row.ad or row.stok_kodu} · sürüm {row.surum}"}


def _current_stock(c: Any, tenant: str, row: Any) -> tuple[Optional[float], str]:
    b = c.execute(sa.select(BOOKS.c.stok_bakiye, BOOKS.c.baski_adedi).where(
        BOOKS.c.tenant_id == tenant, BOOKS.c.stok_kodu == row.stok_kodu)).first()
    if b is not None and b.stok_bakiye is not None:
        return float(b.stok_bakiye), "stok"
    if row.stok_bakiye is not None:
        return float(row.stok_bakiye), "stok"
    base = (b.baski_adedi if b is not None else None) or row.baski_adedi
    return (float(base), "baski") if base is not None else (None, "yok")


def _check_stock(c: Any, tenant: str, row: Any) -> None:
    total = _sync_total(c, row.id)
    stock, kind = _current_stock(c, tenant, row)
    if total <= 0:
        raise DistError("Planda dağıtılacak adet yok.")
    if stock is not None and total + (row.rezerv_adet or 0) > stock:
        what = "stok bakiyesini" if kind == "stok" else "baskı adedini (stok Logo'da görünmüyor)"
        raise DistError(f"Dağılım {_tr(total)} + rezerv {_tr(row.rezerv_adet)} = {_tr(total + (row.rezerv_adet or 0))} adet, "
                        f"{what} ({_tr(stock)}) aşıyor. Adetleri ya da rezervi azaltın.", 409)


def submit(engine: sa.engine.Engine, tenant: str, user: str, plan_id: str) -> dict[str, Any]:
    with engine.begin() as c:
        row = _plan_row(c, tenant, plan_id, lock=True)
        if row.durum != "taslak":
            raise DistError("Yalnız taslak onaya gönderilir.", 409)
        _check_stock(c, tenant, row)
        now = _now()
        c.execute(PLANS.update().where(PLANS.c.id == row.id).values(
            durum="onayda", submitted_by=user, submitted_at=now, decision_note=None, updated_by=user, updated_at=now))
    return plan_detail(engine, tenant, plan_id)


def withdraw(engine: sa.engine.Engine, tenant: str, user: str, plan_id: str) -> dict[str, Any]:
    with engine.begin() as c:
        row = _plan_row(c, tenant, plan_id, lock=True)
        if row.durum != "onayda":
            raise DistError("Plan onayda değil.", 409)
        c.execute(PLANS.update().where(PLANS.c.id == row.id).values(durum="taslak", updated_by=user, updated_at=_now()))
    return plan_detail(engine, tenant, plan_id)


def decide(engine: sa.engine.Engine, tenant: str, user: str, plan_id: str, approve: bool, note: Any = None) -> dict[str, Any]:
    """Lojistik onayı ya da geri gönderme. Gönderen onaylayamaz (iki göz); onayda stok değişmezi yeniden denetlenir,
    kitabın önceki onaylı planı arşive geçer."""
    note_t = _text(note, 2000)
    if not approve and not note_t:
        raise DistError("Geri gönderme gerekçesi yazın.")
    archived = None
    with engine.begin() as c:
        row = _plan_row(c, tenant, plan_id, lock=True)
        if row.durum != "onayda":
            raise DistError("Plan onay beklemiyor.", 409)
        if (row.submitted_by or "").lower() == user.lower():
            raise DistError("Onaya gönderen kişi aynı planı onaylayamaz; başka bir yetkili onaylamalı.", 409)
        now = _now()
        if approve:
            _check_stock(c, tenant, row)
            prev = c.execute(sa.select(PLANS).where(PLANS.c.tenant_id == tenant, PLANS.c.stok_kodu == row.stok_kodu,
                                                    PLANS.c.durum == "onayli", PLANS.c.id != row.id)).first()
            if prev:
                c.execute(PLANS.update().where(PLANS.c.id == prev.id).values(durum="arsiv", updated_by=user, updated_at=now))
                archived = prev.id
            c.execute(PLANS.update().where(PLANS.c.id == row.id).values(
                durum="onayli", decided_by=user, decided_at=now, decision_note=note_t, updated_by=user, updated_at=now))
        else:
            c.execute(PLANS.update().where(PLANS.c.id == row.id).values(
                durum="taslak", decided_by=user, decided_at=now, decision_note=note_t, updated_by=user, updated_at=now))
    out = plan_detail(engine, tenant, plan_id)
    out["archived"] = archived
    return out


def revise(engine: sa.engine.Engine, tenant: str, user: str, plan_id: str, reason: Any) -> dict[str, Any]:
    why = _text(reason, 2000)
    if not why:
        raise DistError("Revizyon gerekçesi yazın.")
    with engine.begin() as c:
        row = _plan_row(c, tenant, plan_id)
        if row.durum != "onayli":
            raise DistError("Yalnız onaylı plan revize edilir.", 409)
        if c.execute(sa.select(PLANS.c.id).where(PLANS.c.tenant_id == tenant, PLANS.c.stok_kodu == row.stok_kodu,
                                                 PLANS.c.durum.in_(("taslak", "onayda")))).first():
            raise DistError("Bu kitabın açık bir revizyonu zaten var.", 409)
        now = _now()
        pid = uuid.uuid4().hex
        v = _next_version(c, tenant, row.stok_kodu)
        vals = {k: getattr(row, k) for k in ("stok_kodu", "ad", "depo_giris_tarihi", "baski_adedi", "stok_bakiye", "toplam_adet",
                                             "onerilen_toplam", "rezerv_adet", "hedef_plan_id", "hedef_json", "basis_json",
                                             "gerekce", "gerekce_kaynak", "note")}
        c.execute(PLANS.insert().values(id=pid, tenant_id=tenant, surum=v, durum="taslak", revision_of=row.id,
                                        revision_reason=why, created_by=user, created_at=now, updated_by=user,
                                        updated_at=now, **vals))
        for t in (LINES, COMPS):
            for r in c.execute(sa.select(t).where(t.c.plan_id == row.id)).all():
                c.execute(t.insert().values(**{**dict(r._mapping), "plan_id": pid}))
    return plan_detail(engine, tenant, pid)


# ------------------------------------------------------------------ okuma


def _line_dict(r: Any, track: Optional[dict[str, float]] = None) -> dict[str, Any]:
    out = {"no": r.line_no, "cariKodu": r.logo_cari_kodu, "crmId": r.crm_account_id, "unvan": r.cari_unvan, "il": r.il,
           "bolge": r.bolge, "kanal": r.kanal, "bmt": r.bmt_ad, "bmtHesap": r.bmt_hesap, "dagilimCarisi": r.dagilim_carisi,
           "pay": r.pay, "gecmisNet": r.gecmis_net, "gecmisIadeOrani": r.gecmis_iade_orani, "onerilen": r.onerilen_adet,
           "adet": r.adet, "elle": r.elle_duzeltildi, "gerekce": r.gerekce, "editedBy": r.edited_by,
           "editedAt": _iso(r.edited_at)}
    if track is not None:
        out["takip"] = track
    return out


def _plan_dict(row: Any) -> dict[str, Any]:
    return {
        "id": row.id, "stokKodu": row.stok_kodu, "ad": row.ad, "surum": row.surum, "durum": row.durum,
        "durumEtiket": STATUSES.get(row.durum, row.durum), "depoGiris": row.depo_giris_tarihi, "baskiAdedi": row.baski_adedi,
        "stok": row.stok_bakiye, "toplam": row.toplam_adet, "onerilenToplam": row.onerilen_toplam, "rezerv": row.rezerv_adet,
        "hedef": _j(row.hedef_json, {}), "basis": _j(row.basis_json, {}), "gerekce": row.gerekce,
        "gerekceKaynak": row.gerekce_kaynak, "note": row.note, "revisionOf": row.revision_of,
        "revisionReason": row.revision_reason, "createdBy": row.created_by, "createdAt": _iso(row.created_at),
        "updatedBy": row.updated_by, "updatedAt": _iso(row.updated_at), "submittedBy": row.submitted_by,
        "submittedAt": _iso(row.submitted_at), "decidedBy": row.decided_by, "decidedAt": _iso(row.decided_at),
        "decisionNote": row.decision_note,
    }


def _comp_dict(r: Any) -> dict[str, Any]:
    return {"stokKodu": r.comp_stok_kodu, "ad": r.ad, "yazar": r.yazar, "yayinevi": r.yayinevi,
            "pencere": [r.pencere_bas, r.pencere_bit] if r.pencere_bas else None, "puan": r.puan,
            "benzerlik": r.benzerlik_olasiligi, "modelKarar": r.model_karar, "gerekce": r.gerekce, "secildi": r.secildi,
            "agirlik": r.agirlik, "satis": r.satis_adet, "iade": r.iade_adet,
            "net": None if r.satis_adet is None else (r.satis_adet or 0) - (r.iade_adet or 0), "netCiro": r.net_ciro,
            "musteri": r.musteri, "kanallar": _j(r.kanallar_json, {}), "crmDagilimAdet": r.crm_dagilim_adet,
            "crmDagilimSiparis": r.crm_dagilim_siparis}


def _visible(q: Any, bmt: Optional[str]) -> Any:
    return q.where(sa.func.lower(LINES.c.bmt_hesap) == bmt.lower()) if bmt else q


def plan_detail(engine: sa.engine.Engine, tenant: str, plan_id: str, bmt: Optional[str] = None) -> dict[str, Any]:
    """Plan başlığı, bölge × kanal matrisi (bmt verilirse yalnız o kişinin carileri), benzer kitaplar."""
    with engine.connect() as c:
        row = _plan_row(c, tenant, plan_id)
        lines = c.execute(_visible(sa.select(LINES.c.bolge, LINES.c.kanal, LINES.c.adet, LINES.c.onerilen_adet)
                                   .where(LINES.c.plan_id == row.id), bmt)).all()
        comps = c.execute(sa.select(COMPS).where(COMPS.c.plan_id == row.id)).all()
        stock, kind = _current_stock(c, tenant, row)
        n_lines = len(lines)
        edited = c.execute(sa.select(sa.func.count()).select_from(LINES).where(
            LINES.c.plan_id == row.id, LINES.c.elle_duzeltildi.is_(True))).scalar()
    out = _plan_dict(row)
    out["matris"] = matrix([dict(r._mapping) for r in lines])
    out["benzerler"] = sorted((_comp_dict(r) for r in comps), key=lambda x: (not x["secildi"], -(x["agirlik"] or 0), -(x["puan"] or 0)))
    out["satirSayisi"] = n_lines
    out["musteri"] = sum(1 for r in lines if r.adet > 0)
    out["elleSatir"] = int(edited or 0)
    out["guncelStok"] = {"adet": stock, "tur": kind}
    out["asim"] = bool(stock is not None and row.toplam_adet + (row.rezerv_adet or 0) > stock)
    out["kapsam"] = "kendi" if bmt else "hepsi"
    end = data_end(engine, tenant)
    out["veriSonu"] = end.isoformat() if end else None
    return out


def list_lines(engine: sa.engine.Engine, tenant: str, plan_id: str, *, q: str = "", bolge: str = "", kanal: str = "",
               yalniz: str = "", page: int = 0, bmt: Optional[str] = None) -> dict[str, Any]:
    """Müşteri satırları, sayfalı (sayı tavanı yok; `total` hepsini söyler). yalniz: adetli | elle | dagilim."""
    with engine.connect() as c:
        row = _plan_row(c, tenant, plan_id)
        cond = [LINES.c.plan_id == row.id]
        if bolge:
            cond.append(LINES.c.bolge == bolge)
        if kanal:
            cond.append(LINES.c.kanal == kanal)
        if yalniz == "adetli":
            cond.append(LINES.c.adet > 0)
        elif yalniz == "elle":
            cond.append(LINES.c.elle_duzeltildi.is_(True))
        elif yalniz == "dagilim":
            cond.append(LINES.c.dagilim_carisi.is_(True))
        rows = c.execute(_visible(sa.select(LINES).where(*cond), bmt).order_by(LINES.c.adet.desc(), LINES.c.pay.desc(),
                                                                              LINES.c.cari_unvan)).all()
        track = _track_by_client(c, row.id) if row.durum in ("onayli", "arsiv") else {}
    needle = fold(q)
    if needle:
        rows = [r for r in rows if needle in fold(f"{r.cari_unvan} {r.logo_cari_kodu or ''} {r.il or ''} {r.bmt_ad or ''}")]
    page = max(0, int(page or 0))
    part = rows[page * PAGE_SIZE:(page + 1) * PAGE_SIZE]
    return {"items": [_line_dict(r, track.get(r.logo_cari_kodu) if track else None) for r in part], "total": len(rows),
            "page": page, "pageSize": PAGE_SIZE}


def get_plan_row(engine: sa.engine.Engine, tenant: str, plan_id: str) -> Any:
    with engine.connect() as c:
        return _plan_row(c, tenant, plan_id)


def plans_of(engine: sa.engine.Engine, tenant: str, code: str) -> list[dict[str, Any]]:
    with engine.connect() as c:
        rows = c.execute(sa.select(PLANS).where(PLANS.c.tenant_id == tenant, PLANS.c.stok_kodu == code)
                         .order_by(PLANS.c.surum.desc())).all()
    return [_plan_dict(r) for r in rows]


# ------------------------------------------------------------------ sevk listesi (Excel)


def export_xlsx(engine: sa.engine.Engine, tenant: str, plan_id: str) -> tuple[bytes, str]:
    """Onaylı planın sevk listesi. CRM «Dağılım» siparişine elle yazmaya uygun kolonlar; Logo'ya ve CRM'e yazılmaz."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    with engine.connect() as c:
        row = _plan_row(c, tenant, plan_id)
        if row.durum != "onayli":
            raise DistError("Sevk listesi yalnız onaylı plandan alınır.", 409)
        lines = c.execute(sa.select(LINES).where(LINES.c.plan_id == row.id, LINES.c.adet > 0)
                          .order_by(LINES.c.bolge, LINES.c.il, LINES.c.cari_unvan)).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Sevk listesi"
    head = Font(bold=True, size=14)
    ws["A1"] = f"İlk dağılım sevk listesi — {row.ad or row.stok_kodu}"
    ws["A1"].font = head
    ws["A2"] = (f"Stok kodu {row.stok_kodu} · sürüm {row.surum} · onaylayan {row.decided_by or '-'} · "
                f"{(row.decided_at.astimezone(TZ) if row.decided_at and row.decided_at.tzinfo else row.decided_at or _now()).strftime('%d.%m.%Y %H:%M')}")
    ws["A3"] = f"Toplam {_tr(row.toplam_adet)} adet · {len(lines)} müşteri · rezerv {_tr(row.rezerv_adet)} adet · Sipariş tipi: Dağılım"
    cols = ["Sıra", "Sipariş tipi", "Cari kodu", "Cari unvanı", "İl", "Bölge", "Kanal", "BMT", "Stok kodu", "Kitap", "Adet", "Not"]
    hr = 5
    for i, h in enumerate(cols, 1):
        cell = ws.cell(row=hr, column=i, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4B3FA8")
    for n, ln in enumerate(lines, 1):
        vals = [n, "Dağılım", ln.logo_cari_kodu, ln.cari_unvan, ln.il, ln.bolge, ln.kanal, ln.bmt_ad, row.stok_kodu,
                row.ad, int(ln.adet), ln.gerekce]
        for i, v in enumerate(vals, 1):
            ws.cell(row=hr + n, column=i, value=v)
    last = hr + max(1, len(lines))
    if lines:
        ws.add_table(Table(displayName="SevkListesi", ref=f"A{hr}:{get_column_letter(len(cols))}{last}",
                           tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))
    tot = ws.cell(row=last + 1, column=10, value="Toplam")
    tot.font = Font(bold=True)
    ws.cell(row=last + 1, column=11, value=int(sum(ln.adet for ln in lines))).font = Font(bold=True)
    for i, w in enumerate([6, 11, 14, 42, 14, 18, 20, 22, 18, 36, 9, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{hr + 1}"
    mx = matrix([{"bolge": ln.bolge, "kanal": ln.kanal, "adet": ln.adet, "onerilen_adet": ln.onerilen_adet} for ln in lines])
    s2 = wb.create_sheet("Özet")
    s2["A1"] = "Bölge × kanal"
    s2["A1"].font = head
    kan = [k["kanal"] for k in mx["kanallar"]]
    s2.cell(row=3, column=1, value="Bölge").font = Font(bold=True)
    for j, k in enumerate(kan, 2):
        s2.cell(row=3, column=j, value=k).font = Font(bold=True)
        s2.cell(row=3, column=j).alignment = Alignment(wrap_text=True)
    s2.cell(row=3, column=len(kan) + 2, value="Toplam").font = Font(bold=True)
    cells = {(h["bolge"], h["kanal"]): h["adet"] for h in mx["hucreler"]}
    for i, b in enumerate(mx["bolgeler"], 4):
        s2.cell(row=i, column=1, value=b["bolge"])
        for j, k in enumerate(kan, 2):
            v = cells.get((b["bolge"], k))
            s2.cell(row=i, column=j, value=int(v) if v else None)
        s2.cell(row=i, column=len(kan) + 2, value=int(b["adet"])).font = Font(bold=True)
    s2.column_dimensions["A"].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    name = re.sub(r"[^0-9A-Za-z._-]+", "-", f"sevk-listesi-{row.stok_kodu}-s{row.surum}").strip("-") + ".xlsx"
    return buf.getvalue(), name


# ------------------------------------------------------------------ takip


def _track_by_client(c: Any, plan_id: str, upto_week: Optional[int] = None) -> dict[str, dict[str, float]]:
    q = sa.select(TRACKING.c.logo_cari_kodu, sa.func.sum(TRACKING.c.sevk_adet), sa.func.sum(TRACKING.c.satis_adet),
                  sa.func.sum(TRACKING.c.iade_adet)).where(TRACKING.c.plan_id == plan_id)
    if upto_week:
        q = q.where(TRACKING.c.hafta <= int(upto_week))
    out = {}
    for r in c.execute(q.group_by(TRACKING.c.logo_cari_kodu)):
        out[r[0]] = {"sevk": float(r[1] or 0), "fatura": float(r[2] or 0), "iade": float(r[3] or 0)}
    return out


def refresh_tracking(engine: sa.engine.Engine, tenant: str, S: Sources, row: Any) -> dict[str, Any]:
    """Onaylı planın Logo takibi: onay gününden itibaren takip haftaları, cari × hafta."""
    p = params()
    if not row.decided_at:
        return {"planId": row.id, "satir": 0}
    start = (row.decided_at.astimezone(TZ) if row.decided_at.tzinfo else row.decided_at).date()
    end = start + timedelta(days=7 * p["takipHafta"])
    rows = S.logo().tracking(row.stok_kodu, start, end)
    merged: dict[tuple[str, int], dict[str, float]] = {}
    for r in rows:
        if not r["cari_kodu"] or r["hafta"] < 1:
            continue
        cur = merged.setdefault((r["cari_kodu"], r["hafta"]), {"sevk": 0.0, "fatura": 0.0, "iade": 0.0})
        for k in ("sevk", "fatura", "iade"):
            cur[k] += r[k]
    now = _now()
    with engine.begin() as c:
        c.execute(TRACKING.delete().where(TRACKING.c.plan_id == row.id))
        if merged:
            c.execute(TRACKING.insert(), [dict(plan_id=row.id, logo_cari_kodu=k[:60], hafta=h, tenant_id=tenant,
                                               sevk_adet=v["sevk"], satis_adet=v["fatura"], iade_adet=v["iade"], asof=now)
                                          for (k, h), v in merged.items()])
    return {"planId": row.id, "satir": len(merged)}


def _elapsed_end(engine: sa.engine.Engine, tenant: str) -> date:
    """Değerlendirme günü: bugün ile Logo verisinin bittiği günden erken olanı (donmuş kopyada sahte uyarı olmasın)."""
    end = data_end(engine, tenant)
    return min(today(), end) if end else today()


def tracking_summary(engine: sa.engine.Engine, tenant: str, row: Any, bmt: Optional[str] = None) -> dict[str, Any]:
    p = params()
    with engine.connect() as c:
        lines = c.execute(_visible(sa.select(LINES).where(LINES.c.plan_id == row.id, LINES.c.adet > 0), bmt)).all()
        track = _track_by_client(c, row.id)
    plan_q = sum(ln.adet for ln in lines)
    keys = {ln.logo_cari_kodu for ln in lines}
    t_in = [v for k, v in track.items() if k in keys]
    sevk = sum(v["sevk"] for v in t_in)
    fatura = sum(v["fatura"] for v in t_in)
    iade = sum(v["iade"] for v in t_in)
    outside = {k: v for k, v in track.items() if k not in keys} if not bmt else {}
    start = (row.decided_at.astimezone(TZ) if row.decided_at and row.decided_at.tzinfo else row.decided_at)
    start_d = start.date() if start else None
    asof = _elapsed_end(engine, tenant)
    week = max(0, min(p["takipHafta"], ((asof - start_d).days // 7 + 1) if start_d and asof >= start_d else 0))
    return {"planId": row.id, "stokKodu": row.stok_kodu, "ad": row.ad, "surum": row.surum, "onay": _iso(row.decided_at),
            "hafta": week, "takipHafta": p["takipHafta"], "plan": plan_q, "sevk": sevk, "fatura": fatura, "iade": iade,
            "net": fatura - iade, "sevkOrani": (sevk / plan_q) if plan_q else None,
            "iadeOrani": (iade / fatura) if fatura else None, "musteri": len(lines),
            "planDisi": {"musteri": len(outside), "sevk": sum(v["sevk"] for v in outside.values())},
            "degerlendirmeGunu": asof.isoformat(), "veriBitti": bool(start_d and asof < start_d)}


def tracking(engine: sa.engine.Engine, tenant: str, code: str = "", upto_week: int = 0, bmt: Optional[str] = None) -> dict[str, Any]:
    """Bir kitabın (ya da izlenen bütün kitapların) takibi: özet, bölge, hafta ve cari ayrıntısı."""
    p = params()
    with engine.connect() as c:
        q = sa.select(PLANS).where(PLANS.c.tenant_id == tenant, PLANS.c.durum == "onayli")
        if code:
            q = q.where(PLANS.c.stok_kodu == code)
        plans = c.execute(q.order_by(PLANS.c.decided_at.desc())).all()
    out = []
    for row in plans:
        s = tracking_summary(engine, tenant, row, bmt)
        if code:
            with engine.connect() as c:
                lines = c.execute(_visible(sa.select(LINES).where(LINES.c.plan_id == row.id), bmt)).all()
                track = _track_by_client(c, row.id, upto_week or None)
                weeks = c.execute(sa.select(TRACKING.c.hafta, sa.func.sum(TRACKING.c.sevk_adet), sa.func.sum(TRACKING.c.satis_adet),
                                            sa.func.sum(TRACKING.c.iade_adet)).where(TRACKING.c.plan_id == row.id)
                                  .group_by(TRACKING.c.hafta).order_by(TRACKING.c.hafta)).all()
            regions: dict[str, dict[str, float]] = {}
            for ln in lines:
                t = track.get(ln.logo_cari_kodu) or {"sevk": 0.0, "fatura": 0.0, "iade": 0.0}
                rg = regions.setdefault(ln.bolge, {"plan": 0.0, "sevk": 0.0, "fatura": 0.0, "iade": 0.0})
                rg["plan"] += ln.adet
                for k in ("sevk", "fatura", "iade"):
                    rg[k] += t[k]
            s["bolgeler"] = [{"bolge": b, **v, "net": v["fatura"] - v["iade"]} for b, v in sorted(
                regions.items(), key=lambda x: (REGION_ORDER.index(x[0]) if x[0] in REGION_ORDER else 99, x[0]))]
            s["haftalar"] = [{"hafta": int(w[0]), "sevk": float(w[1] or 0), "fatura": float(w[2] or 0), "iade": float(w[3] or 0)}
                             for w in weeks]
            s["cariler"] = [{**_line_dict(ln), "takip": track.get(ln.logo_cari_kodu) or {"sevk": 0.0, "fatura": 0.0, "iade": 0.0}}
                            for ln in sorted(lines, key=lambda x: -x.adet) if ln.adet > 0]
        out.append(s)
        if code:
            break
    return {"items": out, "veriSonu": (lambda d: d.isoformat() if d else None)(data_end(engine, tenant)),
            "takipHafta": p["takipHafta"]}


def my_region(engine: sa.engine.Engine, tenant: str, user: Optional[str]) -> dict[str, Any]:
    """BMT görünümü: onaylı planlarda kişinin carilerine düşen kitaplar (user None = herkes)."""
    p = params()
    limit_day = today() - timedelta(days=7 * p["takipHafta"] + 7)
    with engine.connect() as c:
        plans = c.execute(sa.select(PLANS).where(PLANS.c.tenant_id == tenant, PLANS.c.durum == "onayli")
                          .order_by(PLANS.c.decided_at.desc())).all()
        items = []
        for row in plans:
            d = row.decided_at.astimezone(TZ).date() if row.decided_at and row.decided_at.tzinfo else (row.decided_at.date() if row.decided_at else None)
            if d and d < limit_day:
                continue
            lines = c.execute(_visible(sa.select(LINES).where(LINES.c.plan_id == row.id, LINES.c.adet > 0), user)
                              .order_by(LINES.c.adet.desc())).all()
            if not lines:
                continue
            track = _track_by_client(c, row.id)
            items.append({
                "planId": row.id, "stokKodu": row.stok_kodu, "ad": row.ad, "onay": _iso(row.decided_at),
                "depoGiris": row.depo_giris_tarihi, "adet": sum(ln.adet for ln in lines), "musteri": len(lines),
                "sevk": sum((track.get(ln.logo_cari_kodu) or {}).get("sevk", 0.0) for ln in lines),
                "cariler": [{"cariKodu": ln.logo_cari_kodu, "unvan": ln.cari_unvan, "il": ln.il, "adet": ln.adet,
                             "bmt": ln.bmt_ad, **(track.get(ln.logo_cari_kodu) or {"sevk": 0.0, "fatura": 0.0, "iade": 0.0})}
                            for ln in lines]})
    return {"items": items, "kapsam": "kendi" if user else "hepsi",
            "veriSonu": (lambda d: d.isoformat() if d else None)(data_end(engine, tenant))}


def refresh_owners(engine: sa.engine.Engine, S: Sources, plan_ids: list[str]) -> int:
    """Onaylı planlardaki carilerin BMT'sini CRM'deki güncel sahiple tazeler (sahiplik değişirse telefon görünümü doğru kalsın)."""
    if not plan_ids:
        return 0
    by_code = {a["cari_kodu"]: a for a in S.accounts() if a.get("cari_kodu")}
    n = 0
    with engine.begin() as c:
        for ln in c.execute(sa.select(LINES.c.plan_id, LINES.c.line_no, LINES.c.logo_cari_kodu, LINES.c.bmt_hesap)
                            .where(LINES.c.plan_id.in_(plan_ids))).all():
            a = by_code.get(ln.logo_cari_kodu or "")
            if a and (a.get("bmt_hesap") or None) != (ln.bmt_hesap or None):
                c.execute(LINES.update().where(LINES.c.plan_id == ln.plan_id, LINES.c.line_no == ln.line_no).values(
                    bmt_hesap=a.get("bmt_hesap"), bmt_ad=a.get("bmt_ad")))
                n += 1
    return n


# ------------------------------------------------------------------ uyarılar


def _open_alert(c: Any, tenant: str, now: datetime, *, tur: str, anahtar: str, etiket: str, plan_id: Optional[str],
                stok: Optional[str], detay: dict[str, Any], bmt: Optional[str] = None, bilgi: bool = False) -> bool:
    cur = c.execute(sa.select(ALERTS).where(ALERTS.c.tenant_id == tenant, ALERTS.c.tur == tur, ALERTS.c.anahtar == anahtar,
                                           ALERTS.c.durum.in_(("acik", "bilgi")))).first()
    if cur:
        c.execute(ALERTS.update().where(ALERTS.c.id == cur.id).values(son_zaman=now, etiket=etiket[:400], detay_json=_dump(detay)))
        return False
    c.execute(ALERTS.insert().values(id=uuid.uuid4().hex, tenant_id=tenant, plan_id=plan_id, stok_kodu=stok, tur=tur,
                                     anahtar=anahtar[:200], etiket=etiket[:400], detay_json=_dump(detay), bmt_hesap=bmt,
                                     durum="bilgi" if bilgi else "acik", ilk_zaman=now, son_zaman=now))
    return True


def evaluate_alerts(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    """Uyarıları açar ve koşulu kalkanı kapatır. Süreler Logo verisinin bittiği güne göre."""
    p = params()
    now = _now()
    asof = _elapsed_end(engine, tenant)
    want: set[tuple[str, str]] = set()
    opened = 0
    with engine.begin() as c:
        books = c.execute(sa.select(BOOKS).where(BOOKS.c.tenant_id == tenant)).all()
        plans = _latest_plans(c, tenant)
        for b in books:
            d = _day(b.depo_giris_tarihi)
            if b.stok_kodu in plans or not d or d >= today():
                continue
            key = f"plan_yok:{b.stok_kodu}"
            want.add(("plan_yok", key))
            opened += _open_alert(c, tenant, now, tur="plan_yok", anahtar=key, plan_id=None, stok=b.stok_kodu,
                                  etiket=f"{b.ad or b.stok_kodu} {d.strftime('%d.%m.%Y')} depoya girdi, dağılım planı yok",
                                  detay={"depoGiris": b.depo_giris_tarihi, "baskiAdedi": b.baski_adedi})
        for row in (r for r in plans.values() if r.durum == "onayli"):
            start = (row.decided_at.astimezone(TZ) if row.decided_at.tzinfo else row.decided_at).date() if row.decided_at else None
            if not start or asof < start:
                continue
            lines = c.execute(sa.select(LINES).where(LINES.c.plan_id == row.id, LINES.c.adet > 0)).all()
            track = _track_by_client(c, row.id)
            if business_days_between(start, asof) >= p["sevkGun"]:
                for ln in lines:
                    t = track.get(ln.logo_cari_kodu) or {}
                    if ln.logo_cari_kodu and t.get("sevk", 0) <= 0:
                        key = f"sevk_gecikti:{row.id}:{ln.logo_cari_kodu}"
                        want.add(("sevk_gecikti", key))
                        opened += _open_alert(c, tenant, now, tur="sevk_gecikti", anahtar=key, plan_id=row.id, stok=row.stok_kodu,
                                              etiket=f"{row.ad or row.stok_kodu}: {ln.cari_unvan} ({_tr(ln.adet)} adet) {p['sevkGun']} iş gününde sevk edilmedi",
                                              detay={"cari": ln.logo_cari_kodu, "adet": ln.adet, "onay": start.isoformat()},
                                              bmt=ln.bmt_hesap)
            if (asof - start).days >= p["satisGun"]:
                regions: dict[str, dict[str, float]] = {}
                for ln in lines:
                    t = track.get(ln.logo_cari_kodu) or {}
                    rg = regions.setdefault(ln.bolge, {"sevk": 0.0, "net": 0.0})
                    rg["sevk"] += t.get("sevk", 0.0)
                    rg["net"] += t.get("fatura", 0.0) - t.get("iade", 0.0)
                for reg, v in regions.items():
                    if v["sevk"] > 0 and v["net"] <= 0:
                        key = f"hic_satmadi:{row.id}:{reg}"
                        want.add(("hic_satmadi", key))
                        opened += _open_alert(c, tenant, now, tur="hic_satmadi", anahtar=key, plan_id=row.id, stok=row.stok_kodu,
                                              etiket=f"{row.ad or row.stok_kodu}: {reg} bölgesine {_tr(v['sevk'])} adet sevk edildi, {p['satisGun']} günde faturalanan net satış yok",
                                              detay={"bolge": reg, **v})
            for ln in lines:
                t = track.get(ln.logo_cari_kodu) or {}
                if t.get("sevk", 0) > 0 and t.get("fatura", 0) - t.get("iade", 0) >= p["tukenmePay"] * t["sevk"]:
                    key = f"tukendi:{row.id}:{ln.logo_cari_kodu}"
                    want.add(("tukendi", key))
                    opened += _open_alert(c, tenant, now, tur="tukendi", anahtar=key, plan_id=row.id, stok=row.stok_kodu,
                                          etiket=f"{row.ad or row.stok_kodu}: {ln.cari_unvan} sevk edilenin {_pct(p['tukenmePay'])}'ından fazlasını faturaladı",
                                          detay={"cari": ln.logo_cari_kodu, **t}, bmt=ln.bmt_hesap, bilgi=True)
        closed = 0
        for a in c.execute(sa.select(ALERTS).where(ALERTS.c.tenant_id == tenant, ALERTS.c.durum.in_(("acik", "bilgi")))).all():
            if (a.tur, a.anahtar) not in want:
                c.execute(ALERTS.update().where(ALERTS.c.id == a.id).values(durum="kapandi", kapanis=now))
                closed += 1
    return {"acilan": opened, "kapanan": closed, "degerlendirmeGunu": asof.isoformat()}


def _alert_dict(a: Any) -> dict[str, Any]:
    return {"id": a.id, "planId": a.plan_id, "stokKodu": a.stok_kodu, "tur": a.tur, "turEtiket": ALERT_KINDS.get(a.tur, a.tur),
            "anahtar": a.anahtar, "etiket": a.etiket, "detay": _j(a.detay_json, {}), "durum": a.durum,
            "ilk": _iso(a.ilk_zaman), "son": _iso(a.son_zaman), "kapanis": _iso(a.kapanis), "bildirim": _iso(a.bildirim)}


def alerts(engine: sa.engine.Engine, tenant: str, *, durum: str = "acik", tur: str = "", page: int = 0,
           bmt: Optional[str] = None) -> dict[str, Any]:
    with engine.connect() as c:
        cond = [ALERTS.c.tenant_id == tenant]
        if durum == "acik":
            cond.append(ALERTS.c.durum.in_(("acik", "bilgi")))
        elif durum:
            cond.append(ALERTS.c.durum == durum)
        if tur:
            cond.append(ALERTS.c.tur == tur)
        if bmt:
            cond.append(sa.func.lower(ALERTS.c.bmt_hesap) == bmt.lower())
        total = c.execute(sa.select(sa.func.count()).select_from(ALERTS).where(*cond)).scalar()
        rows = c.execute(sa.select(ALERTS).where(*cond).order_by(ALERTS.c.son_zaman.desc())
                         .offset(max(0, page) * PAGE_SIZE).limit(PAGE_SIZE)).all()
        counts = {t: n for t, n in c.execute(sa.select(ALERTS.c.tur, sa.func.count()).where(
            ALERTS.c.tenant_id == tenant, ALERTS.c.durum.in_(("acik", "bilgi"))).group_by(ALERTS.c.tur))}
    return {"items": [_alert_dict(a) for a in rows], "total": int(total or 0), "page": page, "pageSize": PAGE_SIZE,
            "sayilar": counts, "turler": ALERT_KINDS}


def notify(engine: sa.engine.Engine, tenant: str, recipients: list[str], link: str,
           send: Callable[[str, str, list[str]], str]) -> dict[str, Any]:
    """Bildirilmemiş açık uyarılar tek özet e-postayla. SMTP/alıcı yoksa uyarı kayıtta bekler."""
    with engine.connect() as c:
        rows = c.execute(sa.select(ALERTS).where(ALERTS.c.tenant_id == tenant, ALERTS.c.durum == "acik",
                                                 ALERTS.c.bildirim.is_(None)).order_by(ALERTS.c.tur, ALERTS.c.ilk_zaman)).all()
    if not rows:
        return {"status": "nothing"}
    if not recipients:
        return {"status": "no_recipient", "bekleyen": len(rows)}
    by: dict[str, list[str]] = {}
    for a in rows:
        by.setdefault(ALERT_KINDS.get(a.tur, a.tur), []).append(a.etiket or a.anahtar)
    text = "İlk dağılım uyarıları\n\n" + "\n\n".join(f"{k} ({len(v)})\n" + "\n".join(f"- {x}" for x in v) for k, v in by.items())
    if link:
        text += f"\n\nAyrıntı: {link}"
    status = send(f"İlk dağılım: {len(rows)} yeni uyarı", text, recipients)
    if status == "sent":
        with engine.begin() as c:
            c.execute(ALERTS.update().where(ALERTS.c.id.in_([a.id for a in rows])).values(bildirim=_now()))
    return {"status": status, "adet": len(rows)}
