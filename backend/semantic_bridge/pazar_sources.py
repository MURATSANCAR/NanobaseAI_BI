"""M39 Pazar araştırması ve rekabet: kaynakların okunması (yalnız okuma — CRM'e ve Logo'ya yazılmaz).

**CRM (`Timas_MSCRM`, .28 prod):**

- Rakip katalog = `new_rakipkitapBase`, `statecode = 0` (2026-09-09 kataloğunda 30.415 satır). Kaynağı, toplanma
  yöntemi ve «Satış Adedi / Satış Adedi 2»nin anlamı bilinmiyor (analiz §10 soru 1, **ölçülecek**): iki alan ham
  olarak kopyalanır, hiçbir hesapta kullanılmaz, ekranda «anlamı bilinmiyor» notuyla görünür.
- Tanıtım metni üçüncü kişinin eseridir (analiz §8): toplu okumada yalnız ilk `PAZAR_BLURB_CHARS` karakteri
  alınır (SQL tarafında `LEFT`), cümle sınırında kısaltılır ve yalnız emsal benzerliğinde kullanılır.
- Emsal bağı = `new_new_kitap_new_rakipkitapBase` (TİMAŞ kitabı ↔ rakip kitap; katalogda 214).
- TİMAŞ kitapları = `new_kitapBase`, `statecode = 0 AND new_Tip = 1` (Tip 1 = kitap; set, dergi, e-kitap hariç):
  KDV dahil fiyat, sayfa, ciltleme şekli, ebat, Kitaplık, yayınevi (marka), web kategorisi, arka kapağın ilk
  karakterleri (emsal için).
- Kitaplık listesi = `new_kitaplikBase` (kategori kaynağı «kitaplık» iken).

**Logo:** TİMAŞ'ın kendi satışı — satır tanımı M46 ile aynı (`budget_sources`): `STLINE`, `CANCELLED = 0`,
`LINETYPE = 0`, `INVOICEREF <> 0` (faturalı), `TRCODE 7/8/9` satış, `2/3` iade (eksi); ciro = `LINENET`. Yayınevi =
`ITEMS.SPECODE` (kayıtlı SQL «yayınevi bazında net ciro»), kanal = `CLCARD.SPECODE2`. Yıl → firma eşlemesi
`L_CAPIPERIOD`'dan. Bu satış **sell-in**'dir (kitapçıya ve dağıtıcıya satış), okura satış ya da pazar payı değildir;
ekran bunu her yerde yazar.

**Yüklenen dosya:** PDF'in metin katmanı sayfa sayfa (`pypdf`), Excel'de her çalışma sayfası, CSV/metin tek parça.
Taranmış (metinsiz) PDF sayfası yüklemede «metin yok» diye sayılır; rakam çıkarımında ortak belge okuma hattıyla
(`doc_read`, OCR) okunur (`read_scanned`), rakam yine OCR metninde birebir aranır ve ekranda «OCR» etiketiyle görünür.
"""
from __future__ import annotations

import csv
import io
import logging
import re
from datetime import date, datetime, timedelta
from typing import Any, Callable, Optional

from semantic_bridge import budget_sources as bsrc
from semantic_bridge.editorial import _prefix

log = logging.getLogger("semantic.pazar.sources")

SourceError = bsrc.SourceError
Runner = Callable[[str], list[dict[str, Any]]]
runner = bsrc.runner
firms_by_year = bsrc.firms_by_year
read_data_end = bsrc.read_data_end

#: CRM `new_ciltlemesekli` seçenekleri (table_descriptions.json, 2026-09-09).
CILT = {"1": "Amerikan Cilt", "2": "Ciltli", "3": "Fleksi Cilt", "4": "Bez Cilt"}

_TAG = re.compile(r"<[^>]+>")
_WS = re.compile(r"\s+")


def clean(v: Any) -> Optional[str]:
    if v is None:
        return None
    import html as _html

    t = _WS.sub(" ", _html.unescape(_TAG.sub(" ", str(v)))).strip()
    return t or None


def short(v: Any, n: int) -> Optional[str]:
    """Metnin başı, en çok `n` karakter, mümkünse cümle sonunda kesilir (üçüncü kişinin metni kopyalanmaz)."""
    t = clean(v)
    if not t or n <= 0:
        return None
    if len(t) <= n:
        return t
    cut = t[:n]
    end = max(cut.rfind(". "), cut.rfind("! "), cut.rfind("? "))
    return (cut[: end + 1] if end >= n // 3 else cut.rstrip() + "…").strip()


def num(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    return n if n == n else None


def day(v: Any) -> Optional[str]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.replace(tzinfo=None).isoformat(timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    s = str(v)
    return s[:19] if s[:4].isdigit() else None


def guid(v: Any) -> Optional[str]:
    s = str(v or "").strip().strip("{}").upper()
    return s or None


# ------------------------------------------------------------------ CRM SQL


def competitors_sql(schema: str, blurb_chars: int) -> str:
    p = _prefix(schema)
    n = max(0, int(blurb_chars)) * 2 or 1   # cümle sınırına pay; Python tarafında `short` ile kısalır
    return (
        "SELECT r.new_rakipkitapId AS id, r.new_name AS ad, r.new_Yaynevi AS yayinevi, r.new_Yazarlar AS yazarlar,"
        " r.new_Isnb AS isbn, r.new_ListeFiyat AS fiyat, r.new_SayfaSays AS sayfa, r.new_CiltTipi AS cilt,"
        " r.new_KagitBoyutu AS kagit, r.new_BaskSays AS baski, r.new_Kategoriler AS kategori, r.new_Dil AS dil,"
        " r.new_SatisAdedi AS satis1, r.new_SatAdedi2 AS satis2, r.new_SatisDurumu AS durum,"
        f" LEFT(CAST(r.new_TanitimMetni AS nvarchar(max)), {n}) AS tanitim,"
        " r.CreatedOn AS olusturma, r.ModifiedOn AS degisme"
        f" FROM {p}new_rakipkitapBase r WHERE r.statecode = 0"
    )


def links_sql(schema: str) -> str:
    p = _prefix(schema)
    return f"SELECT l.new_kitapid AS kitap, l.new_rakipkitapid AS rakip FROM {p}new_new_kitap_new_rakipkitapBase l"


def own_books_sql(schema: str, blurb_chars: int) -> str:
    p = _prefix(schema)
    n = max(0, int(blurb_chars)) * 2 or 1
    return (
        "SELECT k.new_kitapId AS id, k.new_StokKodu AS stok, k.new_name AS ad, k.new_yazartext AS yazar,"
        " k.new_isbn13 AS isbn, k.new_yayineviid AS marka_id, m.new_name AS marka,"
        " k.new_kitaplikid AS kitaplik_id, kl.new_name AS kitaplik, k.new_kdvdahilfiyat AS fiyat,"
        " k.new_sayfasayisi AS sayfa, k.new_ciltlemesekli AS cilt, k.new_Ebat AS ebat,"
        " k.new_webkategorileritext AS web, k.new_ilkyayintarihi AS ilk_yayin,"
        f" LEFT(CAST(k.new_ozet AS nvarchar(max)), {n}) AS ozet"
        f" FROM {p}new_kitapBase k"
        f" LEFT JOIN {p}new_kitaplikBase kl ON kl.new_kitaplikId = k.new_kitaplikid"
        f" LEFT JOIN {p}new_markaBase m ON m.new_markaId = k.new_yayineviid"
        " WHERE k.statecode = 0 AND k.new_Tip = 1"
    )


def kitaplik_sql(schema: str) -> str:
    p = _prefix(schema)
    return f"SELECT new_kitaplikId AS id, new_name AS ad FROM {p}new_kitaplikBase WHERE statecode = 0"


# ------------------------------------------------------------------ CRM okuma


def read_competitors(schema: str, run: Runner, blurb_chars: int) -> list[dict[str, Any]]:
    out = []
    for r in run(competitors_sql(schema, blurb_chars)):
        rid = guid(r.get("id"))
        if not rid:
            continue
        out.append({
            "crm_id": rid, "ad": (clean(r.get("ad")) or "")[:400], "yayinevi": (clean(r.get("yayinevi")) or "")[:200] or None,
            "yazarlar": (clean(r.get("yazarlar")) or "")[:400] or None, "isbn": (clean(r.get("isbn")) or "")[:40] or None,
            "liste_fiyat": num(r.get("fiyat")), "sayfa": _int(r.get("sayfa")), "cilt": (clean(r.get("cilt")) or "")[:100] or None,
            "kagit": (clean(r.get("kagit")) or "")[:100] or None, "baski_sayisi": _int(r.get("baski")),
            "kategori_ham": norm_category(r.get("kategori")), "dil": (clean(r.get("dil")) or "")[:60] or None,
            "satis_adedi_ham": _int(r.get("satis1")), "satis_adedi2_ham": (clean(r.get("satis2")) or "")[:60] or None,
            "satis_durumu": (clean(r.get("durum")) or "")[:100] or None,
            "tanitim_kisa": short(r.get("tanitim"), blurb_chars),
            "crm_created": day(r.get("olusturma")), "crm_modified": day(r.get("degisme")),
        })
    return out


def read_own_books(schema: str, run: Runner, blurb_chars: int) -> list[dict[str, Any]]:
    out = []
    for r in run(own_books_sql(schema, blurb_chars)):
        bid = guid(r.get("id"))
        if not bid:
            continue
        cilt = r.get("cilt")
        ilk = day(r.get("ilk_yayin"))
        if ilk and ilk[:4] < "1950":        # boş tarih (1899/1900) tarih sayılmaz
            ilk = None
        out.append({
            "crm_id": bid, "stok_kodu": (clean(r.get("stok")) or "")[:60] or None, "ad": (clean(r.get("ad")) or "")[:400],
            "yazar": (clean(r.get("yazar")) or "")[:400] or None, "isbn": (clean(r.get("isbn")) or "")[:40] or None,
            "marka_id": guid(r.get("marka_id")), "marka": (clean(r.get("marka")) or "")[:200] or None,
            "kitaplik_id": guid(r.get("kitaplik_id")), "kitaplik": (clean(r.get("kitaplik")) or "")[:200] or None,
            "fiyat": num(r.get("fiyat")), "sayfa": _int(r.get("sayfa")),
            "cilt": CILT.get(str(int(cilt))) if num(cilt) is not None else None,
            "ebat": (clean(r.get("ebat")) or "")[:60] or None, "web": (clean(r.get("web")) or "")[:300] or None,
            "ilk_yayin": ilk[:10] if ilk else None, "ozet_kisa": short(r.get("ozet"), blurb_chars),
        })
    return out


def read_links(schema: str, run: Runner) -> list[tuple[str, str]]:
    out = []
    for r in run(links_sql(schema)):
        k, rk = guid(r.get("kitap")), guid(r.get("rakip"))
        if k and rk:
            out.append((k, rk))
    return out


def read_kitaplik(schema: str, run: Runner) -> list[dict[str, Any]]:
    return [{"id": guid(r["id"]), "ad": clean(r.get("ad")) or "—"} for r in run(kitaplik_sql(schema)) if guid(r.get("id"))]


def norm_category(v: Any) -> Optional[str]:
    """Rakip kategori metni eşleme birimi: HTML/boşluk temizliği, en çok 300 karakter. Ayırıcı biçimi (virgül, «>»,
    noktalı virgül) ölçülmedi; metin bütün olarak eşlenir, bölünmez."""
    t = clean(v)
    return t[:300] if t else None


def _int(v: Any) -> Optional[int]:
    n = num(v)
    return int(n) if n is not None else None


# ------------------------------------------------------------------ Logo SQL (TİMAŞ'ın kendi satışı)

_SALE = "S.CANCELLED = 0 AND S.LINETYPE = 0 AND S.INVOICEREF <> 0 AND S.TRCODE IN (2,3,7,8,9)"
_SIGN = "(CASE WHEN S.TRCODE IN (7,8,9) THEN 1 ELSE -1 END)"


def cut_day(year: int, end: date) -> date:
    """Yıl içindeki aynı dönemin (1 Ocak – veri sonu) bitişinin ertesi günü. 29 Şubat 28'e düşer."""
    try:
        d = date(year, end.month, end.day)
    except ValueError:
        d = date(year, end.month, 28)
    return d + timedelta(days=1)


def item_sales_sql(firm: str, year: int, cut: date) -> str:
    return f"""
-- Faturalı satış satırları; iade eksi. ytd_* = 1 Ocak – veri sonu (aynı dönem karşılaştırması için).
SELECT I.CODE AS stok, MAX(I.SPECODE) AS yayinevi,
  SUM(CASE WHEN S.DATE_ < '{cut.isoformat()}' THEN {_SIGN} * S.AMOUNT ELSE 0 END) AS ytd_adet,
  SUM(CASE WHEN S.DATE_ < '{cut.isoformat()}' THEN {_SIGN} * S.LINENET ELSE 0 END) AS ytd_ciro,
  SUM({_SIGN} * S.AMOUNT) AS adet, SUM({_SIGN} * S.LINENET) AS ciro
FROM dbo.LG_{firm}_01_STLINE AS S
JOIN dbo.LG_{firm}_ITEMS AS I ON I.LOGICALREF = S.STOCKREF
WHERE {_SALE} AND S.DATE_ >= '{year}-01-01' AND S.DATE_ < '{year + 1}-01-01'
GROUP BY I.CODE""".strip()


def channel_sales_sql(firm: str, year: int, cut: date) -> str:
    return f"""
-- Kanal = cari kartın özel kodu 2 (CLCARD.SPECODE2); boş kod «(boş)».
SELECT COALESCE(NULLIF(C.SPECODE2, ''), '(boş)') AS kanal,
  SUM(CASE WHEN S.DATE_ < '{cut.isoformat()}' THEN {_SIGN} * S.AMOUNT ELSE 0 END) AS ytd_adet,
  SUM(CASE WHEN S.DATE_ < '{cut.isoformat()}' THEN {_SIGN} * S.LINENET ELSE 0 END) AS ytd_ciro,
  SUM({_SIGN} * S.AMOUNT) AS adet, SUM({_SIGN} * S.LINENET) AS ciro
FROM dbo.LG_{firm}_01_STLINE AS S
LEFT JOIN dbo.LG_{firm}_CLCARD AS C ON C.LOGICALREF = S.CLIENTREF
WHERE {_SALE} AND S.DATE_ >= '{year}-01-01' AND S.DATE_ < '{year + 1}-01-01'
GROUP BY COALESCE(NULLIF(C.SPECODE2, ''), '(boş)')""".strip()


def read_own_sales(run: Runner, years: int) -> dict[str, Any]:
    """Son `years` yılın (veri sonunun yılı dahil) stok kodu ve kanal kırılımında satışı. Yılı Logo'da dönemi
    olmayan yıl atlanır ve `missing`'e yazılır (sessiz değil)."""
    firms = firms_by_year(run)
    if not firms:
        raise SourceError("Logo'da dönem bulunamadı.")
    end = read_data_end(run, firms)
    if end is None:
        raise SourceError("Logo'da satış satırı bulunamadı.")
    rows: list[dict[str, Any]] = []
    missing: list[int] = []
    for y in range(end.year - max(1, years) + 1, end.year + 1):
        firm = firms.get(y)
        if not firm:
            missing.append(y)
            continue
        cut = cut_day(y, end)
        for r in run(item_sales_sql(firm, y, cut)):
            code = str(r.get("stok") or "").strip()
            if not code:
                continue
            rows.append({"yil": y, "boyut": "stok", "anahtar": code[:120], "yayinevi": (clean(r.get("yayinevi")) or "(boş)")[:120],
                         "ytd_adet": num(r.get("ytd_adet")) or 0.0, "ytd_ciro": num(r.get("ytd_ciro")) or 0.0,
                         "adet": num(r.get("adet")) or 0.0, "ciro": num(r.get("ciro")) or 0.0})
        for r in run(channel_sales_sql(firm, y, cut)):
            rows.append({"yil": y, "boyut": "kanal", "anahtar": (clean(r.get("kanal")) or "(boş)")[:120], "yayinevi": None,
                         "ytd_adet": num(r.get("ytd_adet")) or 0.0, "ytd_ciro": num(r.get("ytd_ciro")) or 0.0,
                         "adet": num(r.get("adet")) or 0.0, "ciro": num(r.get("ciro")) or 0.0})
    return {"rows": rows, "dataEnd": end.isoformat(), "missingYears": missing,
            "firms": {str(y): f for y, f in firms.items()}}


# ------------------------------------------------------------------ yüklenen dosyanın metni

ALLOWED = {"pdf": "application/pdf", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
           "csv": "text/csv", "txt": "text/plain"}
MAGIC = {"pdf": (b"%PDF",), "xlsx": (b"PK",)}


def ext_of(name: str) -> str:
    return (name.rsplit(".", 1)[-1] if "." in (name or "") else "").lower()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_scanned(filename: str, data: bytes, pages: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Metni olmayan PDF sayfalarını ortak belge okuma hattıyla (OCR) doldurur. Dönen: (sayfalar, ilerleme ekleri)
    — `ocrSayfa` (OCR'la okunan sayfalar), `ocrGuven` (sayfa → güven), `okumaHatasi` (servis yoksa ya da düştüyse).
    Okunamayan sayfa boş metinle kalır ve çıkarımda «metinsiz» sayılır; sessizce atlanmaz."""
    from semantic_bridge import doc_read as DR

    try:
        reading = DR.read(filename, data, allowed=("pdf",))
    except DR.ReadError as e:
        return pages, {"okumaHatasi": str(e)}
    by = {p["sayfa"]: p for p in reading.pages}
    out = []
    for p in pages:
        r = by.get(p["sayfa"])
        if not (p["metin"] or "").strip() and r and r["okuma"] == DR.OCR:
            out.append({**p, "metin": r["metin"], "okuma": DR.OCR, "guven": r.get("guven")})
        else:
            out.append(p)
    extra: dict[str, Any] = {"ocrSayfa": reading.ocr_pages,
                             "ocrGuven": {p["sayfa"]: p.get("guven") for p in reading.pages if p["okuma"] == DR.OCR}}
    if reading.errors:
        extra["okumaHatasi"] = " ".join(reading.errors)
    return out, extra


def pages_of(filename: str, data: bytes) -> list[dict[str, Any]]:
    """Dosyanın sayfaları: `[{"sayfa": "12", "metin": "…"}]`. PDF'te sayfa numarası (1'den), Excel'de «Sayfa adı»,
    CSV/metinde «1». Metni olmayan PDF sayfası boş metinle döner (taranmış sayfa)."""
    ext = ext_of(filename)
    if ext == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError:  # pragma: no cover
            raise SourceError("PDF okuyucu bu kurulumda yok.") from None
        try:
            reader = PdfReader(io.BytesIO(data))
            return [{"sayfa": str(i + 1), "metin": (p.extract_text() or "").strip()} for i, p in enumerate(reader.pages)]
        except Exception as e:  # noqa: BLE001
            raise SourceError(f"PDF okunamadı: {str(e)[:160]}") from None
    if ext == "xlsx":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise SourceError(f"Excel okunamadı: {str(e)[:160]}") from None
        out = []
        for ws in wb.worksheets:
            lines = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(x.strip() for x in cells):
                    lines.append("\t".join(cells).rstrip())
            out.append({"sayfa": f"Excel: {ws.title}"[:60], "metin": "\n".join(lines)})
        return out
    if ext == "csv":
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
            text = "\n".join("\t".join(r) for r in csv.reader(io.StringIO(text), dialect))
        except csv.Error:
            pass
        return [{"sayfa": "1", "metin": text}]
    if ext == "txt":
        return [{"sayfa": "1", "metin": _decode(data)}]
    raise SourceError("Rapor PDF, Excel (.xlsx), CSV ya da metin olmalı.")
