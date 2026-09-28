"""M20 yansıma raporunun dışa aktarımı: PDF (yönetime) ve Excel (liste). Sayılar `pr.report`tan; erişim/tiraj yok.

PDF, projedeki sunucu tarafı PDF yoluyla aynı yazı tipi çözümünü kullanır (`editorial_export._resolve_fonts`, fpdf2).
Metinde teknoloji adı geçmez; ürün adı «Zeki AI».
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any, Optional

from semantic_bridge import pr as PR

PRODUCT = "Zeki AI"


def _d(v: Optional[str]) -> str:
    if not v:
        return "—"
    x = date.fromisoformat(v[:10])
    return f"{x.day:02d}.{x.month:02d}.{x.year}"


def _pct(v: Optional[float]) -> str:
    return "—" if v is None else f"%{v * 100:.0f}"


def report_pdf(rep: dict[str, Any], comment: Optional[str], user: str = "") -> bytes:
    try:
        from fpdf import FPDF
    except ImportError as e:  # pragma: no cover
        raise PR.PrError("PDF üretici sunucuda kurulu değil.", 503) from e
    from semantic_bridge.editorial_export import _fold, _resolve_fonts

    regular, bold = _resolve_fonts()
    fam = "Body" if regular else "Helvetica"
    T = (lambda s: str(s or "")) if regular else (lambda s: _fold(str(s or "")))
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_title(T(f"Basın yansıma raporu {_d(rep['from'])} – {_d(rep['to'])}"))
    pdf.set_author(PRODUCT)
    pdf.set_creator(PRODUCT)
    pdf.set_margins(14, 14, 14)
    pdf.set_auto_page_break(auto=True, margin=14)
    if regular:
        pdf.add_font(fam, "", str(regular))
        pdf.add_font(fam, "B", str(bold or regular))
    pdf.add_page()
    W = pdf.w - pdf.l_margin - pdf.r_margin

    def para(s: str, size: float = 9.5, style: str = "", gap: float = 1.5) -> None:
        pdf.set_font(fam, style, size)
        pdf.multi_cell(W, size * 0.5, T(s), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    def head(s: str) -> None:
        pdf.ln(2)
        para(s, 12, "B", 1)

    def row(cells: list[tuple[str, float]], style: str = "") -> None:
        pdf.set_font(fam, style, 8.5)
        for txt, w in cells:
            t = T(txt)
            while pdf.get_string_width(t) > w - 1.5 and len(t) > 1:
                t = t[:-2] + "…" if len(t) > 2 else t[:-1]
            pdf.cell(w, 5.5, t, border="B")
        pdf.ln(5.5)

    s, c = rep["sends"], rep["coverage"]
    para("Timaş Yayınları · Basın ve halkla ilişkiler", 8.5)
    para(f"Yansıma raporu: {_d(rep['from'])} – {_d(rep['to'])}", 15, "B", 1)
    if comment:
        head("Dönemin özeti (Zeki AI yorumu)")
        para(comment)
    head("Gönderimler")
    para(f"Gönderim: {s['total']} · dönüş (cevap ya da haber): {s['answered']} · dönüş oranı: {_pct(s['answerRate'])}")
    for k, v in sorted(s["byStatus"].items(), key=lambda x: -x[1]):
        para(f"{PR.SEND_STATUSES.get(k, k)}: {v}", 9, gap=0.3)
    head("Yansımalar")
    para(f"Kayıtlı yansıma: {c['total']} · CRM haber arşivinde bu dönemde: {rep['archive']['total']} · onay bekleyen aday: "
         f"{rep['pendingCandidates']}")
    para("Ton: " + ", ".join(f"{PR.TONES.get(k, 'belirsiz')} {v}" for k, v in c["byTone"].items()) if c["byTone"] else "Ton: —", 9)
    para("Mecra türü: " + ", ".join(f"{PR.OUTLET_TYPES.get(k, 'belirsiz')} {v}" for k, v in c["byOutletType"].items())
         if c["byOutletType"] else "Mecra türü: —", 9)
    if c["books"]:
        head("En çok haber alan kitaplar")
        row([("Kitap", 150), ("Yansıma", 32)], "B")
        for b in c["books"]:
            row([(b["title"] or "—", 150), (str(b["count"]), 32)])
    if c["outlets"]:
        head("Mecralar")
        row([("Mecra", 150), ("Yansıma", 32)], "B")
        for o in c["outlets"]:
            row([(o["name"], 150), (str(o["count"]), 32)])
    if c["items"]:
        head("Yansıma listesi")
        row([("Tarih", 22), ("Mecra", 42), ("Başlık", 90), ("Ton", 28)], "B")
        for r in c["items"]:
            row([(_d(r["publishedAt"]), 22), (r["outlet"] or "—", 42), (r["title"], 90), (r["toneLabel"] or "—", 28)])
    pdf.ln(3)
    para("Erişim ve tiraj rakamı yer almaz: bu bilgi hiçbir kaynakta tutulmuyor. Yansıma sayıları portalda kayıtlı "
         "(elle girilen ve kabul edilen tarama adayları) kayıtlardır; CRM haber arşivi ayrı sayılır.", 8)
    para(f"Hazırlayan: {user or '—'} · {PRODUCT}", 8)
    return bytes(pdf.output())


def report_xlsx(rep: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    s, c = rep["sends"], rep["coverage"]
    rows = [("Dönem", f"{rep['from']} – {rep['to']}"), ("Gönderim", s["total"]), ("Dönüş (cevap ya da haber)", s["answered"]),
            ("Dönüş oranı", s["answerRate"]), ("Kayıtlı yansıma", c["total"]), ("CRM haber arşivi (dönem)", rep["archive"]["total"]),
            ("Onay bekleyen aday", rep["pendingCandidates"])]
    rows += [(f"Durum: {PR.SEND_STATUSES.get(k, k)}", v) for k, v in s["byStatus"].items()]
    rows += [(f"Ton: {PR.TONES.get(k, 'belirsiz')}", v) for k, v in c["byTone"].items()]
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 34
    ws2 = wb.create_sheet("Yansımalar")
    head = ["Tarih", "Mecra", "Mecra türü", "Başlık", "Kitap", "Yazar", "Ton", "Kaynak", "Bağlantı"]
    ws2.append(head)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for r in c["items"]:
        ws2.append([r["publishedAt"] or "", r["outlet"] or "", PR.OUTLET_TYPES.get(r["outletType"] or "", ""), r["title"],
                    r["bookTitle"] or "", r["authorName"] or "", r["toneLabel"] or "", r["sourceLabel"], r["url"] or ""])
    ws3 = wb.create_sheet("Kitaplar")
    ws3.append(["Kitap", "Yansıma"])
    for b in c["books"]:
        ws3.append([b["title"] or "", b["count"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
