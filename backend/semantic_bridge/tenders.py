"""M33 Okul, kütüphane ve kamu ihale takibi: ihale kaydı, şartname kalemlerinin katalogla eşleştirilmesi, teklif fiyat
tablosu, belge kontrol listesi ve şirket belge arşivi, karar/onay akışı (iki göz), sonuç ve kamu satış özeti.

**İlk sürüm kapsamı:** ilan kaynağı yok ve müşteride dış tarama kapalı; ilan elle girilir, şartname dosyası yüklenir.
Resmî kaynaktan ilan içe alma ikinci sürümdür ve yalnız `TENDER_WATCH_ENABLED=1` ile açılır (varsayılan kapalı).
Portal hiçbir kuruma teklif göndermez, e-imza kullanmaz; CRM'e ve Logo'ya yazmaz. Her yazma `semantic_audit`'e düşer.

**Kalem–katalog eşleştirme sırası:** (1) ISBN/barkod birebir (ISBN-10 → ISBN-13 çevrilir, sağlama basamağı denetlenir);
(2) ad (Türkçe harf ve noktalama sadeleşmiş) birebir, birden çok kitap varsa yazarla ayıklanır (soyadı tutmalı; yazarı
doğrulanan kayıt yazarsız kaydın önüne geçer); (2b) çekirdek ad (cilt yazımı, baskı/boy ifadesi ayrılmış) tek kayıtta;
kesin farklı ürün (başka cilt, set/tek kitap, e-kitap/basılı, başka yazar) hiçbir yolda otomatik eşleşmez, ayırt
edilemeyen ikiz kayıt ya da tek tarafta yazılı cilt/baskı en çok «öneri» olur (sessiz yanlış eşleme yok); (3) kalanlarda ad ve
yazar sözcüklerinin seyrekliğiyle ağırlıklı benzerlikten adaylar, Zeki AI'a kapalı küme seçimi («aynı eser hangisi,
hiçbiri») — `QueuedLlm.choose`, seçim + olasılık. Eşikler ayardır: otomatik kabul (varsayılan p ≥ 0,90 ve marj ≥ 0,50),
öneri (p ≥ 0,70 ve marj ≥ 0,30; insan onaylar), altı «emin değil» (adaylar gösterilir, insan seçer). Rakamı model üretmez:
stok, fiyat ve maliyet SQL'den gelir.

**Teklif fiyatı (K2 öneri, K4 karar):** önerilen birim fiyat = liste fiyatı (KDV hariç) × fiyat oranı. Fiyat oranı
geçmiş sonuçlardaki «kazanan fiyat ÷ liste toplamı» ortancasıdır (aynı kurum türünde en az `TENDER_PRICE_HISTORY_MIN`
sonuç varsa o türden, yoksa bütün sonuçlardan, o da yoksa 1 = liste fiyatı); gerekçesi ekranda. Karar onayı teklif
fiyatının da onayıdır; öneren onaylayamaz.

**Uygunluk puanı (K2):** eşleşen kalem oranı, stoğu yeten eşleşme oranı, hazır zorunlu belge oranı ve son teklif
tarihine kalan süre; ağırlıklar ayardır (`TENDER_SCORE_WEIGHTS`), ölçülemeyen parça dışarıda kalır ve ağırlıklar
yeniden dağıtılır. Parçalar ekranda ayrı ayrı yazılır.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import math
import os
import re
import statistics
import threading
import unicodedata
import uuid
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Optional
from xml.etree import ElementTree
from zoneinfo import ZoneInfo

import sqlalchemy as sa

log = logging.getLogger("semantic.tenders")
TZ = ZoneInfo("Europe/Istanbul")
_md = sa.MetaData()

TENDERS = sa.Table(
    "semantic_tenders", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("kaynak", sa.String(10), nullable=False),              # elle | dosya | resmi
    sa.Column("kaynak_no", sa.String(80)),                          # ihale kayıt numarası
    sa.Column("kurum", sa.String(300), nullable=False),
    sa.Column("kurum_turu", sa.String(20), nullable=False),
    sa.Column("il", sa.String(60)),
    sa.Column("konu", sa.Text, nullable=False),
    sa.Column("usul", sa.String(20)),
    sa.Column("yaklasik_tutar", sa.Float),
    sa.Column("ilan_tarihi", sa.String(10)),
    sa.Column("son_teklif_tarihi", sa.String(16)),                   # YYYY-MM-DD ya da YYYY-MM-DDTHH:MM (yerel)
    sa.Column("teslim_suresi", sa.String(200)),
    sa.Column("teminat_tutari", sa.Float),
    sa.Column("teminat_iade_tarihi", sa.String(10)),
    sa.Column("yetkili", sa.String(300)),                           # ilandaki kurum yetkilisi; yalnız bu kayıtta (KVKK)
    sa.Column("durum", sa.String(20), nullable=False),
    sa.Column("sorumlu", sa.String(120)),
    sa.Column("fiyat_orani", sa.Float),                             # teklif birim fiyatı ÷ liste fiyatı (KDV hariç)
    sa.Column("fiyat_orani_kaynak", sa.String(300)),
    sa.Column("uygunluk_puani", sa.Float),
    sa.Column("uygunluk_json", sa.Text),
    sa.Column("kitap_ilani_json", sa.Text),                         # Zeki: «kitap/yayın alımı mı» + olasılık
    sa.Column("ozet_json", sa.Text),                                # Zeki şartname özeti (kaynak cümleli)
    sa.Column("karar_metni", sa.Text),
    sa.Column("notlar", sa.Text),
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("updated_by", sa.String(120)),
    sa.Column("updated_at", sa.DateTime(timezone=True)),
    sa.Index("ix_semantic_tenders_state", "tenant_id", "durum"),
)
FILES = sa.Table(
    "semantic_tender_files", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tender_id", sa.String(32), nullable=False, index=True),
    sa.Column("tur", sa.String(10), nullable=False),                 # sartname | ek | belge
    sa.Column("ad", sa.String(300), nullable=False),
    sa.Column("dosya_yolu", sa.String(500), nullable=False),
    sa.Column("mime", sa.String(120), nullable=False),
    sa.Column("boyut", sa.Integer, nullable=False),
    sa.Column("sha256", sa.String(64), nullable=False),
    sa.Column("yukleyen", sa.String(120), nullable=False),
    sa.Column("zaman", sa.DateTime(timezone=True), nullable=False),
)
ITEMS = sa.Table(
    "semantic_tender_items", _md,
    sa.Column("tender_id", sa.String(32), primary_key=True),
    sa.Column("sira", sa.Integer, primary_key=True),
    sa.Column("sartname_metni", sa.Text, nullable=False),
    sa.Column("ad", sa.String(500)),
    sa.Column("yazar", sa.String(300)),
    sa.Column("yayinevi", sa.String(200)),
    sa.Column("adet", sa.Float),
    sa.Column("isbn", sa.String(20)),
    sa.Column("eslesen_stok_kodu", sa.String(60)),
    sa.Column("eslesen_ad", sa.String(500)),
    sa.Column("eslesme_yontemi", sa.String(10)),                    # isbn | ad | zeki | elle
    sa.Column("eslesme_durumu", sa.String(10), nullable=False),     # bekliyor | eslesti | oneri | belirsiz | yok
    sa.Column("olasilik", sa.Float),
    sa.Column("aday_json", sa.Text),
    sa.Column("stok", sa.Float),
    sa.Column("liste_fiyati", sa.Float),                            # KDV dahil (CRM ya da Logo; kaynak ayrı)
    sa.Column("fiyat_kaynagi", sa.String(200)),
    sa.Column("kdv_orani", sa.Float),                               # oran (0,10 = %10)
    sa.Column("logo_fiyati", sa.Float),
    sa.Column("logo_fiyat_notu", sa.String(200)),
    sa.Column("onerilen_fiyat", sa.Float),                          # KDV hariç birim teklif fiyatı
    sa.Column("fiyat_elle", sa.Boolean, nullable=False, default=False),
    sa.Column("tahmini_maliyet", sa.Float),
    sa.Column("maliyet_kaynagi", sa.String(200)),
    sa.Column("onaylayan", sa.String(120)),
    sa.Column("onay_zamani", sa.DateTime(timezone=True)),
    sa.Column("notu", sa.String(500)),
)
CHECKLIST = sa.Table(
    "semantic_tender_checklist", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tender_id", sa.String(32), nullable=False, index=True),
    sa.Column("sira", sa.Integer, nullable=False),
    sa.Column("kalem", sa.String(500), nullable=False),
    sa.Column("belge_turu", sa.String(20)),
    sa.Column("zorunlu", sa.Boolean, nullable=False, default=True),
    sa.Column("durum", sa.String(10), nullable=False),             # var | eksik | gecersiz
    sa.Column("belge_id", sa.String(32)),
    sa.Column("gecerlilik_tarihi", sa.String(10)),
    sa.Column("kaynak", sa.String(10)),                            # elle | ozet (şartname özetinden)
    sa.Column("kaynak_cumle", sa.Text),
    sa.Column("notu", sa.String(500)),
)
DOCUMENTS = sa.Table(
    "semantic_tender_documents", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("ad", sa.String(300), nullable=False),
    sa.Column("tur", sa.String(20), nullable=False),
    sa.Column("gecerlilik_tarihi", sa.String(10)),
    sa.Column("dosya_adi", sa.String(300)),
    sa.Column("dosya_yolu", sa.String(500)),
    sa.Column("mime", sa.String(120)),
    sa.Column("boyut", sa.Integer),
    sa.Column("sha256", sa.String(64)),
    sa.Column("notu", sa.String(500)),
    sa.Column("yukleyen", sa.String(120), nullable=False),
    sa.Column("zaman", sa.DateTime(timezone=True), nullable=False),
)
DECISIONS = sa.Table(
    "semantic_tender_decisions", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tender_id", sa.String(32), nullable=False, index=True),
    sa.Column("karar", sa.String(10), nullable=False),              # basvur | basvurma
    sa.Column("durum", sa.String(12), nullable=False),              # onayda | onaylandi | reddedildi | geri_cekildi
    sa.Column("teklif_toplami", sa.Float),
    sa.Column("fiyat_orani", sa.Float),
    sa.Column("gerekce", sa.Text),
    sa.Column("ozet_json", sa.Text),                                # öneri anındaki karar özeti (rakamlar)
    sa.Column("oneren", sa.String(120), nullable=False),
    sa.Column("oneri_zamani", sa.DateTime(timezone=True), nullable=False),
    sa.Column("onaylayan", sa.String(120)),
    sa.Column("onay_zamani", sa.DateTime(timezone=True)),
    sa.Column("onay_notu", sa.Text),
)
RESULTS = sa.Table(
    "semantic_tender_results", _md,
    sa.Column("tender_id", sa.String(32), primary_key=True),
    sa.Column("sonuc", sa.String(12), nullable=False),              # kazanildi | kaybedildi | iptal
    sa.Column("kazanan", sa.String(300)),
    sa.Column("kazanan_fiyat", sa.Float),                           # KDV hariç toplam
    sa.Column("bizim_fiyat", sa.Float),
    sa.Column("liste_toplami", sa.Float),                           # aynı kalemlerin KDV hariç liste toplamı (oran için)
    sa.Column("neden", sa.Text),
    sa.Column("kaynak", sa.String(300)),
    sa.Column("kaydeden", sa.String(120), nullable=False),
    sa.Column("zaman", sa.DateTime(timezone=True), nullable=False),
)
JOBS = sa.Table(
    "semantic_tender_jobs", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tender_id", sa.String(32), nullable=False, index=True),
    sa.Column("tur", sa.String(12), nullable=False),                # ozet | eslestirme
    sa.Column("durum", sa.String(12), nullable=False),              # sirada | calisiyor | bitti | hata | kesildi
    sa.Column("ilerleme", sa.Integer, nullable=False, default=0),
    sa.Column("toplam", sa.Integer, nullable=False, default=0),
    sa.Column("sonuc_json", sa.Text),
    sa.Column("hata", sa.Text),
    sa.Column("baslatan", sa.String(120), nullable=False),
    sa.Column("baslangic", sa.DateTime(timezone=True), nullable=False),
    sa.Column("bitis", sa.DateTime(timezone=True)),
)
REMINDERS = sa.Table(
    "semantic_tender_reminders", _md,
    sa.Column("key", sa.String(160), primary_key=True),             # aynı hatırlatma bir kez gider
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("gonderildi", sa.DateTime(timezone=True), nullable=False),
)

STATUSES = {"yeni": "Yeni", "inceleniyor": "İnceleniyor", "basvurulacak": "Başvurulacak", "basvurulmayacak": "Başvurulmayacak",
            "teklif_verildi": "Teklif verildi", "kazanildi": "Kazanıldı", "kaybedildi": "Kaybedildi", "iptal": "İptal"}
OPEN_STATUSES = ("yeni", "inceleniyor", "basvurulacak", "teklif_verildi")
#: Elle (PATCH) geçilebilen durumlar. Başvuru kararı karar akışıyla, kazanıldı/kaybedildi sonuç kaydıyla gelir.
MANUAL_STATUSES = ("yeni", "inceleniyor", "teklif_verildi", "iptal")
KURUM_TURLERI = {"okul": "Okul / okul kütüphanesi", "mem": "İl / ilçe milli eğitim", "halk_kutuphanesi": "Halk kütüphanesi",
                 "universite": "Üniversite", "belediye": "Belediye", "bakanlik": "Bakanlık / genel müdürlük",
                 "diger_kamu": "Diğer kamu kurumu"}
USULLER = {"acik": "Açık ihale", "pazarlik": "Pazarlık usulü", "dogrudan_temin": "Doğrudan temin",
           "yayin_alimi": "Yayın alımı başvurusu", "diger": "Diğer"}
KAYNAKLAR = {"elle": "Elle girildi", "dosya": "Dosyadan", "resmi": "Resmî kaynak"}
MATCH_STATES = {"bekliyor": "Eşleştirilmedi", "eslesti": "Eşleşti", "oneri": "Zeki AI önerisi — onay bekliyor",
                "belirsiz": "Emin değil — adaylardan seçin", "yok": "Katalogda yok"}
MATCH_METHODS = {"isbn": "ISBN/barkod", "ad": "Ad", "zeki": "Zeki AI", "elle": "Elle"}
DOC_TYPES = {"vergi": "Vergi borcu yoktur yazısı", "sgk": "SGK borcu yoktur yazısı", "imza_sirkuleri": "İmza sirküleri",
             "ticaret_sicil": "Ticaret sicil gazetesi / kayıt belgesi", "faaliyet": "Oda kayıt / faaliyet belgesi",
             "teminat": "Geçici / kesin teminat mektubu", "is_deneyim": "İş deneyim belgesi",
             "yetki": "Yetki belgesi / vekâletname", "diger": "Diğer"}
CHECK_STATES = {"var": "Var", "eksik": "Eksik", "gecersiz": "Süresi dolmuş / geçersiz"}
RESULT_STATES = {"kazanildi": "Kazanıldı", "kaybedildi": "Kaybedildi", "iptal": "İptal edildi"}
DECISIONS_LABEL = {"basvur": "Başvur", "basvurma": "Başvurma"}

#: Belge kalemi → belge türü (şartname özetindeki belge adlarını arşivdeki belgeyle bulmak için). Sözcükler Türkçe
#: harfleri sadeleşmiş hâldedir; eşleşmeyen kalem Zeki AI'a kapalı küme seçimiyle sorulur, o da bilemezse «diğer».
DOC_KEYWORDS: list[tuple[str, tuple[str, ...]]] = [
    ("sgk", ("sgk", "sosyal guvenlik", "prim borcu")),
    ("vergi", ("vergi borcu", "vergi dairesi", "vergi levhasi")),
    ("imza_sirkuleri", ("imza sirkuler", "imza beyannamesi")),
    ("teminat", ("teminat",)),
    ("ticaret_sicil", ("ticaret sicil", "sicil gazetesi")),
    ("faaliyet", ("oda kayit", "faaliyet belgesi", "odasina kayitli", "sanayi odasi", "ticaret odasi")),
    ("is_deneyim", ("is deneyim", "is bitirme")),
    ("yetki", ("vekaletname", "yetki belgesi", "yetkili bayi", "yetkili satici")),
]

_lock = threading.Lock()
_ready: set[int] = set()


class TenderError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def ensure(engine: sa.engine.Engine) -> None:
    key = id(engine)
    with _lock:
        if key in _ready:
            return
        _md.create_all(engine, checkfirst=True)
        _ready.add(key)


# ------------------------------------------------------------------ ayarlar


def _conf(key: str, default: str) -> str:
    try:
        from semantic_bridge import admin as admin_mod
        v = admin_mod.conf(key)
    except Exception:  # noqa: BLE001 — testte admin ayarı yok
        v = ""
    return v if v not in (None, "") else os.environ.get(key, default)


def _conf_float(key: str, default: float) -> float:
    try:
        return float(str(_conf(key, str(default))).replace(",", "."))
    except ValueError:
        return default


def settings() -> dict[str, Any]:
    """Modülün ayarları (ekran > ortam > varsayılan). Ölçülmemiş varsayımlar buradadır, kodda sabit değil."""
    try:
        weights = json.loads(_conf("TENDER_SCORE_WEIGHTS", "") or "{}")
    except ValueError:
        weights = {}
    w = {"eslesme": 40.0, "stok": 25.0, "belge": 20.0, "sure": 15.0}
    w.update({k: float(v) for k, v in weights.items() if k in w and isinstance(v, (int, float))})
    return {
        "autoProb": _conf_float("TENDER_MATCH_AUTO_PROB", 0.90), "autoMargin": _conf_float("TENDER_MATCH_AUTO_MARGIN", 0.50),
        "suggestProb": _conf_float("TENDER_MATCH_SUGGEST_PROB", 0.70), "suggestMargin": _conf_float("TENDER_MATCH_SUGGEST_MARGIN", 0.30),
        "candidates": max(2, int(_conf_float("TENDER_MATCH_CANDIDATES", 8))),
        "minScore": _conf_float("TENDER_MATCH_MIN_SCORE", 0.35),
        "priceSource": (_conf("TENDER_PRICE_SOURCE", "crm") or "crm").strip().lower(),
        "defaultVat": _conf_float("TENDER_DEFAULT_VAT", 0.0),
        "historyMin": max(1, int(_conf_float("TENDER_PRICE_HISTORY_MIN", 3))),
        "weights": w, "fullDays": max(1.0, _conf_float("TENDER_SCORE_FULL_DAYS", 14)),
        "summaryChunk": max(4000, int(_conf_float("TENDER_SUMMARY_CHUNK_CHARS", 24000))),
        "remindDays": [int(x) for x in re.findall(r"\d+", _conf("TENDER_REMIND_DAYS", "7,2"))] or [7, 2],
        "docWarnDays": int(_conf_float("TENDER_DOC_WARN_DAYS", 30)),
        "fileMaxMb": max(1, int(_conf_float("TENDER_FILE_MAX_MB", 50))),
        "publicChannel": (_conf("TENDER_PUBLIC_CHANNEL", "KURUM") or "KURUM").strip(),
        "watchEnabled": _conf("TENDER_WATCH_ENABLED", "0").strip().lower() in ("1", "true", "evet", "on"),
    }


def files_root() -> Path:
    return Path(os.environ.get("TENDER_DIR", "/data/nanobaseai/bi/var/tenders"))


# ------------------------------------------------------------------ küçük yardımcılar


def _now() -> datetime:
    return datetime.now(timezone.utc)


def today() -> date:
    return datetime.now(TZ).date()


def _iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    if v.tzinfo is None:
        v = v.replace(tzinfo=timezone.utc)
    return v.isoformat()


def _j(v: Optional[str], default: Any) -> Any:
    try:
        out = json.loads(v or "")
    except (TypeError, ValueError):
        return default
    return out if isinstance(out, type(default)) else default


def _dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _text(v: Any, limit: int) -> Optional[str]:
    s = str(v or "").strip()
    if not s:
        return None
    if len(s) > limit:
        raise TenderError(f"Metin {limit} karakteri aşıyor.")
    return s


def _num(v: Any, label: str, *, allow_none: bool = True, minimum: Optional[float] = 0.0) -> Optional[float]:
    if v is None or v == "":
        if allow_none:
            return None
        raise TenderError(f"{label} girilmeli.")
    if isinstance(v, str):
        s = v.strip().replace(" ", "").replace("₺", "")
        s = s.replace(".", "").replace(",", ".") if "," in s else s
        v = s
    try:
        n = float(v)
    except (TypeError, ValueError):
        raise TenderError(f"{label} sayı olmalı.") from None
    if not math.isfinite(n) or (minimum is not None and n < minimum):
        raise TenderError(f"{label} geçersiz.")
    return n


def _day(v: Any, label: str) -> Optional[str]:
    if v in (None, ""):
        return None
    s = str(v).strip()[:10]
    try:
        return date.fromisoformat(s).isoformat()
    except ValueError:
        raise TenderError(f"{label} YYYY-AA-GG biçiminde olmalı.") from None


def _deadline(v: Any) -> Optional[str]:
    if v in (None, ""):
        return None
    s = str(v).strip().replace(" ", "T")
    try:
        if len(s) <= 10:
            return date.fromisoformat(s).isoformat()
        return datetime.fromisoformat(s[:16]).strftime("%Y-%m-%dT%H:%M")
    except ValueError:
        raise TenderError("Son teklif tarihi YYYY-AA-GG ya da YYYY-AA-GG SS:DD biçiminde olmalı.") from None


def days_left(deadline: Optional[str], asof: Optional[date] = None) -> Optional[int]:
    if not deadline:
        return None
    try:
        d = date.fromisoformat(deadline[:10])
    except ValueError:
        return None
    return (d - (asof or today())).days


_APOSTROPHES = re.compile(r"(?<=\w)['’‘`´ʼ](?=\w)")


def fold(s: Any) -> str:
    """Karşılaştırma biçimi: Türkçe küçük harf (İ/I/ı/i), aksansız, noktalamasız, tek boşluk. Sözcük içi kesme işareti
    sözcüğü bölmez («Kur'an» → «kuran», «Emre'nin» → «emrenin»); tire ve öbür noktalama boşluk olur."""
    t = _APOSTROPHES.sub("", str(s or ""))
    t = t.replace("İ", "i").replace("I", "ı").lower()
    t = t.replace("ı", "i").replace("ş", "s").replace("ğ", "g").replace("ü", "u").replace("ö", "o").replace("ç", "c")
    t = unicodedata.normalize("NFKD", t)
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    t = re.sub(r"[^0-9a-z]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


_STOP = {"ve", "ile", "bir", "de", "da", "ya", "the", "of", "and", "a", "an", "icin"}


def tokens(s: Any) -> list[str]:
    return [t for t in fold(s).split() if len(t) > 1 and t not in _STOP]


# ------------------------------------------------------------------ ad çözümleme: çekirdek ad + cilt/set/biçim

_ROMAN = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "vi": 6, "vii": 7, "viii": 8, "ix": 9, "x": 10, "xi": 11, "xii": 12,
          "xiii": 13, "xiv": 14, "xv": 15, "xvi": 16, "xvii": 17, "xviii": 18, "xix": 19, "xx": 20}
_ROMAN_RX = "|".join(sorted(_ROMAN, key=len, reverse=True))
#: Sıra sayılı cilt (noktalı yazım; «3. Kitap» cilttir, «3 Kitap» set sayısıdır): «1. Cilt», «II. Kısım», «3.Kitap».
_VOL_ORDINAL = re.compile(rf"(?<![0-9a-z])(\d{{1,3}}|{_ROMAN_RX})\s*\.\s*(cilt|cild|kitap|kisim|bolum|sayi|seri)\b")
#: Noktasız cilt: «2 Cilt», «Cilt 2», «Cilt: II», «C. 3».
_VOL_BEFORE = re.compile(rf"\b(\d{{1,3}}|{_ROMAN_RX}) (cilt|cild|kisim|bolum|sayi)\b")
_VOL_AFTER = re.compile(rf"\b(cilt|cild|c|kisim|bolum|sayi|vol|volume) (\d{{1,3}}|{_ROMAN_RX})\b")
_VOL_TRAIL = re.compile(r" (\d{1,2})$")
_RAW_ROMAN_TRAIL = re.compile(r"\s([IVX]{1,5})\.?\s*\)?\s*$")      # ham metinde sonda büyük harf Romen rakamı
#: Set/paket: «Seti», «Takımı», «5 Kitap», «10 Kitaplık Set», «Kutu Set».
_SET_RX = re.compile(r"\b(\d{1,3} kitap(lik)?|kutu set|set|seti|takim|takimi)\b")
_DIGITAL_RX = re.compile(r"\b(e kitap|ekitap|e book|ebook|sesli kitap)\b")
#: Baskı/biçim ifadeleri (aynı eser, farklı baskı): çekirdek addan çıkarılır, ayrıca tutulur.
_FORMAT_RX = re.compile(r"\b(ciltli|ciltsiz|karton kapak|sert kapak|bez cilt|somizli|cep boy|buyuk boy|orta boy|kucuk boy|"
                        r"kutulu|ozel baski|yeni baski|\d{1,3} baski|tipki basim|genisletilmis baski)\b")
#: Yazar bilinmiyor sayılan yazımlar.
_NO_AUTHOR = {"kolektif", "komisyon", "anonim", "derleme", "cesitli", "yazarlar"}


def _fold_dots(s: str) -> str:
    """`fold` gibi, ama nokta kalır (sıra sayısı «3.» ile adet «3»ü ayırmak için)."""
    t = _APOSTROPHES.sub("", s).replace("İ", "i").replace("I", "ı").lower()
    t = t.replace("ı", "i").replace("ş", "s").replace("ğ", "g").replace("ü", "u").replace("ö", "o").replace("ç", "c")
    t = "".join(ch for ch in unicodedata.normalize("NFKD", t) if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", re.sub(r"[^0-9a-z.]+", " ", t)).strip()


def title_key(s: Any) -> dict[str, Any]:
    """Adın karşılaştırma anahtarı: çekirdek ad (biçim/cilt/set sözcükleri çıkmış), cilt no, set mi, dijital mi, biçimler.
    Kitaba özel kural yok: yalnız Türkçe yayıncılıkta genel yazımlar (cilt, set, baskı/boy, e-kitap)."""
    raw = str(s or "")
    vol: Optional[int] = None
    dotted = _fold_dots(raw)
    m = _VOL_ORDINAL.search(dotted)
    if m:
        vol = int(m.group(1)) if m.group(1).isdigit() else _ROMAN.get(m.group(1))
        dotted = dotted[:m.start()] + " " + dotted[m.end():]
    f = fold(dotted)
    if vol is None:
        m = _VOL_BEFORE.search(f) or _VOL_AFTER.search(f)
        if m:
            num_ = m.group(1) if m.re is _VOL_BEFORE else m.group(2)
            vol = int(num_) if num_.isdigit() else _ROMAN.get(num_)
            f = (f[:m.start()] + " " + f[m.end():]).strip()
    digital = bool(_DIGITAL_RX.search(f))
    f = _DIGITAL_RX.sub(" ", f)
    is_set = bool(_SET_RX.search(f))
    f = _SET_RX.sub(" ", f)
    fmts = frozenset(x.group(1) for x in _FORMAT_RX.finditer(f))
    f = re.sub(r"\s+", " ", _FORMAT_RX.sub(" ", f)).strip()
    if vol is None and len(f.split()) >= 2:
        t = _VOL_TRAIL.search(f)
        r = _RAW_ROMAN_TRAIL.search(raw)
        roman = r.group(1).lower() if r else ""
        if t:
            vol, f = int(t.group(1)), f[:t.start()].strip()
        elif roman in _ROMAN and f.endswith(" " + roman):
            vol, f = _ROMAN[roman], f[: -len(roman) - 1].strip()
    core = " ".join(w for w in f.split() if w not in _STOP)
    return {"core": core, "vol": vol, "set": is_set, "digital": digital, "fmt": fmts}


def _surnames(s: Any) -> tuple[set[str], set[str]]:
    """(bütün sözcükler, soyadları). Birden çok yazar virgül, «ve», «&», «/», «;» ile ayrılır; her yazarın son sözcüğü
    soyadı sayılır («Tanpınar, Ahmet Hamdi» biçiminde ilk sözcük de)."""
    txt = re.sub(r"\((.*?)\)", " ", str(s or ""))           # «(Çev. …)», «(Haz. …)» yazar değil
    parts = [p for p in re.split(r"[,;&/]| ve | and ", txt, flags=re.I) if p.strip()]
    words: set[str] = set()
    last: set[str] = set()
    for p in parts:
        tk = [t for t in tokens(p) if t not in _NO_AUTHOR]
        if tk:
            words |= set(tk)
            last.add(tk[-1])
    if "," in txt and len(parts) == 2 and len(tokens(parts[0])) == 1:     # «Soyad, Ad»
        last.add(tokens(parts[0])[0])
    return words, last


def author_compat(a: Any, b: Any) -> Optional[bool]:
    """İki yazar yazımı aynı kişi(ler) mi: None = bilinmiyor (biri boş/«kolektif»), True/False. Yalnız ortak ad yetmez
    («Ahmet Ümit» ≠ «Ahmet Hamdi Tanpınar»): bir tarafın soyadı öbüründe geçmeli ya da en az iki sözcük ortak olmalı."""
    wa, la = _surnames(a)
    wb, lb = _surnames(b)
    if not wa or not wb:
        return None
    return bool(la & wb) or bool(lb & wa) or len(wa & wb) >= 2


def key_conflict(a: dict[str, Any], b: dict[str, Any]) -> Optional[str]:
    """Kesin farklı eser/ürün: iki tarafta da cilt no var ve farklı; biri set öbürü tek kitap; biri dijital."""
    if a["vol"] is not None and b["vol"] is not None and a["vol"] != b["vol"]:
        return f"cilt farklı ({a['vol']} / {b['vol']})"
    if a["set"] != b["set"]:
        return "biri set, öbürü tek kitap"
    if a["digital"] != b["digital"]:
        return "biri dijital, öbürü basılı"
    return None


def key_doubt(a: dict[str, Any], b: dict[str, Any]) -> Optional[str]:
    """Aynı eser olabilir ama emin olunamaz: cilt bir tarafta var öbüründe yok; baskı/biçim ifadesi farklı."""
    if (a["vol"] is None) != (b["vol"] is None):
        return "cilt numarası yalnız bir tarafta yazılı"
    if a["fmt"] != b["fmt"] and (a["fmt"] or b["fmt"]):
        return "baskı/biçim ifadesi farklı (" + (", ".join(sorted(a["fmt"] ^ b["fmt"]))) + ")"
    return None


def _title_tail(part: str) -> bool:
    """Tireyle ayrılmış parça adın devamı mı (cilt no, Romen rakamı, set, e-kitap, baskı/boy) — yazar değil."""
    if re.fullmatch(rf"(\d{{1,3}}|{_ROMAN_RX})", fold(part)):
        return True
    k = title_key(part)
    return not k["core"] and (k["vol"] is not None or k["set"] or k["digital"] or bool(k["fmt"]))


# ------------------------------------------------------------------ ISBN


def _isbn13_ok(d: str) -> bool:
    if len(d) != 13 or not d.isdigit():
        return False
    s = sum(int(c) * (1 if i % 2 == 0 else 3) for i, c in enumerate(d[:12]))
    return (10 - s % 10) % 10 == int(d[12])


def _isbn10_ok(d: str) -> bool:
    if len(d) != 10 or not d[:9].isdigit() or not (d[9].isdigit() or d[9] == "X"):
        return False
    s = sum((10 - i) * int(c) for i, c in enumerate(d[:9])) + (10 if d[9] == "X" else int(d[9]))
    return s % 11 == 0


def norm_isbn(v: Any, *, strict: bool = False) -> Optional[str]:
    """ISBN-10/13 ya da EAN-13 barkod → 13 haneli anahtar. `strict` ise sağlama basamağı tutmayan reddedilir
    (serbest metinden çıkarırken telefon/tarih gibi sayıları ayıklamak için). Katalog alanlarında gevşek: kayıtlı
    değer ne ise o."""
    d = re.sub(r"[^0-9Xx]", "", str(v or "")).upper()
    if len(d) == 10:
        if strict and not _isbn10_ok(d):
            return None
        core = "978" + d[:9]
        s = sum(int(c) * (1 if i % 2 == 0 else 3) for i, c in enumerate(core))
        return core + str((10 - s % 10) % 10) if core.isdigit() else None
    if len(d) == 13 and d.isdigit():
        if strict and not _isbn13_ok(d):
            return None
        return d
    return None


_ISBN_RX = re.compile(r"(?<![0-9])(?:97[89][\s\-]?)?(?:[0-9][\s\-]?){9}[0-9Xx](?![0-9])")


def find_isbn(text: str) -> Optional[str]:
    for m in _ISBN_RX.finditer(text or ""):
        k = norm_isbn(m.group(0), strict=True)
        if k:
            return k
    return None


# ------------------------------------------------------------------ şartname kalem listesi okuma


_HEADERS: dict[str, tuple[str, ...]] = {
    "sira": ("sira", "sira no", "s no", "no", "sn"),
    "ad": ("kitap adi", "kitabin adi", "eser adi", "eserin adi", "adi", "ad", "kitap", "eser", "malzeme", "malzemenin adi",
           "urun adi", "urun", "cinsi", "mal hizmet", "mal hizmet adi", "aciklama", "yayin adi", "kitap ismi", "isim"),
    "yazar": ("yazar", "yazari", "yazar adi", "yazarin adi"),
    "yayinevi": ("yayinevi", "yayin evi", "yayinevi adi"),
    "isbn": ("isbn", "isbn no", "barkod", "barkod no", "barcode", "ean"),
    "adet": ("adet", "miktar", "miktari", "miktar adet", "sayi", "sayisi", "adedi", "istenen adet"),
}


def _header_map(row: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, cell in enumerate(row):
        f = fold(cell)
        if not f:
            continue
        for key, names in _HEADERS.items():
            if key in out:
                continue
            if f in names or any(f.startswith(n + " ") for n in names if len(n) > 3):
                out[key] = i
                break
    return out if len(out) >= 2 and ("ad" in out or "isbn" in out) else {}


def _to_int(v: Any) -> Optional[float]:
    s = str(v or "").strip()
    if not s:
        return None
    s = re.sub(r"(?i)\s*(adet|ad\.?|tane|nüsha|nusha)\s*$", "", s)
    s = s.replace(".", "") if re.fullmatch(r"\d{1,3}(\.\d{3})+", s) else s
    s = s.replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        return None
    return n if n >= 0 and math.isfinite(n) else None


_QTY_RX = re.compile(r"(?i)(\d[\d.]*)\s*(?:adet|ad\.|tane|nüsha|nusha)\b|\b[xX]\s*(\d[\d.]*)\s*$")
_SIRA_RX = re.compile(r"^\s*(\d{1,5})\s*[.)\-–]\s+")


def _row_from_cells(cells: list[str]) -> dict[str, Any]:
    """Başlıksız satır: ISBN hücresi, adet (son tam sayı), en uzun metin ad, ikinci uzun metin yazar."""
    cells = [str(c).strip() for c in cells if str(c or "").strip()]
    isbn = None
    rest: list[str] = []
    for c in cells:
        k = find_isbn(c) if isbn is None else None
        if k and len(re.sub(r"[^0-9Xx]", "", c)) in (10, 13):
            isbn = k
            continue
        rest.append(c)
    adet = None
    texts: list[str] = []
    for i, c in enumerate(rest):
        n = _to_int(c)
        if n is not None and re.fullmatch(r"[\d.,]+(\s*(adet|ad\.?|tane))?", c, flags=re.I):
            if i == 0 and len(rest) > 2:   # ilk hücre sıra numarası
                continue
            adet = n
        else:
            texts.append(c)
    texts.sort(key=len, reverse=True)
    return {"ad": texts[0] if texts else None, "yazar": texts[1] if len(texts) > 1 else None, "yayinevi": None,
            "isbn": isbn, "adet": adet}


def _row_from_line(line: str) -> dict[str, Any]:
    """Tek hücreli serbest satır: «12. Küçük Prens - Antoine de Saint-Exupéry 978… 25 adet»."""
    s = _SIRA_RX.sub("", line).strip()
    isbn = find_isbn(s)
    if isbn:
        s = _ISBN_RX.sub(" ", s)
        s = re.sub(r"(?i)\bisbn\b[:\s]*", " ", s)
    adet = None
    m = None
    for m in _QTY_RX.finditer(s):
        pass
    if m:
        adet = _to_int(m.group(1) or m.group(2))
        s = s[:m.start()] + s[m.end():]
    else:
        m2 = re.search(r"\s(\d{1,6})\s*$", s)
        if m2:
            adet = _to_int(m2.group(1))
            s = s[:m2.start()]
    s = re.sub(r"\s+", " ", s).strip(" -–,;:/")
    ad, yazar = s, None
    parts = [p.strip() for p in re.split(r"\s[-–/]\s", s) if p.strip()]
    # «Osmanlı Tarihi - 2», «Sefiller - II. Cilt», «Nutuk - Ciltli»: tireden sonraki parça yazar değil, adın parçası.
    while len(parts) >= 2 and _title_tail(parts[1]):
        parts[0:2] = [f"{parts[0]} {parts[1]}"]
    if len(parts) >= 2:
        ad, yazar = parts[0], parts[1]
    else:
        ad = parts[0] if parts else ad
    return {"ad": ad or None, "yazar": yazar, "yayinevi": None, "isbn": isbn, "adet": adet}


def parse_rows(rows: list[list[Any]]) -> dict[str, Any]:
    """Tablo satırları (Excel/CSV/yapıştırılmış sekmeli metin) → kalemler. Başlık satırı ilk 15 satırda aranır;
    bulunamazsa her satır kendi hücrelerinden okunur. Okunamayan satır nedeniyle listelenir; hiçbir satır kesilmez."""
    rows = [[("" if c is None else str(c)) for c in r] for r in rows]
    header: dict[str, int] = {}
    start = 0
    for i, r in enumerate(rows[:15]):
        h = _header_map(r)
        if h:
            header, start = h, i + 1
            break
    items: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    for n, r in enumerate(rows[start:], start=start + 1):
        raw = " | ".join(c.strip() for c in r if c.strip())
        if not raw:
            continue
        if header:
            get = lambda k: r[header[k]].strip() if k in header and header[k] < len(r) else ""  # noqa: E731
            isbn = norm_isbn(get("isbn")) if get("isbn") else find_isbn(raw)
            row = {"ad": get("ad") or None, "yazar": get("yazar") or None, "yayinevi": get("yayinevi") or None,
                   "isbn": isbn, "adet": _to_int(get("adet"))}
            if not row["ad"] and not row["isbn"]:
                fb = _row_from_cells(r)
                row.update({k: row[k] or fb[k] for k in row})
        elif len([c for c in r if c.strip()]) > 1:
            row = _row_from_cells(r)
        else:
            row = _row_from_line(raw)
        if not row["ad"] and not row["isbn"]:
            skipped.append({"satir": n, "metin": raw[:300], "neden": "Kitap adı ya da ISBN bulunamadı"})
            continue
        if row["ad"] and not tokens(row["ad"]) and not row["isbn"]:
            skipped.append({"satir": n, "metin": raw[:300], "neden": "Başlık ya da toplam satırı gibi görünüyor"})
            continue
        items.append({**row, "metin": raw})
    return {"items": items, "skipped": skipped, "header": sorted(header)}


def parse_text(text: str) -> dict[str, Any]:
    """Yapıştırılmış metin: sekme, noktalı virgül ya da dikey çizgi ayrılmış tablo ya da satır satır liste."""
    lines = [ln for ln in (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    rows: list[list[str]] = []
    for ln in lines:
        if "\t" in ln:
            rows.append(ln.split("\t"))
        elif ln.count(";") >= 1 and ln.count(";") >= ln.count("|"):
            rows.append(ln.split(";"))
        elif ln.count("|") >= 1:
            rows.append(ln.split("|"))
        else:
            rows.append([ln])
    return parse_rows(rows)


def extract_rows(filename: str, data: bytes) -> list[list[Any]]:
    """Dosyadan tablo satırları: xlsx (bütün sayfalar art arda), csv/txt, docx (tablolar ve paragraflar), pdf (metin)."""
    ext = _ext(filename)
    if ext == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out: list[list[Any]] = []
        for ws in wb.worksheets:
            out.extend([list(r) for r in ws.iter_rows(values_only=True)])
        return out
    if ext in ("csv", "txt"):
        text = _decode(data)
        if ext == "csv":
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
                return [r for r in csv.reader(io.StringIO(text), dialect)]
            except csv.Error:
                pass
        return [[ln] for ln in text.splitlines()]
    if ext == "docx":
        return [r for r in _docx_rows(data)]
    if ext == "pdf":
        return [[ln] for ln in pdf_text(data).splitlines()]
    raise TenderError("Kalem listesi Excel (.xlsx), CSV, metin, Word (.docx) ya da PDF olmalı.")


def document_reading(filename: str, data: bytes):
    """Şartname özetinin girdisi, ortak belge okuma hattıyla (`doc_read`): PDF'in metin katmanı, taranmış sayfalar OCR;
    görüntü dosyası (JPG/PNG/TIFF) da okunur. Dönen: (düz metin, okuma) — okuma sayfa/OCR/güven bilgisini taşır."""
    from semantic_bridge import doc_read as DR

    ext = _ext(filename)
    if ext in ("pdf",) + tuple(DR.IMAGES):
        try:
            reading = DR.read(filename, data, allowed=("pdf",) + tuple(DR.IMAGES))
        except DR.ReadError as e:
            raise TenderError(str(e), e.status) from None
        return reading.text(), reading
    text = document_text(filename, data)
    return text, DR.Reading(filename=filename, pages=[{"sayfa": "1", "metin": text,
                                                        "okuma": DR.METIN if text.strip() else DR.YOK, "guven": None}])


def document_text(filename: str, data: bytes) -> str:
    """Şartname özetinin girdisi: dosyanın düz metni (yalnız metin katmanı; taranmış sayfa için `document_reading`)."""
    ext = _ext(filename)
    if ext == "pdf":
        return pdf_text(data)
    if ext == "docx":
        return "\n".join("\t".join(r) for r in _docx_rows(data))
    if ext in ("txt", "csv"):
        return _decode(data)
    if ext == "xlsx":
        return "\n".join("\t".join("" if c is None else str(c) for c in r) for r in extract_rows(filename, data))
    raise TenderError("Bu dosya türünden metin okunamıyor (PDF, Word, Excel, metin).")


def pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover
        raise TenderError("PDF okuyucu bu kurulumda yok.", 503) from None
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception as e:  # noqa: BLE001
        raise TenderError(f"PDF okunamadı: {str(e)[:120]}") from None


def _docx_rows(data: bytes) -> list[list[str]]:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = ElementTree.fromstring(z.read("word/document.xml"))
    except (KeyError, zipfile.BadZipFile, ElementTree.ParseError):
        raise TenderError("Word dosyası okunamadı.") from None
    body = root.find("w:body", ns)
    out: list[list[str]] = []
    for el in list(body) if body is not None else []:
        tag = el.tag.split("}")[-1]
        if tag == "tbl":
            for tr in el.iter(f"{{{ns['w']}}}tr"):
                out.append(["".join(t.text or "" for t in tc.iter(f"{{{ns['w']}}}t")) for tc in tr.findall("w:tc", ns)])
        elif tag == "p":
            out.append(["".join(t.text or "" for t in el.iter(f"{{{ns['w']}}}t"))])
    return out


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _ext(filename: str) -> str:
    return re.sub(r"[^a-z0-9]", "", filename.lower().rsplit(".", 1)[-1])[:8] if "." in (filename or "") else ""


# ------------------------------------------------------------------ katalog ve eşleştirme


class Catalog:
    """Eşleştirme dizini: ISBN → kitaplar, sadeleşmiş ad → kitaplar, sözcük → kitaplar (seyreklik ağırlığıyla)."""

    def __init__(self, books: list[dict[str, Any]], notes: Optional[list[str]] = None):
        self.books = books
        self.notes = list(notes or [])
        self.by_code: dict[str, dict[str, Any]] = {}
        self.by_isbn: dict[str, list[int]] = {}
        self.by_title: dict[str, list[int]] = {}
        self.index: dict[str, set[int]] = {}
        self.toks: list[set[str]] = []
        self.author_toks: list[set[str]] = []
        self.keys: list[dict[str, Any]] = []
        self.by_core: dict[str, list[int]] = {}
        for i, b in enumerate(books):
            key = title_key(b.get("ad"))
            self.keys.append(key)
            if key["core"]:
                self.by_core.setdefault(key["core"], []).append(i)
            if b.get("kod"):
                self.by_code.setdefault(b["kod"], b)
            keys = {k for k in (norm_isbn(x) for x in (b.get("isbn_ham") or [])) if k}
            b["isbn"] = sorted(keys)
            for k in keys:
                self.by_isbn.setdefault(k, []).append(i)
            f = fold(b.get("ad"))
            if f:
                self.by_title.setdefault(f, []).append(i)
            t = set(tokens(b.get("ad")))
            self.toks.append(t)
            self.author_toks.append(set(tokens(b.get("yazar"))))
            for w in t:
                self.index.setdefault(w, set()).add(i)
        n = max(1, len(books))
        self.idf = {w: math.log(1 + n / len(ids)) for w, ids in self.index.items()}

    def label(self, i: int) -> str:
        b = self.books[i]
        parts = [b.get("ad") or "(adsız)", b.get("yazar"), b.get("yayinevi"), (b.get("isbn") or [None])[0], b.get("kod")]
        return " · ".join(str(p) for p in parts if p)

    def candidates(self, ad: Optional[str], yazar: Optional[str], k: int, min_score: float) -> list[tuple[int, float]]:
        """Adın sözcükleri (seyreklikle ağırlıklı) iki yönlü örtüşme; yazar sözcüğü tutan adaya ek puan."""
        qt = set(tokens(ad))
        if not qt:
            return []
        at = set(tokens(yazar))
        # Çok yaygın sözcük («kitap», «hikaye») aday havuzunu şişirir; havuz önce seyrek sözcüklerden kurulur.
        common_limit = max(50, len(self.books) // 5)
        rare = [w for w in qt if 0 < len(self.index.get(w, ())) <= common_limit]
        pool: set[int] = set()
        for w in (rare or qt):
            pool |= self.index.get(w, set())
        qw = sum(self.idf.get(w, math.log(1 + len(self.books))) for w in qt)
        scored = []
        for i in pool:
            common = qt & self.toks[i]
            cw = sum(self.idf.get(w, 0.0) for w in common)
            bw = sum(self.idf.get(w, 0.0) for w in self.toks[i]) or 1.0
            recall, precision = cw / qw, cw / bw
            s = 0.0 if not cw else 2 * recall * precision / (recall + precision)
            if at and self.author_toks[i]:
                s = min(1.0, s + 0.15 * (len(at & self.author_toks[i]) / len(at)))
            if s >= min_score:
                scored.append((i, round(s, 4)))
        scored.sort(key=lambda x: (-x[1], self.books[x[0]].get("ad") or ""))
        return scored[:k]


NONE_CHOICE = "Hiçbiri — katalogda bu eser yok"


def _author_ok(cat: Catalog, i: int, yazar: Optional[str]) -> bool:
    """Yazar çelişmiyor: bilinmiyor (bir taraf boş ya da «kolektif») ya da aynı kişi(ler) (`author_compat`)."""
    return author_compat(yazar, cat.books[i].get("yazar")) is not False


def _narrow_by_author(cat: Catalog, ids: list[int], yazar: Optional[str]) -> list[int]:
    """Aynı adlı kayıtlardan yazarı DOĞRULANAN tek kayıt varsa o kalır: yazarı boş kayıt (ör. CRM'de olmayan Logo
    malzeme kartı) yazarı tutan kaydın önüne geçmez. Birden çok doğrulanan ya da hiç yoksa liste olduğu gibi döner."""
    if len(ids) <= 1 or not yazar:
        return ids
    confirmed = [i for i in ids if author_compat(yazar, cat.books[i].get("yazar")) is True]
    return confirmed if len(confirmed) == 1 else ids


def _twins(cat: Catalog, i: int, others: list[int], yazar: Optional[str]) -> list[int]:
    """Kalemin bilgisiyle `i`'den ayırt edilemeyen kayıtlar: aynı çekirdek ad ve cilt/set/biçim, kalemin yazarıyla
    ilişkisi aynı (ikisi de bilinmiyor ya da ikisi de tutuyor). Aynı eserin iki baskısı/kaydı ya da yazarı verilmemiş
    kalemde aynı adlı iki eser — model hangisini seçerse seçsin insan onayı gerekir."""
    ki, rel = cat.keys[i], author_compat(yazar, cat.books[i].get("yazar"))
    out = []
    for j in others:
        kj = cat.keys[j]
        if j == i or kj["core"] != ki["core"] or key_conflict(ki, kj) or key_doubt(ki, kj):
            continue
        if author_compat(yazar, cat.books[j].get("yazar")) == rel:
            out.append(j)
    return out


def match_one(item: dict[str, Any], cat: Catalog, choose: Optional[Callable[[str, list[str]], Any]],
              cfg: dict[str, Any]) -> dict[str, Any]:
    """Tek kalemi eşleştirir. Dönen: durum, yöntem, stok kodu, olasılık, adaylar (ekranda seçim için).

    Sessiz yanlış eşleme olmasın diye (kabul 2026-09-28: adla verilen 4 kalemden 1'i yanlış kitaba): kesin farklı ürün
    (başka cilt, set/tek kitap, dijital/basılı, başka yazar) otomatik eşleşmez ve modele aday olarak gitmez (insan
    listede görür); ayırt edilemeyen ikiz kayıt ya da tek tarafta yazılı cilt/baskı ifadesi varsa sonuç en çok «öneri»dir."""
    out: dict[str, Any] = {"eslesme_durumu": "yok", "eslesme_yontemi": None, "eslesen_stok_kodu": None, "eslesen_ad": None,
                           "olasilik": None, "adaylar": [], "not": None}

    def hit(i: int, method: str, state: str = "eslesti", p: Optional[float] = None) -> dict[str, Any]:
        b = cat.books[i]
        out.update(eslesme_durumu=state, eslesme_yontemi=method, eslesen_stok_kodu=b.get("kod"), eslesen_ad=b.get("ad"), olasilik=p)
        return out

    def cand_list(ids: list[tuple[int, Optional[float]]]) -> list[dict[str, Any]]:
        return [{"stokKodu": cat.books[i].get("kod"), "ad": cat.books[i].get("ad"), "yazar": cat.books[i].get("yazar"),
                 "yayinevi": cat.books[i].get("yayinevi"), "isbn": (cat.books[i].get("isbn") or [None])[0], "benzerlik": s}
                for i, s in ids]

    yazar = item.get("yazar")
    ikey = title_key(item.get("ad"))
    # 1. ISBN / barkod
    isbn = norm_isbn(item.get("isbn")) if item.get("isbn") else None
    pool: list[int] = []
    if isbn and isbn in cat.by_isbn:
        pool = [i for i in cat.by_isbn[isbn] if cat.books[i].get("kod")] or cat.by_isbn[isbn]
        if len(pool) == 1:
            return hit(pool[0], "isbn", p=1.0)
        out["not"] = "ISBN birden çok katalog kaydında var; adla ayıklandı."
    # 2. Ad birebir (yazarla ayıklanır). ISBN birden çok kayıtta ise ad yalnız o kayıtlar arasında aranır.
    f = fold(item.get("ad"))
    if pool:
        same = [i for i in pool if f and fold(cat.books[i].get("ad")) == f]
    else:
        same = list(cat.by_title.get(f, [])) if f else []
    same = _narrow_by_author(cat, [i for i in same if cat.books[i].get("kod") and _author_ok(cat, i, yazar)], yazar)
    if len(same) == 1:
        return hit(same[0], "isbn" if pool else "ad", p=1.0)
    # 2b. Çekirdek ad: noktalama, «1. Cilt»/«Cilt 1», baskı/boy ifadesi farkı. Yalnız tek ve çelişmeyen kayıtta; tek
    #     tarafta yazılı cilt ya da farklı baskı ifadesi «öneri» olur (insan onaylar).
    if not pool and not same and ikey["core"]:
        core = [i for i in cat.by_core.get(ikey["core"], [])
                if cat.books[i].get("kod") and _author_ok(cat, i, yazar) and not key_conflict(ikey, cat.keys[i])]
        core = _narrow_by_author(cat, core, yazar)
        if len(core) == 1:
            doubt = key_doubt(ikey, cat.keys[core[0]])
            if doubt is None:
                hit(core[0], "ad", p=1.0)
                out["not"] = "Ad eşleşti (yazım farkı yok sayıldı)."
                return out
            hit(core[0], "ad", "oneri")
            out["adaylar"] = cand_list([(core[0], None)])
            out["not"] = f"Ad eşleşti ama {doubt}; onaylayın ya da başkasını seçin."
            return out
    # 3. Adaylar → Zeki AI
    if pool:
        cands: list[tuple[int, Optional[float]]] = [(i, None) for i in pool]
    elif len(same) > 1:
        cands = [(i, 1.0) for i in same]
        out["not"] = "Aynı adla birden çok katalog kaydı var."
    else:
        cands = list(cat.candidates(item.get("ad"), yazar, cfg["candidates"], cfg["minScore"]))
    cands = [(i, s) for i, s in cands if cat.books[i].get("kod")]
    out["adaylar"] = cand_list(cands)          # insan bütün adayları görür (çelişenler dahil)
    if not cands:
        out["not"] = out["not"] or "Katalogda benzer ad bulunamadı."
        return out
    ok = [(i, s) for i, s in cands if not key_conflict(ikey, cat.keys[i]) and _author_ok(cat, i, yazar)]
    if not ok:
        out["eslesme_durumu"] = "belirsiz"
        why = sorted({key_conflict(ikey, cat.keys[i]) or "yazar farklı" for i, _ in cands})
        out["not"] = f"Emin değil: adaylar kalemden farklı ({'; '.join(why)}). Elle seçin ya da «katalogda yok» işaretleyin."
        return out
    if choose is None:
        out["eslesme_durumu"] = "belirsiz"
        out["not"] = "Zeki AI bu kurulumda yok; adaylardan seçin."
        return out
    labels = [cat.label(i) for i, _ in ok]
    seen: dict[str, int] = {}
    for n, lb in enumerate(labels):  # seçenekler benzersiz olmalı
        if lb in seen:
            labels[n] = f"{lb} (#{n + 1})"
        seen[labels[n]] = n
    prompt = ("Kamu ihalesi teknik şartnamesindeki bir kitap kalemi ile yayınevinin kataloğundaki adaylar aşağıda. "
              "Şartnamedeki kalem hangi katalog kaydıyla AYNI ESERDİR? Baskı/kapak farkı önemsizdir; farklı eser, farklı "
              "cilt, set ile tek kitap ya da farklı yazar aynı eser değildir. Hiçbiri aynı eser değilse son seçeneği seçin.\n\n"
              f"Şartname kalemi: {item.get('metin') or item.get('ad')}\n"
              f"Ad: {item.get('ad') or '-'}\nYazar: {yazar or '-'}\nYayınevi: {item.get('yayinevi') or '-'}\n"
              f"ISBN: {item.get('isbn') or '-'}")
    try:
        r = choose(prompt, labels + [NONE_CHOICE])
    except Exception as e:  # noqa: BLE001 — model yoksa kalem insan seçimine kalır
        log.warning("ihale eşleştirme: model cevap vermedi: %s", e)
        out["eslesme_durumu"] = "belirsiz"
        out["not"] = "Zeki AI'a ulaşılamadı; adaylardan seçin ya da eşleştirmeyi yeniden başlatın."
        return out
    p = getattr(r, "probability", None)
    auto = r.confident(cfg["autoProb"], cfg["autoMargin"]) if hasattr(r, "confident") else False
    sugg = r.confident(cfg["suggestProb"], cfg["suggestMargin"]) if hasattr(r, "confident") else False
    if r.choice == NONE_CHOICE:
        out["olasilik"] = p
        out["eslesme_durumu"] = "yok" if sugg else "belirsiz"
        out["not"] = "Zeki AI: katalogda aynı eser yok." if sugg else "Zeki AI emin değil; adaylara bakın."
        return out
    if r.choice is None or r.choice not in seen:
        out["eslesme_durumu"] = "belirsiz"
        out["not"] = "Zeki AI seçim yapamadı; adaylardan seçin."
        return out
    i = ok[seen[r.choice]][0]
    # Model «aynı eser» dedi; kural denetimi: ayırt edilemeyen ikiz ya da tek tarafta yazılı cilt/baskı → en çok öneri.
    doubts = [d for d in (key_doubt(ikey, cat.keys[i]),) if d]
    if _twins(cat, i, [j for j, _ in ok], yazar):
        doubts.append("aynı ad ve bilgiyle birden çok katalog kaydı var (baskı/boy ya da yazar kalemde yok)")
    if auto and not doubts:
        return hit(i, "zeki", "eslesti", p)
    if auto or sugg:
        hit(i, "zeki", "oneri", p)
        if doubts:
            out["not"] = "Zeki AI aynı eser dedi ama " + "; ".join(doubts) + ". Onaylayın ya da başkasını seçin."
        return out
    hit(i, "zeki", "belirsiz", p)
    out["not"] = "Zeki AI emin değil; önerilen aday işaretli, onaylayın ya da başkasını seçin."
    return out


# ------------------------------------------------------------------ fiyat ve toplamlar


def vat_rate(raw: Optional[float], default: float) -> float:
    """CRM KDV oranı yüzde (10) ya da oran (0,10) tutulabilir (ölçülecek); ikisi de orana çevrilir."""
    if raw is None:
        return default
    return raw / 100.0 if raw > 1 else raw


def net_list(item: dict[str, Any]) -> Optional[float]:
    """KDV hariç liste fiyatı."""
    lp = item.get("liste_fiyati")
    if lp is None:
        return None
    return lp / (1 + (item.get("kdv_orani") or 0.0))


def suggested_price(item: dict[str, Any], ratio: float) -> Optional[float]:
    n = net_list(item)
    return None if n is None else round(n * ratio, 2)


def price_ratio(engine: sa.engine.Engine, tenant: str, kurum_turu: Optional[str], minimum: int) -> dict[str, Any]:
    """Geçmiş sonuçlardan fiyat oranı önerisi: kazanan fiyat ÷ aynı kalemlerin liste toplamı (KDV hariç) ortancası."""
    with engine.connect() as c:
        rows = c.execute(sa.select(RESULTS.c.kazanan_fiyat, RESULTS.c.liste_toplami, TENDERS.c.kurum_turu)
                         .join(TENDERS, TENDERS.c.id == RESULTS.c.tender_id)
                         .where(TENDERS.c.tenant_id == tenant, RESULTS.c.sonuc.in_(("kazanildi", "kaybedildi")),
                                RESULTS.c.kazanan_fiyat.is_not(None), RESULTS.c.liste_toplami > 0)).all()
    same = [r.kazanan_fiyat / r.liste_toplami for r in rows if r.kurum_turu == kurum_turu]
    every = [r.kazanan_fiyat / r.liste_toplami for r in rows]
    if kurum_turu and len(same) >= minimum:
        return {"oran": round(statistics.median(same), 4), "n": len(same),
                "kaynak": f"{len(same)} geçmiş {KURUM_TURLERI.get(kurum_turu, kurum_turu).lower()} ihalesinde kazanan fiyatın liste fiyatına oranı (ortanca)"}
    if len(every) >= minimum:
        return {"oran": round(statistics.median(every), 4), "n": len(every),
                "kaynak": f"{len(every)} geçmiş ihalede kazanan fiyatın liste fiyatına oranı (ortanca)"}
    return {"oran": 1.0, "n": len(every),
            "kaynak": f"Liste fiyatı (geçmiş sonuç {len(every)}; oran önerisi için en az {minimum} sonuç gerekir)"}


def totals(items: list[dict[str, Any]]) -> dict[str, Any]:
    """Teklif tablosu toplamları. Toplama yalnız eşleşmiş ve adedi, fiyatı olan kalem girer; dışarıda kalanlar sayılır.
    Tutar kuruşa yuvarlanır (satır başına), KDV ayrı satır."""
    states = {k: 0 for k in MATCH_STATES}
    ara = kdv = liste = 0.0
    priced = excluded = short = 0
    cost_lines = 0
    cost = cost_revenue = 0.0
    for it in items:
        states[it["eslesme_durumu"]] = states.get(it["eslesme_durumu"], 0) + 1
        if it["eslesme_durumu"] != "eslesti":
            continue
        if it.get("stok") is not None and it.get("adet") is not None and it["stok"] < it["adet"]:
            short += 1
        if it.get("adet") is None or it.get("onerilen_fiyat") is None:
            excluded += 1
            continue
        line = round(it["adet"] * it["onerilen_fiyat"], 2)
        ara += line
        kdv += round(line * (it.get("kdv_orani") or 0.0), 2)
        nl = net_list(it)
        if nl is not None:
            liste += round(it["adet"] * nl, 2)
        priced += 1
        if it.get("tahmini_maliyet") is not None:
            cost_lines += 1
            cost += it["adet"] * it["tahmini_maliyet"]
            cost_revenue += line
    marj = (cost_revenue - cost) / cost_revenue if cost_lines and cost_revenue > 0 else None
    return {"kalem": len(items), "durumlar": states, "stokYetersiz": short, "fiyatli": priced, "disarida": excluded,
            "araToplam": round(ara, 2), "kdv": round(kdv, 2), "genelToplam": round(ara + kdv, 2),
            "listeToplami": round(liste, 2), "maliyet": round(cost, 2) if cost_lines else None,
            "maliyetKapsam": cost_lines, "marj": marj,
            "maliyetNotu": None if cost_lines else "Maliyet bilinmiyor: birim maliyet kaynağı bu kurulumda bağlı değil."}


def score(items: list[dict[str, Any]], checklist: list[dict[str, Any]], deadline: Optional[str], cfg: dict[str, Any],
          asof: Optional[date] = None) -> dict[str, Any]:
    """Uygunluk puanı ve parçaları (0–100). Ölçülemeyen parça dışarıda; ağırlıklar kalanlara göre yeniden dağılır."""
    parts: dict[str, Optional[float]] = {"eslesme": None, "stok": None, "belge": None, "sure": None}
    if items:
        matched = [i for i in items if i["eslesme_durumu"] == "eslesti"]
        parts["eslesme"] = len(matched) / len(items)
        with_stock = [i for i in matched if i.get("stok") is not None and i.get("adet") is not None]
        if with_stock:
            parts["stok"] = sum(1 for i in with_stock if i["stok"] >= i["adet"]) / len(with_stock)
    req = [c for c in checklist if c.get("zorunlu")]
    if req:
        parts["belge"] = sum(1 for c in req if c.get("durum") == "var") / len(req)
    left = days_left(deadline, asof)
    if left is not None:
        parts["sure"] = max(0.0, min(1.0, left / cfg["fullDays"]))
    w = cfg["weights"]
    used = {k: v for k, v in parts.items() if v is not None and w.get(k, 0) > 0}
    total_w = sum(w[k] for k in used)
    value = round(100 * sum(w[k] * v for k, v in used.items()) / total_w, 1) if total_w else None
    return {"puan": value, "parcalar": {k: (round(v, 4) if v is not None else None) for k, v in parts.items()},
            "agirliklar": w, "kalanGun": left}


# ------------------------------------------------------------------ satır biçimleri


def _tender_row(c: Any, tenant: str, tid: str) -> Any:
    row = c.execute(sa.select(TENDERS).where(TENDERS.c.id == tid, TENDERS.c.tenant_id == tenant)).first()
    if row is None:
        raise TenderError("İhale bulunamadı.", 404)
    return row


def _item_dict(r: Any) -> dict[str, Any]:
    d = dict(r._mapping)
    d["adaylar"] = _j(d.pop("aday_json", None), [])
    return d


def _item_out(d: dict[str, Any]) -> dict[str, Any]:
    nl = net_list(d)
    line = round(d["adet"] * d["onerilen_fiyat"], 2) if d.get("adet") is not None and d.get("onerilen_fiyat") is not None else None
    return {"sira": d["sira"], "metin": d["sartname_metni"], "ad": d["ad"], "yazar": d["yazar"], "yayinevi": d["yayinevi"],
            "adet": d["adet"], "isbn": d["isbn"], "stokKodu": d["eslesen_stok_kodu"], "eslesenAd": d["eslesen_ad"],
            "yontem": d["eslesme_yontemi"], "yontemAdi": MATCH_METHODS.get(d["eslesme_yontemi"] or "", None),
            "durum": d["eslesme_durumu"], "durumAdi": MATCH_STATES.get(d["eslesme_durumu"]), "olasilik": d["olasilik"],
            "adaylar": d.get("adaylar") or [], "stok": d["stok"],
            "stokYetersiz": d["stok"] is not None and d["adet"] is not None and d["stok"] < d["adet"],
            "listeFiyati": d["liste_fiyati"], "listeFiyatiKdvHaric": round(nl, 2) if nl is not None else None,
            "fiyatKaynagi": d["fiyat_kaynagi"], "kdvOrani": d["kdv_orani"], "logoFiyati": d["logo_fiyati"],
            "logoFiyatNotu": d["logo_fiyat_notu"], "onerilenFiyat": d["onerilen_fiyat"], "fiyatElle": bool(d["fiyat_elle"]),
            "tutar": line, "tahminiMaliyet": d["tahmini_maliyet"], "maliyetKaynagi": d["maliyet_kaynagi"],
            "marj": ((d["onerilen_fiyat"] - d["tahmini_maliyet"]) / d["onerilen_fiyat"])
            if d.get("tahmini_maliyet") is not None and d.get("onerilen_fiyat") else None,
            "onaylayan": d["onaylayan"], "onayZamani": _iso(d["onay_zamani"]), "not": d["notu"]}


def _check_out(r: Any, docs: dict[str, dict[str, Any]], asof: date) -> dict[str, Any]:
    d = dict(r._mapping)
    doc = docs.get(d["belge_id"] or "")
    valid = d["gecerlilik_tarihi"] or (doc or {}).get("gecerlilik_tarihi")
    state = d["durum"]
    if state == "var" and valid and valid < asof.isoformat():
        state = "gecersiz"
    return {"id": d["id"], "sira": d["sira"], "kalem": d["kalem"], "belgeTuru": d["belge_turu"],
            "belgeTuruAdi": DOC_TYPES.get(d["belge_turu"] or "", None), "zorunlu": bool(d["zorunlu"]), "durum": state,
            "durumAdi": CHECK_STATES.get(state), "belgeId": d["belge_id"], "belgeAdi": (doc or {}).get("ad"),
            "gecerlilik": valid, "kaynak": d["kaynak"], "kaynakCumle": d["kaynak_cumle"], "not": d["notu"]}


def _doc_out(r: Any, asof: date, warn_days: int) -> dict[str, Any]:
    d = dict(r._mapping)
    v = d["gecerlilik_tarihi"]
    left = (date.fromisoformat(v) - asof).days if v else None
    state = "suresiz" if left is None else ("doldu" if left < 0 else ("yaklasti" if left <= warn_days else "gecerli"))
    return {"id": d["id"], "ad": d["ad"], "tur": d["tur"], "turAdi": DOC_TYPES.get(d["tur"], d["tur"]), "gecerlilik": v,
            "kalanGun": left, "durum": state, "dosyaAdi": d["dosya_adi"], "boyut": d["boyut"], "dosyaVar": bool(d["dosya_yolu"]),
            "not": d["notu"], "yukleyen": d["yukleyen"], "zaman": _iso(d["zaman"])}


def _decision_out(r: Any) -> dict[str, Any]:
    d = dict(r._mapping)
    return {"id": d["id"], "karar": d["karar"], "kararAdi": DECISIONS_LABEL.get(d["karar"]), "durum": d["durum"],
            "teklifToplami": d["teklif_toplami"], "fiyatOrani": d["fiyat_orani"], "gerekce": d["gerekce"],
            "ozet": _j(d["ozet_json"], {}), "oneren": d["oneren"], "oneriZamani": _iso(d["oneri_zamani"]),
            "onaylayan": d["onaylayan"], "onayZamani": _iso(d["onay_zamani"]), "onayNotu": d["onay_notu"]}


def _result_out(r: Any) -> Optional[dict[str, Any]]:
    if r is None:
        return None
    d = dict(r._mapping)
    ratio = d["kazanan_fiyat"] / d["liste_toplami"] if d["kazanan_fiyat"] is not None and d["liste_toplami"] else None
    return {"sonuc": d["sonuc"], "sonucAdi": RESULT_STATES.get(d["sonuc"]), "kazanan": d["kazanan"],
            "kazananFiyat": d["kazanan_fiyat"], "bizimFiyat": d["bizim_fiyat"], "listeToplami": d["liste_toplami"],
            "kazananListeOrani": ratio, "neden": d["neden"], "kaynak": d["kaynak"], "kaydeden": d["kaydeden"],
            "zaman": _iso(d["zaman"])}


def _tender_out(r: Any, asof: Optional[date] = None) -> dict[str, Any]:
    d = dict(r._mapping)
    return {"id": d["id"], "kaynak": d["kaynak"], "kaynakAdi": KAYNAKLAR.get(d["kaynak"]), "kaynakNo": d["kaynak_no"],
            "kurum": d["kurum"], "kurumTuru": d["kurum_turu"], "kurumTuruAdi": KURUM_TURLERI.get(d["kurum_turu"]),
            "il": d["il"], "konu": d["konu"], "usul": d["usul"], "usulAdi": USULLER.get(d["usul"] or ""),
            "yaklasikTutar": d["yaklasik_tutar"], "ilanTarihi": d["ilan_tarihi"], "sonTeklifTarihi": d["son_teklif_tarihi"],
            "kalanGun": days_left(d["son_teklif_tarihi"], asof), "teslimSuresi": d["teslim_suresi"],
            "teminatTutari": d["teminat_tutari"], "teminatIadeTarihi": d["teminat_iade_tarihi"], "yetkili": d["yetkili"],
            "durum": d["durum"], "durumAdi": STATUSES.get(d["durum"]), "sorumlu": d["sorumlu"],
            "fiyatOrani": d["fiyat_orani"], "fiyatOraniKaynak": d["fiyat_orani_kaynak"],
            "uygunlukPuani": d["uygunluk_puani"], "uygunluk": _j(d["uygunluk_json"], {}),
            "kitapIlani": _j(d["kitap_ilani_json"], {}), "ozet": _j(d["ozet_json"], {}), "kararMetni": d["karar_metni"],
            "notlar": d["notlar"], "olusturan": d["created_by"], "olusturma": _iso(d["created_at"]),
            "guncelleyen": d["updated_by"], "guncelleme": _iso(d["updated_at"])}


# ------------------------------------------------------------------ ihale kaydı


_FIELDS = {  # gövde anahtarı → (kolon, dönüştürücü)
    "kaynakNo": ("kaynak_no", lambda v: _text(v, 80)),
    "kurum": ("kurum", lambda v: _text(v, 300)),
    "il": ("il", lambda v: _text(v, 60)),
    "konu": ("konu", lambda v: _text(v, 4000)),
    "yaklasikTutar": ("yaklasik_tutar", lambda v: _num(v, "Yaklaşık tutar")),
    "ilanTarihi": ("ilan_tarihi", lambda v: _day(v, "İlan tarihi")),
    "sonTeklifTarihi": ("son_teklif_tarihi", _deadline),
    "teslimSuresi": ("teslim_suresi", lambda v: _text(v, 200)),
    "teminatTutari": ("teminat_tutari", lambda v: _num(v, "Teminat tutarı")),
    "teminatIadeTarihi": ("teminat_iade_tarihi", lambda v: _day(v, "Teminat iade tarihi")),
    "yetkili": ("yetkili", lambda v: _text(v, 300)),
    "sorumlu": ("sorumlu", lambda v: _text(v, 120)),
    "notlar": ("notlar", lambda v: _text(v, 8000)),
}


def _choice(v: Any, allowed: dict[str, str], label: str) -> str:
    s = str(v or "").strip()
    if s not in allowed:
        raise TenderError(f"{label} geçersiz: {', '.join(allowed)} olmalı.")
    return s


def create(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    vals: dict[str, Any] = {}
    for key, (col, conv) in _FIELDS.items():
        if key in body:
            vals[col] = conv(body.get(key))
    if not vals.get("kurum"):
        raise TenderError("Kurum adı girilmeli.")
    if not vals.get("konu"):
        raise TenderError("İhale konusu girilmeli.")
    vals["kurum_turu"] = _choice(body.get("kurumTuru"), KURUM_TURLERI, "Kurum türü")
    if body.get("usul"):
        vals["usul"] = _choice(body.get("usul"), USULLER, "Usul")
    kaynak = str(body.get("kaynak") or "elle")
    if kaynak not in ("elle", "dosya"):
        raise TenderError("İlan elle ya da dosyadan girilir; resmî kaynaktan içe alma bu sürümde yok.")
    tid = uuid.uuid4().hex
    now = _now()
    cfg = settings()
    ratio = price_ratio(engine, tenant, vals["kurum_turu"], cfg["historyMin"])
    owner = vals.pop("sorumlu", None) or user
    with engine.begin() as c:
        c.execute(TENDERS.insert().values(
            id=tid, tenant_id=tenant, kaynak=kaynak, durum="yeni", sorumlu=owner,
            fiyat_orani=ratio["oran"], fiyat_orani_kaynak=ratio["kaynak"], created_by=user, created_at=now,
            updated_by=user, updated_at=now, **vals))
    rescore(engine, tenant, tid)
    return detail(engine, tenant, tid)


def update(engine: sa.engine.Engine, tenant: str, user: str, tid: str, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    vals: dict[str, Any] = {}
    for key, (col, conv) in _FIELDS.items():
        if key in body:
            vals[col] = conv(body.get(key))
    if "kurum" in vals and not vals["kurum"]:
        raise TenderError("Kurum adı boş olamaz.")
    if "konu" in vals and not vals["konu"]:
        raise TenderError("İhale konusu boş olamaz.")
    if "kurumTuru" in body:
        vals["kurum_turu"] = _choice(body.get("kurumTuru"), KURUM_TURLERI, "Kurum türü")
    if "usul" in body:
        vals["usul"] = _choice(body.get("usul"), USULLER, "Usul") if body.get("usul") else None
    reprice = False
    if "fiyatOrani" in body:
        r = _num(body.get("fiyatOrani"), "Fiyat oranı", allow_none=False)
        if not 0 < r <= 5:
            raise TenderError("Fiyat oranı 0 ile 5 arasında olmalı (1 = liste fiyatı).")
        vals["fiyat_orani"] = round(r, 4)
        vals["fiyat_orani_kaynak"] = f"Elle girildi ({user})"
        reprice = True
    with engine.begin() as c:
        row = _tender_row(c, tenant, tid)
        if "durum" in body:
            st = _choice(body.get("durum"), STATUSES, "Durum")
            if st != row.durum:
                if st not in MANUAL_STATUSES:
                    raise TenderError("Başvuru kararı karar akışıyla, kazanıldı/kaybedildi sonuç kaydıyla girilir.", 409)
                if st == "teklif_verildi" and row.durum != "basvurulacak":
                    raise TenderError("Teklif verildi ancak onaylanmış başvuru kararından sonra işaretlenir.", 409)
                if row.durum in ("kazanildi", "kaybedildi") and st != "iptal":
                    raise TenderError("Sonuçlanmış ihalenin durumu sonuç kaydından değişir.", 409)
                vals["durum"] = st
        if reprice and c.execute(sa.select(DECISIONS.c.id).where(DECISIONS.c.tender_id == tid, DECISIONS.c.durum == "onayda")).first():
            raise TenderError("Onay bekleyen karar varken fiyat oranı değişmez; önce kararı geri çekin.", 409)
        diff = {k: [getattr(row, k), v] for k, v in vals.items() if getattr(row, k) != v}
        if vals:
            c.execute(TENDERS.update().where(TENDERS.c.id == tid).values(updated_by=user, updated_at=_now(), **vals))
        if reprice:
            _reprice(c, tid, vals["fiyat_orani"])
    rescore(engine, tenant, tid)
    return detail(engine, tenant, tid), diff


def _reprice(c: Any, tid: str, ratio: float) -> None:
    for r in c.execute(sa.select(ITEMS).where(ITEMS.c.tender_id == tid, ITEMS.c.fiyat_elle.is_(False))).all():
        c.execute(ITEMS.update().where(ITEMS.c.tender_id == tid, ITEMS.c.sira == r.sira)
                  .values(onerilen_fiyat=suggested_price(dict(r._mapping), ratio)))


def delete(engine: sa.engine.Engine, tenant: str, tid: str) -> dict[str, Any]:
    """Yalnız karar ve sonuç kaydı olmayan ihale silinir (yanlış giriş). Dosyaları da silinir."""
    with engine.begin() as c:
        row = _tender_row(c, tenant, tid)
        if c.execute(sa.select(DECISIONS.c.id).where(DECISIONS.c.tender_id == tid)).first() or \
                c.execute(sa.select(RESULTS.c.tender_id).where(RESULTS.c.tender_id == tid)).first():
            raise TenderError("Karar ya da sonuç kaydı olan ihale silinmez; durumunu «İptal» yapın.", 409)
        paths = [r.dosya_yolu for r in c.execute(sa.select(FILES.c.dosya_yolu).where(FILES.c.tender_id == tid))]
        for t, col in ((FILES, FILES.c.tender_id), (ITEMS, ITEMS.c.tender_id), (CHECKLIST, CHECKLIST.c.tender_id),
                       (JOBS, JOBS.c.tender_id)):
            c.execute(t.delete().where(col == tid))
        c.execute(TENDERS.delete().where(TENDERS.c.id == tid))
    for p in paths:
        try:
            Path(p).unlink(missing_ok=True)
        except OSError as e:
            log.warning("ihale dosyası silinemedi (%s): %s", p, e)
    return {"id": tid, "kurum": row.kurum}


def rescore(engine: sa.engine.Engine, tenant: str, tid: str) -> None:
    cfg = settings()
    with engine.begin() as c:
        row = _tender_row(c, tenant, tid)
        items = [_item_dict(r) for r in c.execute(sa.select(ITEMS).where(ITEMS.c.tender_id == tid))]
        docs = _docs_map(c, tenant)
        checks = [_check_out(r, docs, today()) for r in c.execute(sa.select(CHECKLIST).where(CHECKLIST.c.tender_id == tid))]
        s = score(items, checks, row.son_teklif_tarihi, cfg)
        c.execute(TENDERS.update().where(TENDERS.c.id == tid).values(uygunluk_puani=s["puan"], uygunluk_json=_dump(s)))


def _docs_map(c: Any, tenant: str) -> dict[str, dict[str, Any]]:
    return {r.id: dict(r._mapping) for r in c.execute(sa.select(DOCUMENTS).where(DOCUMENTS.c.tenant_id == tenant))}


def list_tenders(engine: sa.engine.Engine, tenant: str, *, durum: str = "acik", il: str = "", kurum_turu: str = "",
                 q: str = "", son: str = "") -> dict[str, Any]:
    """Liste: süzgeç durum (acik | kapali | hepsi | tek durum), il, kurum türü, arama, son teklif tarihi üst sınırı.
    Sayı tavanı yok; son teklif tarihine göre sıralı (tarihsiz en sonda)."""
    stmt = sa.select(TENDERS).where(TENDERS.c.tenant_id == tenant)
    if durum == "acik":
        stmt = stmt.where(TENDERS.c.durum.in_(OPEN_STATUSES))
    elif durum == "kapali":
        stmt = stmt.where(TENDERS.c.durum.not_in(OPEN_STATUSES))
    elif durum in STATUSES:
        stmt = stmt.where(TENDERS.c.durum == durum)
    if il:
        stmt = stmt.where(TENDERS.c.il == il)
    if kurum_turu:
        stmt = stmt.where(TENDERS.c.kurum_turu == kurum_turu)
    if son:
        stmt = stmt.where(TENDERS.c.son_teklif_tarihi <= _day(son, "Son tarih") + "T23:59")
    with engine.connect() as c:
        rows = c.execute(stmt).all()
        counts = {r.tender_id: r for r in c.execute(
            sa.select(ITEMS.c.tender_id, sa.func.count().label("n"),
                      sa.func.sum(sa.case((ITEMS.c.eslesme_durumu == "eslesti", 1), else_=0)).label("ok"))
            .group_by(ITEMS.c.tender_id)).all()}
        pending = {r.tender_id for r in c.execute(sa.select(DECISIONS.c.tender_id).where(DECISIONS.c.durum == "onayda"))}
        all_rows = c.execute(sa.select(TENDERS.c.il, TENDERS.c.durum).where(TENDERS.c.tenant_id == tenant)).all()
    qf = fold(q)
    out = []
    asof = today()
    for r in rows:
        if qf and qf not in fold(" ".join(str(x or "") for x in (r.kurum, r.konu, r.kaynak_no, r.il))):
            continue
        d = _tender_out(r, asof)
        d.pop("ozet", None)
        n = counts.get(r.id)
        d["kalem"] = int(n.n) if n else 0
        d["eslesen"] = int(n.ok or 0) if n else 0
        d["onayBekliyor"] = r.id in pending
        out.append(d)
    out.sort(key=lambda d: (d["sonTeklifTarihi"] is None, d["sonTeklifTarihi"] or "", d["kurum"]))
    statuses: dict[str, int] = {}
    for r in all_rows:
        statuses[r.durum] = statuses.get(r.durum, 0) + 1
    return {"items": out, "total": len(out), "iller": sorted({r.il for r in all_rows if r.il}), "durumSayilari": statuses}


def detail(engine: sa.engine.Engine, tenant: str, tid: str) -> dict[str, Any]:
    asof = today()
    cfg = settings()
    with engine.connect() as c:
        row = _tender_row(c, tenant, tid)
        items = [_item_dict(r) for r in c.execute(sa.select(ITEMS).where(ITEMS.c.tender_id == tid).order_by(ITEMS.c.sira))]
        docs = _docs_map(c, tenant)
        checks = [_check_out(r, docs, asof) for r in
                  c.execute(sa.select(CHECKLIST).where(CHECKLIST.c.tender_id == tid).order_by(CHECKLIST.c.sira))]
        files = [{"id": f.id, "tur": f.tur, "ad": f.ad, "boyut": f.boyut, "mime": f.mime, "yukleyen": f.yukleyen,
                  "zaman": _iso(f.zaman)} for f in c.execute(sa.select(FILES).where(FILES.c.tender_id == tid).order_by(FILES.c.zaman))]
        decisions = [_decision_out(r) for r in c.execute(
            sa.select(DECISIONS).where(DECISIONS.c.tender_id == tid).order_by(DECISIONS.c.oneri_zamani.desc()))]
        result = _result_out(c.execute(sa.select(RESULTS).where(RESULTS.c.tender_id == tid)).first())
        jobs = [_job_out(r) for r in c.execute(sa.select(JOBS).where(JOBS.c.tender_id == tid).order_by(JOBS.c.baslangic.desc()))]
    out = _tender_out(row, asof)
    tot = totals(items)
    out.update(kalemler=[_item_out(i) for i in items], toplamlar=tot, kontrolListesi=checks, dosyalar=files,
               kararlar=decisions, sonuc=result, isler=_live_jobs(jobs),
               kararOzeti=brief_facts(out, tot, checks, history(engine, tenant, row.kurum, row.kurum_turu)),
               ayarlar={k: cfg[k] for k in ("autoProb", "autoMargin", "suggestProb", "suggestMargin", "candidates", "priceSource")})
    return out


def history(engine: sa.engine.Engine, tenant: str, kurum: str, kurum_turu: str) -> dict[str, Any]:
    """Aynı kurumun ve aynı kurum türünün geçmiş sonuçları: sayı, kazanılan, kazanan fiyat/liste oranı ortancası."""
    with engine.connect() as c:
        rows = c.execute(sa.select(RESULTS, TENDERS.c.kurum, TENDERS.c.kurum_turu)
                         .join(TENDERS, TENDERS.c.id == RESULTS.c.tender_id).where(TENDERS.c.tenant_id == tenant)).all()

    def summ(rs: list[Any]) -> dict[str, Any]:
        ratios = [r.kazanan_fiyat / r.liste_toplami for r in rs if r.kazanan_fiyat is not None and r.liste_toplami]
        return {"sonuc": len(rs), "kazanilan": sum(1 for r in rs if r.sonuc == "kazanildi"),
                "kaybedilen": sum(1 for r in rs if r.sonuc == "kaybedildi"),
                "kazananOranOrtanca": round(statistics.median(ratios), 4) if ratios else None, "oranSayisi": len(ratios)}

    return {"kurum": summ([r for r in rows if fold(r.kurum) == fold(kurum)]),
            "kurumTuru": summ([r for r in rows if r.kurum_turu == kurum_turu])}


def brief_facts(t: dict[str, Any], tot: dict[str, Any], checks: list[dict[str, Any]], hist: dict[str, Any]) -> dict[str, Any]:
    """Tek sayfa karar özetinin rakamları (model yok). Zeki AI metni yalnız bu rakamları kullanır."""
    missing = [c["kalem"] for c in checks if c["zorunlu"] and c["durum"] != "var"]
    risks = []
    if t.get("kalanGun") is not None and t["kalanGun"] < 0:
        risks.append("Son teklif tarihi geçti.")
    elif t.get("kalanGun") is not None and t["kalanGun"] <= 2:
        risks.append(f"Son teklif tarihine {t['kalanGun']} gün kaldı.")
    if missing:
        risks.append(f"{len(missing)} zorunlu belge eksik ya da geçersiz.")
    if tot["stokYetersiz"]:
        risks.append(f"{tot['stokYetersiz']} eşleşen kalemde stok istenen adetten az.")
    if tot["durumlar"].get("oneri") or tot["durumlar"].get("belirsiz") or tot["durumlar"].get("bekliyor"):
        risks.append("Onay ya da seçim bekleyen eşleştirme var.")
    if tot["disarida"]:
        risks.append(f"{tot['disarida']} eşleşen kalemin adedi ya da fiyatı eksik; toplama girmedi.")
    return {"yaklasikTutar": t.get("yaklasikTutar"), "uygunlukPuani": t.get("uygunlukPuani"), "kalem": tot["kalem"],
            "eslesen": tot["durumlar"].get("eslesti", 0), "katalogdaYok": tot["durumlar"].get("yok", 0),
            "stokYetersiz": tot["stokYetersiz"], "teklifAraToplam": tot["araToplam"], "teklifGenelToplam": tot["genelToplam"],
            "listeToplami": tot["listeToplami"], "fiyatOrani": t.get("fiyatOrani"), "marj": tot["marj"],
            "maliyetKapsam": tot["maliyetKapsam"], "maliyetNotu": tot["maliyetNotu"], "eksikBelgeler": missing,
            "teminatTutari": t.get("teminatTutari"), "kalanGun": t.get("kalanGun"), "gecmis": hist, "riskler": risks}


# ------------------------------------------------------------------ kalemler


def import_items(engine: sa.engine.Engine, tenant: str, user: str, tid: str, parsed: dict[str, Any], mode: str) -> dict[str, Any]:
    if mode not in ("replace", "append"):
        raise TenderError("İçe aktarma biçimi replace ya da append olmalı.")
    if not parsed["items"]:
        raise TenderError("Listede kitap kalemi bulunamadı." + (f" Okunamayan satır: {len(parsed['skipped'])}." if parsed["skipped"] else ""))
    with engine.begin() as c:
        row = _tender_row(c, tenant, tid)
        if row.durum not in ("yeni", "inceleniyor"):
            raise TenderError("Kalem listesi yalnız karar verilmeden önce değişir.", 409)
        if c.execute(sa.select(DECISIONS.c.id).where(DECISIONS.c.tender_id == tid, DECISIONS.c.durum == "onayda")).first():
            raise TenderError("Onay bekleyen karar varken kalem listesi değişmez.", 409)
        if _running(c, tid):
            raise TenderError("Eşleştirme sürüyor; bitince yeniden deneyin.", 409)
        start = 0
        if mode == "replace":
            c.execute(ITEMS.delete().where(ITEMS.c.tender_id == tid))
        else:
            start = c.execute(sa.select(sa.func.coalesce(sa.func.max(ITEMS.c.sira), 0)).where(ITEMS.c.tender_id == tid)).scalar() or 0
        for n, it in enumerate(parsed["items"], start=start + 1):
            c.execute(ITEMS.insert().values(
                tender_id=tid, sira=n, sartname_metni=(it.get("metin") or it.get("ad") or "")[:4000],
                ad=(it.get("ad") or None) and it["ad"][:500], yazar=(it.get("yazar") or None) and it["yazar"][:300],
                yayinevi=(it.get("yayinevi") or None) and it["yayinevi"][:200], adet=it.get("adet"), isbn=it.get("isbn"),
                eslesme_durumu="bekliyor", fiyat_elle=False))
        if row.durum == "yeni":
            c.execute(TENDERS.update().where(TENDERS.c.id == tid).values(durum="inceleniyor", updated_by=user, updated_at=_now()))
    rescore(engine, tenant, tid)
    return {"eklenen": len(parsed["items"]), "okunamayan": parsed["skipped"], "baslik": parsed.get("header") or []}


def update_item(engine: sa.engine.Engine, tenant: str, user: str, tid: str, sira: int, body: dict[str, Any],
                enrich: Callable[[list[str]], dict[str, dict[str, Any]]]) -> tuple[dict[str, Any], dict[str, Any]]:
    """İnsan düzeltmesi: eşleşen kitabı seçme/kaldırma, öneriyi onaylama, adet, birim teklif fiyatı, not."""
    with engine.begin() as c:
        row = _tender_row(c, tenant, tid)
        if row.durum not in ("yeni", "inceleniyor"):
            raise TenderError("Kalemler yalnız karar verilmeden önce düzeltilir.", 409)
        if c.execute(sa.select(DECISIONS.c.id).where(DECISIONS.c.tender_id == tid, DECISIONS.c.durum == "onayda")).first():
            raise TenderError("Onay bekleyen karar varken kalem değişmez; önce kararı geri çekin.", 409)
        it = c.execute(sa.select(ITEMS).where(ITEMS.c.tender_id == tid, ITEMS.c.sira == sira)).first()
        if it is None:
            raise TenderError("Kalem bulunamadı.", 404)
        cur = dict(it._mapping)
        vals: dict[str, Any] = {}
        if "stokKodu" in body:
            code = str(body.get("stokKodu") or "").strip()
            if not code:
                vals.update(eslesen_stok_kodu=None, eslesen_ad=None, eslesme_yontemi="elle", eslesme_durumu="yok", olasilik=None,
                            stok=None, liste_fiyati=None, fiyat_kaynagi=None, logo_fiyati=None, logo_fiyat_notu=None,
                            onerilen_fiyat=None, tahmini_maliyet=None, maliyet_kaynagi=None, kdv_orani=None)
            else:
                info = enrich([code]).get(code)
                if not info:
                    raise TenderError(f"{code} stok kodu katalogda yok.")
                keep_method = code == cur["eslesen_stok_kodu"] and cur["eslesme_yontemi"]
                vals.update(eslesen_stok_kodu=code, eslesen_ad=info.get("ad"), eslesme_durumu="eslesti",
                            eslesme_yontemi=keep_method or "elle", **_enriched_cols(info))
                vals["onerilen_fiyat"] = cur["onerilen_fiyat"] if cur["fiyat_elle"] else suggested_price({**cur, **vals}, row.fiyat_orani or 1.0)
            vals.update(onaylayan=user, onay_zamani=_now())
        if body.get("onayla"):
            if cur["eslesme_durumu"] not in ("oneri", "belirsiz") or not cur["eslesen_stok_kodu"]:
                raise TenderError("Onaylanacak eşleştirme önerisi yok.", 409)
            vals.update(eslesme_durumu="eslesti", onaylayan=user, onay_zamani=_now())
        if "adet" in body:
            vals["adet"] = _num(body.get("adet"), "Adet")
        if "onerilenFiyat" in body:
            p = _num(body.get("onerilenFiyat"), "Birim teklif fiyatı")
            vals["onerilen_fiyat"] = round(p, 2) if p is not None else suggested_price({**cur, **vals}, row.fiyat_orani or 1.0)
            vals["fiyat_elle"] = p is not None
        if "not" in body:
            vals["notu"] = _text(body.get("not"), 500)
        if not vals:
            raise TenderError("Değişiklik yok.")
        diff = {k: [cur.get(k), v] for k, v in vals.items() if cur.get(k) != v and k not in ("onay_zamani",)}
        c.execute(ITEMS.update().where(ITEMS.c.tender_id == tid, ITEMS.c.sira == sira).values(**vals))
        out = _item_out(_item_dict(c.execute(sa.select(ITEMS).where(ITEMS.c.tender_id == tid, ITEMS.c.sira == sira)).first()))
    rescore(engine, tenant, tid)
    return out, diff


def _enriched_cols(info: dict[str, Any]) -> dict[str, Any]:
    return {"stok": info.get("stok"), "liste_fiyati": info.get("liste_fiyati"), "fiyat_kaynagi": info.get("fiyat_kaynagi"),
            "kdv_orani": info.get("kdv_orani"), "logo_fiyati": info.get("logo_fiyati"), "logo_fiyat_notu": info.get("logo_fiyat_notu"),
            "tahmini_maliyet": info.get("maliyet"), "maliyet_kaynagi": info.get("maliyet_kaynagi")}


def enrich_codes(cat: Catalog, codes: list[str], stock: dict[str, float], prices: dict[str, dict[str, Any]],
                 costs: dict[str, dict[str, Any]], cfg: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Stok kodu → ekranda gösterilecek stok, liste fiyatı (seçilen kaynak yazılı), KDV, Logo fiyatı, maliyet."""
    out = {}
    for code in codes:
        b = cat.by_code.get(code)
        if b is None:
            continue
        lp = prices.get(code) or {}
        crm_price = b.get("liste_fiyati")
        vat = vat_rate(b.get("kdv_orani"), cfg["defaultVat"])
        logo_price = lp.get("fiyat")
        logo_gross = None if logo_price is None else (logo_price if lp.get("kdvDahil") in (True, None) else logo_price * (1 + vat))
        if cfg["priceSource"] == "logo" and logo_gross is not None:
            price, src = logo_gross, f"Logo {lp.get('liste')}" + (" (KDV eklendi)" if lp.get("kdvDahil") is False else "")
        elif crm_price is not None:
            price, src = crm_price, "CRM KDV dahil liste fiyatı"
        elif logo_gross is not None:
            price, src = logo_gross, f"Logo {lp.get('liste')} (CRM'de fiyat yok)"
        else:
            price, src = None, "Fiyat bulunamadı"
        cost = costs.get(code)
        out[code] = {"ad": b.get("ad"), "stok": stock.get(code), "liste_fiyati": price, "fiyat_kaynagi": src, "kdv_orani": vat,
                     "logo_fiyati": logo_price,
                     "logo_fiyat_notu": (f"{lp.get('liste')}; geçerli liste {lp.get('gecerliListe')}" if lp else None),
                     "maliyet": cost["maliyet"] if cost else None, "maliyet_kaynagi": cost["kaynak"] if cost else None}
    return out


# ------------------------------------------------------------------ uzun işler (özet, eşleştirme)


_active: set[str] = set()   # bu süreçte koşan iş kimlikleri


def _job_out(r: Any) -> dict[str, Any]:
    d = dict(r._mapping)
    return {"id": d["id"], "tur": d["tur"], "durum": d["durum"], "ilerleme": d["ilerleme"], "toplam": d["toplam"],
            "sonuc": _j(d["sonuc_json"], {}), "hata": d["hata"], "baslatan": d["baslatan"],
            "baslangic": _iso(d["baslangic"]), "bitis": _iso(d["bitis"])}


def _live_jobs(jobs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Köprü yeniden başladıysa yarım kalan iş «kesildi» görünür (kayıt yazılmaz; yeni iş açılabilir)."""
    for j in jobs:
        if j["durum"] in ("sirada", "calisiyor") and j["id"] not in _active:
            j["durum"] = "kesildi"
            j["hata"] = j["hata"] or "İş yarıda kaldı (hizmet yeniden başladı); yeniden başlatın."
    return jobs


def _running(c: Any, tid: str) -> bool:
    ids = [r.id for r in c.execute(sa.select(JOBS.c.id).where(JOBS.c.tender_id == tid, JOBS.c.durum.in_(("sirada", "calisiyor"))))]
    return any(i in _active for i in ids)


def start_job(engine: sa.engine.Engine, tenant: str, user: str, tid: str, kind: str,
              work: Callable[[Callable[[int, int], None]], dict[str, Any]]) -> dict[str, Any]:
    """İşi arka plan iş parçacığında başlatır. Aynı ihalede aynı anda tek iş."""
    jid = uuid.uuid4().hex
    with engine.begin() as c:
        _tender_row(c, tenant, tid)
        if _running(c, tid):
            raise TenderError("Bu ihalede bir iş sürüyor; bitmesini bekleyin.", 409)
        c.execute(JOBS.insert().values(id=jid, tender_id=tid, tur=kind, durum="calisiyor", ilerleme=0, toplam=0,
                                       baslatan=user, baslangic=_now()))
    _active.add(jid)

    def progress(done: int, total: int) -> None:
        with engine.begin() as c:
            c.execute(JOBS.update().where(JOBS.c.id == jid).values(ilerleme=done, toplam=total))

    def run() -> None:
        try:
            res = work(progress)
            with engine.begin() as c:
                c.execute(JOBS.update().where(JOBS.c.id == jid).values(durum="bitti", sonuc_json=_dump(res), bitis=_now()))
        except Exception as e:  # noqa: BLE001 — iş kaydına düz cümle
            log.exception("ihale işi %s (%s) hata verdi", jid, kind)
            msg = str(e) if isinstance(e, TenderError) else f"İş tamamlanamadı: {str(e)[:300]}"
            with engine.begin() as c:
                c.execute(JOBS.update().where(JOBS.c.id == jid).values(durum="hata", hata=msg, bitis=_now()))
        finally:
            _active.discard(jid)

    threading.Thread(target=run, name=f"tender-{kind}-{jid[:6]}", daemon=True).start()
    with engine.connect() as c:
        return _job_out(c.execute(sa.select(JOBS).where(JOBS.c.id == jid)).first())


def job(engine: sa.engine.Engine, tenant: str, tid: str, jid: str) -> dict[str, Any]:
    with engine.connect() as c:
        _tender_row(c, tenant, tid)
        r = c.execute(sa.select(JOBS).where(JOBS.c.id == jid, JOBS.c.tender_id == tid)).first()
    if r is None:
        raise TenderError("İş bulunamadı.", 404)
    return _live_jobs([_job_out(r)])[0]


def run_match(engine: sa.engine.Engine, tenant: str, tid: str, cat: Catalog, choose: Optional[Callable[[str, list[str]], Any]],
              enrich: Callable[[list[str]], dict[str, dict[str, Any]]], progress: Callable[[int, int], None],
              only_pending: bool = False) -> dict[str, Any]:
    """Eşleştirme işi: kalem kalem eşleştirir (her kalem bitince yazılır), sonra eşleşen kodların stok/fiyat/maliyetini okur.
    Elle eşleştirilmiş ve onaylanmış kalemlere dokunulmaz."""
    cfg = settings()
    with engine.connect() as c:
        row = _tender_row(c, tenant, tid)
        items = [_item_dict(r) for r in c.execute(sa.select(ITEMS).where(ITEMS.c.tender_id == tid).order_by(ITEMS.c.sira))]
    todo = [i for i in items if not (i["eslesme_yontemi"] == "elle" or (i["onaylayan"] and i["eslesme_durumu"] == "eslesti"))]
    if only_pending:
        todo = [i for i in todo if i["eslesme_durumu"] in ("bekliyor", "belirsiz")]
    total = len(todo)
    progress(0, total)
    stats = {k: 0 for k in MATCH_STATES}
    for n, it in enumerate(todo, start=1):
        m = match_one({"metin": it["sartname_metni"], "ad": it["ad"], "yazar": it["yazar"], "yayinevi": it["yayinevi"],
                       "isbn": it["isbn"]}, cat, choose, cfg)
        stats[m["eslesme_durumu"]] += 1
        with engine.begin() as c:
            c.execute(ITEMS.update().where(ITEMS.c.tender_id == tid, ITEMS.c.sira == it["sira"]).values(
                eslesme_durumu=m["eslesme_durumu"], eslesme_yontemi=m["eslesme_yontemi"],
                eslesen_stok_kodu=m["eslesen_stok_kodu"], eslesen_ad=m["eslesen_ad"], olasilik=m["olasilik"],
                aday_json=_dump(m["adaylar"]), notu=m["not"], onaylayan=None, onay_zamani=None))
        if n % 5 == 0 or n == total:
            progress(n, total)
    with engine.connect() as c:
        rows = [_item_dict(r) for r in c.execute(sa.select(ITEMS).where(ITEMS.c.tender_id == tid))]
    codes = sorted({r["eslesen_stok_kodu"] for r in rows if r["eslesen_stok_kodu"]})
    info = enrich(codes)
    with engine.begin() as c:
        for r in rows:
            code = r["eslesen_stok_kodu"]
            vals = _enriched_cols(info.get(code, {})) if code else _enriched_cols({})
            if not r["fiyat_elle"]:
                vals["onerilen_fiyat"] = suggested_price({**r, **vals}, row.fiyat_orani or 1.0) if code else None
            c.execute(ITEMS.update().where(ITEMS.c.tender_id == tid, ITEMS.c.sira == r["sira"]).values(**vals))
    rescore(engine, tenant, tid)
    return {"kalem": total, "durumlar": stats, "kodlar": len(codes), "katalogNotlari": cat.notes}


_NUM_RX = re.compile(r"\d+(?:[.,]\d+)*")


def _nums(s: str) -> set[str]:
    return {re.sub(r"[.,]", "", m) for m in _NUM_RX.findall(s or "")}


def _quote_ok(value: str, quote: str, text_fold: str) -> bool:
    """Özet maddesi kabulü: alıntı cümle belgede geçmeli, değerdeki her sayı alıntıda bulunmalı."""
    qf = fold(quote)
    if len(qf) < 8 or qf not in text_fold:
        return False
    from semantic_bridge import zeki_text as Z

    return Z.numbers_ok(value, quote)          # tek sayı denetçisi: değerdeki her sayı alıntıda


def summarize_text(text: str, chat: Callable[[list[dict[str, str]]], str], cfg: dict[str, Any],
                   progress: Callable[[int, int], None], reading: Any = None) -> dict[str, Any]:
    """Şartname özeti: metin parçalara bölünür (sessiz kesme yok), her parça için kaynak cümleli JSON istenir; alıntısı
    belgede bulunmayan ya da alıntıda olmayan sayı içeren madde atılır ve sayılır. `reading` (ortak belge okuma)
    verilirse her maddenin alıntısının bulunduğu sayfa ve okuması (metin/OCR) maddeye yazılır; OCR'dan gelen madde
    ekranda «OCR» etiketiyle görünür, alıntı denetimi OCR metninde aynen uygulanır."""
    text = (text or "").strip()
    if not text:
        why = " ".join(getattr(reading, "errors", None) or [])
        raise TenderError("Şartname dosyasından metin okunamadı. " + (why or "Taranmış görüntü olabilir; belge okuma "
                                                                          "servisi bağlıysa yeniden deneyin."))
    size = cfg["summaryChunk"]
    chunks = [text[i:i + size] for i in range(0, len(text), size)]
    tf = fold(text)
    out: dict[str, Any] = {"konu": None, "teslimSuresi": None, "teminat": None, "belgeler": [], "kosullar": [],
                           "atilan": 0, "parca": len(chunks), "karakter": len(text)}
    progress(0, len(chunks))
    sys_msg = ("Kamu ihalesi şartnamelerini okuyan bir yardımcısın. Yalnız verilen metinde yazanı çıkar; tahmin etme, "
               "sayı uydurma. Her madde için metinden AYNEN alıntılanmış kaynak cümleyi ver. Sadece JSON yaz.")
    for n, part in enumerate(chunks, start=1):
        prompt = ("Aşağıdaki şartname bölümünden şu JSON'u çıkar (bulunmayan alan null ya da boş liste):\n"
                  '{"konu": {"deger": "", "kaynak": ""}, "teslim_suresi": {"deger": "", "kaynak": ""}, '
                  '"teminat": {"deger": "", "kaynak": ""}, "belgeler": [{"deger": "istenen belge adı", "kaynak": ""}], '
                  '"kosullar": [{"deger": "kritik koşul (ceza, teslim yeri, numune, yerli malı vb.)", "kaynak": ""}]}\n\n'
                  f"Bölüm {n}/{len(chunks)}:\n{part}")
        raw = chat([{"role": "system", "content": sys_msg}, {"role": "user", "content": prompt}])
        data = _json_from(raw)
        for key, name in (("konu", "konu"), ("teslim_suresi", "teslimSuresi"), ("teminat", "teminat")):
            v = data.get(key)
            if isinstance(v, dict) and v.get("deger"):
                if _quote_ok(str(v["deger"]), str(v.get("kaynak") or ""), tf):
                    out[name] = out[name] or {"deger": str(v["deger"])[:500], "kaynak": str(v["kaynak"])[:600]}
                else:
                    out["atilan"] += 1
        for key, name in (("belgeler", "belgeler"), ("kosullar", "kosullar")):
            for v in data.get(key) or []:
                if not isinstance(v, dict) or not v.get("deger"):
                    continue
                if not _quote_ok(str(v["deger"]), str(v.get("kaynak") or ""), tf):
                    out["atilan"] += 1
                    continue
                if any(fold(x["deger"]) == fold(v["deger"]) for x in out[name]):
                    continue
                out[name].append({"deger": str(v["deger"])[:300], "kaynak": str(v["kaynak"])[:600]})
        progress(n, len(chunks))
    if reading is not None:
        from semantic_bridge import doc_read as DR

        def mark(item: Optional[dict[str, Any]]) -> None:
            if not item:
                return
            hit = DR.find_quote(item.get("kaynak"), reading)
            if hit:
                item["sayfa"], item["okuma"] = hit["sayfa"], hit["okuma"]
                if hit["okuma"] == DR.OCR:
                    item["guven"] = hit.get("guven")

        for key in ("konu", "teslimSuresi", "teminat"):
            mark(out[key])
        for key in ("belgeler", "kosullar"):
            for item in out[key]:
                mark(item)
        out["okuma"] = reading.summary()
    return out


def _json_from(raw: str) -> dict[str, Any]:
    s = (raw or "").strip()
    s = re.sub(r"^```(?:json)?|```$", "", s, flags=re.M).strip()
    m = re.search(r"\{.*\}", s, flags=re.S)
    try:
        v = json.loads(m.group(0) if m else s)
    except ValueError:
        return {}
    return v if isinstance(v, dict) else {}


def doc_type(name: str, choose: Optional[Callable[[str, list[str]], Any]] = None) -> str:
    f = fold(name)
    for key, words in DOC_KEYWORDS:
        if any(w in f for w in words):
            return key
    if choose is not None:
        labels = list(DOC_TYPES.values())
        try:
            r = choose(f"Kamu ihalesinde istenen şu belge hangi türdendir?\n\nBelge: {name}", labels)
            if r.choice and r.confident(0.70, 0.30):
                return list(DOC_TYPES)[labels.index(r.choice)]
        except Exception as e:  # noqa: BLE001
            log.warning("ihale belge türü: model cevap vermedi: %s", e)
    return "diger"


def apply_summary(engine: sa.engine.Engine, tenant: str, user: str, tid: str, summary: dict[str, Any],
                  classify: Optional[dict[str, Any]], choose: Optional[Callable[[str, list[str]], Any]]) -> dict[str, Any]:
    """Özeti ihaleye yazar; belge listesini kontrol listesine ekler (aynı kalem varsa eklenmez), arşivde geçerli aynı
    türden belge varsa «var» işaretler. Teslim süresi boşsa özetten doldurulur."""
    added = 0
    with engine.begin() as c:
        row = _tender_row(c, tenant, tid)
        docs = [d for d in _docs_map(c, tenant).values()]
        existing = {fold(r.kalem) for r in c.execute(sa.select(CHECKLIST.c.kalem).where(CHECKLIST.c.tender_id == tid))}
        n = c.execute(sa.select(sa.func.coalesce(sa.func.max(CHECKLIST.c.sira), 0)).where(CHECKLIST.c.tender_id == tid)).scalar() or 0
        for b in summary.get("belgeler") or []:
            if fold(b["deger"]) in existing:
                continue
            kind = doc_type(b["deger"], choose)
            doc = _best_doc(docs, kind)
            n += 1
            c.execute(CHECKLIST.insert().values(
                id=uuid.uuid4().hex, tender_id=tid, sira=n, kalem=b["deger"][:500], belge_turu=kind, zorunlu=True,
                durum="var" if doc else "eksik", belge_id=doc["id"] if doc else None, kaynak="ozet", kaynak_cumle=b["kaynak"]))
            existing.add(fold(b["deger"]))
            added += 1
        vals: dict[str, Any] = {"ozet_json": _dump(summary), "updated_by": user, "updated_at": _now()}
        if classify is not None:
            vals["kitap_ilani_json"] = _dump(classify)
        if not row.teslim_suresi and summary.get("teslimSuresi"):
            vals["teslim_suresi"] = summary["teslimSuresi"]["deger"][:200]
        c.execute(TENDERS.update().where(TENDERS.c.id == tid).values(**vals))
    rescore(engine, tenant, tid)
    return {"kontrolListesineEklenen": added}


def _best_doc(docs: list[dict[str, Any]], kind: str) -> Optional[dict[str, Any]]:
    if kind == "diger":
        return None
    asof = today().isoformat()
    ok = [d for d in docs if d["tur"] == kind and (not d["gecerlilik_tarihi"] or d["gecerlilik_tarihi"] >= asof)]
    ok.sort(key=lambda d: d["gecerlilik_tarihi"] or "9999", reverse=True)
    return ok[0] if ok else None


def classify_notice(konu: str, choose: Optional[Callable[[str, list[str]], Any]]) -> Optional[dict[str, Any]]:
    """İlan konusu kitap/yayın alımı mı (evet / hayır / belirsiz) ve olasılığı."""
    if choose is None or not (konu or "").strip():
        return None
    labels = ["Evet — kitap ya da yayın alımı", "Hayır — kitap alımı değil", "Belirsiz"]
    try:
        r = choose("Aşağıdaki kamu ihalesi ilanının konusu kitap, yayın ya da kütüphane materyali alımı mı?\n\n"
                   f"Konu: {konu}", labels)
    except Exception as e:  # noqa: BLE001
        log.warning("ihale ilan sınıflaması: model cevap vermedi: %s", e)
        return None
    key = {labels[0]: "evet", labels[1]: "hayir", labels[2]: "belirsiz"}.get(r.choice or "", "belirsiz")
    return {"sonuc": key, "olasilik": r.probability, "yontem": r.method}


# ------------------------------------------------------------------ kontrol listesi


def checklist(engine: sa.engine.Engine, tenant: str, tid: str) -> dict[str, Any]:
    with engine.connect() as c:
        _tender_row(c, tenant, tid)
        docs = _docs_map(c, tenant)
        items = [_check_out(r, docs, today()) for r in
                 c.execute(sa.select(CHECKLIST).where(CHECKLIST.c.tender_id == tid).order_by(CHECKLIST.c.sira))]
    return {"items": items, "turler": DOC_TYPES, "durumlar": CHECK_STATES}


def update_checklist(engine: sa.engine.Engine, tenant: str, user: str, tid: str, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """Toplu düzenleme: `items` listesi; `id` yoksa yeni kalem, `sil: true` siler. Belge bağlanırsa durum «var»."""
    rows = body.get("items")
    if not isinstance(rows, list) or not rows:
        raise TenderError("items listesi gerekli.")
    changes = {"eklenen": 0, "guncellenen": 0, "silinen": 0}
    with engine.begin() as c:
        _tender_row(c, tenant, tid)
        docs = _docs_map(c, tenant)
        n = c.execute(sa.select(sa.func.coalesce(sa.func.max(CHECKLIST.c.sira), 0)).where(CHECKLIST.c.tender_id == tid)).scalar() or 0
        for r in rows:
            if not isinstance(r, dict):
                raise TenderError("Kontrol listesi satırı geçersiz.")
            cid = r.get("id")
            if cid and r.get("sil"):
                changes["silinen"] += c.execute(CHECKLIST.delete().where(CHECKLIST.c.id == cid, CHECKLIST.c.tender_id == tid)).rowcount
                continue
            vals: dict[str, Any] = {}
            if "kalem" in r:
                vals["kalem"] = _text(r.get("kalem"), 500)
                if not vals["kalem"]:
                    raise TenderError("Belge kalemi boş olamaz.")
            if "belgeTuru" in r:
                vals["belge_turu"] = _choice(r.get("belgeTuru"), DOC_TYPES, "Belge türü") if r.get("belgeTuru") else None
            if "zorunlu" in r:
                vals["zorunlu"] = bool(r.get("zorunlu"))
            if "durum" in r:
                vals["durum"] = _choice(r.get("durum"), CHECK_STATES, "Durum")
            if "gecerlilik" in r:
                vals["gecerlilik_tarihi"] = _day(r.get("gecerlilik"), "Geçerlilik tarihi")
            if "not" in r:
                vals["notu"] = _text(r.get("not"), 500)
            if "belgeId" in r:
                bid = r.get("belgeId") or None
                if bid and bid not in docs:
                    raise TenderError("Bağlanan belge arşivde yok.", 404)
                vals["belge_id"] = bid
                if bid and "durum" not in r:
                    vals["durum"] = "var"
            if cid:
                if not vals:
                    continue
                res = c.execute(CHECKLIST.update().where(CHECKLIST.c.id == cid, CHECKLIST.c.tender_id == tid).values(**vals))
                if not res.rowcount:
                    raise TenderError("Kontrol listesi kalemi bulunamadı.", 404)
                changes["guncellenen"] += 1
            else:
                if not vals.get("kalem"):
                    raise TenderError("Yeni kalemin adı girilmeli.")
                n += 1
                c.execute(CHECKLIST.insert().values(id=uuid.uuid4().hex, tender_id=tid, sira=n, kaynak="elle",
                                                    zorunlu=vals.pop("zorunlu", True), durum=vals.pop("durum", "eksik"), **vals))
                changes["eklenen"] += 1
    rescore(engine, tenant, tid)
    return checklist(engine, tenant, tid), changes


# ------------------------------------------------------------------ dosyalar ve belge arşivi


_MIME = {"pdf": "application/pdf", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         "xls": "application/vnd.ms-excel", "csv": "text/csv", "txt": "text/plain",
         "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "doc": "application/msword",
         "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png"}
_MAGIC = {"pdf": (b"%PDF",), "xlsx": (b"PK",), "docx": (b"PK",), "png": (b"\x89PNG",), "jpg": (b"\xff\xd8",),
          "jpeg": (b"\xff\xd8",), "xls": (b"\xd0\xcf\x11\xe0",), "doc": (b"\xd0\xcf\x11\xe0",)}


def _save(folder: str, filename: str, data: bytes, max_mb: int) -> tuple[str, str, str, str]:
    name = re.sub(r"[\\/\x00-\x1f]", "_", (filename or "").strip())[:200]
    ext = _ext(name)
    mime = _MIME.get(ext)
    if not mime:
        raise TenderError("Dosya PDF, Word, Excel, CSV, metin ya da görsel (JPG/PNG) olmalı.")
    if not data:
        raise TenderError("Dosya boş.")
    if len(data) > max_mb * 1024 * 1024:
        raise TenderError(f"Dosya {max_mb} MB sınırını aşıyor.", 413)
    magic = _MAGIC.get(ext)
    if magic and not any(data.startswith(m) for m in magic):
        raise TenderError("Dosyanın içeriği uzantısıyla uyuşmuyor.")
    digest = hashlib.sha256(data).hexdigest()
    path = files_root() / folder / f"{uuid.uuid4().hex}.{ext}"
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    except OSError as e:
        log.error("ihale dosyası yazılamadı (%s): %s", path, e)
        raise TenderError("Dosya sunucuya kaydedilemedi; yöneticiye bildirin.", 503) from e
    return name, str(path), mime, digest


def add_file(engine: sa.engine.Engine, tenant: str, user: str, tid: str, filename: str, data: bytes, kind: str) -> dict[str, Any]:
    kind = _choice(kind or "sartname", {"sartname": "", "ek": "", "belge": ""}, "Dosya türü")
    with engine.connect() as c:
        _tender_row(c, tenant, tid)
    name, path, mime, digest = _save(tid, filename, data, settings()["fileMaxMb"])
    fid = uuid.uuid4().hex
    with engine.begin() as c:
        c.execute(FILES.insert().values(id=fid, tender_id=tid, tur=kind, ad=name, dosya_yolu=path, mime=mime, boyut=len(data),
                                        sha256=digest, yukleyen=user, zaman=_now()))
    return {"id": fid, "tur": kind, "ad": name, "boyut": len(data), "mime": mime}


def file_of(engine: sa.engine.Engine, tenant: str, tid: str, fid: str) -> tuple[str, str, str]:
    with engine.connect() as c:
        _tender_row(c, tenant, tid)
        r = c.execute(sa.select(FILES).where(FILES.c.id == fid, FILES.c.tender_id == tid)).first()
    if r is None or not Path(r.dosya_yolu).exists():
        raise TenderError("Dosya bulunamadı.", 404)
    return r.dosya_yolu, r.ad, r.mime


def delete_file(engine: sa.engine.Engine, tenant: str, tid: str, fid: str) -> dict[str, Any]:
    with engine.begin() as c:
        _tender_row(c, tenant, tid)
        r = c.execute(sa.select(FILES).where(FILES.c.id == fid, FILES.c.tender_id == tid)).first()
        if r is None:
            raise TenderError("Dosya bulunamadı.", 404)
        c.execute(FILES.delete().where(FILES.c.id == fid))
    Path(r.dosya_yolu).unlink(missing_ok=True)
    return {"id": fid, "ad": r.ad}


def documents(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    cfg = settings()
    with engine.connect() as c:
        rows = c.execute(sa.select(DOCUMENTS).where(DOCUMENTS.c.tenant_id == tenant)).all()
    items = [_doc_out(r, today(), cfg["docWarnDays"]) for r in rows]
    items.sort(key=lambda d: (d["gecerlilik"] is None, d["gecerlilik"] or "", d["ad"]))
    return {"items": items, "turler": DOC_TYPES, "uyariGun": cfg["docWarnDays"]}


def add_document(engine: sa.engine.Engine, tenant: str, user: str, meta: dict[str, Any], filename: str, data: bytes) -> dict[str, Any]:
    ad = _text(meta.get("ad"), 300)
    if not ad:
        raise TenderError("Belge adı girilmeli.")
    tur = _choice(meta.get("tur"), DOC_TYPES, "Belge türü")
    valid = _day(meta.get("gecerlilik"), "Geçerlilik tarihi")
    path = mime = digest = name = None
    if data:
        name, path, mime, digest = _save("belgeler", filename, data, settings()["fileMaxMb"])
    did = uuid.uuid4().hex
    with engine.begin() as c:
        c.execute(DOCUMENTS.insert().values(id=did, tenant_id=tenant, ad=ad, tur=tur, gecerlilik_tarihi=valid, dosya_adi=name,
                                            dosya_yolu=path, mime=mime, boyut=len(data) if data else None, sha256=digest,
                                            notu=_text(meta.get("not"), 500), yukleyen=user, zaman=_now()))
        r = c.execute(sa.select(DOCUMENTS).where(DOCUMENTS.c.id == did)).first()
    return _doc_out(r, today(), settings()["docWarnDays"])


def update_document(engine: sa.engine.Engine, tenant: str, did: str, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    vals: dict[str, Any] = {}
    if "ad" in body:
        vals["ad"] = _text(body.get("ad"), 300)
        if not vals["ad"]:
            raise TenderError("Belge adı boş olamaz.")
    if "tur" in body:
        vals["tur"] = _choice(body.get("tur"), DOC_TYPES, "Belge türü")
    if "gecerlilik" in body:
        vals["gecerlilik_tarihi"] = _day(body.get("gecerlilik"), "Geçerlilik tarihi")
    if "not" in body:
        vals["notu"] = _text(body.get("not"), 500)
    if not vals:
        raise TenderError("Değişiklik yok.")
    with engine.begin() as c:
        r = c.execute(sa.select(DOCUMENTS).where(DOCUMENTS.c.id == did, DOCUMENTS.c.tenant_id == tenant)).first()
        if r is None:
            raise TenderError("Belge bulunamadı.", 404)
        diff = {k: [getattr(r, k), v] for k, v in vals.items() if getattr(r, k) != v}
        c.execute(DOCUMENTS.update().where(DOCUMENTS.c.id == did).values(**vals))
        r = c.execute(sa.select(DOCUMENTS).where(DOCUMENTS.c.id == did)).first()
    return _doc_out(r, today(), settings()["docWarnDays"]), diff


def delete_document(engine: sa.engine.Engine, tenant: str, did: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = c.execute(sa.select(DOCUMENTS).where(DOCUMENTS.c.id == did, DOCUMENTS.c.tenant_id == tenant)).first()
        if r is None:
            raise TenderError("Belge bulunamadı.", 404)
        c.execute(CHECKLIST.update().where(CHECKLIST.c.belge_id == did).values(belge_id=None, durum="eksik"))
        c.execute(DOCUMENTS.delete().where(DOCUMENTS.c.id == did))
    if r.dosya_yolu:
        Path(r.dosya_yolu).unlink(missing_ok=True)
    return {"id": did, "ad": r.ad}


def document_file(engine: sa.engine.Engine, tenant: str, did: str) -> tuple[str, str, str]:
    with engine.connect() as c:
        r = c.execute(sa.select(DOCUMENTS).where(DOCUMENTS.c.id == did, DOCUMENTS.c.tenant_id == tenant)).first()
    if r is None or not r.dosya_yolu or not Path(r.dosya_yolu).exists():
        raise TenderError("Belgenin dosyası yok.", 404)
    return r.dosya_yolu, r.dosya_adi or r.ad, r.mime or "application/octet-stream"


# ------------------------------------------------------------------ karar ve sonuç


def submit_decision(engine: sa.engine.Engine, tenant: str, user: str, tid: str, body: dict[str, Any]) -> dict[str, Any]:
    karar = _choice(body.get("karar"), DECISIONS_LABEL, "Karar")
    gerekce = _text(body.get("gerekce"), 4000)
    if karar == "basvurma" and not gerekce:
        raise TenderError("Başvurmama kararının gerekçesini yazın.")
    d = detail(engine, tenant, tid)
    if d["durum"] not in ("yeni", "inceleniyor"):
        raise TenderError("Karar yalnız inceleme aşamasındaki ihalede önerilir.", 409)
    tot = d["toplamlar"]
    if karar == "basvur":
        if not tot["fiyatli"]:
            raise TenderError("Teklif tablosunda fiyatlı eşleşmiş kalem yok.")
        pend = sum(tot["durumlar"].get(k, 0) for k in ("oneri", "belirsiz", "bekliyor"))
        if pend:
            raise TenderError(f"{pend} kalemde eşleştirme onay ya da işlem bekliyor; önce onlara karar verin.", 409)
    with engine.begin() as c:
        if c.execute(sa.select(DECISIONS.c.id).where(DECISIONS.c.tender_id == tid, DECISIONS.c.durum == "onayda")).first():
            raise TenderError("Onay bekleyen bir karar zaten var.", 409)
        did = uuid.uuid4().hex
        c.execute(DECISIONS.insert().values(
            id=did, tender_id=tid, karar=karar, durum="onayda",
            teklif_toplami=tot["araToplam"] if karar == "basvur" else None, fiyat_orani=d["fiyatOrani"], gerekce=gerekce,
            ozet_json=_dump(d["kararOzeti"]), oneren=user, oneri_zamani=_now()))
        r = c.execute(sa.select(DECISIONS).where(DECISIONS.c.id == did)).first()
    return _decision_out(r)


def _pending(c: Any, tid: str) -> Any:
    r = c.execute(sa.select(DECISIONS).where(DECISIONS.c.tender_id == tid, DECISIONS.c.durum == "onayda")).first()
    if r is None:
        raise TenderError("Onay bekleyen karar yok.", 409)
    return r


def withdraw_decision(engine: sa.engine.Engine, tenant: str, user: str, tid: str) -> dict[str, Any]:
    with engine.begin() as c:
        _tender_row(c, tenant, tid)
        r = _pending(c, tid)
        if r.oneren.lower() != user.lower():
            raise TenderError("Kararı yalnız öneren geri çeker.", 409)
        c.execute(DECISIONS.update().where(DECISIONS.c.id == r.id).values(durum="geri_cekildi"))
        r = c.execute(sa.select(DECISIONS).where(DECISIONS.c.id == r.id)).first()
    return _decision_out(r)


def decide(engine: sa.engine.Engine, tenant: str, user: str, tid: str, approve: bool, note: Any = None) -> dict[str, Any]:
    """Onay ya da geri gönderme. Öneren onaylayamaz (iki göz). Onay: başvur → «başvurulacak» ve eşleşen kalemlerin
    teklif fiyatları onaylanmış sayılır; başvurma → «başvurulmayacak»."""
    note_t = _text(note, 2000)
    if not approve and not note_t:
        raise TenderError("Geri gönderme gerekçesi yazın.")
    now = _now()
    with engine.begin() as c:
        _tender_row(c, tenant, tid)
        r = _pending(c, tid)
        if r.oneren.lower() == user.lower():
            raise TenderError("Kararı öneren kişi aynı kararı onaylayamaz; başka bir yetkili onaylamalı.", 409)
        c.execute(DECISIONS.update().where(DECISIONS.c.id == r.id).values(
            durum="onaylandi" if approve else "reddedildi", onaylayan=user, onay_zamani=now, onay_notu=note_t))
        if approve:
            c.execute(TENDERS.update().where(TENDERS.c.id == tid).values(
                durum="basvurulacak" if r.karar == "basvur" else "basvurulmayacak", updated_by=user, updated_at=now))
            if r.karar == "basvur":
                c.execute(ITEMS.update().where(ITEMS.c.tender_id == tid, ITEMS.c.eslesme_durumu == "eslesti")
                          .values(onaylayan=user, onay_zamani=now))
        r = c.execute(sa.select(DECISIONS).where(DECISIONS.c.id == r.id)).first()
    return _decision_out(r)


def record_result(engine: sa.engine.Engine, tenant: str, user: str, tid: str, body: dict[str, Any]) -> dict[str, Any]:
    sonuc = _choice(body.get("sonuc"), RESULT_STATES, "Sonuç")
    d = detail(engine, tenant, tid)
    if sonuc in ("kazanildi", "kaybedildi") and d["durum"] not in ("teklif_verildi", "kazanildi", "kaybedildi"):
        raise TenderError("Kazanıldı/kaybedildi ancak teklif verildikten sonra girilir.", 409)
    kazanan_fiyat = _num(body.get("kazananFiyat"), "Kazanan fiyat")
    bizim = _num(body.get("bizimFiyat"), "Bizim teklif") if "bizimFiyat" in body else None
    if bizim is None and sonuc != "iptal":
        last = next((k for k in d["kararlar"] if k["durum"] == "onaylandi" and k["karar"] == "basvur"), None)
        bizim = last["teklifToplami"] if last else None
    if sonuc == "kaybedildi" and not _text(body.get("neden"), 4000):
        raise TenderError("Kayıp nedenini yazın.")
    kazanan = _text(body.get("kazanan"), 300)
    if sonuc == "kazanildi" and not kazanan:
        kazanan = "TİMAŞ"
        kazanan_fiyat = kazanan_fiyat if kazanan_fiyat is not None else bizim
    vals = dict(sonuc=sonuc, kazanan=kazanan, kazanan_fiyat=kazanan_fiyat, bizim_fiyat=bizim,
                liste_toplami=d["toplamlar"]["listeToplami"] or None, neden=_text(body.get("neden"), 4000),
                kaynak=_text(body.get("kaynak"), 300), kaydeden=user, zaman=_now())
    with engine.begin() as c:
        if c.execute(sa.select(RESULTS.c.tender_id).where(RESULTS.c.tender_id == tid)).first():
            c.execute(RESULTS.update().where(RESULTS.c.tender_id == tid).values(**vals))
        else:
            c.execute(RESULTS.insert().values(tender_id=tid, **vals))
        c.execute(TENDERS.update().where(TENDERS.c.id == tid).values(durum=sonuc, updated_by=user, updated_at=_now()))
        r = c.execute(sa.select(RESULTS).where(RESULTS.c.tender_id == tid)).first()
    return _result_out(r)


def results(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    """Sonuçlar ekranı: bütün sonuçlanmış ihaleler ve kurum türüne göre kazanma özeti."""
    with engine.connect() as c:
        rows = c.execute(sa.select(RESULTS, TENDERS.c.kurum, TENDERS.c.kurum_turu, TENDERS.c.konu, TENDERS.c.il)
                         .join(TENDERS, TENDERS.c.id == RESULTS.c.tender_id).where(TENDERS.c.tenant_id == tenant)
                         .order_by(RESULTS.c.zaman.desc())).all()
    items = []
    by_type: dict[str, dict[str, Any]] = {}
    for r in rows:
        o = _result_out(r)
        o.update(id=r.tender_id, kurum=r.kurum, kurumTuru=r.kurum_turu, kurumTuruAdi=KURUM_TURLERI.get(r.kurum_turu),
                 konu=r.konu, il=r.il)
        items.append(o)
        t = by_type.setdefault(r.kurum_turu, {"kurumTuru": r.kurum_turu, "kurumTuruAdi": KURUM_TURLERI.get(r.kurum_turu),
                                              "sonuc": 0, "kazanilan": 0, "oranlar": []})
        t["sonuc"] += 1
        t["kazanilan"] += 1 if r.sonuc == "kazanildi" else 0
        if o["kazananListeOrani"] is not None:
            t["oranlar"].append(o["kazananListeOrani"])
    summary = []
    for t in by_type.values():
        ratios = t.pop("oranlar")
        summary.append({**t, "kazananOranOrtanca": round(statistics.median(ratios), 4) if ratios else None, "oranSayisi": len(ratios)})
    return {"items": items, "ozet": sorted(summary, key=lambda x: -x["sonuc"])}


# ------------------------------------------------------------------ takvim ve hatırlatma


def calendar(engine: sa.engine.Engine, tenant: str, days: int = 90, asof: Optional[date] = None) -> dict[str, Any]:
    """Son teklif tarihleri, belge geçerlilik bitişleri (arşiv + ihaleye özel), teminat iade tarihleri. Geçmişte kalmış
    ama hâlâ açık olanlar (süresi dolmuş belge, tarihi geçen açık ihale) de listelenir."""
    asof = asof or today()
    end = (asof + timedelta(days=max(0, days))).isoformat()
    events: list[dict[str, Any]] = []
    with engine.connect() as c:
        tenders = c.execute(sa.select(TENDERS).where(TENDERS.c.tenant_id == tenant)).all()
        docs = c.execute(sa.select(DOCUMENTS).where(DOCUMENTS.c.tenant_id == tenant)).all()
        checks = c.execute(sa.select(CHECKLIST, TENDERS.c.kurum, TENDERS.c.durum.label("ihale_durum")).join(TENDERS, TENDERS.c.id == CHECKLIST.c.tender_id)
                           .where(TENDERS.c.tenant_id == tenant, CHECKLIST.c.gecerlilik_tarihi.is_not(None))).all()
        pending = {r.tender_id for r in c.execute(sa.select(DECISIONS.c.tender_id).where(DECISIONS.c.durum == "onayda"))}
    for t in tenders:
        if t.son_teklif_tarihi and t.durum in OPEN_STATUSES and t.durum != "teklif_verildi" and t.son_teklif_tarihi[:10] <= end:
            events.append({"tarih": t.son_teklif_tarihi, "tur": "son_teklif", "baslik": f"Son teklif: {t.kurum}",
                           "ayrinti": t.konu[:200], "ihaleId": t.id, "kalanGun": days_left(t.son_teklif_tarihi, asof),
                           "onayBekliyor": t.id in pending})
        if t.teminat_iade_tarihi and t.teminat_iade_tarihi <= end and t.teminat_iade_tarihi >= (asof - timedelta(days=30)).isoformat():
            events.append({"tarih": t.teminat_iade_tarihi, "tur": "teminat_iade", "baslik": f"Teminat iadesi: {t.kurum}",
                           "ayrinti": f"{_fmt_tr(t.teminat_tutari)} ₺" if t.teminat_tutari else None,
                           "ihaleId": t.id, "kalanGun": days_left(t.teminat_iade_tarihi, asof)})
    for d in docs:
        if d.gecerlilik_tarihi and d.gecerlilik_tarihi <= end:
            events.append({"tarih": d.gecerlilik_tarihi, "tur": "belge", "baslik": f"Belge geçerliliği: {d.ad}",
                           "ayrinti": DOC_TYPES.get(d.tur), "belgeId": d.id, "kalanGun": days_left(d.gecerlilik_tarihi, asof)})
    for ch in checks:
        if ch.ihale_durum in OPEN_STATUSES and ch.gecerlilik_tarihi <= end:
            events.append({"tarih": ch.gecerlilik_tarihi, "tur": "ihale_belgesi", "baslik": f"{ch.kalem} ({ch.kurum})",
                           "ayrinti": "İhaleye özel belge geçerliliği", "ihaleId": ch.tender_id,
                           "kalanGun": days_left(ch.gecerlilik_tarihi, asof)})
    events.sort(key=lambda e: e["tarih"])
    return {"items": events, "gun": days, "bugun": asof.isoformat()}


def due_reminders(engine: sa.engine.Engine, tenant: str, cfg: dict[str, Any], asof: Optional[date] = None) -> list[dict[str, Any]]:
    """Gönderilmemiş hatırlatmalar: son teklif tarihine `remindDays` gün kala (ve geçmişse bir kez), belge geçerliliği
    `docWarnDays` içinde biterken, teminat iade tarihi 7 gün içinde, onay bekleyen karar (günde bir)."""
    asof = asof or today()
    cal = calendar(engine, tenant, max(cfg["docWarnDays"], max(cfg["remindDays"] or [0]), 7), asof)["items"]
    with engine.connect() as c:
        sent = {r.key for r in c.execute(sa.select(REMINDERS.c.key).where(REMINDERS.c.tenant_id == tenant))}
        pend = c.execute(sa.select(DECISIONS, TENDERS.c.kurum).join(TENDERS, TENDERS.c.id == DECISIONS.c.tender_id)
                         .where(TENDERS.c.tenant_id == tenant, DECISIONS.c.durum == "onayda")).all()
    out = []
    for e in cal:
        left = e.get("kalanGun")
        if left is None:
            continue
        key = None
        if e["tur"] == "son_teklif":
            step = next((d for d in sorted(cfg["remindDays"]) if left <= d), None) if left >= 0 else -1
            if step is not None:
                key = f"son:{e['ihaleId']}:{e['tarih'][:10]}:{step}"
        elif e["tur"] in ("belge", "ihale_belgesi"):
            key = f"belge:{e.get('belgeId') or e.get('ihaleId')}:{e['tarih']}:{e['baslik'][:40]}" + (":doldu" if left < 0 else "")
        elif e["tur"] == "teminat_iade" and left <= 7:
            key = f"teminat:{e['ihaleId']}:{e['tarih']}"
        if key and key not in sent:
            out.append({**e, "key": key})
    for p in pend:
        key = f"onay:{p.id}:{asof.isoformat()}"
        if key not in sent:
            out.append({"tur": "onay", "tarih": asof.isoformat(), "baslik": f"Karar onayı bekliyor: {p.kurum}",
                        "ayrinti": f"{DECISIONS_LABEL.get(p.karar)} — öneren {p.oneren}", "ihaleId": p.tender_id, "key": key})
    return out


def mark_sent(engine: sa.engine.Engine, tenant: str, keys: Iterable[str]) -> None:
    now = _now()
    with engine.begin() as c:
        for k in keys:
            if not c.execute(sa.select(REMINDERS.c.key).where(REMINDERS.c.key == k[:160])).first():
                c.execute(REMINDERS.insert().values(key=k[:160], tenant_id=tenant, gonderildi=now))


def reminder_text(items: list[dict[str, Any]], link: str) -> str:
    lines = ["İhale takibi hatırlatmaları", ""]
    labels = {"son_teklif": "Son teklif tarihi", "belge": "Belge geçerliliği", "ihale_belgesi": "İhale belgesi",
              "teminat_iade": "Teminat iadesi", "onay": "Onay bekleyen karar"}
    for kind, label in labels.items():
        part = [i for i in items if i["tur"] == kind]
        if not part:
            continue
        lines.append(f"{label} ({len(part)})")
        for i in part:
            left = i.get("kalanGun")
            when = "" if left is None else (" — bugün" if left == 0 else (f" — {left} gün kaldı" if left > 0 else f" — {-left} gün geçti"))
            lines.append(f"  • {i['baslik']} ({i['tarih'][:10]}){when}" + (f": {i['ayrinti']}" if i.get("ayrinti") else ""))
        lines.append("")
    if link:
        lines.append(link)
    lines.append("Bu e-posta portalın ihale takibinden otomatik gelir; kuruma hiçbir gönderim yapılmaz.")
    return "\n".join(lines)


# ------------------------------------------------------------------ Excel teklif tablosu


def pricing_xlsx(d: dict[str, Any]) -> bytes:
    """Teklif fiyat tablosu: eşleşen kalemler (toplama girenler), ardından toplama girmeyenler ayrı sayfada."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Teklif fiyat tablosu"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EDE9FE")
    ws["A1"] = f"{d['kurum']} — {d['konu'][:120]}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"İhale no: {d.get('kaynakNo') or '—'} · Son teklif: {d.get('sonTeklifTarihi') or '—'} · "
                f"Fiyat oranı: {d.get('fiyatOrani') or 1:.4f} ({d.get('fiyatOraniKaynak') or ''})")
    ws["A3"] = "Taslaktır: teklif fiyatı karar onayıyla kesinleşir; portal kuruma gönderim yapmaz."
    cols = ["Sıra", "Şartname kalemi", "Eşleşen kitap", "Stok kodu", "ISBN", "Adet", "Stok", "Liste fiyatı (KDV dahil)",
            "Fiyat kaynağı", "KDV oranı", "Birim teklif fiyatı (KDV hariç)", "Tutar (KDV hariç)", "KDV", "Birim maliyet",
            "Marj", "Eşleşme", "Onaylayan"]
    start = 5
    for i, h in enumerate(cols, start=1):
        cell = ws.cell(row=start, column=i, value=h)
        cell.font, cell.fill = bold, head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    rows = [k for k in d["kalemler"] if k["durum"] == "eslesti" and k["adet"] is not None and k["onerilenFiyat"] is not None]
    rest = [k for k in d["kalemler"] if k not in rows]
    r = start
    for k in rows:
        r += 1
        vat = k["kdvOrani"] or 0.0
        line = round(k["adet"] * k["onerilenFiyat"], 2)
        vals = [k["sira"], k["metin"], k["eslesenAd"], k["stokKodu"], k["isbn"], k["adet"], k["stok"], k["listeFiyati"],
                k["fiyatKaynagi"], vat, k["onerilenFiyat"], line, round(line * vat, 2), k["tahminiMaliyet"], k["marj"],
                k["yontemAdi"], k["onaylayan"]]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)
        for col, fmt in ((8, "#,##0.00"), (10, "0%"), (11, "#,##0.00"), (12, "#,##0.00"), (13, "#,##0.00"), (14, "#,##0.00"), (15, "0.0%")):
            ws.cell(row=r, column=col).number_format = fmt
    tot = d["toplamlar"]
    r += 2
    for label, value in (("Ara toplam (KDV hariç)", tot["araToplam"]), ("KDV", tot["kdv"]), ("Genel toplam (KDV dahil)", tot["genelToplam"])):
        ws.cell(row=r, column=11, value=label).font = bold
        c = ws.cell(row=r, column=12, value=value)
        c.font, c.number_format = bold, "#,##0.00"
        r += 1
    ws.cell(row=r + 1, column=1, value=(f"Marj: {tot['marj']:.1%} ({tot['maliyetKapsam']} kalemin maliyeti biliniyor)"
                                        if tot["marj"] is not None else tot["maliyetNotu"]))
    widths = [6, 48, 40, 16, 16, 8, 10, 14, 26, 8, 16, 16, 12, 12, 8, 14, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    if rest:
        ws2 = wb.create_sheet("Toplama girmeyenler")
        heads = ["Sıra", "Şartname kalemi", "Adet", "Durum", "Neden"]
        for i, h in enumerate(heads, start=1):
            c = ws2.cell(row=1, column=i, value=h)
            c.font, c.fill = bold, head_fill
        for n, k in enumerate(rest, start=2):
            why = k["durumAdi"] if k["durum"] != "eslesti" else ("Adet yok" if k["adet"] is None else "Fiyat yok")
            for i, v in enumerate([k["sira"], k["metin"], k["adet"], k["durumAdi"], why], start=1):
                ws2.cell(row=n, column=i, value=v)
        for i, w in enumerate([6, 60, 8, 30, 30], start=1):
            ws2.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ karar özeti metni (Zeki AI)


def _fmt_tr(v: float, digits: int = 2) -> str:
    s = f"{v:,.{digits}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def facts_text(f: dict[str, Any]) -> tuple[str, set[str]]:
    """Karar özeti rakamlarının model girdisi ve metinde izin verilen sayılar (biçimlenmiş hâlleriyle)."""
    lines = []
    allowed: set[str] = set()

    def add(label: str, v: Any, kind: str = "num") -> None:
        if v is None:
            return
        if kind == "money":
            s = _fmt_tr(float(v)) + " TL"
        elif kind == "pct":
            s = "%" + _fmt_tr(float(v) * 100, 1)
        elif kind == "int":
            s = _fmt_tr(float(v), 0)
        else:
            s = str(v)
        lines.append(f"{label}: {s}")
        allowed.update(_nums(s))

    add("Yaklaşık tutar", f.get("yaklasikTutar"), "money")
    add("Uygunluk puanı (100 üzerinden)", f.get("uygunlukPuani"))
    add("Şartname kalemi", f.get("kalem"), "int")
    add("Katalogla eşleşen kalem", f.get("eslesen"), "int")
    add("Katalogda olmayan kalem", f.get("katalogdaYok"), "int")
    add("Stoğu yetmeyen kalem", f.get("stokYetersiz"), "int")
    add("Teklif ara toplamı (KDV hariç)", f.get("teklifAraToplam"), "money")
    add("Fiyat oranı (liste fiyatına göre)", f.get("fiyatOrani"), "pct")
    add("Tahmini marj", f.get("marj"), "pct")
    add("Teminat", f.get("teminatTutari"), "money")
    add("Son teklif tarihine kalan gün", f.get("kalanGun"), "int")
    h = (f.get("gecmis") or {}).get("kurumTuru") or {}
    add("Aynı kurum türünde geçmiş sonuç", h.get("sonuc"), "int")
    add("Bunlardan kazanılan", h.get("kazanilan"), "int")
    if f.get("eksikBelgeler"):
        lines.append("Eksik belgeler: " + "; ".join(f["eksikBelgeler"]))
        for b in f["eksikBelgeler"]:
            allowed.update(_nums(b))
    if f.get("maliyetNotu"):
        lines.append(f["maliyetNotu"])
    for r in f.get("riskler") or []:
        lines.append("Risk: " + r)
        allowed.update(_nums(r))
    return "\n".join(lines), allowed


def brief_text(f: dict[str, Any], chat: Optional[Callable[[list[dict[str, str]]], str]]) -> dict[str, Any]:
    """Müdür için 5 cümlelik karar özeti. Rakamlar yalnız verilenlerden; olgu dışı sayı içeren metin atılır ve kural
    metni kullanılır. Teklif fiyatı ya da başvuru kararı önermez (K4)."""
    facts, _allowed = facts_text(f)
    fallback = facts.replace("\n", ". ") + "."
    if chat is None:
        return {"metin": fallback, "kaynak": "kural"}
    prompt = ("Aşağıdaki ihale bilgileriyle yönetici için en çok 5 cümlelik Türkçe bir karar özeti yaz. Yalnız verilen "
              "rakamları aynen kullan, yeni sayı yazma, yuvarlama yapma. Başvuru kararı ya da fiyat önerme; yalnız "
              "durumu, güçlü ve zayıf yanları söyle.\n\n" + facts)
    try:
        text = str(chat([{"role": "user", "content": prompt}]) or "").strip()
    except Exception as e:  # noqa: BLE001
        log.warning("ihale karar özeti: model cevap vermedi: %s", e)
        return {"metin": fallback, "kaynak": "kural", "not": "Zeki AI'a ulaşılamadı; rakamlar kural metniyle yazıldı."}
    from semantic_bridge import zeki_text as Z

    # Tek sayı denetçisi: olgu metninde (değerler biçimlenmiş hâlleriyle) olmayan sayı; 1–5 serbest (sıra, cümle).
    extra = Z.unsupported(text, facts, free_upto=5)
    if not text or extra:
        return {"metin": fallback, "kaynak": "kural",
                "not": "Zeki AI metni verilmeyen bir sayı içerdiği için kullanılmadı." if extra else None}
    return {"metin": text, "kaynak": "zeki"}


def set_brief(engine: sa.engine.Engine, tenant: str, user: str, tid: str, text: str) -> None:
    with engine.begin() as c:
        _tender_row(c, tenant, tid)
        c.execute(TENDERS.update().where(TENDERS.c.id == tid).values(karar_metni=text, updated_by=user, updated_at=_now()))


def meta() -> dict[str, Any]:
    cfg = settings()
    return {"durumlar": STATUSES, "acikDurumlar": list(OPEN_STATUSES), "elleDurumlar": list(MANUAL_STATUSES),
            "kurumTurleri": KURUM_TURLERI, "usuller": USULLER, "kaynaklar": KAYNAKLAR, "eslesmeDurumlari": MATCH_STATES,
            "eslesmeYontemleri": MATCH_METHODS, "belgeTurleri": DOC_TYPES, "kontrolDurumlari": CHECK_STATES,
            "sonuclar": RESULT_STATES, "kararlar": DECISIONS_LABEL,
            "ayarlar": {k: cfg[k] for k in ("autoProb", "autoMargin", "suggestProb", "suggestMargin", "candidates", "minScore",
                                             "priceSource", "defaultVat", "historyMin", "weights", "fullDays", "remindDays",
                                             "docWarnDays", "fileMaxMb", "publicChannel", "watchEnabled")}}
