"""M36 Dijital yayın ve e-kitap yönetimi: dijital katalog, hak kararı (e-kitap / sesli), fırsat listesi, platform durumu,
platform satış raporu yükleme ve kitaba eşleme, CRM'e işlenecekler.

**Kaynaklar** (`dijital_sources`, yalnız okuma): CRM kitap kartı, Telif Alış sözleşme hakları, e-kitap üretim aşaması,
kitap geçmişi; Logo faturalı satış (son 12 ay); Kitap Tasarım Stüdyosu e-kitap durumu (varsa). Gece işi (`refresh`)
okumaları köprünün `semantic_dijital_*` tablolarına yazar; ekranlar bu tablolardan okur.

**Yazma yok:** CRM'e, Logo'ya, T-soft'a ve dijital platformlara hiçbir şey gitmez. Platform durumu, hak kararı, dijital
fiyat kararı ve rapor satırları yalnız bu tablolarda durur; CRM'e işlenmesi gereken değerler (e-ISBN, E-Pub durumu,
e-kitap stok kodu) `semantic_dijital_crm_pending`'de «CRM'e işlenecek» diye listelenir. Platforma yükleme insanın işidir.

**Hak kararı (kitap × biçim, kurallı):** kitaba bağlı yürürlükteki bütün Telif Alış sözleşmelerinde biçimin bayrağı
(`new_EKitap` / `new_SesliKitapHakki`) varsa `var`; biri eksikse `eksik` (taraf adıyla); yürürlükte sözleşme yoksa `yok`
(koruma dışı eserse `koruma_disi`); bayraklar tam ama serbest metinli hak notu varsa `incele`. `incele` kaydına telif
biriminin kararı (`uygun` → `var`, `kismi` → `kismi`, `uygun_degil` → `yok`) yazılır ve **notun metni değişmedikçe**
bir daha sorulmaz (not özeti kararla saklanır). Kitap olmayan ürün `kitap_degil`, kendi sözleşmesi olmayan set `set`
(SEO/GEO ile aynı süzgeç). Hak notu olan sözleşme hiçbir zaman kendiliğinden «hak var» sayılmaz. M54 hak haritası
bağlanırsa karar oradan gelir (`dijital_sources.register_rights_provider`).

**Dijital sürüm var mı:** e-kitap = CRM'de e-kitap stok kodu dolu ya da bir platformda «yüklendi/yayında» işaretli;
sesli = aynı adlı «SesliKitap» tipli CRM kartı ya da sesli platformda «yüklendi/yayında». Platform durumu yalnız
kullanıcının girdiği ya da onaylı rapordan gelen bilgidir; bilinmeyen «bilinmiyor» yazar.

**Fırsat (kurallı, rakam SQL'den):** kitap tipi kart ∧ hak `var`/`kismi` ∧ dijital sürüm yok ∧ son 12 ay basılı net adet ≥
eşik (`DIJITAL_OPP_MIN_QTY`) ∧ yayın durumu satış dışı değil. Puan = basılı satışta kitabın yüzdelik sırası (0–100).
Sesli adaylarda ayrıca tür süzgeci (`DIJITAL_AUDIO_GENRES`, boşsa hepsi).

**Hak riski:** dijitalde görünen (e-kitap stok kodu / platform / Logo'da e-kitap satışı) ama hakkı `eksik`, `yok` ya da
`incele` olan kitap. Hak notu M54 telif ile **ortak** sınıflanır (`rights_notes`, tek tablo); «dijitali kısıtlıyor /
kısıtlamıyor / belirsiz» okuması o sınıftan türetilir (yalnız sıralama);
karar telif biriminindir.

**Rapor yükleme:** Excel/CSV satırlarının **hiçbiri atılmaz**: boş satır dışında her satır `semantic_dijital_sales`'e
yazılır; eşleşmeyen satır açık iş olarak kalır, «toplam» satırı işaretlenir ve toplamlara girmez. Kişisel veri kolonu
(e-posta, telefon, adres, müşteri adı…) yüklemede atılır ve saklanmaz; atılan kolon adları kayıtta yazılır. Kurallı
eşleme: e-ISBN, e-kitap barkodu, ISBN, barkod, e-kitap stok kodu, stok kodu. Kalan satırlara Zeki AI aday önerir
(kapalı küme seçim + olasılık); onayı insan verir. Döviz kuru onay sırasında kullanıcıdan alınır, modül kur varsaymaz.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import math
import re
import threading
import time
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, Iterable, Optional
from zoneinfo import ZoneInfo

import sqlalchemy as sa

from semantic_bridge import dijital_sources as src
from semantic_bridge import rights_map as RM
from semantic_bridge import rights_notes as RN
from semantic_bridge.seo_geo import crm as seo_crm

log = logging.getLogger("semantic.dijital")
TZ = ZoneInfo("Europe/Istanbul")
PAGE = 50
_md = sa.MetaData()

TITLES = sa.Table(
    "semantic_dijital_titles", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("kitap_id", sa.String(40), primary_key=True),          # CRM new_kitapId (büyük harf)
    sa.Column("ad", sa.String(500)),
    sa.Column("yazar", sa.String(500)),
    sa.Column("stok_kodu", sa.String(60)),
    sa.Column("isbn", sa.String(40)),
    sa.Column("ean", sa.String(40)),
    sa.Column("e_isbn", sa.String(40)),
    sa.Column("ekitap_stok_kodu", sa.String(60)),
    sa.Column("ekitap_barkod", sa.String(40)),
    sa.Column("tip", sa.Integer),
    sa.Column("tip_adi", sa.String(80)),
    sa.Column("yayin_durumu", sa.String(200)),
    sa.Column("hedef_kitle", sa.String(80)),
    sa.Column("turler", sa.String(500)),
    sa.Column("ilk_yayin", sa.String(10)),
    sa.Column("epub_durumu_crm", sa.Integer),                         # 1 Evet, 0 Hayır
    sa.Column("uretim_durumu", sa.String(120)),                       # CRM üretim kartı e-kitap aşaması
    sa.Column("uretim_tarih", sa.String(10)),
    sa.Column("studio_is", sa.String(80)),
    sa.Column("studio_epub_durumu", sa.String(20)),                   # yok|is_var|hazirlaniyor|hazir|eski|hata
    sa.Column("studio_denetim", sa.String(10)),                       # OK|WARN|FAIL
    sa.Column("studio_e_isbn", sa.String(40)),
    sa.Column("hak_ekitap", sa.String(20)),
    sa.Column("hak_ekitap_gerekce", sa.Text),
    sa.Column("hak_sesli", sa.String(20)),
    sa.Column("hak_sesli_gerekce", sa.Text),
    sa.Column("hak_kaynak", sa.String(20)),                           # crm|hak-haritasi
    sa.Column("hak_notu_var", sa.Boolean),
    sa.Column("sozlesme_json", sa.Text),
    sa.Column("basili_fiyat", sa.Float),                              # CRM KDV dahil liste fiyatı
    sa.Column("basili_12ay_adet", sa.Float),
    sa.Column("basili_12ay_ciro", sa.Float),
    sa.Column("logo_dijital_12ay_adet", sa.Float),                    # e-kitap stok koduyla Logo faturaları
    sa.Column("logo_dijital_12ay_ciro", sa.Float),
    sa.Column("ekitap_var", sa.Boolean),
    sa.Column("sesli_var", sa.Boolean),
    sa.Column("firsat_puani", sa.Float),
    sa.Column("firsat_gerekcesi", sa.Text),
    sa.Column("sesli_firsat_puani", sa.Float),
    sa.Column("sesli_firsat_gerekcesi", sa.Text),
    sa.Column("baski_degisim_tarih", sa.String(10)),
    sa.Column("baski_degisim_tur", sa.String(40)),
    sa.Column("okundu_at", sa.DateTime(timezone=True)),
    # kullanıcı alanları: gece okuması bunlara dokunmaz
    sa.Column("dijital_fiyat", sa.Float),
    sa.Column("fiyat_gerekce", sa.Text),
    sa.Column("fiyat_onaylayan", sa.String(120)),
    sa.Column("fiyat_tarih", sa.DateTime(timezone=True)),
)
PLATFORMS = sa.Table(
    "semantic_dijital_platforms", _md,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("ad", sa.String(120), nullable=False),
    sa.Column("tur", sa.String(12), nullable=False),                  # ekitap|sesli|abonelik
    sa.Column("dagitim", sa.String(12), nullable=False),              # dogrudan|dagitici
    sa.Column("para_birimi", sa.String(8), nullable=False),
    sa.Column("rapor_gunu", sa.Integer),                              # ay kapanışından kaç gün sonra rapor beklenir
    sa.Column("aktif", sa.Boolean, nullable=False, default=True),
    sa.Column("olusturan", sa.String(120)),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
)
LISTINGS = sa.Table(
    "semantic_dijital_listings", _md,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),  # yalnız eklenir; son satır geçerli durumdur
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("kitap_id", sa.String(40), nullable=False),
    sa.Column("platform_id", sa.Integer, nullable=False),
    sa.Column("durum", sa.String(16), nullable=False),
    sa.Column("tarih", sa.String(10)),
    sa.Column("kaynak", sa.String(12), nullable=False),               # kullanici|rapor
    sa.Column("notlar", sa.Text),
    sa.Column("fiyat", sa.Float),
    sa.Column("yazan", sa.String(120)),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
)
DECISIONS = sa.Table(
    "semantic_dijital_rights_decisions", _md,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("kitap_id", sa.String(40), nullable=False),
    sa.Column("sozlesme_id", sa.String(40), nullable=False),
    sa.Column("bicim", sa.String(8), nullable=False),                 # ekitap|sesli
    sa.Column("karar", sa.String(16), nullable=False),                # uygun|uygun_degil|kismi
    sa.Column("gerekce", sa.Text, nullable=False),
    sa.Column("not_ozeti", sa.String(64)),                            # kararın verildiği hak notunun özeti
    sa.Column("yazan", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
)
IMPORTS = sa.Table(
    "semantic_dijital_imports", _md,
    sa.Column("id", sa.String(20), primary_key=True),                 # DR-<yıl>-<sıra>
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("platform_id", sa.Integer, nullable=False),
    sa.Column("donem", sa.String(7), nullable=False),                 # YYYY-AA
    sa.Column("dosya_adi", sa.String(300)),
    sa.Column("yukleyen", sa.String(120)),
    sa.Column("satir", sa.Integer, nullable=False, default=0),
    sa.Column("eslesen", sa.Integer, nullable=False, default=0),
    sa.Column("eslesmeyen", sa.Integer, nullable=False, default=0),
    sa.Column("durum", sa.String(12), nullable=False),                # onizleme|onaylandi|iptal
    sa.Column("kolonlar_json", sa.Text),
    sa.Column("baslik_json", sa.Text),
    sa.Column("atilan_json", sa.Text),
    sa.Column("ham_json", sa.Text),
    sa.Column("bos_satir", sa.Integer),
    sa.Column("eslestirme_json", sa.Text),                            # Zeki AI önerisi işinin durumu
    sa.Column("kur_json", sa.Text),
    sa.Column("onaylayan", sa.String(120)),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("committed_at", sa.DateTime(timezone=True)),
)
SALES = sa.Table(
    "semantic_dijital_sales", _md,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("import_id", sa.String(20), nullable=False, index=True),
    sa.Column("sira", sa.Integer, nullable=False),                    # dosyadaki satır numarası (1'den)
    sa.Column("platform_id", sa.Integer, nullable=False),
    sa.Column("donem_ay", sa.String(7), nullable=False),
    sa.Column("tur", sa.String(8), nullable=False, default="satir"),  # satir|ozet (toplam satırı: toplama girmez)
    sa.Column("kimlik_ham", sa.String(120)),
    sa.Column("baslik_ham", sa.String(500)),
    sa.Column("yazar_ham", sa.String(300)),
    sa.Column("kitap_id", sa.String(40)),
    sa.Column("eslesme", sa.String(8)),                               # kural|zeki|elle
    sa.Column("eslesme_anahtar", sa.String(40)),
    sa.Column("olasilik", sa.Float),
    sa.Column("aday_json", sa.Text),
    sa.Column("adet", sa.Float),
    sa.Column("brut", sa.Float),
    sa.Column("net", sa.Float),
    sa.Column("para_birimi", sa.String(8)),
    sa.Column("kur", sa.Float),
    sa.Column("net_tl", sa.Float),
)
PENDING = sa.Table(
    "semantic_dijital_crm_pending", _md,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("kitap_id", sa.String(40), nullable=False),
    sa.Column("alan", sa.String(40), nullable=False),
    sa.Column("deger", sa.String(200), nullable=False),
    sa.Column("kaynak", sa.String(200)),
    sa.Column("durum", sa.String(10), nullable=False),                # acik|kapandi
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("closed_at", sa.DateTime(timezone=True)),
)
META = sa.Table(
    "semantic_dijital_meta", _md,
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("key", sa.String(60), primary_key=True),
    sa.Column("value_json", sa.Text),
    sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
)

RIGHTS = {"var": "Hak var", "kismi": "Kısıtlı hak", "incele": "İncelenmeli", "eksik": "Eksik", "yok": "Hak yok",
          "koruma_disi": "Koruma dışı eser", "set": "Set (içindeki kitaplara bağlı)", "kitap_degil": "Kitap değil"}
OK_RIGHTS = ("var", "kismi", "koruma_disi")
RISK_RIGHTS = ("eksik", "yok", "incele")
DECISIONS_KINDS = {"uygun": "Dijital dağıtıma uygun", "kismi": "Kısıtlı (bölge/platform)", "uygun_degil": "Uygun değil"}
FORMATS = {"ekitap": "E-kitap", "sesli": "Sesli kitap"}
LISTING_STATES = {"bilinmiyor": "Bilinmiyor", "hazirlaniyor": "Hazırlanıyor", "yuklendi": "Yüklendi", "yayinda": "Yayında",
                  "reddedildi": "Reddedildi", "kaldirildi": "Kaldırıldı"}
LIVE_STATES = ("yuklendi", "yayinda")
PLATFORM_KINDS = {"ekitap": "E-kitap", "sesli": "Sesli kitap", "abonelik": "Abonelik"}
DISTRIBUTION = {"dogrudan": "Doğrudan", "dagitici": "Dağıtıcı üzerinden"}
STUDIO_STATES = {"yok": "Stüdyoda iş yok", "is_var": "Stüdyoda iş var, e-kitap üretilmedi", "hazirlaniyor": "E-kitap hazırlanıyor",
                 "hazir": "E-kitap hazır", "eski": "E-kitap eski (içerik değişti)", "hata": "E-kitap üretimi hata verdi"}
NOTE_CHOICES = {"kisitliyor": "dijitali kısıtlıyor", "kisitlamiyor": "dijitali kısıtlamıyor", "belirsiz": "belirsiz"}
PENDING_FIELDS = {"new_ekitapisbn": "E-Kitap ISBN", "new_EPubDurumu": "E-Pub Durumu", "new_EKitapStokKodu": "E-Kitap Stok Kodu"}
#: Yayın durumu etiketinin kodu (YS05 …) satış dışıysa fırsat listesine girmez (SEO/GEO ile aynı liste).
OUT_OF_SALE = set(seo_crm.STATUS_FLAGS)
#: CRM `new_Tip` 8 = «Ekitap» kartı (sesli kitap kartı ayarla: `DIJITAL_AUDIO_TYPES`).
EBOOK_TYPE = 8

DEFAULTS: dict[str, str] = {
    "CRM_SCHEMA": "Timas_MSCRM.dbo",
    "DIJITAL_OPP_MIN_QTY": "1000",
    "DIJITAL_AUDIO_MIN_QTY": "1000",
    "DIJITAL_AUDIO_GENRES": "",
    "DIJITAL_BOOK_TYPES": "1",
    "DIJITAL_AUDIO_TYPES": "9",
    "DIJITAL_MATCH_MIN_PROB": "0.6",
    "DIJITAL_MATCH_CANDIDATES": "8",
    "DIJITAL_EDITION_DAYS": "90",
    "DIJITAL_WEEKLY_DAY": "1",
    "DIJITAL_NOTE_BUDGET_SEC": "1800",
    "DIJITAL_RIGHTS_RECIPIENTS": "",
    "DIJITAL_ALERT_RECIPIENTS": "",
    "DIJITAL_FINANCE_RECIPIENTS": "",
    "DIJITAL_IMPORT_MAX_MB": "40",
}

_lock = threading.Lock()
_ready: set[int] = set()


class DigitalError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def ensure(engine: sa.engine.Engine) -> None:
    key = id(engine)
    with _lock:
        if key in _ready:
            return
        _md.create_all(engine, checkfirst=True)
        _ready.add(key)


# ------------------------------------------------------------------ küçük yardımcılar


def _now() -> datetime:
    return datetime.now(timezone.utc)


def today() -> date:
    return datetime.now(TZ).date()


def _iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    return (v if v.tzinfo else v.replace(tzinfo=timezone.utc)).isoformat()


def _j(v: Optional[str], default: Any) -> Any:
    if not v:
        return default
    try:
        return json.loads(v)
    except ValueError:
        return default


def _dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _text(v: Any, limit: int) -> Optional[str]:
    s = re.sub(r"\s+", " ", str(v or "")).strip()
    return s[:limit] or None


def fold(s: Any) -> str:
    t = str(s or "").translate(str.maketrans("İIıŞşĞğÜüÖöÇçÂâÎîÛû", "iiissgguuooccaaiiuu")).lower()
    return re.sub(r"\s+", " ", t).strip()


def tokens(s: Any) -> list[str]:
    return [t for t in re.split(r"[^0-9a-z]+", fold(s)) if len(t) > 1]


def note_hash(note: Any) -> str:
    return hashlib.sha1(fold(note).encode("utf-8")).hexdigest()[:16]


def int_list(raw: Any) -> list[int]:
    return [int(x) for x in re.split(r"[,; ]+", str(raw or "")) if x.strip().lstrip("-").isdigit()]


def settings_from(conf: Callable[[str, str], str]) -> dict[str, Any]:
    def g(key: str) -> str:
        v = conf(key, DEFAULTS[key])
        return DEFAULTS[key] if v is None or str(v).strip() == "" and DEFAULTS[key] else str(v)

    def gi(key: str, lo: int, hi: int) -> int:
        try:
            return max(lo, min(hi, int(float(g(key)))))
        except ValueError:
            return int(DEFAULTS[key])

    def gf(key: str) -> float:
        try:
            return float(g(key).replace(",", "."))
        except ValueError:
            return float(DEFAULTS[key])

    return {
        "schema": g("CRM_SCHEMA"),
        "oppMinQty": gf("DIJITAL_OPP_MIN_QTY"),
        "audioMinQty": gf("DIJITAL_AUDIO_MIN_QTY"),
        "audioGenres": [fold(x) for x in re.split(r"[,;]", g("DIJITAL_AUDIO_GENRES")) if x.strip()],
        "bookTypes": int_list(g("DIJITAL_BOOK_TYPES")) or [1],
        "audioTypes": int_list(g("DIJITAL_AUDIO_TYPES")) or [9],
        "matchMinProb": min(1.0, max(0.0, gf("DIJITAL_MATCH_MIN_PROB"))),
        "matchCandidates": gi("DIJITAL_MATCH_CANDIDATES", 2, 25),
        "editionDays": gi("DIJITAL_EDITION_DAYS", 1, 3650),
        "weeklyDay": gi("DIJITAL_WEEKLY_DAY", 1, 7),
        "noteBudgetSec": gi("DIJITAL_NOTE_BUDGET_SEC", 0, 86_400),
        "rightsRecipients": _mails(g("DIJITAL_RIGHTS_RECIPIENTS")),
        "alertRecipients": _mails(g("DIJITAL_ALERT_RECIPIENTS")),
        "financeRecipients": _mails(g("DIJITAL_FINANCE_RECIPIENTS")),
        "importMaxMb": gi("DIJITAL_IMPORT_MAX_MB", 1, 500),
    }


def _mails(raw: str) -> list[str]:
    return [x.strip() for x in re.split(r"[,;\s]+", raw or "") if "@" in x]


def meta_get(engine: sa.engine.Engine, tenant: str, key: str) -> dict[str, Any]:
    with engine.connect() as c:
        row = c.execute(sa.select(META).where(META.c.tenant_id == tenant, META.c.key == key)).first()
    if not row:
        return {}
    return {**_j(row.value_json, {}), "_at": _iso(row.updated_at)}


def meta_set(engine: sa.engine.Engine, tenant: str, key: str, value: dict[str, Any]) -> None:
    now = _now()
    value = {k: v for k, v in value.items() if k != "_at"}
    with engine.begin() as c:
        cond = (META.c.tenant_id == tenant) & (META.c.key == key)
        if c.execute(sa.select(META.c.key).where(cond)).first():
            c.execute(META.update().where(cond).values(value_json=_dump(value), updated_at=now))
        else:
            c.execute(META.insert().values(tenant_id=tenant, key=key, value_json=_dump(value), updated_at=now))


# ------------------------------------------------------------------ hak kararı


def rights_for(contracts: list[dict[str, Any]], ref: date, bicim: str,
               decisions: Optional[dict[tuple[str, str], dict[str, Any]]] = None) -> tuple[str, str, bool]:
    """(karar, gerekçe, hak notu var mı). `contracts`: kitaba bağlı Telif Alış sözleşmeleri (`seo_geo.crm.contract_sql`
    satırları + `parties`). `decisions`: (sozlesme_id, bicim) → son telif kararı (`karar`, `gerekce`, `not_ozeti`)."""
    flag = "ebook" if bicim == "ekitap" else "audiobook"
    what = FORMATS[bicim].lower()
    decisions = decisions or {}
    live = [c for c in contracts if seo_crm.in_force(c, ref)]
    noted = any(c.get("rights_note") for c in live)
    if not live:
        if any(c.get("public_domain") for c in contracts):
            return "koruma_disi", "Koruma dışı eser; telif sözleşmesi gerekmez.", False
        if contracts:
            return "yok", "Telif alış sözleşmelerinin hiçbiri yürürlükte değil (süresi dolmuş ya da feshedilmiş).", False
        return "yok", "Kitaba bağlı telif alış sözleşmesi yok.", False
    missing = [c for c in live if not c.get(flag)]
    if missing:
        who = ", ".join(sorted({x for c in missing for x in c.get("parties") or []})) or "bir taraf"
        return "eksik", f"{FORMATS[bicim]} hakkı şu sözleşmede yok: {who}.", noted
    open_notes, kinds, reasons = [], [], []
    for c in live:
        note = c.get("rights_note")
        if not note:
            continue
        d = decisions.get((str(c.get("id") or "").upper(), bicim))
        if not d or d.get("not_ozeti") != note_hash(note):
            open_notes.append(c)
            continue
        kinds.append(d["karar"])
        reasons.append(str(d.get("gerekce") or ""))
    if open_notes:
        who = ", ".join(sorted({x for c in open_notes for x in c.get("parties") or []})) or "bir sözleşme"
        return "incele", f"Hak notu var ({who}); telif birimi {what} için karar vermeli.", True
    if "uygun_degil" in kinds:
        return "yok", "Telif birimi kararı: uygun değil. " + " ".join(r for r in reasons if r), True
    if "kismi" in kinds:
        return "kismi", "Telif birimi kararı: kısıtlı hak. " + " ".join(r for r in reasons if r), True
    if kinds:
        return "var", f"Hak notu telif biriminde incelendi, {what} için uygun.", True
    return "var", f"Yürürlükteki bütün telif alış sözleşmelerinde {what} hakkı var.", False


def decisions_map(engine: sa.engine.Engine, tenant: str) -> dict[tuple[str, str], dict[str, Any]]:
    """(sozlesme_id, bicim) → son karar."""
    out: dict[tuple[str, str], dict[str, Any]] = {}
    with engine.connect() as c:
        rows = c.execute(sa.select(DECISIONS).where(DECISIONS.c.tenant_id == tenant).order_by(DECISIONS.c.id)).all()
    for r in rows:
        out[(r.sozlesme_id, r.bicim)] = {"karar": r.karar, "gerekce": r.gerekce, "not_ozeti": r.not_ozeti, "yazan": r.yazan,
                                         "tarih": _iso(r.created_at), "kitap_id": r.kitap_id}
    return out


# ------------------------------------------------------------------ fırsat


def percentile_ranks(values: dict[str, float]) -> dict[str, float]:
    """Kimlik → 0–100 yüzdelik sıra (yalnız pozitif değerler arasında; eşitler aynı sırayı alır)."""
    pos = sorted(v for v in values.values() if v > 0)
    n = len(pos)
    if not n:
        return {}
    out = {}
    for k, v in values.items():
        if v <= 0:
            continue
        lo = _bisect_left(pos, v)
        out[k] = round(100.0 * lo / n, 1) if n > 1 else 100.0
    return out


def _bisect_left(a: list[float], x: float) -> int:
    lo, hi = 0, len(a)
    while lo < hi:
        mid = (lo + hi) // 2
        if a[mid] < x:
            lo = mid + 1
        else:
            hi = mid
    return lo


def _status_code(label: Optional[str]) -> str:
    return (label or "").split(" ", 1)[0].upper()


def _fmt_int(v: float) -> str:
    return f"{int(round(v)):,}".replace(",", ".")


def opportunity(t: dict[str, Any], bicim: str, pct: Optional[float], st: dict[str, Any]) -> Optional[str]:
    """Fırsatsa gerekçe cümlesi (rakamlar okumadan), değilse None."""
    if t.get("tip") not in st["bookTypes"] or _status_code(t.get("yayin_durumu")) in OUT_OF_SALE:
        return None
    hak = t.get("hak_ekitap" if bicim == "ekitap" else "hak_sesli")
    if hak not in ("var", "kismi"):
        return None
    if t.get("ekitap_var" if bicim == "ekitap" else "sesli_var"):
        return None
    qty = float(t.get("basili_12ay_adet") or 0)
    if qty < (st["oppMinQty"] if bicim == "ekitap" else st["audioMinQty"]):
        return None
    if bicim == "sesli" and st["audioGenres"]:
        hay = fold(f"{t.get('turler') or ''} {t.get('hedef_kitle') or ''}")
        if not any(g in hay for g in st["audioGenres"]):
            return None
    parts = [f"Son 12 ayda {_fmt_int(qty)} adet basılı satış"
             + (f" (kitapların %{int(pct)}'inden fazla)" if pct is not None and pct >= 1 else "")]
    parts.append(f"{FORMATS[bicim].lower()} hakkı {'kısıtlı' if hak == 'kismi' else 'var'}")
    parts.append("e-kitap kaydı yok" if bicim == "ekitap" else "sesli kitap kaydı yok")
    if bicim == "ekitap" and t.get("studio_epub_durumu") == "hazir":
        parts.append("stüdyoda e-kitap dosyası hazır")
    return "; ".join(parts) + "."


# ------------------------------------------------------------------ yeni baskı / kapak


def edition_change(history: list[dict[str, Any]], since: date, after: Optional[str]) -> Optional[tuple[str, str]]:
    """Pencere içinde baskı sayısı artışı ya da kapak adresi değişikliği → (tarih, tür). `after`: dijital sürümün en son
    güncellendiği gün (platform kaydı); ondan önceki değişiklik sayılmaz."""
    prev_print, prev_cover, hit = None, None, None
    for h in history:
        kinds = []
        if h.get("baski") is not None:
            if prev_print is not None and h["baski"] > prev_print:
                kinds.append("yeni baskı")
            prev_print = h["baski"]
        if h.get("kapak"):
            if prev_cover is not None and h["kapak"] != prev_cover:
                kinds.append("kapak değişti")
            prev_cover = h["kapak"]
        if kinds and h["tarih"] >= since.isoformat() and (not after or h["tarih"] > after):
            hit = (h["tarih"], ", ".join(kinds))
    return hit


# ------------------------------------------------------------------ gece okuması


def _listing_state(engine: sa.engine.Engine, tenant: str) -> dict[tuple[str, int], Any]:
    """(kitap, platform) → son platform kaydı."""
    with engine.connect() as c:
        rows = c.execute(sa.select(LISTINGS).where(LISTINGS.c.tenant_id == tenant).order_by(LISTINGS.c.id)).all()
    out = {}
    for r in rows:
        out[(r.kitap_id, r.platform_id)] = r
    return out


def _platforms(engine: sa.engine.Engine, tenant: str) -> dict[int, Any]:
    with engine.connect() as c:
        return {r.id: r for r in c.execute(sa.select(PLATFORMS).where(PLATFORMS.c.tenant_id == tenant)).all()}


def live_formats(engine: sa.engine.Engine, tenant: str) -> dict[str, set[str]]:
    """Kitap → platformda yüklendi/yayında olduğu biçimler (platform türü sesli ise sesli, değilse e-kitap)."""
    plats = _platforms(engine, tenant)
    out: dict[str, set[str]] = {}
    for (kid, pid), r in _listing_state(engine, tenant).items():
        if r.durum in LIVE_STATES and pid in plats:
            out.setdefault(kid, set()).add("sesli" if plats[pid].tur == "sesli" else "ekitap")
    return out


def last_listing_day(engine: sa.engine.Engine, tenant: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for (kid, _pid), r in _listing_state(engine, tenant).items():
        d = r.tarih or (r.created_at.date().isoformat() if r.created_at else None)
        if d and d > out.get(kid, ""):
            out[kid] = d
    return out


def build_titles(data: dict[str, Any], st: dict[str, Any], ref: date, *, decisions: dict[tuple[str, str], dict[str, Any]],
                 live: dict[str, set[str]], last_listing: dict[str, str], sales12: Optional[dict[str, dict[str, float]]],
                 studio: Optional[dict[str, dict[str, Any]]], history: Optional[dict[str, list[dict[str, Any]]]],
                 external: Optional[dict[str, dict[str, tuple[str, str]]]] = None) -> list[dict[str, Any]]:
    """CRM okuması + Logo 12 ay satışı + stüdyo + portal kayıtları → kitap satırları. `sales12` ya da `studio` None ise o
    kaynak okunamadı demektir: ilgili kolon çağıranın elindeki eski değerle doldurulur (`merge_previous`)."""
    books = data["books"]
    audio_names = {fold(b["ad"]) for b in books if b.get("tip") in st["audioTypes"] and b.get("ad")}
    audio_names |= {re.sub(r"\s*\(?sesli kitap\)?\s*$", "", n).strip() for n in audio_names}
    since = ref - timedelta(days=st["editionDays"])
    rows = []
    for b in books:
        kid = b["kitap_id"]
        contracts = data["contracts"].get(kid, [])
        hak_e, why_e, noted = rights_for(contracts, ref, "ekitap", decisions)
        hak_s, why_s, _ = rights_for(contracts, ref, "sesli", decisions)
        hak_e, why_e = seo_crm.by_kind(hak_e, why_e, b.get("tip_adi"), b.get("yayin_durumu"), bool(contracts))
        hak_s, why_s = seo_crm.by_kind(hak_s, why_s, b.get("tip_adi"), b.get("yayin_durumu"), bool(contracts))
        source = "crm"
        if external is not None:
            if kid in (external.get("ekitap") or {}):
                hak_e, why_e = external["ekitap"][kid]
                source = "hak-haritasi"
            if kid in (external.get("sesli") or {}):
                hak_s, why_s = external["sesli"][kid]
                source = "hak-haritasi"
        fmts = live.get(kid, set())
        s12 = (sales12 or {}).get(src.code_key(b.get("stok_kodu"))) if b.get("stok_kodu") else None
        e12 = (sales12 or {}).get(src.code_key(b.get("ekitap_stok_kodu"))) if b.get("ekitap_stok_kodu") else None
        stu = (studio or {}).get(kid) if studio is not None else None
        prod = data.get("production", {}).get(kid) or {}
        name = fold(b.get("ad"))
        row = {
            "kitap_id": kid, "ad": b.get("ad"), "yazar": b.get("yazar"), "stok_kodu": b.get("stok_kodu"), "isbn": b.get("isbn"),
            "ean": b.get("ean"), "e_isbn": b.get("e_isbn"), "ekitap_stok_kodu": b.get("ekitap_stok_kodu"),
            "ekitap_barkod": b.get("ekitap_barkod"), "tip": b.get("tip"), "tip_adi": b.get("tip_adi"),
            "yayin_durumu": b.get("yayin_durumu"), "hedef_kitle": b.get("hedef_kitle"), "turler": b.get("turler"),
            "ilk_yayin": b.get("ilk_yayin"), "epub_durumu_crm": int(b.get("epub_crm") or 0),
            "uretim_durumu": prod.get("durum"), "uretim_tarih": prod.get("tarih"),
            "hak_ekitap": hak_e, "hak_ekitap_gerekce": why_e, "hak_sesli": hak_s, "hak_sesli_gerekce": why_s,
            "hak_kaynak": source, "hak_notu_var": bool(noted),
            "sozlesme_json": _dump([{"id": str(c.get("id") or "").upper(), "ad": c.get("name"), "taraflar": c.get("parties") or [],
                                     "yururlukte": seo_crm.in_force(c, ref), "bitis": str(src._day(c.get("ends")) or "") or None,
                                     "suresiz": bool(c.get("open_ended")), "ekitap": bool(c.get("ebook")),
                                     "sesli": bool(c.get("audiobook")), "zkitap": bool(c.get("zbook")),
                                     "iletim": bool(c.get("internet")), "korumaDisi": bool(c.get("public_domain")),
                                     "not": seo_crm.clean(c.get("rights_note"))} for c in contracts]),
            "basili_fiyat": b.get("basili_fiyat"),
            "ekitap_var": bool(b.get("ekitap_stok_kodu")) or "ekitap" in fmts,
            "sesli_var": (bool(name) and name in audio_names and b.get("tip") not in st["audioTypes"]) or "sesli" in fmts,
            "okundu_at": _now(),
        }
        if sales12 is not None:
            row.update(basili_12ay_adet=(s12 or {}).get("adet", 0.0), basili_12ay_ciro=(s12 or {}).get("ciro", 0.0),
                       logo_dijital_12ay_adet=(e12 or {}).get("adet", 0.0) if b.get("ekitap_stok_kodu") else None,
                       logo_dijital_12ay_ciro=(e12 or {}).get("ciro", 0.0) if b.get("ekitap_stok_kodu") else None)
        if studio is not None:
            row.update(studio_is=(stu or {}).get("is"), studio_epub_durumu=(stu or {}).get("durum") or "yok",
                       studio_denetim=(stu or {}).get("denetim"), studio_e_isbn=(stu or {}).get("e_isbn"))
        if history is not None:
            ch = edition_change(history.get(kid, []), since, last_listing.get(kid)) if row["ekitap_var"] or row["sesli_var"] else None
            row.update(baski_degisim_tarih=ch[0] if ch else None, baski_degisim_tur=ch[1] if ch else None)
        rows.append(row)
    return rows


def score(rows: list[dict[str, Any]], st: dict[str, Any]) -> None:
    """Fırsat puanı ve gerekçesi (yerinde). Puan: kitap tipli kartlar arasında basılı satışın yüzdelik sırası."""
    pct = percentile_ranks({r["kitap_id"]: float(r.get("basili_12ay_adet") or 0) for r in rows if r.get("tip") in st["bookTypes"]})
    for r in rows:
        p = pct.get(r["kitap_id"])
        e = opportunity(r, "ekitap", p, st)
        s = opportunity(r, "sesli", p, st)
        r.update(firsat_puani=p if e else None, firsat_gerekcesi=e, sesli_firsat_puani=p if s else None, sesli_firsat_gerekcesi=s)


USER_COLS = ("dijital_fiyat", "fiyat_gerekce", "fiyat_onaylayan", "fiyat_tarih")
KEEP_IF_MISSING = {
    "sales": ("basili_12ay_adet", "basili_12ay_ciro", "logo_dijital_12ay_adet", "logo_dijital_12ay_ciro"),
    "studio": ("studio_is", "studio_epub_durumu", "studio_denetim", "studio_e_isbn"),
    "history": ("baski_degisim_tarih", "baski_degisim_tur"),
}


def previous(engine: sa.engine.Engine, tenant: str) -> dict[str, dict[str, Any]]:
    with engine.connect() as c:
        return {r.kitap_id: dict(r._mapping) for r in c.execute(sa.select(TITLES).where(TITLES.c.tenant_id == tenant)).all()}


def merge_previous(rows: list[dict[str, Any]], prev: dict[str, dict[str, Any]], missing: Iterable[str]) -> None:
    """Okunamayan kaynağın kolonları bir önceki okumadan kalır (ekrana «bir önceki okuma» notu düşülür)."""
    cols = [c for m in missing for c in KEEP_IF_MISSING[m]]
    for r in rows:
        old = prev.get(r["kitap_id"]) or {}
        for c in cols:
            r[c] = old.get(c)


def write_titles(engine: sa.engine.Engine, tenant: str, rows: list[dict[str, Any]]) -> dict[str, int]:
    """Kitap satırlarını yazar: yeni kitap eklenir, var olanın CRM/Logo kolonları güncellenir (kullanıcı kolonları
    korunur), CRM'de artık etkin olmayan kitap silinir."""
    cols = [c.name for c in TITLES.columns if c.name not in USER_COLS and c.name not in ("tenant_id",)]
    with engine.begin() as c:
        have = {r[0] for r in c.execute(sa.select(TITLES.c.kitap_id).where(TITLES.c.tenant_id == tenant)).all()}
        new = [r for r in rows if r["kitap_id"] not in have]
        upd = [r for r in rows if r["kitap_id"] in have]
        if new:
            c.execute(TITLES.insert(), [{"tenant_id": tenant, **{k: r.get(k) for k in cols}} for r in new])
        if upd:
            # bindparam adı kolon adıyla aynı olamaz (SQLAlchemy SET için ayırır): «v_» önekli.
            stmt = TITLES.update().where(TITLES.c.tenant_id == sa.bindparam("w_t"), TITLES.c.kitap_id == sa.bindparam("w_k")) \
                .values({k: sa.bindparam("v_" + k) for k in cols if k != "kitap_id"})
            c.execute(stmt, [{"w_t": tenant, "w_k": r["kitap_id"], **{"v_" + k: r.get(k) for k in cols if k != "kitap_id"}} for r in upd])
        gone = have - {r["kitap_id"] for r in rows}
        for part in _chunks(sorted(gone), 500):
            c.execute(TITLES.delete().where(TITLES.c.tenant_id == tenant, TITLES.c.kitap_id.in_(part)))
    return {"yeni": len(new), "guncellenen": len(upd), "silinen": len(gone)}


def _chunks(items: list[Any], n: int) -> Iterable[list[Any]]:
    for i in range(0, len(items), n):
        yield items[i:i + n]


def pending_from(rows: list[dict[str, Any]], live: dict[str, set[str]]) -> set[tuple[str, str, str, str]]:
    """CRM'e işlenecekler: (kitap, alan, değer, kaynak)."""
    out = set()
    for r in rows:
        kid = r["kitap_id"]
        se = r.get("studio_e_isbn")
        if se and src.isbn_key(se) and src.isbn_key(se) != src.isbn_key(r.get("e_isbn")):
            out.add((kid, "new_ekitapisbn", se, "Stüdyoda üretilen e-kitaptaki e-ISBN"))
        if r.get("studio_epub_durumu") == "hazir" and r.get("studio_denetim") in ("OK", "WARN") and int(r.get("epub_durumu_crm") or 0) != 1:
            out.add((kid, "new_EPubDurumu", "Evet", "Stüdyoda e-kitap hazır, denetimden geçti"))
        if "ekitap" in live.get(kid, set()) and not r.get("ekitap_stok_kodu"):
            out.add((kid, "new_EKitapStokKodu", "açılmalı", "Kitap bir e-kitap platformunda yüklendi/yayında"))
    return out


def write_pending(engine: sa.engine.Engine, tenant: str, want: set[tuple[str, str, str, str]]) -> dict[str, int]:
    now = _now()
    added = closed = 0
    with engine.begin() as c:
        open_rows = c.execute(sa.select(PENDING).where(PENDING.c.tenant_id == tenant, PENDING.c.durum == "acik")).all()
        have = {(r.kitap_id, r.alan, r.deger): r.id for r in open_rows}
        keys = {(k, a, d) for k, a, d, _ in want}
        for k, a, d, why in sorted(want):
            if (k, a, d) not in have:
                c.execute(PENDING.insert().values(tenant_id=tenant, kitap_id=k, alan=a, deger=d, kaynak=why, durum="acik", created_at=now))
                added += 1
        for key, rid in have.items():
            if key not in keys:
                c.execute(PENDING.update().where(PENDING.c.id == rid).values(durum="kapandi", closed_at=now))
                closed += 1
    return {"eklenen": added, "kapanan": closed}


def sales_by_code(rows: list[dict[str, Any]]) -> tuple[dict[str, dict[str, float]], dict[str, dict[str, dict[str, float]]]]:
    """Logo satır listesi → (kod → 12 ay toplamı, kod → ay → {adet, ciro})."""
    total: dict[str, dict[str, float]] = {}
    monthly: dict[str, dict[str, dict[str, float]]] = {}
    for r in rows:
        k = src.code_key(r["stok_kodu"])
        t = total.setdefault(k, {"adet": 0.0, "ciro": 0.0})
        t["adet"] += r["adet"]
        t["ciro"] += r["ciro"]
        m = monthly.setdefault(k, {}).setdefault(r["ay"], {"adet": 0.0, "ciro": 0.0})
        m["adet"] += r["adet"]
        m["ciro"] += r["ciro"]
    return total, monthly


class Refresher:
    """Gece okuması (ve «Yenile» düğmesi): kaynakları okur, `semantic_dijital_titles`'ı ve CRM'e işlenecekleri yazar.
    Aynı anda tek okuma; ikinci istek «sürüyor» döner."""

    def __init__(self, engine: Callable[[], Any], tenant: Callable[[], str], crm: Callable[[], src.Runner],
                 logo: Callable[[], src.Runner], studio_jobs: Optional[Callable[[], Any]],
                 studio_view: Optional[Callable[[str], dict]], settings: Callable[[], dict[str, Any]]):
        self.engine, self.tenant, self.crm, self.logo = engine, tenant, crm, logo
        self.studio_jobs, self.studio_view, self.settings = studio_jobs, studio_view, settings
        self._lock = threading.Lock()
        self._running: Optional[float] = None

    def status(self) -> dict[str, Any]:
        engine, tenant = self.engine(), self.tenant()
        ensure(engine)
        info = meta_get(engine, tenant, "refresh")
        return {**info, "running": self._running is not None, "since": self._running}

    def start(self) -> dict[str, Any]:
        if self._running is not None:
            return {"started": False, **self.status()}
        threading.Thread(target=self._safe_run, name="dijital-refresh", daemon=True).start()
        return {"started": True}

    def _safe_run(self) -> None:
        try:
            self.run()
        except Exception as e:  # noqa: BLE001 — hata ekrana ve loga; önceki katalog kalır
            log.warning("dijital okuması hata verdi: %s", e)

    def run(self) -> dict[str, Any]:
        if not self._lock.acquire(blocking=False):
            raise DigitalError("Okuma zaten sürüyor.", 409)
        self._running = time.time()
        engine, tenant = self.engine(), self.tenant()
        ensure(engine)
        st = self.settings()
        t0 = time.monotonic()
        notes: list[str] = []
        try:
            try:
                crm_run = self.crm()
                data = src.read_crm(crm_run, st["schema"])
            except src.SourceError as e:
                info = {"ok": False, "error": f"CRM okunamadı: {e}", "at": _iso(_now())}
                meta_set(engine, tenant, "refresh", info)
                raise
            ref = today()
            missing: list[str] = []
            history = None
            try:
                history = src.read_history(crm_run, st["schema"], ref - timedelta(days=st["editionDays"]))
            except src.SourceError as e:
                notes.append(f"Kitap geçmişi okunamadı ({e}); yeni baskı işaretleri bir önceki okumadan.")
                missing.append("history")
            sales12 = None
            logo_meta: dict[str, Any] = {}
            try:
                run = self.logo()
                firms = src.firms_by_year(run)
                end = src.data_end(run, firms)
                if end is None:
                    raise src.SourceError("Logo'da satış satırı yok.")
                start, stop = src.window12(end)
                rows = src.read_sales(run, firms, start, stop)
                sales12, monthly = sales_by_code(rows)
                ebook_codes = {src.code_key(b["ekitap_stok_kodu"]): b["kitap_id"] for b in data["books"] if b.get("ekitap_stok_kodu")}
                logo_meta = {"veriSonu": end.isoformat(), "pencere": [start.isoformat(), (stop - timedelta(days=1)).isoformat()],
                             "ekitapSatis": {code: {"kitap_id": kid, "aylar": monthly[code]} for code, kid in ebook_codes.items() if code in monthly}}
            except src.SourceError as e:
                notes.append(f"Logo okunamadı ({e}); basılı satış ve Logo e-kitap satışı bir önceki okumadan.")
                missing.append("sales")
            studio = src.read_studio(self.studio_jobs, self.studio_view)
            if studio is None:
                notes.append("Kitap Tasarım Stüdyosu okunamadı; stüdyo e-kitap durumu bir önceki okumadan.")
                missing.append("studio")
            decisions = decisions_map(engine, tenant)
            live = live_formats(engine, tenant)
            ids = [b["kitap_id"] for b in data["books"]]
            ext_e, ext_s = src.external_rights(ids, "ekitap"), src.external_rights(ids, "sesli")
            external = {"ekitap": ext_e or {}, "sesli": ext_s or {}} if (ext_e is not None or ext_s is not None) else None
            titles = build_titles(data, st, ref, decisions=decisions, live=live, last_listing=last_listing_day(engine, tenant),
                                  sales12=sales12, studio=studio, history=history, external=external)
            if missing:
                merge_previous(titles, previous(engine, tenant), missing)
            score(titles, st)
            written = write_titles(engine, tenant, titles)
            pend = write_pending(engine, tenant, pending_from(titles, live))
            if logo_meta:
                meta_set(engine, tenant, "logo", logo_meta)
            meta_set(engine, tenant, "crm_counts", {**data["counts"], "okundu": _iso(_now())})
            info = {"ok": True, "at": _iso(_now()), "sure": round(time.monotonic() - t0, 1), "kitap": len(titles),
                    "yazilan": written, "crmIslenecek": pend, "notlar": notes,
                    "hakKaynagi": "telif hak haritası" if external is not None else "CRM sözleşmeleri"}
            meta_set(engine, tenant, "refresh", info)
            return info
        finally:
            self._running = None
            self._lock.release()


# ------------------------------------------------------------------ hak notu ön okuması (Zeki AI)


def read_notes(engine: sa.engine.Engine, tenant: str, llm: Any, budget_sec: int) -> dict[str, Any]:
    """Yürürlükteki sözleşmelerin okunmamış (ya da metni değişmiş) hak notlarını **ortak sınıflamaya** gönderir
    (`rights_notes`: M54 telif ile tek soru, tek tablo; telifin sorduğu ya da onayladığı not yeniden sorulmaz). Süre
    bütçesi biterse kalan sonraki geceye kalır (sessiz tavan değil: kalan sayısı döner)."""
    if llm is None or not hasattr(llm, "choose"):
        return {"okunan": 0, "kalan": None, "model": False}
    with engine.connect() as c:
        rows = c.execute(sa.select(TITLES.c.ad, TITLES.c.sozlesme_json)
                         .where(TITLES.c.tenant_id == tenant, TITLES.c.hak_notu_var.is_(True))).all()
    crm_rows = [{"id": ct["id"], "no": ct.get("ad"), "kitap": r.ad, "metin": ct["not"]}
                for r in rows for ct in _j(r.sozlesme_json, []) if ct.get("not") and ct.get("id") and ct.get("yururlukte")]
    todo = RN.pending(engine, tenant, crm_rows)
    out = RN.classify(engine, tenant, todo, llm, budget_sec=budget_sec, stop_on_error=True)
    return {"okunan": out["okunan"], "kalan": out["kalan"], "model": True}


def note_reads(engine: sa.engine.Engine, tenant: str) -> dict[str, dict[str, Any]]:
    """Ortak sınıflamadan dijital okuma; `ozet` bu modülün not özetiyle karşılaştırılır (metin değiştiyse eşleşmez)."""
    return {k.upper(): {"sonuc": v["dijital"], "olasilik": v["olasilik"], "ozet": note_hash(v["metin"]), "sinif": v["sinif"],
                        "durum": v["durum"]} for k, v in RN.reads(engine, tenant).items()}


# ------------------------------------------------------------------ okuma uçları


def _title_dict(r: Any, plats: dict[int, Any], listings: dict[tuple[str, int], Any], *, full: bool = False) -> dict[str, Any]:
    chips = []
    for pid, p in sorted(plats.items(), key=lambda x: x[1].ad):
        l = listings.get((r.kitap_id, pid))
        if l is None and not p.aktif:
            continue
        chips.append({"platformId": pid, "platform": p.ad, "tur": p.tur, "durum": l.durum if l else "bilinmiyor",
                      "durumAdi": LISTING_STATES[l.durum if l else "bilinmiyor"], "tarih": l.tarih if l else None})
    out = {
        "kitapId": r.kitap_id, "ad": r.ad, "yazar": r.yazar, "stokKodu": r.stok_kodu, "isbn": r.isbn, "eIsbn": r.e_isbn,
        "ekitapStokKodu": r.ekitap_stok_kodu, "ekitapBarkod": r.ekitap_barkod, "tip": r.tip, "tipAdi": r.tip_adi,
        "yayinDurumu": r.yayin_durumu, "hedefKitle": r.hedef_kitle, "epubCrm": bool(r.epub_durumu_crm),
        "uretimDurumu": r.uretim_durumu, "studioDurumu": r.studio_epub_durumu or "yok",
        "studioDurumuAdi": STUDIO_STATES.get(r.studio_epub_durumu or "yok"), "studioDenetim": r.studio_denetim,
        "hakEkitap": r.hak_ekitap, "hakEkitapAdi": RIGHTS.get(r.hak_ekitap or "", r.hak_ekitap), "hakSesli": r.hak_sesli,
        "hakSesliAdi": RIGHTS.get(r.hak_sesli or "", r.hak_sesli), "hakNotu": bool(r.hak_notu_var),
        "ekitapVar": bool(r.ekitap_var), "sesliVar": bool(r.sesli_var), "basiliFiyat": r.basili_fiyat,
        "dijitalFiyat": r.dijital_fiyat, "basili12Adet": r.basili_12ay_adet, "logoDijital12Adet": r.logo_dijital_12ay_adet,
        "firsatPuani": r.firsat_puani, "sesliFirsatPuani": r.sesli_firsat_puani,
        "baskiDegisim": {"tarih": r.baski_degisim_tarih, "tur": r.baski_degisim_tur} if r.baski_degisim_tarih else None,
        "platformlar": chips,
    }
    if full:
        out.update({
            "ean": r.ean, "turler": r.turler, "ilkYayin": r.ilk_yayin, "uretimTarih": r.uretim_tarih, "studioIs": r.studio_is,
            "studioEIsbn": r.studio_e_isbn, "hakEkitapGerekce": r.hak_ekitap_gerekce, "hakSesliGerekce": r.hak_sesli_gerekce,
            "hakKaynak": "telif hak haritası" if r.hak_kaynak == "hak-haritasi" else "CRM sözleşmeleri",
            "sozlesmeler": _j(r.sozlesme_json, []), "basili12Ciro": r.basili_12ay_ciro, "logoDijital12Ciro": r.logo_dijital_12ay_ciro,
            "firsatGerekcesi": r.firsat_gerekcesi, "sesliFirsatGerekcesi": r.sesli_firsat_gerekcesi,
            "fiyatGerekce": r.fiyat_gerekce, "fiyatOnaylayan": r.fiyat_onaylayan, "fiyatTarih": _iso(r.fiyat_tarih),
            "okundu": _iso(r.okundu_at),
        })
    return out


def _book_filter(tur: str, st: dict[str, Any]):
    if tur == "hepsi":
        return sa.true()
    if tur == "dijital":
        return TITLES.c.tip.in_([*st["audioTypes"], EBOOK_TYPE])
    return TITLES.c.tip.in_(st["bookTypes"])


def list_titles(engine: sa.engine.Engine, tenant: str, st: dict[str, Any], *, hak: str = "", durum: str = "", q: str = "",
                tur: str = "kitap", platform: int = 0, page: int = 0) -> dict[str, Any]:
    """Dijital katalog. `hak`: e-kitap hak kararı; `durum`: dijitalde|dijitalde-yok|risk|firsat|epub-hazir|crm-islenecek|
    yeni-baski; `platform`: o platformdaki durum süzgeci için (durum ile birlikte `platform-<durum>`). Sayfa 50 satır;
    toplam ayrıca döner (tavan değil, sayfalama)."""
    cond = [TITLES.c.tenant_id == tenant, _book_filter(tur, st)]
    if hak:
        cond.append(TITLES.c.hak_ekitap == hak)
    if q.strip():
        like = f"%{q.strip()}%"
        cond.append(sa.or_(TITLES.c.ad.ilike(like), TITLES.c.yazar.ilike(like), TITLES.c.stok_kodu.ilike(like),
                           TITLES.c.isbn.ilike(like), TITLES.c.e_isbn.ilike(like), TITLES.c.ekitap_stok_kodu.ilike(like)))
    if durum == "dijitalde":
        cond.append(TITLES.c.ekitap_var.is_(True))
    elif durum == "dijitalde-yok":
        cond.append(sa.and_(sa.or_(TITLES.c.ekitap_var.is_(False), TITLES.c.ekitap_var.is_(None)), TITLES.c.hak_ekitap.in_(OK_RIGHTS)))
    elif durum == "risk":
        cond.append(sa.or_(sa.and_(TITLES.c.hak_ekitap.in_(RISK_RIGHTS), sa.or_(TITLES.c.ekitap_var.is_(True), TITLES.c.logo_dijital_12ay_adet > 0)),
                           sa.and_(TITLES.c.hak_sesli.in_(RISK_RIGHTS), TITLES.c.sesli_var.is_(True))))
    elif durum == "firsat":
        cond.append(TITLES.c.firsat_puani.is_not(None))
    elif durum == "epub-hazir":
        cond.append(TITLES.c.studio_epub_durumu == "hazir")
    elif durum == "yeni-baski":
        cond.append(TITLES.c.baski_degisim_tarih.is_not(None))
    elif durum == "crm-islenecek":
        sub = sa.select(PENDING.c.kitap_id).where(PENDING.c.tenant_id == tenant, PENDING.c.durum == "acik")
        cond.append(TITLES.c.kitap_id.in_(sub))
    listings = _listing_state(engine, tenant)
    if platform:
        ids = [k for (k, p), r in listings.items() if p == platform and r.durum in LIVE_STATES]
        cond.append(TITLES.c.kitap_id.in_(ids or ["-"]))
    page = max(0, int(page))
    with engine.connect() as c:
        total = c.execute(sa.select(sa.func.count()).select_from(TITLES).where(*cond)).scalar() or 0
        rows = c.execute(sa.select(TITLES).where(*cond)
                         .order_by(sa.desc(sa.func.coalesce(TITLES.c.basili_12ay_adet, 0)), TITLES.c.ad)
                         .offset(page * PAGE).limit(PAGE)).all()
    plats = _platforms(engine, tenant)
    return {"items": [_title_dict(r, plats, listings) for r in rows], "total": total, "page": page, "pageSize": PAGE}


def _title_row(engine: sa.engine.Engine, tenant: str, kitap_id: str) -> Any:
    with engine.connect() as c:
        r = c.execute(sa.select(TITLES).where(TITLES.c.tenant_id == tenant, TITLES.c.kitap_id == str(kitap_id).upper())).first()
    if r is None:
        raise DigitalError("Kitap dijital katalogda yok (CRM'de etkin değil ya da henüz okunmadı).", 404)
    return r


def get_title(engine: sa.engine.Engine, tenant: str, kitap_id: str, *, with_sales: bool) -> dict[str, Any]:
    r = _title_row(engine, tenant, kitap_id)
    plats = _platforms(engine, tenant)
    out = _title_dict(r, plats, _listing_state(engine, tenant), full=True)
    reads = note_reads(engine, tenant)
    maps = RM.for_keys(engine, tenant, [ct["id"] for ct in out["sozlesmeler"] if ct.get("not")])
    for ct in out["sozlesmeler"]:
        # Yapılandırılmış hak haritası (M54 ile ortak; alanlar alıntılı, onay telif biriminde).
        ct["hakHaritasi"] = RM.digital_view(RM.matching(maps, ct["id"], ct["not"])) if ct.get("not") else None
        rd = reads.get(ct["id"]) if ct.get("not") else None
        ct["notOkuma"] = ({"sonuc": rd["sonuc"], "sonucAdi": NOTE_CHOICES.get(rd["sonuc"] or "", "okunamadı"), "olasilik": rd["olasilik"],
                           "sinif": rd["sinif"], "sinifAdi": RN.RY.NOTE_CLASSES.get(rd["sinif"] or ""),
                           "onayli": rd["durum"] == "onayli"}
                          if rd and rd["ozet"] == note_hash(ct["not"]) else None)
    decisions = decisions_map(engine, tenant)
    out["kararlar"] = [{"sozlesmeId": k[0], "bicim": k[1], "bicimAdi": FORMATS[k[1]], **v, "kararAdi": DECISIONS_KINDS.get(v["karar"])}
                       for k, v in decisions.items() if v.get("kitap_id") == r.kitap_id]
    with engine.connect() as c:
        hist = c.execute(sa.select(LISTINGS).where(LISTINGS.c.tenant_id == tenant, LISTINGS.c.kitap_id == r.kitap_id)
                         .order_by(sa.desc(LISTINGS.c.id))).all()
        pend = c.execute(sa.select(PENDING).where(PENDING.c.tenant_id == tenant, PENDING.c.kitap_id == r.kitap_id)
                         .order_by(sa.desc(PENDING.c.id))).all()
    out["platformGecmisi"] = [{"platform": plats[h.platform_id].ad if h.platform_id in plats else str(h.platform_id),
                               "durum": h.durum, "durumAdi": LISTING_STATES.get(h.durum), "tarih": h.tarih, "kaynak": h.kaynak,
                               "not": h.notlar, "fiyat": h.fiyat, "yazan": h.yazan, "zaman": _iso(h.created_at)} for h in hist]
    out["crmIslenecek"] = [_pending_dict(p, None) for p in pend]
    if with_sales:
        out["satis"] = title_sales(engine, tenant, r.kitap_id)
    return out


# ------------------------------------------------------------------ yazma uçları


def set_listing(engine: sa.engine.Engine, tenant: str, user: str, kitap_id: str, platform_id: int, body: dict[str, Any],
                *, kaynak: str = "kullanici") -> dict[str, Any]:
    r = _title_row(engine, tenant, kitap_id)
    plats = _platforms(engine, tenant)
    if int(platform_id) not in plats:
        raise DigitalError("Platform bulunamadı.", 404)
    durum = str(body.get("durum") or "")
    if durum not in LISTING_STATES:
        raise DigitalError("Platform durumu geçersiz.")
    tarih = str(body.get("tarih") or "")[:10] or today().isoformat()
    try:
        date.fromisoformat(tarih)
    except ValueError:
        raise DigitalError("Tarih YYYY-AA-GG biçiminde olmalı.") from None
    fiyat = body.get("fiyat")
    fiyat = _money(fiyat, "Platform fiyatı") if fiyat not in (None, "") else None
    p = plats[int(platform_id)]
    hak = r.hak_sesli if p.tur == "sesli" else r.hak_ekitap
    warn = None
    if durum in ("hazirlaniyor", "yuklendi", "yayinda") and hak not in OK_RIGHTS:
        warn = (f"Dikkat: bu kitabın {('sesli kitap' if p.tur == 'sesli' else 'e-kitap')} hakkı «{RIGHTS.get(hak or '', hak)}». "
                "Kayıt alındı; kitap hak riski listesine düşer.")
    with engine.begin() as c:
        c.execute(LISTINGS.insert().values(tenant_id=tenant, kitap_id=r.kitap_id, platform_id=int(platform_id), durum=durum,
                                           tarih=tarih, kaynak=kaynak, notlar=_text(body.get("not"), 2000), fiyat=fiyat,
                                           yazan=user, created_at=_now()))
        fmt = "sesli_var" if p.tur == "sesli" else "ekitap_var"
        if durum in LIVE_STATES:
            c.execute(TITLES.update().where(TITLES.c.tenant_id == tenant, TITLES.c.kitap_id == r.kitap_id).values({fmt: True}))
    out = get_title(engine, tenant, r.kitap_id, with_sales=False)
    out["uyari"] = warn
    return out


def _money(v: Any, label: str) -> float:
    if isinstance(v, (int, float)):
        n = float(v)
    else:
        n = parse_number(v)
        if n is None:
            raise DigitalError(f"{label} sayı olmalı.")
    if n < 0 or math.isinf(n) or n != n:
        raise DigitalError(f"{label} geçersiz.")
    return round(n, 4)


def decide_rights(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    """Telif biriminin «incele» kaydına kararı. Karar sözleşmenin o anki hak notuna bağlanır; not değişirse yeniden sorulur."""
    r = _title_row(engine, tenant, str(body.get("kitapId") or ""))
    bicim = str(body.get("bicim") or "")
    karar = str(body.get("karar") or "")
    gerekce = _text(body.get("gerekce"), 2000)
    if bicim not in FORMATS:
        raise DigitalError("Biçim e-kitap ya da sesli kitap olmalı.")
    if karar not in DECISIONS_KINDS:
        raise DigitalError("Karar uygun, kısıtlı ya da uygun değil olmalı.")
    if not gerekce:
        raise DigitalError("Gerekçe yazılmalı.")
    sid = str(body.get("sozlesmeId") or "").upper()
    contract = next((c for c in _j(r.sozlesme_json, []) if c.get("id") == sid), None)
    if contract is None:
        raise DigitalError("Sözleşme bu kitaba bağlı değil.", 404)
    if not contract.get("not"):
        raise DigitalError("Bu sözleşmede hak notu yok; karar bayraklardan verilir.", 409)
    with engine.begin() as c:
        c.execute(DECISIONS.insert().values(tenant_id=tenant, kitap_id=r.kitap_id, sozlesme_id=sid, bicim=bicim, karar=karar,
                                            gerekce=gerekce, not_ozeti=note_hash(contract["not"]), yazan=user, created_at=_now()))
    recompute_rights(engine, tenant, r.kitap_id)
    return get_title(engine, tenant, r.kitap_id, with_sales=False)


def recompute_rights(engine: sa.engine.Engine, tenant: str, kitap_id: str) -> None:
    """Karardan sonra kitabın hak kararını saklı sözleşmelerden yeniden hesaplar (gece okumasını beklemeden)."""
    r = _title_row(engine, tenant, kitap_id)
    if r.hak_kaynak == "hak-haritasi":
        return
    contracts = [{"id": c["id"], "status": 100000000 if c.get("yururlukte") else 0, "open_ended": 1, "ends": None,
                  "terminated": None, "ebook": c.get("ekitap"), "audiobook": c.get("sesli"), "public_domain": c.get("korumaDisi"),
                  "rights_note": c.get("not"), "parties": c.get("taraflar") or []} for c in _j(r.sozlesme_json, [])]
    dec = decisions_map(engine, tenant)
    he, we, _ = rights_for(contracts, today(), "ekitap", dec)
    hs, ws, _ = rights_for(contracts, today(), "sesli", dec)
    he, we = seo_crm.by_kind(he, we, r.tip_adi, r.yayin_durumu, bool(contracts))
    hs, ws = seo_crm.by_kind(hs, ws, r.tip_adi, r.yayin_durumu, bool(contracts))
    with engine.begin() as c:
        c.execute(TITLES.update().where(TITLES.c.tenant_id == tenant, TITLES.c.kitap_id == r.kitap_id)
                  .values(hak_ekitap=he, hak_ekitap_gerekce=we, hak_sesli=hs, hak_sesli_gerekce=ws))


def set_price(engine: sa.engine.Engine, tenant: str, user: str, kitap_id: str, body: dict[str, Any]) -> dict[str, Any]:
    """Dijital fiyat kararı (açıkça verilen yetkiyle). Platforma gönderilmez; kayıt ve CRM/platform girişinin dayanağıdır."""
    r = _title_row(engine, tenant, kitap_id)
    raw = body.get("fiyat")
    fiyat = None if raw in (None, "") else _money(raw, "Dijital fiyat")
    gerekce = _text(body.get("gerekce"), 2000)
    if fiyat is not None and not gerekce:
        raise DigitalError("Gerekçe yazılmalı.")
    with engine.begin() as c:
        c.execute(TITLES.update().where(TITLES.c.tenant_id == tenant, TITLES.c.kitap_id == r.kitap_id)
                  .values(dijital_fiyat=fiyat, fiyat_gerekce=gerekce if fiyat is not None else None,
                          fiyat_onaylayan=user if fiyat is not None else None, fiyat_tarih=_now() if fiyat is not None else None))
    return get_title(engine, tenant, r.kitap_id, with_sales=False)


# ------------------------------------------------------------------ platformlar


def _platform_dict(p: Any) -> dict[str, Any]:
    return {"id": p.id, "ad": p.ad, "tur": p.tur, "turAdi": PLATFORM_KINDS.get(p.tur), "dagitim": p.dagitim,
            "dagitimAdi": DISTRIBUTION.get(p.dagitim), "paraBirimi": p.para_birimi, "raporGunu": p.rapor_gunu,
            "aktif": bool(p.aktif)}


def list_platforms(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    plats = sorted(_platforms(engine, tenant).values(), key=lambda p: (not p.aktif, p.ad))
    return {"items": [_platform_dict(p) for p in plats]}


def _platform_values(body: dict[str, Any], partial: bool) -> dict[str, Any]:
    v: dict[str, Any] = {}
    if not partial or "ad" in body:
        ad = _text(body.get("ad"), 120)
        if not ad:
            raise DigitalError("Platform adı yazılmalı.")
        v["ad"] = ad
    if not partial or "tur" in body:
        tur = str(body.get("tur") or "ekitap")
        if tur not in PLATFORM_KINDS:
            raise DigitalError("Platform türü geçersiz.")
        v["tur"] = tur
    if not partial or "dagitim" in body:
        dg = str(body.get("dagitim") or "dogrudan")
        if dg not in DISTRIBUTION:
            raise DigitalError("Dağıtım biçimi geçersiz.")
        v["dagitim"] = dg
    if not partial or "paraBirimi" in body:
        pb = re.sub(r"[^A-Z]", "", str(body.get("paraBirimi") or "TRY").upper())[:8]
        if len(pb) != 3:
            raise DigitalError("Para birimi üç harfli kod olmalı (TRY, USD, EUR…).")
        v["para_birimi"] = pb
    if "raporGunu" in body:
        rg = body.get("raporGunu")
        if rg in (None, ""):
            v["rapor_gunu"] = None
        else:
            try:
                v["rapor_gunu"] = max(0, min(120, int(rg)))
            except (TypeError, ValueError):
                raise DigitalError("Rapor günü sayı olmalı.") from None
    if "aktif" in body:
        v["aktif"] = bool(body.get("aktif"))
    return v


def create_platform(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    v = _platform_values(body, partial=False)
    if any(fold(p.ad) == fold(v["ad"]) for p in _platforms(engine, tenant).values()):
        raise DigitalError("Bu adla bir platform zaten var.", 409)
    with engine.begin() as c:
        pid = c.execute(PLATFORMS.insert().values(tenant_id=tenant, olusturan=user, created_at=_now(), aktif=v.pop("aktif", True), **v)).inserted_primary_key[0]
    return _platform_dict(_platforms(engine, tenant)[pid])


def update_platform(engine: sa.engine.Engine, tenant: str, pid: int, body: dict[str, Any]) -> dict[str, Any]:
    plats = _platforms(engine, tenant)
    if int(pid) not in plats:
        raise DigitalError("Platform bulunamadı.", 404)
    v = _platform_values(body, partial=True)
    if "ad" in v and any(fold(p.ad) == fold(v["ad"]) and p.id != int(pid) for p in plats.values()):
        raise DigitalError("Bu adla bir platform zaten var.", 409)
    if v:
        with engine.begin() as c:
            c.execute(PLATFORMS.update().where(PLATFORMS.c.id == int(pid)).values(**v))
    return _platform_dict(_platforms(engine, tenant)[int(pid)])


# ------------------------------------------------------------------ göstergeler, fırsat, risk, CRM'e işlenecek


def overview(engine: sa.engine.Engine, tenant: str, st: dict[str, Any]) -> dict[str, Any]:
    books = TITLES.c.tip.in_(st["bookTypes"])
    t = TITLES.c
    with engine.connect() as c:
        def n(*cond) -> int:
            return int(c.execute(sa.select(sa.func.count()).select_from(TITLES).where(t.tenant_id == tenant, *cond)).scalar() or 0)

        kpi = {
            "kitap": n(books),
            "dijitalde": n(books, t.ekitap_var.is_(True)),
            "hakliDijitalYok": n(books, t.hak_ekitap.in_(OK_RIGHTS), sa.or_(t.ekitap_var.is_(False), t.ekitap_var.is_(None))),
            "firsat": n(t.firsat_puani.is_not(None)),
            "sesliFirsat": n(t.sesli_firsat_puani.is_not(None)),
            "hakRiski": n(sa.or_(sa.and_(t.hak_ekitap.in_(RISK_RIGHTS), sa.or_(t.ekitap_var.is_(True), t.logo_dijital_12ay_adet > 0)),
                                 sa.and_(t.hak_sesli.in_(RISK_RIGHTS), t.sesli_var.is_(True)))),
            "incele": n(sa.or_(t.hak_ekitap == "incele", t.hak_sesli == "incele")),
            "ekitapKaydi": n(t.tip == EBOOK_TYPE),
            "sesliKaydi": n(t.tip.in_(st["audioTypes"])),
            "epubCrmEvet": n(t.epub_durumu_crm == 1),
            "epubStudyoHazir": n(t.studio_epub_durumu == "hazir"),
            "yeniBaski": n(t.baski_degisim_tarih.is_not(None)),
            "crmIslenecek": int(c.execute(sa.select(sa.func.count()).select_from(PENDING)
                                          .where(PENDING.c.tenant_id == tenant, PENDING.c.durum == "acik")).scalar() or 0),
        }
        hak = {k: v for k, v in c.execute(sa.select(t.hak_ekitap, sa.func.count()).where(t.tenant_id == tenant, books)
                                          .group_by(t.hak_ekitap)).all() if k}
        hak_s = {k: v for k, v in c.execute(sa.select(t.hak_sesli, sa.func.count()).where(t.tenant_id == tenant, books)
                                            .group_by(t.hak_sesli)).all() if k}
        last = c.execute(sa.select(IMPORTS.c.donem, IMPORTS.c.platform_id).where(IMPORTS.c.tenant_id == tenant, IMPORTS.c.durum == "onaylandi")
                         .order_by(sa.desc(IMPORTS.c.donem))).first()
    plats = _platforms(engine, tenant)
    return {
        "kpi": kpi, "hakDagilimi": {"ekitap": hak, "sesli": hak_s}, "hakAdlari": RIGHTS,
        "sonRapor": {"donem": last.donem, "platform": plats[last.platform_id].ad if last.platform_id in plats else None} if last else None,
        "sozlesme": meta_get(engine, tenant, "crm_counts"), "logo": {k: v for k, v in meta_get(engine, tenant, "logo").items() if k != "ekitapSatis"},
        "okuma": meta_get(engine, tenant, "refresh"),
        "ayarlar": {"oppMinQty": st["oppMinQty"], "audioMinQty": st["audioMinQty"], "audioGenres": st["audioGenres"]},
    }


def opportunities(engine: sa.engine.Engine, tenant: str, *, tur: str = "ekitap", q: str = "", page: int = 0,
                  all_rows: bool = False) -> dict[str, Any]:
    col = TITLES.c.firsat_puani if tur != "sesli" else TITLES.c.sesli_firsat_puani
    why = TITLES.c.firsat_gerekcesi if tur != "sesli" else TITLES.c.sesli_firsat_gerekcesi
    cond = [TITLES.c.tenant_id == tenant, col.is_not(None)]
    if q.strip():
        like = f"%{q.strip()}%"
        cond.append(sa.or_(TITLES.c.ad.ilike(like), TITLES.c.yazar.ilike(like), TITLES.c.stok_kodu.ilike(like)))
    with engine.connect() as c:
        total = int(c.execute(sa.select(sa.func.count()).select_from(TITLES).where(*cond)).scalar() or 0)
        stmt = sa.select(TITLES, why.label("sec_gerekce")).where(*cond).order_by(sa.desc(TITLES.c.basili_12ay_adet), TITLES.c.ad)
        if not all_rows:
            stmt = stmt.offset(max(0, page) * PAGE).limit(PAGE)
        rows = c.execute(stmt).all()
    plats = _platforms(engine, tenant)
    listings = _listing_state(engine, tenant)
    items = []
    for r in rows:
        d = _title_dict(r, plats, listings)
        d["gerekce"] = r.sec_gerekce
        d["puan"] = r.firsat_puani if tur != "sesli" else r.sesli_firsat_puani
        items.append(d)
    return {"items": items, "total": total, "page": page, "pageSize": PAGE, "tur": tur}


def opportunities_csv(data: dict[str, Any]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Stok kodu", "Kitap", "Yazar", "Hedef kitle", "Son 12 ay basılı adet", "Puan", "E-kitap hakkı", "Sesli hakkı",
                "Stüdyo e-kitap", "Gerekçe"])
    for i in data["items"]:
        w.writerow([i["stokKodu"] or "", i["ad"] or "", i["yazar"] or "", i["hedefKitle"] or "", _tr(i["basili12Adet"], 0),
                    _tr(i["puan"], 1), i["hakEkitapAdi"] or "", i["hakSesliAdi"] or "", i["studioDurumuAdi"] or "", i["gerekce"] or ""])
    return "\ufeff" + buf.getvalue()


def _tr(v: Any, d: int) -> str:
    if v is None:
        return ""
    return f"{float(v):,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def rights_risks(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    """Dijitalde görünüp hakkı eksik/yok/incele olan kitaplar + dijitalde olmayan «incele» kayıtları (karar bekleyen).
    Model ön okuması «kısıtlıyor» diyenler önce."""
    reads = note_reads(engine, tenant)
    t = TITLES.c
    with engine.connect() as c:
        rows = c.execute(sa.select(TITLES).where(t.tenant_id == tenant, sa.or_(t.hak_ekitap.in_(RISK_RIGHTS), t.hak_sesli.in_(RISK_RIGHTS)))).all()
    plats = _platforms(engine, tenant)
    listings = _listing_state(engine, tenant)
    maps = RM.for_keys(engine, tenant, [ct["id"] for r in rows for ct in _j(r.sozlesme_json, []) if ct.get("not")])
    risks, reviews = [], []
    for r in rows:
        on_e = bool(r.ekitap_var) or float(r.logo_dijital_12ay_adet or 0) > 0
        on_s = bool(r.sesli_var)
        risk_e = r.hak_ekitap in RISK_RIGHTS and on_e
        risk_s = r.hak_sesli in RISK_RIGHTS and on_s
        contracts = _j(r.sozlesme_json, [])
        flags = []
        for ct in contracts:
            rd = reads.get(ct["id"]) if ct.get("not") else None
            if rd and rd["ozet"] == note_hash(ct["not"]) and rd["sonuc"]:
                flags.append(rd["sonuc"])
        d = _title_dict(r, plats, listings)
        d.update(hakEkitapGerekce=r.hak_ekitap_gerekce, hakSesliGerekce=r.hak_sesli_gerekce,
                 notOkuma="kisitliyor" if "kisitliyor" in flags else ("belirsiz" if "belirsiz" in flags else ("kisitlamiyor" if flags else None)),
                 notluSozlesmeler=[{"id": ct["id"], "ad": ct.get("ad"), "taraflar": ct.get("taraflar"), "not": ct["not"],
                                    "hakHaritasi": RM.digital_view(RM.matching(maps, ct["id"], ct["not"]))}
                                   for ct in contracts if ct.get("not") and ct.get("yururlukte")],
                 riskBicim=[b for b, x in (("ekitap", risk_e), ("sesli", risk_s)) if x])
        if risk_e or risk_s:
            risks.append(d)
        elif "incele" in (r.hak_ekitap, r.hak_sesli):
            reviews.append(d)
    order = {"kisitliyor": 0, "belirsiz": 1, None: 2, "kisitlamiyor": 3}
    key = lambda d: (order.get(d["notOkuma"], 2), -(d["basili12Adet"] or 0))  # noqa: E731
    return {"risk": sorted(risks, key=key), "incele": sorted(reviews, key=key), "kararlar": DECISIONS_KINDS,
            "notAdlari": NOTE_CHOICES}


def _pending_dict(p: Any, titles: Optional[dict[str, Any]]) -> dict[str, Any]:
    t = (titles or {}).get(p.kitap_id)
    return {"id": p.id, "kitapId": p.kitap_id, "ad": getattr(t, "ad", None), "stokKodu": getattr(t, "stok_kodu", None),
            "alan": p.alan, "alanAdi": PENDING_FIELDS.get(p.alan, p.alan), "deger": p.deger, "kaynak": p.kaynak,
            "durum": p.durum, "acildi": _iso(p.created_at), "kapandi": _iso(p.closed_at)}


def crm_pending(engine: sa.engine.Engine, tenant: str, durum: str = "acik") -> dict[str, Any]:
    with engine.connect() as c:
        cond = [PENDING.c.tenant_id == tenant]
        if durum in ("acik", "kapandi"):
            cond.append(PENDING.c.durum == durum)
        rows = c.execute(sa.select(PENDING).where(*cond).order_by(sa.desc(PENDING.c.id))).all()
        ids = sorted({r.kitap_id for r in rows})
        titles = {}
        for part in _chunks(ids, 500):
            titles.update({t.kitap_id: t for t in c.execute(sa.select(TITLES).where(TITLES.c.tenant_id == tenant, TITLES.c.kitap_id.in_(part))).all()})
    return {"items": [_pending_dict(p, titles) for p in rows], "alanlar": PENDING_FIELDS}


# ------------------------------------------------------------------ rapor yükleme: dosya → satırlar


ROLES: dict[str, tuple[str, ...]] = {
    "kimlik": ("e-isbn", "eisbn", "e isbn", "isbn", "isbn13", "isbn 13", "ean", "ean13", "barkod", "barcode", "asin", "sku",
               "stok kodu", "stokkodu", "urun kodu", "product id", "identifier", "gtin"),
    "baslik": ("baslik", "title", "book title", "kitap adi", "kitap", "urun adi", "urun", "product", "product name", "eser", "eser adi"),
    "yazar": ("yazar", "author", "authors", "contributor", "yazarlar"),
    "adet": ("adet", "miktar", "quantity", "qty", "units", "net units", "units sold", "satis adedi", "satilan adet", "sales units"),
    "brut": ("brut", "brut tutar", "gross", "gross sales", "gross revenue", "list price total", "liste fiyati toplami"),
    "net": ("net", "net tutar", "net revenue", "net sales", "publisher revenue", "earnings", "payable", "royalty", "hakedis",
            "yayinci payi", "tutar", "amount", "revenue"),
    "para": ("para birimi", "currency", "doviz", "doviz cinsi", "curr"),
    "tarih": ("tarih", "date", "sale date", "donem", "period", "month", "ay", "satis tarihi"),
}
PERSONAL = ("e-posta", "eposta", "email", "e-mail", "mail", "telefon", "phone", "gsm", "adres", "address", "musteri", "customer",
            "alici", "buyer", "ad soyad", "adi soyadi", "name surname", "full name", "tc kimlik", "tckn", "kimlik no", "okur", "reader")
TOTAL_WORDS = ("toplam", "genel toplam", "total", "grand total", "sum")


def _header_role(h: Any) -> Optional[str]:
    f = fold(h).replace("_", " ").strip(" :.")
    if not f:
        return None
    if any(f == p or f.startswith(p + " ") or f.endswith(" " + p) for p in PERSONAL):
        return "kisisel"
    for role, aliases in ROLES.items():
        if f in aliases:
            return role
    for role, aliases in ROLES.items():
        if any(re.search(rf"(^|\s){re.escape(a)}($|\s)", f) for a in aliases if len(a) > 2):
            return role
    return None


def detect_columns(rows: list[list[Any]]) -> dict[str, Any]:
    """Başlık satırı (ilk 30 satırda en çok rol tanınan) ve rol → kolon eşlemesi; kişisel veri kolonları ayrıca."""
    best, best_n = None, 1
    for i, r in enumerate(rows[:30]):
        roles = [_header_role(c) for c in r]
        n = len({x for x in roles if x and x != "kisisel"})
        if n > best_n:
            best, best_n = i, n
    if best is None:
        raise DigitalError("Dosyada başlık satırı bulunamadı: en az iki kolon (kitap kimliği ya da adı, adet ya da tutar) adıyla tanınmalı.")
    header = ["" if c is None else str(c).strip() for c in rows[best]]
    mapping: dict[str, Optional[int]] = {k: None for k in ROLES}
    personal = []
    for j, h in enumerate(header):
        role = _header_role(h)
        if role == "kisisel":
            personal.append(j)
        elif role and mapping[role] is None:
            mapping[role] = j
    return {"header_row": best, "header": header, "mapping": mapping, "personal": personal}


def parse_number(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and v != v) else float(v)
    s = re.sub(r"[^\d,.\-()]", "", str(v).strip())
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if not s or not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s) is None else s.replace(",", "")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if neg else n


def _month_of(v: Any, default: str) -> str:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m")
    if isinstance(v, date):
        return v.strftime("%Y-%m")
    s = str(v or "").strip()
    m = re.match(r"^(\d{4})[-./](\d{1,2})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{1,2})[-./](\d{1,2})[-./](\d{4})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{int(m.group(3)):04d}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{1,2})[-./](\d{4})$", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return f"{int(m.group(2)):04d}-{int(m.group(1)):02d}"
    return default


def extract_rows(filename: str, data: bytes) -> list[list[Any]]:
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext in ("xlsx", "xlsm") or data[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise DigitalError(f"Excel dosyası okunamadı: {str(e)[:120]}") from None
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if ext in ("csv", "txt", "tsv", ""):
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=";,\t|")
            return [r for r in csv.reader(io.StringIO(text), dialect)]
        except csv.Error:
            return [r for r in csv.reader(io.StringIO(text), delimiter=";")]
    raise DigitalError("Rapor Excel (.xlsx) ya da CSV olmalı.")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _cell(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return v
    return v if isinstance(v, (int, float)) or v is None else str(v)


def table_rows(rows: list[list[Any]], det: dict[str, Any]) -> tuple[list[dict[str, Any]], int]:
    """Başlıktan sonraki satırlar: boş satır sayılır ama yazılmaz; kişisel kolonlar hücreden çıkarılır."""
    drop = set(det["personal"])
    out, empty = [], 0
    for i, r in enumerate(rows[det["header_row"] + 1:], start=det["header_row"] + 2):
        cells = [_cell(c) for j, c in enumerate(r) if j not in drop]
        if all(c is None or str(c).strip() == "" for c in cells):
            empty += 1
            continue
        out.append({"sira": i, "hucreler": [_cell(c) if j not in drop else None for j, c in enumerate(r)]})
    return out, empty


def _index(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    """Kurallı eşleme dizinleri ve aday arama için kitap listesi."""
    with engine.connect() as c:
        rows = c.execute(sa.select(TITLES.c.kitap_id, TITLES.c.ad, TITLES.c.yazar, TITLES.c.stok_kodu, TITLES.c.isbn, TITLES.c.ean,
                                   TITLES.c.e_isbn, TITLES.c.ekitap_barkod, TITLES.c.ekitap_stok_kodu)
                         .where(TITLES.c.tenant_id == tenant)).all()
    keys: dict[str, dict[str, set[str]]] = {}
    for r in rows:
        for label, v, kind in (("e-ISBN", r.e_isbn, "isbn"), ("e-kitap barkodu", r.ekitap_barkod, "isbn"), ("ISBN", r.isbn, "isbn"),
                               ("barkod", r.ean, "isbn"), ("e-kitap stok kodu", r.ekitap_stok_kodu, "code"), ("stok kodu", r.stok_kodu, "code")):
            k = src.isbn_key(v) if kind == "isbn" else src.code_key(v)
            if k:
                keys.setdefault(label, {}).setdefault(k, set()).add(r.kitap_id)
    toks: dict[str, set[str]] = {}
    books = {}
    for r in rows:
        books[r.kitap_id] = r
        for t in set(tokens(r.ad)):
            toks.setdefault(t, set()).add(r.kitap_id)
    return {"keys": keys, "tokens": toks, "books": books}


KEY_ORDER = ("e-ISBN", "e-kitap barkodu", "ISBN", "barkod", "e-kitap stok kodu", "stok kodu")


def rule_match(ident: Any, idx: dict[str, Any]) -> tuple[Optional[str], Optional[str]]:
    """(kitap, anahtar). Anahtar birden çok kitaba gidiyorsa eşleşme yok (belirsiz; insan seçer)."""
    if ident in (None, ""):
        return None, None
    ik, ck = src.isbn_key(ident), src.code_key(ident)
    for label in KEY_ORDER:
        k = ik if label in ("e-ISBN", "e-kitap barkodu", "ISBN", "barkod") else ck
        hit = idx["keys"].get(label, {}).get(k) if k else None
        if hit and len(hit) == 1:
            return next(iter(hit)), label
    return None, None


def candidates(title: Any, author: Any, idx: dict[str, Any], n: int) -> list[tuple[str, float]]:
    """Ad (ve yazar) kelime örtüşmesiyle en yakın `n` kitap: (kitap, benzerlik 0–1)."""
    qt = set(tokens(title))
    if not qt:
        return []
    pool: set[str] = set()
    for t in qt:
        pool |= idx["tokens"].get(t, set())
    at = set(tokens(author))
    scored = []
    for kid in pool:
        b = idx["books"][kid]
        bt = set(tokens(b.ad))
        s = len(qt & bt) / max(1, len(qt | bt))
        if at and b.yazar:
            s = 0.8 * s + 0.2 * (len(at & set(tokens(b.yazar))) / max(1, len(at)))
        scored.append((kid, round(s, 3)))
    scored.sort(key=lambda x: (-x[1], x[0]))
    return scored[:n]


def _next_import_id(c: Any) -> str:
    prefix = f"DR-{today().year}-"
    ids = [r[0] for r in c.execute(sa.select(IMPORTS.c.id).where(IMPORTS.c.id.like(prefix + "%"))).all()]
    n = max([int(i[len(prefix):]) for i in ids if i[len(prefix):].isdigit()] or [0]) + 1
    return f"{prefix}{n:04d}"


def _sales_rows(table: list[dict[str, Any]], mapping: dict[str, Optional[int]], donem: str, currency: str,
                idx: dict[str, Any]) -> list[dict[str, Any]]:
    def cell(r: dict[str, Any], role: str) -> Any:
        j = mapping.get(role)
        cells = r["hucreler"]
        return cells[j] if j is not None and j < len(cells) else None

    out = []
    for r in table:
        ident, title, author = cell(r, "kimlik"), cell(r, "baslik"), cell(r, "yazar")
        first = next((c for c in r["hucreler"] if c not in (None, "")), "")
        is_total = not ident and fold(first).strip(" :") in TOTAL_WORDS or (not ident and not title and fold(first).startswith(TOTAL_WORDS))
        kid, key = (None, None) if is_total else rule_match(ident, idx)
        out.append({
            "sira": r["sira"], "donem_ay": _month_of(cell(r, "tarih"), donem) if mapping.get("tarih") is not None else donem,
            "tur": "ozet" if is_total else "satir", "kimlik_ham": _text(ident, 120), "baslik_ham": _text(title, 500),
            "yazar_ham": _text(author, 300), "kitap_id": kid, "eslesme": "kural" if kid else None, "eslesme_anahtar": key,
            "adet": parse_number(cell(r, "adet")), "brut": parse_number(cell(r, "brut")), "net": parse_number(cell(r, "net")),
            "para_birimi": (re.sub(r"[^A-Z]", "", str(cell(r, "para") or "").upper())[:3] or currency).replace("TL", "TRY"),
        })
    return out


def create_import(engine: sa.engine.Engine, tenant: str, user: str, platform_id: int, donem: str, filename: str,
                  data: bytes) -> dict[str, Any]:
    plats = _platforms(engine, tenant)
    if int(platform_id) not in plats:
        raise DigitalError("Platform seçilmeli.", 404)
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", donem or ""):
        raise DigitalError("Dönem YYYY-AA biçiminde olmalı.")
    if not data:
        raise DigitalError("Dosya boş.")
    rows = extract_rows(filename or "", data)
    det = detect_columns(rows)
    table, empty = table_rows(rows, det)
    idx = _index(engine, tenant)
    if not idx["books"]:
        raise DigitalError("Dijital katalog henüz okunmadı; önce katalog okunmalı.", 409)
    sales = _sales_rows(table, det["mapping"], donem, plats[int(platform_id)].para_birimi, idx)
    dropped = [det["header"][j] for j in det["personal"]]
    safe_header = [h if j not in det["personal"] else "" for j, h in enumerate(det["header"])]
    now = _now()
    with engine.begin() as c:
        iid = _next_import_id(c)
        c.execute(IMPORTS.insert().values(
            id=iid, tenant_id=tenant, platform_id=int(platform_id), donem=donem, dosya_adi=_text(filename, 300), yukleyen=user,
            satir=len(sales), eslesen=0, eslesmeyen=0, durum="onizleme", kolonlar_json=_dump(det["mapping"]),
            baslik_json=_dump(safe_header), atilan_json=_dump(dropped), ham_json=_dump(table), bos_satir=empty, created_at=now))
        _write_sales(c, iid, int(platform_id), sales)
        _count(c, iid)
    return get_import(engine, tenant, iid)


def _write_sales(c: Any, iid: str, pid: int, sales: list[dict[str, Any]]) -> None:
    c.execute(SALES.delete().where(SALES.c.import_id == iid))
    if sales:
        c.execute(SALES.insert(), [{"import_id": iid, "platform_id": pid, **s} for s in sales])


def _count(c: Any, iid: str) -> None:
    rows = c.execute(sa.select(SALES.c.tur, SALES.c.kitap_id).where(SALES.c.import_id == iid)).all()
    data = [r for r in rows if r.tur == "satir"]
    c.execute(IMPORTS.update().where(IMPORTS.c.id == iid).values(
        satir=len(rows), eslesen=sum(1 for r in data if r.kitap_id), eslesmeyen=sum(1 for r in data if not r.kitap_id)))


def _import_row(engine: sa.engine.Engine, tenant: str, iid: str) -> Any:
    with engine.connect() as c:
        r = c.execute(sa.select(IMPORTS).where(IMPORTS.c.tenant_id == tenant, IMPORTS.c.id == iid)).first()
    if r is None:
        raise DigitalError("Rapor yüklemesi bulunamadı.", 404)
    return r


def _import_dict(r: Any, plats: dict[int, Any]) -> dict[str, Any]:
    return {"id": r.id, "platformId": r.platform_id, "platform": plats[r.platform_id].ad if r.platform_id in plats else None,
            "donem": r.donem, "dosya": r.dosya_adi, "yukleyen": r.yukleyen, "satir": r.satir, "eslesen": r.eslesen,
            "eslesmeyen": r.eslesmeyen, "durum": r.durum, "olusturma": _iso(r.created_at), "onaylayan": r.onaylayan,
            "onay": _iso(r.committed_at), "kurlar": _j(r.kur_json, {})}


def _sale_dict(s: Any, books: dict[str, Any]) -> dict[str, Any]:
    b = books.get(s.kitap_id) if s.kitap_id else None
    return {"id": s.id, "sira": s.sira, "donem": s.donem_ay, "tur": s.tur, "kimlik": s.kimlik_ham, "baslik": s.baslik_ham,
            "yazar": s.yazar_ham, "kitapId": s.kitap_id, "kitapAd": getattr(b, "ad", None), "stokKodu": getattr(b, "stok_kodu", None),
            "eslesme": s.eslesme, "anahtar": s.eslesme_anahtar, "olasilik": s.olasilik, "adaylar": _j(s.aday_json, []),
            "adet": s.adet, "brut": s.brut, "net": s.net, "paraBirimi": s.para_birimi, "kur": s.kur, "netTl": s.net_tl}


def get_import(engine: sa.engine.Engine, tenant: str, iid: str) -> dict[str, Any]:
    r = _import_row(engine, tenant, iid)
    plats = _platforms(engine, tenant)
    with engine.connect() as c:
        sales = c.execute(sa.select(SALES).where(SALES.c.import_id == iid).order_by(SALES.c.sira)).all()
        ids = sorted({s.kitap_id for s in sales if s.kitap_id})
        books = {}
        for part in _chunks(ids, 500):
            books.update({b.kitap_id: b for b in c.execute(sa.select(TITLES.c.kitap_id, TITLES.c.ad, TITLES.c.stok_kodu)
                                                           .where(TITLES.c.tenant_id == tenant, TITLES.c.kitap_id.in_(part))).all()})
    data = [s for s in sales if s.tur == "satir"]
    totals: dict[str, dict[str, float]] = {}
    for s in data:
        t = totals.setdefault(s.para_birimi or "?", {"satir": 0, "adet": 0.0, "net": 0.0, "brut": 0.0})
        t["satir"] += 1
        t["adet"] += s.adet or 0
        t["net"] += s.net or 0
        t["brut"] += s.brut or 0
    out = _import_dict(r, plats)
    out.update({
        "kolonlar": _j(r.kolonlar_json, {}), "baslik": _j(r.baslik_json, []), "atilanKolonlar": _j(r.atilan_json, []),
        "bosSatir": r.bos_satir or 0, "ozetSatir": sum(1 for s in sales if s.tur == "ozet"), "toplamlar": totals,
        "paraBirimleri": sorted({s.para_birimi for s in data if s.para_birimi}),
        "eslestirme": _j(r.eslestirme_json, None), "roller": list(ROLES),
        "satirlar": [_sale_dict(s, books) for s in sales],
    })
    return out


def remap_import(engine: sa.engine.Engine, tenant: str, iid: str, mapping: dict[str, Any]) -> dict[str, Any]:
    """Kolon eşlemesini kullanıcı düzeltir: satırlar saklı tablodan yeniden kurulur (yalnız önizlemede)."""
    r = _import_row(engine, tenant, iid)
    if r.durum != "onizleme":
        raise DigitalError("Onaylanmış raporun kolonları değiştirilemez.", 409)
    header = _j(r.baslik_json, [])
    cur = _j(r.kolonlar_json, {})
    for role in ROLES:
        if role in mapping:
            v = mapping[role]
            if v in (None, "", -1):
                cur[role] = None
            else:
                try:
                    j = int(v)
                except (TypeError, ValueError):
                    raise DigitalError("Kolon numarası geçersiz.") from None
                if not 0 <= j < len(header) or not header[j]:
                    raise DigitalError("Kolon dosyada yok ya da kişisel veri olduğu için atıldı.")
                cur[role] = j
    plats = _platforms(engine, tenant)
    sales = _sales_rows(_j(r.ham_json, []), cur, r.donem, plats[r.platform_id].para_birimi if r.platform_id in plats else "TRY",
                        _index(engine, tenant))
    with engine.begin() as c:
        c.execute(IMPORTS.update().where(IMPORTS.c.id == iid).values(kolonlar_json=_dump(cur), eslestirme_json=None))
        _write_sales(c, iid, r.platform_id, sales)
        _count(c, iid)
    return get_import(engine, tenant, iid)


MATCH_PROMPT = ("Bir e-kitap/sesli kitap platformunun satış raporundaki satırı yayınevinin kataloğundaki kitaplardan "
                "biriyle eşleştiriyorsun. Rapordaki satır:\nBaşlık: {title}\nYazar: {author}\nKimlik: {ident}\n\n"
                "Aynı eser (aynı kitabın e-kitap ya da sesli sürümü dahil) hangisi? Hiçbiri aynı eser değilse «Hiçbiri» seç.")
NONE_CHOICE = "Hiçbiri"


def suggest_matches(engine: sa.engine.Engine, tenant: str, iid: str, llm: Any, st: dict[str, Any],
                    progress: Optional[Callable[[int, int], None]] = None) -> dict[str, Any]:
    """Kurallı eşleşmeyen satırlara aday kitap: ad/yazar benzerliğiyle adaylar, Zeki AI kapalı küme seçimiyle olasılık.
    Kayıt yalnız öneridir (`aday_json`, `olasilik`); eşleşmeyi insan onaylar."""
    idx = _index(engine, tenant)
    with engine.connect() as c:
        todo = c.execute(sa.select(SALES).where(SALES.c.import_id == iid, SALES.c.tur == "satir", SALES.c.kitap_id.is_(None))
                         .order_by(SALES.c.sira)).all()
    done = asked = 0
    for s in todo:
        cands = candidates(s.baslik_ham, s.yazar_ham, idx, st["matchCandidates"])
        adaylar = [{"kitapId": k, "ad": idx["books"][k].ad, "yazar": idx["books"][k].yazar, "stokKodu": idx["books"][k].stok_kodu,
                    "benzerlik": sim, "olasilik": None} for k, sim in cands]
        best, prob = None, None
        if adaylar and llm is not None and hasattr(llm, "choose"):
            labels, seen = [], set()
            for a in adaylar:
                lab = f"{a['ad'] or ''} — {a['yazar'] or 'yazar yok'} ({a['stokKodu'] or a['kitapId'][:8]})"
                while lab in seen:
                    lab += " "
                seen.add(lab)
                labels.append(lab)
            try:
                ch = llm.choose(MATCH_PROMPT.format(title=s.baslik_ham or "", author=s.yazar_ham or "", ident=s.kimlik_ham or ""),
                                labels + [NONE_CHOICE])
                asked += 1
                probs = ch.probs or {}
                for a, lab in zip(adaylar, labels):
                    a["olasilik"] = probs.get(lab)
                if ch.choice and ch.choice != NONE_CHOICE:
                    best = adaylar[labels.index(ch.choice)]["kitapId"]
                    prob = ch.probability
            except Exception as e:  # noqa: BLE001 — model yoksa adaylar benzerlikle kalır
                log.info("dijital: eşleme önerisi alınamadı: %s", e)
                llm = None
        adaylar.sort(key=lambda a: (-(a["olasilik"] or 0), -a["benzerlik"]))
        with engine.begin() as c:
            c.execute(SALES.update().where(SALES.c.id == s.id).values(
                aday_json=_dump([{**a, "oneri": a["kitapId"] == best} for a in adaylar]), olasilik=prob))
        done += 1
        if progress:
            progress(done, len(todo))
    return {"satir": len(todo), "modelSorulan": asked, "model": llm is not None}


def decide_rows(engine: sa.engine.Engine, tenant: str, iid: str, rows: list[dict[str, Any]], st: dict[str, Any]) -> dict[str, Any]:
    """Satır eşleşmesi: `{id, kitapId|null, kaynak: zeki|elle}` listesi. `{"zekiGuclu": true}` ile olasılığı eşiğin
    üstündeki bütün önerilerin toplu onayı ayrı uçtan gelir (`accept_strong`). Onaylı raporda da açık satır eşlenebilir."""
    r = _import_row(engine, tenant, iid)
    if r.durum == "iptal":
        raise DigitalError("İptal edilmiş rapor değiştirilemez.", 409)
    idx = _index(engine, tenant)
    with engine.begin() as c:
        for x in rows:
            sid = int(x.get("id") or 0)
            kid = str(x.get("kitapId") or "").upper() or None
            kaynak = "zeki" if x.get("kaynak") == "zeki" else "elle"
            if kid and kid not in idx["books"]:
                raise DigitalError("Seçilen kitap katalogda yok.", 404)
            res = c.execute(SALES.update().where(SALES.c.id == sid, SALES.c.import_id == iid, SALES.c.tur == "satir").values(
                kitap_id=kid, eslesme=kaynak if kid else None, eslesme_anahtar=None))
            if res.rowcount == 0:
                raise DigitalError("Satır bu rapora ait değil.", 404)
        _count(c, iid)
    return get_import(engine, tenant, iid)


def accept_strong(engine: sa.engine.Engine, tenant: str, iid: str, st: dict[str, Any]) -> dict[str, Any]:
    """Olasılığı `DIJITAL_MATCH_MIN_PROB`'un üstündeki Zeki AI önerilerini onaylar (kullanıcı düğmeye basar)."""
    _import_row(engine, tenant, iid)
    with engine.connect() as c:
        rows = c.execute(sa.select(SALES).where(SALES.c.import_id == iid, SALES.c.tur == "satir", SALES.c.kitap_id.is_(None),
                                                SALES.c.olasilik >= st["matchMinProb"])).all()
    picks = []
    for s in rows:
        best = next((a for a in _j(s.aday_json, []) if a.get("oneri")), None)
        if best:
            picks.append({"id": s.id, "kitapId": best["kitapId"], "kaynak": "zeki"})
    return decide_rows(engine, tenant, iid, picks, st) if picks else get_import(engine, tenant, iid)


def commit_import(engine: sa.engine.Engine, tenant: str, user: str, iid: str, body: dict[str, Any]) -> dict[str, Any]:
    """Önizlemeyi onaylar. TRY dışındaki her para birimi için kur ister (modül kur varsaymaz); aynı platformun aynı dönemi
    onaylıysa önce o iptal edilmeli (çift sayım olmasın). Eşleşmeyen satırlar atılmaz, açık iş olarak kalır."""
    r = _import_row(engine, tenant, iid)
    if r.durum != "onizleme":
        raise DigitalError("Yalnız önizlemedeki rapor onaylanır.", 409)
    rates_in = body.get("kurlar") or {}
    with engine.connect() as c:
        dup = c.execute(sa.select(IMPORTS.c.id).where(IMPORTS.c.tenant_id == tenant, IMPORTS.c.platform_id == r.platform_id,
                                                      IMPORTS.c.donem == r.donem, IMPORTS.c.durum == "onaylandi")).first()
        currencies = {s[0] for s in c.execute(sa.select(SALES.c.para_birimi).where(SALES.c.import_id == iid, SALES.c.tur == "satir")).all()}
    if dup:
        raise DigitalError(f"Bu platformun {r.donem} raporu zaten onaylı ({dup.id}); çift sayılmasın diye önce o iptal edilmeli.", 409)
    rates: dict[str, float] = {"TRY": 1.0}
    for cur in sorted(x for x in currencies if x and x != "TRY"):
        v = rates_in.get(cur)
        n = parse_number(v) if not isinstance(v, (int, float)) else float(v)
        if n is None or n <= 0:
            raise DigitalError(f"{cur} için TL kuru girilmeli (modül kur varsaymaz).")
        rates[cur] = n
    if None in currencies or "" in currencies:
        raise DigitalError("Bazı satırlarda para birimi yok; platformun para birimi tanımlanmalı.")
    with engine.begin() as c:
        for cur, k in rates.items():
            c.execute(SALES.update().where(SALES.c.import_id == iid, SALES.c.para_birimi == cur)
                      .values(kur=k, net_tl=SALES.c.net * k))
        c.execute(IMPORTS.update().where(IMPORTS.c.id == iid).values(durum="onaylandi", onaylayan=user, committed_at=_now(),
                                                                     kur_json=_dump(rates), ham_json=None))
    return get_import(engine, tenant, iid)


def delete_import(engine: sa.engine.Engine, tenant: str, iid: str) -> dict[str, Any]:
    """Önizleme silinir; onaylı rapor «iptal» olur (satırları toplamlardan çıkar, kayıt kalır)."""
    r = _import_row(engine, tenant, iid)
    with engine.begin() as c:
        if r.durum == "onizleme":
            c.execute(SALES.delete().where(SALES.c.import_id == iid))
            c.execute(IMPORTS.delete().where(IMPORTS.c.id == iid))
            return {"id": iid, "silindi": True}
        c.execute(IMPORTS.update().where(IMPORTS.c.id == iid).values(durum="iptal"))
    return {"id": iid, "iptal": True}


def list_imports(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    plats = _platforms(engine, tenant)
    with engine.connect() as c:
        rows = c.execute(sa.select(IMPORTS).where(IMPORTS.c.tenant_id == tenant).order_by(sa.desc(IMPORTS.c.created_at))).all()
    return {"items": [_import_dict(r, plats) for r in rows]}


def set_match_job(engine: sa.engine.Engine, iid: str, state: dict[str, Any]) -> None:
    with engine.begin() as c:
        c.execute(IMPORTS.update().where(IMPORTS.c.id == iid).values(eslestirme_json=_dump(state)))


# ------------------------------------------------------------------ satış panosu


def _committed(tenant: str):
    return sa.select(IMPORTS.c.id).where(IMPORTS.c.tenant_id == tenant, IMPORTS.c.durum == "onaylandi")


def sales(engine: sa.engine.Engine, tenant: str, *, donem: str = "", platform: int = 0, st: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    """Onaylı raporlardan aylık dijital gelir (platform kırılımı), kitap kırılımı, eşleşmeyen açık satırlar; Logo'daki
    e-kitap stok kodu faturaları ayrı sütun (iki kaynak toplanmaz: aynı satış iki yerde olabilir)."""
    cond = [SALES.c.import_id.in_(_committed(tenant)), SALES.c.tur == "satir"]
    if re.fullmatch(r"\d{4}(-\d{2})?", donem or ""):
        cond.append(SALES.c.donem_ay.like(donem + "%"))
    if platform:
        cond.append(SALES.c.platform_id == int(platform))
    plats = _platforms(engine, tenant)
    with engine.connect() as c:
        monthly = c.execute(sa.select(SALES.c.donem_ay, SALES.c.platform_id, sa.func.sum(SALES.c.adet), sa.func.sum(SALES.c.net_tl),
                                      sa.func.count()).where(*cond).group_by(SALES.c.donem_ay, SALES.c.platform_id)
                            .order_by(SALES.c.donem_ay)).all()
        by_book = c.execute(sa.select(SALES.c.kitap_id, sa.func.sum(SALES.c.adet), sa.func.sum(SALES.c.net_tl))
                            .where(*cond, SALES.c.kitap_id.is_not(None)).group_by(SALES.c.kitap_id)).all()
        open_rows = c.execute(sa.select(SALES).where(*cond, SALES.c.kitap_id.is_(None)).order_by(SALES.c.import_id, SALES.c.sira)).all()
        ids = sorted({r[0] for r in by_book})
        titles = {}
        for part in _chunks(ids, 500):
            titles.update({t.kitap_id: t for t in c.execute(sa.select(TITLES).where(TITLES.c.tenant_id == tenant, TITLES.c.kitap_id.in_(part))).all()})
    books = []
    for kid, adet, net in by_book:
        t = titles.get(kid)
        books.append({"kitapId": kid, "ad": getattr(t, "ad", None), "stokKodu": getattr(t, "stok_kodu", None),
                      "hedefKitle": getattr(t, "hedef_kitle", None), "adet": adet or 0, "netTl": net,
                      "basili12Adet": getattr(t, "basili_12ay_adet", None)})
    books.sort(key=lambda b: -(b["netTl"] or 0))
    logo = meta_get(engine, tenant, "logo")
    logo_months: dict[str, dict[str, float]] = {}
    for code, v in (logo.get("ekitapSatis") or {}).items():
        for m, x in v["aylar"].items():
            if donem and not m.startswith(donem):
                continue
            a = logo_months.setdefault(m, {"adet": 0.0, "ciro": 0.0})
            a["adet"] += x["adet"]
            a["ciro"] += x["ciro"]
    ratio: dict[str, dict[str, float]] = {}
    for b in books:
        k = b["hedefKitle"] or "Belirtilmemiş"
        x = ratio.setdefault(k, {"dijitalAdet": 0.0, "basiliAdet": 0.0})
        x["dijitalAdet"] += b["adet"] or 0
        x["basiliAdet"] += b["basili12Adet"] or 0
    return {
        "aylik": [{"donem": m, "platformId": p, "platform": plats[p].ad if p in plats else str(p), "adet": a or 0, "netTl": n,
                   "satir": k} for m, p, a, n, k in monthly],
        "kitaplar": books,
        "eslesmeyen": [{"id": s.id, "importId": s.import_id, "sira": s.sira, "donem": s.donem_ay,
                        "platform": plats[s.platform_id].ad if s.platform_id in plats else None, "kimlik": s.kimlik_ham,
                        "baslik": s.baslik_ham, "yazar": s.yazar_ham, "adet": s.adet, "netTl": s.net_tl,
                        "adaylar": _j(s.aday_json, [])} for s in open_rows],
        "logoEkitap": [{"donem": m, **v} for m, v in sorted(logo_months.items())],
        "logoPencere": logo.get("pencere"), "logoVeriSonu": logo.get("veriSonu"),
        "dijitalBasiliOran": [{"hedefKitle": k, **v, "oran": (v["dijitalAdet"] / v["basiliAdet"]) if v["basiliAdet"] else None}
                              for k, v in sorted(ratio.items())],
    }


def sales_csv(data: dict[str, Any]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Dönem", "Platform", "Satır", "Adet", "Net (TL)"])
    for r in data["aylik"]:
        w.writerow([r["donem"], r["platform"], r["satir"], _tr(r["adet"], 0), _tr(r["netTl"], 2)])
    w.writerow([])
    w.writerow(["Stok kodu", "Kitap", "Dijital adet", "Net (TL)", "Son 12 ay basılı adet"])
    for b in data["kitaplar"]:
        w.writerow([b["stokKodu"] or "", b["ad"] or "", _tr(b["adet"], 0), _tr(b["netTl"], 2), _tr(b["basili12Adet"], 0)])
    return "\ufeff" + buf.getvalue()


def title_sales(engine: sa.engine.Engine, tenant: str, kitap_id: str) -> dict[str, Any]:
    plats = _platforms(engine, tenant)
    with engine.connect() as c:
        rows = c.execute(sa.select(SALES.c.donem_ay, SALES.c.platform_id, sa.func.sum(SALES.c.adet), sa.func.sum(SALES.c.net_tl))
                         .where(SALES.c.import_id.in_(_committed(tenant)), SALES.c.tur == "satir", SALES.c.kitap_id == kitap_id)
                         .group_by(SALES.c.donem_ay, SALES.c.platform_id).order_by(SALES.c.donem_ay)).all()
    logo = meta_get(engine, tenant, "logo")
    lm = next((v["aylar"] for v in (logo.get("ekitapSatis") or {}).values() if v.get("kitap_id") == kitap_id), {})
    return {"platform": [{"donem": m, "platform": plats[p].ad if p in plats else str(p), "adet": a or 0, "netTl": n} for m, p, a, n in rows],
            "logo": [{"donem": m, **v} for m, v in sorted(lm.items())]}


# ------------------------------------------------------------------ zamanlayıcı uyarıları


def due_alerts(engine: sa.engine.Engine, tenant: str, st: dict[str, Any], ref: Optional[date] = None) -> list[dict[str, Any]]:
    """Gönderilmemiş uyarılar: hak riski listesine yeni giren kitap (telif), haftalık yeni baskı/kapak (dijital sorumlusu),
    beklenen aylık rapor gelmedi (finans). Anahtar gönderildikten sonra `notified` meta'sına yazılır."""
    ref = ref or today()
    sent = set(meta_get(engine, tenant, "notified").get("keys") or [])
    out: list[dict[str, Any]] = []
    for r in rights_risks(engine, tenant)["risk"]:
        for b in r["riskBicim"]:
            key = f"risk:{r['kitapId']}:{b}:{r['hakEkitap'] if b == 'ekitap' else r['hakSesli']}"
            if key not in sent:
                out.append({"key": key, "kime": "telif", "metin": f"{r['ad']} ({r['stokKodu'] or '-'}): {FORMATS[b]} dijitalde görünüyor, "
                                                                  f"hak «{r['hakEkitapAdi'] if b == 'ekitap' else r['hakSesliAdi']}»."})
    if ref.isoweekday() == st["weeklyDay"]:
        with engine.connect() as c:
            rows = c.execute(sa.select(TITLES).where(TITLES.c.tenant_id == tenant, TITLES.c.baski_degisim_tarih.is_not(None))).all()
        for t in rows:
            key = f"baski:{t.kitap_id}:{t.baski_degisim_tarih}"
            if key not in sent:
                out.append({"key": key, "kime": "dijital", "metin": f"{t.ad} ({t.stok_kodu or '-'}): {t.baski_degisim_tur} "
                                                                    f"({t.baski_degisim_tarih}); dijital sürüm güncellenmeli mi bakılmalı."})
    first = ref.replace(day=1)
    prev_end = first - timedelta(days=1)
    period = prev_end.strftime("%Y-%m")
    with engine.connect() as c:
        have = {r[0] for r in c.execute(sa.select(IMPORTS.c.platform_id).where(IMPORTS.c.tenant_id == tenant, IMPORTS.c.donem == period,
                                                                               IMPORTS.c.durum == "onaylandi")).all()}
    for p in _platforms(engine, tenant).values():
        if not p.aktif or p.rapor_gunu is None or p.id in have:
            continue
        if ref >= prev_end + timedelta(days=int(p.rapor_gunu)):
            key = f"rapor:{p.id}:{period}"
            if key not in sent:
                out.append({"key": key, "kime": "finans", "metin": f"{p.ad}: {period} satış raporu yüklenmedi "
                                                                   f"(ay kapanışından {p.rapor_gunu} gün sonra bekleniyordu)."})
    return out


def mark_sent(engine: sa.engine.Engine, tenant: str, keys: Iterable[str]) -> None:
    cur = set(meta_get(engine, tenant, "notified").get("keys") or [])
    meta_set(engine, tenant, "notified", {"keys": sorted(cur | set(keys))})


def alert_text(items: list[dict[str, Any]], link: str) -> str:
    lines = ["Dijital yayın ve e-kitap: yeni uyarılar.", ""] + [f"- {i['metin']}" for i in items]
    if link:
        lines += ["", f"Ekran: {link}"]
    lines += ["", "Bu e-posta portalın iç uyarısıdır; platformlara, CRM'e ya da Logo'ya hiçbir şey gönderilmedi."]
    return "\n".join(lines)
