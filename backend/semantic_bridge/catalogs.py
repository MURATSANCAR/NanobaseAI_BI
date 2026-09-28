"""M24 Katalog yönetimi: dönemsel katalog kaydı, gerekçeli kitap önerisi, canlı fiyat/stok uyarısı, tasarımcı paketi.

**Akış:** katalog (tür + dönem + tema + süzgeç) → kural puanlı öneri listesi (her satırda gerekçe) → seçim, sıralama,
öne çıkarma, sayfa ipucu → onaya gönder → onay (açıkça verilen `ozellik:katalog-bulten.onay`; gönderen onaylayamaz) →
yayında → arşiv. Katalog kayıtları yalnız köprünün `semantic_catalog*` tablolarındadır; CRM'e, Logo'ya, T-soft'a yazılmaz.

**Öneri puanı (K2, açıklanabilir):** satış hızı (havuzdaki yüzdelik sırası), stok ay sayısı (hedef aya oranla), yenilik
(ilk baskıdan bu yana geçen ay), özel gün bağı (katalogda seçildiyse). Ağırlıklar `CATALOG_SCORE_WEIGHTS`; ölçülemeyen
parça dışarıda kalır, ağırlıklar yeniden dağıtılır. Hedef kitle, yaş, marka, anahtar sözcük sert süzgeçtir. Satıştan
kalkmış (yayıncılık statüsü iptal/çekildi/geri istendi/devredildi/bizim değil ya da satış durumu kapalı) ve stoku olmayan
kitap öneriye girmez; elenen sayısı nedeniyle ekranda yazar (sessiz eleme yok). Rakamı model üretmez; Zeki AI yalnız bu
olgulardan tek cümlelik gerekçe ve katalog tanıtım metni kısaltması yazar (`marketing.guard` denetiminden geçer).

**Uyarılar (K1):** katalogdaki kitabın bugünkü fiyatı eklendiği/onaylandığı fiyattan farklı; stok ay sayısı
`CATALOG_CRITICAL_STOCK_MONTHS` altında ya da stok yok; satıştan kalktı; CRM'de etkin kart yok (kritik). Hak notu,
kapak bağlantısı ya da tanıtım metni yok (bilgi). «Kabul et» o anki değeri yeni dayanak yapar; değer yeniden değişirse uyarı
geri gelir. Stok verisinin tarihi (Logo'daki son faturalı satış günü) her ekranda ve dışa aktarımda yazar.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import math
import os
import re
import threading
import uuid
import zipfile
from datetime import date, datetime, timezone
from typing import Any, Callable, Iterable, Optional
from zoneinfo import ZoneInfo

import sqlalchemy as sa

log = logging.getLogger("semantic.catalogs")
TZ = ZoneInfo("Europe/Istanbul")
_md = sa.MetaData()
PRODUCT = "Zeki AI"

CATALOGS = sa.Table(
    "semantic_catalogs", _md,
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False, index=True),
    sa.Column("kind", sa.String(16), nullable=False),                 # bayi | okul | fuar | yabanci-hak | e-katalog
    sa.Column("title", sa.String(300), nullable=False),
    sa.Column("season", sa.String(120)),
    sa.Column("theme", sa.String(300)),
    sa.Column("filters_json", sa.Text),                               # hedef, yaş, marka, anahtar, özel gün
    sa.Column("status", sa.String(12), nullable=False),               # taslak | onayda | onayli | yayinda | arsiv
    sa.Column("price_source", sa.String(20), nullable=False),
    sa.Column("stock_as_of", sa.String(10)),
    sa.Column("alerts_at", sa.DateTime(timezone=True)),
    sa.Column("submitted_by", sa.String(120)),
    sa.Column("submitted_at", sa.DateTime(timezone=True)),
    sa.Column("approved_by", sa.String(120)),
    sa.Column("approved_at", sa.DateTime(timezone=True)),
    sa.Column("published_at", sa.DateTime(timezone=True)),
    sa.Column("note", sa.Text),                                       # son geri gönderme gerekçesi
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
)
ITEMS = sa.Table(
    "semantic_catalog_items", _md,
    sa.Column("catalog_id", sa.String(32), primary_key=True),
    sa.Column("crm_book_id", sa.String(40), primary_key=True),
    sa.Column("stok_kodu", sa.String(80)),
    sa.Column("ad", sa.String(500)),
    sa.Column("position", sa.Integer, nullable=False),
    sa.Column("featured", sa.Boolean, nullable=False, default=False),
    sa.Column("page_hint", sa.String(120)),
    sa.Column("reason", sa.Text),                                     # kural gerekçesi (olgular)
    sa.Column("reason_ai", sa.Text),                                  # Zeki AI tek cümle (denetimden geçmiş)
    sa.Column("price_snapshot", sa.Float),
    sa.Column("stock_months_snapshot", sa.Float),
    sa.Column("text_override", sa.Text),                              # katalog tanıtım metni
    sa.Column("text_source", sa.String(12)),                          # zeki | kullanici
    sa.Column("alert_json", sa.Text),
    sa.Column("accepted_json", sa.Text),                              # uyarı türü → kabul edilen değer
    sa.Column("added_by", sa.String(120)),
    sa.Column("added_at", sa.DateTime(timezone=True)),
)
META = sa.Table(
    "semantic_catalog_meta", _md,                                     # kitap havuzu anlık görüntüsü, son zamanlayıcı koşusu
    sa.Column("tenant_id", sa.String(80), primary_key=True),
    sa.Column("key", sa.String(40), primary_key=True),
    sa.Column("value_json", sa.Text, nullable=False),
    sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
)
JOBS = sa.Table(
    "semantic_catalog_jobs", _md,                                     # Zeki AI işleri (katalog metni/gerekçe, bülten taslağı)
    sa.Column("id", sa.String(32), primary_key=True),
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("kind", sa.String(20), nullable=False),                 # katalog-zeki | bulten-taslak
    sa.Column("ref_id", sa.String(32), nullable=False, index=True),
    sa.Column("status", sa.String(12), nullable=False),               # bekliyor | calisiyor | bitti | hata
    sa.Column("step", sa.String(200)),
    sa.Column("result_json", sa.Text),
    sa.Column("error", sa.Text),
    sa.Column("created_by", sa.String(120), nullable=False),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("finished_at", sa.DateTime(timezone=True)),
)

KINDS = {"bayi": "Bayi ve kitapçı", "okul": "Okul ve kurum", "fuar": "Fuar", "yabanci-hak": "Yabancı hak (İngilizce)",
         "e-katalog": "E-katalog"}
STATUSES = {"taslak": "Taslak", "onayda": "Onay bekliyor", "onayli": "Onaylı", "yayinda": "Yayında", "arsiv": "Arşiv"}
PRICE_SOURCES = {
    "crm": "CRM kitap kartı · KDV dahil fiyat",
    "crm-perakende": "CRM kitap kartı · perakende birim fiyat",
    "crm-uzeri": "CRM · kitabın üzerindeki fiyat (Baskı önerisi raporundaki)",
    "logo": "Logo · son B2B/CRM satış faturasındaki birim fiyat",
    "tsoft": "Web sitesi (T-soft) satış fiyatı",
}
PRICE_SOURCE_WHY = (
    "Varsayılan CRM kitap kartındaki KDV dahil fiyattır: kitabın kayıtlı liste fiyatı tek alanda durur ve ihale teklif "
    "tablosu da aynı alanı kullanır. Logo'daki fiyat son satış faturasının birim fiyatıdır (bayi iskontosu içerebilir), web "
    "fiyatı kampanyalı olabilir. Katalog fiyatının hangi kaynaktan alınacağı pazarlama ekibince kesinleşene kadar bu "
    "varsayım geçerlidir; Yönetim → Katalog ve bülten'den ya da katalog başına değiştirilir.")
STOCK_SOURCES = {"crm": "CRM stok adedi (Baskı önerisi raporundaki)", "logo": "Logo depo stoku"}
#: Satıştan kalkmış sayılan CRM yayıncılık statüleri (`seo_geo.crm.STATUS_FLAGS` sınıflaması).
OUT_OF_SALE = {"iptal": "iptal edilmiş", "cekildi": "satıştan çekilmiş", "geri_istendi": "geri istenmiş",
               "devredildi": "devredilmiş", "bizim_degil": "yayınevimizin değil"}
ALERT_KINDS = {"fiyat": ("kritik", "Fiyat değişti"), "stok": ("kritik", "Stok kritik"), "satis": ("kritik", "Satıştan kalktı"),
               "crm-yok": ("kritik", "CRM'de etkin kart yok"), "hak": ("bilgi", "Hak notu var"),
               "kapak": ("bilgi", "Kapak bağlantısı yok"), "metin": ("bilgi", "Tanıtım metni yok")}
DEFAULT_WEIGHTS = {"hiz": 40.0, "stok": 20.0, "yenilik": 20.0, "ozelgun": 20.0}

_lock = threading.Lock()
_ready: set[int] = set()


class CatalogError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def ensure(engine: sa.engine.Engine) -> None:
    with _lock:
        if id(engine) in _ready:
            return
        _md.create_all(engine, checkfirst=True)
        _ready.add(id(engine))


# ------------------------------------------------------------------ ayarlar


def _conf(key: str, default: str) -> str:
    try:
        from semantic_bridge import admin as admin_mod
        v = admin_mod.conf(key)
    except Exception:  # noqa: BLE001 — testte admin ayarı yok
        v = ""
    return v if v not in (None, "") else os.environ.get(key, default)


def _conf_float(key: str, default: float) -> float:
    try:
        return float(str(_conf(key, str(default))).replace(",", "."))
    except ValueError:
        return default


def settings() -> dict[str, Any]:
    """Modülün ayarları (ekran > ortam > varsayılan). Ölçülmemiş varsayımlar buradadır, kodda sabit değil."""
    try:
        w_raw = json.loads(_conf("CATALOG_SCORE_WEIGHTS", "") or "{}")
    except ValueError:
        w_raw = {}
    weights = dict(DEFAULT_WEIGHTS)
    weights.update({k: float(v) for k, v in (w_raw or {}).items() if k in weights and isinstance(v, (int, float)) and v >= 0})
    src = (_conf("CATALOG_PRICE_SOURCE", "crm") or "crm").strip().lower()
    stock = (_conf("CATALOG_STOCK_SOURCE", "crm") or "crm").strip().lower()
    types = [int(x) for x in re.findall(r"\d+", _conf("CATALOG_BOOK_TYPES", "1,4"))] or [1, 4]
    return {
        "priceSource": src if src in PRICE_SOURCES else "crm",
        "stockSource": stock if stock in STOCK_SOURCES else "crm",
        "bookTypes": types,
        "criticalMonths": max(0.0, _conf_float("CATALOG_CRITICAL_STOCK_MONTHS", 1.0)),
        "targetMonths": max(0.5, _conf_float("CATALOG_TARGET_STOCK_MONTHS", 6.0)),
        "newMonths": max(1.0, _conf_float("CATALOG_NEW_MONTHS", 12.0)),
        "weights": weights,
        "textWords": max(10, int(_conf_float("CATALOG_TEXT_WORDS", 60))),
        "pdfPerPage": min(12, max(1, int(_conf_float("CATALOG_PDF_PER_PAGE", 6)))),
        "recipients": [x.strip() for x in (_conf("CATALOG_ALERT_RECIPIENTS", "") or "").replace(";", ",").split(",") if "@" in x],
    }


# ------------------------------------------------------------------ küçük yardımcılar


def now() -> datetime:
    return datetime.now(timezone.utc)


def today() -> date:
    return datetime.now(TZ).date()


def iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    return (v if v.tzinfo else v.replace(tzinfo=timezone.utc)).isoformat()


def dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def load(s: Optional[str], default: Any = None) -> Any:
    if not s:
        return default
    try:
        return json.loads(s)
    except ValueError:
        return default


def uid() -> str:
    return uuid.uuid4().hex[:16]


_FOLD = str.maketrans("İIıŞşĞğÜüÖöÇçÂâÎîÛû", "iiissgguuooccaaiiuu")


def fold(s: Any) -> str:
    return " ".join(str(s or "").translate(_FOLD).lower().split())


def _text(v: Any, limit: int) -> Optional[str]:
    s = str(v or "").strip()
    return s[:limit] if s else None


def tr_num(v: Optional[float], digits: int = 0) -> str:
    if v is None:
        return "—"
    s = f"{v:,.{digits}f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def tr_money(v: Optional[float]) -> str:
    return "—" if v is None else tr_num(v, 2) + " ₺"


def tr_day(v: Optional[str]) -> str:
    if not v:
        return "—"
    d = date.fromisoformat(v[:10])
    return f"{d.day:02d}.{d.month:02d}.{d.year}"


# ------------------------------------------------------------------ kitap havuzu


def build_pool(crm_part: dict[str, Any], stock: dict[str, Any], tsoft: dict[str, dict[str, Any]], days: list[dict[str, Any]],
               rights: set[str], logo_end: Optional[str], notes: list[str], cfg: dict[str, Any],
               read_at: Optional[str] = None) -> dict[str, Any]:
    """Kaynaklardan tek havuz: her kitap için fiyatlar, stok, hız, stok ayı, özel günler, T-soft bağı."""
    by_code = stock.get("byCode") or {}
    day_of: dict[str, list[str]] = {}
    for d in days:
        for bid in d.get("kitaplar") or []:
            day_of.setdefault(bid, []).append(d["key"])
    books = []
    for b in crm_part.get("books") or []:
        s = by_code.get(b.get("stok") or "", {})
        t = tsoft.get(b.get("ean") or "") or {}
        prices = dict(b.get("fiyatlar") or {})
        prices["crm-uzeri"] = s.get("uzeri")
        prices["logo"] = s.get("logoFiyat")
        prices["tsoft"] = t.get("fiyat")
        books.append({**{k: v for k, v in b.items() if k != "fiyatlar"}, "fiyatlar": prices,
                      "stokCrm": s.get("stok"), "depo": s.get("depo"), "hiz": s.get("hiz"), "yillik": s.get("yillik"),
                      "logoFiyatTarihi": s.get("logoFiyatTarihi"),
                      "kapak": b.get("kapak") or t.get("gorsel"), "kapakKaynak": "crm" if b.get("kapak") else ("web" if t.get("gorsel") else None),
                      "webUrl": t.get("url"), "ozelGunler": sorted(day_of.get(b["id"], [])), "hakNotu": b["id"] in rights})
    diff = sum(1 for b in books if len({round(v, 2) for v in (b["fiyatlar"].get(k) for k in PRICE_SOURCES) if v}) > 1)
    return {"books": books, "days": [{k: v for k, v in d.items() if k != "kitaplar"} | {"kitapSayisi": len(d.get("kitaplar") or [])}
                                     for d in days],
            "hedefler": crm_part.get("hedefler") or [], "logoSon": logo_end, "notes": notes,
            "readAt": read_at or iso(now()), "fiyatFarkli": diff}


def stock_of(b: dict[str, Any], cfg: dict[str, Any]) -> dict[str, Any]:
    """Stok ve stok ay sayısı. Ay sayısı = stok ÷ ağırlıklı aylık satış hızı (Baskı Tekrar «Tükenme süresi» ile aynı)."""
    stok = b.get("stokCrm") if cfg["stockSource"] == "crm" else b.get("depo")
    hiz = b.get("hiz")
    if stok is None:
        return {"stok": None, "stokAy": None, "satisYok": False, "bilinmiyor": True}
    if not hiz:
        return {"stok": stok, "stokAy": 0.0 if stok <= 0 else None, "satisYok": stok > 0, "bilinmiyor": False}
    return {"stok": stok, "stokAy": stok / hiz, "satisYok": False, "bilinmiyor": False}


def price_of(b: dict[str, Any], source: str) -> Optional[float]:
    v = (b.get("fiyatlar") or {}).get(source)
    return round(float(v), 2) if v is not None else None


def out_of_sale(b: dict[str, Any]) -> Optional[str]:
    if b.get("flag") in OUT_OF_SALE:
        return f"CRM yayıncılık statüsü: {b.get('yst') or OUT_OF_SALE[b['flag']]}"
    if b.get("satisAcik") is False:
        return "CRM'de satış durumu kapalı"
    return None


def _months_since(iso_day: Optional[str], ref: date) -> Optional[float]:
    if not iso_day:
        return None
    d = date.fromisoformat(iso_day[:10])
    return (ref.year - d.year) * 12 + ref.month - d.month + (ref.day - d.day) / 30.0


# ------------------------------------------------------------------ öneri


def _filters(raw: Any) -> dict[str, Any]:
    f = raw if isinstance(raw, dict) else {}

    def num(v):
        try:
            return float(v) if v not in (None, "") else None
        except (TypeError, ValueError):
            return None

    return {"hedef": [str(x) for x in (f.get("hedef") or []) if str(x).strip()][:10],
            "yasMin": num(f.get("yasMin")), "yasMax": num(f.get("yasMax")),
            "marka": [str(x) for x in (f.get("marka") or []) if str(x).strip()][:50],
            "anahtar": _text(f.get("anahtar"), 200),
            "ozelGun": _text(f.get("ozelGun"), 120), "yalnizOzelGun": bool(f.get("yalnizOzelGun")),
            "yalnizYeni": bool(f.get("yalnizYeni"))}


def _speed_ranks(books: list[dict[str, Any]]) -> dict[str, float]:
    """Satış hızının havuzdaki yüzdelik sırası (0..1); hızı olmayan 0."""
    sp = sorted({b["hiz"] for b in books if b.get("hiz")})
    if not sp:
        return {}
    idx = {v: (i + 1) / len(sp) for i, v in enumerate(sp)}
    return {b["id"]: idx[b["hiz"]] for b in books if b.get("hiz")}


def reason_parts(b: dict[str, Any], st: dict[str, Any], day_name: Optional[str], ref: date, cfg: dict[str, Any]) -> list[str]:
    parts = []
    if b.get("yillik"):
        parts.append(f"son 12 ayda {tr_num(b['yillik'])} adet satış")
    if st.get("stokAy") is not None:
        parts.append(f"stok {tr_num(st['stokAy'], 1)} ay yeter")
    elif st.get("satisYok"):
        parts.append(f"{tr_num(st['stok'])} adet stok, son 12 ayda satışı yok")
    m = _months_since(b.get("ilkYayin"), ref)
    if m is not None and m <= cfg["newMonths"]:
        parts.append(f"ilk baskı {tr_day(b['ilkYayin'])} (yeni)")
    if day_name:
        parts.append(f"«{day_name}» ile CRM'de bağlı")
    if b.get("hedef"):
        yas = ""
        if b.get("yasBas") or b.get("yasBit"):
            yas = f" {tr_num(b.get('yasBas'))}–{tr_num(b.get('yasBit'))} yaş"
        parts.append(f"hedef kitle {b['hedef']}{yas}")
    return parts


def candidates(pool: dict[str, Any], filters: Any, cfg: dict[str, Any], exclude: Iterable[str] = (), ref: Optional[date] = None,
               extra_score: Optional[Callable[[dict[str, Any]], tuple[Optional[float], Optional[str]]]] = None,
               price_source: str = "crm") -> dict[str, Any]:
    """Süzgeçten geçen kitapların puanlı listesi (tamamı; sayfalama ekranda). `extra_score` bülten ilgi eşleşmesi içindir."""
    ref = ref or today()
    f = _filters(filters)
    books = pool.get("books") or []
    ranks = _speed_ranks(books)
    days = {d["key"]: d for d in pool.get("days") or []}
    day = days.get(f["ozelGun"] or "")
    skip = set(exclude)
    anahtar = fold(f["anahtar"]) if f["anahtar"] else None
    hedef = {fold(h) for h in f["hedef"]}
    marka = {fold(m) for m in f["marka"]}
    w = dict(cfg["weights"])
    elenen: dict[str, int] = {}
    out = []

    def drop(why: str) -> None:
        elenen[why] = elenen.get(why, 0) + 1

    for b in books:
        if b["id"] in skip:
            continue
        if hedef and fold(b.get("hedef")) not in hedef:
            continue
        if marka and fold(b.get("marka")) not in marka:
            continue
        if f["yasMin"] is not None or f["yasMax"] is not None:
            lo, hi = b.get("yasBas"), b.get("yasBit")
            if lo is None and hi is None:
                continue
            lo, hi = lo if lo is not None else hi, hi if hi is not None else lo
            if (f["yasMax"] is not None and lo > f["yasMax"]) or (f["yasMin"] is not None and hi < f["yasMin"]):
                continue
        if anahtar and not all(tok in fold(" ".join(str(b.get(k) or "") for k in ("ad", "yazar", "tur", "kitaplik", "web", "marka")))
                               for tok in anahtar.split()):
            continue
        linked = bool(day and day["key"] in (b.get("ozelGunler") or []))
        if f["yalnizOzelGun"] and day and not linked:
            continue
        m = _months_since(b.get("ilkYayin"), ref)
        if f["yalnizYeni"] and (m is None or m > cfg["newMonths"]):
            continue
        why = out_of_sale(b)
        if why:
            drop("satıştan kalkmış")
            continue
        st = stock_of(b, cfg)
        if st["stok"] is not None and st["stok"] <= 0:
            drop("stok yok")
            continue
        parts: dict[str, Optional[float]] = {
            "hiz": ranks.get(b["id"], 0.0) if b.get("hiz") is not None else None,
            "stok": (1.0 if st["satisYok"] else (min(1.0, st["stokAy"] / cfg["targetMonths"]) if st["stokAy"] is not None else None)),
            "yenilik": (max(0.0, 1.0 - 0.5 * m / cfg["newMonths"]) if m is not None and m <= cfg["newMonths"] else (0.0 if m is not None else None)),
            "ozelgun": (1.0 if linked else 0.0) if day else None,
        }
        weights = dict(w)
        extra_why = None
        if extra_score is not None:
            val, extra_why = extra_score(b)
            parts["ilgi"] = val
            weights["ilgi"] = float(cfg.get("interestWeight", 40.0))
        used = {k: weights.get(k, 0.0) for k, v in parts.items() if v is not None and weights.get(k, 0.0) > 0}
        total_w = sum(used.values())
        score = round(100 * sum(parts[k] * wt for k, wt in used.items()) / total_w, 1) if total_w else 0.0
        rp = reason_parts(b, st, day["ad"] if linked else None, ref, cfg)
        if extra_why:
            rp.insert(0, extra_why)
        out.append({"id": b["id"], "stok": b.get("stok"), "ad": b.get("ad"), "yazar": b.get("yazar"), "marka": b.get("marka"),
                    "hedef": b.get("hedef"), "yasBas": b.get("yasBas"), "yasBit": b.get("yasBit"), "ilkYayin": b.get("ilkYayin"),
                    "fiyat": price_of(b, price_source), "stokAdet": st["stok"], "stokAy": st["stokAy"], "satisYok": st["satisYok"],
                    "yillik": b.get("yillik"), "puan": score, "parcalar": {k: (round(v, 3) if v is not None else None) for k, v in parts.items()},
                    "gerekce": "; ".join(rp), "ozelGun": linked, "kapak": b.get("kapak")})
    out.sort(key=lambda x: (-x["puan"], -(x["yillik"] or 0), x["ad"] or ""))
    return {"items": out, "total": len(out), "elenen": elenen, "agirliklar": w, "ozelGun": day}


# ------------------------------------------------------------------ uyarılar


def item_alerts(item: dict[str, Any], b: Optional[dict[str, Any]], price_source: str, cfg: dict[str, Any]) -> list[dict[str, Any]]:
    """Kitabın bugünkü verisine göre uyarılar. `item`: price_snapshot, text_override, accepted (tür → değer)."""
    acc = item.get("accepted") or {}
    out: list[dict[str, Any]] = []

    def add(kind: str, text: str, value: Any = None) -> None:
        if kind in acc and acc[kind] == value and value is not None:
            return
        level, title = ALERT_KINDS[kind]
        out.append({"tur": kind, "seviye": level, "baslik": title, "metin": text, "deger": value})

    if b is None:
        add("crm-yok", "Kitap CRM'de etkin kart olarak bulunamadı (pasife alınmış ya da tipi değişmiş olabilir).", "yok")
        return out
    why = out_of_sale(b)
    if why:
        add("satis", why + ".", b.get("flag") or "kapali")
    price = price_of(b, price_source)
    snap = item.get("price_snapshot")
    if snap is not None and price is not None and abs(price - snap) >= 0.005:
        add("fiyat", f"Fiyat {tr_money(snap)} → {tr_money(price)} ({PRICE_SOURCES.get(price_source, price_source)}).", price)
    elif snap is not None and price is None:
        add("fiyat", f"Fiyat kaynağında artık fiyat yok (eklendiğinde {tr_money(snap)}).", "yok")
    st = stock_of(b, cfg)
    if st["stok"] is not None and st["stok"] <= 0:
        add("stok", "Stok yok.", "0")
    elif st["stokAy"] is not None and st["stokAy"] < cfg["criticalMonths"]:
        add("stok", f"Stok {tr_num(st['stokAy'], 1)} ay yeter ({tr_num(st['stok'])} adet); eşik {tr_num(cfg['criticalMonths'], 1)} ay.",
            round(st["stokAy"], 1))
    if b.get("hakNotu"):
        add("hak", "Telif sözleşmesinde hak notu var; kapak ve tanıtım metninin katalogda kullanımı sözleşmeye bakılarak doğrulanmalı.", "not")
    if not b.get("kapak"):
        add("kapak", "CRM'de ve web sitesinde kapak bağlantısı yok" + (f" (CRM'de dosya adı: {b['kapakDosya']})" if b.get("kapakDosya") else "") + ".", "yok")
    if not b.get("metinVar") and not item.get("text_override"):
        add("metin", "CRM'de kısa bilgi ve özet boş; katalog metni elle yazılmalı.", "yok")
    return out


def critical_count(alerts: Iterable[dict[str, Any]]) -> int:
    return sum(1 for a in alerts if a.get("seviye") == "kritik")


# ------------------------------------------------------------------ katalog kayıtları


def _row(c, tenant: str, cid: str):
    r = c.execute(sa.select(CATALOGS).where(CATALOGS.c.id == cid, CATALOGS.c.tenant_id == tenant)).mappings().first()
    if not r:
        raise CatalogError("Katalog bulunamadı.", 404)
    return r


def _view(r) -> dict[str, Any]:
    return {"id": r["id"], "tur": r["kind"], "turAdi": KINDS.get(r["kind"], r["kind"]), "baslik": r["title"], "donem": r["season"],
            "tema": r["theme"], "suzgec": _filters(load(r["filters_json"], {})), "durum": r["status"],
            "durumAdi": STATUSES.get(r["status"], r["status"]), "fiyatKaynagi": r["price_source"],
            "fiyatKaynagiAdi": PRICE_SOURCES.get(r["price_source"], r["price_source"]), "stokTarihi": r["stock_as_of"],
            "uyariZamani": iso(r["alerts_at"]), "gonderen": r["submitted_by"], "gonderim": iso(r["submitted_at"]),
            "onaylayan": r["approved_by"], "onayZamani": iso(r["approved_at"]), "yayin": iso(r["published_at"]), "not": r["note"],
            "olusturan": r["created_by"], "olusturma": iso(r["created_at"]), "guncelleme": iso(r["updated_at"])}


def create(engine, tenant: str, user: str, body: dict[str, Any], cfg: dict[str, Any]) -> dict[str, Any]:
    kind = str(body.get("tur") or "").strip()
    title = _text(body.get("baslik"), 300)
    if kind not in KINDS:
        raise CatalogError("Katalog türü seçin: " + ", ".join(KINDS.values()) + ".")
    if not title:
        raise CatalogError("Katalog adı gerekli.")
    ps = str(body.get("fiyatKaynagi") or cfg["priceSource"])
    if ps not in PRICE_SOURCES:
        raise CatalogError("Fiyat kaynağı geçersiz.")
    cid, t = uid(), now()
    with engine.begin() as c:
        c.execute(CATALOGS.insert().values(id=cid, tenant_id=tenant, kind=kind, title=title, season=_text(body.get("donem"), 120),
                                           theme=_text(body.get("tema"), 300), filters_json=dump(_filters(body.get("suzgec"))),
                                           status="taslak", price_source=ps, created_by=user, created_at=t, updated_at=t))
    return get(engine, tenant, cid)


def get(engine, tenant: str, cid: str) -> dict[str, Any]:
    with engine.connect() as c:
        return _view(_row(c, tenant, cid))


def list_catalogs(engine, tenant: str, durum: str = "") -> dict[str, Any]:
    with engine.connect() as c:
        q = sa.select(CATALOGS).where(CATALOGS.c.tenant_id == tenant)
        if durum == "acik":
            q = q.where(CATALOGS.c.status != "arsiv")
        elif durum:
            q = q.where(CATALOGS.c.status == durum)
        rows = c.execute(q.order_by(CATALOGS.c.updated_at.desc())).mappings().all()
        counts: dict[str, dict[str, int]] = {}
        for it in c.execute(sa.select(ITEMS.c.catalog_id, ITEMS.c.alert_json, ITEMS.c.featured)
                            .where(ITEMS.c.catalog_id.in_([r["id"] for r in rows] or [""]))).mappings():
            d = counts.setdefault(it["catalog_id"], {"kitap": 0, "kritik": 0, "bilgi": 0, "oneCikan": 0})
            d["kitap"] += 1
            d["oneCikan"] += 1 if it["featured"] else 0
            for a in load(it["alert_json"], []) or []:
                d["kritik" if a.get("seviye") == "kritik" else "bilgi"] += 1
    items = [{**_view(r), **counts.get(r["id"], {"kitap": 0, "kritik": 0, "bilgi": 0, "oneCikan": 0})} for r in rows]
    return {"items": items, "total": len(items)}


def update(engine, tenant: str, cid: str, body: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    vals: dict[str, Any] = {}
    with engine.begin() as c:
        r = _row(c, tenant, cid)
        if r["status"] not in ("taslak",) and set(body) - {"not"}:
            raise CatalogError("Yalnız taslak katalog değiştirilir; önce taslağa geri alın.", 409)
        if "baslik" in body:
            t = _text(body["baslik"], 300)
            if not t:
                raise CatalogError("Katalog adı boş olamaz.")
            vals["title"] = t
        if "tur" in body:
            if body["tur"] not in KINDS:
                raise CatalogError("Katalog türü geçersiz.")
            vals["kind"] = body["tur"]
        if "donem" in body:
            vals["season"] = _text(body["donem"], 120)
        if "tema" in body:
            vals["theme"] = _text(body["tema"], 300)
        if "suzgec" in body:
            vals["filters_json"] = dump(_filters(body["suzgec"]))
        if "fiyatKaynagi" in body:
            if body["fiyatKaynagi"] not in PRICE_SOURCES:
                raise CatalogError("Fiyat kaynağı geçersiz.")
            vals["price_source"] = body["fiyatKaynagi"]
        diff = {k: {"eski": r[k], "yeni": v} for k, v in vals.items() if r[k] != v}
        if vals:
            vals["updated_at"] = now()
            c.execute(CATALOGS.update().where(CATALOGS.c.id == cid).values(**vals))
    return get(engine, tenant, cid), diff


def delete(engine, tenant: str, cid: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = _row(c, tenant, cid)
        if r["status"] != "taslak":
            raise CatalogError("Yalnız taslak katalog silinir; onaylı ya da yayındaki katalog arşive alınır.", 409)
        c.execute(ITEMS.delete().where(ITEMS.c.catalog_id == cid))
        c.execute(JOBS.delete().where(JOBS.c.ref_id == cid))
        c.execute(CATALOGS.delete().where(CATALOGS.c.id == cid))
    return _view(r)


def _items(c, cid: str) -> list[dict[str, Any]]:
    rows = c.execute(sa.select(ITEMS).where(ITEMS.c.catalog_id == cid).order_by(ITEMS.c.position)).mappings().all()
    return [dict(r) for r in rows]


def set_items(engine, tenant: str, user: str, cid: str, items: list[dict[str, Any]], pool: dict[str, Any],
              cfg: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """Katalog kitap listesini (sıra, öne çıkan, sayfa ipucu, gerekçe, metin) yazar. Yeni kitabın fiyat/stok dayanağı
    eklendiği andaki değerdir; var olanınki korunur."""
    books = {b["id"]: b for b in pool.get("books") or []}
    if not isinstance(items, list):
        raise CatalogError("Kitap listesi gerekli.")
    seen: list[str] = []
    for it in items:
        bid = str((it or {}).get("crmKitapId") or "").lower()
        if not bid or bid in seen:
            raise CatalogError("Listede kimliği boş ya da iki kez geçen kitap var.")
        seen.append(bid)
    with engine.begin() as c:
        r = _row(c, tenant, cid)
        if r["status"] != "taslak":
            raise CatalogError("Yalnız taslak katalogda kitap listesi değişir.", 409)
        old = {x["crm_book_id"]: x for x in _items(c, cid)}
        c.execute(ITEMS.delete().where(ITEMS.c.catalog_id == cid))
        added, removed = [x for x in seen if x not in old], [x for x in old if x not in seen]
        for pos, it in enumerate(items, start=1):
            bid = str(it["crmKitapId"]).lower()
            prev = old.get(bid) or {}
            b = books.get(bid)
            if not prev and b is None:
                raise CatalogError(f"Kitap havuzda yok: {bid}. Havuzu yenileyip yeniden deneyin.", 409)
            st = stock_of(b, cfg) if b else {"stokAy": None}
            text = it.get("metin") if "metin" in it else prev.get("text_override")
            text = _text(text, 4000)
            text_src = prev.get("text_source") if text == prev.get("text_override") else ("kullanici" if text else None)
            c.execute(ITEMS.insert().values(
                catalog_id=cid, crm_book_id=bid, stok_kodu=(b or {}).get("stok") or prev.get("stok_kodu"),
                ad=(b or {}).get("ad") or prev.get("ad"), position=pos, featured=bool(it.get("oneCikan")),
                page_hint=_text(it.get("sayfa"), 120),
                reason=prev.get("reason") if prev else _text(it.get("gerekce"), 2000),
                reason_ai=prev.get("reason_ai"),
                price_snapshot=prev.get("price_snapshot") if prev else (price_of(b, r["price_source"]) if b else None),
                stock_months_snapshot=prev.get("stock_months_snapshot") if prev else st.get("stokAy"),
                text_override=text, text_source=text_src, alert_json=prev.get("alert_json"),
                accepted_json=prev.get("accepted_json"), added_by=prev.get("added_by") or user, added_at=prev.get("added_at") or now()))
        c.execute(CATALOGS.update().where(CATALOGS.c.id == cid).values(updated_at=now()))
    return detail(engine, tenant, cid, pool, cfg), {"eklenen": len(added), "cikan": len(removed), "toplam": len(seen)}


def accept(engine, tenant: str, cid: str, bid: str, kind: str, pool: dict[str, Any], cfg: dict[str, Any]) -> dict[str, Any]:
    """Uyarıyı kabul: fiyat için bugünkü fiyat yeni dayanak olur; diğerlerinde o anki değer not edilir."""
    books = {b["id"]: b for b in pool.get("books") or []}
    bid = bid.lower()
    with engine.begin() as c:
        r = _row(c, tenant, cid)
        it = c.execute(sa.select(ITEMS).where(ITEMS.c.catalog_id == cid, ITEMS.c.crm_book_id == bid)).mappings().first()
        if not it:
            raise CatalogError("Kitap bu katalogda yok.", 404)
        b = books.get(bid)
        cur = {a["tur"]: a for a in item_alerts({**dict(it), "accepted": {}}, b, r["price_source"], cfg)}
        if kind not in cur:
            raise CatalogError("Bu kitapta bu türde açık uyarı yok.", 409)
        vals: dict[str, Any] = {}
        acc = load(it["accepted_json"], {}) or {}
        if kind == "fiyat" and b is not None and price_of(b, r["price_source"]) is not None:
            vals["price_snapshot"] = price_of(b, r["price_source"])
            acc.pop("fiyat", None)
        else:
            acc[kind] = cur[kind]["deger"]
        vals["accepted_json"] = dump(acc)
        c.execute(ITEMS.update().where(ITEMS.c.catalog_id == cid, ITEMS.c.crm_book_id == bid).values(**vals))
        c.execute(CATALOGS.update().where(CATALOGS.c.id == cid).values(updated_at=now()))
    return {"tur": kind, "deger": cur[kind]["deger"], "ad": it["ad"]}


def detail(engine, tenant: str, cid: str, pool: Optional[dict[str, Any]], cfg: dict[str, Any]) -> dict[str, Any]:
    """Katalog + kitaplar; uyarılar havuzun bugünkü verisiyle canlı hesaplanır."""
    books = {b["id"]: b for b in (pool or {}).get("books") or []}
    with engine.connect() as c:
        r = _row(c, tenant, cid)
        rows = _items(c, cid)
    out = _view(r)
    items = []
    for it in rows:
        b = books.get(it["crm_book_id"]) if pool else None
        acc = load(it["accepted_json"], {}) or {}
        alerts = item_alerts({**it, "accepted": acc}, b, r["price_source"], cfg) if pool else (load(it["alert_json"], []) or [])
        st = stock_of(b, cfg) if b else {"stok": None, "stokAy": None, "satisYok": False}
        items.append({
            "crmKitapId": it["crm_book_id"], "stokKodu": it["stok_kodu"], "ad": (b or {}).get("ad") or it["ad"],
            "yazar": (b or {}).get("yazar"), "marka": (b or {}).get("marka"), "isbn": (b or {}).get("isbn"),
            "hedef": (b or {}).get("hedef"), "yasBas": (b or {}).get("yasBas"), "yasBit": (b or {}).get("yasBit"),
            "sira": it["position"], "oneCikan": bool(it["featured"]), "sayfa": it["page_hint"], "gerekce": it["reason"],
            "gerekceZeki": it["reason_ai"], "metin": it["text_override"], "metinKaynagi": it["text_source"],
            "metinVar": bool((b or {}).get("metinVar")), "fiyat": price_of(b, r["price_source"]) if b else None,
            "fiyatDayanak": it["price_snapshot"], "stokAdet": st.get("stok"), "stokAy": st.get("stokAy"), "satisYok": st.get("satisYok"),
            "stokAyDayanak": it["stock_months_snapshot"], "hiz": (b or {}).get("hiz"), "yillik": (b or {}).get("yillik"),
            "kapak": (b or {}).get("kapak"), "kapakKaynak": (b or {}).get("kapakKaynak"), "webUrl": (b or {}).get("webUrl"),
            "uyarilar": alerts, "kritik": critical_count(alerts)})
    out["kitaplar"] = items
    out["ozet"] = {"kitap": len(items), "oneCikan": sum(1 for i in items if i["oneCikan"]),
                   "kritik": sum(i["kritik"] for i in items), "uyariliKitap": sum(1 for i in items if i["kritik"]),
                   "bilgi": sum(len(i["uyarilar"]) - i["kritik"] for i in items),
                   "fiyatsiz": sum(1 for i in items if i["fiyat"] is None), "kapaksiz": sum(1 for i in items if not i["kapak"])}
    out["havuz"] = {"okuma": (pool or {}).get("readAt"), "logoSon": (pool or {}).get("logoSon"),
                    "notlar": (pool or {}).get("notes") or []} if pool else None
    return out


def refresh_alerts(engine, tenant: str, pool: dict[str, Any], cfg: dict[str, Any]) -> list[dict[str, Any]]:
    """Açık kataloglarda (arşiv dışı) uyarıları yeniden hesaplar ve yazar; katalog başına kritik uyarı özeti döner."""
    books = {b["id"]: b for b in pool.get("books") or []}
    summary = []
    t = now()
    with engine.begin() as c:
        cats = c.execute(sa.select(CATALOGS).where(CATALOGS.c.tenant_id == tenant, CATALOGS.c.status != "arsiv")).mappings().all()
        for r in cats:
            crit = 0
            lines = []
            for it in _items(c, r["id"]):
                alerts = item_alerts({**it, "accepted": load(it["accepted_json"], {}) or {}}, books.get(it["crm_book_id"]),
                                     r["price_source"], cfg)
                c.execute(ITEMS.update().where(ITEMS.c.catalog_id == r["id"], ITEMS.c.crm_book_id == it["crm_book_id"])
                          .values(alert_json=dump(alerts)))
                k = [a for a in alerts if a["seviye"] == "kritik"]
                crit += len(k)
                lines += [f"{it['ad'] or it['stok_kodu']}: {a['metin']}" for a in k]
            c.execute(CATALOGS.update().where(CATALOGS.c.id == r["id"]).values(alerts_at=t, stock_as_of=pool.get("logoSon")))
            summary.append({"id": r["id"], "baslik": r["title"], "durum": r["status"], "kritik": crit, "satirlar": lines})
    return summary


# ------------------------------------------------------------------ onay akışı


def transition(engine, tenant: str, user: str, cid: str, action: str, note: Optional[str] = None) -> dict[str, Any]:
    """submit (taslak→onayda) · withdraw (onayda→taslak) · approve (onayda→onayli; gönderen onaylayamaz) · reject
    (onayda→taslak, gerekçe şart) · publish (onayli→yayinda) · archive (→arsiv) · reopen (onayli/yayinda/arsiv→taslak)."""
    flows = {"submit": ({"taslak"}, "onayda"), "withdraw": ({"onayda"}, "taslak"), "approve": ({"onayda"}, "onayli"),
             "reject": ({"onayda"}, "taslak"), "publish": ({"onayli"}, "yayinda"),
             "archive": ({"taslak", "onayli", "yayinda"}, "arsiv"), "reopen": ({"onayli", "yayinda", "arsiv"}, "taslak")}
    if action not in flows:
        raise CatalogError("İşlem geçersiz.", 404)
    allowed, target = flows[action]
    t = now()
    with engine.begin() as c:
        r = _row(c, tenant, cid)
        if r["status"] not in allowed:
            raise CatalogError(f"Katalog «{STATUSES.get(r['status'])}» durumundayken bu işlem yapılamaz.", 409)
        vals: dict[str, Any] = {"status": target, "updated_at": t}
        if action == "submit":
            n = c.execute(sa.select(sa.func.count()).select_from(ITEMS).where(ITEMS.c.catalog_id == cid)).scalar() or 0
            if not n:
                raise CatalogError("Boş katalog onaya gönderilemez; önce kitap ekleyin.", 409)
            vals.update(submitted_by=user, submitted_at=t, note=None)
        elif action == "approve":
            if (r["submitted_by"] or "").lower() == user.lower():
                raise CatalogError("Onaya gönderen kişi aynı kataloğu onaylayamaz.", 403)
            vals.update(approved_by=user, approved_at=t)
        elif action == "reject":
            if not (note or "").strip():
                raise CatalogError("Geri gönderme gerekçesi yazın.")
            vals.update(note=note.strip()[:2000], approved_by=None, approved_at=None)
        elif action == "publish":
            vals.update(published_at=t)
        elif action == "reopen":
            vals.update(approved_by=None, approved_at=None, published_at=None)
        c.execute(CATALOGS.update().where(CATALOGS.c.id == cid).values(**vals))
    return get(engine, tenant, cid)


def pending_counts(engine, tenant: str) -> dict[str, int]:
    with engine.connect() as c:
        n = c.execute(sa.select(sa.func.count()).select_from(CATALOGS)
                      .where(CATALOGS.c.tenant_id == tenant, CATALOGS.c.status == "onayda")).scalar() or 0
    return {"katalogOnayda": int(n)}


# ------------------------------------------------------------------ meta (havuz anlık görüntüsü)


def meta_get(engine, tenant: str, key: str) -> tuple[Any, Optional[str]]:
    with engine.connect() as c:
        r = c.execute(sa.select(META).where(META.c.tenant_id == tenant, META.c.key == key)).mappings().first()
    return (load(r["value_json"], None), iso(r["updated_at"])) if r else (None, None)


def meta_set(engine, tenant: str, key: str, value: Any) -> None:
    with engine.begin() as c:
        c.execute(META.delete().where(META.c.tenant_id == tenant, META.c.key == key))
        c.execute(META.insert().values(tenant_id=tenant, key=key, value_json=dump(value), updated_at=now()))


# ------------------------------------------------------------------ işler (Zeki AI)


def job_create(engine, tenant: str, user: str, kind: str, ref_id: str) -> dict[str, Any]:
    with engine.begin() as c:
        running = c.execute(sa.select(JOBS.c.id).where(JOBS.c.ref_id == ref_id, JOBS.c.kind == kind,
                                                       JOBS.c.status.in_(("bekliyor", "calisiyor")))).first()
        if running:
            raise CatalogError("Aynı iş zaten sürüyor; bitince yeniden deneyin.", 409)
        jid = uid()
        c.execute(JOBS.insert().values(id=jid, tenant_id=tenant, kind=kind, ref_id=ref_id, status="bekliyor", created_by=user,
                                       created_at=now()))
    return job_get(engine, tenant, jid)


def job_update(engine, jid: str, **vals: Any) -> None:
    if "result" in vals:
        vals["result_json"] = dump(vals.pop("result"))
    if vals.get("status") in ("bitti", "hata"):
        vals["finished_at"] = now()
    with engine.begin() as c:
        c.execute(JOBS.update().where(JOBS.c.id == jid).values(**vals))


def job_get(engine, tenant: str, jid: str) -> dict[str, Any]:
    with engine.connect() as c:
        j = c.execute(sa.select(JOBS).where(JOBS.c.id == jid, JOBS.c.tenant_id == tenant)).mappings().first()
    if not j:
        raise CatalogError("İş bulunamadı.", 404)
    return {"id": j["id"], "tur": j["kind"], "ref": j["ref_id"], "durum": j["status"], "adim": j["step"],
            "sonuc": load(j["result_json"], None), "hata": j["error"], "olusturma": iso(j["created_at"]), "bitis": iso(j["finished_at"])}


def _words(s: str) -> int:
    return len(re.findall(r"\S+", s or ""))


def zeki_text(chat: Callable[[list[dict[str, str]]], str], book: dict[str, Any], texts: dict[str, Optional[str]], words: int,
              english: bool) -> dict[str, Any]:
    """Katalog tanıtım metni: CRM kısa bilgi/özetinden en çok `words` kelime (yabancı hak kataloğunda İngilizce taslak).
    Yeni bilgi, sayı, alıntı, üstünlük iddiası eklenmez; denetimden geçmeyen cümle düşer."""
    from semantic_bridge.marketing import guard

    src = texts.get("kisa") or texts.get("ozet")
    if not src:
        return {"metin": None, "dusen": [], "not": "CRM'de kısa bilgi ve özet boş."}
    lang = ("İngilizce yaz (yabancı yayınevlerine gidecek hak kataloğu). Kitap adını çevirme, özgün Türkçe adı koru."
            if english else "Türkçe yaz.")
    sys_msg = ("Timaş Yayınları'nın katalog editörüsün. Verilen tanıtım metnini katalog için kısaltırsın. Metinde olmayan "
               "hiçbir bilgi, sayı, ödül, alıntı ya da «en çok satan» gibi üstünlük iddiası ekleme. Yalnız metni döndür.")
    user = (f"Kitap: {book.get('ad')} — {book.get('yazar') or ''}\nEn çok {words} kelime. {lang}\n\nTanıtım metni:\n{src[:5000]}")
    raw = (chat([{"role": "system", "content": sys_msg}, {"role": "user", "content": user}]) or "").strip()
    raw = re.sub(r"^```\w*|```$", "", raw).strip().strip('"').strip()
    chk = guard.check(raw, [src, texts.get("ozet") or "", book.get("ad") or "", book.get("yazar") or ""])
    return {"metin": chk["metin"] or None, "dusen": chk["dusen"], "kelime": _words(chk["metin"])}


def zeki_reason(chat: Callable[[list[dict[str, str]]], str], book: dict[str, Any], facts: str, kind_label: str) -> dict[str, Any]:
    """Kitabın katalogda yer alma gerekçesi: yalnız verilen olgulardan tek cümle."""
    from semantic_bridge.marketing import guard

    sys_msg = ("Katalog hazırlayan pazarlama uzmanına kitabın neden katalogda olduğunu tek cümleyle açıklarsın. Yalnız verilen "
               "olguları kullan; olguda olmayan sayı, iddia ya da bilgi ekleme. Tek cümle döndür.")
    user = f"Katalog: {kind_label}\nKitap: {book.get('ad')} — {book.get('yazar') or ''}\nOlgular: {facts}"
    raw = (chat([{"role": "system", "content": sys_msg}, {"role": "user", "content": user}]) or "").strip()
    chk = guard.check(raw.splitlines()[0] if raw else "", [facts, book.get("ad") or "", book.get("yazar") or ""], [facts])
    return {"metin": chk["metin"] or None, "dusen": chk["dusen"]}


def run_zeki(engine, tenant: str, jid: str, cid: str, chat: Callable[[list[dict[str, str]]], str],
             read_texts: Callable[[list[str]], dict[str, dict[str, Optional[str]]]], pool: dict[str, Any], opts: dict[str, Any],
             cfg: dict[str, Any]) -> None:
    """Katalogdaki kitaplara tanıtım metni ve/veya tek cümle gerekçe. Elle yazılmış metne dokunmaz."""
    try:
        job_update(engine, jid, status="calisiyor", step="Kitaplar okunuyor")
        d = detail(engine, tenant, cid, pool, cfg)
        english = d["tur"] == "yabanci-hak"
        words = int(opts.get("kelime") or cfg["textWords"])
        items = d["kitaplar"]
        texts = read_texts([i["crmKitapId"] for i in items]) if opts.get("metin") else {}
        res = {"metin": 0, "gerekce": 0, "dusen": 0, "atlanan": 0, "bos": 0}
        for n, it in enumerate(items, start=1):
            job_update(engine, jid, step=f"{n}/{len(items)} · {it['ad'] or it['stokKodu']}")
            vals: dict[str, Any] = {}
            if opts.get("metin"):
                if it["metin"] and it["metinKaynagi"] == "kullanici":
                    res["atlanan"] += 1
                elif it["metin"] and not opts.get("yeniden"):
                    res["atlanan"] += 1
                else:
                    out = zeki_text(chat, it, texts.get(it["crmKitapId"], {}), words, english)
                    res["dusen"] += len(out["dusen"])
                    if out["metin"]:
                        vals.update(text_override=out["metin"], text_source="zeki")
                        res["metin"] += 1
                    else:
                        res["bos"] += 1
            if opts.get("gerekce") and it.get("gerekce"):
                out = zeki_reason(chat, it, it["gerekce"], d["turAdi"])
                res["dusen"] += len(out["dusen"])
                if out["metin"]:
                    vals["reason_ai"] = out["metin"]
                    res["gerekce"] += 1
            if vals:
                with engine.begin() as c:
                    c.execute(ITEMS.update().where(ITEMS.c.catalog_id == cid, ITEMS.c.crm_book_id == it["crmKitapId"]).values(**vals))
        job_update(engine, jid, status="bitti", step=None, result=res)
    except Exception as e:  # noqa: BLE001 — iş hatası ekranda görünür
        log.exception("catalog zeki job failed")
        job_update(engine, jid, status="hata", step=None, error=str(e)[:500] or e.__class__.__name__)


# ------------------------------------------------------------------ dışa aktarım


def _text_for(it: dict[str, Any], texts: dict[str, dict[str, Optional[str]]]) -> tuple[str, str]:
    if it.get("metin"):
        return it["metin"], ("Zeki AI kısaltması (denetimli)" if it.get("metinKaynagi") == "zeki" else "elle yazıldı")
    t = texts.get(it["crmKitapId"]) or {}
    if t.get("kisa"):
        return t["kisa"], "CRM kısa bilgi"
    if t.get("ozet"):
        return t["ozet"], "CRM özet"
    return "", "yok"


def _age(it: dict[str, Any]) -> str:
    if it.get("yasBas") or it.get("yasBit"):
        return f"{tr_num(it.get('yasBas'))}–{tr_num(it.get('yasBit'))} yaş"
    return ""


def catalog_xlsx(d: dict[str, Any], texts: dict[str, dict[str, Optional[str]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Katalog"
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="EDE9FE")
    ws["A1"] = f"{d['baslik']} — {d['turAdi']}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"Dönem: {d.get('donem') or '—'} · Tema: {d.get('tema') or '—'} · Durum: {d['durumAdi']} · "
                f"Fiyat: {d['fiyatKaynagiAdi']} · Stok verisi {tr_day((d.get('havuz') or {}).get('logoSon'))} tarihine kadar")
    cols = ["Sıra", "Öne çıkan", "Sayfa", "Stok kodu", "ISBN", "Kitap", "Yazar", "Marka", "Hedef kitle", "Yaş", "Fiyat (₺)",
            "Stok (adet)", "Stok (ay)", "Kritik uyarı", "Tanıtım metni", "Metin kaynağı", "Kapak bağlantısı", "Gerekçe"]
    for i, h in enumerate(cols, start=1):
        cell = ws.cell(row=4, column=i, value=h)
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for n, it in enumerate(d["kitaplar"], start=5):
        text, src = _text_for(it, texts)
        crit = "; ".join(a["metin"] for a in it["uyarilar"] if a["seviye"] == "kritik")
        vals = [it["sira"], "evet" if it["oneCikan"] else "", it.get("sayfa") or "", it.get("stokKodu"), it.get("isbn"), it.get("ad"),
                it.get("yazar"), it.get("marka"), it.get("hedef"), _age(it), it.get("fiyat"), it.get("stokAdet"),
                None if it.get("stokAy") is None else round(it["stokAy"], 1), crit, text, src, it.get("kapak") or "",
                it.get("gerekceZeki") or it.get("gerekce") or ""]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=n, column=i, value=v)
        ws.cell(row=n, column=11).number_format = "#,##0.00"
    for col, width in zip("ABCDEFGHIJKLMNOPQR", (6, 9, 10, 16, 16, 40, 24, 18, 12, 10, 11, 11, 9, 40, 70, 16, 40, 60)):
        ws.column_dimensions[col].width = width
    info = wb.create_sheet("Bilgi")
    rows = [("Katalog", d["baslik"]), ("Tür", d["turAdi"]), ("Durum", d["durumAdi"]), ("Onaylayan", d.get("onaylayan") or "—"),
            ("Fiyat kaynağı", d["fiyatKaynagiAdi"]), ("Fiyat kaynağı gerekçesi", PRICE_SOURCE_WHY),
            ("Stok", "Stok ay sayısı = stok ÷ ağırlıklı aylık satış hızı (Yönetim raporları › Baskı önerisi tanımı)."),
            ("Stok verisi", f"Logo'daki son faturalı satış günü: {tr_day((d.get('havuz') or {}).get('logoSon'))}"),
            ("Kitap havuzu okuma", (d.get("havuz") or {}).get("okuma") or "—"), ("Hazırlayan", PRODUCT)]
    for i, (k, v) in enumerate(rows, start=1):
        info.cell(row=i, column=1, value=k).font = bold
        info.cell(row=i, column=2, value=v)
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 26, 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def package_zip(d: dict[str, Any], texts: dict[str, dict[str, Optional[str]]]) -> bytes:
    """Tasarımcı paketi: Excel + kapak bağlantıları (CSV) + tanıtım metinleri + sayfa brief'i. Görsel indirilmez;
    bağlantısı olmayan kapak listede «eksik» yazar."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("katalog.xlsx", catalog_xlsx(d, texts))
        cb = io.StringIO()
        w = csv.writer(cb, delimiter=";")
        w.writerow(["Sıra", "Stok kodu", "ISBN", "Kitap", "Kapak bağlantısı", "Kaynak", "Durum"])
        for it in d["kitaplar"]:
            w.writerow([it["sira"], it.get("stokKodu") or "", it.get("isbn") or "", it.get("ad") or "", it.get("kapak") or "",
                        {"crm": "CRM kapak bağlantısı", "web": "Web sitesi görseli"}.get(it.get("kapakKaynak") or "", ""),
                        "hazır" if it.get("kapak") else "EKSİK — grafik ekibinden istenmeli"])
        z.writestr("kapaklar.csv", "﻿" + cb.getvalue())
        lines = []
        for it in d["kitaplar"]:
            text, src = _text_for(it, texts)
            lines.append(f"{it['sira']}. {it.get('ad') or ''} — {it.get('yazar') or ''} ({it.get('stokKodu') or ''})\n"
                         f"Kaynak: {src}\n{text or '(metin yok)'}\n")
        z.writestr("metinler.txt", "\n".join(lines))
        feat = [it for it in d["kitaplar"] if it["oneCikan"]]
        pages: dict[str, list[str]] = {}
        for it in d["kitaplar"]:
            pages.setdefault(it.get("sayfa") or "(sayfa belirtilmedi)", []).append(f"  {it['sira']}. {it.get('ad') or ''}"
                                                                                  + (" [öne çıkan]" if it["oneCikan"] else ""))
        brief = [f"{d['baslik']} — {d['turAdi']}", f"Dönem: {d.get('donem') or '—'} · Tema: {d.get('tema') or '—'}",
                 f"Durum: {d['durumAdi']}" + (f" · onaylayan {d['onaylayan']}" if d.get("onaylayan") else ""),
                 f"Fiyat: {d['fiyatKaynagiAdi']} · stok verisi {tr_day((d.get('havuz') or {}).get('logoSon'))} tarihine kadar",
                 f"Kitap: {len(d['kitaplar'])} · öne çıkan: {len(feat)} · kapağı eksik: {d['ozet']['kapaksiz']}", "",
                 "Sayfa düzeni (katalog sorumlusunun sayfa ipuçları):"]
        for page, lst in pages.items():
            brief.append(f"- {page}")
            brief.extend(lst)
        if d["ozet"]["kritik"]:
            brief += ["", f"Dikkat: {d['ozet']['kritik']} kritik uyarı açık (fiyat/stok/satış durumu); basımdan önce kapatılmalı."]
        z.writestr("BENIOKU.txt", "\n".join(brief) + "\n")
    return buf.getvalue()


def preview_pdf(d: dict[str, Any], texts: dict[str, dict[str, Optional[str]]], per_page: int) -> bytes:
    """Basit PDF önizleme: sayfa başına `per_page` kitap (1–12), iki sütun. Kapak görseli bu sürümde basılmaz."""
    try:
        from fpdf import FPDF
    except ImportError as e:  # pragma: no cover
        raise CatalogError("PDF üretici sunucuda kurulu değil.", 503) from e
    from semantic_bridge.editorial_export import _fold, _resolve_fonts

    regular, bold = _resolve_fonts()
    fam = "Body" if regular else "Helvetica"
    T = (lambda s: str(s or "")) if regular else (lambda s: _fold(str(s or "")))
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_title(T(d["baslik"]))
    pdf.set_author(PRODUCT)
    pdf.set_creator(PRODUCT)
    pdf.set_margins(12, 12, 12)
    pdf.set_auto_page_break(auto=False)
    if regular:
        pdf.add_font(fam, "", str(regular))
        pdf.add_font(fam, "B", str(bold or regular))
    per_page = min(12, max(1, int(per_page)))
    cols = 1 if per_page <= 2 else 2
    rows = math.ceil(per_page / cols)
    W = pdf.w - pdf.l_margin - pdf.r_margin
    top, bottom = 30.0, pdf.h - 16
    cell_w, cell_h = W / cols, (bottom - top) / rows
    items = d["kitaplar"]
    pages = [items[i:i + per_page] for i in range(0, len(items), per_page)] or [[]]
    foot = (f"{d['fiyatKaynagiAdi']} · stok verisi {tr_day((d.get('havuz') or {}).get('logoSon'))} tarihine kadar · "
            f"{d['durumAdi']} · {PRODUCT}")

    def fit(s: str, w: float) -> str:
        t = T(s)
        while pdf.get_string_width(t) > w and len(t) > 2:
            t = t[:-2] + "…"
        return t

    for pn, chunk in enumerate(pages, start=1):
        pdf.add_page()
        pdf.set_font(fam, "B", 14)
        pdf.cell(W, 7, fit(d["baslik"], W), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(fam, "", 8.5)
        pdf.cell(W, 5, fit(f"Timaş Yayınları · {d['turAdi']}" + (f" · {d['donem']}" if d.get("donem") else "")
                           + (f" · {d['tema']}" if d.get("tema") else ""), W), new_x="LMARGIN", new_y="NEXT")
        for i, it in enumerate(chunk):
            x = pdf.l_margin + (i % cols) * cell_w
            y = top + (i // cols) * cell_h
            pdf.set_draw_color(220, 215, 235)
            pdf.rect(x + 1, y + 1, cell_w - 2, cell_h - 2)
            iw = cell_w - 8
            pdf.set_xy(x + 4, y + 4)
            pdf.set_font(fam, "B", 10.5)
            pdf.cell(iw, 5.5, fit(("★ " if it["oneCikan"] else "") + (it.get("ad") or ""), iw), new_x="LEFT", new_y="NEXT")
            pdf.set_font(fam, "", 8.5)
            pdf.cell(iw, 4.5, fit(" · ".join(x for x in (it.get("yazar"), it.get("marka")) if x), iw), new_x="LEFT", new_y="NEXT")
            meta = " · ".join(x for x in (f"ISBN {it['isbn']}" if it.get("isbn") else None, _age(it) or None,
                                           tr_money(it.get("fiyat")) if it.get("fiyat") is not None else "fiyat yok") if x)
            pdf.cell(iw, 4.5, fit(meta, iw), new_x="LEFT", new_y="NEXT")
            text, _src = _text_for(it, texts)
            if text:
                pdf.set_font(fam, "", 8)
                room = max(1, int((y + cell_h - 4 - pdf.get_y()) / 3.8))
                words, lines, cur = T(text).split(), [], ""
                for wd in words:
                    trial = (cur + " " + wd).strip()
                    if pdf.get_string_width(trial) <= iw:
                        cur = trial
                    else:
                        lines.append(cur)
                        cur = wd
                    if len(lines) >= room:
                        break
                if cur and len(lines) < room:
                    lines.append(cur)
                if len(lines) >= room and len(words) > sum(len(ln.split()) for ln in lines):
                    lines[-1] = fit(lines[-1] + " …", iw)
                for ln in lines[:room]:
                    pdf.cell(iw, 3.8, ln, new_x="LEFT", new_y="NEXT")
        pdf.set_xy(pdf.l_margin, pdf.h - 12)
        pdf.set_font(fam, "", 7.5)
        pdf.cell(W - 20, 4, fit(foot, W - 22))
        pdf.cell(20, 4, f"{pn}/{len(pages)}", align="R")
    return bytes(pdf.output())
