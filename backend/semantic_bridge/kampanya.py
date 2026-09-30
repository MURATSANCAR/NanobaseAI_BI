"""M35 E-ticaret kampanya ve promosyon yönetimi: kampanya kayıt defteri, kitap bazında indirim simülasyonu ve kontroller,
aday kitaplar, takvim, sonuç (önce/sonra) ve öğrenim kaydı.

**Kaynaklar** (`kampanya_sources`): Logo (malzeme kartı ve KDV, stok bakiyesi, `PRCLIST`, aylık faturalı satış, güncel yılın
brüt farkı, kampanya kitaplarının günlük satışı) ve CRM (kitap kartı, yürürlükteki Telif Alış sözleşmesinin asgari perakende
fiyatı / hesaplama tipi / telif oranı, bayi kampanyaları ve kampanya kodlu siparişler). İkisi de yalnız okunur. Özel günler ve
kitap bağları SEO sezon takviminin tablolarından (`semantic_seo_seasons_*`), sitenin günlük fiyatı SEO modülünün T-soft ürün
tablosundan (`semantic_seo_products`) okunur; T-soft'a hiçbir istek gitmez. Gece işi (`Refresher.run`) okumaları köprünün
`semantic_kampanya_*` tablolarına yazar; ekranlar bu tablolardan okur.

**Yazma yok, dış gönderim yok:** kampanya taslak → onay → insan platform panelinde / T-soft'ta / CRM'de elle kurar. Modül hiçbir
platforma başvuru göndermez ve «gönderildi» durumu yoktur; kişi «elle kurdum» diye işaretler (`kuruldu_*`). Bayi kampanyası
CRM'de elle tanımlanır; portal kaydı CRM'e yazılmaz, ekranda «CRM'e işlenecek» diye durur. Her yazma `semantic_audit`'e düşer.

**Hesap (model yok):**
- *Birim net gelir* = fiyat (KDV dahil) ÷ (1 + KDV) × (1 − kanal kesintisi). Kanal kesintisi kampanyada girilir (pazar yeri
  komisyonu, bayi iskontosu); girilmediyse 0 sayılır ve satırda yazılır.
- *Birim maliyet* M9 sağlayıcısından (onaylı analiz → M9 Logo gerçekleşen), yoksa (`KAMPANYA_COST_SOURCE=m9+logo`) bu modülün
  okuduğu güncel yıl Logo maliyeti (maliyetli satırların Σ maliyet ÷ Σ adet). Yoksa marj **hesaplanamaz**; sıfır yazılmaz.
- *Birim telif*: satıştan ödenen (Satıştan / kademeli / Baskı + satış) yürürlükteki sözleşmelerin oranı; esas CRM «Hesaplama
  Tipi»nden: Toptan satış fiyatı → birim net gelir, Perakende fiyatı → liste fiyatı (indirim telife yansımaz), Perakende oranlı →
  kampanya perakende fiyatı. Hesaplama tipi boşsa telif türü (Brüt → liste, Net → net gelir). Baskıdan/tek ödemede satış fiyatı
  telifi değiştirmez.
- *Marj* = birim net gelir − birim maliyet − birim telif; kampanya öncesi (liste fiyatı) ve kampanyalı fiyatla aynı formül.
  Logo'nun güncel yıl brüt farkı ayrıca «gerçekleşen» olarak gösterilir.
- *Tükenme*: son `KAMPANYA_ADAY_HIZ_AY` ayın günlük hızı × beklenen artış (kampanyada girilir; boşsa aynı kanaldaki öğrenim
  kayıtlarının ortancası, en az `KAMPANYA_LEARN_MIN_N` kayıt; yoksa 1 = artış varsayılmaz).
- *Kontroller*: maliyet yok, asgari fiyat altı, son 30 gün en düşük fiyat kuralı (yalnız site; kendi günlük fiyat kaydımızdan),
  zarar / marj alt sınırı, stok yok / kampanya bitmeden tükenme, internet satış hakkı ve CRM durum işareti.

**Zeki AI (LLM kapısından, `llm_for("kampanya")`):** kampanya metni (başlık / kısa açıklama / banner) ve sonuç özeti — ikisi de
pazarlama çekirdeğinin `marketing.guard` denetiminden geçer (kaynaksız rakam, kanıtsız üstünlük iddiası, teknoloji adı düşer);
CRM bayi kampanyasının türü kurallarla bulunamazsa kapalı küme `QueuedLlm.choose` (olasılık/marj eşiği ayarda). Rakam modelden
gelmez.
"""
from __future__ import annotations

import io
import json
import logging
import math
import re
import statistics
import threading
import time
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, Iterable, Optional
from zoneinfo import ZoneInfo

import sqlalchemy as sa

from semantic_bridge import kampanya_sources as src
from semantic_bridge.marketing import guard as mguard

log = logging.getLogger("semantic.kampanya")
TZ = ZoneInfo("Europe/Istanbul")
_md = sa.MetaData()

CAMPAIGNS = sa.Table(
    "semantic_kampanya_campaigns", _md,
    sa.Column("id", sa.String(20), primary_key=True),                 # KM-<yıl>-<sıra>
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("ad", sa.String(300), nullable=False),
    sa.Column("kanal", sa.String(20), nullable=False),                # site|pazar_yeri|bayi|fuar
    sa.Column("platform", sa.String(120)),
    sa.Column("baslangic", sa.String(10), nullable=False),
    sa.Column("bitis", sa.String(10), nullable=False),
    sa.Column("durum", sa.String(20), nullable=False),
    sa.Column("varsayilan_indirim", sa.Float),                        # toplu indirim oranı (0–1)
    sa.Column("kanal_kesinti", sa.Float),                             # oran (0–1); None = girilmedi
    sa.Column("beklenen_artis", sa.Float),                            # kat; None = öğrenimden ya da 1
    sa.Column("butce", sa.Float),
    sa.Column("hazirlayan", sa.String(120)),
    sa.Column("gonderen", sa.String(120)),
    sa.Column("onaylayan", sa.String(120)),
    sa.Column("onay_notu", sa.Text),
    sa.Column("crm_kampanya_id", sa.String(40)),
    sa.Column("kurulum_notu", sa.Text),
    sa.Column("kuruldu_by", sa.String(120)),
    sa.Column("kuruldu_at", sa.DateTime(timezone=True)),
    sa.Column("metin_json", sa.Text),
    sa.Column("notlar", sa.Text),
    sa.Column("sonuc_ozet", sa.Text),
    sa.Column("sonuc_at", sa.DateTime(timezone=True)),
    sa.Column("sonuc_bildirildi", sa.DateTime(timezone=True)),
    sa.Column("uyari_json", sa.Text),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
    sa.Column("submitted_at", sa.DateTime(timezone=True)),
    sa.Column("decided_at", sa.DateTime(timezone=True)),
)
ITEMS = sa.Table(
    "semantic_kampanya_items", _md,
    sa.Column("campaign_id", sa.String(20), primary_key=True),
    sa.Column("stok_kodu", sa.String(60), primary_key=True),
    sa.Column("sira", sa.Integer, nullable=False, default=0),
    sa.Column("ean", sa.String(20)),
    sa.Column("ad", sa.String(400)),
    sa.Column("yazar", sa.String(300)),
    sa.Column("liste_fiyati", sa.Float),                              # KDV dahil
    sa.Column("liste_elle", sa.Boolean, nullable=False, default=False),
    sa.Column("kampanya_fiyati", sa.Float),                           # KDV dahil
    sa.Column("fiyat_elle", sa.Boolean, nullable=False, default=False),  # True: fiyat girildi (indirim ondan), False: indirim girildi
    sa.Column("indirim_orani", sa.Float),
    sa.Column("kdv", sa.Float),
    sa.Column("stok", sa.Float),
    sa.Column("gunluk_hiz", sa.Float),
    sa.Column("tukenme_tahmini", sa.String(10)),
    sa.Column("birim_maliyet", sa.Float),
    sa.Column("maliyet_kaynak", sa.String(120)),
    sa.Column("net_once", sa.Float),
    sa.Column("net_sonra", sa.Float),
    sa.Column("telif_once", sa.Float),
    sa.Column("telif_sonra", sa.Float),
    sa.Column("telif_etkisi", sa.Float),
    sa.Column("marj_once", sa.Float),
    sa.Column("marj_sonra", sa.Float),
    sa.Column("marj_orani_once", sa.Float),
    sa.Column("marj_orani_sonra", sa.Float),
    sa.Column("maliyet_eksik", sa.Boolean, nullable=False, default=False),
    sa.Column("kontroller_json", sa.Text),
    sa.Column("hesap_json", sa.Text),
    sa.Column("aday_gerekcesi", sa.Text),
    sa.Column("eklenme", sa.DateTime(timezone=True), nullable=False),
    sa.Column("hesaplandi", sa.DateTime(timezone=True)),
)
CALENDAR = sa.Table(
    "semantic_kampanya_calendar", _md,
    sa.Column("id", sa.String(20), primary_key=True),                 # KT-<yıl>-<sıra>
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("tur", sa.String(12), nullable=False),                  # platform|ozel_gun|fuar
    sa.Column("ad", sa.String(300), nullable=False),
    sa.Column("baslangic", sa.String(10), nullable=False),
    sa.Column("bitis", sa.String(10), nullable=False),
    sa.Column("platform", sa.String(120)),
    sa.Column("kaynak", sa.String(12), nullable=False),               # kullanici
    sa.Column("notlar", sa.Text),
    sa.Column("olusturan", sa.String(120)),
    sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
)
RESULTS = sa.Table(
    "semantic_kampanya_results", _md,
    sa.Column("campaign_id", sa.String(20), primary_key=True),
    sa.Column("gun", sa.String(10), primary_key=True),
    sa.Column("stok_kodu", sa.String(60), primary_key=True),
    sa.Column("donem", sa.String(10), nullable=False),                # once|kampanya|sonra
    sa.Column("adet", sa.Float, nullable=False),
    sa.Column("net_tutar", sa.Float, nullable=False),
    sa.Column("iade_adet", sa.Float, nullable=False),
    sa.Column("maliyet", sa.Float, nullable=False),
    sa.Column("maliyetli_tutar", sa.Float, nullable=False),
    sa.Column("kaynak", sa.String(20), nullable=False),               # logo | logo-kanal (cari süzgeçli)
    sa.Column("logo_kesim", sa.String(10)),
    sa.Column("asof", sa.DateTime(timezone=True), nullable=False),
)
LEARNINGS = sa.Table(
    "semantic_kampanya_learnings", _md,
    sa.Column("id", sa.String(20), primary_key=True),                 # KO-<yıl>-<sıra>
    sa.Column("tenant_id", sa.String(80), nullable=False),
    sa.Column("campaign_id", sa.String(20), nullable=False),
    sa.Column("kanal", sa.String(20)),
    sa.Column("ozet", sa.Text, nullable=False),
    sa.Column("tur", sa.String(40)),
    sa.Column("indirim_orani", sa.Float),
    sa.Column("satis_degisimi", sa.Float),                            # kampanya günlük adet ÷ önceki dönem günlük adet
    sa.Column("iade_degisimi", sa.Float),
    sa.Column("marj_degisimi", sa.Float),                             # puan
    sa.Column("yazan", sa.String(120)),
    sa.Column("tarih", sa.DateTime(timezone=True), nullable=False),
)
SNAPS = sa.Table(
    "semantic_kampanya_price_snapshots", _md,
    sa.Column("tarih", sa.String(10), primary_key=True),
    sa.Column("product_key", sa.String(40), primary_key=True),        # barkod (EAN-13)
    sa.Column("kaynak", sa.String(10), primary_key=True),             # tsoft
    sa.Column("stok_kodu", sa.String(60)),
    sa.Column("fiyat", sa.Float),                                     # KDV dahil liste
    sa.Column("indirimli", sa.Float),                                 # KDV dahil indirimli (varsa)
    sa.Column("asof", sa.DateTime(timezone=True), nullable=False),
)
BOOKS = sa.Table(
    "semantic_kampanya_books", _md,
    sa.Column("stok_kodu", sa.String(60), primary_key=True),
    sa.Column("ad", sa.String(400)),
    sa.Column("ean", sa.String(20)),
    sa.Column("crm_id", sa.String(40)),
    sa.Column("crm_tip", sa.Integer),
    sa.Column("yazar", sa.String(300)),
    sa.Column("yayin", sa.String(60)),
    sa.Column("tsoft", sa.Boolean),
    sa.Column("liste_crm", sa.Float),                                 # CRM KDV dahil fiyat
    sa.Column("liste_logo", sa.Float),                                # Logo bugün geçerli satış fiyatı, KDV dahile çevrilmiş
    sa.Column("kdv", sa.Float),                                       # Logo ITEMS.SELLVAT (%)
    sa.Column("stok", sa.Float, nullable=False, default=0.0),
    sa.Column("adet_son", sa.Float, nullable=False, default=0.0),     # son hız penceresi, net (iade düşülmüş)
    sa.Column("adet_onceki", sa.Float, nullable=False, default=0.0),  # ondan önceki eşit pencere
    sa.Column("adet_12", sa.Float, nullable=False, default=0.0),
    sa.Column("tutar_12", sa.Float, nullable=False, default=0.0),
    sa.Column("gunluk_hiz", sa.Float, nullable=False, default=0.0),
    sa.Column("ciro_yil", sa.Float),                                  # güncel yıl, satış satırları Σ LINENET
    sa.Column("maliyet_yil", sa.Float),                               # Σ AMOUNT × OUTCOST
    sa.Column("maliyetli_ciro", sa.Float),
    sa.Column("maliyetli_adet", sa.Float),
    sa.Column("maliyetsiz_satir", sa.Integer),
    sa.Column("satir_yil", sa.Integer),
    sa.Column("marj_yili", sa.Integer),
    sa.Column("asgari_fiyat", sa.Float),
    sa.Column("sozlesme_json", sa.Text),
    sa.Column("hak", sa.String(16)),                                  # SEO CRM özeti: var|eksik|yok|incele|koruma_disi
    sa.Column("durum_bayragi", sa.String(16)),
    sa.Column("sezon_json", sa.Text),                                 # bağlı özel gün anahtarları
    sa.Column("in_logo", sa.Boolean, nullable=False, default=False),
    sa.Column("in_crm", sa.Boolean, nullable=False, default=False),
    sa.Column("asof", sa.DateTime(timezone=True), nullable=False),
)
CRM_TYPES = sa.Table(
    "semantic_kampanya_crm_types", _md,
    sa.Column("crm_id", sa.String(40), primary_key=True),
    sa.Column("tur", sa.String(20), nullable=False),                  # iskonto|vade|hediye|stant|diger|belirsiz
    sa.Column("yontem", sa.String(12), nullable=False),               # kural|model
    sa.Column("olasilik", sa.Float),
    sa.Column("asof", sa.DateTime(timezone=True), nullable=False),
)
META = sa.Table(
    "semantic_kampanya_meta", _md,
    sa.Column("key", sa.String(60), primary_key=True),
    sa.Column("value_json", sa.Text, nullable=False),
    sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
)

KANALLAR = {"site": "Site (timas.com.tr)", "pazar_yeri": "Pazar yeri", "bayi": "Bayi", "fuar": "Fuar"}
DURUMLAR = {"taslak": "Taslak", "onay_bekliyor": "Onay bekliyor", "onaylandi": "Onaylandı", "yurutuluyor": "Yürütülüyor",
            "bitti": "Bitti", "iptal": "İptal"}
DUZENLENIR = ("taslak",)
TAKVIM_TURLERI = {"platform": "Platform dönemi", "ozel_gun": "Özel gün", "fuar": "Fuar"}
KURALLAR = {"stok": "Stok fazlası", "dusus": "Satışı yavaşlamış", "sezon": "Sezona bağlı", "hak": "İnternet satış hakkı uygun",
            "maliyet": "Maliyeti bilinen", "marj": "Marjı indirimi kaldırıyor"}
#: Sinyal kuralları: en az biri tutmalı. Koşul kuralları: seçildiyse hepsi tutmalı.
SINYAL = ("stok", "dusus", "sezon")
KOSUL = ("hak", "maliyet", "marj")
VARSAYILAN_KURALLAR = ("stok", "dusus", "sezon", "hak")
METIN_TURLERI = {"baslik": ("Kampanya başlığı", 60), "aciklama": ("Kısa açıklama", 160), "banner": ("Banner metni", 40)}
CRM_TURLERI = {"iskonto": "İskonto", "vade": "Vade", "hediye": "Hediye", "stant": "Stant", "diger": "Diğer",
               "belirsiz": "Belirsiz"}
SEVIYE = ("kirmizi", "sari", "bilgi")
PAGE_SIZE = 50

DEFAULTS: dict[str, str] = {
    "KAMPANYA_LIST_PRICE_SOURCE": "crm",
    "KAMPANYA_COST_SOURCE": "m9+logo",
    "KAMPANYA_MARJ_MIN_PCT": "",
    "KAMPANYA_KANAL_CARI": "",
    "KAMPANYA_ADAY_HIZ_AY": "3",
    "KAMPANYA_ADAY_STOK_AY": "12",
    "KAMPANYA_ADAY_DUSUS_PCT": "30",
    "KAMPANYA_ADAY_MARJ_MIN_PCT": "",
    "KAMPANYA_SEZON_ONCESI_GUN": "21",
    "KAMPANYA_SONRA_GUN": "14",
    "KAMPANYA_FIYAT_GUN": "30",
    "KAMPANYA_TAKVIM_GUN": "60",
    "KAMPANYA_LEARN_MIN_N": "3",
    "KAMPANYA_ONAY_ALICILARI": "",
    "KAMPANYA_DEPO_ALICILARI": "",
    "KAMPANYA_LLM_BUDGET_SEC": "900",
    "KAMPANYA_LLM_MIN_PROB": "0.70",
    "KAMPANYA_LLM_MIN_MARGIN": "0.30",
}

_ready: set[int] = set()
_lock = threading.Lock()
_cost_provider: Optional[Callable[[list[str]], dict[str, dict[str, Any]]]] = None


class KampanyaError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def register_cost_provider(fn: Optional[Callable[[list[str]], dict[str, dict[str, Any]]]]) -> None:
    """M9 birim maliyet sağlayıcısı (`pricing.cost_provider.Provider.unit_costs`). `None` bağı kaldırır."""
    global _cost_provider
    _cost_provider = fn


def cost_provider_connected() -> bool:
    return _cost_provider is not None


def ensure(engine: sa.engine.Engine) -> None:
    key = id(engine)
    with _lock:
        if key in _ready:
            return
        from semantic_layer.store import schema_stamp
        schema_stamp.create_all(_md, engine)
        _ready.add(key)


# ------------------------------------------------------------------------------------------ küçük yardımcılar


def _now() -> datetime:
    return datetime.now(timezone.utc)


def today() -> date:
    return datetime.now(TZ).date()


def _iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    return (v if v.tzinfo else v.replace(tzinfo=timezone.utc)).isoformat()


def _j(v: Optional[str], default: Any) -> Any:
    try:
        return json.loads(v) if v else default
    except ValueError:
        return default


def _dump(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str)


def _text(v: Any, limit: int) -> Optional[str]:
    s = " ".join(str(v or "").split())
    return s[:limit] or None


def _longtext(v: Any, limit: int) -> Optional[str]:
    s = str(v or "").replace("\r\n", "\n").strip()
    return s[:limit] or None


def _r2(v: Optional[float]) -> Optional[float]:
    return None if v is None else round(float(v), 2)


def _r4(v: Optional[float]) -> Optional[float]:
    return None if v is None else round(float(v), 4)


def fold(s: Any) -> str:
    t = str(s or "").translate(str.maketrans("İIıŞşĞğÜüÖöÇç", "iiissgguuoocc")).lower()
    return re.sub(r"\s+", " ", t).strip()


def _num(v: Any, label: str, *, allow_none: bool = False, minimum: Optional[float] = 0.0,
         maximum: Optional[float] = None) -> Optional[float]:
    if v is None or v == "":
        if allow_none:
            return None
        raise KampanyaError(f"{label} boş olamaz.")
    if isinstance(v, bool):
        raise KampanyaError(f"{label} sayı olmalı.")
    try:
        n = float(str(v).replace(".", "").replace(",", ".")) if isinstance(v, str) and "," in v else float(v)
    except (TypeError, ValueError):
        raise KampanyaError(f"{label} sayı olmalı.") from None
    if math.isnan(n) or math.isinf(n):
        raise KampanyaError(f"{label} sayı olmalı.")
    if minimum is not None and n < minimum:
        raise KampanyaError(f"{label} {minimum:g} değerinden küçük olamaz.")
    if maximum is not None and n > maximum:
        raise KampanyaError(f"{label} {maximum:g} değerinden büyük olamaz.")
    return n


def ratio(v: Any, label: str, *, allow_none: bool = True) -> Optional[float]:
    """İndirim/kesinti oranı: 1'in altı oran, 1 ve üstü yüzde (30 → 0,30; 1 → 0,01). %95'ten büyük olamaz."""
    n = _num(v, label, allow_none=allow_none, minimum=0.0, maximum=95.0)
    if n is None:
        return None
    r = n / 100.0 if n >= 1 else n
    if r > 0.95:
        raise KampanyaError(f"{label} %95'ten büyük olamaz.")
    return round(r, 6)


def _day(v: Any, label: str) -> str:
    s = str(v or "").strip()[:10]
    try:
        return date.fromisoformat(s).isoformat()
    except ValueError:
        raise KampanyaError(f"{label} YYYY-AA-GG biçiminde bir tarih olmalı.") from None


def _d(s: Optional[str]) -> Optional[date]:
    try:
        return date.fromisoformat(str(s)[:10]) if s else None
    except ValueError:
        return None


def _money(v: Optional[float]) -> str:
    if v is None:
        return "—"
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{s} ₺"


def _pct(v: Optional[float], digits: int = 0) -> str:
    if v is None:
        return "—"
    return f"%{v * 100:.{digits}f}".replace(".", ",")


def _day_tr(v: Optional[str]) -> str:
    d = _d(v)
    return d.strftime("%d.%m.%Y") if d else "—"


def _sentence(s: str) -> str:
    """İlk harfi büyüt (Türkçe i → İ); geri kalanına dokunma."""
    if not s:
        return s
    first = "İ" if s[0] == "i" else s[0].upper()
    return first + s[1:]


def _recipients(raw: Any) -> list[str]:
    return [x.strip() for x in re.split(r"[,;\s]+", str(raw or "")) if re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", x.strip())]


# ------------------------------------------------------------------------------------------ ayarlar


def settings_from(conf: Callable[[str, str], str]) -> dict[str, Any]:
    def g(key: str) -> str:
        v = conf(key, DEFAULTS[key])
        return DEFAULTS[key] if v is None else str(v)

    def gi(key: str, lo: int, hi: int) -> int:
        try:
            return max(lo, min(hi, int(float(g(key)))))
        except ValueError:
            return int(DEFAULTS[key])

    def gf(key: str) -> Optional[float]:
        raw = g(key).strip()
        if raw == "":
            return None
        try:
            return float(raw.replace(",", "."))
        except ValueError:
            return float(DEFAULTS[key]) if DEFAULTS[key] else None

    price = g("KAMPANYA_LIST_PRICE_SOURCE").strip().lower()
    cost = g("KAMPANYA_COST_SOURCE").strip().lower()
    return {
        "listPriceSource": price if price in ("crm", "logo") else "crm",
        "costSource": cost if cost in ("m9", "m9+logo", "logo", "yok") else "m9+logo",
        "marginMinPct": gf("KAMPANYA_MARJ_MIN_PCT"),
        "kanalCari": parse_kanal_cari(g("KAMPANYA_KANAL_CARI")),
        "hizAy": gi("KAMPANYA_ADAY_HIZ_AY", 1, 24),
        "stokAy": gf("KAMPANYA_ADAY_STOK_AY") or 12.0,
        "dususPct": gf("KAMPANYA_ADAY_DUSUS_PCT") or 30.0,
        "adayMarjMinPct": gf("KAMPANYA_ADAY_MARJ_MIN_PCT"),
        "sezonOncesiGun": gi("KAMPANYA_SEZON_ONCESI_GUN", 0, 180),
        "sonraGun": gi("KAMPANYA_SONRA_GUN", 0, 120),
        "fiyatGun": gi("KAMPANYA_FIYAT_GUN", 1, 365),
        "takvimGun": gi("KAMPANYA_TAKVIM_GUN", 7, 366),
        "learnMinN": gi("KAMPANYA_LEARN_MIN_N", 1, 1000),
        "onayAlicilari": _recipients(g("KAMPANYA_ONAY_ALICILARI")),
        "depoAlicilari": _recipients(g("KAMPANYA_DEPO_ALICILARI")),
        "llmBudgetSec": gi("KAMPANYA_LLM_BUDGET_SEC", 0, 6 * 3600),
        "llmMinProb": gf("KAMPANYA_LLM_MIN_PROB") or 0.70,
        "llmMinMargin": gf("KAMPANYA_LLM_MIN_MARGIN") or 0.0,
    }


def parse_kanal_cari(raw: str) -> dict[str, list[str]]:
    """«site:120.01.001,120.01.002;pazar_yeri:120.05.010» → {kanal: [cari kodu]}; bilinmeyen kanal atlanır."""
    out: dict[str, list[str]] = {}
    for part in str(raw or "").split(";"):
        k, _, v = part.partition(":")
        k = k.strip()
        if k in KANALLAR:
            codes = [c.strip() for c in v.split(",") if src.code_ok(c.strip())]
            if codes:
                out[k] = codes
    return out


# ------------------------------------------------------------------------------------------ meta


def meta_get(engine: sa.engine.Engine, key: str) -> dict[str, Any]:
    with engine.connect() as c:
        row = c.execute(sa.select(META).where(META.c.key == key)).first()
    if not row:
        return {}
    return {**_j(row.value_json, {}), "_at": _iso(row.updated_at)}


def meta_set(engine: sa.engine.Engine, key: str, value: dict[str, Any]) -> None:
    now = _now()
    with engine.begin() as c:
        if c.execute(sa.select(META.c.key).where(META.c.key == key)).first():
            c.execute(META.update().where(META.c.key == key).values(value_json=_dump(value), updated_at=now))
        else:
            c.execute(META.insert().values(key=key, value_json=_dump(value), updated_at=now))


def data_end(engine: sa.engine.Engine) -> Optional[date]:
    return _d(meta_get(engine, "data_end").get("date"))


def _next_id(c: Any, table: sa.Table, prefix: str) -> str:
    ids = [r[0] for r in c.execute(sa.select(table.c.id).where(table.c.id.like(prefix + "%"))).all()]
    n = max([int(i[len(prefix):]) for i in ids if i[len(prefix):].isdigit()] or [0]) + 1
    return f"{prefix}{n:04d}"


# ------------------------------------------------------------------------------------------ kitaplar, maliyet, fiyat kaydı


def book_dict(b: Any) -> dict[str, Any]:
    return {"stok": b.stok_kodu, "ad": b.ad, "ean": b.ean, "crmId": b.crm_id, "crmTip": b.crm_tip, "yazar": b.yazar,
            "yayin": b.yayin, "tsoft": b.tsoft, "listeCrm": b.liste_crm, "listeLogo": b.liste_logo, "kdv": b.kdv, "stokAdet": b.stok,
            "adetSon": b.adet_son, "adetOnceki": b.adet_onceki, "adet12": b.adet_12, "tutar12": b.tutar_12, "gunlukHiz": b.gunluk_hiz,
            "ciroYil": b.ciro_yil, "maliyetYil": b.maliyet_yil, "maliyetliCiro": b.maliyetli_ciro, "maliyetliAdet": b.maliyetli_adet,
            "maliyetsizSatir": b.maliyetsiz_satir, "satirYil": b.satir_yil, "marjYili": b.marj_yili, "asgariFiyat": b.asgari_fiyat,
            "sozlesmeler": _j(b.sozlesme_json, []), "hak": b.hak, "durumBayragi": b.durum_bayragi, "sezonlar": _j(b.sezon_json, []),
            "inLogo": b.in_logo, "inCrm": b.in_crm}


def books_by_code(engine: sa.engine.Engine, codes: Iterable[str]) -> dict[str, dict[str, Any]]:
    codes = [c for c in dict.fromkeys(codes) if c]
    out: dict[str, dict[str, Any]] = {}
    with engine.connect() as c:
        for i in range(0, len(codes), 900):
            for b in c.execute(sa.select(BOOKS).where(BOOKS.c.stok_kodu.in_(codes[i:i + 900]))).all():
                out[b.stok_kodu] = book_dict(b)
    return out


def list_price(book: dict[str, Any], st: dict[str, Any]) -> tuple[Optional[float], Optional[str]]:
    """(KDV dahil liste fiyatı, kaynak). Seçilen kaynak boşsa diğeri kullanılır ve kaynak adı öyle yazılır."""
    crm, logo = book.get("listeCrm"), book.get("listeLogo")
    order = [("crm", crm), ("logo", logo)] if st["listPriceSource"] == "crm" else [("logo", logo), ("crm", crm)]
    for name, v in order:
        if v and v > 0:
            return round(float(v), 2), name
    return None, None


def logo_unit_cost(book: dict[str, Any]) -> Optional[float]:
    qty, cost = book.get("maliyetliAdet") or 0.0, book.get("maliyetYil")
    if qty > 0 and cost and cost > 0:
        return cost / qty
    return None


def costs_for(st: dict[str, Any], books: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Kod → {"maliyet": float|None, "kaynak": okunur metin}. Kaynak sırası `KAMPANYA_COST_SOURCE`."""
    codes = list(books)
    out: dict[str, dict[str, Any]] = {c: {"maliyet": None, "kaynak": "Maliyet bilinmiyor"} for c in codes}
    mode = st["costSource"]
    if mode == "yok" or not codes:
        return out
    if mode in ("m9", "m9+logo") and _cost_provider is not None:
        try:
            got = _cost_provider(codes) or {}
        except Exception as e:  # noqa: BLE001 — sağlayıcı düşerse Logo'ya (izinliyse) geçilir
            log.warning("kampanya: M9 birim maliyeti okunamadı: %s", e)
            got = {}
        from semantic_bridge.pricing import cost_provider as cp

        for code, hit in got.items():
            if code in out and hit and hit.get("maliyet") is not None:
                out[code] = {"maliyet": float(hit["maliyet"]), "kaynak": cp.label(hit)}
    if mode in ("m9+logo", "logo"):
        for code, b in books.items():
            if out[code]["maliyet"] is None:
                u = logo_unit_cost(b)
                if u is not None:
                    out[code] = {"maliyet": u, "kaynak": f"Logo gerçekleşen maliyet ({b.get('marjYili') or ''})".replace(" ()", "")}
    return out


def price_floor(engine: sa.engine.Engine, eans: Iterable[str], start: date, days: int) -> dict[str, dict[str, Any]]:
    """Barkod → son `days` gündeki (başlangıçtan önce) en düşük site fiyatı, kayıtlı gün sayısı, son fiyat."""
    eans = [e for e in dict.fromkeys(eans) if e]
    a, b = (start - timedelta(days=days)).isoformat(), start.isoformat()
    out: dict[str, dict[str, Any]] = {}
    if not eans:
        return out
    with engine.connect() as c:
        for i in range(0, len(eans), 900):
            rows = c.execute(sa.select(SNAPS.c.product_key, SNAPS.c.tarih, SNAPS.c.fiyat, SNAPS.c.indirimli).where(
                SNAPS.c.product_key.in_(eans[i:i + 900]), SNAPS.c.kaynak == "tsoft", SNAPS.c.tarih >= a, SNAPS.c.tarih < b)
                .order_by(SNAPS.c.tarih)).all()
            for r in rows:
                eff = effective_price(r.fiyat, r.indirimli)
                if eff is None:
                    continue
                cur = out.setdefault(r.product_key, {"min": eff, "gun": 0, "son": eff, "sonGun": r.tarih})
                cur["min"] = min(cur["min"], eff)
                cur["gun"] += 1
                cur["son"], cur["sonGun"] = eff, r.tarih
    return out


def effective_price(price: Optional[float], sale: Optional[float]) -> Optional[float]:
    if sale is not None and sale > 0 and (price is None or sale < price):
        return round(float(sale), 2)
    return round(float(price), 2) if price is not None and price > 0 else None


# ------------------------------------------------------------------------------------------ simülasyon


def telif_units(contracts: list[dict[str, Any]], liste_excl: Optional[float], net: Optional[float],
                retail_excl: Optional[float]) -> tuple[Optional[float], list[str]]:
    """Satıştan ödenen sözleşmelerin birim telifi (toplam) ve esas notları. Oranı eksik satış sözleşmesi varsa None."""
    total, notes, unknown = 0.0, [], False
    for k in contracts:
        if k.get("odeme") not in src.SATIS_TELIF:
            continue
        rate = k.get("oran")
        if rate is None or rate <= 0:
            unknown = True
            notes.append("Satıştan telifli sözleşmede oran yok")
            continue
        mode = telif_esasi(k)
        base = {"liste": liste_excl, "net": net, "perakende": retail_excl}[mode]
        if base is None:
            unknown = True
            continue
        total += base * rate / 100.0
        notes.append(mode)
    return (None if unknown else round(total, 4)), notes


def telif_esasi(k: dict[str, Any]) -> str:
    h, t = k.get("hesaplama"), k.get("tur")
    if h == 1:
        return "net"
    if h == 2:
        return "liste"
    if h == 3:
        return "perakende"
    if t == 1:
        return "liste"
    return "net"


def uplift_for(engine: sa.engine.Engine, tenant: str, kanal: str, st: dict[str, Any]) -> tuple[float, str]:
    """Beklenen satış artışı (kat) ve kaynağı: aynı kanaldaki öğrenim kayıtlarının ortancası ya da 1."""
    with engine.connect() as c:
        vals = [r[0] for r in c.execute(sa.select(LEARNINGS.c.satis_degisimi).where(
            LEARNINGS.c.tenant_id == tenant, LEARNINGS.c.kanal == kanal, LEARNINGS.c.satis_degisimi.isnot(None))).all()]
    vals = [v for v in vals if v and v > 0]
    if len(vals) >= st["learnMinN"]:
        return round(statistics.median(vals), 3), f"aynı kanaldaki {len(vals)} kampanyanın öğrenim kaydı (ortanca)"
    return 1.0, "öğrenim kaydı yetersiz; artış varsayılmadı"


def simulate_item(book: dict[str, Any], item: dict[str, Any], camp: dict[str, Any], st: dict[str, Any],
                  cost: Optional[dict[str, Any]], floor: Optional[dict[str, Any]], ref: date,
                  end: Optional[date], uplift: tuple[float, str]) -> dict[str, Any]:
    """Bir kitabın kampanya hesabı. `item`: liste (elle, KDV dahil) | kampanyaFiyati | indirim. Çıktı ITEMS alanları."""
    liste, liste_src = (item.get("liste"), "elle") if item.get("liste") else list_price(book, st)
    kdv = float(book.get("kdv") or 0.0)
    f = 1.0 + kdv / 100.0
    kf, ind = item.get("kampanyaFiyati"), item.get("indirim")
    if kf is not None and kf > 0:
        ind = (1 - kf / liste) if liste else None
    elif ind is not None and liste:
        kf = round(liste * (1 - ind), 2)
    else:
        kf = None
    kes_raw = camp.get("kanalKesinti")
    kes = float(kes_raw or 0.0)
    net_once = liste / f * (1 - kes) if liste else None
    net_sonra = kf / f * (1 - kes) if kf else None
    contracts = book.get("sozlesmeler") or []
    t_once, notes_once = telif_units(contracts, liste / f if liste else None, net_once, liste / f if liste else None)
    t_sonra, _ = telif_units(contracts, liste / f if liste else None, net_sonra, kf / f if kf else None)
    unit = cost.get("maliyet") if cost else None

    def marj(net: Optional[float], tel: Optional[float]) -> tuple[Optional[float], Optional[float]]:
        if net is None or unit is None or tel is None:
            return None, None
        m = net - unit - tel
        return m, (m / net if net else None)

    m_once, o_once = marj(net_once, t_once)
    m_sonra, o_sonra = marj(net_sonra, t_sonra)

    kontrol: list[dict[str, str]] = []

    def add(kod: str, seviye: str, mesaj: str) -> None:
        kontrol.append({"kod": kod, "seviye": seviye, "mesaj": mesaj})

    kanal = camp.get("kanal")
    if not liste:
        add("fiyat", "kirmizi", "Liste fiyatı yok (CRM ve Logo); kampanya fiyatı hesaplanamaz.")
    elif kf is None:
        add("fiyat", "sari", "Kampanya fiyatı ya da indirim girilmedi.")
    if unit is None:
        add("maliyet", "kirmizi", "Birim maliyet bilinmiyor: marj hesaplanamaz.")
    if (book.get("maliyetsizSatir") or 0) > 0 and unit is not None:
        add("maliyet-logo", "bilgi", f"Logo'da bu yıl {int(book['maliyetsizSatir'])} satış satırında maliyet girilmemiş.")
    floor_price = book.get("asgariFiyat")
    if kanal == "bayi":
        if floor_price:
            add("asgari", "bilgi", f"Sözleşmedeki asgari perakende fiyat {_money(floor_price)}; bayi fiyatına uygulanmaz, bayinin "
                                   "perakende fiyatı ayrıca kontrol edilmeli.")
    elif floor_price:
        if kf is not None and kf < floor_price - 0.005:
            add("asgari", "kirmizi", f"Kampanya fiyatı {_money(kf)}, sözleşmedeki asgari perakende fiyat {_money(floor_price)}.")
    elif contracts:
        add("asgari", "bilgi", "Sözleşmede asgari fiyat yok.")
    else:
        add("asgari", "bilgi", "Yürürlükte Telif Alış sözleşmesi bulunamadı.")
    if kanal == "site":
        need = st["fiyatGun"]
        if not floor or not floor.get("gun"):
            add("fiyat-30", "sari", f"Sitenin günlük fiyat kaydı yok; son {need} gün en düşük fiyat kuralı denetlenemedi.")
        else:
            low = floor["min"]
            if floor["gun"] < need:
                add("fiyat-30", "sari", f"Fiyat kaydı {floor['gun']} günlük; kural {need} gün ister.")
            if kf is not None and kf >= low - 0.005:
                add("fiyat-30", "kirmizi", f"Kampanya fiyatı son {need} günün en düşük fiyatının ({_money(low)}) altında değil; "
                                           "indirim olarak duyurulamaz.")
            elif liste and liste > low + 0.005 and kf is not None:
                add("fiyat-30", "sari", f"İndirim öncesi fiyat olarak son {need} günün en düşüğü {_money(low)} gösterilmeli "
                                        f"(liste {_money(liste)} değil); gerçek indirim {_pct(1 - kf / low)}.")
    else:
        add("fiyat-30", "bilgi", "Son 30 gün fiyat kaydı yalnız site için tutuluyor; platformun kuralı platform panelinde.")
    if m_sonra is not None:
        if m_sonra < 0:
            add("marj", "kirmizi", f"Kampanya fiyatında birim zarar {_money(m_sonra)}.")
        elif st["marginMinPct"] is not None and o_sonra is not None and o_sonra < st["marginMinPct"] / 100.0:
            add("marj", "sari", f"Kampanyalı marj {_pct(o_sonra, 1)}, alt sınır %{st['marginMinPct']:g}.")
    if "liste" in notes_once:
        add("telif", "bilgi", "Telif kapak (liste) fiyatından: indirim telifi düşürmez.")
    if t_once is None and any(k.get("odeme") in src.SATIS_TELIF for k in contracts):
        add("telif", "sari", "Satıştan telifli sözleşmede oran eksik; telif ve marj hesaplanamaz.")
    if kanal != "site" and kes_raw is None:
        add("kesinti", "bilgi", "Kanal kesintisi (komisyon / bayi iskontosu) girilmedi; 0 sayıldı.")
    hak, flag = book.get("hak"), book.get("durumBayragi")
    if flag:
        add("hak", "kirmizi", f"CRM'de kitabın durumu: {flag.replace('_', ' ')}.")
    if hak in ("yok",):
        add("hak", "kirmizi", "Sözleşmede internetten satış (umuma iletim) hakkı yok.")
    elif hak in ("eksik", "incele"):
        add("hak", "sari", "İnternet satış hakkı sözleşmede eksik ya da incelenmeli.")

    # Tükenme: veri sonundan kampanya başına normal hız, kampanyada hız × artış.
    stok = float(book.get("stokAdet") or 0.0)
    hiz = float(book.get("gunlukHiz") or 0.0)
    bas, bit = _d(camp.get("baslangic")), _d(camp.get("bitis"))
    artis, artis_src = (camp["beklenenArtis"], "kampanyada girildi") if camp.get("beklenenArtis") else uplift
    tuk: Optional[str] = None
    if bas and bit:
        start = max(bas, ref)
        base = end or ref
        at_start = stok - hiz * max(0, (start - base).days)
        rate = hiz * artis
        if stok <= 0:
            add("stok", "kirmizi", "Logo'da stok yok.")
            tuk = start.isoformat()
        elif rate > 0:
            d = start + timedelta(days=max(0.0, at_start) / rate)
            tuk = d.isoformat()
            if d <= bit:
                add("stok", "sari", f"Stok kampanya bitmeden tükenebilir (tahmin {_day_tr(tuk)}; günlük hız {hiz:.1f} × {artis:g}).")
    hesap = {"listeKaynak": liste_src, "kdv": kdv, "kanalKesinti": kes, "kanalKesintiGirildi": kes_raw is not None,
             "telifEsaslari": [n for n in notes_once if n in ("liste", "net", "perakende")], "artis": artis, "artisKaynak": artis_src,
             "sonFiyat": floor}
    return {
        "liste_fiyati": _r2(liste), "kampanya_fiyati": _r2(kf), "indirim_orani": _r4(ind), "kdv": kdv, "stok": stok,
        "gunluk_hiz": _r4(hiz), "tukenme_tahmini": tuk, "birim_maliyet": _r4(unit),
        "maliyet_kaynak": (cost or {}).get("kaynak") or "Maliyet bilinmiyor",
        "net_once": _r4(net_once), "net_sonra": _r4(net_sonra), "telif_once": _r4(t_once), "telif_sonra": _r4(t_sonra),
        "telif_etkisi": _r4(t_sonra - t_once) if t_sonra is not None and t_once is not None else None,
        "marj_once": _r4(m_once), "marj_sonra": _r4(m_sonra), "marj_orani_once": _r4(o_once), "marj_orani_sonra": _r4(o_sonra),
        "maliyet_eksik": unit is None, "kontroller_json": _dump(kontrol), "hesap_json": _dump(hesap),
    }


# ------------------------------------------------------------------------------------------ kampanya kaydı


def _item_dict(r: Any) -> dict[str, Any]:
    kontrol = _j(r.kontroller_json, [])
    return {"stok": r.stok_kodu, "sira": r.sira, "ean": r.ean, "ad": r.ad, "yazar": r.yazar, "liste": r.liste_fiyati,
            "listeElle": bool(r.liste_elle), "kampanyaFiyati": r.kampanya_fiyati, "indirim": r.indirim_orani, "kdv": r.kdv,
            "stokAdet": r.stok, "gunlukHiz": r.gunluk_hiz, "tukenme": r.tukenme_tahmini, "birimMaliyet": r.birim_maliyet,
            "maliyetKaynak": r.maliyet_kaynak, "netOnce": r.net_once, "netSonra": r.net_sonra, "telifOnce": r.telif_once,
            "telifSonra": r.telif_sonra, "telifEtkisi": r.telif_etkisi, "marjOnce": r.marj_once, "marjSonra": r.marj_sonra,
            "marjOraniOnce": r.marj_orani_once, "marjOraniSonra": r.marj_orani_sonra, "maliyetEksik": bool(r.maliyet_eksik),
            "kontroller": kontrol, "kirmizi": sum(1 for k in kontrol if k.get("seviye") == "kirmizi"),
            "sari": sum(1 for k in kontrol if k.get("seviye") == "sari"), "hesap": _j(r.hesap_json, {}),
            "adayGerekcesi": r.aday_gerekcesi, "hesaplandi": _iso(r.hesaplandi)}


def totals(items: list[dict[str, Any]]) -> dict[str, Any]:
    """Kampanyanın özeti: kitap sayısı, ortalama indirim, liste ağırlıklı marj oranları, kontrol sayıları."""
    inds = [i["indirim"] for i in items if i.get("indirim") is not None]
    known = [i for i in items if i.get("marjSonra") is not None and i.get("netSonra")]

    def weighted(key_m: str, key_n: str) -> Optional[float]:
        num = sum(i[key_m] for i in known if i.get(key_m) is not None and i.get(key_n))
        den = sum(i[key_n] for i in known if i.get(key_m) is not None and i.get(key_n))
        return round(num / den, 4) if den else None

    return {"kitap": len(items), "ortalamaIndirim": round(sum(inds) / len(inds), 4) if inds else None,
            "marjOraniOnce": weighted("marjOnce", "netOnce"), "marjOraniSonra": weighted("marjSonra", "netSonra"),
            "marjBilinen": len(known), "maliyetEksik": sum(1 for i in items if i.get("maliyetEksik")),
            "kirmizi": sum(i["kirmizi"] for i in items), "sari": sum(i["sari"] for i in items),
            "kirmiziKitap": sum(1 for i in items if i["kirmizi"]),
            "stokRiski": sum(1 for i in items if any(k.get("kod") == "stok" for k in i["kontroller"]))}


def _camp_dict(r: Any, items: Optional[list[dict[str, Any]]] = None, summary: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    out = {"id": r.id, "ad": r.ad, "kanal": r.kanal, "kanalAdi": KANALLAR.get(r.kanal, r.kanal), "platform": r.platform,
           "baslangic": r.baslangic, "bitis": r.bitis, "durum": r.durum, "durumAdi": DURUMLAR.get(r.durum, r.durum),
           "varsayilanIndirim": r.varsayilan_indirim, "kanalKesinti": r.kanal_kesinti, "beklenenArtis": r.beklenen_artis,
           "butce": r.butce, "hazirlayan": r.hazirlayan, "gonderen": r.gonderen, "onaylayan": r.onaylayan, "onayNotu": r.onay_notu,
           "crmKampanyaId": r.crm_kampanya_id, "kurulumNotu": r.kurulum_notu, "kurulduBy": r.kuruldu_by,
           "kurulduAt": _iso(r.kuruldu_at), "metin": _j(r.metin_json, {}), "notlar": r.notlar, "sonucOzet": r.sonuc_ozet,
           "sonucAt": _iso(r.sonuc_at), "uyarilar": _j(r.uyari_json, {}).get("items", []), "createdAt": _iso(r.created_at),
           "updatedAt": _iso(r.updated_at), "submittedAt": _iso(r.submitted_at), "decidedAt": _iso(r.decided_at),
           "crmIslenecek": r.kanal == "bayi" and not r.crm_kampanya_id and r.durum in ("onaylandi", "yurutuluyor")}
    if items is not None:
        out["kitaplar"] = items
        out["ozet"] = totals(items)
    elif summary is not None:
        out["ozet"] = summary
    return out


def _row(c: Any, tenant: str, cid: str, lock: bool = False) -> Any:
    q = sa.select(CAMPAIGNS).where(CAMPAIGNS.c.id == cid, CAMPAIGNS.c.tenant_id == tenant)
    r = c.execute(q.with_for_update() if lock and c.dialect.name != "sqlite" else q).first()
    if not r:
        raise KampanyaError("Kampanya bulunamadı.", 404)
    return r


def _items(c: Any, cid: str) -> list[dict[str, Any]]:
    rows = c.execute(sa.select(ITEMS).where(ITEMS.c.campaign_id == cid).order_by(ITEMS.c.sira, ITEMS.c.stok_kodu)).all()
    return [_item_dict(r) for r in rows]


def get_campaign(engine: sa.engine.Engine, tenant: str, cid: str) -> dict[str, Any]:
    with engine.connect() as c:
        r = _row(c, tenant, cid)
        return _camp_dict(r, _items(c, cid))


def list_campaigns(engine: sa.engine.Engine, tenant: str, *, durum: str = "", kanal: str = "", q: str = "",
                   page: int = 0) -> dict[str, Any]:
    cond = [CAMPAIGNS.c.tenant_id == tenant]
    if durum:
        cond.append(CAMPAIGNS.c.durum.in_([d for d in durum.split(",") if d in DURUMLAR] or ["-"]))
    if kanal:
        cond.append(CAMPAIGNS.c.kanal == kanal)
    if q.strip():
        like = f"%{q.strip()}%"
        cond.append(sa.or_(CAMPAIGNS.c.ad.ilike(like), CAMPAIGNS.c.platform.ilike(like), CAMPAIGNS.c.id.ilike(like)))
    page = max(0, int(page))
    with engine.connect() as c:
        total = c.execute(sa.select(sa.func.count()).select_from(CAMPAIGNS).where(*cond)).scalar() or 0
        rows = c.execute(sa.select(CAMPAIGNS).where(*cond).order_by(CAMPAIGNS.c.baslangic.desc(), CAMPAIGNS.c.id.desc())
                         .offset(page * PAGE_SIZE).limit(PAGE_SIZE)).all()
        ids = [r.id for r in rows]
        items: dict[str, list[dict[str, Any]]] = {i: [] for i in ids}
        for i in range(0, len(ids), 500):
            for it in c.execute(sa.select(ITEMS).where(ITEMS.c.campaign_id.in_(ids[i:i + 500]))).all():
                items[it.campaign_id].append(_item_dict(it))
        counts = {k: v for k, v in c.execute(sa.select(CAMPAIGNS.c.durum, sa.func.count()).where(CAMPAIGNS.c.tenant_id == tenant)
                                            .group_by(CAMPAIGNS.c.durum)).all()}
    return {"items": [_camp_dict(r, summary=totals(items[r.id])) for r in rows], "total": total, "page": page,
            "pageSize": PAGE_SIZE, "sayilar": {k: int(counts.get(k, 0)) for k in DURUMLAR}}


def all_campaigns(engine: sa.engine.Engine, tenant: str, durum: str) -> list[dict[str, Any]]:
    """Durumdaki bütün kampanyalar (sayfa sayfa; sessiz tavan yok)."""
    out: list[dict[str, Any]] = []
    page = 0
    while True:
        got = list_campaigns(engine, tenant, durum=durum, page=page)
        out += got["items"]
        if (page + 1) * PAGE_SIZE >= got["total"]:
            return out
        page += 1


def _clean_head(body: dict[str, Any], base: Optional[Any] = None) -> dict[str, Any]:
    """Kampanya başlık alanları (yalnız gönderilenler; `base` yoksa zorunlular denetlenir)."""
    out: dict[str, Any] = {}
    if base is None or "ad" in body:
        ad = _text(body.get("ad"), 300)
        if not ad:
            raise KampanyaError("Kampanya adı boş olamaz.")
        out["ad"] = ad
    if base is None or "kanal" in body:
        k = str(body.get("kanal") or "")
        if k not in KANALLAR:
            raise KampanyaError("Kanal site, pazar yeri, bayi ya da fuar olmalı.")
        out["kanal"] = k
    for key, label in (("baslangic", "Başlangıç"), ("bitis", "Bitiş")):
        if base is None or key in body:
            out[key] = _day(body.get(key), label)
    if "platform" in body:
        out["platform"] = _text(body.get("platform"), 120)
    if "varsayilanIndirim" in body:
        out["varsayilan_indirim"] = ratio(body.get("varsayilanIndirim"), "Toplu indirim")
    if "kanalKesinti" in body:
        out["kanal_kesinti"] = ratio(body.get("kanalKesinti"), "Kanal kesintisi")
    if "beklenenArtis" in body:
        out["beklenen_artis"] = _num(body.get("beklenenArtis"), "Beklenen satış artışı", allow_none=True, minimum=0.1, maximum=100)
    if "butce" in body:
        out["butce"] = _num(body.get("butce"), "Bütçe", allow_none=True)
    if "notlar" in body:
        out["notlar"] = _longtext(body.get("notlar"), 4000)
    b = out.get("baslangic") or (base.baslangic if base is not None else None)
    e = out.get("bitis") or (base.bitis if base is not None else None)
    if b and e and e < b:
        raise KampanyaError("Bitiş başlangıçtan önce olamaz.")
    return out


def create_campaign(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    vals = _clean_head(body)
    now = _now()
    with engine.begin() as c:
        cid = _next_id(c, CAMPAIGNS, f"KM-{today().year}-")
        c.execute(CAMPAIGNS.insert().values(id=cid, tenant_id=tenant, durum="taslak", hazirlayan=user, created_at=now,
                                            updated_at=now, **vals))
    return get_campaign(engine, tenant, cid)


#: Onaydan sonra da değişebilen alanlar (kampanyanın hesabını değiştirmez).
AFTER_APPROVAL = ("notlar", "metin", "crmKampanyaId", "kurulumNotu")


def update_campaign(engine: sa.engine.Engine, tenant: str, user: str, cid: str, body: dict[str, Any],
                    crm_check: Optional[Callable[[str], bool]] = None) -> tuple[dict[str, Any], dict[str, Any]]:
    with engine.begin() as c:
        r = _row(c, tenant, cid, lock=True)
        head_keys = [k for k in body if k not in AFTER_APPROVAL]
        if head_keys and r.durum not in DUZENLENIR:
            raise KampanyaError("Onaya gönderilmiş kampanyanın hesabı değişmez; önce geri çekin ya da geri gönderilmesini bekleyin.", 409)
        vals = _clean_head(body, base=r)
        if "notlar" in body:
            vals["notlar"] = _longtext(body.get("notlar"), 4000)
        if "metin" in body:
            vals["metin_json"] = _dump(_clean_copy(body.get("metin"), _j(r.metin_json, {})))
        if "crmKampanyaId" in body:
            g = str(body.get("crmKampanyaId") or "").strip().strip("{}").lower()
            if g and not src.guid_ok(g):
                raise KampanyaError("CRM kampanya kimliği geçerli değil.")
            if g and crm_check is not None and not crm_check(g):
                raise KampanyaError("Bu kimlikte bir CRM kampanyası bulunamadı.", 404)
            vals["crm_kampanya_id"] = g or None
        if "kurulumNotu" in body:
            if r.durum not in ("onaylandi", "yurutuluyor", "bitti"):
                raise KampanyaError("Kurulum yalnız onaylanmış kampanyada işaretlenir.", 409)
            note = _longtext(body.get("kurulumNotu"), 2000)
            if note != r.kurulum_notu:
                vals.update(kurulum_notu=note, kuruldu_by=user if note else None, kuruldu_at=_now() if note else None)
        diff = {k: v for k, v in vals.items() if getattr(r, k) != v}
        if diff:
            c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == cid).values(**vals, updated_at=_now()))
    needs_sim = any(k in diff for k in ("kanal", "baslangic", "bitis", "kanal_kesinti", "beklenen_artis"))
    return get_campaign(engine, tenant, cid), {"diff": {k: (v if k != "metin_json" else "düzenlendi") for k, v in diff.items()},
                                               "resimulate": needs_sim}


def _clean_copy(raw: Any, current: dict[str, Any]) -> dict[str, Any]:
    """Metin düzenlemesi: seçilen/düzenlenen metin tür başına; üretilen varyantlar korunur."""
    if not isinstance(raw, dict):
        raise KampanyaError("Metin bilgisi geçersiz.")
    out = dict(current)
    sel = dict(current.get("secili") or {})
    for key, (label, limit) in METIN_TURLERI.items():
        if key in (raw.get("secili") or {}):
            v = _text((raw.get("secili") or {}).get(key), 2000)
            if v and len(v) > limit:
                raise KampanyaError(f"{label} en çok {limit} karakter olmalı ({len(v)}).")
            sel[key] = v
    out["secili"] = sel
    return out


def delete_campaign(engine: sa.engine.Engine, tenant: str, cid: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = _row(c, tenant, cid, lock=True)
        if r.durum != "taslak" or r.submitted_at is not None:
            raise KampanyaError("Yalnız hiç onaya gönderilmemiş taslak silinir; diğerleri iptal edilir.", 409)
        c.execute(ITEMS.delete().where(ITEMS.c.campaign_id == cid))
        c.execute(CAMPAIGNS.delete().where(CAMPAIGNS.c.id == cid))
    return {"id": cid, "ad": r.ad}


def cancel_campaign(engine: sa.engine.Engine, tenant: str, user: str, cid: str, note: Any) -> dict[str, Any]:
    with engine.begin() as c:
        r = _row(c, tenant, cid, lock=True)
        if r.durum in ("bitti", "iptal"):
            raise KampanyaError("Biten ya da iptal edilmiş kampanya iptal edilemez.", 409)
        n = _longtext(note, 2000)
        if not n:
            raise KampanyaError("İptal gerekçesi yazılmalı.")
        c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == cid).values(
            durum="iptal", notlar=((r.notlar + "\n\n") if r.notlar else "") + f"İptal ({user}): {n}", updated_at=_now()))
    return get_campaign(engine, tenant, cid)


# ------------------------------------------------------------------------------------------ kitaplar ve simülasyon


def _item_inputs(r: Any) -> dict[str, Any]:
    if r.fiyat_elle:
        return {"liste": r.liste_fiyati if r.liste_elle else None, "kampanyaFiyati": r.kampanya_fiyati, "indirim": None}
    return {"liste": r.liste_fiyati if r.liste_elle else None, "kampanyaFiyati": None, "indirim": r.indirim_orani}


#: Onaylanmış kampanyada gece yalnız bunlar tazelenir; fiyat, maliyet ve marj onaylanan hâliyle kalır.
STOCK_FIELDS = ("stok", "gunluk_hiz", "tukenme_tahmini")


def simulate(engine: sa.engine.Engine, st: dict[str, Any], tenant: str, cid: str, *, bulk: Optional[float] = None,
             only: Optional[list[str]] = None, save: bool = True, ref: Optional[date] = None,
             stock_only: bool = False) -> dict[str, Any]:
    """Kampanyanın kitaplarını yeniden hesaplar. `bulk` verilirse bütün kitaplara o indirim uygulanır. Hesap yalnız taslakta
    değişir; onaya gönderilmiş kampanyada `stock_only` ile yalnız stok, hız, tükenme ve stok kontrolü tazelenir."""
    ref = ref or today()
    with engine.connect() as c:
        r = _row(c, tenant, cid)
        rows = c.execute(sa.select(ITEMS).where(ITEMS.c.campaign_id == cid)).all()
    if save and not stock_only and r.durum not in DUZENLENIR:
        raise KampanyaError("Onaya gönderilmiş kampanyanın hesabı değişmez; onaycı gönderildiği hâli görür.", 409)
    books = books_by_code(engine, [x.stok_kodu for x in rows])
    costs = costs_for(st, books) if not stock_only else {}
    bas = _d(r.baslangic) or ref
    floors = price_floor(engine, [b.get("ean") for b in books.values()], min(bas, ref + timedelta(days=1)), st["fiyatGun"]) \
        if not stock_only else {}
    camp = {"kanal": r.kanal, "kanalKesinti": r.kanal_kesinti, "beklenenArtis": r.beklenen_artis, "baslangic": r.baslangic,
            "bitis": r.bitis}
    up = uplift_for(engine, tenant, r.kanal, st)
    end = data_end(engine)
    out_items = []
    now = _now()
    updates = []
    for x in rows:
        if only and x.stok_kodu not in only:
            continue
        b = books.get(x.stok_kodu) or {"stok": x.stok_kodu, "ad": x.ad}
        inp = _item_inputs(x)
        if bulk is not None:
            inp = {"liste": inp["liste"], "kampanyaFiyati": None, "indirim": bulk}
        if stock_only:
            inp = {"liste": x.liste_fiyati, "kampanyaFiyati": x.kampanya_fiyati, "indirim": None}
        vals = simulate_item(b, inp, camp, st, costs.get(x.stok_kodu), floors.get(b.get("ean") or ""), ref, end, up)
        if stock_only:
            old = [k for k in _j(x.kontroller_json, []) if k.get("kod") != "stok"]
            new = [k for k in _j(vals["kontroller_json"], []) if k.get("kod") == "stok"]
            vals = {k: vals[k] for k in STOCK_FIELDS} | {"kontroller_json": _dump(old + new)}
        else:
            vals.update(ean=b.get("ean") or x.ean, ad=b.get("ad") or x.ad, yazar=b.get("yazar") or x.yazar, hesaplandi=now,
                        fiyat_elle=inp["kampanyaFiyati"] is not None)
        updates.append((x.stok_kodu, vals))
    if save:
        with engine.begin() as c:
            for code, vals in updates:
                c.execute(ITEMS.update().where(ITEMS.c.campaign_id == cid, ITEMS.c.stok_kodu == code).values(**vals))
            if bulk is not None:
                c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == cid).values(varsayilan_indirim=bulk, updated_at=now))
        return get_campaign(engine, tenant, cid)
    by_code = {x.stok_kodu: x for x in rows}
    for code, vals in updates:
        base = {col.name: getattr(by_code[code], col.name) for col in ITEMS.columns}
        fake = type("Row", (), {**base, **vals})
        out_items.append(_item_dict(fake))
    return {"kitaplar": out_items, "ozet": totals(out_items)}


def add_items(engine: sa.engine.Engine, st: dict[str, Any], tenant: str, cid: str, raw: Any) -> tuple[dict[str, Any], list[str]]:
    """Toplu ekleme: [{stok, indirim?, kampanyaFiyati?, gerekce?}]. Var olan kitap güncellenir. Bilinmeyen kod eklenmez."""
    if not isinstance(raw, list) or not raw:
        raise KampanyaError("Eklenecek kitap listesi boş.")
    with engine.connect() as c:
        r = _row(c, tenant, cid)
        have = {x.stok_kodu: x for x in c.execute(sa.select(ITEMS).where(ITEMS.c.campaign_id == cid)).all()}
        next_sira = max([x.sira for x in have.values()] or [0]) + 1
    if r.durum not in DUZENLENIR:
        raise KampanyaError("Onaya gönderilmiş kampanyaya kitap eklenmez.", 409)
    wanted: dict[str, dict[str, Any]] = {}
    for x in raw:
        if not isinstance(x, dict):
            raise KampanyaError("Kitap bilgisi geçersiz.")
        code = str(x.get("stok") or "").strip()
        if not src.code_ok(code):
            raise KampanyaError(f"Stok kodu «{code}» geçersiz.")
        ind = ratio(x.get("indirim"), f"{code} indirimi") if x.get("indirim") not in (None, "") else None
        kf = _num(x.get("kampanyaFiyati"), f"{code} kampanya fiyatı", allow_none=True, minimum=0.01) if x.get("kampanyaFiyati") not in (None, "") else None
        wanted[code] = {"indirim": ind if kf is None else None, "kampanyaFiyati": kf,
                        "gerekce": _text(x.get("gerekce"), 600)}
    books = books_by_code(engine, list(wanted))
    missing = [c_ for c_ in wanted if c_ not in books]
    now = _now()
    with engine.begin() as c:
        for code, w in wanted.items():
            if code in missing:
                continue
            b = books[code]
            ind = w["indirim"] if w["indirim"] is not None or w["kampanyaFiyati"] is not None else r.varsayilan_indirim
            vals = {"ean": b.get("ean"), "ad": b.get("ad"), "yazar": b.get("yazar"), "indirim_orani": ind,
                    "kampanya_fiyati": w["kampanyaFiyati"], "fiyat_elle": w["kampanyaFiyati"] is not None}
            if w["gerekce"]:
                vals["aday_gerekcesi"] = w["gerekce"]
            if code in have:
                c.execute(ITEMS.update().where(ITEMS.c.campaign_id == cid, ITEMS.c.stok_kodu == code).values(**vals))
            else:
                c.execute(ITEMS.insert().values(campaign_id=cid, stok_kodu=code, sira=next_sira, eklenme=now, liste_elle=False,
                                                maliyet_eksik=False, **vals))
                next_sira += 1
    out = simulate(engine, st, tenant, cid, only=[c_ for c_ in wanted if c_ not in missing])
    return out, missing


def update_item(engine: sa.engine.Engine, st: dict[str, Any], tenant: str, cid: str, code: str, body: dict[str, Any]) -> dict[str, Any]:
    with engine.connect() as c:
        r = _row(c, tenant, cid)
        it = c.execute(sa.select(ITEMS).where(ITEMS.c.campaign_id == cid, ITEMS.c.stok_kodu == code)).first()
    if not it:
        raise KampanyaError("Kitap bu kampanyada yok.", 404)
    if r.durum not in DUZENLENIR:
        raise KampanyaError("Onaya gönderilmiş kampanyada kitap değişmez.", 409)
    vals: dict[str, Any] = {}
    if "kampanyaFiyati" in body and body.get("kampanyaFiyati") not in (None, ""):
        vals.update(kampanya_fiyati=_num(body.get("kampanyaFiyati"), "Kampanya fiyatı", minimum=0.01), indirim_orani=None,
                    fiyat_elle=True)
    elif "indirim" in body:
        vals.update(indirim_orani=ratio(body.get("indirim"), "İndirim"), kampanya_fiyati=None, fiyat_elle=False)
    if "liste" in body:
        lv = _num(body.get("liste"), "Liste fiyatı", allow_none=True, minimum=0.01)
        vals.update(liste_fiyati=lv, liste_elle=lv is not None)
    if "gerekce" in body:
        vals["aday_gerekcesi"] = _text(body.get("gerekce"), 600)
    if not vals:
        raise KampanyaError("Değişecek alan yok.")
    with engine.begin() as c:
        c.execute(ITEMS.update().where(ITEMS.c.campaign_id == cid, ITEMS.c.stok_kodu == code).values(**vals))
    return simulate(engine, st, tenant, cid, only=[code])


def remove_item(engine: sa.engine.Engine, tenant: str, cid: str, code: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = _row(c, tenant, cid, lock=True)
        if r.durum not in DUZENLENIR:
            raise KampanyaError("Onaya gönderilmiş kampanyadan kitap çıkarılmaz.", 409)
        n = c.execute(ITEMS.delete().where(ITEMS.c.campaign_id == cid, ITEMS.c.stok_kodu == code)).rowcount
    if not n:
        raise KampanyaError("Kitap bu kampanyada yok.", 404)
    return get_campaign(engine, tenant, cid)


# ------------------------------------------------------------------------------------------ onay akışı


def submit(engine: sa.engine.Engine, st: dict[str, Any], tenant: str, user: str, cid: str) -> dict[str, Any]:
    out = simulate(engine, st, tenant, cid)                      # onaycı güncel hesabı görsün
    if out["durum"] not in DUZENLENIR:
        raise KampanyaError("Yalnız taslak onaya gönderilir.", 409)
    if not out["kitaplar"]:
        raise KampanyaError("Kampanyada kitap yok.")
    missing = [i["stok"] for i in out["kitaplar"] if i.get("kampanyaFiyati") is None]
    if missing:
        raise KampanyaError(f"{len(missing)} kitapta kampanya fiyatı ya da indirim yok.")
    now = _now()
    with engine.begin() as c:
        c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == cid, CAMPAIGNS.c.durum == "taslak").values(
            durum="onay_bekliyor", gonderen=user, submitted_at=now, updated_at=now, onaylayan=None, decided_at=None))
    return get_campaign(engine, tenant, cid)


def withdraw(engine: sa.engine.Engine, tenant: str, user: str, cid: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = _row(c, tenant, cid, lock=True)
        if r.durum != "onay_bekliyor":
            raise KampanyaError("Yalnız onay bekleyen kampanya geri çekilir.", 409)
        c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == cid).values(durum="taslak", updated_at=_now()))
    return get_campaign(engine, tenant, cid)


def decide(engine: sa.engine.Engine, tenant: str, user: str, cid: str, approve: bool, note: Any) -> dict[str, Any]:
    """Onay ya da geri gönderme. Hazırlayan ve onaya gönderen onaylayamaz (M8 hakediş ilkesi)."""
    n = _longtext(note, 2000)
    with engine.begin() as c:
        r = _row(c, tenant, cid, lock=True)
        if r.durum != "onay_bekliyor":
            raise KampanyaError("Kampanya onay beklemiyor.", 409)
        if user in {r.hazirlayan, r.gonderen}:
            raise KampanyaError("Kampanyayı hazırlayan ya da onaya gönderen onaylayamaz.", 409)
        if not approve and not n:
            raise KampanyaError("Geri gönderme gerekçesi yazılmalı.")
        now = _now()
        c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == cid).values(
            durum="onaylandi" if approve else "taslak", onaylayan=user, onay_notu=n, decided_at=now, updated_at=now))
    return get_campaign(engine, tenant, cid)


def advance(engine: sa.engine.Engine, tenant: str, ref: Optional[date] = None) -> list[dict[str, str]]:
    """Tarih geçişleri: onaylandı → yürütülüyor (başlangıç geldi), onaylandı/yürütülüyor → bitti (bitiş geçti)."""
    ref = (ref or today()).isoformat()
    changed = []
    with engine.begin() as c:
        for r in c.execute(sa.select(CAMPAIGNS.c.id, CAMPAIGNS.c.ad, CAMPAIGNS.c.durum, CAMPAIGNS.c.baslangic, CAMPAIGNS.c.bitis).where(
                CAMPAIGNS.c.tenant_id == tenant, CAMPAIGNS.c.durum.in_(("onaylandi", "yurutuluyor")))).all():
            new = "bitti" if r.bitis < ref else ("yurutuluyor" if r.baslangic <= ref else r.durum)
            if new != r.durum:
                c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == r.id).values(durum=new, updated_at=_now()))
                changed.append({"id": r.id, "ad": r.ad, "durum": new})
    return changed


# ------------------------------------------------------------------------------------------ takvim


def _special_days(engine: sa.engine.Engine, tenant: str, a: date, b: date) -> list[dict[str, Any]]:
    """SEO sezon takviminin özel günleri (CRM + kodda tanımlı hareketli günler), [a, b] ile kesişen aralıklar."""
    from semantic_bridge.seo_geo import seasons as S

    days: list[dict[str, Any]] = []
    try:
        if sa.inspect(engine).has_table(S.DAYS.name):
            with engine.connect() as c:
                rows = c.execute(sa.select(S.DAYS).where(S.DAYS.c.tenant_id == tenant)).mappings().all()
            days = [{"key": r["day_key"], "name": r["name"], "source": r["source"], "weekFrom": r["week_from"],
                     "weekTo": r["week_to"], "fixedDate": r["fixed_date"], "crmBooks": r["crm_books"]} for r in rows]
    except Exception as e:  # noqa: BLE001 — sezon tablosu okunamazsa kodda tanımlı günlerle devam
        log.warning("kampanya: sezon takvimi okunamadı: %s", e)
    if not days:
        days = [{**d, "crmBooks": 0} for d in S.merge_builtin([])]
    out = []
    for d in days:
        how = S.resolve(d)
        for y in range(a.year - 1, b.year + 1):
            for s, e in S.occurrences(how, y):
                if e >= a and s <= b:
                    out.append({"id": f"gun:{d['key']}:{s.isoformat()}", "tur": "ozel_gun", "ad": d["name"], "baslangic": s.isoformat(),
                                "bitis": e.isoformat(), "kaynak": "crm" if d.get("source") == "crm" else "kural",
                                "kesinlik": how["precision"], "neden": how["why"], "gunKey": d["key"],
                                "kitapSayisi": int(d.get("crmBooks") or 0), "silinir": False})
    return out


def calendar(engine: sa.engine.Engine, tenant: str, a: date, b: date) -> dict[str, Any]:
    items = _special_days(engine, tenant, a, b)
    with engine.connect() as c:
        for r in c.execute(sa.select(CALENDAR).where(CALENDAR.c.tenant_id == tenant, CALENDAR.c.bitis >= a.isoformat(),
                                                     CALENDAR.c.baslangic <= b.isoformat())).all():
            items.append({"id": r.id, "tur": r.tur, "ad": r.ad, "baslangic": r.baslangic, "bitis": r.bitis, "platform": r.platform,
                          "kaynak": r.kaynak, "notlar": r.notlar, "olusturan": r.olusturan, "silinir": True})
        camps = c.execute(sa.select(CAMPAIGNS).where(CAMPAIGNS.c.tenant_id == tenant, CAMPAIGNS.c.durum != "iptal",
                                                     CAMPAIGNS.c.bitis >= a.isoformat(), CAMPAIGNS.c.baslangic <= b.isoformat())).all()
    kampanyalar = [{"id": r.id, "ad": r.ad, "kanal": r.kanal, "platform": r.platform, "baslangic": r.baslangic, "bitis": r.bitis,
                    "durum": r.durum, "durumAdi": DURUMLAR.get(r.durum, r.durum)} for r in camps]
    items.sort(key=lambda x: (x["baslangic"], x["ad"]))
    return {"from": a.isoformat(), "to": b.isoformat(), "items": items, "kampanyalar": kampanyalar,
            "cakismalar": overlaps(kampanyalar, items)}


def overlaps(camps: list[dict[str, Any]], days: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Aynı kanal ve platformda tarihleri kesişen iki kampanya; kampanyanın denk geldiği özel gün / platform dönemi."""
    out = []
    for i, x in enumerate(camps):
        for y in camps[i + 1:]:
            if x["kanal"] == y["kanal"] and fold(x.get("platform")) == fold(y.get("platform")) \
                    and x["baslangic"] <= y["bitis"] and y["baslangic"] <= x["bitis"]:
                out.append({"tur": "kampanya", "a": x["id"], "b": y["id"],
                            "mesaj": f"«{x['ad']}» ile «{y['ad']}» aynı kanalda ve tarihlerde çakışıyor."})
        for d in days:
            if x["baslangic"] <= d["bitis"] and d["baslangic"] <= x["bitis"]:
                out.append({"tur": "denk", "a": x["id"], "b": d["id"],
                            "mesaj": f"«{x['ad']}» {TAKVIM_TURLERI.get(d['tur'], d['tur']).lower()} «{d['ad']}» ile aynı günlerde."})
    return out


def add_calendar(engine: sa.engine.Engine, tenant: str, user: str, body: dict[str, Any]) -> dict[str, Any]:
    tur = str(body.get("tur") or "")
    if tur not in TAKVIM_TURLERI:
        raise KampanyaError("Tür platform dönemi, özel gün ya da fuar olmalı.")
    ad = _text(body.get("ad"), 300)
    if not ad:
        raise KampanyaError("Ad boş olamaz.")
    a, b = _day(body.get("baslangic"), "Başlangıç"), _day(body.get("bitis") or body.get("baslangic"), "Bitiş")
    if b < a:
        raise KampanyaError("Bitiş başlangıçtan önce olamaz.")
    with engine.begin() as c:
        kid = _next_id(c, CALENDAR, f"KT-{today().year}-")
        c.execute(CALENDAR.insert().values(id=kid, tenant_id=tenant, tur=tur, ad=ad, baslangic=a, bitis=b,
                                           platform=_text(body.get("platform"), 120), kaynak="kullanici",
                                           notlar=_longtext(body.get("notlar"), 2000), olusturan=user, created_at=_now()))
    return {"id": kid, "tur": tur, "ad": ad, "baslangic": a, "bitis": b}


def delete_calendar(engine: sa.engine.Engine, tenant: str, kid: str) -> dict[str, Any]:
    with engine.begin() as c:
        r = c.execute(sa.select(CALENDAR).where(CALENDAR.c.id == kid, CALENDAR.c.tenant_id == tenant)).first()
        if not r:
            raise KampanyaError("Takvim kaydı bulunamadı (özel günler sezon takviminden gelir, buradan silinmez).", 404)
        c.execute(CALENDAR.delete().where(CALENDAR.c.id == kid))
    return {"id": kid, "ad": r.ad}


# ------------------------------------------------------------------------------------------ adaylar


def parse_rules(raw: str) -> list[str]:
    rules = [r for r in (raw or "").split(",") if r in KURALLAR]
    return rules or list(VARSAYILAN_KURALLAR)


def _season_window(engine: sa.engine.Engine, tenant: str, bas: date, bit: date, lead: int) -> dict[str, str]:
    """Kampanya penceresine (başlangıçtan `lead` gün öncesi dahil) denk gelen özel günler: gün anahtarı → ad."""
    return {d["gunKey"]: d["ad"] for d in _special_days(engine, tenant, bas, bit + timedelta(days=lead)) if d.get("gunKey")}


def candidates(engine: sa.engine.Engine, st: dict[str, Any], tenant: str, *, cid: Optional[str] = None, rules: str = "",
               q: str = "", page: int = 0, ref: Optional[date] = None) -> dict[str, Any]:
    """Kural süzgeci: sinyallerden (stok fazlası, yavaşlama, sezon) en az biri; seçilen koşullar (hak, maliyet, marj) hepsi.
    Gerekçe rakamlarla tek satır. Sayı tavanı yok, sayfalı."""
    ref = ref or today()
    chosen = parse_rules(rules)
    camp = None
    have: set[str] = set()
    if cid:
        with engine.connect() as c:
            camp = _row(c, tenant, cid)
            have = {x[0] for x in c.execute(sa.select(ITEMS.c.stok_kodu).where(ITEMS.c.campaign_id == cid)).all()}
    bas = _d(camp.baslangic) if camp else ref
    bit = _d(camp.bitis) if camp else ref + timedelta(days=30)
    seasons = _season_window(engine, tenant, bas, bit, st["sezonOncesiGun"]) if "sezon" in chosen else {}
    months = st["hizAy"]
    cond = [BOOKS.c.stok > 0]
    if q.strip():
        like = f"%{q.strip()}%"
        cond.append(sa.or_(BOOKS.c.ad.ilike(like), BOOKS.c.stok_kodu.ilike(like), BOOKS.c.yazar.ilike(like)))
    with engine.connect() as c:
        rows = [book_dict(b) for b in c.execute(sa.select(BOOKS).where(*cond)).all()]
    rows = [b for b in rows if b["stok"] not in have and b.get("crmTip") == 1]
    pre = []
    for b in rows:
        reasons, score, signals = [], 0.0, 0
        monthly = (b.get("adetSon") or 0.0) / months
        stok_ay = (b["stokAdet"] / monthly) if monthly > 0 else None
        if "stok" in chosen and (stok_ay is None or stok_ay >= st["stokAy"]):
            signals += 1
            score += min(3.0, (stok_ay or 36.0) / max(st["stokAy"], 1.0))
            reasons.append(f"stok {stok_ay:.0f} ay yeter" if stok_ay is not None else f"son {months} ayda satış yok, stok {b['stokAdet']:.0f}")
        prev = b.get("adetOnceki") or 0.0
        drop = (1 - (b.get("adetSon") or 0.0) / prev) if prev > 0 else None
        if "dusus" in chosen and drop is not None and drop * 100 >= st["dususPct"]:
            signals += 1
            score += min(2.0, drop * 2)
            reasons.append(f"satış %{drop * 100:.0f} düştü (son {months} ay {b.get('adetSon') or 0:.0f}, önceki {prev:.0f} adet)")
        hits = [seasons[k] for k in b.get("sezonlar") or [] if k in seasons]
        if "sezon" in chosen and hits:
            signals += 1
            score += 1.5
            reasons.append(", ".join(dict.fromkeys(hits)) + " ile bağlı")
        if signals == 0:
            continue
        if "hak" in chosen and (b.get("durumBayragi") or b.get("hak") in ("yok", "eksik")):
            continue
        pre.append((b, score, reasons, stok_ay, drop, hits))
    need_cost = "maliyet" in chosen or "marj" in chosen
    costs = costs_for(st, {b["stok"]: b for b, *_ in pre}) if need_cost and pre else {}
    ind = (camp.varsayilan_indirim if camp is not None else None) or 0.0
    out = []
    kes = (camp.kanal_kesinti if camp is not None else None) or 0.0
    for b, score, reasons, stok_ay, drop, hits in pre:
        unit = (costs.get(b["stok"]) or {}).get("maliyet") if need_cost else None
        if "maliyet" in chosen and unit is None:
            continue
        marj_oran = None
        if "marj" in chosen:
            liste, _ = list_price(b, st)
            if not liste or unit is None:
                continue
            f = 1 + float(b.get("kdv") or 0) / 100
            net = liste * (1 - ind) / f * (1 - kes)
            tel, _n = telif_units(b.get("sozlesmeler") or [], liste / f, net, liste * (1 - ind) / f)
            if tel is None or not net:
                continue
            marj_oran = (net - unit - tel) / net
            floor_pct = st["adayMarjMinPct"]
            if marj_oran < (floor_pct / 100.0 if floor_pct is not None else 0.0):
                continue
            reasons.append(f"%{ind * 100:.0f} indirimde marj %{marj_oran * 100:.0f}")
        out.append({"stok": b["stok"], "ad": b["ad"], "yazar": b["yazar"], "ean": b["ean"], "stokAdet": b["stokAdet"],
                    "stokAy": _r2(stok_ay), "dusus": _r4(drop), "adetSon": b.get("adetSon"), "adetOnceki": b.get("adetOnceki"),
                    "sezonlar": list(dict.fromkeys(hits)), "liste": list_price(b, st)[0], "marjOrani": _r4(marj_oran),
                    "hak": b.get("hak"), "puan": round(score, 3), "gerekce": _sentence("; ".join(reasons))})
    out.sort(key=lambda x: (-x["puan"], -(x["stokAy"] or 0), x["stok"]))
    page = max(0, int(page))
    shown = out[page * PAGE_SIZE:(page + 1) * PAGE_SIZE]
    dr_tarih = attach_dr(engine, tenant, shown)
    return {"items": shown, "total": len(out), "page": page, "pageSize": PAGE_SIZE, "drTarih": dr_tarih,
            "kurallar": chosen, "esikler": {"stokAy": st["stokAy"], "dususPct": st["dususPct"], "hizAy": months,
                                            "sezonOncesiGun": st["sezonOncesiGun"], "marjMinPct": st["adayMarjMinPct"],
                                            "indirim": ind},
            "sezonlar": sorted(set(seasons.values())), "dataEnd": meta_get(engine, "data_end").get("date")}


def attach_dr(engine: sa.engine.Engine, tenant: str, rows: list[dict[str, Any]]) -> Optional[str]:
    """Aday satırlarına D&R'nin güncel fiyatı (M39 D&R kataloğu, son görüntü, barkod eşleşmesi): `dr` = {liste, satış
    fiyatı, indirim, site durumu} ya da None; `drUyari` = «D&R zaten %X indirimde» (sitede satışta ve indirim ≥ %1).
    Katalog okunmadıysa ya da okunamazsa satırlar `dr: None` alır, aday listesi düşmez. Dönen: görüntü tarihi."""
    for r in rows:
        r["dr"], r["drUyari"] = None, None
    if not rows:
        return None
    try:
        from semantic_bridge import pazar_dagitim as PD

        got = PD.dr_fiyatlari(engine, tenant, codes=[r["stok"] for r in rows], eans=[r.get("ean") for r in rows])
    except Exception as e:  # noqa: BLE001 — D&R bilgisi aday listesini düşürmez
        log.info("kampanya: D&R fiyatı eklenemedi: %s", e)
        return None
    for r in rows:
        info = got["kod"].get(r["stok"]) or got["ean"].get(PD.barkod(r.get("ean")) or "")
        if not info:
            continue
        r["dr"] = {k: info[k] for k in ("fiyat", "drFiyat", "indirim", "durum", "siteSatista", "katalogda", "son")}
        r["drUyari"] = PD.dr_uyari(info)
    return got["tarih"]


# ------------------------------------------------------------------------------------------ sonuç


def windows(bas: date, bit: date, after: int) -> dict[str, tuple[date, date]]:
    """Dönemler (bitiş dahil): önce = kampanyayla eşit uzunlukta hemen öncesi, sonra = bitişten sonraki `after` gün."""
    n = (bit - bas).days + 1
    return {"once": (bas - timedelta(days=n), bas - timedelta(days=1)), "kampanya": (bas, bit),
            "sonra": (bit + timedelta(days=1), bit + timedelta(days=after))}


def store_results(engine: sa.engine.Engine, cid: str, rows: list[dict[str, Any]], wins: dict[str, tuple[date, date]],
                  end: Optional[date], kaynak: str) -> int:
    now = _now()
    kesim = end.isoformat() if end else None
    batch = []
    for r in rows:
        donem = next((k for k, (a, b) in wins.items() if a <= r["gun"] <= b), None)
        if donem is None:
            continue
        batch.append({"campaign_id": cid, "gun": r["gun"].isoformat(), "stok_kodu": r["stok"], "donem": donem, "adet": r["adet"],
                      "net_tutar": r["tutar"], "iade_adet": r["iade"], "maliyet": r["maliyet"], "maliyetli_tutar": r["maliyetliTutar"],
                      "kaynak": kaynak, "logo_kesim": kesim, "asof": now})
    with engine.begin() as c:
        c.execute(RESULTS.delete().where(RESULTS.c.campaign_id == cid))
        for i in range(0, len(batch), 1000):
            if batch[i:i + 1000]:
                c.execute(RESULTS.insert(), batch[i:i + 1000])
    return len(batch)


def has_results(engine: sa.engine.Engine, cid: str) -> bool:
    with engine.connect() as c:
        return c.execute(sa.select(RESULTS.c.campaign_id).where(RESULTS.c.campaign_id == cid).limit(1)).first() is not None


def results(engine: sa.engine.Engine, st: dict[str, Any], tenant: str, cid: str) -> dict[str, Any]:
    """Önce / kampanya / sonra: adet, iade, net tutar, maliyetli marj; günlük ortalama ve değişim; kitap kitap; günlük seri."""
    with engine.connect() as c:
        r = _row(c, tenant, cid)
        rows = c.execute(sa.select(RESULTS).where(RESULTS.c.campaign_id == cid)).all()
        items = {x.stok_kodu: x for x in c.execute(sa.select(ITEMS).where(ITEMS.c.campaign_id == cid)).all()}
    bas, bit = _d(r.baslangic), _d(r.bitis)
    wins = windows(bas, bit, st["sonraGun"]) if bas and bit else {}
    kesim = max((x.logo_kesim for x in rows if x.logo_kesim), default=None) or meta_get(engine, "data_end").get("date")
    end = _d(kesim)
    per: dict[str, dict[str, float]] = {k: {"adet": 0.0, "iade": 0.0, "tutar": 0.0, "maliyet": 0.0, "maliyetliTutar": 0.0} for k in wins}
    books: dict[str, dict[str, dict[str, float]]] = {}
    series: dict[str, dict[str, float]] = {}
    for x in rows:
        p = per.setdefault(x.donem, {"adet": 0.0, "iade": 0.0, "tutar": 0.0, "maliyet": 0.0, "maliyetliTutar": 0.0})
        for k, v in (("adet", x.adet), ("iade", x.iade_adet), ("tutar", x.net_tutar), ("maliyet", x.maliyet),
                     ("maliyetliTutar", x.maliyetli_tutar)):
            p[k] += v or 0.0
        bk = books.setdefault(x.stok_kodu, {}).setdefault(x.donem, {"adet": 0.0, "iade": 0.0, "tutar": 0.0})
        bk["adet"] += x.adet or 0.0
        bk["iade"] += x.iade_adet or 0.0
        bk["tutar"] += x.net_tutar or 0.0
        s = series.setdefault(x.gun, {"adet": 0.0, "tutar": 0.0, "donem": x.donem})
        s["adet"] += x.adet or 0.0
        s["tutar"] += x.net_tutar or 0.0
    donemler = {}
    for k, (a, b) in wins.items():
        covered = max(0, (min(b, end) - a).days + 1) if end else 0
        p = per.get(k) or {}
        marj = (p["maliyetliTutar"] - p["maliyet"]) if p.get("maliyetliTutar") else None
        donemler[k] = {"bas": a.isoformat(), "bit": b.isoformat(), "gun": covered, "gunToplam": (b - a).days + 1,
                       "adet": _r2(p.get("adet")), "iade": _r2(p.get("iade")), "tutar": _r2(p.get("tutar")),
                       "gunlukAdet": _r2(p["adet"] / covered) if covered and p else None,
                       "iadeOrani": _r4(p["iade"] / p["adet"]) if p and p.get("adet") else None,
                       "marj": _r2(marj), "marjOrani": _r4(marj / p["maliyetliTutar"]) if marj is not None and p["maliyetliTutar"] else None,
                       "maliyetKapsami": _r4(p["maliyetliTutar"] / p["tutar"]) if p and p.get("tutar") else None}
    o, k_ = donemler.get("once") or {}, donemler.get("kampanya") or {}
    degisim = {
        "satis": _r4(k_["gunlukAdet"] / o["gunlukAdet"]) if k_.get("gunlukAdet") is not None and o.get("gunlukAdet") else None,
        "iadePuan": _r4((k_.get("iadeOrani") or 0) - (o.get("iadeOrani") or 0)) if k_.get("iadeOrani") is not None and o.get("iadeOrani") is not None else None,
        "marjPuan": _r4((k_.get("marjOrani") or 0) - (o.get("marjOrani") or 0)) if k_.get("marjOrani") is not None and o.get("marjOrani") is not None else None,
    }
    notlar = []
    if not rows:
        notlar.append("Henüz sonuç okunmadı." if r.durum in ("onaylandi", "yurutuluyor", "bitti") else "Kampanya onaylanınca sonuç okunur.")
    if end and bit and end < bit:
        notlar.append(f"Logo verisi {_day_tr(kesim)} tarihine kadar; kampanya dönemi eksik okunuyor (saatlik ya da canlı değil).")
    if end and bas and end < bas:
        notlar.append("Kampanya başlangıcı Logo verisinin kesim tarihinden sonra: kampanya dönemi satışı henüz yok.")
    kaynak = rows[0].kaynak if rows else None
    if kaynak == "logo":
        notlar.append("Satış bütün kanallardan (Logo faturası); kanal ayrımı için yönetimde kanal cari kodları girilmeli.")
    kitaplar = []
    for code, it in items.items():
        d = books.get(code, {})
        on, ka = d.get("once", {}), d.get("kampanya", {})
        kitaplar.append({"stok": code, "ad": it.ad, "indirim": it.indirim_orani, "onceAdet": on.get("adet", 0.0),
                         "kampanyaAdet": ka.get("adet", 0.0), "sonraAdet": d.get("sonra", {}).get("adet", 0.0),
                         "kampanyaIade": ka.get("iade", 0.0), "kampanyaTutar": ka.get("tutar", 0.0),
                         "degisim": _r4(ka["adet"] / on["adet"]) if ka.get("adet") is not None and on.get("adet") else None})
    kitaplar.sort(key=lambda x: -(x["kampanyaAdet"] or 0))
    return {"id": cid, "durum": r.durum, "logoKesim": kesim, "kaynak": kaynak, "donemler": donemler, "degisim": degisim,
            "kitaplar": kitaplar, "seri": [{"gun": g, **v} for g, v in sorted(series.items())], "notlar": notlar,
            "ozet": r.sonuc_ozet, "ozetAt": _iso(r.sonuc_at), "ogrenimler": learnings(engine, tenant, campaign_id=cid)["items"]}


#: Sorgu bilgisi: gece okumasının Logo/CRM sorguları ve kampanya sonucunu dolduran Logo sorgusu (`semantic_query_origin`).
KOKEN_OKUMA = "kampanya.okuma"


def koken_sonuc(cid: str) -> str:
    return f"kampanya.sonuc.{cid}"


def refresh_results(engine: sa.engine.Engine, st: dict[str, Any], tenant: str, cid: str, run: src.Runner,
                    firms: dict[int, str], end: Optional[date]) -> dict[str, Any]:
    with engine.connect() as c:
        r = _row(c, tenant, cid)
        codes = [x[0] for x in c.execute(sa.select(ITEMS.c.stok_kodu).where(ITEMS.c.campaign_id == cid)).all()]
    if r.durum not in ("onaylandi", "yurutuluyor", "bitti"):
        raise KampanyaError("Sonuç yalnız onaylanmış kampanyada okunur.", 409)
    bas, bit = _d(r.baslangic), _d(r.bitis)
    if not codes or not bas or not bit:
        return {"satir": 0}
    wins = windows(bas, bit, st["sonraGun"])
    a, b = wins["once"][0], wins["sonra"][1]
    if end:
        b = min(b, end)
    cari = st["kanalCari"].get(r.kanal)
    from semantic_bridge import sorgu_yakala as Y

    with Y.yakala() as q:   # sorgu bilgisi: sonuç tablosunu dolduran Logo sorgusu (asıl SQL) saklanır
        rows = src.read_daily(run, firms, codes, a, b + timedelta(days=1), cari) if b >= a else []
    Y.koken_yaz(engine, tenant, koken_sonuc(cid), q)
    n = store_results(engine, cid, rows, wins, end, "logo-kanal" if cari else "logo")
    return {"satir": n, "logoKesim": end.isoformat() if end else None}


# ------------------------------------------------------------------------------------------ öğrenim


def learnings(engine: sa.engine.Engine, tenant: str, *, campaign_id: str = "", kanal: str = "", page: int = 0) -> dict[str, Any]:
    cond = [LEARNINGS.c.tenant_id == tenant]
    if campaign_id:
        cond.append(LEARNINGS.c.campaign_id == campaign_id)
    if kanal:
        cond.append(LEARNINGS.c.kanal == kanal)
    page = max(0, int(page))
    with engine.connect() as c:
        total = c.execute(sa.select(sa.func.count()).select_from(LEARNINGS).where(*cond)).scalar() or 0
        rows = c.execute(sa.select(LEARNINGS, CAMPAIGNS.c.ad.label("kampanya")).select_from(
            LEARNINGS.outerjoin(CAMPAIGNS, CAMPAIGNS.c.id == LEARNINGS.c.campaign_id)).where(*cond)
            .order_by(LEARNINGS.c.tarih.desc()).offset(page * PAGE_SIZE).limit(PAGE_SIZE)).all()
    return {"items": [{"id": r.id, "kampanyaId": r.campaign_id, "kampanya": r.kampanya, "kanal": r.kanal, "ozet": r.ozet, "tur": r.tur,
                       "indirim": r.indirim_orani, "satisDegisimi": r.satis_degisimi, "iadeDegisimi": r.iade_degisimi,
                       "marjDegisimi": r.marj_degisimi, "yazan": r.yazan, "tarih": _iso(r.tarih)} for r in rows],
            "total": total, "page": page, "pageSize": PAGE_SIZE}


def add_learning(engine: sa.engine.Engine, st: dict[str, Any], tenant: str, user: str, cid: str, body: dict[str, Any]) -> dict[str, Any]:
    ozet = _longtext(body.get("ozet"), 4000)
    if not ozet:
        raise KampanyaError("Öğrenim metni boş olamaz.")
    res = results(engine, st, tenant, cid)
    if res["durum"] != "bitti":
        raise KampanyaError("Öğrenim kampanya bitince yazılır.", 409)
    with engine.connect() as c:
        r = _row(c, tenant, cid)
        inds = [x[0] for x in c.execute(sa.select(ITEMS.c.indirim_orani).where(ITEMS.c.campaign_id == cid)).all() if x[0] is not None]
    with engine.begin() as c:
        lid = _next_id(c, LEARNINGS, f"KO-{today().year}-")
        c.execute(LEARNINGS.insert().values(
            id=lid, tenant_id=tenant, campaign_id=cid, kanal=r.kanal, ozet=ozet, tur=_text(body.get("tur"), 40),
            indirim_orani=round(sum(inds) / len(inds), 4) if inds else None, satis_degisimi=res["degisim"]["satis"],
            iade_degisimi=res["degisim"]["iadePuan"], marj_degisimi=res["degisim"]["marjPuan"], yazan=user, tarih=_now()))
    return {"id": lid, "kampanyaId": cid, "ozet": ozet}


def delete_learning(engine: sa.engine.Engine, tenant: str, user: str, lid: str, is_admin: bool) -> dict[str, Any]:
    with engine.begin() as c:
        r = c.execute(sa.select(LEARNINGS).where(LEARNINGS.c.id == lid, LEARNINGS.c.tenant_id == tenant)).first()
        if not r:
            raise KampanyaError("Öğrenim kaydı bulunamadı.", 404)
        if r.yazan != user and not is_admin:
            raise KampanyaError("Öğrenim kaydını yalnız yazan siler.", 403)
        c.execute(LEARNINGS.delete().where(LEARNINGS.c.id == lid))
    return {"id": lid, "kampanyaId": r.campaign_id}


# ------------------------------------------------------------------------------------------ Zeki AI


COPY_PROMPT = """Bir yayınevinin e-ticaret kampanyası için {label} yaz: {n} farklı seçenek, her biri tek satır, en çok {limit} karakter.
Kampanya: {name}. Kanal: {channel}. Dönem: {period}.{season}
Kampanyadaki kitaplar:
{books}

Kurallar: yalnız verilen bilgileri kullan; indirim oranı dışında rakam yazma{discount}; «en çok satan», «bir numara», «rekor» gibi
kanıtsız üstünlük iddiası yok; teknoloji ya da yapay zekâ adı yok; ünlem en çok bir tane. Her seçeneği yeni satırda, başında
«- » ile yaz; başka açıklama ekleme."""

SUMMARY_PROMPT = """Aşağıdaki kampanya sonucunu yönetime 4–6 cümlelik Türkçe bir paragrafla özetle. Yalnız verilen rakamları kullan,
yeni rakam üretme, yorumu rakama dayandır; veri eksikse bunu söyle. Teknoloji adı yazma.

Kampanya: {name} ({channel}{platform}), {bas} – {bit}, {n} kitap, ortalama indirim {ind}.
Logo verisi {kesim} tarihine kadar.
{facts}"""

CRM_TYPE_PROMPT = """Bir yayınevinin CRM'deki bayi kampanyası kaydı aşağıda. Bu kampanyanın türü hangisidir?
Ad: {name}
Açıklama: {desc}
İskonto: {disc}. Vade: {vade}. Hediye adedi: {gift}."""
CRM_TYPE_CHOICES = {"iskonto": "Ek iskonto / indirim", "vade": "Vade (ödeme süresi) kampanyası", "hediye": "Hediye ürün kampanyası",
                    "stant": "Stant / teşhir kampanyası", "diger": "Diğer"}


def _books_text(items: list[dict[str, Any]]) -> str:
    return "\n".join(f"- {i.get('ad') or i['stok']}" + (f" ({i['yazar']})" if i.get("yazar") else "") for i in items) or "-"


def parse_variants(raw: str, limit: int, sources: list[str], facts: list[str]) -> tuple[list[str], int]:
    """«- …» satırları → sınırı aşmayan, denetimden geçen seçenekler; düşen sayısı."""
    out, dropped = [], 0
    for line in (raw or "").splitlines():
        t = re.sub(r"^\s*(?:[-•*]+|\d+[.)])\s*", "", line).strip().strip("«»\"'")
        if not t:
            continue
        g = mguard.check(t, sources, facts)
        text = g["metin"].strip()
        if g["dusenSayisi"] or not text or len(text) > limit:
            dropped += 1
            continue
        if text not in out:
            out.append(text)
    return out, dropped


def draft_copy(engine: sa.engine.Engine, llm: Any, tenant: str, cid: str, kind: str, n: int = 3) -> dict[str, Any]:
    if kind not in METIN_TURLERI:
        raise KampanyaError("Metin türü başlık, kısa açıklama ya da banner olmalı.")
    if llm is None:
        raise KampanyaError("Zeki AI bu kurulumda tanımlı değil.", 503)
    camp = get_campaign(engine, tenant, cid)
    if not camp["kitaplar"]:
        raise KampanyaError("Kampanyada kitap yok.")
    label, limit = METIN_TURLERI[kind]
    ind = camp["ozet"]["ortalamaIndirim"]
    discount = f" (indirim %{round(ind * 100):.0f})" if ind else ""
    seasons = [d["ad"] for d in _special_days(engine, tenant, _d(camp["baslangic"]), _d(camp["bitis"]))]
    season = f" Denk gelen özel gün: {', '.join(dict.fromkeys(seasons))}." if seasons else ""
    prompt = COPY_PROMPT.format(label=label.lower(), n=max(1, min(n, 6)), limit=limit, name=camp["ad"], channel=camp["kanalAdi"],
                                period=f"{_day_tr(camp['baslangic'])} – {_day_tr(camp['bitis'])}", season=season,
                                books=_books_text(camp["kitaplar"]), discount=discount)
    raw = (llm.chat([{"role": "user", "content": prompt}], max_tokens=600, temperature=0.5) or "").strip()
    sources = [camp["ad"], camp.get("platform") or "", *seasons, *(i.get("ad") or "" for i in camp["kitaplar"]),
               *(i.get("yazar") or "" for i in camp["kitaplar"])]
    facts = [f"%{round(ind * 100):.0f}", f"{round(ind * 100):.0f}"] if ind else []
    variants, dropped = parse_variants(raw, limit, sources, facts)
    if not variants:
        raise KampanyaError("Zeki AI kurallara uyan bir metin üretemedi; yeniden deneyin.", 503)
    with engine.begin() as c:
        r = _row(c, tenant, cid, lock=True)
        cur = _j(r.metin_json, {})
        cur[kind] = variants
        cur.setdefault("secili", {})
        cur["secili"].setdefault(kind, variants[0])
        cur["at"] = _now().isoformat()
        c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == cid).values(metin_json=_dump(cur), updated_at=_now()))
    return {"tur": kind, "secenekler": variants, "dusenSayisi": dropped, "sinir": limit}


def summary_facts(camp: dict[str, Any], res: dict[str, Any]) -> tuple[str, list[str]]:
    lines, facts = [], []
    names = {"once": "Kampanyadan önceki eşit dönem", "kampanya": "Kampanya dönemi", "sonra": "Kampanya sonrası"}
    for k, label in names.items():
        d = res["donemler"].get(k)
        if not d:
            continue
        bits = [f"{d['gun']} gün veri", f"satış {d['adet']:.0f} adet", f"günlük {d['gunlukAdet']:.1f} adet" if d.get("gunlukAdet") is not None else None,
                f"iade {d['iade']:.0f} adet", f"net tutar {_money(d['tutar'])}",
                f"marj {_pct(d['marjOrani'], 1)}" if d.get("marjOrani") is not None else "marj hesaplanamadı"]
        line = f"{label}: " + ", ".join(b for b in bits if b)
        lines.append(line)
        facts.append(line)
    dg = res["degisim"]
    if dg.get("satis") is not None:
        s = f"Günlük satış kampanyada önceki döneme göre {dg['satis']:.2f} kat."
        lines.append(s)
        facts.append(s)
    top = [k for k in res["kitaplar"][:5] if k.get("kampanyaAdet")]
    if top:
        s = "En çok satan kampanya kitapları: " + "; ".join(f"{k['ad'] or k['stok']} {k['kampanyaAdet']:.0f} adet" for k in top) + "."
        lines.append(s)
        facts.append(s)
    for n in res["notlar"]:
        lines.append(f"Not: {n}")
        facts.append(n)
    return "\n".join(lines), facts


def draft_summary(engine: sa.engine.Engine, st: dict[str, Any], llm: Any, tenant: str, cid: str) -> dict[str, Any]:
    if llm is None:
        raise KampanyaError("Zeki AI bu kurulumda tanımlı değil.", 503)
    camp = get_campaign(engine, tenant, cid)
    res = results(engine, st, tenant, cid)
    if not res["seri"]:
        raise KampanyaError("Önce sonuç okunmalı.", 409)
    text_facts, facts = summary_facts(camp, res)
    prompt = SUMMARY_PROMPT.format(name=camp["ad"], channel=camp["kanalAdi"], platform=f", {camp['platform']}" if camp.get("platform") else "",
                                   bas=_day_tr(camp["baslangic"]), bit=_day_tr(camp["bitis"]), n=camp["ozet"]["kitap"],
                                   ind=_pct(camp["ozet"]["ortalamaIndirim"]), kesim=_day_tr(res["logoKesim"]), facts=text_facts)
    raw = (llm.chat([{"role": "user", "content": prompt}], max_tokens=700, temperature=0.2) or "").strip()
    sources = [camp["ad"], camp["kanalAdi"], camp.get("platform") or "", *(i.get("ad") or "" for i in camp["kitaplar"])]
    g = mguard.check(raw, sources, facts + [camp["baslangic"], camp["bitis"], _day_tr(camp["baslangic"]), _day_tr(camp["bitis"]),
                                            _day_tr(res["logoKesim"]), str(camp["ozet"]["kitap"]), _pct(camp["ozet"]["ortalamaIndirim"])])
    text = g["metin"][:6000]
    if not text:
        raise KampanyaError("Zeki AI özet yazamadı; yeniden deneyin.", 503)
    with engine.begin() as c:
        c.execute(CAMPAIGNS.update().where(CAMPAIGNS.c.id == cid).values(sonuc_ozet=text, sonuc_at=_now(), updated_at=_now()))
    return {"ozet": text, "dusenSayisi": g["dusenSayisi"]}


def classify_rule(c: dict[str, Any]) -> Optional[str]:
    if (c.get("netIskonto") or 0) > 0 or (c.get("ekIskonto") or 0) > 0:
        return "iskonto"
    if (c.get("hediyeAdet") or 0) > 0:
        return "hediye"
    if c.get("vadeli"):
        return "vade"
    if "stant" in fold(c.get("ad")):
        return "stant"
    return None


def classify_crm(engine: sa.engine.Engine, st: dict[str, Any], llm: Any, camps: list[dict[str, Any]], budget_sec: int) -> dict[str, int]:
    """Bayi kampanyası türü: önce kural; kurala uymayan ve adı olan kayıt kapalı küme seçimle (eşik altı «belirsiz»)."""
    out = {"kural": 0, "model": 0, "belirsiz": 0, "kalan": 0}
    with engine.connect() as c:
        done = {r[0] for r in c.execute(sa.select(CRM_TYPES.c.crm_id)).all()}
    t0 = time.monotonic()
    labels = list(CRM_TYPE_CHOICES.values())
    back = {v: k for k, v in CRM_TYPE_CHOICES.items()}
    for x in camps:
        if x["id"] in done:
            continue
        tur, yontem, prob = classify_rule(x), "kural", None
        if tur is None:
            if llm is None or not hasattr(llm, "choose") or time.monotonic() - t0 >= budget_sec or not x.get("ad"):
                out["kalan"] += 1
                continue
            try:
                r = llm.choose(CRM_TYPE_PROMPT.format(name=x.get("ad") or "-", desc=x.get("aciklama") or "-",
                                                      disc="var" if (x.get("netIskonto") or x.get("ekIskonto")) else "yok",
                                                      vade="var" if x.get("vadeli") else "yok", gift=x.get("hediyeAdet") or 0), labels)
            except Exception as e:  # noqa: BLE001 — model düşerse kalan sonraki geceye
                log.warning("kampanya: CRM kampanya türü seçilemedi: %s", e)
                out["kalan"] += 1
                break
            ok = bool(r.choice and r.confident(st["llmMinProb"], min_margin=st["llmMinMargin"]))
            tur, yontem, prob = (back.get(r.choice, "belirsiz") if ok else "belirsiz"), "model", r.probability
        with engine.begin() as c:
            c.execute(CRM_TYPES.insert().values(crm_id=x["id"], tur=tur, yontem=yontem, olasilik=prob, asof=_now()))
        out["belirsiz" if tur == "belirsiz" else yontem] += 1
    return out


def crm_types(engine: sa.engine.Engine, ids: list[str]) -> dict[str, dict[str, Any]]:
    if not ids:
        return {}
    with engine.connect() as c:
        rows = c.execute(sa.select(CRM_TYPES).where(CRM_TYPES.c.crm_id.in_(ids))).all()
    return {r.crm_id: {"tur": r.tur, "turAdi": CRM_TURLERI.get(r.tur, r.tur), "yontem": r.yontem, "olasilik": r.olasilik} for r in rows}


# ------------------------------------------------------------------------------------------ dışa aktarım


def export_xlsx(camp: dict[str, Any], internal: bool) -> bytes:
    """Brif ve kitap listesi. `internal` = marj, maliyet ve kontroller de (yalnız iç kullanım). Onaysızsa TASLAK yazar."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Brif"
    taslak = camp["durum"] not in ("onaylandi", "yurutuluyor", "bitti")
    sel = (camp.get("metin") or {}).get("secili") or {}
    rows = [("Kampanya", camp["ad"]), ("Kanal", camp["kanalAdi"]), ("Platform", camp.get("platform") or ""),
            ("Başlangıç", _day_tr(camp["baslangic"])), ("Bitiş", _day_tr(camp["bitis"])), ("Durum", camp["durumAdi"]),
            ("Onaylayan", camp.get("onaylayan") or ""), ("Kitap sayısı", camp["ozet"]["kitap"]),
            ("Ortalama indirim", _pct(camp["ozet"]["ortalamaIndirim"])),
            ("Başlık", sel.get("baslik") or ""), ("Kısa açıklama", sel.get("aciklama") or ""), ("Banner", sel.get("banner") or ""),
            ("Notlar", camp.get("notlar") or "")]
    if taslak:
        ws.append(["TASLAK — onaylanmamış kampanya"])
        ws["A1"].font = Font(bold=True, color="B91C1C")
    for k, v in rows:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80
    kt = wb.create_sheet("Kitaplar")
    head = ["Stok kodu", "Barkod", "Kitap", "Yazar", "Liste fiyatı (KDV dahil)", "Kampanya fiyatı (KDV dahil)", "İndirim %"]
    if internal:
        head += ["Stok", "Tükenme tahmini", "Birim maliyet", "Maliyet kaynağı", "Birim telif (önce)", "Birim telif (kampanya)",
                 "Marj % (önce)", "Marj % (kampanya)", "Kontroller"]
    kt.append(head)
    for c in kt[1]:
        c.font = Font(bold=True)
    for i in camp["kitaplar"]:
        row = [i["stok"], i.get("ean") or "", i.get("ad") or "", i.get("yazar") or "", i.get("liste"), i.get("kampanyaFiyati"),
               round(i["indirim"] * 100, 1) if i.get("indirim") is not None else None]
        if internal:
            row += [i.get("stokAdet"), _day_tr(i.get("tukenme")) if i.get("tukenme") else "", _r2(i.get("birimMaliyet")),
                    i.get("maliyetKaynak") or "", _r2(i.get("telifOnce")), _r2(i.get("telifSonra")),
                    round(i["marjOraniOnce"] * 100, 1) if i.get("marjOraniOnce") is not None else "hesaplanamaz",
                    round(i["marjOraniSonra"] * 100, 1) if i.get("marjOraniSonra") is not None else "hesaplanamaz",
                    " | ".join(k["mesaj"] for k in i["kontroller"] if k.get("seviye") in ("kirmizi", "sari"))]
        kt.append(row)
    for col in range(1, len(head) + 1):
        kt.column_dimensions[get_column_letter(col)].width = 18 if col != 3 else 42
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------------------------ okuma (Logo + CRM → köprü)


class Refresher:
    """Gece okuması: Logo + CRM → `semantic_kampanya_books`, sezon bağları, site fiyat kaydı. Aynı anda tek okuma."""

    def __init__(self, engine_fn: Callable[[], sa.engine.Engine], tenant_fn: Callable[[], str], logo_file: Callable[[], str],
                 crm_file: Callable[[], str], schema_fn: Callable[[], str], settings_fn: Callable[[], dict[str, Any]]):
        self._engine, self._tenant, self._logo, self._crm, self._schema, self._settings = (
            engine_fn, tenant_fn, logo_file, crm_file, schema_fn, settings_fn)
        self._thread: Optional[threading.Thread] = None
        self._guard = threading.Lock()
        self.state: dict[str, Any] = {"running": False, "step": None, "startedAt": None, "error": None}

    def status(self) -> dict[str, Any]:
        engine = self._engine()
        ensure(engine)
        return {**self.state, "last": meta_get(engine, "refresh"), "dataEnd": meta_get(engine, "data_end").get("date"),
                "fiyatKaydi": meta_get(engine, "snapshots")}

    def running(self) -> bool:
        return bool(self._thread and self._thread.is_alive()) or bool(self.state.get("running"))

    def start(self) -> bool:
        with self._guard:
            if self.running():
                return False
            self.state.update(running=True, step="Başlıyor", startedAt=time.time(), error=None)
            self._thread = threading.Thread(target=self.run, daemon=True, name="kampanya-refresh")
            self._thread.start()
            return True

    def _step(self, name: str) -> None:
        self.state["step"] = name

    def logo(self) -> tuple[src.Runner, dict[int, str]]:
        from semantic_bridge import sorgu_yakala as Y

        run = Y.izle(src.runner(self._logo()), "logo", Y.db_of(self._logo()))
        return run, src.firms_by_year(run)

    def crm(self) -> src.Runner:
        from semantic_bridge import sorgu_yakala as Y

        return Y.izle(src.runner(self._crm()), "crm", Y.db_of(self._crm()))

    def run(self) -> dict[str, Any]:
        engine, tenant, st = self._engine(), self._tenant(), self._settings()
        ensure(engine)
        done: dict[str, Any] = {}
        warnings: list[str] = []
        t0 = time.monotonic()
        ref = today()
        from semantic_bridge import sorgu_yakala as Y

        q, token = Y.baslat(engine)   # sorgu bilgisi: kitap tablosunu dolduran CRM/Logo sorguları
        try:
            self.state.update(running=True, step="CRM kitap kartları", error=None)
            schema = self._schema()
            # CRM okunamazsa kitap tablosu yarım yazılmasın: okuma durur, eski tablo kalır, hata ekranda.
            crm = self.crm()
            crm_books = src.read_crm_books(crm, schema)
            self._step("CRM sözleşme sınırları")
            contracts = src.read_contracts(crm, schema)
            done.update(crmKitap=len(crm_books), sozlesmeliKitap=len(contracts))

            self._step("Logo dönemleri, malzeme kartları, stok ve fiyat")
            run, firms = self.logo()
            end = src.read_data_end(run, firms) or ref
            meta_set(engine, "data_end", {"date": end.isoformat()})
            items = src.read_items(run, firms)
            stock = src.read_stock(run, firms)
            prices = src.read_prices(run, firms, ref)

            self._step("Logo satış hızı")
            months = st["hizAy"]
            last = end.year * 12 + end.month - 1                   # veri sonu ayı dahil
            first_last = last - months + 1
            first_prev = first_last - months
            first_12 = last - 11
            first = min(first_prev, first_12)
            a = date(first // 12, first % 12 + 1, 1)
            monthly = src.read_monthly(run, firms, a, end + timedelta(days=1))
            agg: dict[str, dict[str, float]] = {}
            for m in monthly:
                ym = m["yil"] * 12 + m["ay"] - 1
                cur = agg.setdefault(m["stok"], {"son": 0.0, "onceki": 0.0, "a12": 0.0, "t12": 0.0})
                net = m["adet"] - m["iade"]
                if first_last <= ym <= last:
                    cur["son"] += net
                elif first_prev <= ym < first_last:
                    cur["onceki"] += net
                if first_12 <= ym <= last:
                    cur["a12"] += net
                    cur["t12"] += m["tutar"]
            win_start = date(first_last // 12, first_last % 12 + 1, 1)
            win_days = max(1, (end - win_start).days + 1)
            done["satisSatiri"] = len(monthly)

            self._step("Logo güncel yıl brüt farkı")
            margins = src.read_margins(run, firms, end.year, end)

            self._step("Sezon bağları ve haklar")
            seasons_by_ean, rights = self._seo_links(engine, tenant)
            Y.bitir(token)
            Y.koken_yaz(engine, tenant, KOKEN_OKUMA, q, portal_tables=("semantic_seo_seasons_books", "semantic_seo_crm_books"))
            self._step("Köprü tablosuna yazılıyor")
            done["kitap"] = self._write_books(engine, items, stock, prices, agg, win_days, margins, end.year, crm_books, contracts,
                                              seasons_by_ean, rights)
            self._step("Site fiyat kaydı")
            try:
                done["fiyatKaydi"] = snapshot_prices(engine, tenant, ref)
            except Exception as e:  # noqa: BLE001 — fiyat kaydı yoksa kural «denetlenemedi» der
                warnings.append(f"Site fiyat kaydı alınamadı: {e}")
            meta_set(engine, "window", {"hizAy": months, "son": [win_start.isoformat(), end.isoformat()], "gun": win_days})
            meta_set(engine, "refresh", {"ok": True, "done": done, "warnings": warnings, "sn": round(time.monotonic() - t0, 1)})
            self.state.update(running=False, step=None, error=None, finishedAt=time.time())
            return {"ok": True, "done": done, "warnings": warnings, "dataEnd": end.isoformat()}
        except Exception as e:  # noqa: BLE001 — eski veriler kalır, hata ekranda
            log.warning("kampanya refresh failed: %s", e)
            Y.bitir(token)
            msg = str(e) if isinstance(e, (src.SourceError, KampanyaError)) else f"Okuma hata verdi: {str(e)[:200]}"
            meta_set(engine, "refresh", {"ok": False, "error": msg, "done": done, "warnings": warnings})
            self.state.update(running=False, step=None, error=msg, finishedAt=time.time())
            return {"ok": False, "error": msg, "done": done}

    @staticmethod
    def _seo_links(engine: sa.engine.Engine, tenant: str) -> tuple[dict[str, list[str]], dict[str, tuple[Optional[str], Optional[str]]]]:
        """SEO modülünün tablolarından: barkod → bağlı özel gün anahtarları; barkod → (hak, durum işareti). Yoksa boş."""
        by_ean: dict[str, list[str]] = {}
        rights: dict[str, tuple[Optional[str], Optional[str]]] = {}
        try:
            from semantic_bridge.seo_geo import seasons as S
            from semantic_bridge.seo_geo.store import CRM_BOOKS

            insp = sa.inspect(engine)
            with engine.connect() as c:
                if insp.has_table(S.BOOKS.name):
                    for day_key, ean in c.execute(sa.select(S.BOOKS.c.day_key, S.BOOKS.c.ean).where(S.BOOKS.c.tenant_id == tenant)).all():
                        by_ean.setdefault(ean, []).append(day_key)
                if insp.has_table(CRM_BOOKS.name):
                    for ean, hak, flag in c.execute(sa.select(CRM_BOOKS.c.ean, CRM_BOOKS.c.rights, CRM_BOOKS.c.status_flag)
                                                    .where(CRM_BOOKS.c.tenant_id == tenant)).all():
                        rights[ean] = (hak, flag)
        except Exception as e:  # noqa: BLE001
            log.warning("kampanya: SEO tabloları okunamadı: %s", e)
        return by_ean, rights

    @staticmethod
    def _write_books(engine: sa.engine.Engine, items: dict[str, dict[str, Any]], stock: dict[str, float],
                     prices: dict[str, dict[str, Any]], agg: dict[str, dict[str, float]], win_days: int,
                     margins: dict[str, dict[str, float]], year: int, crm_books: dict[str, dict[str, Any]],
                     contracts: dict[str, list[dict[str, Any]]], seasons_by_ean: dict[str, list[str]],
                     rights: dict[str, tuple[Optional[str], Optional[str]]]) -> int:
        now = _now()
        codes = set(items) | set(crm_books)
        batch = []
        for code in codes:
            it, cb = items.get(code) or {}, crm_books.get(code) or {}
            kdv = it.get("kdv") or 0.0
            p = prices.get(code)
            liste_logo = None
            if p:
                liste_logo = p["fiyat"] if p["kdvDahil"] else p["fiyat"] * (1 + kdv / 100.0)
            a = agg.get(code) or {}
            m = margins.get(code)
            ks = contracts.get(code) or []
            floors = [k["asgari"] for k in ks if k.get("asgari")]
            ean = cb.get("ean")
            hak, flag = rights.get(ean or "", (None, None))
            batch.append({
                "stok_kodu": code, "ad": cb.get("ad") or it.get("ad"), "ean": ean, "crm_id": cb.get("id"), "crm_tip": cb.get("tip"),
                "yazar": cb.get("yazar"), "yayin": cb.get("yayin"), "tsoft": cb.get("tsoft"), "liste_crm": cb.get("fiyat"),
                "liste_logo": _r2(liste_logo), "kdv": kdv, "stok": stock.get(code, 0.0), "adet_son": a.get("son", 0.0),
                "adet_onceki": a.get("onceki", 0.0), "adet_12": a.get("a12", 0.0), "tutar_12": a.get("t12", 0.0),
                "gunluk_hiz": round(max(0.0, a.get("son", 0.0)) / win_days, 4),
                "ciro_yil": m["ciro"] if m else None, "maliyet_yil": m["maliyet"] if m else None,
                "maliyetli_ciro": m["maliyetli_ciro"] if m else None, "maliyetli_adet": m["maliyetli_adet"] if m else None,
                "maliyetsiz_satir": int(m["maliyetsiz_satir"]) if m else None, "satir_yil": int(m["satir"]) if m else None,
                "marj_yili": year if m else None, "asgari_fiyat": max(floors) if floors else None,
                "sozlesme_json": _dump(ks) if ks else None, "hak": hak, "durum_bayragi": flag,
                "sezon_json": _dump(sorted(set(seasons_by_ean.get(ean or "", [])))) if ean and seasons_by_ean.get(ean) else None,
                "in_logo": code in items, "in_crm": code in crm_books, "asof": now})
        with engine.begin() as c:
            c.execute(BOOKS.delete())
            for i in range(0, len(batch), 1000):
                if batch[i:i + 1000]:
                    c.execute(BOOKS.insert(), batch[i:i + 1000])
        return len(batch)


def snapshot_prices(engine: sa.engine.Engine, tenant: str, ref: date) -> dict[str, Any]:
    """Günün site fiyatı: SEO modülünün T-soft ürün tablosundan (yalnız okunur) barkod başına fiyat ve indirimli fiyat."""
    from semantic_bridge.seo_geo import shopping
    from semantic_bridge.seo_geo.store import PRODUCTS

    if not sa.inspect(engine).has_table(PRODUCTS.name):
        return {"urun": 0, "neden": "site ürün tablosu yok"}
    with engine.connect() as c:
        rows = c.execute(sa.select(PRODUCTS.c.code, PRODUCTS.c.data_json).where(PRODUCTS.c.tenant_id == tenant,
                                                                                PRODUCTS.c.active.is_(True))).all()
        codes_by_ean = {e: s for s, e in c.execute(sa.select(BOOKS.c.stok_kodu, BOOKS.c.ean).where(BOOKS.c.ean.isnot(None))).all()}
    day_ = ref.isoformat()
    now = _now()
    batch: dict[str, dict[str, Any]] = {}
    for code, data in rows:
        p = _j(data, {})
        if not isinstance(p, dict):
            continue
        ean = src.ean_of(shopping.pick(p, shopping.K_BARCODE))
        if not ean:
            continue
        price, sale = shopping.price_of(p)
        if price is None and sale is None:
            continue
        batch[ean] = {"tarih": day_, "product_key": ean, "kaynak": "tsoft", "stok_kodu": codes_by_ean.get(ean) or code,
                      "fiyat": price, "indirimli": sale, "asof": now}
    with engine.begin() as c:
        c.execute(SNAPS.delete().where(SNAPS.c.tarih == day_, SNAPS.c.kaynak == "tsoft"))
        vals = list(batch.values())
        for i in range(0, len(vals), 1000):
            if vals[i:i + 1000]:
                c.execute(SNAPS.insert(), vals[i:i + 1000])
    with engine.connect() as c:
        days = c.execute(sa.select(sa.func.count(sa.distinct(SNAPS.c.tarih))).where(SNAPS.c.kaynak == "tsoft")).scalar() or 0
        first = c.execute(sa.select(sa.func.min(SNAPS.c.tarih)).where(SNAPS.c.kaynak == "tsoft")).scalar()
    info = {"urun": len(batch), "gun": int(days), "ilkGun": first, "sonGun": day_}
    meta_set(engine, "snapshots", info)
    return info


def summary(engine: sa.engine.Engine, tenant: str) -> dict[str, Any]:
    with engine.connect() as c:
        counts = {k: int(v) for k, v in c.execute(sa.select(CAMPAIGNS.c.durum, sa.func.count()).where(CAMPAIGNS.c.tenant_id == tenant)
                                                  .group_by(CAMPAIGNS.c.durum)).all()}
        books = c.execute(sa.select(sa.func.count()).select_from(BOOKS)).scalar() or 0
    return {"sayilar": {k: counts.get(k, 0) for k in DURUMLAR}, "kitapSayisi": int(books)}
