"""M43 Depo ve stok: yeterlilik ve Baskı Öneri eşliği, durum kuralları, Logo–CRM fark sınıfı, SQL yer tutucuları ve
içe aktarılan Baskı Öneri kaynakları, Logo okuması (ambar toplamı, yıl kopyaları), eşik/öneri/not/fotoğraf kayıtları,
aktarım hatası ve Zeki AI sınıflaması, depo hattı ve kişi yetkisi, uçlar (maliyet ve onay yetkisi, Excel, gece işi),
yetki kuralları.

Veriler yapaydır ve yalnız hesap kurallarını sınar; gerçek Logo/CRM kabulü test sunucusunda (`scripts/acceptance/M43/`).
Yerelde koşulmaz (AGENTS.md).
"""
from __future__ import annotations

import io
import math
from datetime import date, datetime, timedelta

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from semantic_bridge import access as A
from semantic_bridge import stock as S
from semantic_bridge import stock_api
from semantic_bridge import stock_sources as src
from semantic_bridge import stock_store as store
from semantic_bridge.management import baski_oneri
from semantic_layer.runtime.llm_choose import Choice
from semantic_layer.store.catalog_store import open_store

T = "t1"
SETTINGS = {"runoutDays": 30, "safetyDays": 15, "leadDays": 40, "excessDays": 730, "deadDays": 365, "excludePlanned": True,
            "excludePrefixes": ["157"], "model": True, "minProb": 0.7, "minMargin": 0.3, "modelBudget": 60, "pickDays": 30,
            "recipients": []}


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    store._ready.discard(id(e))
    store.ensure(e)
    return e


def raw_fixture() -> dict:
    """Yedi kitap, her durumdan biri; 157 ticari ürün listeye girmez."""
    return {
        "at": 1_000_000.0, "dataEnd": "2026-08-17",
        "balances": {"B-BIT": {"ad": "Bitecek", "bakiye": 100.0}, "B-YOK": {"ad": "Stoksuz", "bakiye": -2.0},
                     "B-OLU": {"ad": "Hareketsiz", "bakiye": 500.0}, "B-SATSIZ": {"ad": "Satışsız", "bakiye": 50.0},
                     "B-FAZ": {"ad": "Fazla", "bakiye": 10_000.0}, "B-YET": {"ad": "Yeterli", "bakiye": 600.0},
                     "B-PAS": {"ad": "Pasif", "bakiye": 0.0}, "15701.01.0001": {"ad": "Ajanda", "bakiye": 40.0}},
        "warehouse": {"B-BIT": {"1": 60.0, "2": 40.0}, "B-FAZ": {"1": 10_000.0}},
        "warehouses": {"1": {"ad": "Merkez"}, "2": {"ad": "Perakende"}},
        "speeds": {"B-BIT": {"satis_hizi": 150.0, "yillik_toplam": 1800}, "B-YOK": {"satis_hizi": 20.0},
                   "B-FAZ": {"satis_hizi": 100.0}, "B-YET": {"satis_hizi": 100.0}, "B-OLU": {"satis_hizi": 0},
                   "B-SATSIZ": {"satis_hizi": 0}},
        "movement": {"B-BIT": {"son": "2026-08-10", "netSatis": 1500}, "B-YOK": {"son": "2026-08-01", "netSatis": 200},
                     "B-SATSIZ": {"son": "2026-03-01", "netSatis": 0}, "B-FAZ": {"son": "2026-08-01", "netSatis": 1200},
                     "B-YET": {"son": "2026-08-01", "netSatis": 1200}},
        "turnover": {"B-BIT": 1.09}, "eos": {"B-BIT": 98.0}, "ordersLogo": {"B-BIT": {"adet": 30, "siparis": 2}},
        "ordersCrm": {"B-BIT": 25.0}, "waiting": {"B-YOK": {"adet": 7, "kayit": 3}},
        "books": {"B-BIT": {"ad": "Bitecek Kitap", "yazar": "Yazar A", "yayinevi": "Timaş"}},
        "shelves": [{"stokKodu": "B-BIT", "raf": "A-1", "depo": "Merkez", "depoNo": "1", "adet": 80.0},
                    {"stokKodu": "B-BIT", "raf": "P-1", "depo": "Perakende", "depoNo": "2", "adet": 30.0},
                    {"stokKodu": "B-YET", "raf": "A-2", "depo": "Merkez", "depoNo": "1", "adet": 580.0},
                    {"stokKodu": "B-CRM", "raf": "A-3", "depo": "Merkez", "depoNo": "1", "adet": 5.0}],
        "transferItems": {"B-BIT": {"net": 10.0, "fis": 1}, "B-YET": {"net": -5.0, "fis": 1}},
        "transfers": [{"id": "h1", "fisNo": "MH-1", "fisTarihi": "2026-09-20", "islemTuru": 6, "islemTipi": 2, "durum": 1,
                       "depo": "Merkez", "hata": True, "mesaj": "Fiş 1234: dönem kapalı", "satir": 2, "miktar": 10},
                      {"id": "h2", "fisNo": "MH-2", "fisTarihi": "2026-09-25", "islemTuru": 1, "islemTipi": 1, "durum": 1,
                       "depo": "Merkez", "hata": False, "mesaj": None, "satir": 1, "miktar": 5}],
        "pick": [], "depots": [], "warnings": [], "readMs": 5,
    }


def model(thresholds=None, cards=None, s=None):
    return S.build(raw_fixture(), s or SETTINGS, thresholds or {}, cards or [], {})


# ------------------------------------------------------------------ hesap


def test_days_of_cover_and_baski_oneri_parity():
    assert S.days_of_cover(100, 150) == pytest.approx(20.0)
    assert S.days_of_cover(-3, 10) == 0.0
    assert S.days_of_cover(10, 0) is None and S.days_of_cover(10, None) is None
    # Tükenme süresi ve öneri Baskı Öneri'nin kendi işlevleriyle aynı (içe aktarma, kopya değil).
    for stok, hiz in ((100.0, 150.0), (10.0, 0.0), (0.0, 0.0), (500.0, 20.0)):
        tuk, label = S.baski_label(stok, hiz)
        assert tuk == baski_oneri._out(baski_oneri._dax_div(stok, hiz))
        assert label == baski_oneri._marj_oneri(stok, hiz)[1]
    assert S.baski_label(10.0, 0.0) == (baski_oneri.INF, "Yeterli Stok")


def test_states_and_item_fields():
    m = model()
    by = m["byCode"]
    assert "15701.01.0001" not in by                                   # ticari ürün öneki listede yok
    assert by["B-BIT"]["durum"] == "bitecek" and by["B-BIT"]["gun"] == 20.0
    assert by["B-BIT"]["tukenmeTarihi"] == (date(2026, 8, 17) + timedelta(days=20)).isoformat()
    assert sum(a["adet"] for a in by["B-BIT"]["ambarlar"]) == by["B-BIT"]["bakiye"]
    assert [a["ad"] for a in by["B-BIT"]["ambarlar"]] == ["Merkez", "Perakende"]
    assert by["B-BIT"]["crmRaf"] == 110 and by["B-BIT"]["rafSayisi"] == 2
    assert by["B-BIT"]["bekleyenCrm"] == 25 and by["B-BIT"]["bekleyenLogo"] == 30 and by["B-BIT"]["devirHizi"] == 1.09
    assert by["B-BIT"]["ad"] == "Bitecek Kitap" and by["B-BIT"]["kritikGun"] == 55 and by["B-BIT"]["baskiUyarisi"]
    assert by["B-YOK"]["durum"] == "stoksuz" and by["B-YOK"]["gun"] == 0.0 and by["B-YOK"]["bekleyenUrun"] == 7
    assert by["B-OLU"]["durum"] == "olu" and by["B-OLU"]["netSatis12"] == 0.0
    assert by["B-SATSIZ"]["durum"] == "satissiz"
    assert by["B-FAZ"]["durum"] == "fazla" and by["B-FAZ"]["gun"] == 3000.0
    assert by["B-YET"]["durum"] == "yeterli"
    assert by["B-PAS"]["durum"] == "pasif"
    assert by["B-CRM"]["logoVar"] is False and by["B-CRM"]["bakiye"] == 0.0 and by["B-CRM"]["farkSinif"] == "logo_yok"


def test_movement_unknown_does_not_mark_dead():
    raw = raw_fixture()
    raw["movement"] = None                                          # okunamadı: «hareketsiz» denmez
    m = S.build(raw, SETTINGS, {}, [], {})
    assert m["byCode"]["B-OLU"]["durum"] == "satissiz" and m["byCode"]["B-OLU"]["netSatis12"] is None


def test_threshold_and_lead_change_the_critical_line():
    thr = {"B-YET": {"id": "x", "guvenlikGun": 170, "yenidenSiparisAdet": 700.0, "onaylayan": "a", "onayTarihi": None}}
    m = model(thresholds=thr)
    it = m["byCode"]["B-YET"]                                     # 180 gün yeter; 40 + 170 = 210 → kritik
    assert it["kritikGun"] == 210 and it["durum"] == "bitecek" and it["rspAltinda"] is True
    s = {**SETTINGS, "leadDays": None}
    assert S.build(raw_fixture(), s, {}, [], {})["lead"] == S.FALLBACK_LEAD


def test_open_production_card_and_forecast():
    cards = [{"id": "c1", "stockCode": "B-BIT", "stage": "matbaada", "stageLabel": "Dosya matbaada", "qty": 3000,
              "plan": {"depo": "2026-10-10", "baski": "2026-10-01"}, "created": "2026-09-01", "name": "URT-1"},
             {"id": "c2", "stockCode": "B-BIT", "stage": "tamam", "plan": {}, "created": "2026-01-01"}]
    m = S.build(raw_fixture(), SETTINGS, {}, cards, {"start": "2026-09-01", "p50": {"B-BIT": [100, 120, 130, 90]}})
    it = m["byCode"]["B-BIT"]
    assert it["uretim"]["kartId"] == "c1" and it["uretim"]["depoPlan"] == "2026-10-10" and it["uretimSayisi"] == 1
    assert it["tahmin"] == {"baslangic": "2026-09-01", "g30": 100, "g60": 220, "g90": 350}


def test_lists():
    m = model()
    run = S.running_out(m, 30)
    assert [i["stokKodu"] for i in run] == ["B-YOK", "B-BIT"]
    assert [i["stokKodu"] for i in S.excess(m)] == ["B-FAZ", "B-OLU", "B-SATSIZ"]
    assert [i["stokKodu"] for i in S.excess(m, "olu")] == ["B-OLU"]
    assert S.filter_items(m, q="bitecek kitap")[0]["stokKodu"] == "B-BIT"
    assert [i["stokKodu"] for i in S.filter_items(m, depo="2")] == ["B-BIT"]
    assert all(i["durum"] != "pasif" for i in S.filter_items(m))
    pg = S.page_of(list(range(120)), 2)
    assert pg["items"] == list(range(100, 120)) and pg["total"] == 120


def test_diff_classes():
    assert S.classify_diff(100, 110, 10) == ("aktarim", 10)
    assert S.classify_diff(100, 110, 4) == ("kismen", 10)
    assert S.classify_diff(100, 110, None) == ("sayim", 10)
    assert S.classify_diff(100, None, None) == ("crm_yok", -100)
    assert S.classify_diff(None, 5, None) == ("logo_yok", 5)
    assert S.classify_diff(100, 100.2, None) is None
    m = model()
    assert m["byCode"]["B-BIT"]["farkSinif"] == "aktarim"
    assert m["byCode"]["B-YET"]["farkSinif"] == "kismen" and m["byCode"]["B-YET"]["fark"] == -20   # −5 aktarılmamış
    assert S.diff_rows(m)[0]["stokKodu"] == "B-FAZ"                       # en büyük mutlak fark üstte


def test_threshold_proposal_is_arithmetic():
    it = model()["byCode"]["B-BIT"]
    p = S.threshold_proposal(it, 40, 15)
    assert p["yenidenSiparisAdet"] == math.ceil(150 / 30 * 55) and p["guvenlikGun"] == 15
    assert S.threshold_proposal(model()["byCode"]["B-OLU"], 40, 15) is None


# ------------------------------------------------------------------ SQL ve okuma


def test_render_placeholders_and_imported_sources():
    sql = src.render("logo_bakiye", firm="411", exclude_planned=True)
    assert "LG_411_01_STLINE" in sql and "PRODSTAT, 0) = 1" in sql and "{" not in sql
    off = src.render("logo_bakiye", firm="411", exclude_planned=False)
    assert "ISNULL(F.PRODSTAT" not in off and "STFICHE AS F" not in off and "{" not in off
    assert "FIRMNR = 411" in src.render("logo_ambarlar", firm="411")
    assert "Timas_MSCRM.dbo.new_serilothareketsatiriBase" in src.render("crm_raf_stok", schema="Timas_MSCRM.dbo")
    with pytest.raises(src.SourceError):
        src.render("logo_bakiye", firm="41; DROP")
    with pytest.raises(src.SourceError):
        src.render("crm_raf_stok", schema="x; drop")
    h = src.render("logo_hareket", firm="211", start=date(2025, 8, 17), end=date(2026, 1, 1))
    assert "'2025-08-17'" in h and "'2026-01-01'" in h and "TRCODE <> 14" in h
    from semantic_bridge import management
    for sid in ("logo_satis_hizi", "logo_depo_stok", "crm_bekleyen_siparis"):
        assert src.sql_text(f"baski-oneri:{sid}") == management.sql_text("baski-oneri", sid)
    for sid, *_ in src.SOURCES:
        assert src.sql_text(sid).strip()
    # Kişisel/gizli kolon hiçbir sorguda seçilmez.
    for sid, *_ in src.SOURCES:
        assert "deposifre" not in src.sql_text(sid).lower() and "depokullaniciadi" not in src.sql_text(sid).lower()


class FakeRun:
    def __init__(self):
        self.seen: list[str] = []

    def __call__(self, sql: str):
        self.seen.append(sql)
        if "L_CAPIPERIOD" in sql:
            return [{"FIRMNR": 211, "BEGDATE": date(2021, 1, 1), "ENDDATE": date(2025, 12, 31)},
                    {"FIRMNR": 411, "BEGDATE": date(2026, 1, 1), "ENDDATE": date(2026, 12, 31)}]
        if "SOURCEINDEX AS ambar_no" in sql:
            return [{"stok_kodu": "K1", "ad": "Kitap", "ambar_no": 1, "bakiye": 70},
                    {"stok_kodu": "K1", "ad": "Kitap", "ambar_no": 2, "bakiye": 30}]
        if "son_hareket" in sql:
            year = "211" if "LG_211_" in sql else "411"
            return [{"stok_kodu": "K1", "son_hareket": date(2025, 11, 3) if year == "211" else date(2026, 2, 1), "net_satis": 5}]
        return []


def test_logo_reader_sums_warehouses_and_reads_each_year_copy():
    run = FakeRun()
    lg = src.Logo(run)
    items, wh = lg.balances()
    assert items["K1"]["bakiye"] == 100 and wh["K1"] == {1: 70.0, 2: 30.0}
    mv = lg.movement(date(2025, 8, 17), date(2026, 8, 18))
    assert mv["K1"] == {"son": date(2026, 2, 1), "netSatis": 10.0}
    assert sum("LG_211_01_STLINE" in s for s in run.seen) == 1 and any("'2025-08-17'" in s for s in run.seen)


# ------------------------------------------------------------------ kayıtlar


def test_thresholds_draft_approve_archive_and_reject(engine):
    a = store.save_threshold(engine, T, "ayse", {"stokKodu": "B-BIT", "guvenlikGun": 20, "yenidenSiparisAdet": "1.500"})
    assert a["durum"] == "taslak" and a["yenidenSiparisAdet"] == 1500                  # Türkçe yazım: 1.500
    b = store.save_threshold(engine, T, "ayse", {"stokKodu": "B-BIT", "guvenlikGun": 25})
    assert [t["id"] for t in store.thresholds(engine, T, "taslak")] == [b["id"]]      # taslak yenisiyle değişir
    ok = store.decide_threshold(engine, T, "mehmet", b["id"], True)
    assert ok["durum"] == "onayli" and ok["onaylayan"] == "mehmet"
    c = store.save_threshold(engine, T, "ayse", {"stokKodu": "B-BIT", "guvenlikGun": 30})
    store.decide_threshold(engine, T, "mehmet", c["id"], True)
    assert store.approved_thresholds(engine, T)["B-BIT"]["guvenlikGun"] == 30
    assert [t["durum"] for t in store.thresholds(engine, T, codes=["B-BIT"])].count("arsiv") == 1
    d = store.save_threshold(engine, T, "ayse", {"stokKodu": "B-BIT", "guvenlikGun": 5})
    with pytest.raises(store.StockError):
        store.decide_threshold(engine, T, "mehmet", d["id"], False)                # ret gerekçesiz olmaz
    assert store.decide_threshold(engine, T, "mehmet", d["id"], False, "sezon")["durum"] == "red"
    with pytest.raises(store.StockError) as e:
        store.decide_threshold(engine, T, "mehmet", d["id"], True)
    assert e.value.status == 409
    for bad in ({"stokKodu": "B", "guvenlikGun": -1}, {"stokKodu": "B", "guvenlikGun": "x"}, {"stokKodu": "", "guvenlikGun": 3},
                {"stokKodu": "B", "guvenlikGun": 3, "yenidenSiparisAdet": -4}):
        with pytest.raises(store.StockError):
            store.save_threshold(engine, T, "ayse", bad)


def test_suggestions_decide_and_expire(engine):
    r1 = store.add_suggestion(engine, T, "bitecek", "K1", {"gun": 3}, "neden", "M12")
    store.add_suggestion(engine, T, "bitecek", "K2", {"gun": 9}, "neden", "M12")
    assert set(store.open_suggestions(engine, T, "bitecek")) == {"K1", "K2"}
    assert store.expire_suggestions(engine, T, "bitecek", {"K1"}) == 1
    assert store.list_suggestions(engine, T, durum="gecersiz")["total"] == 1
    with pytest.raises(store.StockError):
        store.decide_suggestion(engine, T, "ayse", r1, "red", "")
    out = store.decide_suggestion(engine, T, "ayse", r1, "kabul", None)
    assert out["durum"] == "kabul" and out["hedef"] == "M12"
    assert store.decided_codes(engine, T, "bitecek", datetime(2000, 1, 1)) == {"K1"}
    assert store.list_suggestions(engine, T, hedef="M12", durum="")["total"] == 2


def test_notes_and_snapshot(engine):
    n = store.add_note(engine, T, "ayse", "Ayşe", "B-BIT", {"not": "Sayımda 5 eksik"})
    with pytest.raises(store.StockError) as e:
        store.delete_note(engine, T, "mehmet", False, n["id"])
    assert e.value.status == 403
    store.delete_note(engine, T, "mehmet", True, n["id"])
    assert store.notes(engine, T, "B-BIT") == []
    rows = [{"stokKodu": "B-BIT", "bakiye": 100.0, "crmRaf": 110.0, "satisHizi": 150.0}]
    store.write_snapshot(engine, T, date(2026, 9, 28), rows)
    store.write_snapshot(engine, T, date(2026, 9, 28), rows)                           # aynı gün tekrar: çift satır yok
    assert len(store.snapshots(engine, T, "B-BIT")) == 1


# ------------------------------------------------------------------ aktarım, Zeki AI, depo hattı


class FakeLlm:
    def __init__(self, pick: str, p: float = 0.9, margin: float = 0.8, fail: bool = False):
        self.pick, self.p, self.margin, self.fail, self.calls = pick, p, margin, fail, 0

    def choose(self, prompt, choices, system=None):
        self.calls += 1
        if self.fail:
            raise RuntimeError("model yok")
        pick = self.pick if self.pick in choices else choices[-1]
        probs = {c: (self.p if c == pick else (1 - self.p) / (len(choices) - 1)) for c in choices}
        return Choice(pick, choices.index(pick), probs, "logprobs", margin=self.margin)

    def chat(self, messages, **kw):
        return "Bugün 99 kitap bitiyor."                                                 # olgu dışı sayı → kural metni


def test_transfer_rows_and_message_classes():
    m = model()
    classes, left = S.classify_messages(FakeLlm("Dönem kapalı"), [t["mesaj"] for t in m["transfers"] if t["hata"]], {}, SETTINGS,
                                        deadline=10**12)
    assert left == 0 and classes[S.message_key("Fiş 9: dönem kapalı")]["sinif"] == "Dönem kapalı"
    rows = S.transfer_rows(m, classes, "", now=date(2026, 9, 28))
    assert [r["fisNo"] for r in rows] == ["MH-1", "MH-2"] and rows[0]["yasGun"] == 8
    assert rows[0]["sinif"] == "Dönem kapalı" and rows[0]["islemTuruEtiket"] == "Depolar arası sevk"
    assert [r["fisNo"] for r in S.transfer_rows(m, classes, "bekliyor")] == ["MH-2"]
    unsure, _ = S.classify_messages(FakeLlm("Diğer", p=0.5, margin=0.1), ["yeni hata"], {}, SETTINGS, deadline=10**12)
    assert unsure[S.message_key("yeni hata")]["sinif"] is None                         # emin değil: sınıf yazılmaz
    _, left = S.classify_messages(FakeLlm("Diğer", fail=True), ["a", "b"], {}, SETTINGS, deadline=10**12)
    assert left == 2


def test_bulletin_rejects_model_numbers():
    m = model()
    ov = S.overview(m, S.transfer_rows(m, {}))
    text, how = S.bulletin(m, ov, FakeLlm("x"), True)
    assert how == "kural" and "99" not in text and "Logo verisinin son günü" in text


def test_pick_line_spans_and_people_gate():
    now = datetime(2026, 9, 28, 12, 0)
    m = model()
    m["pick"] = [
        {"id": "s1", "no": "S1", "durum": 100000011, "depoda": "2026-09-28 06:00:00", "toplayan": "Ali"},
        {"id": "s2", "no": "S2", "durum": 100000000, "depoda": "2026-09-20 08:00:00", "pusula": "2026-09-20 10:00:00",
         "kutulandi": "2026-09-20 14:00:00", "sevk": "2026-09-21 08:00:00", "toplayan": "Ali"},
    ]
    out = S.pick_line(m, SETTINGS, with_people=False, now=now)
    st = {a["key"]: a for a in out["asamalar"]}
    assert st["depoda"]["adet"] == 1 and st["depoda"]["enEskiSaat"] == 6.0
    spans = {x["from"] + ">" + x["to"]: x for x in out["sureler"]}
    assert spans["pusula>kutulandi"]["ortancaSaat"] == 4.0 and spans["depoda>sevk"]["ortancaSaat"] == 24.0
    assert "kisiler" not in out and "toplayan" not in out["acik"][0]
    full = S.pick_line(m, SETTINGS, with_people=True, now=now)
    assert full["kisiler"] == [{"toplayan": "Ali", "siparis": 1, "ortancaSaat": 4.0}]


# ------------------------------------------------------------------ uçlar


class FakeCosts:
    def birim(self, codes):
        return {k: {"birim": 12.5, "tarih": "2026-06-30", "kaynak": "gerceklesen"} for k in codes if k in ("B-BIT", "B-FAZ")}


def _client(engine, monkeypatch, perms: set[str], llm=None):
    audits: list[tuple] = []
    app = FastAPI()
    users = {"a": ("ayse", "Ayşe")}

    def auth(request):
        u = users.get(request.headers.get("x-user", ""))
        if not u:
            from fastapi import HTTPException
            raise HTTPException(status_code=401)
        return engine, T, u[0], u[1]

    deps = {"auth": auth, "require_caller": lambda r: None, "can": lambda u, k: k in perms, "is_admin": lambda u: False,
            "audit": lambda *a: audits.append(a), "conf": lambda k, d="": d, "fresh": lambda: False,
            "engine": lambda: engine, "tenant": lambda: T, "logo_file": lambda: "", "crm_file": lambda: "",
            "llm": lambda priority: llm, "m12": lambda: None, "costs": lambda: FakeCosts()}
    svc = stock_api.register(app, deps)
    monkeypatch.setattr(svc, "raw", lambda fresh=False, **kw: raw_fixture())
    monkeypatch.setattr(svc, "settings", lambda: SETTINGS)
    return TestClient(app), audits, svc


def test_api_reads_and_cost_gate(engine, monkeypatch):
    c, _, _ = _client(engine, monkeypatch, set())
    h = {"x-user": "a"}
    ov = c.get("/api/v1/stock/overview", headers=h).json()
    assert ov["bitecek"] == 2 and ov["aktarimHatasi"] == 1 and "deger" not in ov
    items = c.get("/api/v1/stock/items?durum=bitecek", headers=h).json()
    assert [i["stokKodu"] for i in items["items"]] == ["B-BIT"] and "stokDegeri" not in items["items"][0]
    it = c.get("/api/v1/stock/items/B-BIT", headers=h).json()
    assert [r["raf"] for r in it["raflar"]] == ["A-1", "P-1"] and it["esikOnerisi"]["yenidenSiparisAdet"] == 275
    assert c.get("/api/v1/stock/items/YOK-1", headers=h).status_code == 404
    assert c.get("/api/v1/stock/running-out?gun=30", headers=h).json()["total"] == 2
    assert c.get("/api/v1/stock/excess?tur=x", headers=h).status_code == 400
    assert c.get("/api/v1/stock/diff?sinif=aktarim", headers=h).json()["total"] == 1
    assert c.get("/api/v1/stock/transfer-errors", headers=h).json()["total"] == 1
    assert c.get("/api/v1/stock/names", headers=h).json()["items"]
    c2, _, _ = _client(engine, monkeypatch, {A.page("stok"), stock_api.FEATURE_COST})
    ex = c2.get("/api/v1/stock/excess?tur=fazla", headers=h).json()
    assert ex["items"][0]["stokDegeri"] == 125_000 and ex["deger"] == {"toplam": 125_000, "maliyetli": 1, "maliyetsiz": 0}


def test_api_threshold_needs_explicit_approval(engine, monkeypatch):
    c, audits, _ = _client(engine, monkeypatch, {stock_api.FEATURE_DECIDE})
    h = {"x-user": "a"}
    d = c.post("/api/v1/stock/thresholds", json={"stokKodu": "B-BIT", "guvenlikGun": 20}, headers=h)
    assert d.status_code == 201 and d.json()["durum"] == "taslak"
    assert c.post(f"/api/v1/stock/thresholds/{d.json()['id']}/approve", json={}, headers=h).status_code == 403
    c2, audits2, svc = _client(engine, monkeypatch, {stock_api.FEATURE_DECIDE, stock_api.FEATURE_THRESHOLD})
    ok = c2.post(f"/api/v1/stock/thresholds/{d.json()['id']}/approve", json={}, headers=h)
    assert ok.status_code == 200 and ok.json()["durum"] == "onayli"
    assert [a[2] for a in audits + audits2] == ["create", "approve"]
    assert c2.get("/api/v1/stock/items/B-BIT", headers=h).json()["esik"]["guvenlikGun"] == 20
    props = c2.get("/api/v1/stock/thresholds?durum=oneri", headers=h).json()
    assert "B-BIT" not in [p["stokKodu"] for p in props["items"]]


def test_api_export_and_run_due(engine, monkeypatch):
    c, audits, _ = _client(engine, monkeypatch, set(), llm=FakeLlm("kampanya"))
    h = {"x-user": "a"}
    x = c.get("/api/v1/stock/export/bitecekler.xlsx", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.cell(row=4, column=1).value == "Stok kodu" and ws.max_row == 6                # başlık + 2 kitap
    assert c.get("/api/v1/stock/export/yok.xlsx", headers=h).status_code == 404
    out = c.post("/api/v1/stock/run-due").json()
    assert out["fotograf"] == 6 and out["bitecek"]["yeni"] == 2 and out["fazla"]["yeni"] == 3
    fazla = store.list_suggestions(engine, T, tur="fazla")["items"]
    assert {s["hedef"] for s in fazla} == {"M35"} and all("Zeki AI önerisi" in s["gerekce"] for s in fazla)
    again = c.post("/api/v1/stock/run-due").json()                                            # ikinci koşu çift öneri açmaz
    assert again["bitecek"]["yeni"] == 0 and again["fazla"]["yeni"] == 0
    assert out["bulten"]["durum"] == "alıcı yok"


# ------------------------------------------------------------------ yetki


def test_access_rules_for_stock():
    assert A.rule_for("/api/v1/stock/run-due") == A.SYSTEM
    assert A.rule_for("/api/v1/stock/pick-line") == {"sayfa:stok-depo-hatti"}
    assert A.rule_for("/api/v1/stock/running-out") == {"sayfa:stok", "sayfa:stok-bitecekler"}
    assert "sayfa:stok-fark" in A.rule_for("/api/v1/stock/items/K1")
    f = A.features_for
    assert f("POST", "/api/v1/stock/suggestions/s1/decision") == ["ozellik:stok.oneri-karar"]
    assert f("POST", "/api/v1/stock/thresholds") == ["ozellik:stok.oneri-karar"]
    assert f("POST", "/api/v1/stock/thresholds/t1/approve") == []                 # açıkça verilen yetki ucun içinde
    assert f("GET", "/api/v1/stock/export/fazla.xlsx") == ["ozellik:veri.disa-aktar"]
    assert f("GET", "/api/v1/stock/items") == [] and f("POST", "/api/v1/stock/items/K1/notes") == []
    assert {"ozellik:stok.esik-onay", "ozellik:stok.maliyet", "ozellik:stok.depo-hatti"} <= A.explicit_keys()
    assert "ozellik:stok.oneri-karar" not in A.explicit_keys()


def test_runout_date_survives_a_tiny_sales_speed():
    """Hız çok küçükken yeterlilik günü takvimi aşar; gece işi düşmemeli, tarih boş kalmalı (2026-09-28 kabul)."""
    gun = S.days_of_cover(200000.0, 1e-6)
    assert gun is not None and gun > 1e9
    assert S.runout_date(date(2026, 8, 17), gun) is None
    assert S.runout_date(date(2026, 8, 17), 10.4) == "2026-08-27"
    assert S.runout_date(None, 10.0) is None and S.runout_date(date(2026, 8, 17), None) is None


# ------------------------------------------------------------------ hız (2026-09-29): uç kaynağı beklemez


def raw_live(at: float) -> dict:
    """Canlı okumanın bellekteki biçimi (tarih nesnesi, sayı anahtar) — tablodan dönen JSON biçimiyle aynı rakamı vermeli."""
    import copy

    r = copy.deepcopy(raw_fixture())
    r["at"] = at
    r["dataEnd"] = date(2026, 8, 17)
    r["warehouse"] = {k: {int(n): q for n, q in v.items()} for k, v in r["warehouse"].items()}
    r["warehouses"] = {int(k): v for k, v in r["warehouses"].items()}
    for v in r["movement"].values():
        v["son"] = date.fromisoformat(v["son"])
    for t in r["transfers"]:
        t["fisTarihi"] = date.fromisoformat(t["fisTarihi"])
    r["pick"] = [{"id": "s1", "no": "SP-1", "durum": 100000012, "tip": 1, "oncelik": 1, "siparis": date(2026, 9, 1),
                  "depoda": datetime(2026, 9, 20, 8), "pusula": datetime(2026, 9, 20, 9), "kutulandi": None, "sevk": None,
                  "depo": "Merkez", "toplayan": "Ali", "koli": 2},
                 {"id": "s2", "no": "SP-2", "durum": 3, "tip": 1, "oncelik": 3, "siparis": date(2026, 9, 1),
                  "depoda": datetime(2026, 9, 19, 8), "pusula": datetime(2026, 9, 19, 9), "kutulandi": datetime(2026, 9, 19, 12),
                  "sevk": datetime(2026, 9, 19, 15), "depo": "Merkez", "toplayan": "Ali", "koli": 1}]
    r["runs"] = {"logo_bakiye": {"sql": "SELECT 1", "rows": 8, "ms": 1200, "at": at}}
    return r


def _svc(tmp_path, **kw) -> S.Service:
    return S.Service(lambda: None, lambda: None, lambda: "", lambda: SETTINGS, cache=lambda: tmp_path / "stock-raw.json", **kw)


def _same(m: dict) -> str:
    import json

    return json.dumps({k: v for k, v in m.items() if k != "byCode"}, sort_keys=True, default=str)


def test_table_read_gives_the_same_numbers_as_the_live_read(engine, tmp_path, monkeypatch):
    """Eski hesap = yeni hesap: canlı okumadan kurulan model ile tablodan (köprü yeniden başladıktan sonra) kurulan model,
    özet, listeler ve depo hattı birebir aynı."""
    import json
    import time as _t

    live = raw_live(_t.time())
    svc = _svc(tmp_path)
    reads: list[int] = []
    monkeypatch.setattr(svc, "read", lambda: (reads.append(1), raw_live(live["at"]))[1])
    m1 = svc.model(engine, T)                       # hiç okuma yok: kaynak beklenir, okuma tabloya yazılır
    assert reads == [1] and store.read_at(engine, T) == live["at"]
    svc2 = _svc(tmp_path)                           # köprü yeniden başladı
    reads2: list[int] = []
    monkeypatch.setattr(svc2, "read", lambda: (reads2.append(1), raw_live(live["at"]))[1])
    m2 = svc2.model(engine, T)
    assert reads2 == [] and not svc2.refreshing()   # tablodaki okuma taze: kaynağa gidilmedi
    assert _same(m2) == _same(m1)
    now = date(2026, 9, 29)
    for fn in (lambda m: S.overview(m, S.transfer_rows(m, {}, now=now)), lambda m: S.filter_items(m),
               lambda m: S.running_out(m, 30), lambda m: S.excess(m), lambda m: S.diff_rows(m),
               lambda m: S.transfer_rows(m, {}, "hata", now=now),
               lambda m: S.pick_line(m, SETTINGS, True, now=datetime(2026, 9, 21))):
        assert json.dumps(fn(m2), sort_keys=True, default=str) == json.dumps(fn(m1), sort_keys=True, default=str)


def test_refresh_request_does_not_wait_for_the_source(engine, tmp_path, monkeypatch):
    """«Verileri yenile»: uç eldeki okumayla hemen döner, yeni okuma arka planda; bitince tabloya yazılır ve sonraki istek
    onu görür. Gece turu (`wait`) ise kaynağı bekler."""
    import threading
    import time as _t

    first = raw_live(_t.time())
    store.read_put(engine, T, first)                # gece turu yazmış
    gate = threading.Event()
    svc = _svc(tmp_path)

    def slow_read():
        gate.wait(10)
        r = raw_live(first["at"] + 1)
        r["balances"]["B-BIT"]["bakiye"] = 90.0
        return r
    monkeypatch.setattr(svc, "read", slow_read)
    t0 = _t.monotonic()
    m = svc.model(engine, T, fresh=True)
    assert _t.monotonic() - t0 < 1.0 and m["readAt"] == first["at"] and svc.refreshing()
    assert m["byCode"]["B-BIT"]["bakiye"] == 100.0
    gate.set()
    svc._bg.join(10)
    m2 = svc.model(engine, T)
    assert m2["readAt"] == first["at"] + 1 and m2["byCode"]["B-BIT"]["bakiye"] == 90.0
    assert store.read_at(engine, T) == first["at"] + 1
    # gece turu bekler: tabloda okuma olsa da kaynak okunur
    monkeypatch.setattr(svc, "read", lambda: raw_live(first["at"] + 2))
    assert svc.model(engine, T, fresh=True, wait=True)["readAt"] == first["at"] + 2


def test_stale_read_takes_a_newer_one_from_the_table(engine, tmp_path, monkeypatch):
    """Bellekteki okuma eskiyince önce tabloya bakılır: gece turu (ya da başka süreç) yenisini yazdıysa kaynağa gidilmez."""
    import time as _t

    svc = _svc(tmp_path)
    old = raw_live(_t.time() - S.TTL - 60)
    reads: list[int] = []
    monkeypatch.setattr(svc, "read", lambda: (reads.append(1), old)[1])
    svc.model(engine, T)
    assert reads == [1]
    store.read_put(engine, T, raw_live(_t.time()))   # tur yeni okumayı yazdı
    first = svc.model(engine, T)                     # eldeki model hemen döner, yenisi arkada kurulur (2. tur)
    assert first["readAt"] == old["at"]
    svc.idle()
    m = svc.model(engine, T)
    assert reads == [1] and m["readAt"] > old["at"] and not svc.refreshing()


def test_production_cards_do_not_block_and_arrive_later(engine, tmp_path, monkeypatch):
    """Üretim okuması yoksa stok ucu onu beklemez (kart boş); okuma gelince model yeniden kurulur."""
    import time as _t

    cards = [{"id": "c1", "stockCode": "B-BIT", "stage": "matbaada", "qty": 3000, "plan": {"depo": "2026-10-10"},
              "created": "2026-09-01", "name": "URT-1"}]
    ready = {"v": False}
    svc = _svc(tmp_path, m12_cards=lambda e, t: cards if ready["v"] else [], m12_ready=lambda: ready["v"])
    monkeypatch.setattr(svc, "read", lambda: raw_live(_t.time()))
    assert svc.model(engine, T)["byCode"]["B-BIT"]["uretim"] is None
    ready["v"] = True
    svc.model(engine, T)                  # üretim okuması geldi: model arkada yeniden kurulur, istek eldekini alır
    svc.idle()
    assert svc.model(engine, T)["byCode"]["B-BIT"]["uretim"]["kartId"] == "c1"

    class NoRead:
        class source:
            @staticmethod
            def peek():
                return None

        def cards(self, engine, tenant):
            raise AssertionError("üretim okuması beklenmemeli")

    deps = {"auth": None, "require_caller": None, "can": lambda u, k: False, "is_admin": lambda u: False, "audit": None,
            "conf": lambda k, d="": d, "fresh": lambda: False, "engine": lambda: engine, "tenant": lambda: T,
            "logo_file": lambda: "", "crm_file": lambda: "", "llm": lambda p: None, "m12": lambda: NoRead(),
            "costs": lambda: None}
    api_svc = stock_api.register(FastAPI(), deps)
    assert api_svc.m12_ready() is False and api_svc.m12_cards(engine, T) == []


# ------------------------------------------------------------------ hız, 2. tur (2026-09-29): model istekte kurulmaz


def _counting_svc(tmp_path, builds: list, delay: float = 0.0) -> S.Service:
    """Her model kurulumunda üretim kartları bir kez istenir: sayaç kurulum sayısıdır."""
    import time as _t

    def cards(e, t):
        builds.append(1)
        if delay:
            _t.sleep(delay)
        return []
    return _svc(tmp_path, m12_cards=cards)


def _no_source():
    raise AssertionError("kaynak okunmamalı")


def test_concurrent_requests_share_one_model_build(engine, tmp_path, monkeypatch):
    """Aynı anda gelen istekler (/names + /items) modeli ayrı ayrı kurmaz: tek kurulum, aynı model."""
    import threading
    import time as _t

    store.read_put(engine, T, raw_live(_t.time()))
    builds: list = []
    svc = _counting_svc(tmp_path, builds, delay=0.3)
    monkeypatch.setattr(svc, "read", _no_source)
    got: list = []
    threads = [threading.Thread(target=lambda: got.append(svc.model(engine, T))) for _ in range(5)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(10)
    assert len(got) == 5 and len(builds) == 1 and all(m is got[0] for m in got)


def test_new_read_does_not_block_the_request_and_gives_the_same_model(engine, tmp_path, monkeypatch):
    """Yeni okuma gelince istek eldeki modeli hemen alır; arkada kurulan model, aynı okumadan istek içinde kurulan
    (eski yol) modelle birebir aynı (eski hesap = yeni hesap)."""
    import time as _t

    first = raw_live(_t.time())
    store.read_put(engine, T, first)
    builds: list = []
    svc = _counting_svc(tmp_path, builds)
    monkeypatch.setattr(svc, "read", _no_source)
    m1 = svc.model(engine, T)
    newer = raw_live(first["at"] + 1)
    newer["balances"]["B-BIT"]["bakiye"] = 90.0
    svc._adopt(newer)                                    # gece turu yeni okumayı yazdı, bellekte
    builds_before = len(builds)
    t0 = _t.monotonic()
    again = svc.model(engine, T)
    assert again is m1 and _t.monotonic() - t0 < 1.0      # eldeki model, beklemeden
    svc.idle()
    m2 = svc.model(engine, T)
    assert m2["readAt"] == first["at"] + 1 and m2["byCode"]["B-BIT"]["bakiye"] == 90.0
    assert len(builds) == builds_before + 1
    same = S.build(newer, SETTINGS, store.approved_thresholds(engine, T), [], {})
    same["movementWindow"] = newer.get("movementWindow")
    from semantic_bridge import pazar_dagitim            # modele dağıtımcı bilgisi de eklenir (eski yolda da aynı)
    try:
        same["dagitim"] = pazar_dagitim.attach_stock(engine, T, same["items"])
    except Exception:  # noqa: BLE001
        same["dagitim"] = None
    assert _same(m2) == _same(same)


def test_threshold_approval_rebuilds_before_the_next_answer(engine, tmp_path):
    """Eşik onayı modeli düşürür: sonraki cevap onaylı eşikle gelir (eski modeli görmez)."""
    import time as _t

    store.read_put(engine, T, raw_live(_t.time()))
    svc = _svc(tmp_path)
    assert svc.model(engine, T)["byCode"]["B-BIT"]["esik"] is None
    t = store.save_threshold(engine, T, "ayse", {"stokKodu": "B-BIT", "guvenlikGun": 20})
    store.decide_threshold(engine, T, "ayse", t["id"], True, None)
    svc.invalidate()
    assert svc.model(engine, T)["byCode"]["B-BIT"]["esik"]["guvenlikGun"] == 20


def test_finished_read_is_not_restarted_at_once(engine, tmp_path, monkeypatch):
    """Okuma 5 dakikadan uzun sürünce biter bitmez «eski» sayılıyordu ve ekran açık kaldıkça kaynak durmadan okunuyordu:
    kendiliğinden yeniden okuma, önceki okuma bittikten 5 dakika sonra."""
    import time as _t

    svc = _svc(tmp_path)
    reads: list[int] = []
    monkeypatch.setattr(svc, "read", lambda: (reads.append(1), raw_live(_t.time() - S.TTL - 30))[1])   # uzun okuma
    svc.model(engine, T)
    svc.model(engine, T)
    assert reads == [1] and not svc.refreshing()
    with svc._lock:
        svc._raw["_done"] = _t.time() - S.TTL - 1
    svc.model(engine, T)
    svc.idle()
    assert reads == [1, 1]


def test_warm_prepares_the_model_from_the_table_without_the_source(engine, tmp_path, monkeypatch):
    import time as _t

    builds: list = []
    svc = _counting_svc(tmp_path, builds)
    monkeypatch.setattr(svc, "read", _no_source)
    assert svc.warm(engine, T) is False                 # tabloda okuma yok: kaynağa gidilmez
    store.read_put(engine, T, raw_live(_t.time()))
    assert svc.warm(engine, T) is True and len(builds) == 1
    svc.model(engine, T)
    assert len(builds) == 1                              # istek hazır modeli aldı


def test_names_and_items_answers_are_the_old_calculation(engine, monkeypatch):
    """Arama kutusu listesi ve süzgeçli liste model başına bir kez hesaplanır; cevap eski hesapla birebir."""
    c, _, svc = _client(engine, monkeypatch, set())
    h = {"x-user": "a"}
    m = svc.model(engine, T)
    old_names = [{"value": i["stokKodu"], "label": f"{i['ad'] or i['stokKodu']} · {i['stokKodu']}"}
                 for i in sorted(m["items"], key=lambda i: S.fold(i["ad"] or i["stokKodu"]))]
    for _ in range(2):
        r = c.get("/api/v1/stock/names", headers=h)
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/json")
        assert r.json() == {"items": old_names}
    cases = [{}, {"q": "bit"}, {"q": "TİMAŞ"}, {"q": "yazar a"}, {"durum": "bitecek,fazla"}, {"sira": "ad"},
             {"sira": "bakiye"}, {"sira": "hiz"}, {"yayinevi": "Timaş"}, {"depo": "2"}, {"durum": "pasif"}]
    for p in cases:
        for _ in range(2):                                # ikinci istek hafızadan
            out = c.get("/api/v1/stock/items", params=p, headers=h).json()
            args = {"q": "", "yayinevi": "", "depo": "", "durum": "", "sira": "gun", **p}
            old = S.filter_items(m, **args)               # eski yol: satır başına sadeleştirme
            assert [i["stokKodu"] for i in out["items"]] == [i["stokKodu"] for i in old[:50]], p
            assert out["total"] == len(old), p
            assert out["yayinevleri"] == sorted({i["yayinevi"] for i in m["items"] if i["yayinevi"]}, key=S.fold)
    assert c.get("/api/v1/stock/items", params={"depo": "x"}, headers=h).status_code == 400
    folded = {i["stokKodu"]: S.search_text(i) for i in m["items"]}
    for q in ("bit", "kitap yazar", "b-y", "zzz"):
        assert S.filter_items(m, q=q, folded=folded) == S.filter_items(m, q=q)
