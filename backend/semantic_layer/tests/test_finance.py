"""M45 Finansal raporlama: hesap eşlemesi çözümü, gelir tablosu (kural, dışlama, eşlenmemiş), kapanış, kârlılık
(kesin / yaklaşık, telif), 13 haftalık nakit, vergi takvimi ve hatırlatma, zamanlama, yetki kuralları, SQL tanımları.

Veriler yapaydır ve yalnız hesap kurallarını sınar; gerçek Logo/CRM kabulü test sunucusunda
(`scripts/acceptance/M45/`, doğrudan SQL referanslarıyla).
"""

from __future__ import annotations

from datetime import date, datetime, timedelta

import pytest

from semantic_bridge import access as A
from semantic_bridge import finance as F
from semantic_bridge import finance_sources as S
from semantic_layer.store.catalog_store import open_store

T = "t1"


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    F._ready.discard(id(e))
    F.ensure(e)
    yield e
    F._ready.discard(id(e))


def _actual(y, m, code, borc=0.0, alacak=0.0, kural="dahil", merkez="#YOK", ad=None):
    return {"year": y, "month": m, "hesap_kodu": code, "merkez_kodu": merkez, "kural": kural, "hesap_adi": ad or code,
            "borc": borc, "alacak": alacak, "satir": 1}


def _seed(engine, end=date(2026, 8, 17)):
    rows = [
        _actual(2026, 7, "600.01.001", alacak=1_000_000),
        _actual(2026, 7, "610.01", borc=50_000),
        _actual(2026, 7, "621.01", borc=400_000),
        _actual(2026, 7, "760.10", borc=100_000),
        _actual(2026, 7, "770.01", borc=80_000),
        _actual(2026, 7, "730.38.381", borc=300_000),             # üretim gideri: stoka gider, dışlanır
        _actual(2026, 7, "642.01", alacak=5_000),
        _actual(2026, 7, "656.01", borc=2_000),
        _actual(2026, 7, "795.01", borc=7_000),                   # kuralı yok → eşlenmemiş
        _actual(2026, 7, "631.01", borc=100_000, kural="yansitma"),  # 760'ın yansıtması: sayılmaz
        _actual(2026, 7, "690.01", borc=5, kural="kapanis"),
        _actual(2026, 6, "600.01.001", alacak=900_000),
        _actual(2025, 7, "600.01.001", alacak=700_000),
        _actual(2025, 7, "621.01", borc=300_000),
    ]
    with engine.begin() as c:
        c.execute(F.ACTUALS.insert(), rows)
        c.execute(F.SALES_MONTH.insert(), [
            {"year": 2026, "month": 7, "adet": 10_000, "brut": 1_800_000, "net": 950_000, "maliyet": 300_000,
             "maliyetli_net": 760_000, "maliyetsiz_adet": 2_000, "maliyetsiz_net": 190_000, "satis_net": 1_000_000,
             "iade_net": 50_000, "satis_satir": 1_000, "maliyetsiz_satir": 200},
        ])
    F.meta_set(engine, "data_end", {"date": end.isoformat(), "costed": "2026-06-30"})
    for y in (2025, 2026):
        F.meta_set(engine, f"ledger:{y}", {"firm": "411" if y == 2026 else "211"})
    F.meta_set(engine, "trial:2026", {"months": {"7": {"borc": 1_042_005.0, "alacak": 1_042_005.0}}})


# ------------------------------------------------------------------ eşleme


def test_resolve_walks_from_specific_to_group_and_prefers_decisions():
    assert F._ancestors("600.01.001") == ["600.01.001", "600.01", "600"]
    assert F._ancestors("600") == ["600"]
    r = F.resolve("600.01.001", {})
    assert r == {**r, "satir": "BRUT_SATIS", "durum": "oneri", "kaynak": "hesap-plani", "kayit": "600"}
    assert F.resolve("795.01", {})["durum"] == "yok"
    assert F.resolve("730.38.381", {})["satir"] == F.EXCLUDED

    class Row:
        def __init__(self, **kw):
            self.__dict__.update({"satir_kodu": None, "oneri_satir": None, "oneri_kaynak": None, "oneri_olasilik": None,
                                  "onaylayan": None, "onay_tarihi": None, "not_": None, **kw})

    stored = {"600": Row(durum="onayli", satir_kodu="DIGER_GELIR"), "600.01": Row(durum="oneri", oneri_satir="BRUT_SATIS", oneri_kaynak="zeki")}
    # onaylı karar (grupta bile) öneriden önce gelir
    assert F.resolve("600.01.001", stored)["satir"] == "DIGER_GELIR"
    stored["600.01.001"] = Row(durum="dislandi")
    assert F.resolve("600.01.001", stored)["satir"] == F.EXCLUDED


def test_set_mapping_validates_and_approve_uses_current_suggestion(engine):
    _seed(engine)
    with pytest.raises(F.FinanceError):
        F.set_mapping(engine, T, "muh", "795.01", {"satir": "NET_SATIS"})   # ara toplam hedef olamaz
    with pytest.raises(S.SourceError):
        F.set_mapping(engine, T, "muh", "795'; DROP", {"satir": "GYG"})
    out = F.approve_suggestions(engine, T, "muh", ["600", "795.01", "760.10"])
    assert {a["hesap"] for a in out["approved"]} == {"600", "760.10"}
    assert out["skipped"][0]["hesap"] == "795.01"
    m = F.account_map(engine, T, 2026)
    by = {i["hesap"]: i for i in m["items"]}
    assert by["600.01.001"]["esleme"]["durum"] == "onayli" and by["600.01.001"]["esleme"]["kayit"] == "600"
    assert by["795.01"]["esleme"]["durum"] == "yok"
    F.reset_mapping(engine, T, "600")
    assert F.resolve("600.01.001", F._stored_map(engine, T))["durum"] == "oneri"


def test_model_suggestion_is_closed_set_and_thresholded(engine):
    _seed(engine)

    class Choice:
        def __init__(self, idx, p, margin):
            self.index, self.probability, self.margin = idx, p, margin
            self.choice = "x"

        def confident(self, p, min_margin=0.0):
            return self.probability >= p and self.margin >= min_margin

        def as_dict(self):
            return {"index": self.index, "probability": self.probability}

    leaves = F.leaf_codes(F.lines(engine, T))

    class Llm:
        def __init__(self, answers):
            self.answers = answers
            self.seen = []

        def choose(self, prompt, choices):
            self.seen.append(choices)
            return self.answers.pop(0)

    items = F.suggestion_targets(engine, T)
    assert [i["hesap"] for i in items] == ["795.01"]
    llm = Llm([Choice(leaves.index("GYG"), 0.93, 0.8)])
    res = F.suggest_with_model(engine, T, llm, items, 0.70, 0.30)
    assert res == {"istenen": 1, "oneri": 1, "belirsiz": 0, "hata": 0}
    assert llm.seen[0][-1].startswith("Gelir tablosu dışı")
    r = F.resolve("795.01", F._stored_map(engine, T))
    assert r["satir"] == "GYG" and r["kaynak"] == "zeki" and r["durum"] == "oneri"
    # eşik altı: satır yazılmaz, «belirsiz»
    res2 = F.suggest_with_model(engine, T, Llm([Choice(0, 0.51, 0.05)]), items, 0.70, 0.30)
    assert res2["belirsiz"] == 1
    assert F.resolve("795.01", F._stored_map(engine, T))["durum"] == "yok"


# ------------------------------------------------------------------ gelir tablosu


def test_pnl_counts_only_included_lines_and_keeps_unmapped_in_net(engine):
    _seed(engine)
    d = F.pnl(engine, T, 2026, 7, "ay")
    v = {r["kod"]: r["values"]["donem"] for r in d["rows"]}
    assert v["BRUT_SATIS"] == 1_000_000 and v["SATIS_INDIRIM"] == -50_000 and v["NET_SATIS"] == 950_000
    assert v["SMM"] == -400_000 and v["BRUT_KAR"] == 550_000
    # 7/A: gider 760'ta bir kez; 631 yansıtması sayılmaz
    assert v["PSD"] == -100_000 and v["GYG"] == -80_000 and v["FAALIYET"] == -180_000
    assert v["DIGER_GELIR"] == 5_000 and v["DIGER_GIDER"] == -2_000
    assert v[F.UNMAPPED] == -7_000
    assert v["NET_KAR"] == 550_000 - 180_000 + 5_000 - 2_000 - 7_000
    assert d["dislanan"] == -300_000
    assert d["kurallar"]["yansitma"]["borc"] == 100_000 and d["kurallar"]["kapanis"]["hesap"] == 1
    assert d["esleme"]["eslenmemis"] == 1
    # geçen yıl ve önceki ay sütunları
    assert {r["kod"]: r["values"]["gecenYil"] for r in d["rows"]}["BRUT_KAR"] == 400_000
    assert {r["kod"]: r["values"]["onceki"] for r in d["rows"]}["BRUT_SATIS"] == 900_000
    # bütçe yoksa sütun boş ve nedeni yazılı
    assert d["columns"]["butce"]["plan"] is None
    # maliyet kapsamı: %80 maliyetli → SMM ve brüt kâr «yaklaşık»
    assert d["maliyet"]["maliyetliPay"] == pytest.approx(0.8) and d["maliyet"]["yaklasik"]
    assert next(r for r in d["rows"] if r["kod"] == "SMM")["yaklasik"]
    assert d["faturaNetSatis"] == 950_000 and d["mutabakatFarki"] == 0
    assert d["mizan"]["durum"] == "denk"
    assert d["veriSonu"] == "2026-08-17" and d["maliyetSonu"] == "2026-06-30"


def test_periods():
    assert F.period_months(2026, 8, "ceyrek") == [(2026, 7), (2026, 8), (2026, 9)]
    assert F.previous_period(2026, 1, "ay") == [(2025, 12)]
    assert F.previous_period(2026, 2, "ceyrek") == [(2025, 10), (2025, 11), (2025, 12)]
    assert F.previous_period(2026, 5, "ytd") is None
    assert F.period_label([(2026, 1), (2026, 3)]) == "Ocak–Mart 2026"
    with pytest.raises(F.FinanceError):
        F.period_months(2026, 13, "ay")
    assert F._complete([(2026, 7)], date(2026, 8, 17)) and not F._complete([(2026, 8)], date(2026, 8, 17))


def test_line_accounts_lists_only_that_lines_accounts(engine):
    _seed(engine)
    d = F.line_accounts(engine, T, "FAALIYET", 2026, 7, "ay")
    assert {i["hesap"] for i in d["items"]} == {"760.10", "770.01"}
    assert d["toplam"] == -180_000
    with pytest.raises(F.FinanceError):
        F.line_accounts(engine, T, "YOK", 2026, 7, "ay")


def test_close_stores_values_and_flags_later_change(engine):
    _seed(engine)
    with pytest.raises(F.FinanceError):
        F.close_month(engine, T, "muh", 2026, 8)       # ay veri sonunda tamamlanmadı
    c = F.close_month(engine, T, "muh", 2026, 7, "Temmuz kapandı")
    assert c["durum"] == "kapandi" and c["kapatan"] == "muh"
    assert F.pnl(engine, T, 2026, 7)["kapanis"]["degisti"] is False
    with engine.begin() as conn:
        conn.execute(F.ACTUALS.insert(), [_actual(2026, 7, "770.02", borc=1_000)])
    assert F.pnl(engine, T, 2026, 7)["kapanis"]["degisti"] is True
    with pytest.raises(F.FinanceError):
        F.reopen_month(engine, T, "muh", 2026, 7, "")
    assert F.reopen_month(engine, T, "muh", 2026, 7, "Düzeltme fişi")["durum"] == "acik"


def test_reconciliation_shows_both_sources(engine):
    _seed(engine)
    r = F.reconciliation(engine, T, 2026, 7)
    assert r["muhasebe"]["netSatis"] == 950_000 and r["fatura"]["netSatis"] == 950_000
    assert r["fatura"]["iskonto"] == 850_000 and r["fark"] == 0


# ------------------------------------------------------------------ kârlılık


def _seed_profit(engine):
    with engine.begin() as c:
        c.execute(F.PROFIT_ITEMS.insert(), [
            # K1: tamamı maliyetli
            {"year": 2026, "month": 7, "stok_kodu": "K1", "kanal": "KİTAPÇI", "adet": 100, "brut": 20_000, "net": 10_000,
             "maliyet": 4_000, "maliyetli_net": 10_000, "maliyetsiz_adet": 0, "maliyetsiz_net": 0},
            # K2: yarısı maliyetsiz, M9 birim maliyeti 30
            {"year": 2026, "month": 7, "stok_kodu": "K2", "kanal": "E-TİCARET", "adet": 100, "brut": 16_000, "net": 8_000,
             "maliyet": 1_500, "maliyetli_net": 4_000, "maliyetsiz_adet": 50, "maliyetsiz_net": 4_000},
            # K3: tamamı maliyetsiz, birim maliyeti de bilinmiyor
            {"year": 2026, "month": 7, "stok_kodu": "K3", "kanal": "E-TİCARET", "adet": 10, "brut": 2_000, "net": 1_000,
             "maliyet": 0, "maliyetli_net": 0, "maliyetsiz_adet": 10, "maliyetsiz_net": 1_000},
        ])


def test_profitability_separates_exact_and_approximate(engine):
    _seed(engine)
    _seed_profit(engine)
    snap = {"books": {"K1": {"price": 110, "vat": 10, "royalty": {"rate": 0.10, "basis": "net", "on": "satis"}},
                      "K2": {"price": 220, "vat": 0.10, "royalty": {"rate": 0.05, "basis": "kapak", "on": "satis"}}}}
    d = F.profitability(engine, by="kitap", year=2026, frm=7, to=7, unit_costs=lambda codes: {"K2": {"maliyet": 30.0}},
                        snapshot=snap)
    by = {i["key"]: i for i in d["items"]}
    assert by["K1"]["katkiKesin"] == 6_000 and by["K1"]["marjKesin"] == 0.6
    assert by["K1"]["telif"] == 1_000                          # %10 × net
    assert by["K1"]["katkiYaklasik"] == 10_000 - 4_000 - 1_000
    assert by["K2"]["maliyetTahmini"] == 1_500                 # 50 × 30
    assert by["K2"]["telif"] == pytest.approx(0.05 * 100 * 200)  # kapak (KDV hariç) × adet
    assert by["K2"]["katkiYaklasik"] == pytest.approx(8_000 - 1_500 - 1_500 - 1_000)
    assert by["K2"]["kapsam"] == 1.0
    # birim maliyeti bilinmeyen satış hesaba katılmaz, uydurulmaz
    assert by["K3"]["maliyetBilinmeyenNet"] == 1_000 and by["K3"]["katkiYaklasik"] is None and by["K3"]["kapsam"] == 0
    assert d["toplam"]["net"] == 19_000 and d["toplam"]["maliyetBilinmeyenNet"] == 1_000
    assert d["telif"]["sozlesmeli"] == 2 and d["telif"]["sozlesmesiz"] == 1
    k = F.profitability(engine, by="kanal", year=2026, frm=7, to=7, unit_costs=lambda c: {}, snapshot=None, with_royalty=False)
    assert {i["key"] for i in k["items"]} == {"KİTAPÇI", "E-TİCARET"} and k["telif"]["hesaplandi"] is False
    csv_text = F.profitability_csv(d)
    assert csv_text.startswith("﻿") and "Yaklaşık katkı" in csv_text and "Veri son günü: 2026-08-17" in csv_text
    with pytest.raises(F.FinanceError):
        F.profitability(engine, by="yazar", year=2026)


def test_snapshot_unit_cost_matches_m9_definition():
    snap = {"sales": {"K1": {"2025": {"cogs": 300, "costedQty": 10}, "2026": {"cogs": 0, "costedQty": 0}}}}
    assert F.snapshot_unit_costs(snap, ["K1", "K9"]) == {"K1": {"maliyet": 30.0, "kaynak": "gerceklesen", "yil": 2025}}
    assert F.snapshot_unit_costs(None, ["K1"]) == {}


# ------------------------------------------------------------------ nakit


def test_build_cash_places_items_by_week_and_flags_open_week():
    asof = date(2026, 8, 17)   # pazartesi
    out = F.build_cash({
        "asof": asof, "position": {"100": 1_000, "102": 9_000, "101": 5_000},
        "receivables": [(date(2026, 8, 20), 4_000, 3), (date(2026, 8, 1), 2_500, 1)],   # ikincisi vadesi geçmiş
        "payables": [(date(2026, 8, 25), 20_000, 2)],
        "cheques": [(1, 1, date(2026, 9, 1), 3_000), (3, 9, date(2026, 9, 2), 1_000), (1, 8, date(2026, 9, 3), 999)],
        "crm": [(date(2026, 8, 18), 500)],
        "royalty": [(date(2026, 8, 26), 700, "TRY"), (date(2026, 8, 26), 100, "EUR")],
        "tax": [(date(2026, 8, 26), 1_200, "KDV")],
        "budget": [(date(2026, 8, 1), 3_100)],
    })
    w = out["haftalar"]
    assert out["baslangic"] == "2026-08-17" and len(w) == 13 and out["acilisBakiye"] == 10_000
    assert w[0]["giris"] == 4_500 and w[0]["cikis"] == 0 and w[0]["kapanis"] == 14_500
    assert w[1]["cikis"] == 20_000 + 700 + 1_200 and w[1]["acik"] is True
    assert w[2]["giris"] == 3_000 and w[2]["cikis"] == 1_000      # tahsil edilmiş çek (durum 8) girmez
    assert out["vadesiGecmis"] == {"alacak": 2_500}
    assert out["dovizTelif"] == {"EUR": 100}
    budget = next(k for k in out["kalemler"] if k["kalem"] == "butce")
    assert budget["yon"] == "bilgi" and budget["toplam"] == pytest.approx(3_100 * 15 / 31, abs=0.05)
    # bilgi satırı toplama katılmaz
    assert w[0]["cikis"] == 0


def test_cash_run_roundtrip_and_budget_toggle(engine):
    _seed(engine)
    res = F.build_cash({"asof": date(2026, 8, 17), "position": {"100": 0, "102": 1_000},
                        "payables": [(date(2026, 8, 18), 500, 1)], "budget": [(date(2026, 8, 1), 3_100)]})
    rid = F.save_cash_run(engine, T, "cfo", res)
    c = F.cash(engine, T)
    assert c["run"]["id"] == rid and c["haftalar"][0]["kapanis"] == 500 and not c["acikHafta"]
    cb = F.cash(engine, T, include_budget=True)
    assert cb["haftalar"][0]["cikis"] > 500 and cb["butceDahil"]
    h = F.cash_history(engine, T)
    assert h["items"] == [] and h["not"]


def test_cash_and_summary_schedule(engine):
    tz = F.TZ
    mon_6 = datetime(2026, 9, 28, 6, 0, tzinfo=tz)
    mon_8 = datetime(2026, 9, 28, 8, 0, tzinfo=tz)
    assert F.cash_due(engine, T, mon_6) is True           # hiç kurulmadıysa hemen
    rid = F.save_cash_run(engine, T, None, F.build_cash({"asof": date(2026, 8, 17)}))
    with engine.begin() as c:
        c.execute(F.CASH_RUNS.update().where(F.CASH_RUNS.c.id == rid).values(run_at=datetime(2026, 9, 28, 7, 30, tzinfo=tz)))
    assert F.cash_due(engine, T, mon_8) is False          # bu haftanın tablosu 07:00'den sonra kuruldu
    assert F.cash_due(engine, T, mon_8 + timedelta(days=7)) is True
    assert F.summary_due(engine, datetime(2026, 9, 27, 9, 0, tzinfo=tz)) is False   # pazar
    assert F.summary_due(engine, datetime(2026, 9, 28, 8, 10, tzinfo=tz)) is False
    assert F.summary_due(engine, datetime(2026, 9, 28, 8, 40, tzinfo=tz)) is True
    F.meta_set(engine, "summary_mail", {"day": "2026-09-28"})
    assert F.summary_due(engine, datetime(2026, 9, 28, 10, 0, tzinfo=tz)) is False


def test_cheque_direction_is_configurable(monkeypatch):
    assert S.cheque_direction(1, 1) == "giris" and S.cheque_direction(4, 10) == "cikis" and S.cheque_direction(1, 8) is None
    monkeypatch.setenv("FINANCE_CS_IN_STATUS", "1,4,5")
    assert S.cheque_direction(2, 5) == "giris"
    assert S.cheque_direction("x", 1) is None


# ------------------------------------------------------------------ vergi takvimi


def test_tax_calendar_crud_copy_and_reminders(engine):
    t = F.tax_create(engine, T, "muh", {"beyan": "KDV beyannamesi", "donem": "Eylül 2026", "sonGun": "2026-10-28",
                                        "sorumlu": "Muhasebe", "tutar": "125.000,50"})
    assert t["durum"] == "bekliyor" and t["tutar"] == 125_000.5
    with pytest.raises(F.FinanceError):
        F.tax_create(engine, T, "muh", {"beyan": "", "sonGun": "2026-10-01"})
    with pytest.raises(F.FinanceError):
        F.tax_update(engine, T, "muh", t["id"], {"durum": "bilinmiyor"})
    due = F.tax_due_reminders(engine, T, date(2026, 10, 22))
    assert [(d["id"], d["esik"]) for d in due] == [(t["id"], 7)]
    F.mark_reminded(engine, due)
    assert F.tax_due_reminders(engine, T, date(2026, 10, 23)) == []
    assert F.tax_due_reminders(engine, T, date(2026, 10, 26))[0]["esik"] == 2
    u, diff = F.tax_update(engine, T, "muh", t["id"], {"durum": "verildi"})
    assert u["durum"] == "verildi" and "durum" in diff
    assert F.tax_due_reminders(engine, T, date(2026, 10, 27)) == []
    out = F.tax_copy_year(engine, T, "muh", 2026, 2027)
    assert out == {"kopyalanan": 1, "atlanan": 0}
    nxt = F.tax_list(engine, T, 2027)["items"][0]
    assert nxt["sonGun"] == "2027-10-28" and nxt["durum"] == "bekliyor" and nxt["tutar"] is None
    assert F.tax_copy_year(engine, T, "muh", 2026, 2027)["kopyalanan"] == 0
    F.tax_delete(engine, T, t["id"])
    assert F.tax_list(engine, T, 2026)["items"] == []


def test_sapma_note_upserts_by_key(engine):
    n1 = F.save_note(engine, T, "cfo", {"year": 2026, "anahtar": "gider|B04|760", "metin": "Fuar faturası iki kez"})
    n2 = F.save_note(engine, T, "cfo", {"year": 2026, "anahtar": "gider|B04|760", "metin": "Düzeltme fişi kesildi"})
    assert n1["id"] == n2["id"] and F.notes(engine, T, 2026)[0]["metin"] == "Düzeltme fişi kesildi"
    with pytest.raises(F.FinanceError):
        F.save_note(engine, T, "cfo", {"year": 2026, "anahtar": "", "metin": "x"})


def test_summary_text_has_no_invented_numbers(engine):
    _seed(engine)
    data = {"veriSonu": "2026-08-17", "donem": "Ocak–Ağustos 2026",
            "cards": [{"label": "Net satış", "value": 848_110_178.0, "unit": "₺", "yaklasik": False},
                      {"label": "Brüt kâr", "value": 1_000.0, "unit": "₺", "yaklasik": True}],
            "dikkat": [{"metin": "KDV: son gün 2026-09-28"}]}
    txt = F.summary_text(data, "https://x/timas/finansal-raporlar")
    assert "848.110.178 ₺" in txt and "(yaklaşık)" in txt and "2026-08-17" in txt and "finansal-raporlar" in txt


def test_pnl_excel_opens(engine):
    _seed(engine)
    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(F.pnl_xlsx(F.pnl(engine, T, 2026, 7))))
    ws = wb.active
    assert ws.title == "Gelir tablosu" and "Veri son günü: 2026-08-17" in ws["A2"].value


# ------------------------------------------------------------------ SQL tanımları


def test_sql_uses_catalog_definitions():
    s = S.sales_month_sql("411", 2026)
    assert "S.INVOICEREF <> 0" in s and "S.LINETYPE = 0" in s and "S.TRCODE IN (2,3,7,8,9)" in s and "S.LINENET" in s
    assert "LG_411_01_STLINE" in s and "'20260101'" in s and "'20270101'" in s
    a = S.account_actuals_sql("411", 2026)
    assert "'yansitma'" in a and "'kapanis'" in a and "F.CANCELLED = 0" in a and "L.CANCELLED = 0" in a
    assert "('690','692')" in a
    e = S.entries_sql("411", "760.10", date(2026, 7, 1), date(2026, 8, 1), 2, 100)
    assert "A.CODE = '760.10' OR A.CODE LIKE '760.10.%'" in e and "OFFSET 200 ROWS FETCH NEXT 100 ROWS ONLY" in e
    with pytest.raises(S.SourceError):
        S.entries_sql("411", "760'--", date(2026, 7, 1), date(2026, 8, 1), 0, 100)
    f = S.fifo_due_sql("411", 2026, date(2026, 8, 17), payable=False)
    assert "C.CODE LIKE '120%'" in f and "P.SIGN = 0" in f and "GETDATE" not in f and "'20260818'" in f
    p = S.fifo_due_sql("411", 2026, date(2026, 8, 17), payable=True)
    assert "C.CODE LIKE '320%'" in p and "P.SIGN = 1" in p
    assert "OUTCOST = 0" in S.client_uncosted_sql("411", 2026)


# ------------------------------------------------------------------ yetki


def test_finance_access_rules():
    assert A.rule_for("/api/v1/finance/run-due") == A.SYSTEM
    assert A.rule_for("/api/v1/finance/pnl") == frozenset({"sayfa:finansal-raporlar"})
    assert A.features_for("GET", "/api/v1/finance/cash") == ["ozellik:finans.nakit"]
    assert A.features_for("POST", "/api/v1/finance/cash/rebuild") == ["ozellik:finans.nakit"]
    assert A.features_for("GET", "/api/v1/finance/cash/history") == ["ozellik:finans.nakit"]
    assert A.features_for("GET", "/api/v1/finance/summary") == []
    assert A.features_for("POST", "/api/v1/finance/tax-calendar") == ["ozellik:finans.vergi-takvimi"]
    assert A.features_for("GET", "/api/v1/finance/tax-calendar") == []
    assert A.features_for("POST", "/api/v1/finance/notes") == ["ozellik:finans.sapma-notu"]
    assert A.features_for("GET", "/api/v1/finance/pnl/export.xlsx") == ["ozellik:veri.disa-aktar"]
    assert A.features_for("GET", "/api/v1/finance/profitability/export.csv") == ["ozellik:veri.disa-aktar"]
    # eşleme ve kapanış açıkça verilir, ucun içinde denetlenir
    assert A.features_for("PATCH", "/api/v1/finance/account-map/600") == []
    assert {"ozellik:finans.esleme", "ozellik:finans.kapanis"} <= A.explicit_keys()
    assert "ozellik:finans.nakit" not in A.explicit_keys()
    assert {"sayfa:finansal-raporlar", "ozellik:finans.nakit", "ozellik:finans.vergi-takvimi"} <= A.all_keys()
    page = next(p for p in A.catalog()["pages"] if p["key"] == "sayfa:finansal-raporlar")
    assert page["area"] == "finans" and page.get("explicit") is True
