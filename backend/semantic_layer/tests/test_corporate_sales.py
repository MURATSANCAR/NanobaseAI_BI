"""M32 Kurumsal satış ve B2B: hacim indirimi, fiyat listesi seçimi, teklif hesabı ve onay (iki göz), fırsat görünürlüğü,
paket önerisi, dönemsel hatırlatma, sessiz bayi, maliyet kaynağı, ZEKİ AI kapalı küme seçimi, yetki kuralları.

Veriler yapaydır ve yalnız hesap kurallarını sınar; gerçek Logo/CRM kabulü test sunucusunda
(`scripts/acceptance/M32/`, günlük 2026-09-28 M32).
"""

from __future__ import annotations

import json
from datetime import date, datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest

from semantic_bridge import access as A
from semantic_bridge import corporate_sales as C
from semantic_bridge import corporate_sales_docs as D
from semantic_bridge import corporate_sales_sources as S
from semantic_layer.store.catalog_store import open_store

T = "t1"
ST = C.settings_from(lambda key, default="": C.DEFAULTS.get(key, default))


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    C._ready.discard(id(e))
    C.ensure(e)
    return e


def _books(engine):
    now = datetime.now(timezone.utc)
    rows = [
        # stok, ad, stok miktarı, bu yıl adet, fiyat, yaş
        ("L1", "Liderlik 1", 500, 900, 100.0, None, None),
        ("L2", "Liderlik 2", 500, 700, 80.0, None, None),
        ("L3", "Liderlik 3", 500, 500, 60.0, None, None),
        ("L4", "Liderlik 4", 50, 2000, 50.0, None, None),     # stok az
        ("L5", "Liderlik 5", 500, 300, None, None, None),     # geçerli fiyat yok
        ("L6", "Liderlik 6", 500, 100, 40.0, None, None),
        ("C1", "Çocuk 1", 500, 800, 30.0, 6, 9),
        ("C2", "Çocuk 2", 500, 700, 30.0, 12, 15),
    ]
    with engine.begin() as c:
        for code, ad, stok, yil, fiyat, lo, hi in rows:
            c.execute(C.BOOKS.insert().values(stok_kodu=code, ad=ad, yazar="Yazar", stok=stok, yil_adet=yil, fiyat=fiyat,
                                              fiyat_liste="SL1" if fiyat else None, fiyat_liste_sayisi=1 if fiyat else None,
                                              fiyat_kdv_dahil=True, yas_min=lo, yas_max=hi, bayi_son=0, bayi_onceki=0,
                                              in_logo=True, in_crm=True))
            tema = "Liderlik ve yönetim" if code.startswith("L") else "Çocuk kütüphanesi"
            c.execute(C.THEMES.insert().values(stok_kodu=code, tema=tema, kaynak="crm", durum="onayli", zaman=now))


def _account(engine, ref="K001", name="Örnek Bankası A.Ş."):
    with engine.begin() as c:
        c.execute(C.ACCOUNTS.insert().values(tenant_id=T, ref=ref, logo_code=ref, unvan=name, pasif=False, asof=datetime.now(timezone.utc)))


# ------------------------------------------------------------------ hacim indirimi ve fiyat


def test_tiers_parse_and_setting_wins_over_history():
    assert C.parse_tiers("100:10;300:15, 1000=20%") == [(100, 0.10), (300, 0.15), (1000, 0.20)]
    st = {**ST, "volumeTiers": C.parse_tiers("100:10;300:15")}
    assert C.volume_discount(50, st, [])["indirim"] == 0.0
    assert C.volume_discount(300, st, [])["indirim"] == 0.15
    assert C.volume_discount(300, st, [])["kaynak"] == "ayar"


def test_measured_discount_is_the_bucket_median_and_needs_enough_invoices():
    inv = [{"adet": 150, "brut": 100.0, "net": 100.0 * (1 - r)} for r in (0.10, 0.20, 0.30, 0.40, 0.50)]
    inv += [{"adet": 20, "brut": 100.0, "net": 90.0}]
    b = C.measured_discounts(inv, [1, 100, 300], 5)
    assert b[1]["n"] == 5 and b[1]["indirim"] == pytest.approx(0.30)
    assert b[0]["n"] == 1 and b[0]["indirim"] is None           # 5 faturadan az: öneri yok
    assert C.volume_discount(150, ST, b)["indirim"] == pytest.approx(0.30)
    low = C.volume_discount(20, ST, b)
    assert low["indirim"] == 0.0 and low["kaynak"] == "gecmis-yetersiz"


def test_price_list_prefers_general_list_then_priority_then_newest():
    rows = [
        {"fiyat": 90, "liste": "CARI", "oncelik": 0, "kdv_dahil": 1, "cari_ozel": "", "cari": "120.01", "bas": "2026-01-01"},
        {"fiyat": 100, "liste": "ESKI", "oncelik": 1, "kdv_dahil": 1, "cari_ozel": "", "cari": "", "bas": "2025-01-01"},
        {"fiyat": 110, "liste": "YENI", "oncelik": 1, "kdv_dahil": 1, "cari_ozel": "", "cari": "", "bas": "2026-03-01"},
        {"fiyat": 0, "liste": "SIFIR", "oncelik": 0, "kdv_dahil": 1, "cari_ozel": "", "cari": "", "bas": "2026-05-01"},
    ]
    p = S.pick_price(rows)
    assert p["liste"] == "YENI" and p["fiyat"] == 110 and p["listeSayisi"] == 3 and p["genel"]


def test_firm_ranges_split_a_window_across_year_copies():
    firms = {2024: "211", 2025: "211", 2026: "411"}
    assert S.firm_ranges(firms, date(2025, 8, 1), date(2026, 8, 18)) == [
        ("211", date(2025, 8, 1), date(2026, 1, 1)), ("411", date(2026, 1, 1), date(2026, 8, 18))]


def test_channel_names_are_checked_before_sql():
    assert S.channel_list("BAYI, KITAPCI") == ["BAYI", "KITAPCI"]
    with pytest.raises(S.SourceError):
        S.channel_list("BAYI'; DROP TABLE x --")


# ------------------------------------------------------------------ maliyet


def test_cost_is_unknown_without_the_m9_provider_and_never_invented():
    S.register_cost_provider(None)
    out = S.unit_costs(["L1"], "m9")
    assert out["L1"]["birim"] is None and out["L1"]["kaynak"] == "bilinmiyor"
    S.register_cost_provider(lambda codes: {"L1": {"birim": 40.0, "tarih": "2026-09-01"}})
    try:
        assert S.unit_costs(["L1", "L2"], "m9")["L1"]["birim"] == 40.0
        assert S.unit_costs(["L1", "L2"], "m9")["L2"]["birim"] is None
    finally:
        S.register_cost_provider(None)
    logo = S.unit_costs(["L1"], "logo", {"L1": {"birim": 35.0, "tarih": "2026-06-30"}})
    assert logo["L1"]["birim"] == 35.0 and logo["L1"]["tahmini"] is True


# ------------------------------------------------------------------ teklif hesabı ve onay


def test_totals_ask_for_approval_above_the_discount_threshold_and_margin_needs_cost():
    lines = [{"stok": "A", "adet": 10, "listeTutar": 1000.0, "netTutar": 650.0, "indirim": 0.35, "maliyetBirim": None},
             {"stok": "B", "adet": 10, "listeTutar": 1000.0, "netTutar": 800.0, "indirim": 0.20, "maliyetBirim": None}]
    t = C.quote_totals(lines, ST)
    assert t["onayGerekli"] and [r["kod"] for r in t["onayNedenleri"]] == ["indirim"]
    assert t["marj"] is None and t["maliyetKapsami"] == 0.0 and t["toplamNet"] == 1450.0
    lines[0]["maliyetBirim"] = 70.0                               # 650 − 700 < 0: maliyetin altında
    lines[0]["indirim"] = 0.10
    t = C.quote_totals(lines, ST)
    assert t["marj"] == pytest.approx((650 - 700) / 650, abs=1e-4)
    assert [r["kod"] for r in t["onayNedenleri"]] == ["marj"]
    assert t["maliyetKapsami"] == pytest.approx(650 / 1450, abs=1e-4)


def test_quote_flow_two_eyes_and_acceptance_wins_the_opportunity(engine):
    _books(engine)
    _account(engine)
    o = C.create_opportunity(engine, T, "ali", {"accountRef": "K001", "ad": "Yeni çalışan paketi"})
    assert o["asama"] == "aday" and o["kurum"] == "Örnek Bankası A.Ş."
    q = C.create_quote(engine, ST, T, "ali", False, o["id"], {"kalemler": [{"stok": "L1", "adet": 300, "indirim": 35}]})
    assert q["durum"] == "taslak" and q["kalemler"][0]["netBirim"] == pytest.approx(65.0)
    assert C.opportunity(engine, T, "ali", False, False, o["id"])["asama"] == "teklif"
    q = C.submit_quote(engine, ST, T, "ali", False, q["id"])
    assert q["durum"] == "onayda" and q["onayGerekli"]
    with pytest.raises(C.CorporateError) as e:
        C.decide_quote(engine, T, "ali", q["id"], True, None)       # gönderen onaylayamaz
    assert e.value.status == 409
    with pytest.raises(C.CorporateError):
        C.decide_quote(engine, T, "mudur", q["id"], False, "")      # geri gönderme gerekçe ister
    q = C.decide_quote(engine, T, "mudur", q["id"], True, "uygun")
    assert q["durum"] == "hazir" and q["onaylayan"] == "mudur"
    q = C.mark_sent(engine, T, "ali", False, q["id"])
    assert q["durum"] == "gonderildi" and q["firsat"]["asama"] == "karar"
    q = C.record_result(engine, T, "ali", False, q["id"], {"sonuc": "kabul"})
    opp = C.opportunity(engine, T, "ali", False, False, o["id"])
    assert opp["asama"] == "kazanildi" and opp["deger"] == pytest.approx(19_500.0)


def test_quote_without_threshold_breach_is_ready_without_approval(engine):
    _books(engine)
    o = C.create_opportunity(engine, T, "ali", {"kurum": "Kartsız Vakıf", "ad": "Çocuk kütüphanesi"})
    q = C.create_quote(engine, ST, T, "ali", False, o["id"], {"kalemler": [{"stok": "C1", "adet": 20, "indirim": 10}]})
    assert C.submit_quote(engine, ST, T, "ali", False, q["id"])["durum"] == "hazir"
    with pytest.raises(C.CorporateError):
        C.update_quote(engine, ST, T, "ali", False, q["id"], {"kalemler": [{"stok": "C1", "adet": 30}]})  # taslak değil


def test_quote_refuses_duplicate_books_and_missing_price(engine):
    _books(engine)
    o = C.create_opportunity(engine, T, "ali", {"kurum": "X", "ad": "Y"})
    with pytest.raises(C.CorporateError):
        C.create_quote(engine, ST, T, "ali", False, o["id"], {"kalemler": [{"stok": "L1", "adet": 1}, {"stok": "L1", "adet": 2}]})
    with pytest.raises(C.CorporateError):
        C.create_quote(engine, ST, T, "ali", False, o["id"], {"kalemler": [{"stok": "L5", "adet": 1}]})
    q = C.create_quote(engine, ST, T, "ali", False, o["id"], {"kalemler": [{"stok": "L5", "adet": 1, "listeFiyati": 75}]})
    assert q["kalemler"][0]["fiyatKaynak"] == "elle"


def test_representative_sees_only_own_opportunities_unless_allowed(engine):
    o = C.create_opportunity(engine, T, "ali", {"kurum": "X", "ad": "Y"})
    C.create_opportunity(engine, T, "veli", {"kurum": "Z", "ad": "W"})
    assert [i["kurum"] for i in C.list_opportunities(engine, T, "ali", False, False)["items"]] == ["X"]
    assert len(C.list_opportunities(engine, T, "ali", True, False)["items"]) == 2
    with pytest.raises(C.CorporateError) as e:
        C.opportunity(engine, T, "veli", False, False, o["id"])
    assert e.value.status == 403
    with pytest.raises(C.CorporateError):
        C.update_opportunity(engine, T, "ali", False, o["id"], {"asama": "kaybedildi"})   # neden zorunlu
    out, diff = C.update_opportunity(engine, T, "ali", False, o["id"], {"asama": "kaybedildi", "kayipSinif": "fiyat"})
    assert out["asama"] == "kaybedildi" and out["closedAt"] and "asama" in diff
    assert C.pipeline_summary(engine, T, "ali", False)["nedenler"][0]["kod"] == "fiyat"


# ------------------------------------------------------------------ paket önerisi


def test_package_candidates_need_theme_stock_and_price_and_alternatives_differ(engine):
    _books(engine)
    r = C.suggest_packages(engine, ST, {"temalar": ["liderlik ve yönetim"], "kisi": 100, "kitapSayisi": 2, "indirim": 0})
    assert r["elenen"] == {"stokYetersiz": 1, "fiyatYok": 1, "yasUymuyor": 0}
    first = [l["stok"] for l in r["alternatifler"][0]["kalemler"]]
    assert first == ["L1", "L2"]                                    # bu yılın satışına göre
    second = [l["stok"] for l in r["alternatifler"][1]["kalemler"]]
    assert set(first).isdisjoint(second)
    assert all(a["kalemler"][0]["adet"] == 100 for a in r["alternatifler"])


def test_package_respects_the_per_package_budget_and_age(engine):
    _books(engine)
    r = C.suggest_packages(engine, ST, {"temalar": ["Liderlik ve yönetim"], "kisi": 100, "kitapSayisi": 2, "butce": 130,
                                        "butceTuru": "kisi", "indirim": 0})
    for a in r["alternatifler"]:
        assert a["paketNet"] <= 130 + 1e-6
    kids = C.suggest_packages(engine, ST, {"temalar": ["Çocuk kütüphanesi"], "kisi": 10, "kitapSayisi": 1, "yasMin": 7, "yasMax": 8})
    assert [a["kalemler"][0]["stok"] for a in kids["alternatifler"]] == ["C1"]
    assert kids["elenen"]["yasUymuyor"] == 1
    with pytest.raises(C.CorporateError):
        C.suggest_packages(engine, ST, {"temalar": ["Uydurma tema"], "kisi": 1})


def test_suggested_theme_does_not_enter_packages_until_approved(engine):
    _books(engine)
    with engine.begin() as c:
        c.execute(C.THEMES.insert().values(stok_kodu="C2", tema="Liderlik ve yönetim", kaynak="oneri", durum="onerildi",
                                           zaman=datetime.now(timezone.utc)))
    r = C.suggest_packages(engine, ST, {"temalar": ["Liderlik ve yönetim"], "kisi": 10, "kitapSayisi": 1, "alternatif": 12})
    assert "C2" not in {a["kalemler"][0]["stok"] for a in r["alternatifler"]}
    C.set_theme(engine, ST, "editor", "C2", "Liderlik ve yönetim", "onayla")
    r = C.suggest_packages(engine, ST, {"temalar": ["Liderlik ve yönetim"], "kisi": 10, "kitapSayisi": 1, "alternatif": 12})
    assert "C2" in {a["kalemler"][0]["stok"] for a in r["alternatifler"]}
    with pytest.raises(C.CorporateError):
        C.set_theme(engine, ST, "editor", "L1", "Liderlik ve yönetim", "kaldir")       # CRM bağı burada kaldırılmaz


# ------------------------------------------------------------------ hatırlatma ve bayi


def test_reminders_come_from_last_years_same_month_and_are_not_duplicated(engine):
    _account(engine, "K001", "Belediye")
    with engine.begin() as c:
        c.execute(C.SALES.insert(), [
            {"logo_code": "K001", "year": 2025, "month": 10, "ciro": 12_000.0, "adet": 400, "fatura": 1, "son": "2025-10-12"},
            {"logo_code": "K002", "year": 2025, "month": 12, "ciro": -50.0, "adet": -1, "fatura": 0, "son": None},
        ])
    assert C.reminder_months(date(2026, 9, 28), 45) == ["2026-09", "2026-10", "2026-11"]
    out = C.generate_reminders(engine, T, ST, date(2026, 9, 28))
    assert out["yeni"] == 1
    assert C.generate_reminders(engine, T, ST, date(2026, 9, 28))["yeni"] == 0
    rid = C.list_reminders(engine, T, ST, ay="2026-10")["items"][0]["id"]
    o = C.opportunity_from_reminder(engine, T, "ali", rid, {})
    assert o["kaynak"] == "hatirlatma" and o["accountRef"] == "K001" and o["kararTarihi"] == "2026-10-01"
    with pytest.raises(C.CorporateError):
        C.opportunity_from_reminder(engine, T, "ali", rid, {})


def test_silent_dealers_are_counted_from_the_logo_data_end_and_classed_abc(engine):
    C.meta_set(engine, "data_end", {"date": "2026-08-17"})
    with engine.begin() as c:
        c.execute(C.DEALERS.insert(), [
            {"logo_code": "B1", "unvan": "Büyük", "son_fatura": "2026-08-10", "fatura_12ay": 30, "ciro_12ay": 900.0, "sinif": "A"},
            {"logo_code": "B2", "unvan": "Sessiz", "son_fatura": "2026-05-01", "fatura_12ay": 3, "ciro_12ay": 80.0, "sinif": "B"},
        ])
    d = C.dealer_rows(engine, gun=60, durum="sessiz")
    assert [i["logoKod"] for i in d["items"]] == ["B2"] and d["items"][0]["gun"] == 108
    assert d["counts"] == {"aktif": 1, "sessiz": 1}
    assert C.abc_classes({"a": 800.0, "b": 150.0, "c": 50.0, "d": 0.0}) == {"a": "A", "b": "B", "c": "C", "d": "C"}
    assert C.dealers_csv(d).startswith("﻿Cari kodu;")


# ------------------------------------------------------------------ ZEKİ AI


class _Choice(SimpleNamespace):
    def confident(self, p, min_margin=0.0, min_coverage=0.0):
        return self.probability is not None and self.probability >= p and self.margin >= min_margin


class _ChooseLlm:
    def __init__(self, pick, p):
        self.pick, self.p = pick, p

    def choose(self, prompt, choices):
        assert self.pick in choices
        return _Choice(choice=self.pick, probability=self.p, margin=self.p - (1 - self.p), method="logprobs")


class _TextLlm:
    def __init__(self, answer):
        self.answer = answer

    def chat(self, messages, **kw):
        return self.answer


def test_segment_choice_needs_the_probability_threshold():
    assert C.ask_segment(_ChooseLlm("Kamu kurumu", 0.95), "X Belediyesi", ST)["secim"] == "kamu"
    low = C.ask_segment(_ChooseLlm("Okul", 0.55), "X", ST)
    assert low["secim"] == "okul" and not low["emin"]
    assert C.ask_segment(_TextLlm("Üniversite"), "X Üniversitesi", ST)["secim"] == "universite"
    assert C.ask_segment(_TextLlm("belki okul"), "X", ST)["secim"] is None


def test_theme_none_option_and_loss_class():
    book = SimpleNamespace(ad="Kitap", stok_kodu="K", yazar=None, hedef=None, yaslar=None, yas_min=None, yas_max=None,
                           turler=None, kitaplik=None, ozet=None)
    assert C.ask_theme(_ChooseLlm(C.NONE_THEME, 0.99), book, ["Aile"], ST)["secim"] is None
    assert C.ask_theme(_ChooseLlm("Aile", 0.99), book, ["Aile"], ST)["secim"] == "Aile"
    assert C.ask_loss(_ChooseLlm("Rakip tercih edildi", 0.9), "başka yayınevi daha ucuz verdi", ST)["secim"] == "rakip"


def test_letter_draft_drops_lines_with_numbers():
    llm = _TextLlm("Sayın Yetkili,\n300 adet kitap için %30 indirim.\nKitaplarımızı öneriyoruz.")
    text = C.draft_letter(llm, ST, {"kurum": "K", "ad": "F", "tema": None}, [{"stok": "L1", "ad": "Kitap", "yazar": None}])
    assert "300" not in text and "Kitaplarımızı" in text


def test_model_tasks_skip_without_model(engine):
    assert C.run_model_tasks(engine, T, ST, None, 60)["atlandi"] == "model tanımlı değil"


# ------------------------------------------------------------------ belge ve yetki


def test_excel_has_quote_and_order_sheets_without_cost(engine):
    _books(engine)
    o = C.create_opportunity(engine, T, "ali", {"kurum": "K", "ad": "F"})
    q = C.create_quote(engine, ST, T, "ali", False, o["id"], {"kalemler": [{"stok": "L1", "adet": 2}]})
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(D.xlsx(q, "Timaş Yayınları")))
    assert wb.sheetnames == ["Teklif", "Sipariş satırları"]
    cells = " ".join(str(c.value) for row in wb["Teklif"].iter_rows() for c in row if c.value is not None)
    assert "maliyet" not in cells.lower() and "TASLAK" in cells


def test_access_rules_for_corporate_endpoints():
    assert A.rule_for("/api/v1/corporate/run-due") == A.SYSTEM
    assert A.rule_for("/api/v1/corporate/opportunities") == frozenset({A.page("kurumsal-satis")})
    f = A.features_for
    assert f("POST", "/api/v1/corporate/opportunities") == ["ozellik:kurumsal.teklif"]
    assert f("POST", "/api/v1/corporate/quotes/q1/submit") == ["ozellik:kurumsal.teklif"]
    assert f("POST", "/api/v1/corporate/quotes/q1/approve") == []           # açıkça verilen yetki ucun içinde
    assert f("GET", "/api/v1/corporate/opportunities") == []
    assert f("GET", "/api/v1/corporate/b2b/dealers") == ["ozellik:kurumsal.b2b"]
    assert f("GET", "/api/v1/corporate/b2b/dealers.csv") == ["ozellik:kurumsal.b2b", "ozellik:veri.disa-aktar"]
    assert f("GET", "/api/v1/corporate/quotes/q1/document.pdf") == ["ozellik:veri.disa-aktar"]
    assert f("POST", "/api/v1/corporate/themes/L1/approve") == ["ozellik:kurumsal.tema-onay"]
    assert "ozellik:kurumsal.teklif-onay" in A.explicit_keys()
    cat = json.loads((Path(A.__file__).parent / "access_catalog.json").read_text(encoding="utf-8"))
    assert {"id": "satis", "label": "Satış ve saha"} in cat["areas"]
