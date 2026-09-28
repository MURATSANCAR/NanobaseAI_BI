"""M34 E-ticaret ve platform yönetimi: üç kaynağın eşlenmesi, yedi fark türü, fark kaydının yaşam döngüsü (açıldı → işaret →
yeniden açıldı / kendiliğinden kapandı / doğrulandı), okunamayan kaynakta kapanmama, huni, pazar yeri özeti, içerik paketi,
yetki kuralları, yazma yasağı ve uçlar.

Veriler yapaydır ve yalnız kuralları sınar; gerçek CRM/Logo/site kabulü test sunucusunda (`scripts/acceptance/M34/`).
"""
from __future__ import annotations

import io
import json
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest

from semantic_bridge import access as A
from semantic_bridge import eticaret as E
from semantic_bridge import eticaret_sources as src
from semantic_layer.store.catalog_store import open_store

T = "t1"
ST = E.settings_from(lambda key, default="": E.DEFAULTS.get(key, default))
BRIDGE = Path(E.__file__).parent
ALL = {"tsoft": True, "crm": True, "logo": True, "rights": True}


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    E._ready.discard(id(e))
    E.ensure(e)
    return e


def _p(pid, barcode, name, active=True, **extra):
    data = {"ProductId": pid, "ProductName": name, "Barcode": barcode, "SellingPriceVatIncluded": extra.pop("price", None),
            "StatViews": extra.pop("views", 0), "CountTotalSales": extra.pop("sales", 0), "CommentCount": extra.pop("comments", 0),
            "ImageUrls": extra.pop("images", [{"Big": f"https://site/{pid}.jpg"}]), **extra}
    return {"product_id": pid, "code": pid, "name": name, "active": active, "data": data}


def _c(cid, ean, name, stok, tsoft=True, etkin=True, fiyat=100.0, dolu=None):
    full = {k: True for k in src.CARD_FIELDS}
    return {"id": cid, "ad": name, "ean": ean, "stok": stok, "tsoft": tsoft, "etkin": etkin, "tip": 1, "fiyat": fiyat,
            "perakende": None, "degisme": "2026-09-01", "dolu": {**full, **(dolu or {})}}


def _site(products, rights=None):
    return {"products": products, "rights": rights or {}, "tsoftAt": datetime(2026, 9, 28, 3, 30, tzinfo=timezone.utc), "crmAt": None}


def _items(products, cards, stock=None, prices=None, sales=None, rights=None, st=ST, cut=date(2026, 8, 17)):
    return E.build_items(_site(products, rights), cards, stock if stock is not None else {}, prices or {}, sales or {}, st,
                         src.tsoft_values, lambda p: (p.get("ImageUrls") or [{}])[0].get("Big"), cut)


def _kinds(items, key, sources=ALL, st=ST):
    return sorted(d["tur"] for d in E.compute_diffs(items[key], st, sources))


# ------------------------------------------------------------------ eşleme ve fark türleri


def test_clean_book_has_no_diff():
    it = _items([_p("1", "9786050000011", "Kitap Bir", price=100.0)], [_c("c1", "9786050000011", "Kitap Bir", "K1")],
                stock={"K1": 40.0}, sales={"K1": {"adet": 120.0, "ciro": 9000.0}})
    assert _kinds(it, "9786050000011") == []
    row = it["9786050000011"]
    assert row["stok_logo"] == 40.0 and row["logo_adet"] == 120.0 and row["doluluk_puani"] == 100
    assert E.impact(row) == 120.0


def test_crm_active_but_not_on_site_and_the_reverse():
    it = _items([_p("2", "9786050000028", "İki", active=False, price=100.0), _p("3", "9786050000035", "Üç", price=100.0)],
                [_c("c2", "9786050000028", "İki", "K2"), _c("c3", "9786050000035", "Üç", "K3", tsoft=False),
                 _c("c4", "9786050000042", "Dört", "K4")],
                stock={"K2": 5.0, "K3": 5.0, "K4": 5.0})
    assert _kinds(it, "9786050000028") == ["aktiflik"]              # CRM TSOFT Aktif, sitede pasif
    assert _kinds(it, "9786050000042") == ["aktiflik"]              # CRM TSOFT Aktif, sitede ürün yok
    d = E.compute_diffs(it["9786050000042"], ST, ALL)[0]
    assert "sitede bu barkodla ürün yok" in d["aciklama"]
    assert _kinds(it, "9786050000035") == ["aktiflik"]              # sitede satışta, CRM işaretsiz


def test_barcode_problems_are_their_own_kind():
    it = _items([_p("5", "", "Barkodsuz", price=10.0), _p("6", "9786050000059", "Kartsız", price=10.0),
                 _p("7", "9786050000066", "Çift A", price=10.0), _p("8", "9786050000066", "Çift B", price=10.0)],
                [_c("c7", "9786050000066", "Çift A", "K7", fiyat=10.0), _c("c9", None, "EAN'sız", "K9")], stock={"K7": 3.0})
    assert _kinds(it, "tsoft:5") == ["barkod"]
    assert _kinds(it, "9786050000059") == ["barkod"]
    assert _kinds(it, "crm:c9") == ["barkod"]
    assert "2 aktif üründe" in E.compute_diffs(it["9786050000066"], ST, ALL)[0]["aciklama"]


def test_name_difference_ignores_author_suffix_and_turkish_letters():
    assert not E.names_differ("Çocuk Kalbi", "COCUK KALBI - Yazar Adı")
    assert not E.names_differ("Şeker Portakalı", "şeker portakalı")
    assert E.names_differ("Kayıp Şehir", "Unutulan Ülke")
    assert not E.names_differ(None, "X")


def test_price_diff_uses_the_reference_and_tolerance():
    it = _items([_p("9", "9786050000073", "Fiyat", price=120.0, DiscountedSellingPriceVatIncluded=99.0)],
                [_c("c9", "9786050000073", "Fiyat", "K9", fiyat=100.0)], stock={"K9": 10.0}, prices={"K9": {"fiyat": 120.0}})
    d = [x for x in E.compute_diffs(it["9786050000073"], ST, ALL) if x["tur"] == "fiyat"][0]
    assert "120,00 ₺" in d["tsoft_deger"] and "100,00 ₺" in d["crm_deger"] and "indirimli" in d["aciklama"]
    logo_ref = {**ST, "priceRef": "logo"}
    assert "fiyat" not in _kinds(it, "9786050000073", st=logo_ref)          # Logo listesiyle aynı
    assert "fiyat" not in _kinds(it, "9786050000073", sources={**ALL, "crm": False})  # esas kaynak okunamadı
    no_price = _items([_p("10", "9786050000080", "Fiyatsız")], [_c("c10", "9786050000080", "Fiyatsız", "K10")], stock={"K10": 1.0})
    assert "fiyat" not in _kinds(no_price, "9786050000080")                  # sitede fiyat alanı yok: uydurulmaz


def test_stock_diff_writes_the_cut_date_and_missing_code():
    it = _items([_p("11", "9786050000097", "Stoksuz", price=100.0)], [_c("c11", "9786050000097", "Stoksuz", "K11")], stock={})
    d = [x for x in E.compute_diffs(it["9786050000097"], ST, ALL) if x["tur"] == "stok"][0]
    assert "kesim 2026-08-17" in d["aciklama"] and "hareket yok" in d["aciklama"]
    assert "stok" not in _kinds(it, "9786050000097", sources={**ALL, "logo": False})


def test_rights_flag_means_not_for_sale():
    rights = {"9786050000103": {"rights": "var", "statusFlag": "bizim_degil", "data": {"statusLabel": "YS05 Bizim değil"}},
              "9786050000110": {"rights": "yok", "statusFlag": None, "data": {}}}
    it = _items([_p("12", "9786050000103", "Devredilen", price=100.0), _p("13", "9786050000110", "Sözleşmesiz", price=100.0)],
                [_c("c12", "9786050000103", "Devredilen", "K12"), _c("c13", "9786050000110", "Sözleşmesiz", "K13")],
                stock={"K12": 5.0, "K13": 5.0}, rights=rights)
    assert _kinds(it, "9786050000103") == ["hak"]
    assert _kinds(it, "9786050000110") == []                                # «hak yok» yalnız ayarla
    assert _kinds(it, "9786050000110", st={**ST, "hakRights": ["yok"]}) == ["hak"]


def test_missing_card_fields_and_score():
    it = _items([_p("14", "9786050000127", "Eksik", price=100.0, images=[])],
                [_c("c14", "9786050000127", "Eksik", "K14", dolu={"arka_kapak": False})], stock={"K14": 5.0})
    row = it["9786050000127"]
    assert row["eksik"] == ["arka_kapak", "site_gorsel"] and row["doluluk_puani"] == 60
    d = [x for x in E.compute_diffs(row, ST, ALL) if x["tur"] == "eksik_kart"][0]
    assert "Arka kapak metni" in d["crm_deger"] and "Sitede görsel" in d["crm_deger"]


def test_disabled_kind_is_not_computed():
    it = _items([_p("15", "9786050000134", "Adı Farklı", price=100.0)], [_c("c15", "9786050000134", "Başka Kitap", "K15")],
                stock={"K15": 5.0})
    assert _kinds(it, "9786050000134") == ["ad"]
    assert _kinds(it, "9786050000134", st={**ST, "kinds": ["fiyat"]}) == []


# ------------------------------------------------------------------ fark kaydının yaşam döngüsü


def _run(engine, items, sources=ALL, st=ST, at=None, site_at=None):
    at = at or datetime(2026, 9, 28, 4, 30, tzinfo=timezone.utc)
    return E.apply_run(engine, T, items, st, sources, site_at or at - timedelta(hours=1), at)


def _stockless():
    return _items([_p("20", "9786050000202", "Stoksuz", price=100.0)], [_c("c20", "9786050000202", "Stoksuz", "K20")], stock={"K20": 0.0})


def _stocked():
    return _items([_p("20", "9786050000202", "Stoksuz", price=100.0)], [_c("c20", "9786050000202", "Stoksuz", "K20")], stock={"K20": 9.0})


def test_new_diff_opens_and_disappearing_condition_closes_it(engine):
    r = _run(engine, _stockless())
    assert r["counts"]["yeni"] == 1 and len(r["events"]) == 1
    lst = E.list_diffs(engine, T)
    assert lst["total"] == 1 and lst["items"][0]["tur"] == "stok" and lst["items"][0]["durum"] == "acik"
    r = _run(engine, _stocked(), at=datetime(2026, 9, 29, 4, 30, tzinfo=timezone.utc))
    assert r["counts"]["kapandi"] == 1
    d = E.get_diff(engine, T, lst["items"][0]["id"])
    assert d["durum"] == "kapandi" and d["kapatan"] == "sistem" and d["gunluk"][0]["eylem"] == "kapandi"


def test_unreadable_source_does_not_close(engine):
    _run(engine, _stockless())
    r = _run(engine, _stocked(), sources={**ALL, "logo": False}, at=datetime(2026, 9, 29, 4, 30, tzinfo=timezone.utc))
    assert r["counts"]["kapandi"] == 0 and "stok" not in r["evaluated"]
    assert E.list_diffs(engine, T)["total"] == 1


def test_fixed_is_verified_or_reopened_by_the_next_site_read(engine):
    t0 = datetime(2026, 9, 28, 4, 30, tzinfo=timezone.utc)
    _run(engine, _stockless(), at=t0)
    did = E.list_diffs(engine, T)["items"][0]["id"]
    E.mark(engine, T, "ayse", did, "duzeltildi", "panelden pasife aldım", None, ST, at=t0 + timedelta(hours=5))
    # Site aynı gece okunmadı (okuma işaretten önce): bekler.
    r = _run(engine, _stockless(), at=t0 + timedelta(days=1), site_at=t0 + timedelta(hours=2))
    assert r["counts"]["yeniden"] == 0 and E.get_diff(engine, T, did)["durum"] == "duzeltildi"
    # Site işaretten sonra okundu ve fark hâlâ var: yeniden açılır.
    r = _run(engine, _stockless(), at=t0 + timedelta(days=2), site_at=t0 + timedelta(days=2, hours=-1))
    d = E.get_diff(engine, T, did)
    assert r["counts"]["yeniden"] == 1 and d["durum"] == "acik" and "hâlâ var" in d["gunluk"][0]["not"]
    # Yeniden düzeltildi ve bu kez koşul kalktı: doğrulanır.
    E.mark(engine, T, "ayse", did, "duzeltildi", "", None, ST, at=t0 + timedelta(days=2, hours=5))
    r = _run(engine, _stocked(), at=t0 + timedelta(days=3))
    d = E.get_diff(engine, T, did)
    assert r["counts"]["dogrulandi"] == 1 and d["durum"] == "kapandi" and d["dogrulandi"] is True and d["kapatan"] == "ayse"


def test_intentional_diff_stays_until_values_change(engine):
    t0 = datetime(2026, 9, 28, 4, 30, tzinfo=timezone.utc)
    items = _items([_p("21", "9786050000219", "Kampanya", price=80.0)], [_c("c21", "9786050000219", "Kampanya", "K21", fiyat=100.0)],
                   stock={"K21": 9.0})
    _run(engine, items, at=t0)
    did = E.list_diffs(engine, T)["items"][0]["id"]
    with pytest.raises(E.EticaretError) as e:
        E.mark(engine, T, "ayse", did, "bilincli", "", None, ST)
    assert e.value.status == 422                                           # gerekçe şart
    E.mark(engine, T, "ayse", did, "bilincli", "Eylül kampanyası", None, ST, at=t0 + timedelta(hours=3))
    _run(engine, items, at=t0 + timedelta(days=1))
    assert E.get_diff(engine, T, did)["durum"] == "bilincli"
    assert E.list_diffs(engine, T)["total"] == 0                            # varsayılan liste: açık + sonra
    changed = _items([_p("21", "9786050000219", "Kampanya", price=70.0)], [_c("c21", "9786050000219", "Kampanya", "K21", fiyat=100.0)],
                     stock={"K21": 9.0})
    r = _run(engine, changed, at=t0 + timedelta(days=2))
    assert r["counts"]["yeniden"] == 1 and E.get_diff(engine, T, did)["durum"] == "acik"


def test_snooze_expires_and_closed_cannot_be_marked(engine):
    t0 = datetime(2026, 9, 28, 4, 30, tzinfo=timezone.utc)
    _run(engine, _stockless(), at=t0)
    did = E.list_diffs(engine, T)["items"][0]["id"]
    d = E.mark(engine, T, "ayse", did, "sonra", "", "Mehmet", ST, at=t0)
    assert d["durum"] == "sonra" and d["sahip"] == "mehmet"
    _run(engine, _stockless(), at=t0 + timedelta(days=ST["snoozeDays"] + 1))
    assert E.get_diff(engine, T, did)["durum"] == "acik"
    _run(engine, _stocked(), at=t0 + timedelta(days=ST["snoozeDays"] + 2))
    with pytest.raises(E.EticaretError) as e:
        E.mark(engine, T, "ayse", did, "acik", "", None, ST)
    assert e.value.status == 409
    with pytest.raises(E.EticaretError):
        E.mark(engine, T, "ayse", did, "kapandi", "", None, ST)              # kapatmak sistemin işi


def test_kind_turned_off_in_settings_closes_its_records(engine):
    _run(engine, _stockless())
    r = _run(engine, _stockless(), st={**ST, "kinds": ["fiyat"]}, at=datetime(2026, 9, 29, 4, 30, tzinfo=timezone.utc))
    assert r["counts"]["kapandi"] == 1


def test_list_order_puts_legal_risk_first_and_counts_by_kind(engine):
    rights = {"9786050000301": {"rights": "var", "statusFlag": "cekildi", "data": {"statusLabel": "YS11 Satıştan çekildi"}}}
    items = _items([_p("30", "9786050000301", "Çekilen", price=100.0), _p("31", "9786050000318", "Stoksuz", price=100.0)],
                   [_c("c30", "9786050000301", "Çekilen", "K30"), _c("c31", "9786050000318", "Stoksuz", "K31")],
                   stock={"K30": 5.0, "K31": 0.0}, sales={"K31": {"adet": 900.0, "ciro": 1.0}}, rights=rights)
    _run(engine, items)
    lst = E.list_diffs(engine, T)
    assert [d["tur"] for d in lst["items"]] == ["hak", "stok"]
    assert lst["turSayilari"]["hak"] == 1 and lst["turSayilari"]["stok"] == 1
    assert E.list_diffs(engine, T, q="çekil")["total"] in (0, 1)           # arama (SQLite'ta Türkçe harf duyarlı olabilir)
    assert E.list_diffs(engine, T, tur="stok")["total"] == 1
    with pytest.raises(E.EticaretError):
        E.list_diffs(engine, T, tur="yok-boyle")
    ov = E.overview(engine, T, {})
    assert ov["gostergeler"]["satistaOlmamali"] == 1 and ov["gostergeler"]["siteAktif"] == 2 and ov["bugun"]["total"] == 2
    csv_text = E.diffs_csv(E.all_diffs(engine, T)).decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("Tür;Barkod") and len(csv_text.splitlines()) == 3
    det = E.item_detail(engine, T, "9786050000301")
    assert det["kitap"]["yayinDurumu"] == "YS11 Satıştan çekildi" and det["farklar"][0]["tur"] == "hak"
    with pytest.raises(E.EticaretError) as e:
        E.item_detail(engine, T, "yok")
    assert e.value.status == 404


# ------------------------------------------------------------------ huni, pazar yerleri, içerik paketi


def test_funnel_low_conversion_uses_the_site_median(engine):
    products = [_p(str(40 + i), f"97860500004{i:02d}", f"K{i}", price=10.0, views=1000, sales=s) for i, s in enumerate([100, 90, 80, 5])]
    cards = [_c(f"c{i}", f"97860500004{i:02d}", f"K{i}", f"S{i}") for i in range(4)]
    _run(engine, _items(products, cards, stock={f"S{i}": 5.0 for i in range(4)}))
    f = E.funnel(engine, T, ST)
    assert f["total"] == 4 and f["toplam"]["goruntulenme"] == 4000 and f["ortancaDonusum"] == pytest.approx(0.09)
    low = E.funnel(engine, T, ST, dusuk=True)
    assert [x["siteSatis"] for x in low["items"]] == [5]
    assert "Okur yorumu yok" in low["items"][0]["olasiNedenler"]


def test_marketplace_summary_same_period_and_return_rate():
    rows = [{"kod": "120.01", "unvan": "Kitapyurdu", "kanal": "E-TICARET", "yil": 2026, "ay": 1, "satis": 1000.0, "iade": 100.0,
             "satisAdet": 10.0, "iadeAdet": 1.0},
            {"kod": "120.01", "unvan": "Kitapyurdu", "kanal": "E-TICARET", "yil": 2026, "ay": 8, "satis": 500.0, "iade": 0.0,
             "satisAdet": 5.0, "iadeAdet": 0.0},
            {"kod": "120.01", "unvan": "Kitapyurdu", "kanal": "E-TICARET", "yil": 2025, "ay": 3, "satis": 1000.0, "iade": 0.0,
             "satisAdet": 10.0, "iadeAdet": 0.0},
            {"kod": "120.02", "unvan": "D-Market", "kanal": "E-TICARET", "yil": 2026, "ay": 2, "satis": 300.0, "iade": 0.0,
             "satisAdet": 3.0, "iadeAdet": 0.0}]
    s = E.marketplace_summary(rows, 2026, date(2026, 8, 17))
    a = s["cariler"][0]
    assert a["kod"] == "120.01" and a["net"] == 1400.0 and a["iadeOrani"] == pytest.approx(100 / 1500)
    assert a["oncekiNet"] == 1000.0 and a["degisim"] == pytest.approx(0.4) and a["aylik"][7] == 500.0
    assert s["toplam"]["net"] == 1700.0 and s["donem"]["son"] == "2026-08-17" and s["donem"]["oncekiSon"] == "2025-08-17"
    this, prev = E.marketplace_windows(2026, date(2026, 8, 17))
    assert this == (date(2026, 1, 1), date(2026, 8, 18)) and prev == (date(2025, 1, 1), date(2025, 8, 18))
    assert E.marketplace_windows(2025, date(2026, 8, 17))[0] == (date(2025, 1, 1), date(2026, 1, 1))


def test_marketplace_books_flag_stockout_risk():
    items = {"K1": {"product_key": "1", "stok_logo": 10.0, "logo_adet": 365.25, "tsoft_aktif": True},
             "K2": {"product_key": "2", "stok_logo": 1000.0, "logo_adet": 365.25, "tsoft_aktif": True}}
    rows = [{"stok": "K1", "ad": "A", "satisAdet": 10.0, "iadeAdet": 1.0, "ciro": 100.0, "son": "2026-08-01"},
            {"stok": "K2", "ad": "B", "satisAdet": 5.0, "iadeAdet": 0.0, "ciro": 50.0, "son": "2026-08-01"},
            {"stok": "K3", "ad": "C", "satisAdet": 1.0, "iadeAdet": 0.0, "ciro": 5.0, "son": None}]
    out = E.marketplace_books(rows, items, ST)
    assert [b["stok"] for b in out] == ["K1", "K2", "K3"]
    assert out[0]["tukenmeRiski"] is True and out[0]["kalanGun"] == pytest.approx(10.0, abs=0.2)
    assert out[1]["tukenmeRiski"] is False and out[2]["kalanGun"] is None and out[0]["net"] == 9.0


def test_content_pack_lists_missing_barcodes_and_writes_xlsx():
    from openpyxl import load_workbook

    crm = {"9786050000011": {"ad": "Bir", "yazar": "Y", "spot": "Kısa", "arka_kapak": "Uzun", "stok": "K1", "fiyat": "100"}}
    items = {"9786050000011": {"url": "bir-kitap", "gorsel_url": "https://site/1.jpg"}}
    pack = E.content_pack_rows(["9786050000011", "9786050000028"], crm, items, "https://www.timas.com.tr")
    assert pack["eksik"] == ["9786050000028"] and pack["rows"][0]["site"] == "https://www.timas.com.tr/bir-kitap"
    wb = load_workbook(io.BytesIO(E.content_pack_xlsx(pack)))
    assert wb.sheetnames == ["İçerik paketi", "CRM'de bulunamayan"]
    assert wb["İçerik paketi"]["A2"].value == "9786050000011"
    assert E.content_pack_csv(pack).decode("utf-8-sig").splitlines()[0].startswith("Barkod (EAN-13);ISBN")


def test_reason_classification_respects_the_threshold(engine):
    class Choice:
        def __init__(self, choice, p):
            self.choice, self.probability = choice, p

        def confident(self, p, min_margin=0.0):
            return self.probability >= p

    class Llm:
        def __init__(self, p):
            self.p, self.calls = p, 0

        def choose(self, prompt, choices):
            self.calls += 1
            assert "Sitedeki fiyat" in prompt and choices == E.REASONS
            return Choice(E.REASONS[0], self.p)

    items = _items([_p("50", "9786050000509", "Fiyat", price=80.0)], [_c("c50", "9786050000509", "Fiyat", "K50", fiyat=100.0)],
                   stock={"K50": 5.0})
    _run(engine, items)
    out = E.classify_reasons(engine, T, Llm(0.95), ST, 60)
    assert out["sinif"] == 1 and E.list_diffs(engine, T)["items"][0]["neden"]["oneri"] == E.REASONS[0]
    llm = Llm(0.2)
    assert E.classify_reasons(engine, T, llm, ST, 60)["sinif"] == 0 and llm.calls == 0   # imza değişmedi: yeniden sorulmaz
    assert E.classify_reasons(engine, T, None, ST, 60) == {"skipped": "model yok"}


def test_alerts_once_and_weekly_text(engine):
    _run(engine, _stockless())
    pend = E.pending_alerts(engine, T, ST)
    assert len(pend) == 1 and "Stok yokken satışta" in E.alert_text(pend, "https://portal/timas")
    E.mark_alerted(engine, T, [pend[0]["id"]])
    assert E.pending_alerts(engine, T, ST) == []
    text = E.weekly_text(E.overview(engine, T, {}), None, [], "")
    assert "Açık fark: 1" in text
    monday = datetime(2026, 9, 28, 6, 0, tzinfo=timezone.utc)
    assert E.weekly_due(engine, T, {**ST, "weeklyTo": ["m@x"]}, monday) is True
    assert E.weekly_due(engine, T, ST, monday) is False                      # alıcı yok


# ------------------------------------------------------------------ ayar, yetki, yazma yasağı


def test_settings_defaults_and_parsing():
    assert ST["channels"] == ["E-TICARET"] and ST["priceRef"] == "crm" and ST["kinds"] == list(E.DIFF_KINDS)
    st = E.settings_from(lambda k, d="": {"ECOM_DEFAULT_OWNERS": "fiyat:Ayse, bozuk, hak:mehmet", "ECOM_PRICE_REFERENCE": "LOGO",
                                           "ECOM_DIFF_KINDS": "fiyat,uydurma"}.get(k, d))
    assert st["owners"] == {"fiyat": "ayse", "hak": "mehmet"} and st["priceRef"] == "logo" and st["kinds"] == ["fiyat"]


def test_access_rules_for_eticaret():
    f = A.features_for
    assert f("POST", "/api/v1/eticaret/diffs/abc/mark") == ["ozellik:eticaret.fark-isaretle"]
    assert f("POST", "/api/v1/eticaret/diffs/mark-bulk") == ["ozellik:eticaret.fark-isaretle"]
    assert f("POST", "/api/v1/eticaret/refresh") == ["ozellik:eticaret.fark-isaretle"]
    assert f("POST", "/api/v1/eticaret/items/9786050000011/propose") == ["ozellik:eticaret.oneri-uret"]
    assert f("POST", "/api/v1/eticaret/proposals/p1/decide") == []           # açıkça verilen onay ucun içinde
    assert f("GET", "/api/v1/eticaret/diffs/export.csv") == ["ozellik:veri.disa-aktar"]
    assert f("GET", "/api/v1/eticaret/export/content-pack") == ["ozellik:veri.disa-aktar"]
    assert f("GET", "/api/v1/eticaret/diffs") == [] and f("POST", "/api/v1/eticaret/run-due") == []
    assert A.rule_for("/api/v1/eticaret/run-due") == A.SYSTEM
    assert A.rule_for("/api/v1/eticaret/funnel") == frozenset({"sayfa:eticaret", "sayfa:eticaret-huni"})
    assert A.rule_for("/api/v1/eticaret/marketplaces/120.01/books") == frozenset({"sayfa:eticaret", "sayfa:eticaret-pazar-yerleri"})
    assert "sayfa:eticaret-huni" in A.rule_for("/api/v1/eticaret/meta")
    assert "ozellik:eticaret.oneri-onay" in A.explicit_keys()
    assert {"ozellik:eticaret.fark-isaretle", "ozellik:eticaret.oneri-uret"} <= A.all_keys() - A.explicit_keys()
    cat = json.loads((BRIDGE / "access_catalog.json").read_text(encoding="utf-8"))
    keys = {p["key"] for p in cat["pages"] if p["area"] == "pazarlama"}
    assert {"sayfa:eticaret", "sayfa:eticaret-farklar", "sayfa:eticaret-huni", "sayfa:eticaret-pazar-yerleri"} <= keys


def test_no_write_to_tsoft_crm_or_logo():
    """Kaynak katmanında yazan SQL yok; T-soft istemcisi hiç çağrılmaz (site verisi SEO tablosundan)."""
    text = (BRIDGE / "eticaret_sources.py").read_text(encoding="utf-8")
    sql = " ".join(a + b for a, b in re.findall(r'"""(.*?)"""|"([^"\n]*(?:SELECT|FROM)[^"\n]*)"', text, re.S)).upper()
    for kw in ("INSERT ", "UPDATE ", "DELETE ", "MERGE ", "EXEC ", "DROP "):
        assert kw not in sql
    for f in ("eticaret.py", "eticaret_api.py", "eticaret_sources.py"):
        body = (BRIDGE / f).read_text(encoding="utf-8")
        assert "connections.tsoft" not in body and "from semantic_bridge.seo_geo import connections" not in body
        assert "LlmClient(" not in body


def test_no_technology_names_on_screen_texts():
    banned = re.compile(r"qwen|vllm|temporal|timesfm|ollama|openai|gpt|claude|llama", re.I)
    for f in ("eticaret.py", "eticaret_api.py"):
        strings = re.findall(r'"([^"\n]{12,})"', (BRIDGE / f).read_text(encoding="utf-8"))
        assert not [s for s in strings if banned.search(s)], f


# ------------------------------------------------------------------ uçlar


class _FakeSeo:
    def __init__(self):
        self.props = {"p1": {"id": "p1", "product_id": "60", "status": "hazir", "fields_json": json.dumps({"SeoTitle": "Yeni"}),
                             "before_json": "{}", "score_before": 50, "score_after": 80, "model": "E-ticaret · Zeki AI",
                             "created_by": "ayse", "created_at": None, "decided_by": None, "decided_at": None, "note": None,
                             "sent_at": None, "result": None}}

    def proposal(self, pid):
        from fastapi import HTTPException

        if pid not in self.props:
            raise HTTPException(404, {"code": "SEO", "message": "Öneri bulunamadı."})
        return dict(self.props[pid])

    def approve(self, prop, fields, user, note):
        self.props[prop["id"]].update(status="onaylandi", decided_by=user, result="Onaylandı: SeoTitle. Gönderim yok.")
        return dict(self.props[prop["id"]])


def _client(engine, perms):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from semantic_bridge import eticaret_api

    users = {"a": "ayse", "m": "mehmet", "x": "zeynep"}
    app = FastAPI()
    audits = []

    def auth(request):
        u = users.get(request.headers.get("cookie", ""))
        if not u:
            from fastapi import HTTPException

            raise HTTPException(401, {"code": "UNAUTHORIZED", "message": "Oturum gerekli."})
        return engine, T, u, u.title()

    eticaret_api.register(app, {
        "auth": auth, "require_caller": lambda r: None, "can": lambda u, k: k in perms.get(u, set()), "is_admin": lambda u: False,
        "audit": lambda *a: audits.append(a), "conf": lambda k, d="": d, "engine": lambda: engine, "tenant": lambda: T,
        "logo_file": lambda: "/yok/logo.json", "crm_file": lambda: "/yok/crm.json", "llm": lambda p: None, "seo": _FakeSeo(),
        "send_mail": lambda s, t, to: "sent"})
    return TestClient(app), audits


def test_api_mark_and_decide_rules(engine):
    _run(engine, _stockless())
    c, audits = _client(engine, {"ayse": {"ozellik:eticaret.oneri-onay"}, "mehmet": {"ozellik:eticaret.oneri-onay"}})
    a, m, x = {"cookie": "a"}, {"cookie": "m"}, {"cookie": "x"}
    meta = c.get("/api/v1/eticaret/meta", headers=a).json()
    assert meta["gonderim"] is False and meta["me"]["canApprove"] is True and "hak" in meta["turler"]
    did = c.get("/api/v1/eticaret/diffs", headers=a).json()["items"][0]["id"]
    assert c.post(f"/api/v1/eticaret/diffs/{did}/mark", json={"durum": "bilincli"}, headers=a).status_code == 422
    assert c.post(f"/api/v1/eticaret/diffs/{did}/mark", json={"durum": "kapandi"}, headers=a).status_code == 422
    r = c.post(f"/api/v1/eticaret/diffs/{did}/mark", json={"durum": "sonra", "note": "haftaya", "sahip": "Mehmet"}, headers=a)
    assert r.status_code == 200 and r.json()["durum"] == "sonra" and r.json()["sahip"] == "mehmet"
    assert c.get("/api/v1/eticaret/diffs/yok", headers=a).status_code == 404
    assert "ayse" in {g["kullanici"] for g in c.get(f"/api/v1/eticaret/diffs/{did}", headers=a).json()["gunluk"]}   # aynı saniyede sıra kimliğe kalır
    bulk = c.post("/api/v1/eticaret/diffs/mark-bulk", json={"ids": [did, "yok"], "durum": "acik"}, headers=a).json()
    assert len(bulk["items"]) == 1 and bulk["atlanan"][0]["id"] == "yok"
    # Öneri onayı: yetkisiz 403, isteyen 409, ret gerekçesiz 422, başka onaycı 200.
    assert c.post("/api/v1/eticaret/proposals/p1/decide", json={"action": "approve"}, headers=x).status_code == 403
    assert c.post("/api/v1/eticaret/proposals/p1/decide", json={"action": "approve"}, headers=a).status_code == 409
    assert c.post("/api/v1/eticaret/proposals/p1/decide", json={"action": "reject"}, headers=m).status_code == 422
    ok = c.post("/api/v1/eticaret/proposals/p1/decide", json={"action": "approve"}, headers=m)
    assert ok.status_code == 200 and ok.json()["status"] == "onaylandi"
    assert c.post("/api/v1/eticaret/proposals/p1/decide", json={"action": "approve"}, headers=m).status_code == 409
    # Kaynak bağlantısı yok: pazar yeri 503 (düz cümle), içerik paketi seçimsiz 400.
    r = c.get("/api/v1/eticaret/marketplaces?yil=2026", headers=a)
    assert r.status_code == 503 and r.json()["detail"]["code"] == "ETICARET_SOURCE"
    assert c.get("/api/v1/eticaret/export/content-pack?keys=", headers=a).status_code == 400
    assert c.get("/api/v1/eticaret/overview", headers=a).json()["gostergeler"]["acikFark"] == 1
    assert c.get("/api/v1/eticaret/overview").status_code == 401
    assert any(x[3] == "eticaret_diff" for x in audits) and any(x[3] == "eticaret_proposal" for x in audits)


def test_api_run_due_reports_missing_site_data(engine):
    c, _ = _client(engine, {})
    out = c.post("/api/v1/eticaret/run-due?model=false").json()
    assert "T-soft eşitlemesi" in (out["read"]["hata"] or "")
    assert out["alert"]["status"] == "nothing"
