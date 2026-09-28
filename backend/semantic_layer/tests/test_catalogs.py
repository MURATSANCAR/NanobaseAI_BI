"""M24 Katalog ve bülten: izin kuralı (segment sayısı delinmez, kişi verisi dönmez), öneri puanı ve eleme sayıları,
stok ay sayısı (Baskı önerisi tanımı), fiyat/stok/satış uyarıları ve kabulü, katalog ve bülten onay akışı (iki göz),
dışa aktarım (Excel, tasarımcı paketi, PDF), bülten HTML'i, sonuç dosyası okuma (adres saklanmaz), Zeki AI metninin
denetimi, yetki kuralları ve köprü uçları.

Veriler yapaydır ve yalnız kuralları sınar; gerçek CRM/Logo kabulü test sunucusunda (`scripts/acceptance/M24/`).
"""

from __future__ import annotations

import io
import json
import sqlite3
import zipfile
from datetime import date

import pytest

from semantic_bridge import access as A
from semantic_bridge import catalogs as C
from semantic_bridge import newsletters as N
from semantic_layer.store.catalog_store import open_store

TN = "t1"
REF = date(2026, 9, 28)


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    C._ready.discard(id(e))
    N._ready.discard(id(e))
    C.ensure(e)
    N.ensure(e)
    yield e
    C._ready.discard(id(e))
    N._ready.discard(id(e))


@pytest.fixture(autouse=True)
def _env(monkeypatch):
    for k in ("CATALOG_PRICE_SOURCE", "CATALOG_STOCK_SOURCE", "CATALOG_SCORE_WEIGHTS", "CATALOG_CRITICAL_STOCK_MONTHS",
              "NEWSLETTER_REQUIRE_KVKK", "NEWSLETTER_SUBJECT_OPTIONS", "NEWSLETTER_INTEREST_KEYWORDS"):
        monkeypatch.delenv(k, raising=False)


def _book(i, **kw):
    b = {"id": f"00000000-0000-0000-0000-{i:012d}", "stok": f"15201.{i:04d}", "ad": f"Kitap {i}", "yazar": "Yazar", "marka": "Timaş",
         "isbn": f"978605{i:07d}", "ean": f"978605{i:07d}", "yst": "YS04 Aktif", "flag": None, "satisAcik": True, "tip": 1,
         "hedef": "Yetişkin", "yasBas": None, "yasBit": None, "tur": "Roman", "kitaplik": "Edebiyat", "web": None,
         "ilkYayin": "2020-01-01", "fiyatlar": {"crm": 100.0, "crm-perakende": 90.0, "crm-uzeri": 100.0, "logo": 80.0, "tsoft": 95.0},
         "stokCrm": 600.0, "depo": 500.0, "hiz": 100.0, "yillik": 1200.0, "kapak": "https://ornek.test/k.jpg", "kapakKaynak": "crm",
         "kapakDosya": None, "webUrl": None, "ozelGunler": [], "hakNotu": False, "metinVar": True}
    b.update(kw)
    return b


def _pool(books, days=None):
    return {"books": books, "days": days or [], "readAt": "2026-09-28T07:15:00+00:00", "logoSon": "2026-08-17", "notes": [],
            "ilgiAlanlari": [{"id": "11111111-1111-1111-1111-111111111111", "ad": "Tarih ve Kültür", "etkin": True}]}


# ------------------------------------------------------------------ izin kuralı ve segment


def _contacts() -> sqlite3.Connection:
    c = sqlite3.connect(":memory:")
    c.executescript("""
        CREATE TABLE ContactBase (ContactId TEXT PRIMARY KEY, StateCode INT, DoNotBulkEMail INT, DoNotEMail INT, new_iysonayi INT,
          new_kvkkonayi INT, EMailAddress1 TEXT, new_tarihveakademi INT, new_timasakademi INT, new_sosyalbilimlerincelemearastirma INT,
          new_ElenceliBilgi INT, new_haberdarolmakistiyorum INT, new_DogumTarihi TEXT, BirthDate TEXT);
        CREATE TABLE new_contact_new_kitapilgialanBase (contactid TEXT, new_kitapilgialanid TEXT);
    """)
    rows = [
        # id, state, bulk, email, iys, kvkk, adres, tarih, akademi, dogum
        ("c1", 0, 0, 0, 1, 1, "a@x.test", 1, 0, "1990-05-01"),     # izinli, KVKK var
        ("c2", 0, None, None, 1, 0, "b@x.test", 1, 0, "2010-01-01"),  # izinli (boş izin alanı = izin var sayılmaz değil, 0)
        ("c3", 0, 1, 0, 1, 1, "c@x.test", 1, 0, None),             # toplu e-posta reddi
        ("c4", 0, 0, 1, 1, 1, "d@x.test", 1, 0, None),             # e-posta reddi
        ("c5", 0, 0, 0, 0, 1, "e@x.test", 1, 0, None),             # İYS onayı yok
        ("c6", 0, 0, 0, None, 1, "f@x.test", 1, 0, None),          # İYS boş
        ("c7", 0, 0, 0, 1, 1, "   ", 1, 0, None),                  # adres boş
        ("c8", 1, 0, 0, 1, 1, "g@x.test", 1, 0, None),             # pasif kişi
        ("c9", 0, 0, 0, 1, 1, "h@x.test", 0, 1, None),             # başka ilgi (akademi)
        ("c10", 0, 0, 0, 1, 1, "i@x.test", 0, 0, None),            # ilgisiz; yalnız bağ tablosundan
    ]
    for r in rows:
        c.execute("INSERT INTO ContactBase (ContactId, StateCode, DoNotBulkEMail, DoNotEMail, new_iysonayi, new_kvkkonayi, EMailAddress1,"
                  " new_tarihveakademi, new_timasakademi, new_DogumTarihi) VALUES (?,?,?,?,?,?,?,?,?,?)", r)
    c.execute("INSERT INTO new_contact_new_kitapilgialanBase VALUES ('c10', '11111111-1111-1111-1111-111111111111')")
    return c


def _count(conn, seg, kvkk=False):
    sql = N.segment_sql("", N.normalize_segment(seg), kvkk, REF)
    cur = conn.execute(sql)
    row = dict(zip([d[0] for d in cur.description], cur.fetchone()))
    return N.segment_result(row, kvkk)


def test_permission_rule_is_never_loosened():
    conn = _contacts()
    r = _count(conn, {"ilgiBayraklari": ["new_tarihveakademi"]})
    assert r["izinli"] == 2            # c1, c2 — reddi, İYS'siz, adressiz, pasif olan girmez
    assert r["aday"] == 7              # etkin ve tarih ilgili: c1..c7
    d = r["dagilim"]
    assert (d["topluEpostaReddi"], d["epostaReddi"], d["iysOnayiYok"], d["adresYok"]) == (1, 1, 2, 1)
    assert _count(conn, {"ilgiBayraklari": ["new_tarihveakademi"]}, kvkk=True)["izinli"] == 1   # KVKK şartı: yalnız c1
    # süzgeç kuralı gevşetemez: haberdar olmak isteyen yoksa sayı 0'a iner, izinsiz kişi eklenmez
    assert _count(conn, {"ilgiBayraklari": ["new_tarihveakademi"], "haberdar": True})["izinli"] == 0


def test_interest_flags_and_links_are_or_and_age_filter():
    conn = _contacts()
    r = _count(conn, {"ilgiBayraklari": ["new_timasakademi"], "ilgiAlanlari": ["11111111-1111-1111-1111-111111111111"]})
    assert r["izinli"] == 2            # c9 (bayrak) + c10 (bağ)
    assert _count(conn, {})["izinli"] == 4  # ilgi yok = bütün izinliler: c1, c2, c9, c10
    assert _count(conn, {"ilgiBayraklari": ["new_tarihveakademi"], "yasMin": 18})["izinli"] == 1   # c1 (1990)
    assert _count(conn, {"ilgiBayraklari": ["new_tarihveakademi"], "yasMax": 17})["izinli"] == 1   # c2 (2010)


def test_segment_sql_selects_no_person_column_and_rejects_injection():
    seg = N.normalize_segment({"ilgiBayraklari": ["new_tarihveakademi", "EMailAddress1; DROP TABLE x"],
                               "ilgiAlanlari": ["x' OR 1=1 --", "11111111-1111-1111-1111-111111111111"]})
    assert seg["ilgiBayraklari"] == ["new_tarihveakademi"] and len(seg["ilgiAlanlari"]) == 1
    sql = N.segment_sql("Timas_MSCRM.dbo.", seg, False, REF)
    select = sql.split(" FROM ")[0]
    assert "EMailAddress1 AS" not in select and "FullName" not in select and "c.ContactId," not in select
    assert "new_iysonayi = 1" in N.permit_sql(False) and "DoNotBulkEMail" in N.permit_sql(False)
    res = N.segment_result({"aday": 5, "izinli": 3}, False)
    assert "@" not in json.dumps(res)


# ------------------------------------------------------------------ öneri ve stok


def test_stock_months_follow_the_print_suggestion_definition():
    cfg = C.settings()
    assert C.stock_of(_book(1), cfg)["stokAy"] == 6.0                 # CRM stok 600 ÷ hız 100
    cfg_logo = {**cfg, "stockSource": "logo"}
    assert C.stock_of(_book(1), cfg_logo)["stokAy"] == 5.0            # Logo depo 500 ÷ 100
    st = C.stock_of(_book(2, hiz=0.0), cfg)
    assert st["stokAy"] is None and st["satisYok"] is True
    assert C.stock_of(_book(3, stokCrm=None), cfg)["bilinmiyor"] is True


def test_candidates_exclude_out_of_sale_and_no_stock_and_count_them():
    books = [_book(1), _book(2, flag="cekildi", yst="YS11 Çekildi"), _book(3, satisAcik=False), _book(4, stokCrm=0.0),
             _book(5, hedef="Çocuk", yasBas=6, yasBit=10), _book(6, hiz=300.0, yillik=3600.0)]
    res = C.candidates(_pool(books), {}, C.settings(), ref=REF)
    ids = [x["id"] for x in res["items"]]
    assert _book(2)["id"] not in ids and _book(3)["id"] not in ids and _book(4)["id"] not in ids
    assert res["elenen"] == {"satıştan kalkmış": 2, "stok yok": 1}
    assert ids[0] == _book(6)["id"]                                   # en hızlı satan üstte
    assert "son 12 ayda 3.600 adet satış" in res["items"][0]["gerekce"]
    kids = C.candidates(_pool(books), {"hedef": ["Çocuk"], "yasMin": 7, "yasMax": 9}, C.settings(), ref=REF)
    assert [x["id"] for x in kids["items"]] == [_book(5)["id"]]


def test_special_day_boost_and_new_books():
    day = {"key": "ogretmenler-gunu", "ad": "Öğretmenler Günü", "baslangic": "2026-11-24", "kitapSayisi": 1}
    books = [_book(1, hiz=200.0), _book(2, hiz=100.0, ozelGunler=["ogretmenler-gunu"], ilkYayin="2026-06-01")]
    res = C.candidates(_pool(books, [day]), {"ozelGun": "ogretmenler-gunu"}, C.settings(), ref=REF)
    assert res["items"][0]["id"] == _book(2)["id"]
    assert "«Öğretmenler Günü» ile CRM'de bağlı" in res["items"][0]["gerekce"] and "(yeni)" in res["items"][0]["gerekce"]
    only = C.candidates(_pool(books, [day]), {"ozelGun": "ogretmenler-gunu", "yalnizOzelGun": True}, C.settings(), ref=REF)
    assert only["total"] == 1


def test_weights_renormalize_when_a_part_is_unmeasured():
    books = [_book(1, stokCrm=None, depo=None, hiz=None, yillik=None, ilkYayin=None)]
    res = C.candidates(_pool(books), {}, C.settings(), ref=REF)
    assert res["items"][0]["puan"] == 0.0 and res["items"][0]["parcalar"]["stok"] is None


# ------------------------------------------------------------------ uyarılar


def test_alerts_price_stock_sale_and_missing_card():
    cfg = C.settings()
    b = _book(1, fiyatlar={"crm": 120.0}, stokCrm=50.0, hiz=100.0, flag="iptal", yst="YS01 İptal")
    a = {x["tur"]: x for x in C.item_alerts({"price_snapshot": 100.0}, b, "crm", cfg)}
    assert set(a) >= {"fiyat", "stok", "satis"}
    assert "100,00 ₺ → 120,00 ₺" in a["fiyat"]["metin"] and a["stok"]["seviye"] == "kritik"
    assert C.item_alerts({"price_snapshot": 100.0}, None, "crm", cfg)[0]["tur"] == "crm-yok"
    ok = C.item_alerts({"price_snapshot": 100.0}, _book(2), "crm", cfg)
    assert C.critical_count(ok) == 0


def test_catalog_flow_snapshots_accept_and_two_eyes(engine):
    cfg = C.settings()
    pool = _pool([_book(1), _book(2)])
    cat = C.create(engine, TN, "ayse", {"tur": "bayi", "baslik": "Kış 2026 bayi kataloğu"}, cfg)
    with pytest.raises(C.CatalogError) as e:
        C.transition(engine, TN, "ayse", cat["id"], "submit")
    assert e.value.status == 409                                       # boş katalog
    d, diff = C.set_items(engine, TN, "ayse", cat["id"], [{"crmKitapId": _book(1)["id"], "oneCikan": True, "sayfa": "Kapak içi"},
                                                         {"crmKitapId": _book(2)["id"]}], pool, cfg)
    assert diff == {"eklenen": 2, "cikan": 0, "toplam": 2} and d["kitaplar"][0]["fiyatDayanak"] == 100.0
    # fiyat değişir → kritik uyarı; kabul edince yeni dayanak olur, uyarı kapanır
    pool["books"][0]["fiyatlar"]["crm"] = 110.0
    d = C.detail(engine, TN, cat["id"], pool, cfg)
    assert d["kitaplar"][0]["kritik"] == 1 and d["ozet"]["uyariliKitap"] == 1
    C.accept(engine, TN, cat["id"], _book(1)["id"], "fiyat", pool, cfg)
    d = C.detail(engine, TN, cat["id"], pool, cfg)
    assert d["kitaplar"][0]["kritik"] == 0 and d["kitaplar"][0]["fiyatDayanak"] == 110.0
    # sıra değişir, dayanak korunur
    d, _ = C.set_items(engine, TN, "ayse", cat["id"], [{"crmKitapId": _book(2)["id"]}, {"crmKitapId": _book(1)["id"], "oneCikan": True}],
                       pool, cfg)
    assert [k["sira"] for k in d["kitaplar"]] == [1, 2] and d["kitaplar"][1]["fiyatDayanak"] == 110.0
    C.transition(engine, TN, "ayse", cat["id"], "submit")
    with pytest.raises(C.CatalogError):
        C.set_items(engine, TN, "ayse", cat["id"], [], pool, cfg)      # onaydayken liste değişmez
    with pytest.raises(C.CatalogError) as e:
        C.transition(engine, TN, "ayse", cat["id"], "approve")
    assert e.value.status == 403                                       # gönderen onaylayamaz
    with pytest.raises(C.CatalogError):
        C.transition(engine, TN, "mehmet", cat["id"], "reject", "")    # gerekçesiz geri gönderme yok
    out = C.transition(engine, TN, "mehmet", cat["id"], "approve")
    assert out["durum"] == "onayli" and out["onaylayan"] == "mehmet"
    assert C.transition(engine, TN, "ayse", cat["id"], "publish")["durum"] == "yayinda"
    with pytest.raises(C.CatalogError):
        C.delete(engine, TN, cat["id"])                                # yalnız taslak silinir
    summary = C.refresh_alerts(engine, TN, pool, cfg)
    assert summary[0]["kritik"] == 0
    lst = C.list_catalogs(engine, TN)
    assert lst["items"][0]["kitap"] == 2 and lst["items"][0]["oneCikan"] == 1


def test_exports_have_rows_links_and_brief(engine):
    cfg = C.settings()
    pool = _pool([_book(1), _book(2, kapak=None, kapakKaynak=None)])
    cat = C.create(engine, TN, "ayse", {"tur": "yabanci-hak", "baslik": "Rights 2026"}, cfg)
    C.set_items(engine, TN, "ayse", cat["id"], [{"crmKitapId": _book(1)["id"], "metin": "Kısa metin."},
                                               {"crmKitapId": _book(2)["id"]}], pool, cfg)
    d = C.detail(engine, TN, cat["id"], pool, cfg)
    texts = {_book(2)["id"]: {"kisa": "CRM kısa bilgisi.", "ozet": None}}
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(C.catalog_xlsx(d, texts)))
    ws = wb["Katalog"]
    assert ws.cell(row=5, column=6).value == "Kitap 1" and ws.cell(row=6, column=15).value == "CRM kısa bilgisi."
    z = zipfile.ZipFile(io.BytesIO(C.package_zip(d, texts)))
    assert set(z.namelist()) == {"katalog.xlsx", "kapaklar.csv", "metinler.txt", "BENIOKU.txt"}
    assert "EKSİK" in z.read("kapaklar.csv").decode("utf-8-sig")
    pytest.importorskip("fpdf")
    pdf = C.preview_pdf(d, texts, 4)
    assert pdf[:4] == b"%PDF"


def test_zeki_text_drops_numbers_not_in_source():
    src = {"kisa": "Bir kasabanın hikâyesi. Umut ve dostluk üzerine.", "ozet": None}
    out = C.zeki_text(lambda m: "Bir kasabanın hikâyesi. 100 bin okura ulaştı. En çok satan roman.", {"ad": "X"}, src, 60, False)
    assert out["metin"] == "Bir kasabanın hikâyesi."
    assert {x["neden"] for x in out["dusen"]} == {"kaynaksiz-rakam", "kanitsiz-iddia"}


# ------------------------------------------------------------------ bülten


def test_newsletter_flow_html_and_results(engine):
    pool = _pool([_book(1, webUrl="https://timas.test/kitap-1"), _book(2, ad="<script>x</script>")])
    nl = N.create(engine, TN, "ayse", {"baslik": "Öğretmenler Günü bülteni", "segment": {"ilgiBayraklari": ["new_tarihveakademi"]}})
    d = N.set_items(engine, TN, nl["id"], [{"crmKitapId": _book(1)["id"], "metin": "Sıcak bir öykü."},
                                          {"crmKitapId": _book(2)["id"]}], pool, "crm")
    html = d["html"]
    assert "<script>x</script>" not in html and "&lt;script&gt;" in html and "https://timas.test/kitap-1" in html
    assert "@" not in html
    with pytest.raises(C.CatalogError) as e:
        N.transition(engine, TN, "ayse", nl["id"], "submit")          # konu ve segment sayısı yok
    assert e.value.status == 409
    seg = N.normalize_segment({"ilgiBayraklari": ["new_tarihveakademi"]})
    N.set_segment_count(engine, TN, nl["id"], seg, N.segment_result({"aday": 10, "izinli": 4}, False))
    N.update(engine, TN, nl["id"], {"konu": "Öğretmenlerimize"})
    N.transition(engine, TN, "ayse", nl["id"], "submit")
    with pytest.raises(C.CatalogError):
        N.transition(engine, TN, "ayse", nl["id"], "approve")
    N.transition(engine, TN, "mehmet", nl["id"], "approve")
    with pytest.raises(C.CatalogError):
        N.update(engine, TN, nl["id"], {"giris": "değişmez"})         # onaylı bülten metni değişmez
    N.update(engine, TN, nl["id"], {"crmKampanya": "22222222-2222-2222-2222-222222222222", "gonderimTarihi": "2026-11-20"})
    N.transition(engine, TN, "ayse", nl["id"], "mark-sent")
    r = N.add_result(engine, TN, "ayse", nl["id"], "elle", {"sent": 1000, "opened": 250, "clicked": 40}, None)
    assert r["acilmaOrani"] == 0.25 and r["tiklamaOrani"] == 0.04
    with pytest.raises(C.CatalogError):
        N.add_result(engine, TN, "ayse", nl["id"], "elle", {"sent": 10, "opened": 20}, None)
    assert N.record_crm_result(engine, TN, nl["id"], {"toplam": 900, "okunan": 300, "tiklanan": 30}) is True
    assert N.record_crm_result(engine, TN, nl["id"], {"toplam": 900, "okunan": 300, "tiklanan": 30}) is False
    rep = N.report(engine, TN)
    assert rep["items"][0]["sonuc"]["gonderilen"] == 900 and "@" not in json.dumps(rep)


def test_result_file_counts_rows_without_keeping_addresses():
    summary = N.parse_results("Gönderilen;Açılan;Tıklanan;Abonelikten çıkan\n1.200;300;45;3\n")
    assert (summary["sent"], summary["opened"], summary["clicked"], summary["unsubscribed"]) == (1200, 300, 45, 3)
    per = N.parse_results("E-posta,Açıldı,Tıklandı\na@x.test,evet,hayır\nb@x.test,1,1\nc@x.test,,\n")
    assert (per["sent"], per["opened"], per["clicked"], per["bicim"]) == (3, 2, 1, "kisi")
    assert "@" not in json.dumps(per)
    with pytest.raises(C.CatalogError):
        N.parse_results("Kişi\nx\n")


def test_newsletter_draft_keeps_only_checked_text(engine):
    pool = _pool([_book(1)])
    nl = N.create(engine, TN, "ayse", {"baslik": "Eylül bülteni"})
    N.set_items(engine, TN, nl["id"], [{"crmKitapId": _book(1)["id"]}], pool, "crm")
    job = C.job_create(engine, TN, "ayse", "bulten-taslak", nl["id"])
    reply = json.dumps({"giris": "Eylül bülteninde yeni bir öykü var. Bu kitap 50 bin sattı.",
                        "kitaplar": [{"no": 1, "metin": "Umut ve dostluk üzerine."}],
                        "konular": ["Eylülün öyküsü", "Rekor kıran kitap", "Umut ve dostluk"]}, ensure_ascii=False)
    texts = {_book(1)["id"]: {"kisa": "Umut ve dostluk üzerine. Yeni bir öykü.", "ozet": None}}
    N.run_draft(engine, TN, job["id"], nl["id"], lambda m: reply, lambda ids: texts, pool, "crm", {}, N.settings())
    assert C.job_get(engine, TN, job["id"])["durum"] == "bitti"
    d = N.detail(engine, TN, nl["id"], pool, "crm")
    assert "50" not in (d["giris"] or "") and "Rekor kıran kitap" not in d["konular"]
    assert d["kitaplar"][0]["metin"] == "Umut ve dostluk üzerine." and d["konu"] in d["konular"]


def test_interest_scorer_matches_book_genre():
    seg = N.normalize_segment({"ilgiBayraklari": ["new_tarihveakademi"]})
    score = N.interest_scorer(seg, {}, N.settings()["keywords"])
    assert score(_book(1, tur="Tarih, Biyografi"))[0] == 1.0
    assert score(_book(2, tur="Roman"))[0] == 0.0
    assert N.interest_scorer(N.normalize_segment({}), {}, N.settings()["keywords"]) is None


# ------------------------------------------------------------------ yetki ve köprü


def test_catalog_rules():
    P = "/api/v1/catalog-newsletter"
    assert A.rule_for(P + "/catalogs") == frozenset({"sayfa:katalog-bulten"})
    assert A.rule_for(P + "/run-due") == A.SYSTEM
    f = A.features_for
    assert f("POST", P + "/catalogs") == ["ozellik:katalog.duzenle"]
    assert f("PUT", P + "/catalogs/abc/items") == ["ozellik:katalog.duzenle"]
    assert f("POST", P + "/catalogs/abc/items/x/accept") == ["ozellik:katalog.duzenle"]
    assert f("POST", P + "/catalogs/abc/approve") == [] and f("POST", P + "/newsletters/abc/reject") == []
    assert f("POST", P + "/newsletters/abc/draft") == ["ozellik:bulten.duzenle"]
    assert f("DELETE", P + "/newsletters/abc/results/r1") == ["ozellik:bulten.duzenle"]
    assert f("POST", P + "/segments/count") == []                     # açıkça verilen bulten.segment ucun içinde
    assert f("GET", P + "/catalogs/abc/package.zip") == ["ozellik:veri.disa-aktar"]
    assert f("GET", P + "/newsletters/abc/html") == ["ozellik:veri.disa-aktar"]
    assert f("GET", P + "/catalogs/abc") == []
    assert {"ozellik:katalog-bulten.onay", "ozellik:bulten.segment"} <= A.explicit_keys()
    assert {"sayfa:katalog-bulten", "ozellik:katalog.duzenle", "ozellik:bulten.duzenle"} <= A.all_keys()


def test_api_explicit_permissions_and_no_person_data(monkeypatch, store, settings):
    from fastapi.testclient import TestClient

    from semantic_bridge import admin as admin_mod
    from semantic_bridge import board as board_mod
    from semantic_bridge.app import Runtime, create_app
    from semantic_layer.candidates.llm_client import FakeLlm
    from semantic_layer.tests.conftest import TENANT

    users = {"timas_session=a": "ayse", "timas_session=z": "zekiai"}
    monkeypatch.setattr(board_mod, "_fetch_user", lambda cookie: users.get(cookie))
    monkeypatch.setattr(board_mod, "_fetch_session", lambda cookie: {"username": users.get(cookie), "displayName": "x"})
    monkeypatch.setattr(admin_mod, "admins", lambda: ["zekiai"])
    monkeypatch.delenv("SEMANTIC_CALLER_TOKEN", raising=False)
    A._ready.clear()
    A.invalidate()
    C._ready.discard(id(store.engine))
    N._ready.discard(id(store.engine))
    C.ensure(store.engine)
    C.meta_set(store.engine, TENANT, "pool", _pool([_book(1), _book(2)]))
    app = create_app(Runtime(settings, store=store, llm=FakeLlm([""])))
    client = TestClient(app)
    a, z = {"cookie": "timas_session=a"}, {"cookie": "timas_session=z"}
    P = "/api/v1/catalog-newsletter"
    meta = client.get(P + "/meta", headers=a).json()
    assert meta["havuz"]["kitap"] == 2 and meta["me"]["canApprove"] is False and meta["me"]["canSegment"] is False
    assert client.post(P + "/catalogs", json={}, headers=a).status_code == 400
    cat = client.post(P + "/catalogs", json={"tur": "okul", "baslik": "Okul 2026"}, headers=a).json()
    sug = client.post(P + f"/catalogs/{cat['id']}/suggest", json={}, headers=a).json()
    assert sug["total"] == 2
    d = client.put(P + f"/catalogs/{cat['id']}/items", json={"items": [{"crmKitapId": _book(1)["id"]}]}, headers=a).json()
    assert d["ozet"]["kitap"] == 1
    assert client.post(P + f"/catalogs/{cat['id']}/submit", json={}, headers=a).status_code == 200
    assert client.post(P + f"/catalogs/{cat['id']}/approve", json={}, headers=a).status_code == 403   # açık yetki yok
    assert client.post(P + f"/catalogs/{cat['id']}/approve", json={}, headers=z).json()["durum"] == "onayli"
    assert client.post(P + "/segments/count", json={"segment": {}}, headers=a).status_code == 403
    nl = client.post(P + "/newsletters", json={"baslik": "Deneme bülteni"}, headers=a).json()
    body = client.get(P + f"/newsletters/{nl['id']}", headers=a).text
    assert "@" not in body
    assert client.get(P + f"/newsletters/{nl['id']}/html", headers=a).status_code == 409   # onaysız HTML indirilmez
    # test verisi bırakılmaz
    assert client.post(P + f"/catalogs/{cat['id']}/reopen", json={}, headers=a).status_code == 200
    assert client.delete(P + f"/catalogs/{cat['id']}", headers=a).status_code == 200
    assert client.delete(P + f"/newsletters/{nl['id']}", headers=a).status_code == 200
