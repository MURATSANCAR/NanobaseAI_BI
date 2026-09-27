"""M29 İlk dağılım: yuvarlama, benzer kitap payı, bölge, öneri, iki göz onayı, stok değişmezi, hücre düzeltmesi,
revizyon, sevk listesi, takip, uyarılar (verinin bittiği güne göre), BMT kapsamı, yetki kuralları.

Veriler yapaydır ve yalnız hesap kurallarını sınar; gerçek Logo/CRM kabulü test sunucusunda
(`scripts/acceptance/M29/`, günlük 2026-09-28). Yerelde koşulmaz (AGENTS.md).
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone

import pytest
import sqlalchemy as sa

from semantic_bridge import access as A
from semantic_bridge import budget as B
from semantic_bridge import distribution as D
from semantic_layer.store.catalog_store import open_store

T = "t1"


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    D._ready.discard(id(e))
    B._ready.discard(id(e))
    D.ensure(e)
    B.ensure(e)
    return e


class FakeLogo:
    """Benzer kitap penceresi, cari kartı, takip. İki benzer kitap (C1 çok satan, C2 az satan)."""

    def __init__(self, end=date(2026, 8, 17)):
        self.end = end
        self.track_rows: list[dict] = []

    def first_year(self):
        return 2021

    def data_end(self):
        return self.end

    def first_sales(self, codes, since_year):
        return {c: date(2025, 3, 1) for c in codes if c in ("C1", "C2")}

    def comp_windows(self, windows):
        rows = []
        for k, _, _ in windows:
            if k == "C1":   # 1.000 adet: İstanbul zinciri 600, Ankara kitapçısı 400 (100 iade)
                rows += [dict(comp=k, cari_kodu="120.01", satis=600, iade=0, ciro=60_000),
                         dict(comp=k, cari_kodu="120.02", satis=500, iade=100, ciro=40_000)]
            if k == "C2":   # 100 adet: yalnız İzmir kitapçısı
                rows += [dict(comp=k, cari_kodu="120.03", satis=100, iade=0, ciro=9_000)]
            if k == "K1":   # baskı tekrarı: kendi son penceresi
                rows += [dict(comp=k, cari_kodu="120.01", satis=300, iade=0, ciro=30_000)]
        return rows

    def clients(self, codes):
        base = {"120.01": {"unvan": "Zincir A.Ş.", "il": "İSTANBUL", "kanal": "ZINCIR"},
                "120.02": {"unvan": "Başkent Kitap", "il": "Ankara/Çankaya", "kanal": "KITAPCI"},
                "120.03": {"unvan": "Ege Kitabevi", "il": "IZMIR", "kanal": "KITAPCI"},
                "120.09": {"unvan": "Dağılım Carisi", "il": "Trabzon", "kanal": "DAGITICI"}}
        return {c: base[c] for c in codes if c in base}

    def tracking(self, code, start, end):
        return list(self.track_rows)


class FakeSources(D.Sources):
    def __init__(self, logo=None, verdict_llm=None):
        super().__init__(lambda: None, lambda: None, lambda: "Timas_MSCRM.dbo", llm=lambda: verdict_llm)
        self._fake = logo or FakeLogo()

    def logo(self):
        return self._fake

    def accounts(self):
        return [
            {"id": "a1", "ad": "Zincir A.Ş.", "cari_kodu": "120.01", "dagilim": True, "bmt_ad": "Ali Bmt", "bmt_hesap": "ali", "il": "İstanbul"},
            {"id": "a2", "ad": "Başkent Kitap", "cari_kodu": "120.02", "dagilim": False, "bmt_ad": "Veli Bmt", "bmt_hesap": "veli", "il": "Ankara"},
            {"id": "a9", "ad": "Dağılım Carisi", "cari_kodu": "120.09", "dagilim": True, "bmt_ad": "Veli Bmt", "bmt_hesap": "veli", "il": "Trabzon"},
            {"id": "a0", "ad": "Kodsuz Cari", "cari_kodu": None, "dagilim": True, "bmt_ad": None, "bmt_hesap": None, "il": "Van"},
        ]

    def dist_orders(self, codes):
        return {"C1": {"adet": 900, "siparis": 12, "cari": 12, "ilk": "2025-02-20"}}


def _seed(engine, stok=10_000.0, first_print=True, code="N1"):
    with engine.begin() as c:
        c.execute(D.BOOKS.insert().values(tenant_id=T, stok_kodu=code, ad="Yeni Roman", yayinevi="Timaş Yayınları",
                                          depo_giris_tarihi="2026-08-10", baski_adedi=12_000, baski_no=1 if first_print else 5,
                                          ilk_baski=first_print, kaynak="logo", stok_bakiye=stok, asof=D._now()))
        c.execute(B.BOOKINFO.insert(), [
            dict(stok_kodu="N1", ad="Yeni Roman", yazar="Yazar Bir", yayinevi="Timaş Yayınları", kitaplik="Roman", ilk_yayin="2026-08-01", in_logo=True),
            dict(stok_kodu="C1", ad="Eski Roman", yazar="Yazar Bir", yayinevi="Timaş Yayınları", kitaplik="Roman", ilk_yayin="2025-03-01", in_logo=True),
            dict(stok_kodu="C2", ad="Başka Roman", yazar="Yazar İki", yayinevi="Timaş Yayınları", kitaplik="Roman", ilk_yayin="2025-03-01", in_logo=True),
            dict(stok_kodu="X9", ad="İlgisiz", yazar="Yazar Üç", yayinevi="Başka", kitaplik="Masal", ilk_yayin="2025-03-01", in_logo=True),
        ])
    D.meta_set(engine, T, "logo", {"veriSonu": "2026-08-17"})


def _approve_target(engine, adet=24_000):
    """M46: 2026 yürürlükteki planında N1 için aylık eşit hedef (Ağustos + Eylül = 2 × adet/12)."""
    now = datetime.now(timezone.utc)
    with engine.begin() as c:
        c.execute(B.PLANS.insert().values(id="p46", tenant_id=T, year=2026, scenario="temel", version=1, status="onayli",
                                          title="2026 · Temel", params_json="{}", basis_json="{}", created_by="x", created_at=now,
                                          decided_by="y", decided_at=now))
        c.execute(B.BOOKS.insert().values(plan_id="p46", stok_kodu="N1", ad="Yeni Roman", segment="backlist", adet=adet,
                                          ciro=adet * 100, marj=0.5, oneri_json="{}"))


# ------------------------------------------------------------------ saf hesaplar


def test_allocate_keeps_the_total_exact_and_follows_the_shares():
    out = D.allocate(1000, {"a": 0.5, "b": 0.3333, "c": 0.1667, "d": 0})
    assert sum(out.values()) == 1000 and out["d"] == 0
    assert out["a"] == 500 and out["b"] == 333 and out["c"] == 167
    assert D.allocate(0, {"a": 1}) == {"a": 0}
    assert D.allocate(10, {}) == {}


def test_allocate_minimum_moves_small_lines_to_the_others():
    out = D.allocate(100, {"a": 0.9, "b": 0.08, "c": 0.02}, min_qty=5)
    assert out["c"] == 0 and sum(out.values()) == 100 and out["b"] >= 5


def test_comp_shares_weight_each_comp_by_its_own_total():
    rows = FakeLogo().comp_windows([("C1", None, None), ("C2", None, None)])
    share, stats, totals = D.comp_shares(rows, {"C1": 1.0, "C2": 1.0})
    assert totals == {"C1": 1000, "C2": 100}
    # C1 içindeki paylar 0,6 / 0,4; C2 tek cari → eşit ağırlıkla 0,3 / 0,2 / 0,5 (çok satan kitap öneriyi ele geçirmez)
    assert share == pytest.approx({"120.01": 0.3, "120.02": 0.2, "120.03": 0.5})
    assert stats["120.02"]["iade"] == 100 and abs(sum(share.values()) - 1) < 1e-9


def test_region_of_uses_the_province_and_keeps_unknowns_visible():
    assert D.region_of("İSTANBUL") == "İstanbul"
    assert D.region_of("Ankara/Çankaya") == "İç Anadolu"
    assert D.region_of("IZMIR") == "Ege"
    assert D.region_of("icel") == "Akdeniz"
    assert D.region_of("") == "İli belirsiz"
    assert D.region_of("Berlin") == "İli belirsiz"
    assert D.region_of("Berlin", "YURTDIŞI") == "Yurt dışı"
    assert sum(len(v) for v in D.REGIONS.values()) == 81


def test_model_text_cannot_invent_numbers():
    facts = "Toplam 1.250 adet; İstanbul %38; ilk 56 gün."
    assert D.safe_model_text("İstanbul'a 1.250 adetin %38'i gider, 56 günlük satışa bakıldı.", facts)
    assert not D.safe_model_text("İstanbul'a 1.300 adet gider.", facts)
    assert not D.safe_model_text("", facts)


def test_business_days_skip_the_weekend():
    assert D.business_days_between(date(2026, 9, 25), date(2026, 9, 28)) == 1   # cuma → pazartesi
    assert D.business_days_between(date(2026, 9, 21), date(2026, 9, 28)) == 5


# ------------------------------------------------------------------ öneri


def test_proposal_from_comps_caps_at_stock_minus_reserve_and_keeps_distribution_accounts(engine):
    _seed(engine, stok=1_000)
    plan = D.generate(engine, T, "mudur", "N1", FakeSources())
    b = plan["basis"]
    assert b["yontem"] == "kitap-karti" and b["toplamKaynak"] == "benzer"
    # Benzer ortanca (1.000 ile 100) = 550; stok 1.000 − rezerv %20 = 800 → 550.
    assert plan["rezerv"] == 200 and plan["onerilenToplam"] == 550 and plan["toplam"] == 550
    lines = {x["cariKodu"]: x for x in D.list_lines(engine, T, plan["id"])["items"]}
    assert sum(x["adet"] for x in lines.values()) == 550
    assert lines["120.01"]["bolge"] == "İstanbul" and lines["120.01"]["kanal"] == "Kitabevi zinciri"
    assert lines["120.02"]["bolge"] == "İç Anadolu" and lines["120.02"]["gecmisIadeOrani"] == pytest.approx(100 / 500)
    # Dağılım carisi payı sıfır da olsa listede; Logo kodu olmayan dağılım carisi de (kesilmez).
    assert lines["120.09"]["adet"] == 0 and lines["120.09"]["dagilimCarisi"]
    assert any(x["cariKodu"] is None and x["unvan"] == "Kodsuz Cari" for x in lines.values())
    comps = {c["stokKodu"]: c for c in plan["benzerler"]}
    assert comps["C1"]["secildi"] and comps["C1"]["net"] == 1000 and comps["C1"]["crmDagilimAdet"] == 900
    assert "X9" not in comps                                   # ortak özelliği yok
    assert plan["gerekceKaynak"] == "kural" and "benzer kitabın" in plan["gerekce"]


def test_proposal_uses_the_approved_target_first_two_months(engine):
    _seed(engine, stok=10_000)
    _approve_target(engine, adet=24_000)
    plan = D.generate(engine, T, "mudur", "N1", FakeSources())
    assert plan["basis"]["toplamKaynak"] == "hedef"
    assert plan["hedef"]["ikiAyAdet"] == 4_000 and plan["hedef"]["aylar"] == [8, 9]
    assert plan["toplam"] == 4_000 and plan["hedef"]["planId"] == "p46"


def test_reprint_uses_the_books_own_recent_customers(engine):
    _seed(engine, code="K1", first_print=False)
    plan = D.generate(engine, T, "mudur", "K1", FakeSources())
    assert plan["basis"]["yontem"] == "kendi" and [c["stokKodu"] for c in plan["benzerler"]] == ["K1"]
    assert {x["cariKodu"]: x["adet"] for x in D.list_lines(engine, T, plan["id"], yalniz="adetli")["items"]} == {"120.01": 300}


class ScreenLlm:
    def __init__(self, reply):
        self.reply = reply

    def chat(self, messages, **_):
        if "Adaylar" in messages[-1]["content"]:
            return self.reply
        return "Öneri benzer kitaplardan kuruldu; toplam 9.999 adet."   # olgu dışı sayı → kural metni


def test_model_screen_drops_unlike_comps_and_cannot_write_numbers(engine):
    _seed(engine, stok=1_000)
    plan = D.generate(engine, T, "mudur", "N1", FakeSources(verdict_llm=ScreenLlm("C1: benzer\nC2: değil")))
    comps = {c["stokKodu"]: c for c in plan["benzerler"]}
    # Tek aday kalırsa ayıklama yetersiz sayılır, puan sırası kullanılır.
    assert plan["basis"]["uyarilar"] and comps["C2"]["modelKarar"] is None
    assert plan["gerekceKaynak"] == "kural"


def test_open_plan_blocks_a_second_proposal(engine):
    _seed(engine)
    D.generate(engine, T, "mudur", "N1", FakeSources())
    with pytest.raises(D.DistError) as e:
        D.generate(engine, T, "mudur", "N1", FakeSources())
    assert e.value.status == 409


# ------------------------------------------------------------------ düzeltme ve onay


def test_line_and_cell_edits_are_marked_and_totals_follow(engine):
    _seed(engine, stok=10_000)
    plan = D.generate(engine, T, "mudur", "N1", FakeSources())
    lines = {x["cariKodu"]: x for x in D.list_lines(engine, T, plan["id"])["items"]}
    out, diff = D.update_line(engine, T, "mudur", plan["id"], lines["120.09"]["no"], {"adet": "1.250", "gerekce": "Yazar imza günü"})
    assert out["adet"] == 1250 and out["elle"] and diff["adet"]["yeni"] == 1250
    with pytest.raises(D.DistError):
        D.update_line(engine, T, "mudur", plan["id"], lines["120.09"]["no"], {"adet": -1})
    with pytest.raises(D.DistError):
        D.update_line(engine, T, "mudur", plan["id"], lines["120.09"]["no"], {"adet": 2.5})
    detail, _ = D.update_cell(engine, T, "mudur", plan["id"], {"bolge": "İstanbul", "kanal": "Kitabevi zinciri", "adet": 99})
    cell = next(h for h in detail["matris"]["hucreler"] if h["bolge"] == "İstanbul")
    assert cell["adet"] == 99 and detail["toplam"] == sum(x["adet"] for x in D.list_lines(engine, T, plan["id"])["items"])
    assert detail["elleSatir"] == 2


def test_two_eyes_and_the_stock_invariant(engine):
    _seed(engine, stok=1_000)
    plan = D.generate(engine, T, "mudur", "N1", FakeSources())
    no = D.list_lines(engine, T, plan["id"])["items"][0]["no"]
    D.update_line(engine, T, "mudur", plan["id"], no, {"adet": 900})      # 900 + diğerleri + rezerv 200 > 1.000
    with pytest.raises(D.DistError) as e:
        D.submit(engine, T, "mudur", plan["id"])
    assert e.value.status == 409 and "aşıyor" in str(e.value)
    D.update_plan(engine, T, "mudur", plan["id"], {"rezerv": 0})
    D.update_line(engine, T, "mudur", plan["id"], no, {"adet": 100})
    D.submit(engine, T, "mudur", plan["id"])
    with pytest.raises(D.DistError) as e:
        D.update_line(engine, T, "mudur", plan["id"], no, {"adet": 1})    # onaydayken değişmez
    assert e.value.status == 409
    with pytest.raises(D.DistError) as e:
        D.decide(engine, T, "MUDUR", plan["id"], True)                   # gönderen onaylayamaz
    assert e.value.status == 409
    with pytest.raises(D.DistError):
        D.decide(engine, T, "lojistik", plan["id"], False, "")           # gerekçesiz geri gönderme yok
    out = D.decide(engine, T, "lojistik", plan["id"], True, "tamam")
    assert out["durum"] == "onayli" and out["decidedBy"] == "lojistik"
    rev = D.revise(engine, T, "mudur", plan["id"], "Konya imza günü")
    assert rev["durum"] == "taslak" and rev["surum"] == 2 and rev["revisionOf"] == plan["id"]
    D.submit(engine, T, "mudur", rev["id"])
    D.decide(engine, T, "lojistik", rev["id"], True)
    assert D.plan_detail(engine, T, plan["id"])["durum"] == "arsiv"
    with pytest.raises(D.DistError):
        D.delete_plan(engine, T, rev["id"])                              # yalnız taslak silinir


def test_export_only_from_an_approved_plan(engine):
    from openpyxl import load_workbook

    _seed(engine, stok=10_000)
    plan = D.generate(engine, T, "mudur", "N1", FakeSources())
    with pytest.raises(D.DistError):
        D.export_xlsx(engine, T, plan["id"])
    D.submit(engine, T, "mudur", plan["id"])
    D.decide(engine, T, "lojistik", plan["id"], True)
    data, name = D.export_xlsx(engine, T, plan["id"])
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Sevk listesi"]
    head = [c.value for c in ws[5]]
    assert head[:4] == ["Sıra", "Sipariş tipi", "Cari kodu", "Cari unvanı"] and name.endswith(".xlsx")
    qty = [r[10] for r in ws.iter_rows(min_row=6, values_only=True) if isinstance(r[0], int)]
    assert sum(qty) == plan["toplam"] and all(q > 0 for q in qty)
    assert "Özet" in wb.sheetnames


# ------------------------------------------------------------------ takip, uyarı, kapsam


def _approved(engine, logo=None):
    _seed(engine, stok=10_000)
    S = FakeSources(logo)
    plan = D.generate(engine, T, "mudur", "N1", S)
    D.submit(engine, T, "mudur", plan["id"])
    D.decide(engine, T, "lojistik", plan["id"], True)
    with engine.begin() as c:   # onay: 3 Ağustos 2026 (Logo verisi 17 Ağustos'ta bitiyor)
        c.execute(D.PLANS.update().where(D.PLANS.c.id == plan["id"]).values(decided_at=datetime(2026, 8, 3, 9, tzinfo=timezone.utc)))
    return plan, S


def test_tracking_and_alerts_are_judged_at_the_data_end(engine, monkeypatch):
    logo = FakeLogo()
    logo.track_rows = [dict(cari_kodu="120.01", hafta=1, sevk=200, fatura=200, iade=0),
                       dict(cari_kodu="120.03", hafta=2, sevk=50, fatura=0, iade=0)]
    plan, S = _approved(engine, logo)
    row = D.get_plan_row(engine, T, plan["id"])
    assert D.refresh_tracking(engine, T, S, row)["satir"] == 2
    s = D.tracking(engine, T, "N1")["items"][0]
    assert s["sevk"] == 250 and s["fatura"] == 200 and s["hafta"] == 3      # 17 Ağustos → 3. hafta
    monkeypatch.setattr(D, "today", lambda: date(2026, 9, 28))
    ev = D.evaluate_alerts(engine, T)
    assert ev["degerlendirmeGunu"] == "2026-08-17"
    al = D.alerts(engine, T)["items"]
    kinds = {a["tur"] for a in al}
    # 120.02 sevk edilmedi (10 iş günü); 120.01 tükeniyor (bilgi); Ege'de sevk var satış yok ama 28 gün dolmadı.
    assert "sevk_gecikti" in kinds and "tukendi" in kinds and "hic_satmadi" not in kinds
    assert all(a["detay"]["cari"] != "120.01" for a in al if a["tur"] == "sevk_gecikti")
    # Veri 1 Eylül'e uzansa Ege «hiç satmadı» olur.
    D.meta_set(engine, T, "logo", {"veriSonu": "2026-09-01"})
    D.evaluate_alerts(engine, T)
    assert any(a["tur"] == "hic_satmadi" and a["detay"]["bolge"] == "Ege" for a in D.alerts(engine, T)["items"])
    # Sevk gelince gecikme kapanır.
    logo.track_rows.append(dict(cari_kodu="120.02", hafta=1, sevk=10, fatura=0, iade=0))
    D.refresh_tracking(engine, T, S, row)
    assert D.evaluate_alerts(engine, T)["kapanan"] >= 1


def test_plan_missing_alert_and_notification_waits_without_recipients(engine, monkeypatch):
    _seed(engine)
    monkeypatch.setattr(D, "today", lambda: date(2026, 8, 12))
    D.evaluate_alerts(engine, T)
    assert [a["tur"] for a in D.alerts(engine, T)["items"]] == ["plan_yok"]
    assert D.notify(engine, T, [], "", lambda *a: "sent")["status"] == "no_recipient"
    sent = []
    assert D.notify(engine, T, ["satis@timas.com.tr"], "https://x/timas", lambda s, t, to: sent.append((s, to)) or "sent")["status"] == "sent"
    assert D.notify(engine, T, ["satis@timas.com.tr"], "", lambda *a: "sent")["status"] == "nothing"
    D.generate(engine, T, "mudur", "N1", FakeSources())
    D.evaluate_alerts(engine, T)
    assert D.alerts(engine, T)["items"] == []


def test_bmt_sees_only_own_accounts(engine, monkeypatch):
    monkeypatch.setattr(D, "today", lambda: date(2026, 8, 20))
    plan, _ = _approved(engine)
    mine = D.my_region(engine, T, "veli")["items"][0]
    assert {c["cariKodu"] for c in mine["cariler"]} <= {"120.02", "120.09"}
    assert all(c["cariKodu"] != "120.01" for c in mine["cariler"])
    assert D.my_region(engine, T, "kimse")["items"] == []
    detail = D.plan_detail(engine, T, plan["id"], bmt="ali")
    assert detail["kapsam"] == "kendi" and {h["bolge"] for h in detail["matris"]["hucreler"]} == {"İstanbul"}
    assert all(x["bmtHesap"] == "ali" for x in D.list_lines(engine, T, plan["id"], bmt="ALI")["items"])


def test_books_list_states(engine):
    _seed(engine)
    assert [b["durum"] for b in D.list_books(engine, T)["items"]] == ["yok"]
    plan = D.generate(engine, T, "mudur", "N1", FakeSources())
    out = D.list_books(engine, T, durum="bekleyen")
    assert out["items"][0]["durum"] == "taslak" and out["items"][0]["plan"]["id"] == plan["id"]
    assert D.list_books(engine, T, durum="onayli")["items"] == []


# ------------------------------------------------------------------ yetki


def test_access_rules_for_distribution():
    assert A.rule_for("/api/v1/distribution/books") == frozenset({A.page("ilk-dagilim")})
    assert A.rule_for("/api/v1/distribution/run-due") == A.SYSTEM
    assert A.page("ilk-dagilim") in A.rule_for("/api/v1/budget/targets")
    assert A.features_for("POST", "/api/v1/distribution/plans/generate") == ["ozellik:dagilim.plan"]
    assert A.features_for("PATCH", "/api/v1/distribution/plans/x/lines/3") == ["ozellik:dagilim.plan"]
    assert A.features_for("POST", "/api/v1/distribution/plans/x/approve") == []      # uç içinde, açıkça verilen onay
    assert A.features_for("GET", "/api/v1/distribution/plans/x/export.xlsx") == ["ozellik:veri.disa-aktar"]
    assert A.features_for("GET", "/api/v1/distribution/plans/x") == []
    assert {"ozellik:dagilim.onay", "ozellik:dagilim.herkesinki"} <= A.explicit_keys()
    assert "ozellik:dagilim.plan" not in A.explicit_keys()
