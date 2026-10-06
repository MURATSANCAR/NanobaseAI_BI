"""M27 Fuar, etkinlik ve ödül: kitap/adet önerisi (geçmiş fuar, katsayı, yeni çıkan, stok), öneri ve elle listenin
birlikte yaşaması, kart ve görev akışı (iki göz onayı, bütçe değişince karar yeniden, görev sahibinin işaretlemesi),
gider ve fiş, yazar programı çakışması, fuar sonucu (toplam, geçen yıl, sipariş, gider, veri sonu uyarısı), CRM tip
eşlemesi (öneri ≠ karar), takvim ve Kampüs ajandası (yalnız kendi kaydı), hatırlatmalar (bir kez), SQL kuruluşu
(kanal, faturalı satır, İstanbul günü, enjeksiyon), yetki kuralları ve uçlar.

Veriler yapaydır ve yalnız kuralları sınar; gerçek Logo/CRM kabulü test sunucusunda (scripts/acceptance/M27).
"""

from __future__ import annotations

import base64
import json
from datetime import date, timedelta
from types import SimpleNamespace

import pytest

from semantic_bridge import access as A
from semantic_bridge import events as E
from semantic_bridge import events_sources as S
from semantic_layer.store.catalog_store import open_store

T = "t1"
TY_FUAR = "11111111-0000-0000-0000-000000000001"
TY_ZIY = "11111111-0000-0000-0000-000000000002"
TY_YENI = "11111111-0000-0000-0000-000000000003"
AUTH1 = "22222222-0000-0000-0000-000000000001"
NOW = date(2026, 10, 1)


@pytest.fixture
def engine(tmp_path, monkeypatch):
    monkeypatch.setenv("EVENTS_DIR", str(tmp_path / "events"))
    for k in ("EVENTS_SUGGEST_FACTOR", "EVENTS_NEW_BOOK_FACTOR", "EVENTS_NEW_BOOK_MONTHS", "EVENTS_TASK_TEMPLATE",
              "EVENTS_REMIND_DAYS", "EVENTS_FAIR_CHANNEL"):
        monkeypatch.delenv(k, raising=False)
    e = open_store("sqlite://").engine
    E._ready.discard(id(e))
    E.ensure(e)
    return e


def _fair(engine, user="ayse", **over) -> str:
    body = {"name": "TÜYAP İstanbul Kitap Fuarı", "kind": "stant", "startsOn": "2026-11-07", "endsOn": "2026-11-15",
            "city": "İstanbul", "venue": "Beylikdüzü", "budgetPlanned": "250.000", **over}
    return E.create_fair(engine, T, user, body)


SALES = [
    {"cariKodu": "120.FUAR.01", "cariAdi": "TÜYAP Stant", "stokKodu": "K1", "ad": "Kitap Bir", "adet": 120.0, "ciro": 12000.0},
    {"cariKodu": "120.FUAR.02", "cariAdi": "TÜYAP Stant 2", "stokKodu": "k1", "ad": "Kitap Bir", "adet": -20.0, "ciro": -2000.0},
    {"cariKodu": "120.FUAR.01", "cariAdi": "TÜYAP Stant", "stokKodu": "K2", "ad": "Kitap İki", "adet": 40.0, "ciro": 6000.0},
    {"cariKodu": "120.FUAR.01", "cariAdi": "TÜYAP Stant", "stokKodu": "K3", "ad": "Kitap Üç", "adet": 0.0, "ciro": 0.0},
]
BOOKS = {
    "K1": {"id": "b1", "stokKodu": "K1", "ad": "Kitap Bir", "ilkYayin": "2019-01-01"},
    "K2": {"id": "b2", "stokKodu": "K2", "ad": "Kitap İki", "ilkYayin": "2020-01-01"},
    "N1": {"id": "n1", "stokKodu": "N1", "ad": "Yeni Kitap", "ilkYayin": "2026-06-01"},
    "N2": {"id": "n2", "stokKodu": "N2", "ad": "Stoksuz Yeni", "ilkYayin": "2026-07-01"},
    "N3": {"id": "n3", "stokKodu": "N3", "ad": "Fuardan Sonra Çıkacak", "ilkYayin": "2026-12-01"},
}


# ------------------------------------------------------------------ öneri


def test_suggestion_uses_last_fair_factor_new_books_and_flags_stock():
    st = {**E.settings(), "suggestFactor": 1.1, "newBookFactor": 0.5, "newBookMonths": 12}
    items = E.suggest_books(SALES, {"K1": 50.0, "K2": 500.0, "N1": 30.0, "N2": 0.0}, BOOKS, st, "2026-11-07", "geçen yılın aynı günleri")
    by = {i["stokKodu"]: i for i in items}
    assert by["K1"]["basisQty"] == 100.0 and by["K1"]["qty"] == 110                     # iade düşülür, büyük/küçük harf tek kitap
    assert "önerinin altında" in by["K1"]["reason"] and "önerinin altında" not in by["K2"]["reason"]
    assert by["K2"]["qty"] == 44 and "K3" not in by                                        # satmayan kitap önerilmez
    assert by["N1"]["qty"] == 35 and by["N1"]["kind"] == "yeni"                            # ortanca(100, 40)=70 × 0,5
    assert "N2" not in by and "N3" not in by                                               # stoksuz / fuardan sonra çıkan
    assert [i["stokKodu"] for i in items][:2] == ["K1", "K2"]


def test_suggestion_without_basis_and_without_stock():
    st = {**E.settings(), "newBookMonths": 12}
    items = E.suggest_books([], None, BOOKS, st, "2026-11-07", "x")
    assert {i["stokKodu"] for i in items} == {"N1", "N2"}                                  # stok bilinmiyor: şart aranmaz
    assert all(i["qty"] is None and "elle girilmeli" in i["reason"] for i in items)


def test_saved_suggestions_keep_manual_plan_and_drop_stale_rows(engine):
    fid = _fair(engine)
    st = {**E.settings(), "newBookMonths": 0}
    E.save_suggestions(engine, T, "ayse", fid, E.suggest_books(SALES, {}, BOOKS, st, "2026-11-07", "x"))
    E.put_books(engine, T, "ayse", fid, [{"stokKodu": "K1", "qtyPlanned": "90", "featured": True}, {"stokKodu": "K2"}])
    counts = E.save_suggestions(engine, T, "ayse", fid, E.suggest_books(SALES[2:], {}, BOOKS, st, "2026-11-07", "x"))
    rows = {b["stokKodu"]: b for b in E.fair_detail(engine, T, fid)["bookList"]}
    assert counts == {"added": 0, "updated": 1, "removed": 0}
    assert rows["K1"]["qtyPlanned"] == 90 and rows["K1"]["featured"] and rows["K1"]["qtySuggested"] is None
    E.save_suggestions(engine, T, "ayse", fid, [])
    assert set(b["stokKodu"] for b in E.fair_detail(engine, T, fid)["bookList"]) == {"K1"}


def test_put_books_adds_card_books_and_removes_manual_rows(engine):
    fid = _fair(engine)
    with pytest.raises(E.EventsError):
        E.put_books(engine, T, "ayse", fid, [{"stokKodu": "YOK", "qtyPlanned": 5}], BOOKS)
    out = E.put_books(engine, T, "ayse", fid, [{"stokKodu": "n1", "qtyPlanned": "12"}], BOOKS)
    assert out["diff"][0]["yeni"] is True
    b = E.fair_detail(engine, T, fid)["bookList"][0]
    assert b["stokKodu"] == "N1" and b["ad"] == "Yeni Kitap" and b["source"] == "elle" and b["qtyPlanned"] == 12
    E.put_books(engine, T, "ayse", fid, [], BOOKS)
    assert E.fair_detail(engine, T, fid)["bookList"] == []
    with pytest.raises(E.EventsError):
        E.put_books(engine, T, "ayse", fid, [{"stokKodu": "K1"}, {"stokKodu": "k1"}], BOOKS)


# ------------------------------------------------------------------ kart, onay, görev


def test_fair_creation_validates_and_builds_task_template(engine, monkeypatch):
    with pytest.raises(E.EventsError):
        _fair(engine, startsOn="2026-11-10", endsOn="2026-11-01")
    with pytest.raises(E.EventsError):
        _fair(engine, kind="uzay")
    monkeypatch.setenv("EVENTS_TASK_TEMPLATE", json.dumps([{"title": "Stant", "days": -30}, {"title": "Rapor", "days": 3}]))
    fid = _fair(engine)
    d = E.fair_detail(engine, T, fid, NOW)
    assert d["status"] == "aday" and d["budgetPlanned"] == 250000.0 and d["owner"] == "ayse"
    assert [(t["title"], t["dueOn"]) for t in d["tasks"]] == [("Stant", "2026-10-08"), ("Rapor", "2026-11-10")]
    assert d["phase"] == "hazirlik" and d["daysLeft"] == 37 and d["prep"] == 0


def test_two_eyes_approval_and_budget_change_reopens(engine):
    fid = _fair(engine, user="ayse")
    with pytest.raises(E.EventsError) as e:
        E.approve_fair(engine, T, "ayse", fid)
    assert e.value.status == 403
    E.approve_fair(engine, T, "mehmet", fid, "uygun")
    assert E.fair_detail(engine, T, fid)["status"] == "onayli"
    with pytest.raises(E.EventsError):
        E.approve_fair(engine, T, "mehmet", fid)
    out = E.update_fair(engine, T, "ayse", fid, {"venue": "Salon 9"})
    assert not out["reopened"] and E.fair_detail(engine, T, fid)["status"] == "onayli"
    out = E.update_fair(engine, T, "ayse", fid, {"budgetPlanned": 300000})
    d = E.fair_detail(engine, T, fid)
    assert out["reopened"] and d["status"] == "aday" and d["approvedBy"] is None
    with pytest.raises(E.EventsError):
        E.update_fair(engine, T, "ayse", fid, {"status": "onayli"})
    E.update_fair(engine, T, "ayse", fid, {"status": "iptal"})
    E.update_fair(engine, T, "ayse", fid, {"status": "aday"})
    E.approve_fair(engine, T, "mehmet", fid)
    with pytest.raises(E.EventsError):
        E.delete_fair(engine, T, fid)                                                      # onaylı kart silinmez


def test_task_owner_can_tick_but_not_edit(engine):
    fid = _fair(engine, ownerUser="ayse")
    tid = E.add_task(engine, T, "ayse", fid, {"title": "Sevkiyat", "dueOn": "2026-10-20", "owner": "TIMAS\\Veli"})["id"]
    with pytest.raises(E.EventsError) as e:
        E.patch_task(engine, T, "veli", fid, tid, {"title": "Başka"}, can_edit=False)
    assert e.value.status == 403
    with pytest.raises(E.EventsError):
        E.patch_task(engine, T, "ali", fid, tid, {"done": True}, can_edit=False)
    E.patch_task(engine, T, "veli", fid, tid, {"done": True}, can_edit=False)
    t = next(x for x in E.fair_detail(engine, T, fid)["tasks"] if x["id"] == tid)
    assert t["done"] and t["doneBy"] == "veli" and t["owner"] == "veli"
    E.patch_task(engine, T, "ayse", fid, tid, {"done": False}, can_edit=True)
    assert not next(x for x in E.fair_detail(engine, T, fid)["tasks"] if x["id"] == tid)["done"]


def test_cost_with_receipt_and_limits(engine):
    fid = _fair(engine)
    with pytest.raises(E.EventsError):
        E.add_cost(engine, T, "ayse", fid, {"kind": "stant", "amount": "0"})
    with pytest.raises(E.EventsError) as e:
        E.add_cost(engine, T, "ayse", fid, {"kind": "stant", "amount": 10, "receipt": {"type": "text/html", "dataBase64": "eA=="}})
    assert e.value.status == 415
    big = base64.b64encode(b"x" * (1024 * 1024 + 1)).decode()
    with pytest.raises(E.EventsError) as e:
        E.add_cost(engine, T, "ayse", fid, {"kind": "stant", "amount": 10, "receipt": {"type": "image/png", "dataBase64": big}}, max_mb=1)
    assert e.value.status == 413
    out = E.add_cost(engine, T, "ayse", fid, {"kind": "konaklama", "amount": "12.500,50",
                                              "receipt": {"type": "image/jpeg", "dataBase64": base64.b64encode(b"jpg").decode()}})
    E.add_cost(engine, T, "ayse", fid, {"kind": "konaklama", "amount": 500})
    d = E.fair_detail(engine, T, fid)
    assert d["costTotal"] == 13000.5 and d["costByKind"] == {"konaklama": 13000.5}
    data, ctype, _ = E.receipt(engine, T, fid, out["id"])
    assert data == b"jpg" and ctype == "image/jpeg"
    E.delete_cost(engine, T, fid, out["id"])
    assert not (E.files_root() / fid / (out["id"] + ".jpg")).exists()


def test_author_slot_conflicts_across_fairs(engine):
    f1, f2 = _fair(engine), _fair(engine, name="Ankara Kitap Fuarı")
    E.add_author(engine, T, "ayse", f1, {"name": "Yazar Bir", "contactId": AUTH1, "slotStart": "2026-11-08T14:00", "slotEnd": "2026-11-08T16:00"})
    out = E.add_author(engine, T, "ayse", f2, {"name": "Yazar Bir", "contactId": AUTH1, "slotStart": "2026-11-08T15:00", "slotEnd": "2026-11-08T17:00"})
    assert len(out["conflicts"]) == 1 and out["conflicts"][0]["fairId"] == f1
    free = E.add_author(engine, T, "ayse", f2, {"name": "Yazar Bir", "contactId": AUTH1, "slotStart": "2026-11-09T15:00"})
    assert free["conflicts"] == []
    with pytest.raises(E.EventsError):
        E.add_author(engine, T, "ayse", f2, {"name": "X", "slotStart": "2026-11-09T15:00", "slotEnd": "2026-11-09T14:00"})
    E.update_fair(engine, T, "ayse", f1, {"status": "iptal"})
    assert E.fair_detail(engine, T, f2)["authors"][0]["conflicts"] == []                   # iptal kart çakışma sayılmaz


# ------------------------------------------------------------------ sonuç


def _result(engine, fid, **over):
    fair = E.fair_row(engine, T, fid)
    args = dict(sales={"rows": SALES, "warnings": [], "sql": ["SELECT 1"], "channel": "FUAR"},
                prev_sales={"rows": [SALES[0]], "warnings": [], "sql": []},
                basis=E.basis_window(fair, None),
                orders=[{"tip": 4, "tipAdi": "Fuar", "tutar": 1000.0}, {"tip": 16, "tipAdi": "İmza siparişi", "tutar": 250.0},
                        {"tip": 4, "tipAdi": "Fuar", "tutar": 500.0}],
                crm_events=[{"id": "e1", "ad": "TÜYAP", "baslangic": "2026-11-07", "katilimci": 300, "satilan": 90, "gider": 1000.0,
                             "iptal": False, "durumAdi": "Tamamlandı"},
                            {"id": "e2", "ad": "İptal", "baslangic": "2026-11-08", "katilimci": 50, "gider": 999.0, "iptal": True}],
                costs=E.fair_costs(engine, fid), planned=E.planned_books(engine, fid), data_end="2026-11-12", tail_days=0)
    args.update(over)
    return E.compute_result(fair, **args)


def test_result_totals_comparison_orders_costs_and_warnings(engine):
    fid = _fair(engine, budgetPlanned=5000)
    E.put_books(engine, T, "ayse", fid, [{"stokKodu": "K1", "qtyPlanned": 150}, {"stokKodu": "N1", "qtyPlanned": 10}], BOOKS)
    E.add_cost(engine, T, "ayse", fid, {"kind": "stant", "amount": 3000})
    r = _result(engine, fid)
    assert r["netCiro"] == 16000.0 and r["netAdet"] == 140.0 and r["kitapSayisi"] == 2
    assert r["prev"]["netCiro"] == 12000.0 and round(r["prev"]["degisim"], 4) == round(4000 / 12000, 4)
    assert [(o["tip"], o["adet"], o["tutar"]) for o in r["orders"]] == [(4, 2, 1500.0), (16, 1, 250.0)]
    assert r["costs"] == {"portal": 3000.0, "crm": 1000.0, "byKind": {"Stant kirası": 3000.0}}
    assert r["toplamGider"] == 4000.0 and r["roi"] == 4.0 and r["butceFarki"] == -1000.0
    assert r["katilimci"] == 300                                                          # iptal etkinlik sayılmaz
    assert r["unsoldPlanned"] == [{"stokKodu": "N1", "ad": "Yeni Kitap", "planlanan": 10.0}]
    assert r["sellThrough"] == 100 / 160
    assert any("12.11.2026 tarihine kadar" in w for w in r["warnings"])
    assert any("Logo carisi bağlanmadı" in w for w in r["warnings"])
    assert r["clients"][0]["kod"] == "120.FUAR.01" and r["clients"][0]["ciro"] == 18000.0
    assert len(r["summary"]) >= 4 and "16.000 TL" in r["summary"][0]
    E.save_result(engine, T, fid, r)
    rows = {b["stokKodu"]: b for b in E.fair_detail(engine, T, fid)["bookList"]}
    assert rows["K1"]["qtySold"] == 100.0 and rows["N1"]["qtySold"] == 0.0 and rows["K1"]["sellThrough"] == 100 / 150


def test_result_without_costs_and_with_logo_error(engine):
    fid = _fair(engine, logoClientCodes=["120.FUAR.01"])
    r = _result(engine, fid, sales={"rows": [], "warnings": [], "sql": []}, prev_sales=None, data_end=None,
                crm_events=[], orders=[], logo_error="Veritabanına şu an ulaşılamıyor.")
    assert r["netCiro"] == 0 and r["roi"] is None and r["prev"] is None
    assert any("Logo okunamadı" in w for w in r["warnings"]) and any("Gider girilmedi" in w for w in r["warnings"])
    assert not any("Logo carisi" in w for w in r["warnings"])


def test_basis_uses_previous_fair_card(engine):
    old = _fair(engine, name="TÜYAP 2025", startsOn="2025-11-08", endsOn="2025-11-16", logoClientCodes=["120.OLD"])
    new = _fair(engine, prevFairId=old)
    b = E.basis_window(E.fair_row(engine, T, new), E.fair_row(engine, T, old))
    assert (b["from"], b["to"], b["codes"], b["prevId"]) == ("2025-11-08", "2025-11-16", ["120.OLD"], old)
    b = E.basis_window(E.fair_row(engine, T, new), None)
    assert (b["from"], b["to"]) == ("2025-11-08", "2025-11-16")                           # 364 gün: haftanın aynı günü
    with pytest.raises(E.EventsError):
        E.update_fair(engine, T, "ayse", new, {"prevFairId": new})


# ------------------------------------------------------------------ tip eşlemesi, takvim, ajanda


class FakeChoice(SimpleNamespace):
    pass


def test_type_suggestion_is_not_a_decision(engine):
    types = [{"id": TY_FUAR, "ad": "Kitap Fuarı", "adet": 40, "etkin": True, "son": None},
             {"id": TY_ZIY, "ad": "Cari Ziyareti", "adet": 9000, "etkin": True, "son": None}]
    asked = []

    def choose(prompt, labels):
        asked.append(prompt)
        pick = "Fuar" if "Fuarı" in prompt else "Satış ziyareti"
        return FakeChoice(choice=pick, probability=0.93, margin=0.8, method="logprobs")

    out = E.classify_types(engine, T, types, choose, {TY_FUAR: ["TÜYAP 2025"]})
    assert out == {"asked": 2, "unsure": 0, "total": 2} and "TÜYAP 2025" in asked[0]
    tmap = E.type_map(engine, T)
    assert tmap[TY_FUAR]["suggested"] == "fuar" and tmap[TY_FUAR]["sinif"] is None
    ev = E.classify_events([{"id": "e1", "tipId": TY_FUAR, "baslangic": "2026-03-01", "ad": "x"}], tmap)
    assert ev[0]["sinif"] is None and ev[0]["sinifOneri"] == "fuar"                        # takvim öneriyi kullanmaz
    assert E.classify_types(engine, T, types, choose, {})["asked"] == 0                    # önerisi olan yeniden sorulmaz
    E.set_types(engine, T, "ayse", [{"id": TY_FUAR, "class": "fuar"}], {TY_FUAR: "Kitap Fuarı", TY_ZIY: "Cari Ziyareti"})
    ev = E.classify_events([{"id": "e1", "tipId": TY_FUAR, "baslangic": "2026-03-01", "ad": "x"}], E.type_map(engine, T))
    assert ev[0]["sinif"] == "fuar"
    with pytest.raises(E.EventsError):
        E.set_types(engine, T, "ayse", [{"id": TY_FUAR, "class": "uydurma"}], {TY_FUAR: "x"})
    with pytest.raises(E.EventsError):
        E.set_types(engine, T, "ayse", [{"id": "'; DROP TABLE x --", "class": "fuar"}], {})
    rows = E.type_rows(types, E.type_map(engine, T))
    assert rows[0]["id"] == TY_ZIY and rows[1]["class"] == "fuar"                          # kararı olmayan önce


def test_calendar_hides_sales_visits_but_counts_them(engine):
    fid = _fair(engine, startsOn="2026-10-30", endsOn="2026-11-02")
    fairs = E.list_fairs(engine, T, 2026)
    ev = [{"id": "a", "ad": "Fuar", "baslangic": "2026-03-01", "sinif": "fuar"},
          {"id": "b", "ad": "Ziyaret", "baslangic": "2026-03-02", "sinif": "satis"},
          {"id": "c", "ad": "Bilinmeyen", "baslangic": "2026-03-03", "sinif": None},
          {"id": "d", "ad": "Geçen yıl", "baslangic": "2025-03-03", "sinif": "fuar"}]
    cal = E.calendar(fairs, ev, 2026, ["fuar", "imza"])
    assert [e["id"] for e in cal["events"]] == ["a"]
    assert cal["months"][2]["counts"] == {"fuar": 1, "satis": 1, "yok": 1}
    assert cal["months"][9]["fairs"] == [fid] and cal["months"][10]["fairs"] == [fid]      # iki aya taşan kart
    assert [e["id"] for e in E.calendar(fairs, ev, 2026, ["fuar"], include_unmapped=True)["events"]] == ["a", "c"]


def test_agenda_shows_only_my_records(engine):
    mine = _fair(engine, ownerUser="ayse", startsOn="2026-10-10", endsOn="2026-10-12")
    _fair(engine, ownerUser="mehmet", startsOn="2026-10-10", endsOn="2026-10-12")
    _fair(engine, ownerUser="ayse", startsOn="2027-03-01", endsOn="2027-03-02")                # pencere dışı
    crm = [{"id": "e1", "ad": "İmza günü", "baslangic": "2026-10-05", "saat": "2026-10-05T14:30", "sorumlu": "ayse", "iptal": False},
           {"id": "e2", "ad": "Başkasının", "baslangic": "2026-10-05", "sorumlu": "mehmet", "iptal": False},
           {"id": "e3", "ad": "İptal", "baslangic": "2026-10-06", "sorumlu": "ayse", "iptal": True},
           {"id": "e4", "ad": "Geçmiş", "baslangic": "2026-09-20", "sorumlu": "ayse", "iptal": False}]
    a = E.agenda(engine, T, "AYSE", crm, NOW, days=60)
    kinds = [(i["kind"], i["id"]) for i in a["items"]]
    assert ("crm", "e1") in kinds and ("fuar", mine) in kinds
    assert not any(i in kinds for i in (("crm", "e2"), ("crm", "e3"), ("crm", "e4")))
    assert next(i for i in a["items"] if i["id"] == "e1")["time"] == "14:30"
    assert all(i["kind"] != "gorev" or i["where"] == "TÜYAP İstanbul Kitap Fuarı" for i in a["items"])


def test_reminders_are_written_once(engine, monkeypatch):
    monkeypatch.setenv("EVENTS_REMIND_DAYS", "37,7")
    fid = _fair(engine)
    E.add_task(engine, T, "ayse", fid, {"title": "Geç kaldı", "dueOn": "2026-09-20"})
    aid = E.create_award(engine, T, "editor", {"name": "Çocuk Edebiyatı Ödülü", "deadline": "2026-10-31"})["id"]
    E.put_books(engine, T, "ayse", fid, [{"stokKodu": "K1", "qtyPlanned": 100}], BOOKS)
    with engine.begin() as c:
        c.execute(E.BOOKS.update().values(stock=10))
    first = E.run_reminders(engine, T, NOW)
    assert first["geri-sayim"] == 1 and first["odul"] == 1 and first["stok"] == 1 and first["geciken-gorev"] >= 1
    again = E.run_reminders(engine, T, NOW)
    assert again == {"geri-sayim": 0, "geciken-gorev": 0, "odul": 0, "stok": 0}
    up = E.upcoming(engine, T, NOW)
    assert up["awards"][0]["id"] == aid and up["awards"][0]["daysLeft"] == 30
    assert any(t["title"] == "Geç kaldı" for t in up["lateTasks"]) and len(up["reminders"]) >= 4


def test_fairs_needing_result(engine):
    fid = _fair(engine, startsOn="2026-09-01", endsOn="2026-09-05")
    assert E.fairs_needing_result(engine, T, NOW) == [fid]
    E.save_result(engine, T, fid, {"books": []})
    assert E.fairs_needing_result(engine, T, NOW) == []


# ------------------------------------------------------------------ ödül defteri


def test_award_entries_flow(engine):
    aid = E.create_award(engine, T, "editor", {"name": "Ödül", "deadline": "2026-12-01", "url": "https://ornek.org"})["id"]
    with pytest.raises(E.EventsError):
        E.create_award(engine, T, "editor", {"name": "Kötü", "url": "javascript:alert(1)"})
    with pytest.raises(E.EventsError):
        E.add_entry(engine, T, "editor", aid, {"stokKodu": "YOK"}, BOOKS)
    eid = E.add_entry(engine, T, "editor", aid, {"stokKodu": "k2"}, BOOKS)["id"]
    E.patch_entry(engine, T, "editor", eid, {"status": "gonderildi"})
    e = E.list_awards(engine, T, NOW)["items"][0]["entries"][0]
    assert e["bookName"] == "Kitap İki" and e["stokKodu"] == "K2" and e["submittedAt"] == E.today().isoformat()
    with pytest.raises(E.EventsError):
        E.patch_entry(engine, T, "editor", eid, {"status": "bilinmiyor"})
    assert E.delete_award(engine, T, aid)["entries"] == 1


# ------------------------------------------------------------------ SQL kuruluşu


def test_fair_sales_sql_uses_invoiced_lines_channel_and_codes():
    q = S.fair_sales_sql("411", date(2026, 11, 7), date(2026, 11, 16), "FUAR", ["120.FUAR.01", "x'; DROP--", "120.02"])
    assert "LG_411_01_STLINE" in q and "LG_411_CLCARD" in q and "S.INVOICEREF <> 0" in q and "S.LINETYPE = 0" in q
    assert "S.CANCELLED = 0" in q and "C.SPECODE2 = 'FUAR'" in q and "TRCODE IN (7,8,9) THEN S.VATMATRAH ELSE -S.VATMATRAH" in q
    assert "SH.DATE_ >= '2026-11-07' AND SH.DATE_ < '2026-11-16'" in q
    assert "C.CODE IN (N'120.FUAR.01', N'120.02')" in q and "DROP" not in q
    with pytest.raises(S.SourceError):
        S.fair_sales_sql("41'1", date(2026, 1, 1), date(2026, 1, 2), "FUAR")
    with pytest.raises(S.SourceError):
        S.fair_sales_sql("411", date(2026, 1, 1), date(2026, 1, 2), "FU'AR")


def test_crm_sql_uses_istanbul_day_bounds_and_safe_ids():
    assert S.utc_bound(date(2026, 11, 7)) == "2026-11-06 21:00:00"
    q = S.events_sql("Timas_MSCRM.dbo", date(2026, 1, 1), date(2027, 1, 1))
    assert "e.new_BalangTarihi >= '2025-12-31 21:00:00'" in q and "e.statecode = 0" in q
    o = S.orders_sql("Timas_MSCRM.dbo", date(2026, 11, 7), date(2026, 11, 16), [4, 5, 16], [100000001])
    assert "new_siparistipi IN (4, 5, 16)" in o and "statuscode NOT IN (100000001)" in o
    assert S.guids(["{AAAAAAAA-0000-0000-0000-000000000001}", "x' OR 1=1", None]) == ["aaaaaaaa-0000-0000-0000-000000000001"]
    with pytest.raises(S.SourceError):
        S.author_events_sql("Timas_MSCRM.dbo", "1; DROP", date(2026, 1, 1), date(2027, 1, 1))
    with pytest.raises(S.SourceError):
        S.events_sql("Timas;DROP.dbo", date(2026, 1, 1), date(2026, 2, 1))
    assert S.year_slices(date(2025, 12, 30), date(2026, 1, 3)) == [(2025, date(2025, 12, 30), date(2026, 1, 1)),
                                                                    (2026, date(2026, 1, 1), date(2026, 1, 3))]
    assert S.dayiso("2026-11-06T21:00:00") == "2026-11-07" and S.dayiso("1900-01-01") is None
    assert S.account("TIMAS\\Ayse") == "ayse"


def test_event_row_maps_crm_fields():
    r = S.event_row({"id": "{AAAAAAAA-0000-0000-0000-000000000001}", "ad": " TÜYAP ", "tip_id": TY_FUAR, "baslangic": "2026-11-06T21:00:00",
                     "bitis": "2026-11-01T21:00:00", "durum": 100000000, "sorumlu_hesap": "TIMAS\\Ayse", "gider": "1500"})
    assert r["baslangic"] == "2026-11-07" and r["bitis"] == "2026-11-07" and r["iptal"] and r["sorumlu"] == "ayse"
    assert r["id"] == "aaaaaaaa-0000-0000-0000-000000000001" and r["gider"] == 1500.0 and r["ad"] == "TÜYAP"


# ------------------------------------------------------------------ yetki


def test_access_rules_for_events():
    assert A.rule_for("/api/v1/events/calendar") == {"sayfa:etkinlikler"}
    assert A.rule_for("/api/v1/events/me/agenda") == A.OPEN
    assert A.rule_for("/api/v1/events/run-due") == A.SYSTEM
    f = A.features_for
    assert f("POST", "/api/v1/events/fairs") == ["ozellik:etkinlik.duzenle"]
    assert f("PUT", "/api/v1/events/fairs/f1/books") == ["ozellik:etkinlik.duzenle"]
    assert f("POST", "/api/v1/events/fairs/f1/costs") == ["ozellik:etkinlik.duzenle"]
    assert f("POST", "/api/v1/events/fairs/f1/tasks") == ["ozellik:etkinlik.duzenle"]
    assert f("PATCH", "/api/v1/events/fairs/f1/tasks/t1") == []                           # sahibi işaretler; ucun içinde
    assert f("POST", "/api/v1/events/fairs/f1/approve") == []                             # açıkça verilen onay ucun içinde
    assert f("PUT", "/api/v1/events/type-map") == ["ozellik:etkinlik.duzenle"]
    assert f("POST", "/api/v1/events/awards/a1/entries") == ["ozellik:odul.duzenle"]
    assert f("PATCH", "/api/v1/events/award-entries/e1") == ["ozellik:odul.duzenle"]
    assert f("GET", "/api/v1/events/fairs/f1/result/export.pdf") == ["ozellik:veri.disa-aktar"]
    assert f("GET", "/api/v1/events/fairs/f1/result") == [] and f("GET", "/api/v1/events/calendar") == []
    assert "ozellik:etkinlik.onay" in A.explicit_keys()
    assert {"sayfa:etkinlikler", "ozellik:etkinlik.duzenle", "ozellik:odul.duzenle"} <= A.all_keys()


# ------------------------------------------------------------------ uçlar


class FakeSource:
    def __init__(self):
        self.calls = []

    def events(self, frm, to, fresh=False):
        self.calls.append(("events", frm, to))
        return [{"id": "e1", "ad": "İmza günü", "tipId": TY_FUAR, "baslangic": E.today().isoformat(), "saat": None,
                 "sorumlu": "ayse", "iptal": False, "yer": "Kadıköy", "il": "İstanbul"},
                {"id": "e2", "ad": "Ziyaret", "tipId": TY_ZIY, "baslangic": E.today().isoformat(), "sorumlu": "mehmet", "iptal": False}]

    def window_events(self, frm, to, fresh=False):
        """Gerçek kaynaktaki gibi: takvim yılı okumaları, başlangıcı [frm, to) aralığında olanlar."""
        from datetime import date as _d

        from semantic_bridge.events_sources import year_slices

        out = []
        for y, _a, _b in year_slices(frm, to):
            out += [r for r in self.events(_d(y, 1, 1), _d(y + 1, 1, 1), fresh)
                    if frm.isoformat() <= r["baslangic"] < to.isoformat()]
        return out

    def types(self, fresh=False):
        return [{"id": TY_FUAR, "ad": "Fuar", "adet": 1, "etkin": True, "son": None}]


@pytest.fixture
def client(engine, monkeypatch):
    from fastapi import FastAPI, HTTPException, Request
    from fastapi.testclient import TestClient

    from semantic_bridge import events_api

    users = {"a": "ayse", "m": "mehmet"}
    perms = {"ayse": {"ozellik:etkinlik.duzenle"}, "mehmet": {"ozellik:etkinlik.onay"}}

    def auth(request: Request):
        u = users.get(request.headers.get("x-user", ""))
        if not u:
            raise HTTPException(status_code=401)
        return engine, T, u, u.title()

    fake = FakeSource()
    monkeypatch.setattr(events_api.src, "Source", lambda *a, **k: fake)
    app = FastAPI()
    svc = events_api.register(app, {
        "auth": auth, "can": lambda u, k: k in perms.get(u, set()), "is_admin": lambda u: False,
        "audit": lambda *a, **k: None, "conf": lambda k, d="": d, "fresh": lambda: False,
        "crm_connect": lambda: None, "logo_connect": lambda: None, "llm": lambda p: None,
        "system": lambda: (engine, T), "require_caller": lambda r: None,
    })
    return TestClient(app), svc


def test_endpoints_agenda_and_approval(client):
    c, svc = client
    a, m = {"x-user": "a"}, {"x-user": "m"}
    made = c.post("/api/v1/events/fairs", json={"name": "Kart", "startsOn": E.today().isoformat(), "endsOn": E.today().isoformat()}, headers=a)
    assert made.status_code == 201
    fid = made.json()["id"]
    assert c.post(f"/api/v1/events/fairs/{fid}/approve", json={}, headers=a).status_code == 403   # onay yetkisi yok
    ok = c.post(f"/api/v1/events/fairs/{fid}/approve", json={"note": "tamam"}, headers=m)
    assert ok.status_code == 200 and ok.json()["status"] == "onayli"
    bad = c.post("/api/v1/events/fairs", json={"name": ""}, headers=a)
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "EVENTS"
    ag = c.get("/api/v1/events/me/agenda", headers=a).json()
    ids = {i["id"] for i in ag["items"]}
    assert "e1" in ids and "e2" not in ids and fid in ids
    assert c.get("/api/v1/events/me/agenda").status_code == 401
    meta = c.get("/api/v1/events/meta", headers=a).json()
    assert meta["me"]["canEdit"] and not meta["me"]["canApprove"] and meta["classes"]["fuar"] == "Fuar"
    assert c.post("/api/v1/events/type-map/suggest", headers=a).status_code == 503                # model bağlı değil
    cal = c.get("/api/v1/events/calendar", headers=a).json()
    assert cal["unmappedTypes"] == 2 and cal["events"] == []


def test_book_card_sql_uses_crm_column_spelling():
    """CRM harmanlaması Türkçe (I → ı): new_kitapid yazımı 207 ile düşüyordu, kitap önerisi 502 (2026-09-28 kabul)."""
    sql = S.books_sql("Timas_MSCRM.dbo")
    assert "k.new_kitapId AS id" in sql and "k.new_StokKodu AS stok_kodu" in sql and "new_kitapid" not in sql


# ------------------------------------------------------------------ fuar takvimi (Excel)


def _plan_xlsx() -> bytes:
    """FUARLAR.xlsx'in iki sayfası: Gantt (ad başladığı günün sütununda, yıl satırları) + düz liste ve «netleşmedi» notu."""
    import io

    from openpyxl import Workbook

    wb = Workbook()
    g = wb.active
    g.title = "GÜNCEL FUAR TAKVİMİ"
    g.append([None, *range(1, 32), "Fuar Tarihi", "Fuar Gün sayısı", "Fuar Alanı", "Düzenleyen Fuar Firması", "Katılımcı Firma"])
    g.append([2026])
    pad = lambda col, name: [None] * col + [name] + [None] * (30 - col)  # noqa: E731  (1–31. gün sütunları)
    g.append(["EYLÜL", *pad(24, "Uşak Kitap Fuarı"), "25 EYLÜL-04 EKİM", "10 GÜN", None, None, "bayi"])
    g.append(["EKİM", "Rami", *[None] * 30])                                # tarihsiz kısa ad: atlanır
    g.append([None, *pad(2, "Kocaeli Kitap Fuarı"), "03 EKİM-11 EKİM", "   9 GÜN", "Kocaeli kongre merkezi", "ka2 ajans", "TİMAŞ"])
    g.append([None, *pad(0, "Bozuk Fuar"), "yakında", None, None, None, "bayi"])
    g.append([2027])
    g.append(["OCAK", *pad(8, "Çukurova Kitap Fuarı"), "09 OCAK-17 OCAK", "   9 GÜN", "Adana", "TÜYAP", "TİMAŞ"])
    s = wb.create_sheet("fuar liste")
    s.append([None])
    s.append(["Sıra no", "Fuar Adı", "Fuar Tarihi", "KATILIMCI FİRMA", None, None, None, "2027 fuarlar  tarih netleşmedi"])
    s.append([1, "Uşak kitap fuarı ", "25 EYLÜL-04 EKİM", "BAYİ", None, None, None, "Arnavutköy "])
    s.append([2, "Kocaeli Kitap Fuarı", "03 - 11 EKİM", "TİMAŞ", None, None, None, "Bursa"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_fair_plan_parses_gantt_sheet_years_and_pending():
    from semantic_bridge import fair_plan as FP

    p = FP.parse_workbook(_plan_xlsx(), date(2026, 10, 6))
    assert p["sheet"] == "GÜNCEL FUAR TAKVİMİ"                       # alanı çok olan sayfa seçilir
    names = [i["name"] for i in p["items"]]
    assert names == ["Uşak Kitap Fuarı", "Kocaeli Kitap Fuarı", "Çukurova Kitap Fuarı"]
    u, k, c = p["items"]
    assert (u["startsOn"], u["endsOn"], u["days"]) == ("2026-09-25", "2026-10-04", 10)
    assert k["participant"] == "timas" and k["participantLabel"] == "TİMAŞ" and k["venue"] == "Kocaeli kongre merkezi"
    assert u["participant"] == "bayi" and u["participantLabel"] == "Bayi"
    assert c["startsOn"] == "2027-01-09"                                # 2027 yıl satırı
    assert p["pending"] == {"title": "2027 fuarlar tarih netleşmedi", "items": ["Arnavutköy", "Bursa"]}
    assert any("Bozuk Fuar" in w for w in p["warnings"])
    assert FP.parse_range("02 - 11 EKİM", 2026) == (date(2026, 10, 2), date(2026, 10, 11))
    assert FP.parse_range("28 ARALIK-03 OCAK", 2026) == (date(2026, 12, 28), date(2027, 1, 3))
    with pytest.raises(FP.PlanError):
        FP.parse_workbook(b"not an excel", date(2026, 10, 6))


def test_fair_plan_endpoints(client, monkeypatch):
    c, _svc = client
    monkeypatch.setattr(E, "today", lambda: date(2026, 10, 6))
    a, m = {"x-user": "a"}, {"x-user": "m"}
    empty = c.get("/api/v1/events/me/fair-plan", headers=m).json()
    assert empty["items"] == [] and not empty["canEdit"]
    body = {"fileName": "FUARLAR.xlsx", "dataBase64": base64.b64encode(_plan_xlsx()).decode()}
    assert c.post("/api/v1/events/fair-plan", json=body, headers=m).status_code == 403          # düzenleme yetkisi yok
    bad = c.post("/api/v1/events/fair-plan", json={"dataBase64": "!!"}, headers=a)
    assert bad.status_code == 422
    out = c.post("/api/v1/events/fair-plan", json=body, headers=a).json()
    assert len(out["items"]) == 3 and out["fileName"] == "FUARLAR.xlsx" and out["uploadedBy"] == "ayse"
    got = c.get("/api/v1/events/me/fair-plan", headers=m).json()
    assert [(i["phase"], i["daysLeft"]) for i in got["items"]] == [("bitti", -2), ("suruyor", 5), ("yaklasan", 95)]
    assert A.rule_for("/api/v1/events/me/fair-plan") == A.OPEN
    assert A.features_for("POST", "/api/v1/events/fair-plan") == ["ozellik:etkinlik.duzenle"]


def test_fair_plan_template_round_trip_and_bad_rows():
    """Şablon güncel listeyle iner, aynen geri yüklenince aynı takvim çıkar; eksik/ters tarihli satır atlanır."""
    import io

    from openpyxl import load_workbook

    from semantic_bridge import fair_plan as FP

    today = date(2026, 10, 6)
    src = FP.parse_workbook(_plan_xlsx(), today)
    data = FP.template({"items": src["items"], "pending": src["pending"]})
    back = FP.parse_workbook(data, today)
    keep = ("name", "startsOn", "endsOn", "venue", "organizer", "participant")
    assert back["sheet"] == "Fuarlar" and back["warnings"] == []
    assert [{k: i[k] for k in keep} for i in back["items"]] == [{k: i[k] for k in keep} for i in src["items"]]
    assert back["pending"]["items"] == ["Arnavutköy", "Bursa"]
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Fuarlar"]
    assert [c.value for c in ws[1]] == FP.TEMPLATE_HEAD and ws["B2"].number_format == "DD.MM.YYYY"
    assert any("TİMAŞ,Bayi" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)
    ws.append(["Tarihsiz Fuar", None, None, None, None, "Bayi"])               # tarih yok: satır boş sayılır
    ws.append(["Ters Fuar", date(2026, 11, 5), date(2026, 11, 1), None, None, "Bayi"])
    ws.append(["Metin Tarihli", "05.11.2026", "2026-11-08", None, None, "TİMAŞ"])
    buf = io.BytesIO()
    wb.save(buf)
    p = FP.parse_workbook(buf.getvalue(), today)
    names = [i["name"] for i in p["items"]]
    assert "Metin Tarihli" in names and "Ters Fuar" not in names and "Tarihsiz Fuar" not in names
    assert any("Ters Fuar" in w and "önce" in w for w in p["warnings"])


def test_fair_plan_template_endpoint(client):
    c, _svc = client
    assert A.features_for("GET", "/api/v1/events/fair-plan/template.xlsx") == ["ozellik:etkinlik.duzenle"]
    r = c.get("/api/v1/events/fair-plan/template.xlsx", headers={"x-user": "a"})
    assert r.status_code == 200 and r.content[:2] == b"PK" and "sablon" in r.headers["content-disposition"]
