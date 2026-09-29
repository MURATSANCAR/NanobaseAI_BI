"""CRM'de departmansız kullanıcılar listesi: kök iş birimi, AD birimi eşlemesi, alıcı süzgeci, zamanlayıcı dosyası."""

from __future__ import annotations

import io
import os
from datetime import datetime

from semantic_bridge import crm_unassigned as CU

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))


class FakeDirectory:
    def __init__(self):
        self.sql: list[str] = []

    def _crm_prefix(self):
        return "CRM.dbo."

    def _crm_rows(self, sql):
        self.sql.append(sql)
        if "SystemUserBase" in sql and "BusinessUnitBase" in sql:
            return [
                {"Id": "A", "FullName": "Ali Kaya", "DomainName": "TIMAS\\ali", "InternalEMailAddress": "ali@timas.com.tr",
                 "CreatedOn": datetime(2025, 1, 2), "BirimAdi": "Timaş CRM"},
                {"Id": "B", "FullName": "Eski Hesap", "DomainName": "TIMAS\\eski", "InternalEMailAddress": None,
                 "CreatedOn": None, "BirimAdi": "Timaş CRM"},
            ]
        return [{"Id": "a", "Name": "07-Temel Rol"}, {"Id": "a", "Name": "ACTION_Depo"}]

    def list_people(self):
        return [{"subject": "ali", "label": "Ali Kaya", "hint": "ali", "detail": "Satis"}]


def conf_of(values):
    return lambda key, default="": values.get(key, default)


def test_root_unit_is_read_by_structure_not_by_name():
    sql = CU.users_sql("CRM.dbo.")
    assert "ParentBusinessUnitId IS NULL" in sql and "Timaş" not in sql
    assert "STRING_AGG" not in CU.roles_sql("CRM.dbo.")          # CRM sunucusu bu işlevi tanımıyor


def test_rows_carry_ad_unit_and_split_roles():
    rows = CU.build(FakeDirectory())
    assert [(r["name"], r["adUnit"]) for r in rows] == [("Eski Hesap", CU.NO_AD), ("Ali Kaya", "Satis")]
    ali = rows[1]
    assert ali["mainRoles"] == ["07-Temel Rol"] and ali["roles"] == ["07-Temel Rol", "ACTION_Depo"]
    assert ali["created"] == "2025-01-02"
    sm = CU.summary(rows)
    assert (sm["count"], sm["noAd"], sm["noRole"]) == (2, 1, 1)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(CU.xlsx(rows)))
    assert wb.sheetnames == ["Kişiler", "AD birimine göre"] and wb["Kişiler"].max_row == 3


def test_no_recipient_means_no_crm_read_and_domain_filter_applies():
    d = FakeDirectory()
    out = CU.send(d, conf_of({}), lambda *a: "sent")
    assert out["status"] == "no_recipient" and d.sql == []
    sent = []
    out = CU.send(d, conf_of({"CRM_UNASSIGNED_TO": "a@timas.com.tr; b@gmail.com",
                              "ALERT_RECIPIENT_DOMAINS": "timas.com.tr"}),
                  lambda subject, text, to, att: sent.append((subject, to, att[0][0])) or "sent",
                  now=datetime(2026, 9, 29, 7, 0))
    assert out["ok"] and out["to"] == ["a@timas.com.tr"] and out["skipped"] == ["b@gmail.com"]
    assert sent[0][1] == ["a@timas.com.tr"] and sent[0][2].endswith("2026-09-29-0700.xlsx")
    assert "2 kişi" in sent[0][0]


def _timer(name):
    with open(os.path.join(ROOT, "scripts", "server", name), encoding="utf-8") as f:
        return [ln.strip() for ln in f if ln.strip().startswith("OnCalendar=")]


def test_timer_files_run_at_seven_and_noon():
    for name in ("timas-admin-group.timer", "timas-crm-unassigned.timer"):
        assert _timer(name) == ["OnCalendar=*-*-* 07:00:00", "OnCalendar=*-*-* 12:00:00"], name
