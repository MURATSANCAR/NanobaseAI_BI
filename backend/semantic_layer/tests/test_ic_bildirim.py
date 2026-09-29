"""Bilgi İşlem'e giden beş bildirimin ortak şablonu (ic_bildirim) ve kopma bildiriminin gürültü önlemesi.

Sınananlar: her bildirimin HTML'i ve düz metni üretilir; konu durumla başlar; «Ne oldu / Etkisi / Ne yapmalı»,
bağlantı ve alt not var; teknoloji adı yok; ileti `multipart/alternative` (ekliyse `multipart/mixed`); departmansız
CRM listesinde yalnız ad, kullanıcı adı ve oluşturulma tarihi. Gürültü: kısa kesinti ve köprünün planlı yeniden
başlatması olay açmaz, e-posta göndermez; uzun kesintide tek kopma ve tek düzelme e-postası gider.

Veriler yapaydır ve yalnız kuralları sınar; gerçek tablolarla kabul `scripts/acceptance/bt-eposta/kabul.py`.
"""

from __future__ import annotations

import email
import io
import re
from datetime import datetime, timedelta, timezone

import pytest
import sqlalchemy as sa

from semantic_bridge import crm_unassigned as CU
from semantic_bridge import data_security as D
from semantic_bridge import ic_bildirim as IB
from semantic_bridge import it_ops as I
from semantic_bridge import mailbox as M
from semantic_layer.store.catalog_store import open_store
from semantic_layer.tests.conftest import TENANT

UTC = timezone.utc
LOCAL = timezone(timedelta(hours=3))
T0 = datetime(2026, 9, 29, 6, 42, tzinfo=UTC)          # 09:42 İstanbul
LINK = "https://portal.timas.com.tr/timas/uyarilar"
BT = ["bilgiislem@timas.com.tr"]
ST = {"everySec": 300, "failsToOpen": 2, "outageMin": 10, "restartGraceMin": 5, "resolveMin": 15, "logoStaleDays": 3,
      "crmStaleHours": 24, "remindHours": 0, "staleRemindHours": 0, "weeklyDay": 1, "reportHour": 8}
FRESH = datetime(2026, 9, 29, 6, 0, tzinfo=UTC)                # Logo'dan okunan son fatura (veri okuması şartı)
DIRTY = "OperationalError: FreeTDS pyodbc; uvicorn systemd nginx; vLLM Qwen3.8-27B; Temporal PostgreSQL docker"


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    I._ready.discard(id(e))
    I.ensure(e)
    return e


class Outbox:
    def __init__(self, result: str = "sent"):
        self.result = result
        self.notices: list[IB.Notice] = []

    def __call__(self, notice, to):
        self.notices.append(notice)
        return self.result


def tour(engine, ring, ok, at, st=ST, detail="zaman aşımı", data_end=None):
    if ok and ring in I.STALE_RINGS and data_end is None:
        data_end = FRESH                                         # başarılı Logo/CRM denemesi veri sonunu okur
    if data_end is False:
        data_end = None                                          # «bağlandı ama veri okunamadı» denemesi
    I.record_check(engine, TENANT, ring, ok, detail=detail, at=at, data_end=data_end)
    return I.evaluate(engine, TENANT, [{"ring": ring, "ok": ok, "detail": detail, "data_end": data_end}], st, now=at)


def closed_ids(engine):
    return [x["id"] for x in I.list_incidents(engine, TENANT, state="closed")["items"]]


def assert_template(n: IB.Notice, tag: str, link_part: str) -> tuple[str, str]:
    """Her bildirimin ortak sözleşmesi."""
    html, text = IB.render_html(n), IB.render_text(n)
    assert n.subject.startswith(f"[{tag}] "), n.subject
    assert len(n.subject) <= 200 and n.headline
    for title in ("Ne oldu", "Etkisi", "Ne yapmalı"):
        assert f">{title}</p>" in html, title
    for title in ("NE OLDU", "ETKİSİ", "NE YAPMALI"):
        assert title in text, title
    assert IB.FOOTER in html and IB.FOOTER in text
    assert link_part in n.link and n.link in text and f'href="{n.link}"' in html
    assert "max-width:640px" in html and "<link" not in html and 'lang="tr"' in html
    assert IB.tech_words(html) == [] and IB.tech_words(text) == [], (IB.tech_words(html), IB.tech_words(text))
    assert re.search(r"(parola|password|şifre|anahtar)\w*\s*[:=]", text, re.I) is None   # değer yazılmaz
    return html, text


# ------------------------------------------------------------------ şablon


def test_turkish_suffixes_and_dates():
    at = lambda h, m: datetime(2026, 9, 29, h, m, tzinfo=LOCAL)  # noqa: E731
    assert IB.hm_suffix(at(9, 42)) == "09:42'den" and IB.hm_suffix(at(10, 40)) == "10:40'tan"
    assert IB.hm_suffix(at(9, 43)) == "09:43'ten" and IB.hm_suffix(at(9, 36)) == "09:36'dan"
    assert IB.hm_suffix(at(12, 0), "de") == "12:00'de" and IB.hm_suffix(at(10, 27), "de") == "10:27'de"
    assert IB.hm_suffix(at(9, 30), "de") == "09:30'da" and IB.hm_suffix(at(10, 0)) == "10:00'dan"
    assert IB.suffix("Ağustos") == "tan" and IB.suffix("Eylül") == "den"
    assert IB.day_range(datetime(2026, 9, 22, 8, tzinfo=LOCAL), datetime(2026, 9, 28, 8, tzinfo=LOCAL)) == "22–28 Eylül"
    assert IB.day_range(datetime(2026, 9, 29, tzinfo=LOCAL), datetime(2026, 10, 5, tzinfo=LOCAL)) == "29 Eylül – 5 Ekim"
    assert IB.portal_link(LINK, "sistem-durumu") == "https://portal.timas.com.tr/timas/sistem-durumu"
    assert IB.portal_link("", "sistem-durumu") == ""


def test_plain_removes_technology_and_error_class_names():
    t = IB.plain(DIRTY)
    assert IB.tech_words(t) == [] and "OperationalError" not in t
    assert "Zeki AI modeli" in t and "veritabanı" in t
    assert IB.plain("smtp.gmail.com:587 reddetti") == "smtp.gmail.com:587 reddetti"   # sunucu adı bozulmaz


def test_message_is_multipart_alternative_and_mixed_with_attachment():
    n = IB.Notice(tone="bilgi", subject=IB.subject("Bilgi", "deneme"), headline="Tek cümle.", what=["a"], impact=["b"],
                  actions=["c"], link="https://x/timas/yonetim", at=T0)
    msg = email.message_from_bytes(IB.message(n, "zeki@timas.com.tr", BT).as_bytes())
    assert msg.get_content_type() == "multipart/alternative"
    assert [p.get_content_type() for p in msg.get_payload()] == ["text/plain", "text/html"]
    msg = email.message_from_bytes(IB.message(n, "zeki@timas.com.tr", BT, [("a.xlsx", b"x", CU.XLSX_MIME)]).as_bytes())
    assert msg.get_content_type() == "multipart/mixed"
    parts = [p.get_content_type() for p in msg.walk()]
    assert "multipart/alternative" in parts and CU.XLSX_MIME in parts


def test_html_is_escaped():
    n = IB.Notice(tone="kesinti", subject="[Kesinti] x", headline="<script>alert(1)</script>", what=["<b>"], at=T0)
    html = IB.render_html(n)
    assert "<script>" not in html and "&lt;script&gt;" in html


# ------------------------------------------------------------------ 1) kopma ve düzelme — gürültü önleme


def test_short_outage_sends_nothing(engine):
    out = Outbox()
    tour(engine, "logo", False, T0)
    tour(engine, "logo", False, T0 + timedelta(minutes=5))           # iki deneme ama 5 dk < 10 dk
    assert not I.list_incidents(engine, TENANT)["items"]
    tour(engine, "logo", True, T0 + timedelta(minutes=8))
    I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=8))
    assert out.notices == [] and not I.list_incidents(engine, TENANT, state="all")["items"]


def test_planned_restart_is_not_an_outage(engine):
    out = Outbox()
    boot = T0 + timedelta(minutes=6)
    assert I.note_boot(engine, TENANT, boot) and not I.note_boot(engine, TENANT, boot)   # süreç başına bir kez
    # Durma (açılıştan önce) ve açılış (sonra) sırasındaki başarısız denemeler sayılmaz.
    for m in (4, 8, 10):
        tour(engine, "model", False, T0 + timedelta(minutes=m), detail="kapı durdu")
    I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=10))
    assert out.notices == [] and not I.list_incidents(engine, TENANT, state="all")["items"]
    assert I.in_restart_window(T0 + timedelta(minutes=2), [boot], 5) and not I.in_restart_window(T0, [boot], 5)
    # Açılış kaydı yalnız son BOOT_KEEP_DAYS günü tutar.
    I.note_boot(engine, TENANT, boot + timedelta(days=I.BOOT_KEEP_DAYS + 1))
    assert I.boots(engine, TENANT) == [boot + timedelta(days=I.BOOT_KEEP_DAYS + 1)]


def test_long_outage_sends_one_down_and_one_fixed(engine):
    out = Outbox()
    for m in (0, 5, 10, 15, 20, 25):
        tour(engine, "logo", False, T0 + timedelta(minutes=m), detail=DIRTY)
        I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=m))
    assert len(out.notices) == 1                                    # açıldı (10. dk), sonra ikinci e-posta yok
    down = out.notices[0]
    assert down.subject == "[Kesinti] Logo bağlantısı 09:42'den beri yanıt vermiyor"
    html, text = assert_template(down, "Kesinti", "/timas/sistem-durumu")
    assert "Satış, ciro" in text and "3 kez art arda" in text     # etkisi ve 10. dk'ya dek deneme sayısı
    # «Otomatik düzelir» izlenimi yok; ne olacağı tek cümle.
    assert "kendiliğinden" not in text and "yapmanız gerekmez" not in text
    assert "Bağlantı 15 dakika kesintisiz çalışıp veri okununca «Düzeldi» e-postası gelir." in text
    assert "kendiliğinden kapanır" not in I.RING_BY_ID["logo"]["recipe"]
    # İlk başarılı deneme olayı kapatmaz: 15 dk kesintisiz çalışma + veri okuması gerekir.
    for m in (45, 50, 55):
        tour(engine, "logo", True, T0 + timedelta(minutes=m))
        I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=m))
    assert len(out.notices) == 1 and not closed_ids(engine)
    st = I.status(engine, TENANT, ST, now=T0 + timedelta(minutes=55))
    assert "doğrulama sürüyor" in st["summary"]["text"]
    tour(engine, "logo", True, T0 + timedelta(minutes=60))
    I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=60))
    I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=65))
    assert len(out.notices) == 2
    fixed = out.notices[1]
    assert fixed.subject == "[Düzeldi] Logo bağlantısı 10:27'de geri geldi · 45 dk sürdü" and fixed.tone == "duzeldi"
    _, ftext = assert_template(fixed, "Düzeldi", "/timas/sistem-durumu")
    assert "Kesintisiz çalışma: Sağlandı: 10:27'den bu yana 15 dk" in ftext
    assert "Veri okuması: Sağlandı: Logo'dan son fatura tarihi okundu (29 Eylül 2026)" in ftext
    assert "sık çalışan zamanlanmış iş yok; düzelme ilk iki şartla doğrulandı" in ftext
    assert IB.TONES["duzeldi"]["bar"] in IB.render_html(fixed) and IB.TONES["kesinti"]["bar"] in html


def test_failure_during_recovery_keeps_the_same_incident(engine):
    out = Outbox()
    plan = [(0, False), (5, False), (10, False), (15, True), (20, False), (25, True), (30, True), (35, True), (40, True)]
    for m, ok in plan:
        tour(engine, "logo", ok, T0 + timedelta(minutes=m))
        I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=m))
    assert [n.badge for n in out.notices] == ["Kesinti", "Düzeldi"]          # yeni «Kesinti» yok
    assert out.notices[1].subject == "[Düzeldi] Logo bağlantısı 10:07'de geri geldi · 25 dk sürdü"
    assert len(I.list_incidents(engine, TENANT, state="all")["items"]) == 1


def test_no_data_read_or_no_job_run_keeps_the_incident_open(engine):
    out = Outbox()
    for m in (0, 5, 10):
        tour(engine, "logo", False, T0 + timedelta(minutes=m))
    I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=10))
    # (b) bağlandı ama veri sonu okunamadı: 20 dk başarılı deneme olsa da kapanmaz
    for m in (15, 20, 25, 30, 35):
        tour(engine, "logo", True, T0 + timedelta(minutes=m), data_end=False)
    assert not closed_ids(engine)
    # (c) bağlantıyı kullanan sık iş var ama kurtulmadan sonra koşmadı → bekler; koşunca kapanır
    I.upsert_job(engine, TENANT, "kopru-saglik", label="Sorgu motoru sağlık denetimi", source="watchdog",
                 last_at=T0 + timedelta(minutes=2), last_ok=True)
    for m in (40, 45, 50):
        tour(engine, "logo", True, T0 + timedelta(minutes=m))
    assert not closed_ids(engine)
    I.upsert_job(engine, TENANT, "kopru-saglik", label="Sorgu motoru sağlık denetimi", source="watchdog",
                 last_at=T0 + timedelta(minutes=52), last_ok=True)
    tour(engine, "logo", True, T0 + timedelta(minutes=55))
    assert closed_ids(engine)
    I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(minutes=55))
    text = IB.render_text(out.notices[-1])
    assert "Zamanlanmış iş: Sağlandı: «Sorgu motoru sağlık denetimi»" in text


def test_stale_data_notice_says_since_when(engine):
    end = datetime(2026, 8, 17, 18, 0, tzinfo=UTC)
    I.record_check(engine, TENANT, "logo", True, data_end=end, at=T0)
    I.evaluate(engine, TENANT, [{"ring": "logo", "ok": True, "data_end": end}], ST, now=T0)
    out = Outbox()
    I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0)
    n = out.notices[0]
    assert n.subject == "[Uyarı] Logo verisi 17 Ağustos'tan beri güncellenmiyor" and n.tone == "uyari"
    _, text = assert_template(n, "Uyarı", "/timas/sistem-durumu")
    # Somut BT adımı; «BT ile görüşün» yok (e-posta zaten BT'ye gider); sunucu adı/adresi yok.
    assert "BT ile görüşün" not in text and "canlı Logo sunucusunda okuma yetkisi verin" in text
    assert "Yönetim → Ayarlar → Logo bağlantısında canlı Logo sunucusunu seçin" in text
    assert re.search(r"\b\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}\b", text) is None
    I.notify(engine, TENANT, ST, out, BT, link=LINK, now=T0 + timedelta(days=3))
    assert len(out.notices) == 1                                   # hatırlatma varsayılanda kapalı


def test_jobs_digest_lists_only_it_actions():
    jobs = [{"label": "Planlı raporlar", "lastAt": T0.isoformat(), "lastError": "Logo'ya ulaşılamadı. " + DIRTY,
             "every": "5 dk"},
            {"label": "Stok gece okuması", "lastAt": T0.isoformat(), "lastError": "Sayı dönüşümü taştı", "every": "gece"}]
    n = I.jobs_digest_notice(jobs, T0, LINK)
    assert n.subject == "[Uyarı] Zamanlanmış işler · 1 iş bağlantı ya da yetki hatası verdi · 29 Eylül"
    _, text = assert_template(n, "Uyarı", "/timas/sistem-durumu")
    assert "Stok gece okuması" not in text and "1 zamanlanmış iş uygulama kaynaklı hata verdi" in text
    assert I.split_jobs(jobs) == ([jobs[0]], [jobs[1]])


# ------------------------------------------------------------------ 2) haftalık sağlık özeti


def test_weekly_notice(engine):
    for m in (0, 5, 10):
        tour(engine, "vpn", False, T0 - timedelta(days=2) + timedelta(minutes=m))
    for m in (40, 45, 50, 55):                                        # 15 dk kesintisiz → kapanır, kesinti 40 dk
        tour(engine, "vpn", True, T0 - timedelta(days=2) + timedelta(minutes=m))
    now = datetime(2026, 9, 29, 5, 0, tzinfo=UTC)                   # 08:00 İstanbul; pencere 22–28 Eylül
    jobs = [{"label": "Planlı raporlar", "lastAt": T0.isoformat(), "lastError": "CRM oturum açma reddedildi"},
            {"label": "Stok gece okuması", "lastAt": T0.isoformat(), "lastError": "Sayı dönüşümü taştı"},
            {"label": "Kampanya özeti", "lastAt": T0.isoformat(), "lastError": "Beklenmeyen boş sonuç"}]
    n = I.weekly_notice(engine, TENANT, now, jobs, LINK)
    assert n.subject == "[Haftalık] Sistem sağlığı · 22–28 Eylül" and n.tone == "bilgi"
    html, text = assert_template(n, "Haftalık", "/timas/sistem-durumu")
    assert "Şirket ağı bağlantısı · 1 · 40 dk" in text and "toplam kesinti 40 dk" in n.headline
    # «Ne yapmalı» yalnız BT'nin işi; uygulama hataları «Bilgi için» altında sayıyla.
    assert "kök neden" not in " ".join(n.actions) and "Planlı raporlar" in text
    assert "Stok gece okuması" not in text and "Kampanya özeti" not in text
    assert n.info == ["2 zamanlanmış iş uygulama kaynaklı hata verdi; uygulama ekibi ilgileniyor, sizden bir işlem beklenmiyor."]
    assert ">Bilgi için: uygulama ekibi ilgileniyor</p>" in html and "BİLGİ İÇİN: UYGULAMA EKİBİ İLGİLENİYOR" in text
    quiet = I.weekly_notice(engine, TENANT, now + timedelta(days=14), [], LINK)
    assert "kesintisiz" in quiet.headline


# ------------------------------------------------------------------ 3) güvenlik


def test_security_alert_and_digest_notices():
    a = {"id": 1, "rule": "hatali_giris", "username": "ahmet", "severity": "kritik",
         "summary": "«ahmet» hesabına 7 hatalı giriş denendi; kaynak adres: 10.0.4.27."}
    n = D.alert_notice([a], IB.portal_link(LINK, "veri-guvenligi"), T0)
    assert n.subject == "[Güvenlik] Art arda hatalı giriş · «ahmet» hesabı · 29 Eylül 09:42" and n.tone == "guvenlik"
    assert_template(n, "Güvenlik", "/timas/veri-guvenligi")
    two = D.alert_notice([a, dict(a, id=2, rule="yeni_yonetici", username="ayse", summary="«ayse» yönetici oldu.")],
                         IB.portal_link(LINK, "veri-guvenligi"), T0)
    assert two.subject.startswith("[Güvenlik] 2 yeni güvenlik uyarısı") and len(two.impact) == 2
    d = D.digest_notice_from({"ayse": {"forbidden": 4, "not_permitted": 1}}, T0, IB.portal_link(LINK, "veri-guvenligi"))
    assert d.subject == "[Günlük] Portal erişim özeti · 1 kişi · 29 Eylül"
    assert_template(d, "Günlük", "/timas/veri-guvenligi")
    assert D.digest_notice_from({}, T0) is None                    # kayıt yoksa e-posta yok


# ------------------------------------------------------------------ 4) kurumsal e-posta kutusu


def test_mailbox_connection_notices():
    since, now = T0 - timedelta(minutes=30), T0 + timedelta(minutes=3)
    n = M.connection_down_notice("timas@timas.com.tr", since, now, "Gmail API: invalid_grant", LINK)
    assert n.subject == "[Kesinti] Kurumsal e-posta kutusu 09:12'den beri okunamıyor"
    html, text = assert_template(n, "Kesinti", "/timas/kurumsal-eposta")
    assert "timas@timas.com.tr kutusu" in text
    back = T0 + timedelta(minutes=95)
    f = M.connection_fixed_notice("timas@timas.com.tr", since, back, back + timedelta(minutes=16), LINK,
                                  steady_min=16, need_min=15, listed_at=back + timedelta(minutes=15))
    assert f.subject == "[Düzeldi] Kurumsal e-posta kutusu 11:17'de yeniden okundu · 2 sa 5 dk sürdü"
    _, ftext = assert_template(f, "Düzeldi", "/timas/kurumsal-eposta")
    assert "Kesintisiz çalışma: Sağlandı: 11:17'den bu yana 16 dk" in ftext
    assert "Veri okuması: Sağlandı: kutudaki ileti listesi okundu" in ftext
    assert "Kutuyu kullanan ayrı bir zamanlanmış iş yok; düzelme ilk iki şartla doğrulandı" in ftext


def test_mailbox_connection_resolves_only_when_really_fixed():
    d = M.connection_decision
    sent = {"eposta": "sent", "since": (T0 - timedelta(minutes=30)).isoformat()}
    back = T0
    # 30 dk okunamadı → bir kez uyarı; gönderilmişse ikinci uyarı yok
    assert d(T0, T0 - timedelta(minutes=30), None, None, None, 30, 15) == "down"
    assert d(T0, T0 - timedelta(minutes=45), sent, None, None, 30, 15) is None
    assert d(T0, T0 - timedelta(minutes=45), {"eposta": "failed"}, None, None, 30, 15) == "down"
    # okuma geldi ama 15 dk dolmadı ya da liste okunmadı → olay sürer
    assert d(back + timedelta(minutes=10), back + timedelta(minutes=10), sent, back, back + timedelta(minutes=10), 30, 15) == "wait"
    assert d(back + timedelta(minutes=20), back + timedelta(minutes=20), sent, back, None, 30, 15) == "wait"
    assert d(back + timedelta(minutes=20), back + timedelta(minutes=20), sent, back, back + timedelta(minutes=20), 30, 15) == "resolved"
    # uyarı hiç gitmediyse düzelme de sessiz
    assert d(back + timedelta(minutes=20), back + timedelta(minutes=20), None, back, back, 30, 15) is None


# ------------------------------------------------------------------ 5) departmansız CRM kullanıcıları


class FakeDirectory:
    def _crm_prefix(self):
        return "CRM.dbo."

    def _crm_rows(self, sql):
        if "SystemUserBase" in sql:
            return [{"Id": "A", "FullName": "Ali Kaya", "DomainName": "TIMAS\\ali", "InternalEMailAddress": "ali@timas.com.tr",
                     "CreatedOn": datetime(2025, 1, 2), "BirimAdi": "Timaş"},
                    {"Id": "B", "FullName": "Eski Hesap", "DomainName": "TIMAS\\eski", "InternalEMailAddress": None,
                     "CreatedOn": None, "BirimAdi": "Timaş"}]
        return [{"Id": "a", "Name": "07-Temel Rol"}]

    def list_people(self):
        return [{"subject": "ali", "label": "Ali Kaya", "hint": "ali", "detail": "Satis"}]


def test_crm_unassigned_mail_carries_only_needed_fields():
    from openpyxl import load_workbook

    sent = []
    out = CU.send(FakeDirectory(), lambda k, d="": {"CRM_UNASSIGNED_TO": "bilgiislem@timas.com.tr", "ALERT_LINK": LINK}.get(k, d),
                  lambda n, to, att: sent.append((n, to, att)) or "sent", now=datetime(2026, 9, 29, 7, 0))
    assert out["status"] == "sent"
    n, to, att = sent[0]
    assert n.subject == "[Bilgi] CRM'de departmanı olmayan 2 kullanıcı · 29 Eylül 07:00" and n.tone == "bilgi"
    html, text = assert_template(n, "Bilgi", "/timas/yonetim")
    assert "ali@timas.com.tr" not in html + text and "Ali Kaya" not in text   # gövdede kişi yok; birim sayıları var
    assert "1 kişinin şirket dizininde etkin hesabı yok" in text
    wb = load_workbook(io.BytesIO(att[0][1]))
    ws = wb["Kişiler"]
    assert [c.value for c in ws[1]] == ["#", "Ad Soyad", "Kullanıcı adı", "CRM'de oluşturulma"]
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "ali@timas.com.tr" not in cells and "07-Temel Rol" not in cells and "ali" in cells
    # Yetkiler ekranındaki tam liste değişmedi.
    full = load_workbook(io.BytesIO(CU.xlsx(CU.build(FakeDirectory()))))["Kişiler"]
    assert full.max_column == 9


def test_all_five_recipient_settings_exist():
    from semantic_bridge import admin as admin_mod

    keys = {s["key"] for s in admin_mod.SPEC}
    assert {"ITOPS_RECIPIENTS", "ITOPS_WEEKLY_TO", "SECURITY_ALERT_RECIPIENTS", "MAIL_CONNECTION_ALERT_TO",
            "CRM_UNASSIGNED_TO", "ITOPS_OUTAGE_MIN", "ITOPS_RESTART_GRACE_MIN", "ITOPS_REMIND_HOURS",
            "ITOPS_RESOLVE_MIN"} <= keys
    st = I.settings(lambda k, d="": d)
    assert (st["outageMin"], st["restartGraceMin"], st["remindHours"], st["resolveMin"]) == (10, 5, 0, 15)
