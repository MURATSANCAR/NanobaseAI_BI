"""CSV indiren her uca Excel eşi (`semantic_bridge/csv_excel.py`).

Denetlenen: (1) CSV okuma (ayraç, BOM, «sep=» satırı, tırnaklı çok satırlı hücre); (2) hücre türü yalnız tartışmasız
olduğunda sayı/tarih — kod/ISBN/telefon/baştaki sıfır metin, Türkçe ve nokta ondalık kolon kolon, «=» formül değil;
(3) Excel'in satırları CSV'nin satırlarıyla birebir (sayı, sıra, metin); (4) ara katman yalnız `bicim=xlsx` + 200 +
text/csv cevabı çevirir, başlıkları (satır sayısı vb.) korur, dosya adını .xlsx yapar, GET ve POST'ta çalışır;
(5) tarayıcıda üretilen CSV'nin çeviri ucu; (6) gerçek köprüde en dış katman, sayfa kapısı ve yetki aynen işler.
Kabul (gerçek veriyle, her uç): `scripts/acceptance/excel-indirme/kabul.py`.
"""
from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from fastapi import FastAPI
from fastapi.responses import JSONResponse, Response
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from semantic_bridge import csv_excel as X

BOM = "﻿"


def _sheet(data: bytes):
    wb = load_workbook(io.BytesIO(data))
    return wb.active


def _rows(ws) -> list[list]:
    return [[c.value for c in r] for r in ws.iter_rows()]


# ------------------------------------------------------------------ okuma


def test_parse_semicolon_bom_and_quoted_newline():
    text = BOM + 'Ad;Not\r\n"Ali; Veli";"iki\nsatır"\r\nAyşe;\r\n'
    assert X.parse(text) == [["Ad", "Not"], ["Ali; Veli", "iki\nsatır"], ["Ayşe", ""]]


def test_parse_comma_tab_and_sep_line():
    assert X.parse("a,b\n1,2\n") == [["a", "b"], ["1", "2"]]
    assert X.parse("a\tb\n1\t2\n") == [["a", "b"], ["1", "2"]]
    assert X.parse("sep=;\na;b\n1;2\n") == [["a", "b"], ["1", "2"]]


def test_decode_cp1254_fallback():
    assert X.decode("Şube;Tutar\n".encode("cp1254")).startswith("Şube")


# ------------------------------------------------------------------ hücre türü


def _convert(csv_text: str) -> list[list]:
    data, _ = X.to_xlsx(csv_text, "t")
    return _rows(_sheet(data))


def test_turkish_and_dot_decimals_column_by_column():
    rows = _convert("Tutar;Oran;Adet\n1.234,56;0.1234;12\n-7,5;0.5;-3\n")
    assert rows[1] == [1234.56, 0.1234, 12]
    assert rows[2] == [-7.5, 0.5, -3]


def test_grouped_value_follows_file_convention_or_stays_text():
    # Dosya Türkçe ondalık kullanıyor: «12.500» binlik ayraçlı 12500.
    assert _convert("Tutar;Adet\n1,5;12.500\n")[1] == [1.5, 12500]
    # Dosya nokta ondalık kullanıyor: «4.333» ondalık 4,333 (Python çıktısında binlik ayraç olmaz).
    assert _convert("Oran;Ortalama\n0.5;4.333\n1.25;3.667\n")[1:] == [[0.5, 4.333], [1.25, 3.667]]
    # Dosyada ipucu yok: «12.500» iki türlü okunur → metin kalır (CSV'de nasıl görünüyorsa).
    assert _convert("Adet\n12.500\n")[1] == ["12.500"]


def test_codes_isbn_phone_and_leading_zero_stay_text():
    rows = _convert("Cari kodu;ISBN;Cep;Posta;okur_no;Değer\n120.01.001;9786050812345;5321234567;01234;42;0012\n")
    assert rows[1] == ["120.01.001", "9786050812345", "5321234567", "01234", "42", "0012"]


def test_long_integers_stay_text_even_without_code_header():
    assert _convert("Değer\n9786050812345\n42\n")[1:] == [["9786050812345"], [42]]


def test_dates_datetimes_and_percent():
    data, _ = X.to_xlsx("Tarih;Zaman;TR tarih;Pay\n2026-09-29;2026-09-29T14:30:05+03:00;29.09.2026;%12,5\n", "t")
    ws = _sheet(data)
    r = [c.value for c in ws[2]]
    assert r[0] == datetime(2026, 9, 29) and ws["A2"].number_format == "dd.mm.yyyy"
    assert r[1] == datetime(2026, 9, 29, 14, 30, 5)      # yazılan saat; dilim çevrilmez
    assert r[2] == datetime(2026, 9, 29)
    assert r[3] == pytest.approx(0.125) and ws["D2"].number_format == "0.0%"


def test_formula_like_text_is_not_a_formula():
    data, _ = X.to_xlsx("Not\n=HYPERLINK(\"x\")\n+90 532\n@ali\n", "t")
    ws = _sheet(data)
    assert [ws.cell(row=i, column=1).data_type for i in (2, 3, 4)] == ["s", "s", "s"]
    assert ws["A2"].value == '=HYPERLINK("x")'


def test_mixed_decimal_styles_keep_ambiguous_values_as_text():
    rows = _convert("Değer\n1,5\n2.25\n7\n")
    assert rows[1:] == [["1,5"], ["2.25"], [7]]


def test_control_characters_long_cells_and_row_limit(monkeypatch):
    # Excel'in kabul etmediği kontrol karakteri atılır (dosya düşmez); sekme ve satır sonu kalır.
    data, _ = X.to_xlsx('Not\n"hata\x01 metni\tsekme\nsatır"\n', "t")
    assert _rows(_sheet(data))[1] == ["hata metni\tsekme\nsatır"]
    # 32.767 karakteri aşan hücre «…» ile biter.
    data, _ = X.to_xlsx("Not\n" + "a" * 40000 + "\n", "t")
    v = _rows(_sheet(data))[1][0]
    assert len(v) == X.MAX_CELL and v.endswith("…")
    # Sayfa sınırını aşan liste sessizce kesilmez.
    monkeypatch.setattr(X, "MAX_ROWS", 3)
    with pytest.raises(X.TooManyRows):
        X.to_xlsx("A\n1\n2\n3\n", "t")


def test_styles_and_escaping():
    data, _ = X.to_xlsx('Ad & <Soyad>;Tutar;Tarih;Pay\n"Ali ""Veli"" & <b>";1.234,5;2026-09-29 08:05;%12\n', "a'b")
    ws = _sheet(data)
    assert ws.title == "a'b" and ws["A1"].font.b and ws["A1"].value == "Ad & <Soyad>"
    assert ws["A2"].value == 'Ali "Veli" & <b>'
    assert (ws["B2"].value, ws["B2"].number_format) == (1234.5, "#,##0.0")
    assert (ws["C2"].value, ws["C2"].number_format) == (datetime(2026, 9, 29, 8, 5), "dd.mm.yyyy hh:mm")
    assert (ws["D2"].value, ws["D2"].number_format) == (pytest.approx(0.12), "0%")


def test_every_row_and_value_survives():
    lines = ["Ad;Tutar;Tarih;Kod"] + [f"Kişi {i};{i},{i % 10}0;2026-01-{(i % 28) + 1:02d};K{i:04d}" for i in range(2500)]
    data, n = X.to_xlsx(BOM + "\r\n".join(lines), "t")
    assert n == 2500
    ws = _sheet(data)
    assert ws.max_row == 2501 and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:D2501"
    assert [c.value for c in ws[1]] == ["Ad", "Tutar", "Tarih", "Kod"]
    last = [c.value for c in ws[2501]]
    assert last == ["Kişi 2499", pytest.approx(2499.9), datetime(2026, 1, (2499 % 28) + 1), "K2499"]


def test_empty_and_header_only():
    data, n = X.to_xlsx(BOM + "A;B\r\n", "t")
    assert n == 0 and _rows(_sheet(data)) == [["A", "B"]]
    data, n = X.to_xlsx("", "t")
    assert n == 0


def test_sheet_title_and_names():
    assert X.sheet_title("a/b:c*?.csv") == "a b c"
    assert len(X.sheet_title("x" * 50)) == 31
    assert X.xlsx_name('attachment; filename="rapor.csv"', "") == "rapor.xlsx"
    assert X.xlsx_name("attachment; filename*=UTF-8''terim-bankas%C4%B1.csv", "") == "terim-bankası.xlsx"
    assert X.xlsx_name("", "export.csv") == "export.xlsx"
    d = X.disposition("terim-bankası.xlsx")
    assert 'filename="terim-bankasi.xlsx"' in d and "filename*=UTF-8''terim-bankas%C4%B1.xlsx" in d


# ------------------------------------------------------------------ ara katman


@pytest.fixture
def client():
    app = FastAPI()

    @app.get("/x/export.csv")
    def csv_get() -> Response:
        return Response((BOM + "Ad;Tutar\r\nAli;1.234,50\r\n").encode("utf-8"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": 'attachment; filename="liste.csv"', "X-Readers-Count": "1"})

    @app.post("/x/export")
    def csv_post() -> Response:
        return Response(b"a;b\n1;2\n", media_type="text/csv", headers={"X-Readers-Count": "7"})

    @app.get("/x/json")
    def js() -> dict:
        return {"ok": True}

    @app.get("/x/denied.csv")
    def denied() -> JSONResponse:
        return JSONResponse({"detail": "yok"}, status_code=403)

    X.register(app)
    return TestClient(app)


def test_middleware_converts_csv_only_on_request(client):
    plain = client.get("/x/export.csv")
    assert plain.headers["content-type"].startswith("text/csv") and plain.content.startswith(BOM.encode())
    r = client.get("/x/export.csv?bicim=xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == X.XLSX_MEDIA
    assert 'filename="liste.xlsx"' in r.headers["content-disposition"]
    assert r.headers["x-readers-count"] == "1" and int(r.headers["content-length"]) == len(r.content)
    assert _rows(_sheet(r.content)) == [["Ad", "Tutar"], ["Ali", 1234.5]]


def test_middleware_post_and_fallback_name(client):
    r = client.post("/x/export?bicim=xlsx", json={})
    assert r.headers["content-type"] == X.XLSX_MEDIA and r.headers["x-readers-count"] == "7"
    assert 'filename="export.xlsx"' in r.headers["content-disposition"]
    assert _rows(_sheet(r.content)) == [["a", "b"], [1, 2]]


def test_middleware_leaves_errors_and_non_csv_alone(client):
    assert client.get("/x/json?bicim=xlsx").json() == {"ok": True}
    r = client.get("/x/denied.csv?bicim=xlsx")
    assert r.status_code == 403 and r.json() == {"detail": "yok"}
    assert client.get("/x/export.csv?bicim=csv").headers["content-type"].startswith("text/csv")


def test_client_side_csv_endpoint(client):
    r = client.post("/api/v1/export/xlsx", json={"csv": BOM + "Ad Soyad;Dahili\r\nAli;101\r\n", "filename": "dahili-rehber.csv"})
    assert r.status_code == 200 and r.headers["content-type"] == X.XLSX_MEDIA
    assert 'filename="dahili-rehber.xlsx"' in r.headers["content-disposition"]
    assert _rows(_sheet(r.content)) == [["Ad Soyad", "Dahili"], ["Ali", 101]]
    assert client.post("/api/v1/export/xlsx", json={"csv": "  "}).status_code == 400


# ------------------------------------------------------------------ gerçek köprü


def test_real_bridge_outermost_and_gate_still_applies(monkeypatch, store, settings):
    from semantic_bridge import access as A
    from semantic_layer.tests.test_access import _app

    app, client = _app(monkeypatch, store, settings)
    assert app.user_middleware[0].cls is X.CsvToXlsx          # en dış katman: kapının ve önbelleğin dışında
    assert A.rule_for("/api/v1/export/xlsx") == A.OPEN
    csv_routes = sorted(r.path for r in app.routes if getattr(r, "path", "").endswith(".csv"))
    assert len(csv_routes) >= 28
    missing = [p for p in csv_routes if A.rule_for(p) is None]
    assert missing == []
    admin = {"cookie": "timas_session=z", "x-semantic-caller": "test"}
    r = client.get("/api/v1/admin/prompts/export.csv?bicim=xlsx", headers=admin)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == X.XLSX_MEDIA
    assert 'filename="promtlar.xlsx"' in r.headers["content-disposition"]
    assert _rows(_sheet(r.content))[0][:3] == ["createdAt", "username", "answerType"]
    # Yönetici olmayan: uç Excel'e çevrilmeden hata döner (JSON), dosya inmez.
    r = client.get("/api/v1/admin/prompts/export.csv?bicim=xlsx", headers={"cookie": "timas_session=a"})
    assert r.status_code in (401, 403) and "json" in r.headers["content-type"]
