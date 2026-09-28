"""M44 Lojistik ve kargo: kimlik bilgisi kolonu yasağı (SQL dosyaları ve üretilen her SQL), metin sayı/tarih çevrimi (Kural C19),
kargo kaydı dizini, firma karnesi (desi başı = Σ tutar ÷ Σ desi), teslim bekleyen yaş kovaları, entegrasyon hatası süzgeci,
mutabakat (mükerrer takip no, Logo cari eşlemesi, Logo sevk ↔ CRM sevkiyat), Zeki AI sınıflama ve taslakta olgu dışı sayı
yasağı, iş eşikleri, karar kaydı, yetki kuralları ve uçlar.

Veriler yapaydır ve yalnız kuralları sınar; gerçek CRM/Logo kabulü test sunucusunda (scripts/acceptance/M44).
"""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta

import pytest

from semantic_bridge import access as A
from semantic_bridge import shipping as S
from semantic_bridge import shipping_sources as src
from semantic_layer.store.catalog_store import open_store

TN = "t1"


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    S._ready.discard(id(e))
    S.ensure(e)
    return e


def cfg(**over):
    base = S.settings_from(lambda key, default="": default)
    base.update(over)
    return base


def cargo(i, *, firma="ARAS KARGO", irs="01.09.2026", teslim="", tutar="100,50", desi="10", sevk="2", sehir="İstanbul",
          iade="", takip=None, sube="Esenyurt", cod=""):
    return {"id": f"00000000-0000-0000-0000-{i:012d}", "takip_no": takip or f"T{i}", "musteri_irs_no": f"I{i}", "kargo_irs_no": None,
            "firma": firma, "cikis_sube": sube, "varis_sube": "Merkez", "sehir": sehir, "irs_tarihi": irs, "teslim_tarihi": teslim,
            "teslim_saati": None, "iade_durumu": iade, "tahsilatli": cod, "sevk_adeti": sevk, "desi": desi, "agirlik": "3,2",
            "tutar": tutar, "kanal": "B2B", "olusturma": datetime(2026, 9, 1, 8, 0)}


# ------------------------------------------------------------------ kimlik bilgisi yasağı


def test_sql_files_never_touch_credential_columns_or_select_star():
    files = src.all_sql_files()
    assert len(files) >= 10
    for f in files:
        text = f.read_text(encoding="utf-8").lower()
        for col in src.FORBIDDEN_COLUMNS:
            assert not re.search(rf"(?<![a-z0-9_]){col}(?![a-z0-9_])", text), f"{f.name}: {col}"
        assert not re.search(r"select\s+(top\s+\d+\s+)?(\w+\.)?\*", text), f.name


def test_guard_blocks_forbidden_columns_anywhere():
    for bad in ("SELECT f.new_sifre FROM x f", "SELECT a FROM x WHERE new_Token IS NOT NULL", "SELECT a FROM x ORDER BY new_clientsecret",
                "SELECT * FROM x", "SELECT f.* FROM x f"):
        with pytest.raises(src.SourceError):
            src.guard(bad)
    assert src.guard("SELECT COUNT(*) AS n FROM x") == "SELECT COUNT(*) AS n FROM x"


def test_every_builder_passes_guard_and_carrier_query_has_three_columns():
    sch = "Timas_MSCRM.dbo"
    carrier = src.sql("crm_kargo_firma", p=src.prefix(sch))
    head = carrier.split("FROM")[0].lower()
    assert head.count(" as ") == 3 and "new_name" in head and "new_kargokodu" in head
    built = [
        src.order_sql(sch, src.errors_where(date(2026, 9, 1))),
        src.order_sql(sch, src.untracked_where(date(2026, 9, 1), (100000000,), (16,))),
        src.order_sql(sch, src.boxed_where()),
        src.order_sql(sch, src.search_where("SP-12'3_x").replace("{p}", src.prefix(sch))),
        src.sql("crm_sevk_sayisi", p=src.prefix(sch), durumlar="100000000", bas="2026-09-01 00:00:00", bugun_utc="2026-09-28 00:00:00"),
        src.sql("crm_kargo_bilgisi", p=src.prefix(sch), kisisel="", kosul=""),
        src.sql("crm_kargo_bilgisi", p=src.prefix(sch), kisisel=src.PERSONAL_COLUMNS, kosul=" AND b.new_kargobilgisiId IN ('x')"),
        src.sql("logo_sevk", f="411", bas="2026-08-01", bit="2026-09-01"),
        src.sql("logo_kargo_fatura", f="411", cariler="'320.01'", bas="2026-08-01", bit="2026-09-01"),
    ]
    for s in built:
        assert src.guard(s) == s
        assert "--" not in s                    # açıklama satırları atıldı
    # Arama değeri kaçışlanır: tırnak ikilenir, LIKE'ta alt çizgi ve yüzde harf olarak aranır, eşitlikte olduğu gibi kalır.
    w = src.search_where("D'Or_12")
    assert "= N'D''Or_12'" in w and "LIKE N'%D''Or[_]12%'" in w
    with pytest.raises(src.SourceError):
        src.search_where("100%")
    with pytest.raises(src.SourceError):
        src.search_where("x")
    with pytest.raises(src.SourceError):
        src.search_where("a{p}b")


def test_comment_placeholders_do_not_become_code():
    text = src.sql_text("crm_siparis_asama")
    assert not any(line.lstrip().startswith("--") for line in text.splitlines())
    assert text.count("{kosul}") == 1


# ------------------------------------------------------------------ metin sayı ve tarih (C19)


def test_cargo_number_matches_try_cast_replace():
    assert src.cargo_number("12,5") == 12.5
    assert src.cargo_number("12.5") == 12.5
    assert src.cargo_number("1.234,56") is None           # SQL'de de TRY_CAST('1.234.56') NULL
    assert src.cargo_number("") is None and src.cargo_number(None) is None and src.cargo_number("abc") is None


def test_cargo_date_formats_and_empty_sentinel():
    fmts = cfg()["dateFormats"]
    assert src.cargo_date("05.09.2026", fmts) == date(2026, 9, 5)
    assert src.cargo_date("2026-09-05 14:30:00", fmts) == date(2026, 9, 5)
    assert src.cargo_date("05/09/2026", fmts) == date(2026, 9, 5)
    assert src.cargo_date("01.01.1900", fmts) is None
    assert src.cargo_date("dün", fmts) is None


def test_crm_utc_day_boundary():
    # CRM UTC saklar: 27.09 21:30 UTC = 28.09 İstanbul.
    assert src.crm_day(datetime(2026, 9, 27, 21, 30)) == date(2026, 9, 28)
    assert src.utc_bound(date(2026, 9, 28)) == "2026-09-27 21:00:00"


# ------------------------------------------------------------------ kargo kaydı dizini ve karne


def test_index_parses_counts_unreadable_and_unifies_firm_names():
    rows = [cargo(1), cargo(2, firma="Aras Kargo "), cargo(3, firma="aras  kargo", tutar="1.234,5"), cargo(4, irs="?")]
    idx = S.CargoIndex(rows, cfg())
    assert len({r["firmaKey"] for r in idx.rows}) == 1
    assert idx.unreadable["tutar"] == 1 and idx.unreadable["irsTarihi"] == 1 and idx.undated == 1
    assert idx.data_end == date(2026, 9, 1)
    assert idx.by_tracking["T1"][0]["id"].endswith("1")


def test_scorecard_cost_per_desi_is_sum_over_sum_and_hidden_without_right():
    rows = [cargo(1, tutar="100", desi="10", teslim="03.09.2026"), cargo(2, tutar="50", desi="", teslim="05.09.2026"),
            cargo(3, firma="MNG KARGO", tutar="30", desi="3", iade="İade edildi"),
            cargo(4, irs="01.08.2026")]                                      # dönem dışı
    idx = S.CargoIndex(rows, cfg())
    card = S.scorecard(idx, date(2026, 9, 1), date(2026, 9, 2), cost=True)
    aras = next(i for i in card["items"] if i["firma"] == "ARAS KARGO")
    assert aras["gonderi"] == 2 and aras["tutar"] == 150 and aras["desi"] == 10 and aras["desiBasi"] == 15.0
    assert aras["sevkBasi"] == 37.5 and aras["ortancaGun"] == 3.0 and aras["teslim"] == 2
    mng = next(i for i in card["items"] if i["firma"] == "MNG KARGO")
    assert mng["iade"] == 1 and mng["iadeOrani"] == 1.0 and mng["bekleyen"] == 0
    hidden = S.scorecard(idx, date(2026, 9, 1), date(2026, 9, 2), cost=False)
    assert all("tutar" not in i and "desiBasi" not in i for i in hidden["items"]) and "tutar" not in hidden["toplam"]
    with pytest.raises(S.ShippingError):
        S.scorecard(idx, date(2026, 9, 1), date(2026, 9, 2), group="kanal")


def test_scorecard_city_target_only_where_given():
    rows = [cargo(1, teslim="02.09.2026"), cargo(2, teslim="06.09.2026"), cargo(3, sehir="Van", teslim="09.09.2026")]
    idx = S.CargoIndex(rows, cfg())
    card = S.scorecard(idx, date(2026, 9, 1), date(2026, 9, 2), group="sehir", targets={"istanbul": 2})
    ist = next(i for i in card["items"] if i["sehir"] == "İstanbul")
    van = next(i for i in card["items"] if i["sehir"] == "Van")
    assert ist["hedefGun"] == 2 and ist["hedefiAsan"] == 1 and ist["hedefiAsanOrani"] == 0.5
    assert "hedefGun" not in van                                            # hedef yoksa oran uydurulmaz


def test_waiting_buckets_threshold_and_returns_excluded():
    asof = date(2026, 9, 28)
    rows = [cargo(1, irs="27.09.2026"), cargo(2, irs="20.09.2026"), cargo(3, irs="01.08.2026"), cargo(4, irs="10.09.2026", teslim="12.09.2026"),
            cargo(5, irs="01.09.2026", iade="İade"), cargo(6, irs="")]
    idx = S.CargoIndex(rows, cfg())
    w = S.waiting(idx, asof, 5)
    assert w["toplam"] == 4 and w["esikUstu"] == 2
    buckets = {b["kova"]: b["adet"] for b in w["kovalar"]}
    assert buckets["0–2 gün"] == 1 and buckets["6–14 gün"] == 1 and buckets["30+ gün"] == 1 and buckets["Tarihsiz"] == 1
    assert [i["yas"] for i in w["items"]] == [58, 8]
    assert "tutar" not in w["items"][0]


def test_freshness_flags_stale_cargo_data():
    idx = S.CargoIndex([cargo(1, irs="01.08.2026")], cfg())
    f = idx.freshness(cfg(), date(2026, 9, 28))
    assert f["eski"] and "güncellenmemiş" in f["not"]


# ------------------------------------------------------------------ entegrasyon hatası


def test_integration_failures_skip_success_values():
    ok = cfg()["okValues"]
    row = {"aras_sonuc": "Başarılı", "aras_mesaj": None, "ups_sonuc": "HATA", "ups_mesaj": "Adres bulunamadı",
           "mng_sonuc": None, "mng_mesaj": "", "akademi_sonuc": None, "akademi_mesaj": "Servis zaman aşımı"}
    f = S.integration_failures(row, ok)
    assert [x["entegrasyon"] for x in f] == ["UPS", "Akademi"]
    assert f[0]["mesajHash"] == S.message_hash("UPS", "Adres bulunamadı")


def test_sanitize_removes_phone_and_email():
    t = S.sanitize_for_model("Alıcı 0532 123 45 67 ahmet@ornek.com adresi hatalı, takip 1234567890")
    assert "0532" not in t and "ornek" not in t and "1234567890" not in t


# ------------------------------------------------------------------ mutabakat


def test_reconcile_duplicates_mapping_and_shipment_match():
    rows = [cargo(1, takip="A1", tutar="100"), cargo(2, takip="A1", tutar="100"), cargo(3, takip="A2", tutar="50"),
            cargo(4, firma="YURTİÇİ", takip="Y1", tutar="20")]
    idx = S.CargoIndex(rows, cfg())
    codes = S.parse_carrier_codes("aras kargo=320.01.001, 320.01.002")
    inv = [{"cari": "320.01.001", "kdv_dahil": 300.0, "kdv": 50.0}, {"cari": "999", "kdv_dahil": 1.0, "kdv": 0}]
    rec = S.reconcile(idx, date(2026, 9, 1), date(2026, 10, 1), carrier_codes=codes, invoices=inv,
                      logo_shipments=[{"irsaliye": 1, "fatura_no": "F1", "satir": 3, "adet": 10}, {"irsaliye": 2, "fatura_no": None, "satir": 1, "adet": 1}],
                      crm_shipments=[{"fatura_no": "F1"}, {"fatura_no": "F2"}, {"fatura_no": None}], notes=[])
    aras = next(i for i in rec["items"] if i["firma"] == "ARAS KARGO")
    assert aras["crmTutar"] == 250 and aras["mukerrer"] == 1 and aras["mukerrerTutar"] == 100
    assert aras["logoKdvHaric"] == 250 and aras["farkKdvHaric"] == 0 and aras["logoKdvDahil"] == 300
    yk = next(i for i in rec["items"] if i["firma"] == "YURTİÇİ")
    assert not yk["logoEslendi"] and yk["farkKdvHaric"] is None
    assert rec["sevk"]["eslesen"] == 1 and rec["sevk"]["crmSevkiyat"] == 3 and rec["sevk"]["logoFaturali"] == 1
    assert "F2" in rec["sevk"]["eslesmeyen"] and rec["toplam"]["eslenmeyenFirma"] == 1


def test_month_bounds_and_period():
    assert S.month_bounds("2026-12") == (date(2026, 12, 1), date(2027, 1, 1))
    with pytest.raises(S.ShippingError):
        S.month_bounds("2026-13")
    assert S.period("", "", date(2026, 9, 10)) == (date(2026, 8, 12), date(2026, 9, 11))
    with pytest.raises(S.ShippingError):
        S.period("2026-09-10", "2026-09-01", date(2026, 9, 10))


# ------------------------------------------------------------------ Zeki AI


class FakeChoice:
    def __init__(self, choice, p, margin):
        self.choice, self.probability, self.margin, self.method = choice, p, margin, "logprobs"

    def confident(self, min_prob, min_margin=0.0, min_coverage=0.0):
        return self.probability is not None and self.probability >= min_prob and self.margin >= min_margin


def test_classify_once_per_message_and_unsure_below_threshold(engine):
    errs = [{"id": "o1", "hatalar": [{"entegrasyon": "UPS", "mesaj": "Adres bulunamadı 05321234567", "sonuc": "HATA",
                                       "mesajHash": S.message_hash("UPS", "Adres bulunamadı 05321234567")}]},
            {"id": "o2", "hatalar": [{"entegrasyon": "MNG", "mesaj": "?", "sonuc": "Hata", "mesajHash": S.message_hash("MNG", "?")}]}]
    prompts = []

    def choose(prompt, choices):
        prompts.append(prompt)
        return FakeChoice("Adres bilgisi", 0.95, 0.8) if "Adres" in prompt else FakeChoice("Diğer", 0.4, 0.1)

    out = S.classify_pending(engine, TN, errs, choose, cfg())
    assert out["sinif"] == 2 and all("05321234567" not in p for p in prompts)
    cls = S.load_classes(engine, TN)
    assert cls[errs[0]["hatalar"][0]["mesajHash"]]["sinif"] == "Adres bilgisi"
    assert cls[errs[1]["hatalar"][0]["mesajHash"]]["sinif"] == S.UNSURE
    assert S.classify_pending(engine, TN, errs, choose, cfg())["sinif"] == 0          # önbellekte
    assert S.classify_pending(engine, TN, errs, None, cfg())["sinif"] == 0


def _order():
    return {"id": "0f000000-0000-0000-0000-000000000001", "no": "SP-77", "durumAdi": "Sevk edildi", "firma": "ARAS KARGO",
            "takipNo": "T9", "musteri": "Ayşe Yılmaz", "asamalar": {"siparis": "2026-09-20T10:00", "depoda": None, "pusula": None,
                                                                    "kutulandi": "2026-09-21T09:00", "sevk": "2026-09-22T15:00", "tamamlandi": None}}


def test_draft_facts_have_no_personal_data_and_invented_numbers_fall_back():
    facts = S.draft_facts(_order(), [])
    assert "Ayşe" not in str(facts) and facts["siparisNo"] == "SP-77"
    text, src_ = S.draft_text("gecikme", facts, lambda m: "Sayın müşterimiz, siparişiniz 3 gün içinde teslim edilecek.")
    assert src_ == "kural" and "SP-77" in text                                          # «3 gün» olgularda yok
    ok, src2 = S.draft_text("gecikme", facts, lambda m: "Sayın müşterimiz, SP-77 numaralı siparişiniz 22.09.2026'da yola çıktı.")
    assert src2 == "zeki" and "22.09.2026" in ok
    assert S.draft_text("iade", facts, None)[1] == "kural"


def test_drafts_and_decisions_roundtrip(engine):
    d = S.create_draft(engine, TN, "ayse", _order(), "ozur", "metin", "kural")
    assert S.drafts_for(engine, TN, _order()["id"])[0]["id"] == d["id"]
    out, diff = S.update_draft(engine, TN, "ayse", d["id"], {"durum": "kullanildi"})
    assert out["durum"] == "kullanildi" and diff
    with pytest.raises(S.ShippingError):
        S.update_draft(engine, TN, "ayse", d["id"], {"durum": "gonderildi"})
    S.delete_draft(engine, TN, d["id"])
    assert S.drafts_for(engine, TN, _order()["id"]) == []
    with pytest.raises(S.ShippingError):
        S.create_decision(engine, TN, "mudur", {"tur": "kurye", "karar": ""})
    k = S.create_decision(engine, TN, "mudur", {"tur": "bolge", "karar": "Doğu illerinde MNG", "kapsam": {"sehir": "Van", "sifre": "x"}})
    assert k["kapsam"] == {"sehir": "Van"} and S.list_decisions(engine, TN)[0]["id"] == k["id"]
    S.delete_decision(engine, TN, k["id"])
    assert S.list_decisions(engine, TN) == []


def test_ops_settings_defaults_validation_and_meta(engine):
    c = cfg()
    assert S.ops_settings(engine, TN, c)["bekleyenGun"] == 5
    out, diff = S.save_ops_settings(engine, TN, "mudur", {"bekleyenGun": 7, "bolgeHedef": {"Van": 4}}, c)
    assert out["bekleyenGun"] == 7 and out["bolgeHedef"] == {"Van": 4} and "bekleyenGun" in diff
    for bad in ({"bekleyenGun": 0}, {"bolgeHedef": {"Van": "x"}}, {}):
        with pytest.raises(S.ShippingError):
            S.save_ops_settings(engine, TN, "mudur", bad, c)
    S.meta_set(engine, TN, "gunluk", "2026-09-28")
    assert S.meta_get(engine, TN, "gunluk") == "2026-09-28"
    assert "_gunluk" not in S.ops_settings(engine, TN, c)["guncelleyen"]


def test_due_once_per_day_weekday_and_monthday():
    now = datetime(2026, 9, 28, 7, 0, tzinfo=S.TZ)         # pazartesi
    assert S.due(now, "06:45", None) and not S.due(now, "06:45", "2026-09-28")
    assert not S.due(now, "08:00", None, weekday=0) and S.due(now.replace(hour=8, minute=5), "08:00", None, weekday=0)
    assert not S.due(now, "06:45", None, monthday=3)


def test_xlsx_writes_all_rows():
    from openpyxl import load_workbook
    import io

    data = S.xlsx("Liste", [("a", "A"), ("b", "B")], [{"a": i, "b": True} for i in range(250)])
    ws = load_workbook(io.BytesIO(data)).active
    assert ws.max_row == 4 + 250 and ws.cell(5, 2).value == "Evet"


# ------------------------------------------------------------------ yetki ve uçlar


def test_access_rules_and_catalog():
    assert A.rule_for("/api/v1/shipping/overview") == frozenset({"sayfa:kargo"})
    assert A.rule_for("/api/v1/shipping/carriers") == frozenset({"sayfa:kargo-firmalar"})
    assert A.rule_for("/api/v1/shipping/reconcile/summary") == frozenset({"sayfa:kargo-mutabakat"})
    assert A.rule_for("/api/v1/shipping/export/bekleyen.xlsx") == frozenset({"sayfa:kargo", "sayfa:kargo-firmalar", "sayfa:kargo-mutabakat"})
    assert A.rule_for("/api/v1/shipping/run-due") == A.SYSTEM
    f = A.features_for
    assert f("POST", "/api/v1/shipping/drafts") == ["ozellik:kargo.taslak"]
    assert f("PATCH", "/api/v1/shipping/drafts/abc") == ["ozellik:kargo.taslak"]
    assert f("GET", "/api/v1/shipping/export/firmalar.xlsx") == ["ozellik:veri.disa-aktar"]
    assert f("POST", "/api/v1/shipping/decisions") == []                   # açıkça verilen kargo.karar ucun içinde
    assert {"ozellik:kargo.maliyet", "ozellik:kargo.alici", "ozellik:kargo.karar"} <= A.explicit_keys()
    assert "ozellik:kargo.taslak" not in A.explicit_keys()
    assert {"sayfa:kargo", "sayfa:kargo-firmalar", "sayfa:kargo-mutabakat", "ozellik:kargo.taslak"} <= A.all_keys()


def test_api_module_imports_request_at_module_level():
    from semantic_bridge import shipping_api

    assert hasattr(shipping_api, "Request")


def test_api_decision_needs_explicit_right_settings_validate_and_run_due(monkeypatch, store, settings):
    from fastapi.testclient import TestClient

    from semantic_bridge import admin as admin_mod
    from semantic_bridge import board as board_mod
    from semantic_bridge.app import Runtime, create_app
    from semantic_layer.candidates.llm_client import FakeLlm

    users = {"timas_session=a": "ayse", "timas_session=z": "zekiai"}
    monkeypatch.setattr(board_mod, "_fetch_user", lambda cookie: users.get(cookie))
    monkeypatch.setattr(board_mod, "_fetch_session", lambda cookie: {"username": users.get(cookie), "displayName": "x"})
    monkeypatch.setattr(admin_mod, "admins", lambda: ["zekiai"])
    monkeypatch.delenv("SEMANTIC_CALLER_TOKEN", raising=False)
    A._ready.clear()
    A.invalidate()
    client = TestClient(create_app(Runtime(settings, store=store, llm=FakeLlm([""]))))
    a, z = {"cookie": "timas_session=a"}, {"cookie": "timas_session=z"}
    m = client.get("/api/v1/shipping/meta", headers=a).json()
    assert m["me"]["maliyet"] is False and m["me"]["karar"] is False and m["me"]["taslak"] is True
    assert client.get("/api/v1/shipping/meta", headers=z).json()["me"]["maliyet"] is True
    assert client.post("/api/v1/shipping/decisions", json={"tur": "kurye", "karar": "x"}, headers=a).status_code == 403
    r = client.post("/api/v1/shipping/decisions", json={"tur": "kurye", "karar": "Aras ile devam"}, headers=z)
    assert r.status_code == 201
    assert client.delete(f"/api/v1/shipping/decisions/{r.json()['id']}", headers=z).status_code == 200   # test verisi bırakılmaz
    assert client.put("/api/v1/shipping/settings", json={"bekleyenGun": 3}, headers=a).status_code == 403
    assert client.put("/api/v1/shipping/settings", json={"bekleyenGun": 0}, headers=z).status_code == 400
    assert client.get("/api/v1/shipping/reconcile?ay=2026-08", headers=a).status_code == 403
    assert client.post("/api/v1/shipping/run-due?gorev=yok").status_code == 400
    assert client.get("/api/v1/shipping/export/yok.xlsx", headers=z).status_code == 404
    assert client.post("/api/v1/shipping/drafts", json={"tur": "x"}, headers=a).status_code == 400
