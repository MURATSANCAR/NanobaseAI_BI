"""Excel indirme — test sunucusunda kaydı olmadığı için gerçek veriyle sınanamayan 8 CSV ucu (2026-09-29).

Her uç için modülün kendi test kurgusuyla (yardımcıları içe aktarılır ya da aynı yol izlenir) SQLite üzerinde birden
çok satırlı, gerçekçi veri kurulur (Türkçe karakter, ondalıklı tutar, tarih, kod/numara kolonu); gerçek köprü
(`create_app`, `test_access._app`) üzerinden önce CSV, sonra `?bicim=xlsx` alınır ve denetlenir:
  · ikisi de 200; CSV `text/csv`, Excel `application/vnd.openxmlformats-…sheet`;
  · Excel'in dosya adı CSV'nin adıyla aynı gövde + `.xlsx`;
  · dosya openpyxl ile açılır; satır sayısı ≥ 3 (başlık + en az iki veri satırı);
  · başlık birebir ve her hücre CSV'deki değerle aynı — karşılaştırıcı `csv_excel`'in kendi okuyucusu DEĞİL, aşağıdaki
    bağımsız çözümleyici (`scripts/acceptance/excel-indirme/kabul.py` → `compare` ile aynı mantık, buraya kopyalandı).
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, timedelta, timezone
from typing import Any, Optional
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

from semantic_bridge import access as A
from semantic_bridge import csv_excel as X
from semantic_layer.tests.conftest import TENANT

ADMIN = {"cookie": "timas_session=z", "x-semantic-caller": "test"}


# ------------------------------------------------------------------ bağımsız karşılaştırıcı (kabul.py ile aynı mantık)


def _csv_rows(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    first = text.split("\n", 1)[0]
    delim = max((";", ",", "\t"), key=first.count) if any(d in first for d in ";,\t") else ";"
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    while rows and not any(c.strip() for c in rows[-1]):
        rows.pop()
    return rows


def _numbers_of(s: str) -> list[float]:
    t = s.strip().replace(" ", "")
    pct = t.startswith("%") or t.endswith("%")
    t = t.strip("% ")
    out = []
    for cand in (t.replace(".", "").replace(",", "."), t, t.replace(",", "")):
        try:
            v = float(cand)
        except ValueError:
            continue
        out.append(v / 100 if pct else v)
    return out


def _date_parts(s: str) -> Optional[tuple[int, ...]]:
    t = s.strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2})(?::(\d{2}))?)?", t)
    if m:
        return tuple(int(x or 0) for x in m.groups())
    m = re.match(r"^(\d{2})\.(\d{2})\.(\d{4})(?: (\d{2}):(\d{2})(?::(\d{2}))?)?$", t)
    if m:
        d, mo, y, hh, mm, ss = m.groups()
        return (int(y), int(mo), int(d), int(hh or 0), int(mm or 0), int(ss or 0))
    return None


def _same(raw: str, val: Any) -> bool:
    if val is None:
        return raw.strip() == ""
    if isinstance(val, str):
        return val == raw
    if isinstance(val, (datetime, date)):
        p = _date_parts(raw)
        v = (val.year, val.month, val.day, getattr(val, "hour", 0), getattr(val, "minute", 0), getattr(val, "second", 0))
        return p is not None and p == v
    if isinstance(val, bool):
        return False
    if isinstance(val, (int, float)):
        return any(abs(c - float(val)) <= 1e-9 * max(1.0, abs(c)) for c in _numbers_of(raw))
    return False


def compare(csv_bytes: bytes, xlsx_bytes: bytes) -> dict[str, Any]:
    exp = _csv_rows(csv_bytes)
    ws = load_workbook(io.BytesIO(xlsx_bytes)).active
    got = [[c.value for c in r] for r in ws.iter_rows()]
    while got and all(v is None for v in got[-1]):
        got.pop()
    rep: dict[str, Any] = {"csvRows": len(exp), "xlsxRows": len(got), "bad": [], "cells": 0, "num": 0, "date": 0,
                           "text": 0, "empty": 0, "numLookingText": [], "sheet": ws.title}
    if len(exp) != len(got):
        rep["bad"].append(f"satır sayısı CSV {len(exp)} ≠ Excel {len(got)}")
    for i, (er, gr) in enumerate(zip(exp, got)):
        for j in range(max(len(er), len(gr))):
            raw = er[j] if j < len(er) else ""
            val = gr[j] if j < len(gr) else None
            rep["cells"] += 1
            if i == 0:
                if (val or "") != raw:
                    rep["bad"].append(f"başlık {j + 1}: {raw!r} ≠ {val!r}")
                continue
            head = exp[0][j] if j < len(exp[0]) else "?"
            if not _same(raw, val):
                rep["bad"].append(f"satır {i + 1} kolon {j + 1} ({head}): CSV {raw!r} ≠ Excel {val!r}")
            elif val is None:
                rep["empty"] += 1
            elif isinstance(val, (int, float)) and not isinstance(val, bool):
                rep["num"] += 1
            elif isinstance(val, (datetime, date)):
                rep["date"] += 1
            elif isinstance(val, str):
                rep["text"] += 1
                if _numbers_of(raw) and re.fullmatch(r"-?[\d.,]+%?", raw.strip()):
                    rep["numLookingText"].append(f"{head}={raw}")
    return rep


def _filename(disposition: str) -> str:
    """Content-Disposition → dosya adı (filename* öncelikli). csv_excel.xlsx_name'den bağımsız."""
    m = re.search(r"filename\*\s*=\s*UTF-8''([^;]+)", disposition or "", re.I)
    if m:
        return unquote(m.group(1).strip().strip('"'))
    m = re.search(r'filename\s*=\s*"([^"]*)"', disposition or "") or re.search(r"filename\s*=\s*([^;]+)", disposition or "")
    return m.group(1).strip() if m else ""


def _both(client, path: str, label: str, min_rows: int = 3) -> dict[str, Any]:
    rc = client.get(path, headers=ADMIN)
    assert rc.status_code == 200, f"{label} CSV: {rc.status_code} {rc.text[:300]}"
    assert rc.headers["content-type"].startswith("text/csv"), rc.headers["content-type"]
    rx = client.get(path + ("&" if "?" in path else "?") + "bicim=xlsx", headers=ADMIN)
    assert rx.status_code == 200, f"{label} Excel: {rx.status_code} {rx.text[:300]}"
    assert rx.headers["content-type"] == X.XLSX_MEDIA
    csv_name = _filename(rc.headers.get("content-disposition", ""))
    xlsx_name = _filename(rx.headers.get("content-disposition", ""))
    assert csv_name.lower().endswith(".csv"), csv_name
    assert xlsx_name == csv_name[:-4] + ".xlsx", (csv_name, xlsx_name)
    assert rx.content[:2] == b"PK"
    rep = compare(rc.content, rx.content)
    assert rep["bad"] == [], f"{label}: " + " | ".join(rep["bad"][:12])
    assert rep["csvRows"] >= min_rows and rep["xlsxRows"] == rep["csvRows"], rep
    print(f"\nÖZET {label}: dosya={csv_name}→{xlsx_name} satır={rep['csvRows']} (veri {rep['csvRows'] - 1}) "
          f"hücre={rep['cells']} sayı={rep['num']} tarih={rep['date']} metin={rep['text']} boş={rep['empty']} "
          f"sayıGörünümlüMetin={rep['numLookingText']}")
    return rep


@pytest.fixture
def bridge(monkeypatch, store, settings):
    from semantic_layer.tests.test_access import _app

    app, client = _app(monkeypatch, store, settings)
    assert app.user_middleware[0].cls is X.CsvToXlsx
    yield store.engine, client
    A._ready.clear()
    A.invalidate()


def _grant(engine, *keys: str) -> None:
    """Duyarlı İK anahtarları yöneticiye yalnız rolle gelir (hr_core.who_from): zekiai'ye rol bağla."""
    A.ensure(engine, TENANT)
    rid = A.save_role(engine, TENANT, "zekiai", {"name": "İK dışa aktarım testi", "perms": list(keys)})["id"]
    A.add_binding(engine, TENANT, "zekiai", rid, {"type": "user", "subject": "zekiai"})
    A.invalidate()


# ------------------------------------------------------------------ 1. İK bağlılık anketi


def test_hr_engagement_survey_export(bridge):
    from semantic_bridge import hr_core as H
    from semantic_bridge import hr_engagement as E
    from semantic_layer.tests import test_hr_engagement as TE

    engine, client = bridge
    H._ready.discard(engine)
    E._ready.discard(engine)
    E.ensure(engine)
    tn = TE.TN
    ed, _ = H.save_unit(engine, tn, "ik", {"name": "Editörya"})
    cocuk, _ = H.save_unit(engine, tn, "ik", {"name": "Çocuk", "parentId": ed["id"]})
    satis, _ = H.save_unit(engine, tn, "ik", {"name": "Satış"})
    for n in range(1, 14):
        H.save_employee(engine, tn, "ik", {"displayName": f"Çalışan{n} Öztürk", "username": f"u{n}",
                                           "unitId": (cocuk if n <= 7 else satis)["id"]})
    s = TE._survey(engine, minGroup=3)
    likert = ["gurur", "kalma", "anlam", "yonetici_gb", "takdir", "gelisim", "is_yuku", "araclar", "iletisim", "adalet", "guven"]
    enps = [10, 9, 8, 6, 3, 10, 9, 7, 5, 10, 9, 2, 8]
    for n in range(1, 14):
        ans: dict[str, Any] = {"enps": enps[n - 1]}
        for k, key in enumerate(likert):
            ans[key] = 1 + (n * 3 + k * 2) % 5
        if n % 4 == 0:
            ans["acik_iyi"] = "Ekip arkadaşlarım; «güven» ortamı"
        TE._answer(engine, s["id"], f"u{n}", ans)
    E.close_survey(engine, tn, "ikuzman", s["id"])
    _grant(engine, E.F_EXPORT)
    rep = _both(client, f"/api/v1/hr/engagement/surveys/{s['id']}/export.csv", "İK anket")
    assert rep["num"] > 0


# ------------------------------------------------------------------ 2. İK değerlendirme dönemi


def test_hr_performance_cycle_export(bridge):
    from semantic_bridge import hr_core as H
    from semantic_bridge import hr_performance as P
    from semantic_layer.tests import test_hr_performance as TP

    engine, client = bridge
    H._ready.discard(engine)
    P._ready.discard(engine)
    P.ensure(engine)
    tn = TP.TN
    unit, _ = H.save_unit(engine, tn, "ik", {"name": "Editörya — Çocuk Kitapları"})
    other, _ = H.save_unit(engine, tn, "ik", {"name": "Satış"})
    ids: dict[str, str] = {}
    for key, name, user, mgr, u in (("gm", "Genel Müdür", "gm", None, unit), ("mdr", "Müdür Şükrü Çağlar", "mudur", "gm", unit),
                                    ("ayse", "Ayşe Yılmaz", "ayse", "mdr", unit), ("ali", "Ali Kaya", "ali", "mdr", unit),
                                    ("gul", "Gülşen İnce", "gulsen", "mdr", unit),
                                    ("ik", "İK Uzmanı", "ikuzman", None, other)):
        e, _ = H.save_employee(engine, tn, "ik", {"displayName": name, "username": user, "unitId": u["id"],
                                                  "managerId": ids.get(mgr) if mgr else None})
        ids[key] = e["id"]
    H.save_unit(engine, tn, "ik", {"managerEmployeeId": ids["mdr"]}, unit["id"])
    ids["unit"], ids["other"] = unit["id"], other["id"]
    hr = TP.sc(engine, "ikuzman", TP.HR_KEYS)
    cy = TP._cycle(engine, ids, hr)
    people = {p["employeeId"]: p["reviewId"] for p in P.cycle_status(engine, tn, hr, cy["id"])["people"]}
    ayse, mdr, gm = TP.sc(engine, "ayse"), TP.sc(engine, "mudur", TP.ALL), TP.sc(engine, "gm", TP.ALL)
    P.save_review(engine, tn, ayse, people[ids["ayse"]], {"self": {"ratings": {"is_kalitesi": 4}, "notes": {"guclu": "Terminler"}}}, submit="self")
    for who, score in (("ayse", 4), ("ali", 3), ("gul", 5)):
        P.save_review(engine, tn, mdr, people[ids[who]], {"manager": {"overall": score, "ratings": {"zaman": 4}, "notes": {"guclu": "iyi"}}},
                      submit="manager")
    P.save_review(engine, tn, gm, people[ids["mdr"]], {"manager": {"overall": 2, "ratings": {"zaman": 2}, "notes": {"guclu": "x"}}},
                  submit="manager")
    P.review_action(engine, tn, mdr, people[ids["ayse"]], "share", {"meetingAt": "2026-12-20T10:00:00+03:00"})
    P.review_action(engine, tn, ayse, people[ids["ayse"]], "comment", {"comment": "Katılmıyorum", "objection": True})
    _grant(engine, P.F_EXPORT)
    rep = _both(client, f"/api/v1/hr/performance/cycles/{cy['id']}/export.csv", "İK değerlendirme")
    assert rep["csvRows"] == 1 + 5 and rep["num"] >= 4


# ------------------------------------------------------------------ 3. Okur yüklemesi → CRM'e işlenecek yeni kişiler


def test_readers_import_crm_csv(bridge, monkeypatch):
    from semantic_bridge import readers as R
    from semantic_bridge import readers_imports as I
    from semantic_layer.tests import test_readers as TRd

    engine, client = bridge
    monkeypatch.setenv("READERS_HASH_SALT", TRd.SALT)
    for k in ("READERS_EXPORT_ENABLED", "READERS_REQUIRE_KVKK", "READERS_MINOR_EXPORT", "READERS_CRM_FLAG_RULES"):
        monkeypatch.delenv(k, raising=False)
    R._ready.discard(id(engine))
    R.ensure(engine)
    TRd._profiles_for_segments(engine)          # a@x.com CRM'de var → «eşleşti»; ötekiler «yeni»
    data = ("Ad Soyad;E-posta;Cep telefonu;Şehir;Doğum tarihi;KVKK onayı\n"
            "Ayşe Yılmaz;a@x.com;0532 111 22 33;Ankara;1985;evet\n"
            "Şükrü Öztürk;Sukru.Ozturk@Ornek.com.tr;0 (533) 444 55 66;İstanbul;12.03.1990;evet\n"
            "Gülşen Çağlar;gulsen@ornek.com;+90 544 777 88 99;İzmir;;hayır\n"
            "İsmail Ağaoğlu;ismail@ornek.com;;Bursa;2001;evet\n"
            "Özge Işık;;0555 012 34 56;Muğla;1978;\n").encode("utf-8")
    imp = I.create(engine, TENANT, "zekiai", "fuar-katilimcilari.csv", data)
    out = I.confirm(engine, TENANT, imp["id"], "zekiai", {"eventName": "İstanbul Kitap Fuarı", "eventDate": "2026-09-20"})
    assert out["matched"] == 1 and out["new"] == 4
    rep = _both(client, f"/api/v1/readers/imports/{imp['id']}/crm.csv", "Okur CRM")
    assert rep["csvRows"] == 1 + 4 and rep["date"] >= 4


# ------------------------------------------------------------------ 4. Çeviri kalite raporu


def test_translation_quality_csv(bridge, monkeypatch, tmp_path):
    from semantic_bridge import editorial_desk as desk
    from semantic_bridge import editorial_translation as T

    engine, client = bridge
    monkeypatch.setenv("EDITORIAL_DIR", str(tmp_path))
    T._ready.discard(id(engine))
    desk._ready.discard(id(engine))
    T.ensure(engine)
    desk.ensure(engine)
    jid = T.create_job(engine, TENANT, "editor", {"title": "Kırmızı Pazartesi", "sourceLang": "en", "targetLang": "tr",
                                                  "translator": "ayse", "reviewer": "mehmet"})["id"]
    src = (b"CHAPTER ONE\n\nThe road was long. He paid 12 coins.\n\nIt cost 1250.50 dollars in 2019.\n\n"
           b"CHAPTER TWO\n\nNobody came to the Grand Vizier.\n\nShe waited until 9 pm.")
    T.upload_source(engine, TENANT, "editor", False, jid, "kirmizi.txt", src)
    segs = T.segments(engine, TENANT, "ayse", False, jid, None, "hepsi")["items"]
    by = {s["source"]: s for s in segs}
    targets = {"The road was long.": "Yol uzundu.", "He paid 12 coins.": "13 sikke ödedi",
               "It cost 1250.50 dollars in 2019.": "2018'de 1.250,50 dolara mal oldu.",
               "Nobody came to the Grand Vizier.": "- Sadrazam'a kimse gelmedi; «hiç kimse».",
               "She waited until 9 pm.": "=Akşam 9'a kadar bekledi."}
    for source, target in targets.items():
        T.save_segment(engine, TENANT, "ayse", False, by[source]["id"], {"target": target, "status": "cevrildi"})
    first = by["The road was long."]["id"]
    T.review_segment(engine, TENANT, "mehmet", False, first, {"action": "onayla", "target": "Yol uzundu."})
    T.add_error(engine, TENANT, "mehmet", False, first, {"category": "uslup", "severity": "kucuk",
                                                         "note": "«uzundu» yerine «uzun sürdü»; ağır"})
    viz = by["Nobody came to the Grand Vizier."]["id"]
    T.review_segment(engine, TENANT, "mehmet", False, viz, {"action": "onayla"})
    T.add_error(engine, TENANT, "mehmet", False, viz, {"category": "terim", "severity": "buyuk", "note": "Sadrazam → Veziriazam"})
    rep = _both(client, f"/api/v1/editorial/translation/jobs/{jid}/quality.csv", "Çeviri kalite")
    assert rep["csvRows"] >= 4


# ------------------------------------------------------------------ 5. Serbest çalışan hakedişi


def test_freelance_payout_csv(bridge, monkeypatch, tmp_path):
    from semantic_bridge import freelance as F
    from semantic_layer.tests import test_freelance as TF

    engine, client = bridge
    monkeypatch.setattr(F, "_today", lambda: TF.TODAY)
    monkeypatch.setattr(F, "_now", lambda: TF.NOW)
    monkeypatch.setenv("FREELANCE_DIR", str(tmp_path))
    F._ready.discard(id(engine))
    F.ensure(engine)
    p = TF._person(engine, name="Şule Çağlayan", roles=("kapak",), logoCard="320.01.001")
    TF._accepted(engine, p, "Kapak — ön yüz", price=4000)
    TF._accepted(engine, p, "İç resim; bölüm başları", units=12, price=1250.555)
    TF._accepted(engine, p, 'Arka kapak "tanıtım" düzeni', units=3, price=333.33)
    out = F.create_payout(engine, TENANT, "editor1", {"personId": p["id"]})
    F.payout_action(engine, TENANT, "editor1", out["id"], "submit", {}, False)
    F.payout_action(engine, TENANT, "mudur", out["id"], "approve", {}, True)
    F.payout_action(engine, TENANT, "mudur", out["id"], "pay", {"paidOn": "2026-09-25", "paidRef": "Havale 4471"}, True)
    rep = _both(client, f"/api/v1/editorial/freelance/payouts/{out['id']}/export.csv", "Serbest hakediş")
    assert rep["csvRows"] >= 1 + 5 + 3 and rep["num"] >= 9


# ------------------------------------------------------------------ 6. Telif koşusu: ödeme listesi ve stopaj özeti


def test_royalty_payments_and_withholding_csv(bridge):
    from semantic_bridge import contracts as C
    from semantic_bridge import royalty as RY
    from semantic_layer.tests import test_royalty as TR

    engine, client = bridge
    C._ready.discard(id(engine))
    RY._ready.discard(id(engine))
    RY.ensure(engine)
    P3 = "cccccccc-cccc-cccc-cccc-cccccccccccc"
    items = [TR._item(TR.G1, title="Gökyüzü Atlası", books=[{"title": "Gökyüzü Atlası", "stockCode": "K1", "format": "karton"}]),
             TR._item(TR.G2, parties=[{"name": "Ali Çınar", "share": 50, "contactId": TR.P2},
                                      {"name": "Ayşe Yazar", "share": 50, "contactId": TR.P1}],
                      books=[{"title": "İki Kalem", "stockCode": "K2", "format": "karton"}]),
             TR._item(TR.G3, parties=[{"name": "Yıldız Ajans Ltd. Şti.", "share": 100, "accountId": P3}],
                      books=[{"title": "Üç Nokta", "stockCode": "K3", "format": "karton"}], start="2025-06-01")]
    rows = (TR._rows("K1", 100, 1000.0) + TR._rows("K2", 237, 3181.35, ret_qty=5, ret_net=75.5, month=4)
            + TR._rows("K3", 40, 777.77, month=5))
    run = RY.create_run(engine, TENANT, "acan", {"periodStart": TR.A_, "periodEnd": TR.B_})
    RY.mark_computing(engine, TENANT, "hazirlayan", run["id"])
    run = RY.compute_run(engine, TENANT, run["id"], "hazirlayan", TR.FakeSrc(items, rows),
                         {"statuses": [100000000], "paymentCodes": [2, 7], "withholdingPct": 17.0})
    assert run["summary"]["counts"].get("istisna", 0) == 0, run["summary"]
    RY.submit(engine, TENANT, "gonderen", run["id"], {"acceptDataEnd": True, "note": "veri 17 Ağustos'a kadar"})
    RY.mark_approving(engine, TENANT, "onaylayan", run["id"])
    assert RY.approve_run(engine, TENANT, run["id"], "onaylayan")["status"] == "onayli"
    pay = _both(client, f"/api/v1/royalty/runs/{run['id']}/payments.csv", "Telif ödeme listesi")
    assert pay["csvRows"] == 1 + 4 and pay["date"] >= 4
    wh = _both(client, f"/api/v1/royalty/runs/{run['id']}/withholding.csv", "Telif stopaj özeti")
    assert wh["csvRows"] >= 1 + 3 and wh["num"] >= 6


# ------------------------------------------------------------------ 7. Bütçe planı: kitap hedefleri


def test_budget_plan_export(bridge):
    from semantic_bridge import budget as B
    from semantic_layer.tests import test_budget as TB

    engine, client = bridge
    B._ready.discard(id(engine))
    B.ensure(engine)
    TB._seed(engine)
    plan = TB._approve(engine)
    rep = _both(client, f"/api/v1/budget/plans/{plan['id']}/export.csv", "Bütçe hedefleri")
    assert rep["csvRows"] >= 1 + 4 and rep["num"] > 10 and rep["date"] >= 3


# ------------------------------------------------------------------ 8. Zorunlu eğitim: dolacak/dolmuş listesi


def test_hr_learning_expiring_csv(bridge):
    from semantic_bridge import hr_core as H
    from semantic_bridge import hr_learning as L
    from semantic_layer.tests import test_hr_learning as TL

    engine, client = bridge
    H._ready.discard(engine)
    L._ready.discard(engine)
    L.ensure(engine)
    u1, u2 = TL.unit(engine, "Depo ve Sevkiyat"), TL.unit(engine, "Editörya")
    a, b, c_, d, e = (TL.emp(engine, n, u, u1) for n, u in (("Ayşe Öztürk", "ayse"), ("Ali Çelik", "ali"), ("Can Şahin", "can"),
                                                          ("Deniz Ünal", "deniz"), ("Ece Ağaoğlu", "ece")))
    TL.emp(engine, "Fatma İnce", "fatma", u2)
    k = TL.course(engine, requiredUnits=[u1])
    k2 = TL.course(engine, title="Yangın söndürme; tatbikat", validityDays=730)
    t = L.today()
    TL.cert(engine, a, k, t - timedelta(days=400), t - timedelta(days=35))       # doldu
    TL.cert(engine, c_, k, t - timedelta(days=10), t + timedelta(days=355))      # geçerli (listede yok)
    TL.cert(engine, d, k, t - timedelta(days=355), t + timedelta(days=10))       # dolacak
    TL.cert(engine, e, k2, t - timedelta(days=700), t + timedelta(days=20))      # dolacak
    sid = H.new_id("otr")
    with engine.begin() as c:
        c.execute(L.SESSIONS.insert().values(id=sid, tenant_id=TL.TN, course_id=k, starts_at=datetime.now(timezone.utc) + timedelta(days=7),
                                             location="Ana bina, 3. kat", trainer="Mehmet Güneş", capacity=12, state="planli",
                                             created_by="ik", created_at=H.now(), updated_at=H.now()))
        c.execute(L.ENROLLMENTS.insert().values(id=H.new_id("kat"), tenant_id=TL.TN, session_id=sid, employee_id=b,
                                                approval="onaylandi", requested_by="ik", requested_at=H.now()))
    _grant(engine, L.F_MANAGE, L.F_EXPORT)
    rep = _both(client, "/api/v1/hr/learning/export/expiring.csv?days=30", "Zorunlu eğitim")
    assert rep["csvRows"] >= 1 + 8 and rep["date"] >= 3
