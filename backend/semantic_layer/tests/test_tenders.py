"""M33 İhale takibi: ISBN ve kalem listesi okuma, ISBN → ad → Zeki AI eşleştirme sırası ve eşikleri, teklif tablosu
toplamları (kuruş, KDV ayrı), uygunluk puanı, karar akışı (iki göz), sonuçtan fiyat oranı önerisi, kontrol listesi ve
belge arşivi, şartname özetinin kaynak cümle denetimi, karar özeti metninde olgu dışı sayı, hatırlatmalar, yetki
kuralları ve ilan kaynağı kapalıyken içe alma.

Veriler yapaydır ve yalnız kuralları sınar; gerçek Logo/CRM kabulü test sunucusunda (scripts/acceptance/M33).
"""

from __future__ import annotations

import json
from datetime import date, timedelta
from types import SimpleNamespace

import pytest
import sqlalchemy as sa

from semantic_bridge import access as A
from semantic_bridge import tenders as T
from semantic_layer.store.catalog_store import open_store

TN = "t1"


@pytest.fixture
def engine():
    e = open_store("sqlite://").engine
    T._ready.discard(id(e))
    T.ensure(e)
    return e


@pytest.fixture(autouse=True)
def _files(tmp_path, monkeypatch):
    monkeypatch.setenv("TENDER_DIR", str(tmp_path / "ihale"))
    for k in ("TENDER_WATCH_ENABLED", "TENDER_MATCH_AUTO_PROB", "TENDER_PRICE_SOURCE", "TENDER_DEFAULT_VAT"):
        monkeypatch.delenv(k, raising=False)


BOOKS = [
    {"kod": "15201.01.0001", "ad": "Küçük Prens", "yazar": "Antoine de Saint-Exupéry", "yayinevi": "Timaş Çocuk",
     "isbn_ham": ["978-605-08-1234-8"], "liste_fiyati": 110.0, "kdv_orani": 10.0, "kaynak": "crm"},
    {"kod": "15201.01.0002", "ad": "Çalıkuşu", "yazar": "Reşat Nuri Güntekin", "yayinevi": "Timaş Yayınları",
     "isbn_ham": [], "liste_fiyati": 220.0, "kdv_orani": 0.10, "kaynak": "crm"},
    {"kod": "15201.01.0003", "ad": "Masal Kitabı", "yazar": "Ayşe Yılmaz", "yayinevi": "Timaş Çocuk", "isbn_ham": [],
     "liste_fiyati": 55.0, "kdv_orani": None, "kaynak": "crm"},
    {"kod": "15201.01.0004", "ad": "Masal Kitabı", "yazar": "Mehmet Kaya", "yayinevi": "Timaş Çocuk", "isbn_ham": [],
     "liste_fiyati": 60.0, "kdv_orani": None, "kaynak": "crm"},
    {"kod": "15201.01.0005", "ad": "Uzay Yolculuğu Serisi 1 Aya Gidiyoruz", "yazar": "Zeynep Demir", "yayinevi": "Genç Timaş",
     "isbn_ham": [], "liste_fiyati": 90.0, "kdv_orani": 10.0, "kaynak": "crm"},
]


def _cat():
    return T.Catalog([dict(b) for b in BOOKS])


CFG = {"autoProb": 0.90, "autoMargin": 0.50, "suggestProb": 0.70, "suggestMargin": 0.30, "candidates": 8, "minScore": 0.35}


class FakeChoice:
    def __init__(self, choice, p, margin):
        self.choice, self.probability, self.margin, self.method = choice, p, margin, "logprobs"

    def confident(self, min_prob, min_margin=0.0, min_coverage=0.0):
        return self.probability is not None and self.probability >= min_prob and self.margin >= min_margin


def chooser(pick: str, p: float, margin: float, calls: list):
    def choose(prompt, choices):
        calls.append((prompt, list(choices)))
        target = next((c for c in choices if c.startswith(pick)), pick)
        return FakeChoice(target, p, margin)
    return choose


# ------------------------------------------------------------------ ISBN ve liste okuma


def test_isbn_normalization_and_checksum():
    assert T.norm_isbn("978-605-08-1234-8") == "9786050812348"
    assert T.norm_isbn("0-306-40615-2") == "9780306406157"          # ISBN-10 → 13
    assert T.norm_isbn("0306406153", strict=True) is None            # sağlama tutmuyor
    assert T.find_isbn("Tel 0212 555 44 33 · ISBN 978-0-306-40615-7 · 25 adet") == "9780306406157"
    assert T.find_isbn("Sipariş no 1234567890123") is None           # sağlamasız 13 hane ISBN sayılmaz


def test_parse_pasted_table_with_header_keeps_every_row_and_reports_the_unreadable():
    text = ("Sıra\tKitap Adı\tYazarı\tISBN\tAdet\n"
            "1\tKüçük Prens\tAntoine de Saint-Exupéry\t9786050812348\t25\n"
            "2\tÇalıkuşu\tReşat Nuri Güntekin\t\t1.200\n"
            "3\t\t\t\t\n"
            "4\t---\t\t\t3\n")
    out = T.parse_text(text)
    assert [i["ad"] for i in out["items"]] == ["Küçük Prens", "Çalıkuşu"]
    assert out["items"][0]["isbn"] == "9786050812348" and out["items"][0]["adet"] == 25
    assert out["items"][1]["adet"] == 1200                             # Türkçe binlik ayırıcı
    # yalnız sıra numarası olan satır ve «---» nedeniyle listelenir; tamamen boş satır sessizce geçilir (sondaki satır sonu)
    assert [s["satir"] for s in out["skipped"]] == [4, 5]
    assert set(out["header"]) >= {"ad", "yazar", "isbn", "adet"}


def test_parse_free_lines():
    out = T.parse_text("12. Çalıkuşu - Reşat Nuri Güntekin 15 adet\nKüçük Prens x 4\nAya Gidiyoruz 978-0-306-40615-7")
    a, b, c = out["items"]
    assert (a["ad"], a["yazar"], a["adet"]) == ("Çalıkuşu", "Reşat Nuri Güntekin", 15)
    assert (b["ad"], b["adet"]) == ("Küçük Prens", 4)
    assert c["isbn"] == "9780306406157" and c["ad"] == "Aya Gidiyoruz"


# ------------------------------------------------------------------ eşleştirme


def test_isbn_then_exact_title_match_without_asking_the_model():
    cat, calls = _cat(), []
    m = T.match_one({"ad": "Bambaşka ad", "isbn": "9786050812348"}, cat, chooser("x", 1, 1, calls), CFG)
    assert (m["eslesme_durumu"], m["eslesme_yontemi"], m["eslesen_stok_kodu"]) == ("eslesti", "isbn", "15201.01.0001")
    m = T.match_one({"ad": "CALIKUSU", "yazar": "Reşat Nuri"}, cat, chooser("x", 1, 1, calls), CFG)
    assert (m["eslesme_durumu"], m["eslesme_yontemi"], m["eslesen_stok_kodu"]) == ("eslesti", "ad", "15201.01.0002")
    assert calls == []


def test_same_title_is_separated_by_author_or_sent_to_the_model():
    cat, calls = _cat(), []
    m = T.match_one({"ad": "Masal Kitabı", "yazar": "Mehmet Kaya"}, cat, chooser("x", 1, 1, calls), CFG)
    assert m["eslesen_stok_kodu"] == "15201.01.0004" and calls == []
    m = T.match_one({"ad": "Masal Kitabı"}, cat, chooser("Masal Kitabı · Ayşe", 0.95, 0.9, calls), CFG)
    assert m["eslesme_yontemi"] == "zeki" and m["eslesen_stok_kodu"] == "15201.01.0003"
    assert calls and calls[0][1][-1] == T.NONE_CHOICE and len(calls[0][1]) == 3


def test_model_thresholds_auto_suggest_unsure_and_none():
    item = {"ad": "Aya Gidiyoruz", "yazar": "Zeynep Demir"}
    for p, margin, state in ((0.95, 0.9, "eslesti"), (0.8, 0.4, "oneri"), (0.6, 0.1, "belirsiz")):
        m = T.match_one(item, _cat(), chooser("Uzay Yolculuğu", p, margin, []), CFG)
        assert m["eslesme_durumu"] == state, (p, state)
        assert m["eslesen_stok_kodu"] == "15201.01.0005" and m["olasilik"] == p
    m = T.match_one(item, _cat(), chooser(T.NONE_CHOICE, 0.9, 0.8, []), CFG)
    assert m["eslesme_durumu"] == "yok" and m["eslesen_stok_kodu"] is None and m["adaylar"]


def test_model_down_or_missing_leaves_the_item_to_a_human():
    def down(prompt, choices):
        raise RuntimeError("LLM HTTP 503")
    for choose in (down, None):
        m = T.match_one({"ad": "Aya Gidiyoruz"}, _cat(), choose, CFG)
        assert m["eslesme_durumu"] == "belirsiz" and m["adaylar"][0]["stokKodu"] == "15201.01.0005"
    m = T.match_one({"ad": "Kuantum Fiziğine Giriş"}, _cat(), None, CFG)
    assert m["eslesme_durumu"] == "yok" and m["adaylar"] == []


# ------------------------------------------------------------------ ad eşleştirme sağlamlığı (kabul 2026-09-28: 4 adlı kalemin 1'i yanlış)


def test_fold_turkish_case_and_apostrophes():
    assert T.fold("KUR'AN-I KERİM") == T.fold("Kur’an-ı Kerim") == "kuran i kerim"
    assert T.fold("IŞIK") == T.fold("ışık") == "isik"
    assert T.fold("İstanbul") == T.fold("ISTANBUL") == "istanbul"
    assert T.fold("Yunus Emre'nin Divanı") == "yunus emrenin divani"


def test_title_key_reads_volume_set_digital_and_format():
    k = T.title_key
    assert (k("Osmanlı Tarihi 3. Cilt")["vol"], k("Osmanlı Tarihi 3. Cilt")["core"]) == (3, "osmanli tarihi")
    assert k("OSMANLI TARİHİ 3.CİLT")["vol"] == 3
    assert k("Osmanlı Tarihi Cilt: II")["vol"] == 2
    assert (k("Osmanlı Tarihi II")["vol"], k("Osmanlı Tarihi II")["core"]) == (2, "osmanli tarihi")
    assert k("Osmanlı Tarihi - 2")["vol"] == 2
    assert k("Osmanlı Tarihi 3. Kitap")["vol"] == 3 and not k("Osmanlı Tarihi 3. Kitap")["set"]   # noktalı: cilt
    assert k("Çocuk Klasikleri 10 Kitap")["set"] and k("Çocuk Klasikleri 10 Kitap")["vol"] is None  # noktasız: set
    assert k("Osmanlı Tarihi Seti (2 Kitap)")["core"] == "osmanli tarihi"
    assert k("Nutuk (Ciltli)") == {"core": "nutuk", "vol": None, "set": False, "digital": False, "fmt": frozenset({"ciltli"})}
    assert k("Sefiller E-Kitap")["digital"] and k("Sefiller E-Kitap")["core"] == "sefiller"
    assert k("Fahrenheit 451")["vol"] is None and k("1984")["core"] == "1984"
    assert k("Uzay Yolculuğu Serisi 1 Aya Gidiyoruz")["vol"] is None


def test_author_compat_needs_a_surname_not_just_a_first_name():
    c = T.author_compat
    assert c("Reşat Nuri", "Reşat Nuri Güntekin") is True
    assert c("A. H. Tanpınar", "Ahmet Hamdi Tanpınar") is True
    assert c("Tanpınar, Ahmet Hamdi", "Ahmet Hamdi Tanpınar") is True
    assert c("Victor Hugo (Çev. Ali Veli)", "Victor Hugo") is True
    assert c("Ahmet Ümit", "Ahmet Hamdi Tanpınar") is False                 # yalnız ad ortak: başka yazar
    assert c("", "Ahmet Ümit") is None and c("Kolektif", "Ahmet Ümit") is None


ROBUST = [
    {"kod": "V1", "ad": "Osmanlı Tarihi 1. Cilt", "yazar": "İlber Ortaylı", "isbn_ham": [], "kaynak": "crm"},
    {"kod": "V2", "ad": "Osmanlı Tarihi 2. Cilt", "yazar": "İlber Ortaylı", "isbn_ham": [], "kaynak": "crm"},
    {"kod": "S1", "ad": "Osmanlı Tarihi Seti (2 Kitap)", "yazar": "İlber Ortaylı", "isbn_ham": [], "kaynak": "crm"},
    {"kod": "K1", "ad": "Kayıp Gül", "yazar": "Serdar Özkan", "isbn_ham": [], "kaynak": "crm"},
    {"kod": "K2", "ad": "KAYIP GÜL", "yazar": None, "isbn_ham": [], "kaynak": "logo"},       # CRM'de olmayan Logo kartı
    {"kod": "N1", "ad": "Nutuk", "yazar": "Mustafa Kemal Atatürk", "isbn_ham": [], "kaynak": "crm"},
    {"kod": "N2", "ad": "Nutuk", "yazar": "Mustafa Kemal Atatürk", "isbn_ham": [], "kaynak": "crm"},  # ikinci baskı kaydı
    {"kod": "H1", "ad": "Huzur", "yazar": "Ahmet Hamdi Tanpınar", "isbn_ham": [], "kaynak": "crm"},
    {"kod": "H2", "ad": "Huzur", "yazar": "Ahmet Ümit", "isbn_ham": [], "kaynak": "crm"},
]


def _robust():
    return T.Catalog([dict(b) for b in ROBUST])


def _states(m):
    return m["eslesme_durumu"], m["eslesen_stok_kodu"]


def test_volume_spelling_matches_the_same_volume_without_the_model():
    for ad in ("Osmanlı Tarihi Cilt 2", "Osmanlı Tarihi II", "OSMANLI TARİHİ 2.CİLT", "Osmanlı Tarihi - 2"):
        calls: list = []
        m = T.match_one({"ad": ad, "yazar": "İlber Ortaylı"}, _robust(), chooser("x", 1, 1, calls), CFG)
        assert _states(m) == ("eslesti", "V2") and m["eslesme_yontemi"] == "ad" and calls == [], ad


def test_other_volume_or_set_never_matches_and_is_not_offered_to_the_model():
    calls: list = []
    m = T.match_one({"ad": "Osmanlı Tarihi 5. Cilt"}, _robust(), chooser("Osmanlı Tarihi 1", 0.99, 0.99, calls), CFG)
    assert m["eslesme_durumu"] == "belirsiz" and m["eslesen_stok_kodu"] is None and "Emin değil" in m["not"]
    assert m["adaylar"] and calls == []                                        # insan listeyi görür, model çağrılmaz
    calls = []
    m = T.match_one({"ad": "Osmanlı Tarihi"}, _robust(), chooser("Osmanlı Tarihi 1. Cilt", 0.99, 0.95, calls), CFG)
    assert _states(m) == ("oneri", "V1") and "cilt" in m["not"]                 # cilt yalnız katalogda: en çok öneri
    assert not any("Seti" in c for c in calls[0][1])                          # set tek kitaba aday değil
    assert any(a["stokKodu"] == "S1" for a in m["adaylar"])


def test_author_confirmed_record_beats_one_without_author():
    calls: list = []
    m = T.match_one({"ad": "Kayıp Gül", "yazar": "Serdar Özkan"}, _robust(), chooser("x", 1, 1, calls), CFG)
    assert _states(m) == ("eslesti", "K1") and calls == []
    m = T.match_one({"ad": "Huzur", "yazar": "Ahmet Ümit"}, _robust(), chooser("x", 1, 1, calls), CFG)
    assert _states(m) == ("eslesti", "H2")                                      # «Ahmet» ortak ama soyadı başka
    m = T.match_one({"ad": "Huzur", "yazar": "A. H. Tanpınar"}, _robust(), chooser("x", 1, 1, calls), CFG)
    assert _states(m) == ("eslesti", "H1") and calls == []


def test_indistinguishable_twins_are_at_most_a_suggestion():
    for item, pick in (({"ad": "Nutuk", "yazar": "Mustafa Kemal Atatürk"}, "Nutuk · Mustafa Kemal Atatürk · N2"),
                       ({"ad": "Kayıp Gül"}, "Kayıp Gül"), ({"ad": "Huzur"}, "Huzur · Ahmet Hamdi")):
        m = T.match_one(item, _robust(), chooser(pick, 0.99, 0.95, []), CFG)
        assert m["eslesme_durumu"] == "oneri" and "birden çok katalog kaydı" in m["not"], item
        assert len(m["adaylar"]) >= 2


def test_format_word_is_ignored_but_needs_confirmation():
    m = T.match_one({"ad": "Kayıp Gül (Ciltli)", "yazar": "Serdar Özkan"}, _robust(), chooser("x", 1, 1, []), CFG)
    assert _states(m) == ("oneri", "K1") and "ciltli" in m["not"]


def test_free_line_dash_volume_is_part_of_the_title_not_the_author():
    out = T.parse_text("Osmanlı Tarihi - 2 - İlber Ortaylı 3 adet\nNutuk - Ciltli 5 adet")
    a, b = out["items"]
    assert (a["ad"], a["yazar"], a["adet"]) == ("Osmanlı Tarihi 2", "İlber Ortaylı", 3)
    assert (b["ad"], b["yazar"], b["adet"]) == ("Nutuk Ciltli", None, 5)


# ------------------------------------------------------------------ fiyat, toplam, puan


def test_totals_round_each_line_and_keep_vat_separate():
    items = [
        {"eslesme_durumu": "eslesti", "adet": 3, "onerilen_fiyat": 33.333, "kdv_orani": 0.10, "liste_fiyati": 110.0, "stok": 1, "tahmini_maliyet": 20.0},
        {"eslesme_durumu": "eslesti", "adet": None, "onerilen_fiyat": 10.0, "kdv_orani": 0.0, "liste_fiyati": 10.0, "stok": 5, "tahmini_maliyet": None},
        {"eslesme_durumu": "oneri", "adet": 2, "onerilen_fiyat": 50.0, "kdv_orani": 0.0, "liste_fiyati": 50.0, "stok": 0, "tahmini_maliyet": None},
    ]
    t = T.totals(items)
    assert t["araToplam"] == 100.0 and t["kdv"] == 10.0 and t["genelToplam"] == 110.0
    assert t["fiyatli"] == 1 and t["disarida"] == 1 and t["stokYetersiz"] == 1
    assert t["listeToplami"] == 300.0 and t["maliyet"] == 60.0 and abs(t["marj"] - 0.4) < 1e-9
    assert T.totals([])["maliyetNotu"].startswith("Maliyet bilinmiyor")


def test_score_skips_unmeasured_parts_and_renormalizes():
    cfg = {"weights": {"eslesme": 40.0, "stok": 25.0, "belge": 20.0, "sure": 15.0}, "fullDays": 14}
    items = [{"eslesme_durumu": "eslesti", "stok": 5, "adet": 2}, {"eslesme_durumu": "yok", "stok": None, "adet": 1}]
    s = T.score(items, [], None, cfg)
    assert s["parcalar"]["belge"] is None and s["parcalar"]["sure"] is None
    assert s["puan"] == round(100 * (40 * 0.5 + 25 * 1.0) / 65, 1)
    deadline = (date.today() + timedelta(days=7)).isoformat()
    s = T.score([], [{"zorunlu": True, "durum": "var"}, {"zorunlu": True, "durum": "eksik"}], deadline, cfg, date.today())
    assert s["puan"] == round(100 * (20 * 0.5 + 15 * 0.5) / 35, 1)


def test_vat_rate_accepts_percent_or_fraction():
    assert T.vat_rate(10.0, 0) == 0.10 and T.vat_rate(0.10, 0) == 0.10 and T.vat_rate(None, 0.18) == 0.18


# ------------------------------------------------------------------ akış


def _enrich(cat):
    def enrich(codes):
        return T.enrich_codes(cat, codes, {c: 10.0 for c in codes}, {}, {}, T.settings())
    return enrich


def _tender(engine, user="ayse", **kw):
    body = {"kurum": "Üsküdar İlçe MEM", "kurumTuru": "mem", "konu": "Okul kütüphaneleri için kitap alımı", "il": "İstanbul",
            "sonTeklifTarihi": (date.today() + timedelta(days=10)).isoformat(), **kw}
    return T.create(engine, TN, user, body)


def test_create_validates_and_refuses_official_source(engine):
    with pytest.raises(T.TenderError):
        T.create(engine, TN, "ayse", {"kurum": "", "kurumTuru": "mem", "konu": "x"})
    with pytest.raises(T.TenderError):
        T.create(engine, TN, "ayse", {"kurum": "A", "kurumTuru": "yok", "konu": "x"})
    with pytest.raises(T.TenderError):
        T.create(engine, TN, "ayse", {"kurum": "A", "kurumTuru": "mem", "konu": "x", "kaynak": "resmi"})
    t = _tender(engine)
    assert t["durum"] == "yeni" and t["fiyatOrani"] == 1.0 and t["sorumlu"] == "ayse"


def test_match_job_writes_items_prices_and_keeps_human_choices(engine):
    t = _tender(engine)
    T.import_items(engine, TN, "ayse", t["id"], T.parse_text("Kitap Adı\tISBN\tAdet\nKüçük Prens\t9786050812348\t5\n"
                                                           "Çalıkuşu\t\t20\nKuantum Fiziğine Giriş\t\t1"), "replace")
    cat = _cat()
    res = T.run_match(engine, TN, t["id"], cat, None, _enrich(cat), lambda d, n: None)
    assert res["durumlar"]["eslesti"] == 2 and res["durumlar"]["yok"] == 1
    d = T.detail(engine, TN, t["id"])
    k1 = d["kalemler"][0]
    assert k1["stok"] == 10.0 and k1["kdvOrani"] == 0.10 and k1["fiyatKaynagi"] == "CRM KDV dahil liste fiyatı"
    assert k1["onerilenFiyat"] == 100.0 and k1["tutar"] == 500.0
    assert d["toplamlar"]["stokYetersiz"] == 1                        # Çalıkuşu: stok 10 < 20
    assert d["durum"] == "inceleniyor"
    # insan üçüncü kalemi elle eşler; yeniden eşleştirme dokunmaz
    T.update_item(engine, TN, "ayse", t["id"], 3, {"stokKodu": "15201.01.0005"}, _enrich(cat))
    T.run_match(engine, TN, t["id"], cat, None, _enrich(cat), lambda d, n: None)
    k3 = T.detail(engine, TN, t["id"])["kalemler"][2]
    assert (k3["durum"], k3["yontem"], k3["stokKodu"]) == ("eslesti", "elle", "15201.01.0005")


def _ready_tender(engine):
    t = _tender(engine)
    T.import_items(engine, TN, "ayse", t["id"], T.parse_text("Kitap Adı\tAdet\nKüçük Prens\t5\nÇalıkuşu\t2"), "replace")
    cat = _cat()
    T.run_match(engine, TN, t["id"], cat, None, _enrich(cat), lambda d, n: None)
    return t


def test_two_eyes_decision_then_offer_and_result_feed_the_next_price_ratio(engine):
    t = _ready_tender(engine)
    with pytest.raises(T.TenderError):
        T.submit_decision(engine, TN, "ayse", t["id"], {"karar": "basvurma"})      # gerekçe yok
    dec = T.submit_decision(engine, TN, "ayse", t["id"], {"karar": "basvur", "gerekce": "stok yeterli"})
    assert dec["durum"] == "onayda" and dec["teklifToplami"] == 500.0 + 400.0
    with pytest.raises(T.TenderError) as e:
        T.decide(engine, TN, "AYSE", t["id"], True)                               # öneren onaylayamaz
    assert e.value.status == 409
    with pytest.raises(T.TenderError):
        T.update_item(engine, TN, "ayse", t["id"], 1, {"adet": 9}, _enrich(_cat()))   # onayda kalem değişmez
    T.decide(engine, TN, "mehmet", t["id"], True, "uygun")
    d = T.detail(engine, TN, t["id"])
    assert d["durum"] == "basvurulacak" and all(k["onaylayan"] == "mehmet" for k in d["kalemler"] if k["durum"] == "eslesti")
    with pytest.raises(T.TenderError):
        T.record_result(engine, TN, "ayse", t["id"], {"sonuc": "kaybedildi", "neden": "fiyat"})   # teklif verilmedi
    T.update(engine, TN, "ayse", t["id"], {"durum": "teklif_verildi"})
    r = T.record_result(engine, TN, "ayse", t["id"], {"sonuc": "kaybedildi", "kazanan": "Rakip AŞ", "kazananFiyat": 600, "neden": "fiyat"})
    assert r["bizimFiyat"] == 900.0 and r["listeToplami"] == 900.0 and abs(r["kazananListeOrani"] - 600 / 900) < 1e-9
    assert T.detail(engine, TN, t["id"])["durum"] == "kaybedildi"
    ratio = T.price_ratio(engine, TN, "mem", 1)
    assert abs(ratio["oran"] - round(600 / 900, 4)) < 1e-9 and "ortanca" in ratio["kaynak"]
    assert T.price_ratio(engine, TN, "mem", 3)["oran"] == 1.0


def test_status_changes_outside_the_flow_are_refused(engine):
    t = _tender(engine)
    for st in ("basvurulacak", "kazanildi"):
        with pytest.raises(T.TenderError):
            T.update(engine, TN, "ayse", t["id"], {"durum": st})
    with pytest.raises(T.TenderError):
        T.update(engine, TN, "ayse", t["id"], {"durum": "teklif_verildi"})
    out, diff = T.update(engine, TN, "ayse", t["id"], {"durum": "iptal"})
    assert out["durum"] == "iptal" and diff["durum"] == ["yeni", "iptal"]


def test_price_ratio_change_reprices_only_unedited_lines(engine):
    t = _ready_tender(engine)
    T.update_item(engine, TN, "ayse", t["id"], 2, {"onerilenFiyat": "150,50"}, _enrich(_cat()))
    T.update(engine, TN, "ayse", t["id"], {"fiyatOrani": 0.8})
    ks = T.detail(engine, TN, t["id"])["kalemler"]
    assert ks[0]["onerilenFiyat"] == 80.0 and ks[1]["onerilenFiyat"] == 150.5 and ks[1]["fiyatElle"]


def test_summary_links_archive_documents_and_expired_ones_do_not_count(engine):
    t = _tender(engine)
    ok = T.add_document(engine, TN, "fin", {"ad": "Vergi borcu yoktur", "tur": "vergi",
                                            "gecerlilik": (date.today() + timedelta(days=40)).isoformat()}, "", b"")
    T.add_document(engine, TN, "fin", {"ad": "SGK", "tur": "sgk", "gecerlilik": (date.today() - timedelta(days=1)).isoformat()}, "", b"")
    summary = {"belgeler": [{"deger": "Vergi borcu olmadığına dair yazı", "kaynak": "x"},
                            {"deger": "SGK prim borcu yoktur yazısı", "kaynak": "y"},
                            {"deger": "Numune kitap", "kaynak": "z"}]}
    out = T.apply_summary(engine, TN, "ayse", t["id"], summary, {"sonuc": "evet", "olasilik": 0.97}, None)
    assert out["kontrolListesineEklenen"] == 3
    cl = {c["kalem"]: c for c in T.checklist(engine, TN, t["id"])["items"]}
    assert cl["Vergi borcu olmadığına dair yazı"]["durum"] == "var" and cl["Vergi borcu olmadığına dair yazı"]["belgeId"] == ok["id"]
    assert cl["SGK prim borcu yoktur yazısı"]["durum"] == "eksik"              # süresi dolmuş belge bağlanmaz
    assert cl["Numune kitap"]["belgeTuru"] == "diger"
    T.apply_summary(engine, TN, "ayse", t["id"], summary, None, None)
    assert len(T.checklist(engine, TN, t["id"])["items"]) == 3                 # aynı kalem ikinci kez eklenmez
    T.delete_document(engine, TN, ok["id"])
    assert cl and T.checklist(engine, TN, t["id"])["items"][0]["durum"] == "eksik"


def test_summarize_drops_items_whose_quote_is_not_in_the_document():
    text = "Teklif mektubu ile birlikte geçici teminat olarak teklif bedelinin %3'ü verilecektir. Teslim süresi 30 gündür."
    reply = {"konu": None,
             "teminat": {"deger": "%3 geçici teminat", "kaynak": "geçici teminat olarak teklif bedelinin %3'ü verilecektir"},
             "teslim_suresi": {"deger": "45 gün", "kaynak": "Teslim süresi 30 gündür"},
             "belgeler": [{"deger": "İmza sirküleri", "kaynak": "noter onaylı imza sirküleri"}], "kosullar": []}
    out = T.summarize_text(text, lambda msgs: "```json\n" + json.dumps(reply, ensure_ascii=False) + "\n```",
                           {"summaryChunk": 4000}, lambda d, n: None)
    assert out["teminat"]["deger"] == "%3 geçici teminat"
    assert out["teslimSuresi"] is None and out["belgeler"] == [] and out["atilan"] == 2


def test_brief_text_rejects_numbers_that_are_not_in_the_facts():
    facts = {"yaklasikTutar": 125000.0, "uygunlukPuani": 72.5, "kalem": 312, "eslesen": 188, "katalogdaYok": 40,
             "stokYetersiz": 21, "teklifAraToplam": 90000.0, "teklifGenelToplam": 99000.0, "listeToplami": 100000.0,
             "fiyatOrani": 0.9, "marj": None, "maliyetKapsam": 0, "maliyetNotu": "Maliyet bilinmiyor.", "eksikBelgeler": [],
             "teminatTutari": None, "kalanGun": 6, "gecmis": {"kurum": {}, "kurumTuru": {"sonuc": 0}}, "riskler": []}
    good = T.brief_text(facts, lambda m: "312 kalemin 188'i eşleşti; 21 kalemde stok yetersiz. Son teklife 6 gün var.")
    assert good["kaynak"] == "zeki"
    bad = T.brief_text(facts, lambda m: "312 kalemin 190'ı eşleşti.")
    assert bad["kaynak"] == "kural" and "312" in bad["metin"]
    assert T.brief_text(facts, None)["kaynak"] == "kural"


def test_reminders_go_once_per_step(engine):
    t = _tender(engine, sonTeklifTarihi=(date.today() + timedelta(days=5)).isoformat())
    T.add_document(engine, TN, "fin", {"ad": "İmza sirküleri", "tur": "imza_sirkuleri",
                                       "gecerlilik": (date.today() + timedelta(days=20)).isoformat()}, "", b"")
    cfg = T.settings()
    due = T.due_reminders(engine, TN, cfg)
    kinds = sorted(i["tur"] for i in due)
    assert kinds == ["belge", "son_teklif"]
    T.mark_sent(engine, TN, [i["key"] for i in due])
    assert T.due_reminders(engine, TN, cfg) == []
    later = T.due_reminders(engine, TN, cfg, date.today() + timedelta(days=4))   # 1 gün kala: 2 günlük adım
    assert [i["tur"] for i in later if i["tur"] == "son_teklif"] == ["son_teklif"]
    assert "kuruma hiçbir gönderim" in T.reminder_text(due, "")
    assert t["id"]


def test_pricing_xlsx_has_the_offer_lines_and_totals(engine):
    from openpyxl import load_workbook
    import io

    t = _ready_tender(engine)
    wb = load_workbook(io.BytesIO(T.pricing_xlsx(T.detail(engine, TN, t["id"]))))
    ws = wb["Teklif fiyat tablosu"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Ara toplam (KDV hariç)" in values and 900.0 in values


def test_file_upload_checks_type_and_content(engine):
    t = _tender(engine)
    with pytest.raises(T.TenderError):
        T.add_file(engine, TN, "ayse", t["id"], "sartname.exe", b"MZ", "sartname")
    with pytest.raises(T.TenderError):
        T.add_file(engine, TN, "ayse", t["id"], "sartname.pdf", b"not a pdf", "sartname")
    f = T.add_file(engine, TN, "ayse", t["id"], "liste.csv", "Kitap Adı;Adet\nKüçük Prens;3\n".encode("cp1254"), "ek")
    path, name, _ = T.file_of(engine, TN, t["id"], f["id"])
    parsed = T.parse_rows(T.extract_rows(name, open(path, "rb").read()))
    assert parsed["items"][0]["ad"] == "Küçük Prens" and parsed["items"][0]["adet"] == 3


def test_delete_only_without_decision(engine):
    t = _ready_tender(engine)
    T.submit_decision(engine, TN, "ayse", t["id"], {"karar": "basvur"})
    with pytest.raises(T.TenderError):
        T.delete(engine, TN, t["id"])
    t2 = _tender(engine)
    T.delete(engine, TN, t2["id"])
    with pytest.raises(T.TenderError):
        T.detail(engine, TN, t2["id"])


# ------------------------------------------------------------------ yetki


def test_tender_rules():
    assert A.rule_for("/api/v1/tenders") == frozenset({"sayfa:ihale"})
    assert A.rule_for("/api/v1/tenders/abc/items/match") == frozenset({"sayfa:ihale"})
    assert A.rule_for("/api/v1/tenders/run-due") == A.SYSTEM
    f = A.features_for
    assert f("POST", "/api/v1/tenders") == ["ozellik:ihale.duzenle"]
    assert f("PATCH", "/api/v1/tenders/abc/items/3") == ["ozellik:ihale.duzenle"]
    assert f("POST", "/api/v1/tenders/abc/decision/submit") == ["ozellik:ihale.duzenle"]
    assert f("POST", "/api/v1/tenders/abc/decision/approve") == []           # açıkça verilen ihale.karar ucun içinde
    assert f("POST", "/api/v1/tenders/documents") == ["ozellik:ihale.belge"]
    assert f("DELETE", "/api/v1/tenders/documents/d1") == ["ozellik:ihale.belge"]
    assert f("GET", "/api/v1/tenders/abc/pricing.xlsx") == ["ozellik:veri.disa-aktar"]
    assert f("GET", "/api/v1/tenders/abc") == [] and f("POST", "/api/v1/tenders/run-due") == []
    assert f("POST", "/api/v1/tenders/watch/import") == []
    assert {"ozellik:ihale.karar", "ozellik:ihale.kaynak-yonet"} <= A.explicit_keys()
    assert {"sayfa:ihale", "ozellik:ihale.duzenle", "ozellik:ihale.belge"} <= A.all_keys()


def test_api_decision_needs_explicit_permission_and_watch_is_off(monkeypatch, store, settings):
    from fastapi.testclient import TestClient

    from semantic_bridge import admin as admin_mod
    from semantic_bridge import board as board_mod
    from semantic_bridge.app import Runtime, create_app
    from semantic_layer.candidates.llm_client import FakeLlm
    from semantic_layer.tests.conftest import TENANT

    users = {"timas_session=a": "ayse", "timas_session=z": "zekiai"}
    monkeypatch.setattr(board_mod, "_fetch_user", lambda cookie: users.get(cookie))
    monkeypatch.setattr(board_mod, "_fetch_session", lambda cookie: {"username": users.get(cookie), "displayName": "x"})
    monkeypatch.setattr(admin_mod, "admins", lambda: ["zekiai"])
    monkeypatch.delenv("SEMANTIC_CALLER_TOKEN", raising=False)
    A._ready.clear()
    A.invalidate()
    client = TestClient(create_app(Runtime(settings, store=store, llm=FakeLlm([""]))))
    a, z = {"cookie": "timas_session=a"}, {"cookie": "timas_session=z"}
    assert client.post("/api/v1/tenders", json={}, headers=a).status_code == 400
    t = client.post("/api/v1/tenders", json={"kurum": "Kadıköy Belediyesi", "kurumTuru": "belediye", "konu": "Kütüphane kitap alımı"},
                    headers=a).json()
    assert client.post(f"/api/v1/tenders/{t['id']}/decision/approve", json={}, headers=a).status_code == 403
    assert client.post(f"/api/v1/tenders/{t['id']}/decision/approve", json={}, headers=z).status_code == 409   # onay bekleyen yok
    assert client.post("/api/v1/tenders/watch/import", json={}, headers=a).status_code == 403
    r = client.post("/api/v1/tenders/watch/import", json={}, headers=z)
    assert r.status_code == 409 and "kapalı" in r.json()["detail"]["message"]
    assert client.get("/api/v1/tenders/watch/status", headers=a).json()["enabled"] is False
    assert client.get("/api/v1/tenders/yok-boyle", headers=a).status_code == 404
    lst = client.get("/api/v1/tenders", headers=a).json()
    assert lst["total"] == 1 and lst["items"][0]["kurum"] == "Kadıköy Belediyesi"
    # test verisi bırakılmaz
    assert client.delete(f"/api/v1/tenders/{t['id']}", headers=a).status_code == 200
    assert TENANT
