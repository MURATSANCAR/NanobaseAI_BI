"""Budget envelope bulk import — Excel/CSV template and row upsert."""

from __future__ import annotations

import csv
import io
import re
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.engine import Engine

from nanobase_api.budget_lines import save_commitment, save_periods
from nanobase_api.budgets import list_budgets, upsert_budget

_EXPORTS_DIR = Path("/tmp/nanobase-budget-exports")


def _exports_dir() -> Path:
    try:
        _EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        return _EXPORTS_DIR
    except Exception:
        return Path(tempfile.mkdtemp(prefix="nanobase-budget-exports-"))

# Template headers (TR — matches portal budget terminology)
TEMPLATE_HEADERS = [
    "Kalem adı",
    "Mali yıl",
    "Senaryo",
    "Bütçe türü",
    "Maliyet merkezi",
    "Planlanan",
    "Para birimi",
    "Taahhüt",
    "Durum",
    "Sorumlu",
    "Notlar",
    "Gerçekleşen SQL",
]

PERIOD_SHEET_NAME = "Aylık"
COMMITMENT_SHEET_NAME = "Taahhütler"

PERIOD_HEADERS = [
    "Kalem adı",
    "Mali yıl",
    "Senaryo",
    "Bütçe türü",
    "Maliyet merkezi",
    *[f"A{i}" for i in range(1, 13)],
]

COMMITMENT_HEADERS = [
    "Kalem adı",
    "Mali yıl",
    "Senaryo",
    "Bütçe türü",
    "Maliyet merkezi",
    "Açıklama",
    "Tutar",
    "Para birimi",
    "Durum",
    "Vade",
    "ID",
]

_SCENARIO_MAP = {
    "base": "base",
    "temel": "base",
    "baz": "base",
    "baseline": "base",
    "optimistic": "optimistic",
    "iyimser": "optimistic",
    "optimist": "optimistic",
    "pessimistic": "pessimistic",
    "kötümser": "pessimistic",
    "kotumser": "pessimistic",
    "pessimist": "pessimistic",
}

# Normalized field → accepted header aliases (lowered, stripped)
_HEADER_ALIASES: dict[str, str] = {}
for _field, _aliases in {
    "name": ("kalem adı", "kalem adi", "name", "envelope name", "budget name", "başlık", "baslik"),
    "fiscal_year": ("mali yıl", "mali yil", "fiscal year", "fiscal_year", "year", "yıl", "yil"),
    "scenario": ("senaryo", "scenario", "plan", "versiyon", "version"),
    "kind": ("bütçe türü", "butce turu", "kind", "type", "tür", "tur", "budget type"),
    "cost_center": ("maliyet merkezi", "cost center", "cost_center", "departman", "department", "cc"),
    "allocated": ("planlanan", "allocated", "allocation", "tahsis", "bütçe tutarı", "butce tutari", "amount"),
    "currency": ("para birimi", "currency", "ccy", "para"),
    "committed": ("taahhüt", "taahhut", "committed", "reserved", "rezerv", "commitment"),
    "status": ("durum", "status", "state"),
    "owner": ("sorumlu", "owner", "owner_name", "sahip"),
    "notes": ("notlar", "notes", "note"),
    "actuals_sql": ("gerçekleşen sql", "gerceklesen sql", "actuals_sql", "actuals sql", "sql"),
    "id": ("id", "budget_id", "kalem id"),
    "description": ("açıklama", "aciklama", "description", "desc", "title"),
    "amount": ("tutar", "amount", "value", "bedel"),
    "due_date": ("vade", "due", "due_date", "due date", "son tarih"),
}.items():
    for a in _aliases:
        _HEADER_ALIASES[a] = _field

_KIND_MAP = {
    "opex": "opex",
    "operating": "opex",
    "işletme": "opex",
    "isletme": "opex",
    "işletme giderleri": "opex",
    "isletme giderleri": "opex",
    "capex": "capex",
    "capital": "capex",
    "yatırım": "capex",
    "yatirim": "capex",
    "yatırım harcamaları": "capex",
    "yatirim harcamalari": "capex",
    "other": "other",
    "diğer": "other",
    "diger": "other",
}

_STATUS_MAP = {
    "draft": "draft",
    "taslak": "draft",
    "approved": "approved",
    "onaylı": "approved",
    "onayli": "approved",
    "onaylandı": "approved",
    "onaylandi": "approved",
    "closed": "closed",
    "kapalı": "closed",
    "kapali": "closed",
    "kapatıldı": "closed",
    "kapatildi": "closed",
}

_COMMITMENT_STATUS_MAP = {
    "open": "open",
    "açık": "open",
    "acik": "open",
    "aktif": "open",
    "released": "released",
    "serbest": "released",
    "serbest bırakıldı": "released",
    "serbest birakildi": "released",
    "cancelled": "cancelled",
    "canceled": "cancelled",
    "iptal": "cancelled",
    "iptal edildi": "cancelled",
}

_ENVELOPE_SHEET_ALIASES = {
    "bütçe kalemleri",
    "butce kalemleri",
    "kalemler",
    "lines",
    "budgets",
    "envelopes",
    "budget lines",
}
_PERIOD_SHEET_ALIASES = {
    "aylık",
    "aylik",
    "monthly",
    "periods",
    "dönem",
    "donem",
    "dönemler",
    "donemler",
}
_COMMITMENT_SHEET_ALIASES = {
    "taahhütler",
    "taahhutler",
    "taahhüt",
    "taahhut",
    "commitments",
    "commitment",
}


def _norm_header(value: Any) -> str:
    s = str(value or "").strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _parse_kind(raw: Any, default: str = "opex") -> str:
    key = _norm_header(raw)
    if not key:
        return default
    if key in _KIND_MAP:
        return _KIND_MAP[key]
    for token, mapped in _KIND_MAP.items():
        if token in key:
            return mapped
    return key if key in ("opex", "capex", "other") else default


def _parse_scenario(raw: Any, default: str = "base") -> str:
    key = _norm_header(raw)
    if not key:
        return default
    mapped = _SCENARIO_MAP.get(key)
    if mapped:
        return mapped
    if key in ("base", "optimistic", "pessimistic"):
        return key
    return default


def _parse_status(raw: Any, default: str = "draft") -> str:
    key = _norm_header(raw)
    if not key:
        return default
    return _STATUS_MAP.get(key, key if key in ("draft", "approved", "closed") else default)


def _parse_commitment_status(raw: Any, default: str = "open") -> str:
    key = _norm_header(raw)
    if not key:
        return default
    return _COMMITMENT_STATUS_MAP.get(key, key if key in ("open", "released", "cancelled") else default)


def _parse_float(raw: Any, default: float = 0.0) -> float:
    if raw is None or raw == "":
        return default
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip().replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return default


def _parse_int(raw: Any, default: int | None = None) -> int | None:
    if raw is None or raw == "":
        return default
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        return int(raw)
    s = str(raw).strip()
    try:
        return int(float(s.replace(",", ".")))
    except ValueError:
        return default


def parse_month_column(header: Any) -> int | None:
    """Map A1/M1/Ay 1/Month 1 style headers to period_index 1–12."""
    h = _norm_header(header)
    if not h:
        return None
    h = re.sub(r"^(?:planlanan|allocated)\s+", "", h)
    m = re.match(r"^(?:a|m|o|м)\s*(\d{1,2})$", h)
    if m:
        n = int(m.group(1))
        return n if 1 <= n <= 12 else None
    m = re.match(r"^(?:ay|month|period|dönem|donem)\s*(\d{1,2})$", h)
    if m:
        n = int(m.group(1))
        return n if 1 <= n <= 12 else None
    if re.fullmatch(r"\d{1,2}", h):
        n = int(h)
        return n if 1 <= n <= 12 else None
    return None


def _sheet_kind(title: str) -> str | None:
    key = _norm_header(title)
    if key in _PERIOD_SHEET_ALIASES:
        return "periods"
    if key in _COMMITMENT_SHEET_ALIASES:
        return "commitments"
    if key in _ENVELOPE_SHEET_ALIASES:
        return "envelopes"
    return None


def build_import_template(*, fiscal_year: int | None = None) -> Path:
    """Write a branded .xlsx template with envelopes + monthly + commitments sheets."""
    year = fiscal_year or datetime.now(timezone.utc).year
    wb = Workbook()
    ws = wb.active
    ws.title = "Bütçe kalemleri"

    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    samples = [
        {
            "Kalem adı": "Bulut altyapı (AWS/Azure/GCP IaaS-PaaS)",
            "Mali yıl": year,
            "Senaryo": "base",
            "Bütçe türü": "OPEX",
            "Maliyet merkezi": "IT-OPEX",
            "Planlanan": 2880000,
            "Para birimi": "TRY",
            "Taahhüt": 0,
            "Durum": "Taslak",
            "Sorumlu": "",
            "Notlar": "ERP alis_faturalari — Sync from source ile eşlenir",
            "Gerçekleşen SQL": (
                f"SELECT COALESCE(SUM(genel_toplam),0) AS amount FROM alis_faturalari "
                f"WHERE butce_kodu = 'IT-CLOUD' AND EXTRACT(YEAR FROM fatura_tarihi) = {year}"
            ),
        },
        {
            "Kalem adı": "Siber güvenlik (EDR, SIEM, SOC, WAF)",
            "Mali yıl": year,
            "Senaryo": "base",
            "Bütçe türü": "OPEX",
            "Maliyet merkezi": "IT-OPEX",
            "Planlanan": 864000,
            "Para birimi": "TRY",
            "Taahhüt": 0,
            "Durum": "Taslak",
            "Sorumlu": "",
            "Notlar": "MSSP + Fortinet aylık/çeyrek faturaları",
            "Gerçekleşen SQL": (
                f"SELECT COALESCE(SUM(genel_toplam),0) AS amount FROM alis_faturalari "
                f"WHERE butce_kodu = 'IT-SEC' AND EXTRACT(YEAR FROM fatura_tarihi) = {year}"
            ),
        },
        {
            "Kalem adı": "Sunucu / storage / network donanımı",
            "Mali yıl": year,
            "Senaryo": "base",
            "Bütçe türü": "CAPEX",
            "Maliyet merkezi": "IT-CAPEX",
            "Planlanan": 3360000,
            "Para birimi": "TRY",
            "Taahhüt": 0,
            "Durum": "Taslak",
            "Sorumlu": "",
            "Notlar": "CAPEX — Dell/Cisco alış faturaları",
            "Gerçekleşen SQL": (
                f"SELECT COALESCE(SUM(genel_toplam),0) AS amount FROM alis_faturalari "
                f"WHERE butce_kodu = 'IT-HW' AND EXTRACT(YEAR FROM fatura_tarihi) = {year}"
            ),
        },
    ]
    for ri, sample in enumerate(samples, start=2):
        for ci, header in enumerate(TEMPLATE_HEADERS, start=1):
            ws.cell(row=ri, column=ci, value=sample.get(header, ""))

    widths = [28, 10, 12, 12, 16, 14, 12, 12, 12, 16, 32, 36]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Monthly periods sheet
    monthly = wb.create_sheet(PERIOD_SHEET_NAME)
    for col, header in enumerate(PERIOD_HEADERS, start=1):
        cell = monthly.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    # Even split samples across 12 months
    monthly_samples = [
        ("Bulut altyapı (AWS/Azure/GCP IaaS-PaaS)", "OPEX", "IT-OPEX", 2880000),
        ("Siber güvenlik (EDR, SIEM, SOC, WAF)", "OPEX", "IT-OPEX", 864000),
        ("Sunucu / storage / network donanımı", "CAPEX", "IT-CAPEX", 3360000),
    ]
    for ri, (name, kind, cc, total) in enumerate(monthly_samples, start=2):
        monthly.cell(row=ri, column=1, value=name)
        monthly.cell(row=ri, column=2, value=year)
        monthly.cell(row=ri, column=3, value="base")
        monthly.cell(row=ri, column=4, value=kind)
        monthly.cell(row=ri, column=5, value=cc)
        per = round(total / 12, 2)
        for mi in range(12):
            monthly.cell(row=ri, column=6 + mi, value=per)
    for i in range(1, 18):
        monthly.column_dimensions[get_column_letter(i)].width = 12 if i > 5 else 22

    # Commitments sheet
    commits = wb.create_sheet(COMMITMENT_SHEET_NAME)
    for col, header in enumerate(COMMITMENT_HEADERS, start=1):
        cell = commits.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    commits.append(
        [
            "Bulut altyapı (AWS/Azure/GCP IaaS-PaaS)",
            year,
            "base",
            "OPEX",
            "IT-OPEX",
            "AWS / Azure reserved capacity commitment",
            480000,
            "TRY",
            "Açık",
            f"{year}-12-31",
            "",
        ]
    )
    for i, w in enumerate([28, 10, 12, 12, 14, 28, 12, 10, 10, 12, 10], start=1):
        commits.column_dimensions[get_column_letter(i)].width = w

    guide = wb.create_sheet("Açıklama")
    guide["A1"] = "Sayfa / Sütun"
    guide["B1"] = "Açıklama"
    guide["A1"].font = Font(bold=True)
    guide["B1"].font = Font(bold=True)
    guide_rows = [
        ("Bütçe kalemleri", "Yıllık kalem listesi — zorunlu sayfa."),
        ("Kalem adı", "Zorunlu. Bütçe kaleminin adı (en az 2 karakter)."),
        ("Mali yıl", f"Boş bırakılırsa {year} kullanılır."),
        ("Senaryo", "base / optimistic / pessimistic (veya Temel / İyimser / Kötümser). Boş = base."),
        ("Bütçe türü", "OPEX, CAPEX veya Diğer."),
        ("Maliyet merkezi", "İsteğe bağlı maliyet merkezi kodu."),
        ("Planlanan", "Zorunlu. Yıllık planlanan tutar (≥ 0)."),
        ("Para birimi", "Varsayılan TRY."),
        ("Taahhüt", "Özet taahhüt; satır detayı için Taahhütler sayfası."),
        ("Durum", "Taslak, Onaylandı veya Kapatıldı."),
        ("Aylık", "İsteğe bağlı. A1–A12 aylık plan dağılımı; kalem adı + senaryo ile eşleşir."),
        ("Taahhütler", "İsteğe bağlı. Açık taahhüt satırları; içe aktarımda kalem taahhüdü güncellenir."),
        ("Gerçekleşen SQL", "İsteğe bağlı; yüklemede çalıştırılmaz."),
    ]
    for i, (col, desc) in enumerate(guide_rows, start=2):
        guide.cell(row=i, column=1, value=col)
        guide.cell(row=i, column=2, value=desc)
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 64

    path = _exports_dir() / f"budget-import-template-{year}-{uuid.uuid4().hex[:8]}.xlsx"
    wb.save(path)
    return path


def _map_headers(raw_headers: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for idx, h in enumerate(raw_headers):
        field = _HEADER_ALIASES.get(_norm_header(h))
        if field:
            mapping[idx] = field
    return mapping


def _map_period_headers(raw_headers: list[Any]) -> tuple[dict[int, str], dict[int, int]]:
    """Return (field_map, month_col_index → period_index)."""
    field_map: dict[int, str] = {}
    month_map: dict[int, int] = {}
    for idx, h in enumerate(raw_headers):
        month = parse_month_column(h)
        if month is not None:
            month_map[idx] = month
            continue
        field = _HEADER_ALIASES.get(_norm_header(h))
        if field and field not in ("allocated", "amount", "committed"):
            field_map[idx] = field
        elif field in ("name", "fiscal_year", "kind", "cost_center", "id"):
            field_map[idx] = field
    return field_map, month_map


def _row_dict(mapping: dict[int, str], values: list[Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for idx, field in mapping.items():
        if idx < len(values):
            out[field] = values[idx]
    return out


def _is_empty_row(row: dict[str, Any]) -> bool:
    name = _cell_str(row.get("name"))
    allocated = row.get("allocated")
    if name:
        return False
    if allocated not in (None, "", 0, 0.0):
        return False
    return True


def _match_existing(
    existing: list[dict[str, Any]],
    *,
    name: str,
    fiscal_year: int,
    kind: str,
    cost_center: str,
    scenario: str = "base",
    loose: bool = False,
) -> dict[str, Any] | None:
    name_l = name.strip().lower()
    cc_l = (cost_center or "").strip().lower()
    scen_l = (scenario or "base").strip().lower() or "base"
    exact: dict[str, Any] | None = None
    loose_hit: dict[str, Any] | None = None
    for item in existing:
        if (
            str(item.get("name") or "").strip().lower() == name_l
            and int(item.get("fiscal_year") or 0) == int(fiscal_year)
            and str(item.get("kind") or "").lower() == kind
            and str(item.get("scenario") or "base").strip().lower() == scen_l
        ):
            item_cc = str(item.get("cost_center") or "").strip().lower()
            if item_cc == cc_l:
                exact = item
                break
            if loose and loose_hit is None:
                loose_hit = item
    if exact:
        return exact
    if loose:
        return loose_hit
    return None


def _resolve_budget_id(
    existing: list[dict[str, Any]],
    raw_row: dict[str, Any],
    *,
    year_default: int,
) -> str | None:
    row_id = _cell_str(raw_row.get("id"))
    if row_id:
        matched = next((e for e in existing if str(e.get("id")) == row_id), None)
        if matched:
            return str(matched.get("id"))
    name = _cell_str(raw_row.get("name"))
    if len(name) < 2:
        return None
    fiscal_year = _parse_int(raw_row.get("fiscal_year"), year_default) or year_default
    kind = _parse_kind(raw_row.get("kind"))
    cost_center = _cell_str(raw_row.get("cost_center"))
    scenario = _parse_scenario(raw_row.get("scenario"))
    matched = _match_existing(
        existing,
        name=name,
        fiscal_year=fiscal_year,
        kind=kind,
        cost_center=cost_center,
        scenario=scenario,
        loose=True,
    )
    return str(matched.get("id")) if matched and matched.get("id") else None


def _iter_sheet_rows(ws) -> list[list[Any]]:
    out: list[list[Any]] = []
    for values in ws.iter_rows(values_only=True):
        if values is None:
            continue
        out.append(list(values))
    return out


def _parse_envelope_sheet_rows(rows: list[list[Any]]) -> list[dict[str, Any]]:
    if not rows:
        raise ValueError("bi_budget_import_empty")
    mapping = _map_headers(rows[0])
    if "name" not in mapping.values() or "allocated" not in mapping.values():
        raise ValueError("bi_budget_import_invalid_file")
    return [_row_dict(mapping, values) for values in rows[1:]]


def _parse_period_sheet_rows(rows: list[list[Any]]) -> list[dict[str, Any]]:
    if not rows:
        return []
    field_map, month_map = _map_period_headers(rows[0])
    if "name" not in field_map.values() or not month_map:
        return []
    out: list[dict[str, Any]] = []
    for values in rows[1:]:
        base = _row_dict(field_map, values)
        if not _cell_str(base.get("name")):
            continue
        months: dict[int, float] = {}
        for idx, period_index in month_map.items():
            if idx < len(values):
                months[period_index] = _parse_float(values[idx], 0.0)
        if not months:
            continue
        base["months"] = months
        out.append(base)
    return out


def _parse_commitment_sheet_rows(rows: list[list[Any]]) -> list[dict[str, Any]]:
    if not rows:
        return []
    mapping = _map_headers(rows[0])
    # description column: prefer description; fall back to notes alias if present
    if "description" not in mapping.values() and "notes" in mapping.values():
        # remap notes → description for this sheet
        mapping = {idx: ("description" if f == "notes" else f) for idx, f in mapping.items()}
    if "name" not in mapping.values() or "amount" not in mapping.values():
        # amount may be mapped as allocated if header was "Planlanan" — accept allocated too
        if "name" not in mapping.values():
            return []
        if "amount" not in mapping.values() and "allocated" not in mapping.values():
            return []
    out: list[dict[str, Any]] = []
    for values in rows[1:]:
        row = _row_dict(mapping, values)
        if "amount" not in row and "allocated" in row:
            row["amount"] = row.get("allocated")
        if not _cell_str(row.get("name")):
            continue
        if _parse_float(row.get("amount"), -1) < 0:
            continue
        out.append(row)
    return out


def _parse_xlsx_workbook(raw: bytes) -> dict[str, Any]:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("bi_budget_import_invalid_file") from exc

    envelopes: list[dict[str, Any]] = []
    periods: list[dict[str, Any]] = []
    commitments: list[dict[str, Any]] = []
    envelope_done = False

    for title in wb.sheetnames:
        kind = _sheet_kind(title)
        if kind == "periods":
            periods.extend(_parse_period_sheet_rows(_iter_sheet_rows(wb[title])))
        elif kind == "commitments":
            commitments.extend(_parse_commitment_sheet_rows(_iter_sheet_rows(wb[title])))
        elif kind == "envelopes" or (kind is None and not envelope_done and title == wb.sheetnames[0]):
            # Skip guide sheet
            if _norm_header(title) in {"açıklama", "aciklama", "guide", "readme", "help"}:
                continue
            try:
                envelopes = _parse_envelope_sheet_rows(_iter_sheet_rows(wb[title]))
                envelope_done = True
            except ValueError:
                if kind == "envelopes":
                    raise
                continue

    if not envelopes:
        # last resort: first sheet with name+allocated
        for title in wb.sheetnames:
            if _sheet_kind(title) in ("periods", "commitments"):
                continue
            if _norm_header(title) in {"açıklama", "aciklama", "guide", "readme", "help"}:
                continue
            try:
                envelopes = _parse_envelope_sheet_rows(_iter_sheet_rows(wb[title]))
                break
            except ValueError:
                continue

    if not envelopes:
        raise ValueError("bi_budget_import_invalid_file")

    return {"envelopes": envelopes, "periods": periods, "commitments": commitments}


def _parse_csv_rows(raw: bytes) -> list[dict[str, Any]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    reader = csv.reader(io.StringIO(text))
    try:
        header_row = next(reader)
    except StopIteration as exc:
        raise ValueError("bi_budget_import_empty") from exc
    mapping = _map_headers(header_row)
    if "name" not in mapping.values() or "allocated" not in mapping.values():
        raise ValueError("bi_budget_import_invalid_file")
    out: list[dict[str, Any]] = []
    for values in reader:
        out.append(_row_dict(mapping, values))
    return out


def parse_budget_file(filename: str, raw: bytes) -> dict[str, Any]:
    """Parse upload into envelopes (+ optional periods/commitments for xlsx)."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return {"envelopes": _parse_csv_rows(raw), "periods": [], "commitments": []}
    if name.endswith(".xlsx") or name.endswith(".xlsm") or raw[:2] == b"PK":
        return _parse_xlsx_workbook(raw)
    return {"envelopes": _parse_csv_rows(raw), "periods": [], "commitments": []}


def import_budget_rows(
    engine: Engine,
    rows: list[dict[str, Any]] | dict[str, Any],
    *,
    tenant_id: str | None = None,
    default_fiscal_year: int | None = None,
) -> dict[str, Any]:
    """Upsert parsed envelopes; optionally apply monthly periods and commitment lines."""
    year_default = default_fiscal_year or datetime.now(timezone.utc).year
    tid = str(tenant_id or "default")
    if isinstance(rows, dict):
        envelopes = list(rows.get("envelopes") or [])
        period_rows = list(rows.get("periods") or [])
        commitment_rows = list(rows.get("commitments") or [])
    else:
        envelopes = list(rows)
        period_rows = []
        commitment_rows = []

    existing = list_budgets(engine, tenant_id=tid)
    created = 0
    updated = 0
    skipped = 0
    periods_updated = 0
    commitments_upserted = 0
    errors: list[dict[str, Any]] = []

    if not envelopes and not period_rows and not commitment_rows:
        raise ValueError("bi_budget_import_empty")

    for idx, raw_row in enumerate(envelopes, start=2):  # 1-based header → data from row 2
        if _is_empty_row(raw_row):
            skipped += 1
            continue
        name = _cell_str(raw_row.get("name"))
        if len(name) < 2:
            errors.append({"row": idx, "code": "bi_budget_name_required", "detail": name, "sheet": "envelopes"})
            continue

        fiscal_year = _parse_int(raw_row.get("fiscal_year"), year_default) or year_default
        kind = _parse_kind(raw_row.get("kind"))
        cost_center = _cell_str(raw_row.get("cost_center"))
        scenario = _parse_scenario(raw_row.get("scenario"))
        allocated = _parse_float(raw_row.get("allocated"), default=-1.0)
        if allocated < 0:
            errors.append(
                {
                    "row": idx,
                    "code": "bi_budget_allocated_invalid",
                    "detail": str(raw_row.get("allocated")),
                    "sheet": "envelopes",
                }
            )
            continue

        payload: dict[str, Any] = {
            "name": name,
            "fiscal_year": fiscal_year,
            "kind": kind,
            "cost_center": cost_center,
            "scenario": scenario,
            "allocated": allocated,
            "currency": _cell_str(raw_row.get("currency")) or "TRY",
            "committed": max(0.0, _parse_float(raw_row.get("committed"))),
            "status": _parse_status(raw_row.get("status")),
            "owner": _cell_str(raw_row.get("owner")),
            "notes": _cell_str(raw_row.get("notes")),
            "actuals_sql": _cell_str(raw_row.get("actuals_sql")),
        }

        row_id = _cell_str(raw_row.get("id"))
        matched: dict[str, Any] | None = None
        if row_id:
            matched = next((e for e in existing if str(e.get("id")) == row_id), None)
            if matched:
                payload["id"] = row_id
        if not matched:
            matched = _match_existing(
                existing,
                name=name,
                fiscal_year=fiscal_year,
                kind=kind,
                cost_center=cost_center,
                scenario=scenario,
            )
            if matched and matched.get("id"):
                payload["id"] = matched["id"]

        try:
            saved = upsert_budget(engine, payload, tenant_id=tid, actor="import")
            if matched:
                updated += 1
                existing = [e for e in existing if e.get("id") != saved.get("id")] + [saved]
            else:
                created += 1
                existing.append(saved)
        except ValueError as exc:
            errors.append({"row": idx, "code": str(exc), "detail": name, "sheet": "envelopes"})
        except Exception as exc:
            errors.append(
                {
                    "row": idx,
                    "code": "bi_budget_import_row_invalid",
                    "detail": str(exc)[:200],
                    "sheet": "envelopes",
                }
            )

    # Refresh cache after envelope upserts
    existing = list_budgets(engine, tenant_id=tid)

    # Monthly periods
    for idx, prow in enumerate(period_rows, start=2):
        bid = _resolve_budget_id(existing, prow, year_default=year_default)
        if not bid:
            errors.append(
                {
                    "row": idx,
                    "code": "bi_budget_not_found",
                    "detail": _cell_str(prow.get("name")),
                    "sheet": "periods",
                }
            )
            continue
        months = prow.get("months") or {}
        periods_payload = [
            {"period_index": i, "allocated": float(months.get(i) or 0)}
            for i in range(1, 13)
        ]
        try:
            save_periods(engine, bid, periods_payload, tenant_id=tid)
            periods_updated += 1
        except Exception as exc:
            errors.append(
                {
                    "row": idx,
                    "code": "bi_budget_import_row_invalid",
                    "detail": str(exc)[:200],
                    "sheet": "periods",
                }
            )

    # Commitment lines (save_commitment already syncs envelope committed)
    for idx, crow in enumerate(commitment_rows, start=2):
        bid = _resolve_budget_id(existing, crow, year_default=year_default)
        if not bid:
            errors.append(
                {
                    "row": idx,
                    "code": "bi_budget_not_found",
                    "detail": _cell_str(crow.get("name")),
                    "sheet": "commitments",
                }
            )
            continue
        amount = _parse_float(crow.get("amount"), -1.0)
        if amount < 0:
            errors.append(
                {
                    "row": idx,
                    "code": "bi_commitment_amount_invalid",
                    "detail": str(crow.get("amount")),
                    "sheet": "commitments",
                }
            )
            continue
        entry = {
            "id": _cell_str(crow.get("id")) or None,
            "budget_id": bid,
            "description": _cell_str(crow.get("description")) or _cell_str(crow.get("notes")),
            "amount": amount,
            "currency": _cell_str(crow.get("currency")) or "TRY",
            "status": _parse_commitment_status(crow.get("status")),
            "due_date": _cell_str(crow.get("due_date")) or None,
        }
        try:
            save_commitment(engine, bid, entry, tenant_id=tid)
            commitments_upserted += 1
        except ValueError as exc:
            errors.append(
                {
                    "row": idx,
                    "code": str(exc),
                    "detail": _cell_str(crow.get("description")),
                    "sheet": "commitments",
                }
            )
        except Exception as exc:
            errors.append(
                {
                    "row": idx,
                    "code": "bi_budget_import_row_invalid",
                    "detail": str(exc)[:200],
                    "sheet": "commitments",
                }
            )

    if created == 0 and updated == 0 and periods_updated == 0 and commitments_upserted == 0 and not errors:
        raise ValueError("bi_budget_import_empty")

    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "periods_updated": periods_updated,
        "commitments_upserted": commitments_upserted,
        "errors": errors,
    }
