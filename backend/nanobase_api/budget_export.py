"""Budget board-pack export — PDF + multi-sheet Excel."""

from __future__ import annotations

import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy.engine import Engine

from nanobase_api.budget_fx import convert_amount, list_fx_rates
from nanobase_api.budget_lines import list_commitments, list_periods
from nanobase_api.budgets import budget_summary, list_budgets

_EXPORTS_DIR = Path("/tmp/nanobase-budget-exports")

# NanobaseAI brand palette (portal accent + BI status tokens)
_BRAND = {
    "navy": (15, 23, 42),  # #0F172A
    "violet": (124, 58, 237),  # #7C3AED
    "violet_muted": (139, 92, 246),  # #8B5CF6
    "blue": (37, 99, 235),  # #2563EB
    "indigo": (79, 70, 229),  # #4F46E5
    "surface": (245, 243, 255),  # #F5F3FF
    "white": (255, 255, 255),
    "slate": (71, 85, 105),  # #475569
    "slate_dark": (30, 41, 59),  # #1E293B
    "line": (226, 232, 240),  # #E2E8F0
    "zebra": (248, 250, 252),  # #F8FAFC
    "ok": (5, 150, 105),  # #059669
    "ok_bg": (209, 250, 229),  # #D1FAE5
    "watch": (217, 119, 6),  # #D97706
    "watch_bg": (254, 243, 199),  # #FEF3C7
    "over": (220, 38, 38),  # #DC2626
    "over_bg": (254, 226, 226),  # #FEE2E2
    "amber_bg": (255, 251, 235),  # #FFFBEB
    "amber_border": (251, 191, 36),  # #FBBF24
}

_LABELS: dict[str, dict[str, str]] = {
    "en": {
        "brand": "NanobaseAI",
        "brand_tag": "Business Intelligence",
        "title": "Budget board pack",
        "subtitle": "Fiscal year {year}",
        "meta": "Generated {generated} · {count} lines · {scenario}{currency_note}",
        "meta_generated": "Generated",
        "meta_lines": "{count} lines",
        "meta_scenario": "Scenario",
        "mixed_currency": "Mixed currencies — totals may be incomplete without FX.",
        "fx_missing": "{count} line(s) omitted due to missing FX rates.",
        "board_note": "Board pack export — read-only snapshot.",
        "kpi_allocated": "Allocated",
        "kpi_actual": "Actual",
        "kpi_remaining": "Remaining",
        "kpi_watch": "Watch",
        "section_by_kind": "By kind",
        "section_lines": "Budget lines",
        "empty": "No budget lines for this filter.",
        "col_name": "Name",
        "col_kind": "Kind",
        "col_cc": "Cost center",
        "col_currency": "Currency",
        "col_allocated": "Allocated",
        "col_committed": "Committed",
        "col_actual": "Actual",
        "col_remaining": "Remaining",
        "col_used": "Used %",
        "col_health": "Health",
        "col_status": "Status",
        "col_description": "Description",
        "col_amount": "Amount",
        "col_due": "Due",
        "col_month": "M{month}",
        "kind_opex": "OPEX",
        "kind_capex": "CAPEX",
        "kind_other": "Other",
        "health_ok": "OK",
        "health_watch": "Watch",
        "health_over": "Over",
        "health_na": "—",
        "sheet_summary": "Summary",
        "sheet_lines": "Lines",
        "sheet_kinds": "By kind",
        "sheet_monthly": "Monthly",
        "sheet_commitments": "Commitments",
        "sheet_scenario": "Scenario: {scenario}",
        "footer": "NanobaseAI · Confidential board pack · Page {page}",
        "showing_of": "Showing {shown} of {total} lines",
    },
    "tr": {
        "brand": "NanobaseAI",
        "brand_tag": "İş Zekâsı",
        "title": "Bütçe yönetim paketi",
        "subtitle": "Mali yıl {year}",
        "meta": "Oluşturulma {generated} · {count} kalem · {scenario}{currency_note}",
        "meta_generated": "Oluşturulma",
        "meta_lines": "{count} kalem",
        "meta_scenario": "Senaryo",
        "mixed_currency": "Karma para birimleri — FX olmadan toplamlar eksik olabilir.",
        "fx_missing": "{count} kalem eksik kur nedeniyle atlandı.",
        "board_note": "Yönetim paketi dışa aktarımı — salt okunur anlık görüntü.",
        "kpi_allocated": "Planlanan",
        "kpi_actual": "Gerçekleşen",
        "kpi_remaining": "Kalan",
        "kpi_watch": "İzleme",
        "section_by_kind": "Türe göre",
        "section_lines": "Bütçe kalemleri",
        "empty": "Bu filtre için bütçe kalemi yok.",
        "col_name": "Kalem",
        "col_kind": "Tür",
        "col_cc": "Maliyet merkezi",
        "col_currency": "Para birimi",
        "col_allocated": "Planlanan",
        "col_committed": "Taahhüt",
        "col_actual": "Gerçekleşen",
        "col_remaining": "Kalan",
        "col_used": "Kullanım %",
        "col_health": "Durum",
        "col_status": "Statü",
        "col_description": "Açıklama",
        "col_amount": "Tutar",
        "col_due": "Vade",
        "col_month": "A{month}",
        "kind_opex": "OPEX",
        "kind_capex": "CAPEX",
        "kind_other": "Diğer",
        "health_ok": "İyi",
        "health_watch": "İzle",
        "health_over": "Aşım",
        "health_na": "—",
        "sheet_summary": "Özet",
        "sheet_lines": "Kalemler",
        "sheet_kinds": "Türler",
        "sheet_monthly": "Aylık",
        "sheet_commitments": "Taahhütler",
        "sheet_scenario": "Senaryo: {scenario}",
        "footer": "NanobaseAI · Gizli yönetim paketi · Sayfa {page}",
        "showing_of": "{total} kalemden {shown} gösteriliyor",
    },
}


def _exports_dir() -> Path:
    _EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    return _EXPORTS_DIR


def _export_path(ext: str) -> Path:
    name = f"budget-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}.{ext}"
    return _exports_dir() / name


def _norm_locale(locale: str | None) -> str:
    loc = (locale or "en").strip().lower()[:2] or "en"
    return loc if loc in _LABELS else "en"


def _label(locale: str, key: str, **kwargs: Any) -> str:
    tpl = _LABELS.get(locale, _LABELS["en"]).get(key) or _LABELS["en"].get(key) or key
    try:
        return tpl.format(**kwargs)
    except Exception:
        return tpl


def _money(n: Any) -> str:
    try:
        return f"{float(n):,.0f}"
    except (TypeError, ValueError):
        return "—"


def _health_label(locale: str, health: str | None) -> str:
    if health == "ok":
        return _label(locale, "health_ok")
    if health == "watch":
        return _label(locale, "health_watch")
    if health == "over":
        return _label(locale, "health_over")
    return _label(locale, "health_na")


def _kind_label(locale: str, kind: str) -> str:
    k = (kind or "other").lower()
    if k == "opex":
        return _label(locale, "kind_opex")
    if k == "capex":
        return _label(locale, "kind_capex")
    return _label(locale, "kind_other")


def _norm_scenario(scenario: str | None) -> str:
    s = (scenario or "base").strip().lower() or "base"
    if s not in ("base", "optimistic", "pessimistic"):
        return "base"
    return s


def display_budget_lines(
    engine: Engine,
    *,
    fiscal_year: int,
    tenant_id: str | None,
    scenario: str = "base",
    reporting_currency: str | None = None,
) -> list[dict[str, Any]]:
    """Scenario-filtered lines; when reporting_currency set, FX-convert or omit missing rates."""
    scen = _norm_scenario(scenario)
    tid = str(tenant_id or "default")
    items = list_budgets(engine, fiscal_year=fiscal_year, scenario=scen, tenant_id=tid)
    report_ccy = (reporting_currency or "").strip().upper() or None
    if not report_ccy:
        return items

    rates = list_fx_rates(engine, tenant_id=tid)
    out: list[dict[str, Any]] = []
    for item in items:
        ccy = str(item.get("currency") or "TRY").upper()
        allocated = float(item.get("allocated") or 0)
        committed = float(item.get("committed") or 0)
        actual = (
            float(item.get("actual"))
            if item.get("actual") is not None and not item.get("actual_error")
            else None
        )
        remaining = item.get("remaining")
        remaining_n = float(remaining) if remaining is not None else None

        a = convert_amount(allocated, from_currency=ccy, to_currency=report_ccy, rates=rates)
        c = convert_amount(committed, from_currency=ccy, to_currency=report_ccy, rates=rates)
        act = (
            convert_amount(actual, from_currency=ccy, to_currency=report_ccy, rates=rates)
            if actual is not None
            else None
        )
        rem = (
            convert_amount(remaining_n, from_currency=ccy, to_currency=report_ccy, rates=rates)
            if remaining_n is not None
            else None
        )
        if a is None or c is None or (actual is not None and act is None):
            continue
        row = dict(item)
        row["allocated"] = a
        row["committed"] = c
        row["actual"] = act
        row["remaining"] = rem
        row["currency"] = report_ccy
        out.append(row)
    return out


def _health_colors(health: str | None) -> tuple[tuple[int, int, int], tuple[int, int, int]]:
    if health == "ok":
        return _BRAND["ok_bg"], _BRAND["ok"]
    if health == "watch":
        return _BRAND["watch_bg"], _BRAND["watch"]
    if health == "over":
        return _BRAND["over_bg"], _BRAND["over"]
    return _BRAND["zebra"], _BRAND["slate"]


def _resolve_pdf_fonts() -> tuple[str, Path | None, Path | None]:
    """Return (family, regular_ttf, bold_ttf). family is Helvetica when no TTF found."""
    regular = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    bold = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    if not regular.is_file():
        regular = Path("/usr/share/fonts/dejavu/DejaVuSans.ttf")
        bold = Path("/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf")
    if not regular.is_file():
        for candidate in (
            Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
            Path("/Library/Fonts/Arial Unicode.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        ):
            if candidate.is_file():
                return "DejaVu", candidate, candidate
        return "Helvetica", None, None
    return "DejaVu", regular, bold if bold.is_file() else regular


def _pdf_text(value: Any, *, font_family: str) -> str:
    s = str(value or "")
    if font_family != "Helvetica":
        return s
    return (
        s.replace("ı", "i")
        .replace("İ", "I")
        .replace("ğ", "g")
        .replace("Ğ", "G")
        .replace("ü", "u")
        .replace("Ü", "U")
        .replace("ş", "s")
        .replace("Ş", "S")
        .replace("ö", "o")
        .replace("Ö", "O")
        .replace("ç", "c")
        .replace("Ç", "C")
        .replace("·", "-")
        .encode("latin-1", "replace")
        .decode("latin-1")
    )


def build_budget_pdf(
    engine: Engine,
    *,
    fiscal_year: int,
    locale: str | None = None,
    tenant_id: str | None = None,
    scenario: str = "base",
    reporting_currency: str | None = None,
) -> Path:
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise ValueError("bi_budget_export_pdf_unavailable") from exc

    loc = _norm_locale(locale)
    scen = _norm_scenario(scenario)
    tid = str(tenant_id or "default")
    report_ccy = (reporting_currency or "").strip().upper() or None
    items = display_budget_lines(
        engine,
        fiscal_year=fiscal_year,
        tenant_id=tid,
        scenario=scen,
        reporting_currency=report_ccy,
    )
    summary = budget_summary(
        engine,
        fiscal_year=fiscal_year,
        tenant_id=tid,
        scenario=scen,
        reporting_currency=report_ccy,
    )
    totals = summary.get("totals") or {}
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    currency_note = f" · {report_ccy}" if report_ccy else ""

    font_family, regular, bold = _resolve_pdf_fonts()

    def _txt(value: Any) -> str:
        return _pdf_text(value, font_family=font_family)

    class BudgetPdf(FPDF):
        def header(self) -> None:
            # Soft page wash
            self.set_fill_color(*_BRAND["surface"])
            self.rect(0, 0, 210, 297, style="F")
            # Top brand bar
            self.set_fill_color(*_BRAND["navy"])
            self.rect(0, 0, 210, 28, style="F")
            # Accent ribbon (violet → indigo → blue)
            band_w = 210 / 3
            for i, color in enumerate((_BRAND["violet"], _BRAND["indigo"], _BRAND["blue"])):
                self.set_fill_color(*color)
                self.rect(i * band_w, 28, band_w + 0.2, 3.2, style="F")
            # Brand pill
            self.set_fill_color(*_BRAND["violet"])
            self.rect(12, 7.5, 38, 13, style="F", round_corners=True, corner_radius=3)
            self.set_xy(12, 9.5)
            self.set_text_color(*_BRAND["white"])
            self.set_font(font_family, "B", 9)
            self.cell(38, 9, _txt(_label(loc, "brand")), align="C")
            # Tagline + year chip on the right
            self.set_xy(54, 8)
            self.set_font(font_family, "", 8)
            self.set_text_color(203, 213, 225)
            self.cell(80, 6, _txt(_label(loc, "brand_tag")), align="L")
            self.set_fill_color(*_BRAND["blue"])
            self.rect(158, 8, 40, 12, style="F", round_corners=True, corner_radius=3)
            self.set_xy(158, 9.5)
            self.set_text_color(*_BRAND["white"])
            self.set_font(font_family, "B", 9)
            self.cell(40, 9, _txt(str(fiscal_year)), align="C")
            self.set_y(38)
            self.set_text_color(*_BRAND["slate_dark"])

        def footer(self) -> None:
            self.set_y(-16)
            self.set_draw_color(*_BRAND["violet"])
            self.set_line_width(0.4)
            self.line(12, self.get_y(), 198, self.get_y())
            self.set_y(-13)
            self.set_font(font_family, "", 7.5)
            self.set_text_color(*_BRAND["slate"])
            self.cell(
                0,
                8,
                _txt(_label(loc, "footer", page=self.page_no())),
                align="C",
            )

    pdf = BudgetPdf(format="A4")
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.set_margins(12, 38, 12)

    # Register Unicode fonts before the first page so header/footer can use them.
    if regular is not None:
        try:
            pdf.add_font("DejaVu", "", str(regular))
            pdf.add_font("DejaVu", "B", str(bold or regular))
            font_family = "DejaVu"
        except Exception:
            font_family = "Helvetica"

    pdf.add_page()

    # —— Hero title block ——
    pdf.set_fill_color(*_BRAND["white"])
    pdf.set_draw_color(*_BRAND["line"])
    pdf.set_line_width(0.3)
    pdf.rect(12, 38, 186, 28, style="DF", round_corners=True, corner_radius=4)
    pdf.set_xy(16, 41)
    pdf.set_font(font_family, "B", 18)
    pdf.set_text_color(*_BRAND["navy"])
    pdf.cell(0, 9, _txt(_label(loc, "title")), ln=True)
    pdf.set_x(16)
    pdf.set_font(font_family, "", 10)
    pdf.set_text_color(*_BRAND["violet"])
    pdf.cell(0, 6, _txt(_label(loc, "subtitle", year=fiscal_year)), ln=True)
    pdf.set_x(16)
    pdf.set_font(font_family, "", 8)
    pdf.set_text_color(*_BRAND["slate"])
    pdf.cell(
        0,
        5,
        _txt(
            _label(
                loc,
                "meta",
                generated=generated,
                count=len(items),
                scenario=scen,
                currency_note=currency_note,
            )
        ),
        ln=True,
    )
    pdf.set_y(72)

    # —— Warning banners ——
    def _warn_banner(message: str) -> None:
        y = pdf.get_y()
        pdf.set_fill_color(*_BRAND["amber_bg"])
        pdf.set_draw_color(*_BRAND["amber_border"])
        pdf.rect(12, y, 186, 9, style="DF", round_corners=True, corner_radius=2.5)
        pdf.set_xy(16, y + 1.5)
        pdf.set_font(font_family, "", 8)
        pdf.set_text_color(*_BRAND["watch"])
        pdf.cell(178, 6, _txt(message), ln=True)
        pdf.set_y(y + 11)

    if summary.get("mixed_currency"):
        _warn_banner(_label(loc, "mixed_currency"))
    fx_missing = summary.get("fx_missing") or []
    if fx_missing:
        _warn_banner(_label(loc, "fx_missing", count=len(fx_missing)))

    # —— KPI cards ——
    kpi_specs = (
        ("kpi_allocated", totals.get("allocated"), _BRAND["violet"], _money),
        ("kpi_actual", totals.get("actual"), _BRAND["blue"], _money),
        ("kpi_remaining", totals.get("remaining"), _BRAND["ok"], _money),
        ("kpi_watch", summary.get("budget_watch_count") or 0, _BRAND["watch"], lambda v: str(int(v or 0))),
    )
    card_w, card_h, gap = 44.5, 24, 2.7
    start_x, start_y = 12.0, pdf.get_y() + 1
    for i, (key, val, accent, fmt) in enumerate(kpi_specs):
        x = start_x + i * (card_w + gap)
        pdf.set_fill_color(*_BRAND["white"])
        pdf.set_draw_color(*_BRAND["line"])
        pdf.rect(x, start_y, card_w, card_h, style="DF", round_corners=True, corner_radius=3.5)
        # Left accent bar
        pdf.set_fill_color(*accent)
        pdf.rect(x, start_y, 2.2, card_h, style="F", round_corners=("TOP_LEFT", "BOTTOM_LEFT"), corner_radius=3.5)
        # Top accent chip
        pdf.set_fill_color(*accent)
        pdf.rect(x + 6, start_y + 3.5, 14, 2.2, style="F", round_corners=True, corner_radius=1)
        pdf.set_xy(x + 5, start_y + 7)
        pdf.set_font(font_family, "", 7.5)
        pdf.set_text_color(*_BRAND["slate"])
        pdf.cell(card_w - 8, 5, _txt(_label(loc, key)), ln=True)
        pdf.set_x(x + 5)
        pdf.set_font(font_family, "B", 12)
        pdf.set_text_color(*accent)
        pdf.cell(card_w - 8, 8, _txt(fmt(val)), ln=False)
    pdf.set_y(start_y + card_h + 8)

    # —— Section header ——
    y = pdf.get_y()
    pdf.set_fill_color(*_BRAND["violet"])
    pdf.rect(12, y + 1, 2.5, 8, style="F", round_corners=True, corner_radius=1)
    pdf.set_xy(17, y)
    pdf.set_font(font_family, "B", 12)
    pdf.set_text_color(*_BRAND["navy"])
    pdf.cell(0, 10, _txt(_label(loc, "section_lines")), ln=True)
    pdf.ln(1)

    # —— Table ——
    headers = [
        _label(loc, "col_name"),
        _label(loc, "col_kind"),
        _label(loc, "col_allocated"),
        _label(loc, "col_actual"),
        _label(loc, "col_health"),
    ]
    widths = [72, 24, 32, 32, 26]
    row_h = 7.2

    def _draw_table_header() -> None:
        pdf.set_fill_color(*_BRAND["navy"])
        pdf.set_text_color(*_BRAND["white"])
        pdf.set_font(font_family, "B", 8)
        x0 = pdf.get_x()
        y0 = pdf.get_y()
        # Full header background with rounded top
        pdf.rect(x0, y0, sum(widths), row_h + 0.4, style="F", round_corners=("TOP_LEFT", "TOP_RIGHT"), corner_radius=2.5)
        # Violet underline under header
        pdf.set_fill_color(*_BRAND["violet"])
        pdf.rect(x0, y0 + row_h + 0.2, sum(widths), 1.1, style="F")
        pdf.set_xy(x0, y0)
        for h, w in zip(headers, widths):
            pdf.cell(w, row_h, _txt(h)[:22], align="C")
        pdf.ln(row_h + 1.4)

    _draw_table_header()

    if not items:
        pdf.set_fill_color(*_BRAND["white"])
        pdf.set_draw_color(*_BRAND["line"])
        y = pdf.get_y()
        pdf.rect(12, y, sum(widths), 12, style="DF", round_corners=True, corner_radius=2)
        pdf.set_xy(16, y + 3)
        pdf.set_font(font_family, "", 9)
        pdf.set_text_color(*_BRAND["slate"])
        pdf.cell(0, 6, _txt(_label(loc, "empty")), ln=True)
    else:
        shown = items[:80]
        for idx, b in enumerate(shown):
            if pdf.get_y() > 265:
                pdf.add_page()
                _draw_table_header()

            health = b.get("health")
            health_bg, health_fg = _health_colors(health if isinstance(health, str) else None)
            zebra = _BRAND["zebra"] if idx % 2 == 0 else _BRAND["white"]
            vals = [
                (str(b.get("name") or "")[:38], "L", zebra, _BRAND["slate_dark"]),
                (_kind_label(loc, str(b.get("kind") or "")), "C", zebra, _BRAND["slate"]),
                (_money(b.get("allocated")), "R", zebra, _BRAND["slate_dark"]),
                (_money(b.get("actual")), "R", zebra, _BRAND["slate_dark"]),
                (_health_label(loc, health if isinstance(health, str) else None), "C", health_bg, health_fg),
            ]
            x0 = pdf.l_margin
            y0 = pdf.get_y()
            # Row background
            pdf.set_fill_color(*zebra)
            pdf.rect(x0, y0, sum(widths), row_h, style="F")
            # Kind tint chip background strip
            kind = str(b.get("kind") or "").lower()
            kind_accent = {
                "opex": _BRAND["violet_muted"],
                "capex": _BRAND["blue"],
                "other": _BRAND["slate"],
            }.get(kind, _BRAND["slate"])
            # Draw cells
            pdf.set_xy(x0, y0)
            pdf.set_font(font_family, "", 7.5)
            for i, (v, align, bg, fg) in enumerate(vals):
                w = widths[i]
                if i == 4:
                    # Health badge pill
                    pdf.set_fill_color(*bg)
                    badge_w = w - 4
                    badge_x = pdf.get_x() + 2
                    pdf.rect(badge_x, y0 + 1.2, badge_w, row_h - 2.4, style="F", round_corners=True, corner_radius=2)
                    pdf.set_text_color(*fg)
                    pdf.set_font(font_family, "B", 7)
                    pdf.set_xy(badge_x, y0 + 1.2)
                    pdf.cell(badge_w, row_h - 2.4, _txt(v)[:16], align="C")
                    pdf.set_xy(x0 + sum(widths[: i + 1]), y0)
                    pdf.set_font(font_family, "", 7.5)
                elif i == 1:
                    pdf.set_text_color(*kind_accent)
                    pdf.set_font(font_family, "B", 7)
                    pdf.cell(w, row_h, _txt(v)[:18], align=align)
                    pdf.set_font(font_family, "", 7.5)
                else:
                    pdf.set_text_color(*fg)
                    pdf.cell(w, row_h, _txt(v)[:40], align=align)
            pdf.ln(row_h)
            # Subtle bottom rule
            pdf.set_draw_color(*_BRAND["line"])
            pdf.set_line_width(0.15)
            pdf.line(x0, pdf.get_y(), x0 + sum(widths), pdf.get_y())

        if len(items) > len(shown):
            pdf.ln(3)
            pdf.set_font(font_family, "", 8)
            pdf.set_text_color(*_BRAND["slate"])
            pdf.cell(
                0,
                6,
                _txt(_label(loc, "showing_of", shown=len(shown), total=len(items))),
                ln=True,
            )

    # —— Board note ——
    pdf.ln(6)
    y = pdf.get_y()
    if y < 270:
        pdf.set_fill_color(237, 233, 254)  # soft violet
        pdf.set_draw_color(*_BRAND["violet_muted"])
        pdf.rect(12, y, 186, 10, style="DF", round_corners=True, corner_radius=2.5)
        pdf.set_xy(16, y + 2)
        pdf.set_font(font_family, "", 8)
        pdf.set_text_color(*_BRAND["violet"])
        pdf.cell(178, 6, _txt(_label(loc, "board_note")), ln=True)

    path = _export_path("pdf")
    out = pdf.output()
    path.write_bytes(bytes(out) if not isinstance(out, (bytes, bytearray)) else bytes(out))
    return path


def build_budget_xlsx(
    engine: Engine,
    *,
    fiscal_year: int,
    locale: str | None = None,
    tenant_id: str | None = None,
    scenario: str = "base",
    reporting_currency: str | None = None,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    loc = _norm_locale(locale)
    scen = _norm_scenario(scenario)
    tid = str(tenant_id or "default")
    report_ccy = (reporting_currency or "").strip().upper() or None
    items = display_budget_lines(
        engine,
        fiscal_year=fiscal_year,
        tenant_id=tid,
        scenario=scen,
        reporting_currency=report_ccy,
    )
    summary = budget_summary(
        engine,
        fiscal_year=fiscal_year,
        tenant_id=tid,
        scenario=scen,
        reporting_currency=report_ccy,
    )
    totals = summary.get("totals") or {}

    wb = Workbook()
    ws = wb.active
    ws.title = _label(loc, "sheet_summary")[:31]
    navy = PatternFill("solid", fgColor="0F172A")
    white = Font(bold=True, color="FFFFFF", size=14)
    ws["A1"] = _label(loc, "title")
    ws["A1"].fill = navy
    ws["A1"].font = white
    ws.merge_cells("A1:D1")
    ws["A2"] = _label(loc, "subtitle", year=fiscal_year)
    ws["A3"] = _label(loc, "sheet_scenario", scenario=scen)
    if report_ccy:
        ws["B3"] = report_ccy
    if summary.get("mixed_currency"):
        ws["A4"] = _label(loc, "mixed_currency")
    fx_missing = summary.get("fx_missing") or []
    if fx_missing:
        ws["A4"] = _label(loc, "fx_missing", count=len(fx_missing))
    ws["A5"] = _label(loc, "kpi_allocated")
    ws["B5"] = float(totals.get("allocated") or 0) if totals.get("allocated") is not None else None
    ws["A6"] = _label(loc, "kpi_actual")
    ws["B6"] = float(totals.get("actual") or 0) if totals.get("actual") is not None else None
    ws["A7"] = _label(loc, "kpi_remaining")
    ws["B7"] = float(totals.get("remaining") or 0) if totals.get("remaining") is not None else None
    ws["A8"] = _label(loc, "kpi_watch")
    ws["B8"] = int(summary.get("budget_watch_count") or 0)
    ws["A9"] = _label(loc, "board_note")
    ws.merge_cells("A9:D9")
    ws["A9"].alignment = Alignment(wrap_text=True)
    for r in range(5, 8):
        if ws[f"B{r}"].value is not None:
            ws[f"B{r}"].number_format = "#,##0"

    lines = wb.create_sheet(_label(loc, "sheet_lines")[:31])
    headers = [
        _label(loc, "col_name"),
        _label(loc, "col_cc"),
        _label(loc, "col_kind"),
        _label(loc, "col_currency"),
        _label(loc, "col_allocated"),
        _label(loc, "col_committed"),
        _label(loc, "col_actual"),
        _label(loc, "col_remaining"),
        _label(loc, "col_used"),
        _label(loc, "col_health"),
        _label(loc, "col_status"),
    ]
    for ci, h in enumerate(headers, start=1):
        cell = lines.cell(1, ci, h)
        cell.fill = navy
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True)
    lines.freeze_panes = "A2"
    for ri, b in enumerate(items, start=2):
        vals = [
            b.get("name"),
            b.get("cost_center"),
            _kind_label(loc, str(b.get("kind") or "")),
            b.get("currency") or report_ccy or "TRY",
            float(b.get("allocated") or 0),
            float(b.get("committed") or 0),
            float(b.get("actual") or 0) if b.get("actual") is not None else None,
            float(b.get("remaining") or 0) if b.get("remaining") is not None else None,
            float(b.get("used_pct") or 0) if b.get("used_pct") is not None else None,
            _health_label(loc, b.get("health")),
            b.get("status"),
        ]
        for ci, v in enumerate(vals, start=1):
            cell = lines.cell(ri, ci, v)
            if ci in (5, 6, 7, 8) and isinstance(v, (int, float)):
                cell.number_format = "#,##0"
            if ci == 9 and isinstance(v, (int, float)):
                cell.number_format = "0.0"
        health = b.get("health")
        fill = None
        if health == "over":
            fill = PatternFill("solid", fgColor="FEE2E2")
        elif health == "watch":
            fill = PatternFill("solid", fgColor="FEF3C7")
        elif health == "ok":
            fill = PatternFill("solid", fgColor="D1FAE5")
        if fill:
            lines.cell(ri, 10).fill = fill

    for i in range(1, 12):
        lines.column_dimensions[get_column_letter(i)].width = 14
    lines.column_dimensions["A"].width = 28

    kinds = wb.create_sheet(_label(loc, "sheet_kinds")[:31])
    kinds.append(
        [
            _label(loc, "col_kind"),
            _label(loc, "col_allocated"),
            _label(loc, "col_actual"),
            _label(loc, "col_remaining"),
            "#",
        ]
    )
    for cell in kinds[1]:
        cell.fill = navy
        cell.font = Font(bold=True, color="FFFFFF")
    by_kind = summary.get("by_kind") or {}
    for kind in ("opex", "capex", "other"):
        bucket = by_kind.get(kind) or {}
        kinds.append(
            [
                _kind_label(loc, kind),
                float(bucket.get("allocated") or 0),
                float(bucket.get("actual") or 0),
                float(bucket.get("remaining") or 0),
                int(bucket.get("count") or 0),
            ]
        )

    monthly = wb.create_sheet(_label(loc, "sheet_monthly")[:31])
    month_headers = [_label(loc, "col_name"), _label(loc, "col_kind")] + [
        _label(loc, "col_month", month=i) for i in range(1, 13)
    ]
    monthly.append(month_headers)
    for cell in monthly[1]:
        cell.fill = navy
        cell.font = Font(bold=True, color="FFFFFF")
    try:
        for b in items:
            periods = list_periods(engine, str(b.get("id")), tenant_id=tid)
            by_idx = {int(p.get("period_index") or 0): p for p in periods}
            monthly.append(
                [
                    b.get("name"),
                    _kind_label(loc, str(b.get("kind") or "")),
                    *[float((by_idx.get(i) or {}).get("allocated") or 0) for i in range(1, 13)],
                ]
            )
    except Exception:
        pass

    commits = wb.create_sheet(_label(loc, "sheet_commitments")[:31])
    commits.append(
        [
            _label(loc, "col_name"),
            _label(loc, "col_kind"),
            _label(loc, "col_cc"),
            _label(loc, "col_description"),
            _label(loc, "col_amount"),
            _label(loc, "col_currency"),
            _label(loc, "col_status"),
            _label(loc, "col_due"),
            "ID",
        ]
    )
    for cell in commits[1]:
        cell.fill = navy
        cell.font = Font(bold=True, color="FFFFFF")
    try:
        for b in items:
            for c in list_commitments(engine, str(b.get("id")), tenant_id=tid):
                commits.append(
                    [
                        b.get("name"),
                        _kind_label(loc, str(b.get("kind") or "")),
                        b.get("cost_center") or "",
                        c.get("description") or "",
                        float(c.get("amount") or 0),
                        c.get("currency") or b.get("currency") or report_ccy or "TRY",
                        c.get("status") or "",
                        c.get("due_date") or "",
                        c.get("id") or "",
                    ]
                )
    except Exception:
        pass

    path = _export_path("xlsx")
    wb.save(path)
    return path


def export_budget_board_pack(
    engine: Engine,
    *,
    format: str,
    fiscal_year: int,
    locale: str | None = None,
    tenant_id: str | None = None,
    scenario: str = "base",
    reporting_currency: str | None = None,
) -> Path:
    fmt = (format or "pdf").lower().strip()
    if fmt == "xlsx":
        return build_budget_xlsx(
            engine,
            fiscal_year=fiscal_year,
            locale=locale,
            tenant_id=tenant_id,
            scenario=scenario,
            reporting_currency=reporting_currency,
        )
    if fmt == "pdf":
        return build_budget_pdf(
            engine,
            fiscal_year=fiscal_year,
            locale=locale,
            tenant_id=tenant_id,
            scenario=scenario,
            reporting_currency=reporting_currency,
        )
    raise ValueError("bi_budget_export_format_invalid")
