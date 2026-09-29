"""Excel indirme kabulü — CSV indiren her ekranın Excel eşi (test sunucusunda, aday ağaçla açılan yan köprüye karşı).

Her uç için gerçek veriyle:
  E1  CSV (eski adres) hâlâ CSV döner (text/csv).
  E2  Excel adresi (`bicim=xlsx`) 200, `application/vnd.openxmlformats-…sheet`, dosya adı CSV'nin adıyla aynı, uzantı .xlsx.
  E3  Dosya gerçek Excel: openpyxl açar (varsa LibreOffice ile de açılır ve aynı satır sayısını verir).
  E4  Satır ve kolon sayısı CSV ile birebir; başlık birebir.
  E5  Her hücre CSV'deki değerle aynı: metin aynen; sayı hücresi CSV'deki yazımın sayısal değeri; tarih hücresi aynı gün/saat;
      boş hücre boş. (Karşılaştırma `csv_excel`'in kendi okuyucusuyla değil, bağımsız bir çözümleyiciyle yapılır.)
  E6  Sayısal görünümlü hücrelerin kaçı sayı olarak yazıldı, kaçı metin kaldı (bilgi; kod kolonları bilerek metin kalır).
  E7  Sayfa kapısı aynı: yetkisiz oturumla Excel adresi dosya vermez (COOKIE2 verilirse).
Tarayıcıda üretilen CSV'ler (baskı öneri, dahili rehber) ön yüzdeki biçimle yeniden üretilir ve `POST /api/v1/export/xlsx` ile aynı
denetimden geçer. Liste indirme POST'ları (e-ticaret hedef grubu, okur listesi) dışa aktarım kaydı yazdığı için burada
çağrılmaz; POST çevirisi birim testinde (`test_csv_excel.py`).

Ortam: BASE (ör. http://127.0.0.1:8796), COOKIE (timasai kısa oturumu, «timas_session=…»), isteğe bağlı COOKIE2 (yetkisiz),
SEMANTIC_CALLER_TOKEN (köprü istiyorsa), PYTHONPATH=<aday ağaç>/backend. Kabul hiçbir yere yazmaz (yalnız GET ve çeviri ucu).
Kullanım: python kabul.py [--out kanit.json] [--only anahtar1,anahtar2]
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime
from typing import Any, Callable, Optional

from openpyxl import load_workbook

from semantic_bridge import csv_excel as X

BASE = os.environ.get("BASE", "http://127.0.0.1:8796").rstrip("/")
COOKIE = os.environ.get("COOKIE", "")
COOKIE2 = os.environ.get("COOKIE2", "")
CALLER = os.environ.get("SEMANTIC_CALLER_TOKEN", "")
XLSX = X.XLSX_MEDIA

results: list[dict[str, Any]] = []
# `yerinde.py` köprüyü süreç içinde kurar ve istekleri buradan geçirir (HTTP yerine); None ise BASE'e HTTP.
TRANSPORT: Optional[Callable[[str, dict, str, Optional[bytes]], tuple[int, dict, bytes]]] = None


def check(name: str, ok: Optional[bool], detail: str = "") -> None:
    state = "GEÇTİ" if ok is True else ("UYARI" if ok is None else "KALDI")
    results.append({"ad": name, "durum": state, "ayrinti": detail})
    print(f"{state} {name}" + (f" — {detail}" if detail else ""), flush=True)


def request(path: str, *, cookie: Optional[str] = None, method: str = "GET", body: Any = None, timeout: int = 900):
    h = {"Cookie": COOKIE if cookie is None else cookie}
    if CALLER:
        h["X-Semantic-Caller"] = CALLER
    data = None
    if body is not None:
        data = json.dumps(body).encode()
        h["Content-Type"] = "application/json"
    if TRANSPORT is not None:
        return TRANSPORT(path, h, method, data)
    req = urllib.request.Request(BASE + path, headers=h, method=method, data=data)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, dict(r.headers), r.read()
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()


def get_json(path: str) -> Any:
    st, _, b = request(path)
    if st != 200:
        raise RuntimeError(f"{path}: {st} {b[:200]!r}")
    return json.loads(b or b"{}")


def with_xlsx(path: str) -> str:
    return path + ("&" if "?" in path else "?") + "bicim=xlsx"


# ------------------------------------------------------------------ bağımsız karşılaştırma


def _csv_rows(data: bytes) -> list[list[str]]:
    import csv

    text = data.decode("utf-8-sig", errors="replace")
    first = text.split("\n", 1)[0]
    delim = max((";", ",", "\t"), key=first.count) if any(d in first for d in ";,\t") else ";"
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    while rows and not any(c.strip() for c in rows[-1]):
        rows.pop()
    return rows


def _numbers_of(s: str) -> list[float]:
    """CSV yazımının olası sayısal değerleri (Türkçe ve nokta ondalık, yüzde)."""
    t = s.strip().replace(" ", "")
    pct = t.startswith("%") or t.endswith("%")
    t = t.strip("% ")
    out = []
    for cand in (t.replace(".", "").replace(",", "."), t, t.replace(",", "")):
        try:
            v = float(cand)
        except ValueError:
            continue
        out.append(v / 100 if pct else v)
    return out


def _date_parts(s: str) -> Optional[tuple[int, ...]]:
    t = s.strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2})(?::(\d{2}))?)?", t)
    if m:
        return tuple(int(x or 0) for x in m.groups())
    m = re.match(r"^(\d{2})\.(\d{2})\.(\d{4})(?: (\d{2}):(\d{2})(?::(\d{2}))?)?$", t)
    if m:
        d, mo, y, H, M, S = m.groups()
        return (int(y), int(mo), int(d), int(H or 0), int(M or 0), int(S or 0))
    return None


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def same(raw: str, val: Any) -> bool:
    if val is None:
        return raw.strip() == ""
    if isinstance(val, str):
        # Excel'in kabul etmediği kontrol karakterleri atılır; 32.767'yi aşan hücre «…» ile biter (csv_excel kuralı).
        r = _CTRL.sub("", raw)
        return val == r or (len(val) == 32767 and val.endswith("…") and r.startswith(val[:-1]))
    if isinstance(val, (datetime, date)):
        p = _date_parts(raw)
        v = (val.year, val.month, val.day, getattr(val, "hour", 0), getattr(val, "minute", 0), getattr(val, "second", 0))
        return p is not None and p == v
    if isinstance(val, bool):
        return False
    if isinstance(val, (int, float)):
        return any(abs(c - float(val)) <= 1e-9 * max(1.0, abs(c)) for c in _numbers_of(raw))
    return False


def compare(csv_bytes: bytes, xlsx_bytes: bytes) -> dict[str, Any]:
    exp = _csv_rows(csv_bytes)
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    got = [[c.value for c in r] for r in ws.iter_rows()]
    while got and all(v is None for v in got[-1]):
        got.pop()
    rep: dict[str, Any] = {"csvRows": len(exp), "xlsxRows": len(got), "bad": [], "num": 0, "date": 0, "text": 0,
                           "numLookingText": 0, "sheet": ws.title}
    if len(exp) != len(got):
        rep["bad"].append(f"satır sayısı CSV {len(exp)} ≠ Excel {len(got)}")
    for i, (er, gr) in enumerate(zip(exp, got)):
        ncol = max(len(er), len(gr))
        for j in range(ncol):
            raw = er[j] if j < len(er) else ""
            val = gr[j] if j < len(gr) else None
            if i == 0:
                if (val or "") != raw:
                    rep["bad"].append(f"başlık {j + 1}: {raw!r} ≠ {val!r}")
                continue
            if not same(raw, val):
                if len(rep["bad"]) < 12:
                    rep["bad"].append(f"satır {i + 1} kolon {j + 1} ({exp[0][j] if j < len(exp[0]) else '?'}): CSV {raw!r} ≠ Excel {val!r}")
                else:
                    rep.setdefault("more", 0)
                    rep["more"] += 1
            elif isinstance(val, (int, float)) and not isinstance(val, bool):
                rep["num"] += 1
            elif isinstance(val, (datetime, date)):
                rep["date"] += 1
            elif isinstance(val, str):
                rep["text"] += 1
                if _numbers_of(raw) and re.fullmatch(r"-?[\d.,]+%?", raw.strip()):
                    rep["numLookingText"] += 1
    return rep


def structure_ok(xlsx_bytes: bytes) -> tuple[bool, str]:
    """Dosya yapısı: zip bütün, Office içerik türleri ve çalışma sayfası var, her XML parçası iyi biçimli."""
    import zipfile
    from xml.etree import ElementTree as ET

    try:
        z = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
        bad = z.testzip()
        if bad:
            return False, f"bozuk parça {bad}"
        names = z.namelist()
        for need in ("[Content_Types].xml", "xl/workbook.xml", "xl/worksheets/sheet1.xml"):
            if need not in names:
                return False, f"{need} yok"
        for n in names:
            if n.endswith((".xml", ".rels")):
                ET.fromstring(z.read(n))
        return True, f"{len(names)} parça"
    except Exception as e:  # noqa: BLE001
        return False, str(e)[:160]


_CALC: Optional[bool] = None


def libreoffice_rows(xlsx_bytes: bytes) -> Optional[int]:
    """LibreOffice Calc ile açılış (kuruluysa). Calc yoksa (yalnız yazı modülü) None — sınama atlanır."""
    global _CALC
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    if _CALC is None:
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.append(["a", 1])
        buf = io.BytesIO()
        wb.save(buf)
        _CALC = False
        _CALC = _convert_rows(soffice, buf.getvalue()) == 1
        if not _CALC:
            check("LibreOffice Calc", None, "sunucuda Calc kurulu değil (sade bir Excel dosyasını da açamıyor); bu sınama atlandı")
    return _convert_rows(soffice, xlsx_bytes) if _CALC else None


def _convert_rows(soffice: str, xlsx_bytes: bytes) -> int:
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "a.xlsx")
        open(src, "wb").write(xlsx_bytes)
        subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{d}/profil", "--convert-to", "csv:Text - txt - csv (StarCalc):59,34,76",
                        "--outdir", d, src], capture_output=True, timeout=180)
        out = os.path.join(d, "a.csv")
        if not os.path.exists(out):
            return -1
        return len(_csv_rows(open(out, "rb").read()))


SAVE_DIR = os.environ.get("SAVE_DIR", "")


def save(key: str, name: str, data: bytes) -> None:
    if SAVE_DIR:
        os.makedirs(SAVE_DIR, exist_ok=True)
        open(os.path.join(SAVE_DIR, f"{name}.xlsx"), "wb").write(data)


def filename_of(headers: dict[str, str]) -> str:
    cd = next((v for k, v in headers.items() if k.lower() == "content-disposition"), "")
    return re.sub(r"\.(csv|xlsx)$", "", X.xlsx_name(cd, "")[: -len(".xlsx")], flags=re.I) if cd else ""


# ------------------------------------------------------------------ bir ucu sına


def test_endpoint(key: str, path: str, *, need_rows: bool = False) -> Optional[dict[str, Any]]:
    t0 = time.time()
    st, h, csv_b = request(path)
    ctype = next((v for k, v in h.items() if k.lower() == "content-type"), "")
    if st == 403:
        # Kişinin rolünde yoksa CSV de inmez; Excel adresinin de aynı kapıda durduğu denetlenir.
        st2, h2, _ = request(with_xlsx(path))
        c2 = next((v for k, v in h2.items() if k.lower() == "content-type"), "")
        check(f"{key} yetki", None, f"test kişisinin rolünde yok (CSV 403); Excel adresi de {st2} {c2.split(';')[0]} — dosya inmedi")
        return None
    if st != 200:
        check(f"{key} E1 CSV", False, f"{path} → {st} {csv_b[:200]!r}")
        return None
    check(f"{key} E1 CSV hâlâ CSV", ctype.startswith("text/csv"), ctype)
    st2, h2, xb = request(with_xlsx(path))
    ctype2 = next((v for k, v in h2.items() if k.lower() == "content-type"), "")
    name_csv, name_x = filename_of(h), filename_of(h2)
    cd2 = next((v for k, v in h2.items() if k.lower() == "content-disposition"), "")
    ok2 = st2 == 200 and ctype2 == XLSX and ".xlsx" in cd2 and (not name_csv or name_csv == name_x)
    check(f"{key} E2 Excel cevabı", ok2, f"{st2} {ctype2} «{name_x}.xlsx» (CSV «{name_csv}.csv»)")
    if not ok2:
        return None
    try:
        rep = compare(csv_b, xb)
    except Exception as e:  # noqa: BLE001
        check(f"{key} E3 Excel açılır", False, str(e)[:200])
        return None
    check(f"{key} E3 Excel açılır", True, f"sayfa «{rep['sheet']}», {len(xb):,} bayt")
    so, sd = structure_ok(xb)
    check(f"{key} E3 dosya yapısı", so, sd)
    save(key, name_x or key, xb)
    lo = libreoffice_rows(xb)
    if lo is not None:
        check(f"{key} E3 LibreOffice açar", lo == rep["xlsxRows"], f"LibreOffice {lo} satır, Excel {rep['xlsxRows']}")
    body_rows = rep["csvRows"] - 1
    check(f"{key} E4 satır/başlık birebir", rep["csvRows"] == rep["xlsxRows"] and not [b for b in rep["bad"] if b.startswith("başlık")],
          f"{body_rows:,} veri satırı")
    check(f"{key} E5 hücreler CSV ile aynı", not rep["bad"],
          "; ".join(rep["bad"]) + (f" (+{rep['more']} fark)" if rep.get("more") else "") if rep["bad"] else
          f"{rep['num']:,} sayı, {rep['date']:,} tarih, {rep['text']:,} metin hücresi")
    if rep["numLookingText"]:
        check(f"{key} E6 sayı gibi görünüp metin kalan", None, f"{rep['numLookingText']:,} hücre (kod/numara kolonları bilerek metin)")
    if need_rows and body_rows == 0:
        check(f"{key} veri satırı", None, "uç boş liste döndü; hücre karşılaştırması yalnız başlıkta yapıldı")
    if COOKIE2:
        st3, h3, b3 = request(with_xlsx(path), cookie=COOKIE2)
        c3 = next((v for k, v in h3.items() if k.lower() == "content-type"), "")
        check(f"{key} E7 yetkisiz kişiye Excel inmez", st3 in (401, 403) and c3 != XLSX, f"{st3} {c3}")
    return {"key": key, "path": path, "rows": body_rows, "ms": int((time.time() - t0) * 1000), **{k: rep[k] for k in ("num", "date", "text")}}


def first_id(list_path: str, pick: Callable[[dict], bool] = lambda _: True, items_key: str = "items") -> list[str]:
    try:
        d = get_json(list_path)
    except Exception as e:  # noqa: BLE001
        print(f"   liste okunamadı {list_path}: {e}")
        return []
    items = d.get(items_key) if isinstance(d, dict) else d
    return [str(i["id"]) for i in (items or []) if isinstance(i, dict) and i.get("id") and pick(i)]


def test_with_ids(key: str, ids: list[str], make: Callable[[str], str], limit: int = 6) -> None:
    """Kimlikli uç: veri satırı olan ilk kaydı arar (en çok `limit` kayıt); hiçbiri satır vermezse başlıkla sınar."""
    if not ids:
        check(f"{key} kayıt", None, "koşula uyan kayıt yok; uç sınanamadı")
        return
    tried = None
    for i in ids[:limit]:
        st, _, b = request(make(i))
        if st == 200 and len(_csv_rows(b)) > 1:
            test_endpoint(key, make(i))
            return
        tried = tried or i
    test_endpoint(key, make(tried), need_rows=True)


# ------------------------------------------------------------------ tarayıcıda üretilen CSV'ler


def _js_num(v: Any) -> str:
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def _esc(v: Any) -> str:
    s = "" if v is None else _js_num(v)
    return f'"{s.replace(chr(34), chr(34) * 2)}"' if re.search(r'[;"\n]', s) else s


def baski_csv() -> tuple[str, str, int]:
    """BaskiOneri.tsx `csvOf` ile aynı biçim (ilk görünüm, süzgeçsiz)."""
    snap = get_json("/api/v1/management/reports/baski-oneri")
    view = snap["data"]["views"][0]
    lines = [";".join(_esc(c["label"]) for c in view["columns"])]
    for r in view["rows"]:
        cells = []
        for i, c in enumerate(view["columns"]):
            v = r[i] if i < len(r) else None
            if c.get("format") in ("text", "oneri", "date") or not isinstance(v, (int, float)) or isinstance(v, bool):
                cells.append(_esc(v))
            else:
                cells.append(_js_num(v).replace(".", ",", 1))
        lines.append(";".join(cells))
    return "﻿" + "\r\n".join(lines), f"baski-oneri-{view['id']}-{snap['data'].get('asOf') or 'rapor'}.csv", len(view["rows"])


def rehber_csv() -> tuple[str, str, int]:
    """KampusPage.tsx `directoryCsv` ile aynı biçim."""
    people = get_json("/api/v1/people")["items"]
    cols = [("Ad Soyad", "name"), ("Ünvan", "title"), ("Birim", "unit"), ("Dahili", "extension"), ("Kat", "floor"),
            ("Masa", "desk"), ("Cep", "mobile"), ("Telefon", "phone"), ("E-posta", "email")]

    def cell(v: Any) -> str:
        return '"' + str(v or "").replace('"', '""') + '"'

    lines = [";".join(cell(h) for h, _ in cols)] + [";".join(cell(p.get(k)) for _, k in cols) for p in people]
    return "﻿" + "\r\n".join(lines), f"dahili-rehber-{date.today().isoformat()}.csv", len(people)


def test_client_csv(key: str, make: Callable[[], tuple[str, str, int]]) -> None:
    try:
        text, name, n = make()
    except Exception as e:  # noqa: BLE001
        check(f"{key} veri", None, f"ekran verisi okunamadı: {str(e)[:200]}")
        return
    st, h, xb = request("/api/v1/export/xlsx", method="POST", body={"csv": text, "filename": name})
    ctype = next((v for k, v in h.items() if k.lower() == "content-type"), "")
    xname = filename_of(h)
    ok = st == 200 and ctype == XLSX and xname == name[:-4]
    check(f"{key} E2 Excel cevabı", ok, f"{st} {ctype} «{xname}.xlsx»")
    if not ok:
        return
    rep = compare(text.encode("utf-8"), xb)
    so, sd = structure_ok(xb)
    check(f"{key} E3 dosya yapısı", so, sd)
    lo = libreoffice_rows(xb)
    if lo is not None:
        check(f"{key} E3 LibreOffice açar", lo == rep["xlsxRows"], f"LibreOffice {lo} satır")
    check(f"{key} E4 satır/başlık birebir", rep["csvRows"] == rep["xlsxRows"] == n + 1, f"{n:,} veri satırı")
    check(f"{key} E5 hücreler CSV ile aynı", not rep["bad"],
          "; ".join(rep["bad"]) if rep["bad"] else f"{rep['num']:,} sayı, {rep['date']:,} tarih, {rep['text']:,} metin hücresi")


# ------------------------------------------------------------------ uç listesi


def plain() -> list[tuple[str, str]]:
    return [
        ("pazar-matris", "/api/v1/pazar/matrix/export.csv"),
        ("dijital-firsatlar-ekitap", "/api/v1/dijital/opportunities/export.csv?tur=ekitap"),
        ("dijital-firsatlar-sesli", "/api/v1/dijital/opportunities/export.csv?tur=sesli"),
        ("dijital-satislar", "/api/v1/dijital/sales/export.csv"),
        ("ik-zorunlu-egitim", "/api/v1/hr/learning/export/expiring.csv"),
        ("eticaret-farklar", "/api/v1/eticaret/diffs/export.csv"),
        ("kurumsal-sessiz-bayiler", "/api/v1/corporate/b2b/dealers.csv"),
        ("kurumsal-one-cikanlar", "/api/v1/corporate/b2b/highlights.csv"),
        ("bayi-listesi", "/api/v1/dealers/list/export.csv"),
        ("ceviri-terim-bankasi", "/api/v1/editorial/translation/terms/export.csv"),
        ("yonetim-promtlar", "/api/v1/admin/prompts/export.csv?days=7"),
        ("finans-karlilik-kitap", "/api/v1/finance/profitability/export.csv"),
        ("finans-karlilik-kanal", "/api/v1/finance/profitability/export.csv?by=kanal"),
        ("seo-gunbatimi", "/api/v1/seo-geo/sunset/export.csv"),
        ("seo-anahtar-haritasi", "/api/v1/seo-geo/keymap/export.csv"),
        ("seo-benzer", "/api/v1/seo-geo/similar/export.csv"),
        ("seo-yonlendirmeler", "/api/v1/seo-geo/redirects/export.csv"),
        ("seo-is-listesi", "/api/v1/seo-geo/worklist/export.csv"),
        ("musteri-hesaplar", "/api/v1/musteri/accounts/export.csv"),
        ("musteri-veri-sagligi", "/api/v1/musteri/health/export.csv"),
        ("backlist-firsatlar", "/api/v1/marketing/backlist/export.csv"),
    ]


def with_ids() -> list[tuple[str, Callable[[], list[str]], Callable[[str], str]]]:
    q = urllib.parse.quote
    return [
        ("ik-anket", lambda: first_id("/api/v1/hr/engagement/surveys",
                                      lambda s: s.get("state") == "kapandi" and s.get("minGroup") and (s.get("responded") or 0) >= s["minGroup"]),
         lambda i: f"/api/v1/hr/engagement/surveys/{q(i)}/export.csv"),
        ("ik-degerlendirme", lambda: first_id("/api/v1/hr/performance/cycles"),
         lambda i: f"/api/v1/hr/performance/cycles/{q(i)}/export.csv"),
        ("okur-crm-yeni-kisiler", lambda: first_id("/api/v1/readers/imports", lambda x: not x.get("purgedAt")),
         lambda i: f"/api/v1/readers/imports/{q(i)}/crm.csv"),
        ("ceviri-kalite", lambda: first_id("/api/v1/editorial/translation/jobs"),
         lambda i: f"/api/v1/editorial/translation/jobs/{q(i)}/quality.csv"),
        ("serbest-hakedis", lambda: first_id("/api/v1/editorial/freelance/payouts"),
         lambda i: f"/api/v1/editorial/freelance/payouts/{q(i)}/export.csv"),
        ("telif-odeme-listesi", lambda: first_id("/api/v1/royalty/runs", lambda r: r.get("status") == "onayli"),
         lambda i: f"/api/v1/royalty/runs/{q(i)}/payments.csv"),
        ("telif-stopaj", lambda: first_id("/api/v1/royalty/runs", lambda r: r.get("status") == "onayli"),
         lambda i: f"/api/v1/royalty/runs/{q(i)}/withholding.csv"),
        ("set-kart-listesi", lambda: first_id("/api/v1/marketing/sets?durum=kart-bekliyor,satista",
                                              lambda s: s.get("durum") in ("kart-bekliyor", "satista")),
         lambda i: f"/api/v1/marketing/sets/{q(i)}/card-todo.csv"),
        ("pazarlama-plani", lambda: first_id("/api/v1/marketing/plans"),
         lambda i: f"/api/v1/marketing/plans/{q(i)}/export.csv"),
        ("pazarlama-crm-listesi", lambda: first_id("/api/v1/marketing/plans"),
         lambda i: f"/api/v1/marketing/plans/{q(i)}/crm-todo.csv"),
        ("butce-hedefleri", lambda: first_id("/api/v1/budget/plans"),
         lambda i: f"/api/v1/budget/plans/{q(i)}/export.csv"),
    ]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="")
    ap.add_argument("--only", default="")
    a = ap.parse_args()
    only = {x for x in a.only.split(",") if x}
    t0 = datetime.now().isoformat(timespec="seconds")
    for key, path in plain():
        if not only or key in only:
            print(f"== {key}", flush=True)
            try:
                test_endpoint(key, path, need_rows=True)
            except Exception as e:  # noqa: BLE001
                check(f"{key}", False, str(e)[:300])
    for key, ids, make in with_ids():
        if not only or key in only:
            print(f"== {key}", flush=True)
            try:
                test_with_ids(key, ids(), make)
            except Exception as e:  # noqa: BLE001
                check(f"{key}", False, str(e)[:300])
    for key, make in (("baski-oneri", baski_csv), ("dahili-rehber", rehber_csv)):
        if not only or key in only:
            print(f"== {key} (tarayıcı CSV'si)", flush=True)
            try:
                test_client_csv(key, make)
            except Exception as e:  # noqa: BLE001
                check(f"{key}", False, str(e)[:300])
    ok = sum(1 for r in results if r["durum"] == "GEÇTİ")
    bad = sum(1 for r in results if r["durum"] == "KALDI")
    warn = sum(1 for r in results if r["durum"] == "UYARI")
    print(f"== {ok} geçti, {bad} kaldı, {warn} uyarı")
    if a.out:
        json.dump({"t0": t0, "base": BASE, "sonuclar": results}, open(a.out, "w"), ensure_ascii=False, indent=1)
    raise SystemExit(1 if bad else 0)


if __name__ == "__main__":
    main()
