#!/usr/bin/env python3
"""TİMAŞ «Kitap Maliyet Formu» Excel'lerini (BASIM EXCELLER) okur: her sayfanın girdileri, sayfanın kendi tarifesi
ve Excel'in hesapladığı sonuçlar. Ürün özelliği değildir (ekran Excel yüklemez); iki iş için:

1. Varsayılan tarifeyi üretmek: `--tarife-yaz backend/semantic_bridge/pricing/form_tariff.json <dosya.xlsx>`
2. Kabul: `kabul.py` bu çıktıyı M9 maliyet formu ucuna verip Excel sonucuyla karşılaştırır.

    python3 excel_oku.py --out /tmp/claude-<oturum>/formlar.json <klasör ya da dosyalar>

Hücre adresleri şablondan (222 satır; iki sayfada ebat tablosuna satır eklenmiş, tablo sonları başlıktan bulunur).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

FIRE = re.compile(r"\$?F\$?5\s*\+\s*(\d+(?:\.\d+)?)")
J33 = re.compile(r"^=\s*\$?([A-Z]{1,3})\$?(\d+)\s*/\s*(\d+(?:\.\d+)?)\s*$")

#: Tarife fiyat satırları (L/O/P sütunları) → anahtar.
PRICE_ROWS = {
    42: "tekRenkKalip", 43: "renkliKalip", 45: "Kabartma Lak-50x70", 46: "Lokal Lak-50x70", 47: "Simli Lak-50x70",
    48: "DİSPERSİYON LAK", 49: "SELOFAN", 50: "harmanlama", 51: "gofre", 52: "yaldiz", 53: "sogukBaski",
    54: "ayracIscilik", 55: "somizIscilik", 56: "tekliVakum", 57: "gren", 58: "ozelKesimRadius", 59: "ozelKesimTbk",
    60: "sertKapakTakma", 61: "sertKapakTaslama", 62: "sertKapakIplikDikis", 63: "sertKapakFormaHarman",
    64: "amerikanCiltForma", 65: "telDikisForma", 66: "iplikDikis", 67: "boardPenceresiz", 68: "boardPencereli",
    69: "boardBristolMukavva", 70: "setMikroOluklu", 71: "setBristolKutu", 72: "setVakum", 73: "setPoset",
    75: "setKoli", 76: "setShirink",
}
EXTRA_ROWS = {"yanKagit": 11, "somiz": 12, "ayrac": 13, "mukavva": 14, "ciltBezi": 15, "digerKagit": 16}
RESULT_CELLS = {"forma": "F4", "iskonto": "I5", "satisFiyati": "J5", "matbaaToplam": "J25", "digerToplam": "J37",
                "toplam": "J40", "dolayli": "J41", "genelToplam": "J42", "birimMaliyet": "J46", "karAdet": "J48",
                "toplamKar": "J49", "karYuzde": "J50", "kitapMaliyeti": "D39", "kagitAdet": "D40", "matbaaAdet": "D41",
                "telifAdet": "D42"}


def is_form(ws) -> bool:
    return str(ws["A2"].value or "").strip().lower().startswith("kitab") and "maliyet" in str(ws["F1"].value or "").lower()


def is_dijital(ws) -> bool:
    """«TBK dijital» şablonu (59 satır): diğer giderler 29. satırda, iç baskı sayfa fiyatıyla (Excel A28 kapak+cilt)."""
    return is_form(ws) and str(ws["A29"].value or "").strip().upper().startswith("DİĞER GİDER")


#: Dijital formun sonuç hücreleri (ofset formundaki aynı anlamlı anahtarlarla).
RESULT_CELLS_DIJITAL = {"forma": "F4", "iskonto": "I5", "satisFiyati": "J5", "matbaaToplam": "J19", "digerToplam": "J29",
                        "toplam": "J32", "dolayli": "J33", "genelToplam": "J34", "birimMaliyet": "J38", "karAdet": "J40",
                        "toplamKar": "J41", "karYuzde": "J42", "kitapMaliyeti": "D31", "kagitAdet": "D32",
                        "matbaaAdet": "D33", "telifAdet": "D34"}


def tariff_dijital(wv, base: dict) -> dict:
    """Dijital Excel'in kendi fiyatları (M37:Z52 ebat × kâğıt tablosu, O18:P32 yayınevi iskontosu, G21–G27 ek işlemler)
    varsayılan tarifenin `dijital` bölümüne yazılır; ofset kısmı (kur vb.) varsayılandan."""
    cols = [s(wv.cell(37, c).value) for c in range(14, 27)]
    tablo = []
    for r in range(38, 53):
        e = s(wv[f"M{r}"].value)
        if e:
            tablo.append({"ebat": e, "fiyat": {c: num(wv.cell(r, 14 + i).value) or 0.0 for i, c in enumerate(cols) if c}})
    pubs = {}
    for r in range(18, 40):
        k = s(wv[f"O{r}"].value)
        if not k:
            break
        pubs[k] = num(wv[f"P{r}"].value)
    d = dict(base.get("dijital") or {})
    kal = {k: dict(v) for k, v in (d.get("kalemler") or {}).items()}
    for key, cell in (("ayracBaski", "G22"), ("kapakBaski", "G25"), ("kulakli", "G24"), ("selofan", "G26")):
        kal.setdefault(key, {})["m"] = num(wv[cell].value) or 0.0
    kal.setdefault("gofre", {})["m"] = num(wv["G23"].value) or 0.0
    kal.setdefault("lokalLak", {})["m"] = num(wv["O54"].value) or 0.0
    d.update(tablo=tablo, kagitlar=[c for c in cols[:11] if c], kapakKagitlari=[c for c in cols[11:] if c],
             publishers=pubs, kalemler=kal, pay=d.get("pay", 25.0))
    return {**base, "dijital": d}


def inputs_dijital(wf, wv) -> dict:
    tarih = wv["J1"].value
    return {
        "baski": "dijital",
        "yayinevi": s(wv["A1"].value), "kitap": s(wv["F2"].value), "yazar": s(wv["F3"].value),
        "tarih": tarih.date().isoformat() if isinstance(tarih, datetime) else None,
        "simdikiFiyat": num(wv["J2"].value), "fiyat": num(wv["J3"].value), "ozelIskonto": num(wv["I4"].value),
        "matbaaAyar": num(wv["I1"].value), "sayfa": num(wv["G4"].value), "adet": num(wv["F5"].value), "ebat": s(wv["G5"].value),
        "dijital": {
            "icKagit": s(wv["D18"].value), "renkliSayfa": num(wv["B17"].value), "renkliKagit": s(wv["D17"].value),
            "kapakKagit": s(wv["D28"].value), "kapakGr": num(wv["B28"].value),
            # G17 boşsa Excel ek pay uygulamaz (H17 = F17).
            "pay": num(wv["G17"].value) or 0.0,
            "ekler": {"ayracAtma": x(wv["B21"].value), "ayracRenk": num(wv["B22"].value), "gofre": x(wv["B23"].value),
                      "kulakli": x(wv["B24"].value), "kapakRenk": num(wv["B25"].value), "selofan": x(wv["B26"].value),
                      "lokalLak": (s(wv["A27"].value) or "").upper() == "LOKAL LAK"},
        },
        "diger": {"yanKagit": num(wv["J20"].value), "nakliyeAdet": num(wv["I21"].value), "hediye": num(wv["J22"].value),
                  "zayiat": num(wv["J23"].value), "reklam": num(wv["J24"].value), "diger": num(wv["J28"].value)},
        "telif": num(wv["I27"].value), "dolayli": num(wv["I33"].value) or 0.0,
    }


def x(v) -> bool:
    return v is not None and str(v).strip() != ""


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def s(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def fire(ws_f, cell: str, default: float) -> float:
    m = FIRE.search(str(ws_f[cell].value or ""))
    return float(m.group(1)) if m else default


def tariff(wf, wv) -> dict:
    papers = []
    for r in range(9, 36):
        name = s(wv[f"M{r}"].value)
        if not name:
            continue
        qf = str(wf[f"Q{r}"].value or "")
        papers.append({"name": name, "base": num(wv[f"N{r}"].value), "cur": "USD" if "$N$7" in qf else "EUR",
                       "unit": "ton" if "/1000" in qf.replace(" ", "") else "adet"})
    prices = {}
    for r, key in PRICE_ROWS.items():
        prices[key] = {"label": s(wv[f"L{r}"].value), "m": num(wv[f"O{r}"].value) or 0.0, "n": num(wv[f"P{r}"].value) or 0.0,
                       "eski": [num(wv[f"W{r}"].value), num(wv[f"X{r}"].value)]}
    trims, r = [], 79
    while wv[f"L{r}"].value is not None and r < 400:
        trims.append({"ebat": str(wv[f"L{r}"].value).strip(), "taslama": num(wv[f"O{r}"].value), "kapakTakma": num(wv[f"P{r}"].value),
                      "mukavvaVerim": num(wv[f"S{r}"].value), "icVerim": num(wv[f"T{r}"].value),
                      "icEn": num(wv[f"U{r}"].value), "icBoy": num(wv[f"V{r}"].value)})
        r += 1
    cliche = {}
    for r in range(1, 7):
        k = s(wv[f"N{r}"].value)
        if k:
            cliche[k] = {"pieces": num(wv[f"O{r}"].value), "price": num(wv[f"P{r}"].value)}
    pubs = {}
    for r in range(7, 60):
        k = s(wv[f"S{r}"].value)
        if not k:
            break
        pubs[k] = num(wv[f"T{r}"].value)
    return {"kur": {"USD": num(wv["J56"].value), "EUR": num(wv["J57"].value)},
            "vade": {"oran": num(wv["P7"].value), "ay": num(wv["P8"].value)},
            "papers": papers, "prices": prices, "trims": trims, "cliche": cliche, "publishers": pubs,
            "fire": {"ic": fire(wf, "F7", 800), "renkli": fire(wf, "F8", 800), "kapak": fire(wf, "F10", 1000),
                     "yanKagit": fire(wf, "F11", 800), "somiz": fire(wf, "F12", 800), "ayrac": fire(wf, "F13", 1000),
                     "mukavva": fire(wf, "F14", 700), "ciltBezi": fire(wf, "F15", 200), "digerKagit": fire(wf, "F16", 1000)},
            "dolayli": 90.0, "kapakBolen": 3.0}


def part(wf, wv, row: int, **extra) -> dict:
    return {"kagit": s(wv[f"A{row}"].value), "en": num(wv[f"B{row}"].value), "boy": num(wv[f"C{row}"].value),
            "gr": num(wv[f"D{row}"].value), "verim": num(wv[f"E{row}"].value) if not str(wf[f"E{row}"].value or "").startswith("=") else None,
            **extra}


ROW_KEY = {7: "icKagit", 10: "kapakKagit", 11: "yanKagit", 12: "somiz", 13: "ayrac", 14: "mukavva", 15: "ciltBezi", 16: "digerKagit"}


def kesim_key(wv):
    """Özel kesimin (B27) hangi kâğıt satırına yapıldığı: Excel VLOOKUP(B27, A7:F17) ile ilk eşleşen satır."""
    name = s(wv["B27"].value)
    if not name:
        return None
    for r in range(7, 18):
        if s(wv[f"A{r}"].value) == name:
            return ROW_KEY.get(r)
    return None


def inputs(wf, wv) -> dict:
    j33f = wf["J33"].value
    ucret, bolen = None, 1.0
    if isinstance(j33f, str) and j33f.startswith("="):
        m = J33.match(j33f.replace(" ", ""))
        if m:
            ucret, bolen = num(wv[f"{m.group(1)}{m.group(2)}"].value), float(m.group(3))
    elif j33f is not None:
        ucret = num(j33f)
    f34 = str(wf["F34"].value or "")
    tarih = wv["J1"].value
    ekler = {k: part(wf, wv, r, fire=fire(wf, f"F{r}", {11: 800, 12: 800, 13: 1000, 14: 700, 15: 200, 16: 1000}[r]))
             for k, r in EXTRA_ROWS.items()}
    ekler["yanKagit"]["renk"] = num(wv["B32"].value)
    ekler["somiz"].update(renk=num(wv["B31"].value), iscilik=x(wv["I28"].value))
    ekler["ayrac"].update(renk=num(wv["B30"].value), iscilik=x(wv["I31"].value))
    ekler.update(kenarBoyama=num(wv["E17"].value), vakum=x(wv["I29"].value),
                 icSelofan={"var": x(wv["B26"].value), "tur": s(wv["A26"].value) or "SELOFAN"})
    return {
        "yayinevi": s(wv["A1"].value), "kitap": s(wv["F2"].value), "yazar": s(wv["F3"].value),
        "tarih": tarih.date().isoformat() if isinstance(tarih, datetime) else (tarih.isoformat() if isinstance(tarih, date) else None),
        "simdikiFiyat": num(wv["J2"].value), "fiyat": num(wv["J3"].value), "ozelIskonto": num(wv["I4"].value),
        "matbaaAyar": num(wv["I1"].value), "sayfa": num(wv["G4"].value), "adet": num(wv["F5"].value), "ebat": s(wv["G5"].value),
        "ic": {**part(wf, wv, 7), "renk": num(wv["B24"].value), "fire": fire(wf, "F7", 800)},
        "renkli": {**part(wf, wv, 8), "sayfa": num(wv["D23"].value), "renk": num(wv["B23"].value), "fire": fire(wf, "F8", 800)},
        "kapak": {**part(wf, wv, 10), "fire": fire(wf, "F10", 1000), "renk": num(wv["B33"].value), "bolen": num(wv["D33"].value),
                  "selofan": {"var": x(wv["B34"].value), "tur": s(wv["A34"].value) or "SELOFAN", "bolen": num(wv["D34"].value),
                              "ciftYuz": 'D32="X"' in f34 and str(wv["D32"].value or "").strip().upper() == "X"},
                  "lak": {"var": x(wv["B35"].value), "tur": s(wv["A35"].value)},
                  "yaldiz": x(wv["B29"].value), "gofre": x(wv["B28"].value),
                  "gren": int(x(wv["I27"].value)) + int(x(wv["I30"].value)),
                  "klise": {"adet": num(wv["I26"].value), "ebat": s(wv["D27"].value)},
                  "ozelKesim": kesim_key(wv)},
        "ekler": ekler,
        "cilt": {"tur": s(wv["A36"].value), "birim": num(wv["D36"].value)},
        "diger": {"kapakUcreti": ucret, "kapakBolen": bolen, "kapakEtiket": s(wv["H33"].value), "nakliye": num(wv["J32"].value),
                  "mizanpaj": num(wv["J34"].value), "diger": num(wv["J36"].value)},
        "telif": num(wv["I35"].value), "dolayli": num(wv["I41"].value),
    }


def default_tariff() -> dict:
    """Depodaki varsayılan tarife (dijital formun ofset dışı alanları için)."""
    p = Path(__file__).resolve().parents[3] / "backend/semantic_bridge/pricing/form_tariff.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def read(path: Path) -> list[dict]:
    wbf = load_workbook(path, data_only=False)
    wbv = load_workbook(path, data_only=True)
    out = []
    for wf in wbf.worksheets:
        if not is_form(wf):
            continue
        wv = wbv[wf.title]
        if is_dijital(wf):
            out.append({"file": path.name, "sheet": wf.title, "inputs": inputs_dijital(wf, wv),
                        "tariff": tariff_dijital(wv, default_tariff()),
                        "excel": {k: num(wv[c].value) for k, c in RESULT_CELLS_DIJITAL.items()}})
            continue
        out.append({"file": path.name, "sheet": wf.title, "inputs": inputs(wf, wv), "tariff": tariff(wf, wv),
                    "excel": {k: num(wv[c].value) for k, c in RESULT_CELLS.items()}})
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--out", help="bütün formlar JSON")
    ap.add_argument("--tarife-yaz", help="ilk dosyanın ilk sayfasının tarifesini bu dosyaya yaz")
    a = ap.parse_args()
    files: list[Path] = []
    for p in map(Path, a.paths):
        files += sorted(f for f in p.glob("*.xlsx") if not f.name.startswith(("~$", "._"))) if p.is_dir() else [p]
    forms = [f for path in files for f in read(path)]
    print(f"{len(files)} dosya, {len(forms)} form sayfası", file=sys.stderr)
    if a.tarife_yaz:
        t = forms[0]["tariff"]
        Path(a.tarife_yaz).write_text(json.dumps(t, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    if a.out:
        Path(a.out).write_text(json.dumps(forms, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
