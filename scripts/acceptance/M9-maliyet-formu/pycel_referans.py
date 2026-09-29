#!/usr/bin/env python3
"""Excel formüllerini bağımsız hesaplayan referans (pycel): Excel'de hiç kullanılmamış seçeneklerin kabulü için.

İki kip:
- `--denetim <klasör>`: özgün Excel'lerde pycel'in hesabı Excel'in kendi kaydettiği sonuçla aynı mı (referansın güvenilirliği).
- `--senaryo <klasör> --taban <özgün.xlsx>`: `senaryo.py`'nin girdisi değiştirilmiş dosyaları; girdiler dosyadan, fiyatlar
  (sağdaki tablolar) özgün dosyadan, sonuçlar pycel'den → `kabul.py`'nin okuduğu biçimde JSON.

Sistem bu hesaba hiç karışmaz: referans yalnız Excel'in formül metninden hesaplanır.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pycel import ExcelCompiler

sys.path.insert(0, str(Path(__file__).resolve().parent))
import excel_oku as E  # noqa: E402

TOL = 1e-6


def explicit_intersection(path: Path) -> Path:
    """J27–J31 formülleri tek hücre yerine `$H$27:$H$31` aralığını karşılaştırır; Excel bunu formülün kendi satırıyla
    kesiştirir (örtük kesişim: J29'da H29). pycel aynı metinli formüllerde satırı karıştırıyor (2026-09-29: J29'da
    vakum yerine gren). Excel'in anlamı açık yazılır (her satırda `$H$<satır>`). Düzeltme dosyanın XML'inde yapılır:
    openpyxl ile kaydetmek bazı dosyalarda (ör. Pusula) pycel'in düşey aramasını bozuyordu."""
    out = path.with_name(path.stem + ".kesisim.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                x = data.decode("utf-8")
                for r in range(27, 32):
                    x = re.sub(r'(<c r="J%d"[^>]*>\s*<f>)([^<]*)' % r,
                               lambda m, r=r: m.group(1) + m.group(2).replace("$H$27:$H$31", f"$H${r}"), x)
                data = x.encode("utf-8")
            zout.writestr(item, data)
    return out


def pycel_results(path: Path, sheet: str) -> dict:
    xl = ExcelCompiler(filename=str(explicit_intersection(path)))
    out = {}
    for k, cell in E.RESULT_CELLS.items():
        v = xl.evaluate(f"'{sheet}'!{cell}")
        out[k] = E.num(v) if not isinstance(v, str) else None
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--denetim")
    ap.add_argument("--senaryo")
    ap.add_argument("--taban")
    ap.add_argument("--out")
    a = ap.parse_args()
    if a.denetim:
        ok = bad = 0
        for f in sorted(p for p in Path(a.denetim).glob("*.xlsx") if not p.name.endswith(".kesisim.xlsx")):
            for form in E.read(f):
                try:
                    got = pycel_results(f, form["sheet"])
                except Exception as e:  # noqa: BLE001
                    print(f"[HATA] {f.name} [{form['sheet']}] {type(e).__name__}: {str(e)[:200]}")
                    bad += 1
                    continue
                diff = {k: (form["excel"][k], got[k]) for k in got
                        if not (form["excel"][k] is None and got[k] in (None, 0)) and
                        (form["excel"][k] is None or got[k] is None or abs(got[k] - form["excel"][k]) > (TOL if k == "karYuzde" else 0.01))}
                ok += not diff
                bad += bool(diff)
                print(f"[{'OK' if not diff else 'FARK'}] {f.name} [{form['sheet']}]" + (f" {diff}" if diff else ""))
        print(f"\npycel = Excel: {ok}/{ok + bad}")
        return 0 if not bad else 1
    base = E.read(Path(a.taban))[0]
    forms = []
    for f in sorted(p for p in Path(a.senaryo).glob("*.xlsx") if not p.name.endswith(".kesisim.xlsx")):
        wf = load_workbook(f, data_only=False).worksheets[0]
        wv = load_workbook(f, data_only=True).worksheets[0]
        try:
            res = pycel_results(f, wf.title)
        except Exception as e:  # noqa: BLE001
            print(f"[HATA] {f.name}: {type(e).__name__}: {str(e)[:200]}")
            continue
        forms.append({"file": f.name, "sheet": wf.title, "inputs": E.inputs(wf, wv), "tariff": base["tariff"], "excel": res})
        print(f"{f.name}: birim {res['birimMaliyet']}")
    Path(a.out).write_text(json.dumps(forms, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(forms)} senaryo → {a.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
