#!/usr/bin/env python3
"""Planlı rapor Excel taslağı: gerçek köprü API'si + gerçek TİMAŞ Logo veritabanı ile kabul.

Akış kullanıcının ekrandaki akışıdır: cümle → 50 satırlık önizleme → kolon düzeltmeleri (anında) →
veri değiştiren düzeltme (model + yeniden sorgu) → onay (kayıt) → çalıştırma → Excel indirme.

Excel'in içeriği uygulamanın SQL'i yeniden çalıştırılarak değil, bu betikte ayrıca yazılmış bir referans
sorgusuyla (farklı yapı: CTE + kendi gruplaması) doğrudan veritabanından karşılaştırılır.

Sunucuda çalışır (Mac'te değil):
  sudo -n env $(sudo -n grep ^SEMANTIC_CALLER_TOKEN= /etc/nanobase/semantic-bridge.env) \\
    /data/nanobaseai/bi/semantic-venv/bin/python scripts/server/accept-report-excel-draft.py --base http://127.0.0.1:8795

Kanıt: /data/nanobaseai/bi/logs/accept-excel-draft-<zaman>.json (+ indirilen .xlsx). Oluşturulan test raporu sonda silinir.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import sys
import time
import traceback
import urllib.error
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

QUESTION_TEXT = "her pazartesi 08:30'da 2026 yılında müşteri bazında net ciro ve fatura sayısı excel olarak hazırla"
SECRETS = Path("/data/nanobaseai/bi/secrets/logo-mssql-connection.json")
LOGS = Path("/data/nanobaseai/bi/logs")

# Referans: net ciro'nun sertifikalı tanımı (satış faturası 7/8/9 artı, iade 2/3 eksi, iptal hariç) ve fatura
# sayısı (aynı kapsamdaki fatura adedi; kokpit KPI 73.660), 2026. Uygulamanın SQL'inden bağımsız yazıldı:
# önce fatura başına işaretli tutar, sonra cari adına göre toplam.
REFERENCE_SQL = """
WITH f AS (
    SELECT i.CLIENTREF,
           CASE WHEN i.TRCODE IN (7, 8, 9) THEN i.NETTOTAL
                WHEN i.TRCODE IN (2, 3) THEN -i.NETTOTAL END AS signed_net,
           i.TRCODE
    FROM dbo.LG_411_01_INVOICE AS i
    WHERE i.CANCELLED = 0
      AND i.DATE_ >= '20260101' AND i.DATE_ < '20270101'
      AND i.TRCODE IN (2, 3, 7, 8, 9)
)
SELECT c.DEFINITION_ AS name, SUM(f.signed_net) AS net, COUNT(*) AS invoices
FROM f LEFT JOIN dbo.LG_411_CLCARD AS c ON c.LOGICALREF = f.CLIENTREF
GROUP BY c.DEFINITION_
"""


class Api:
    def __init__(self, base: str, token: str):
        self.base = base.rstrip("/")
        self.headers = {"X-Semantic-Caller": token, "Cookie": "timas_session=acceptance"}

    def call(self, method: str, path: str, body: Any = None, *, raw: bool = False) -> tuple[int, Any, float]:
        data = None if body is None else json.dumps(body).encode()
        h = dict(self.headers)
        if body is not None:
            h["Content-Type"] = "application/json"
        req = urllib.request.Request(self.base + path, data=data, headers=h, method=method)
        t = time.perf_counter()
        try:
            with urllib.request.urlopen(req, timeout=900) as res:
                payload = res.read()
                code = res.status
        except urllib.error.HTTPError as e:
            payload = e.read()
            code = e.code
        dt = time.perf_counter() - t
        if raw:
            return code, payload, dt
        try:
            return code, json.loads(payload.decode() or "null"), dt
        except ValueError:
            return code, payload.decode(errors="replace"), dt


class Run:
    def __init__(self) -> None:
        self.steps: list[dict[str, Any]] = []
        self.failed = 0

    def check(self, step: str, name: str, ok: bool, **evidence: Any) -> bool:
        status = "PASS" if ok else "FAIL"
        if not ok:
            self.failed += 1
        self.steps.append({"step": step, "check": name, "status": status, **evidence})
        print(f"[{status}] {step} · {name}" + (f" · {json.dumps(evidence, ensure_ascii=False, default=str)[:300]}" if evidence and not ok else ""))
        return ok

    def note(self, step: str, name: str, status: str, **evidence: Any) -> None:
        self.steps.append({"step": step, "check": name, "status": status, **evidence})
        print(f"[{status}] {step} · {name} · {json.dumps(evidence, ensure_ascii=False, default=str)[:400]}")


def close(a: Any, b: Any, tol: float = 0.01) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return abs(float(a) - float(b)) <= tol


def _name(v: Any) -> Any:
    # xlsx satır sonundaki \r'yi saklamaz (XML satır sonu normalleşmesi); SQL Server da GROUP BY'da
    # sondaki boşluğu yok sayar ("AD" ile "AD " tek gruptur). İkisinde de ad aynı ad.
    return v.replace("\r\n", "\n").replace("\r", "\n").rstrip(" ") if isinstance(v, str) else v


def reference_rows() -> dict[str | None, dict[str, Any]]:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "backend"))
    from scripts.export_crm_metadata import connect

    cn = connect(SECRETS, json.loads(SECRETS.read_text())["database"])
    try:
        cur = cn.cursor()
        cur.execute(REFERENCE_SQL)
        out = {}
        for name, net, invoices in cur.fetchall():
            out[_name(name)] = {"net": float(net) if net is not None else None, "invoices": int(invoices)}
        return out
    finally:
        cn.close()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="http://127.0.0.1:8795")
    ap.add_argument("--keep", action="store_true", help="test raporunu silme")
    args = ap.parse_args()
    token = os.environ.get("SEMANTIC_CALLER_TOKEN", "")
    api = Api(args.base, token)
    run = Run()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    evidence: dict[str, Any] = {"base": args.base, "startedAt": stamp, "question": QUESTION_TEXT}
    rid = None
    try:
        # ------------------------------------------------------------ 1. cümle → önizleme
        code, d, dt = api.call("POST", "/api/v1/reports/parse", {"text": QUESTION_TEXT})
        evidence["parse"] = {"status": code, "seconds": round(dt, 1), "question": d.get("question") if isinstance(d, dict) else None,
                             "sql": d.get("sql") if isinstance(d, dict) else None, "rowCount": d.get("rowCount") if isinstance(d, dict) else None}
        if not run.check("1 önizleme", "parse 200 döner", code == 200, body=d if code != 200 else None):
            raise SystemExit
        total = int(d["rowCount"] or 0)
        run.check("1 önizleme", "sonuç 50 satırdan büyük (örnek uyarısı gerekli)", total > 50, rowCount=total)
        run.check("1 önizleme", "önizleme tam 50 satır taşır", len(d["records"]) == 50, shown=len(d["records"]))
        names = [c["name"] for c in d["columns"]]
        run.check("1 önizleme", "kolon düzeni kaynak kolonlarla aynı sırada", [c["key"] for c in d["layout"]] == names, layout=d["layout"], columns=names)
        run.check("1 önizleme", "varsayılan düzende gizli kolon yok, ad = kaynak ad",
                  all((not c["hidden"]) and c["label"] == c["key"] for c in d["layout"]))
        run.check("1 önizleme", "plan alanları çözüldü (haftalık pazartesi 08:30)",
                  d["recurrence"] == "weekly" and d["weekday"] == 0 and d["at"] == "08:30", recurrence=d["recurrence"], weekday=d["weekday"], at=d["at"])
        question, sql, layout, preview = d["question"], d["sql"], d["layout"], d["records"]

        num = [c["name"] for c in d["columns"] if str(c["type"]).lower() in ("float", "decimal", "number", "double", "money")]
        txt = [c["name"] for c in d["columns"] if str(c["type"]).lower() in ("str", "string", "text", "varchar")]
        cnt = [c["name"] for c in d["columns"] if str(c["type"]).lower() in ("int", "integer", "bigint")]
        if not (num and txt and cnt):
            run.note("1 önizleme", "beklenen kolon tipleri (metin + ondalık + tam sayı) bulunamadı", "DOĞRULANAMADI", columns=d["columns"])
            raise SystemExit
        k_name, k_net, k_cnt = txt[0], num[0], cnt[0]

        # ------------------------------------------------------------ 2. kolon komutları (model yok, veri yok)
        instr = f"{k_cnt} kolonunu gizle, {k_name} adını Müşteri yap, {k_net} adını Net Ciro yap, Net Ciro para olarak göster"
        code, r2, dt = api.call("POST", "/api/v1/reports/refine", {"question": question, "instruction": instr, "columns": layout})
        evidence["refineRules"] = {"status": code, "seconds": round(dt, 2), "instruction": instr, "body": r2}
        if run.check("2 kolon komutu", "refine 200 döner", code == 200, body=r2 if code != 200 else None):
            by = {c["key"]: c for c in r2["layout"]}
            run.check("2 kolon komutu", "modele gitmeden çözüldü (via=rules)", r2["via"] == "rules", via=r2["via"])
            run.check("2 kolon komutu", "veri yeniden çekilmedi (requery=false, soru aynı)", r2["requery"] is False and r2["question"] == question)
            run.check("2 kolon komutu", "anında yanıt (< 2 sn)", dt < 2, seconds=round(dt, 2))
            run.check("2 kolon komutu", f"{k_cnt} gizlendi", by[k_cnt]["hidden"] is True)
            run.check("2 kolon komutu", "adlar değişti", by[k_name]["label"] == "Müşteri" and by[k_net]["label"] == "Net Ciro", layout=r2["layout"])
            run.check("2 kolon komutu", "Net Ciro biçimi para", by[k_net]["format"] == "money")
            run.check("2 kolon komutu", "her değişiklik kullanıcıya madde olarak döndü", len(r2["changes"]) == 4, changes=r2["changes"])
            layout = r2["layout"]

        code, r3, dt = api.call("POST", "/api/v1/reports/refine", {"question": question, "instruction": "Net Ciro kolonunu başa al", "columns": layout})
        evidence["refineReorder"] = {"status": code, "body": r3}
        if run.check("2 kolon komutu", "sıra komutu 200", code == 200, body=r3 if code != 200 else None):
            run.check("2 kolon komutu", "Net Ciro ilk kolon oldu", r3["layout"][0]["key"] == k_net, order=[c["key"] for c in r3["layout"]])
            layout = r3["layout"]

        # ------------------------------------------------------------ 3. hatalı istekler kullanıcıya düz Türkçe döner
        dup = [dict(c, label="Aynı", hidden=False) for c in layout]
        code, e1, _ = api.call("POST", "/api/v1/reports/preview", {"question": question, "columns": dup})
        run.check("3 geçersiz düzen", "aynı adlı iki görünür kolon 422", code == 422 and "aynı adı" in json.dumps(e1, ensure_ascii=False), status=code, body=e1)
        hid = [dict(c, hidden=True) for c in layout]
        code, e2, _ = api.call("POST", "/api/v1/reports/refine", {"question": question, "instruction": "Net Ciro adını Ciro yap", "columns": hid})
        run.check("3 geçersiz düzen", "tüm kolonlar gizli 422", code == 422 and "görünür" in json.dumps(e2, ensure_ascii=False), status=code, body=e2)
        code, e3, _ = api.call("POST", "/api/v1/reports/refine", {"question": question, "instruction": "   ", "columns": layout})
        run.check("3 geçersiz düzen", "boş düzeltme cümlesi 422", code == 422, status=code, body=e3)
        code, e4, _ = api.call("POST", "/api/v1/reports", {"title": "x", "question": question, "recurrence": "weekly", "at": "08:30",
                                                              "weekday": 0, "recipients": [], "fmt": "xlsx", "columns": dup})
        run.check("3 geçersiz düzen", "kaydetme de yinelenen adı reddeder", code == 422, status=code, body=e4)

        # ------------------------------------------------------------ 4. veri değiştiren düzeltme (model + yeniden sorgu)
        code, r4, dt = api.call("POST", "/api/v1/reports/refine",
                                {"question": question, "instruction": "yalnız 2026 ağustos ayını göster", "columns": layout})
        evidence["refineModel"] = {"status": code, "seconds": round(dt, 1),
                                   "body": {k: v for k, v in (r4 or {}).items() if k != "records"} if isinstance(r4, dict) else r4}
        if code == 200:
            run.check("4 model düzeltmesi", "soru değişti ve veri yeniden çekildi", r4["requery"] is True and r4["question"] != question,
                      newQuestion=r4.get("question"))
            if r4["requery"]:
                run.check("4 model düzeltmesi", "yeni önizleme ≤ 50 satır, toplam ayrı bildirildi",
                          len(r4["records"]) <= 50 and r4["rowCount"] is not None, shown=len(r4["records"]), rowCount=r4["rowCount"])
                run.check("4 model düzeltmesi", "ağustos sonucu yıl sonucundan küçük", (r4["rowCount"] or 0) < total, aug=r4["rowCount"], year=total)
                source_keys = {c["name"] for c in r4["columns"]}
                kept = [c for c in layout if c["key"] in source_keys]
                new_by = {c["key"]: c for c in r4["layout"]}
                preserved = all(new_by[c["key"]]["label"] == c["label"] and new_by[c["key"]]["hidden"] == c["hidden"] for c in kept)
                run.check("4 model düzeltmesi", "kaynak adı aynı kalan kolonlarda kişinin adı/gizlemesi korundu", preserved,
                          kept=[c["key"] for c in kept], layout=r4["layout"])
                lost = [c["label"] for c in layout if c["key"] not in source_keys]
                run.check("4 model düzeltmesi", "kaybolan kolon sessiz değil (dropped listesinde)", sorted(lost) == sorted(r4["dropped"]),
                          lost=lost, dropped=r4["dropped"], added=r4["added"])
        else:
            run.note("4 model düzeltmesi", "model/yeniden sorgu yanıtı", "FAIL" if code != 503 else "DOĞRULANAMADI", status=code, body=r4)
            if code != 503:
                run.failed += 1

        # ------------------------------------------------------------ 5. onay → kayıt → çalıştırma → Excel
        body = {"title": f"KABUL excel taslağı {stamp}", "prompt": QUESTION_TEXT, "question": question, "sql": sql,
                "recurrence": "weekly", "at": "08:30", "weekday": 0, "recipients": [], "fmt": "xlsx", "columns": layout}
        code, rep, _ = api.call("POST", "/api/v1/reports", body)
        if not run.check("5 onay", "rapor kolon düzeniyle kaydedildi", code == 200 and rep.get("columns") == layout, status=code, body=rep):
            raise SystemExit
        rid = rep["id"]
        code, lst, _ = api.call("GET", "/api/v1/reports")
        saved = next((x for x in lst.get("reports", []) if x["id"] == rid), None)
        run.check("5 onay", "liste düzeni geri okur (sunucuda kalıcı)", bool(saved) and saved["columns"] == layout)

        code, ran, dt = api.call("POST", f"/api/v1/reports/{rid}/run", {})
        evidence["run"] = {"status": code, "seconds": round(dt, 1), "body": ran}
        run.check("5 onay", "çalıştırma dosya üretti (alıcı yok → no_recipient)", code == 200 and ran["lastStatus"] == "no_recipient" and ran["hasFile"],
                  status=code, lastStatus=ran.get("lastStatus"), lastError=ran.get("lastError"))
        run.check("5 onay", "dosya satır sayısı = önizlemenin bildirdiği toplam", ran.get("lastRows") == total, lastRows=ran.get("lastRows"), rowCount=total)
        run.check("5 onay", "çalıştırma notu yok (düzendeki tüm kolonlar sonuçta)", not ran.get("lastError"), lastError=ran.get("lastError"))

        code, blob, _ = api.call("GET", f"/api/v1/reports/{rid}/file", raw=True)
        if not run.check("5 onay", "Excel indirildi", code == 200 and blob[:2] == b"PK", status=code):
            raise SystemExit
        LOGS.mkdir(parents=True, exist_ok=True)
        xlsx_path = LOGS / f"accept-excel-draft-{stamp}.xlsx"
        xlsx_path.write_bytes(blob)
        evidence["xlsx"] = str(xlsx_path)

        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(blob)).active
        header = [c.value for c in ws[1]]
        expected_header = [c["label"] for c in layout if not c["hidden"]]
        run.check("6 Excel", "başlık satırı = görünür kolon adları, ekrandaki sırayla", header == expected_header, header=header, expected=expected_header)
        run.check("6 Excel", "gizli kolon dosyada yok", not any(c["label"] in header for c in layout if c["hidden"]))
        data = list(ws.iter_rows(min_row=2, values_only=True))
        run.check("6 Excel", "veri satırı sayısı = toplam (50 değil)", len(data) == total, rows=len(data), total=total)
        i_net, i_name = header.index("Net Ciro"), header.index("Müşteri")
        cnt_label = next(c["label"] for c in layout if c["key"] == k_cnt)
        fmt = ws.cell(row=2, column=i_net + 1).number_format
        run.check("6 Excel", "Net Ciro hücreleri para biçiminde", "₺" in fmt, numberFormat=fmt)
        run.check("6 Excel", "Net Ciro hücreleri sayı (metin değil)", all(isinstance(r[i_net], (int, float)) or r[i_net] is None for r in data))
        same50 = all(close(data[i][i_net], preview[i][k_net]) and data[i][i_name] == preview[i][k_name] for i in range(50))
        run.check("6 Excel", "ilk 50 satır ekrandaki önizlemeyle aynı (değer değişmedi)", same50)

        # ------------------------------------------------------------ 7. bağımsız referans (doğrudan DB)
        t = time.perf_counter()
        ref = reference_rows()
        evidence["reference"] = {"seconds": round(time.perf_counter() - t, 1), "groups": len(ref), "sql": REFERENCE_SQL.strip()}
        run.check("7 referans", "müşteri grubu sayısı = Excel satır sayısı", len(ref) == len(data), reference=len(ref), excel=len(data))
        mism = []
        for r in data:
            want = ref.get(_name(r[i_name]))
            if want is None or not close(r[i_net], want["net"]):
                mism.append({"name": r[i_name], "excel": r[i_net], "reference": None if want is None else want["net"]})
        run.check("7 referans", "her müşterinin net cirosu referansla aynı (±0,01 ₺)", not mism, mismatches=len(mism), sample=mism[:5])
        # Fatura sayısı Excel'de gizli; önizlemenin ilk 50 satırındaki değer referansla karşılaştırılır.
        cnt_mism = [{"name": r[k_name], "preview": r[k_cnt], "reference": (ref.get(_name(r[k_name])) or {}).get("invoices")}
                    for r in preview if (ref.get(_name(r[k_name])) or {}).get("invoices") != r[k_cnt]]
        run.check("7 referans", f"ilk 50 müşterinin fatura sayısı (satış + iade faturası) referansla aynı ({cnt_label})", not cnt_mism,
                  mismatches=len(cnt_mism), sample=cnt_mism[:5])
        excel_sum = sum(float(r[i_net] or 0) for r in data)
        ref_sum = sum(float(v["net"] or 0) for v in ref.values())
        run.check("7 referans", "Excel toplam net ciro = referans toplam", close(excel_sum, ref_sum, 0.05), excel=round(excel_sum, 2), reference=round(ref_sum, 2))
        evidence["totals"] = {"excelNet": round(excel_sum, 2), "referenceNet": round(ref_sum, 2)}
    except SystemExit:
        pass
    except Exception as e:  # noqa: BLE001
        run.failed += 1
        run.note("hata", type(e).__name__, "FAIL", error=str(e), trace=traceback.format_exc()[-1500:])
    finally:
        if rid and not args.keep:
            code, _, _ = api.call("DELETE", f"/api/v1/reports/{rid}")
            run.check("temizlik", "test raporu silindi", code == 200, status=code)
        passed = sum(1 for s in run.steps if s["status"] == "PASS")
        unverified = sum(1 for s in run.steps if s["status"] == "DOĞRULANAMADI")
        evidence.update(steps=run.steps, passed=passed, failed=run.failed, unverified=unverified,
                        finishedAt=datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"))
        LOGS.mkdir(parents=True, exist_ok=True)
        out = LOGS / f"accept-excel-draft-{stamp}.json"
        out.write_text(json.dumps(evidence, ensure_ascii=False, indent=2, default=str))
        print(f"\nSONUÇ: {passed} geçti · {run.failed} kaldı · {unverified} doğrulanamadı · kanıt {out}")
    return 1 if run.failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
